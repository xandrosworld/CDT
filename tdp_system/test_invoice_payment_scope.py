import io
import json
import tempfile
import unittest
import zipfile
from pathlib import Path

from openpyxl import load_workbook

try:
    from . import server
except ImportError:  # pragma: no cover - direct invocation
    import server


class InvoicePaymentScopeTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp_dir = tempfile.TemporaryDirectory()
        cls.original_db_path = server.DB_PATH
        cls.original_data_dir = server.DATA_DIR
        cls.original_testing = server.app.config.get("TESTING", False)
        server.DATA_DIR = Path(cls.temp_dir.name) / "data"
        server.DATA_DIR.mkdir(parents=True, exist_ok=True)
        server.DB_PATH = server.DATA_DIR / "invoice_payment_scope.sqlite3"
        server.app.config["TESTING"] = True
        server.init_database()
        cls.client = server.app.test_client()

    @classmethod
    def tearDownClass(cls):
        server.DB_PATH = cls.original_db_path
        server.DATA_DIR = cls.original_data_dir
        server.app.config["TESTING"] = cls.original_testing
        cls.temp_dir.cleanup()

    def setUp(self):
        with server.db() as conn:
            conn.execute('DELETE FROM outgoing_buyer_profiles')
            conn.execute("DELETE FROM outgoing_substitution_actions")
            conn.execute("DELETE FROM invoice_inventory_ledger")
            conn.execute("DELETE FROM invoice_inventory_confirmations")
            conn.execute("DELETE FROM outgoing_source_invoice_items")
            conn.execute("DELETE FROM outgoing_source_events")
            conn.execute("DELETE FROM outgoing_source_invoices")
            conn.execute("DELETE FROM outgoing_invoice_lines")
            conn.execute("DELETE FROM outgoing_invoice_drafts")
            conn.execute("DELETE FROM orders")
            conn.execute("DELETE FROM batches")
            conn.execute("DELETE FROM audit_log")
            conn.execute(
                """INSERT OR REPLACE INTO contractors(code,name,price_group,pricing_mode)
                   VALUES('NT-A','Nhà thầu A','NT-A','group')"""
            )
            conn.execute(
                """INSERT OR REPLACE INTO contractors(code,name,price_group,pricing_mode)
                   VALUES('NT-B','Nhà thầu B','NT-B','group')"""
            )
            conn.execute(
                """INSERT OR REPLACE INTO products(
                       code,name,unit,tax,supplier,buy_price,purchase_list,seller,cccd
                   ) VALUES('P-1','Hàng 1','kg','8%','NCC',50,0,'','')"""
            )

    @staticmethod
    def add_invoice(
        conn,
        *,
        status="issued",
        invoice_date="2026-09-03",
        invoice_number="0000001",
        snapshot_suffix="",
        product_code="P-1",
        draft_kind="standard",
        contractor="NT-A",
    ):
        timestamp = server.now_iso()
        batch_id = conn.execute(
            """INSERT INTO batches(work_date,source_name,status,created_at,approved_at)
               VALUES(?,'invoice-source.xlsx','approved',?,?)""",
            (invoice_date, timestamp, timestamp),
        ).lastrowid
        order_id = conn.execute(
            """INSERT INTO orders(
                   batch_id,work_date,contractor,kitchen,product_code,product_name,qty,
                   actual_received,actual_delivered,unit,supplier,buy_price,sell_price,tax,
                   purchase_list,errors,warnings,updated_at
               ) VALUES(?,?,?,'BẾP A','P-1','Hàng gốc',2,2,2,'kg','NCC',50,100,
                        '8%',0,'[]','[]',?)""",
            (batch_id, invoice_date, contractor, timestamp),
        ).lastrowid
        issued_at = timestamp if status == "issued" else None
        draft_id = conn.execute(
            """INSERT INTO outgoing_invoice_drafts(
                   batch_id,contractor,invoice_date,status,subtotal,tax_amount,total_amount,
                   created_at,issued_at,round_no,draft_kind,issued_invoice_number,
                   issued_invoice_series,issued_invoice_date,buyer_name_snapshot,
                   buyer_tax_code_snapshot,buyer_address_snapshot,company_name_snapshot,
                   company_tax_code_snapshot,company_address_snapshot,payment_requester_snapshot,
                   payment_bank_name_snapshot,payment_bank_account_snapshot
               ) VALUES(?,?,?, ?,200,16,216,?,?,1,?,?, '1C26TDP',?, ?,?,?,?,?,?,?,?,?)""",
            (
                batch_id, contractor, invoice_date, status, timestamp, issued_at, draft_kind,
                invoice_number, invoice_date,
                "CÔNG TY MUA" + snapshot_suffix, "0200000001", "Địa chỉ mua",
                "CÔNG TY TĐP", "0202265016", "Địa chỉ TĐP", "VŨ THỊ THỤY",
                "Vietcombank", "1052787580",
            ),
        ).lastrowid
        conn.execute(
            """INSERT INTO outgoing_invoice_lines(
                   draft_id,order_id,product_code,product_name,qty,unit,unit_price,tax,
                   invoice_nature,amount
               ) VALUES(?,?,?,?,2,'kg',100,'8%','1',200)""",
            (draft_id, order_id, product_code, "Hàng xuất hóa đơn"),
        )
        return draft_id, batch_id, order_id

    @staticmethod
    def add_source(
        conn,
        *,
        status_class="issued",
        sync_status="synced",
        stock_status="posted",
        total=216,
        buyer_tax_code="0200000001",
        invoice_number="0000001",
        invoice_date="2026-09-03",
        source="minvoice",
    ):
        timestamp = server.now_iso()
        return conn.execute(
            """INSERT INTO outgoing_source_invoices(
                   tenant,source,identity_key,remote_id,business_key,buyer_tax_code,buyer_name,
                   invoice_number,invoice_series,invoice_date,subtotal,tax_amount,total_amount,
                   source_status_raw,source_status_class,source_status_field,relation_reference,
                   sync_status,stock_status,raw_json,error_message,synced_at,created_at,updated_at
               ) VALUES('default',?,?,?,?,?,?,?,'1C26TDP',?,200,16,?, ?,?,'status','',
                        ?,?,'{}','',?,?,?)""",
            (
                source, f"ID-{invoice_number}-{status_class}", f"REMOTE-{invoice_number}-{status_class}",
                f"BIZ-{invoice_number}", buyer_tax_code, "CÔNG TY MUA", invoice_number,
                invoice_date, total, status_class, status_class, sync_status, stock_status,
                timestamp, timestamp, timestamp,
            ),
        ).lastrowid

    def scope(self, contractor="NT-A", date_from="2026-09-01", date_to="2026-09-30"):
        return self.client.get(
            "/api/outgoing-invoices/payment-scope/" + contractor +
            f"?from={date_from}&to={date_to}"
        )

    def test_operational_order_and_unissued_or_cancelled_drafts_never_enter_scope(self):
        with server.db() as conn:
            self.add_invoice(conn, status="draft", invoice_number="DRAFT-1")
            self.add_invoice(conn, status="cancelled", invoice_number="CANCEL-1")
        response = self.scope()
        self.assertEqual(response.status_code, 404)
        self.assertEqual(response.get_json()["code"], "issued_invoice_scope_empty")

    def test_local_issued_scope_preview_and_official_payment_bundle(self):
        with server.db() as conn:
            self.add_invoice(conn)
        preview_response = self.scope()
        self.assertEqual(preview_response.status_code, 200, preview_response.get_data(as_text=True))
        preview = preview_response.get_json()
        with server.db() as conn:
            before = conn.execute('SELECT count(*) FROM audit_log').fetchone()[0]
        inline = self.client.post('/api/documents/preview',json={'kind':'payment','contractor':'NT-A',
            'from':'2026-09-01','to':'2026-09-30','scope_id':preview['scope_id']})
        self.assertEqual(200,inline.status_code,inline.get_json())
        self.assertEqual(2,inline.get_json()['sheet_count'])
        self.assertIn('ĐỀ NGHỊ THANH TOÁN',str(inline.get_json()['sheets']))
        with server.db() as conn:
            self.assertEqual(before,conn.execute('SELECT count(*) FROM audit_log').fetchone()[0])
        self.assertTrue(preview["official_template_ready"])
        self.assertEqual(preview["template_status"], "official_customer_xlsx")
        self.assertEqual(preview["totals"], {"subtotal": 200, "tax_amount": 16, "total_amount": 216})
        self.assertEqual(preview["invoices"][0]["verification_source"], "local_issued_confirmation")
        self.assertNotIn("lines", preview)
        bundle = self.client.get(
            "/api/export/invoice-payment-bundle/NT-A?from=2026-09-01&to=2026-09-30&scope_id=" +
            preview["scope_id"]
        )
        self.assertEqual(bundle.status_code, 200)
        self.assertEqual(bundle.headers["X-TDP-Template-Status"], "official-customer-xlsx")
        self.assertEqual(bundle.headers["X-TDP-Invoice-Scope"], preview["scope_id"])
        with zipfile.ZipFile(io.BytesIO(bundle.data)) as archive:
            names = archive.namelist()
            self.assertEqual(len(names), 3)
            self.assertTrue(any(name.startswith("De_nghi_thanh_toan") for name in names))
            self.assertTrue(any(name.startswith("Bang_tong_hop_giao_nhan") for name in names))
            manifest = archive.read("THONG_TIN_DOI_CHIEU.txt").decode("utf-8")
            self.assertIn("biểu mẫu chính thức", manifest)
            self.assertNotIn("CHƯA PHẢI MẪU", manifest)
            self.assertIn(preview["scope_id"], manifest)
            payment_name = next(name for name in names if name.startswith("De_nghi_thanh_toan"))
            payment = load_workbook(io.BytesIO(archive.read(payment_name)), data_only=True, keep_links=False)
            try:
                self.assertEqual("ĐỀ NGHỊ THANH TOÁN", payment["Đề nghị thanh toán"]["A6"].value)
                self.assertEqual("0000001", payment["Đề nghị thanh toán"]["C15"].value)
                self.assertEqual(216, payment["Đề nghị thanh toán"]["F15"].value)
            finally:
                payment.close()
            statement_name = next(name for name in names if name.startswith("Bang_tong_hop_giao_nhan"))
            statement = load_workbook(io.BytesIO(archive.read(statement_name)), data_only=True, keep_links=False)
            try:
                self.assertIn("BẢNG TỔNG HỢP GIAO NHẬN", statement["Bảng kê giao hàng"]["A5"].value)
                self.assertEqual(216, statement["Bảng kê giao hàng"]["J11"].value)
                self.assertEqual(216, statement["Bảng kê giao hàng"]["J12"].value)
                self.assertEqual(statement["Đối chiếu hóa đơn"].cell(4, 10).value, 0)
            finally:
                statement.close()

    def test_synced_issued_source_is_provenance_and_bad_source_states_are_blocked(self):
        with server.db() as conn:
            self.add_invoice(conn)
            self.add_source(conn)
        synced = self.scope()
        self.assertEqual(synced.status_code, 200, synced.get_data(as_text=True))
        self.assertEqual(synced.get_json()["invoices"][0]["verification_source"], "synced_issued_source")

        for source_class, sync_status, stock_status in (
            ("cancelled", "reconcile_required", "reversal_required"),
            ("replaced", "reconcile_required", "reversal_required"),
            ("adjusted", "reconcile_required", "reversal_required"),
            ("unknown", "review_required", "blocked"),
        ):
            with self.subTest(source_class=source_class):
                self.setUp()
                with server.db() as conn:
                    self.add_invoice(conn)
                    self.add_source(
                        conn, status_class=source_class,
                        sync_status=sync_status, stock_status=stock_status,
                    )
                blocked = self.scope()
                self.assertEqual(blocked.status_code, 409)
                self.assertEqual(blocked.get_json()["code"], "invoice_source_not_payable")

    def test_legacy_msmi_output_is_not_treated_as_minvoice_provenance(self):
        with server.db() as conn:
            self.add_invoice(conn)
            self.add_source(conn, source="msmi")
        preview = self.scope()
        self.assertEqual(200, preview.status_code, preview.get_data(as_text=True))
        self.assertEqual(
            "local_issued_confirmation",
            preview.get_json()["invoices"][0]["verification_source"],
        )

    def test_official_statement_filters_one_contractor_and_reconciles_multiple_rounds(self):
        with server.db() as conn:
            self.add_invoice(conn, invoice_number="0000001", invoice_date="2026-09-03")
            self.add_invoice(conn, invoice_number="0000002", invoice_date="2026-09-04")
            self.add_invoice(
                conn, contractor="NT-B", invoice_number="0000003", invoice_date="2026-09-05",
            )
        preview = self.scope().get_json()
        self.assertEqual(len(preview["invoices"]), 2)
        response = self.client.get(
            "/api/outgoing-invoices/delivery-statement/NT-A"
            "?from=2026-09-01&to=2026-09-30&scope_id=" + preview["scope_id"]
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers["X-TDP-Invoice-Scope"], preview["scope_id"])
        self.assertIn("Bang_tong_hop_giao_nhan_NT-A", response.headers["Content-Disposition"])
        workbook = load_workbook(io.BytesIO(response.data), data_only=False, keep_links=False)
        try:
            detail = workbook["Bảng kê giao hàng"]
            reconcile = workbook["Đối chiếu hóa đơn"]
            self.assertEqual(detail.max_row, 17)
            self.assertEqual(reconcile.max_row, 5)
            self.assertEqual([detail.cell(row, 10).value for row in (11, 12)], [216, 216])
            self.assertEqual(detail.cell(13, 10).value, 432)
            self.assertEqual([reconcile.cell(row, 10).value for row in (4, 5)], [0, 0])
            self.assertTrue(all("0000003" not in str(cell.value) for row in reconcile for cell in row))
            self.assertNotIn("Q-008", detail["A2"].value)
        finally:
            workbook.close()
        with server.db() as conn:
            event = conn.execute(
                "SELECT metadata_json FROM audit_log "
                "WHERE event_type='outgoing.invoice_delivery_statement' "
                "ORDER BY id DESC LIMIT 1"
            ).fetchone()
        self.assertIsNotNone(event)
        self.assertEqual(json.loads(event["metadata_json"])["invoices"], 2)

    def test_source_total_mismatch_and_invoice_line_mismatch_fail_closed(self):
        with server.db() as conn:
            self.add_invoice(conn)
            self.add_source(conn, total=999)
        source_mismatch = self.scope()
        self.assertEqual(source_mismatch.status_code, 409)
        self.assertEqual(source_mismatch.get_json()["code"], "invoice_source_total_mismatch")
        self.setUp()
        with server.db() as conn:
            self.add_invoice(conn)
            conn.execute("UPDATE outgoing_invoice_lines SET amount=199")
        line_mismatch = self.scope()
        self.assertEqual(line_mismatch.status_code, 409)
        self.assertEqual(line_mismatch.get_json()["code"], "invoice_line_amount_mismatch")
        self.setUp()
        with server.db() as conn:
            self.add_invoice(conn)
            self.add_source(conn, buyer_tax_code="0200999999")
        buyer_mismatch = self.scope()
        self.assertEqual(buyer_mismatch.status_code, 409)
        self.assertEqual(buyer_mismatch.get_json()["code"], "invoice_source_buyer_mismatch")

    def test_snapshot_conflict_forces_separate_period(self):
        with server.db() as conn:
            self.add_invoice(conn, invoice_number="0000001")
            self.add_invoice(conn, invoice_number="0000002", snapshot_suffix=" KHÁC")
        conflict = self.scope()
        self.assertEqual(conflict.status_code, 409)
        self.assertEqual(conflict.get_json()["code"], "issued_invoice_snapshot_conflict")

    def test_scope_id_stops_stale_download_and_replacement_line_is_kept(self):
        with server.db() as conn:
            self.add_invoice(conn, product_code="SUB-1", draft_kind="substitution")
        preview = self.scope().get_json()
        with server.db() as conn:
            self.add_invoice(conn, invoice_number="0000002", invoice_date="2026-09-04")
        stale = self.client.get(
            "/api/export/invoice-payment-bundle/NT-A?from=2026-09-01&to=2026-09-30&scope_id=" +
            preview["scope_id"]
        )
        self.assertEqual(stale.status_code, 409)
        self.assertEqual(stale.get_json()["code"], "stale_invoice_payment_scope")
        stale_statement = self.client.get(
            "/api/outgoing-invoices/delivery-statement/NT-A"
            "?from=2026-09-01&to=2026-09-30&scope_id=" + preview["scope_id"]
        )
        self.assertEqual(stale_statement.status_code, 409)
        self.assertEqual(stale_statement.get_json()["code"], "stale_invoice_payment_scope")
        current = self.scope().get_json()
        self.assertEqual(len(current["invoices"]), 2)

    def test_ui_labels_and_build_include_official_payment_documents(self):
        root = Path(__file__).resolve().parent.parent
        script = (root / "tdp_system" / "static" / "app.js").read_text(encoding="utf-8")
        page = (root / "tdp_system" / "static" / "index.html").read_text(encoding="utf-8")
        build = (root / "BUILD_PORTABLE.ps1").read_text(encoding="utf-8")
        for marker in (
            "Hồ sơ đề nghị thanh toán từ hóa đơn đỏ",
            "Mẫu chính thức", "Tải Đề nghị thanh toán + bảng kê",
            "/api/outgoing-invoices/payment-scope/",
            "stale_invoice_payment_scope",
            "download-invoice-payment-control",
            "download-invoice-delivery-statement",
            "/api/outgoing-invoices/delivery-statement/",
        ):
            self.assertIn(marker, script if marker != "stale_invoice_payment_scope" else
                          (root / "tdp_system" / "contract_modules.py").read_text(encoding="utf-8"))
        self.assertIn("/static/app.js?v=20260906-16", page)
        self.assertIn("invoice_payment_scope.py", build)
        self.assertIn("--hidden-import invoice_payment_scope", build)
        self.assertIn("invoice_delivery_statement.py", build)
        self.assertIn("--hidden-import invoice_delivery_statement", build)
        self.assertIn("invoice_payment_documents.py", build)
        self.assertIn("--hidden-import invoice_payment_documents", build)


    def add_direct_source(self, conn, **kwargs):
        source_id = self.add_source(conn, **kwargs)
        conn.execute("INSERT OR REPLACE INTO outgoing_buyer_profiles(contractor,tax_code,address,updated_at) VALUES('NT-A','0200000001','Current address','now')")
        conn.execute('UPDATE outgoing_source_invoices SET raw_json=? WHERE id=?',
                     (json.dumps({'inv_buyerAddressLine': 'Địa chỉ mua'}), source_id))
        conn.execute("INSERT INTO outgoing_source_invoice_items(invoice_id,line_index,source_item_name,source_unit,qty,unit_price,amount,tax_rate) VALUES(?,1,'Hàng VAT','kg',2,100,200,'8%')", (source_id,))
        for key, value in {'company': 'CÔNG TY TĐP', 'company_tax_code': '0202265016',
                           'company_address': 'Địa chỉ TĐP', 'payment_requester': 'VŨ THỊ THỤY',
                           'payment_bank_name': 'Vietcombank', 'payment_bank_account': '1052787580'}.items():
            server.setting_set(conn, key, value)
        return source_id

    def test_synced_vat_without_local_draft_has_payment_pack_without_invented_delivery(self):
        with server.db() as conn:
            self.add_direct_source(conn)
        response = self.scope()
        self.assertEqual(response.status_code, 200, response.json)
        scope = response.json
        self.assertEqual(scope['totals']['total_amount'], 216)
        self.assertEqual(scope['statement_kind'], 'invoices')
        self.assertIsNone(scope['invoices'][0]['draft_id'])
        bundle = self.client.get('/api/export/invoice-payment-bundle/NT-A?from=2026-09-01&to=2026-09-30&scope_id=' + scope['scope_id'])
        self.assertEqual(bundle.status_code, 200, bundle.get_data()[:200])
        with zipfile.ZipFile(io.BytesIO(bundle.data)) as archive:
            statement = next(n for n in archive.namelist() if n.startswith('Bang_ke_hoa_don_VAT'))
            book = load_workbook(io.BytesIO(archive.read(statement)))
            self.assertEqual(book.worksheets[0]['G4'].value, 216)
            self.assertEqual(book.worksheets[0]['G5'].value, 216)
            self.assertEqual(book.worksheets[1]['G4'].value, 200)
            self.assertNotIn('Bảng kê giao hàng', book.sheetnames)
            book.close()
        with server.db() as conn:
            self.assertEqual(conn.execute('SELECT COUNT(*) FROM outgoing_invoice_drafts').fetchone()[0], 0)
            self.assertEqual(conn.execute('SELECT COUNT(*) FROM invoice_inventory_ledger').fetchone()[0], 0)
            conn.execute('UPDATE outgoing_source_invoice_items SET amount=201')
        self.assertEqual(self.client.get('/api/export/invoice-payment-bundle/NT-A?from=2026-09-01&to=2026-09-30&scope_id=' + scope['scope_id']).status_code, 409)

    def test_local_and_synced_duplicate_count_once_and_source_only_added(self):
        with server.db() as conn:
            self.add_invoice(conn)
            self.add_direct_source(conn)
            self.add_direct_source(conn, invoice_number='0000002')
        result = self.scope()
        self.assertEqual(result.status_code, 200, result.json)
        self.assertEqual(len(result.json['invoices']), 2)
        self.assertEqual(result.json['totals']['total_amount'], 432)

    def test_buyer_profile_can_be_configured_before_any_local_draft(self):
        body = {'legal_name': 'Công ty A', 'tax_code': '0209999999', 'address': 'Địa chỉ mua'}
        response = self.client.put('/api/outgoing-buyers/NT-A', json=body)
        self.assertEqual(response.status_code, 200)
        listing = self.client.get('/api/outgoing-invoices').get_json()
        self.assertEqual(listing['items'], [])
        self.assertEqual(listing['buyer_profiles']['NT-A']['tax_code'], body['tax_code'])
        self.assertEqual(listing['buyer_profiles']['NT-A']['address'], body['address'])
        with server.db() as conn:
            self.assertEqual(conn.execute('SELECT COUNT(*) FROM outgoing_source_invoices').fetchone()[0], 0)
            self.assertEqual(conn.execute('SELECT COUNT(*) FROM invoice_inventory_ledger').fetchone()[0], 0)

    def test_unissued_source_draft_does_not_inflate_or_block_issued_payment(self):
        with server.db() as conn:
            self.add_direct_source(conn)
            self.add_direct_source(conn, invoice_number='0000002', status_class='draft',
                                   sync_status='review_required', stock_status='blocked')
        response = self.scope()
        self.assertEqual(response.status_code, 200)
        scope = response.get_json()
        self.assertEqual(len(scope['invoices']), 1)
        self.assertEqual(scope['totals']['total_amount'], 216)
        self.assertEqual(scope['excluded_draft_count'], 1)

    def test_source_scope_blocks_uncertain_status_buyer_and_changed_totals(self):
        with server.db() as conn:
            source_id = self.add_direct_source(conn)
        for assignment, code in [
            ("source_status_class='cancelled'", 'invoice_source_not_payable'),
            ("source_status_class='issued', total_amount=999", 'invoice_source_total_mismatch'),
            ("total_amount=216, raw_json='{}'", 'invoice_source_buyer_incomplete'),
        ]:
            with server.db() as conn:
                conn.execute('UPDATE outgoing_source_invoices SET ' + assignment + ' WHERE id=?', (source_id,))
            result = self.scope()
            self.assertEqual(result.status_code, 409)
            self.assertEqual(result.json['code'], code)


if __name__ == "__main__":
    unittest.main()
