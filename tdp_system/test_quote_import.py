from __future__ import annotations

import io
import json
import tempfile
import unittest
import zipfile
from pathlib import Path

from openpyxl import Workbook, load_workbook

try:
    from . import quote_import, server
except ImportError:  # pragma: no cover - direct file invocation
    import quote_import
    import server


class QuoteImportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        cls.original_db_path = server.DB_PATH
        cls.original_data_dir = server.DATA_DIR
        cls.original_master_source = server.MASTER_SOURCE
        cls.original_testing = server.app.config.get("TESTING", False)
        server.DATA_DIR = Path(cls.temp.name) / "data"
        server.DATA_DIR.mkdir(parents=True, exist_ok=True)
        server.DB_PATH = server.DATA_DIR / "quotes.sqlite3"
        server.MASTER_SOURCE = Path(cls.temp.name) / "master-disabled.xlsx"
        server.app.config["TESTING"] = True
        server.init_database()
        cls.client = server.app.test_client()
        with server.db() as conn:
            conn.executemany(
                "INSERT INTO contractors(code,name,price_group,pricing_mode) VALUES(?,?,?,?)",
                [
                    ("C1", "Contractor 1", "C1", "group"),
                    ("C2", "Contractor 2", "C2", "group"),
                    ("ATV", "ATV", "ATV", "group"),
                    ("HATRAN", "HATRAN", "HATRAN", "group"),
                    ("BIADAUVOI", "BIADAUVOI", "BIADAUVOI", "group"),
                    ("NGUYENGIA", "NGUYENGIA", "NGUYENGIA", "group"),
                    ("SUPPY", "SUPPY", "SUPPY", "group"),
                    ("TOYOTA", "TOYOTA", "TOYOTA", "group"),
                    ("NHUAHAIPHONG", "NHUAHAIPHONG", "NHUAHAIPHONG", "group"),
                    ("GIANHAPTAY", "GIANHAPTAY", "GIANHAPTAY", "daily"),
                    ("YLKHAN", "YLKHAN", "YLKHAN", "daily"),
                ],
            )
            conn.execute("INSERT INTO kitchens(code,contractor,name) VALUES('K1','C1','Kitchen 1')")
            conn.execute("INSERT INTO suppliers(code,name) VALUES('S1','Supplier 1')")
            conn.executemany(
                """INSERT INTO products(
                       code,name,unit,tax,supplier,buy_price,purchase_list,seller,cccd
                   ) VALUES(?,?,?,?,?,?,0,'','')""",
                [
                    ("P1", "Product 1", "kg", "8%", "S1", 8_000),
                    ("P2", "Product 2", "kg", "8%", "S1", 7_000),
                    ("P3", "Product 3", "kg", "8%", "S1", 6_000),
                ],
            )
            conn.executemany(
                "INSERT INTO product_prices(product_code,price_group,price_text,price_value) VALUES(?,?,?,?)",
                [
                    ("P1", "C1", "9000", 9_000), ("P1", "C2", "9500", 9_500),
                    ("P2", "C1", "8000", 8_000), ("P2", "C2", "8500", 8_500),
                ],
            )

    @classmethod
    def tearDownClass(cls):
        with quote_import.QUOTE_IMPORT_LOCK:
            quote_import.PENDING_QUOTE_IMPORTS.clear()
        server.DB_PATH = cls.original_db_path
        server.DATA_DIR = cls.original_data_dir
        server.MASTER_SOURCE = cls.original_master_source
        server.app.config["TESTING"] = cls.original_testing
        cls.temp.cleanup()

    def setUp(self):
        with quote_import.QUOTE_IMPORT_LOCK:
            quote_import.PENDING_QUOTE_IMPORTS.clear()
        with server.db() as conn:
            conn.execute("DROP TRIGGER IF EXISTS fail_quote_price")
            conn.execute("DELETE FROM batches")
            conn.execute("DELETE FROM quote_version_prices")
            conn.execute("DELETE FROM quote_version_products")
            conn.execute("DELETE FROM quote_versions")
            conn.execute("DELETE FROM audit_log WHERE event_type='quote.import.confirm'")
            conn.execute("DELETE FROM settings WHERE key LIKE 'quote_recipient_%'")

    @staticmethod
    def workbook_bytes(*, p1_buy=10_000, p1_c1=12_000, p1_c2=13_000,
                       p2_buy=None, p2_c1=11_000, p2_c2="x") -> bytes:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "BÁO GIÁ"
        sheet.append(["BẢNG BÁO GIÁ TỔNG"])
        sheet.append([
            "STT", "MÃ THAM CHIẾU", "MÃ HÀNG", "TÊN THÀNH ĐẠT PHÁT",
            "Giá mua", "NCC", None, "bk", "Tên làm bảng kê", "ĐVT", "THUẾ",
            "C1", "C2", "Thêm", "=L2",
        ])
        sheet.append([None, 1, 2, 3, 4, 5, None, 6, 7, 8, 9, 10, 11, 12, 13])
        sheet.append([1, "P1S1", "P1", "Product 1", p1_buy, "S1", None, None, None, "kg", "8%", p1_c1, p1_c2, None, "=L4/E4"])
        sheet.append([2, "P2S1", "P2", "Product 2", p2_buy, "S1", None, None, None, "kg", "8%", p2_c1, p2_c2, None, "=L5/E5"])
        stream = io.BytesIO()
        workbook.save(stream)
        workbook.close()
        return stream.getvalue()

    @staticmethod
    def matrix_bytes(headers, rows) -> bytes:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "BÁO GIÁ"
        sheet.append(["BẢNG BÁO GIÁ TỔNG"])
        sheet.append([
            "STT", "MÃ THAM CHIẾU", "MÃ HÀNG", "TÊN THÀNH ĐẠT PHÁT",
            "Giá mua", "NCC", None, "bk", "Tên làm bảng kê", "ĐVT", "THUẾ",
            *headers, "Thêm", "=L2",
        ])
        sheet.append([None, 1, 2, 3, 4, 5, None, 6, 7, 8, 9])
        for index, row in enumerate(rows, 1):
            sheet.append([
                index, f"{row['code']}S1", row["code"], row.get("name", row["code"]),
                row.get("buy"), row.get("supplier", "S1"), None, None, None,
                row.get("unit", "kg"), row.get("tax", "8%"), *row["prices"], None, "=1+1",
            ])
        stream = io.BytesIO()
        workbook.save(stream)
        workbook.close()
        return stream.getvalue()

    def preview(self, period: str, payload: bytes | None = None, filename="quote.xlsx"):
        return self.client.post(
            "/api/quotes/import/preview",
            data={
                "effective_period": period,
                "file": (io.BytesIO(payload or self.workbook_bytes()), filename),
            },
            content_type="multipart/form-data",
        )

    def confirm(self, preview: dict, state_hash: str | None = None):
        return self.client.post(
            "/api/quotes/import/confirm",
            json={
                "token": preview["token"],
                "confirmed": True,
                "state_hash": preview["stateHash"] if state_hash is None else state_hash,
            },
        )

    def import_quote(self, period: str, payload: bytes | None = None):
        preview_response = self.preview(period, payload)
        self.assertEqual(preview_response.status_code, 200, preview_response.get_data(as_text=True))
        preview = preview_response.get_json()
        self.assertTrue(preview["canConfirm"], preview)
        response = self.confirm(preview)
        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        return response.get_json()

    def test_inline_single_and_all_quotes_use_confirmed_version_without_writes(self):
        self.import_quote('2026-09')
        with server.db() as conn:
            before=conn.execute('SELECT count(*) FROM audit_log').fetchone()[0]
        for kind in ('quote','quotes'):
            response=self.client.post('/api/documents/preview',json={'kind':kind,'contractor':'C1','period':'2026-09'})
            self.assertEqual(200,response.status_code,response.get_json())
            self.assertGreaterEqual(response.get_json()['sheet_count'],1)
            self.assertIn('12,000',str(response.get_json()['sheets']))
        with server.db() as conn:
            self.assertEqual(before,conn.execute('SELECT count(*) FROM audit_log').fetchone()[0])

    def create_order(self, work_date: str, *, product="P1", buy_price=99_999, sell_price=88_888):
        batch_response = self.client.post("/api/batches", json={"work_date": work_date})
        self.assertEqual(batch_response.status_code, 200, batch_response.get_data(as_text=True))
        batch_id = batch_response.get_json()["batch"]["id"]
        response = self.client.post("/api/orders", json={
            "batch_id": batch_id,
            "contractor": "C1",
            "kitchen": "K1",
            "product_code": product,
            "product_name": f"Product {product[-1]}",
            "qty": 2,
            "supplier": "S1",
            "buy_price": buy_price,
            "sell_price": sell_price,
            "tax": "8%",
        })
        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        order = next(item for item in response.get_json()["orders"] if item["id"] == response.get_json()["id"])
        return batch_id, order

    def test_two_periods_and_versions_drive_new_orders_without_rewriting_old_snapshots(self):
        august = self.import_quote("2026-08", self.workbook_bytes(p1_buy=10_000, p1_c1=12_000))
        september_v1 = self.import_quote("2026-09", self.workbook_bytes(p1_buy=11_000, p1_c1=14_000))
        self.assertEqual((august["versionNo"], september_v1["versionNo"]), (1, 1))

        august_batch, august_order = self.create_order("2026-08-25")
        september_batch, september_order = self.create_order("2026-09-01")
        self.assertEqual((august_order["buy_price"], august_order["sell_price"]), (10_000, 12_000))
        self.assertEqual((september_order["buy_price"], september_order["sell_price"]), (11_000, 14_000))

        september_v2 = self.import_quote(
            "2026-09", self.workbook_bytes(p1_buy=11_500, p1_c1=15_000),
        )
        self.assertEqual(september_v2["versionNo"], 2)
        _, new_september_order = self.create_order("2026-09-02")
        self.assertEqual((new_september_order["buy_price"], new_september_order["sell_price"]), (11_500, 15_000))
        with server.db() as conn:
            old_august = conn.execute(
                "SELECT buy_price,sell_price FROM orders WHERE batch_id=?", (august_batch,)
            ).fetchone()
            old_september = conn.execute(
                "SELECT buy_price,sell_price FROM orders WHERE batch_id=?", (september_batch,)
            ).fetchone()
        self.assertEqual(tuple(old_august), (10_000, 12_000))
        self.assertEqual(tuple(old_september), (11_000, 14_000))

    def test_replay_is_idempotent_and_versions_are_queryable(self):
        payload = self.workbook_bytes()
        first = self.import_quote("2026-09", payload)
        replay_preview = self.preview("2026-09", payload, "renamed-same-bytes.xlsx").get_json()
        self.assertTrue(replay_preview["replay"])
        self.assertEqual(replay_preview["proposedVersion"], 1)
        replay = self.confirm(replay_preview)
        self.assertEqual(replay.status_code, 200, replay.get_data(as_text=True))
        self.assertTrue(replay.get_json()["idempotent"])
        versions = self.client.get("/api/quotes/versions?period=2026-09").get_json()["items"]
        self.assertEqual(len(versions), 1)
        self.assertEqual((versions[0]["id"], versions[0]["version_no"]), (first["versionId"], 1))
        with server.db() as conn:
            self.assertEqual(conn.execute("SELECT COUNT(*) FROM quote_version_products").fetchone()[0], 2)
            self.assertEqual(conn.execute("SELECT COUNT(*) FROM quote_version_prices").fetchone()[0], 4)
            self.assertEqual(conn.execute(
                "SELECT COUNT(*) FROM audit_log WHERE event_type='quote.import.confirm'"
            ).fetchone()[0], 1)

    def test_stale_preview_and_concurrent_version_are_rejected_without_partial_write(self):
        preview = self.preview("2026-09").get_json()
        wrong = self.confirm(preview, "0" * 64)
        self.assertEqual(wrong.status_code, 409, wrong.get_data(as_text=True))
        self.assertEqual(wrong.get_json()["code"], "stale_preview")
        self.assertEqual(self.confirm(preview).status_code, 410)

        first = self.preview("2026-09", self.workbook_bytes(p1_c1=12_100)).get_json()
        competing = self.preview("2026-09", self.workbook_bytes(p1_c1=12_200)).get_json()
        self.assertEqual(self.confirm(first).status_code, 200)
        stale = self.confirm(competing)
        self.assertEqual(stale.status_code, 409, stale.get_data(as_text=True))
        self.assertEqual(stale.get_json()["code"], "stale_database")
        with server.db() as conn:
            self.assertEqual(conn.execute("SELECT COUNT(*) FROM quote_versions").fetchone()[0], 1)

    def test_confirmation_failure_rolls_back_version_lines_and_audit(self):
        with server.db() as conn:
            conn.execute(
                """CREATE TRIGGER fail_quote_price BEFORE INSERT ON quote_version_prices
                   WHEN NEW.product_code='P2'
                   BEGIN SELECT RAISE(ABORT,'forced quote failure'); END"""
            )
        preview = self.preview("2026-09").get_json()
        failed = self.confirm(preview)
        self.assertEqual(failed.status_code, 500, failed.get_data(as_text=True))
        with server.db() as conn:
            conn.execute("DROP TRIGGER fail_quote_price")
            counts = tuple(conn.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0] for table in (
                "quote_versions", "quote_version_products", "quote_version_prices",
            ))
            audit_count = conn.execute(
                "SELECT COUNT(*) FROM audit_log WHERE event_type='quote.import.confirm'"
            ).fetchone()[0]
        self.assertEqual(counts, (0, 0, 0))
        self.assertEqual(audit_count, 0)

    def test_buy_price_priority_and_blank_only_uses_actual_purchase_value(self):
        self.import_quote("2026-09", self.workbook_bytes(p1_buy=10_000, p2_buy=None))
        _, quoted = self.create_order("2026-09-01", product="P1", buy_price=99_999)
        self.assertEqual(quoted["buy_price"], 10_000)
        _, actual = self.create_order("2026-09-02", product="P2", buy_price=6_500)
        self.assertEqual(actual["buy_price"], 6_500)
        _, missing = self.create_order("2026-09-03", product="P2", buy_price=0)
        self.assertEqual(missing["buy_price"], 0)
        errors = missing["errors"] if isinstance(missing["errors"], list) else json.loads(missing["errors"])
        self.assertIn("Thiếu giá mua", errors)
        self.assertNotEqual(missing["buy_price"], 7_000)  # catalogue value is from another/no period

    def test_preview_validates_period_sheet_and_formula_price_without_writes(self):
        invalid_period = self.preview("2026-13")
        self.assertEqual(invalid_period.status_code, 400)
        self.assertEqual(invalid_period.get_json()["code"], "invalid_period")

        workbook = Workbook()
        workbook.active.title = "Orders"
        stream = io.BytesIO()
        workbook.save(stream)
        workbook.close()
        missing_sheet = self.preview("2026-09", stream.getvalue())
        self.assertEqual(missing_sheet.status_code, 400)
        self.assertEqual(missing_sheet.get_json()["code"], "quote_sheet_missing")

        formula = self.preview("2026-09", self.workbook_bytes(p1_c1="=1+1")).get_json()
        self.assertFalse(formula["canConfirm"])
        self.assertEqual(formula["counts"]["rows_with_errors"], 1)
        blocked = self.confirm(formula)
        self.assertEqual(blocked.status_code, 400)
        self.assertEqual(blocked.get_json()["code"], "preview_has_errors")
        with server.db() as conn:
            self.assertEqual(conn.execute("SELECT COUNT(*) FROM quote_versions").fetchone()[0], 0)

    def test_contractor_headers_not_positions_select_toyota_q_and_other_group(self):
        headers = ["ATV", "HATRAN", "BIADAUVOI", "NGUYENGIA", "SUPPY", "TOYOTA", "NHUAHAIPHONG"]
        payload = self.matrix_bytes(headers, [{
            "code": "P1", "name": "Product 1", "buy": 10_000,
            "prices": [11_000, 12_000, 13_000, 14_000, 15_000, 16_000, 17_000],
        }])
        preview = self.preview("2026-09", payload).get_json()
        toyota_mapping = next(item for item in preview["priceColumns"] if item["priceGroup"] == "TOYOTA")
        self.assertEqual(toyota_mapping, {
            "priceGroup": "TOYOTA", "sourceColumn": 17, "sourceHeader": "TOYOTA",
        })
        self.assertEqual(self.confirm(preview).status_code, 200)
        toyota = self.client.get("/api/quotes?contractor=TOYOTA&period=2026-09").get_json()
        suppy = self.client.get("/api/quotes?contractor=SUPPY&period=2026-09").get_json()
        self.assertEqual(toyota["items"][0]["sell_price"], 16_000)
        self.assertEqual(toyota["items"][0]["source_column"], 17)
        self.assertEqual(suppy["items"][0]["sell_price"], 15_000)
        self.assertEqual(suppy["items"][0]["source_column"], 16)

        reordered = self.matrix_bytes(["C2", "C1"], [{
            "code": "P1", "name": "Product 1", "buy": 10_000, "prices": [22_000, 11_000],
        }])
        self.import_quote("2026-10", reordered)
        c1 = self.client.get("/api/quotes?contractor=C1&period=2026-10").get_json()
        c2 = self.client.get("/api/quotes?contractor=C2&period=2026-10").get_json()
        self.assertEqual((c1["items"][0]["sell_price"], c2["items"][0]["sell_price"]), (11_000, 22_000))

    def test_x_blank_are_omitted_but_zero_remains_visible_and_exported(self):
        payload = self.matrix_bytes(["TOYOTA"], [
            {"code": "P1", "name": "Product 1", "buy": 10_000, "prices": ["x"]},
            {"code": "P2", "name": "Product 2", "buy": 9_000, "prices": [None]},
            {"code": "P3", "name": "Product 3", "buy": 8_000, "prices": [0]},
        ])
        self.import_quote("2026-09", payload)
        response = self.client.get("/api/quotes?contractor=TOYOTA&period=2026-09")
        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        data = response.get_json()
        self.assertEqual([item["product_code"] for item in data["items"]], ["P3"])
        self.assertEqual((data["items"][0]["sell_price"], data["items"][0]["price_state"]), (0, "zero"))
        self.assertTrue(data["items"][0]["exportable"])
        self.assertEqual((data["excludedCount"], data["outputCount"]), (2, 1))

        exported = self.client.get("/api/export/quote/TOYOTA?period=2026-09")
        self.assertEqual(exported.status_code, 200)
        workbook = load_workbook(io.BytesIO(exported.data), data_only=True, keep_links=False)
        try:
            self.assertEqual(workbook.sheetnames, ["all"])
            sheet = workbook["all"]
            self.assertEqual(sheet.cell(8, 1).value, "STT")
            self.assertEqual(sheet.cell(9, 3).value, "CHẤT TẨY RỬA-GIẤY CÁC LOẠI")
            self.assertEqual(sheet.cell(10, 2).value, "P3")
            self.assertEqual(sheet.cell(10, 5).value, 0)
            self.assertEqual(sheet.cell(10, 6).value, 0.08)
            note_row = next(cell.row for cell in sheet["A"] if cell.value == "Báo giá trên chưa bao gồm VAT!")
            self.assertIsNone(sheet.cell(note_row, 4).value)
            self.assertIn("TOYOTA · kỳ 2026-09 · phiên bản 1", workbook.properties.subject)
        finally:
            workbook.close()

    def test_toyota_export_requires_confirmed_period_and_names_source_version(self):
        missing = self.client.get("/api/export/quote/TOYOTA?period=2026-09")
        self.assertEqual(missing.status_code, 409, missing.get_data(as_text=True))
        self.assertEqual(missing.get_json()["code"], "quote_period_not_confirmed")

        self.import_quote("2026-09", self.matrix_bytes(["TOYOTA"], [{
            "code": "P1", "name": "Product 1", "buy": 10_000, "prices": [12_000],
        }]))
        exported = self.client.get("/api/export/quote/TOYOTA?period=2026-09")
        self.assertEqual(exported.status_code, 200)
        self.assertIn(
            "BAO_GIA_TOYOTA_T09-2026_V1.xlsx",
            exported.headers["Content-Disposition"],
        )

    def test_three_group_contractors_use_one_golden_engine_without_price_leakage(self):
        contractors = ["ATV", "TOYOTA", "NHUAHAIPHONG"]
        prices = [11_000, 22_000, 33_000]
        self.import_quote("2026-09", self.matrix_bytes(contractors, [{
            "code": "P1", "name": "Product 1", "buy": 10_000, "prices": prices,
        }]))
        expected_recipients = {
            "ATV": "CÔNG TY CỔ PHẦN SUẤT ĂN CÔNG NGHIỆP ATV",
            "TOYOTA": "CÔNG TY TNHH TOYOTA NANKAI HẢI PHÒNG",
            "NHUAHAIPHONG": "CÔNG TY CỔ PHẦN NHỰA VÀ CƠ KHÍ HẢI PHÒNG",
        }
        observed = {}
        for contractor, expected_price in zip(contractors, prices):
            with self.subTest(contractor=contractor):
                api_rows = self.client.get(
                    f"/api/quotes?contractor={contractor}&period=2026-09"
                ).get_json()
                self.assertEqual(api_rows["recipient"], expected_recipients[contractor])
                self.assertEqual(api_rows["items"][0]["sell_price"], expected_price)
                exported = self.client.get(
                    f"/api/export/quote/{contractor}?period=2026-09"
                )
                self.assertEqual(exported.status_code, 200, exported.status)
                self.assertIn(
                    f"BAO_GIA_{contractor}_T09-2026_V1.xlsx",
                    exported.headers["Content-Disposition"],
                )
                workbook = load_workbook(io.BytesIO(exported.data), data_only=False, keep_links=False)
                try:
                    sheet = workbook["all"]
                    self.assertEqual(sheet["A6"].value, f"KÍNH GỬI: {expected_recipients[contractor]}")
                    self.assertEqual(sheet["B10"].value, "P1")
                    observed[contractor] = sheet["E10"].value
                    note_row = next(
                        cell.row for cell in sheet["A"]
                        if cell.value == "Báo giá trên chưa bao gồm VAT!"
                    )
                    self.assertIsNone(sheet.cell(note_row, 4).value)
                    self.assertIn(
                        f"{contractor} · kỳ 2026-09 · phiên bản 1",
                        workbook.properties.subject,
                    )
                    self.assertFalse(any(
                        cell.data_type == "f"
                        for cells in sheet.iter_rows()
                        for cell in cells
                    ))
                finally:
                    workbook.close()
        self.assertEqual(observed, dict(zip(contractors, prices)))

        with server.db() as conn:
            conn.execute(
                "INSERT INTO settings(key,value) VALUES(?,?)",
                ("quote_recipient_ATV", "CÔNG TY ATV THEO HỢP ĐỒNG"),
            )
        overridden = self.client.get("/api/export/quote/ATV?period=2026-09")
        workbook = load_workbook(io.BytesIO(overridden.data), data_only=True, keep_links=False)
        try:
            self.assertEqual(workbook["all"]["A6"].value, "KÍNH GỬI: CÔNG TY ATV THEO HỢP ĐỒNG")
            self.assertEqual(workbook["all"]["E10"].value, 11_000)
        finally:
            workbook.close()

    def test_historical_and_all_downloads_keep_exact_version_and_contractor_prices(self):
        version_1 = self.import_quote("2026-09", self.matrix_bytes(["C1", "C2"], [{
            "code": "P1", "name": "Product 1", "buy": 10_000,
            "prices": [11_000, 21_000],
        }]))
        version_2 = self.import_quote("2026-09", self.matrix_bytes(["C1", "C2"], [{
            "code": "P1", "name": "Product 1", "buy": 10_500,
            "prices": [12_000, 22_000],
        }]))

        historical = self.client.get(
            f"/api/export/quote/C1?period=2026-09&version_id={version_1['versionId']}"
        )
        self.assertEqual(historical.status_code, 200)
        historical_book = load_workbook(
            io.BytesIO(historical.data), data_only=True, keep_links=False,
        )
        try:
            self.assertEqual(historical_book["all"]["E10"].value, 11_000)
            self.assertIn(
                "BAO_GIA_C1_T09-2026_V1.xlsx",
                historical.headers["Content-Disposition"],
            )
        finally:
            historical_book.close()

        current = self.client.get("/api/export/quote/C1?period=2026-09")
        current_book = load_workbook(io.BytesIO(current.data), data_only=True, keep_links=False)
        try:
            self.assertEqual(current_book["all"]["E10"].value, 12_000)
            self.assertIn("BAO_GIA_C1_T09-2026_V2.xlsx", current.headers["Content-Disposition"])
        finally:
            current_book.close()

        archived = self.client.get(
            f"/api/export/quotes/all?period=2026-09&version_id={version_1['versionId']}"
        )
        self.assertEqual(archived.status_code, 200)
        self.assertEqual(archived.mimetype, "application/zip")
        self.assertIn(
            "BAO_GIA_TAT_CA_T09-2026_V1.zip",
            archived.headers["Content-Disposition"],
        )
        with zipfile.ZipFile(io.BytesIO(archived.data)) as bundle:
            names = sorted(bundle.namelist())
            self.assertEqual(names, [
                "BAO_GIA_C1_T09-2026_V1.xlsx",
                "BAO_GIA_C2_T09-2026_V1.xlsx",
            ])
            self.assertTrue(all("/" not in name and "\\" not in name for name in names))
            observed = {}
            for name in names:
                book = load_workbook(io.BytesIO(bundle.read(name)), data_only=True, keep_links=False)
                try:
                    observed[name] = book["all"]["E10"].value
                finally:
                    book.close()
        self.assertEqual(observed, {
            "BAO_GIA_C1_T09-2026_V1.xlsx": 11_000,
            "BAO_GIA_C2_T09-2026_V1.xlsx": 21_000,
        })

        latest_bundle = self.client.get("/api/export/quotes/all?period=2026-09")
        self.assertEqual(latest_bundle.status_code, 200)
        self.assertIn(
            f"BAO_GIA_TAT_CA_T09-2026_V{version_2['versionNo']}.zip",
            latest_bundle.headers["Content-Disposition"],
        )
        with zipfile.ZipFile(io.BytesIO(latest_bundle.data)) as bundle:
            c1 = load_workbook(
                io.BytesIO(bundle.read("BAO_GIA_C1_T09-2026_V2.xlsx")),
                data_only=True,
                keep_links=False,
            )
            try:
                self.assertEqual(c1["all"]["E10"].value, 12_000)
            finally:
                c1.close()

        wrong_period = self.client.get(
            f"/api/export/quote/C1?period=2026-10&version_id={version_1['versionId']}"
        )
        self.assertEqual(wrong_period.status_code, 404)
        self.assertEqual(wrong_period.get_json()["code"], "quote_version_not_found")

        invalid_version = self.client.get(
            "/api/export/quotes/all?period=2026-09&version_id=khong-hop-le"
        )
        self.assertEqual(invalid_version.status_code, 400)
        self.assertEqual(invalid_version.get_json()["code"], "quote_version_invalid")

        other_period_batch = self.client.post(
            "/api/batches", json={"work_date": "2026-08-31"},
        ).get_json()["batch"]["id"]
        mismatch = self.client.get(
            "/api/export/quotes/all?period=2026-09"
            f"&version_id={version_1['versionId']}&batch_id={other_period_batch}"
        )
        self.assertEqual(mismatch.status_code, 409)
        self.assertEqual(mismatch.get_json()["code"], "daily_batch_period_mismatch")

    def test_daily_contractor_requires_batch_deduplicates_safe_rows_and_blocks_conflict(self):
        missing = self.client.get("/api/export/quote/GIANHAPTAY")
        self.assertEqual(missing.status_code, 409, missing.get_data(as_text=True))
        self.assertEqual(missing.get_json()["code"], "daily_batch_required")
        with server.db() as conn:
            cursor = conn.execute(
                "INSERT INTO batches(work_date,source_name,status,created_at) VALUES(?,?,?,?)",
                ("2026-09-03", "don-ngay.xlsx", "draft", "2026-09-03 08:00:00"),
            )
            batch_id = cursor.lastrowid
            rows = [
                ("K1", "P1", "Product 1", 12_000, 4),
                ("K2", "P1", "Product 1", 12_000, 8),
                ("K1", "P2", "Product 2", 0, 5),
            ]
            conn.executemany(
                """INSERT INTO orders(
                       batch_id,work_date,contractor,kitchen,product_code,product_name,
                       qty,actual_received,actual_delivered,unit,supplier,buy_price,
                       sell_price,tax,source_sheet,source_row,updated_at
                   ) VALUES(?,?,?,?,?,?,1,1,1,'kg','S1',10000,?,'8%','03.09.2026',?,?)""",
                [
                    (batch_id, "2026-09-03", "GIANHAPTAY", kitchen, code, name,
                     price, source_row, "2026-09-03 08:00:00")
                    for kitchen, code, name, price, source_row in rows
                ],
            )
        preview = self.client.get(f"/api/quotes?contractor=GIANHAPTAY&batch_id={batch_id}")
        self.assertEqual(preview.status_code, 200, preview.get_data(as_text=True))
        payload = preview.get_json()
        self.assertEqual(payload["dailySource"]["batch_id"], batch_id)
        self.assertEqual([item["product_code"] for item in payload["items"]], ["P1", "P2"])
        p1 = next(item for item in payload["items"] if item["product_code"] == "P1")
        self.assertEqual((p1["sell_price"], p1["source_rows"], p1["duplicate_count"]), (12_000, [4, 8], 1))
        exported = self.client.get(
            f"/api/export/quote/GIANHAPTAY?batch_id={batch_id}"
        )
        self.assertEqual(exported.status_code, 200, exported.status)
        self.assertIn(
            f"BAO_GIA_GIANHAPTAY_2026-09-03_PHIEN_{batch_id}.xlsx",
            exported.headers["Content-Disposition"],
        )
        workbook = load_workbook(io.BytesIO(exported.data), data_only=True, keep_links=False)
        try:
            sheet = workbook["all"]
            self.assertEqual(sheet["A5"].value, "BẢNG BÁO GIÁ NGÀY 03/09/2026")
            self.assertEqual(sheet["A6"].value, "KÍNH GỬI: GIANHAPTAY")
            self.assertEqual(
                [cell.value for cell in sheet["B"] if cell.value], ["MÃ", "P1", "P2"],
            )
        finally:
            workbook.close()

        with server.db() as conn:
            conn.execute(
                """INSERT INTO orders(
                       batch_id,work_date,contractor,kitchen,product_code,product_name,
                       qty,actual_received,actual_delivered,unit,supplier,buy_price,
                       sell_price,tax,source_sheet,source_row,updated_at
                   ) VALUES(?,?,?,?,?,?,1,1,1,'kg','S1',10000,13000,'8%','03.09.2026',9,?)""",
                (batch_id, "2026-09-03", "GIANHAPTAY", "K3", "P1", "Product 1", "2026-09-03 08:00:00"),
            )
        conflicted = self.client.get(
            f"/api/quotes?contractor=GIANHAPTAY&batch_id={batch_id}"
        ).get_json()
        self.assertEqual(conflicted["conflicts"][0]["product_code"], "P1")
        self.assertIn("giá bán theo ngày khác nhau", conflicted["conflicts"][0]["reasons"])
        blocked = self.client.get(
            f"/api/export/quote/GIANHAPTAY?batch_id={batch_id}"
        )
        self.assertEqual(blocked.status_code, 409, blocked.get_data(as_text=True))
        self.assertEqual(blocked.get_json()["code"], "quote_duplicate_conflict")

        unknown = self.client.get("/api/export/quote/KHONGTONTAI?period=2026-09")
        self.assertEqual(unknown.status_code, 404)
        self.assertEqual(unknown.get_json()["code"], "contractor_not_found")

    def test_same_duplicate_deduplicates_but_different_price_or_buy_blocks_confirm(self):
        same = self.matrix_bytes(["TOYOTA"], [
            {"code": "P1", "name": "Product 1", "buy": 10_000, "supplier": "S1", "prices": [12_000]},
            {"code": "P1", "name": "Product 1", "buy": 10_000, "supplier": "OTHER", "prices": [12_000]},
        ])
        same_preview = self.preview("2026-09", same).get_json()
        self.assertEqual(
            (same_preview["counts"]["duplicate_codes"], same_preview["counts"]["safe_duplicate_codes"], same_preview["counts"]["conflicts"]),
            (1, 1, 0),
        )
        self.assertTrue(same_preview["canConfirm"])
        self.assertEqual(self.confirm(same_preview).status_code, 200)
        rows = self.client.get("/api/quotes?contractor=TOYOTA&period=2026-09").get_json()["items"]
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["source_rows"], [4, 5])

        with server.db() as conn:
            conn.execute("DELETE FROM quote_version_prices")
            conn.execute("DELETE FROM quote_version_products")
            conn.execute("DELETE FROM quote_versions")
        different = self.matrix_bytes(["TOYOTA"], [
            {"code": "P1", "name": "Product 1", "buy": 10_000, "prices": [12_000]},
            {"code": "P1", "name": "Product 1", "buy": 11_000, "prices": [13_000]},
        ])
        conflict = self.preview("2026-09", different).get_json()
        self.assertFalse(conflict["canConfirm"])
        self.assertEqual((conflict["counts"]["conflict_codes"], conflict["counts"]["conflicts"]), (1, 1))
        self.assertEqual(conflict["conflicts"][0]["sourceRows"], [4, 5])
        self.assertIn("giá/trạng thái khác nhau", conflict["conflicts"][0]["reasons"])
        blocked = self.confirm(conflict)
        self.assertEqual(blocked.status_code, 400, blocked.get_data(as_text=True))
        with server.db() as conn:
            self.assertEqual(conn.execute("SELECT COUNT(*) FROM quote_versions").fetchone()[0], 0)

        buy_conflict = self.matrix_bytes(["TOYOTA"], [
            {"code": "P1", "name": "Product 1", "buy": 10_000, "prices": [12_000]},
            {"code": "P1", "name": "Product 1", "buy": 11_000, "prices": [12_000]},
        ])
        buy_preview = self.preview("2026-09", buy_conflict).get_json()
        self.assertFalse(buy_preview["canConfirm"])
        self.assertIn("giá mua khác nhau", buy_preview["conflicts"][0]["reasons"])


if __name__ == "__main__":
    unittest.main()
