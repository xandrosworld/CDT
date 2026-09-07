import json
import sqlite3
import unittest
from contextlib import contextmanager
from copy import deepcopy
from unittest.mock import patch

from flask import Flask
from openpyxl import load_workbook

from .contract_modules import upsert_msmi_invoice
from .invoice_expenses import ExpenseError, expense_token, set_expenses, register_expense_routes
from .invoice_mapping import save_mapping, apply_saved_mappings
from .invoice_receipt import create_input_receipt, InvoiceReceiptError
from .invoice_workbench_listing import invoice_range_payload, range_workbook
from .test_invoice_input_sync import init_test_database, now_iso
from .test_msmi_sync import remote_invoice


class InvoiceExpenseTests(unittest.TestCase):
    def setUp(self):
        self.conn=sqlite3.connect(':memory:')
        init_test_database(self.conn)
        self.conn.execute("INSERT INTO products(code,name,unit) VALUES('P','Hàng thử','kg')")
        self.remote=remote_invoice(1)
        self.remote['hdhhdvu']=[dict(ma=code,ten='Hàng '+code,dvtinh='kg',sluong=2,dgia=5000,thtien=10000,tsuat='8',tchat='1') for code in ('A','B')]
        self.remote.update(tgtcthue=20000,tgtthue=1600,tgtttbso=21600)
        self.invoice_id=upsert_msmi_invoice(self.conn,self.remote,'INPUT_ELECTRONIC_INVOICE','TDP',now_iso())[0]
        self.ids=[r[0] for r in self.conn.execute('SELECT id FROM msmi_invoice_items ORDER BY line_index')]

    def tearDown(self):
        self.conn.close()

    def change(self, expense=True, ids=None, **extra):
        return set_expenses(self.conn,**dict(tenant='TDP',invoice_id=self.invoice_id,
            expense=expense,item_ids=ids,expected=expense_token(self.conn,self.invoice_id),now=now_iso(),**extra))

    def data(self, **kw):
        return invoice_range_payload(self.conn,tenant='TDP',invoice_type='input',date_from='2026-08-01',date_to='2026-08-31',**kw)

    def source(self):
        return [tuple(r) for r in self.conn.execute('SELECT source_item_code,source_item_name,source_unit,qty,unit_price,amount FROM msmi_invoice_items ORDER BY line_index')]

    def sync(self):
        upsert_msmi_invoice(self.conn,self.remote,'INPUT_ELECTRONIC_INVOICE','TDP',now_iso())
        apply_saved_mappings(self.conn,'input',self.invoice_id)

    def test_whole_invoice_keeps_money_and_removes_code_requirement_and_is_reversible(self):
        before=self.source()
        self.assertEqual(2,self.change()['changed_lines'])
        data=self.data()
        self.assertEqual('not_inventory',data['items'][0]['receipt_status'])
        self.assertEqual(0,data['totals']['issue_count'])
        self.assertEqual(2,len(self.data(line_filter='expense')['lines']))
        self.assertEqual(20000,data['totals']['line_amount'])
        self.assertEqual(21600,data['totals']['invoice_amount'])
        self.assertEqual(before,self.source())
        wb=load_workbook(range_workbook(data))
        self.assertEqual('Chi phí · không nhập kho',wb.active.cell(3,13).value)
        saved=list(self.conn.iterdump())
        self.assertEqual(0,self.change()['changed_lines'])
        self.assertEqual(saved,list(self.conn.iterdump()))
        with self.assertRaises(InvoiceReceiptError):
            create_input_receipt(self.conn,self.invoice_id,now_iso)
        self.change(False)
        self.assertEqual(2,self.data()['totals']['issue_count'])
        self.assertEqual(before,self.source())

    def test_mixed_invoice_posts_only_goods_and_resync_keeps_posted_source_valid(self):
        save_mapping(self.conn,direction='input',item_id=self.ids[0],product_code='P',now_iso=now_iso)
        self.change(ids=[self.ids[1]])
        self.assertEqual('ready',self.data()['items'][0]['receipt_status'])
        result=create_input_receipt(self.conn,self.invoice_id,now_iso)
        self.assertEqual(1,result['inventory_lines'])
        ledger=[tuple(r) for r in self.conn.execute('SELECT * FROM invoice_inventory_ledger')]
        self.sync()
        self.assertEqual('synced',self.conn.execute('SELECT sync_status FROM msmi_invoices').fetchone()[0])
        self.assertEqual(ledger,[tuple(r) for r in self.conn.execute('SELECT * FROM invoice_inventory_ledger')])
        with self.assertRaises(ExpenseError):
            self.change(False)

    def test_resync_keeps_expense_and_does_not_reapply_mapping(self):
        save_mapping(self.conn,direction='input',item_id=self.ids[0],product_code='P',now_iso=now_iso)
        self.change(ids=[self.ids[0]])
        self.sync()
        line=self.conn.execute('SELECT * FROM msmi_invoice_items WHERE line_index=1').fetchone()
        self.assertEqual((0,'','not_inventory'),(line['inventory_eligible'],line['product_code'],line['mapping_status']))
        self.change(False,ids=[line['id']])
        self.assertEqual('P',self.conn.execute('SELECT product_code FROM msmi_invoice_items WHERE line_index=1').fetchone()[0])

    def test_changed_source_requires_review_instead_of_silent_stock_or_expense(self):
        self.change()
        self.remote['hdhhdvu'][0]['ten']='Nguồn đã sửa tên'
        self.sync()
        data=self.data()
        self.assertEqual(1,data['items'][0]['expense_review_count'])
        self.assertGreater(data['totals']['issue_count'],0)
        with self.assertRaises(InvoiceReceiptError):
            create_input_receipt(self.conn,self.invoice_id,now_iso)
        self.change()
        self.assertEqual(0,self.data()['totals']['issue_count'])

    def test_removed_source_line_is_not_silently_forgotten(self):
        self.change()
        self.remote['hdhhdvu'].pop()
        self.remote.update(tgtcthue=10000,tgtthue=800,tgtttbso=10800)
        self.sync()
        self.assertEqual(1,self.data()['items'][0]['expense_review_count'])
        self.change(False)
        self.assertEqual(0,self.data()['items'][0]['expense_review_count'])
        self.assertEqual(0,self.conn.execute('SELECT COUNT(*) FROM invoice_input_expense_choices').fetchone()[0])

    def test_tenant_source_and_stale_guards(self):
        token=expense_token(self.conn,self.invoice_id)
        self.change(ids=[self.ids[0]])
        for tenant,expected in [('OTHER',expense_token(self.conn,self.invoice_id)),('TDP',token)]:
            with self.assertRaises(ExpenseError):
                set_expenses(self.conn,tenant=tenant,invoice_id=self.invoice_id,expense=True,item_ids=None,expected=expected,now=now_iso())
        self.conn.execute("UPDATE msmi_invoices SET sync_status='review_required'")
        with self.assertRaises(ExpenseError):
            self.change()

    def test_route_rolls_back_after_partial_write(self):
        self.conn.commit()
        @contextmanager
        def db():
            with self.conn:
                yield self.conn
        app=Flask(__name__);app.config['TESTING']=True
        register_expense_routes(app,{'db':db,'setting_get':lambda c,k,d:d,'now_iso':now_iso})
        client=app.test_client()
        before=list(self.conn.iterdump())
        body={'expense':True,'expected':expense_token(self.conn,self.invoice_id)}
        path=f'/api/invoice-workbench/input-invoices/{self.invoice_id}/expense'
        with patch('tdp_system.invoice_expenses.restore_expenses',side_effect=RuntimeError('test rollback')):
            with self.assertRaises(RuntimeError):
                client.post(path,json=body)
        self.assertEqual(before,list(self.conn.iterdump()))
        self.assertEqual(200,client.post(path,json=body).status_code)
        self.assertEqual(409,client.post(path,json=body).status_code)

    def test_expense_splits_display_group_and_undo_keeps_original_mapping(self):
        from .invoice_line_groups import preview_group,create_group
        for item in self.ids:
            save_mapping(self.conn,direction='input',item_id=item,product_code='P',now_iso=now_iso)
        preview=preview_group(self.conn,'TDP',self.ids)
        create_group(self.conn,'TDP',self.ids,preview['token'],now_iso())
        self.change(ids=[self.ids[1]])
        self.assertEqual(0,self.conn.execute('SELECT SUM(active) FROM invoice_input_line_groups').fetchone()[0])
        self.assertEqual('P',self.conn.execute('SELECT product_code FROM msmi_invoice_items WHERE id=?',(self.ids[0],)).fetchone()[0])
        self.change(False,ids=[self.ids[1]])
        self.assertEqual('ready',self.data()['items'][0]['receipt_status'])

    def test_existing_stock_blocks_classification_even_with_wrong_header_status(self):
        for item in self.ids:
            save_mapping(self.conn,direction='input',item_id=item,product_code='P',now_iso=now_iso)
        create_input_receipt(self.conn,self.invoice_id,now_iso)
        self.conn.execute("UPDATE msmi_invoices SET receipt_status='ready'")
        with self.assertRaises(ExpenseError):self.change()


if __name__=='__main__':
    unittest.main()
