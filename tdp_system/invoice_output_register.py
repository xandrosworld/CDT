"""Source invoice amounts and quantities, independent of stock posting."""

try:
    from .document_preview import white_print_style
except ImportError:
    from document_preview import white_print_style

from decimal import Decimal
from io import BytesIO
try:
    from .invoice_line_tax import complete_sum
except ImportError:
    from invoice_line_tax import complete_sum


def output_register_summary(invoices, date_from, date_to):
    period = [i for i in invoices if date_from <= i['invoice_date'] <= date_to]
    def total(field):
        return float(sum((Decimal(str(i.get(field) or 0)) for i in period), Decimal(0)))
    lines = [line for invoice in period for line in invoice.get('items', [])]
    goods = [line for line in lines if line.get('inventory_eligible')]
    mapped = sum(bool(line.get('product_code')) and line.get('mapping_status') == 'mapped'
                 and not line.get('identity_warning') for line in goods)
    detail = sum((Decimal(str(line.get('amount') or 0)) for line in lines), Decimal(0))
    subtotal = sum((Decimal(str(i.get('subtotal') or 0)) for i in period), Decimal(0))
    return {'date_from': date_from, 'date_to': date_to, 'invoice_count': len(period),
            'line_count': len(lines), 'goods_line_count': len(goods), 'mapped_line_count': mapped,
            'unmapped_line_count': len(goods) - mapped, 'subtotal': float(subtotal),
            'tax_amount': total('tax_amount'), 'total_amount': total('total_amount'),
            'detail_amount': float(detail), 'detail_difference': float(subtotal - detail),
            'detail_tax_amount': complete_sum(lines, 'line_tax_amount'),
            'detail_total_amount': complete_sum(lines, 'amount_with_tax'),
            'missing_tax_lines': sum(line.get('line_tax_amount') is None for line in lines)}


def output_sales_workbook(payload):
    """One sheet of all source lines, with header totals counted once per invoice."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    assert payload['direction'] == 'output' and payload['scope'] == 'period'
    assert payload['status'] == 'all' and payload['line_filter'] == 'all'
    summary = payload['output_summary']
    wb = Workbook(); ws = wb.active; ws.title = 'Dau ra M-Invoice'
    ws.append(['HÓA ĐƠN ĐẦU RA · M-INVOICE']); ws.merge_cells('A1:N1')
    ws.append([f"{payload['date_from']} → {payload['date_to']} · {summary['invoice_count']} hóa đơn · {summary['line_count']} dòng"])
    ws.merge_cells('A2:N2')
    for label, field in [('TIỀN HÀNG TRÊN HÓA ĐƠN (CHƯA THUẾ)', 'subtotal'),
                         ('TIỀN THUẾ TRÊN HÓA ĐƠN', 'tax_amount'),
                         ('TỔNG THANH TOÁN TRÊN HÓA ĐƠN', 'total_amount')]:
        ws.append([label, None, None, None, None, None, None, None, summary[field]])
        ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=8)
    ws.append(['Giữ nguyên lượng, đơn vị và tiền nguồn. Tổng hóa đơn tính mỗi hóa đơn một lần; gồm cả hóa đơn chưa ghi kho.'])
    ws.merge_cells('A6:N6')
    ws.append(['Ngày', 'Ký hiệu / Số HĐ', 'Khách hàng', 'Mã hàng', 'Tên trên hóa đơn',
               'ĐVT', 'Số lượng', 'Đơn giá bán', 'Tiền dòng (chưa thuế)', 'Khớp mã', 'Thuế suất', 'Tiền thuế', 'Tiền gồm thuế', 'Đối chiếu thuế'])
    headers = {invoice['id']: invoice for invoice in payload['items']}
    for line in sorted(payload['lines'], key=lambda r: (headers[r['invoice_id']]['invoice_date'], r['invoice_id'], r.get('line_index') or 0)):
        h = headers[line['invoice_id']]
        mapped = line.get('product_code') and line.get('mapping_status') == 'mapped' and not line.get('identity_warning')
        status = ('Thiếu chi tiết nguồn' if not line.get('id') else 'Đã khớp mã' if mapped
                  else 'Không cần ghép mã' if not line.get('inventory_eligible') else 'Cần khớp mã')
        ws.append([h['invoice_date'], h['invoice_series'] + ' / ' + h['invoice_number'], h.get('buyer_name', ''),
                   line.get('product_code', ''), line.get('source_item_name', ''), line.get('source_unit', ''),
                   line.get('qty'), line.get('unit_price'), line.get('amount'), status,
                   line.get('tax_rate'), line.get('line_tax_amount'), line.get('amount_with_tax'), line.get('tax_note', 'Chưa có tiền thuế nguồn')])
    last_detail = ws.max_row
    ws.append(['CỘNG TIỀN CÁC DÒNG CHI TIẾT', None, None, None, None, None, None, None, summary['detail_amount'], None, None,
               summary['detail_tax_amount'], summary['detail_total_amount'],
               'Chưa đủ tiền thuế từng dòng' if summary['missing_tax_lines'] else ''])
    ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=8)
    tax_difference = summary['tax_amount'] - summary['detail_tax_amount'] if summary['detail_tax_amount'] is not None else None
    total_difference = summary['total_amount'] - summary['detail_total_amount'] if summary['detail_total_amount'] is not None else None
    if (abs(summary['detail_difference']) > 1 or tax_difference is not None and abs(tax_difference) > 1
            or total_difference is not None and abs(total_difference) > 1):
        ws.append(['TỔNG HÓA ĐƠN − CỘNG CHI TIẾT', None, None, None, None, None, None, None, summary['detail_difference'], None, None, tax_difference, total_difference])
        ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=8)
        for review in payload.get('amount_reviews', []):
            parts = ['%s: %s đ' % (dict(subtotal='Tiền hàng', tax='Thuế', total='Thanh toán')[r['kind']], format(r['difference'], ',.0f'))
                     for r in review['comparisons'] if abs(r['difference']) > 1]
            ws.append([f"HĐ {review['invoice_series']} / {review['invoice_number']}: tổng hóa đơn trừ chi tiết · " + '; '.join(parts)])
            ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=14)
    ws.append(['TỔNG THEO M-INVOICE', None, None, None, None, None, None, None,
               summary['subtotal'], None, None, summary['tax_amount'], summary['total_amount']])
    ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=8)
    widths = [14, 24, 34, 16, 48, 12, 16, 19, 24, 20, 12, 19, 22, 44]
    edge = Side(style='thin', color='D6DEE5')
    for cells in ws:
        for cell in cells:
            if isinstance(cell.value, str): cell.data_type = 's'
            cell.font = Font(name='Arial', size=10, bold=cell.row <= 7 or cell.row > last_detail,
                             color='FFFFFF' if cell.row == 7 else '17324A')
            cell.alignment = Alignment(vertical='center', wrap_text=True)
            cell.border = Border(bottom=edge)
            if cell.row == 7: cell.fill = PatternFill('solid', fgColor='17324A')
            if cell.column in (7, 8, 9, 12, 13): cell.number_format = '#,##0.######'
        ws.row_dimensions[cells[0].row].height = 42 if 8 <= cells[0].row <= last_detail else 32
    for index, width in enumerate(widths, 1): ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = 'E8'; ws.auto_filter.ref = f'A7:N{max(7, last_detail)}'
    ws.print_title_rows = '1:7'; ws.print_area = f'A1:N{ws.max_row}'
    ws.page_setup.orientation = 'landscape'; ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    output = BytesIO(); white_print_style(wb).save(output); output.seek(0)
    return output
