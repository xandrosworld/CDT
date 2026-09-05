import hashlib
import io
import sqlite3
import unittest
import zipfile

from docx import Document
from openpyxl import load_workbook

try:
    from .xcom_payment_documents import (
        XcomPaymentError,
        assign_payment_scope,
        build_meal_payment_summary,
        build_simple_payment_request_docx,
        consume_payment_preview,
        create_payment_document,
        create_payment_preview,
        init_xcom_payment_schema,
        upsert_meal_tariff,
        upsert_payment_profile,
    )
except ImportError:
    from xcom_payment_documents import (
        XcomPaymentError,
        assign_payment_scope,
        build_meal_payment_summary,
        build_simple_payment_request_docx,
        consume_payment_preview,
        create_payment_document,
        create_payment_preview,
        init_xcom_payment_schema,
        upsert_meal_tariff,
        upsert_payment_profile,
    )


FIXED_NOW = lambda: "2026-08-31T20:00:00"


class XcomPaymentDocumentTests(unittest.TestCase):
    def setUp(self):
        self.conn = sqlite3.connect(":memory:")
        self.conn.row_factory = sqlite3.Row
        self.conn.executescript(
            """
            PRAGMA foreign_keys=ON;
            CREATE TABLE kitchen_units(
                kitchen_code TEXT PRIMARY KEY,
                unit_code TEXT NOT NULL,
                updated_at TEXT NOT NULL
            );
            CREATE TABLE meal_attendance(
                work_date TEXT NOT NULL,
                kitchen TEXT NOT NULL,
                shift TEXT NOT NULL,
                actual_count REAL NOT NULL DEFAULT 0,
                ordered_count REAL NOT NULL DEFAULT 0,
                source_type TEXT NOT NULL DEFAULT 'MONTHLY_WORKBOOK',
                source_file TEXT,
                source_sheet TEXT,
                source_column TEXT,
                updated_at TEXT NOT NULL,
                PRIMARY KEY(work_date,kitchen,shift)
            );
            """
        )
        init_xcom_payment_schema(self.conn)
        self.simple_profile = {
            "profile_code": "SIMPLE",
            "document_type": "MEAL_SIMPLE",
            "issuer_name": "CÔNG TY TNHH DỊCH VỤ HÀ TRÂN",
            "recipient_name": "CÔNG TY KHÁCH HÀNG",
            "beneficiary_name": "CÔNG TY TNHH DỊCH VỤ HÀ TRÂN",
            "bank_account": "0031000318858",
            "bank_name": "Vietcombank CN Hải Phòng",
            "requester": "HOÀNG THỊ TUYẾT",
            "vat_rate": "10",
        }
        self.bot_profile = {
            **self.simple_profile,
            "profile_code": "BOT",
            "document_type": "MEAL_BOT_BUNDLE",
            "issuer_tax_code": "0201853470",
            "issuer_address": "Thuận Thiên, Kiến Thụy, Hải Phòng",
            "recipient_name": "CÔNG TY CỔ PHẦN BOT CẦU BẠCH ĐẰNG",
            "recipient_tax_code": "0201650135",
            "recipient_address": "Hải Phòng",
            "contract_no": "01062025/HĐDV/BOT-HATRAN",
            "contract_date": "2025-05-27",
            "seller_signer_name": "VŨ ANH TUẤN",
            "seller_signer_title": "GIÁM ĐỐC",
            "buyer_signer_name": "VĂN THÀNH TÂM",
            "buyer_signer_title": "PHÓ TỔNG GIÁM ĐỐC",
        }

    def tearDown(self):
        self.conn.close()

    def _insert_scope_data(self, profile, scope_type="XCOM", scope_code="UNIT-A"):
        upsert_payment_profile(self.conn, profile, now_iso=FIXED_NOW)
        assign_payment_scope(
            self.conn, profile["profile_code"], scope_type, scope_code, now_iso=FIXED_NOW
        )
        self.conn.execute(
            "INSERT OR REPLACE INTO kitchen_units VALUES('BEP-A','UNIT-A',?)", (FIXED_NOW(),)
        )

    def _insert_attendance(
        self,
        work_date="2026-08-01",
        kitchen="BEP-A",
        shift="Sáng",
        actual=1,
        ordered=999,
    ):
        self.conn.execute(
            """INSERT INTO meal_attendance(
                   work_date,kitchen,shift,actual_count,ordered_count,updated_at
               ) VALUES(?,?,?,?,?,?)""",
            (work_date, kitchen, shift, actual, ordered, FIXED_NOW()),
        )

    def test_actual_count_exact_period_and_half_up_vnd(self):
        self._insert_scope_data(self.simple_profile)
        self._insert_attendance(actual=1, ordered=999)
        upsert_meal_tariff(self.conn, "SIMPLE", "2026-08", "Sáng", "105", now_iso=FIXED_NOW)

        result = build_meal_payment_summary(self.conn, "SIMPLE", "2026-08-01", "2026-08-31")

        self.assertEqual(result["actual_count"], 1)
        self.assertEqual(result["subtotal"], 105)
        self.assertEqual(result["vat_amount"], 11)  # 10.5 rounds HALF_UP, never banker's-rounds.
        self.assertEqual(result["total"], 116)
        self.assertNotEqual(result["actual_count"], 999)

    def test_profile_scope_and_tariff_upserts_are_idempotent(self):
        first_profile = upsert_payment_profile(self.conn, self.simple_profile, now_iso=FIXED_NOW)
        second_profile = upsert_payment_profile(self.conn, self.simple_profile, now_iso=FIXED_NOW)
        first_scope = assign_payment_scope(
            self.conn, "SIMPLE", "XCOM", "UNIT-A", now_iso=FIXED_NOW
        )
        second_scope = assign_payment_scope(
            self.conn, "SIMPLE", "XCOM", "UNIT-A", now_iso=FIXED_NOW
        )
        first_tariff = upsert_meal_tariff(
            self.conn, "SIMPLE", "2026-08", "Sáng", "25000", now_iso=FIXED_NOW
        )
        second_tariff = upsert_meal_tariff(
            self.conn, "SIMPLE", "2026-08", "Sáng", "25000.4", now_iso=FIXED_NOW
        )

        self.assertTrue(first_profile["created"])
        self.assertFalse(second_profile["created"])
        self.assertFalse(second_profile["changed"])
        self.assertTrue(first_scope["changed"])
        self.assertFalse(second_scope["changed"])
        self.assertTrue(first_tariff["changed"])
        self.assertFalse(second_tariff["changed"])

    def test_missing_exact_period_tariff_is_blocked(self):
        self._insert_scope_data(self.simple_profile)
        self._insert_attendance(work_date="2026-09-01")
        upsert_meal_tariff(self.conn, "SIMPLE", "2026-08", "Sáng", 25_000, now_iso=FIXED_NOW)

        with self.assertRaises(XcomPaymentError) as caught:
            build_meal_payment_summary(self.conn, "SIMPLE", "2026-09-01", "2026-09-30")

        self.assertEqual(caught.exception.code, "missing_tariff")
        self.assertEqual(caught.exception.details, ["2026-09 / Sáng"])

    def test_fractional_actual_meal_count_is_blocked_instead_of_rounded(self):
        self._insert_scope_data(self.simple_profile)
        self._insert_attendance(actual=0.5)
        upsert_meal_tariff(self.conn, "SIMPLE", "2026-08", "Sáng", 1, now_iso=FIXED_NOW)

        with self.assertRaises(XcomPaymentError) as caught:
            build_meal_payment_summary(self.conn, "SIMPLE", "2026-08-01", "2026-08-31")

        self.assertEqual(caught.exception.code, "invalid_actual_count")
        self.assertEqual(caught.exception.details["actual_count"], "0.5")

    def test_payment_preview_token_expires_after_ttl(self):
        self._insert_scope_data(self.simple_profile)
        self._insert_attendance()
        upsert_meal_tariff(self.conn, "SIMPLE", "2026-08", "Sáng", 25_000, now_iso=FIXED_NOW)
        preview = create_payment_preview(
            self.conn,
            "SIMPLE",
            "2026-08-01",
            "2026-08-31",
            "2026-08-31",
            now_iso=FIXED_NOW,
            token_factory=lambda: "fixed-preview-token-0123456789",
        )
        self.assertEqual(preview["expires_at"], "2026-08-31T20:15:00")

        with self.assertRaises(XcomPaymentError) as caught:
            consume_payment_preview(
                self.conn,
                preview["preview_token"],
                "SIMPLE",
                "2026-08-01",
                "2026-08-31",
                "2026-08-31",
                now_iso=lambda: "2026-08-31T20:15:00",
            )

        self.assertEqual(caught.exception.code, "preview_expired")

    def test_tax_code_period_and_date_range_validation(self):
        invalid_tax = {**self.bot_profile, "recipient_tax_code": "0201650135-12"}
        with self.assertRaises(XcomPaymentError) as caught:
            upsert_payment_profile(self.conn, invalid_tax, now_iso=FIXED_NOW)
        self.assertEqual(caught.exception.code, "invalid_tax_code")
        valid_branch = {**self.bot_profile, "profile_code": "BOT-BRANCH", "issuer_tax_code": "0201853470-001"}
        self.assertTrue(
            upsert_payment_profile(self.conn, valid_branch, now_iso=FIXED_NOW)["created"]
        )

        self._insert_scope_data(self.simple_profile)
        with self.assertRaises(XcomPaymentError) as caught:
            upsert_meal_tariff(self.conn, "SIMPLE", "2026-13", "Sáng", 25_000, now_iso=FIXED_NOW)
        self.assertEqual(caught.exception.code, "invalid_period")
        self._insert_attendance()
        upsert_meal_tariff(self.conn, "SIMPLE", "2026-08", "Sáng", 25_000, now_iso=FIXED_NOW)
        with self.assertRaises(XcomPaymentError) as caught:
            build_meal_payment_summary(self.conn, "SIMPLE", "2026-08-31", "2026-08-01")
        self.assertEqual(caught.exception.code, "invalid_date_range")
        with self.assertRaises(XcomPaymentError) as caught:
            build_meal_payment_summary(self.conn, "SIMPLE", "2026-02-30", "2026-08-31")
        self.assertEqual(caught.exception.code, "invalid_date")

    def test_missing_bank_or_bot_legal_signer_config_is_blocked(self):
        profile = {**self.simple_profile, "bank_account": ""}
        self._insert_scope_data(profile)
        self._insert_attendance()
        upsert_meal_tariff(self.conn, "SIMPLE", "2026-08", "Sáng", 25_000, now_iso=FIXED_NOW)
        with self.assertRaises(XcomPaymentError) as caught:
            build_meal_payment_summary(self.conn, "SIMPLE", "2026-08-01", "2026-08-31")
        self.assertEqual(caught.exception.code, "missing_config")
        self.assertIn("bank_account", caught.exception.details)

        incomplete_bot = {**self.bot_profile, "buyer_signer_name": ""}
        upsert_payment_profile(self.conn, incomplete_bot, now_iso=FIXED_NOW)
        assign_payment_scope(self.conn, "BOT", "KITCHEN", "BEP-B", now_iso=FIXED_NOW)
        self.conn.execute("INSERT INTO kitchen_units VALUES('BEP-B','BOT','x')")
        self._insert_attendance(kitchen="BEP-B", shift="Trưa")
        upsert_meal_tariff(self.conn, "BOT", "2026-08", "Trưa", 35_000, now_iso=FIXED_NOW)
        with self.assertRaises(XcomPaymentError) as caught:
            build_meal_payment_summary(self.conn, "BOT", "2026-08-01", "2026-08-31")
        self.assertIn("buyer_signer_name", caught.exception.details)

    def test_kitchen_scope_overrides_xcom_scope_without_double_count(self):
        self._insert_scope_data(self.simple_profile)
        other = {**self.simple_profile, "profile_code": "OTHER"}
        upsert_payment_profile(self.conn, other, now_iso=FIXED_NOW)
        assign_payment_scope(self.conn, "OTHER", "KITCHEN", "BEP-A", now_iso=FIXED_NOW)
        self._insert_attendance(actual=12)
        upsert_meal_tariff(self.conn, "SIMPLE", "2026-08", "Sáng", 10, now_iso=FIXED_NOW)
        upsert_meal_tariff(self.conn, "OTHER", "2026-08", "Sáng", 20, now_iso=FIXED_NOW)

        with self.assertRaises(XcomPaymentError) as caught:
            build_meal_payment_summary(self.conn, "SIMPLE", "2026-08-01", "2026-08-31")
        self.assertEqual(caught.exception.code, "no_actual_attendance")
        other_result = build_meal_payment_summary(self.conn, "OTHER", "2026-08-01", "2026-08-31")
        self.assertEqual(other_result["actual_count"], 12)
        self.assertEqual(other_result["subtotal"], 240)

    def test_docx_is_deterministic_and_preserves_customer_payment_form(self):
        self._insert_scope_data(self.simple_profile)
        self._insert_attendance()
        upsert_meal_tariff(self.conn, "SIMPLE", "2026-08", "Sáng", 25_000, now_iso=FIXED_NOW)
        summary = build_meal_payment_summary(self.conn, "SIMPLE", "2026-08-01", "2026-08-31")

        first = build_simple_payment_request_docx(summary, "2026-08-31")
        second = build_simple_payment_request_docx(summary, "2026-08-31")

        self.assertEqual(hashlib.sha256(first).digest(), hashlib.sha256(second).digest())
        document = Document(io.BytesIO(first))
        paragraph_text = "\n".join(p.text for p in document.paragraphs)
        self.assertIn("GIẤY ĐỀ NGHỊ THANH TOÁN", paragraph_text)
        self.assertIn("Số tiền: 27.500 VNĐ", paragraph_text)
        self.assertIn("từ 01/08/2026 đến 31/08/2026", paragraph_text)
        self.assertEqual(len(document.tables), 2)
        self.assertEqual(len(document.tables[0].columns), 2)
        self.assertEqual(len(document.tables[1].columns), 3)
        all_text = paragraph_text + "\n" + "\n".join(
            cell.text for table in document.tables for row in table.rows for cell in row.cells
        )
        self.assertIn(self.simple_profile["bank_account"], all_text)
        self.assertNotIn("999", all_text)

    def test_bot_bundle_is_deterministic_formula_backed_and_reconciled(self):
        self._insert_scope_data(self.bot_profile, "KITCHEN", "BEP-A")
        self._insert_attendance(actual=2, ordered=700)
        self._insert_attendance(work_date="2026-08-02", shift="Trưa", actual=3, ordered=700)
        upsert_meal_tariff(self.conn, "BOT", "2026-08", "Sáng", 25_000, now_iso=FIXED_NOW)
        upsert_meal_tariff(self.conn, "BOT", "2026-08", "Trưa", 35_000, now_iso=FIXED_NOW)

        first = create_payment_document(
            self.conn, "BOT", "2026-08-01", "2026-08-31", "2026-08-31"
        )
        second = create_payment_document(
            self.conn, "BOT", "2026-08-01", "2026-08-31", "2026-08-31"
        )

        self.assertEqual(first["file_sha256"], second["file_sha256"])
        self.assertEqual(first["payload"], second["payload"])
        workbook = load_workbook(io.BytesIO(first["payload"]), data_only=False)
        self.assertEqual(
            workbook.sheetnames,
            ["BBĐC", "ĐNTT", "suất ăn"],
        )
        reconcile = workbook["BBĐC"]
        self.assertEqual(reconcile["C19"].value, 2)
        self.assertEqual(reconcile["D20"].value, 3)
        self.assertNotEqual(reconcile["C19"].value, 700)
        self.assertEqual(workbook["suất ăn"]["C9"].value, "='BBĐC'!C19")
        self.assertEqual(workbook["ĐNTT"]["D13"].value, "='BBĐC'!C50")
        self.assertEqual(workbook["ĐNTT"]["F18"].value, "=F16+F17")
        self.assertEqual(first["summary"]["subtotal"], 155_000)
        self.assertEqual(first["summary"]["vat_amount"], 15_500)
        self.assertEqual(first["summary"]["total"], 170_500)
        with zipfile.ZipFile(io.BytesIO(first["payload"])) as archive:
            self.assertTrue(all(info.date_time == (1980, 1, 1, 0, 0, 0) for info in archive.infolist()))

    def test_bot_bundle_blocks_cross_month_period_that_cannot_fit_customer_form(self):
        self._insert_scope_data(self.bot_profile, "KITCHEN", "BEP-A")
        self._insert_attendance(work_date="2026-08-31", shift="Sáng", actual=2)
        self._insert_attendance(work_date="2026-09-01", shift="Sáng", actual=3)
        upsert_meal_tariff(self.conn, "BOT", "2026-08", "Sáng", 25_000, now_iso=FIXED_NOW)
        upsert_meal_tariff(self.conn, "BOT", "2026-09", "Sáng", 30_000, now_iso=FIXED_NOW)

        with self.assertRaises(XcomPaymentError) as caught:
            create_payment_document(
                self.conn, "BOT", "2026-08-31", "2026-09-01", "2026-09-02"
            )
        self.assertEqual(caught.exception.code, "invalid_bot_period")


if __name__ == "__main__":
    unittest.main()
