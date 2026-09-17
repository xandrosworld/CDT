"""Versioned document selections. Never changes orders, debt or stock postings."""
import hashlib
import json
from collections import defaultdict
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP

from flask import jsonify, request
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

try:
    from .purchase_summary_export import collect_purchase_summary_rows, PurchaseSummaryError
    from .receipt_export import enrich_receipt_identity_rows, group_receipt_rows, RECEIPT_MAX_DAILY_AMOUNT
except ImportError:
    from purchase_summary_export import collect_purchase_summary_rows, PurchaseSummaryError
    from receipt_export import enrich_receipt_identity_rows, group_receipt_rows, RECEIPT_MAX_DAILY_AMOUNT


class SelectionError(PurchaseSummaryError):
    def __init__(self, message):
        super().__init__(message, code='purchase_document_selection_required', status=409)

    status = 409
    code = 'purchase_document_selection_required'


def init_schema(conn):
    conn.execute('''CREATE TABLE IF NOT EXISTS purchase_document_selections (
        work_date TEXT PRIMARY KEY, revision INTEGER NOT NULL, source_hash TEXT NOT NULL,
        quantities_json TEXT NOT NULL, actor TEXT NOT NULL, reason TEXT NOT NULL,
        updated_at TEXT NOT NULL)''')
    conn.execute('''CREATE TABLE IF NOT EXISTS purchase_document_selection_history (
        work_date TEXT NOT NULL, revision INTEGER NOT NULL, source_hash TEXT NOT NULL,
        quantities_json TEXT NOT NULL, source_json TEXT NOT NULL, actor TEXT NOT NULL,
        reason TEXT NOT NULL, created_at TEXT NOT NULL, PRIMARY KEY(work_date,revision))''')


def digest(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True, ensure_ascii=False,
                                    default=str, separators=(',', ':')).encode()).hexdigest()


def decimal(value):
    if isinstance(value, bool):
        raise SelectionError('Số lượng không hợp lệ')
    try:
        result = Decimal(str(value))
        if not result.is_finite() or result < 0 or result.as_tuple().exponent < -6:
            raise ValueError()
        return result
    except (InvalidOperation, ValueError, TypeError):
        raise SelectionError('Số lượng phải không âm, tối đa 6 số lẻ') from None


def amount(row, qty):
    # Preserve the original whole-line amount, including any rounding allocation.
    if qty == Decimal(str(row['quantity'])):
        return Decimal(str(row['amount']))
    return (Decimal(str(row['amount'])) * qty / Decimal(str(row['quantity']))).quantize(
        Decimal('0.01'), rounding=ROUND_HALF_UP)


def day_source(conn, work_date):
    rows = []
    for batch in conn.execute("SELECT * FROM batches WHERE work_date=? AND status='approved' ORDER BY id", (work_date,)):
        batch = dict(batch)
        orders = [dict(r) for r in conn.execute('SELECT * FROM orders WHERE batch_id=? ORDER BY id', (batch['id'],))]
        try:
            candidates = collect_purchase_summary_rows(conn, batch, orders)
        except PurchaseSummaryError as exc:
            if exc.code in ('no_purchase_summary_rows', 'excluded_sellers_no_rows'):
                continue
            raise
        rows.extend(enrich_receipt_identity_rows(conn, candidates))
    if any(row['work_date'] != work_date for row in rows):
        raise SelectionError('Ngày mua trên dòng hàng khác ngày đơn; cần kiểm tra ngày nguồn trước khi chọn bảng kê')
    rows.sort(key=lambda r: r['selection_key'])
    return rows, digest(rows)


def saved_plan(conn, work_date):
    if not conn.execute("SELECT 1 FROM sqlite_master WHERE name='purchase_document_selections'").fetchone():
        return None
    row = conn.execute('SELECT * FROM purchase_document_selections WHERE work_date=?', (work_date,)).fetchone()
    return dict(row) if row else None


def split_rows(rows, quantities):
    if not isinstance(quantities, dict) or set(quantities) - {r['selection_key'] for r in rows}:
        raise SelectionError('Lựa chọn có dòng không còn thuộc nguồn; tải lại để kiểm tra')
    selected, pending = [], []
    for row in rows:
        full = decimal(row['quantity'])
        chosen = decimal(quantities.get(row['selection_key'], 0))
        if chosen > full:
            raise SelectionError('Số lượng chọn vượt lượng mua: ' + row['product_name'])
        part = amount(row, chosen)
        if chosen and part <= 0:
            raise SelectionError('Lượng đã chọn quá nhỏ để lập thành tiền: ' + row['product_name'])
        if chosen:
            selected.append({**row, 'quantity': float(chosen), 'amount': float(part)})
        if chosen < full:
            pending.append({**row, 'quantity': float(full - chosen),
                            'amount': float(Decimal(str(row['amount'])) - part)})
    # Validate the entire day, across batches and aliases of the same identity.
    if selected:
        group_receipt_rows(selected)
    return selected, pending


def group_totals(rows):
    groups = {}
    for row in rows:
        key = row['cccd']
        group = groups.setdefault(key, {'seller': row['seller'], 'identity': key, 'amount': Decimal(0)})
        group['amount'] += Decimal(str(row['amount']))
    return groups


def scope(conn, start, end):
    try:
        a, b = date.fromisoformat(start), date.fromisoformat(end)
        if a > b or (b - a).days > 31:
            raise ValueError()
    except (ValueError, TypeError):
        raise SelectionError('Chọn khoảng ngày hợp lệ, tối đa 32 ngày') from None
    days = []
    for value in conn.execute("SELECT DISTINCT work_date FROM batches WHERE status='approved' AND work_date BETWEEN ? AND ? ORDER BY work_date", (start, end)):
        work_date = value['work_date']
        rows, source_hash = day_source(conn, work_date)
        if not rows:
            continue
        plan = saved_plan(conn, work_date)
        totals = group_totals(rows)
        # Never silently choose a subset of an over-limit seller's rows.
        quantities = json.loads(plan['quantities_json']) if plan else {
            row['selection_key']: str(row['quantity']) for row in rows
            if totals[row['cccd']]['amount'] <= RECEIPT_MAX_DAILY_AMOUNT}
        stale = bool(plan and plan['source_hash'] != source_hash)
        selected_amounts = defaultdict(Decimal)
        public_rows = []
        for row in rows:
            qty = quantities.get(row['selection_key'], '0')
            selected_amounts[row['cccd']] += amount(row, decimal(qty))
            public_rows.append({k: row[k] for k in (
                'selection_key', 'batch_id', 'seller', 'cccd', 'product_name', 'unit',
                'quantity', 'buy_price', 'amount', 'source_ref', 'kitchen')})
            public_rows[-1]['selected_quantity'] = qty
        days.append({'date': work_date, 'source_hash': source_hash,
                     'revision': plan['revision'] if plan else 0, 'stale': stale,
                     'rows': public_rows, 'groups': [
                         {'seller': g['seller'], 'identity': key, 'total': float(g['amount']),
                          'selected': float(selected_amounts[key]),
                          'pending': float(g['amount'] - selected_amounts[key])}
                         for key, g in totals.items()]})
    token = digest([{'date': d['date'], 'source_hash': d['source_hash'], 'revision': d['revision']} for d in days])
    return {'from': start, 'to': end, 'days': days, 'state_hash': token,
            'limit': int(RECEIPT_MAX_DAILY_AMOUNT)}


def save(conn, body, timestamp, audit):
    current = scope(conn, body.get('from'), body.get('to'))
    actor, reason = str(body.get('actor') or '').strip(), str(body.get('reason') or '').strip()
    if not actor or len(actor) > 120 or len(reason) > 500:
        raise SelectionError('Ghi người lưu (tối đa 120 ký tự), lý do tối đa 500 ký tự')
    submitted = body.get('days')
    if not isinstance(submitted, list) or not submitted:
        raise SelectionError('Chưa có ngày được chọn')
    requested = {item.get('date'): item.get('quantities') for item in submitted if isinstance(item, dict)}
    if len(requested) != len(submitted) or set(requested) != {d['date'] for d in current['days']}:
        raise SelectionError('Danh sách ngày đã thay đổi; tải lại để kiểm tra')
    prepared = []
    for day in current['days']:
        rows, source_hash = day_source(conn, day['date'])
        if day['stale']:
            raise SelectionError('Nguồn ngày ' + day['date'] + ' đã đổi sau khi lập bảng kê; cần đối chiếu bản đã lập trước khi thay lựa chọn')
        split_rows(rows, requested[day['date']])
        quantities = {k: format(decimal(v).normalize(), 'f') for k, v in requested[day['date']].items() if decimal(v)}
        prior = saved_plan(conn, day['date'])
        unchanged = bool(prior and json.loads(prior['quantities_json']) == quantities)
        if prior and not unchanged and not reason:
            raise SelectionError('Ghi lý do thay lựa chọn; bản mới thay thế bản cũ, không phải bảng kê cộng thêm')
        prepared.append((day, rows, quantities, unchanged))
    if current['state_hash'] != body.get('state_hash'):
        if all(item[3] for item in prepared):
            return {**current, 'ok': True, 'unchanged': True}
        raise SelectionError('Có người vừa sửa lựa chọn hoặc nguồn đã đổi; tải lại trước khi lưu')
    for day, rows, quantities, unchanged in prepared:
        if unchanged:
            continue
        revision = day['revision'] + 1
        values = (day['date'], revision, day['source_hash'], json.dumps(quantities, sort_keys=True), actor, reason, timestamp)
        conn.execute('''INSERT INTO purchase_document_selections VALUES(?,?,?,?,?,?,?)
            ON CONFLICT(work_date) DO UPDATE SET revision=excluded.revision,source_hash=excluded.source_hash,
            quantities_json=excluded.quantities_json,actor=excluded.actor,reason=excluded.reason,updated_at=excluded.updated_at''', values)
        conn.execute('INSERT INTO purchase_document_selection_history VALUES(?,?,?,?,?,?,?,?)',
                     values[:4] + (json.dumps(rows, ensure_ascii=False, default=str),) + values[4:])
        audit(conn, 'purchase_document.selection', entity_type='purchase_document_day', entity_id=day['date'],
              metadata={'revision': revision, 'actor': actor, 'reason': reason, 'selected_rows': len(quantities),
                        'source_hash': day['source_hash']})
    return {**scope(conn, body['from'], body['to']), 'ok': True}


def pending_queue(conn, through, after=''):
    """Read current remainders without creating another purchase or daily allowance.

    Paginate source dates, not just pending dates: even invalid/stale days remain
    visible, and a quiet page must not imply that older/later pages are settled.
    """
    try:
        date.fromisoformat(through)
        if after:
            date.fromisoformat(after)
            if after > through:
                raise ValueError()
    except (ValueError, TypeError):
        raise SelectionError('Ngày theo dõi phần chờ không hợp lệ') from None
    dates = {r[0] for r in conn.execute(
        "SELECT DISTINCT work_date FROM batches WHERE status='approved' AND work_date>? AND work_date<=?",
        (after, through))}
    if conn.execute("SELECT 1 FROM sqlite_master WHERE name='purchase_document_selections'").fetchone():
        dates.update(r[0] for r in conn.execute(
            'SELECT work_date FROM purchase_document_selections WHERE work_date>? AND work_date<=?',
            (after, through)))
    dates = sorted(dates)
    page, entries = dates[:20], []
    for work_date in page:
        try:
            state = scope(conn, work_date, work_date)
            if not state['days']:
                if saved_plan(conn, work_date):
                    raise SelectionError('Nguồn của lựa chọn đã lưu không còn; cần đối chiếu bản đã lập')
                continue
            day = state['days'][0]
            if day['stale']:
                raise SelectionError('Nguồn đã đổi sau khi lưu; cần đối chiếu bản đã lập, chưa xác định lại phần chờ')
            rows, _ = day_source(conn, work_date)
            _, pending = split_rows(rows, {r['selection_key']: r['selected_quantity'] for r in day['rows']})
            if not pending:
                continue
            groups = []
            for group in day['groups']:
                remainder = [r for r in pending if r['cccd'] == group['identity']]
                if remainder:
                    groups.append({**group, 'remaining_capacity': max(0, state['limit'] - group['selected']),
                                   'rows': remainder})
            entries.append({'date': work_date, 'revision': day['revision'], 'groups': groups,
                            'pending': float(sum((Decimal(str(r['amount'])) for r in pending), Decimal(0))),
                            'status': 'saved' if day['revision'] else 'needs_selection'})
        except (ValueError, TypeError) as exc:
            entries.append({'date': work_date, 'status': 'needs_review', 'error': str(exc)})
    return {'through': through, 'entries': entries, 'scanned_dates': len(page),
            'next_after': page[-1] if len(dates) > len(page) else None}


def export_scope(conn, batch, rows):
    """A saved date-wide choice is shared by every export path and batch."""
    plan = saved_plan(conn, batch['work_date'])
    if batch['status'] != 'approved':
        if plan:
            raise SelectionError('Ngày này đã lưu lựa chọn; duyệt đơn rồi kiểm tra lại trước khi xuất')
        return rows, [], None
    all_rows, source_hash = day_source(conn, batch['work_date'])
    if not plan:
        try:
            group_receipt_rows(all_rows)
        except ValueError as exc:
            if getattr(exc, 'code', '') == 'receipt_daily_limit_exceeded':
                raise SelectionError('Có người bán vượt hạn mức lựa chọn 5.000.000đ/ngày. Bấm “Chọn hàng lập bảng kê / phần chờ” để chọn và lưu.') from None
            raise
        return rows, [], None
    if plan['source_hash'] != source_hash:
        raise SelectionError('Nguồn bảng kê đã thay đổi sau khi lưu lựa chọn; cần đối chiếu trước khi xuất lại')
    selected, pending = split_rows(all_rows, json.loads(plan['quantities_json']))
    batch_id = int(batch['id'])
    return ([r for r in selected if r['batch_id'] == batch_id],
            [r for r in pending if r['batch_id'] == batch_id],
            {'revision': plan['revision'], 'date': batch['work_date'], 'all_rows': all_rows,
             'selected_rows': selected, 'pending_rows': pending})


def annotate_workbook(book, selection, pending):
    """Keep selection reconciliation in its annex, off individual receipts."""
    if not selection:
        return book
    note = (f"Lựa chọn ngày {selection['date']} · bản {selection['revision']} (thay thế bản lựa chọn trước). "
            "Chỉ gồm phần đã chọn; xem Đối chiếu lựa chọn để biết tổng mua và phần chờ trong ngày.")
    all_groups = group_totals(selection['all_rows'])
    chosen = group_totals(selection['selected_rows'])
    for ws in book:
        receipt = ws.title.startswith('biên nhận')
        if receipt:
            # A documentation selection is not evidence of a cash payment.
            # Leave the payment method/date for the actual supporting record.
            for row in ws:
                for cell in row:
                    if str(cell.value or '').strip() == 'Bên mua thanh toán tiền mặt ngay sau khi nhận đủ hàng':
                        cell.value = 'Hình thức / ngày thanh toán: ........................................'
            # Customer-facing receipts end at the signatures. Full-day totals,
            # revisions and pending amounts remain in the reconciliation sheets.
            continue
        full = sum(g['amount'] for g in all_groups.values())
        part = sum((g['amount'] for g in chosen.values()), Decimal(0))
        visible_note = (f"Lựa chọn ngày {selection['date']} · bản {selection['revision']} (thay thế bản trước). "
                        f"Tổng mua cả ngày: {full:,.2f}đ; đã chọn cả ngày: {part:,.2f}đ; "
                        f"chờ bổ sung: {full - part:,.2f}đ.")
        bottom, left, right = ws.max_row + 2, 1, ws.max_column
        ws.cell(bottom, left, visible_note)
        ws.merge_cells(start_row=bottom, start_column=left, end_row=bottom, end_column=right)
        ws.cell(bottom, left).alignment = Alignment(wrap_text=True, vertical='center')
        ws.cell(bottom, left).font = Font(name='Times New Roman', size=9, italic=True)
        ws.row_dimensions[bottom].height = 48
        ws.print_area = f'{get_column_letter(left)}1:{get_column_letter(right)}{bottom}'
    ws = book.create_sheet('Đối chiếu lựa chọn')
    ws.append(['Ngày', 'Người bán', 'Tổng mua trong ngày', 'Phần đã chọn', 'Chờ bổ sung chứng từ'])
    for identity, item in all_groups.items():
        part = chosen.get(identity, {}).get('amount', Decimal(0))
        ws.append([selection['date'], item['seller'], float(item['amount']), float(part), float(item['amount'] - part)])
    ws.append(['TỔNG', '', *[sum(float(ws.cell(r, c).value) for r in range(2, ws.max_row + 1)) for c in (3, 4, 5)]])
    ws.append(['Phần chờ không làm giảm tiền mua/công nợ, không tự tạo hàng âm. Tổng mua thực tế trong ngày giữ nguyên.'])
    ws.append([note])
    pending_sheet = book.create_sheet('Chờ bổ sung chứng từ')
    pending_sheet.append(['Ngày mua', 'Người bán', 'Tên hàng', 'ĐVT', 'Số lượng chờ', 'Giá mua BK', 'Tiền chờ', 'Bếp', 'Dòng nguồn'])
    for r in pending:
        pending_sheet.append([r['work_date'], r['seller'], r['product_name'], r['unit'],
                              r['quantity'], r['buy_price'], r['amount'], r['kitchen'], r['source_ref']])
    pending_sheet.append(['TỔNG', '', '', '', '', '', sum(r['amount'] for r in pending)])
    for sheet in (ws, pending_sheet):
        for col in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(col)].width = 24 if col != 3 else 36
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
                cell.font = Font(name='Times New Roman', size=11)
                cell.alignment = Alignment(wrap_text=True, vertical='top')
                if isinstance(cell.value, (float, int)):
                    cell.number_format = '#,##0.######'
        sheet.freeze_panes = 'A2'
        for cell in sheet[1]:
            cell.font = Font(name='Times New Roman', size=11, bold=True)
            cell.fill = PatternFill('solid', fgColor='E9F2F1')
        if sheet == ws:
            for r in (sheet.max_row - 1, sheet.max_row):
                sheet.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
                sheet.row_dimensions[r].height = 36
        sheet.print_area = f'A1:{get_column_letter(sheet.max_column)}{sheet.max_row}'
        sheet.page_setup.orientation = 'landscape'
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
    return book


def register_routes(app, ctx):
    @app.get('/api/purchase-document-selection/pending')
    def purchase_document_pending():
        try:
            # Server business date, never the browser's date or the selected range.
            through = datetime.now(timezone(timedelta(hours=7))).date().isoformat()
            with ctx['db']() as conn:
                conn.execute('BEGIN')
                return jsonify(ok=True, **pending_queue(conn, through, request.args.get('after', '')))
        except (ValueError, TypeError) as exc:
            return jsonify(ok=False, error=str(exc)), 409

    @app.route('/api/purchase-document-selection', methods=['GET', 'POST'])
    def purchase_document_selection():
        try:
            with ctx['db']() as conn:
                if request.method == 'GET':
                    conn.execute('BEGIN')
                    return jsonify(ok=True, **scope(conn, request.args.get('from'), request.args.get('to')))
                body = request.get_json(silent=True)
                if not isinstance(body, dict):
                    raise SelectionError('Nội dung lưu không hợp lệ')
                conn.execute('BEGIN IMMEDIATE')
                return jsonify(save(conn, body, ctx['now_iso'](), ctx['audit_event']))
        except (ValueError, TypeError) as exc:
            return jsonify(ok=False, error=str(exc), code=getattr(exc, 'code', 'invalid_document_selection')), getattr(exc, 'status', 409)
