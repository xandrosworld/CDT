from __future__ import annotations

from collections import defaultdict
from pathlib import Path
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
import server


SOURCE = Path.home() / "Downloads" / "Đơn hàng 29.08.xlsx"
OUTPUT = Path(__file__).resolve().parents[1] / "CAN_XAC_NHAN_DON_29-08.xlsx"
TODAY_SHEET = "29.08"
WORK_DATE = "2026-08-29"
HISTORY_SHEETS = [
    "28.08", "27.08", "26.08", "25.08", "24.08", "23.08", "22.08",
    "21.08", "20.08", "19.08", "18.08", "17.08", "16.8", "15.08",
    "14.08", "13.8", "12.08",
]


def raw_rows(ws):
    header_row, mapping = server.detect_header(ws)
    if not header_row:
        return []
    rows = []
    for row_number in range(header_row + 1, ws.max_row + 1):
        qty = server.number_value(ws.cell(row_number, mapping["qty"]).value)
        if qty <= 0:
            continue
        row = {field: ws.cell(row_number, col).value for field, col in mapping.items()}
        rows.append({
            "sheet": ws.title,
            "product_code": server.clean_text(row.get("product_code")).upper(),
            "contractor": server.clean_text(row.get("contractor")).upper(),
            "kitchen": server.clean_text(row.get("kitchen")).upper(),
            "buy_price": server.number_value(row.get("buy_price")),
            "sell_price": server.number_value(row.get("sell_price")),
            "cccd": server.clean_text(row.get("cccd")),
        })
    return rows


def first_positive(rows, field):
    for row in rows:
        if field == "cccd":
            if row[field]:
                return row[field], row["sheet"]
        elif row[field] > 0:
            return row[field], row["sheet"]
    return "", ""


def historical_suggestion(item, candidates, field):
    same_contractor = [row for row in candidates if row["contractor"] == item["contractor"]]
    same_kitchen = [row for row in same_contractor if row["kitchen"] == item["kitchen"]]
    value, sheet = first_positive(same_kitchen, field)
    if value != "":
        return value, f"{sheet} · cùng bếp"
    value, sheet = first_positive(same_contractor, field)
    if value != "":
        return value, f"{sheet} · cùng nhà thầu, khác bếp"
    return "", "Không có lịch sử đủ an toàn"


def apply_table_style(ws, widths):
    header_fill = PatternFill("solid", fgColor="16324F")
    warning_fill = PatternFill("solid", fgColor="FFF2CC")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 34
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        if row[10].value:
            for cell in row:
                cell.fill = warning_fill
    for column, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(column)].width = width


def main():
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    server.init_database()
    current, _ = server.parse_workbook(SOURCE, WORK_DATE, [TODAY_SHEET])

    source_book = load_workbook(SOURCE, data_only=True, read_only=False)
    history = defaultdict(list)
    for sheet_name in HISTORY_SHEETS:
        if sheet_name not in source_book.sheetnames:
            continue
        for row in raw_rows(source_book[sheet_name]):
            history[row["product_code"]].append(row)

    blocked = [item for item in current if item["errors"]]
    book = Workbook()
    summary = book.active
    summary.title = "TỔNG HỢP"
    summary.append(["KIỂM TRA ĐƠN 29.08", "KẾT QUẢ"])
    summary_rows = [
        ("Tổng dòng đơn", len(current)),
        ("Dòng đạt kiểm tra", len(current) - len(blocked)),
        ("Dòng cần xác nhận", len(blocked)),
        ("GIANHAPTAY thiếu giá bán theo ngày", sum(
            item["contractor"] == "GIANHAPTAY" and item["sell_price"] <= 0 for item in current
        )),
        ("Dòng còn thiếu giá mua", sum(item["buy_price"] <= 0 for item in current)),
        ("Dòng thiếu CCCD hàng bảng kê", sum(
            any("CCCD" in error for error in item["errors"]) for item in current
        )),
        ("Kết luận", "Chưa được duyệt/xuất hóa đơn cho tới khi điền xong các ô xác nhận."),
        ("Lưu ý", "Giá lịch sử chỉ để tham khảo; GIANHAPTAY là giá theo ngày, không tự lấy giá hôm trước."),
    ]
    for row in summary_rows:
        summary.append(row)
    summary.column_dimensions["A"].width = 43
    summary.column_dimensions["B"].width = 82
    summary.freeze_panes = "A2"
    for cell in summary[1]:
        cell.fill = PatternFill("solid", fgColor="16324F")
        cell.font = Font(color="FFFFFF", bold=True)
    for row in summary.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    detail = book.create_sheet("CẦN XÁC NHẬN")
    detail.append([
        "STT", "Dòng Excel", "Nhà thầu", "Mã bếp", "Mã hàng", "Tên hàng", "SL", "NCC",
        "Giá mua trong file", "Giá bán trong file", "Lỗi cần xử lý",
        "Giá mua lịch sử (tham khảo)", "Giá bán lịch sử (tham khảo)", "Nguồn tham khảo",
        "GIÁ MUA KHÁCH XÁC NHẬN", "GIÁ BÁN KHÁCH XÁC NHẬN", "CCCD BỔ SUNG", "Ghi chú/quyết định",
    ])
    for index, item in enumerate(blocked, start=1):
        candidates = history.get(item["product_code"], [])
        suggested_buy, buy_source = historical_suggestion(item, candidates, "buy_price")
        suggested_sell, sell_source = historical_suggestion(item, candidates, "sell_price")
        cccd, cccd_source = historical_suggestion(item, candidates, "cccd")
        sources = []
        if suggested_buy != "":
            sources.append(f"Giá mua: {buy_source}")
        if suggested_sell != "":
            sources.append(f"Giá bán: {sell_source}")
        if cccd:
            sources.append(f"CCCD: {cccd_source}")
        if not sources:
            sources.append("Không có lịch sử đủ an toàn")
        detail.append([
            index, item["source_row"], item["contractor"], item["kitchen"], item["product_code"],
            item["product_name"], item["qty"], item["supplier"], item["buy_price"] or "",
            item["sell_price"] or "", "; ".join(item["errors"]), suggested_buy, suggested_sell,
            "\n".join(sources), "", "", cccd if any("CCCD" in e for e in item["errors"]) else "", "",
        ])
    apply_table_style(detail, [7, 11, 17, 18, 13, 30, 9, 17, 16, 16, 32, 20, 20, 30, 22, 22, 18, 26])
    for row in range(2, detail.max_row + 1):
        for col in (9, 10, 12, 13, 15, 16):
            detail.cell(row, col).number_format = "#,##0"

    book.save(OUTPUT)
    print(f"OUTPUT={OUTPUT}")
    print(f"TOTAL={len(current)} BLOCKED={len(blocked)}")


if __name__ == "__main__":
    main()
