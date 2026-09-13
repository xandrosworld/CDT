import io
import unittest
from openpyxl import load_workbook
from . import server
from . import test_daily_workbook_import as fixtures


class SupplierPlanSourceTests(unittest.TestCase):
    setUpClass = classmethod(fixtures.DailyWorkbookImportTests.setUpClass.__func__)
    tearDownClass = classmethod(fixtures.DailyWorkbookImportTests.tearDownClass.__func__)
    setUp = fixtures.DailyWorkbookImportTests.setUp

    def file(self, qty=5, price=0, *, bad_date=False, missing=False):
        wb = load_workbook(io.BytesIO(fixtures.DailyWorkbookImportTests.workbook_bytes(quantities=(2, 3))))
        ws = wb['đặt hàng']
        ws.cell(3, 4, qty)
        ws.cell(3, 6, 'S1')
        ws.cell(3, 8, price)
        ws.cell(3, 13, qty)
        ws.cell(3, 14, price * qty)
        if bad_date:
            ws.cell(2, 15, 'Ngày')
            ws.cell(3, 15, '2026-09-02')
        if missing:
            wb.remove(ws)
        stream = io.BytesIO()
        wb.save(stream)
        wb.close()
        return stream.getvalue()

    def daily(self, data):
        preview = self.client.post('/api/import/analyze', data={
            'file': (io.BytesIO(data), 'Đơn hàng 01.09.2026.xlsx'), 'continuous': '1',
        }).get_json()
        response = self.client.post('/api/import/confirm', json={
            'token': preview['token'], 'state_hash': preview['stateHash'],
            'sheets': ['01.09'], 'work_date': '2026-09-01',
        })
        self.assertEqual(response.status_code, 200, response.get_json())
        return response.get_json()['batch']['id']

    def preview(self, bid, data, plan=True):
        result = self.client.post('/api/purchase-orders/import/preview', data={
            'batch_id': str(bid), 'plan_only': '1' if plan else '0',
            'file': (io.BytesIO(data), 'Đơn hàng 01.09.2026.xlsx'),
        })
        self.assertEqual(result.status_code, 200, result.get_json())
        self.assertEqual(result.get_json()['batch_id'], bid)
        return result.get_json()

    def confirm(self, preview, status=200):
        response = self.client.post('/api/purchase-orders/import/confirm',
                                    json={'token': preview['token'], 'confirmed': True})
        self.assertEqual(response.status_code, status, response.get_json())
        return response.get_json()

    def needs(self, bid):
        return self.client.get(f'/api/supplier-needs/{bid}').get_json()

    def protected(self):
        with server.db() as conn:
            return {t: [tuple(r) for r in conn.execute('SELECT * FROM ' + t + ' ORDER BY 1')]
                    for t in ('orders', 'batches', 'purchase_workbook_lines', 'purchase_order_lines',
                              'payable_ledger_lines', 'receivable_ledger_lines', 'invoice_inventory_ledger')}

    def test_first_upload_uses_purchase_sheet_and_zero_price_does_not_post_debt(self):
        bid = self.daily(self.file(7))
        plan = self.needs(bid)
        self.assertTrue(plan['send_available'])
        self.assertEqual(plan['format'], 'supplier_sheet_plan')
        self.assertEqual([(r['source_sheet'], r['source_row'], r['order_qty']) for r in plan['rows']], [('đặt hàng', 3, 7)])
        with server.db() as conn:
            self.assertEqual([r['qty'] for r in conn.execute('SELECT qty FROM orders ORDER BY id')], [2, 3])
            self.assertEqual(conn.execute('SELECT count(*) FROM purchase_workbook_lines').fetchone()[0], 0)
            self.assertEqual(conn.execute('SELECT count(*) FROM payable_ledger_lines WHERE batch_id=?', (bid,)).fetchone()[0], 0)

    def test_same_file_replay_and_recovery_do_not_duplicate(self):
        data = self.file()
        bid = self.daily(data)
        before = self.protected()
        self.assertEqual(self.daily(data), bid)
        self.assertEqual(self.protected(), before)
        with server.db() as conn:
            self.assertEqual(conn.execute('SELECT count(*) FROM supplier_plan_history').fetchone()[0], 1)
            # Simulate a batch imported before this fix.
            conn.execute('DELETE FROM supplier_plan_sources')
            conn.execute('DELETE FROM supplier_plan_history')
        self.assertFalse(self.needs(bid)['send_available'])
        self.assertEqual(self.daily(data), bid)
        self.assertTrue(self.needs(bid)['send_available'])
        self.assertEqual(self.protected(), before)

    def test_plan_repair_preserves_sales_accounting_and_reopens_changed_supplier(self):
        original = self.file()
        bid = self.daily(original)
        p = self.needs(bid)
        status = self.client.put(f'/api/supplier-order-status/{bid}/S1', json={
            'status': 'ordered', 'revision': 0, 'plan_hash': p['plan_hash']})
        self.assertEqual(status.status_code, 200)
        before = self.protected()
        self.confirm(self.preview(bid, self.file(9)))
        self.assertEqual(self.protected(), before)
        p = self.needs(bid)
        self.assertEqual(p['rows'][0]['order_qty'], 9)
        self.assertEqual(p['checklist'][0]['status'], 'reopened')
        self.confirm(self.preview(bid, original), status=409)
        self.assertEqual(self.needs(bid)['rows'][0]['order_qty'], 9)

    def test_prices_remain_required_for_confirmed_purchases(self):
        bid = self.daily(self.file())
        before = self.protected()
        self.assertTrue(self.preview(bid, self.file())['can_confirm'])
        purchase = self.preview(bid, self.file(), plan=False)
        # Daily fixture has a quoted price; remove it to test genuinely absent cost.
        with server.db() as conn:
            conn.execute('UPDATE orders SET buy_price=0')
        purchase = self.preview(bid, self.file(), plan=False)
        self.assertFalse(purchase['can_confirm'])
        self.confirm(purchase, status=400)
        self.assertEqual(before['purchase_workbook_lines'], self.protected()['purchase_workbook_lines'])

    def test_missing_or_wrong_date_sheet_cannot_fall_back_to_sales(self):
        bid = self.daily(self.file(bad_date=True))
        p = self.needs(bid)
        self.assertFalse(p['send_available'])
        self.assertEqual(p['groups'], [])
        self.assertEqual(p['rows'], [])
        self.assertTrue(p['source_issues'])
        self.assertEqual(self.client.get(f'/api/export/suppliers/{bid}').status_code, 409)
        with server.db() as conn:
            conn.execute('DELETE FROM supplier_plan_sources')
        p = self.needs(bid)
        self.assertFalse(p['send_available'])
        self.assertEqual(p['groups'], [])
        self.assertEqual(self.client.get(f'/api/export/suppliers/{bid}').status_code, 409)
        response = self.client.put(f'/api/supplier-order-status/{bid}/S1', json={'status': 'ordered', 'revision': 0})
        self.assertEqual(response.status_code, 409)
        bad = self.client.post('/api/purchase-orders/import/preview', data={
            'batch_id': str(bid), 'plan_only': '1', 'file': (io.BytesIO(self.file(missing=True)), 'source.xlsx')})
        self.assertEqual(bad.status_code, 400)

    def test_stale_preview_cannot_overwrite_new_plan(self):
        bid = self.daily(self.file())
        first = self.preview(bid, self.file(8))
        second = self.preview(bid, self.file(9))
        self.confirm(first)
        self.confirm(second, status=409)
        self.assertEqual(self.needs(bid)['rows'][0]['order_qty'], 8)

    def test_confirming_purchase_replaces_the_plan_once(self):
        bid = self.daily(self.file())
        self.confirm(self.preview(bid, self.file(8, price=10000), plan=False))
        p = self.needs(bid)
        self.assertEqual(p['format'], 'customer_canonical')
        self.assertEqual(p['rows'][0]['order_qty'], 8)
        self.assertTrue(p['rows'][0]['confirmed'])

    def test_changed_daily_upload_updates_the_supplier_sheet(self):
        bid = self.daily(self.file(7))
        self.assertEqual(self.daily(self.file(11)), bid)
        self.assertEqual(self.needs(bid)['rows'][0]['order_qty'], 11)

    def test_supplier_name_unit_and_adjustments_come_from_purchase_sheet(self):
        wb = load_workbook(io.BytesIO(self.file(7)))
        ws = wb['đặt hàng']
        ws.cell(3, 3, 'Tên hàng đặt riêng')
        ws.cell(3, 5, 'Túi')
        ws.cell(3, 6, 'NCC khác')
        ws.cell(3, 9, 1)
        ws.cell(3, 10, 2)
        ws.cell(3, 13, 8)
        stream = io.BytesIO(); wb.save(stream); wb.close()
        bid = self.daily(stream.getvalue())
        row = self.needs(bid)['rows'][0]
        self.assertEqual((row['supplier'], row['product_name'], row['unit'], row['order_qty']),
                         ('NCC khác', 'Tên hàng đặt riêng', 'Túi', 8))

    def test_multiple_purchase_sheets_are_rejected_without_overwriting(self):
        data = self.file()
        bid = self.daily(data)
        before = self.needs(bid)['plan_hash']
        wb = load_workbook(io.BytesIO(data))
        wb.copy_worksheet(wb['đặt hàng']).title = 'DAT HANG'
        stream = io.BytesIO(); wb.save(stream); wb.close()
        result = self.client.post('/api/purchase-orders/import/preview', data={
            'batch_id': str(bid), 'plan_only': '1', 'file': (io.BytesIO(stream.getvalue()), 'ambiguous.xlsx')})
        self.assertEqual(result.status_code, 400)
        self.assertEqual(self.needs(bid)['plan_hash'], before)


if __name__ == '__main__':
    unittest.main()
