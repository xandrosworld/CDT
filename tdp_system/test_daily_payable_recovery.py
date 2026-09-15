"""Historical purchase recovery must never reopen or rewrite posted sales/stock."""
import io
import unittest
from openpyxl import load_workbook
from . import server
from .test_supplier_plan_source import SupplierPlanSourceTests as Fixture


class DailyPayableRecoveryTests(unittest.TestCase):
    setUpClass = classmethod(Fixture.setUpClass.__func__)
    tearDownClass = classmethod(Fixture.tearDownClass.__func__)
    file = Fixture.file
    daily = Fixture.daily
    preview = Fixture.preview
    confirm = Fixture.confirm

    def setUp(self):
        with server.db() as conn:
            conn.execute('DELETE FROM batch_bk_approvals')
        Fixture.setUp(self)

    def lock(self, bid):
        # Missing linked document is also an immutable BK history, preventing
        # accidental stock recreation; snapshot integration covers posted docs.
        with server.db() as conn:
            conn.execute("INSERT INTO batch_bk_approvals VALUES(?,999999,'posted-source','2026-09-01')", (bid,))
            conn.execute("UPDATE batches SET status='approved' WHERE id=?", (bid,))
            server.sync_payable_ledger(conn, timestamp=server.now_iso())

    def analyze(self, data):
        r = self.client.post('/api/import/analyze', data={
            'file': (io.BytesIO(data), 'Đơn hàng 01.09.2026.xlsx'), 'continuous': '1',
        })
        self.assertEqual(r.status_code, 200, r.get_json())
        return r.get_json()

    def apply(self, p, sheets=('đặt hàng',), status=200):
        r = self.client.post('/api/import/confirm', json={
            'token': p['token'], 'state_hash': p['stateHash'],
            'sheets': list(sheets), 'work_date': '2026-09-01',
        })
        self.assertEqual(r.status_code, status, r.get_json())
        return r.get_json()

    def protected(self):
        with server.db() as conn:
            return {t: [tuple(r) for r in conn.execute('SELECT * FROM ' + t + ' ORDER BY 1')]
                    for t in ('orders', 'batches', 'purchase_workbook_lines', 'purchase_order_lines',
                              'receivable_ledger_lines', 'invoice_inventory_ledger', 'inventory_transactions',
                              'batch_bk_approvals', 'bk_import_documents', 'bk_import_lines',
                              'daily_workdays', 'daily_import_versions', 'daily_import_scopes', 'order_import_receipts')}

    def test_exact_original_file_recovers_missing_purchase_source_on_locked_day(self):
        data = self.file(7, 9000)
        bid = self.daily(data)
        self.lock(bid)
        with server.db() as conn:
            conn.execute('DELETE FROM supplier_plan_sources')
            conn.execute('DELETE FROM supplier_plan_history')
            server.sync_payable_ledger(conn, timestamp=server.now_iso())
        before = self.protected()
        p = self.analyze(data)
        self.assertTrue(p['purchasePlanOnly'])
        self.assertEqual(p['batchId'], bid)
        self.assertFalse(next(s for s in p['sheets'] if s['scope']=='customer_orders')['confirmAvailable'])
        self.assertTrue(next(s for s in p['sheets'] if s['scope']=='purchase_orders')['confirmAvailable'])
        r = self.apply(p)
        self.assertTrue(r['purchasePlanOnly'])
        self.assertEqual(before, self.protected())
        with server.db() as conn:
            self.assertEqual(conn.execute("SELECT SUM(amount) FROM payable_ledger_lines WHERE status!='reversed'").fetchone()[0], 63000)
        self.assertTrue(self.apply(self.analyze(data))['idempotent'])
        self.assertEqual(before, self.protected())

    def test_purchase_revision_does_not_reopen_sales_or_rewrite_canonical_purchases(self):
        bid = self.daily(self.file(7, 9000))
        self.confirm(self.preview(bid, self.file(7, 9000), plan=False))
        self.lock(bid)
        before = self.protected()
        self.apply(self.analyze(self.file(8, 11000)))
        self.assertEqual(before, self.protected())
        with server.db() as conn:
            self.assertEqual(conn.execute("SELECT SUM(amount) FROM payable_ledger_lines WHERE status!='reversed'").fetchone()[0], 88000)

    def test_locked_sales_selection_and_stale_plan_are_rejected(self):
        bid = self.daily(self.file(7, 9000)); self.lock(bid)
        before = self.protected()
        self.apply(self.analyze(self.file(8, 9000)), sheets=('01.09',), status=400)
        first, second = self.analyze(self.file(8, 9000)), self.analyze(self.file(9, 9000))
        self.apply(first)
        self.apply(second, status=409)
        self.assertEqual(before, self.protected())

    def test_paid_correction_rolls_back_purchase_source_and_protected_tables(self):
        bid = self.daily(self.file(7, 9000)); self.lock(bid)
        with server.db() as conn:
            conn.execute("UPDATE payable_ledger_lines SET paid_amount=1000,status='partially_paid' WHERE status='open'")
            source = tuple(conn.execute('SELECT * FROM supplier_plan_sources').fetchone())
        before = self.protected()
        self.apply(self.analyze(self.file(8, 9000)), status=409)
        self.assertEqual(before, self.protected())
        with server.db() as conn:
            self.assertEqual(source, tuple(conn.execute('SELECT * FROM supplier_plan_sources').fetchone()))

    def test_lock_added_after_preview_requires_new_preview(self):
        bid = self.daily(self.file(7, 9000))
        p = self.analyze(self.file(8, 9000))
        self.lock(bid)
        self.apply(p, status=409)

    def test_wrong_purchase_date_stays_blocked(self):
        bid = self.daily(self.file(7, 9000)); self.lock(bid)
        p = self.analyze(self.file(8, 9000, bad_date=True))
        purchase = next(s for s in p['sheets'] if s['scope']=='purchase_orders')
        self.assertFalse(purchase['confirmAvailable'])
        self.assertEqual(purchase['errorRows'], 1)
        self.apply(p, status=400)

    def catalog_file(self, *, units=('chai',), name='Product 1', explicit=None):
        w = load_workbook(io.BytesIO(self.file(7, 9000)))
        w['đặt hàng'].cell(3,5,explicit)
        if explicit is None:
            w['đặt hàng'].cell(3,5).value = None
        s = w.create_sheet('danh mục hh')
        s.append(['Mã hàng','Tên hàng','ĐVT','Thuế'])
        for unit in units:
            s.append(['P1',name,unit,'8%'])
        out = io.BytesIO(); w.save(out); w.close()
        return out.getvalue()

    def test_blank_unit_uses_unique_code_and_name_from_same_workbook_only(self):
        bid = self.daily(self.file(7, 9000))
        with server.db() as conn:
            before = [tuple(r) for r in conn.execute('SELECT * FROM products')]
        p = self.preview(bid, self.catalog_file())
        self.assertTrue(p['can_confirm'])
        self.assertEqual(p['issues'][0]['unit'], 'chai')
        self.assertTrue(any('danh mục hh' in x for x in p['issues'][0]['warnings']))
        for data in (self.catalog_file(units=('chai','lon')), self.catalog_file(name='Different')):
            self.assertFalse(self.preview(bid, data)['can_confirm'])
        p = self.preview(bid, self.catalog_file(explicit='kg'))
        self.confirm(p)
        with server.db() as conn:
            import json
            self.assertEqual(json.loads(conn.execute('SELECT items_json FROM supplier_plan_sources').fetchone()[0])[0]['unit'], 'kg')
        with server.db() as conn:
            self.assertEqual(before, [tuple(r) for r in conn.execute('SELECT * FROM products')])
