from __future__ import annotations

import io
import unittest

from openpyxl import load_workbook

try:
    from contract_modules import (
        invoice_delivery_statement_workbook,
    )
    from invoice_payment_documents import (
        InvoicePaymentDocumentError,
        invoice_payment_request_workbook,
    )
except ImportError:  # pragma: no cover
    from .contract_modules import (
        invoice_delivery_statement_workbook,
    )
    from .invoice_payment_documents import (
        InvoicePaymentDocumentError,
        invoice_payment_request_workbook,
    )


class InvoicePaymentDocumentTests(unittest.TestCase):
    @staticmethod
    def drafts(total=216):
        return [{
            "id": 1,
            "contractor": "NT-A",
            "status": "issued",
            "issued_invoice_date": "2026-08-29",
            "issued_invoice_series": "1C26TDP",
            "issued_invoice_number": "00001234",
            "subtotal": 200,
            "tax_amount": 16,
            "total_amount": total,
        }]

    @staticmethod
    def lines():
        base = {
            "draft_id": 1,
            "work_date": "2026-08-29",
            "issued_invoice_date": "2026-08-29",
            "issued_invoice_series": "1C26TDP",
            "issued_invoice_number": "00001234",
            "kitchen": "POT",
            "product_code": "I000060",
            "product_name": "Hành tây",
            "unit": "Kg",
            "qty": 1,
            "unit_price": 100,
            "amount": 100,
            "tax": "8%",
        }
        return [dict(base), {**base, "product_code": "I000061", "product_name": "Hành khô"}]

    @staticmethod
    def payment_scope(total=216):
        return {
            "contractor": "NT-A", "date_from": "2026-08-01", "date_to": "2026-08-31",
            "scope_id": "SCOPE-OFFICIAL",
            "snapshot": {
                "company_name_snapshot": "CÔNG TY TNHH THỰC PHẨM THÀNH ĐẠT PHÁT",
                "company_tax_code_snapshot": "0202265016",
                "company_address_snapshot": "Hải Phòng",
                "buyer_name_snapshot": "CÔNG TY QC BUYER",
                "buyer_tax_code_snapshot": "0200000000",
                "buyer_address_snapshot": "Địa chỉ QC",
                "payment_requester_snapshot": "VŨ THỊ THỤY",
                "payment_bank_name_snapshot": "Vietcombank",
                "payment_bank_account_snapshot": "1052787580",
            },
            "invoices": [{
                "draft_id": 1, "invoice_date": "2026-08-29", "invoice_series": "1C26TDP",
                "invoice_number": "00001234", "subtotal": 200,
                "tax_amount": 16, "total_amount": total,
                "verification_source": "synced_issued_source", "source_invoice_id": 11,
            }],
            "totals": {"subtotal": 200, "tax_amount": 16, "total_amount": 216},
        }

    def test_official_xlsx_matches_approved_six_column_payment_form(self):
        workbook = invoice_payment_request_workbook(self.payment_scope())
        stream = io.BytesIO()
        try:
            workbook.save(stream)
        finally:
            workbook.close()
        stream.seek(0)
        loaded = load_workbook(stream, data_only=False, keep_links=False)
        try:
            self.assertEqual(["Đề nghị thanh toán", "Đối chiếu hóa đơn"], loaded.sheetnames)
            sheet = loaded["Đề nghị thanh toán"]
            self.assertEqual("ĐỀ NGHỊ THANH TOÁN", sheet["A5"].value)
            self.assertIn("CÔNG TY QC BUYER", sheet["A6"].value)
            self.assertEqual("STT", sheet["B10"].value)
            self.assertEqual("Tổng tiền thanh toán", sheet["G10"].value.strip())
            self.assertTrue(all(
                sheet.cell(10, column).fill.fill_type is None for column in range(1, 7)
            ))
            self.assertEqual("00001234", sheet["D11"].value)
            self.assertEqual(216, sheet["G11"].value)
            all_text = "\n".join(
                str(cell.value or "") for ws in loaded.worksheets
                for row in ws.iter_rows() for cell in row
            )
            for required in (
                "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", "CÔNG TY QC BUYER",
                "1052787580", "Hai trăm mười sáu đồng",
            ):
                self.assertIn(required, all_text)
            self.assertNotIn("Q-008", all_text)
            self.assertNotIn("VnTools", all_text)
            self.assertEqual("portrait", sheet.page_setup.orientation)
            self.assertEqual("hidden", loaded["Đối chiếu hóa đơn"].sheet_state)
            self.assertFalse(any(
                cell.data_type == "f" or cell.hyperlink is not None
                for ws in loaded.worksheets for row in ws.iter_rows() for cell in row
            ))
        finally:
            loaded.close()

        with self.assertRaises(InvoicePaymentDocumentError) as mismatch:
            invoice_payment_request_workbook(self.payment_scope(total=218))
        self.assertEqual("payment_document_total_mismatch", mismatch.exception.code)

    def test_customer_template_layout_and_variable_invoice_count(self):
        from copy import deepcopy, copy
        from .document_preview import white_print_style
        scope = self.payment_scope()
        for count in (1, 6, 40):
            current = deepcopy(scope)
            current['invoices'] = [{**scope['invoices'][0], 'invoice_number': str(100 + i)} for i in range(count)]
            current['totals'] = {key: value * count for key, value in scope['totals'].items()}
            book = invoice_payment_request_workbook(current, contract_no='HD-2026', contract_date='2026-06-01')
            ws = book.active
            self.assertIn('Độc lập', ws['E1'].value)
            self.assertIn('HD-2026', ws['A7'].value)
            self.assertEqual('STT', ws['B10'].value)
            self.assertEqual('Tổng cộng', ws.cell(11 + count, 2).value)
            self.assertEqual(216 * count, ws.cell(11 + count, 7).value)
            self.assertIn('1052787580', ws.cell(15 + count, 1).value)
            self.assertIn('ĐẠI DIỆN CÔNG TY', ws.cell(19 + count, 5).value)
            title_font = copy(ws['A5'].font)
            white_print_style(book)
            self.assertEqual(title_font, copy(ws['A5'].font))
            self.assertFalse(any(c.data_type=='f' for row in ws for c in row))
            book.close()

    def test_payment_statement_displays_special_tax_without_changing_source(self):
        from copy import deepcopy, copy
        from .invoice_payment_documents import synced_invoice_statement_workbook
        scope = self.payment_scope()
        scope.update(warning='', source_lines=[dict(invoice_date='2026-08-29', invoice_series='1C26TDP',
            invoice_number='00001234', source_item_name='Hàng', source_unit='Kg', qty=1,
            unit_price=100, amount=100, tax_rate=rate, validation_note='') for rate in ('-2','-2.0','KKKNT','8%','-1')])
        before = deepcopy(scope)
        book = synced_invoice_statement_workbook(scope)
        self.assertEqual(['KKKNT','KKKNT','KKKNT','8%','KCT'], [book['Chi tiết hóa đơn'].cell(r,8).value for r in range(4,9)])
        self.assertEqual(before, scope)
        book.close()

    def test_delivery_statement_reconciles_and_blocks_mismatch(self):
        workbook = invoice_delivery_statement_workbook(self.lines(), self.drafts(), lambda _: 1.08)
        try:
            self.assertEqual(workbook.sheetnames, ["Bảng kê giao hàng", "Đối chiếu hóa đơn"])
            self.assertIn("BẢNG TỔNG HỢP GIAO NHẬN", workbook["Bảng kê giao hàng"]["A5"].value)
            self.assertEqual("STT", workbook["Bảng kê giao hàng"]["A10"].value)
            self.assertEqual("Thanh toán", workbook["Bảng kê giao hàng"]["J10"].value)
            self.assertEqual("dd/mm/yyyy", workbook["Bảng kê giao hàng"]["B11"].number_format)
            self.assertEqual(108, workbook["Bảng kê giao hàng"]["J11"].value)
            self.assertEqual(216, workbook["Bảng kê giao hàng"]["J13"].value)
            reconcile = workbook["Đối chiếu hóa đơn"]
            self.assertEqual(reconcile.cell(4, 10).value, 0)
        finally:
            workbook.close()
        with self.assertRaisesRegex(ValueError, "lệch bảng kê"):
            invoice_delivery_statement_workbook(self.lines(), self.drafts(total=218), lambda _: 1.08)


if __name__ == "__main__":
    unittest.main()
