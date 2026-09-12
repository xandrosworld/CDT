"""Per-contractor operational-receivable Excel exports.

One workbook always belongs to exactly one contractor.  Its first worksheet
reconciles the contractor account and the following worksheets are generated
from the kitchens that actually have active operational delivery lines in the
selected period.  Issued VAT invoices are deliberately outside this export.
"""

from __future__ import annotations

try:
    from .document_preview import white_print_style
except ImportError:
    from document_preview import white_print_style

try:
    from document_totals import quantity_totals, quantity_text, quantity_cell
except ImportError:
    from .document_totals import quantity_totals, quantity_text, quantity_cell

import io
import re
import unicodedata
import zipfile
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from typing import Any, Callable

from flask import jsonify, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins


FORBIDDEN_SHEET_CHARACTERS = re.compile(r"[\\/*?:\[\]]")
FORBIDDEN_FILENAME_CHARACTERS = re.compile(r"[<>:\"/\\|?*\x00-\x1f]")
THIN_GRAY = Side(style="thin", color="000000")
TABLE_BORDER = Border(left=THIN_GRAY, right=THIN_GRAY, top=THIN_GRAY, bottom=THIN_GRAY)
TITLE_FILL = PatternFill(fill_type=None)
HEADER_FILL = PatternFill(fill_type=None)
TOTAL_FILL = PatternFill(fill_type=None)


class ReceivableExportError(ValueError):
    def __init__(self, message: str, *, code: str = "invalid_receivable_export", status: int = 400):
        super().__init__(message)
        self.code = code
        self.status = status


def _clean(value: Any) -> str:
    return " ".join(str(value or "").replace("\u00a0", " ").split())


def _mapping_key(value: Any) -> str:
    normalized = unicodedata.normalize("NFD", _clean(value).casefold().replace("đ", "d"))
    return "".join(character for character in normalized if character.isalnum())


def _strict_date(value: Any, label: str) -> str:
    text = _clean(value)
    try:
        return datetime.strptime(text, "%Y-%m-%d").date().isoformat()
    except ValueError:
        raise ReceivableExportError(
            f"{label} phải là ngày YYYY-MM-DD hợp lệ", code="invalid_period",
        ) from None


def _vnd(value: Any) -> int:
    try:
        number = Decimal(str(value or 0))
    except (InvalidOperation, TypeError, ValueError):
        raise ReceivableExportError(
            "Dữ liệu tiền trong sổ phải thu không hợp lệ",
            code="invalid_receivable_export_money", status=500,
        ) from None
    if not number.is_finite():
        raise ReceivableExportError(
            "Dữ liệu tiền trong sổ phải thu không hữu hạn",
            code="invalid_receivable_export_money", status=500,
        )
    return int(number.quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def _number(value: Any) -> float:
    try:
        number = Decimal(str(value or 0))
    except (InvalidOperation, TypeError, ValueError):
        raise ReceivableExportError(
            "Dữ liệu số lượng trong sổ phải thu không hợp lệ",
            code="invalid_receivable_export_quantity", status=500,
        ) from None
    if not number.is_finite():
        raise ReceivableExportError(
            "Dữ liệu số lượng trong sổ phải thu không hữu hạn",
            code="invalid_receivable_export_quantity", status=500,
        )
    return float(number)


def _excel_date(value: Any) -> date | str:
    try:
        return datetime.strptime(str(value), "%Y-%m-%d").date()
    except (TypeError, ValueError):
        return _clean(value)


def _display_date(value: Any) -> str:
    parsed = _excel_date(value)
    return parsed.strftime("%d/%m/%Y") if isinstance(parsed, date) else str(parsed)


def _party_catalog(conn) -> tuple[dict[str, str], dict[str, str], dict[str, list[str]]]:
    """Return case-folded codes, display names and exact normalized name aliases."""
    codes: dict[str, str] = {}
    names: dict[str, str] = {}
    aliases: dict[str, list[str]] = defaultdict(list)

    def remember(raw_code: Any, raw_name: Any = "", *, master: bool = False) -> None:
        code = _clean(raw_code)
        if not code:
            return
        folded = code.casefold()
        canonical = codes.setdefault(folded, code)
        name = _clean(raw_name)
        if name and (master or canonical not in names):
            names[canonical] = name
        key = _mapping_key(name)
        if key and canonical not in aliases[key]:
            aliases[key].append(canonical)

    for row in conn.execute("SELECT code,name FROM contractors ORDER BY code"):
        remember(row["code"], row["name"], master=True)
    for row in conn.execute(
        """SELECT contractor_code,contractor_snapshot FROM receivable_ledger_lines
           ORDER BY id"""
    ):
        remember(row["contractor_code"], row["contractor_snapshot"])
    for table in ("balances", "debt_adjustments"):
        for row in conn.execute(
            f"SELECT DISTINCT party_code FROM {table} WHERE party_type='contractor' ORDER BY party_code"
        ):
            remember(row["party_code"])
    for row in conn.execute(
        """SELECT DISTINCT party_code FROM payments
            WHERE party_type='contractor' AND kind='receipt' ORDER BY party_code"""
    ):
        remember(row["party_code"])
    return codes, names, aliases


def _resolve_contractor(conn, requested: Any) -> tuple[str, dict[str, str], dict[str, str]]:
    codes, names, aliases = _party_catalog(conn)
    value = _clean(requested)
    if not value:
        return "", codes, names
    direct = codes.get(value.casefold())
    if direct:
        return direct, codes, names
    matches = aliases.get(_mapping_key(value), [])
    if len(matches) == 1:
        return matches[0], codes, names
    if len(matches) > 1:
        raise ReceivableExportError(
            "Tên nhà thầu khớp nhiều mã; hãy chọn đúng mã",
            code="ambiguous_contractor",
        )
    raise ReceivableExportError(
        "Nhà thầu không tồn tại", code="invalid_contractor",
    )


def _canonical_account_values(
    raw_accounts: dict[str, dict[str, Any]], codes: dict[str, str],
) -> dict[str, dict[str, int]]:
    result: dict[str, dict[str, int]] = {}
    fields = ("opening", "period_charge", "period_adjustment", "period_paid", "closing")
    for raw_code, raw_values in raw_accounts.items():
        code = codes.get(_clean(raw_code).casefold(), _clean(raw_code))
        target = result.setdefault(code, {field: 0 for field in fields})
        for field in fields:
            target[field] += _vnd(raw_values.get(field))
    return result


def receivable_export_data(
    conn, *, date_from: Any, date_to: Any, contractor: Any,
    debt_period_payload: Callable, tax_factor: Callable,
) -> dict[str, Any]:
    """Build a read-only snapshot for one or more separate contractor files."""
    safe_from = _strict_date(date_from, "Từ ngày")
    safe_to = _strict_date(date_to, "Đến ngày")
    if safe_from > safe_to:
        raise ReceivableExportError(
            "Từ ngày không được lớn hơn Đến ngày", code="invalid_period",
        )
    contractor_code, codes, contractor_names = _resolve_contractor(conn, contractor)
    account_payload = debt_period_payload(conn, safe_from, safe_to, tax_factor)
    accounts = _canonical_account_values(account_payload.get("contractors", {}), codes)

    kitchen_names = {
        _clean(row["code"]): _clean(row["name"])
        for row in conn.execute("SELECT code,name FROM kitchens ORDER BY code")
        if _clean(row["code"])
    }
    rows = []
    for row in conn.execute(
        """SELECT * FROM receivable_ledger_lines
            WHERE status='active' AND work_date>=? AND work_date<=?
            ORDER BY contractor_code,kitchen_code,work_date,product_name,id""",
        (safe_from, safe_to),
    ):
        item = dict(row)
        code = codes.get(_clean(item["contractor_code"]).casefold(), _clean(item["contractor_code"]))
        if contractor_code and code != contractor_code:
            continue
        item["contractor_code"] = code
        item["contractor_name"] = (
            contractor_names.get(code) or _clean(item["contractor_snapshot"]) or code
        )
        item["kitchen_code"] = _clean(item["kitchen_code"])
        item["kitchen_name"] = (
            kitchen_names.get(item["kitchen_code"])
            or _clean(item["kitchen_snapshot"])
            or item["kitchen_code"]
            or "Chưa xác định"
        )
        for field in ("ordered_qty", "actual_delivered", "customer_return_qty", "delivered_qty", "sell_price", "tax_percent"):
            item[field] = _number(item[field])
        for field in ("subtotal", "tax_amount", "amount"):
            item[field] = _vnd(item[field])
        rows.append(item)

    selected_codes: set[str]
    if contractor_code:
        selected_codes = {contractor_code}
    else:
        selected_codes = {row["contractor_code"] for row in rows if row["contractor_code"]}
        selected_codes.update(
            code for code, values in accounts.items() if any(values.values())
        )

    exports = []
    empty_account = {
        "opening": 0, "period_charge": 0, "period_adjustment": 0,
        "period_paid": 0, "closing": 0,
    }
    for code in sorted(selected_codes, key=lambda value: value.casefold()):
        contractor_rows = [row for row in rows if row["contractor_code"] == code]
        name = contractor_names.get(code)
        if not name and contractor_rows:
            name = contractor_rows[0]["contractor_name"]
        account = dict(accounts.get(code, empty_account))
        expected_closing = (
            account["opening"] + account["period_charge"]
            + account["period_adjustment"] - account["period_paid"]
        )
        if expected_closing != account["closing"]:
            raise ReceivableExportError(
                f"Số dư cuối kỳ nhà thầu {code} không đối soát",
                code="receivable_export_balance_mismatch", status=409,
            )
        detail_charge = sum(row["amount"] for row in contractor_rows)
        if detail_charge != account["period_charge"]:
            raise ReceivableExportError(
                f"Phát sinh nhà thầu {code} lệch giữa sổ tài khoản và chi tiết giao hàng",
                code="receivable_export_reconciliation_failed", status=409,
            )

        kitchen_groups: dict[str, dict[str, Any]] = {}
        for row in contractor_rows:
            kitchen_key = row["kitchen_code"] or "CHUA_XAC_DINH"
            group = kitchen_groups.setdefault(kitchen_key, {
                "code": row["kitchen_code"] or "CHƯA XÁC ĐỊNH",
                "name": row["kitchen_name"], "rows": [], "line_count": 0,
                "ordered_qty": 0.0, "actual_delivered": 0.0,
                "customer_return_qty": 0.0, "delivered_qty": 0.0,
                "subtotal": 0, "tax_amount": 0, "amount": 0,
            })
            group["rows"].append(row)
            group["line_count"] += 1
            for field in ("ordered_qty", "actual_delivered", "customer_return_qty", "delivered_qty"):
                group[field] += row[field]
            for field in ("subtotal", "tax_amount", "amount"):
                group[field] += row[field]
        kitchens = sorted(
            kitchen_groups.values(), key=lambda item: (item["code"].casefold(), item["name"].casefold()),
        )
        if sum(item["amount"] for item in kitchens) != account["period_charge"]:
            raise ReceivableExportError(
                f"Tổng các bếp của nhà thầu {code} không khớp sheet tổng",
                code="receivable_export_kitchen_mismatch", status=409,
            )
        exports.append({
            "date_from": safe_from,
            "date_to": safe_to,
            "contractor_code": code,
            "contractor_name": name or code,
            "account": account,
            "kitchens": kitchens,
            "source_of_truth": "approved operational delivery lines; not issued VAT invoices",
        })
    return {
        "date_from": safe_from,
        "date_to": safe_to,
        "contractor": contractor_code or None,
        "contractors": exports,
        "source_of_truth": "approved operational delivery lines; not issued VAT invoices",
    }


def _sheet_name(kitchen: dict[str, Any], used: set[str]) -> str:
    label = " - ".join(
        value for value in (_clean(kitchen.get("code")), _clean(kitchen.get("name"))) if value
    ) or "Bếp chưa xác định"
    base = FORBIDDEN_SHEET_CHARACTERS.sub("-", label).strip(" '") or "Bếp chưa xác định"
    candidate = base[:31]
    serial = 2
    while candidate.casefold() in used:
        suffix = f" ({serial})"
        candidate = base[:31 - len(suffix)] + suffix
        serial += 1
    used.add(candidate.casefold())
    return candidate


def _filename_token(value: Any) -> str:
    token = FORBIDDEN_FILENAME_CHARACTERS.sub("_", _clean(value)).strip(" .")
    return token[:60] or "NHA_THAU"


def _style_title(ws, *, title: str, subtitle: str, end_col: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_col)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end_col)
    ws.cell(1, 1, title)
    ws.cell(2, 1, subtitle)
    ws.cell(1, 1).font = Font(name="Times New Roman", size=16, bold=True)
    ws.cell(1, 1).fill = TITLE_FILL
    ws.cell(1, 1).alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(2, 1).font = Font(name="Times New Roman", size=11, italic=True)
    ws.cell(2, 1).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 26
    ws.row_dimensions[2].height = 32


def _style_header(ws, row_number: int, end_col: int) -> None:
    for cell in ws[row_number][:end_col]:
        cell.font = Font(name="Times New Roman", size=11, bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = TABLE_BORDER
    ws.row_dimensions[row_number].height = 32


def _style_body(ws, *, start_row: int, end_col: int) -> None:
    for row in ws.iter_rows(min_row=start_row, max_row=ws.max_row, max_col=end_col):
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"
            cell.font = Font(name="Times New Roman", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = TABLE_BORDER


def _print_setup(ws, *, header_row: int, end_col: int, print_end_col: int | None = None) -> None:
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(end_col)}{max(ws.max_row, header_row)}"
    ws.print_title_rows = f"{header_row}:{header_row}"
    printable = print_end_col or end_col
    ws.print_area = f"A1:{get_column_letter(printable)}{max(ws.max_row, header_row)}"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.scale = None
    ws.page_margins = PageMargins(left=0.25, right=0.25, top=0.5, bottom=0.5, header=0.2, footer=0.2)
    ws.print_options.horizontalCentered = True
    ws.oddFooter.center.text = "Trang &P / &N"
    ws.sheet_view.showGridLines = False


def receivable_workbook(data: dict[str, Any]) -> Workbook:
    """Create one static workbook for exactly one contractor."""
    account = data["account"]
    period = f"Từ {_display_date(data['date_from'])} đến {_display_date(data['date_to'])}"
    workbook = Workbook()
    workbook.properties.title = f"Công nợ phải thu {data['contractor_code']} {period}"
    workbook.properties.subject = "Công nợ vận hành theo lượng thực giao và giá bán giao dịch"
    workbook.properties.creator = "Thành Đạt Phát"

    ws = workbook.active
    ws.title = "Tổng nhà thầu"
    account_headers = [
        "Số dư đầu kỳ", "Phát sinh phải thu", "Điều chỉnh", "Đã thu", "Số dư cuối kỳ",
        "Nguồn", "Từ ngày", "Đến ngày",
    ]
    ws.append([])
    ws.append([])
    ws.append(account_headers)
    ws.append([
        account["opening"], account["period_charge"], account["period_adjustment"],
        account["period_paid"], account["closing"], "Thực giao đã duyệt (không phải hóa đơn đỏ)",
        _excel_date(data["date_from"]), _excel_date(data["date_to"]),
    ])
    ws.append([])
    kitchen_headers = [
        "Mã bếp", "Tên bếp", "Số dòng", "SL đặt", "Thực giao", "Khách trả",
        "Giao ròng", "Tiền trước thuế", "Tiền thuế", "Phát sinh phải thu",
    ]
    ws.append(kitchen_headers)
    for kitchen in data["kitchens"]:
        ws.append([
            kitchen["code"], kitchen["name"], kitchen["line_count"], quantity_cell(kitchen["rows"], "ordered_qty"),
            quantity_cell(kitchen["rows"], "actual_delivered"), quantity_cell(kitchen["rows"], "customer_return_qty"), quantity_cell(kitchen["rows"], "delivered_qty"),
            kitchen["subtotal"], kitchen["tax_amount"], kitchen["amount"],
        ])
    total_row = ws.max_row + 1
    if data["kitchens"]:
        ws.append([
            "TỔNG", data["contractor_name"], sum(item["line_count"] for item in data["kitchens"]),
            quantity_cell([row for item in data["kitchens"] for row in item["rows"]], "ordered_qty"),
            quantity_cell([row for item in data["kitchens"] for row in item["rows"]], "actual_delivered"),
            quantity_cell([row for item in data["kitchens"] for row in item["rows"]], "customer_return_qty"),
            quantity_cell([row for item in data["kitchens"] for row in item["rows"]], "delivered_qty"),
            sum(item["subtotal"] for item in data["kitchens"]),
            sum(item["tax_amount"] for item in data["kitchens"]),
            sum(item["amount"] for item in data["kitchens"]),
        ])
    _style_title(
        ws,
        title=f"CÔNG NỢ PHẢI THU – {data['contractor_code']}",
        subtitle=(
            f"{data['contractor_name']} · {period} · Công nợ vận hành; "
            "không phải đề nghị thanh toán/hóa đơn đỏ"
        ),
        end_col=len(kitchen_headers),
    )
    _style_body(ws, start_row=4, end_col=len(kitchen_headers))
    _style_header(ws, 3, len(account_headers))
    _style_header(ws, 6, len(kitchen_headers))
    # The account block has eight columns while the kitchen table below has
    # ten.  Do not draw two empty cells (or a fully boxed spacer row) beside
    # the account balance; they looked like missing fields in the printout.
    for column in range(9, len(kitchen_headers) + 1):
        ws.cell(4, column).border = Border()
    for column in range(1, len(kitchen_headers) + 1):
        ws.cell(5, column).border = Border()
    for column in (1, 2, 3, 4, 5):
        ws.cell(4, column).number_format = "#,##0"
    for column in (7, 8):
        ws.cell(4, column).number_format = "dd/mm/yyyy"
    for row_number in range(7, ws.max_row + 1):
        for column in (4, 5, 6, 7):
            ws.cell(row_number, column).number_format = "#,##0.######"
        for column in (8, 9, 10):
            ws.cell(row_number, column).number_format = "#,##0"
    if data["kitchens"]:
        for cell in ws[total_row]:
            cell.font = Font(name="Times New Roman", size=11, bold=True)
            cell.fill = TOTAL_FILL
    for index, width in enumerate([16, 28, 12, 14, 14, 14, 14, 18, 16, 20], 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A7"
    _print_setup(ws, header_row=6, end_col=len(kitchen_headers))

    used_names = {"tổng nhà thầu"}
    detail_headers = [
        "Ngày", "Mã bếp", "Tên bếp", "Mã hàng", "Tên hàng", "SL đặt", "Thực giao",
        "Khách trả", "Giao ròng", "ĐVT", "Giá bán giao dịch", "Thuế suất (%)",
        "Tiền trước thuế", "Tiền thuế", "Phát sinh phải thu", "Mã dòng", "Lần cập nhật", "Nguồn",
    ]
    for kitchen in data["kitchens"]:
        ws = workbook.create_sheet(_sheet_name(kitchen, used_names))
        ws.append(detail_headers)
        for row in kitchen["rows"]:
            ws.append([
                _excel_date(row["work_date"]), row["kitchen_code"] or "CHƯA XÁC ĐỊNH",
                row["kitchen_name"], row["product_code"], row["product_name"],
                row["ordered_qty"], row["actual_delivered"], row["customer_return_qty"],
                row["delivered_qty"], row["unit"], row["sell_price"], row["tax_percent"],
                row["subtotal"], row["tax_amount"], row["amount"], row["id"],
                row["revision"], row["source_ref"],
            ])
        total_row = ws.max_row + 1
        ws.append([
            "TỔNG", kitchen["code"], kitchen["name"], "", "", quantity_cell(kitchen["rows"], "ordered_qty"),
            quantity_cell(kitchen["rows"], "actual_delivered"), quantity_cell(kitchen["rows"], "customer_return_qty"), quantity_cell(kitchen["rows"], "delivered_qty"),
            "", "", "", kitchen["subtotal"], kitchen["tax_amount"], kitchen["amount"], "", "", "",
        ])
        ws.insert_rows(1, 2)
        total_row += 2
        _style_title(
            ws,
            title=f"CHI TIẾT PHẢI THU – {kitchen['code']}",
            subtitle=(
                f"{data['contractor_code']} · {kitchen['name']} · {period} · "
                f"Phát sinh {kitchen['amount']:,} VND"
            ),
            end_col=len(detail_headers),
        )
        _style_header(ws, 3, len(detail_headers))
        _style_body(ws, start_row=4, end_col=len(detail_headers))
        for row_number in range(4, ws.max_row + 1):
            ws.cell(row_number, 1).number_format = "dd/mm/yyyy"
            for column in (6, 7, 8, 9):
                ws.cell(row_number, column).number_format = "#,##0.######"
            for column in (11, 13, 14, 15):
                ws.cell(row_number, column).number_format = "#,##0"
            ws.cell(row_number, 12).number_format = "0.###"
        for cell in ws[total_row]:
            cell.font = Font(name="Times New Roman", size=11, bold=True)
            cell.fill = TOTAL_FILL
        for index, width in enumerate(
            [13, 15, 26, 14, 30, 12, 12, 12, 12, 9, 18, 15, 18, 16, 20, 12, 11, 28], 1,
        ):
            ws.column_dimensions[get_column_letter(index)].width = width
        ws.freeze_panes = "A4"
        _print_setup(ws, header_row=3, end_col=len(detail_headers), print_end_col=15)

    workbook.active = 0
    return workbook


def _workbook_bytes(workbook: Workbook) -> bytes:
    stream = io.BytesIO()
    white_print_style(workbook).save(stream)
    workbook.close()
    return stream.getvalue()


def register_receivable_export_routes(app, ctx: dict[str, Any]) -> None:
    db_factory: Callable = ctx["db"]
    debt_period_payload: Callable = ctx["debt_period_payload"]
    tax_factor: Callable = ctx["tax_factor"]

    @app.get("/api/debts/receivables/export")
    def api_export_receivables():
        today = date.today()
        try:
            requested_contractor = request.args.get("contractor") or ""
            with db_factory() as conn:
                conn.execute("BEGIN")
                data = receivable_export_data(
                    conn,
                    date_from=request.args.get("from") or today.replace(day=1).isoformat(),
                    date_to=request.args.get("to") or today.isoformat(),
                    contractor=requested_contractor,
                    debt_period_payload=debt_period_payload,
                    tax_factor=tax_factor,
                )
                if requested_contractor:
                    contractor_data = data["contractors"][0]
                    payload = _workbook_bytes(receivable_workbook(contractor_data))
                    return send_file(
                        io.BytesIO(payload), as_attachment=True,
                        download_name=(
                            f"Cong_no_phai_thu_{_filename_token(contractor_data['contractor_code'])}_"
                            f"{data['date_from']}_{data['date_to']}.xlsx"
                        ),
                        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )

                bundle = io.BytesIO()
                used_files: set[str] = set()
                with zipfile.ZipFile(bundle, "w", zipfile.ZIP_DEFLATED) as archive:
                    for index, contractor_data in enumerate(data["contractors"], 1):
                        base = (
                            f"Cong_no_phai_thu_{_filename_token(contractor_data['contractor_code'])}_"
                            f"{data['date_from']}_{data['date_to']}"
                        )
                        filename = f"{base}.xlsx"
                        serial = 2
                        while filename.casefold() in used_files:
                            filename = f"{base}_{serial}.xlsx"
                            serial += 1
                        used_files.add(filename.casefold())
                        archive.writestr(filename, _workbook_bytes(receivable_workbook(contractor_data)))
                    if not data["contractors"]:
                        archive.writestr(
                            "KHONG_PHAT_SINH.txt",
                            (
                                f"Không có công nợ phải thu vận hành từ {data['date_from']} "
                                f"đến {data['date_to']}.\n"
                                "Nguồn chỉ gồm lượng thực giao đã duyệt; không dùng hóa đơn đỏ.\n"
                            ).encode("utf-8"),
                        )
                bundle.seek(0)
                return send_file(
                    bundle, as_attachment=True,
                    download_name=f"Cong_no_phai_thu_{data['date_from']}_{data['date_to']}.zip",
                    mimetype="application/zip",
                )
        except ReceivableExportError as error:
            return jsonify({"ok": False, "error": str(error), "code": error.code}), error.status
