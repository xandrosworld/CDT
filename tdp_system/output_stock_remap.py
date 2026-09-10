"""Excel round trip for internal inventory codes; issued invoice fields are immutable."""
import hashlib
import io
import json
import sqlite3
import time
import uuid
import zipfile
from collections import defaultdict

from flask import jsonify, request, send_file
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from xml.etree.ElementTree import ParseError

try:
    from .invoice_monthly_valuation import monthly_average_report
    from .stock_tax_policy import kkknt_codes
    from .invoice_line_tax import annotate_invoice_tax
    from .invoice_output_editing import output_amount_review
    from .invoice_inventory import _minimum_balance_from
    from .inventory_tax_review import compare_tax, display_tax, format_review_sheet
except ImportError:
    from invoice_monthly_valuation import monthly_average_report
    from stock_tax_policy import kkknt_codes
    from invoice_line_tax import annotate_invoice_tax
    from invoice_output_editing import output_amount_review
    from invoice_inventory import _minimum_balance_from
    from inventory_tax_review import compare_tax, display_tax, format_review_sheet

HEADERS = ['ID dòng kho', 'Ngày', 'Ký hiệu / Số HĐ', 'Mã trên HĐ', 'Tên trên HĐ',
           'ĐVT trên HĐ', 'Số lượng HĐ', 'Đơn giá bán', 'Tiền hàng', 'Thuế suất',
           'Tiền thuế', 'Mã nội bộ hiện tại', 'Tên nội bộ hiện tại', 'ĐVT kho',
           'Lượng trừ kho', 'Tồn cuối kỳ', 'Mã nội bộ mới', 'Tên nội bộ mới']
# Keep the original layout readable for files already being edited by customers.
NEW_HEADERS = ['Mã hàng muốn chuyển sang', 'Tên hàng muốn chuyển sang',
               'Mã nội bộ đang trừ kho', 'Tên hàng đang trừ kho', 'Lượng chuyển',
               'ĐVT kho', 'Tồn cuối kỳ', 'Ngày hóa đơn', 'Ký hiệu / Số HĐ',
               'Mã trên hóa đơn gốc (có thể trống)', 'Tên trên hóa đơn gốc',
               'ĐVT trên HĐ', 'Số lượng HĐ', 'Đơn giá bán', 'Tiền hàng',
               'Thuế suất', 'Tiền thuế', 'ID dòng kho']
NEW_COLUMN_ORDER = (16, 17, 11, 12, 14, 13, 15, 1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 0)
SCHEMA = '''
CREATE TABLE IF NOT EXISTS output_stock_remaps (
    ledger_id INTEGER PRIMARY KEY REFERENCES invoice_inventory_ledger(id),
    product_code TEXT NOT NULL REFERENCES products(code),
    revision INTEGER NOT NULL, updated_at TEXT NOT NULL, unit_cost REAL
);
CREATE TABLE IF NOT EXISTS output_stock_excel_sessions (
    token TEXT PRIMARY KEY, kind TEXT NOT NULL, payload TEXT NOT NULL,
    created REAL NOT NULL, result TEXT
);
'''


def init_schema(conn):
    conn.executescript(SCHEMA)
    if 'unit_cost' not in {r['name'] for r in conn.execute('PRAGMA table_info(output_stock_remaps)')}:
        conn.execute('ALTER TABLE output_stock_remaps ADD COLUMN unit_cost REAL')
    columns = [r['name'] for r in conn.execute('PRAGMA table_info(invoice_inventory_ledger)')]
    select = ','.join(f'COALESCE(m.{col},l.{col}) AS {col}'
                      if col in {'product_code','unit_cost'} else 'l.' + col for col in columns)
    conn.execute('DROP VIEW IF EXISTS invoice_inventory_effective_ledger')
    conn.execute('CREATE VIEW IF NOT EXISTS invoice_inventory_effective_ledger AS SELECT ' + select +
                 ' FROM invoice_inventory_ledger l LEFT JOIN output_stock_remaps m ON m.ledger_id=l.id')


class RemapError(ValueError):
    pass


def _json(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(',', ':'))


def _hash(value):
    return hashlib.sha256(_json(value).encode()).hexdigest()


def _excel_value(value):
    # Excel saves numbers to 15 significant digits. Compare the exported
    # representation, while retaining the full source precision in the snapshot
    # and database. This is not a quantity or monetary edit from the workbook.
    if isinstance(value, float):
        return float(format(value, '.15g'))
    return '' if value is None else value


def _snapshot(conn, start, end):
    report = monthly_average_report(conn, date_from=start, date_to=end, include_zero=True)
    stock = {r['product_code']: r for r in report['items']}
    rows = []
    raw_hashes = {}
    source_allowed = {}
    for r in conn.execute('''SELECT l.*,i.invoice_series,i.invoice_number,i.stock_status,
            i.source_status_class,i.sync_status,i.source,i.error_message,i.invoice_date,
            i.subtotal,i.total_amount,i.tax_amount header_tax,i.buyer_name,i.buyer_tax_code,
            li.source_item_code,li.source_item_name,
            li.source_unit,li.qty,li.unit_price,li.amount,li.tax_rate,i.raw_json,
            p.name,p.unit,COALESCE(m.revision,0) remap_revision
          FROM invoice_inventory_effective_ledger l
          JOIN outgoing_source_invoices i ON i.id=l.source_invoice_id
          JOIN outgoing_source_invoice_items li ON li.id=l.source_line_id AND li.invoice_id=i.id
          JOIN products p ON p.code=l.product_code
          LEFT JOIN output_stock_remaps m ON m.ledger_id=l.id
          WHERE l.source_invoice_table='outgoing_source_invoices' AND l.direction='output'
            AND l.event_type='POST' AND i.source='minvoice' AND i.stock_status='posted'
            AND i.source_status_class='issued'
            AND l.txn_date>=? AND l.txn_date<=? ORDER BY l.txn_date,l.id''', (start, end)):
        item = dict(r)
        invoice_id = r['source_invoice_id']
        if invoice_id not in source_allowed:
            source_allowed[invoice_id] = r['sync_status'] == 'synced' or output_amount_review(item) is not None
        if not source_allowed[invoice_id]:
            continue
        tax_line = dict(item, line_index=r['source_line_index'])
        annotate_invoice_tax({'items': [tax_line]}, r['raw_json'])
        if invoice_id not in raw_hashes:
            raw_hashes[invoice_id] = hashlib.sha256(r['raw_json'].encode()).hexdigest()
        item.pop('raw_json')
        item['source_json_hash'] = raw_hashes[invoice_id]
        item['closing_qty'] = stock[r['product_code']]['closing_qty']
        item['cells'] = [r['id'], r['txn_date'], str(r['invoice_series'] or '') + ' / ' + str(r['invoice_number'] or ''),
                         r['source_item_code'] or '', r['source_item_name'] or '', r['source_unit'] or '',
                         r['qty'], r['unit_price'], r['amount'], tax_line['tax_rate'], tax_line['line_tax_amount'],
                         r['product_code'], r['name'], r['unit'], -r['qty_delta'], item['closing_qty']]
        rows.append(item)
    return {'from': start, 'to': end, 'rows': rows, 'stock': report['items'],
            'catalog': [dict(r) for r in conn.execute('SELECT code,name,unit,tax FROM products ORDER BY code')],
            'closures': [dict(r) for r in conn.execute('SELECT period,status,revision FROM inventory_period_closures ORDER BY period')],
            'holds': [dict(r) for r in conn.execute("SELECT id,product_code,qty_out,status FROM inventory_transactions WHERE source_type='OUTGOING_DRAFT' ORDER BY id")],
            'local_invoices': [dict(r) for r in conn.execute("SELECT id,status,issued_invoice_number,issued_invoice_series,issued_invoice_date FROM outgoing_invoice_drafts WHERE status!='cancelled' ORDER BY id")],
            'openings': [dict(r) for r in conn.execute("SELECT id,txn_date,product_code,qty_in,qty_out,unit_cost FROM inventory_transactions WHERE source_type='OPENING' AND status='posted' ORDER BY id")]}


def _editable_period(conn, start):
    if conn.execute("SELECT 1 FROM inventory_period_closures WHERE period>=? AND status='closed' LIMIT 1", (start[:7],)).fetchone():
        raise RemapError('Kỳ này hoặc kỳ sau đã chốt. Mở lại các kỳ liên quan trước khi đổi mã xuất kho.')
    if conn.execute("SELECT 1 FROM inventory_transactions WHERE source_type='OPENING' AND status='posted' AND txn_date>? LIMIT 1", (start,)).fetchone():
        raise RemapError('Đã có tồn đầu kỳ sau. Cần mở lại kỳ đã chuyển tồn trước khi đổi mã xuất kho.')


def _store(conn, kind, payload):
    conn.execute('DELETE FROM output_stock_excel_sessions WHERE created<?', (time.time() - 7 * 86400,))
    token = uuid.uuid4().hex
    conn.execute('INSERT INTO output_stock_excel_sessions(token,kind,payload,created) VALUES(?,?,?,?)',
                 (token, kind, _json(payload), time.time()))
    return token


def _load(conn, token, kind, ttl):
    row = conn.execute('SELECT * FROM output_stock_excel_sessions WHERE token=? AND kind=?', (token, kind)).fetchone()
    if not row or time.time() - row['created'] > ttl:
        raise RemapError('File hoặc lượt xem trước đã hết hạn. Tải Excel mới và kiểm tra lại.')
    return row, json.loads(row['payload'])


def export_workbook(conn, start, end):
    if not conn.in_transaction:
        conn.execute('BEGIN')
    snapshot = _snapshot(conn, start, end)
    _editable_period(conn, snapshot['from'])
    token = _store(conn, 'export', snapshot)
    wb = Workbook(); ws = wb.active; ws.title = 'Doi ma xuat kho'
    for area, text in [
        ('A1:G1', 'CHỈ SỬA 2 CỘT VÀNG A–B: Mã hàng muốn chuyển sang và Tên hàng muốn chuyển sang.'),
        ('H1:Q1', f"{len(snapshot['rows'])} dòng xuất / {len({r['product_code'] for r in snapshot['rows']})} mã hàng. Một mã có thể lặp nhiều dòng; không cộng lặp cột Tồn cuối kỳ."),
        ('A2:G2', 'Sao chép cặp mã + tên từ sheet Danh muc ma hang, rồi tải file lên hệ thống. Giữ nguyên dòng không cần đổi.'),
        ('H2:Q2', 'Để đối chiếu với báo cáo tồn, mở sheet Tổng hợp hàng âm: mỗi mã một dòng, lọc Thuế danh mục. Thuế trên HĐ có thể khác danh mục.'),
        ('A3:B3', '1. SỬA HÀNG MUỐN CHUYỂN SANG'),
        ('C3:G3', '2. ĐỐI CHIẾU HÀNG ĐANG TRỪ KHO'),
        ('H3:Q3', '3. THÔNG TIN NGUỒN · KHÔNG SỬA'),
    ]:
        ws.merge_cells(area)
        cell = ws[area.split(':')[0]]; cell.value = text
        cell.alignment = Alignment(wrap_text=True, vertical='center')
        cell.font = Font(bold=cell.row != 2, color='17354A', size=11)
        cell.fill = PatternFill('solid', fgColor='FFF2CC' if cell.column == 1 else 'EDF2F5')
    ws.append(NEW_HEADERS)
    for r in snapshot['rows']:
        values = [_excel_value(v) for v in r['cells']] + [r['product_code'], r['name']]
        ws.append([values[i] for i in NEW_COLUMN_ORDER])
    ws.freeze_panes = 'E5'; ws.auto_filter.ref = f'A4:R{max(ws.max_row,4)}'
    ws.sheet_view.showGridLines = False
    for row in ws.iter_rows(min_row=4):
        for cell in row:
            if isinstance(cell.value, str): cell.data_type = 's'
            editable = cell.column <= 2
            source = cell.column >= 8
            cell.alignment = Alignment(wrap_text=True, vertical='center')
            cell.fill = PatternFill('solid', fgColor='FFF2CC' if editable else 'F2F4F6' if source else 'EFF7FB')
            cell.font = Font(bold=cell.row == 4, color='17354A' if not source else '4B5563', size=11)
            if editable and cell.row > 4:
                cell.protection = Protection(locked=False)
            if cell.row > 4 and cell.column in (5,7,13,14,15,17):
                cell.number_format = '#,##0.######;[Red]-#,##0.######'
    for index, width in enumerate((25,38,23,38,14,12,16,16,23,29,38,14,16,20,20,14,20,14),1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.column_dimensions['R'].hidden = True
    for row, height in ((1,34),(2,34),(3,26),(4,46)): ws.row_dimensions[row].height = height
    ws['J4'].comment = Comment('Đây là mã từ hóa đơn nguồn, có thể trống. Mã nội bộ dùng trừ kho nằm ở cột C. Không điền mã nội bộ vào cột này.', 'TDP')
    ws['A4'].comment = Comment('Sao chép mã và tên cùng một hàng trong sheet Danh muc ma hang vào hai cột A–B. Mỗi dòng chuyển toàn bộ lượng ở cột E.', 'TDP')
    ws['G4'].comment = Comment('Số tồn của mã ở cột C, lặp lại ở các dòng xuất cùng mã. Không cộng cột này; xem sheet Tổng hợp hàng âm để đếm theo mã.', 'TDP')
    ws['P4'].comment = Comment('Thuế từ dòng hóa đơn gốc. Báo cáo tồn ưu tiên thuế danh mục; xem cả hai nguồn ở sheet Tổng hợp hàng âm.', 'TDP')
    summary = wb.create_sheet('Tổng hợp hàng âm')
    catalog_by_code = {p['code']:p for p in snapshot['catalog']}
    by_code = defaultdict(list)
    first_rows = {}
    for index, row in enumerate(snapshot['rows'], 5):
        by_code[row['product_code']].append(row)
        first_rows.setdefault(row['product_code'], index)
    negatives = [r for r in snapshot['stock'] if r['closing_qty'] < -1e-9]
    summary.append([f'{len(negatives)} mã tồn âm · Mỗi mã chỉ một dòng. Lọc cột Thuế danh mục để so với báo cáo tồn.'])
    summary.merge_cells('A1:I1')
    summary.append(['Mã nội bộ','Tên hàng','Thuế danh mục (như báo cáo tồn)','Thuế HĐ trong file đổi mã',
                    'ĐVT','Tồn cuối kỳ (mỗi mã một lần)','Số dòng xuất có thể đổi','Đối chiếu thuế','Mở dòng xuất'])
    for item in negatives:
        code = item['product_code']; lines = by_code[code]
        review = compare_tax(catalog_by_code[code], [{'tax_rate':r['cells'][9]} for r in lines])
        summary.append([code, catalog_by_code[code]['name'], review['catalog'], review['source'], item['unit'],
                        item['closing_qty'], len(lines), review['note'],
                        'Đến dòng xuất đầu tiên; có thể lọc mã ở cột C' if lines else 'Không có dòng xuất trong file đổi mã'])
        if lines:
            summary.cell(summary.max_row,9).hyperlink = f"#'Doi ma xuat kho'!A{first_rows[code]}"
        if review['differs']:
            for cell in summary[summary.max_row]: cell.fill = PatternFill('solid',fgColor='FFF2CC')
        summary.cell(summary.max_row,6).number_format = '#,##0.######;[Red]-#,##0.######'
    format_review_sheet(summary, 2, (18,40,26,26,12,25,20,62,48))
    catalog = wb.create_sheet('Danh muc ma hang'); catalog.append(['Mã nội bộ', 'Tên nội bộ', 'ĐVT', 'Thuế'])
    for p in snapshot['catalog']: catalog.append([p[k] for k in ('code','name','unit')] + [display_tax(p['tax'])])
    catalog.auto_filter.ref = catalog.dimensions; catalog.freeze_panes = 'A2'; catalog.column_dimensions['B'].width = 48
    guide = wb.create_sheet('Huong dan')
    for line in ['Chỉ sửa hai cột vàng A–B: Mã hàng muốn chuyển sang và Tên hàng muốn chuyển sang, theo sheet Danh muc ma hang.',
                 'Cột C–G là hàng đang trừ kho và tồn để đối chiếu. Cột H–Q là thông tin hóa đơn gốc, không sửa.',
                 'Mã trên hóa đơn gốc có thể trống do nguồn M-Invoice không có mã; không có nghĩa là thiếu mã nội bộ.',
                 'File đổi mã giữ từng dòng xuất hóa đơn. Một mã có thể xuất nhiều dòng; không cộng lặp Tồn cuối kỳ.',
                 'Sheet Tổng hợp hàng âm gom mỗi mã một dòng. Lọc Thuế danh mục để so với báo cáo tồn; Thuế HĐ là nguồn riêng, có thể khác danh mục.',
                 'Một dòng tương ứng toàn bộ lượng trừ kho của một dòng hóa đơn. Không đổi lượng, tiền, thuế hay nội dung hóa đơn.',
                 'Có thể lọc các dòng Tồn cuối kỳ âm; giữ nguyên các dòng không cần đổi. Không xóa dòng hoặc cột.',
                 'Tải file lên để xem trước. Hệ thống kiểm tra đơn vị, tồn mã nhận, kỳ chốt và dữ liệu đã thay đổi.',
                 'KKKNT được lập bảng kê và chuyển nguyên tồn âm. KCT và 0% không thuộc ngoại lệ này.',
                 'Mã âm từ đầu kỳ không có dòng xuất: đối chiếu tồn đầu; KKKNT có thể lập bảng kê mua vào bổ sung theo nguồn thực tế.']:
        guide.append([line])
    guide.column_dimensions['A'].width = 130
    meta = wb.create_sheet('_meta'); meta.append(['token', token]); meta.sheet_state = 'veryHidden'
    for sheet in (catalog, guide):
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str): cell.data_type = 's'
    output = io.BytesIO(); wb.save(output); return output.getvalue()


def _set_overrides(conn, changes, timestamp):
    for c in changes:
        conn.execute('''INSERT INTO output_stock_remaps(ledger_id,product_code,revision,updated_at,unit_cost) VALUES(?,?,1,?,?)
                       ON CONFLICT(ledger_id) DO UPDATE SET product_code=excluded.product_code,
                       revision=output_stock_remaps.revision+1,updated_at=excluded.updated_at,unit_cost=excluded.unit_cost''',
                     (c['ledger_id'], c['new_code'], timestamp, c.get('new_unit_cost')))


def _evaluate(conn, snapshot, changes):
    _editable_period(conn, snapshot['from'])
    conn.execute('SAVEPOINT remap_preview')
    try:
        _set_overrides(conn, changes, 'preview')
        report = monthly_average_report(conn, date_from=snapshot['from'], date_to=snapshot['to'], include_zero=True)
        stocks = {r['product_code']: r for r in report['items']}
        exempt = kkknt_codes(conn)
        try:
            from .outgoing_readiness import canonical_available_stock
        except ImportError:
            from outgoing_readiness import canonical_available_stock
        available = canonical_available_stock(conn)
        # A later receipt must not conceal a shortage in an earlier month-after
        # checkpoint. Check each recipient once, including all active holds.
        future_available = {}
        for code in {c['new_code'] for c in changes} - exempt:
            holds = available.get(code, {})
            future_available[code] = (_minimum_balance_from(conn, code, snapshot['to'])
                                      - holds.get('reserved_qty', 0)
                                      - holds.get('pending_sync_issued_qty', 0))
        errors = []
        for c in changes:
            target = stocks[c['new_code']]
            c['old_closing_after'] = stocks[c['old_code']]['closing_qty']
            c['new_closing_after'] = target['closing_qty']
            c['new_unit_cost'] = target['average_unit_cost']
            if c['new_unit_cost'] < 0:
                errors.append(f"Mã nhận {c['new_code']} có giá tồn không hợp lệ; đối chiếu tồn đầu và nhập trước khi chuyển.")
            if c['new_code'] not in exempt and target['closing_qty'] < -1e-9:
                errors.append(f"Mã nhận {c['new_code']} sẽ âm {target['closing_qty']:g}; chọn mã khác hoặc giảm số dòng chuyển.")
            elif c['new_code'] not in exempt and future_available[c['new_code']] < -1e-9:
                errors.append(f"Mã nhận {c['new_code']} không đủ tồn tại một thời điểm từ cuối kỳ trở đi, đã trừ dự thảo đang giữ. Chọn mã khác hoặc kiểm tra giao dịch tháng sau và dự thảo.")
        return sorted(set(errors))
    finally:
        conn.execute('ROLLBACK TO remap_preview'); conn.execute('RELEASE remap_preview')


def preview_workbook(conn, data):
    if not conn.in_transaction:
        conn.execute('BEGIN')
    if len(data) > 10 * 1024 * 1024:
        raise RemapError('File vượt 10 MB.')
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            if sum(i.file_size for i in z.infolist()) > 80 * 1024 * 1024:
                raise RemapError('File Excel giải nén quá lớn.')
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
        token = wb['_meta']['B1'].value
        _, snapshot = _load(conn, token, 'export', 7 * 86400)
        if _hash(_snapshot(conn, snapshot['from'], snapshot['to'])) != _hash(snapshot):
            raise RemapError('Dữ liệu kho, danh mục hoặc hóa đơn đã thay đổi từ lúc tải file. Tải Excel mới để sửa.')
        ws = wb['Doi ma xuat kho']
        if ws.max_row > 30000 or ws.max_column != len(HEADERS):
            raise RemapError('Bố cục file không hợp lệ.')
        if [c.value for c in next(ws.iter_rows())] == HEADERS:
            first_data_row, column_order = 2, tuple(range(len(HEADERS)))
        elif [c.value for c in next(ws.iter_rows(min_row=4,max_row=4))] == NEW_HEADERS:
            first_data_row, column_order = 5, NEW_COLUMN_ORDER
        else:
            raise RemapError('Không được đổi tiêu đề hoặc thứ tự cột.')
        originals = {r['id']: r for r in snapshot['rows']}
        products = {p['code']: p for p in snapshot['catalog']}
        seen = set(); changes = []; errors = []
        for excel_row, cells in enumerate(ws.iter_rows(min_row=first_data_row), first_data_row):
            values = [None] * len(HEADERS)
            for cell, original_column in zip(cells,column_order):
                values[original_column] = cell.value
            if all(v is None for v in values): continue
            if any(c.data_type == 'f' for c in cells):
                raise RemapError(f'Dòng {excel_row}: không nhận công thức; dán giá trị vào hai cột vàng.')
            key = values[0]
            if key not in originals or key in seen:
                raise RemapError(f'Dòng {excel_row}: ID không thuộc file hoặc bị lặp.')
            seen.add(key); old = originals[key]
            if [_excel_value(v) for v in values[:16]] != [_excel_value(v) for v in old['cells']]:
                raise RemapError(f'Dòng {excel_row}: đã sửa cột gốc. Chỉ được đổi mã và tên nội bộ mới.')
            code, name = str(values[16] or '').strip(), str(values[17] or '').strip()
            p = products.get(code)
            if not p or name != p['name']:
                errors.append(f'Dòng {excel_row}: mã/tên mới không khớp danh mục. Sao chép cặp mã và tên từ sheet Danh muc ma hang.'); continue
            if code == old['product_code']: continue
            if str(p['unit'] or '').strip().casefold() != str(old['unit'] or '').strip().casefold():
                errors.append(f'Dòng {excel_row}: đơn vị mã nhận {p["unit"]} khác đơn vị kho {old["unit"]}.'); continue
            changes.append({'ledger_id': key, 'line_id': old['source_line_id'], 'invoice_id': old['source_invoice_id'],
                            'old_code': old['product_code'], 'old_name': old['name'], 'new_code': code,
                            'new_name': name, 'qty': -old['qty_delta'], 'unit': old['unit'], 'excel_row': excel_row})
        if seen != set(originals): raise RemapError('File bị thiếu dòng. Giữ nguyên các dòng không cần đổi mã.')
        if not changes and not errors: errors.append('Chưa có mã nội bộ nào thay đổi.')
        if changes and not errors: errors.extend(_evaluate(conn, snapshot, changes))
        preview = {'snapshot': snapshot, 'changes': changes, 'errors': errors}
        preview_token = _store(conn, 'preview', preview)
        return {'token': preview_token, 'changes': changes, 'errors': errors, 'can_confirm': bool(changes) and not errors,
                'from': snapshot['from'], 'to': snapshot['to']}
    except (KeyError, zipfile.BadZipFile, InvalidFileException, ParseError, ValueError, TypeError) as exc:
        if isinstance(exc, RemapError): raise
        raise RemapError('File không phải mẫu đổi mã đã tải từ hệ thống hoặc chứa dữ liệu không hợp lệ.') from exc
    finally:
        if 'wb' in locals(): wb.close()


def confirm_preview(conn, token, actor, timestamp):
    row, preview = _load(conn, token, 'preview', 900)
    if row['result']: return {**json.loads(row['result']), 'idempotent': True}
    if not actor or len(actor) > 100: raise RemapError('Cần tên người xác nhận, tối đa 100 ký tự.')
    if preview['errors'] or not preview['changes']: raise RemapError('File còn lỗi hoặc chưa có thay đổi.')
    snapshot, changes = preview['snapshot'], preview['changes']
    if _hash(_snapshot(conn, snapshot['from'], snapshot['to'])) != _hash(snapshot):
        raise RemapError('Dữ liệu đã thay đổi sau khi xem trước. Tải Excel mới để đối chiếu.')
    errors = _evaluate(conn, snapshot, changes)
    if errors: raise RemapError('; '.join(errors))
    conn.execute('SAVEPOINT apply_stock_remap')
    try:
        _set_overrides(conn, changes, timestamp)
        for c in changes:
            # These are local mapping fields only. Source names, codes, amounts,
            # quantities, tax, raw source JSON and issued invoice headers stay intact.
            conn.execute('UPDATE outgoing_source_invoice_items SET product_code=? WHERE id=? AND invoice_id=?',
                         (c['new_code'], c['line_id'], c['invoice_id']))
        conn.execute("INSERT INTO audit_log(event_type,entity_type,entity_id,status,message,metadata_json,created_at) VALUES('inventory.output.remap','excel',?,'ok',?,?,?)",
                     (token, 'Đổi mã nội bộ theo Excel', _json({'actor': actor, 'from': snapshot['from'], 'to': snapshot['to'], 'changes': changes}), timestamp))
        result = {'changed_lines': len(changes), 'idempotent': False}
        conn.execute('UPDATE output_stock_excel_sessions SET result=? WHERE token=?', (_json(result), token))
        conn.execute('RELEASE apply_stock_remap')
        return result
    except Exception:
        conn.execute('ROLLBACK TO apply_stock_remap'); conn.execute('RELEASE apply_stock_remap'); raise


def register_routes(app, ctx):
    @app.get('/api/inventory/output-remap/export')
    def output_remap_export():
        try:
            with ctx['db']() as conn:
                data = export_workbook(conn, request.args.get('from'), request.args.get('to'))
            return send_file(io.BytesIO(data), mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                             as_attachment=True, download_name='Doi_ma_noi_bo_xuat_kho.xlsx')
        except ValueError as exc: return jsonify(ok=False, error=str(exc)), 400
        except sqlite3.Error: return jsonify(ok=False, error='Dữ liệu đang được cập nhật. Thử tải Excel lại sau ít giây.'), 409

    @app.post('/api/inventory/output-remap/preview')
    def output_remap_preview():
        try:
            file = request.files.get('file')
            if not file: raise RemapError('Chọn file Excel đã sửa.')
            with ctx['db']() as conn:
                result = preview_workbook(conn, file.read(10 * 1024 * 1024 + 1))
            return jsonify(ok=True, **result)
        except ValueError as exc: return jsonify(ok=False, error=str(exc)), 400
        except sqlite3.Error: return jsonify(ok=False, error='Dữ liệu đang được cập nhật. Chọn lại file để kiểm tra.'), 409

    @app.post('/api/inventory/output-remap/confirm')
    def output_remap_confirm():
        try:
            body = request.get_json(silent=True) or {}
            if body.get('confirmed') is not True: raise RemapError('Cần xác nhận sau khi xem trước.')
            with ctx['db']() as conn:
                conn.execute('BEGIN IMMEDIATE')
                result = confirm_preview(conn, body.get('token'), str(body.get('actor') or '').strip(), ctx['now_iso']())
            return jsonify(ok=True, **result)
        except ValueError as exc: return jsonify(ok=False, error=str(exc)), 409
        except sqlite3.Error: return jsonify(ok=False, error='Chưa lưu được đổi mã; dữ liệu đã hoàn tác. Kiểm tra lại rồi xác nhận.'), 409
