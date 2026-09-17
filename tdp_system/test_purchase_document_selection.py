import json
import sqlite3
import unittest
from decimal import Decimal
from unittest.mock import patch

from openpyxl import Workbook

from . import purchase_document_selection as selection
from .test_receipt_export import receipt_row
from .receipt_export import build_purchase_documents_workbook, configure_receipt_paper
from .test_receipt_export import GOLDEN
from .print_bundle import workbook_sections


class PurchaseDocumentSelectionTests(unittest.TestCase):
    def setUp(self):
        self.conn = sqlite3.connect(':memory:')
        self.conn.row_factory = sqlite3.Row
        self.conn.execute('CREATE TABLE batches(id INTEGER, work_date TEXT, status TEXT)')
        self.conn.executemany('INSERT INTO batches VALUES(?,?,?)', [
            (1, '2026-09-14', 'approved'), (2, '2026-09-14', 'approved')])
        selection.init_schema(self.conn)
        self.rows = [receipt_row(work_date='2026-09-14', batch_id=i, selection_key=str(i),
                                 product_name='Hàng ' + str(i), quantity=10, amount=3000000,
                                 buy_price=300000) for i in (1, 2)]
        self.mock = patch.object(selection, 'day_source', side_effect=lambda *_:
                                 (self.rows, selection.digest(self.rows)))
        self.mock.start()
        self.addCleanup(self.mock.stop)
        self.addCleanup(self.conn.close)
        self.audit = lambda *a, **kw: None

    def body(self, quantities):
        state = selection.scope(self.conn, '2026-09-14', '2026-09-14')
        return {**state, 'actor': 'Kiểm thử', 'days': [
            {'date': '2026-09-14', 'quantities': quantities}]}

    def save(self, body):
        return selection.save(self.conn, body, '2026-09-16T19:00:00', self.audit)

    def test_over_limit_across_batches_blocked_and_not_chosen_automatically(self):
        state = selection.scope(self.conn, '2026-09-14', '2026-09-14')
        self.assertEqual(6000000, state['days'][0]['groups'][0]['pending'])
        self.assertTrue(all(r['selected_quantity'] == '0' for r in state['days'][0]['rows']))
        with self.assertRaisesRegex(ValueError, '5.000.000'):
            self.save(self.body({'1': '10', '2': '10'}))
        self.assertEqual(0, self.conn.execute('SELECT COUNT(*) FROM purchase_document_selections').fetchone()[0])
        with self.assertRaises(selection.SelectionError):
            selection.export_scope(self.conn, {'id': 1, 'work_date': '2026-09-14', 'status': 'approved'}, self.rows[:1])

    def test_partial_reconciles_and_all_batch_exports_share_same_plan(self):
        result = self.save(self.body({'1': '10', '2': '5'}))
        self.assertEqual(4500000, result['days'][0]['groups'][0]['selected'])
        self.assertEqual(1500000, result['days'][0]['groups'][0]['pending'])
        selected, pending, plan = selection.export_scope(self.conn,
            {'id': 2, 'work_date': '2026-09-14', 'status': 'approved'}, self.rows[1:])
        self.assertEqual(5, selected[0]['quantity'])
        self.assertEqual(1500000, pending[0]['amount'])
        book = selection.annotate_workbook(Workbook(), plan, pending)
        self.assertIn('Chờ bổ sung chứng từ', book.sheetnames)
        self.assertEqual(6000000, book['Đối chiếu lựa chọn']['C2'].value)
        self.assertIn('thay thế', book.worksheets[0]['A3'].value)

    def test_empty_choice_retains_everything_as_pending(self):
        result = self.save(self.body({}))
        self.assertEqual(6000000, result['days'][0]['groups'][0]['pending'])

    def test_individual_receipt_omits_internal_note_but_keeps_reconciliation_and_payment_field(self):
        self.save(self.body({'1': '10', '2': '5'}))
        selected, pending, plan = selection.export_scope(self.conn,
            {'id': 1, 'work_date': '2026-09-14', 'status': 'approved'}, self.rows[:1])
        book = selection.annotate_workbook(build_purchase_documents_workbook(selected, template_path=GOLDEN), plan, pending)
        receipt = book['biên nhận']
        self.assertFalse(any(str(c.value or '').startswith('Lựa chọn ngày ') for row in receipt for c in row))
        for paper in ('A4', 'A5'):
            # Older immutable preview workbooks may still carry the old note.
            footer = receipt.max_row + 2
            receipt.cell(footer, 3, 'Lựa chọn ngày 2026-09-14 · Tổng mua cả ngày: 6,000,000đ')
            configure_receipt_paper(receipt, paper)
            self.assertIsNone(receipt.cell(footer, 3).value)
            self.assertFalse(str(receipt.print_area).endswith('$G$' + str(footer)))
        self.assertEqual(6000000, book['Đối chiếu lựa chọn']['C2'].value)
        self.assertEqual(1500000, book['Đối chiếu lựa chọn']['E2'].value)
        values = [str(c.value or '') for r in receipt for c in r]
        self.assertFalse(any('thanh toán tiền mặt' in s for s in values))
        self.assertTrue(any('Hình thức / ngày thanh toán' in s for s in values))
        sections = workbook_sections('purchases', book)
        self.assertEqual(4, len(sections))
        self.assertFalse(any('Lựa chọn ngày ' in str(n) for n in sections[1]['notes']))
        book.close()

    def test_retries_revisions_and_concurrency(self):
        body = self.body({'1': '10.0'})
        self.save(body)
        self.assertTrue(self.save(body)['unchanged'])
        replacement = self.body({'2': '10'})
        with self.assertRaisesRegex(selection.SelectionError, 'lý do'):
            self.save(replacement)
        replacement['reason'] = 'Đổi hàng đã chọn sau đối chiếu'
        self.assertEqual(2, self.save(replacement)['days'][0]['revision'])
        body['reason'] = 'Trình duyệt cũ'
        with self.assertRaisesRegex(selection.SelectionError, 'vừa sửa'):
            self.save(body)
        self.assertEqual(2, self.conn.execute('SELECT COUNT(*) FROM purchase_document_selection_history').fetchone()[0])

    def test_changed_source_blocks_old_document_and_silent_replacement(self):
        self.save(self.body({'1': '10'}))
        self.rows[0]['amount'] += 100
        self.assertTrue(selection.scope(self.conn, '2026-09-14', '2026-09-14')['days'][0]['stale'])
        with self.assertRaisesRegex(selection.SelectionError, 'thay đổi'):
            selection.export_scope(self.conn, {'id': 1, 'work_date': '2026-09-14', 'status': 'approved'}, self.rows)
        with self.assertRaisesRegex(selection.SelectionError, 'đã đổi'):
            self.save(self.body({'1': '10'}))

    def test_invalid_quantities_and_unknown_lines_rejected(self):
        for quantity in ('-1', '11', 'NaN', 'Infinity', '0.0000001', True):
            with self.subTest(quantity=quantity), self.assertRaises(selection.SelectionError):
                selection.split_rows(self.rows, {'1': quantity})
        with self.assertRaises(selection.SelectionError):
            selection.split_rows(self.rows, {'unknown': 1})

    def test_fractional_money_conserved_without_rounding_whole_lines(self):
        self.rows[0].update(amount=100.25, quantity=3)
        chosen, pending = selection.split_rows(self.rows[:1], {'1': 1})
        self.assertEqual(Decimal('100.25'), Decimal(str(chosen[0]['amount'])) + Decimal(str(pending[0]['amount'])))
        chosen, pending = selection.split_rows(self.rows[:1], {'1': 3})
        self.assertEqual(100.25, chosen[0]['amount'])
        self.assertFalse(pending)

    def test_pending_survives_next_day_and_does_not_reset_allowance_or_write(self):
        self.save(self.body({'1': '10', '2': '5'}))
        before = self.conn.total_changes
        today = selection.pending_queue(self.conn, '2026-09-14')
        tomorrow = selection.pending_queue(self.conn, '2026-09-15')
        self.assertEqual(today['entries'], tomorrow['entries'])
        day = tomorrow['entries'][0]
        self.assertEqual('2026-09-14', day['date'])
        self.assertEqual('saved', day['status'])
        self.assertEqual(1500000, day['pending'])
        self.assertEqual(500000, day['groups'][0]['remaining_capacity'])
        self.assertEqual(5, day['groups'][0]['rows'][0]['quantity'])
        self.assertEqual('2026-09-14', day['groups'][0]['rows'][0]['work_date'])
        self.assertEqual(before, self.conn.total_changes)
        self.assertFalse(selection.pending_queue(self.conn, '2026-09-13')['entries'])

    def test_pending_unsaved_overlimit_is_visible_and_completed_selection_disappears(self):
        queue = selection.pending_queue(self.conn, '2026-09-15')
        self.assertEqual('needs_selection', queue['entries'][0]['status'])
        self.rows[0]['amount'] = 1000000
        self.save(self.body({'1': '10'}))
        self.assertEqual(3000000, selection.pending_queue(self.conn, '2026-09-15')['entries'][0]['pending'])
        replacement = self.body({'1': '10', '2': '10'})
        replacement['reason'] = 'Đã đối chiếu đủ hàng'
        self.save(replacement)
        self.assertFalse(selection.pending_queue(self.conn, '2026-09-15')['entries'])

    def test_pending_keeps_stale_and_missing_sources_visible(self):
        self.save(self.body({'1': '10'}))
        self.rows[0]['quantity'] = 1  # Old saved quantity must never imply a negative remainder.
        day = selection.pending_queue(self.conn, '2026-09-15')['entries'][0]
        self.assertEqual('needs_review', day['status'])
        self.assertNotIn('pending', day)
        self.conn.execute("UPDATE batches SET status='draft'")
        day = selection.pending_queue(self.conn, '2026-09-15')['entries'][0]
        self.assertEqual('needs_review', day['status'])
        self.assertIn('không còn', day['error'])

    def test_pending_one_invalid_day_does_not_hide_other_days(self):
        self.conn.execute("INSERT INTO batches VALUES(3,'2026-09-15','approved')")
        def source(conn, work_date):
            if work_date == '2026-09-14':
                raise selection.SelectionError('Thiếu CCCD')
            rows = [{**r, 'work_date': work_date} for r in self.rows]
            return rows, selection.digest(rows)
        with patch.object(selection, 'day_source', side_effect=source):
            queue = selection.pending_queue(self.conn, '2026-09-15')
        self.assertEqual(['needs_review', 'needs_selection'], [d['status'] for d in queue['entries']])

    def test_pending_pagination_never_silently_truncates_old_days(self):
        self.conn.execute('DELETE FROM batches')
        self.conn.executemany('INSERT INTO batches VALUES(?,?,?)',
            [(i, f'2026-08-{i:02}', 'approved') for i in range(1, 26)])
        def source(conn, work_date):
            rows = [{**r, 'work_date': work_date} for r in self.rows]
            return rows, selection.digest(rows)
        with patch.object(selection, 'day_source', side_effect=source):
            first = selection.pending_queue(self.conn, '2026-09-15')
            second = selection.pending_queue(self.conn, '2026-09-15', first['next_after'])
        self.assertEqual(20, len(first['entries']))
        self.assertEqual(5, len(second['entries']))
        self.assertIsNone(second['next_after'])
        self.assertEqual(25, len({d['date'] for d in first['entries'] + second['entries']}))

    def test_pending_rejects_invalid_cursor(self):
        for through, after in [('wrong', ''), ('2026-09-15', 'bad'), ('2026-09-15', '2026-09-16')]:
            with self.subTest(through=through, after=after), self.assertRaises(selection.SelectionError):
                selection.pending_queue(self.conn, through, after)


if __name__ == '__main__':
    unittest.main()
