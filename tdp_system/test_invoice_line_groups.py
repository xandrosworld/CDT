"""Selected groups preserve invoice evidence and posted stock value."""
import sqlite3
import unittest
import tempfile
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from contextlib import contextmanager, closing
from unittest.mock import patch

from flask import Flask
from openpyxl import load_workbook

from . import invoice_line_groups as groups
from .invoice_receipt import create_input_receipt
from .invoice_valuation import moving_average_report
from .inventory_export import collect_inventory_export_model, build_movement_workbook, _group_input_movements
from .inventory_preview import workbook_preview
from .invoice_workbench_listing import invoice_range_payload, range_workbook
from .test_invoice_input_sync import init_test_database, now_iso
from .test_invoice_receipt_summary import seed_promotion


class SelectedGroupTests(unittest.TestCase):
    def setUp(self):
        self.conn=sqlite3.connect(':memory:')
        init_test_database(self.conn)
        self.fixture=seed_promotion(self.conn)
        self.ids=[self.fixture['promotion_lines'][key] for key in ('QA-OIL-PAID','QA-OIL-FREE')]
        self.conn.commit()
        @contextmanager
        def db():
            with self.conn:
                yield self.conn
        self.app=Flask(__name__)
        groups.register_group_routes(self.app,{'db':db,'now_iso':now_iso})
        self.client=self.app.test_client()

    def tearDown(self):
        self.conn.close()

    def payload(self):
        return invoice_range_payload(self.conn,tenant='TDP',invoice_type='input',date_from='2026-08-01',date_to='2026-08-31')

    def create(self):
        preview=self.client.post('/api/invoice-workbench/input-groups/preview',json={'item_ids':self.ids})
        self.assertEqual(200,preview.status_code,preview.json)
        result=self.client.post('/api/invoice-workbench/input-groups',json={'item_ids':self.ids,'token':preview.json['token']})
        self.assertEqual(200,result.status_code,result.json)
        return result.json

    def test_selected_pair_one_row_excel_split_and_post(self):
        before=[tuple(r) for r in self.conn.execute('SELECT * FROM msmi_invoice_items')]
        changes=self.conn.total_changes
        preview=self.client.post('/api/invoice-workbench/input-groups/preview',json={'item_ids':self.ids}).json
        self.assertEqual((30,1288889),(preview['qty'],preview['amount']))
        self.assertAlmostEqual(1288889/30,preview['unit_cost'])
        self.assertEqual(changes,self.conn.total_changes)
        saved=self.create()
        payload=self.payload()
        self.assertEqual([],payload['group_warnings'])
        self.assertEqual((3,4),(len(payload['lines']),payload['source_line_count']))
        merged=next(r for r in payload['lines'] if r.get('group_id'))
        self.assertEqual((30,1288889,2),(merged['qty'],merged['amount'],len(merged['group_members'])))
        wb=load_workbook(range_workbook(payload))
        main=list(wb['Hoa don'].values)
        self.assertEqual(1,sum(r[9]=='QA-OIL' for r in main))
        self.assertEqual(3,wb['Dong goc da gop'].max_row)
        self.assertEqual(4553704,payload['totals']['line_amount'])
        self.assertEqual(200,self.client.post(f"/api/invoice-workbench/input-groups/{saved['id']}/split").status_code)
        self.assertEqual(4,len(self.payload()['lines']))
        self.create()
        create_input_receipt(self.conn,self.fixture['promotion_invoice'],now_iso)
        create_input_receipt(self.conn,self.fixture['promotion_invoice'],now_iso)
        self.assertEqual(4,self.conn.execute('SELECT COUNT(*) FROM invoice_inventory_ledger').fetchone()[0])
        report=moving_average_report(self.conn,date_from='2026-08-01',date_to='2026-08-31')
        oil=next(r for r in report['items'] if r['product_code']=='QA-OIL')
        self.assertEqual((30,1288889),(oil['closing_qty'],oil['closing_value']))
        after=[tuple(r) for r in self.conn.execute('SELECT * FROM msmi_invoice_items')]
        self.assertEqual(before,after)
        self.assertEqual(3,len(self.payload()['lines']))

    def selection(self, factor=1):
        body={'item_ids':self.ids,'product_code':'QA-OIL','factors':{str(i):factor for i in self.ids}}
        response=self.client.post('/api/invoice-workbench/input-groups/selection-preview',json=body)
        self.assertEqual(200,response.status_code,response.json)
        return {**body,'token':response.json['token']},response.json

    def test_posted_selected_groups_match_input_workbook_and_preview_without_ledger_changes(self):
        self.create()
        self.ids=[self.fixture['promotion_lines'][key] for key in ('QA-CHILI-PAID','QA-CHILI-FREE')]
        self.create()
        create_input_receipt(self.conn,self.fixture['promotion_invoice'],now_iso)
        before=list(self.conn.iterdump())
        model=collect_inventory_export_model(self.conn,date_from='2026-08-01',date_to='2026-08-31')
        rows=model['input_rows'];self.assertEqual(2,len(rows))
        self.assertEqual([(30,1288889),(107,3264815)],[(r['quantity'],r['amount']) for r in rows])
        self.assertAlmostEqual(1288889/30,rows[0]['unit_cost'],places=5)
        self.assertAlmostEqual(3264815/107,rows[1]['unit_cost'],places=5)
        self.assertEqual(4,sum(len(r['group_members']) for r in rows))
        wb=build_movement_workbook(model,direction='input')
        self.assertEqual((30,1288889,107,3264815),(wb.active['K5'].value,wb.active['M5'].value,wb.active['K6'].value,wb.active['M6'].value))
        preview=workbook_preview(wb,'input.xlsx')
        cells=preview['sheets'][preview['sheetOrder'][0]]['cellData']
        self.assertEqual((30,1288889,107,3264815),(cells[4][10]['v'],cells[4][12]['v'],cells[5][10]['v'],cells[5][12]['v']))
        wb.close()
        self.assertEqual(before,list(self.conn.iterdump()))
        self.assertEqual(4,self.conn.execute('SELECT COUNT(*) FROM invoice_inventory_ledger').fetchone()[0])

    def test_input_report_does_not_merge_unselected_or_split_rows(self):
        group=self.create()
        create_input_receipt(self.conn,self.fixture['promotion_invoice'],now_iso)
        model=collect_inventory_export_model(self.conn,date_from='2026-08-01',date_to='2026-08-31')
        self.assertEqual(3,len(model['input_rows']))
        totals=model['totals']
        groups.split_group(self.conn,'TDP',group['id'],now_iso())
        after=collect_inventory_export_model(self.conn,date_from='2026-08-01',date_to='2026-08-31')
        self.assertEqual(4,len(after['input_rows']));self.assertEqual(totals,after['totals'])

    def test_input_report_rejects_stale_partial_or_cross_confirmation_groups(self):
        self.create();create_input_receipt(self.conn,self.fixture['promotion_invoice'],now_iso)
        model=collect_inventory_export_model(self.conn,date_from='2026-08-01',date_to='2026-08-31')
        merged=next(r for r in model['input_rows'] if r.get('group_id'))
        members=merged['group_members']
        self.assertEqual(1,len(_group_input_movements(self.conn,members)))
        self.assertEqual(1,len(_group_input_movements(self.conn,members[:1])))
        modified=[dict(r) for r in members];modified[1]['confirmation_id']+=1
        self.assertEqual(2,len(_group_input_movements(self.conn,modified)))
        self.conn.execute('UPDATE msmi_invoice_items SET source_item_name=source_item_name||? WHERE id=?',(' changed',self.ids[0]))
        self.assertEqual(2,len(_group_input_movements(self.conn,members)))

    def test_select_unmapped_first_cancel_then_atomic_mapping_merge_retry_split(self):
        self.conn.execute("UPDATE msmi_invoice_items SET product_code=NULL,mapping_status='unmapped',conversion_factor=NULL,stock_qty=0 WHERE id IN (?,?)",self.ids)
        self.conn.commit()
        before=list(self.conn.iterdump())
        first=self.client.post('/api/invoice-workbench/input-groups/selection-preview',json={'item_ids':self.ids})
        self.assertEqual(200,first.status_code)
        self.assertFalse(first.json['ready'])
        body,p=self.selection()
        self.assertEqual((30,1288889),(p['qty'],p['amount']))
        self.assertEqual(before,list(self.conn.iterdump()),'Opening/preview/cancel must not save mappings')
        response=self.client.post('/api/invoice-workbench/input-groups/save-selection',json=body)
        self.assertEqual(200,response.status_code,response.json)
        self.assertEqual(3,len(self.payload()['lines']))
        merged=next(r for r in self.payload()['lines'] if r.get('group_id'))
        self.assertEqual((30,1288889),(merged['qty'],merged['amount']))
        self.assertEqual(0,self.conn.execute('SELECT COUNT(*) FROM invoice_inventory_ledger').fetchone()[0])
        changes=self.conn.total_changes
        retry=self.client.post('/api/invoice-workbench/input-groups/save-selection',json=body)
        self.assertTrue(retry.json['idempotent'])
        self.assertEqual(changes,self.conn.total_changes)
        self.client.post(f"/api/invoice-workbench/input-groups/{response.json['id']}/split")
        self.assertEqual(4,len(self.payload()['lines']))
        self.assertEqual(409,self.client.post('/api/invoice-workbench/input-groups/save-selection',json=body).status_code)

    def test_select_first_fraction_conversion_rollback_and_stale_preview(self):
        self.conn.execute("UPDATE msmi_invoice_items SET source_unit='Thùng',product_code=NULL,mapping_status='unmapped' WHERE id IN (?,?)",self.ids)
        self.conn.commit()
        body,p=self.selection(.5)
        self.assertEqual(15,p['qty'])
        before=list(self.conn.iterdump())
        with patch.object(groups,'_audit',side_effect=RuntimeError('QA rollback after mapping')):
            r=self.client.post('/api/invoice-workbench/input-groups/save-selection',json=body)
        self.assertEqual(500,r.status_code)
        self.assertEqual(before,list(self.conn.iterdump()))
        self.conn.execute('UPDATE msmi_invoice_items SET amount=amount+1 WHERE id=?',(self.ids[0],))
        self.conn.commit()
        self.assertEqual(409,self.client.post('/api/invoice-workbench/input-groups/save-selection',json=body).status_code)
        self.assertEqual(0,self.conn.execute('SELECT COUNT(*) FROM invoice_input_group_choices').fetchone()[0])

    def test_select_first_validates_factors_product_tax_and_posted(self):
        body,p=self.selection()
        for factor in [0,-1,'NaN','Infinity','',None,'1e999',True]:
            invalid={**body,'factors':{str(i):factor for i in self.ids}}
            self.assertEqual(409,self.client.post('/api/invoice-workbench/input-groups/selection-preview',json=invalid).status_code,factor)
        self.assertEqual(409,self.client.post('/api/invoice-workbench/input-groups/selection-preview',json={**body,'product_code':'MISSING'}).status_code)
        self.conn.execute('UPDATE msmi_invoice_items SET tax_rate=99 WHERE id=?',(self.ids[0],))
        self.conn.commit()
        self.assertEqual(409,self.client.post('/api/invoice-workbench/input-groups/selection-preview',json=body).status_code)
        self.conn.execute('UPDATE msmi_invoice_items SET tax_rate=8 WHERE id IN (?,?)',self.ids)
        self.conn.execute("UPDATE msmi_invoices SET receipt_status='posted' WHERE id=?",(self.fixture['promotion_invoice'],))
        self.conn.commit()
        self.assertEqual(409,self.client.post('/api/invoice-workbench/input-groups/save-selection',json=body).status_code)

    def test_selected_mapping_survives_resync_but_explicit_edit_can_replace_it(self):
        from .invoice_mapping import apply_saved_mappings, _update_line_snapshot
        body,p=self.selection(2)
        response=self.client.post('/api/invoice-workbench/input-groups/save-selection',json=body)
        self.assertEqual(200,response.status_code,response.json)
        self.conn.execute("UPDATE msmi_invoice_items SET product_code=NULL,mapping_status='unmapped',conversion_factor=NULL,stock_qty=0 WHERE id IN (?,?)",self.ids)
        apply_saved_mappings(self.conn,'input',self.fixture['promotion_invoice'])
        self.assertEqual(2,self.conn.execute('SELECT conversion_factor FROM msmi_invoice_items WHERE id=?',(self.ids[0],)).fetchone()[0])
        self.assertEqual(3,len(self.payload()['lines']))
        _update_line_snapshot(self.conn,'msmi_invoice_items',self.ids[0],'QA-OIL','mapped',1)
        self.assertEqual(1,self.conn.execute('SELECT COUNT(*) FROM invoice_input_group_choices').fetchone()[0])
        self.assertEqual(4,len(self.payload()['lines']))

    def test_select_first_concurrent_confirmation_and_posted_totals(self):
        body,preview=self.selection(2)
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'selected.sqlite3'
            with closing(sqlite3.connect(path)) as dest:self.conn.backup(dest)
            def save():
                with closing(sqlite3.connect(path,timeout=10)) as conn, conn:
                    conn.row_factory=sqlite3.Row
                    conn.execute('BEGIN IMMEDIATE')
                    return groups.save_selection(conn,'TDP',body,now_iso())
            with ThreadPoolExecutor(max_workers=2) as executor:results=list(executor.map(lambda _:save(),range(2)))
            self.assertEqual(1,len({r['id'] for r in results}))
            self.assertEqual(1,sum(bool(r['idempotent']) for r in results))
            with closing(sqlite3.connect(path)) as conn,conn:
                conn.row_factory=sqlite3.Row
                create_input_receipt(conn,self.fixture['promotion_invoice'],now_iso)
                create_input_receipt(conn,self.fixture['promotion_invoice'],now_iso)
                report=moving_average_report(conn,date_from='2026-08-01',date_to='2026-08-31')
                oil=next(r for r in report['items'] if r['product_code']=='QA-OIL')
                self.assertEqual((60,1288889),(oil['closing_qty'],oil['closing_value']))
                self.assertEqual(4,conn.execute('SELECT COUNT(*) FROM invoice_inventory_ledger').fetchone()[0])

    def test_split_selected_group_can_edit_conversion_and_reject_stale_source(self):
        from .invoice_mapping import save_conversion, validated_input_stock_snapshot, InvoiceMappingError
        body,_=self.selection(2)
        result=self.client.post('/api/invoice-workbench/input-groups/save-selection',json=body)
        self.client.post(f"/api/invoice-workbench/input-groups/{result.json['id']}/split")
        changed=save_conversion(self.conn,direction='input',item_id=self.ids[0],conversion_factor=3,now_iso=now_iso)
        self.assertEqual(1,changed['applied_lines'])
        snapshot=validated_input_stock_snapshot(self.conn,self.ids[0])
        self.assertEqual(72,snapshot['stock_qty'])
        other=validated_input_stock_snapshot(self.conn,self.ids[1])
        self.assertEqual(12,other['stock_qty'])
        self.conn.execute('UPDATE msmi_invoice_items SET amount=amount+1 WHERE id=?',(self.ids[0],))
        with self.assertRaises(InvoiceMappingError):validated_input_stock_snapshot(self.conn,self.ids[0])

    def test_retry_overlap_and_stale_source(self):
        saved=self.create()
        retry=self.client.post('/api/invoice-workbench/input-groups',json={'item_ids':self.ids,'token':saved['token']})
        self.assertTrue(retry.json['idempotent'])
        self.assertEqual(1,self.conn.execute('SELECT COUNT(*) FROM invoice_input_line_groups').fetchone()[0])
        with self.assertRaises(groups.GroupError):groups.preview_group(self.conn,'TDP',self.ids)
        self.conn.execute('UPDATE msmi_invoice_items SET amount=amount+1 WHERE id=?',(self.ids[0],))
        self.conn.commit()
        retry=self.client.post('/api/invoice-workbench/input-groups',json={'item_ids':self.ids,'token':saved['token']})
        self.assertEqual(409,retry.status_code)
        payload=self.payload()
        self.assertEqual(4,len(payload['lines']))
        self.assertEqual(1,len(payload['group_warnings']))

    def test_reject_changed_preview_and_audit_rollback(self):
        p=groups.preview_group(self.conn,'TDP',self.ids)
        self.conn.execute('UPDATE msmi_invoice_items SET stock_qty=stock_qty+1 WHERE id=?',(self.ids[0],))
        self.conn.commit()
        response=self.client.post('/api/invoice-workbench/input-groups',json={'item_ids':self.ids,'token':p['token']})
        self.assertEqual(409,response.status_code)
        p=groups.preview_group(self.conn,'TDP',self.ids)
        with patch.object(groups,'_audit',side_effect=RuntimeError('QA audit failed')):
            response=self.client.post('/api/invoice-workbench/input-groups',json={'item_ids':self.ids,'token':p['token']})
        self.assertEqual(500,response.status_code)
        self.assertEqual(0,self.conn.execute('SELECT COUNT(*) FROM invoice_input_line_groups').fetchone()[0])

    def test_wrong_product_tenant_tax_conversion_and_posted_are_blocked(self):
        with self.assertRaises(groups.GroupError):groups.preview_group(self.conn,'OTHER',self.ids)
        wrong=[self.ids[0],self.fixture['promotion_lines']['QA-CHILI-FREE']]
        with self.assertRaises(groups.GroupError):groups.preview_group(self.conn,'TDP',wrong)
        for column,value in [('mapping_status','unit_review'),('tax_rate',99),('stock_qty',0),('amount',-1)]:
            old=self.conn.execute(f'SELECT {column} FROM msmi_invoice_items WHERE id=?',(self.ids[1],)).fetchone()[0]
            self.conn.execute(f'UPDATE msmi_invoice_items SET {column}=? WHERE id=?',(value,self.ids[1]))
            with self.assertRaises(groups.GroupError):groups.preview_group(self.conn,'TDP',self.ids)
            self.conn.execute(f'UPDATE msmi_invoice_items SET {column}=? WHERE id=?',(old,self.ids[1]))
        create_input_receipt(self.conn,self.fixture['promotion_invoice'],now_iso)
        with self.assertRaises(groups.GroupError):groups.preview_group(self.conn,'TDP',self.ids)

    def test_conversion_uses_total_converted_quantity_and_text_safe_export(self):
        self.conn.execute("UPDATE msmi_invoice_items SET source_unit='Thùng',conversion_factor=30,stock_qty=qty*30,source_item_name='=BAD()' WHERE id IN (?,?)",self.ids)
        self.conn.commit()
        self.create()
        merged=next(r for r in self.payload()['lines'] if r.get('group_id'))
        self.assertEqual(900,merged['qty'])
        self.assertAlmostEqual(1288889/900,merged['unit_price'])
        wb=load_workbook(range_workbook(self.payload()))
        self.assertEqual('s',wb['Dong goc da gop']['D2'].data_type)

    def test_concurrent_confirmations_only_one_group(self):
        p=groups.preview_group(self.conn,'TDP',self.ids)
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'group.sqlite3'
            with closing(sqlite3.connect(path)) as target:self.conn.backup(target)
            def save():
                with closing(sqlite3.connect(path,timeout=10)) as conn, conn:
                    conn.row_factory=sqlite3.Row
                    conn.execute('BEGIN IMMEDIATE')
                    return groups.create_group(conn,'TDP',self.ids,p['token'],now_iso())
            with ThreadPoolExecutor(max_workers=2) as executor:results=list(executor.map(lambda _:save(),range(2)))
            self.assertEqual(1,len({r['id'] for r in results}))
            self.assertEqual(1,sum(bool(r['idempotent']) for r in results))

    def test_different_invoice_and_invalid_selection(self):
        old=dict(self.conn.execute('SELECT * FROM msmi_invoices WHERE id=?',(self.fixture['promotion_invoice'],)).fetchone())
        old.pop('id');old['remote_id']='QA-OTHER-INVOICE';old['invoice_number']='OTHER'
        new=self.conn.execute('INSERT INTO msmi_invoices('+','.join(old)+') VALUES('+','.join('?' for _ in old)+')',tuple(old.values())).lastrowid
        self.conn.execute('UPDATE msmi_invoice_items SET invoice_id=? WHERE id=?',(new,self.ids[1]))
        with self.assertRaisesRegex(groups.GroupError,'cùng một hóa đơn'):groups.preview_group(self.conn,'TDP',self.ids)
        for ids in [None,[],[self.ids[0]],[self.ids[0]]*2,['1','2'],[True,2]]:
            with self.assertRaises(groups.GroupError):groups.preview_group(self.conn,'TDP',ids)

    def test_identical_resync_keeps_group_but_different_mapping_restores_originals(self):
        self.create()
        row=dict(self.conn.execute('SELECT * FROM msmi_invoice_items WHERE id=?',(self.ids[0],)).fetchone())
        self.conn.execute('DELETE FROM msmi_invoice_items WHERE id=?',(row.pop('id'),))
        self.conn.execute('INSERT INTO msmi_invoice_items('+','.join(row)+') VALUES('+','.join('?' for _ in row)+')',tuple(row.values()))
        self.assertEqual(3,len(self.payload()['lines']))
        self.conn.execute("UPDATE msmi_invoice_items SET product_code='QA-CHILI' WHERE id=?",(self.ids[1],))
        self.assertEqual(4,len(self.payload()['lines']))
        self.assertEqual(1,len(self.payload()['group_warnings']))


if __name__=='__main__':unittest.main()
