"""Historical combined-approval fixture.

The old public API contract (approval posts BK) was retired on 24 September.
Its replacement is tested in test_order_approval_without_bk. Explicit BK
validation/posting remains covered by test_bk_import and test_bk_supplement.
"""
import unittest
import io
from openpyxl import load_workbook
from . import server, batch_bk_approval as approval
from .test_bk_import import BKImportTests, NOW


class BatchBKApprovalTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls): BKImportTests.setUpClass.__func__(cls)
    @classmethod
    def tearDownClass(cls): BKImportTests.tearDownClass.__func__(cls)

    def setUp(self):
        with server.db() as conn: conn.execute('DELETE FROM batch_bk_approvals')
        BKImportTests.setUp(self)
        with server.db() as conn:
            conn.execute('DELETE FROM people')
            conn.execute("INSERT INTO people(name,cccd,address) VALUES('Người bán BK','123456789','Địa chỉ thử nghiệm')")
            conn.execute("UPDATE products SET seller='Người bán BK',cccd='123456789'")
            self.batch=conn.execute("INSERT INTO batches(work_date,source_name,status,created_at) VALUES('2026-09-03','test.xlsx','draft',?)",(NOW,)).lastrowid
            self.order=self.add_line(conn)

    def add_line(self, conn, code='BK-P1', qty=2):
        name,unit=('Hàng BK','kg') if code=='BK-P1' else ('Hàng BK 2','chai')
        return conn.execute('''INSERT INTO orders(batch_id,work_date,contractor,kitchen,product_code,
            product_name,qty,actual_received,actual_delivered,unit,supplier,buy_price,sell_price,tax,
            purchase_list,source_row,errors,warnings,updated_at)
            VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)''',
            (self.batch,'2026-09-03','C1','K1',code,name,qty,qty,qty,unit,'BK-S1',111,200,'8%',1,27,'[]','[]',NOW)).lastrowid

    def post(self, **extra):
        # Reconstruct historical linked documents for read/export protections;
        # the public approval endpoint no longer creates these receipts.
        from . import batch_bk_approval as approval
        with server.db() as conn:
            preview = approval.prepare(conn, self.batch)
            approval.approve(conn, self.batch, {
                'source_hash': preview['sourceHash'], 'confirm_bk': True, **extra
            }, NOW, server.audit_event)
        return self.client.post(f'/api/batches/{self.batch}/approve', json={})

    def test_download_after_approval_cannot_double_import(self):
        self.assertEqual(self.post().status_code,200)
        payload=self.client.get(f'/api/bk-import/template?batch_id={self.batch}').data
        with server.db() as conn:
            p=approval.bk.parse_bk_preview(conn,payload)
            self.assertTrue(p['canConfirm']);self.assertTrue(p['alreadyPosted'])


    def test_selected_input_export_uses_posted_cost_without_writing_stock(self):
        self.assertEqual(self.post().status_code, 200)
        with server.db() as conn:
            before = conn.serialize()
            conn.execute("UPDATE settings SET value='0.8' WHERE key='purchase_rate'")
        with server.db() as conn:
            before = conn.serialize()
        for _ in range(2):
            response = self.client.post('/api/bk-import/export-approved', json={'batch_ids':[self.batch, self.batch]})
            self.assertEqual(response.status_code, 200, response.get_json(silent=True))
            wb = load_workbook(io.BytesIO(response.data), data_only=True)
            try:
                self.assertEqual(wb.active['A1'].value, 'BẢNG KÊ ĐẦU VÀO ĐÃ GHI NHẬP KHO')
                self.assertEqual(wb.active.max_row, 4)
                self.assertEqual([wb.active.cell(4,c).value for c in (8,9,10)], [2,190,380])
            finally:
                wb.close()
        with server.db() as conn:
            self.assertEqual(before, conn.serialize(), 'Download must not write stock or business data')


    def test_selected_input_export_blocks_unapproved_or_unposted(self):
        response = self.client.post('/api/bk-import/export-approved', json={'batch_ids':[self.batch]})
        self.assertEqual(response.status_code, 409)
        with server.db() as conn:
            conn.execute("UPDATE batches SET status='approved' WHERE id=?", (self.batch,))
        response = self.client.post('/api/bk-import/export-approved', json={'batch_ids':[self.batch]})
        self.assertEqual(response.status_code, 409)
        self.assertEqual(response.get_json()['code'], 'bk_not_posted')


    def test_selected_input_export_multiple_days_in_one_file(self):
        self.assertEqual(self.post().status_code, 200)
        first = self.batch
        with server.db() as conn:
            self.batch = conn.execute("INSERT INTO batches(work_date,source_name,status,created_at) VALUES('2026-09-02','second.xlsx','draft',?)", (NOW,)).lastrowid
            second_order = self.add_line(conn, qty=3)
            conn.execute("UPDATE orders SET work_date='2026-09-02' WHERE id=?", (second_order,))
        self.assertEqual(self.post().status_code, 200)
        response = self.client.post('/api/bk-import/export-approved', json={'batch_ids':[first, self.batch]})
        self.assertEqual(response.status_code, 200, response.get_json(silent=True))
        wb = load_workbook(io.BytesIO(response.data), data_only=True)
        try:
            self.assertEqual(wb.active.max_row, 5)
            self.assertEqual([wb.active.cell(r, 1).value.day for r in (4,5)], [2,3])
            self.assertEqual(sum(wb.active.cell(r, 10).value for r in (4,5)), 950)
        finally:
            wb.close()


    def test_selected_input_export_validates_selection(self):
        for ids in ([], [True], [1.5], ['1'], [-1], list(range(1,102))):
            self.assertEqual(self.client.post('/api/bk-import/export-approved', json={'batch_ids':ids}).status_code, 400)
        with server.db() as conn:
            conn.execute('UPDATE orders SET purchase_list=0 WHERE id=?', (self.order,))
        self.assertEqual(self.post().status_code, 200)
        response = self.client.post('/api/bk-import/export-approved', json={'batch_ids':[self.batch]})
        self.assertEqual(response.status_code, 409)
        self.assertEqual(response.get_json()['code'], 'no_posted_bk')


    def test_catalog_change_after_post_does_not_duplicate_or_reprice_receipt(self):
        self.assertEqual(self.post().status_code,200)
        with server.db() as conn:
            conn.execute("UPDATE settings SET value='0.90' WHERE key='purchase_rate'")
            conn.execute("UPDATE people SET address='Địa chỉ mới'")
        r=self.post();self.assertEqual(r.status_code,200,r.get_json())
        with server.db() as conn:
            self.assertEqual(tuple(conn.execute('SELECT count(*),sum(qty_delta),max(unit_cost) FROM invoice_inventory_ledger').fetchone()),(1,2,190))


    def test_purchase_workbook_import_cannot_replace_posted_bk_source(self):
        from .contract_modules import apply_purchase_order_preview, PurchaseOrderApplyError
        self.assertEqual(self.post().status_code,200)
        with server.db() as conn:
            with self.assertRaises(PurchaseOrderApplyError):
                apply_purchase_order_preview(conn,batch_id=self.batch,items=[],source_hash='other',source_name='new.xlsx',format_name='test',now_iso=server.now_iso)
