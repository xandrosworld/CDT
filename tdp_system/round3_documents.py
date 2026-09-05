"""Read-only report/debt views and non-destructive customer receipts."""
import hashlib
import json
import re
from datetime import date
from decimal import Decimal, InvalidOperation

from flask import jsonify, request
from openpyxl import Workbook

try:
    from document_totals import quantity_totals, quantity_text
    from report_export import ReportExportError
    from receivable_ledger import receivable_ledger_payload, ReceivableLedgerError
    from payable_export import _style_sheet
except ImportError:
    from .document_totals import quantity_totals, quantity_text
    from .report_export import ReportExportError
    from .receivable_ledger import receivable_ledger_payload, ReceivableLedgerError
    from .payable_export import _style_sheet


def customer_receipt(conn, body, ctx):
    """Called inside BEGIN IMMEDIATE; all amounts are integer VND."""
    clean = ctx["clean_text"]
    actor = clean(body.get("actor"))
    key = clean(body.get("request_id"))
    if not actor or len(actor) > 120:
        raise ValueError("Cần ghi rõ người thao tác (tối đa 120 ký tự)")
    if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9._:-]{7,127}", key):
        raise ValueError("Thiếu mã chống ghi nhận trùng; tải lại màn hình")
    try:
        amount = Decimal(str(body.get("amount")))
        if not amount.is_finite() or amount <= 0 or amount != amount.to_integral_value() or amount > 9_000_000_000_000:
            raise ValueError("Số tiền phải là số nguyên đồng dương trong giới hạn cho phép")
    except InvalidOperation:
        raise ValueError("Số tiền không hợp lệ") from None
    code = ctx["canonical_party_code"](conn, "contractor", clean(body.get("party_code")))
    payment_date = ctx["valid_iso_date"](body.get("payment_date"), "Ngày nhận tiền")
    note = clean(body.get("note"))
    values = dict(party_code=code, amount=int(amount), actor=actor, payment_date=payment_date, note=note)
    digest = hashlib.sha256(json.dumps(values, ensure_ascii=False, sort_keys=True).encode()).hexdigest()
    prior = conn.execute("SELECT * FROM payments WHERE request_key=?", (key,)).fetchone()
    if prior:
        if prior["request_hash"] != digest or prior["kind"] != "receipt":
            raise ValueError("Mã chống gửi trùng đã dùng cho nội dung khác; kiểm tra lịch sử trước khi tạo mới")
        return {"id": prior["id"], "idempotent": True, "status": prior["status"]}
    stamp = ctx["now_iso"]()
    cursor = conn.execute(
        """INSERT INTO payments(payment_date,kind,party_type,party_code,amount,note,
           created_at,updated_at,request_key,request_hash,created_by)
           VALUES(?,'receipt','contractor',?,?,?,?,?,?,?,?)""",
        (payment_date, code, int(amount), note, stamp, stamp, key, digest, actor),
    )
    ctx["audit_event"](conn, "payment.create", entity_type="contractor", entity_id=code,
                         metadata={**values, "payment_id": cursor.lastrowid})
    return {"id": cursor.lastrowid, "idempotent": False, "status": "posted"}


def register_round3_routes(app, ctx):
    db = ctx["db"]

    @app.get("/api/reports/monthly")
    @app.get("/api/reports/monthly/export")
    def monthly_report():
        period = request.args.get("period", "")
        try:
            if not re.fullmatch(r"\d{4}-\d{2}", period):
                raise ValueError("Chọn tháng báo cáo hợp lệ")
            date.fromisoformat(period + "-01")
            with db() as conn:
                conn.execute("BEGIN")
                # This monthly screen is approved-only. No arbitrary draft batch
                # is included merely because it happened to be selected elsewhere.
                batch = {"id": -1, "work_date": period + "-01", "status": "approved"}
                book = ctx["export_report"](conn, batch, [])
                draft_count = conn.execute(
                    "SELECT COUNT(*) FROM batches WHERE substr(work_date,1,7)=? AND status!='approved'", (period,)
                ).fetchone()[0]
            if request.path.endswith("/export"):
                try:
                    return ctx["send_xlsx"](book, f"Bao_cao_tong_hop_{period}.xlsx")
                finally:
                    book.close()
            try:
                ws = book.active
                headers = [ws.cell(2, col).value or ("Nhóm" if col == 1 else "") for col in range(1, 8)]
                rows = [[ws.cell(row, col).value for col in range(1, 8)] for row in range(3, ws.max_row + 1)]
                return jsonify(ok=True, period=period, headers=headers, rows=rows, draft_count=draft_count)
            finally:
                book.close()
        except (ValueError, ReportExportError) as error:
            return jsonify(ok=False, error=str(error)), 400

    @app.get("/api/debts/receivables/lines/export")
    def filtered_receivable_export():
        try:
            with db() as conn:
                conn.execute("BEGIN")
                kwargs = dict(date_from=request.args.get("from"), date_to=request.args.get("to"),
                              contractor=request.args.get("contractor", ""), kitchen=request.args.get("kitchen", ""),
                              statuses=request.args.get("status", "active"), limit=20000)
                payload = receivable_ledger_payload(conn, **kwargs)
                rows = payload["rows"]
                while len(rows) < payload["pagination"]["total"]:
                    page = receivable_ledger_payload(conn, **kwargs, offset=len(rows))
                    rows.extend(page["rows"])
                ws = (book := Workbook()).active
                ws.title = "Chi tiết phải thu"
                headers = ["Ngày", "Nhà thầu", "Bếp", "Hàng", "Số đặt", "Thực giao", "Khách trả",
                           "Giao ròng", "ĐVT", "Giá bán", "Trước thuế", "Thuế", "Phải thu", "Trạng thái", "Mã dòng"]
                ws.append(headers)
                for row in rows:
                    ws.append([date.fromisoformat(row["work_date"]), row["contractor"]["code"], row["kitchen"]["code"],
                               row["product_name"], row["ordered_qty"], row["actual_delivered"], row["customer_return_qty"],
                               row["delivered_qty"], row["unit"], row["sell_price"], row["subtotal"], row["tax_amount"],
                               row["amount"], "Hiệu lực" if row["status"] == "active" else "Đã đảo", row["id"]])
                summary = payload["summary"]
                ws.append(["TỔNG THEO BỘ LỌC", "", "", "", "", "", "",
                           quantity_text(summary["filtered_quantities_by_unit"]), "", "",
                           summary["filtered_subtotal"], summary["filtered_tax_amount"], summary["filtered_amount"]])
                _style_sheet(ws, title="CHI TIẾT PHẢI THU THEO BỘ LỌC",
                             subtitle=f"{payload['date_from']} – {payload['date_to']} · Dòng đã đảo chỉ để tra cứu, không tính số dư còn thu.",
                             headers=headers, widths=[14, 16, 16, 38, 16, 16, 16, 25, 10, 18, 20, 20, 20, 16, 12],
                             money_columns=(10, 11, 12, 13), quantity_columns=(5, 6, 7, 8), date_columns=(1,), total_row=ws.max_row + 2)
                for cells in ws.iter_rows(min_row=4):
                    for cell in cells[4:8]:
                        cell.number_format = "#,##0.######"
                ws.auto_filter.ref = f"A3:O{max(3, ws.max_row - 1)}"
            try:
                return ctx["send_xlsx"](book, f"Chi_tiet_phai_thu_{payload['date_from']}_{payload['date_to']}.xlsx")
            finally:
                book.close()
        except ReceivableLedgerError as error:
            return jsonify(ok=False, error=str(error)), error.status

    @app.get("/api/debts/receipts")
    def receipt_history():
        try:
            start = ctx["valid_iso_date"](request.args.get("from"), "Từ ngày")
            end = ctx["valid_iso_date"](request.args.get("to"), "Đến ngày")
            if start > end:
                raise ValueError("Từ ngày không được lớn hơn đến ngày")
            with db() as conn:
                rows = [dict(r) for r in conn.execute(
                    """SELECT id,payment_date,party_code,amount,note,status,revision,created_by,reversed_by,reversal_reason
                       FROM payments WHERE kind='receipt' AND party_type='contractor' AND payment_date>=? AND payment_date<=?
                       ORDER BY payment_date,id""", (start, end))]
            return jsonify(ok=True, rows=rows)
        except ValueError as error:
            return jsonify(ok=False, error=str(error)), 400

    @app.post("/api/debts/receipts/<int:payment_id>/reverse")
    def reverse_receipt(payment_id):
        body = request.get_json(silent=True) or {}
        actor = ctx["clean_text"](body.get("actor"))
        reason = ctx["clean_text"](body.get("reason"))
        if not actor or len(actor) > 120 or not reason or len(reason) > 500:
            return jsonify(ok=False, error="Cần tên người thao tác và lý do hoàn tác (tối đa 120/500 ký tự)"), 400
        with db() as conn:
            conn.execute("BEGIN IMMEDIATE")
            row = conn.execute("SELECT * FROM payments WHERE id=? AND kind='receipt' AND party_type='contractor'", (payment_id,)).fetchone()
            if not row:
                return jsonify(ok=False, error="Không tìm thấy khoản thu"), 404
            expected = body.get("expected_revision")
            if row["status"] == "reversed":
                if row["reversal_reason"] == reason and row["reversed_by"] == actor and expected in (row["revision"], row["revision"] - 1):
                    return jsonify(ok=True, idempotent=True)
                return jsonify(ok=False, error="Khoản thu đã hoàn tác; tải lại lịch sử"), 409
            if expected != row["revision"]:
                return jsonify(ok=False, error="Giao dịch đã thay đổi; tải lại lịch sử"), 409
            stamp = ctx["now_iso"]()
            conn.execute("""UPDATE payments SET status='reversed',revision=revision+1,updated_at=?,reversed_at=?,
                            reversal_reason=?,reversed_by=? WHERE id=?""", (stamp, stamp, reason, actor, payment_id))
            ctx["audit_event"](conn, "payment.reverse", entity_type="contractor", entity_id=row["party_code"],
                                 metadata={"payment_id": payment_id, "actor": actor, "reason": reason, "amount": row["amount"]})
        return jsonify(ok=True, idempotent=False)
