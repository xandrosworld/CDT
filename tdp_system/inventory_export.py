"""Stock valuation helpers and customer reports with source sales and line VAT.

Customer routes use M-Invoice revenue for Xuất and quantity-based NXT.
Internal valuation workbooks remain available to the stock reconciliation code.
"""

from __future__ import annotations

import hashlib
import io
import json
import re
import zipfile
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any, Mapping

from flask import jsonify, request, send_file
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.views import Selection

try:
    from invoice_line_groups import _fingerprint as group_fingerprint, _rows as group_source_rows
    from invoice_valuation import InvoiceValuationError, moving_average_report
    from document_totals import quantity_totals
    from template_workbook import (
        TemplateWorkbookError,
        assert_workbook_safe,
        clone_template_workbook,
        copy_row_layout,
        safe_workbook_bytes,
        write_literal,
    )
except ImportError:  # pragma: no cover - package invocation
    from .invoice_line_groups import _fingerprint as group_fingerprint, _rows as group_source_rows
    from .invoice_valuation import InvoiceValuationError, moving_average_report
    from .document_totals import quantity_totals
    from .template_workbook import (
        TemplateWorkbookError,
        assert_workbook_safe,
        clone_template_workbook,
        copy_row_layout,
        safe_workbook_bytes,
        write_literal,
    )


OPENING_TEMPLATE_SHEET = "Ton 7 (2)"
OPENING_TEMPLATE_SHA256 = "36DF2BA86D13307F96BB5944FCECB19A4A81C093B4AC6A98EA71330D68204DA6"
OFFICIAL_LAYOUT_ID = "TDP_TDK_NHAP_XUAT_NXT_V1"
EXPORT_KINDS = ("opening", "input", "output", "nxt")
MONEY_SCALE = Decimal("0.01")
QTY_SCALE = Decimal("0.000001")
EPSILON = Decimal("0.0000005")

GREEN = "70AD47"
PALE_GREEN = "E2F0D9"
PALE_YELLOW = "FFF2CC"
WHITE = "FFFFFF"
GRID = Side(style="thin", color="808080")


def _quantity_total(model, field):
    groups = [r for r in quantity_totals(model['items'], field) if abs(r['quantity']) > float(EPSILON)]
    if not groups:
        return 0
    return groups[0]['quantity'] if len(groups) == 1 else f"{len(groups)} ĐVT · xem Tổng ĐVT"


def _add_unit_totals(workbook, model):
    fields = ('opening_qty', 'input_qty', 'output_qty', 'closing_qty')
    grouped = {field: quantity_totals(model['items'], field) for field in fields}
    units = sorted({r['unit'] for rows in grouped.values() for r in rows if abs(r['quantity']) > float(EPSILON)})
    if len(units) < 2:
        return
    sheet = workbook.create_sheet('Tổng ĐVT')
    sheet.append(['ĐVT', 'Tồn đầu', 'Nhập', 'Xuất', 'Tồn cuối'])
    values = {field: {r['unit']: r['quantity'] for r in rows} for field, rows in grouped.items()}
    for unit in units:
        sheet.append([unit, *[values[field].get(unit, 0) for field in fields]])
        _put(sheet, f'A{sheet.max_row}', unit)
    for row in sheet:
        for cell in row:
            cell.font = Font(name='Times New Roman', size=11, bold=cell.row == 1)
            cell.border = Border(bottom=GRID)
            cell.alignment = Alignment(horizontal='left' if cell.column == 1 else 'right')
            if cell.column > 1 and cell.row > 1: cell.number_format = '#,##0.######'
    for col in 'ABCDE': sheet.column_dimensions[col].width = 22
    sheet.freeze_panes = 'B2'
    _configure_print(sheet, last_row=sheet.max_row, last_column='E', title_rows='$1:$1')


class InventoryExportError(ValueError):
    def __init__(
        self, message: str, *, code: str = "invalid_inventory_export", status: int = 422,
    ):
        super().__init__(message)
        self.code = code
        self.status = status


def _plain(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip())


def _decimal(value: Any, label: str) -> Decimal:
    if isinstance(value, bool):
        raise InventoryExportError(f"{label} phải là số hữu hạn")
    try:
        number = Decimal(str(value or 0))
    except (InvalidOperation, TypeError, ValueError) as error:
        raise InventoryExportError(f"{label} phải là số hữu hạn") from error
    if not number.is_finite():
        raise InventoryExportError(f"{label} phải là số hữu hạn")
    return number


def _excel_number(value: Any, scale: Decimal | None = None) -> int | float:
    number = _decimal(value, "Giá trị xuất Excel")
    if scale is not None:
        number = number.quantize(scale, rounding=ROUND_HALF_UP)
    if number == number.to_integral_value():
        return int(number)
    return float(number)


def _average(value: Any, quantity: Any) -> int | float:
    qty = _decimal(quantity, "Số lượng")
    if abs(qty) <= EPSILON:
        return 0
    return _excel_number(_decimal(value, "Giá trị") / qty, QTY_SCALE)


def _table_exists(conn: Any, name: str) -> bool:
    return conn.execute(
        "SELECT 1 FROM sqlite_master WHERE type='table' AND name=?", (name,),
    ).fetchone() is not None


def _column_exists(conn: Any, table: str, column: str) -> bool:
    return any(str(row["name"]) == column for row in conn.execute(f"PRAGMA table_info({table})"))


def _product_metadata(conn: Any) -> dict[str, dict[str, Any]]:
    if not _table_exists(conn, "products"):
        return {}
    outgoing_join = (
        "LEFT JOIN outgoing_product_names n ON n.product_code=p.code"
        if _table_exists(conn, "outgoing_product_names")
        else ""
    )
    invoice_expression = "COALESCE(NULLIF(n.invoice_name,''),p.name)" if outgoing_join else "p.name"
    tax_expression = "COALESCE(p.tax,'')" if _column_exists(conn, "products", "tax") else "''"
    rows = conn.execute(
        f"""SELECT p.code,p.name,COALESCE(p.unit,'') unit,{tax_expression} tax,
                   {invoice_expression} invoice_name
              FROM products p {outgoing_join} ORDER BY p.code"""
    ).fetchall()
    return {str(row["code"]): dict(row) for row in rows}


def _source_trace(conn: Any, event: Mapping[str, Any], tax_cache=None) -> dict[str, Any]:
    source_id = int(event["source_invoice_id"])
    line_id = int(event["source_line_id"])
    if event["source_invoice_table"] == "msmi_invoices":
        row = conn.execute(
            """SELECT i.invoice_date,i.invoice_series,i.invoice_number,
                      COALESCE(i.seller_tax_code,'') party_tax_code,
                      COALESCE(i.seller_name,'') party_name,
                      COALESCE(li.source_item_code,'') source_item_code,
                      COALESCE(li.source_item_name,'') source_item_name,
                      COALESCE(li.source_unit,'') source_unit
                 FROM msmi_invoice_items li
                 JOIN msmi_invoices i ON i.id=li.invoice_id
                WHERE i.id=? AND li.id=?""",
            (source_id, line_id),
        ).fetchone()
    elif event["source_invoice_table"] == "outgoing_source_invoices":
        row = conn.execute(
            """SELECT i.invoice_date,i.invoice_series,i.invoice_number,
                      COALESCE(i.buyer_tax_code,'') party_tax_code,
                      COALESCE(i.buyer_name,'') party_name,
                      COALESCE(li.source_item_code,'') source_item_code,
                      COALESCE(li.source_item_name,'') source_item_name,
                      COALESCE(li.source_unit,'') source_unit
                 FROM outgoing_source_invoice_items li
                 JOIN outgoing_source_invoices i ON i.id=li.invoice_id
                WHERE i.id=? AND li.id=?""",
            (source_id, line_id),
        ).fetchone()
    elif event["source_invoice_table"] == "bk_import_documents":
        row = conn.execute(
            """SELECT li.document_date invoice_date,'BK' invoice_series,
                      li.source_reference invoice_number,'' party_tax_code,
                      li.source_party party_name,li.product_code source_item_code,
                      li.product_name_snapshot source_item_name,
                      li.unit_snapshot source_unit
                 FROM bk_import_lines li
                 JOIN bk_import_documents d ON d.id=li.document_id
                WHERE d.id=? AND li.id=?""",
            (source_id, line_id),
        ).fetchone()
    else:
        raise InventoryExportError(
            "Sổ kho tham chiếu loại nguồn hóa đơn không được hỗ trợ",
            code="unsupported_source_trace",
        )
    if not row:
        raise InventoryExportError(
            "Sổ kho thiếu dòng hóa đơn nguồn để lập file đối chiếu",
            code="source_trace_missing",
        )
    trace = dict(row)
    if event['source_invoice_table'] in {'msmi_invoices', 'outgoing_source_invoices'}:
        table = event['source_invoice_table']
        lines_table = 'msmi_invoice_items' if table == 'msmi_invoices' else 'outgoing_source_invoice_items'
        if tax_cache is None:
            tax_cache = {}
        cache_key = (table, source_id)
        if cache_key not in tax_cache:
            header = dict(conn.execute(f'SELECT raw_json,tax_amount FROM {table} WHERE id=?', (source_id,)).fetchone())
            header['items'] = [dict(r) for r in conn.execute(f'SELECT line_index,amount,tax_rate,id FROM {lines_table} WHERE invoice_id=?', (source_id,))]
            try:
                from .invoice_line_tax import annotate_invoice_tax
            except ImportError:
                from invoice_line_tax import annotate_invoice_tax
            annotate_invoice_tax(header, header['raw_json'])
            tax_cache[cache_key] = ({r['id']: r for r in header['items']}, header['detail_tax_difference'])
        source_lines, difference = tax_cache[cache_key]
        line = source_lines[line_id]
        trace.update({key: line[key] for key in ('tax_rate', 'line_tax_amount', 'tax_note')})
        if difference is not None and abs(difference) > 1:
            trace['tax_note'] += f"; tổng thuế HĐ trừ chi tiết: {difference:,.0f}"
    else:
        trace.update(tax_rate='', line_tax_amount=None, tax_note='BK chưa có tiền thuế nguồn')
    return trace


def _verify_item_balance(item: Mapping[str, Any]) -> None:
    quantity_balance = (
        _decimal(item["opening_qty"], "Tồn đầu")
        + _decimal(item["input_qty"], "Nhập")
        - _decimal(item["output_qty"], "Xuất")
        - _decimal(item["closing_qty"], "Tồn cuối")
    )
    value_balance = (
        _decimal(item["opening_value"], "Giá trị tồn đầu")
        + _decimal(item["input_value"], "Giá trị nhập")
        - _decimal(item["output_value"], "Giá trị xuất")
        - _decimal(item["closing_value"], "Giá trị tồn cuối")
    )
    if abs(quantity_balance) > QTY_SCALE:
        raise InventoryExportError(
            f"Mã {item['product_code']} không khớp TĐK + Nhập - Xuất = Tồn cuối",
            code="quantity_reconciliation_failed",
        )
    if abs(value_balance) > MONEY_SCALE:
        raise InventoryExportError(
            f"Mã {item['product_code']} không khớp giá trị kho",
            code="value_reconciliation_failed",
        )


def _group_input_movements(conn: Any, rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Project only explicitly chosen, still-valid groups; never rewrite ledger events."""
    if not _table_exists(conn, "invoice_input_line_groups"):
        return rows
    groups = conn.execute("""SELECT g.* FROM invoice_input_line_groups g
        JOIN msmi_invoices i ON i.id=g.invoice_id AND i.tenant=g.tenant
        WHERE g.active=1 ORDER BY g.id""").fetchall()
    result = list(rows)
    for group in groups:
        indices = json.loads(group["member_indices"])
        positions = [i for i, row in enumerate(result)
                     if row["source_invoice_table"] == "msmi_invoices"
                     and row["source_invoice_id"] == group["invoice_id"]
                     and row["source_line_index"] in indices and row["movement_label"] == "Nhập"
                     and not row.get("group_id")]
        if len(positions) != len(indices):
            continue
        members = [result[i] for i in positions]
        source = group_source_rows(conn, [row["source_line_id"] for row in members])
        if len(source) != len(indices) or group_fingerprint(source) != group["fingerprint"]:
            continue
        if len({(r["product_code"], r["unit"], r["txn_date"], r["confirmation_id"]) for r in members}) != 1:
            continue
        # Historical posted quantities/codes must agree with the choice being displayed.
        by_id = {r["id"]: r for r in source}
        if any(r["product_code"] != by_id[r["source_line_id"]]["product_code"]
               or r["unit"] != by_id[r["source_line_id"]]["product_unit"]
               or abs(_decimal(r["quantity"], "Lượng nhập") - _decimal(by_id[r["source_line_id"]]["stock_qty"], "Lượng quy đổi")) > QTY_SCALE
               for r in members):
            continue
        quantity = sum((_decimal(r["quantity"], "Lượng nhập") for r in members), Decimal(0))
        amount = sum((_decimal(r["amount"], "Tiền nhập") for r in members), Decimal(0))
        if quantity <= 0:
            continue
        merged = dict(members[0], group_id=int(group["id"]), group_members=members,
                      invoice_name=members[0]["product_name"], quantity=_excel_number(quantity, QTY_SCALE),
                      amount=_excel_number(amount, MONEY_SCALE), unit_cost=_average(amount, quantity))
        selected = set(positions)
        result = [merged if i == positions[0] else row for i, row in enumerate(result)
                  if i == positions[0] or i not in selected]
    return result


def collect_inventory_export_model(
    conn: Any, *, date_from: Any, date_to: Any,
) -> dict[str, Any]:
    """Return one immutable model used by all four workbooks."""

    try:
        report = moving_average_report(
            conn,
            date_from=date_from,
            date_to=date_to,
            include_zero=False,
            include_events=True,
        )
    except InvoiceValuationError as error:
        raise InventoryExportError(str(error), code=error.code, status=error.status) from None

    metadata = _product_metadata(conn)
    items: list[dict[str, Any]] = []
    for source in report["items"]:
        item = dict(source)
        product = metadata.get(str(item["product_code"]), {})
        warehouse_codes = list(dict.fromkeys(
            _plain(value) for value in item.get("warehouse_codes", []) if _plain(value)
        ))
        item.update({
            "product_name": _plain(product.get("name") or item.get("product_name")),
            "invoice_name": _plain(product.get("invoice_name") or item.get("product_name")),
            "unit": _plain(product.get("unit") or item.get("unit")),
            "tax": _plain(product.get("tax")),
            # MÃ KHO is a distinct identity in the customer's opening workbook.
            # Keep every source value when several warehouse rows were grouped;
            # old snapshots safely fall back to MÃ TĐP.
            "warehouse_code": " / ".join(warehouse_codes) or str(item["product_code"]),
            "opening_unit_cost": _average(item["opening_value"], item["opening_qty"]),
            "input_unit_cost": _average(item["input_value"], item["input_qty"]),
            "output_unit_cost": _average(item["output_value"], item["output_qty"]),
            "closing_unit_cost": _average(item["closing_value"], item["closing_qty"]),
        })
        _verify_item_balance(item)
        items.append(item)

    input_rows: list[dict[str, Any]] = []
    output_rows: list[dict[str, Any]] = []
    input_by_product: dict[str, list[Decimal]] = {}
    output_by_product: dict[str, list[Decimal]] = {}
    event_identity: list[int] = []
    tax_traces = {}
    tax_cache = {}
    for event in report["events"]:
        trace = _source_trace(conn, event, tax_cache)
        tax_traces[int(event["ledger_event_id"])] = trace
        code = str(event["product_code"])
        product = metadata.get(code)
        if not product:
            raise InventoryExportError(
                f"Sổ kho tham chiếu mã {code} không còn trong danh mục",
                code="product_trace_missing",
            )
        movement = str(event["movement"])
        qty_delta = _decimal(event["qty_delta"], "Số lượng phát sinh")
        amount = _decimal(event["valuation_amount"], "Giá trị phát sinh")
        if movement == "input":
            quantity = qty_delta
            signed_amount = amount
            movement_label = "Nhập"
            target = input_rows
            accumulator = input_by_product
        elif movement == "input_reversal":
            quantity = qty_delta
            signed_amount = amount
            movement_label = "Hoàn tác nhập BK"
            target = input_rows
            accumulator = input_by_product
        elif movement == "output":
            quantity = -qty_delta
            signed_amount = amount
            movement_label = "Xuất"
            target = output_rows
            accumulator = output_by_product
        elif movement == "reversal":
            quantity = -qty_delta
            signed_amount = -amount
            movement_label = "Hoàn tác xuất"
            target = output_rows
            accumulator = output_by_product
        else:
            raise InventoryExportError(
                "Projection có loại phát sinh không được hỗ trợ",
                code="unsupported_movement",
            )
        values = accumulator.setdefault(code, [Decimal(0), Decimal(0)])
        values[0] += quantity
        values[1] += signed_amount
        target.append({
            "ledger_event_id": int(event["ledger_event_id"]),
            "txn_date": str(event["txn_date"]),
            "invoice_series": _plain(trace.get("invoice_series")),
            "invoice_number": _plain(trace.get("invoice_number")),
            "party_tax_code": _plain(trace.get("party_tax_code")),
            "party_name": _plain(trace.get("party_name")),
            "product_code": code,
            "product_name": _plain(product.get("name")),
            "invoice_name": _plain(trace.get("source_item_name") or product.get("invoice_name")),
            "source_item_code": _plain(trace.get("source_item_code")),
            "source_unit": _plain(trace.get("source_unit")),
            "unit": _plain(product.get("unit")),
            "quantity": _excel_number(quantity, QTY_SCALE),
            "unit_cost": _excel_number(event["valuation_unit_cost"], QTY_SCALE),
            "amount": _excel_number(signed_amount, MONEY_SCALE),
            "movement_label": movement_label,
            "source_invoice_id": int(event["source_invoice_id"]),
            "source_line_id": int(event["source_line_id"]),
            "source_line_index": int(event["source_line_index"]),
            "source_invoice_table": str(event["source_invoice_table"]),
            "confirmation_id": int(event["confirmation_id"]),
        })
        event_identity.append(int(event["ledger_event_id"]))

    for row in input_rows + output_rows:
        trace = tax_traces[row['ledger_event_id']]
        row['tax_rate'] = trace.get('tax_rate', '')
        row['tax_note'] = trace.get('tax_note', 'Chưa có tiền thuế nguồn')
        tax = trace.get('line_tax_amount')
        # Source tax belongs to the invoice line, never to its moving-average cost.
        if row['movement_label'].startswith('Hoàn tác'):
            tax = None
            row['tax_note'] = 'Hoàn tác kho; đối chiếu thuế trên hóa đơn nguồn'
        row['line_tax_amount'] = tax
        row['amount_with_tax'] = float(_decimal(row['amount'], 'Tiền') + _decimal(tax, 'Thuế')) if tax is not None else None
    input_rows = _group_input_movements(conn, input_rows)
    try:
        from .invoice_line_tax import grouped_tax
    except ImportError:
        from invoice_line_tax import grouped_tax
    for row in input_rows:
        if row.get('group_members'):
            row.update(grouped_tax(row['group_members']))

    item_map = {str(item["product_code"]): item for item in items}
    for code, item in item_map.items():
        actual_input = input_by_product.get(code, [Decimal(0), Decimal(0)])
        actual_output = output_by_product.get(code, [Decimal(0), Decimal(0)])
        checks = (
            (actual_input[0], _decimal(item["input_qty"], "SL nhập"), QTY_SCALE),
            (actual_input[1], _decimal(item["input_value"], "GT nhập"), MONEY_SCALE),
            (actual_output[0], _decimal(item["output_qty"], "SL xuất"), QTY_SCALE),
            (actual_output[1], _decimal(item["output_value"], "GT xuất"), MONEY_SCALE),
        )
        if any(abs(left - right) > tolerance for left, right, tolerance in checks):
            raise InventoryExportError(
                f"Chi tiết phát sinh mã {code} không khớp bảng NXT",
                code="movement_reconciliation_failed",
            )

    total_fields = (
        "opening_qty", "opening_value", "input_qty", "input_value",
        "output_qty", "output_value", "closing_qty", "closing_value",
    )
    totals = {
        field: _excel_number(sum((_decimal(item[field], field) for item in items), Decimal(0)),
                             MONEY_SCALE if field.endswith("value") else QTY_SCALE)
        for field in total_fields
    }
    signature = {
        "layout_id": OFFICIAL_LAYOUT_ID,
        "date_from": report["date_from"],
        "date_to": report["date_to"],
        "opening_period": report["opening_period"],
        "items": [
            {key: item[key] for key in ("product_code", "warehouse_code") + total_fields}
            for item in items
        ],
        "ledger_event_ids": event_identity,
        "input_groups": [{"id": r["group_id"], "events": [m["ledger_event_id"] for m in r["group_members"]]}
                         for r in input_rows if r.get("group_id")],
    }
    contract_id = hashlib.sha256(
        json.dumps(signature, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    ).hexdigest().upper()
    return {
        "date_from": report["date_from"],
        "date_to": report["date_to"],
        "opening_period": report["opening_period"],
        "opening_date": report["opening_date"],
        "rounding": dict(report["rounding"]),
        "same_day_order": list(report["same_day_order"]),
        "items": items,
        "input_rows": input_rows,
        "output_rows": output_rows,
        "totals": totals,
        "contract_id": contract_id,
        "read_only": True,
        "layout_id": OFFICIAL_LAYOUT_ID,
        "visual_status": "official_customer_tdk_derived_layout",
    }


def monthly_customer_model(conn, model):
    """Apply one monthly valuation to NXT and the corresponding opening export."""
    if model.get('valuation_method') == 'monthly_weighted_average':
        return model
    try:
        from .invoice_monthly_valuation import monthly_average_report
    except ImportError:
        from invoice_monthly_valuation import monthly_average_report
    try:
        report = monthly_average_report(conn, date_from=model['date_from'], date_to=model['date_to'])
    except InvoiceValuationError as error:
        raise InventoryExportError(str(error), code=error.code, status=error.status) from None
    metadata = _product_metadata(conn)
    old = {r['product_code']: r for r in model['items']}
    items = []
    for row in report['items']:
        code = row['product_code']
        item = dict(old.get(code, {}), **row)
        product = metadata.get(code, {})
        item.update(product_name=product.get('name', row['product_name']),
                    invoice_name=product.get('invoice_name', row['product_name']),
                    warehouse_code=' / '.join(row.get('warehouse_codes') or []) or code,
                    tax=product.get('tax', ''),
                    opening_unit_cost=_average(row['opening_value'], row['opening_qty']),
                    input_unit_cost=_average(row['input_value'], row['input_qty']),
                    output_unit_cost=row['average_unit_cost'], closing_unit_cost=row['average_unit_cost'])
        items.append(item)
    totals = {key: _excel_number(sum((_decimal(r[key], key) for r in items), Decimal(0)),
                   QTY_SCALE if key.endswith('_qty') else MONEY_SCALE) for key in model['totals']}
    signature = json.dumps([model['contract_id'], report['valuation_method'], items], ensure_ascii=False, sort_keys=True)
    return dict(model, items=items, totals=totals, rounding=report['rounding'],
                valuation_method=report['valuation_method'],
                contract_id=hashlib.sha256(signature.encode()).hexdigest().upper())


def _excel_literal(value: Any) -> Any:
    if isinstance(value, str) and value.startswith("="):
        return "'" + value
    return value


def _put(sheet: Any, coordinate: str, value: Any) -> None:
    write_literal(sheet, coordinate, _excel_literal(value))


def _period_text(model: Mapping[str, Any]) -> str:
    start = datetime.strptime(str(model["date_from"]), "%Y-%m-%d").strftime("%d/%m/%Y")
    end = datetime.strptime(str(model["date_to"]), "%Y-%m-%d").strftime("%d/%m/%Y")
    return f"Từ ngày {start} đến ngày {end}"


def _set_core_properties(workbook: Any, model: Mapping[str, Any], title: str) -> None:
    workbook.properties.title = title
    workbook.properties.subject = _period_text(model)
    workbook.properties.identifier = str(model["contract_id"])
    workbook.properties.creator = "Thành Đạt Phát"
    workbook.properties.lastModifiedBy = "Thành Đạt Phát"


def _add_control_sheet(workbook: Any, model: Mapping[str, Any], kind: str) -> None:
    sheet = workbook.create_sheet("_ĐỐI_CHIẾU")
    rows = [
        ("CONTRACT_ID", model["contract_id"]),
        ("LOẠI_FILE", kind),
        ("MẪU_BIỂU", model["layout_id"]),
        ("TRẠNG_THÁI_BIỂU_MẪU", "CHÍNH THỨC"),
        ("TỪ_NGÀY", model["date_from"]),
        ("ĐẾN_NGÀY", model["date_to"]),
        ("KỲ_TỒN_ĐẦU", model["opening_period"]),
        ("SL_TỒN_ĐẦU", model["totals"]["opening_qty"]),
        ("GT_TỒN_ĐẦU", model["totals"]["opening_value"]),
        ("SL_NHẬP", model["totals"]["input_qty"]),
        ("GT_NHẬP", model["totals"]["input_value"]),
        ("SL_XUẤT", model["totals"]["output_qty"]),
        ("GT_XUẤT", model["totals"]["output_value"]),
        ("SL_TỒN_CUỐI", model["totals"]["closing_qty"]),
        ("GT_TỒN_CUỐI", model["totals"]["closing_value"]),
        ("CÔNG_THỨC", "TĐK + Nhập - Xuất = Tồn cuối"),
        ("PHƯƠNG_PHÁP_GIÁ", "Bình quân gia quyền cả tháng" if model.get("valuation_method") == "monthly_weighted_average" else "Bình quân gia quyền di động"),
        ("NGUỒN_MẪU_TĐK_SHA256", OPENING_TEMPLATE_SHA256),
        ("KIỂU_DỮ_LIỆU", "Giá trị tĩnh; không công thức; không liên kết ngoài"),
        ("ROUNDING", model["rounding"]["method"]),
    ]
    for row_index, (label, value) in enumerate(rows, start=1):
        _put(sheet, f"A{row_index}", label)
        _put(sheet, f"B{row_index}", value)
    sheet.sheet_state = "veryHidden"


def _configure_print(sheet: Any, *, last_row: int, last_column: str, title_rows: str) -> None:
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.orientation = sheet.ORIENTATION_LANDSCAPE
    # The customer's opening template carries a legacy fixed Scale value.
    # Excel treats Scale together with FitToPagesWide/Tall as a corrupt page
    # setup and can refuse to open the generated TĐK workbook.  Fit-to-page is
    # the approved behavior for all four inventory exports, so explicitly
    # clear the mutually exclusive fixed scale.
    sheet.page_setup.scale = None
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_margins = PageMargins(
        left=0.2, right=0.2, top=0.35, bottom=0.35, header=0.15, footer=0.15,
    )
    sheet.print_area = f"$A$1:${last_column}${max(last_row, 1)}"
    sheet.print_title_rows = title_rows
    sheet.sheet_view.showGridLines = False
    sheet.print_options.horizontalCentered = True
    sheet.oddFooter.center.text = "Trang &P / &N"
    sheet.oddFooter.center.size = 9
    sheet.oddFooter.center.font = "Times New Roman"


def build_opening_workbook(
    model: Mapping[str, Any], *, template_path: str | Path,
    expected_sha256: str = OPENING_TEMPLATE_SHA256,
) -> Any:
    cloned = clone_template_workbook(
        template_path,
        sheet_names=[OPENING_TEMPLATE_SHEET],
        expected_sha256=expected_sha256,
    )
    workbook = cloned.workbook
    sheet = workbook[OPENING_TEMPLATE_SHEET]
    try:
        sheet.title = "TĐK"
        for row in range(1, 5):
            for column in range(1, 10):
                sheet.cell(row, column).value = None
        for row in range(7, max(sheet.max_row, 7) + 1):
            for column in range(1, 10):
                sheet.cell(row, column).value = None
        for merged in ("A1:I1", "A2:I2"):
            sheet.merge_cells(merged)

        _put(sheet, "A1", "TỒN ĐẦU KỲ")
        _put(sheet, "A2", _period_text(model))
        sheet["A1"].font = Font(name="Times New Roman", size=16, bold=True)
        sheet["A1"].alignment = Alignment(horizontal="center")
        sheet["A2"].font = Font(name="Times New Roman", size=11, italic=True)
        sheet["A2"].alignment = Alignment(horizontal="center")
        _put(sheet, "F4", "TỔNG")
        _put(sheet, "G4", _quantity_total(model, "opening_qty"))
        _put(sheet, "I4", model["totals"]["opening_value"])
        sheet['G4'].alignment = Alignment(wrap_text=True, vertical='center')
        sheet.row_dimensions[4].height = 42
        sheet['I4'].number_format = '#,##0'
        for coordinate in ("F4", "G4", "I4"):
            sheet[coordinate].font = Font(name="Times New Roman", size=12, bold=True)
        _put(sheet, "A5", "MÃ TĐP")
        _put(sheet, "B5", "TÊN TĐP")
        _put(sheet, "C5", "Tên trên HĐ")
        _put(sheet, "D5", "MÃ KHO")
        _put(sheet, "E5", "T/Suất")
        _put(sheet, "F5", "ĐVT")
        _put(sheet, "G5", "TỒN ĐẦU KỲ")
        _put(sheet, "G6", "Số lượng")
        _put(sheet, "H6", "Đơn giá")
        _put(sheet, "I6", "Thành tiền")
        for row in range(5, 7):
            for column in range(1, 10):
                sheet.cell(row, column).alignment = Alignment(
                    horizontal="center", vertical="center", wrap_text=True,
                    shrink_to_fit=True,
                )
        sheet.row_dimensions[5].height = max(sheet.row_dimensions[5].height or 0, 28)
        sheet.row_dimensions[6].height = max(sheet.row_dimensions[6].height or 0, 24)

        rows = [
            item for item in model["items"]
            if abs(_decimal(item["opening_qty"], "SL tồn đầu")) > EPSILON
            or abs(_decimal(item["opening_value"], "GT tồn đầu")) > MONEY_SCALE
        ]
        for offset, item in enumerate(rows):
            row = 7 + offset
            if row != 7:
                copy_row_layout(sheet, 7, row, include_values=False)
            values = (
                item["product_code"], item["product_name"], item["invoice_name"],
                item["warehouse_code"], item["tax"], item["unit"], item["opening_qty"],
                item["opening_unit_cost"], item["opening_value"],
            )
            for column, value in enumerate(values, start=1):
                _put(sheet, f"{chr(64 + column)}{row}", value)
            sheet[f"G{row}"].number_format = "#,##0.######"
            sheet[f"H{row}"].number_format = "#,##0"
            sheet[f"I{row}"].number_format = "#,##0"
        last_row = max(6, 6 + len(rows))
        # The customer TĐK sheet was frozen at I43 and carries four pane
        # selections.  Assigning A7 directly leaves stale/duplicate selections
        # in the OOXML; desktop Excel then refuses to open the generated file.
        # Reset the inherited pane state before installing the new split.
        sheet.sheet_view.pane = None
        sheet.sheet_view.selection = [Selection(activeCell="A1", sqref="A1")]
        sheet.freeze_panes = "A7"
        _configure_print(sheet, last_row=last_row, last_column="I", title_rows="$5:$6")
        _set_core_properties(workbook, model, "Tồn đầu kỳ")
        _add_unit_totals(workbook, model)
        _add_control_sheet(workbook, model, "opening")
        workbook.active = 0
        assert_workbook_safe(workbook)
        return workbook
    except Exception:
        workbook.close()
        raise


def _new_document(model: Mapping[str, Any], title: str, subtitle: str) -> tuple[Any, Any]:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title[:31]
    sheet.merge_cells("A1:N1")
    sheet.merge_cells("A2:N2")
    _put(sheet, "A1", title)
    _put(sheet, "A2", f"{_period_text(model)} · {subtitle}")
    sheet["A1"].font = Font(name="Times New Roman", size=16, bold=True)
    sheet["A1"].alignment = Alignment(horizontal="center")
    sheet["A2"].font = Font(name="Times New Roman", size=10, italic=True)
    sheet["A2"].alignment = Alignment(horizontal="center")
    _set_core_properties(workbook, model, title)
    return workbook, sheet


def build_movement_workbook(model: Mapping[str, Any], *, direction: str) -> Any:
    if direction not in {"input", "output"}:
        raise InventoryExportError("Chiều file nhập/xuất không hợp lệ", status=400)
    is_input = direction == "input"
    title = "NHẬP TRONG KỲ" if is_input else "XUẤT KHO · GIÁ VỐN"
    subtitle = (
        "Gồm hóa đơn đầu vào và BK đã xác nhận; hoàn tác nhập thể hiện âm"
        if is_input else "Giá vốn hàng xuất; hoàn tác thể hiện âm. Tiền bán xem báo cáo Xuất · giá bán hóa đơn."
    )
    if not is_input and any(r.get('valuation_status') != 'ok' for r in model['items']):
        subtitle += ' Giá vốn còn cần đối chiếu.'
    workbook, sheet = _new_document(model, title, subtitle)
    rows = model["input_rows"] if is_input else model["output_rows"]
    headers = (
        "STT", "Ngày HĐ", "Ký hiệu", "Số HĐ", "MST", "Đối tác", "Mã TĐP",
        "Tên TĐP", "Tên trên HĐ", "ĐVT kho", "Số lượng", "Đơn giá vốn",
        "Thành tiền" if is_input else "Thành tiền giá vốn", "Loại phát sinh",
    )
    if is_input:
        headers = (*headers[:11], 'Đơn giá nhập', 'Tiền trước thuế', headers[13], 'Thuế suất', 'Tiền thuế', 'Tiền gồm thuế', 'Đối chiếu thuế')
        sheet.unmerge_cells('A1:N1'); sheet.unmerge_cells('A2:N2')
        sheet.merge_cells('A1:R1'); sheet.merge_cells('A2:R2')
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(4, column, header)
        cell.font = Font(name="Times New Roman", size=10, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=GREEN)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=GRID, right=GRID, top=GRID, bottom=GRID)
    for offset, item in enumerate(rows, start=1):
        row = 4 + offset
        values = (
            offset,
            datetime.strptime(item["txn_date"], "%Y-%m-%d").date(),
            item["invoice_series"], item["invoice_number"], item["party_tax_code"],
            item["party_name"], item["product_code"], item["product_name"],
            item["invoice_name"], item["unit"], item["quantity"], item["unit_cost"],
            item["amount"], item["movement_label"],
        )
        if is_input:
            values += (item.get('tax_rate'), item.get('line_tax_amount'), item.get('amount_with_tax'), item.get('tax_note'))
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row, column)
            _put(sheet, cell.coordinate, value)
            cell.font = Font(name="Times New Roman", size=10)
            cell.border = Border(left=GRID, right=GRID, top=GRID, bottom=GRID)
            cell.alignment = Alignment(vertical="top", wrap_text=column in {6, 8, 9})
        sheet.cell(row, 2).number_format = "dd/mm/yyyy"
        sheet.cell(row, 11).number_format = "#,##0.######"
        sheet.cell(row, 12).number_format = "#,##0"
        sheet.cell(row, 13).number_format = "#,##0"
        if item["movement_label"].startswith("Hoàn tác"):
            for column in range(1, 15):
                sheet.cell(row, column).fill = PatternFill("solid", fgColor=PALE_YELLOW)

    total_row = 5 + len(rows)
    _put(sheet, f"A{total_row}", "TỔNG" if is_input else "TỔNG GIÁ VỐN HÀNG XUẤT")
    sheet.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=10)
    quantity_field = "input_qty" if is_input else "output_qty"
    value_field = "input_value" if is_input else "output_value"
    _put(sheet, f"K{total_row}", _quantity_total(model, quantity_field))
    sheet[f'K{total_row}'].alignment = Alignment(wrap_text=True, vertical='center')
    sheet.row_dimensions[total_row].height = 42
    _put(sheet, f"M{total_row}", model["totals"][value_field])
    for column in range(1, 15):
        cell = sheet.cell(total_row, column)
        cell.font = Font(name="Times New Roman", size=10, bold=True)
        cell.fill = PatternFill("solid", fgColor=PALE_GREEN)
        cell.border = Border(left=GRID, right=GRID, top=GRID, bottom=GRID)
    sheet[f"K{total_row}"].number_format = "#,##0.######"
    sheet[f"M{total_row}"].number_format = "#,##0"
    if is_input:
        try:
            from .invoice_line_tax import complete_sum
        except ImportError:
            from invoice_line_tax import complete_sum
        for col, field in ((16, 'line_tax_amount'), (17, 'amount_with_tax')):
            total = complete_sum(rows, field)
            _put(sheet, f'{chr(64 + col)}{total_row}', total if total is not None else 'Chưa đủ thuế nguồn')
            for row_cells in sheet.iter_rows(min_row=5, min_col=col, max_col=col, max_row=total_row):
                row_cells[0].number_format = '#,##0.##'
    widths = (6, 12, 13, 12, 16, 28, 14, 28, 32, 11, 14, 16, 18, 16) + ((12, 18, 20, 44) if is_input else ())
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + column)].width = width
    sheet.freeze_panes = "A5"
    last_column = "R" if is_input else "N"
    sheet.auto_filter.ref = f"A4:{last_column}{max(4, total_row - 1)}"
    _configure_print(sheet, last_row=total_row, last_column=last_column, title_rows="$1:$4")
    _add_unit_totals(workbook, model)
    _add_control_sheet(workbook, model, direction)
    workbook.active = 0
    assert_workbook_safe(workbook)
    return workbook


def build_nxt_workbook(model: Mapping[str, Any]) -> Any:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "NXT"
    sheet.merge_cells("A1:S1")
    sheet.merge_cells("A2:S2")
    sheet.merge_cells("A3:S3")
    _put(sheet, "A1", "BẢNG NHẬP – XUẤT – TỒN")
    _put(sheet, "A2", _period_text(model))
    _put(
        sheet, "A3",
        "TỒN CUỐI = TỒN ĐẦU + NHẬP − XUẤT · Giá vốn: bình quân gia quyền di động",
    )
    sheet["A1"].font = Font(name="Times New Roman", size=16, bold=True)
    sheet["A1"].alignment = Alignment(horizontal="center")
    sheet["A2"].font = Font(name="Times New Roman", size=10, italic=True)
    sheet["A2"].alignment = Alignment(horizontal="center")
    sheet["A3"].font = Font(name="Times New Roman", size=10, bold=True, color="548235")
    sheet["A3"].alignment = Alignment(horizontal="center")
    fixed = (("A4", "A5", "Mã TĐP"), ("B4", "B5", "Tên TĐP"),
             ("C4", "C5", "Tên trên HĐ"), ("D4", "D5", "Mã kho"),
             ("E4", "E5", "T/Suất"), ("F4", "F5", "ĐVT"),
             ("S4", "S5", "Đối chiếu"))
    for start, end, label in fixed:
        sheet.merge_cells(f"{start}:{end}")
        _put(sheet, start, label)
    groups = (
        ("G4", "I4", "TỒN ĐẦU KỲ", ("Số lượng", "Đơn giá", "Thành tiền")),
        ("J4", "L4", "NHẬP TRONG KỲ", ("Số lượng", "Đơn giá BQ", "Thành tiền")),
        ("M4", "O4", "XUẤT TRONG KỲ", ("Số lượng", "Đơn giá BQ", "Thành tiền")),
        ("P4", "R4", "TỒN CUỐI KỲ", ("Số lượng", "Đơn giá BQ", "Thành tiền")),
    )
    for start, end, label, subheaders in groups:
        sheet.merge_cells(f"{start}:{end}")
        _put(sheet, start, label)
        first_column = ord(start[0]) - 64
        for offset, subheader in enumerate(subheaders):
            _put(sheet, f"{chr(64 + first_column + offset)}5", subheader)
    for row in (4, 5):
        for column in range(1, 20):
            cell = sheet.cell(row, column)
            cell.font = Font(name="Times New Roman", size=10, bold=True, color=WHITE)
            cell.fill = PatternFill("solid", fgColor=GREEN)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(left=GRID, right=GRID, top=GRID, bottom=GRID)

    for offset, item in enumerate(model["items"]):
        row = 6 + offset
        status = "OK" if item["valuation_status"] == "ok" else "CẦN KIỂM TRA"
        values = (
            item["product_code"], item["product_name"], item["invoice_name"],
            item["warehouse_code"], item["tax"], item["unit"],
            item["opening_qty"], item["opening_unit_cost"], item["opening_value"],
            item["input_qty"], item["input_unit_cost"], item["input_value"],
            item["output_qty"], item["output_unit_cost"], item["output_value"],
            item["closing_qty"], item["closing_unit_cost"], item["closing_value"], status,
        )
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row, column)
            _put(sheet, cell.coordinate, value)
            cell.font = Font(name="Times New Roman", size=10)
            cell.border = Border(left=GRID, right=GRID, top=GRID, bottom=GRID)
            cell.alignment = Alignment(vertical="top", wrap_text=column in {2, 3, 19})
        for column in (7, 10, 13, 16):
            sheet.cell(row, column).number_format = "#,##0.######"
        for column in (8, 11, 14, 17):
            sheet.cell(row, column).number_format = "#,##0"
        for column in (9, 12, 15, 18):
            sheet.cell(row, column).number_format = "#,##0"
        if status != "OK":
            for column in range(1, 20):
                sheet.cell(row, column).fill = PatternFill("solid", fgColor=PALE_YELLOW)

    total_row = 6 + len(model["items"])
    _put(sheet, f"A{total_row}", "TỔNG")
    sheet.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=6)
    total_columns = {
        7: "opening_qty", 9: "opening_value", 10: "input_qty", 12: "input_value",
        13: "output_qty", 15: "output_value", 16: "closing_qty", 18: "closing_value",
    }
    for column, field in total_columns.items():
        _put(sheet, f"{chr(64 + column)}{total_row}", _quantity_total(model, field) if field.endswith('_qty') else model["totals"][field])
        sheet.cell(total_row,column).alignment = Alignment(wrap_text=True, vertical='center')
    sheet.row_dimensions[total_row].height = 42
    _put(sheet, f"S{total_row}", "CẦN KIỂM TRA" if any(r['valuation_status'] != 'ok' for r in model['items']) else "KHỚP")
    for column in range(1, 20):
        cell = sheet.cell(total_row, column)
        cell.font = Font(name="Times New Roman", size=10, bold=True)
        cell.fill = PatternFill("solid", fgColor=PALE_GREEN)
        cell.border = Border(left=GRID, right=GRID, top=GRID, bottom=GRID)
    for column in (7, 10, 13, 16):
        sheet.cell(total_row, column).number_format = "#,##0.######"
    for column in (9, 12, 15, 18):
        sheet.cell(total_row, column).number_format = "#,##0"
    widths = (
        14, 25, 28, 13, 10, 10,
        13, 13, 17, 13, 13, 17, 13, 13, 17, 13, 13, 17, 15,
    )
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + column)].width = width
    sheet.freeze_panes = "G6"
    sheet.auto_filter.ref = f"A5:S{max(5, total_row - 1)}"
    _configure_print(sheet, last_row=total_row, last_column="S", title_rows="$1:$5")
    _set_core_properties(workbook, model, "Bảng Nhập – Xuất – Tồn")
    _add_unit_totals(workbook, model)
    _add_control_sheet(workbook, model, "nxt")
    workbook.active = 0
    assert_workbook_safe(workbook)
    return workbook


def build_inventory_workbook(
    model: Mapping[str, Any], kind: str, *, template_path: str | Path,
    expected_sha256: str = OPENING_TEMPLATE_SHA256,
) -> Any:
    if kind == "opening":
        return build_opening_workbook(
            model, template_path=template_path, expected_sha256=expected_sha256,
        )
    if kind in {"input", "output"}:
        return build_movement_workbook(model, direction=kind)
    if kind == "nxt":
        return build_nxt_workbook(model)
    raise InventoryExportError("Loại file TĐK–NXT không hợp lệ", code="invalid_export_kind", status=404)


def _filenames(model: Mapping[str, Any]) -> dict[str, str]:
    suffix = f"{model['date_from']}_den_{model['date_to']}"
    return {
        "opening": f"TDK_{suffix}.xlsx",
        "input": f"Nhap_trong_ky_{suffix}.xlsx",
        "output": f"Xuat_gia_von_{suffix}.xlsx",
        "nxt": f"NXT_{suffix}.xlsx",
    }


def inventory_workbook_bytes(
    model: Mapping[str, Any], kind: str, *, template_path: str | Path,
    expected_sha256: str = OPENING_TEMPLATE_SHA256,
) -> bytes:
    workbook = build_inventory_workbook(
        model, kind, template_path=template_path, expected_sha256=expected_sha256,
    )
    try:
        return safe_workbook_bytes(workbook)
    finally:
        workbook.close()


def inventory_archive_bytes(
    model: Mapping[str, Any], *, template_path: str | Path,
    expected_sha256: str = OPENING_TEMPLATE_SHA256,
) -> bytes:
    output = io.BytesIO()
    names = _filenames(model)
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for kind in EXPORT_KINDS:
            info = zipfile.ZipInfo(names[kind])
            safe_day = datetime.strptime(model["date_to"], "%Y-%m-%d")
            info.date_time = (safe_day.year, safe_day.month, safe_day.day, 0, 0, 0)
            info.compress_type = zipfile.ZIP_DEFLATED
            archive.writestr(
                info,
                inventory_workbook_bytes(
                    model,
                    kind,
                    template_path=template_path,
                    expected_sha256=expected_sha256,
                ),
            )
    return output.getvalue()


def register_inventory_export_routes(app: Any, ctx: Mapping[str, Any]) -> None:
    db_factory = ctx["db"]
    template_path = ctx["opening_template_path"]
    expected_sha256 = ctx.get("opening_template_sha256", OPENING_TEMPLATE_SHA256)

    def sales_document(conn):
        # The sales file must not depend on inventory valuation or posting.
        try:
            from .invoice_workbench_listing import invoice_range_payload
            from .invoice_output_register import output_sales_workbook
            from .invoice_workbench import InvoiceWorkbenchError
        except ImportError:
            from invoice_workbench_listing import invoice_range_payload
            from invoice_output_register import output_sales_workbook
            from invoice_workbench import InvoiceWorkbenchError
        try:
            tenant = conn.execute("SELECT value FROM settings WHERE key='tenant_code'").fetchone()
            payload = invoice_range_payload(conn, tenant=str(tenant[0] if tenant else 'TDP').strip() or 'TDP',
                invoice_type='output', date_from=request.args.get('from'), date_to=request.args.get('to'))
            name = f"Xuat_gia_ban_M-Invoice_{payload['date_from']}_den_{payload['date_to']}.xlsx"
            return payload, name, output_sales_workbook(payload).getvalue()
        except InvoiceWorkbenchError as error:
            raise InventoryExportError(str(error), code="invalid_period", status=400) from None

    def error_response(error: Exception):
        if isinstance(error, (InventoryExportError, TemplateWorkbookError)):
            status = getattr(error, "status", 422)
            code = getattr(error, "code", "invalid_inventory_export")
            return jsonify({"ok": False, "error": str(error), "code": code}), status
        raise error

    def customer_document(conn, kind, stock=None, sales=None):
        if not conn.in_transaction:
            conn.execute('BEGIN')
        if kind == 'output_sales':
            kind = 'output'
        if kind not in EXPORT_KINDS:
            raise InventoryExportError("Loại báo cáo không hợp lệ", code="invalid_export_kind", status=404)
        if kind in {'output', 'nxt'}:
            sales = sales or sales_document(conn)
        if kind == 'output':
            return sales
        stock = stock or collect_inventory_export_model(conn,
            date_from=request.args.get('from'), date_to=request.args.get('to'))
        if kind in {'nxt', 'opening'}:
            stock = monthly_customer_model(conn, stock)
        if kind == 'nxt':
            try:
                from .inventory_customer_report import customer_nxt_workbook
            except ImportError:
                from inventory_customer_report import customer_nxt_workbook
            book = customer_nxt_workbook(stock, sales[0])
            try:
                content = safe_workbook_bytes(book)
            finally:
                book.close()
        else:
            content = inventory_workbook_bytes(stock, kind, template_path=template_path,
                                               expected_sha256=expected_sha256)
        return stock, _filenames(stock)[kind], content

    @app.get("/api/invoice-valuation/export")
    def api_inventory_export_archive():
        try:
            with db_factory() as conn:
                if not conn.in_transaction:
                    conn.execute('BEGIN')
                stock = collect_inventory_export_model(conn,
                    date_from=request.args.get('from'), date_to=request.args.get('to'))
                stock = monthly_customer_model(conn, stock)
                sales = sales_document(conn)
                documents = [customer_document(conn, kind, stock, sales) for kind in EXPORT_KINDS]
            stream = io.BytesIO()
            with zipfile.ZipFile(stream, 'w', compression=zipfile.ZIP_DEFLATED) as archive:
                for _, name, content in documents:
                    archive.writestr(name, content)
            stream.seek(0)
            return send_file(stream, as_attachment=True,
                download_name=f"TDK_NXT_{stock['date_from']}_den_{stock['date_to']}.zip", mimetype="application/zip")
        except (InventoryExportError, TemplateWorkbookError) as error:
            return error_response(error)

    @app.get("/api/invoice-valuation/export/<kind>")
    def api_inventory_export_one(kind: str):
        try:
            with db_factory() as conn:
                _, filename, content = customer_document(conn, kind)
            return send_file(io.BytesIO(content), as_attachment=True, download_name=filename,
                mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        except (InventoryExportError, TemplateWorkbookError) as error:
            return error_response(error)

    @app.get("/api/invoice-valuation/preview/<kind>")
    def api_inventory_preview(kind: str):
        try:
            try:
                from inventory_preview import workbook_preview
            except ImportError:
                from .inventory_preview import workbook_preview
            with db_factory() as conn:
                payload, filename, content = customer_document(conn, kind)
            workbook = load_workbook(io.BytesIO(content))
            try:
                preview = workbook_preview(workbook, filename)
            finally:
                workbook.close()
            return jsonify(ok=True, read_only=True, kind=kind, date_from=payload['date_from'],
                           date_to=payload['date_to'], workbook=preview)
        except (InventoryExportError, TemplateWorkbookError) as error:
            return error_response(error)
        except ValueError as error:
            return jsonify(ok=False, error=str(error)), 422


__all__ = [
    "EXPORT_KINDS",
    "InventoryExportError",
    "OFFICIAL_LAYOUT_ID",
    "OPENING_TEMPLATE_SHA256",
    "build_inventory_workbook",
    "build_movement_workbook",
    "build_nxt_workbook",
    "build_opening_workbook",
    "collect_inventory_export_model",
    "inventory_archive_bytes",
    "inventory_workbook_bytes",
    "register_inventory_export_routes",
]
