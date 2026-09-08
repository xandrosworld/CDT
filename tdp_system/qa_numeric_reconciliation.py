"""Reconcile screen payloads and exports on a disposable production snapshot.

No network connectors are used. Private artifacts must stay outside the repository.
Matching a stored source does not certify that the source is business-correct.
"""
import argparse
from collections import Counter, defaultdict
from decimal import Decimal, ROUND_HALF_UP
import hashlib
from html.parser import HTMLParser
import io
import json
import os
from pathlib import Path
import sqlite3
import traceback
from urllib.parse import urlencode

from openpyxl import load_workbook


def dec(value):
    return Decimal(str(value or 0))


def vnd(value):
    return int(dec(value).quantize(Decimal('1'), rounding=ROUND_HALF_UP))


def visible_cells(html):
    """Compare text and merged-cell topology, independent of XML style normalization."""
    class Cells(HTMLParser):
        def __init__(self):
            super().__init__(); self.cells=[]; self.current=None
        def handle_starttag(self, tag, attrs):
            if tag=='td':
                attrs=dict(attrs)
                self.current=[attrs.get('rowspan','1'),attrs.get('colspan','1'),'']
        def handle_data(self, text):
            if self.current is not None: self.current[2]+=text
        def handle_endtag(self, tag):
            if tag=='td' and self.current is not None:
                self.cells.append(self.current); self.current=None
    parser=Cells(); parser.feed(html); return parser.cells


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--snapshot', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    output = args.output.resolve()
    output.mkdir(parents=True, exist_ok=False)
    source = args.snapshot.resolve()
    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    target = output / 'copy.sqlite3'
    with sqlite3.connect(source.as_uri()+'?mode=ro', uri=True) as original, sqlite3.connect(target) as copy:
        original.backup(copy)
    os.environ.update(TDP_DB_PATH=str(target), TDP_DATA_DIR=str(output/'data'),
                      TDP_EXPORT_DIR=str(output/'exports'), MSMI_API_BASE_URL='http://127.0.0.1:9',
                      MINVOICE_API_BASE_URL='http://127.0.0.1:9')
    from . import server
    assert server.DB_PATH.resolve() == target, 'Audit must use only the disposable database'
    server.connector_config_paths = lambda: []
    client = server.app.test_client()
    report = {'checks': [], 'failures': [], 'source_issues': []}

    def data(url):
        response = client.get(url)
        assert response.status_code == 200, (url, response.status_code, response.get_json())
        return response.get_json()

    def workbook(url, name):
        response = client.get(url)
        assert response.status_code == 200, (url, response.status_code, response.get_json())
        (output/(name+'.xlsx')).write_bytes(response.data)
        return load_workbook(io.BytesIO(response.data), data_only=False)

    def check(name, fn):
        try:
            detail = fn()
            report['checks'].append({'name': name, 'detail': detail})
            print(json.dumps({'check': name, 'ok': True}), flush=True)
        except Exception as error:
            report['failures'].append({'name': name, 'error': str(error), 'trace': traceback.format_exc()})
            print(json.dumps({'check': name, 'ok': False, 'error': str(error)}, ensure_ascii=True), flush=True)
        (output/'summary.json').write_text(json.dumps(report, ensure_ascii=False, indent=2, default=str), encoding='utf-8')

    def orders():
        p = data('/api/bootstrap')
        (output/'bootstrap.json').write_text(json.dumps(p, ensure_ascii=False), encoding='utf-8')
        sums = defaultdict(int)
        with server.db() as conn:
            originals = {r['id']: dict(r) for r in conn.execute('SELECT * FROM orders')}
        for row in p['orders']:
            raw = originals[row['id']]
            for key in ('qty','actual_received','actual_delivered','buy_price','sell_price','damaged_qty','supplier_return_qty','customer_return_qty'):
                assert dec(row[key]) == dec(raw[key]), (row['id'], key)
            received = max(Decimal(0), dec(raw['actual_received'])-dec(raw['damaged_qty'])-dec(raw['supplier_return_qty']))
            delivered = max(Decimal(0), dec(raw['actual_delivered'])-dec(raw['customer_return_qty']))
            revenue, cost = vnd(delivered*dec(raw['sell_price'])), vnd(received*dec(raw['buy_price']))
            tax = dec(str(raw['tax']).replace('%','')) if str(raw['tax']).replace('%','').replace('.','',1).isdigit() else Decimal(0)
            if '%' not in str(raw['tax']) and 0 < tax < 1: tax *= 100
            expected = dict(revenue=revenue, cost=cost, profit=revenue-cost, total=revenue+vnd(dec(revenue)*tax/100))
            for key, value in expected.items():
                assert row[key] == value, (row['id'], key, row[key], value)
                sums[key] += value
        for key, value in sums.items():
            assert p['summary']['totals'][key] == value, (key, value)
        report['source_issues'].append({'kind':'orders', 'rows':len(p['orders']),
            'errors':sum(bool(r['errors']) for r in p['orders']), 'warnings':sum(bool(r['warnings']) for r in p['orders'])})
        return dict(rows=len(p['orders']), **sums)

    check('orders_database_calculation_and_screen', orders)

    def invoices(direction, start, end):
        query=urlencode(dict(invoice_type=direction, **{'from':start,'to':end,'status':'all','line_filter':'all'}))
        p=data('/api/invoice-workbench/invoices?'+query)
        name=f'invoices-{direction}-{start}'
        (output/(name+'.json')).write_text(json.dumps(p,ensure_ascii=False),encoding='utf-8')
        wb=workbook('/api/invoice-workbench/invoices/export?'+query,name)
        ws=wb.active
        columns = {cell.value: cell.column for cell in ws[2]}
        fields = {'Tên hàng':'source_item_name','ĐVT':'source_unit','Số lượng':'qty',
                  'Đơn giá':'unit_price','Thành tiền dòng (chưa thuế)':'amount',
                  'Mã kho':'product_code','Lượng kho':'stock_qty','ĐVT kho':'product_unit'}
        for index,line in enumerate(p['lines'],3):
            for label,field in fields.items():
                column = columns[label]
                actual=ws.cell(index,column).value
                expected=line.get(field)
                if field in ('qty','unit_price','amount','stock_qty'): assert abs(dec(actual)-dec(expected))<Decimal('0.000001'),(index,field,actual,expected)
                else: assert (actual or '') == (expected or ''),(index,field)
        line_amount=sum((dec(line.get('amount')) for line in p['lines']),Decimal(0))
        invoice_amount=sum((dec(invoice.get('total_amount')) for invoice in p['items']),Decimal(0))
        assert abs(dec(p['totals']['line_amount'])-line_amount)<Decimal('0.000001')
        assert abs(dec(p['totals']['invoice_amount'])-invoice_amount)<Decimal('0.000001')
        amount_column = columns['Thành tiền dòng (chưa thuế)']
        assert abs(dec(ws.cell(len(p['lines'])+3,amount_column).value)-line_amount)<Decimal('0.000001')
        assert abs(dec(ws.cell(ws.max_row,amount_column).value)-invoice_amount)<Decimal('0.000001')
        with server.db() as conn:
            table='msmi_invoices' if direction=='input' else 'outgoing_source_invoices'
            sql='SELECT * FROM '+table+' WHERE invoice_date BETWEEN ? AND ?'
            params=[start,end]
            if direction=='input': sql+=" AND invoice_type='INPUT_ELECTRONIC_INVOICE'"
            else:
                tenant=conn.execute("SELECT value FROM settings WHERE key='tenant_code'").fetchone()
                sql+=" AND source='minvoice' AND tenant=?"
                params.append(tenant[0] if tenant else 'TDP')
                excluded=conn.execute("SELECT source,COUNT(*) FROM outgoing_source_invoices WHERE invoice_date BETWEEN ? AND ? AND source!='minvoice' GROUP BY source",(start,end)).fetchall()
                report['source_issues'].append({'kind':'excluded_nonproduction_output_sources','period':start,'sources':[dict(source=r[0],invoices=r[1]) for r in excluded]})
            rows=[dict(r) for r in conn.execute(sql,params)]
        assert {r['id'] for r in rows} == {r['id'] for r in p['items']}
        for row in rows:
            api_row=next(r for r in p['items'] if r['id']==row['id'])
            for field in ('subtotal','tax_amount','total_amount'):
                assert dec(row[field])==dec(api_row[field]),(row['id'],field)
        report['source_issues'].append({'kind':name,'counts':p.get('counts')})
        wb.close()
        return {'invoices':len(rows),'lines':len(p['lines']),'line_amount':line_amount,'invoice_amount':invoice_amount}

    for direction in ('input','output'):
        for start,end in [('2026-08-01','2026-08-31'),('2026-09-01','2026-09-06')]:
            check('invoice_rows_and_export_'+direction+'_'+start,lambda d=direction,s=start,e=end:invoices(d,s,e))

    def inventory(start,end):
        query=urlencode({'from':start,'to':end})
        p=data('/api/invoice-valuation?'+query)
        name='nxt-'+start
        (output/(name+'.json')).write_text(json.dumps(p,ensure_ascii=False),encoding='utf-8')
        books={kind:workbook('/api/invoice-valuation/export/'+kind+'?'+query,name+'-'+kind) for kind in ('opening','input','output','nxt')}
        ws=books['nxt']['NXT']
        api={r['product_code']:r for r in p['items']}
        seen=set()
        for values in ws.iter_rows(min_row=6,values_only=True):
            code=values[0]
            if code not in api: continue
            seen.add(code); row=api[code]
            for column,field in [(7,'opening_qty'),(9,'opening_value'),(10,'input_qty'),(12,'input_value'),(13,'output_qty'),(15,'output_value'),(16,'closing_qty'),(18,'closing_value')]:
                assert abs(dec(values[column-1])-dec(row[field]))<Decimal('0.00001'),(code,field,values[column-1],row[field])
            for suffix in ('qty','value'):
                balance=dec(row['opening_'+suffix])+dec(row['input_'+suffix])-dec(row['output_'+suffix])
                assert abs(balance-dec(row['closing_'+suffix]))<Decimal('0.00001'),(code,suffix)
        assert seen==set(api), ('missing',set(api)-seen)
        opening=books['opening'].worksheets[0]
        opening_rows={opening.cell(i,1).value:i for i in range(7,opening.max_row+1) if opening.cell(i,1).value in api}
        for code,index in opening_rows.items():
            assert dec(opening.cell(index,7).value)==dec(api[code]['opening_qty'])
            assert abs(dec(opening.cell(index,9).value)-dec(api[code]['opening_value']))<Decimal('0.00001')
            assert opening.cell(index,9).number_format=='#,##0'
        assert abs(dec(opening['I4'].value)-sum(dec(r['opening_value']) for r in api.values()))<Decimal('0.00001')
        unit_sums=defaultdict(lambda:[Decimal(0)]*4)
        for row in api.values():
            for i,field in enumerate(('opening_qty','input_qty','output_qty','closing_qty')):
                unit_sums[row['unit'].casefold()][i]+=dec(row[field])
        for book in books.values():
            if len(unit_sums)>1:
                assert 'Tổng ĐVT' in book.sheetnames
                actual={str(r[0]).casefold():list(r[1:]) for r in book['Tổng ĐVT'].iter_rows(min_row=2,values_only=True)}
                for unit,values in unit_sums.items():
                    if not any(values):continue
                    assert unit in actual, unit
                    assert all(abs(dec(a)-b)<Decimal('0.000001') for a,b in zip(actual[unit],values)),unit
        assert ws.cell(ws.max_row,19).value==('CẦN KIỂM TRA' if any(r['valuation_status']!='ok' for r in api.values()) else 'KHỚP')
        # Save all four files for independent trace and investigate any source exceptions.
        report['source_issues'].append({'kind':name,'needs_review':[{'code':r['product_code'],'quantity':r['closing_qty'],'status':r['valuation_status']} for r in api.values() if r['valuation_status']!='ok']})
        for book in books.values():book.close()
        return {'rows':len(api),'closing_value':sum(dec(r['closing_value']) for r in api.values())}

    for start,end in [('2026-08-01','2026-08-31'),('2026-09-01','2026-09-06')]:
        check('nxt_screen_and_four_exports_'+start,lambda s=start,e=end:inventory(s,e))

    def monthly(period):
        p=data('/api/reports/monthly?period='+period)
        wb=workbook('/api/reports/monthly/export?period='+period,'monthly-'+period)
        actual=[list(row) for row in wb.active.iter_rows(min_row=3,values_only=True)]
        assert actual==p['rows']
        with server.db() as conn:
            approved=conn.execute("SELECT COUNT(*) FROM batches WHERE substr(work_date,1,7)=? AND status='approved'",(period,)).fetchone()[0]
        if not approved:
            assert all(not any(dec(v) for v in row[3:7]) for row in actual), 'Draft data leaked into approved monthly report'
        wb.close()
        return {'rows':len(actual),'draft_count':p['draft_count'],'total':actual[-1] if actual else None}
    for period in ('2026-08','2026-09'):
        check('monthly_report_screen_and_excel_'+period,lambda p=period:monthly(p))

    def debts(kind,status):
        query=urlencode({'from':'2026-08-01','to':'2026-09-06','status':status,'limit':20000})
        p=data('/api/debts/'+kind+'/ledger?'+query)
        rows=p['rows']
        assert len(rows)==p['pagination']['total']
        suffix='/lines/export' if kind=='receivables' else '/export'
        wb=workbook('/api/debts/'+kind+suffix+'?'+query,kind+'-'+status)
        ws=wb.active
        assert ws.max_row-4==len(rows)
        amount=sum(dec(r['amount']) for r in rows)
        assert amount==dec(p['summary']['filtered_amount'])
        assert amount==dec(ws.cell(ws.max_row,13 if kind=='receivables' else 14).value)
        if kind=='receivables':
            by_id={r['id']:r for r in rows}
            for cells in ws.iter_rows(min_row=4,max_row=ws.max_row-1,values_only=True):
                row=by_id[cells[14]]
                for column,field in [(5,'ordered_qty'),(6,'actual_delivered'),(7,'customer_return_qty'),(8,'delivered_qty'),(10,'sell_price'),(11,'subtotal'),(12,'tax_amount'),(13,'amount')]:
                    assert abs(dec(cells[column-1])-dec(row[field]))<Decimal('0.000001'),(row['id'],field)
                assert row['subtotal']==vnd(dec(row['delivered_qty'])*dec(row['sell_price']))
                assert row['amount']==row['subtotal']+row['tax_amount']
        else:
            actual=Counter((str(r[1]),str(r[3]),dec(r[12]),str(r[5]),str(r[6]),dec(r[7]),dec(r[13])) for r in ws.iter_rows(min_row=4,max_row=ws.max_row-1,values_only=True))
            expected=Counter((str(r['kitchen']),r['product_name'],dec(r['actual_qty']),r['unit'],r['supplier']['code'],dec(r['buy_price']),dec(r['amount'])) for r in rows)
            assert actual==expected,'Payable detail cells differ'
            for row in rows:
                assert row['amount']==vnd(dec(row['actual_qty'])*dec(row['buy_price'])),row['id']
            assert ws.max_column==14
        wb.close()
        return {'rows':len(rows),'amount':amount,'statuses':p['summary']['status_counts']}
    for kind in ('receivables','payables'):
        for status in ('all','active' if kind=='receivables' else 'open,partially_paid'):
            check(kind+'_rows_and_excel_'+status,lambda k=kind,s=status:debts(k,s))

    def suppliers_and_deliveries():
        from .document_preview import formula_value
        p=data('/api/bootstrap');batch=p['batch'];rows=p['orders']
        with server.db() as conn: purchase=server.purchase_order_payload(conn,batch['id'])
        wb=workbook('/api/export/suppliers/'+str(batch['id']),'suppliers')
        ws=wb.active
        for index,row in enumerate(purchase['rows'],3):
            assert (ws.cell(index,2).value or '')==(row['product_code'] or '')
            assert (ws.cell(index,8).value or '')==(row['supplier'] or '')
            assert abs(dec(ws.cell(index,6).value)-dec(row.get('demand_qty',row.get('order_qty',0))))<Decimal('0.000001')
            actual_qty=dec(ws.cell(index,6).value)+dec(ws.cell(index,12).value)-dec(ws.cell(index,11).value)-dec(ws.cell(index,13).value)-dec(ws.cell(index,14).value)
            assert abs(dec(formula_value(ws,ws.cell(index,15)))-actual_qty)<Decimal('0.000001')
            assert vnd(formula_value(ws,ws.cell(index,16)))==vnd(actual_qty*dec(row['buy_price']))
        wb.close()
        delivery=workbook('/api/export/deliveries/'+str(batch['id']),'deliveries')
        expected=Counter((r['product_name'].strip(),max(dec(r['actual_delivered'])-dec(r['customer_return_qty']),Decimal(0)),r['unit'].strip()) for r in rows if dec(r['actual_delivered'])>dec(r['customer_return_qty']))
        found=Counter();priced=0
        for ws in delivery.worksheets:
            if ws.sheet_state!='visible':continue
            for i in range(11,ws.max_row+1):
                if not isinstance(ws.cell(i,3).value,int) or not isinstance(ws.cell(i,5).value,(int,float)):continue
                qty=dec(ws.cell(i,5).value)
                found[(ws.cell(i,4).value,qty,ws.cell(i,6).value)]+=1
                if ws.cell(i,8).value is not None:
                    assert ws.cell(i,9).value==vnd(qty*dec(ws.cell(i,8).value));priced+=1
        assert found==expected, 'Delivery rows differ from net delivered source'
        delivery.close()
        return {'supplier_rows':len(purchase['rows']),'delivery_rows':sum(found.values()),'priced_delivery_rows':priced}
    check('supplier_formulas_and_delivery_rows',suppliers_and_deliveries)

    def physical():
        p=data('/api/physical-stock')
        for row in p['items']:
            if row['opening'] is None:
                assert row['balance_qty'] is None
            else:
                assert abs(dec(row['balance_qty'])-(dec(row['opening_qty'])+dec(row['received_qty'])-dec(row['committed_qty'])))<Decimal('0.000001')
        report['source_issues'].append({'kind':'physical','totals':p['totals']})
        return p['totals']
    check('physical_stock_source_and_uninitialized_state',physical)

    def other_scopes():
        ops=data('/api/operations/bootstrap?as_of=2026-09-06&date=2026-09-06&month=2026-09&active_only=1')
        (output/'operations.json').write_text(json.dumps(ops,ensure_ascii=False),encoding='utf-8')
        totals=data('/api/debts?from=2026-08-01&to=2026-09-06')
        (output/'debt-totals.json').write_text(json.dumps(totals,ensure_ascii=False),encoding='utf-8')
        wb=workbook('/api/export/debts?from=2026-08-01&to=2026-09-06','debt-totals')
        for sheet,section in [('Phải thu','contractors'),('Phải trả','suppliers')]:
            rows=totals[section]
            actual={r[0]:r[1:6] for r in wb[sheet].iter_rows(min_row=4,values_only=True)}
            assert set(actual)==set(rows),section
            for code,row in rows.items():
                expected=[row[k] for k in ('opening','period_charge','period_adjustment','period_paid','closing')]
                assert list(actual[code])==expected,(section,code)
                assert dec(row['closing'])==dec(row['opening'])+dec(row['period_charge'])+dec(row['period_adjustment'])-dec(row['period_paid'])
        wb.close()
        with server.db() as conn: contractors=[r[0] for r in conn.execute('SELECT code FROM contractors ORDER BY code')]
        scopes=[]
        for contractor in contractors:
            url='/api/quotes?'+urlencode({'contractor':contractor,'period':'2026-08','batch_id':1})
            response=client.get(url);quote=response.get_json()
            assert response.status_code==200,quote
            scope={'contractor':contractor,'status':response.status_code,'rows':len(quote['items'])}
            if quote['conflicts']:
                rejected=client.get('/api/export/quote/'+contractor+'?period=2026-08&batch_id=1')
                assert rejected.status_code==409 and rejected.get_json()['code']=='quote_duplicate_conflict'
                scope['blocked_conflicts']=len(quote['conflicts'])
            elif quote['mode']!='daily' and not quote['version']:
                rejected=client.get('/api/export/quote/'+contractor+'?period=2026-08&batch_id=1')
                assert rejected.status_code==409 and rejected.get_json()['code']=='quote_period_not_confirmed'
                scope['blocked_reason']='quote_period_not_confirmed'
                with server.db() as conn:
                    prices={r['product_code']:dict(r) for r in conn.execute('SELECT * FROM product_prices WHERE price_group=?',(quote['price_group'],))}
                for row in quote['items']:
                    assert dec(row['sell_price'])==dec(prices[row['product_code']]['price_value'])
            else:
                book=workbook('/api/export/quote/'+contractor+'?period=2026-08&batch_id=1','quote-'+contractor)
                expected={r['product_code']:r for r in quote['items'] if r['exportable'] and r['price_state'] in ('numeric','zero')}
                found={r[1]:r for r in book.active.iter_rows(min_row=9,values_only=True) if isinstance(r[0],int)}
                assert set(found)==set(expected),contractor
                for code,row in expected.items():
                    cells=found[code]
                    assert cells[2]==' '.join(row['product_name'].split()) and (cells[3] or '')==(row['unit'] or '')
                    assert dec(cells[4])==vnd(row['sell_price']),(contractor,code)
                scope['exported_rows']=len(found)
                book.close()
            scopes.append(scope)
        report['source_issues'].append({'kind':'quotes_available_scope','items':scopes})
        return {'contractors':len(contractors),'quote_export_rows':sum(r.get('exported_rows',0) for r in scopes),
                'blocked_quote_scopes':sum('blocked_conflicts' in r or 'blocked_reason' in r for r in scopes),
                'supplier_opening':sum(dec(r['opening']) for r in totals['suppliers'].values()),'debt_sections':list(totals)}
    check('catalog_quote_and_debt_available_scopes',other_scopes)

    def documents():
        from .document_preview import sheet_preview
        outcomes=[]
        batch=data('/api/bootstrap')['batch']
        with server.db() as conn:
            kitchens=[r[0] for r in conn.execute('SELECT DISTINCT kitchen FROM orders WHERE batch_id=? ORDER BY kitchen',(batch['id'],))]
        for kind in ('deliveries','suppliers','purchases','report','outgoing-statement'):
            selections=[{'batch_id':batch['id'],'kitchen':kitchen} for kitchen in kitchens] if kind=='deliveries' else [{'batch_id':batch['id']}]
            response=client.post('/api/documents/preview',json={'kind':kind,'selections':selections})
            payload=response.get_json()
            if response.status_code in (409,422):
                assert kind=='purchases', (kind,payload)
                outcomes.append({'kind':kind,'blocked':True,'status':response.status_code,'error':payload.get('error')})
                continue
            assert response.status_code==200,(kind,payload)
            for index,preview in enumerate(payload['sheets']):
                book=workbook('/api/documents/'+payload['token']+'/excel?sheets='+str(index),'document-'+kind+'-'+str(index))
                visible=[s for s in book if s.sheet_state=='visible']
                assert len(visible)==1
                exported=sheet_preview(visible[0])['html']
                assert visible_cells(exported)==visible_cells(preview['html']),(kind,index,'Visible cells differ from exported snapshot')
                book.close()
            outcomes.append({'kind':kind,'sheets_compared':len(payload['sheets'])})
        with server.db() as conn: contractors=[r[0] for r in conn.execute('SELECT code FROM contractors')]
        for contractor in contractors:
            response=client.get('/api/outgoing-invoices/payment-scope/'+contractor+'?from=2026-08-01&to=2026-08-31')
            payload=response.get_json()
            assert response.status_code in (200,404,409,422),(contractor,payload)
            outcomes.append({'kind':'vat_payment','contractor':contractor,'status':response.status_code,'code':payload.get('code'),
                'blocked':response.status_code!=200})
        report['source_issues'].append({'kind':'document_scope_results','items':outcomes})
        return {'compared_sheets':sum(r.get('sheets_compared',0) for r in outcomes),'blocked_scopes':sum(r.get('blocked',False) for r in outcomes)}
    check('document_preview_cells_and_downloaded_excel',documents)

    report['source_unchanged']=hashlib.sha256(source.read_bytes()).hexdigest()==source_hash
    report['ok']=not report['failures'] and report['source_unchanged']
    (output/'summary.json').write_text(json.dumps(report,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
    return 0 if report['ok'] else 1


if __name__ == '__main__':
    raise SystemExit(main())
