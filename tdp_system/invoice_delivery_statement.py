"""Static delivery statement reconciled to issued outgoing invoices."""

from __future__ import annotations

try:
    from document_totals import quantity_cell
    from document_preview import white_print_style
except ImportError:
    from .document_totals import quantity_cell
    from .document_preview import white_print_style

import math
import re
from collections import defaultdict
from datetime import datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

try:
    from invoice_payment_documents import number_to_vietnamese
except ImportError:  # pragma: no cover - package invocation
    from .invoice_payment_documents import number_to_vietnamese


COMPANY_PHONE = "0904.495.655"


class InvoiceDeliveryStatementError(ValueError):
    def __init__(self, message: str, *, code: str = "invalid_invoice_delivery_statement", status: int = 409):
        super().__init__(message)
        self.code = code
        self.status = status


def _plain(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip())


def _excel_text(value: Any) -> str:
    text = _plain(value)
    return "'" + text if text.startswith(("=", "+", "-", "@")) else text


def _excel_date(value: Any, label: str):
    text = _plain(value)
    try:
        return datetime.strptime(text, "%Y-%m-%d").date()
    except ValueError:
        raise InvoiceDeliveryStatementError(f"{label} không phải ngày YYYY-MM-DD hợp lệ") from None


def _number(value: Any, label: str) -> float:
    if isinstance(value, bool):
        raise InvoiceDeliveryStatementError(f"{label} không hợp lệ")
    try:
        result = float(value)
    except (TypeError, ValueError):
        raise InvoiceDeliveryStatementError(f"{label} không phải số hữu hạn") from None
    if not math.isfinite(result):
        raise InvoiceDeliveryStatementError(f"{label} không phải số hữu hạn")
    return result


def _vnd(*values: Any) -> int:
    try:
        result = Decimal("1")
        for value in values:
            result *= Decimal(str(value))
        return int(result.quantize(Decimal("1"), rounding=ROUND_HALF_UP))
    except (InvalidOperation, TypeError, ValueError):
        raise InvoiceDeliveryStatementError("Không thể làm tròn số tiền bảng kê") from None


def _tax_percent(value: Any) -> float:
    text = _plain(value).upper().replace(" ", "")
    if text in {"", "KKKNT", "KHÔNGKÊKHAI", "KHONGKEKHAI", "-2", "-2.0"}:
        return -2 if text else 0
    if text in {"KCT", "KHÔNGCHỊUTHUẾ", "KHONGCHIUTHUE", "-1", "-1.0"}:
        return -1
    if text.endswith("%"):
        text = text[:-1]
    try:
        result = float(text)
    except ValueError:
        raise InvoiceDeliveryStatementError("Thuế suất bảng kê không hợp lệ") from None
    if 0 < abs(result) < 1:
        result *= 100
    if not math.isfinite(result) or result < 0 or result > 100:
        raise InvoiceDeliveryStatementError("Thuế suất bảng kê không hợp lệ")
    return round(result, 4)


def _style_sheet(worksheet, title: str, subtitle: str, headers: list[str], *, money_columns=()) -> None:
    end_column = len(headers)
    worksheet.insert_rows(1, 2)
    worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_column)
    worksheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end_column)
    worksheet.cell(1, 1, title)
    worksheet.cell(2, 1, subtitle)
    worksheet.cell(1, 1).font = Font(name="Arial", size=16, bold=True, color="FFFFFF")
    worksheet.cell(1, 1).fill = PatternFill("solid", fgColor="17324D")
    worksheet.cell(1, 1).alignment = Alignment(horizontal="center", vertical="center")
    worksheet.cell(2, 1).font = Font(name="Arial", italic=True, color="475569")
    worksheet.cell(2, 1).alignment = Alignment(horizontal="center", vertical="center")
    for cell in worksheet[3]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="087F73")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in worksheet.iter_rows(min_row=4):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for column in money_columns:
        for row in range(4, worksheet.max_row + 1):
            worksheet.cell(row, column).number_format = "#,##0"
    worksheet.freeze_panes = "A4"
    worksheet.auto_filter.ref = f"A3:{get_column_letter(end_column)}{worksheet.max_row}"
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.paperSize = worksheet.PAPERSIZE_A4
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.print_title_rows = "3:3"
    worksheet.print_area = f"A1:{get_column_letter(end_column)}{worksheet.max_row}"
    worksheet.page_margins = PageMargins(left=0.25, right=0.25, top=0.45, bottom=0.45)
    worksheet.oddFooter.center.text = "Trang &P / &N"
    for column in range(1, end_column + 1):
        values = [str(worksheet.cell(row, column).value or "") for row in range(3, worksheet.max_row + 1)]
        worksheet.column_dimensions[get_column_letter(column)].width = min(
            max([len(value) for value in values] + [10]) + 2, 42,
        )


def invoice_delivery_statement_workbook(
    lines: list[dict[str, Any]],
    drafts: list[dict[str, Any]],
    tax_factor=None,
    *,
    contractor: str = "",
    scope_id: str = "",
    period_from: str = "",
    period_to: str = "",
) -> Workbook:
    """Build an issued-invoice statement and block every scope/total mismatch.

    ``tax_factor`` remains accepted for the historical call contract, but tax is
    parsed from each frozen invoice line so no external operational calculation
    can alter the statement.
    """
    del tax_factor
    if not drafts:
        raise InvoiceDeliveryStatementError("Không có hóa đơn đã phát hành để lập bảng kê")
    draft_map = {int(draft["id"]): draft for draft in drafts}
    if len(draft_map) != len(drafts):
        raise InvoiceDeliveryStatementError("Phạm vi có hóa đơn trùng ID")
    parties = {_plain(draft.get("contractor")).upper() for draft in drafts}
    if "" in parties or len(parties) != 1:
        raise InvoiceDeliveryStatementError(
            "Bảng kê chỉ được chứa đúng một nhà thầu",
            code="delivery_statement_contractor_conflict",
        )
    party = next(iter(parties))
    if contractor and party != _plain(contractor).upper():
        raise InvoiceDeliveryStatementError(
            "Nhà thầu trong hóa đơn không khớp phạm vi được chọn",
            code="delivery_statement_contractor_conflict",
        )
    for draft in drafts:
        if draft.get("status") != "issued":
            raise InvoiceDeliveryStatementError(
                "Bảng kê chính thức chỉ nhận hóa đơn đã phát hành",
                code="delivery_statement_unissued_invoice",
            )
        if not _plain(draft.get("issued_invoice_number")) or not _plain(draft.get("issued_invoice_date")):
            raise InvoiceDeliveryStatementError(
                "Hóa đơn thiếu số/ngày phát hành",
                code="delivery_statement_invoice_identity_missing",
            )

    all_dates = [
        _plain(item.get("work_date") or item.get("issued_invoice_date")) for item in lines
        if _plain(item.get("work_date") or item.get("issued_invoice_date"))
    ]
    safe_from = _plain(period_from) or (min(all_dates) if all_dates else "")
    safe_to = _plain(period_to) or (max(all_dates) if all_dates else "")
    try:
        from_date = datetime.strptime(safe_from, "%Y-%m-%d").date()
        to_date = datetime.strptime(safe_to, "%Y-%m-%d").date()
    except ValueError:
        raise InvoiceDeliveryStatementError("Khoảng ngày bảng tổng hợp giao nhận không hợp lệ") from None
    if from_date > to_date:
        raise InvoiceDeliveryStatementError("Từ ngày không được lớn hơn Đến ngày")

    visible_headers = [
        "STT", "Ngày", "Tên hàng", "ĐVT", "Số lượng", "Đơn giá", "Thành tiền",
        "Thuế suất", "Tiền thuế", "Thanh toán",
    ]
    audit_headers = [
        "Ngày HĐ", "Ký hiệu HĐ", "Số HĐ", "Bếp", "Mã hàng", "ID hóa đơn",
        "ID dòng", "ID đơn gốc", "Nguồn xác minh",
    ]
    detail_headers = visible_headers + audit_headers
    workbook = Workbook()
    detail = workbook.active
    detail.title = "Bảng kê giao hàng"
    computed = defaultdict(lambda: {"subtotal": 0, "tax": 0, "total": 0, "lines": 0})
    seen_line_ids = set()
    ordered_lines = sorted(lines, key=lambda row: (
        _plain(row.get("issued_invoice_date")), int(row.get("draft_id") or 0),
        _plain(row.get("work_date")), int(row.get("id") or row.get("line_id") or 0),
    ))
    for index, line in enumerate(ordered_lines, start=1):
        draft_id = int(line.get("draft_id") or 0)
        if draft_id not in draft_map:
            workbook.close()
            raise InvoiceDeliveryStatementError("Dòng giao hàng nằm ngoài phạm vi hóa đơn")
        line_id = int(line.get("id") or line.get("line_id") or 0)
        if line_id and line_id in seen_line_ids:
            workbook.close()
            raise InvoiceDeliveryStatementError("Dòng hóa đơn bị lặp trong bảng kê")
        if line_id:
            seen_line_ids.add(line_id)
        qty = _number(line.get("qty"), "Số lượng")
        unit_price = _number(line.get("unit_price"), "Đơn giá")
        amount = _vnd(line.get("amount"))
        if qty <= 0 or unit_price < 0 or amount < 0 or _vnd(qty, unit_price) != amount:
            workbook.close()
            raise InvoiceDeliveryStatementError(
                "Dòng bảng kê không thỏa Số lượng × Đơn giá = Thành tiền",
                code="delivery_statement_line_amount_mismatch",
            )
        vat_percent = _tax_percent(line.get("tax"))
        tax_amount = 0 if vat_percent <= 0 else _vnd(amount, vat_percent / 100)
        total_amount = amount + tax_amount
        computed[draft_id]["subtotal"] += amount
        computed[draft_id]["tax"] += tax_amount
        computed[draft_id]["total"] += total_amount
        computed[draft_id]["lines"] += 1
        draft = draft_map[draft_id]
        detail.append([
            index,
            _excel_date(line.get("work_date") or draft.get("issued_invoice_date"), "Ngày giao"),
            _excel_text(line.get("product_name")),
            _excel_text(line.get("unit")),
            qty,
            _vnd(unit_price),
            amount,
            _excel_text(line.get("tax")),
            tax_amount,
            total_amount,
            _excel_date(draft.get("issued_invoice_date"), "Ngày hóa đơn"),
            _excel_text(draft.get("issued_invoice_series")),
            _excel_text(draft.get("issued_invoice_number")),
            _excel_text(line.get("kitchen")),
            _excel_text(line.get("product_code")),
            draft_id,
            line_id or "",
            int(line.get("order_id") or 0) or "",
            _excel_text(draft.get("verification_source") or "local_issued_confirmation"),
        ])
    # Insert the approved customer-facing heading above the already validated
    # literal lines.  Audit columns remain in the workbook but are hidden and
    # excluded from the print area.
    detail.insert_rows(1, 10)
    first_draft = drafts[0]
    company = _plain(first_draft.get("company_name_snapshot")) or "CÔNG TY TNHH THỰC PHẨM THÀNH ĐẠT PHÁT"
    company_address = _plain(first_draft.get("company_address_snapshot"))
    company_tax = _plain(first_draft.get("company_tax_code_snapshot"))
    buyer = _plain(first_draft.get("buyer_name_snapshot")) or party
    buyer_address = _plain(first_draft.get("buyer_address_snapshot"))
    buyer_tax = _plain(first_draft.get("buyer_tax_code_snapshot"))
    same_month = from_date.year == to_date.year and from_date.month == to_date.month
    title = (
        f"BẢNG TỔNG HỢP GIAO NHẬN THÁNG {to_date.month:02d} NĂM {to_date.year}"
        if same_month else "BẢNG TỔNG HỢP GIAO NHẬN"
    )

    def merge_write(cell_range: str, value: str, *, size=11, bold=False, italic=False,
                    horizontal="left", color="000000"):
        detail.merge_cells(cell_range)
        cell = detail[cell_range.split(":", 1)[0]]
        cell.value = _excel_text(value)
        cell.font = Font(name="Times New Roman", size=size, bold=bold, italic=italic, color=color)
        cell.alignment = Alignment(horizontal=horizontal, vertical="center", wrap_text=True)

    merge_write("A1:J1", company.upper(), size=12, bold=True)
    merge_write("A2:J2", f"Địa chỉ: {company_address}" if company_address else "", italic=True)
    merge_write("A3:J3", f"ĐT: {COMPANY_PHONE}", italic=True)
    merge_write("A4:J4", f"MST: {company_tax}" if company_tax else "", italic=True)
    merge_write("A5:J5", title, size=16, bold=True, horizontal="center")
    merge_write(
        "A6:J6", f"(Từ ngày {from_date.strftime('%d/%m/%Y')} – {to_date.strftime('%d/%m/%Y')})",
        size=11, italic=True, horizontal="center",
    )
    merge_write("A7:J7", f"Khách hàng: {buyer.upper()}", size=11, bold=True)
    merge_write("A8:J8", f"Địa chỉ: {buyer_address}" if buyer_address else "", size=11, italic=True)
    merge_write("A9:J9", f"MST: {buyer_tax}" if buyer_tax else "", size=11, italic=True)
    for column, header in enumerate(detail_headers, 1):
        cell = detail.cell(10, column, header)
        cell.font = Font(name="Times New Roman", size=10, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=Side(style="thin"), right=Side(style="thin"),
                             top=Side(style="thin"), bottom=Side(style="thin"))
    detail.row_dimensions[5].height = 27
    detail.row_dimensions[6].height = 20
    detail.row_dimensions[10].height = 34
    data_start = 11
    data_end = 10 + len(ordered_lines)
    for row in detail.iter_rows(min_row=data_start, max_row=data_end, max_col=len(detail_headers)):
        for cell in row:
            cell.font = Font(name="Times New Roman", size=10)
            cell.alignment = Alignment(
                horizontal="right" if cell.column in (5, 6, 7, 9, 10) else "left",
                vertical="center", wrap_text=True,
            )
            cell.border = Border(left=Side(style="thin"), right=Side(style="thin"),
                                 top=Side(style="thin"), bottom=Side(style="thin"))
        row[0].alignment = Alignment(horizontal="center", vertical="center")
        row[1].alignment = Alignment(horizontal="center", vertical="center")
        row[3].alignment = Alignment(horizontal="center", vertical="center")
        row[7].alignment = Alignment(horizontal="center", vertical="center")
        for column in (6, 7, 9, 10):
            detail.cell(row[0].row, column).number_format = '#,##0;[Red]-#,##0;"-"'
        detail.cell(row[0].row, 5).number_format = "#,##0.######"
        detail.cell(row[0].row, 2).number_format = "dd/mm/yyyy"
        detail.cell(row[0].row, 11).number_format = "dd/mm/yyyy"

    total_row = data_end + 1
    detail.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=4)
    detail.cell(total_row, 1, "Tổng cộng:")
    detail.cell(total_row, 5, quantity_cell([{'quantity':detail.cell(r,5).value, 'unit':detail.cell(r,4).value}
        for r in range(data_start,data_end+1)], 'quantity'))
    detail.cell(total_row, 5).number_format = '#,##0.######'
    detail.cell(total_row, 5).alignment = Alignment(horizontal='right', vertical='center', wrap_text=True)
    detail.row_dimensions[total_row].height = max(24, 16 * math.ceil(len(str(detail.cell(total_row,5).value)) / 10))
    detail.cell(total_row, 1).alignment = Alignment(horizontal="center", vertical="center")
    detail.cell(total_row, 7, sum(values["subtotal"] for values in computed.values()))
    detail.cell(total_row, 9, sum(values["tax"] for values in computed.values()))
    detail.cell(total_row, 10, sum(values["total"] for values in computed.values()))
    for column in range(1, 11):
        cell = detail.cell(total_row, column)
        cell.font = Font(name="Times New Roman", size=11, bold=True)
        cell.border = Border(left=Side(style="thin"), right=Side(style="thin"),
                             top=Side(style="thin"), bottom=Side(style="thin"))
        if column in (7, 9, 10):
            cell.number_format = '#,##0;[Red]-#,##0;"-"'
    words_row = total_row + 1
    merge_write(
        f"A{words_row}:J{words_row}",
        "Bằng chữ:    " + number_to_vietnamese(detail.cell(total_row, 10).value) + "./.",
        size=11, bold=True, italic=True,
    )
    date_row = words_row + 1
    merge_write(
        f"G{date_row}:J{date_row}",
        f"Hải Phòng, ngày {to_date.day:02d} tháng {to_date.month:02d} năm {to_date.year}",
        size=11, italic=True, horizontal="center",
    )
    signature_row = date_row + 1
    merge_write(f"A{signature_row}:E{signature_row}", "ĐẠI DIỆN BÊN MUA", size=11, bold=True, horizontal="center")
    merge_write(f"F{signature_row}:J{signature_row}", "ĐẠI DIỆN BÊN BÁN", size=11, bold=True, horizontal="center")
    merge_write(f"A{signature_row + 1}:E{signature_row + 1}", "(Ký, ghi rõ họ và tên)", size=11, italic=True, horizontal="center")
    merge_write(f"F{signature_row + 1}:J{signature_row + 1}", "(Ký, ghi rõ họ và tên)", size=11, italic=True, horizontal="center")

    widths = (7, 13, 31, 9, 11, 15, 17, 11, 16, 18)
    for column, width in enumerate(widths, 1):
        detail.column_dimensions[get_column_letter(column)].width = width
    for column in range(11, len(detail_headers) + 1):
        detail.column_dimensions[get_column_letter(column)].hidden = True
    detail.freeze_panes = "A11"
    detail.auto_filter.ref = f"A10:J{data_end}"
    detail.print_title_rows = "10:10"
    detail.print_area = f"A1:J{signature_row + 5}"
    detail.sheet_properties.pageSetUpPr.fitToPage = True
    detail.page_setup.orientation = "portrait"
    detail.page_setup.paperSize = detail.PAPERSIZE_A4
    detail.page_setup.fitToWidth = 1
    detail.page_setup.fitToHeight = 0
    detail.page_setup.scale = None
    detail.page_margins = PageMargins(
        left=0.25, right=0.25, top=0.4, bottom=0.75, header=0.15, footer=0.2,
    )
    detail.print_options.horizontalCentered = True
    detail.oddFooter.center.text = "Trang &P / &N"
    detail.sheet_view.showGridLines = False

    reconcile_headers = [
        "Ngày HĐ", "Ký hiệu", "Số HĐ", "Tiền HĐ trước thuế", "Tiền bảng kê trước thuế",
        "Thuế HĐ", "Thuế bảng kê", "Tổng HĐ", "Tổng bảng kê", "Chênh lệch",
        "Nguồn xác minh", "ID hóa đơn nguồn", "ID draft",
    ]
    reconcile = workbook.create_sheet("Đối chiếu hóa đơn")
    reconcile.append(reconcile_headers)
    mismatches = []
    for draft in sorted(drafts, key=lambda row: (
        _plain(row.get("issued_invoice_date")), _plain(row.get("issued_invoice_series")),
        _plain(row.get("issued_invoice_number")), int(row["id"]),
    )):
        values = computed[int(draft["id"])]
        invoice_subtotal = _vnd(draft["subtotal"])
        invoice_tax = _vnd(draft["tax_amount"])
        invoice_total = _vnd(draft["total_amount"])
        difference = invoice_total - values["total"]
        if values["lines"] == 0 or any(abs(value) > 1 for value in (
            invoice_subtotal - values["subtotal"], invoice_tax - values["tax"], difference,
        )):
            mismatches.append(_plain(draft.get("issued_invoice_number")) or str(draft["id"]))
        reconcile.append([
            _excel_date(draft.get("issued_invoice_date"), "Ngày hóa đơn"),
            _excel_text(draft.get("issued_invoice_series")),
            _excel_text(draft.get("issued_invoice_number")),
            invoice_subtotal,
            values["subtotal"],
            invoice_tax,
            values["tax"],
            invoice_total,
            values["total"],
            difference,
            _excel_text(draft.get("verification_source") or "local_issued_confirmation"),
            draft.get("source_invoice_id") or "",
            int(draft["id"]),
        ])
    if mismatches:
        workbook.close()
        raise InvoiceDeliveryStatementError(
            "Tổng hóa đơn lệch bảng kê giao hàng: " + ", ".join(mismatches),
            code="delivery_statement_total_mismatch",
        )
    _style_sheet(
        reconcile,
        "ĐỐI CHIẾU HÓA ĐƠN ĐỎ VÀ BẢNG KÊ – " + _excel_text(party),
        "Ngày/ký hiệu/số và ba tổng phải khớp; chênh lệch cho phép tối đa 1 VNĐ",
        reconcile_headers,
        money_columns=(4, 5, 6, 7, 8, 9, 10),
    )
    for row in range(4, reconcile.max_row + 1):
        reconcile.cell(row, 1).number_format = "dd/mm/yyyy"
    for column in (12, 13):
        reconcile.column_dimensions[get_column_letter(column)].hidden = True
    # This sheet is retained as machine/audit evidence, like the hidden proof
    # sheet in the payment request.  It is not part of the customer-facing
    # official form and must not be printed accidentally.
    reconcile.sheet_state = "hidden"
    workbook.properties.title = "Bảng kê giao hàng theo hóa đơn đỏ"
    workbook.properties.subject = f"Nhà thầu {party} · scope {scope_id or 'không cấp'}"
    workbook.calculation.fullCalcOnLoad = False
    workbook.calculation.forceFullCalc = False
    workbook.calculation.calcMode = "manual"
    white_print_style(workbook)
    return workbook


__all__ = ["InvoiceDeliveryStatementError", "invoice_delivery_statement_workbook"]
