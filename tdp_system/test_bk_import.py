import io
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

try:
    from . import inventory_export, server
    from .bk_import import BK_IMPORT_COLUMNS, BK_IMPORT_SOURCE_TYPE
except ImportError:  # pragma: no cover - direct invocation
    import inventory_export
    import server
    from bk_import import BK_IMPORT_COLUMNS, BK_IMPORT_SOURCE_TYPE


NOW = "2026-09-03T08:00:00"


class BKImportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        cls.original_db_path = server.DB_PATH
        cls.original_data_dir = server.DATA_DIR
        cls.original_master_source = server.MASTER_SOURCE
        cls.original_opening_source = server.OPENING_TEMPLATE_SOURCE
        cls.original_testing = server.app.config.get("TESTING", False)
        server.DATA_DIR = Path(cls.temp.name) / "data"
        server.DATA_DIR.mkdir(parents=True, exist_ok=True)
        server.DB_PATH = server.DATA_DIR / "bk-import.sqlite3"
        server.MASTER_SOURCE = Path(cls.temp.name) / "master-disabled.xlsx"
        server.OPENING_TEMPLATE_SOURCE = Path(cls.temp.name) / "opening-disabled.xlsx"
        server.app.config["TESTING"] = True
        server.init_database()
        cls.client = server.app.test_client()

    @classmethod
    def tearDownClass(cls):
        server.DB_PATH = cls.original_db_path
        server.DATA_DIR = cls.original_data_dir
        server.MASTER_SOURCE = cls.original_master_source
        server.OPENING_TEMPLATE_SOURCE = cls.original_opening_source
        server.app.config["TESTING"] = cls.original_testing
        cls.temp.cleanup()

    def setUp(self):
        with server.db() as conn:
            conn.execute("DROP TRIGGER IF EXISTS fail_bk_second_line")
            conn.execute("DELETE FROM invoice_inventory_ledger")
            conn.execute("DELETE FROM invoice_inventory_confirmations")
            conn.execute("DELETE FROM bk_import_lines")
            conn.execute("DELETE FROM bk_import_documents")
            conn.execute("DELETE FROM audit_log")
            conn.execute("DELETE FROM purchase_workbook_lines")
            conn.execute("DELETE FROM orders")
            conn.execute("DELETE FROM batches")
            conn.execute("DELETE FROM inventory_transactions")
            conn.execute("DELETE FROM products")
            conn.execute("DELETE FROM suppliers")
            conn.execute("DELETE FROM kitchens")
            conn.execute("DELETE FROM contractors")
            conn.execute(
                "INSERT INTO contractors(code,name,price_group,pricing_mode) "
                "VALUES('C1','Nhà thầu C1','C1','group')"
            )
            conn.execute(
                "INSERT INTO kitchens(code,contractor,name) VALUES('K1','C1','Bếp K1')"
            )
            conn.execute("INSERT INTO suppliers(code,name) VALUES('BK-S1','NCC BK')")
            conn.executemany(
                """INSERT INTO products(
                       code,name,unit,tax,supplier,buy_price,purchase_list,seller,cccd
                   ) VALUES(?,?,?,?,?,?,1,'','')""",
                (
                    ("BK-P1", "Hàng BK", "kg", "8%", "BK-S1", 100),
                    ("BK-P2", "Hàng BK 2", "chai", "8%", "BK-S1", 50),
                ),
            )
            conn.execute(
                "INSERT INTO settings(key,value) VALUES('purchase_rate','0.95') "
                "ON CONFLICT(key) DO UPDATE SET value='0.95'"
            )

    @staticmethod
    def valid_row(*, reference="BK-2026-0001", line=1, code="BK-P1"):
        if code == "BK-P2":
            return [
                "2026-08-31", "", reference, line, code, "Hàng BK 2", "chai",
                3, 50, 150, "BK-S1", "mua không hóa đơn",
            ]
        return [
            "2026-08-31", "", reference, line, code, "Hàng BK", "kg",
            2, 100, 200, "BK-S1", "mua không hóa đơn",
        ]

    def template_with_rows(self, rows) -> bytes:
        response = self.client.get("/api/bk-import/template")
        self.assertEqual(response.status_code, 200, response.status)
        workbook = load_workbook(io.BytesIO(response.data), data_only=False, keep_links=True)
        try:
            sheet = workbook["BK_IMPORT"]
            for values in rows:
                sheet.append(values)
            output = io.BytesIO()
            workbook.save(output)
            return output.getvalue()
        finally:
            workbook.close()

    def preview(self, payload: bytes, filename="bk.xlsx"):
        return self.client.post(
            "/api/bk-import/preview",
            data={"file": (io.BytesIO(payload), filename)},
            content_type="multipart/form-data",
        )

    def confirm(self, preview):
        return self.client.post(
            "/api/bk-import/confirm",
            json={
                "confirmed": True,
                "token": preview["token"],
                "previewId": preview["previewId"],
            },
        )

    @staticmethod
    def counts():
        with server.db() as conn:
            return {
                name: conn.execute(f"SELECT COUNT(*) FROM {name}").fetchone()[0]
                for name in (
                    "bk_import_documents", "bk_import_lines",
                    "invoice_inventory_confirmations", "invoice_inventory_ledger",
                    "inventory_transactions", "audit_log",
                )
            }

    def post_one(self):
        payload = self.template_with_rows([self.valid_row()])
        preview_response = self.preview(payload)
        self.assertEqual(preview_response.status_code, 200, preview_response.get_data(as_text=True))
        preview = preview_response.get_json()
        response = self.confirm(preview)
        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        return payload, preview, response.get_json()

    def test_official_template_is_formula_free_fixed_source_and_single_sheet(self):
        response = self.client.get("/api/bk-import/template")
        self.assertEqual(response.status_code, 200, response.status)
        self.assertIn(
            "MAU_NHAP_BK_HANG_MUA_KHONG_HOA_DON.xlsx",
            response.headers["Content-Disposition"],
        )
        workbook = load_workbook(io.BytesIO(response.data), data_only=False, keep_links=True)
        try:
            self.assertEqual(workbook.sheetnames, ["BK_IMPORT"])
            sheet = workbook["BK_IMPORT"]
            self.assertIn("HÀNG MUA VÀO KHÔNG CÓ HÓA ĐƠN", sheet["A1"].value)
            self.assertIn(BK_IMPORT_SOURCE_TYPE, sheet["A2"].value)
            self.assertIn("95%", sheet["A2"].value)
            self.assertEqual(
                [sheet.cell(3, column).value for column in range(1, len(BK_IMPORT_COLUMNS) + 1)],
                [label for _field, label in BK_IMPORT_COLUMNS],
            )
            self.assertFalse(workbook._external_links)
            self.assertFalse(any(
                cell.data_type == "f" for row in sheet.iter_rows() for cell in row
            ))
            self.assertEqual("landscape", sheet.page_setup.orientation)
            self.assertEqual("9", str(sheet.page_setup.paperSize))
            self.assertEqual(1, sheet.page_setup.fitToWidth)
            self.assertEqual(0, sheet.page_setup.fitToHeight)
            self.assertEqual("'BK_IMPORT'!$A$1:$L$3", str(sheet.print_area))
            self.assertEqual("$1:$3", str(sheet.print_title_rows))
        finally:
            workbook.close()

    def test_approved_batch_template_prefills_bk_rows_at_95_percent_sales_price(self):
        with server.db() as conn:
            batch_id = conn.execute(
                "INSERT INTO batches(work_date,source_name,status,created_at,approved_at) "
                "VALUES('2026-09-03','orders.xlsx','approved',?,?)", (NOW, NOW),
            ).lastrowid
            order_id = conn.execute(
                """INSERT INTO orders(
                       batch_id,work_date,contractor,kitchen,product_code,product_name,qty,
                       actual_received,actual_delivered,unit,supplier,buy_price,sell_price,tax,
                       purchase_list,note,source_row,errors,warnings,updated_at
                   ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (
                    batch_id, "2026-09-03", "C1", "K1", "BK-P1", "Hàng BK", 2,
                    2, 2, "kg", "BK-S1", 111, 200, "8%", 1, "ghi chú", 27,
                    "[]", "[]", NOW,
                ),
            ).lastrowid
        response = self.client.get(f"/api/bk-import/template?batch_id={batch_id}")
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(io.BytesIO(response.data), data_only=False)
        try:
            sheet = workbook["BK_IMPORT"]
            values = [sheet.cell(4, column).value for column in range(1, 13)]
            self.assertEqual(values[0].strftime("%d/%m/%Y"), "03/09/2026")
            self.assertEqual(sheet["A4"].number_format, "dd/mm/yyyy")
            self.assertEqual(values[1], BK_IMPORT_SOURCE_TYPE)
            self.assertEqual(values[2], f"TDP-BATCH-{batch_id}")
            self.assertEqual(values[3], order_id)
            self.assertEqual((values[4], values[7]), ("BK-P1", 2))
            self.assertEqual((values[8], values[9]), (190, 380))
            self.assertIn("95.00% giá bán", values[11])
        finally:
            workbook.close()

    def test_batch_template_rejects_draft_and_empty_bk_scope(self):
        with server.db() as conn:
            draft_id = conn.execute(
                "INSERT INTO batches(work_date,source_name,status,created_at) "
                "VALUES('2026-09-03','draft.xlsx','draft',?)", (NOW,),
            ).lastrowid
            empty_id = conn.execute(
                "INSERT INTO batches(work_date,source_name,status,created_at,approved_at) "
                "VALUES('2026-09-03','empty.xlsx','approved',?,?)", (NOW, NOW),
            ).lastrowid
        draft = self.client.get(f"/api/bk-import/template?batch_id={draft_id}")
        empty = self.client.get(f"/api/bk-import/template?batch_id={empty_id}")
        self.assertEqual((draft.status_code, draft.get_json()["code"]), (409, "batch_not_approved"))
        self.assertEqual((empty.status_code, empty.get_json()["code"]), (409, "batch_has_no_bk_rows"))

    def test_batch_template_never_substitutes_buy_cost_for_missing_sales_price(self):
        with server.db() as conn:
            batch_id = conn.execute(
                "INSERT INTO batches(work_date,source_name,status,created_at,approved_at) "
                "VALUES('2026-09-03','missing-price.xlsx','approved',?,?)", (NOW, NOW),
            ).lastrowid
            conn.execute(
                """INSERT INTO orders(
                       batch_id,work_date,contractor,kitchen,product_code,product_name,qty,
                       actual_received,actual_delivered,unit,supplier,buy_price,sell_price,tax,
                       purchase_list,note,source_row,errors,warnings,updated_at
                   ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (
                    batch_id, "2026-09-03", "C1", "K1", "BK-P1", "Hàng BK", 2,
                    2, 2, "kg", "BK-S1", 999, 0, "8%", 1, "", 31,
                    "[]", "[]", NOW,
                ),
            )
        response = self.client.get(f"/api/bk-import/template?batch_id={batch_id}")
        self.assertEqual(
            (response.status_code, response.get_json()["code"]),
            (409, "bk_sales_price_missing"),
        )
        self.assertIn("nguồn 31", response.get_json()["error"])
        self.assertNotIn("999", response.get_json()["error"])

    def test_preview_is_stable_read_only_and_source_is_not_msmi(self):
        payload = self.template_with_rows([self.valid_row()])
        before = self.counts()
        first = self.preview(payload).get_json()
        second = self.preview(payload).get_json()
        self.assertEqual(
            (first["sourceHash"], first["contentHash"], first["previewId"], first["rows"][0]["sourceKey"]),
            (second["sourceHash"], second["contentHash"], second["previewId"], second["rows"][0]["sourceKey"]),
        )
        self.assertTrue(first["canConfirm"])
        self.assertTrue(first["previewOnly"])
        self.assertFalse(first["writesInventory"])
        self.assertEqual(first["sourcePolicy"], BK_IMPORT_SOURCE_TYPE)
        self.assertEqual(first["rows"][0]["sourceType"], BK_IMPORT_SOURCE_TYPE)
        self.assertEqual(self.counts(), before)

        bad = self.valid_row(reference="BK-MSMI")
        bad[1] = "mSMI"
        rejected = self.preview(self.template_with_rows([bad])).get_json()
        self.assertFalse(rejected["canConfirm"])
        self.assertTrue(any("không dùng mSMI" in error for error in rejected["rows"][0]["errors"]))

    def test_validation_blocks_formula_duplicate_unknown_name_unit_cost_and_amount(self):
        good = self.valid_row()
        duplicate = list(good)
        formula = self.valid_row(reference="BK-FORMULA", line=2)
        formula[4] = "UNKNOWN"
        formula[7] = "=1+1"
        bad_snapshot = self.valid_row(reference="BK-BAD", line=3)
        bad_snapshot[5] = "Sai tên"
        bad_snapshot[6] = "thùng"
        bad_snapshot[8] = -100
        bad_snapshot[9] = 999
        response = self.preview(self.template_with_rows([good, duplicate, formula, bad_snapshot]))
        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        preview = response.get_json()
        self.assertFalse(preview["canConfirm"])
        self.assertEqual(preview["counts"]["duplicateGroups"], 1)
        self.assertEqual(preview["counts"]["errorRows"], 4)
        self.assertTrue(any("công thức" in value for value in preview["rows"][2]["errors"]))
        self.assertTrue(any("chưa có trong danh mục" in value for value in preview["rows"][2]["errors"]))
        self.assertTrue(any("Tên hàng không khớp" in value for value in preview["rows"][3]["errors"]))
        self.assertTrue(any("ĐVT không khớp" in value for value in preview["rows"][3]["errors"]))
        self.assertTrue(any("Đơn giá vốn phải lớn hơn 0" in value for value in preview["rows"][3]["errors"]))

    def test_confirm_is_atomic_canonical_audited_and_unlocks_readiness(self):
        _payload, _preview, result = self.post_one()
        self.assertFalse(result["idempotent"])
        self.assertEqual((result["newInventoryLines"], result["inventoryLines"]), (1, 1))
        counts = self.counts()
        self.assertEqual(counts["bk_import_documents"], 1)
        self.assertEqual(counts["bk_import_lines"], 1)
        self.assertEqual(counts["invoice_inventory_confirmations"], 1)
        self.assertEqual(counts["invoice_inventory_ledger"], 1)
        self.assertEqual(counts["inventory_transactions"], 0)
        self.assertEqual(counts["audit_log"], 1)
        with server.db() as conn:
            event = conn.execute("SELECT * FROM invoice_inventory_ledger").fetchone()
            self.assertEqual(
                (event["direction"], event["event_type"], event["source_invoice_table"]),
                ("input", "POST", "bk_import_documents"),
            )
            self.assertIsNone(event["mapping_revision_id"])
            audit = conn.execute("SELECT event_type,metadata_json FROM audit_log").fetchone()
            self.assertEqual(audit["event_type"], "bk_import.inventory_post")
            self.assertEqual(json.loads(audit["metadata_json"])["source_kind"], BK_IMPORT_SOURCE_TYPE)
            batch_id = conn.execute(
                "INSERT INTO batches(work_date,source_name,status,created_at) "
                "VALUES('2026-09-01','out.xlsx','approved',?)", (NOW,),
            ).lastrowid
            conn.execute(
                """INSERT INTO orders(
                       batch_id,work_date,contractor,kitchen,product_code,product_name,qty,actual_received,
                       actual_delivered,unit,buy_price,sell_price,tax,purchase_list,
                       errors,warnings,updated_at
                   ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (batch_id, "2026-09-01", "C1", "K1", "BK-P1", "Hàng BK", 3, 3, 3, "kg",
                 100, 150, "8%", 0, "[]", "[]", NOW),
            )
        readiness = self.client.get(f"/api/outgoing-invoices/readiness/{batch_id}")
        self.assertEqual(readiness.status_code, 200, readiness.get_data(as_text=True))
        self.assertEqual(
            (readiness.get_json()["invoiceable_qty"], readiness.get_json()["pending_qty"]),
            (2, 1),
        )

    def test_same_token_and_same_content_repreview_are_idempotent(self):
        payload = self.template_with_rows([self.valid_row()])
        preview = self.preview(payload).get_json()
        first = self.confirm(preview)
        second = self.confirm(preview)
        self.assertEqual((first.status_code, second.status_code), (200, 200))
        self.assertFalse(first.get_json()["idempotent"])
        self.assertTrue(second.get_json()["idempotent"])

        repeated_preview = self.preview(payload).get_json()
        self.assertTrue(repeated_preview["alreadyPosted"])
        self.assertTrue(repeated_preview["canConfirm"])
        third = self.confirm(repeated_preview)
        self.assertEqual(third.status_code, 200, third.get_data(as_text=True))
        self.assertTrue(third.get_json()["idempotent"])
        self.assertEqual(self.counts()["invoice_inventory_ledger"], 1)

    def test_same_rows_in_different_order_keep_one_canonical_document(self):
        first_row = self.valid_row()
        second_row = self.valid_row(reference="BK-2026-0002", line=2, code="BK-P2")
        first_payload = self.template_with_rows([first_row, second_row])
        first_preview = self.preview(first_payload).get_json()
        first = self.confirm(first_preview)
        self.assertEqual(first.status_code, 200, first.get_data(as_text=True))

        reordered_payload = self.template_with_rows([second_row, first_row])
        reordered_preview = self.preview(reordered_payload).get_json()
        self.assertEqual(reordered_preview["contentHash"], first_preview["contentHash"])
        repeated = self.confirm(reordered_preview)
        self.assertEqual(repeated.status_code, 200, repeated.get_data(as_text=True))
        self.assertTrue(repeated.get_json()["idempotent"])
        self.assertEqual(self.counts()["bk_import_documents"], 1)
        self.assertEqual(self.counts()["invoice_inventory_ledger"], 2)

    def test_confirm_blocks_bk_backdated_before_a_newer_opening_snapshot(self):
        preview = self.preview(self.template_with_rows([self.valid_row()])).get_json()
        with server.db() as conn:
            conn.execute(
                """INSERT INTO inventory_transactions(
                       txn_date,product_code,qty_in,qty_out,unit_cost,source_type,source_id,
                       source_line,status,note,created_at,updated_at
                   ) VALUES('2026-09-01','BK-P1',10,0,100,'OPENING','2026-09',
                            'BK-P1','posted','new opening',?,?)""",
                (NOW, NOW),
            )
        response = self.confirm(preview)
        self.assertEqual(
            (response.status_code, response.get_json()["code"]),
            (409, "backdated_before_opening_snapshot"),
        )
        self.assertEqual(self.counts()["bk_import_documents"], 0)
        self.assertEqual(self.counts()["invoice_inventory_ledger"], 0)

    def test_preview_shows_backdated_opening_conflict_before_confirmation(self):
        with server.db() as conn:
            conn.execute(
                """INSERT INTO inventory_transactions(
                       txn_date,product_code,qty_in,qty_out,unit_cost,source_type,source_id,
                       source_line,status,note,created_at,updated_at
                   ) VALUES('2026-09-01','BK-P1',10,0,100,'OPENING','2026-09',
                            'BK-P1','posted','new opening',?,?)""",
                (NOW, NOW),
            )
        preview = self.preview(self.template_with_rows([self.valid_row()])).get_json()
        self.assertFalse(preview["canConfirm"])
        self.assertTrue(any(
            "trước một kỳ tồn đầu" in message for message in preview["rows"][0]["errors"]
        ))

    def test_same_business_source_with_changed_snapshot_is_rejected(self):
        self.post_one()
        changed = self.valid_row()
        changed[7] = 3
        changed[9] = 300
        preview = self.preview(self.template_with_rows([changed])).get_json()
        self.assertFalse(preview["canConfirm"])
        self.assertTrue(any("không được ghi đè" in value for value in preview["rows"][0]["errors"]))
        blocked = self.confirm(preview)
        self.assertEqual((blocked.status_code, blocked.get_json()["code"]), (400, "bk_rows_invalid"))
        self.assertEqual(self.counts()["invoice_inventory_ledger"], 1)

    def test_catalog_change_after_preview_makes_token_stale(self):
        preview = self.preview(self.template_with_rows([self.valid_row()])).get_json()
        with server.db() as conn:
            conn.execute("UPDATE products SET unit='tấn' WHERE code='BK-P1'")
        response = self.confirm(preview)
        self.assertEqual((response.status_code, response.get_json()["code"]), (409, "bk_preview_stale"))
        self.assertEqual(self.counts()["invoice_inventory_ledger"], 0)

    def test_second_line_failure_rolls_back_document_confirmation_ledger_and_audit(self):
        rows = [
            self.valid_row(),
            self.valid_row(reference="BK-2026-0002", line=2, code="BK-P2"),
        ]
        preview = self.preview(self.template_with_rows(rows)).get_json()
        self.assertTrue(preview["canConfirm"])
        with server.db() as conn:
            conn.execute(
                """CREATE TRIGGER fail_bk_second_line BEFORE INSERT ON invoice_inventory_ledger
                   WHEN NEW.source_invoice_table='bk_import_documents'
                    AND NEW.source_line_index=2
                   BEGIN SELECT RAISE(ABORT,'forced'); END"""
            )
        response = self.confirm(preview)
        self.assertEqual((response.status_code, response.get_json()["code"]), (409, "bk_source_conflict"))
        counts = self.counts()
        self.assertEqual(counts["bk_import_documents"], 0)
        self.assertEqual(counts["bk_import_lines"], 0)
        self.assertEqual(counts["invoice_inventory_confirmations"], 0)
        self.assertEqual(counts["invoice_inventory_ledger"], 0)
        self.assertEqual(counts["audit_log"], 0)

    def test_reversal_requires_reason_is_atomic_audited_and_idempotent(self):
        _payload, _preview, posted = self.post_one()
        document_id = posted["documentId"]
        missing = self.client.post(
            f"/api/bk-import/documents/{document_id}/reversal",
            json={"confirmed": True, "reversalDate": "2026-09-03"},
        )
        self.assertEqual((missing.status_code, missing.get_json()["code"]), (400, "reversal_reason_required"))
        body = {"confirmed": True, "reversalDate": "2026-09-03", "reason": "Nhập nhầm bộ BK"}
        first = self.client.post(f"/api/bk-import/documents/{document_id}/reversal", json=body)
        second = self.client.post(f"/api/bk-import/documents/{document_id}/reversal", json=body)
        self.assertEqual((first.status_code, second.status_code), (200, 200))
        self.assertFalse(first.get_json()["idempotent"])
        self.assertTrue(second.get_json()["idempotent"])
        conflicting_repeat = self.client.post(
            f"/api/bk-import/documents/{document_id}/reversal",
            json={**body, "reason": "Lý do khác"},
        )
        self.assertEqual(
            (conflicting_repeat.status_code, conflicting_repeat.get_json()["code"]),
            (409, "bk_reversal_conflict"),
        )
        with server.db() as conn:
            document = conn.execute("SELECT * FROM bk_import_documents").fetchone()
            events = conn.execute(
                "SELECT event_type,qty_delta,reverses_event_key FROM invoice_inventory_ledger ORDER BY id"
            ).fetchall()
            self.assertEqual(document["status"], "reversed")
            self.assertEqual(document["reversal_reason"], "Nhập nhầm bộ BK")
            self.assertEqual([(row["event_type"], row["qty_delta"]) for row in events], [
                ("POST", 2), ("REVERSAL", -2),
            ])
            self.assertEqual(events[1]["reverses_event_key"], conn.execute(
                "SELECT event_key FROM invoice_inventory_ledger WHERE event_type='POST'"
            ).fetchone()[0])
            audit_types = [row[0] for row in conn.execute("SELECT event_type FROM audit_log ORDER BY id")]
            self.assertEqual(audit_types, ["bk_import.inventory_post", "bk_import.inventory_reversal"])
        stock = self.client.get("/api/invoice-inventory?as_of=2026-09-03&include_zero=1").get_json()
        item = next(row for row in stock["items"] if row["product_code"] == "BK-P1")
        self.assertEqual(item["closing_qty"], 0)

    def test_reversal_cannot_make_any_later_stock_point_negative(self):
        _payload, _preview, posted = self.post_one()
        with server.db() as conn:
            confirmation_id = conn.execute(
                """INSERT INTO invoice_inventory_confirmations(
                       confirmation_key,direction,source_invoice_table,source_invoice_id,
                       action,confirmed,note,created_at
                   ) VALUES('OUT-BK-TEST','output','outgoing_source_invoices',999,
                            'post',1,'test output',?)""", (NOW,),
            ).lastrowid
            conn.execute(
                """INSERT INTO invoice_inventory_ledger(
                       event_key,direction,event_type,source_invoice_table,source_invoice_id,
                       source_line_id,source_line_index,product_code,txn_date,qty_delta,unit_cost,
                       mapping_revision_id,confirmation_id,reverses_event_key,status,created_at
                   ) VALUES('OUT-BK-EVENT','output','POST','outgoing_source_invoices',999,
                            999,1,'BK-P1','2026-09-01',-1,100,NULL,?,'','posted',?)""",
                (confirmation_id, NOW),
            )
        response = self.client.post(
            f"/api/bk-import/documents/{posted['documentId']}/reversal",
            json={"confirmed": True, "reversalDate": "2026-09-03", "reason": "Thử hoàn tác"},
        )
        self.assertEqual(
            (response.status_code, response.get_json()["code"]),
            (409, "bk_reversal_negative_stock"),
        )
        with server.db() as conn:
            document = conn.execute("SELECT status FROM bk_import_documents").fetchone()
            reversals = conn.execute(
                "SELECT COUNT(*) FROM invoice_inventory_ledger WHERE event_type='REVERSAL'"
            ).fetchone()[0]
            reversal_confirmations = conn.execute(
                "SELECT COUNT(*) FROM invoice_inventory_confirmations WHERE action='reversal'"
            ).fetchone()[0]
        self.assertEqual(document["status"], "posted")
        self.assertEqual((reversals, reversal_confirmations), (0, 0))

    def test_reversal_cannot_be_hidden_behind_a_newer_opening_snapshot(self):
        _payload, _preview, posted = self.post_one()
        with server.db() as conn:
            conn.execute(
                """INSERT INTO inventory_transactions(
                       txn_date,product_code,qty_in,qty_out,unit_cost,source_type,source_id,
                       source_line,status,note,created_at,updated_at
                   ) VALUES('2026-09-05','BK-P1',2,0,100,'OPENING','2026-09',
                            'BK-P1','posted','new opening',?,?)""",
                (NOW, NOW),
            )
        response = self.client.post(
            f"/api/bk-import/documents/{posted['documentId']}/reversal",
            json={"confirmed": True, "reversalDate": "2026-09-03", "reason": "Hủy BK"},
        )
        self.assertEqual(
            (response.status_code, response.get_json()["code"]),
            (409, "bk_reversal_before_opening_snapshot"),
        )
        with server.db() as conn:
            self.assertEqual(
                conn.execute("SELECT status FROM bk_import_documents").fetchone()[0], "posted"
            )
            self.assertEqual(
                conn.execute(
                    "SELECT COUNT(*) FROM invoice_inventory_ledger WHERE event_type='REVERSAL'"
                ).fetchone()[0],
                0,
            )

    def test_nxt_export_traces_bk_post_and_input_reversal(self):
        _payload, _preview, posted = self.post_one()
        with server.db() as conn:
            model = inventory_export.collect_inventory_export_model(
                conn, date_from="2026-08-01", date_to="2026-08-31",
            )
        self.assertEqual(len(model["input_rows"]), 1)
        row = model["input_rows"][0]
        self.assertEqual((row["invoice_series"], row["invoice_number"]), ("BK", "BK-2026-0001"))
        self.assertEqual((row["party_name"], row["movement_label"]), ("BK-S1", "Nhập"))

        reversed_response = self.client.post(
            f"/api/bk-import/documents/{posted['documentId']}/reversal",
            json={"confirmed": True, "reversalDate": "2026-09-03", "reason": "Hủy BK"},
        )
        self.assertEqual(reversed_response.status_code, 200, reversed_response.get_data(as_text=True))
        with server.db() as conn:
            model = inventory_export.collect_inventory_export_model(
                conn, date_from="2026-08-01", date_to="2026-09-03",
            )
        labels = [row["movement_label"] for row in model["input_rows"]]
        self.assertEqual(labels, ["Nhập", "Hoàn tác nhập BK"])
        self.assertEqual([row["quantity"] for row in model["input_rows"]], [2, -2])

    def test_file_contract_identity_redaction_and_ui_exposes_confirm(self):
        invalid_type = self.preview(b"not xlsx", filename="bk.xlsm")
        corrupt = self.preview(b"not xlsx")
        self.assertEqual((invalid_type.status_code, invalid_type.get_json()["code"]), (400, "invalid_file_type"))
        self.assertEqual((corrupt.status_code, corrupt.get_json()["code"]), (400, "invalid_workbook"))
        row = self.valid_row()
        row[10] = "012345678901"
        row[11] = "CMND 012345678"
        response = self.preview(self.template_with_rows([row]))
        response_text = response.get_data(as_text=True)
        self.assertNotIn("012345678901", response_text)
        self.assertNotIn("012345678", response_text)
        self.assertFalse(response.get_json()["canConfirm"])

        root = Path(__file__).resolve().parent
        index = (root / "static" / "index.html").read_text(encoding="utf-8")
        script = (root / "static" / "app.js").read_text(encoding="utf-8")
        self.assertIn('id="bkWorkbookInput" accept=".xlsx"', index)
        self.assertIn('/api/bk-import/template', script)
        self.assertIn('/api/bk-import/preview', script)
        self.assertIn('/api/bk-import/confirm', script)
        self.assertIn('data-action="confirm-bk-import"', script)
        self.assertNotIn("đang khóa Q-004", script)


if __name__ == "__main__":
    unittest.main()
