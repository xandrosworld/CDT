import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from pypdf import PdfReader
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen.canvas import Canvas

from .excel_print_renderer import build_excel_pdf_bundle, _keep_delivery_footer_with_items


class DeliveryPrintSidesTests(unittest.TestCase):
    def render(self, counts, sides='auto'):
        temp=tempfile.TemporaryDirectory();self.addCleanup(temp.cleanup)
        root=Path(temp.name);source=root/'delivery.xlsx'
        wb=Workbook();wb.active.title='Order 0';wb.active['A1']='Delivery'
        for i in range(1,len(counts)):wb.create_sheet(f'Order {i}')['A1']='Delivery'
        wb.save(source);wb.close()
        def artwork(sources, *, paper, render_dir):
            result=[]
            for i,count in enumerate(counts):
                file=render_dir/f'{i}.pdf';canvas=Canvas(str(file),pagesize=A4)
                for p in range(count):
                    canvas.drawString(36,700,f'ORDER_{i}_PAGE_{p+1}');canvas.showPage()
                canvas.save()
                result.append({'path':file,'document_type':'deliveries','title':f'Order {i}',
                               'sheet':f'Order {i}','workbook':'delivery.xlsx','pages':count})
            return result
        with patch('tdp_system.excel_print_renderer._export_visible_sheets',side_effect=artwork):
            result=build_excel_pdf_bundle([{'path':source,'document_type':'deliveries'}],root/'result.pdf',paper='A4',duplex=sides)
        return result,PdfReader(root/'result.pdf')

    def test_one_page_orders_default_to_simplex_without_blank_pages(self):
        result,pdf=self.render([1,1,1])
        self.assertFalse(result['duplex']);self.assertEqual(result['pages'],3)
        self.assertEqual(result['page_layout']['blank_pages'],[])
        self.assertEqual(pdf.trailer['/Root']['/ViewerPreferences']['/Duplex'],'/Simplex')

    def test_mixed_odd_and_even_orders_start_on_separate_sheets(self):
        result,pdf=self.render([1,2,3,1,2])
        self.assertTrue(result['duplex'])
        self.assertEqual(result['page_layout']['blank_pages'],[2,8,10])
        for i,item in enumerate(result['page_layout']['sections']):
            self.assertEqual(item['start_page']%2,1)
            for j,p in enumerate(range(item['start_page'],item['end_page']+1)):
                self.assertIn(f'ORDER_{i}_PAGE_{j+1}',pdf.pages[p-1].extract_text())
        for p in result['page_layout']['blank_pages']:
            self.assertFalse(pdf.pages[p-1].extract_text().strip())
        prefs=pdf.trailer['/Root']['/ViewerPreferences']
        self.assertEqual(prefs['/Duplex'],'/DuplexFlipLongEdge')
        self.assertEqual(prefs['/PrintScaling'],'/None')
        self.assertTrue(prefs['/PickTrayByPDFSize'])
        for p in pdf.pages:
            self.assertAlmostEqual(float(p.mediabox.width),A4[0],places=2)
            self.assertAlmostEqual(float(p.mediabox.height),A4[1],places=2)

    def test_single_long_order_has_no_trailing_padding(self):
        result,pdf=self.render([3])
        self.assertTrue(result['duplex']);self.assertEqual(len(pdf.pages),3)
        self.assertEqual(result['page_layout']['blank_pages'],[])

    def test_explicit_simplex_overrides_auto_without_padding(self):
        result,pdf=self.render([1,2,1],False)
        self.assertFalse(result['duplex']);self.assertEqual(len(pdf.pages),4)
        self.assertEqual(result['page_layout']['blank_pages'],[])

    def test_signature_only_last_page_moves_two_items_without_changing_source(self):
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp);source=root/'delivery.xlsx'
            book=Workbook();sheet=book.active;sheet.title='Kitchen'
            for r in range(11,35):
                sheet.cell(r,3,r-10);sheet.cell(r,4,f'Product {r-10}')
            sheet.cell(35,3,'TOTAL');book.save(source);book.close()
            original=source.read_bytes()
            pdf=root/'before.pdf';canvas=Canvas(str(pdf),pagesize=A4)
            canvas.drawString(36,700,'Product 24');canvas.showPage()
            canvas.drawString(36,700,'TOTAL - SIGNATURES ONLY');canvas.save()
            entry={'path':pdf,'document_type':'deliveries','sheet':'Kitchen','workbook':source.name,'pages':2}
            def render(sources, *, paper, render_dir):
                adjusted=load_workbook(sources[0]['path'])
                self.assertEqual([b.id for b in adjusted.active.row_breaks.brk],[32])
                self.assertEqual([adjusted.active.cell(r,4).value for r in range(11,35)],
                                 [f'Product {i}' for i in range(1,25)])
                adjusted.close()
                result=render_dir/'after.pdf';canvas=Canvas(str(result),pagesize=A4)
                canvas.drawString(36,700,'Product 1 to Product 22');canvas.showPage()
                canvas.drawString(36,700,'Product 23 - Product 24 - TOTAL - SIGNATURES');canvas.save()
                return [{**entry,'path':result}]
            with patch('tdp_system.excel_print_renderer._export_visible_sheets',side_effect=render):
                result=_keep_delivery_footer_with_items([entry],[{'path':source,'document_type':'deliveries'}],paper='A4',render_dir=root)
            self.assertIn('Product 23',PdfReader(result[0]['path']).pages[-1].extract_text())
            self.assertEqual(source.read_bytes(),original)
