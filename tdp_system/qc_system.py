from __future__ import annotations

import io
import hashlib
import json
import sys
import zipfile
from collections import Counter
from pathlib import Path
from shutil import rmtree
from urllib.parse import quote

from docx import Document as DocxDocument
from openpyxl import Workbook, load_workbook

APP_DIR = Path(__file__).resolve().parent
ROOT = APP_DIR.parent
sys.path.insert(0, str(APP_DIR))
import server  # noqa: E402
import contract_modules  # noqa: E402
import quote_import  # noqa: E402
import quote_export  # noqa: E402
import template_workbook  # noqa: E402
import delivery_export  # noqa: E402
import purchase_summary_export  # noqa: E402
import receipt_export  # noqa: E402
import report_export  # noqa: E402
import print_bundle  # noqa: E402
import bk_import  # noqa: E402
from msmi_client import MsmiClient, MsmiConfig  # noqa: E402


def clean_test_db(path: Path):
    for suffix in ("", "-wal", "-shm"):
        target = Path(str(path) + suffix)
        if target.exists():
            target.unlink()


def run_template_workbook_qc():
    expected_hash = "66808CD910F63F4D543A72CC51DDAEBA1C60B147E0457FEAFC5BAD65EAD26AD3"
    source = ROOT / "Em Thành.xlsx"
    before = hashlib.sha256(source.read_bytes()).hexdigest().upper()
    assert before == expected_hash
    clones = [
        template_workbook.clone_template_workbook(
            source, sheet_names=["biên nhận"], expected_sha256=expected_hash,
        )
        for _ in range(2)
    ]
    try:
        signatures = [
            template_workbook.workbook_topology_signature(clone.workbook)
            for clone in clones
        ]
        assert signatures[0] == signatures[1]
        for clone in clones:
            assert clone.formula_report == {
                "preserved_formulas": 3,
                "replaced_external_reference": 40,
                "replaced_omitted_sheet_reference": 0,
                "replaced_unsafe_formula": 0,
                "missing_cached_values": 0,
                "external_hyperlinks_removed": 0,
                "defined_names_removed": 0,
            }
            assert len(clone.workbook["biên nhận"].merged_cells.ranges) == 11
            assert clone.workbook["biên nhận"]["F22"].value == "=SUM(F15:F21)"
            payload = clone.to_bytes()
            reopened = load_workbook(
                io.BytesIO(payload), data_only=False, keep_links=False,
            )
            try:
                assert not reopened._external_links
                template_workbook.assert_workbook_safe(reopened)
            finally:
                reopened.close()
    finally:
        for clone in clones:
            clone.close()
    assert hashlib.sha256(source.read_bytes()).hexdigest().upper() == before
    return "real receipt topology / 40 external formulas cached / 3 local formulas preserved / repeat deterministic"


def run_delivery_template_qc():
    source = ROOT / "Em Thành.xlsx"
    before = hashlib.sha256(source.read_bytes()).hexdigest().upper()
    assert before == delivery_export.EM_THANH_SHA256
    workbook = delivery_export.build_delivery_workbook(
        [
            {
                "kitchen": "NHUAHP",
                "contractor": "NHUAHAIPHONG",
                "recipient": "CÔNG TY CỔ PHẦN NHỰA VÀ CƠ KHÍ HẢI PHÒNG",
                "address": "Km 104 + 200 QL5, Phường Đông Hải, TP Hải Phòng, Việt Nam",
                "items": [
                    {
                        "product_name": "Bí xanh sơ chế (gọt vỏ)",
                        "quantity": 2,
                        "unit": "Kg",
                        "sell_price": 14000,
                        "note": "",
                    }
                ],
            },
            {
                "kitchen": "POT",
                "contractor": "HATRAN",
                "recipient": "BẾP POT",
                "address": "476 Phạm Văn Đồng, Hải Phòng",
                "items": [
                    {
                        "product_name": "Cà rốt",
                        "quantity": 3,
                        "unit": "Kg",
                        "sell_price": 987654321,
                        "note": "",
                    }
                ],
            },
        ],
        work_date="2026-09-02",
        template_path=source,
    )
    try:
        priced = workbook["NHUAHP"]
        hidden = workbook["POT"]
        assert priced["H10"].value == "Đơn giá" and priced["I11"].value == 28000
        assert hidden.column_dimensions["H"].hidden
        assert hidden.column_dimensions["I"].hidden
        assert all(hidden.cell(row, column).value is None for row in range(10, 40) for column in (7, 8, 9))
        assert all(len(sheet._images) == 0 for sheet in workbook.worksheets)
        assert priced["C6"].value == "Ngày 02 tháng 09 năm 2026"
        assert priced["C6"].alignment.horizontal == "center"
        assert all(priced[coordinate].font.bold for coordinate in ("C1", "C2", "C3"))
        assert priced["C9"].value == "Hình thức thanh toán: TM/CK"
        assert priced["C13"].value == "Ngày ..... tháng ..... năm ........"
        assert hidden["C12"].value == "Ngày ..... tháng ..... năm ........"
        assert "Người nhận hàng" in priced["C14"].value
        assert "Người nhận hàng" in hidden["C13"].value
        assert priced["C15"].value.count("Ký và ghi rõ họ tên") == 4
        assert hidden["C14"].value.count("Ký và ghi rõ họ tên") == 4
        assert all(str(sheet.page_setup.paperSize) == "9" for sheet in workbook.worksheets)
        assert all(sheet.page_setup.orientation == "portrait" for sheet in workbook.worksheets)
        assert all(sheet.print_options.horizontalCentered for sheet in workbook.worksheets)
        assert all(sheet.print_title_rows == "$10:$10" for sheet in workbook.worksheets)
        assert all(len(sheet.conditional_formatting) == 0 for sheet in workbook.worksheets)
        assert all(sheet.column_dimensions["D"].width == 43 for sheet in workbook.worksheets)
        assert all(sheet.column_dimensions["J"].width == 24 for sheet in workbook.worksheets)
        assert all(sheet["D11"].font.name == "Arial" for sheet in workbook.worksheets)
        assert all((sheet["D11"].font.sz or 0) >= 16 for sheet in workbook.worksheets)
        assert priced["D37"].value is None and hidden["D37"].value is None
        sections = print_bundle.workbook_sections("deliveries", workbook)
        assert [len(section["signatures"]) for section in sections] == [4, 4]
        assert [column["label"] for column in sections[0]["columns"]] == [
            "STT", "Tên hàng", "SL", "ĐVT", "Đơn giá", "Thành tiền", "GC",
        ]
        assert [column["label"] for column in sections[1]["columns"]] == [
            "STT", "Tên hàng", "SL", "ĐVT", "GC",
        ]
        candidate = template_workbook.safe_workbook_bytes(workbook)
    finally:
        workbook.close()
    reopened = load_workbook(io.BytesIO(candidate), data_only=False, keep_links=False)
    try:
        assert reopened.sheetnames == ["NHUAHP", "POT"]
        assert [len(sheet._images) for sheet in reopened.worksheets] == [0, 0]
        assert not reopened._external_links
    finally:
        reopened.close()
    assert hashlib.sha256(source.read_bytes()).hexdigest().upper() == before
    assert delivery_export.delivery_prices_visible("NHUAHP", "NHUAHAIPHONG")
    assert not delivery_export.delivery_prices_visible("BẾP NHỰA", "NHUAHAIPHONG")
    return (
        "approved centred title-subtitle-date header / bold supplier identity / payment method / "
        "full-width balanced pages / repeated header / "
        "native editable 4-signature cells / exact NHUAHP+NHUAHAIPHONG price gate / "
        "no hidden price leakage / A4"
    )


def run_purchase_summary_template_qc():
    source = ROOT / "Em Thành.xlsx"
    before = hashlib.sha256(source.read_bytes()).hexdigest().upper()
    assert before == purchase_summary_export.EM_THANH_SHA256
    fake_identity = "0" * 12
    rows = [
        {
            "work_date": "2026-09-02",
            "seller": "Người bán kiểm thử",
            "address": "Địa chỉ kiểm thử",
            "cccd": fake_identity,
            "product_name": "Cà rốt",
            "unit": "Kg",
            "quantity": 2,
            "amount": 200,
            "supplier": "NCC A",
            "kitchen": "BẾP A",
            "source_ref": 3,
        },
        {
            "work_date": "2026-09-02",
            "seller": "Người bán kiểm thử",
            "address": "Địa chỉ kiểm thử",
            "cccd": fake_identity,
            "product_name": "Cà rốt",
            "unit": "kg",
            "quantity": 3,
            "amount": 360,
            "supplier": "NCC B",
            "kitchen": "BẾP B",
            "source_ref": 4,
        },
        {
            "work_date": "2026-09-02",
            "seller": "Người bán kiểm thử",
            "address": "Địa chỉ kiểm thử",
            "cccd": fake_identity,
            "product_name": "Khoai tây",
            "unit": "Kg",
            "quantity": 1,
            "amount": 150,
            "supplier": "NCC A",
            "kitchen": "BẾP A",
            "source_ref": 5,
        },
    ]
    grouped = purchase_summary_export.aggregate_purchase_summary_rows(rows)
    assert len(grouped) == 2
    carrot = next(item for item in grouped if item["product_name"] == "Cà rốt")
    assert carrot["quantity"] == 5 and carrot["unit_price"] == 112 and carrot["amount"] == 560
    assert carrot["suppliers"] == ("NCC A", "NCC B")
    assert carrot["kitchens"] == ("BẾP A", "BẾP B")

    workbook = purchase_summary_export.build_purchase_summary_workbook(
        rows, template_path=source,
    )
    try:
        sheet = workbook["bảng kê tổng"]
        assert workbook.sheetnames == ["bảng kê tổng"]
        assert sheet["A2"].value == "Từ ngày 02/09/2026 đến ngày 02/09/2026"
        assert sheet["G11"].value == 5 and sheet["H11"].value == 112 and sheet["I11"].value == 560
        assert sheet["A13"].value == "TỔNG CỘNG" and sheet["I13"].value == 710
        assert all(sheet.cell(7, column).value is None for column in range(1, 11))
        assert sheet["C49"].value is None and sheet["C50"].value is None
        assert str(sheet.print_area) == "'bảng kê tổng'!$A$1:$J$23"
        assert sheet.page_setup.orientation == "landscape"
        assert str(sheet.page_setup.paperSize) == "9"
        sections = print_bundle.workbook_sections("purchases", workbook)
        assert len(sections) == 1 and len(sections[0]["rows"]) == 2
        assert sections[0]["summary"][1]["value"] == 710
        candidate = template_workbook.safe_workbook_bytes(workbook)
    finally:
        workbook.close()
    reopened = load_workbook(io.BytesIO(candidate), data_only=False, keep_links=False)
    try:
        assert reopened.sheetnames == ["bảng kê tổng"]
        assert not reopened._external_links
        assert reopened["bảng kê tổng"]["C50"].value is None
    finally:
        reopened.close()
    assert hashlib.sha256(source.read_bytes()).hexdigest().upper() == before
    return "confirmed BK only / legal seller item merge / weighted average / exact total / no sample identity / A4"


def run_receipt_template_qc():
    source = ROOT / "Em Thành.xlsx"
    before = hashlib.sha256(source.read_bytes()).hexdigest().upper()
    assert before == purchase_summary_export.EM_THANH_SHA256
    fake_identity = "0" * 12
    rows = [
        {
            "work_date": "2026-09-03",
            "seller": "Người bán kiểm thử",
            "address": "Địa chỉ kiểm thử",
            "cccd": fake_identity,
            "issue_date": "02/01/2020",
            "issue_place": "Nơi cấp kiểm thử",
            "product_name": "Cà rốt",
            "unit": "Kg",
            "quantity": 2,
            "amount": 200,
            "supplier": "NCC A",
            "kitchen": "BẾP A",
            "source_ref": 3,
        },
        {
            "work_date": "2026-09-03",
            "seller": "Người bán kiểm thử",
            "address": "Địa chỉ kiểm thử",
            "cccd": fake_identity,
            "issue_date": "02/01/2020",
            "issue_place": "Nơi cấp kiểm thử",
            "product_name": "Cà rốt",
            "unit": "kg",
            "quantity": 3,
            "amount": 360,
            "supplier": "NCC B",
            "kitchen": "BẾP B",
            "source_ref": 4,
        },
    ]
    workbook = receipt_export.build_purchase_documents_workbook(
        rows,
        template_path=source,
        buyer_name="Người mua kiểm thử",
        buyer_title="Nhân viên thu mua",
        company_name="CÔNG TY KIỂM THỬ",
        company_address="Địa chỉ công ty kiểm thử",
    )
    try:
        assert workbook.sheetnames == ["bảng kê tổng", "biên nhận"]
        sheet = workbook["biên nhận"]
        assert sheet["D10"].value == fake_identity
        assert sheet["F16"].value == 5 and sheet["G16"].value == 560
        assert sheet["C17"].value.startswith("Số tiền bằng chữ:")
        assert sheet.max_row == 26
        assert str(sheet.print_area) == "'biên nhận'!$C$1:$G$26"
        assert sheet.page_setup.orientation == "portrait"
        assert str(sheet.page_setup.paperSize) == "9"
        assert sheet.print_options.horizontalCentered
        assert "D8:G8" in {str(value) for value in sheet.merged_cells.ranges}
        assert "C22:D22" in {str(value) for value in sheet.merged_cells.ranges}
        assert sheet["E21"].value == "Hải Phòng, ngày 03 tháng 09 năm 2026"
        assert sheet["E21"].alignment.horizontal == "center"
        assert sheet["E22"].alignment.horizontal == "center"
        assert sheet["E26"].alignment.horizontal == "center"
        assert sheet["C15"].font.name == "Times New Roman"
        assert sheet["C15"].font.sz == 12
        sections = print_bundle.workbook_sections("purchases", workbook)
        assert len(sections) == 2
        assert sections[1]["document_type"] == "purchase_receipt"
        assert sections[1]["title"] == "GIẤY BIÊN NHẬN"
        assert len(sections[1]["rows"]) == 1
        candidate = template_workbook.safe_workbook_bytes(workbook)
    finally:
        workbook.close()
    reopened = load_workbook(io.BytesIO(candidate), data_only=False, keep_links=False)
    try:
        assert reopened.sheetnames == ["bảng kê tổng", "biên nhận"]
        assert not reopened._external_links
        assert reopened["biên nhận"]["D10"].value == fake_identity
    finally:
        reopened.close()
    assert hashlib.sha256(source.read_bytes()).hexdigest().upper() == before
    return (
        "one golden receipt per legal seller/day / issue identity / 5m hard cap / "
        "dynamic rows / aligned identity+table+two signatures / A4 / manifest-safe title"
    )


def run_monthly_report_template_qc():
    source = ROOT / "Em Thành.xlsx"
    before = hashlib.sha256(source.read_bytes()).hexdigest().upper()
    assert before == purchase_summary_export.EM_THANH_SHA256
    rows = [
        {
            "contractor": "ATV", "kitchen": "LSVINA",
            "revenue": 100, "cost": 60, "profit": 40, "total": 108,
            "source_ref": 1,
        },
        {
            "contractor": "ATV", "kitchen": "BẾP-MỚI",
            "revenue": 200, "cost": 100, "profit": 100, "total": 216,
            "source_ref": 2,
        },
        {
            "contractor": "HATRAN", "kitchen": "POT",
            "revenue": 300, "cost": 200, "profit": 100, "total": 324,
            "source_ref": 3,
        },
    ]
    workbook = report_export.build_monthly_report_workbook(
        rows,
        period="2026-09",
        template_path=source,
        configured_groups={"BẾP-MỚI": "CHỊ TÚ"},
    )
    try:
        sheet = workbook["báo cáo tổng hợp"]
        assert workbook.sheetnames == ["báo cáo tổng hợp"]
        assert sheet["C3"].value == "LSVINA"
        assert sheet["C4"].value == "BẾP-MỚI"
        assert (sheet["D5"].value, sheet["E5"].value, sheet["F5"].value, sheet["G5"].value) == (300, 160, 140, 324)
        assert sheet["C6"].value == "POT"
        assert sheet["A7"].value == "TỔNG THÁNG"
        assert (sheet["D7"].value, sheet["E7"].value, sheet["F7"].value, sheet["G7"].value) == (600, 360, 240, 648)
        assert sheet["A5"].fill.fgColor.rgb == "FF92D050"
        assert sheet["A7"].fill.fgColor.rgb == "FF92D050"
        assert str(sheet.print_area) == "'báo cáo tổng hợp'!$A$2:$G$7"
        assert sheet.print_title_rows == "$2:$2"
        assert str(sheet.page_setup.paperSize) == "9"
        assert sheet.page_setup.orientation == "landscape"
        sections = print_bundle.workbook_sections("report", workbook)
        assert len(sections) == 1 and sections[0]["title"] == "BÁO CÁO TỔNG HỢP"
        assert sections[0]["summary"][-1]["value"] == 648
        candidate = template_workbook.safe_workbook_bytes(workbook)
    finally:
        workbook.close()
    reopened = load_workbook(io.BytesIO(candidate), data_only=False, keep_links=False)
    try:
        assert reopened.sheetnames == ["báo cáo tổng hợp"]
        assert reopened.active["A7"].value == "TỔNG THÁNG"
        assert not reopened._external_links
    finally:
        reopened.close()
    assert hashlib.sha256(source.read_bytes()).hexdigest().upper() == before
    return "golden 7 columns / monthly approved scope / dynamic group+kitchen / group totals / TỔNG THÁNG / A4"


def run_toyota_quote_golden_qc():
    expected_hash = "C3C16615F8DA1AB80843DE3DD611CB8A472B401B8C7B93B177781EDEB203E148"
    candidates = list((Path.home() / "Downloads").glob("*TOYOTA*T09-2026.xlsx"))
    source = next(
        (
            path for path in candidates
            if hashlib.sha256(path.read_bytes()).hexdigest().upper() == expected_hash
        ),
        None,
    )
    assert source is not None, "Thiếu golden báo giá Toyota T09-2026 đúng hash"
    golden = load_workbook(source, data_only=True, keep_links=False)
    try:
        assert golden.sheetnames == ["all"] and not golden._external_links
        sheet = golden["all"]
        reference_rows = []
        reference_groups = []
        reference_zero_count = 0
        for row_number in range(9, sheet.max_row + 1):
            code = sheet.cell(row_number, 2).value
            group_name = sheet.cell(row_number, 3).value
            if not code:
                if group_name and row_number < 539:
                    reference_groups.append(str(group_name).strip())
                continue
            raw_price = sheet.cell(row_number, 5).value
            if raw_price == 0:
                reference_zero_count += 1
            # Text dashes in the old golden are presentation-only placeholders.
            # They become zero only in this layout fixture; production rows still
            # come exclusively from the confirmed period import and omit X/blank.
            price = raw_price if isinstance(raw_price, (int, float)) else 0
            reference_rows.append({
                "product_code": code,
                "product_name": sheet.cell(row_number, 3).value,
                "unit": sheet.cell(row_number, 4).value,
                "tax": sheet.cell(row_number, 6).value,
                "sell_price": price,
                "price_state": "zero" if price == 0 else "numeric",
                "exportable": True,
            })
        assert len(reference_rows) == 513
        assert len({item["product_code"] for item in reference_rows}) == 513
        assert reference_zero_count == 5
        assert reference_groups == list(quote_export.QUOTE_GROUPS.values())
        assert not any(
            cell.data_type == "f"
            for cells in sheet.iter_rows()
            for cell in cells
        )
        golden_dimensions = {
            column: sheet.column_dimensions[column].width for column in "ABCDEFG"
        }
        golden_margins = tuple(
            getattr(sheet.page_margins, field)
            for field in ("left", "right", "top", "bottom", "header", "footer")
        )
    finally:
        golden.close()

    candidate = quote_export.build_toyota_quote_workbook(
        reference_rows,
        period="2026-09",
        version={"version_no": 1, "source_hash": expected_hash},
    )
    try:
        assert candidate.sheetnames == ["all"] and not candidate._external_links
        sheet = candidate["all"]
        assert sheet.max_row == 542 and sheet.max_column == 6
        assert [sheet.cell(8, column).value for column in range(1, 7)] == [
            "STT", "MÃ", "TÊN THÀNH ĐẠT PHÁT", "ĐVT", "Giá chưa VAT", "Thuế",
        ]
        assert sheet.print_title_rows == "$8:$8"
        assert (
            sheet.page_setup.orientation, str(sheet.page_setup.paperSize),
            sheet.page_setup.scale, sheet.page_setup.fitToHeight,
        ) == ("landscape", "9", 87, 0)
        assert {
            column: sheet.column_dimensions[column].width for column in "ABCDEFG"
        } == golden_dimensions
        assert tuple(
            getattr(sheet.page_margins, field)
            for field in ("left", "right", "top", "bottom", "header", "footer")
        ) == golden_margins
        group_rows = [
            cell.row for cell in sheet["C"] if cell.value in quote_export.QUOTE_GROUPS.values()
        ]
        assert group_rows == [9, 40, 59, 72, 110, 119, 127, 149, 159, 282, 309, 327, 333, 458, 483, 534]
        assert sheet["A8"].fill.fgColor.rgb == "00FFFF00"
        assert all(sheet.cell(row_number, 3).fill.fgColor.rgb == "0092D050" for row_number in group_rows)
        assert all(sheet.cell(row_number, 3).font.bold for row_number in group_rows)
        for start, end in zip(group_rows, group_rows[1:] + [538]):
            codes = [
                sheet.cell(row_number, 2).value
                for row_number in range(start + 1, end)
                if sheet.cell(row_number, 2).value
            ]
            assert codes == sorted(codes, key=str.casefold)
        assert sheet["A539"].value == "Báo giá trên chưa bao gồm VAT!"
        assert sheet["D539"].value is None
        assert "TOYOTA · kỳ 2026-09 · phiên bản 1" in candidate.properties.subject
        assert sheet["C541"].value == "XÁC NHẬN CỦA BÊN BÁN"
        assert not any(
            cell.data_type == "f"
            for cells in sheet.iter_rows()
            for cell in cells
        )
    finally:
        candidate.close()
    return "1 sheet all / 513-row golden scale / 16 green groups / A-Z / A4 landscape / repeat header / static values"


def run_quote_import_qc():
    quote_db = APP_DIR / "data" / "qc_quote_import.sqlite3"
    clean_test_db(quote_db)
    source = ROOT / "Em Thành.xlsx"
    try:
        assert source.exists(), "Thiếu nguồn ma trận báo giá Em Thành.xlsx"
        server.DB_PATH = quote_db
        server.MASTER_SOURCE = source
        server.init_database()
        client = server.app.test_client()
        source_bytes = source.read_bytes()

        def preview(period, payload=source_bytes, filename=source.name):
            response = client.post(
                "/api/quotes/import/preview",
                data={
                    "effective_period": period,
                    "file": (io.BytesIO(payload), filename),
                },
                content_type="multipart/form-data",
            )
            assert response.status_code == 200, response.get_data(as_text=True)
            return response.get_json()

        first = preview("2026-09")
        assert first["sheet"] == "BÁO GIÁ" and first["headerRow"] == 2
        assert first["priceGroups"] == [
            "ATV", "HATRAN", "BIADAUVOI", "NGUYENGIA", "SUPPY", "TOYOTA", "NHUAHAIPHONG",
        ]
        assert first["counts"] == {
            "products": 899,
            "price_groups": 7,
            "price_cells": 6293,
            "rows_with_errors": 0,
            "rows_with_warnings": 161,
            "duplicate_codes": 28,
            "safe_duplicate_codes": 27,
            "conflict_codes": 1,
            "conflicts": 5,
        }
        assert first["canConfirm"] is False and first["proposedVersion"] == 1
        assert next(item for item in first["priceColumns"] if item["priceGroup"] == "TOYOTA") == {
            "priceGroup": "TOYOTA", "sourceColumn": 17, "sourceHeader": "TOYOTA",
        }
        assert {item["productCode"] for item in first["conflicts"]} == {"A000045"}
        blocked = client.post(
            "/api/quotes/import/confirm",
            json={"token": first["token"], "confirmed": True, "state_hash": first["stateHash"]},
        )
        assert blocked.status_code == 400 and blocked.get_json()["code"] == "preview_has_errors"

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "BÁO GIÁ"
        sheet.append(["BẢNG BÁO GIÁ QC"])
        sheet.append([
            "STT", "MÃ THAM CHIẾU", "MÃ HÀNG", "TÊN THÀNH ĐẠT PHÁT", "Giá mua", "NCC",
            None, "bk", "Tên làm bảng kê", "ĐVT", "THUẾ", "ATV", "HATRAN", "BIADAUVOI",
            "NGUYENGIA", "SUPPY", "TOYOTA", "NHUAHAIPHONG", "Thêm",
        ])
        sheet.append([None, 1, 2, 3, 4, 5, None, 6, 7, 8, 9, 10, 11, 12, 13, 14, 15, 16, 17])
        sheet.append([1, "L000002ánh", "L000002", "Gạo Bắc Hương", 16000, "ánh", None, None, None,
                      "Kg", "KKKNT", 17000, 17000, 19000, 19000, 19000, 17500, 17000])
        sheet.append([2, "L000004ánh", "L000004", "Gạo tám", 18000, "ánh", None, None, None,
                      "Kg", "KKKNT", "x", 20000, None, 21000, 0, 20500, 20000])
        output = io.BytesIO()
        workbook.save(output)
        workbook.close()
        clean_payload = output.getvalue()

        clean = preview("2026-09", clean_payload, "quote-qc.xlsx")
        assert clean["canConfirm"] is True and clean["counts"]["conflicts"] == 0
        confirmed = client.post(
            "/api/quotes/import/confirm",
            json={"token": clean["token"], "confirmed": True, "state_hash": clean["stateHash"]},
        )
        assert confirmed.status_code == 200, confirmed.get_data(as_text=True)
        contractor_exports = {
            "ATV": {
                "recipient": "CÔNG TY CỔ PHẦN SUẤT ĂN CÔNG NGHIỆP ATV",
                "codes": ["L000002"], "prices": [17000],
            },
            "SUPPY": {
                "recipient": "CÔNG TY TNNN SUPPLY",
                "codes": ["L000002", "L000004"], "prices": [19000, 0],
            },
            "TOYOTA": {
                "recipient": "CÔNG TY TNHH TOYOTA NANKAI HẢI PHÒNG",
                "codes": ["L000002", "L000004"], "prices": [17500, 20500],
            },
        }
        observed_first_prices = {}
        for contractor, expectation in contractor_exports.items():
            response = client.get(f"/api/export/quote/{contractor}?period=2026-09")
            assert response.status_code == 200, response.status
            assert (
                f"BAO_GIA_{contractor}_T09-2026_V1.xlsx"
                in response.headers["Content-Disposition"]
            )
            exported_book = load_workbook(
                io.BytesIO(response.data), data_only=False, keep_links=False,
            )
            try:
                exported_sheet = exported_book["all"]
                assert exported_book.sheetnames == ["all"] and not exported_book._external_links
                assert exported_sheet["A6"].value == f"KÍNH GỬI: {expectation['recipient']}"
                assert exported_sheet["C9"].value == "GẠO"
                product_rows = [
                    row for row in range(9, exported_sheet.max_row + 1)
                    if exported_sheet.cell(row, 2).value
                ]
                assert [exported_sheet.cell(row, 2).value for row in product_rows] == expectation["codes"]
                prices = [exported_sheet.cell(row, 5).value for row in product_rows]
                assert prices == expectation["prices"]
                observed_first_prices[contractor] = prices[0]
                note_row = next(
                    cell.row for cell in exported_sheet["A"]
                    if cell.value == "Báo giá trên chưa bao gồm VAT!"
                )
                assert exported_sheet.cell(note_row, 4).value is None
                assert (
                    f"{contractor} · kỳ 2026-09 · phiên bản 1"
                    in exported_book.properties.subject
                )
                assert not any(
                    cell.data_type == "f"
                    for cells in exported_sheet.iter_rows()
                    for cell in cells
                )
            finally:
                exported_book.close()
        assert observed_first_prices == {"ATV": 17000, "SUPPY": 19000, "TOYOTA": 17500}

        replay = preview("2026-09", clean_payload, "renamed-same-quotation.xlsx")
        assert replay["replay"] is True and replay["proposedVersion"] == 1
        replay_confirm = client.post(
            "/api/quotes/import/confirm",
            json={"token": replay["token"], "confirmed": True, "state_hash": replay["stateHash"]},
        )
        assert replay_confirm.status_code == 200 and replay_confirm.get_json()["idempotent"] is True

        next_period = preview("2026-10", clean_payload, "quote-next-period.xlsx")
        next_confirm = client.post(
            "/api/quotes/import/confirm",
            json={
                "token": next_period["token"], "confirmed": True,
                "state_hash": next_period["stateHash"],
            },
        )
        assert next_confirm.status_code == 200 and next_confirm.get_json()["versionNo"] == 1
        with server.db() as conn:
            assert conn.execute("PRAGMA integrity_check").fetchone()[0] == "ok"
            table_counts = [row[0] for row in conn.execute(
                "SELECT COUNT(*) FROM quote_versions UNION ALL "
                "SELECT COUNT(*) FROM quote_version_products UNION ALL "
                "SELECT COUNT(*) FROM quote_version_prices"
            )]
            assert table_counts == [2, 4, 28]
            toyota = quote_import.quote_rows_for_contractor(conn, "TOYOTA", "2026-09")
            assert {item["product_code"]: item["sell_price"] for item in toyota["items"]} == {
                "L000002": 17500, "L000004": 20500,
            }
            suppy = quote_import.quote_rows_for_contractor(conn, "SUPPY", "2026-09")
            assert any(item["product_code"] == "L000004" and item["price_state"] == "zero" for item in suppy["items"])
            audit_rows = [json.loads(row[0]) for row in conn.execute(
                "SELECT metadata_json FROM audit_log WHERE event_type='quote.import.confirm'"
            )]
            assert len(audit_rows) == 2
            assert all("source_hash" in row and "effective_period" in row for row in audit_rows)
            assert all("Em Thành.xlsx" not in json.dumps(row, ensure_ascii=False) for row in audit_rows)
        return (
            "Toyota=Q / ATV+SUPPY+TOYOTA isolated exports / legal recipients / "
            "real duplicate conflict surfaced / X+blank omitted / zero retained / replay safe"
        )
    finally:
        with quote_import.QUOTE_IMPORT_LOCK:
            quote_import.PENDING_QUOTE_IMPORTS.clear()
        clean_test_db(quote_db)


def mapping_workbook_bytes(headers, rows, header_row=3, sheet_name="Mapping"):
    workbook = Workbook()
    notes = workbook.active
    notes.title = "Hướng dẫn"
    notes["A1"] = "Sheet này không phải dữ liệu"
    sheet = workbook.create_sheet(sheet_name)
    sheet.cell(header_row - 1, 1, "Danh sách khách xác nhận")
    for column, header in enumerate(headers, start=1):
        sheet.cell(header_row, column, header)
    for row_index, values in enumerate(rows, start=header_row + 1):
        for column, value in enumerate(values, start=1):
            sheet.cell(row_index, column, value)
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def preview_mapping(client, mapping_type, payload, filename="mapping.xlsx"):
    return client.post(
        "/api/mappings/import/preview",
        data={"mapping_type": mapping_type, "file": (io.BytesIO(payload), filename)},
        content_type="multipart/form-data",
    )


def catalog_workbook_bytes(rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "danh mục hàng hóa"
    headers = ["STT", "MÃ HÀNG", "Nhóm hàng", "TÊN THÀNH ĐẠT PHÁT", "TÊN XUẤT HÓA ĐƠN", "ĐVT", "Thuế"]
    for column, value in enumerate(headers, start=1):
        sheet.cell(1, column, value)
    for row_index, values in enumerate(rows, start=2):
        sheet.cell(row_index, 1, row_index - 1)
        for column, value in enumerate(values, start=2):
            sheet.cell(row_index, column, value)
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def preview_catalog(client, payload, filename="danh-muc.xlsx"):
    return client.post(
        "/api/catalog/import/preview",
        data={"file": (io.BytesIO(payload), filename)},
        content_type="multipart/form-data",
    )


def kitchen_workbook_bytes(product_code="A000047", duplicate_first=False):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "T2"
    sheet.cell(3, 5, "VINA")
    headers = ["Nhà thầu", "Mã hàng", "Mã bếp", "Món ăn", "SỐ SUẤT ĂN CA SÁNG",
               10, "số lượng", None, "Giá bán", "Thành tiền"]
    for column, value in enumerate(headers, start=1):
        sheet.cell(4, column, value)
    first = ["HATRAN", product_code, "XCOM", "THỊT LUỘC", "Thịt nách heo",
             100, 1, "Kg", 105000, 105000]
    for column, value in enumerate(first, start=1):
        sheet.cell(5, column, value)
    next_row = 6
    if duplicate_first:
        duplicate = list(first)
        duplicate[6] = 0.5
        duplicate[9] = 52500
        for column, value in enumerate(duplicate, start=1):
            sheet.cell(next_row, column, value)
        next_row += 1
    placeholder = ["HATRAN", "-", "VINA", None, None, None, 0, "-", 0, 0]
    for column, value in enumerate(placeholder, start=1):
        sheet.cell(next_row, column, value)
    next_row += 1
    sheet.cell(next_row, 5, "CA SÁNG")
    sheet.cell(next_row, 6, 4)
    next_row += 1
    second = ["HATRAN", "H000007", "MAZDA", "TRỨNG ỐP LA", "Trứng gà CN",
              1000, 8, "Quả", 3000, 24000]
    for column, value in enumerate(second, start=1):
        sheet.cell(next_row, column, value)
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def preview_kitchen(client, payload, work_date="2026-09-02", filename="xưởng cơm QC.xlsx"):
    return client.post(
        "/api/kitchen/import/preview",
        data={"work_date": work_date, "file": (io.BytesIO(payload), filename)},
        content_type="multipart/form-data",
    )


def opening_workbook_bytes(rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Tồn cuối tháng"
    headers = ["MÃ TĐP", "TÊN TDP", "Tên trên HĐ", "MÃ KHO", "T/Suất", "ĐVT"]
    for column, value in enumerate(headers, start=1):
        sheet.cell(5, column, value)
    sheet.cell(5, 7, "Tồn cuối kỳ")
    sheet.cell(6, 7, "Số lượng")
    sheet.cell(6, 8, "Đơn giá")
    sheet.cell(6, 9, "Thành tiền")
    for row_index, values in enumerate(rows, start=7):
        for column, value in enumerate(values, start=1):
            sheet.cell(row_index, column, value)
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def preview_opening(client, payload, period="2026-08", filename="tồn đầu kỳ.xlsx"):
    return client.post(
        "/api/inventory/opening/import/preview",
        data={"period": period, "file": (io.BytesIO(payload), filename)},
        content_type="multipart/form-data",
    )


def run_tax_template_golden_qc():
    """Lock the four customer tax templates to the exact audited source files."""
    source_dir = ROOT / "bosung.30.8.26"
    assert source_dir.is_dir(), (
        f"Thiếu thư mục mẫu thuế chuẩn của khách: {source_dir}"
    )

    headers = [
        "Mã hàng", "Tên hàng", "Đơn vị tính", "Số lượng", "Đơn giá",
        "Cộng tiền hàng", "%CK", "Tiền CK", "Tiền trước thuế", "% VAT",
        "Tiền thuế GTGT", "Tổng tiền", "Tính chất",
    ]
    specs = {
        "thue 0.xlsx": {
            "size": 15438,
            "sha256": "82642C993BF26E52B54762E397ED5A1C3F5886F308834582DBCD01C4FBEA0619",
            "sheets": ["Sheet2", "Sheet3", "Sheet1"],
            "active": "Sheet2",
            "dimension": "A1:M10",
            "rows": 9,
            "quantity": 55.8,
            "goods": 4161900,
            "pretax": 4161900,
            "tax": 0,
            "grand": 4161900,
            "vat": -2,
            "nature": {"1": 9},
        },
        "thue 8.xlsx": {
            "size": 10649,
            "sha256": "892C57A688E28AFADDB58E2070E122945A926FB6EF1DC417714A004885D4C573",
            "sheets": ["Sheet 1 (2)"],
            "active": "Sheet 1 (2)",
            "dimension": "A1:M13",
            "rows": 12,
            "quantity": 6125,
            "goods": 28333000,
            "pretax": 28333000,
            "tax": 2266640,
            "grand": 30599640,
            "vat": 8,
            "nature": {"1": 12},
        },
        "thue 10.xlsx": {
            "size": 9566,
            "sha256": "D4548D308CDD236A854972A629DB63B4CEA2FE1B65816EAFF6AEBE57040D0DE9",
            "sheets": ["Sheet 1 (2)"],
            "active": "Sheet 1 (2)",
            "dimension": "A1:M4",
            "rows": 2,
            "quantity": 105,
            "goods": 18060000,
            "pretax": 18060000,
            "tax": 1806000,
            "grand": 19866000,
            "vat": 10,
            "nature": {"1": 2},
        },
        "thue 10 có khuyến mại.xlsx": {
            "size": 9706,
            "sha256": "243EC60B25C3235F447A7C773E66C7948F36B3AA5BE3F3AFD0A4F06A78F0691B",
            "sheets": ["Sheet 1 (2)"],
            "active": "Sheet 1 (2)",
            "dimension": "A1:M5",
            "rows": 3,
            "quantity": 109,
            "goods": 18060000,
            "pretax": 18060000,
            "tax": 1806000,
            "grand": 19866000,
            "vat": 10,
            "nature": {"1": 2, "2": 1},
        },
    }

    audited_rows = 0
    for filename, expected in specs.items():
        source = source_dir / filename
        assert source.is_file(), f"Thiếu mẫu thuế chuẩn của khách: {source}"
        actual_size = source.stat().st_size
        assert actual_size == expected["size"], (
            f"Sai dung lượng mẫu thuế {filename}: {actual_size}, "
            f"cần {expected['size']} byte"
        )
        actual_hash = hashlib.sha256(source.read_bytes()).hexdigest().upper()
        assert actual_hash == expected["sha256"], (
            f"Sai SHA-256 mẫu thuế {filename}: {actual_hash}; "
            f"cần {expected['sha256']}"
        )

        workbook = load_workbook(
            source, read_only=True, data_only=True, keep_links=False,
        )
        try:
            assert workbook.sheetnames == expected["sheets"], (
                f"Sai danh sách/thứ tự sheet {filename}: {workbook.sheetnames}; "
                f"cần {expected['sheets']}"
            )
            assert workbook.active.title == expected["active"], (
                f"Sai sheet active {filename}: {workbook.active.title}; "
                f"cần {expected['active']}"
            )
            sheet = workbook[expected["active"]]
            assert sheet.calculate_dimension() == expected["dimension"], (
                f"Sai vùng dữ liệu {filename}/{sheet.title}: "
                f"{sheet.calculate_dimension()}, cần {expected['dimension']}"
            )
            actual_headers = [
                sheet.cell(1, column).value for column in range(1, 14)
            ]
            assert actual_headers == headers, (
                f"Sai tiêu đề mẫu thuế {filename}/{sheet.title}: {actual_headers}"
            )

            data_rows = [
                (row_number, list(values))
                for row_number, values in enumerate(
                    sheet.iter_rows(min_row=2, max_col=13, values_only=True),
                    start=2,
                )
                if any(value not in (None, "") for value in values)
            ]
            assert len(data_rows) == expected["rows"], (
                f"Sai số dòng dữ liệu {filename}/{sheet.title}: "
                f"{len(data_rows)}, cần {expected['rows']}"
            )
            audited_rows += len(data_rows)

            def total(column_index):
                return sum(float(values[column_index] or 0) for _, values in data_rows)

            totals = {
                "quantity": total(3),
                "goods": total(5),
                "pretax": total(8),
                "tax": total(10),
                "grand": total(11),
            }
            for key, actual in totals.items():
                tolerance = 1e-6 if key == "quantity" else 0.01
                assert abs(actual - expected[key]) <= tolerance, (
                    f"Sai tổng {key} mẫu thuế {filename}: {actual}, "
                    f"cần {expected[key]}"
                )

            vat_values = {values[9] for _, values in data_rows}
            assert vat_values == {expected["vat"]}, (
                f"Sai giá trị % VAT mẫu thuế {filename}: {vat_values}; "
                f"cần {{{expected['vat']}}}"
            )
            nature = Counter(str(values[12]) for _, values in data_rows)
            assert dict(nature) == expected["nature"], (
                f"Sai Tính chất mẫu thuế {filename}: {dict(nature)}; "
                f"cần {expected['nature']}"
            )

            if filename == "thue 0.xlsx":
                assert 0 not in vat_values and vat_values == {-2}, (
                    "thue 0.xlsx phải giữ mã -2 (KKKNT), không được diễn giải "
                    "thành thuế suất VAT 0%"
                )
            elif filename == "thue 10 có khuyến mại.xlsx":
                promo = next(
                    (values for row_number, values in data_rows if row_number == 4),
                    None,
                )
                assert promo is not None, (
                    f"Thiếu dòng khuyến mại nguồn 4 trong {filename}"
                )
                for column_index, column_name in (
                    (4, "E/Đơn giá"), (5, "F/Cộng tiền hàng"), (6, "G/%CK"),
                    (7, "H/Tiền CK"), (8, "I/Tiền trước thuế"),
                    (10, "K/Tiền thuế GTGT"), (11, "L/Tổng tiền"),
                ):
                    assert promo[column_index] is None, (
                        f"Dòng khuyến mại {filename}!{column_name} phải để trống, "
                        f"đang là {promo[column_index]!r}"
                    )
                assert str(promo[12]) == "2", (
                    f"Dòng khuyến mại {filename}!M4 phải có Tính chất=2"
                )
        finally:
            workbook.close()

    assert audited_rows == 26
    return "4 canonical templates / 26 rows / KKKNT(-2), VAT8, VAT10 / promo blanks passed"


def run_catalog_import_qc():
    catalog_db = APP_DIR / "data" / "qc_catalog_test.sqlite3"
    clean_test_db(catalog_db)
    actual_source = next((ROOT / "_HANDOFF" / "EXTERNAL_INPUTS").glob("08410fad*.xlsx"), None)
    try:
        server.DB_PATH = catalog_db
        server.init_database()
        client = server.app.test_client()

        conflict_file = catalog_workbook_bytes([
            ["QC-CAT-01", "QC", "Hàng QC A", "Tên hóa đơn A", "Kg", "8%"],
            ["QC-CAT-01", "QC", "Hàng QC B", "Tên hóa đơn B", "Kg", "8%"],
        ])
        conflict_preview = preview_catalog(client, conflict_file).get_json()
        assert conflict_preview["can_confirm"] is False
        assert conflict_preview["counts"]["error"] == 2
        assert client.post(
            "/api/catalog/import/confirm",
            json={"token": conflict_preview["token"], "confirmed": True},
        ).status_code == 400

        if actual_source is None or not actual_source.exists():
            return "synthetic conflict/atomic block passed; customer workbook not present"

        source_hash = hashlib.sha256(actual_source.read_bytes()).hexdigest().upper()
        assert source_hash == "635308AF3F203BA51456BB4F7984BE3034D18FBBE10E005B2827C394CFA4C305"
        with server.db() as conn:
            protected_before = dict(conn.execute(
                """SELECT code,supplier,buy_price,purchase_list,seller,cccd
                   FROM products WHERE code='A000001'"""
            ).fetchone())
            price_before = [tuple(row) for row in conn.execute(
                "SELECT price_group,price_text,price_value FROM product_prices WHERE product_code='A000001' ORDER BY price_group"
            )]

        with actual_source.open("rb") as handle:
            preview_response = client.post(
                "/api/catalog/import/preview",
                data={"file": (handle, actual_source.name)},
                content_type="multipart/form-data",
            )
        assert preview_response.status_code == 200, preview_response.get_data(as_text=True)
        preview = preview_response.get_json()
        assert preview["source_hash"] == source_hash
        assert preview["sheet"] == "danh mục hàng hóa" and preview["header_row"] == 1
        assert preview["can_confirm"] is True
        assert preview["counts"] == {
            "total": 1249,
            "unique_products": 1249,
            "new_products": 380,
            "update_products": 869,
            "unchanged_products": 0,
            "retained_products": 2,
            "new_names": 1249,
            "update_names": 0,
            "unchanged_names": 0,
            "duplicate": 0,
            "error": 0,
        }
        assert client.post(
            "/api/catalog/import/confirm", json={"token": preview["token"]},
        ).status_code == 400
        confirmed = client.post(
            "/api/catalog/import/confirm",
            json={"token": preview["token"], "confirmed": True},
        )
        assert confirmed.status_code == 200, confirmed.get_data(as_text=True)
        confirmed_data = confirmed.get_json()
        assert confirmed_data["processed"] == 1249
        assert confirmed_data["inserted_products"] == 380
        assert confirmed_data["inserted_names"] == 1249
        assert client.post(
            "/api/catalog/import/confirm",
            json={"token": preview["token"], "confirmed": True},
        ).status_code == 410

        with server.db() as conn:
            assert conn.execute("PRAGMA integrity_check").fetchone()[0] == "ok"
            assert conn.execute("SELECT COUNT(*) FROM products").fetchone()[0] == 1251
            assert conn.execute("SELECT COUNT(*) FROM outgoing_product_names").fetchone()[0] == 1249
            assert conn.execute(
                "SELECT invoice_name FROM outgoing_product_names WHERE product_code='I000164'"
            ).fetchone()[0] == "Bắp cải"
            assert conn.execute("SELECT unit FROM products WHERE code='D000021'").fetchone()[0] == "Kg"
            assert conn.execute("SELECT unit FROM products WHERE code='I000090'").fetchone()[0] == "Kg"
            assert conn.execute("SELECT product_group FROM products WHERE code='A000001'").fetchone()[0] == "A"
            assert conn.execute("SELECT 1 FROM products WHERE code='G000074'").fetchone()
            assert conn.execute("SELECT 1 FROM products WHERE code='I000189'").fetchone()
            protected_after = dict(conn.execute(
                """SELECT code,supplier,buy_price,purchase_list,seller,cccd
                   FROM products WHERE code='A000001'"""
            ).fetchone())
            price_after = [tuple(row) for row in conn.execute(
                "SELECT price_group,price_text,price_value FROM product_prices WHERE product_code='A000001' ORDER BY price_group"
            )]
            assert protected_after == protected_before
            assert price_after == price_before

        with actual_source.open("rb") as handle:
            repeat_response = client.post(
                "/api/catalog/import/preview",
                data={"file": (handle, "danh-muc-doi-ten.xlsx")},
                content_type="multipart/form-data",
            )
        repeat = repeat_response.get_json()
        assert repeat["counts"]["new_products"] == 0
        assert repeat["counts"]["update_products"] == 0
        assert repeat["counts"]["unchanged_products"] == 1249
        assert repeat["counts"]["unchanged_names"] == 1249
        repeat_confirm = client.post(
            "/api/catalog/import/confirm",
            json={"token": repeat["token"], "confirmed": True},
        )
        assert repeat_confirm.status_code == 200
        assert repeat_confirm.get_json()["unchanged_names"] == 1249

        with actual_source.open("rb") as handle:
            stale_response = client.post(
                "/api/catalog/import/preview",
                data={"file": (handle, actual_source.name)},
                content_type="multipart/form-data",
            )
        stale = stale_response.get_json()
        with server.db() as conn:
            conn.execute("UPDATE products SET unit='QC-CHANGED' WHERE code='A000001'")
        stale_confirm = client.post(
            "/api/catalog/import/confirm",
            json={"token": stale["token"], "confirmed": True},
        )
        assert stale_confirm.status_code == 409
        assert "thay đổi" in stale_confirm.get_json()["error"]
        return "1,249 products / 380 new / 1,249 confirmed invoice names / repeat safe"
    finally:
        contract_modules.PENDING_CATALOG_IMPORTS.clear()
        clean_test_db(catalog_db)


def main():
    tax_template_result = run_tax_template_golden_qc()
    catalog_import_result = run_catalog_import_qc()
    template_engine_result = run_template_workbook_qc()
    delivery_template_result = run_delivery_template_qc()
    purchase_summary_result = run_purchase_summary_template_qc()
    receipt_template_result = run_receipt_template_qc()
    monthly_report_result = run_monthly_report_template_qc()
    toyota_quote_result = run_toyota_quote_golden_qc()
    quote_import_result = run_quote_import_qc()
    qc_db = APP_DIR / "data" / "qc_test.sqlite3"
    clean_test_db(qc_db)
    server.DB_PATH = qc_db
    server.init_database()
    client = server.app.test_client()

    health = client.get("/health")
    assert health.status_code == 200 and health.get_json()["ok"]
    bootstrap = client.get("/api/bootstrap").get_json()
    assert bootstrap["master"]["product_count"] >= 800
    assert len(bootstrap["master"]["kitchens"]) >= 20
    assert bootstrap["master"]["settings"]["payment_requester"] == "VŨ THỊ THỤY"
    assert bootstrap["master"]["settings"]["payment_bank_account"] == "1052787580"
    assert bootstrap["master"]["settings"]["payment_bank_name"] == "Ngân hàng TMCP Ngoại Thương Việt Nam"

    class CapturingMsmiClient(MsmiClient):
        def _get(self, path, params=None):
            self.captured_path = path
            self.captured_params = params
            return {"listInvoice": [{} for _ in range(params["size"])]}

    paging_client = CapturingMsmiClient(MsmiConfig("https://example.invalid", "not-a-real-token"))
    paging_result = paging_client.list_invoices(size=200)
    assert paging_client.captured_path == "api/qlhd-api/invoices"
    assert paging_client.captured_params["size"] == 199 and paging_result["has_more"] is True

    with server.db() as conn:
        mapping_products = [dict(row) for row in conn.execute(
            """SELECT code,name FROM products
               WHERE code!='I000060' AND name IN (
                   SELECT name FROM products WHERE trim(name)!='' GROUP BY name HAVING COUNT(*)=1
               ) ORDER BY code LIMIT 4"""
        )]
        mapping_kitchens = [dict(row) for row in conn.execute(
            """SELECT code,name FROM kitchens
               WHERE code!='POT' AND trim(COALESCE(name,''))!='' AND name IN (
                   SELECT name FROM kitchens WHERE trim(COALESCE(name,''))!='' GROUP BY name HAVING COUNT(*)=1
               ) ORDER BY code LIMIT 2"""
        )]
    assert len(mapping_products) == 4 and len(mapping_kitchens) == 2

    invoice_mapping_file = mapping_workbook_bytes(
        ["Mã hàng", "Tên Thành Đạt Phát", "Tên xuất hóa đơn"],
        [
            [mapping_products[0]["code"], "", "TÊN HÓA ĐƠN QC 01"],
            ["", mapping_products[1]["name"], "Tên hóa đơn QC 02"],
            [mapping_products[0]["code"], "", "TÊN HÓA ĐƠN QC 01"],
        ],
        header_row=3,
        sheet_name="Tên đầu ra",
    )
    mapping_preview = preview_mapping(client, "invoice_names", invoice_mapping_file)
    assert mapping_preview.status_code == 200, mapping_preview.get_data(as_text=True)
    mapping_preview_data = mapping_preview.get_json()
    assert mapping_preview_data["sheet"] == "Tên đầu ra"
    assert mapping_preview_data["header_row"] == 3
    assert mapping_preview_data["can_confirm"] is True
    assert mapping_preview_data["counts"] == {
        "total": 3, "new": 2, "update": 0, "unchanged": 0, "duplicate": 1, "error": 0,
    }
    no_confirmation = client.post(
        "/api/mappings/import/confirm", json={"token": mapping_preview_data["token"]}
    )
    assert no_confirmation.status_code == 400
    mapping_confirm = client.post(
        "/api/mappings/import/confirm",
        json={"token": mapping_preview_data["token"], "confirmed": True},
    )
    assert mapping_confirm.status_code == 200, mapping_confirm.get_data(as_text=True)
    assert mapping_confirm.get_json()["inserted"] == 2
    assert mapping_confirm.get_json()["processed"] == 2
    assert client.post(
        "/api/mappings/import/confirm",
        json={"token": mapping_preview_data["token"], "confirmed": True},
    ).status_code == 410

    mapping_repeat = preview_mapping(client, "invoice_names", invoice_mapping_file).get_json()
    assert mapping_repeat["counts"]["unchanged"] == 2 and mapping_repeat["counts"]["duplicate"] == 1
    mapping_repeat_confirm = client.post(
        "/api/mappings/import/confirm", json={"token": mapping_repeat["token"], "confirmed": True}
    )
    assert mapping_repeat_confirm.status_code == 200
    assert mapping_repeat_confirm.get_json()["unchanged"] == 2

    update_mapping_file = mapping_workbook_bytes(
        ["Mã HH", "Tên XHĐ"],
        [[mapping_products[0]["code"], "TÊN HÓA ĐƠN QC 01 - ĐÃ SỬA"]],
    )
    update_preview = preview_mapping(client, "invoice_names", update_mapping_file).get_json()
    assert update_preview["can_confirm"] is True and update_preview["counts"]["update"] == 1
    update_confirm = client.post(
        "/api/mappings/import/confirm", json={"token": update_preview["token"], "confirmed": True}
    )
    assert update_confirm.status_code == 200 and update_confirm.get_json()["updated"] == 1

    invalid_mapping_file = mapping_workbook_bytes(
        ["Mã SP", "Tên xuất hóa đơn"],
        [
            [mapping_products[2]["code"], "DÒNG HỢP LỆ KHÔNG ĐƯỢC GHI DỞ"],
            ["MA-KHONG-TON-TAI", "DÒNG LỖI"],
        ],
    )
    invalid_preview_response = preview_mapping(client, "invoice_names", invalid_mapping_file)
    assert invalid_preview_response.status_code == 200
    invalid_preview = invalid_preview_response.get_json()
    assert invalid_preview["can_confirm"] is False and invalid_preview["counts"]["error"] == 1
    assert client.post(
        "/api/mappings/import/confirm",
        json={"token": invalid_preview["token"], "confirmed": True},
    ).status_code == 400
    with server.db() as conn:
        assert not conn.execute(
            "SELECT 1 FROM outgoing_product_names WHERE product_code=?", (mapping_products[2]["code"],)
        ).fetchone(), "Không được ghi một phần khi file còn lỗi"

    conflict_mapping_file = mapping_workbook_bytes(
        ["Mã vật tư", "Invoice name"],
        [
            [mapping_products[2]["code"], "TÊN A"],
            [mapping_products[2]["code"], "TÊN B"],
        ],
    )
    conflict_preview = preview_mapping(client, "invoice_names", conflict_mapping_file).get_json()
    assert conflict_preview["can_confirm"] is False and conflict_preview["counts"]["error"] == 2

    kitchen_mapping_file = mapping_workbook_bytes(
        ["Tên bếp", "Mã bếp", "XCOM (xưởng cơm)"],
        [
            [mapping_kitchens[0]["name"], "", "XCOM-QC-01"],
            ["", mapping_kitchens[1]["code"], "xcom-qc-02"],
            [mapping_kitchens[0]["name"], "", "XCOM-QC-01"],
        ],
        sheet_name="Bếp - XCOM",
    )
    kitchen_preview_response = preview_mapping(client, "kitchen_units", kitchen_mapping_file)
    assert kitchen_preview_response.status_code == 200, kitchen_preview_response.get_data(as_text=True)
    kitchen_preview = kitchen_preview_response.get_json()
    assert kitchen_preview["can_confirm"] is True and kitchen_preview["counts"]["duplicate"] == 1
    kitchen_confirm = client.post(
        "/api/mappings/import/confirm", json={"token": kitchen_preview["token"], "confirmed": True}
    )
    assert kitchen_confirm.status_code == 200 and kitchen_confirm.get_json()["inserted"] == 2
    kitchen_list = client.get("/api/mappings/kitchen_units").get_json()["items"]
    assert any(item["code"] == mapping_kitchens[1]["code"] and item["target_value"] == "XCOM-QC-02" for item in kitchen_list)
    legacy_unit_file = mapping_workbook_bytes(
        ["Mã bếp", "Unit"], [[mapping_kitchens[1]["code"], "XCOM-QC-02"]], sheet_name="File cũ",
    )
    legacy_unit_preview = preview_mapping(client, "kitchen_units", legacy_unit_file)
    assert legacy_unit_preview.status_code == 200 and legacy_unit_preview.get_json()["can_confirm"] is True

    # Force a catalogue change between preview and confirmation to prove the SQL transaction rolls back.
    with server.db() as conn:
        conn.execute("INSERT INTO products(code,name) VALUES('QC-MAP-A','Hàng tạm QC A')")
        conn.execute("INSERT INTO products(code,name) VALUES('QC-MAP-B','Hàng tạm QC B')")
    rollback_file = mapping_workbook_bytes(
        ["Mã hàng", "Tên hóa đơn"],
        [["QC-MAP-A", "Tên tạm A"], ["QC-MAP-B", "Tên tạm B"]],
    )
    rollback_preview = preview_mapping(client, "invoice_names", rollback_file).get_json()
    assert rollback_preview["can_confirm"] is True
    with server.db() as conn:
        conn.execute("DELETE FROM products WHERE code='QC-MAP-B'")
    rollback_confirm = client.post(
        "/api/mappings/import/confirm", json={"token": rollback_preview["token"], "confirmed": True}
    )
    assert rollback_confirm.status_code == 409
    with server.db() as conn:
        assert not conn.execute(
            "SELECT 1 FROM outgoing_product_names WHERE product_code='QC-MAP-A'"
        ).fetchone(), "Transaction phải rollback toàn bộ nếu danh mục thay đổi giữa chừng"
        conn.execute("DELETE FROM products WHERE code IN ('QC-MAP-A','QC-MAP-B')")
        batch_audits = conn.execute(
            "SELECT COUNT(*) count FROM audit_log WHERE event_type='mapping.bulk_import' AND status='ok'"
        ).fetchone()["count"]
    assert batch_audits == 4
    print("MAPPING_IMPORT", json.dumps({
        "invoice": mapping_confirm.get_json(),
        "repeat": mapping_repeat_confirm.get_json(),
        "update": update_confirm.get_json(),
        "kitchen": kitchen_confirm.get_json(),
        "invalidBlocked": True,
        "rollbackVerified": True,
    }, ensure_ascii=False, indent=2))

    master_source = ROOT / "Em Thành.xlsx"
    assert client.post("/api/import").status_code == 410
    orphan_pending = APP_DIR / "data" / "pending_qc_orphan.xlsx"
    orphan_pending.write_bytes(b"orphaned preview")
    with master_source.open("rb") as handle:
        cancelled_analysis = client.post(
            "/api/import/analyze",
            data={"file": (handle, master_source.name)},
            content_type="multipart/form-data",
        )
    assert cancelled_analysis.status_code == 200
    assert not orphan_pending.exists()
    cancelled_token = cancelled_analysis.get_json()["token"]
    cancelled_path = APP_DIR / "data" / f"pending_{cancelled_token}.xlsx"
    assert cancelled_path.exists()
    assert client.post("/api/import/cancel", json={"token": cancelled_token}).status_code == 200
    assert not cancelled_path.exists()
    assert client.post("/api/import/cancel", json={"token": cancelled_token}).get_json()["idempotent"]
    assert client.post("/api/import/confirm", json={
        "token": cancelled_token, "work_date": "2026-08-27", "sheets": ["đơn hàng27.08 "],
    }).status_code == 400

    with master_source.open("rb") as handle:
        analyzed = client.post(
            "/api/import/analyze",
            data={"file": (handle, master_source.name)},
            content_type="multipart/form-data",
        )
    assert analyzed.status_code == 200, analyzed.get_data(as_text=True)
    analysis = analyzed.get_json()
    detected_names = {item["name"] for item in analysis["sheets"]}
    assert "đơn hàng27.08 " in detected_names
    selected = client.post("/api/import/confirm", json={
        "token": analysis["token"],
        "work_date": "2026-08-27",
        "sheets": ["đơn hàng27.08 "],
        "state_hash": analysis.get("stateHash", ""),
    })
    assert selected.status_code == 200, selected.get_data(as_text=True)
    # Ignore the single formula/formatted row whose calculated quantity is zero.
    assert len(selected.get_json()["orders"]) == 323

    source = ROOT / "Tách212223.xlsx"
    with source.open("rb") as handle:
        source_analysis_response = client.post(
            "/api/import/analyze",
            data={"file": (handle, source.name)},
            content_type="multipart/form-data",
        )
    assert source_analysis_response.status_code == 200
    source_analysis = source_analysis_response.get_json()
    qc_selected_sheets = [
        "ATV ko điều chỉnh", "NHUA không điều chỉnh", "Bảng chưa xuất", "ATV 25",
    ]
    assert set(qc_selected_sheets).issubset({item["name"] for item in source_analysis["sheets"]})
    response = client.post("/api/import/confirm", json={
        "token": source_analysis["token"],
        "work_date": "2026-08-28",
        "sheets": qc_selected_sheets,
    })
    assert response.status_code == 200, response.get_data(as_text=True)
    payload = response.get_json()
    batch_id = payload["batch"]["id"]
    orders = payload["orders"]
    assert orders, "Không đọc được dòng đơn nào"

    sheets = Counter(item["source_sheet"] for item in orders)
    error_types = Counter(error for item in orders for error in item["errors"])
    warning_types = Counter(warning for item in orders for warning in item["warnings"])
    assert sum(bool(item["errors"]) for item in orders) == 0
    assert sum(bool(item["warnings"]) for item in orders) == 27
    assert sum(any("95%" in warning for warning in item["warnings"]) for item in orders) == 12
    assert sum(any("bán lỗ" in warning for warning in item["warnings"]) for item in orders) == 15
    print("IMPORT", json.dumps({
        "orders": len(orders),
        "sheets": sheets,
        "errorRows": sum(bool(item["errors"]) for item in orders),
        "errorTypes": error_types,
        "warningRows": sum(bool(item["warnings"]) for item in orders),
        "warningTypes": warning_types,
    }, ensure_ascii=False, default=dict, indent=2))
    approved_old_file = client.post(f"/api/batches/{batch_id}/approve")
    assert approved_old_file.status_code == 200, approved_old_file.get_data(as_text=True)
    with server.db() as conn:
        bk_auto_rows = conn.execute(
            "SELECT COUNT(*) n FROM inventory_transactions "
            "WHERE source_type='BK_INPUT' AND source_id=?",
            (str(batch_id),),
        ).fetchone()["n"]
    assert bk_auto_rows == 0

    mapped_code = orders[0]["product_code"]
    mapped_invoice_name = "TEN XUAT HOA DON QC"
    mapped_name_response = client.put(
        f"/api/outgoing-product-names/{mapped_code}", json={"invoice_name": mapped_invoice_name}
    )
    assert mapped_name_response.status_code == 200
    export_sizes = {}
    blocked_invoice_export = client.get(f"/api/export/invoices/{batch_id}")
    assert blocked_invoice_export.status_code == 409
    assert "tạo dự thảo" in blocked_invoice_export.get_json()["error"]
    for kind in ("suppliers", "deliveries", "report", "purchases", "outgoing-statement"):
        result = client.get(f"/api/export/{kind}/{batch_id}")
        if kind == "purchases":
            # This broad legacy fixture assigns more than 5m to one seller in
            # one day.  The golden receipt explicitly forbids that case, so
            # the combined Bảng kê & biên nhận export must fail closed.  A
            # valid positive export is exercised by run_receipt_template_qc.
            assert result.status_code == 422, result.get_data(as_text=True)
            assert result.get_json()["code"] == "receipt_daily_limit_exceeded"
            export_sizes[kind] = "blocked>5m"
            continue
        assert result.status_code == 200, (kind, result.get_data(as_text=True))
        assert len(result.data) > 1000, kind
        export_sizes[kind] = len(result.data)
        wb = load_workbook(io.BytesIO(result.data), data_only=True)
        assert wb.sheetnames
        if kind == "outgoing-statement":
            assert wb.sheetnames[0] == "Tổng hợp"
            assert wb["Tổng hợp"]["A1"].value == "BẢNG KÊ HÀNG HÓA ĐẦU RA"
            assert wb["Tổng hợp"]["A4"].value == "Nhà thầu"
        wb.close()
    mapped_order = dict(orders[0])
    mapped_order["product_name"] = mapped_invoice_name
    mapped_book = server.invoice_workbook([mapped_order])
    assert mapped_book.active.cell(2, 2).value == mapped_invoice_name
    mapped_book.close()

    # VND must use financial HALF_UP rounding and explicit nature 2 must blank
    # every money column in the customer's exact 13-column layout.
    rounding_book = server.invoice_workbook([{
        "product_code": "QC-ROUND", "product_name": "QC rounding", "unit": "kg",
        "actual_delivered": 0.5, "customer_return_qty": 0, "sell_price": 1,
        "tax": "0%", "invoice_nature": "1", "note": "",
    }])
    rounding_sheet = rounding_book.active
    assert rounding_sheet.cell(2, 6).value == 1
    rounding_book.close()
    promotion_book = server.invoice_workbook([{
        "product_code": "QC-PROMO", "product_name": "QC promo", "unit": "chai",
        "actual_delivered": 1, "customer_return_qty": 0, "sell_price": 0,
        "tax": "10%", "invoice_nature": "2", "note": "",
    }])
    promotion_sheet = promotion_book.active
    assert promotion_sheet.cell(2, 13).value == "2"
    assert all(promotion_sheet.cell(2, col).value is None for col in (5, 6, 8, 9, 11, 12))
    promotion_book.close()

    # Manual-flow test: new batch, add, edit, approve.
    new_batch = client.post("/api/batches", json={"work_date": "2026-08-29"}).get_json()
    manual_id = new_batch["batch"]["id"]
    add = client.post("/api/orders", json={
        "batch_id": manual_id,
        "work_date": "2026-08-29",
        "kitchen": "POT",
        "product_code": "I000060",
        "qty": 3,
        "actual_received": 3,
        "actual_delivered": 2.8,
        "buy_price": 14000,
        "sell_price": 16000,
        "tax": "KKKNT",
        "supplier": "kho",
    })
    assert add.status_code == 200, add.get_data(as_text=True)
    item = add.get_json()["orders"][0]
    assert item["contractor"] == "HATRAN"
    assert item["product_name"]
    update = client.put(f"/api/orders/{item['id']}", json={"actual_delivered": 3})
    assert update.status_code == 200
    bulk = client.post("/api/orders/bulk", json={
        "batch_id": manual_id,
        "text": "Mã bếp\tMã hàng\tSố lượng\tNCC\tGiá mua\tGiá bán\tThuế\n"
                "POT\tI000060\t2\tkho\t14000\t16000\tKKKNT",
    })
    assert bulk.status_code == 200, bulk.get_data(as_text=True)
    assert bulk.get_json()["inserted"] == 1
    promotion = client.post("/api/orders", json={
        "batch_id": manual_id,
        "work_date": "2026-08-29",
        "kitchen": "POT",
        "product_code": "I000060",
        "qty": 1,
        "actual_received": 1,
        "actual_delivered": 1,
        "buy_price": 0,
        "sell_price": 0,
        "tax": "10%",
        "invoice_nature": "2",
        "supplier": "kho",
        "note": "Explicit promotion QC",
    })
    assert promotion.status_code == 200, promotion.get_data(as_text=True)
    promotion_row = next(
        row for row in promotion.get_json()["orders"]
        if row["invoice_nature"] == "2" and not row["errors"]
    )
    rejected_grid = client.put("/api/orders/bulk-update", json={
        "batch_id": manual_id,
        "items": [{"id": item["id"], "sell_price": 15555}, {"id": 999999999, "sell_price": 1}],
    })
    assert rejected_grid.status_code == 409
    unchanged_after_reject = client.get(f"/api/bootstrap?batch_id={manual_id}").get_json()
    assert next(row for row in unchanged_after_reject["orders"] if row["id"] == item["id"])["sell_price"] == 16000
    saved_grid = client.put("/api/orders/bulk-update", json={
        "batch_id": manual_id,
        "items": [
            {"id": item["id"], "actual_delivered": 3, "sell_price": 16000},
            {"id": promotion_row["id"], "actual_delivered": 1, "sell_price": 0,
             "invoice_nature": "2", "tax": "0.10"},
        ],
    })
    assert saved_grid.status_code == 200, saved_grid.get_data(as_text=True)
    assert saved_grid.get_json()["updated"] == 2 and saved_grid.get_json()["error_rows"] == 0
    approved = client.post(f"/api/batches/{manual_id}/approve")
    assert approved.status_code == 200, approved.get_data(as_text=True)
    assert approved.get_json()["receivable_ledger"]["active"] >= 2
    receivable_ledger = client.get(
        "/api/debts/receivables/ledger?from=2026-08-29&to=2026-08-29"
        "&contractor=HATRAN&kitchen=POT&status=all"
    )
    assert receivable_ledger.status_code == 200, receivable_ledger.get_data(as_text=True)
    receivable_payload = receivable_ledger.get_json()
    assert receivable_payload["source_of_truth"] == (
        "approved operational delivery lines; not issued VAT invoices"
    )
    assert receivable_payload["summary"] == {
        "source_rows": 3,
        "active_rows": 2,
        "status_counts": {"active": 2, "reversed": 1},
        "subtotal": 80000,
        "tax_amount": 0,
        "charge_amount": 80000,
        "quantity": 5,
        "filtered_quantity": 6,
        "filtered_subtotal": 80000,
        "filtered_tax_amount": 0,
        "filtered_amount": 80000,
    }
    assert sorted(row["delivered_qty"] for row in receivable_payload["rows"]) == [1, 2, 3]
    assert next(
        row for row in receivable_payload["rows"] if row["source"]["id"] == promotion_row["id"]
    )["reversal_reason"] == "non_chargeable_line"

    payment = client.post("/api/payments", json={
        "payment_date": "2026-08-29", "kind": "receipt", "party_type": "contractor",
        "party_code": "HATRAN", "amount": 10000, "note": "QC",
    })
    assert payment.status_code == 200
    balance = client.post("/api/balances", json={
        "party_type": "contractor", "party_code": "HATRAN", "opening": 5000,
    })
    assert balance.status_code == 200
    checked = client.get(f"/api/bootstrap?batch_id={manual_id}").get_json()
    assert checked["summary"]["scope"] == "batch_only"
    assert checked["summary"]["contractors"]["HATRAN"]["opening"] == 0
    assert checked["summary"]["contractors"]["HATRAN"]["paid"] == 0
    debt_adjustment = client.post("/api/debt-adjustments", json={
        "adjustment_date": "2026-08-29", "party_type": "contractor",
        "party_code": "HATRAN", "amount": 2500, "note": "QC điều chỉnh",
    })
    assert debt_adjustment.status_code == 200
    debt_period = client.get("/api/debts?from=2026-08-01&to=2026-08-31")
    assert debt_period.status_code == 200
    assert debt_period.get_json()["contractors"]["HATRAN"]["opening"] == 5000
    assert debt_period.get_json()["contractors"]["HATRAN"]["period_paid"] == 10000
    assert debt_period.get_json()["contractors"]["HATRAN"]["period_adjustment"] == 2500
    assert client.get("/api/debts?from=2026-09-01&to=2026-08-31").status_code == 400
    debt_export = client.get("/api/export/debts?from=2026-08-01&to=2026-08-31")
    assert debt_export.status_code == 200 and len(debt_export.data) > 1000
    debt_book = load_workbook(io.BytesIO(debt_export.data), data_only=True)
    assert debt_book.sheetnames == [
        "Phải thu", "Phải trả", "Thu chi", "Điều chỉnh", "Chi tiết phải trả cũ",
    ]
    assert debt_book["Phải thu"]["A3"].value == "Đối tượng"
    assert any(row[0].value == "HATRAN" and row[3].value == 2500 for row in debt_book["Phải thu"].iter_rows(min_row=4))
    debt_book.close()
    receivable_export_response = client.get(
        "/api/debts/receivables/export?from=2026-08-29&to=2026-08-29&contractor=HATRAN"
    )
    assert receivable_export_response.status_code == 200
    receivable_export = load_workbook(
        io.BytesIO(receivable_export_response.data), data_only=False, keep_links=False,
    )
    assert receivable_export.sheetnames[0] == "Tổng nhà thầu"
    assert len(receivable_export.sheetnames) == 2
    receivable_summary = receivable_export["Tổng nhà thầu"]
    opening_receivable = receivable_summary.cell(4, 1).value
    assert [receivable_summary.cell(4, column).value for column in range(2, 5)] == [
        80000, 2500, 10000,
    ]
    assert receivable_summary.cell(4, 5).value == opening_receivable + 80000 + 2500 - 10000
    receivable_detail = receivable_export.worksheets[1]
    assert receivable_summary.cell(receivable_summary.max_row, 10).value == 80000
    assert receivable_detail.cell(receivable_detail.max_row, 15).value == 80000
    assert not any(
        isinstance(cell.value, str) and cell.value.startswith("=")
        for worksheet in receivable_export.worksheets
        for row in worksheet.iter_rows()
        for cell in row
    )
    receivable_export.close()
    receivable_export_result = "one workbook per contractor / summary + kitchen totals reconciled"

    # Returns/damage: cost and revenue must use net received/net delivered.
    first_item = checked["orders"][0]
    returned = client.put(f"/api/orders/{first_item['id']}", json={
        "actual_received": 3, "damaged_qty": 0.2, "supplier_return_qty": 0.3,
        "actual_delivered": 3, "customer_return_qty": 0.4,
    })
    assert returned.status_code == 200, returned.get_data(as_text=True)
    returned_item = next(row for row in returned.get_json()["orders"] if row["id"] == first_item["id"])
    assert round(returned_item["cost"]) == round(2.5 * returned_item["buy_price"])
    assert round(returned_item["revenue"]) == round(2.6 * returned_item["sell_price"])
    reapproved = client.post(f"/api/batches/{manual_id}/approve")
    assert reapproved.status_code == 200
    receivable_after_return = client.get(
        "/api/debts/receivables/ledger?from=2026-08-29&to=2026-08-29"
        "&contractor=HATRAN&kitchen=POT&status=all"
    ).get_json()
    returned_ledger_line = next(
        row for row in receivable_after_return["rows"] if row["source"]["id"] == first_item["id"]
    )
    assert returned_ledger_line["customer_return_qty"] == 0.4
    assert returned_ledger_line["delivered_qty"] == 2.6
    assert returned_ledger_line["amount"] == 41600
    assert receivable_after_return["summary"]["charge_amount"] == 73600
    receivable_qc_result = (
        "approved net delivery / transaction price / return / tax / revision / period filters passed"
    )

    # Accounting inventory never auto-subtracts the NCC order. Physical cabinet
    # stock is entered by the user in the separate purchase-order roundtrip.
    opening = client.post("/api/inventory/opening", json={
        "period": "2026-08", "items": [{"product_code": "I000060", "qty": 50, "unit_cost": 14000}],
    })
    assert opening.status_code == 200, opening.get_data(as_text=True)
    needs = client.get(f"/api/supplier-needs/{manual_id}")
    assert needs.status_code == 200, needs.get_data(as_text=True)
    need_data = needs.get_json()
    assert need_data["required_qty"] == need_data["ordered_qty"]
    assert need_data["physical_stock_used"] == 0
    locked_supplier_rule = client.put(
        "/api/supplier-rules/kho", json={"combine_kitchens": True}
    )
    assert locked_supplier_rule.status_code == 409
    supplier_rule = client.put(
        "/api/supplier-rules/huong", json={"combine_kitchens": True}
    )
    assert supplier_rule.status_code == 200

    # Outgoing invoice drafts reserve stock, but only user confirmation posts output.
    assert client.put(
        "/api/outgoing-product-names/I000060", json={"invoice_name": "HÀNH TÂY XUẤT HĐ QC"}
    ).status_code == 200
    drafted = client.post(f"/api/outgoing-invoices/draft/{manual_id}")
    assert drafted.status_code == 200, drafted.get_data(as_text=True)
    draft_id = drafted.get_json()["drafts"][0]["id"]
    drafted_again = client.post(f"/api/outgoing-invoices/draft/{manual_id}")
    assert drafted_again.status_code == 200, drafted_again.get_data(as_text=True)
    assert drafted_again.get_json()["drafts"][0]["id"] == draft_id
    locked_edit = client.put(f"/api/orders/{first_item['id']}", json={"qty": 99})
    assert locked_edit.status_code == 409 and "dự thảo" in locked_edit.get_json()["error"].lower()
    with server.db() as conn:
        assert conn.execute(
            "SELECT COUNT(*) n FROM outgoing_invoice_lines WHERE draft_id=? AND product_code='I000060' AND product_name='HÀNH TÂY XUẤT HĐ QC'",
            (draft_id,),
        ).fetchone()["n"] >= 1
        assert conn.execute(
            "SELECT COUNT(*) n FROM inventory_transactions WHERE source_type='OUTGOING_DRAFT' AND source_id=? AND status='reserved'",
            (str(draft_id),),
        ).fetchone()["n"] == 3
        promotion_line = conn.execute(
            "SELECT invoice_nature,amount FROM outgoing_invoice_lines WHERE draft_id=? AND invoice_nature='2'",
            (draft_id,),
        ).fetchone()
        assert promotion_line and promotion_line["amount"] == 0

    invoice_export = client.get(f"/api/export/invoices/{manual_id}")
    assert invoice_export.status_code == 200, invoice_export.get_data(as_text=True)
    assert len(invoice_export.data) > 1000
    export_sizes["invoices"] = len(invoice_export.data)
    with zipfile.ZipFile(io.BytesIO(invoice_export.data)) as archive:
        invoice_files = [name for name in archive.namelist() if name.endswith(".xlsx")]
        assert invoice_files
        promotion_seen = False
        for name in invoice_files:
            assert "_lan_1_" in name, name
            invoice_book = load_workbook(io.BytesIO(archive.read(name)), data_only=False, keep_links=False)
            invoice_sheet = invoice_book.active
            assert len(invoice_book.sheetnames) == 1
            if "_KKKNT_" in name:
                assert invoice_sheet.title == "Sheet2"
            elif "_VAT8_" in name or "_VAT10" in name:
                assert invoice_sheet.title == "Sheet 1 (2)"
            assert [invoice_sheet.cell(1, col).value for col in range(1, 14)] == server.INVOICE_HEADERS
            for row in range(2, invoice_sheet.max_row + 1):
                assert invoice_sheet.cell(row, 13).value in {"1", "2"}
                if invoice_sheet.cell(row, 13).value == "2":
                    promotion_seen = True
                    assert all(invoice_sheet.cell(row, col).value is None for col in (5, 6, 7, 8, 9, 11, 12))
                else:
                    assert round(invoice_sheet.cell(row, 6).value) == round(
                        invoice_sheet.cell(row, 4).value * invoice_sheet.cell(row, 5).value
                    )
                assert not any(
                    invoice_sheet.cell(row, col).data_type == "f" or invoice_sheet.cell(row, col).hyperlink
                    for col in range(1, 14)
                )
            invoice_book.close()
        assert promotion_seen
        invoice_tax_export_result = (
            "draft-line quantities / contractor+round+tax split / locked golden styles / promo blanks passed"
        )

    # Exercise the exact UI/API draft payload without any remote write.  The
    # production client still performs all validation and payload construction.
    assert client.put("/api/outgoing-buyers/HATRAN", json={
        "display_name": "QC Buyer", "legal_name": "CÔNG TY QC BUYER",
        "tax_code": "0200000000", "address": "QC address",
    }).status_code == 200
    captured_minvoice = {}

    class DryRunMinvoice:
        def create_draft(self, payload, dry_run=True, confirm_remote_write=False):
            assert dry_run is True and confirm_remote_write is False
            captured_minvoice["source"] = payload
            built = server.MinvoiceClient(
                server.MinvoiceConfig("https://example.invalid", "qc", "qc")
            ).build_draft_payload(payload)
            captured_minvoice["built"] = built
            return {"ok": True, "dry_run": True, "remote_write": False, "payload": built}

    server.app.config["MINVOICE_CLIENT_FACTORY"] = DryRunMinvoice
    try:
        minvoice_dry_run = client.post(
            f"/api/minvoice/drafts/{draft_id}",
            json={"series": "1C26TDP", "dry_run": True},
        )
    finally:
        server.app.config.pop("MINVOICE_CLIENT_FACTORY", None)
    assert minvoice_dry_run.status_code == 200, minvoice_dry_run.get_data(as_text=True)
    built_lines = captured_minvoice["built"]["data"][0]["details"][0]["data"]
    assert any(line["tchat"] == 2 and line["inv_TotalAmount"] == 0 for line in built_lines)
    assert all(line["tchat"] in {1, 2} for line in built_lines)
    cancelled = client.post(f"/api/outgoing-invoices/{draft_id}/cancel", json={"confirmed": True})
    assert cancelled.status_code == 200
    assert client.post(
        f"/api/outgoing-invoices/{draft_id}/confirm-issued", json={"confirmed": True}
    ).status_code == 409
    with server.db() as conn:
        assert conn.execute(
            "SELECT COUNT(*) n FROM inventory_transactions WHERE source_type='OUTGOING_DRAFT' AND source_id=? AND status='reserved'",
            (str(draft_id),),
        ).fetchone()["n"] == 0
    recreated = client.post(f"/api/outgoing-invoices/draft/{manual_id}")
    assert recreated.status_code == 200 and recreated.get_json()["drafts"][0]["status"] == "draft"
    # Cancelling a round keeps its audit record. Recreating now opens a new
    # round, so all following actions must target the newly-created draft.
    draft_id = recreated.get_json()["drafts"][0]["id"]
    assert client.post(
        f"/api/outgoing-invoices/{draft_id}/confirm-issued", json={"confirmed": True}
    ).status_code == 400
    issued_payload = {
        "confirmed": True, "invoice_number": "00001234",
        "invoice_series": "1C26TDP", "invoice_date": "2026-08-29",
    }
    issued = client.post(f"/api/outgoing-invoices/{draft_id}/confirm-issued", json=issued_payload)
    assert issued.status_code == 200
    issued_repeat = client.post(f"/api/outgoing-invoices/{draft_id}/confirm-issued", json=issued_payload)
    assert issued_repeat.status_code == 200 and issued_repeat.get_json()["idempotent"] is True
    invoice_payment_scope = client.get(
        "/api/outgoing-invoices/payment-scope/HATRAN?from=2026-08-29&to=2026-08-29"
    )
    assert invoice_payment_scope.status_code == 200, invoice_payment_scope.get_data(as_text=True)
    invoice_payment_scope_data = invoice_payment_scope.get_json()
    assert invoice_payment_scope_data["template_status"] == "official_customer_xlsx"
    assert invoice_payment_scope_data["official_template_ready"] is True
    assert invoice_payment_scope_data["totals"]["total_amount"] > 0
    invoice_delivery_statement = client.get(
        "/api/outgoing-invoices/delivery-statement/HATRAN?from=2026-08-29&to=2026-08-29&scope_id="
        + invoice_payment_scope_data["scope_id"]
    )
    assert invoice_delivery_statement.status_code == 200
    assert invoice_delivery_statement.headers["X-TDP-Invoice-Scope"] == invoice_payment_scope_data["scope_id"]
    delivery_statement_book = load_workbook(
        io.BytesIO(invoice_delivery_statement.data), data_only=False, keep_links=False,
    )
    assert delivery_statement_book.sheetnames == ["Bảng kê giao hàng", "Đối chiếu hóa đơn"]
    assert "BẢNG TỔNG HỢP GIAO NHẬN" in delivery_statement_book["Bảng kê giao hàng"]["A5"].value
    assert delivery_statement_book["Bảng kê giao hàng"]["A10"].value == "STT"
    delivery_differences = [
        row[9].value for row in delivery_statement_book["Đối chiếu hóa đơn"].iter_rows(min_row=4)
        if row[2].value
    ]
    assert delivery_differences and all(abs(value or 0) <= 1 for value in delivery_differences)
    assert not any(
        cell.data_type == "f" or cell.hyperlink is not None
        for sheet in delivery_statement_book.worksheets
        for row in sheet.iter_rows() for cell in row
    )
    delivery_statement_book.close()
    invoice_payment_bundle = client.get(
        "/api/export/invoice-payment-bundle/HATRAN?from=2026-08-29&to=2026-08-29&scope_id="
        + invoice_payment_scope_data["scope_id"]
    )
    assert invoice_payment_bundle.status_code == 200, invoice_payment_bundle.get_data(as_text=True)
    assert invoice_payment_bundle.headers["X-TDP-Template-Status"] == "official-customer-xlsx"
    with zipfile.ZipFile(io.BytesIO(invoice_payment_bundle.data)) as archive:
        names = archive.namelist()
        payment_name = next(name for name in names if name.startswith("De_nghi_thanh_toan"))
        statement_name = next(name for name in names if name.startswith("Bang_tong_hop_giao_nhan"))
        assert "THONG_TIN_DOI_CHIEU.txt" in names
        manifest = archive.read("THONG_TIN_DOI_CHIEU.txt").decode("utf-8")
        assert "biểu mẫu chính thức" in manifest
        assert "CHƯA PHẢI MẪU" not in manifest
        payment_book = load_workbook(io.BytesIO(archive.read(payment_name)), data_only=True, keep_links=False)
        assert payment_book.sheetnames == ["Đề nghị thanh toán", "Đối chiếu hóa đơn"]
        assert payment_book["Đề nghị thanh toán"]["A6"].value == "ĐỀ NGHỊ THANH TOÁN"
        assert payment_book["Đề nghị thanh toán"]["C15"].value == "00001234"
        assert payment_book["Đề nghị thanh toán"]["F15"].value > 0
        assert payment_book["Đối chiếu hóa đơn"].sheet_state == "hidden"
        payment_book.close()
        statement_book = load_workbook(io.BytesIO(archive.read(statement_name)), data_only=True)
        assert statement_book.sheetnames == ["Bảng kê giao hàng", "Đối chiếu hóa đơn"]
        assert "BẢNG TỔNG HỢP GIAO NHẬN" in statement_book["Bảng kê giao hàng"]["A5"].value
        differences = [
            row[9].value for row in statement_book["Đối chiếu hóa đơn"].iter_rows(min_row=4)
            if row[2].value
        ]
        assert differences and all(abs(value or 0) <= 1 for value in differences)
        statement_book.close()
    assert client.post(f"/api/outgoing-invoices/{draft_id}/cancel", json={"confirmed": True}).status_code == 409

    # mSMI incremental/idempotent sync with a mock response; no external write.
    class DummyMsmi:
        def list_invoices(self, **kwargs):
            if kwargs.get("page", 0) > 0:
                return {"items": [], "has_more": False}
            return {"items": [{
                "_id": "QC-MSMI-001", "mstNban": "0200000001", "tenNban": "NCC QC",
                "shdon": "1", "khhdon": "C26TQC", "tdlap": "2026-08-29",
                "tgtcthue": 28000, "tgtthue": 0, "tgtttbso": 28000,
                "hdhhdvu": [{"ma": "SRC-01", "ten": "Hành tây QC", "dvtinh": "kg",
                               "sluong": 2, "dgia": 14000, "thtien": 28000, "tsuat": "KKKNT"}],
            }], "has_more": False}

    with server.db() as conn:
        synced_once = contract_modules.sync_msmi(conn, DummyMsmi(), server.now_iso, tenant="TDP")
    with server.db() as conn:
        synced_twice = contract_modules.sync_msmi(conn, DummyMsmi(), server.now_iso, tenant="TDP")
    assert synced_once["new_invoices"] == 1
    assert synced_twice["new_invoices"] == 0 and synced_twice["known_invoices"] == 1
    msmi_list = client.get("/api/msmi/invoices").get_json()["items"]
    msmi_invoice = next(row for row in msmi_list if row["remote_id"] == "QC-MSMI-001")
    msmi_item_id = msmi_invoice["items"][0]["id"]
    mapped = client.put(f"/api/msmi/items/{msmi_item_id}/mapping", json={"product_code": "I000060"})
    assert mapped.status_code == 200, mapped.get_data(as_text=True)
    receipt_once = client.post(f"/api/msmi/invoices/{msmi_invoice['id']}/receipt")
    receipt_twice = client.post(f"/api/msmi/invoices/{msmi_invoice['id']}/receipt")
    assert receipt_once.status_code == 200 and receipt_once.get_json()["new_inventory_lines"] == 1
    assert receipt_twice.status_code == 200 and receipt_twice.get_json()["idempotent"]

    # Customer-style kitchen workbook: preview, explicit confirmation, exact quantities and idempotent re-import.
    kitchen_file = kitchen_workbook_bytes()
    kitchen_file_preview = preview_kitchen(client, kitchen_file)
    assert kitchen_file_preview.status_code == 200, kitchen_file_preview.get_data(as_text=True)
    kitchen_file_data = kitchen_file_preview.get_json()
    assert kitchen_file_data["can_confirm"] is True
    assert kitchen_file_data["counts"]["plans"] == 2 and kitchen_file_data["counts"]["items"] == 2
    assert [(plan["kitchen"], plan["xcom_code"]) for plan in kitchen_file_data["plans"]] == [
        ("VINA", "XCOM"), ("MAZDA", "MAZDA"),
    ]
    assert client.post(
        "/api/kitchen/import/confirm", json={"token": kitchen_file_data["token"]}
    ).status_code == 400
    kitchen_file_confirm = client.post(
        "/api/kitchen/import/confirm",
        json={"token": kitchen_file_data["token"], "confirmed": True},
    )
    assert kitchen_file_confirm.status_code == 200, kitchen_file_confirm.get_data(as_text=True)
    assert kitchen_file_confirm.get_json()["inserted"] == 2
    imported_plans = client.get("/api/kitchen/plans?date=2026-09-02").get_json()["items"]
    assert len(imported_plans) == 2
    mazda_egg = next(
        item for plan in imported_plans if plan["kitchen"] == "MAZDA"
        for item in plan["items"] if item["product_code"] == "H000007"
    )
    assert mazda_egg["required_qty"] == 8 and mazda_egg["norm_qty"] == 2
    kitchen_file_po = client.get("/api/kitchen/po?date=2026-09-02")
    assert kitchen_file_po.status_code == 200 and len(kitchen_file_po.data) > 1000
    kitchen_file_po_book = load_workbook(io.BytesIO(kitchen_file_po.data), data_only=True)
    assert any("XCOM " in str(sheet.cell(2, 4).value) for sheet in kitchen_file_po_book.worksheets)
    kitchen_file_po_book.close()
    kitchen_repeat_data = preview_kitchen(client, kitchen_file).get_json()
    assert kitchen_repeat_data["counts"]["new"] == 0 and kitchen_repeat_data["counts"]["update"] == 2
    kitchen_repeat_confirm = client.post(
        "/api/kitchen/import/confirm",
        json={"token": kitchen_repeat_data["token"], "confirmed": True},
    )
    assert kitchen_repeat_confirm.status_code == 200 and kitchen_repeat_confirm.get_json()["updated"] == 2
    assert len(client.get("/api/kitchen/plans?date=2026-09-02").get_json()["items"]) == 2
    invalid_kitchen_data = preview_kitchen(
        client, kitchen_workbook_bytes("MA-KHONG-TON-TAI"), work_date="2026-09-03"
    ).get_json()
    assert invalid_kitchen_data["can_confirm"] is False and invalid_kitchen_data["counts"]["errors"] == 1
    assert client.post(
        "/api/kitchen/import/confirm",
        json={"token": invalid_kitchen_data["token"], "confirmed": True},
    ).status_code == 400
    assert not client.get("/api/kitchen/plans?date=2026-09-03").get_json()["items"]
    assert preview_kitchen(client, kitchen_file, work_date="2026-02-30").status_code == 400
    assert preview_kitchen(client, b"not-an-xlsx", work_date="2026-09-03").status_code == 400
    duplicate_kitchen = preview_kitchen(
        client, kitchen_workbook_bytes(duplicate_first=True), work_date="2026-09-03"
    ).get_json()
    assert duplicate_kitchen["can_confirm"] is True
    assert duplicate_kitchen["counts"]["errors"] == 0
    assert duplicate_kitchen["counts"]["warnings"] >= 1
    assert duplicate_kitchen["counts"]["items"] == 3

    # Opening-import is an authoritative period snapshot: a revised file that
    # removes a code must remove the old OPENING line instead of leaving stock.
    opening_snapshot_a = opening_workbook_bytes([
        ["I000060", "Hành tây", "", "I000060", "KKKNT", "kg", 2, 14000, 28000],
        ["H000007", "Trứng gà", "", "H000007", "KKKNT", "quả", 3, 3000, 9000],
    ])
    opening_a = preview_opening(client, opening_snapshot_a, period="2026-06").get_json()
    assert opening_a["can_confirm"] is True
    assert client.post("/api/inventory/opening/import/confirm", json={
        "token": opening_a["token"], "confirmed": True,
    }).status_code == 200
    opening_snapshot_b = opening_workbook_bytes([
        ["I000060", "Hành tây", "", "I000060", "KKKNT", "kg", 5, 14000, 70000],
    ])
    opening_b = preview_opening(
        client, opening_snapshot_b, period="2026-06", filename="revised-renamed-opening.xlsx"
    ).get_json()
    opening_b_confirm = client.post("/api/inventory/opening/import/confirm", json={
        "token": opening_b["token"], "confirmed": True,
    })
    assert opening_b_confirm.status_code == 200
    assert opening_b_confirm.get_json()["deleted_stale"] == 1
    with server.db() as conn:
        june_opening = conn.execute(
            "SELECT product_code,qty_in FROM inventory_transactions "
            "WHERE source_type='OPENING' AND source_id='2026-06'"
        ).fetchall()
        assert [(row["product_code"], row["qty_in"]) for row in june_opening] == [("I000060", 5)]

    with server.db() as conn:
        conn.execute(
            """INSERT INTO products(code,name,unit,supplier,buy_price)
               VALUES('QC-XCOM-ROLLBACK','Hàng rollback xưởng cơm','kg','QC',10000)"""
        )
        conn.execute(
            """INSERT INTO product_prices(product_code,price_group,price_text,price_value)
               VALUES('QC-XCOM-ROLLBACK','HATRAN','10000',10000)"""
        )
    rollback_kitchen = preview_kitchen(
        client, kitchen_workbook_bytes("QC-XCOM-ROLLBACK"), work_date="2026-09-04"
    ).get_json()
    assert rollback_kitchen["can_confirm"] is True
    with server.db() as conn:
        conn.execute("DELETE FROM products WHERE code='QC-XCOM-ROLLBACK'")
    rollback_kitchen_confirm = client.post(
        "/api/kitchen/import/confirm",
        json={"token": rollback_kitchen["token"], "confirmed": True},
    )
    assert rollback_kitchen_confirm.status_code == 409
    assert not client.get("/api/kitchen/plans?date=2026-09-04").get_json()["items"]

    actual_kitchen_source = ROOT / "_HANDOFF" / "EXTERNAL_INPUTS" / "xưởng cơm.xlsx"
    actual_kitchen_result = "not_present"
    if actual_kitchen_source.exists():
        assert hashlib.sha256(actual_kitchen_source.read_bytes()).hexdigest().upper() == (
            "4CCF1957C53DE932EF87EFEAEBBD398A3F45A60161FC04409BC6A471664ACE8B"
        )
        with actual_kitchen_source.open("rb") as handle:
            actual_preview_response = client.post(
                "/api/kitchen/import/preview",
                data={"work_date": "2026-09-01", "file": (handle, actual_kitchen_source.name)},
                content_type="multipart/form-data",
            )
        assert actual_preview_response.status_code == 200, actual_preview_response.get_data(as_text=True)
        actual_preview = actual_preview_response.get_json()
        assert actual_preview["can_confirm"] is True
        assert actual_preview["counts"]["plans"] == 5 and actual_preview["counts"]["items"] == 54
        assert [(plan["kitchen"], plan["xcom_code"], plan["meal_count"]) for plan in actual_preview["plans"]] == [
            ("VINA", "XCOM", 28), ("SUNBY", "XCOM", 32), ("DAINAM", "XCOM", 48),
            ("MAZDA", "MAZDA", 46), ("TTS", "TTS", 40),
        ]
        mazda_preview = next(plan for plan in actual_preview["plans"] if plan["kitchen"] == "MAZDA")
        assert mazda_preview["menu_count"] == 2 and mazda_preview["servings_per_menu"] == 23
        assert mazda_preview["meal_price"] == 35000
        actual_confirm = client.post(
            "/api/kitchen/import/confirm",
            json={"token": actual_preview["token"], "confirmed": True},
        )
        assert actual_confirm.status_code == 200 and actual_confirm.get_json()["inserted"] == 5
        actual_plans = client.get("/api/kitchen/plans?date=2026-09-01").get_json()["items"]
        assert len(actual_plans) == 5 and sum(len(plan["items"]) for plan in actual_plans) == 54
        assert all(
            item["price_source"] == "HATRAN 2026-09 · giá kỳ đã khóa"
            for plan in actual_plans for item in plan["items"]
        ), "Nguồn giá sau xác nhận phải được canonical hóa ngay từ lần nhập đầu"
        actual_by_kitchen = {plan["kitchen"]: plan for plan in actual_plans}
        expected_financials = {
            "VINA": (28, 30000, 503580, 297962.962962963),
            "SUNBY": (32, 25000, 488000, 338814.814814815),
            "DAINAM": (48, 25000, 710880, 502222.222222222),
            "MAZDA": (46, 35000, 821100, 576538.461538462),
            "TTS": (40, 40000, 1327640, 190909.090909091),
        }
        for kitchen, (meals, price, food_cost, other_cost) in expected_financials.items():
            plan = actual_by_kitchen[kitchen]
            assert plan["meal_count"] == meals and plan["meal_price"] == price
            assert abs(plan["food_cost"] - food_cost) < 0.01
            assert abs(plan["other_cost"] - other_cost) < 0.01
            assert abs(plan["revenue"] - meals * price) < 0.01
            assert abs(plan["total_cost"] - (food_cost + other_cost)) < 0.01
            assert abs(plan["profit"] - (meals * price - food_cost - other_cost)) < 0.01
        assert client.put("/api/kitchen-units/VINA", json={"xcom_code": "XCOM-NEW"}).status_code == 200
        historic_vina = next(
            plan for plan in client.get("/api/kitchen/plans?date=2026-09-01").get_json()["items"]
            if plan["kitchen"] == "VINA"
        )
        assert historic_vina["mapped_unit"] == "XCOM", "Kế hoạch cũ phải giữ XCOM tại thời điểm nhập"
        assert client.put("/api/kitchen-units/VINA", json={"xcom_code": "XCOM"}).status_code == 200
        mazda_items = actual_by_kitchen["MAZDA"]["items"]
        assert {round(item["applicable_meal_count"]) for item in mazda_items} == {23, 46}
        assert next(item for item in mazda_items if item["product_code"] == "H000007")["required_qty"] == 46

        actual_po_response = client.get("/api/kitchen/po?date=2026-09-01")
        assert actual_po_response.status_code == 200 and len(actual_po_response.data) > 1000
        assert "PO_xuong_com_NHAP_2026-09-01.xlsx" in actual_po_response.headers["Content-Disposition"]
        actual_po_book = load_workbook(io.BytesIO(actual_po_response.data), data_only=True)
        assert set(actual_po_book.sheetnames) == {"XCOM", "MAZDA", "TTS"}
        expected_headers = [
            "STT", "Bếp / ca", "Mã hàng", "Tên nguyên liệu", "ĐVT", "Suất áp dụng",
            "Định lượng/1.000 suất", "Số lượng cần", "Đơn giá", "Thành tiền", "NCC", "Nguồn giá",
        ]
        po_rows = []
        for sheet in actual_po_book.worksheets:
            assert sheet.cell(1, 1).value == "PO XƯỞNG CƠM – BẢN NHÁP"
            assert [sheet.cell(3, col).value for col in range(1, 13)] == expected_headers
            assert sheet.cell(2, 4).value == f"XCOM {sheet.title}"
            row = 4
            while sheet.cell(row, 1).value is not None:
                po_rows.append(tuple(sheet.cell(row, col).value for col in range(1, 13)))
                row += 1
            assert sheet.cell(row + 1, 1).value == "TỔNG HỢP SUẤT ĂN / COST"
        actual_po_book.close()
        assert len(po_rows) == 54
        assert abs(sum(float(row[9] or 0) for row in po_rows) - 3851200) < 0.01
        assert all(row[7] and row[8] and row[9] and row[10] and "HATRAN" in row[11] for row in po_rows)
        with actual_kitchen_source.open("rb") as handle:
            repeat_actual_response = client.post(
                "/api/kitchen/import/preview",
                data={"work_date": "2026-09-01", "file": (handle, "renamed-kitchen-file.xlsx")},
                content_type="multipart/form-data",
            )
        repeat_actual = repeat_actual_response.get_json()
        assert repeat_actual["counts"]["new"] == 0 and repeat_actual["counts"]["update"] == 5
        repeat_actual_confirm = client.post(
            "/api/kitchen/import/confirm",
            json={"token": repeat_actual["token"], "confirmed": True},
        )
        assert repeat_actual_confirm.status_code == 200 and repeat_actual_confirm.get_json()["updated"] == 5
        repeated_actual_plans = client.get("/api/kitchen/plans?date=2026-09-01").get_json()["items"]
        assert len(repeated_actual_plans) == 5
        assert all(
            item["price_source"] == "HATRAN 2026-09 · giá kỳ đã khóa"
            for plan in repeated_actual_plans for item in plan["items"]
        ), "Nhập lại file xưởng cơm không được đổi provenance giá"
        actual_kitchen_result = "5 plans / 54 items / 3 XCOM sheets / financial reconciliation passed"

    # Actual monthly meal-attendance workbook: keep actual meals separate from PO planned meals.
    meal_attendance_source = (
        ROOT / "bosung.30.8.26" / "CHẤM CÔNG+ SUẤT ĂN  2026"
        / "SUẤT ĂN XƯỞNG CƠM 2026" / "SUẤT ĂN T8.2026.xlsx"
    )
    assert hashlib.sha256(meal_attendance_source.read_bytes()).hexdigest().upper() == (
        "1117786C6356E75CA75329926E1F3717BD8EADAE3857F8E5BED3B36B1C9E1EBF"
    )
    with meal_attendance_source.open("rb") as handle:
        meal_attendance_preview_response = client.post(
            "/api/kitchen/attendance/import/preview",
            data={"file": (handle, meal_attendance_source.name)},
            content_type="multipart/form-data",
        )
    assert meal_attendance_preview_response.status_code == 200, meal_attendance_preview_response.get_data(as_text=True)
    meal_attendance_preview = meal_attendance_preview_response.get_json()
    assert meal_attendance_preview["periods"] == ["2026-08"]
    assert meal_attendance_preview["counts"] == {
        "items": 205, "dates": 28, "kitchens": 6, "new": 205,
        "update": 0, "unchanged": 0, "warning": 0, "error": 0,
    }
    assert meal_attendance_preview["totals"] == {"actual": 6020, "ordered": 0}
    assert client.post(
        "/api/kitchen/attendance/import/confirm",
        json={"token": meal_attendance_preview["token"]},
    ).status_code == 400
    meal_attendance_confirm = client.post(
        "/api/kitchen/attendance/import/confirm",
        json={"token": meal_attendance_preview["token"], "confirmed": True},
    )
    assert meal_attendance_confirm.status_code == 200, meal_attendance_confirm.get_data(as_text=True)
    assert meal_attendance_confirm.get_json()["inserted"] == 205
    with server.db() as conn:
        meal_totals_by_kitchen = {
            row["kitchen"]: row["actual"]
            for row in conn.execute(
                """SELECT kitchen,SUM(actual_count) actual FROM meal_attendance
                   WHERE substr(work_date,1,7)='2026-08' GROUP BY kitchen"""
            )
        }
    assert meal_totals_by_kitchen == {
        "BOT": 1313, "DAINAM": 1433, "SUNBY": 874,
        "THACO": 1063, "TTS": 627, "VINA": 710,
    }
    with meal_attendance_source.open("rb") as handle:
        repeat_meal_attendance_response = client.post(
            "/api/kitchen/attendance/import/preview",
            data={"file": (handle, "renamed-meal-attendance.xlsx")},
            content_type="multipart/form-data",
        )
    repeat_meal_attendance = repeat_meal_attendance_response.get_json()
    assert repeat_meal_attendance["counts"]["unchanged"] == 205
    repeat_meal_confirm = client.post(
        "/api/kitchen/attendance/import/confirm",
        json={"token": repeat_meal_attendance["token"], "confirmed": True},
    )
    assert repeat_meal_confirm.status_code == 200
    assert repeat_meal_confirm.get_json()["unchanged"] == 205
    assert client.post(
        "/api/kitchen/attendance/import/preview",
        data={"file": (io.BytesIO(b"not-an-xlsx"), "bad.xlsx")},
        content_type="multipart/form-data",
    ).status_code == 400

    # Xưởng cơm payment documents: exact actual meals + exact-period tariffs.
    payment_profile = client.put("/api/kitchen/payment-profiles/QC-BOT", json={
        "document_type": "MEAL_SIMPLE",
        "issuer_name": "CÔNG TY TNHH DỊCH VỤ HÀ TRÂN QC",
        "recipient_name": "CÔNG TY TNHH BOT QC",
        "beneficiary_name": "CÔNG TY TNHH DỊCH VỤ HÀ TRÂN QC",
        "bank_account": "0000000001",
        "bank_name": "NGÂN HÀNG QC",
        "requester": "NGƯỜI LẬP QC",
        "vat_rate": 8,
    })
    assert payment_profile.status_code == 200, payment_profile.get_data(as_text=True)
    payment_scope = client.put(
        "/api/kitchen/payment-profiles/QC-BOT/scopes/KITCHEN/BOT", json={}
    )
    assert payment_scope.status_code == 200, payment_scope.get_data(as_text=True)
    with server.db() as conn:
        bot_shifts = [
            row["shift"]
            for row in conn.execute(
                """SELECT DISTINCT shift FROM meal_attendance
                   WHERE kitchen='BOT' AND substr(work_date,1,7)='2026-08'
                   ORDER BY shift"""
            )
        ]
    assert bot_shifts
    for shift in bot_shifts:
        tariff = client.put(
            f"/api/kitchen/payment-profiles/QC-BOT/tariffs/2026-08/{quote(shift, safe='')}",
            json={"unit_price": 25000},
        )
        assert tariff.status_code == 200, tariff.get_data(as_text=True)
    payment_preview = client.post("/api/kitchen/payment-documents/preview", json={
        "profile_code": "QC-BOT",
        "date_from": "2026-08-01",
        "date_to": "2026-08-31",
        "issue_date": "2026-08-31",
    })
    assert payment_preview.status_code == 200, payment_preview.get_data(as_text=True)
    payment_preview_data = payment_preview.get_json()
    payment_summary = payment_preview_data["summary"]
    assert payment_summary["actual_count"] == 1313
    assert payment_summary["subtotal"] == 32825000
    assert payment_summary["vat_amount"] == 2626000
    assert payment_summary["total"] == 35451000
    payment_export_url = "/api/kitchen/payment-documents/export"
    payment_export_body = {
        "profile_code": "QC-BOT",
        "date_from": "2026-08-01",
        "date_to": "2026-08-31",
        "issue_date": "2026-08-31",
    }
    first_payment_export = client.post(payment_export_url, json={
        **payment_export_body, "preview_token": payment_preview_data["preview_token"],
    })
    reused_payment_preview = client.post(payment_export_url, json={
        **payment_export_body, "preview_token": payment_preview_data["preview_token"],
    })
    assert reused_payment_preview.status_code == 409
    assert reused_payment_preview.get_json()["code"] == "preview_used"
    second_payment_preview = client.post("/api/kitchen/payment-documents/preview", json={
        "profile_code": "QC-BOT",
        "date_from": "2026-08-01",
        "date_to": "2026-08-31",
        "issue_date": "2026-08-31",
    }).get_json()
    second_payment_export = client.post(payment_export_url, json={
        **payment_export_body, "preview_token": second_payment_preview["preview_token"],
    })
    assert first_payment_export.status_code == 200 and len(first_payment_export.data) > 10000
    assert first_payment_export.data == second_payment_export.data
    assert "wordprocessingml.document" in first_payment_export.content_type
    xcom_payment_document = DocxDocument(io.BytesIO(first_payment_export.data))
    xcom_payment_text = "\n".join(
        [paragraph.text for paragraph in xcom_payment_document.paragraphs]
        + [
            cell.text for table in xcom_payment_document.tables
            for row in table.rows for cell in row.cells
        ]
    )
    for required_text in (
        "GIẤY ĐỀ NGHỊ THANH TOÁN", "CÔNG TY TNHH BOT QC",
        "35.451.000", "0000000001", "NGƯỜI LẬP QC",
    ):
        assert required_text in xcom_payment_text, required_text
    assert len(xcom_payment_document.tables) == 2
    assert [len(table.columns) for table in xcom_payment_document.tables] == [2, 3]
    xcom_payment_result = (
        "1,313 actual meals / exact tariffs / customer-form DOCX / single-use preview passed"
    )

    # Normalized kitchen/menu/cost/XCOM/PO manual flow remains backward compatible.
    assert client.put("/api/kitchen-units/POT", json={"xcom_code": "XCOM-POT"}).status_code == 200
    assert client.put("/api/dated-prices/I000060", json={
        "period": "2026-08", "price_group": "HATRAN", "price_value": 16000,
    }).status_code == 200
    meal = client.post("/api/kitchen/plans", json={
        "work_date": "2026-08-30", "kitchen": "POT", "shift": "Trưa", "meal_count": 100,
        "items": [{"dish_name": "Canh", "product_code": "I000060", "norm_qty": 0.08}],
    })
    assert meal.status_code == 200, meal.get_data(as_text=True)
    meal_id = meal.get_json()["id"]
    missing_meal_price = client.post(f"/api/kitchen/plans/{meal_id}/approve")
    assert missing_meal_price.status_code == 400 and "đơn giá suất ăn" in missing_meal_price.get_json()["error"]
    meal = client.post("/api/kitchen/plans", json={
        "id": meal_id, "work_date": "2026-08-30", "kitchen": "POT", "shift": "Trưa",
        "meal_count": 100, "meal_price": 25000,
        "items": [{"dish_name": "Canh", "product_code": "I000060", "norm_qty": 0.08}],
    })
    assert meal.status_code == 200, meal.get_data(as_text=True)
    assert client.post(f"/api/kitchen/plans/{meal_id}/approve").status_code == 200
    kitchen_po = client.get("/api/kitchen/po?date=2026-08-30")
    assert kitchen_po.status_code == 200 and len(kitchen_po.data) > 1000
    assert "PO_xuong_com_DA_DUYET_2026-08-30.xlsx" in kitchen_po.headers["Content-Disposition"]
    approved_po_book = load_workbook(io.BytesIO(kitchen_po.data), data_only=True)
    assert approved_po_book.active.cell(1, 1).value == "PO XƯỞNG CƠM – ĐÃ DUYỆT"
    approved_po_book.close()
    with server.db() as conn:
        conn.execute(
            """INSERT INTO kitchen_labor_costs(work_date,kitchen,amount,source,updated_at)
               VALUES('2026-08-30','POT',100000,'QC-LABOR',?)""",
            (server.now_iso(),),
        )
        costed_plan = contract_modules.meal_plan_payload(conn, "2026-08-30")[0]
    assert costed_plan["labor_cost"] == 100000
    assert costed_plan["total_cost"] == costed_plan["food_cost"] + costed_plan["other_cost"] + 100000
    assert client.put("/api/dated-prices/I000060", json={
        "period": "2026-13", "price_group": "HATRAN", "price_value": 16000,
    }).status_code == 400
    assert client.put("/api/dated-prices/I000060", json={
        "period": "2026-10", "price_group": "HATRAN", "price_value": 0,
    }).status_code == 400
    # A timeless catalogue HATRAN value must not be borrowed into a new month.
    missing_period_plan = client.post("/api/kitchen/plans", json={
        "work_date": "2026-10-01", "kitchen": "POT", "shift": "Trưa",
        "meal_count": 10, "meal_price": 25000,
        "items": [{"dish_name": "Canh", "product_code": "I000060", "norm_qty": 0.08}],
    })
    assert missing_period_plan.status_code == 200
    missing_period_id = missing_period_plan.get_json()["id"]
    assert any("đúng kỳ 2026-10" in warning for warning in missing_period_plan.get_json()["warnings"])
    assert client.post(f"/api/kitchen/plans/{missing_period_id}/approve").status_code == 400
    assert client.put("/api/dated-prices/I000060", json={
        "period": "2026-10", "price_group": "HATRAN", "price_value": 17000,
    }).status_code == 200
    repriced = client.post("/api/kitchen/plans", json={
        "id": missing_period_id, "work_date": "2026-10-01", "kitchen": "POT", "shift": "Trưa",
        "meal_count": 10, "meal_price": 25000,
        "items": [{"dish_name": "Canh", "product_code": "I000060", "norm_qty": 0.08}],
    })
    assert repriced.status_code == 200
    assert client.post(f"/api/kitchen/plans/{missing_period_id}/approve").status_code == 200
    unmapped_plan = client.post("/api/kitchen/plans", json={
        "work_date": "2026-08-31", "kitchen": "CHUA-GHEP", "shift": "Trưa",
        "meal_count": 10, "meal_price": 25000,
        "items": [{"dish_name": "Canh", "product_code": "I000060", "norm_qty": 0.08}],
    })
    assert unmapped_plan.status_code == 200
    unmapped_approval = client.post(
        f"/api/kitchen/plans/{unmapped_plan.get_json()['id']}/approve"
    )
    assert unmapped_approval.status_code == 400 and "XCOM" in unmapped_approval.get_json()["error"]

    # Attendance/payroll manual flow and real legacy workbook import.
    assert client.post("/api/staff", json={
        "employee_code": "NVQC", "full_name": "Nhân viên QC", "role_name": "Bếp",
        "kitchen": "POT", "base_salary": 5200000,
    }).status_code == 200
    assert client.post("/api/attendance", json={
        "employee_code": "NVQC", "work_date": "2026-08-01", "normal_hours": 8,
        "overtime_hours": 2,
    }).status_code == 200
    assert client.put("/api/payroll-adjustments/NVQC/2026-08", json={
        "allowance": 300000, "responsibility": 200000, "advance": 100000,
    }).status_code == 200
    payroll = client.get("/api/payroll?month=2026-08")
    assert payroll.status_code == 200
    qc_pay = next(row for row in payroll.get_json()["items"] if row["employee_code"] == "NVQC")
    assert qc_pay["net_salary"] > 0
    attendance_source = ROOT / "bosung.30.8.26" / "CHẤM CÔNG+ SUẤT ĂN  2026" / "CHẤM CÔNG T8.2026.xlsx"
    locked_attendance_import = client.post("/api/attendance/import")
    assert locked_attendance_import.status_code == 410
    with attendance_source.open("rb") as handle:
        attendance_preview = client.post(
            "/api/attendance/import/preview",
            data={"period": "2026-08", "file": (handle, attendance_source.name)},
            content_type="multipart/form-data",
        )
    assert attendance_preview.status_code == 200, attendance_preview.get_data(as_text=True)
    attendance_preview_data = attendance_preview.get_json()
    assert attendance_preview_data["can_confirm"] is True
    assert attendance_preview_data["month"] == "2026-08"
    assert attendance_preview_data["counts"]["attendance_entries"] > 0
    assert attendance_preview_data["counts"]["labor_cost_entries"] > 0
    imported_attendance = client.post("/api/attendance/import/confirm", json={
        "token": attendance_preview_data["token"], "confirmed": True,
    })
    assert imported_attendance.status_code == 200, imported_attendance.get_data(as_text=True)
    assert imported_attendance.get_json()["attendance_entries"] > 0
    assert imported_attendance.get_json()["labor_cost_entries"] > 0
    assert imported_attendance.get_json()["payroll_overrides"] >= 7
    imported_payroll = client.get("/api/payroll?month=2026-08").get_json()["items"]
    legacy_payroll = [row for row in imported_payroll if row["employee_code"] != "NVQC"]
    assert max(row["normal_hours"] for row in legacy_payroll) <= 31 * 24
    assert round(next(row for row in legacy_payroll if row["employee_code"] == "TRIEN")["net_salary"]) == 8000000
    assert round(sum(row["net_salary"] for row in legacy_payroll)) == 14379021
    payroll_export = client.get("/api/export/payroll?month=2026-08")
    assert payroll_export.status_code == 200 and len(payroll_export.data) > 1000
    payroll_book = load_workbook(io.BytesIO(payroll_export.data), data_only=True)
    assert payroll_book.sheetnames == ["LƯƠNG XƯỞNG", "Chi phí theo bếp"]
    assert payroll_book["LƯƠNG XƯỞNG"]["A3"].value == "STT"
    assert any(
        row[1].value == "TRIỂN" and round(row[14].value) == 8000000
        for row in payroll_book["LƯƠNG XƯỞNG"].iter_rows(min_row=4)
    )
    payroll_book.close()
    assert client.get("/api/export/payroll?month=2026-13").status_code == 400

    # Opening inventory workbook: safe preview/confirm, signed quantities, aggregation and rollback.
    with server.db() as conn:
        conn.execute(
            """INSERT INTO products(code,name,unit,supplier,buy_price)
               VALUES('QC-OPEN-ROLLBACK','Hàng rollback tồn đầu','kg','QC',10000)"""
        )
    rollback_opening_file = opening_workbook_bytes([
        ["QC-OPEN-NEW", "Hàng mới phải rollback", "", "QC-OPEN-NEW", "KKKNT", "kg", 2, 20000, 40000],
        ["QC-OPEN-ROLLBACK", "Hàng rollback tồn đầu", "", "QC-OPEN-ROLLBACK", "KKKNT", "kg", 1, 10000, 10000],
    ])
    rollback_opening = preview_opening(client, rollback_opening_file, period="2026-07").get_json()
    assert rollback_opening["can_confirm"] is True and rollback_opening["counts"]["new_products"] == 1
    with server.db() as conn:
        conn.execute("DELETE FROM products WHERE code='QC-OPEN-ROLLBACK'")
    rollback_opening_confirm = client.post(
        "/api/inventory/opening/import/confirm",
        json={"token": rollback_opening["token"], "confirmed": True},
    )
    assert rollback_opening_confirm.status_code == 409
    with server.db() as conn:
        assert not conn.execute("SELECT 1 FROM products WHERE code='QC-OPEN-NEW'").fetchone()
        assert conn.execute(
            "SELECT COUNT(*) n FROM inventory_transactions WHERE source_type='OPENING' AND source_id='2026-07'"
        ).fetchone()["n"] == 0

    actual_opening_source = ROOT / "_HANDOFF" / "EXTERNAL_INPUTS" / "TĐK T8-2026.xlsx thụy.xlsx"
    actual_opening_result = "not_present"
    if actual_opening_source.exists():
        assert hashlib.sha256(actual_opening_source.read_bytes()).hexdigest().upper() == (
            "36DF2BA86D13307F96BB5944FCECB19A4A81C093B4AC6A98EA71330D68204DA6"
        )
        assert preview_opening(client, b"not-an-xlsx").status_code == 400
        assert preview_opening(client, actual_opening_source.read_bytes(), period="2026-13").status_code == 400
        with actual_opening_source.open("rb") as handle:
            actual_opening_response = client.post(
                "/api/inventory/opening/import/preview",
                data={"period": "2026-08", "file": (handle, actual_opening_source.name)},
                content_type="multipart/form-data",
            )
        assert actual_opening_response.status_code == 200, actual_opening_response.get_data(as_text=True)
        actual_opening = actual_opening_response.get_json()
        print("OPENING_PREVIEW", json.dumps(actual_opening["counts"], ensure_ascii=False))
        assert actual_opening["can_confirm"] is True
        assert actual_opening["counts"] == {
            "source_rows": 395, "items": 334, "new_products": 113,
            "new": 334, "update": 0, "negative": 4, "warning": 181, "error": 0,
        }
        assert abs(actual_opening["totals"]["qty"] - 101383.26) < 0.0001
        assert abs(actual_opening["totals"]["amount"] - 2419360717.80439) < 0.01
        c000021 = next(row for row in actual_opening["rows"] if row["product_code"] == "C000021")
        assert c000021["source_row_count"] == 3 and abs(c000021["qty"] - 79.97) < 0.0001
        d000056 = next(row for row in actual_opening["rows"] if row["product_code"] == "D000056")
        assert d000056["qty"] == -1.5 and any("âm" in warning for warning in d000056["warnings"])
        m000277 = next(row for row in actual_opening["rows"] if row["product_code"] == "M000277")
        assert any("Thành tiền lệch" in warning for warning in m000277["warnings"])
        k000066 = next(row for row in actual_opening["rows"] if row["product_code"] == "K000066")
        assert any("Tên trong file khác danh mục" in warning for warning in k000066["warnings"])
        assert client.post(
            "/api/inventory/opening/import/confirm", json={"token": actual_opening["token"]}
        ).status_code == 400
        actual_opening_confirm = client.post(
            "/api/inventory/opening/import/confirm",
            json={"token": actual_opening["token"], "confirmed": True},
        )
        assert actual_opening_confirm.status_code == 200, actual_opening_confirm.get_data(as_text=True)
        assert actual_opening_confirm.get_json()["inserted_products"] == 113
        assert actual_opening_confirm.get_json()["inserted"] == 334
        assert actual_opening_confirm.get_json()["updated"] == 0
        with server.db() as conn:
            opening_totals = conn.execute(
                """SELECT COUNT(*) n,SUM(qty_in-qty_out) qty,
                          SUM((qty_in-qty_out)*unit_cost) amount,
                          SUM(CASE WHEN qty_out>0 THEN 1 ELSE 0 END) negative
                   FROM inventory_transactions
                   WHERE source_type='OPENING' AND source_id='2026-08'
                     AND note LIKE 'Tồn đầu kỳ 2026-08 từ %'"""
            ).fetchone()
            assert opening_totals["n"] == 334 and opening_totals["negative"] == 4
            assert abs(opening_totals["qty"] - 101383.26) < 0.0001
            assert abs(opening_totals["amount"] - 2419360717.80439) < 0.01
            assert conn.execute("SELECT COUNT(*) n FROM products").fetchone()["n"] == 984
            assert conn.execute("SELECT name FROM products WHERE code='K000066'").fetchone()["name"] == "Thịt vịt xông khói"
            added_product = conn.execute(
                "SELECT name,unit,supplier FROM products WHERE code='HT00241'"
            ).fetchone()
            assert added_product and added_product["supplier"] == ""
        with actual_opening_source.open("rb") as handle:
            repeat_opening_response = client.post(
                "/api/inventory/opening/import/preview",
                data={"period": "2026-08", "file": (handle, "renamed-opening.xlsx")},
                content_type="multipart/form-data",
            )
        repeat_opening = repeat_opening_response.get_json()
        assert repeat_opening["counts"]["new_products"] == 0
        assert repeat_opening["counts"]["new"] == 0 and repeat_opening["counts"]["update"] == 334
        repeat_opening_confirm = client.post(
            "/api/inventory/opening/import/confirm",
            json={"token": repeat_opening["token"], "confirmed": True},
        )
        assert repeat_opening_confirm.status_code == 200
        assert repeat_opening_confirm.get_json()["inserted"] == 0
        assert repeat_opening_confirm.get_json()["updated"] == 334
        actual_opening_result = "395 rows / 334 TDP codes / signed quantity and VND value reconciled"

    actual_payables_source = ROOT / "bosung.30.8.26" / "Công nợ phải trả Thành Đạt Phát.xlsx"
    actual_payables_result = "not_present"
    payable_payment_result = "not_present"
    payable_export_result = "not_present"
    if actual_payables_source.exists():
        assert hashlib.sha256(actual_payables_source.read_bytes()).hexdigest().upper() == (
            "BC54C62E7B7141BA1846EB5BD7EEC9E11B9CC8DA68FB138D0EA0F0CDD7BAF473"
        )
        with server.db() as conn:
            conn.execute(
                """INSERT INTO historical_payable_lines(
                       purchase_date,kitchen,item_name,qty,unit,supplier,buy_price,actual_qty,
                       amount,source_file,source_sheet,source_row,source_hash,created_at,updated_at
                   ) VALUES('2026-01-01','QC','Dòng snapshot cũ',1,'kg','QC',1,1,1,
                            'old-renamed.xlsx','Data',2,'OLDHASH',?,?)""",
                (server.now_iso(), server.now_iso()),
            )
        with actual_payables_source.open("rb") as handle:
            actual_payables_response = client.post(
                "/api/debts/payables/import/preview",
                data={"file": (handle, actual_payables_source.name)},
                content_type="multipart/form-data",
            )
        assert actual_payables_response.status_code == 200, actual_payables_response.get_data(as_text=True)
        actual_payables = actual_payables_response.get_json()
        assert actual_payables["can_confirm"] is True
        assert actual_payables["counts"] == {
            "total": 9976, "ready": 9975, "skipped": 1, "errors": 0,
            "warnings": 51, "suppliers": 28,
        }
        assert len(actual_payables["issues"]) == 51
        assert any(item["source_row"] > 200 for item in actual_payables["issues"])
        assert abs(actual_payables["totals"]["amount"] - 2327247386.1) < 0.01
        assert client.post(
            "/api/debts/payables/import/confirm", json={"token": actual_payables["token"]},
        ).status_code == 400
        actual_payables_confirm = client.post(
            "/api/debts/payables/import/confirm",
            json={"token": actual_payables["token"], "confirmed": True},
        )
        assert actual_payables_confirm.status_code == 200
        assert actual_payables_confirm.get_json()["inserted"] == 9975
        assert actual_payables_confirm.get_json()["replaced"] == 1
        with server.db() as conn:
            assert conn.execute("SELECT COUNT(*) n FROM historical_payable_lines").fetchone()["n"] == 9975
            assert not conn.execute(
                "SELECT 1 FROM historical_payable_lines WHERE source_hash='OLDHASH'"
            ).fetchone()
            historical_ledger_first = conn.execute(
                "SELECT COUNT(*) n FROM payable_ledger_lines WHERE source_type='historical_import'"
            ).fetchone()["n"]
            historical_revisions_first = conn.execute(
                """SELECT COUNT(*) n FROM payable_ledger_revisions r
                   JOIN payable_ledger_lines l ON l.id=r.ledger_line_id
                   WHERE l.source_type='historical_import'"""
            ).fetchone()["n"]
            assert historical_ledger_first == 9975
            assert historical_revisions_first == 9975
            assert conn.execute(
                """SELECT COUNT(*) n FROM payable_ledger_lines
                   WHERE source_type='historical_import' AND status!='open'"""
            ).fetchone()["n"] == 0
        with actual_payables_source.open("rb") as handle:
            repeat_payables = client.post(
                "/api/debts/payables/import/preview",
                data={"file": (handle, "renamed-payables.xlsx")},
                content_type="multipart/form-data",
            ).get_json()
        repeat_payables_confirm = client.post(
            "/api/debts/payables/import/confirm",
            json={"token": repeat_payables["token"], "confirmed": True},
        )
        assert repeat_payables_confirm.status_code == 200
        assert repeat_payables_confirm.get_json()["inserted"] == 0
        assert repeat_payables_confirm.get_json()["unchanged"] == 9975
        with server.db() as conn:
            assert conn.execute(
                "SELECT COUNT(*) n FROM payable_ledger_lines WHERE source_type='historical_import'"
            ).fetchone()["n"] == historical_ledger_first
            assert conn.execute(
                """SELECT COUNT(*) n FROM payable_ledger_revisions r
                   JOIN payable_ledger_lines l ON l.id=r.ledger_line_id
                   WHERE l.source_type='historical_import'"""
            ).fetchone()["n"] == historical_revisions_first
            payable_line = dict(conn.execute(
                """SELECT * FROM payable_ledger_lines
                   WHERE source_type='historical_import' AND status='open'
                     AND amount>=1 AND TRIM(supplier_code)!=''
                   ORDER BY work_date,id LIMIT 1"""
            ).fetchone())
        allocated_payment = client.post("/api/debts/payables/payments", json={
            "request_id": "QC-PAYABLE-ALLOCATION-0001",
            "payment_date": "2026-08-31",
            "party_code": payable_line["supplier_code"],
            "amount": 1,
            "method": "Chuyển khoản QC",
            "reference_code": "QC-UNC-001",
            "note": "QC phân bổ dòng phải trả",
            "allocations": [{
                "ledger_line_id": payable_line["id"],
                "amount": 1,
                "expected_revision": payable_line["revision"],
            }],
        })
        assert allocated_payment.status_code == 201, allocated_payment.get_data(as_text=True)
        payment_id = allocated_payment.get_json()["id"]
        payment_history = client.get(
            "/api/debts/payables/payments?from=2026-08-31&to=2026-08-31&status=posted"
        )
        assert payment_history.status_code == 200
        assert payment_history.get_json()["summary"]["posted_amount"] == 1
        reversed_payment = client.post(
            f"/api/debts/payables/payments/{payment_id}/reverse",
            json={"expected_revision": 1, "reason": "QC hoàn tác giao dịch thử"},
        )
        assert reversed_payment.status_code == 200, reversed_payment.get_data(as_text=True)
        with server.db() as conn:
            payment_state = conn.execute(
                "SELECT status,revision FROM payments WHERE id=?", (payment_id,)
            ).fetchone()
            allocation_state = conn.execute(
                "SELECT status FROM payable_payment_allocations WHERE payment_id=?", (payment_id,)
            ).fetchone()
            ledger_state = conn.execute(
                "SELECT status,paid_amount FROM payable_ledger_lines WHERE id=?",
                (payable_line["id"],),
            ).fetchone()
            assert tuple(payment_state) == ("reversed", 2)
            assert allocation_state["status"] == "reversed"
            assert tuple(ledger_state) == ("open", 0)
            assert conn.execute(
                "SELECT COUNT(*) n FROM payable_payment_revisions WHERE payment_id=?",
                (payment_id,),
            ).fetchone()["n"] == 2
        payable_export_response = client.get(
            "/api/debts/payables/export?from=2026-01-01&to=2026-07-31"
        )
        assert payable_export_response.status_code == 200
        payable_export = load_workbook(io.BytesIO(payable_export_response.data), data_only=False)
        assert payable_export.sheetnames == ["Công nợ phải trả"]
        worksheet = payable_export.active
        assert [worksheet.cell(3, column).value for column in range(1, 15)] == [
            "Tháng", "Tên bếp", "Ngày, tháng", "Tên hàng", "Số lượng", "ĐVT", "NCC",
            "Giá mua", "Hỏng", "Thêm", "Giảm", "Thiếu", "SL thực tế", "Thành tiền",
        ]
        assert worksheet.max_column == 14
        assert worksheet.freeze_panes == "A4"
        assert worksheet.page_setup.orientation == "landscape"
        assert worksheet.page_setup.fitToWidth == 1
        assert worksheet.cell(worksheet.max_row, 1).value == "TỔNG"
        detail_rows = list(worksheet.iter_rows(min_row=4, max_row=worksheet.max_row - 1, values_only=True))
        assert abs(
            sum(row[4] or 0 for row in detail_rows) - worksheet.cell(worksheet.max_row, 5).value
        ) < 1e-6
        assert abs(
            sum(row[12] or 0 for row in detail_rows) - worksheet.cell(worksheet.max_row, 13).value
        ) < 1e-6
        assert sum(row[13] or 0 for row in detail_rows) == worksheet.cell(worksheet.max_row, 14).value
        assert not any(
            isinstance(cell.value, str) and cell.value.startswith("=")
            for worksheet in payable_export.worksheets
            for row in worksheet.iter_rows()
            for cell in row
        )
        payable_export.close()
        payable_payment_result = "explicit line allocation / reversal / history passed"
        payable_export_result = "1 customer sheet / exact 14 columns / total quantity + amount / static values"
        actual_payables_result = "9,976 rows / 9,975 payable lines / 0 errors / repeat safe"

    payable_ui_script = (APP_DIR / "static" / "app.js").read_text(encoding="utf-8")
    payable_ui_script += (APP_DIR / "static" / "invoice-workbench.js").read_text(encoding="utf-8")
    payable_ui_page = (APP_DIR / "static" / "index.html").read_text(encoding="utf-8")
    payable_ui_css = (APP_DIR / "static" / "real.css").read_text(encoding="utf-8")
    for required_ui_contract in (
        'id="payablePaymentForm"', 'class="payable-line-select"',
        'class="payable-allocation-input"', '"/api/debts/payables/payments"',
        '"/reverse"', "expected_revision", "tdp.payableFilters",
        'id="receiptForm"', "Excel phải trả nhà cung cấp",
        "esc(state.debtTo || todayIso)",
    ):
        assert required_ui_contract in payable_ui_script, required_ui_contract
    assert '<option value="payment">Trả nhà cung cấp</option>' not in payable_ui_script
    assert "payable-line-select:disabled" in payable_ui_css
    debt_ui_combined = payable_ui_page + payable_ui_script
    for required_debt_ui in (
        'data-view="reports"><span>↗</span> Báo cáo tổng hợp',
        'data-view="debts"><span>▤</span> Công nợ',
        'data-action="open-debt-section"',
        "Công nợ phải thu (bếp)",
        "Công nợ phải thu (tổng)",
        "Công nợ phải trả",
        'data-action="back-debt-overview"',
        'if (section === "receivable-kitchen")',
        'if (section === "receivable-total")',
    ):
        assert required_debt_ui in debt_ui_combined, required_debt_ui
    assert "Báo cáo & công nợ" not in debt_ui_combined
    assert ".debt-menu-grid" in payable_ui_css
    assert "/static/app.js?v=20260904-27" in payable_ui_page
    assert "/static/real.css?v=20260904-12" in payable_ui_page
    payable_ui_result = "separate summary-first debt menu / filter / explicit manual allocation / reversal / persisted refresh passed"
    for required_receivable_ui in (
        'id="receivableFilterForm"', 'id="receivableContractor"',
        'id="receivableKitchen"', 'id="receivableStatus"',
        'id="receivableExportSelected"', '"/api/debts/receivables/ledger?"',
        '"/api/debts/receivables/ledger/"',
        '"/api/debts/receivables/export?from="',
        'data-action="toggle-receivable-history"', "tdp.receivableFilters",
        "Sổ phải thu vận hành chi tiết", "không phải đề nghị thanh toán/hóa đơn đỏ",
    ):
        assert required_receivable_ui in payable_ui_script, required_receivable_ui
    for required_receivable_css in (
        ".receivable-filter-grid", ".receivable-ledger-table",
        ".receivable-line-reversed", ".receivable-revision-panel",
    ):
        assert required_receivable_css in payable_ui_css, required_receivable_css
    receivable_ui_result = "period / contractor / kitchen / revision / export / persisted filters passed"

    for required_inventory_ui in (
        'id="inventoryFrom"', 'id="inventoryTo"',
        '"/api/invoice-valuation?from="',
        "/api/invoice-valuation/export",
        "/api/invoice-valuation/export/opening",
        "/api/invoice-valuation/export/input",
        "/api/invoice-valuation/export/output",
        "/api/invoice-valuation/export/nxt",
        "Báo cáo vật tư hàng hóa", "Tải đủ 4 file ZIP", "Xem chi tiết và nhập dữ liệu",
        "/api/bk-import/template",
        "/api/bk-import/preview", "/api/bk-import/confirm",
        "Xác nhận nhập bảng kê vào kho", "Hoàn tác bảng kê",
    ):
        assert required_inventory_ui in payable_ui_script, required_inventory_ui
    assert "Báo cáo vật tư hàng hóa" in payable_ui_page
    assert "Kho hóa đơn" not in payable_ui_page + payable_ui_script
    inventory_ui_result = "customer-named material report / compact period projection / four downloads passed"

    for required_unified_ui in (
        "Bảng kê từ hóa đơn đỏ",
        "Được xuất và chưa được xuất theo nhà thầu",
        "File tải hóa đơn",
        "function invoiceReceiptStatusText(value)",
        "function invoiceStockStatusText(value)",
        "function invoiceSyncStatusText(value)",
        "function mealPlanStatusText(value)",
        'event.target.closest(".unit-conversion-input")',
        "Enter để lưu mã hoặc quy đổi",
        "Thông tin thanh toán mặc định",
        "Được khóa cùng hóa đơn và dùng để lập Đề nghị thanh toán chính thức theo nhà thầu",
        "invoice-issue-text",
        "hãy kiểm tra và tải tiếp",
    ):
        assert required_unified_ui in payable_ui_script, required_unified_ui
    for forbidden_unified_ui in (
        "esc(invoice.receipt_status)", "esc(invoice.stock_status)",
        "esc(sync.last_status)", "esc(plan.status)",
        "Xác nhận reversal", "Đã tạo reversal", "Đã lấy từ mẫu TĐP khách cung cấp",
    ):
        assert forbidden_unified_ui not in payable_ui_script, forbidden_unified_ui
    unified_ui_result = (
        "customer screen names / red-invoice statement boundary / Vietnamese statuses / "
        "Enter mapping and conversion / explicit errors passed"
    )

    home_source = payable_ui_script.split("function renderHome()", 1)[1].split(
        "function renderOrders()", 1
    )[0]
    assert home_source.count('class="daily-action-card') == 4
    for required_home_ui in (
        "Tạo phiếu đặt hàng", "In đơn hàng", "Bảng kê và biên nhận", "Duyệt đơn",
        'id="homeFrom"', 'id="homeTo"', "phiếu theo bếp",
    ):
        assert required_home_ui in home_source, required_home_ui
    for required_compact_ui in (
        'quoteDetailsOpen: false', 'quoteHistoryOpen: false',
        'inventoryDetailsOpen: true', 'documentDetailsOpen: false',
        "Báo giá tổng", "Báo giá chi tiết", "tổng hợp theo nhà thầu và từng bếp trong tháng",
        "Sao chép ảnh” sẽ tự ghi nhận đã đặt", 'body: JSON.stringify({ status: "ordered"',
        "Chọn tất cả hoặc từng bếp", "Chọn bếp cần in phiếu giao", "delivery_selections",
        'payload.phase === "finalization"', "Đã dùng file cuối cùng làm bản chuẩn và tự chốt đơn",
    ):
        assert required_compact_ui in payable_ui_script, required_compact_ui
    assert "Đánh dấu đã đặt" not in payable_ui_script
    compact_customer_ui_result = "exact four-item home / summary-first modules / per-kitchen reprint / latest-file authority passed"

    invoice_input_export_source = (APP_DIR / "invoice_input_export.py").read_text(encoding="utf-8")
    portable_build_source = (ROOT / "BUILD_PORTABLE.ps1").read_text(encoding="utf-8")
    for required_input_export in (
        "/api/invoice-workbench/batches/", "/export-input-xlsx", "Excel đúng bộ lọc",
    ):
        assert required_input_export in payable_ui_script, required_input_export
    for required_input_export_contract in (
        "INPUT_ELECTRONIC_INVOICE", "invoice_input_batch_empty",
        "DANH SÁCH HÓA ĐƠN ĐẦU VÀO ĐÃ TẢI", 'response.headers["X-TDP-Inventory-Effect"] = "none"',
        "Raw connector payloads, remote identifiers and credentials are never",
    ):
        assert required_input_export_contract in invoice_input_export_source, required_input_export_contract
    assert "invoice_input_export.py" in portable_build_source
    assert "--hidden-import invoice_input_export" in portable_build_source
    invoice_input_export_result = "download immediately after read-only pull / no inventory effect / no raw connector payload passed"

    outgoing_readiness_source = (APP_DIR / "outgoing_readiness.py").read_text(encoding="utf-8")
    for required_outgoing_ui in (
        "Danh sách còn thiếu theo kỳ", "Đã dự thảo", "Đã phát hành",
        'data-action="load-outgoing-shortages"',
        'data-action="download-outgoing-shortages"',
        "/api/outgoing-invoices/shortages/export?",
    ):
        assert required_outgoing_ui in payable_ui_script, required_outgoing_ui
    for required_readiness_contract in (
        "invoice_stock_rows(conn, include_zero=True)",
        "pending_sync_issued_qty", "missing_product_code", "allocated_over_demand",
        "period_shortage_payload", "shortage_workbook_bytes",
        "OPENING + invoice_inventory_ledger",
    ):
        assert required_readiness_contract in outgoing_readiness_source, required_readiness_contract
    assert "inventory_lookup(conn)" not in outgoing_readiness_source
    outgoing_readiness_result = (
        "canonical invoice stock / multi-round contractor split / period shortage export passed"
    )

    outgoing_substitution_source = (APP_DIR / "outgoing_substitution.py").read_text(encoding="utf-8")
    for required_substitution_ui in (
        "Luân chuyển / mặt hàng thay thế có xác nhận",
        "Mã hàng thay thế (tự chọn)", "Xem trước, chưa ghi",
        "Xác nhận đúng mã thay thế này", "Dữ liệu đã thay đổi sau lần kiểm tra",
        "/api/outgoing-substitutions/preview", "/api/outgoing-substitutions/confirm",
    ):
        assert required_substitution_ui in payable_ui_script, required_substitution_ui
    for required_substitution_contract in (
        "canonical_available_stock(conn)", "quote_sell_price",
        "approved_override", "SUBSTITUTION_PREVIEW_TTL_SECONDS",
        "stale_substitution_preview", "substitute_unit_mismatch",
        "issued_substitution_requires_invoice_reversal", "draft_kind",
        "outgoing.substitution.confirm", "outgoing.substitution.reverse",
    ):
        assert required_substitution_contract in outgoing_substitution_source, required_substitution_contract
    assert "inventory_lookup(conn)" not in outgoing_substitution_source
    assert "Gợi ý mã thay thế" not in payable_ui_script
    outgoing_substitution_result = (
        "manual code / preview token / period-contractor sell price / canonical stock / reversal audit passed"
    )

    invoice_payment_scope_source = (APP_DIR / "invoice_payment_scope.py").read_text(encoding="utf-8")
    invoice_delivery_statement_source = (APP_DIR / "invoice_delivery_statement.py").read_text(encoding="utf-8")
    for required_payment_scope_ui in (
        "Hồ sơ đề nghị thanh toán từ hóa đơn đỏ",
        "Xem và tải bảng kê", "Mẫu chính thức",
        "Tải Đề nghị thanh toán + bảng kê", "/api/outgoing-invoices/payment-scope/",
        'data-action="download-invoice-payment-control"',
        'data-action="download-invoice-delivery-statement"',
        "/api/outgoing-invoices/delivery-statement/",
    ):
        assert required_payment_scope_ui in payable_ui_script, required_payment_scope_ui
    for required_payment_scope_contract in (
        "status='issued'", "source_status_class", "invoice_source_not_payable",
        "invoice_source_total_mismatch", "issued_invoice_snapshot_conflict",
        "invoice_delivery_total_mismatch", '"template_status": "official_customer_xlsx"',
        '"official_template_ready": True',
    ):
        assert required_payment_scope_contract in invoice_payment_scope_source, required_payment_scope_contract
    assert "receivable_ledger" not in invoice_payment_scope_source
    invoice_payment_scope_result = (
        "issued invoice only / source cancellation guard / frozen snapshot / official customer XLSX passed"
    )
    for required_delivery_statement_contract in (
        "delivery_statement_unissued_invoice", "delivery_statement_contractor_conflict",
        "delivery_statement_total_mismatch", "Số lượng × Đơn giá = Thành tiền",
        "Nguồn xác minh", 'detail.print_title_rows = "10:10"',
        "PAPERSIZE_A4", "_excel_text", "scope_id",
    ):
        assert required_delivery_statement_contract in invoice_delivery_statement_source, required_delivery_statement_contract
    assert "Q-008" not in invoice_delivery_statement_source
    invoice_delivery_statement_result = (
        "one contractor / issued red invoices / multiple rounds / static A4 / exact totals passed"
    )

    # Multi-period debts, payment request and approval-gated print dry-run.
    debts = client.get("/api/debts?from=2026-08-01&to=2026-08-31")
    assert debts.status_code == 200 and "HATRAN" in debts.get_json()["contractors"]
    blank_settings = client.put("/api/document-settings", json={"payment_bank_account": ""})
    assert blank_settings.status_code == 200
    frozen_request = client.get("/api/export/payment-request/HATRAN?from=2026-08-01&to=2026-08-31")
    assert frozen_request.status_code == 200, frozen_request.get_data(as_text=True)
    restored_settings = client.put("/api/document-settings", json={
        "payment_requester": "VŨ THỊ THỤY",
        "payment_bank_name": "Ngân hàng TMCP Ngoại Thương Việt Nam",
        "payment_bank_account": "1052787580",
    })
    assert restored_settings.status_code == 200
    payment_request = client.get("/api/export/payment-request/HATRAN?from=2026-08-01&to=2026-08-31")
    assert payment_request.status_code == 200 and len(payment_request.data) > 1000
    assert payment_request.content_type.startswith("application/zip")
    with zipfile.ZipFile(io.BytesIO(payment_request.data)) as payment_bundle:
        bundle_names = payment_bundle.namelist()
        payment_xlsx_name = next(
            name for name in bundle_names if name.startswith("De_nghi_thanh_toan")
        )
        statement_xlsx_name = next(
            name for name in bundle_names if name.startswith("Bang_tong_hop_giao_nhan")
        )
        assert "THONG_TIN_DOI_CHIEU.txt" in bundle_names
        payment_xlsx_bytes = payment_bundle.read(payment_xlsx_name)
        statement_xlsx_bytes = payment_bundle.read(statement_xlsx_name)
    payment_document = load_workbook(io.BytesIO(payment_xlsx_bytes), data_only=True, keep_links=False)
    payment_text = "\n".join(
        str(cell.value or "") for sheet in payment_document.worksheets
        for row in sheet.iter_rows() for cell in row
    )
    for required_text in (
        "ĐỀ NGHỊ THANH TOÁN", "CÔNG TY TNHH THỰC PHẨM THÀNH ĐẠT PHÁT",
        "1052787580", "Ngân hàng TMCP Ngoại Thương Việt Nam",
        "Tổng cộng", "Bằng chữ", "ĐẠI DIỆN CÔNG TY",
    ):
        assert required_text in payment_text, required_text
    assert "................................" not in payment_text
    assert payment_document["Đối chiếu hóa đơn"].sheet_state == "hidden"
    payment_document.close()
    statement_document = load_workbook(io.BytesIO(statement_xlsx_bytes), data_only=True, keep_links=False)
    assert "BẢNG TỔNG HỢP GIAO NHẬN" in statement_document["Bảng kê giao hàng"]["A5"].value
    statement_document.close()
    prepared = client.post(f"/api/print/prepare/{manual_id}")
    assert prepared.status_code == 200, prepared.get_data(as_text=True)
    prepared_body = prepared.get_json()
    assert prepared_body["prepared"] == ["delivery_pdf", "other_pdf"]
    assert [item["paper"] for item in prepared_body["documents"]] == ["A4", "A5"]
    assert client.post(f"/api/print/approve/{manual_id}").status_code == 200
    print_dry_run = client.post(f"/api/print/run/{manual_id}", json={"dry_run": True})
    assert print_dry_run.status_code == 200 and print_dry_run.get_json()["jobs"] == 2
    assert print_dry_run.get_json()["status"] == "verified"
    for document_type in ("delivery_pdf", "other_pdf"):
        print_pdf = client.get(f"/api/print/pdf/{manual_id}/{document_type}")
        assert print_pdf.status_code == 200 and print_pdf.data.startswith(b"%PDF-")
        assert print_pdf.headers["Content-Type"].startswith("application/pdf")
        # ``send_file`` keeps the Windows file handle open until the test response is
        # explicitly closed. Release it before removing the isolated print folder.
        print_pdf.close()

    operations = client.get("/api/operations/bootstrap?as_of=2026-08-31&month=2026-08&date=2026-08-30")
    assert operations.status_code == 200
    assert operations.get_json()["msmi"]["invoices"]
    assert operations.get_json()["meal_plans"]
    assert operations.get_json()["meal_attendance_totals"] == {
        "actual": 6020, "ordered": 0, "rows": 205, "kitchens": 6,
    }
    assert operations.get_json()["document_settings"] == {
        "requester": "VŨ THỊ THỤY",
        "bank_name": "Ngân hàng TMCP Ngoại Thương Việt Nam",
        "bank_account": "1052787580",
    }

    inventory_files = client.get(
        "/api/invoice-valuation/export?from=2026-08-01&to=2026-08-31"
    )
    assert inventory_files.status_code == 200, inventory_files.get_data(as_text=True)
    assert inventory_files.content_type.startswith("application/zip")
    inventory_contract_ids = set()
    inventory_control_totals = set()
    with zipfile.ZipFile(io.BytesIO(inventory_files.data)) as inventory_bundle:
        inventory_names = inventory_bundle.namelist()
        assert len(inventory_names) == 4
        assert {name.split("_", 1)[0] for name in inventory_names} == {
            "TDK", "Nhap", "Xuat", "NXT",
        }
        for inventory_name in inventory_names:
            inventory_book = load_workbook(
                io.BytesIO(inventory_bundle.read(inventory_name)), data_only=False, keep_links=False,
            )
            control = inventory_book["_ĐỐI_CHIẾU"]
            inventory_manifest = {
                row[0].value: row[1].value
                for row in control.iter_rows(min_col=1, max_col=2)
            }
            inventory_contract_ids.add(control["B1"].value)
            inventory_control_totals.add(tuple(
                inventory_manifest[key]
                for key in (
                    "SL_TỒN_ĐẦU", "GT_TỒN_ĐẦU", "SL_NHẬP", "GT_NHẬP",
                    "SL_XUẤT", "GT_XUẤT", "SL_TỒN_CUỐI", "GT_TỒN_CUỐI",
                )
            ))
            assert inventory_manifest["MẪU_BIỂU"] == "TDP_TDK_NHAP_XUAT_NXT_V1"
            assert inventory_manifest["TRẠNG_THÁI_BIỂU_MẪU"] == "CHÍNH THỨC"
            assert inventory_manifest["PHƯƠNG_PHÁP_GIÁ"] == "Bình quân gia quyền di động"
            assert control.sheet_state == "veryHidden"
            assert not getattr(inventory_book, "_external_links", [])
            assert not any(
                cell.data_type == "f"
                for worksheet in inventory_book.worksheets
                for row in worksheet.iter_rows()
                for cell in row
            )
            inventory_book.close()
    assert len(inventory_contract_ids) == 1 and len(inventory_control_totals) == 1
    inventory_nxt = client.get(
        "/api/invoice-valuation/export/nxt?from=2026-08-01&to=2026-08-31"
    )
    assert inventory_nxt.status_code == 200
    inventory_nxt_book = load_workbook(io.BytesIO(inventory_nxt.data), data_only=True)
    inventory_nxt_sheet = inventory_nxt_book["NXT"]
    assert "bình quân gia quyền di động" in inventory_nxt_sheet["A3"].value.lower()
    assert "Q-005" not in inventory_nxt_sheet["A2"].value
    assert inventory_nxt_sheet["G4"].value == "TỒN ĐẦU KỲ"
    assert inventory_nxt_sheet["J4"].value == "NHẬP TRONG KỲ"
    assert inventory_nxt_sheet["M4"].value == "XUẤT TRONG KỲ"
    assert inventory_nxt_sheet["P4"].value == "TỒN CUỐI KỲ"
    assert inventory_nxt_sheet.page_setup.orientation == "landscape"
    assert str(inventory_nxt_sheet.page_setup.paperSize) == "9"
    inventory_nxt_book.close()
    assert client.get(
        "/api/invoice-valuation/export?from=2026-08-31&to=2026-08-01"
    ).status_code == 400
    inventory_export_result = "4 XLSX / one contract id / source traces and totals reconciled"

    bk_template_response = client.get("/api/bk-import/template")
    assert bk_template_response.status_code == 200
    bk_book = load_workbook(io.BytesIO(bk_template_response.data), data_only=False, keep_links=True)
    assert bk_book.sheetnames == [bk_import.BK_IMPORT_SHEET]
    bk_sheet = bk_book[bk_import.BK_IMPORT_SHEET]
    assert "HÀNG MUA VÀO KHÔNG CÓ HÓA ĐƠN" in bk_sheet["A1"].value
    assert bk_import.BK_IMPORT_SOURCE_TYPE in bk_sheet["A2"].value
    assert [bk_sheet.cell(3, column).value for column in range(1, 13)] == [
        label for _field, label in bk_import.BK_IMPORT_COLUMNS
    ]
    assert not any(
        cell.data_type == "f" for row in bk_sheet.iter_rows() for cell in row
    )
    with server.db() as conn:
        bk_product = conn.execute(
            "SELECT code,name,unit FROM products WHERE code=?", (mapped_code,),
        ).fetchone()
        conn.execute(
            """INSERT OR REPLACE INTO products(
                   code,name,unit,tax,supplier,buy_price,purchase_list,seller,cccd
               ) VALUES('QC-BK-LEGACY','Legacy BK','kg','8%','',0,1,'','')"""
        )
        conn.execute(
            """INSERT OR REPLACE INTO products(
                   code,name,unit,tax,supplier,buy_price,purchase_list,seller,cccd
               ) VALUES('QC-BK-CANONICAL','Canonical BK','kg','8%','QC-SOURCE',100,1,'','')"""
        )
        conn.execute(
            """INSERT INTO inventory_transactions(
                   txn_date,product_code,qty_in,qty_out,unit_cost,source_type,source_id,
                   source_line,status,note,created_at,updated_at
               ) VALUES('2026-08-31','QC-BK-LEGACY',99,0,100,'BK_INPUT','qc-legacy','1',
                        'posted','must not count',?,?)""",
            (server.now_iso(), server.now_iso()),
        )
        bk_ledger_before = (
            conn.execute("SELECT COUNT(*) FROM inventory_transactions").fetchone()[0],
            conn.execute("SELECT COUNT(*) FROM invoice_inventory_ledger").fetchone()[0],
        )
    bk_sheet.append([
        "2026-09-03", "", "BK-QC-0001", 1,
        "QC-BK-CANONICAL", "Canonical BK", "kg",
        2, 100, 200, "QC-SOURCE", "fixture",
    ])
    bk_stream = io.BytesIO()
    bk_book.save(bk_stream)
    bk_book.close()
    bk_payload = bk_stream.getvalue()

    def preview_bk():
        return client.post(
            "/api/bk-import/preview",
            data={"file": (io.BytesIO(bk_payload), "bk-qc.xlsx")},
            content_type="multipart/form-data",
        )

    bk_first_response = preview_bk()
    bk_second_response = preview_bk()
    assert bk_first_response.status_code == 200 and bk_second_response.status_code == 200
    bk_first = bk_first_response.get_json()
    bk_second = bk_second_response.get_json()
    assert bk_first["canConfirm"] is True
    assert bk_first["sourcePolicy"] == bk_import.BK_IMPORT_SOURCE_TYPE
    assert bk_first["previewOnly"] is True and bk_first["writesInventory"] is False
    assert bk_first["counts"]["readyRows"] == 1 and bk_first["counts"]["errorRows"] == 0
    assert (
        bk_first["sourceHash"], bk_first["contentHash"], bk_first["previewId"],
        bk_first["rows"][0]["sourceKey"],
    ) == (
        bk_second["sourceHash"], bk_second["contentHash"], bk_second["previewId"],
        bk_second["rows"][0]["sourceKey"],
    )
    bk_confirm = client.post("/api/bk-import/confirm", json={
        "confirmed": True, "token": bk_first["token"], "previewId": bk_first["previewId"],
    })
    assert bk_confirm.status_code == 200, bk_confirm.get_data(as_text=True)
    bk_confirm_body = bk_confirm.get_json()
    assert bk_confirm_body["newInventoryLines"] == 1
    bk_repeat = client.post("/api/bk-import/confirm", json={
        "confirmed": True, "token": bk_first["token"], "previewId": bk_first["previewId"],
    })
    assert bk_repeat.status_code == 200 and bk_repeat.get_json()["idempotent"] is True
    with server.db() as conn:
        bk_ledger_posted = (
            conn.execute("SELECT COUNT(*) FROM inventory_transactions").fetchone()[0],
            conn.execute("SELECT COUNT(*) FROM invoice_inventory_ledger").fetchone()[0],
        )
        legacy_gate = dict(server.inventory_lookup(conn).get("QC-BK-LEGACY") or {})
        bk_event = conn.execute(
            """SELECT direction,event_type,source_invoice_table,product_code,qty_delta
                 FROM invoice_inventory_ledger WHERE source_invoice_table='bk_import_documents'
                 ORDER BY id"""
        ).fetchone()
    assert bk_ledger_posted == (bk_ledger_before[0], bk_ledger_before[1] + 1)
    assert tuple(bk_event) == (
        "input", "POST", "bk_import_documents", "QC-BK-CANONICAL", 2,
    )
    assert legacy_gate.get("accounting_qty", 0) == 0
    bk_reversal = client.post(
        f"/api/bk-import/documents/{bk_confirm_body['documentId']}/reversal",
        json={
            "confirmed": True, "reversalDate": "2026-09-03",
            "reason": "QC hoàn tác có kiểm soát",
        },
    )
    assert bk_reversal.status_code == 200, bk_reversal.get_data(as_text=True)
    with server.db() as conn:
        assert conn.execute(
            """SELECT COUNT(*) FROM invoice_inventory_ledger
                WHERE source_invoice_table='bk_import_documents' AND event_type='REVERSAL'"""
        ).fetchone()[0] == 1
        assert conn.execute(
            "SELECT status FROM bk_import_documents WHERE id=?",
            (bk_confirm_body["documentId"],),
        ).fetchone()[0] == "reversed"
    bk_import_result = (
        "official template / fixed non-invoice purchase source / stable keys / "
        "canonical post + idempotency + audited safe reversal"
    )

    print("QC", json.dumps({
        "ok": True,
        "masterProducts": bootstrap["master"]["product_count"],
        "selectedSheetOrders": len(selected.get_json()["orders"]),
        "importedOrders": len(orders),
        "importWarnings": sum(bool(item["warnings"]) for item in orders),
        "oldFileApproval": "passed_with_warnings",
        "exportSizes": export_sizes,
        "manualAndPasteFlow": "passed",
        "debtFlow": "passed",
        "inventoryAndReturns": "passed",
        "msmiIdempotentSyncMappingReceipt": "passed",
        "outgoingDraftStockGate": "passed",
        "kitchenMenuCostPo": "passed",
        "kitchenWorkbookImport": "passed",
        "actualCustomerKitchenWorkbook": actual_kitchen_result,
        "actualMealAttendanceWorkbook": "205 daily kitchen/shift rows / 6,020 meals / repeat safe",
        "xcomPaymentDocuments": xcom_payment_result,
        "actualOpeningInventoryWorkbook": actual_opening_result,
        "actualCustomerCatalogWorkbook": catalog_import_result,
        "templatePreservingEngine": template_engine_result,
        "approvedDeliveryTemplate": delivery_template_result,
        "approvedPurchaseSummary": purchase_summary_result,
        "approvedPurchaseReceipt": receipt_template_result,
        "approvedMonthlyReport": monthly_report_result,
        "periodQuoteImport": quote_import_result,
        "toyotaGoldenQuote": toyota_quote_result,
        "actualCustomerTaxTemplates": tax_template_result,
        "outgoingInvoiceTaxExport": invoice_tax_export_result,
        "actualHistoricalPayablesWorkbook": actual_payables_result,
        "payablePaymentAllocation": payable_payment_result,
        "payableExcelExport": payable_export_result,
        "payableUiContract": payable_ui_result,
        "receivableLedger": receivable_qc_result,
        "receivableExcelExport": receivable_export_result,
        "receivableUiContract": receivable_ui_result,
        "invoiceInventoryFourFileExport": inventory_export_result,
        "invoiceInventoryUiContract": inventory_ui_result,
        "unifiedUiBusinessLanguage": unified_ui_result,
        "customerCompactUiContract": compact_customer_ui_result,
        "inputInvoiceImmediateExcelExport": invoice_input_export_result,
        "bkImportSafetyContract": bk_import_result,
        "outgoingReadinessMultiRound": outgoing_readiness_result,
        "outgoingSubstitutionConfirmation": outgoing_substitution_result,
        "invoicePaymentOfficialDocuments": invoice_payment_scope_result,
        "issuedInvoiceDeliveryStatement": invoice_delivery_statement_result,
        "attendancePayrollLegacyImport": "passed",
        "paymentRequestAndPrintApproval": "passed",
    }, ensure_ascii=False, indent=2))
    print_dir = APP_DIR / "data" / "print_jobs" / str(manual_id)
    if print_dir.exists() and print_dir.resolve().is_relative_to((APP_DIR / "data" / "print_jobs").resolve()):
        rmtree(print_dir)
    payment_qc = APP_DIR / "data" / "qc_payment_request.docx"
    if payment_qc.exists() and payment_qc.parent.resolve() == (APP_DIR / "data").resolve():
        payment_qc.unlink()
    payment_render_qc = APP_DIR / "data" / "qc_payment_request_render"
    if payment_render_qc.exists() and payment_render_qc.resolve().is_relative_to((APP_DIR / "data").resolve()):
        rmtree(payment_render_qc)
    clean_test_db(qc_db)


if __name__ == "__main__":
    main()
