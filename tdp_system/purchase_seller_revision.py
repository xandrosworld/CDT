"""Versioned seller corrections on purchase documents, independent of stock/debt."""
import hashlib
import io
import json
import threading
import time
import uuid
from zipfile import BadZipFile, ZipFile
from collections import defaultdict
from decimal import Decimal, ROUND_HALF_UP

from flask import jsonify, request
from openpyxl import load_workbook

try:
    from .purchase_summary_export import PurchaseSummaryError, _iso_date, _key
    from .contract_modules import invoice_tax_percent
except ImportError:
    from purchase_summary_export import PurchaseSummaryError, _iso_date, _key
    from contract_modules import invoice_tax_percent

PENDING = {}
LOCK = threading.Lock()


def digest(value):
    return hashlib.sha256(json.dumps(value,sort_keys=True,ensure_ascii=False,default=str).encode()).hexdigest()


def init_schema(c):
    c.execute('''CREATE TABLE IF NOT EXISTS purchase_seller_revisions(
        batch_id INTEGER NOT NULL, revision INTEGER NOT NULL, source_hash TEXT NOT NULL,
        source_json TEXT NOT NULL, rows_json TEXT NOT NULL, file_hash TEXT NOT NULL,
        filename TEXT NOT NULL, actor TEXT NOT NULL, reason TEXT NOT NULL, created_at TEXT NOT NULL,
        PRIMARY KEY(batch_id,revision))''')


def latest(c,batch_id):
    if not c.execute("SELECT 1 FROM sqlite_master WHERE name='purchase_seller_revisions'").fetchone():return None
    return c.execute('SELECT * FROM purchase_seller_revisions WHERE batch_id=? ORDER BY revision DESC LIMIT 1',(batch_id,)).fetchone()


def apply_revision(c,batch,rows):
    saved=latest(c,int(batch['id']))
    if not saved:return rows
    if saved['source_hash']!=digest(rows):
        raise PurchaseSummaryError('Nguồn mua đã đổi sau khi cập nhật người bán; cần đối chiếu bản bảng kê đã lưu.',code='seller_revision_stale')
    return json.loads(saved['rows_json'])


def number(value):
    try:
        n=Decimal(str(value))
        if isinstance(value,bool) or not n.is_finite() or n<0:raise ValueError()
        return n
    except Exception:
        raise ValueError('Số lượng/đơn giá trong file không hợp lệ.') from None


def business_key(party,code,kitchen,unit,price,tax):
    return (str(party).strip().upper(),str(code).strip().upper(),str(kitchen).strip().upper(),
            str(unit).strip().casefold(),number(price),invoice_tax_percent(tax))


def read_rows(blob,work_date):
    try:
        with ZipFile(io.BytesIO(blob)) as archive:
            entries=archive.infolist()
            if len(entries)>5000 or sum(e.file_size for e in entries)>100*1024*1024:
                raise ValueError('File Excel vượt giới hạn dung lượng giải nén.')
        book=load_workbook(io.BytesIO(blob),read_only=True,data_only=True,keep_links=False)
    except (BadZipFile, OSError, KeyError, ValueError):
        raise ValueError('Không đọc được file Excel. Hãy mở và lưu lại dưới dạng .xlsx.') from None
    try:
        found=[]
        for sheet in book:
            header=next(sheet.iter_rows(min_row=2,max_row=2,max_col=20,values_only=True),())
            if len(header)<20 or str(header[4]).strip().casefold()!='mã hàng' or str(header[19]).strip().casefold()!='cccd':continue
            if sheet.max_row>100000:raise ValueError('Sheet đơn hàng vượt giới hạn 100.000 dòng.')
            rows=[]
            for index,row in enumerate(sheet.iter_rows(min_row=3,max_col=20,values_only=True),3):
                if not isinstance(row[4],str) or not row[4].strip() or row[4].strip() in ('-','#N/A'):continue
                if row[8] in (None,'','-'):continue
                try:day=_iso_date(row[6])
                except ValueError:continue
                if day!=work_date:continue
                qty=number(row[8]);key=business_key(row[0],row[4],row[5],row[9],row[11],row[13])
                rows.append({'key':key,'qty':qty,'seller':str(row[18] or '').strip(),
                             'cccd':row[19],'row':index,'sheet':sheet.title})
            if rows:found.append(rows)
        if len(found)!=1:raise ValueError('Cần đúng một sheet đơn hàng của ngày đang chọn; chưa xác định duy nhất sheet để lấy người bán.')
        return found[0]
    finally:book.close()


def build_preview(c,blob,work_date,filename):
    try:
        from .purchase_summary_export import collect_purchase_summary_rows
        from .receipt_export import enrich_receipt_identity_rows, RECEIPT_MAX_DAILY_AMOUNT
        from .purchase_document_selection import day_source, saved_plan, group_totals
    except ImportError:
        from purchase_summary_export import collect_purchase_summary_rows
        from receipt_export import enrich_receipt_identity_rows, RECEIPT_MAX_DAILY_AMOUNT
        from purchase_document_selection import day_source, saved_plan, group_totals
    work_date=_iso_date(work_date)
    batches=c.execute("SELECT * FROM batches WHERE work_date=? AND status='approved'",(work_date,)).fetchall()
    if len(batches)!=1:raise ValueError('Ngày này phải có đúng một đơn đã duyệt để cập nhật người bán từ file.')
    batch=dict(batches[0]);bid=batch['id']
    orders=[dict(r) for r in c.execute('SELECT * FROM orders WHERE batch_id=? ORDER BY id',(bid,))]
    incoming=read_rows(blob,work_date)
    expected=defaultdict(Decimal);received=defaultdict(Decimal);by_group=defaultdict(list)
    for o in orders:
        key=business_key(o['contractor'],o['product_code'],o['kitchen'],o['unit'],o['sell_price'],o['tax'])
        expected[key]+=number(o['qty'])
    for r in incoming:
        received[r['key']]+=r['qty'];by_group[r['key']].append(r)
    if set(expected)!=set(received) or any(abs(v-received[k])>Decimal('.000001') for k,v in expected.items()):
        raise ValueError('File thay đổi mặt hàng, bếp, lượng bán, giá bán hoặc thuế so với đơn đã duyệt. Luồng này chỉ cập nhật người bán/tách lượng giữa người bán; không thay đơn.')
    base=collect_purchase_summary_rows(c,batch,orders,apply_seller_updates=False)
    previous=latest(c,bid);revision=(previous['revision'] if previous else 0)+1
    current=apply_revision(c,batch,base)
    source_day,source_day_hash=day_source(c,work_date)
    by_order={o['id']:o for o in orders}; people=defaultdict(list)
    for p in c.execute('SELECT * FROM people'):people[_key(p['name'])].append(dict(p))
    proposed=[];changes=[]
    base_keys=[]
    for row in base:
        _,table,source_id=row['selection_key'].split(':')
        if table=='purchase_workbook_lines':
            source=c.execute('SELECT * FROM purchase_workbook_lines WHERE id=?',(source_id,)).fetchone()
            order=by_order.get(source['order_id'])
        else:order=by_order.get(int(source_id))
        if not order:raise ValueError('Có dòng mua chưa liên kết duy nhất với đơn gốc; cần đối chiếu trước khi cập nhật người bán.')
        key=business_key(order['contractor'],order['product_code'],order['kitchen'],order['unit'],order['sell_price'],order['tax'])
        base_keys.append(key)
    for row,key in zip(base,base_keys):
        parts=by_group[key]
        names={_key(r['seller']) for r in parts}
        old_parts=[r for r in current if r['selection_key']==row['selection_key'] or r['selection_key'].startswith(row['selection_key']+':seller:')]
        for part in parts:
            matches=people[_key(part['seller'])]
            if len(matches)!=1:raise ValueError('Người bán '+part['seller']+' chưa có định danh duy nhất trong danh mục.')
            identity=str(matches[0]['cccd'] or '').strip()
            supplied=str(part['cccd'] or '').strip()
            if isinstance(part['cccd'],(int,float)) and float(part['cccd']).is_integer():supplied=str(int(part['cccd'])).zfill(len(identity))
            if not identity or supplied!=identity:raise ValueError('CCCD trong file không khớp danh mục của '+part['seller']+'.')
        # Unchanged identities do not need quantity redistribution or new keys.
        if names=={_key(row['seller'])}:
            new_parts=[row]
        else:
            if base_keys.count(key)!=1:raise ValueError('Nhiều dòng mua cùng mặt hàng/bếp cần tách người bán; chưa đủ căn cứ phân bổ tự động.')
            if abs(sum((r['qty'] for r in parts),Decimal(0))-number(row['quantity']))>Decimal('.000001'):
                raise ValueError('Lượng trong file không khớp lượng mua đã ghi kho của '+row['product_name']+'. Không tự chia lại số lượng mua.')
            allocations=defaultdict(Decimal);identities={}
            for part in parts:
                if not part['qty']:continue
                person=people[_key(part['seller'])][0];identity=str(person['cccd']).strip()
                allocations[identity]+=part['qty'];identities[identity]=person
            remaining=number(row['amount']);new_parts=[]
            for index,(identity,qty) in enumerate(sorted(allocations.items()),1):
                person=identities[identity]
                value=remaining if index==len(allocations) else (number(row['amount'])*qty/number(row['quantity'])).quantize(Decimal('.01'),rounding=ROUND_HALF_UP)
                remaining-=value
                new_parts.append({**row,'seller':person['name'],'cccd':identity,'address':person['address'],
                   'quantity':float(qty),'amount':float(value),
                   'selection_key':row['selection_key']+f':seller:{revision}:{index}'})
        summarize=lambda rs:sorted((r['seller'],r['cccd'],float(r['quantity']),float(r['amount'])) for r in rs)
        if summarize(old_parts)!=summarize(new_parts):
            changes.append({'product_name':row['product_name'],'kitchen':row['kitchen'],'source_ref':row['source_ref'],
                            'before':[{'seller':r['seller'],'quantity':r['quantity'],'amount':r['amount']} for r in old_parts],
                            'after':[{'seller':r['seller'],'quantity':r['quantity'],'amount':r['amount']} for r in new_parts]})
            proposed.extend(new_parts)
        else:
            proposed.extend(old_parts)
    if not changes:
        return {'unchanged':True,'date':work_date,'changes':[]}
    enriched=enrich_receipt_identity_rows(c,proposed)
    totals=group_totals(enriched)
    groups=[{'seller':r['seller'],'amount':float(r['amount']),'over_limit':r['amount']>RECEIPT_MAX_DAILY_AMOUNT} for r in totals.values()]
    prior_plan=saved_plan(c,work_date)
    if prior_plan and prior_plan['source_hash']!=source_day_hash:
        raise ValueError('Lựa chọn bảng kê đã cũ so với nguồn. Cần đối chiếu bản đã lưu trước khi cập nhật người bán.')
    token=digest([digest(base),dict(previous) if previous else None,source_day_hash,prior_plan,dict(batch),orders,enriched])
    return {'date':work_date,'batch_id':bid,'revision':revision,'changes':changes,'groups':groups,
            'total':float(sum(number(r['amount']) for r in base)),'state_hash':token,
            'source_hash':digest(base),'base_rows':base,'rows':proposed,'filename':filename,
            'file_hash':hashlib.sha256(blob).hexdigest(),'previous_plan':prior_plan}


def save_preview(c,preview,blob,actor,reason,timestamp):
    try:
        from .purchase_document_selection import day_source, init_schema as init_selection, group_totals
        from .receipt_export import RECEIPT_MAX_DAILY_AMOUNT
    except ImportError:
        from purchase_document_selection import day_source, init_schema as init_selection, group_totals
        from receipt_export import RECEIPT_MAX_DAILY_AMOUNT
    actor=str(actor or '').strip();reason=str(reason or '').strip()
    if not actor or len(actor)>120 or not reason or len(reason)>500:raise ValueError('Điền người lưu và lý do cập nhật người bán (tối đa 120/500 ký tự).')
    current=build_preview(c,blob,preview['date'],preview['filename'])
    if current.get('unchanged'):return current
    if current['state_hash']!=preview['state_hash']:raise ValueError('Nguồn hoặc bảng kê vừa thay đổi. Xem trước lại file trước khi lưu.')
    init_schema(c)
    c.execute('INSERT INTO purchase_seller_revisions VALUES(?,?,?,?,?,?,?,?,?,?)',
              (current['batch_id'],current['revision'],current['source_hash'],json.dumps(current['base_rows'],ensure_ascii=False),
               json.dumps(current['rows'],ensure_ascii=False),current['file_hash'],current['filename'],actor,reason,timestamp))
    prior=current['previous_plan']
    init_selection(c)
    rows,source_hash=day_source(c,current['date'])
    if prior:
        # Preserve unchanged selections; corrected/split lines need a new selection.
        old=json.loads(prior['quantities_json'])
        quantities={r['selection_key']:old[r['selection_key']] for r in rows if r['selection_key'] in old}
    else:
        totals=group_totals(rows)
        quantities={r['selection_key']:str(r['quantity']) for r in rows if totals[r['cccd']]['amount']<=RECEIPT_MAX_DAILY_AMOUNT}
    values=(current['date'],prior['revision']+1 if prior else 1,source_hash,json.dumps(quantities,sort_keys=True),actor,reason,timestamp)
    c.execute('INSERT OR REPLACE INTO purchase_document_selections VALUES(?,?,?,?,?,?,?)',values)
    c.execute('INSERT INTO purchase_document_selection_history VALUES(?,?,?,?,?,?,?,?)',values[:4]+(json.dumps(rows,ensure_ascii=False),)+values[4:])
    return {'ok':True,'date':current['date'],'revision':current['revision'],'changes':current['changes']}


def register_routes(app,ctx):
    @app.post('/api/purchase-sellers/preview')
    def preview_sellers():
        try:
            upload=request.files.get('file')
            if not upload or not upload.filename.lower().endswith('.xlsx'):raise ValueError('Chọn file Excel .xlsx đã sửa người bán.')
            blob=upload.read(20*1024*1024+1)
            if len(blob)>20*1024*1024:raise ValueError('File vượt quá 20 MB.')
            with ctx['db']() as c:
                c.execute('BEGIN')
                p=build_preview(c,blob,request.form.get('date'),upload.filename)
            if p.get('unchanged'):return jsonify(ok=True,**p)
            token=uuid.uuid4().hex
            with LOCK:
                for k in list(PENDING):
                    if PENDING[k]['expires']<time.time():PENDING.pop(k)
                if len(PENDING)>=10:raise ValueError('Có nhiều bản xem trước đang mở. Hãy đóng và thử lại sau.')
                PENDING[token]={'preview':p,'blob':blob,'expires':time.time()+900}
            return jsonify(ok=True,token=token,**{k:p[k] for k in ('date','revision','changes','groups','total')})
        except (ValueError,TypeError,KeyError) as exc:return jsonify(ok=False,error=str(exc)),409

    @app.post('/api/purchase-sellers/confirm')
    def confirm_sellers():
        try:
            body=request.get_json(silent=True) or {}
            if not isinstance(body,dict):raise ValueError('Nội dung xác nhận không hợp lệ.')
            if body.get('confirmed') is not True:raise ValueError('Xác nhận thông tin người bán đúng thực tế trước khi lưu.')
            with LOCK:p=PENDING.get(str(body.get('token') or ''))
            if not p or p['expires']<time.time():raise ValueError('Bản xem trước hết hạn. Chọn lại file.')
            with ctx['db']() as c:
                c.execute('BEGIN IMMEDIATE')
                result=save_preview(c,p['preview'],p['blob'],body.get('actor'),body.get('reason'),ctx['now_iso']())
                ctx['audit_event'](c,'purchase_document.sellers',entity_type='batch',entity_id=str(p['preview']['batch_id']),
                                   metadata={'revision':result.get('revision'),'file_hash':p['preview']['file_hash'],'actor':body.get('actor'),'reason':body.get('reason')})
            return jsonify(ok=True,**{k:v for k,v in result.items() if k!='ok'})
        except (ValueError,TypeError,KeyError) as exc:return jsonify(ok=False,error=str(exc)),409
