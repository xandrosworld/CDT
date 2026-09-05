import io
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

try:
    from . import contract_modules, server
except ImportError:  # pragma: no cover - direct file invocation
    import contract_modules
    import server


class PurchaseSalesBoundaryRegressionTests(unittest.TestCase):
    """Lock purchasing/NCC actions away from sales, receivables and invoice stock."""

    @classmethod
    def setUpClass(cls):
        cls.temp_dir = tempfile.TemporaryDirectory()
        cls.original_db_path = server.DB_PATH
        cls.original_data_dir = server.DATA_DIR
        cls.original_testing = server.app.config.get("TESTING", False)
        server.DATA_DIR = Path(cls.temp_dir.name) / "data"
        server.DATA_DIR.mkdir(parents=True, exist_ok=True)
        server.DB_PATH = server.DATA_DIR / "purchase_sales_boundary.sqlite3"
        server.app.config["TESTING"] = True
        server.init_database()
        cls.client = server.app.test_client()

    @classmethod
    def tearDownClass(cls):
        server.DB_PATH = cls.original_db_path
        server.DATA_DIR = cls.original_data_dir
        server.app.config["TESTING"] = cls.original_testing
        cls.temp_dir.cleanup()

    def setUp(self):
        pending = getattr(contract_modules, "PENDING_PURCHASE_ORDER_IMPORTS", None)
        if pending is not None:
            pending.clear()
        with server.db() as conn:
            conn.execute("DELETE FROM receivable_ledger_revisions")
            conn.execute("DELETE FROM receivable_ledger_lines")
            for table in (
                "invoice_inventory_ledger", "invoice_inventory_confirmations",
                "outgoing_invoice_lines", "outgoing_invoice_drafts",
                "inventory_transactions", "supplier_order_statuses",
                "purchase_workbook_line_revisions", "purchase_workbook_lines",
                "purchase_order_imports", "purchase_order_lines",
                "order_import_receipts", "payments", "balances", "debt_adjustments",
                "orders", "batches", "supplier_rules", "audit_log",
            ):
                conn.execute(f"DELETE FROM {table}")
            conn.execute(
                "INSERT OR REPLACE INTO contractors(code,name,price_group,pricing_mode) "
                "VALUES('HATRAN','Hà Trân','HATRAN','group')"
            )
            for code in ("POT", "BIA"):
                conn.execute(
                    """INSERT OR REPLACE INTO kitchens(code,contractor,name,address,show_price)
                       VALUES(?,'HATRAN',?,'',0)""",
                    (code, f"Bếp {code}"),
                )
            conn.execute("INSERT OR REPLACE INTO suppliers(code,name) VALUES('DUNG','Nhà Dũng')")
            conn.execute("INSERT OR REPLACE INTO suppliers(code,name) VALUES('KHO','Kho')")
            conn.execute(
                """INSERT OR REPLACE INTO products(
                       code,name,unit,tax,supplier,buy_price,purchase_list,seller,cccd
                   ) VALUES('BOUND-P1','Cà rốt','kg','0%','DUNG',0,0,'','')"""
            )
            conn.execute(
                """INSERT OR REPLACE INTO products(
                       code,name,unit,tax,supplier,buy_price,purchase_list,seller,cccd
                   ) VALUES('BOUND-P2','Khoai tây','kg','0%','KHO',0,0,'','')"""
            )

    @staticmethod
    def _insert_order(conn, batch_id, *, kitchen, product_code, product_name, qty,
                      supplier, sell_price):
        timestamp = server.now_iso()
        return int(conn.execute(
            """INSERT INTO orders(
                   batch_id,work_date,contractor,kitchen,product_code,product_name,qty,
                   actual_received,actual_delivered,unit,supplier,buy_price,sell_price,tax,
                   purchase_list,note,errors,warnings,updated_at
               ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,0,'','[]','[]',?)""",
            (
                batch_id, "2026-09-01", "HATRAN", kitchen, product_code,
                product_name, qty, qty, qty, "kg", supplier, 0, sell_price,
                "0%", timestamp,
            ),
        ).lastrowid)

    def _seed_boundary_fixture(self):
        with server.db() as conn:
            timestamp = server.now_iso()
            batch_id = int(conn.execute(
                """INSERT INTO batches(work_date,source_name,status,created_at,approved_at)
                   VALUES('2026-09-01','boundary-customer-order.xlsx','approved',?,?)""",
                (timestamp, timestamp),
            ).lastrowid)
            self._insert_order(
                conn, batch_id, kitchen="POT", product_code="BOUND-P1",
                product_name="Cà rốt", qty=7, supplier="DUNG", sell_price=30000,
            )
            self._insert_order(
                conn, batch_id, kitchen="BIA", product_code="BOUND-P1",
                product_name="cà rốt", qty=3, supplier="Nhà Dũng", sell_price=30000,
            )
            self._insert_order(
                conn, batch_id, kitchen="POT", product_code="BOUND-P2",
                product_name="Khoai tây", qty=4, supplier="KHO", sell_price=40000,
            )

            conn.execute(
                """INSERT INTO inventory_transactions(
                       txn_date,product_code,qty_in,qty_out,unit_cost,source_type,
                       source_id,source_line,status,note,created_at,updated_at
                   ) VALUES('2026-08-31','BOUND-P1',12,0,10000,'OPENING',
                            'BOUND-OPENING','1','posted','invoice opening',?,?)""",
                (timestamp, timestamp),
            )
            confirmation_id = int(conn.execute(
                """INSERT INTO invoice_inventory_confirmations(
                       confirmation_key,direction,source_invoice_table,source_invoice_id,
                       action,confirmed,note,created_at
                   ) VALUES('boundary-input-post','input','msmi_invoices',99001,
                            'post',1,'fixture',?)""",
                (timestamp,),
            ).lastrowid)
            conn.execute(
                """INSERT INTO invoice_inventory_ledger(
                       event_key,direction,event_type,source_invoice_table,source_invoice_id,
                       source_line_id,source_line_index,product_code,txn_date,qty_delta,
                       unit_cost,mapping_revision_id,confirmation_id,reverses_event_key,
                       status,created_at
                   ) VALUES('boundary-ledger-post','input','POST','msmi_invoices',99001,
                            99101,1,'BOUND-P1','2026-09-01',40,12000,NULL,?,'','posted',?)""",
                (confirmation_id, timestamp),
            )
            server.sync_receivable_ledger(conn, timestamp=timestamp)
        return batch_id

    @staticmethod
    def _rows(conn, sql, params=()):
        return [tuple(row) for row in conn.execute(sql, params)]

    def _protected_snapshot(self, batch_id):
        with server.db() as conn:
            database = {
                "batch": self._rows(conn, "SELECT * FROM batches WHERE id=?", (batch_id,)),
                "orders": self._rows(
                    conn, "SELECT * FROM orders WHERE batch_id=? ORDER BY id", (batch_id,),
                ),
                "payments": self._rows(conn, "SELECT * FROM payments ORDER BY id"),
                "balances": self._rows(
                    conn, "SELECT * FROM balances ORDER BY party_type,party_code",
                ),
                "debt_adjustments": self._rows(
                    conn, "SELECT * FROM debt_adjustments ORDER BY id",
                ),
                "receivable_ledger": self._rows(
                    conn, "SELECT * FROM receivable_ledger_lines ORDER BY id",
                ),
                "receivable_revisions": self._rows(
                    conn, "SELECT * FROM receivable_ledger_revisions ORDER BY id",
                ),
                "legacy_invoice_stock": self._rows(
                    conn, "SELECT * FROM inventory_transactions ORDER BY id",
                ),
                "invoice_confirmations": self._rows(
                    conn, "SELECT * FROM invoice_inventory_confirmations ORDER BY id",
                ),
                "invoice_ledger": self._rows(
                    conn, "SELECT * FROM invoice_inventory_ledger ORDER BY id",
                ),
                "outgoing_drafts": self._rows(
                    conn, "SELECT * FROM outgoing_invoice_drafts ORDER BY id",
                ),
                "outgoing_lines": self._rows(
                    conn, "SELECT * FROM outgoing_invoice_lines ORDER BY id",
                ),
            }
        bootstrap = self.client.get(f"/api/bootstrap?batch_id={batch_id}")
        self.assertEqual(bootstrap.status_code, 200, bootstrap.get_data(as_text=True))
        bootstrap_payload = bootstrap.get_json()
        debts = self.client.get("/api/debts?from=2026-09-01&to=2026-09-01")
        self.assertEqual(debts.status_code, 200, debts.get_data(as_text=True))
        invoice_stock = self.client.get("/api/invoice-inventory?as_of=2026-09-01")
        self.assertEqual(invoice_stock.status_code, 200, invoice_stock.get_data(as_text=True))
        return {
            "database": database,
            "sales_summary": {
                "totals": bootstrap_payload["summary"]["totals"],
                "contractors": bootstrap_payload["summary"]["contractors"],
                "kitchens": bootstrap_payload["summary"]["kitchens"],
            },
            "receivables": debts.get_json()["contractors"],
            "invoice_stock": invoice_stock.get_json(),
        }

    @staticmethod
    def _edited_purchase_file(payload, *, pot_price, bia_price):
        workbook = load_workbook(io.BytesIO(payload))
        try:
            sheet = workbook["đặt hàng"]
            seen = set()
            for row_index in range(3, sheet.max_row + 1):
                kitchen = str(sheet.cell(row_index, 3).value or "").strip().upper()
                supplier = contract_modules.supplier_merge_key(sheet.cell(row_index, 8).value)
                if supplier == "dung":
                    price = pot_price if kitchen == "POT" else bia_price
                    sheet.cell(row_index, 10, price)
                    if kitchen == "POT":
                        sheet.cell(row_index, 11, 1)  # hỏng
                        sheet.cell(row_index, 12, 2)  # thêm; actual = 8
                    seen.add((supplier, kitchen))
                elif supplier == "kho":
                    sheet.cell(row_index, 10, 0)
                    seen.add((supplier, kitchen))
            if seen != {("dung", "POT"), ("dung", "BIA"), ("kho", "POT")}:
                raise AssertionError(f"Unexpected supplier export rows: {seen!r}")
            stream = io.BytesIO()
            workbook.save(stream)
            return stream.getvalue()
        finally:
            workbook.close()

    def _import_purchase_file(self, batch_id, payload, filename):
        preview_response = self.client.post(
            "/api/purchase-orders/import/preview",
            data={
                "batch_id": str(batch_id),
                "file": (io.BytesIO(payload), filename),
            },
            content_type="multipart/form-data",
        )
        self.assertEqual(
            preview_response.status_code, 200, preview_response.get_data(as_text=True),
        )
        preview = preview_response.get_json()
        self.assertTrue(preview["can_confirm"], preview["issues"])
        self.assertEqual(preview["error_rows"], 0)
        confirm = self.client.post(
            "/api/purchase-orders/import/confirm",
            json={"token": preview["token"], "confirmed": True},
        )
        self.assertEqual(confirm.status_code, 200, confirm.get_data(as_text=True))
        return preview, confirm.get_json()

    def test_full_ncc_sequence_never_mutates_sales_receivables_or_invoice_stock(self):
        batch_id = self._seed_boundary_fixture()
        protected = self._protected_snapshot(batch_id)
        self.assertEqual(protected["sales_summary"]["totals"]["revenue"], 460000)
        self.assertEqual(protected["receivables"]["HATRAN"]["period_charge"], 460000)
        self.assertEqual(protected["invoice_stock"]["items"][0]["closing_qty"], 52)

        exported = self.client.get(f"/api/export/suppliers/{batch_id}")
        self.assertEqual(exported.status_code, 200, exported.content_type)
        first_file = self._edited_purchase_file(
            exported.data, pot_price=50000, bia_price=50000,
        )
        first_preview, first_confirm = self._import_purchase_file(
            batch_id, first_file, "dat-hang-da-sua.xlsx",
        )
        self.assertEqual((first_preview["count"], first_confirm["processed"]), (3, 3))
        self.assertEqual(first_preview["total_amount"], 550000)
        self.assertEqual(self._protected_snapshot(batch_id), protected)

        # Reading/merging is the server-side half of image generation. The
        # browser smoke captures the canvas and downloads the actual PNG.
        needs = self.client.get(f"/api/supplier-needs/{batch_id}")
        self.assertEqual(needs.status_code, 200, needs.get_data(as_text=True))
        need_payload = needs.get_json()
        dung = next(group for group in need_payload["groups"] if group["supplier_key"] == "dung")
        self.assertEqual((dung["raw_line_count"], dung["presented_line_count"]), (2, 1))
        self.assertEqual((dung["items"][0]["order_qty"], dung["items"][0]["amount"]), (11, 550000))
        self.assertEqual(
            [column["key"] for column in need_payload["image_contract"]["columns"]],
            ["kitchen", "work_date", "product_name", "order_qty", "unit", "supplier", "note"],
        )
        self.assertEqual(self._protected_snapshot(batch_id), protected)

        ordered = self.client.put(
            f"/api/supplier-order-status/{batch_id}/dung",
            json={"status": "ordered", "revision": 0},
        )
        self.assertEqual(ordered.status_code, 200, ordered.get_data(as_text=True))
        self.assertEqual(self._protected_snapshot(batch_id), protected)
        reopened = self.client.put(
            f"/api/supplier-order-status/{batch_id}/dung",
            json={"status": "reopened", "revision": 1},
        )
        self.assertEqual(reopened.status_code, 200, reopened.get_data(as_text=True))
        self.assertEqual(self._protected_snapshot(batch_id), protected)

        second_export = self.client.get(f"/api/export/suppliers/{batch_id}")
        self.assertEqual(second_export.status_code, 200, second_export.content_type)
        second_file = self._edited_purchase_file(
            second_export.data, pot_price=60000, bia_price=50000,
        )
        second_preview, second_confirm = self._import_purchase_file(
            batch_id, second_file, "dat-hang-gia-lan-hai.xlsx",
        )
        self.assertEqual((second_preview["error_rows"], second_confirm["processed"]), (0, 1))
        self.assertEqual(second_preview["total_amount"], 630000)
        self.assertEqual(self._protected_snapshot(batch_id), protected)

        debts = self.client.get("/api/debts?from=2026-09-01&to=2026-09-01").get_json()
        self.assertEqual(debts["contractors"]["HATRAN"]["period_charge"], 460000)
        self.assertEqual(sum(row["period_charge"] for row in debts["suppliers"].values()), 630000)
        with server.db() as conn:
            kho = conn.execute(
                """SELECT actual_qty,buy_price,amount FROM purchase_workbook_lines
                   WHERE batch_id=? AND LOWER(supplier)='kho'""",
                (batch_id,),
            ).fetchone()
            revisions = conn.execute(
                "SELECT COUNT(*) FROM purchase_workbook_line_revisions WHERE batch_id=?",
                (batch_id,),
            ).fetchone()[0]
        self.assertEqual(tuple(kho), (4, 0, 0))
        self.assertEqual(revisions, 4)


if __name__ == "__main__":
    unittest.main()
