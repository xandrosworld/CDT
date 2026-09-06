"""Read-only, date-scoped invoice rows shared by the screen and its Excel export."""
from __future__ import annotations

from decimal import Decimal
from io import BytesIO

try:
    from invoice_date_migration import input_date_repair_report
except ImportError:
    from .invoice_date_migration import input_date_repair_report

try:
    from invoice_workbench import INPUT_INVOICE, InvoiceWorkbenchError, normalize_invoice_type, validate_date_range
    from invoice_input_sync import input_invoice_payload
    from invoice_output_sync import output_invoice_payload
except ImportError:
    from .invoice_workbench import INPUT_INVOICE, InvoiceWorkbenchError, normalize_invoice_type, validate_date_range
    from .invoice_input_sync import input_invoice_payload
    from .invoice_output_sync import output_invoice_payload


STATUS_LABELS = {
    "all": "Tất cả hóa đơn", "needs_mapping": "Chưa ghép đủ mã / đơn vị",
    "ready": "Sẵn sàng ghi kho", "posted": "Đã ghi kho",
    "error": "Cần kiểm tra", "reversed": "Đã hoàn tác kho", "not_inventory": "Không ghi kho",
}
LINE_LABELS = {
    "all": "Tất cả dòng", "needs_attention": "Dòng cần xử lý",
    "unmapped": "Chưa ghép mã", "unit_review": "Cần quy đổi đơn vị",
    "error": "Lỗi / bất thường", "mapped": "Đã ghép mã",
}


def invoice_state(invoice, direction):
    stock = invoice.get("receipt_status" if direction == "input" else "stock_status", "")
    # A changed/cancelled source needs review even if there is an older posting.
    if stock == "reversal_required" or invoice.get("sync_status") != "synced" or invoice.get("error_message"):
        return "error"
    if stock in {"posted", "reversed"}:
        return stock
    if direction == "output" and invoice.get("source_status_class") != "issued":
        return "error"
    if stock == "ready":
        return "ready"
    if stock == "not_inventory":
        return "not_inventory"
    if any(item.get("inventory_eligible") and item.get("mapping_status") != "mapped"
           for item in invoice.get("items", [])):
        return "needs_mapping"
    return "error"


def line_issue(item, invoice, status):
    if status == "error":
        return invoice.get("error_message") or item.get("validation_note") or "Kiểm tra trạng thái hóa đơn trước khi ghi kho"
    if item.get("inventory_eligible") and item.get("mapping_status") != "mapped":
        return "Cần quy đổi đơn vị" if item.get("mapping_status") == "unit_review" else "Chưa ghép đúng mã hàng"
    return ""


def invoice_range_payload(conn, *, tenant, invoice_type, date_from, date_to, status="all", line_filter="all"):
    safe_type = normalize_invoice_type(invoice_type)
    start, end = validate_date_range(date_from, date_to)
    if status not in STATUS_LABELS or line_filter not in LINE_LABELS:
        raise InvoiceWorkbenchError("Bộ lọc hóa đơn không hợp lệ")
    direction = "input" if safe_type == INPUT_INVOICE else "output"
    table, links, source = (
        ("msmi_invoices", "invoice_sync_batch_invoices", "msmi") if direction == "input"
        else ("outgoing_source_invoices", "invoice_sync_batch_output_invoices", "minvoice")
    )
    configured_tenant = conn.execute("SELECT value FROM settings WHERE key='tenant_code'").fetchone()
    local_tenant = str(configured_tenant[0] if configured_tenant else "TDP").strip() or "TDP"
    # Pre-workbench imports have no batch link. They belong to this local
    # database's configured company, not an arbitrary requested tenant.
    include_legacy = (str(tenant or "TDP").strip() or "TDP") == local_tenant
    source_guard = "i.invoice_type=?" if direction == "input" else "i.source=? AND i.tenant=?"
    source_params = [safe_type] if direction == "input" else [source, str(tenant or "TDP").strip() or "TDP"]
    # EXISTS deduplicates repeated/overlapping downloads. Filter actual invoice
    # dates, NOT the sync window (which may span a whole month or several months).
    ids = [row[0] for row in conn.execute(
        f"""SELECT i.id FROM {table} i WHERE i.invoice_date>=? AND i.invoice_date<=? AND {source_guard}
            AND (EXISTS (SELECT 1 FROM {links} bi JOIN invoice_sync_batches b ON b.id=bi.batch_id
                         WHERE bi.invoice_id=i.id AND b.tenant=? AND b.source=? AND b.invoice_type=?)
                 OR (? AND NOT EXISTS (SELECT 1 FROM {links} old WHERE old.invoice_id=i.id)))""",
        (start, end, *source_params,
         str(tenant or "TDP").strip() or "TDP", source, safe_type, include_legacy),
    )]
    fetch = input_invoice_payload if direction == "input" else output_invoice_payload
    invoices = fetch(conn, invoice_ids=ids)["items"]
    counts = {key: 0 for key in STATUS_LABELS}
    lines, visible_invoices = [], []
    qty_by_unit, line_amount, invoice_amount = {}, Decimal(0), Decimal(0)
    for invoice in invoices:
        state = invoice_state(invoice, direction)
        invoice["workbench_status"] = state
        counts["all"] += 1
        counts[state] += 1
        if status != "all" and status != state:
            continue
        selected = []
        # Preserve visibility of invoices whose source sent no detail rows.
        for item in invoice["items"] or [{"id": None, "source_item_name": "Hóa đơn chưa có chi tiết hàng"}]:
            issue = line_issue(item, invoice, state)
            mapping = item.get("mapping_status", "review")
            eligible = item.get("inventory_eligible")
            visible = (
                line_filter == "all"
                or (line_filter == "needs_attention" and bool(issue))
                or (line_filter == "unmapped" and eligible and mapping == "unmapped")
                or (line_filter == "unit_review" and eligible and mapping == "unit_review")
                or (line_filter == "mapped" and eligible and mapping == "mapped")
                or (line_filter == "error" and (state == "error" or mapping == "review"))
            )
            if not visible:
                continue
            # Confirmation is a next action, not a validation error/red row.
            action_rank = 0 if issue else 1 if state == "ready" else 2
            line = {**item, "invoice_id": invoice["id"], "issue": issue,
                    "action_rank": action_rank, "needs_confirmation": state == "ready"}
            selected.append(line)
            unit = str(item.get("source_unit") or "Không có ĐVT")
            if item.get("id") is not None:
                qty_by_unit[unit] = qty_by_unit.get(unit, Decimal(0)) + Decimal(str(item.get("qty") or 0))
                line_amount += Decimal(str(item.get("amount") or 0))
        if selected:
            visible_invoices.append(invoice)
            lines.extend(selected)
            invoice_amount += Decimal(str(invoice.get("total_amount") or 0))
    # Stable within the original date/invoice order; ALL troublesome lines first.
    lines.sort(key=lambda item: (item["action_rank"], item.get("mapping_status") != "unmapped"))
    return {
        "direction": direction, "source": source, "date_from": start, "date_to": end,
        "status": status, "line_filter": line_filter, "items": visible_invoices,
        "lines": lines, "counts": counts, "status_labels": STATUS_LABELS, "line_labels": LINE_LABELS,
        "totals": {
            "invoice_count": len(visible_invoices), "line_count": sum(item.get("id") is not None for item in lines),
            "issue_count": sum(bool(item["issue"]) for item in lines),
            "qty_by_unit": {unit: float(qty) for unit, qty in sorted(qty_by_unit.items())},
            "line_amount": float(line_amount), "invoice_amount": float(invoice_amount),
        },
        "read_only": True,
        "date_repair": input_date_repair_report(conn, tenant=str(tenant or "TDP").strip() or "TDP",
                                                 date_from=start, date_to=end) if direction == "input" else {},
    }


def range_workbook(payload):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, Side
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = "Hoa don"
    title = "HÓA ĐƠN ĐẦU VÀO" if payload["direction"] == "input" else "HÓA ĐƠN ĐẦU RA"
    ws.append([title + " · " + payload["date_from"] + " → " + payload["date_to"]])
    ws.merge_cells("A1:M1")
    ws.append(["Ngày", "Ký hiệu / Số HĐ", "Đối tác", "Dòng", "Tên hàng", "ĐVT", "Số lượng",
               "Đơn giá", "Thành tiền dòng (chưa thuế)", "Mã kho", "Lượng kho", "ĐVT kho", "Trạng thái / Cần xử lý"])
    invoices = {invoice["id"]: invoice for invoice in payload["items"]}
    for line in payload["lines"]:
        invoice = invoices[line["invoice_id"]]
        ws.append([
            invoice["invoice_date"], (invoice.get("invoice_series", "") + " / " + invoice.get("invoice_number", "")),
            invoice.get("seller_name") or invoice.get("buyer_name") or "", line.get("line_index"),
            line.get("source_item_name", ""), line.get("source_unit", ""), line.get("qty"),
            line.get("unit_price"), line.get("amount"),
            line.get("product_code", ""), line.get("stock_qty"), line.get("product_unit", ""),
            line["issue"] or STATUS_LABELS[invoice["workbench_status"]],
        ])
    total_row = ws.max_row + 1
    ws.append(["TỔNG TIỀN CÁC DÒNG ĐANG LỌC", None, None, None, None, None, None, None, payload["totals"]["line_amount"]])
    for unit, qty in payload["totals"]["qty_by_unit"].items():
        ws.append(["TỔNG LƯỢNG", None, None, None, None, unit, qty])
    ws.append(["TỔNG THANH TOÁN CÁC HÓA ĐƠN CÓ DÒNG ĐANG LỌC (mỗi hóa đơn tính một lần)",
               None, None, None, None, None, None, None, payload["totals"]["invoice_amount"]])
    ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=8)
    widths = [13, 24, 30, 7, 38, 10, 15, 18, 22, 14, 15, 12, 42]
    edge = Side(style="thin", color="B8C2CA")
    border = Border(left=edge, right=edge, top=edge, bottom=edge)
    for cells in ws:
        for cell in cells:
            # Untrusted invoice strings must never become Excel formulas.
            if cell.data_type == "f" or isinstance(cell.value, str):
                cell.data_type = "s"
            cell.font = Font(name="Arial", size=10, bold=cell.row <= 2 or cell.row >= total_row)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = border
            if cell.column in {7, 11}:
                cell.number_format = "#,##0.######"
            if cell.column in {8, 9}:
                cell.number_format = "#,##0"
        ws.row_dimensions[cells[0].row].height = max(36, 14 * max(
            (len(str(cell.value or "")) // max(1, int(widths[cell.column - 1]) - 3) + 1) for cell in cells
        ))
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    repair = payload.get("date_repair") or {}
    if repair.get("blocked_count"):
        review = wb.create_sheet("Ngay can doi chieu")
        review.append(["Hóa đơn cần đối chiếu ngày; sổ kho được giữ nguyên. Không cộng bảng này vào tổng hóa đơn."])
        review.append(["Ký hiệu", "Số hóa đơn", "Bên bán", "Ngày đang lưu", "Ngày theo nguồn VN", "Lý do"])
        for item in repair["items"]:
            review.append([item[key] for key in ("invoice_series", "invoice_number", "seller_name",
                                                "old_date", "new_date", "reason")])
        for cells in review:
            for cell in cells:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                cell.alignment = Alignment(wrap_text=True, vertical="center")
        for col, width in zip("ABCDEF", (18, 18, 42, 18, 22, 70)):
            review.column_dimensions[col].width = width
        review.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:M{max(2, total_row - 1)}"
    ws.freeze_panes = "E3"
    ws.print_title_rows = "1:2"
    ws.print_area = f"A1:M{ws.max_row}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output
