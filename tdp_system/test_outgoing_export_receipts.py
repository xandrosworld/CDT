import unittest
import io
import zipfile
from openpyxl import load_workbook
from . import server
from . import test_outgoing_upload as fixtures


class ExportReceiptTests(unittest.TestCase):
    setUpClass=classmethod(fixtures.OutgoingUploadTests.setUpClass.__func__)
    tearDownClass=classmethod(fixtures.OutgoingUploadTests.tearDownClass.__func__)
    business=fixtures.OutgoingUploadTests.business
    export=fixtures.OutgoingUploadTests.export

    def setUp(self):
        fixtures.OutgoingUploadTests.setUp(self)
        with server.db() as c:c.execute('DELETE FROM outgoing_export_receipts')

    def report(self):
        r=self.client.get('/api/outgoing-invoices/unissued?from=2026-09-01&to=2026-09-30&contractor=NT-A')
        self.assertEqual(r.status_code,200,r.json)
        return r.json

    def test_receipt_only_after_success_and_reopen_preserves_business(self):
        before=self.business()
        plan=self.client.post('/api/outgoing-invoice-upload/stock-preview?to=2026-09-30&contractor=NT-A').json
        response=self.export(plan['token']);self.assertEqual(response.status_code,200)
        token=response.headers['X-Export-Receipt']
        self.assertTrue(all(not r['downloaded_qty'] for r in self.report()['line_choices']))
        ack=self.client.post('/api/outgoing-invoices/export-receipts/'+token+'/received')
        self.assertEqual(ack.status_code,200)
        self.assertEqual(sum(r['downloaded_qty'] for r in self.report()['line_choices']),8)
        self.assertEqual(sum(r['issued_qty'] for r in self.report()['line_choices']),0)
        saved=self.client.get('/api/outgoing-invoices/export-receipts/'+token+'/file')
        self.assertEqual(saved.data,response.data)
        self.assertEqual(before,self.business())
        with server.db() as c:c.execute("UPDATE outgoing_invoice_drafts SET status='cancelled'")
        self.assertTrue(all(not r['downloaded_qty'] for r in self.report()['line_choices']))

    def test_partial_export_keeps_missing_quantity_and_choices(self):
        with server.db() as c:
            c.execute("UPDATE inventory_transactions SET qty_in=4 WHERE source_type='OPENING'")
        before=self.business()
        plan=self.client.post('/api/outgoing-invoice-upload/stock-preview?to=2026-09-30&contractor=NT-A').json
        response=self.export(plan['token']);self.assertEqual(response.status_code,200)
        self.client.post('/api/outgoing-invoices/export-receipts/'+response.headers['X-Export-Receipt']+'/received')
        rows=self.report()['line_choices']
        self.assertEqual(sum(r['downloaded_qty'] for r in rows),4)
        self.assertEqual(sum(r['qty']-r['downloaded_qty'] for r in rows),4)
        self.assertTrue(all(r['enabled'] for r in rows))
        self.assertEqual(before,self.business())

    def test_download_repairs_legacy_package_without_rewriting_receipt_or_business(self):
        before = self.business()
        plan = self.client.post('/api/outgoing-invoice-upload/stock-preview?to=2026-09-30&contractor=NT-A').json
        response = self.export(plan['token'])
        token = response.headers['X-Export-Receipt']
        self.client.post('/api/outgoing-invoices/export-receipts/' + token + '/received')
        legacy = io.BytesIO()
        expected = {}
        with zipfile.ZipFile(io.BytesIO(response.data)) as source, zipfile.ZipFile(legacy, 'w') as target:
            for entry in source.infolist():
                payload = source.read(entry.filename)
                if entry.filename.endswith('.xlsx'):
                    workbook = load_workbook(io.BytesIO(payload))
                    expected[entry.filename] = list(workbook.active.values)
                    stream = io.BytesIO()
                    workbook.save(stream)
                    workbook.close()
                    payload = stream.getvalue()
                target.writestr(entry, payload)
        with server.db() as c:
            c.execute('UPDATE outgoing_export_receipts SET file_data=? WHERE token=?', (legacy.getvalue(), token))
        saved = self.client.get('/api/outgoing-invoices/export-receipts/' + token + '/file')
        self.assertEqual(saved.status_code, 200)
        with zipfile.ZipFile(io.BytesIO(saved.data)) as archive:
            for name, rows in expected.items():
                payload = archive.read(name)
                with zipfile.ZipFile(io.BytesIO(payload)) as xlsx:
                    self.assertIn('xl/sharedStrings.xml', xlsx.namelist())
                workbook = load_workbook(io.BytesIO(payload), data_only=True)
                self.assertEqual(rows, list(workbook.active.values))
                workbook.close()
        with server.db() as c:
            stored = c.execute('SELECT file_data FROM outgoing_export_receipts WHERE token=?', (token,)).fetchone()
            self.assertEqual(stored['file_data'], legacy.getvalue())
        self.assertEqual(before, self.business())
