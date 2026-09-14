import io
import json
import unittest
import zipfile
from unittest.mock import patch

from openpyxl import load_workbook
from . import server, outgoing_upload
from . import test_outgoing_readiness as fixtures
from .outgoing_unissued import issued_allocations
from .outgoing_weights import order_snapshot


class OutgoingUploadTests(unittest.TestCase):
    setUpClass=classmethod(fixtures.OutgoingReadinessTests.setUpClass.__func__)
    tearDownClass=classmethod(fixtures.OutgoingReadinessTests.tearDownClass.__func__)

    def setUp(self):
        with server.db() as c:
            c.execute('DELETE FROM outgoing_weight_exports')
            c.execute("DELETE FROM invoice_mapping_revisions WHERE revision_key LIKE 'actual-weight-export-%'")
            c.execute("DELETE FROM invoice_line_mappings WHERE source='tdp_actual_weight'")
            c.execute("DELETE FROM outgoing_product_units WHERE product_code='HH-01'")
        fixtures.OutgoingReadinessTests.setUp(self)
        with server.db() as c:
            c.execute('DELETE FROM outgoing_upload_jobs')
            c.execute('DELETE FROM outgoing_contractor_choices')
            c.execute('DELETE FROM outgoing_order_choices')
            c.execute('DELETE FROM outgoing_order_weights')
            c.execute('DELETE FROM outgoing_waiting_settlements')
            c.execute("DELETE FROM outgoing_product_units WHERE product_code='HH-01'")
            c.execute("DELETE FROM outgoing_product_names WHERE product_code='HH-01'")
            fixtures.OutgoingReadinessTests.add_opening(c,10)
            self.bid,self.oids=fixtures.OutgoingReadinessTests.add_batch(c,'2026-09-01',[{'qty':3,'sell_price':20000},{'qty':5,'sell_price':30000}])
            c.execute("INSERT INTO outgoing_product_names(product_code,invoice_name,updated_at) VALUES('HH-01','Tên xuất hóa đơn chuẩn','today')")

    def business(self):
        with server.db() as c:
            return {t:[tuple(r) for r in c.execute('SELECT * FROM '+t)] for t in
                ['orders','batches','products','receivable_ledger_lines','payable_ledger_lines',
                 'invoice_inventory_ledger','outgoing_source_invoices','outgoing_source_invoice_items']}

    def workbook(self, edit=None):
        r=self.client.get('/api/outgoing-invoice-upload/template.xlsx?to=2026-09-13&contractor=NT-A')
        self.assertEqual(r.status_code,200,r.get_json(silent=True))
        w=load_workbook(io.BytesIO(r.data));s=w[outgoing_upload.SHEET]
        if edit: edit(s)
        buf=io.BytesIO();w.save(buf);w.close();return buf.getvalue()

    def preview(self, data=None):
        return self.client.post('/api/outgoing-invoice-upload/preview',data={
            'file':(io.BytesIO(data or self.workbook()),'de-nghi.xlsx'),'to':'2026-09-13','contractor':'NT-A'})

    def export(self, token):
        return self.client.post('/api/outgoing-invoice-upload/export',json={'token':token,'confirmed':True})

    def file_rows(self,response):
        self.assertEqual(response.status_code,200,response.get_json(silent=True))
        result=[]
        with zipfile.ZipFile(io.BytesIO(response.data)) as z:
            for name in z.namelist():
                if name.endswith('.xlsx'):
                    w=load_workbook(io.BytesIO(z.read(name)),data_only=True);result.extend(list(w.active.values)[1:]);w.close()
        return result

    def test_template_and_preview_do_not_create_orders_or_holds(self):
        before=self.business()
        r=self.preview();self.assertEqual(r.status_code,200,r.get_json())
        p=r.get_json();self.assertEqual(p['ready_count'],2);self.assertEqual(p['waiting_count'],0)
        self.assertEqual(p['items'][0]['invoice_name'],'Tên xuất hóa đơn chuẩn')
        self.assertEqual(before,self.business())
        with server.db() as c:
            self.assertEqual(c.execute('SELECT COUNT(*) FROM outgoing_invoice_drafts').fetchone()[0],0)
            self.assertEqual(c.execute("SELECT COUNT(*) FROM inventory_transactions WHERE status='reserved'").fetchone()[0],0)

    def test_confirm_redownload_preserves_revenue_prices_and_holds(self):
        before=self.business();p=self.preview().get_json()
        r=self.file_rows(self.export(p['token']))
        self.assertEqual([(row[1],row[3],row[4],row[8]) for row in r],
                         [('Tên xuất hóa đơn chuẩn',3,20000,60000),('Tên xuất hóa đơn chuẩn',5,30000,150000)])
        with server.db() as c:
            self.assertEqual(c.execute("SELECT SUM(qty_out) FROM inventory_transactions WHERE status='reserved'").fetchone()[0],8)
            state=c.serialize()
        self.assertEqual(r,self.file_rows(self.export(p['token'])))
        with server.db() as c:self.assertEqual(state,c.serialize())
        self.assertEqual(before,self.business())

    def test_replaces_own_holds_keeps_other_order_and_contractor(self):
        with server.db() as c:fixtures.OutgoingReadinessTests.add_batch(c,'2026-09-02',[{'qty':2,'contractor':'NT-B'}])
        r=self.client.post('/api/export/order-invoices',json={'scope':'unissued','to':'2026-09-13','contractor':''})
        self.assertEqual(r.status_code,200,r.get_json(silent=True))
        before=self.business()
        data=self.workbook(lambda s:(s.delete_rows(3),setattr(s['H2'],'value',2)))
        p=self.preview(data).get_json();self.assertEqual(p['ready_count'],1)
        self.assertEqual(self.file_rows(self.export(p['token']))[0][3],2)
        with server.db() as c:
            self.assertEqual(c.execute("SELECT SUM(qty_out) FROM inventory_transactions WHERE status='reserved'").fetchone()[0],9)
            quantities={r['order_id']:r['qty'] for r in c.execute("SELECT a.order_id,SUM(a.qty) qty FROM outgoing_order_allocations a JOIN outgoing_invoice_drafts d ON d.id=a.draft_id WHERE d.status='draft' GROUP BY a.order_id")}
            self.assertEqual(quantities[self.oids[1]],5)
        self.assertEqual(before,self.business())

    def test_insufficient_stock_is_partial_and_remainder_stays_pending(self):
        with server.db() as c:c.execute("UPDATE inventory_transactions SET qty_in=4 WHERE source_type='OPENING'")
        p=self.preview().get_json();self.assertEqual(p['ready_count'],2);self.assertEqual(p['waiting_count'],1)
        self.assertEqual([r['ready_qty'] for r in p['items']],[3,1])
        self.assertEqual(sum(r[3] for r in self.file_rows(self.export(p['token']))),4)
        d=self.client.get('/api/outgoing-invoices/unissued?to=2026-09-13&contractor=NT-A').get_json()
        self.assertAlmostEqual(sum(r['unissued_qty'] for r in d['details']),8)

    def test_zero_stock_does_not_create_draft(self):
        with server.db() as c:c.execute("UPDATE inventory_transactions SET qty_in=0 WHERE source_type='OPENING'")
        p=self.preview().get_json();self.assertEqual(p['ready_count'],0)
        self.assertEqual(self.export(p['token']).status_code,409)
        with server.db() as c:self.assertEqual(c.execute('SELECT COUNT(*) FROM outgoing_invoice_drafts').fetchone()[0],0)

    def test_tampered_identity_quantity_duplicate_and_formula_reject_without_holds(self):
        for edit in [lambda s:setattr(s['E2'],'value','OTHER'),lambda s:setattr(s['F2'],'value','Hàng khác'),
                     lambda s:setattr(s['G2'],'value','Gói'),lambda s:setattr(s['H2'],'value',99),
                     lambda s:setattr(s['H2'],'value','=-1'),lambda s:s.append([c.value for c in s[2]])]:
            response=self.preview(self.workbook(edit));self.assertEqual(response.status_code,409,response.get_json())
        with server.db() as c:self.assertEqual(c.execute('SELECT COUNT(*) FROM outgoing_invoice_drafts').fetchone()[0],0)

    def test_expired_preview_and_changed_stock_or_selection_reject(self):
        p=self.preview().get_json()
        with server.db() as c:c.execute('UPDATE outgoing_upload_jobs SET expires_at=0')
        self.assertEqual(self.export(p['token']).status_code,409)
        p=self.preview().get_json()
        with server.db() as c:c.execute("UPDATE inventory_transactions SET qty_in=6 WHERE source_type='OPENING'")
        self.assertEqual(self.export(p['token']).status_code,409)
        p=self.preview().get_json()
        with server.db() as c:c.execute("INSERT INTO outgoing_order_choices VALUES(?,0,'test','today')",(self.oids[0],))
        self.assertEqual(self.export(p['token']).status_code,409)
        with server.db() as c:self.assertEqual(c.execute('SELECT COUNT(*) FROM outgoing_invoice_drafts').fetchone()[0],0)

    def test_signed_after_cutoff_is_deducted_and_existing_file_rejected_after_issue(self):
        with server.db() as c:
            fixtures.OutgoingReadinessTests.add_local_issued_draft(c,self.bid,self.oids[0],number='TEST',invoice_date='2026-09-14',qty=1)
        p=self.preview().get_json()
        self.assertEqual(p['items'][0]['requested_qty'],2)
        self.assertEqual(sum(r[3] for r in self.file_rows(self.export(p['token']))),7)
        with server.db() as c:
            c.execute("UPDATE outgoing_invoice_drafts SET status='issued',issued_invoice_number='TEST2',issued_invoice_series='1C26TDP' WHERE draft_kind='invoice_upload'")
        self.assertEqual(self.export(p['token']).status_code,409)

    def test_weighted_invoice_preserves_amount_and_stock_packages(self):
        with server.db() as c:
            c.execute('DELETE FROM orders WHERE id=?',(self.oids[1],))
            c.execute("UPDATE products SET unit='Gói' WHERE code='HH-01'")
            c.execute("UPDATE orders SET unit='Gói',actual_delivered=14,actual_received=14,qty=14,sell_price=8000")
            c.execute("UPDATE inventory_transactions SET qty_in=20 WHERE source_type='OPENING'")
            c.execute("INSERT INTO outgoing_product_units(product_code,invoice_unit,updated_at) VALUES('HH-01','Kg','today')")
        p=self.preview().get_json();self.assertEqual(p['ready_count'],0)
        with server.db() as c:
            o=c.execute('SELECT * FROM orders WHERE id=?',(self.oids[0],)).fetchone()
            c.execute('INSERT INTO outgoing_order_weights VALUES(?,?,?,?,?,?,?)',(o['id'],order_snapshot(o),14,2.8,'TEST','Cân thực tế','today'))
        p=self.preview().get_json();self.assertEqual(p['items'][0]['invoice_qty'],2.8)
        row=self.file_rows(self.export(p['token']))[0]
        self.assertEqual((row[2],row[3],row[4],row[8]),('Kg',2.8,40000,112000))
        with server.db() as c:self.assertEqual(c.execute("SELECT SUM(qty_out) FROM inventory_transactions WHERE status='reserved'").fetchone()[0],14)

    def test_failed_workbook_rolls_back_all_new_holds(self):
        p=self.preview().get_json();before=self.business()
        with patch('tdp_system.outgoing_upload.build_invoice_workbook',side_effect=ValueError('template broken')):
            self.assertEqual(self.export(p['token']).status_code,409)
        with server.db() as c:
            self.assertEqual(c.execute('SELECT COUNT(*) FROM outgoing_invoice_drafts').fetchone()[0],0)
            self.assertEqual(c.execute('SELECT result_json FROM outgoing_upload_jobs').fetchone()[0],'')
        self.assertEqual(before,self.business())

    def test_saved_remote_draft_remains_unchanged(self):
        self.client.post('/api/export/order-invoices',json={'scope':'unissued','to':'2026-09-13','contractor':'NT-A'})
        with server.db() as c:
            c.execute("UPDATE outgoing_invoice_drafts SET minvoice_status='saved'")
            before=[tuple(r) for r in c.execute('SELECT * FROM outgoing_invoice_drafts')]
        p=self.preview().get_json();self.assertEqual(p['ready_count'],0)
        self.assertEqual(self.export(p['token']).status_code,409)
        with server.db() as c:self.assertEqual(before,[tuple(r) for r in c.execute('SELECT * FROM outgoing_invoice_drafts')])

    def test_preview_amount_matches_merged_line_rounding(self):
        with server.db() as c:
            c.execute('UPDATE orders SET actual_delivered=.25,actual_received=.25,qty=.25,sell_price=3')
        p=self.preview().get_json()
        r=self.file_rows(self.export(p['token']))
        self.assertEqual(p['amount'],sum(row[8] for row in r))

    def test_invoice_signed_after_preview_rejects_old_request(self):
        p=self.preview().get_json()
        with server.db() as c:
            fixtures.OutgoingReadinessTests.add_local_issued_draft(c,self.bid,self.oids[0],number='AFTER',invoice_date='2026-09-14',qty=1)
        response=self.export(p['token']);self.assertEqual(response.status_code,409)
        self.assertIn('sau đối trừ',response.get_json()['error'])
        with server.db() as c:self.assertEqual(c.execute("SELECT COUNT(*) FROM outgoing_invoice_drafts WHERE status='draft'").fetchone()[0],0)

    def test_reupload_same_selection_does_not_double_holds(self):
        data=self.workbook();p=self.preview(data).get_json();self.file_rows(self.export(p['token']))
        p2=self.preview(data).get_json();self.file_rows(self.export(p2['token']))
        with server.db() as c:
            self.assertEqual(c.execute("SELECT SUM(qty_out) FROM inventory_transactions WHERE status='reserved'").fetchone()[0],8)
            self.assertEqual(c.execute("SELECT COUNT(*) FROM outgoing_invoice_drafts WHERE status='draft'").fetchone()[0],1)

    def test_remote_sync_failure_leaves_financial_and_draft_data_unchanged(self):
        p=self.preview().get_json()
        with server.db() as c:server.setting_set(c,'minvoice_active_connection','test-connection')
        try:
            with patch('tdp_system.outgoing_source_refresh.refresh_sources',side_effect=ValueError('offline')):
                response=self.export(p['token'])
            self.assertEqual(response.status_code,409)
            self.assertIn('Chưa cập nhật đủ',response.get_json()['error'])
            with server.db() as c:self.assertEqual(c.execute('SELECT COUNT(*) FROM outgoing_invoice_drafts').fetchone()[0],0)
        finally:
            with server.db() as c:server.setting_set(c,'minvoice_active_connection','')

    def test_external_signed_invoice_is_deducted_once_without_local_issue(self):
        with server.db() as c:
            c.execute("INSERT OR REPLACE INTO outgoing_buyer_profiles(contractor,tax_code,address,updated_at) VALUES('NT-A','0101234567','','today')")
            sid=fixtures.OutgoingReadinessTests.add_posted_source(c,source='minvoice',qty=2,invoice_date='2026-09-14',number='EXTERNAL')
            c.execute("UPDATE outgoing_source_invoices SET buyer_tax_code='0101234567' WHERE id=?",(sid,))
        before=self.business()
        p=self.preview().get_json()
        self.assertEqual([r['requested_qty'] for r in p['items']],[1,5])
        self.assertEqual(sum(r[3] for r in self.file_rows(self.export(p['token']))),6)
        self.assertEqual(sum(r[3] for r in self.file_rows(self.export(p['token']))),6)
        with server.db() as c:
            self.assertEqual(c.execute("SELECT SUM(qty_out) FROM inventory_transactions WHERE status='reserved'").fetchone()[0],6)
            self.assertEqual(c.execute('SELECT external_issued_qty FROM outgoing_waiting_settlements WHERE order_id=?',(self.oids[0],)).fetchone()[0],2)
            self.assertEqual(issued_allocations(c)[0][self.oids[0]],2)
        self.assertEqual(before,self.business())

    def test_malformed_workbook_is_reported_without_server_error(self):
        original=self.workbook();broken=io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(original)) as source, zipfile.ZipFile(broken,'w') as target:
            for name in source.namelist():
                target.writestr(name,b'<worksheet><broken>' if name=='xl/worksheets/sheet1.xml' else source.read(name))
        for data in (b'not an Excel workbook',broken.getvalue()):
            response=self.preview(data)
            self.assertEqual(response.status_code,409,response.get_json())
        with server.db() as c:
            self.assertEqual(c.execute('SELECT COUNT(*) FROM outgoing_upload_jobs').fetchone()[0],0)
            self.assertEqual(c.execute('SELECT COUNT(*) FROM outgoing_invoice_drafts').fetchone()[0],0)

    def test_successful_source_refresh_keeps_preview_valid_and_retry_stable(self):
        self.client.post('/api/export/order-invoices',json={'scope':'unissued','to':'2026-09-13','contractor':'NT-A'})
        p=self.preview().get_json()
        from .outgoing_waiting import refresh_waiting
        def synced(*args):
            with server.db() as c:
                return refresh_waiting(c,server.now_iso(),fill=False)
        with server.db() as c:server.setting_set(c,'minvoice_active_connection','test-connection')
        try:
            with patch('tdp_system.outgoing_source_refresh.refresh_sources',side_effect=synced):
                first=self.file_rows(self.export(p['token']))
                self.assertEqual(first,self.file_rows(self.export(p['token'])))
            with server.db() as c:
                self.assertEqual(c.execute("SELECT SUM(qty_out) FROM inventory_transactions WHERE status='reserved'").fetchone()[0],8)
        finally:
            with server.db() as c:server.setting_set(c,'minvoice_active_connection','')

    def test_upload_all_preserves_existing_buyers_holds_before_new_demand(self):
        with server.db() as c:
            c.execute("UPDATE inventory_transactions SET qty_in=8 WHERE source_type='OPENING'")
            c.execute("UPDATE orders SET contractor='NT-B'")
        response=self.client.post('/api/export/order-invoices',json={'scope':'unissued','to':'2026-09-13','contractor':'NT-B'})
        self.assertEqual(response.status_code,200,response.get_json(silent=True))
        with server.db() as c:
            fixtures.OutgoingReadinessTests.add_batch(c,'2026-09-02',[{'qty':10,'contractor':'NT-A'}])
        template=self.client.get('/api/outgoing-invoice-upload/template.xlsx?to=2026-09-13').data
        response=self.client.post('/api/outgoing-invoice-upload/preview',data={'file':(io.BytesIO(template),'all.xlsx'),'to':'2026-09-13'})
        p=response.get_json();self.assertEqual(response.status_code,200,p)
        self.assertEqual(sum(r['ready_qty'] for r in p['items'] if r['contractor']=='NT-B'),8)
        self.assertEqual(sum(r['ready_qty'] for r in p['items'] if r['contractor']=='NT-A'),0)
        self.assertEqual(sum(row[3] for row in self.file_rows(self.export(p['token']))),8)
