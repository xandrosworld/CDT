import io
import unittest
import zipfile
from openpyxl import Workbook, load_workbook
from .order_workbook_compat import compatible_order_bytes


class OrderWorkbookCompatTests(unittest.TestCase):
    def fixture(self):
        book = Workbook()
        sheet = book.active
        sheet['A1'] = 'Ngày mua'
        sheet['A2'] = '01/10/2026'
        sheet['B2'] = '=2*3'
        sheet.auto_filter.ref = 'A1:B2'
        sheet.auto_filter.add_sort_condition('B2:B2')
        stream = io.BytesIO(); book.save(stream); book.close()
        output = io.BytesIO()
        with zipfile.ZipFile(stream) as src, zipfile.ZipFile(output, 'w') as dst:
            for entry in src.infolist():
                data = src.read(entry.filename)
                if entry.filename == 'xl/worksheets/sheet1.xml':
                    data = data.replace(b'<sortState ref="A1:B2"', b'<sortState ref="1:2"')
                    data = data.replace(b'<f>2*3</f><v></v>', b'<f>2*3</f><v>6</v>')
                dst.writestr(entry, data)
        return output.getvalue()

    def test_recovers_sort_range_without_changing_cells_or_formula_cache(self):
        payload = self.fixture()
        with self.assertRaises(ValueError):
            load_workbook(io.BytesIO(payload))
        fixed = compatible_order_bytes(payload)
        with zipfile.ZipFile(io.BytesIO(payload)) as before, zipfile.ZipFile(io.BytesIO(fixed)) as after:
            self.assertEqual(before.namelist(), after.namelist())
            for name in before.namelist():
                expected = before.read(name)
                if name == 'xl/worksheets/sheet1.xml':
                    expected = expected.replace(b'<sortState ref="1:2"', b'<sortState ref="A1:XFD2"')
                self.assertEqual(expected, after.read(name))
        values = load_workbook(io.BytesIO(fixed), data_only=True)
        formulas = load_workbook(io.BytesIO(fixed), data_only=False)
        self.assertEqual(values.active['B2'].value, 6)
        self.assertEqual(formulas.active['B2'].value, '=2*3')
        self.assertEqual(values.active['A2'].value, '01/10/2026')
        self.assertEqual(values.active.auto_filter.ref, 'A1:B2')
        values.close(); formulas.close()
        self.assertEqual(compatible_order_bytes(fixed), fixed)
