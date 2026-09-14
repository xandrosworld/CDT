import io
import sqlite3
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

try:
    from .pdf_documents import build_pdf_bundle
    from .print_bundle import workbook_sections
    from .receipt_export import (
        ReceiptExportError,
        build_purchase_documents_workbook,
        enrich_receipt_identity_rows,
        group_receipt_rows,
        configure_receipt_paper,
    )
    from .template_workbook import safe_workbook_bytes
except ImportError:  # pragma: no cover - direct file invocation
    from pdf_documents import build_pdf_bundle
    from print_bundle import workbook_sections
    from receipt_export import (
        ReceiptExportError,
        build_purchase_documents_workbook,
        enrich_receipt_identity_rows,
        group_receipt_rows,
        configure_receipt_paper,
    )
    from template_workbook import safe_workbook_bytes


ROOT = Path(__file__).resolve().parent.parent
GOLDEN = ROOT / "Em Thành.xlsx"
FAKE_ID = "0" * 12


def receipt_row(**overrides):
    row = {
        "work_date": "2026-09-03",
        "seller": "Người bán kiểm thử",
        "address": "Địa chỉ kiểm thử",
        "cccd": FAKE_ID,
        "issue_date": "02/01/2020",
        "issue_place": "Nơi cấp kiểm thử",
        "product_name": "Cà rốt",
        "unit": "Kg",
        "quantity": 2,
        "buy_price": 100,
        "amount": 200,
        "supplier": "NCC A",
        "kitchen": "BEP-A",
        "source_ref": 3,
    }
    row.update(overrides)
    return row


class ReceiptExportTests(unittest.TestCase):
    def test_changing_paper_preserves_all_receipt_values_and_restores_a5_layout(self):
        book = build_purchase_documents_workbook([receipt_row()], template_path=GOLDEN)
        self.addCleanup(book.close)
        sheet = book['biên nhận']
        values = lambda: {c.coordinate: c.value for row in sheet for c in row if c.value is not None}
        original = values()
        original_width = sheet.column_dimensions['C'].width
        configure_receipt_paper(sheet, 'A4')
        self.assertEqual(values(), original)
        self.assertGreater(sheet.column_dimensions['C'].width, original_width)
        self.assertEqual(str(sheet.page_setup.paperSize), '9')
        for _ in range(2):
            configure_receipt_paper(sheet, 'A5')
            self.assertEqual(values(), original)
            self.assertEqual(str(sheet.page_setup.paperSize), '11')
            self.assertEqual(sheet.column_dimensions['C'].width, original_width)
            self.assertEqual(sheet.page_setup.fitToHeight, 1)
            self.assertEqual(sheet['C15'].font.sz, 11)

    def test_one_dynamic_golden_receipt_per_legal_seller_and_day(self):
        rows = [
            receipt_row(
                product_name=f"Mặt hàng kiểm thử {index:02d}",
                quantity=index,
                buy_price=1_000,
                amount=index * 1_000,
                source_ref=index,
            )
            for index in range(1, 10)
        ]
        rows.append(
            receipt_row(
                seller="Người bán kiểm thử B",
                address="Địa chỉ kiểm thử B",
                cccd="1" * 12,
                issue_date="03/02/2021",
                issue_place="Nơi cấp kiểm thử B",
                product_name="Khoai tây",
                quantity=1,
                amount=1_000,
                source_ref=20,
            )
        )
        workbook = build_purchase_documents_workbook(
            rows,
            template_path=GOLDEN,
            buyer_name="Người mua kiểm thử",
            buyer_title="Nhân viên thu mua",
            company_name="CÔNG TY KIỂM THỬ",
            company_address="Địa chỉ công ty kiểm thử",
        )
        try:
            self.assertEqual(
                workbook.sheetnames,
                ["bảng kê tổng", "biên nhận", "biên nhận 02"],
            )
            first = workbook["biên nhận"]
            self.assertEqual(first["C3"].value, "GIẤY BIÊN NHẬN")
            self.assertIn("03.09.2026", first["C4"].value)
            self.assertIn('CÔNG TY KIỂM THỬ',first['C4'].value)
            self.assertNotIn('Phát triển',first['C4'].value)
            self.assertEqual(first["C5"].value, "I - Tên đơn vị mua: CÔNG TY KIỂM THỬ")
            self.assertIn("Người mua kiểm thử", first["C6"].value)
            self.assertEqual(first["C7"].value, "Địa chỉ: Địa chỉ công ty kiểm thử")
            self.assertEqual(first["D8"].value, "Người bán kiểm thử")
            self.assertEqual(first["D10"].value, FAKE_ID)
            self.assertEqual(first["D11"].value, "02/01/2020")
            self.assertEqual(first["C23"].value, "Mặt hàng kiểm thử 09")
            self.assertEqual(first["C24"].value, "TỔNG")
            self.assertEqual(first["F24"].value, 45)
            self.assertEqual(first["G24"].value, 45_000)
            self.assertIn("Bốn mươi lăm nghìn đồng", first["C25"].value)
            self.assertEqual(str(first.print_area), "'biên nhận'!$C$1:$G$34")
            self.assertEqual(str(first.page_setup.paperSize), "11")
            self.assertEqual(first.page_setup.orientation, "portrait")
            self.assertEqual(first.page_setup.fitToWidth, 1)
            self.assertEqual(first.page_setup.fitToHeight, 0)
            self.assertTrue(first.print_options.horizontalCentered)
            self.assertEqual(first.max_row, 34)
            self.assertIn("C24:E24", {str(value) for value in first.merged_cells.ranges})
            self.assertIn("D8:G8", {str(value) for value in first.merged_cells.ranges})
            self.assertIn("C30:D30", {str(value) for value in first.merged_cells.ranges})
            self.assertIn("E29:G29", {str(value) for value in first.merged_cells.ranges})
            self.assertIn("E34:G34", {str(value) for value in first.merged_cells.ranges})
            self.assertEqual(first["E29"].value, "Hải Phòng, ngày 03 tháng 09 năm 2026")
            self.assertEqual(first["E29"].alignment.horizontal, "center")
            self.assertEqual(first["E30"].alignment.horizontal, "center")
            self.assertEqual(first["E34"].alignment.horizontal, "center")
            self.assertEqual(first["C15"].font.name, "Times New Roman")
            self.assertEqual(first["C15"].font.sz, 11)
            self.assertGreaterEqual(first.row_dimensions[15].height, 22)
            self.assertFalse(any(
                isinstance(cell.value, str) and cell.value.startswith("=")
                for sheet in workbook.worksheets
                for row in sheet.iter_rows()
                for cell in row
            ))
            payload = safe_workbook_bytes(workbook)
        finally:
            workbook.close()

        reopened = load_workbook(io.BytesIO(payload), data_only=False, keep_links=False)
        try:
            self.assertEqual(reopened.sheetnames, ["bảng kê tổng", "biên nhận", "biên nhận 02"])
            self.assertEqual(reopened["biên nhận 02"]["D10"].value, "1" * 12)
            self.assertLessEqual(reopened["biên nhận 02"]["G16"].value or 0, 5_000_000)
            self.assertFalse(reopened._external_links)
        finally:
            reopened.close()

    def test_same_item_merges_and_daily_limit_fails_without_identity_in_error(self):
        grouped = group_receipt_rows([
            receipt_row(quantity=2, amount=200, supplier="NCC A", kitchen="BEP-A"),
            receipt_row(quantity=3, amount=360, buy_price=120, supplier="NCC B", kitchen="BEP-B"),
        ])
        self.assertEqual(len(grouped), 1)
        self.assertEqual(len(grouped[0]["items"]), 1)
        self.assertEqual(grouped[0]["items"][0]["quantity"], 5)
        self.assertEqual(grouped[0]["items"][0]["unit_price"], 112)
        self.assertEqual(grouped[0]["total_amount"], 560)

        private_identity = "2" * 12
        with self.assertRaises(ReceiptExportError) as caught:
            group_receipt_rows([
                receipt_row(
                    cccd=private_identity,
                    quantity=1,
                    buy_price=5_000_001,
                    amount=5_000_001,
                )
            ])
        self.assertEqual(caught.exception.code, "receipt_daily_limit_exceeded")
        self.assertNotIn(private_identity, str(caught.exception))

        accepted = group_receipt_rows([
            receipt_row(quantity=1, buy_price=5_000_000, amount=5_000_000)
        ])
        self.assertEqual(accepted[0]["total_amount"], 5_000_000)

    def test_identity_enrichment_is_internal_and_never_reads_receivables(self):
        connection = sqlite3.connect(":memory:")
        connection.row_factory = sqlite3.Row
        connection.executescript(
            """
            CREATE TABLE people(
                name TEXT, cccd TEXT, issue_date TEXT, issue_place TEXT, address TEXT
            );
            CREATE TABLE receivable_ledger_lines(id INTEGER, amount REAL);
            """
        )
        connection.execute(
            "INSERT INTO people VALUES(?,?,?,?,?)",
            ("Người bán kiểm thử", FAKE_ID, "02/01/2020", "Nơi cấp kiểm thử", "Địa chỉ kiểm thử"),
        )
        connection.execute("INSERT INTO receivable_ledger_lines VALUES(1,999999999)")
        traced = []
        connection.set_trace_callback(traced.append)
        try:
            enriched = enrich_receipt_identity_rows(connection, [receipt_row()])
            self.assertEqual(enriched[0]["issue_date"], "02/01/2020")
            self.assertEqual(enriched[0]["issue_place"], "Nơi cấp kiểm thử")
            self.assertFalse(any("receivable" in statement.casefold() for statement in traced))

            private_identity = "3" * 12
            with self.assertRaises(ReceiptExportError) as caught:
                enrich_receipt_identity_rows(
                    connection, [receipt_row(cccd=private_identity)]
                )
            self.assertEqual(caught.exception.code, "invalid_receipt_identity")
            self.assertNotIn(private_identity, str(caught.exception))
            self.assertNotIn(FAKE_ID, str(caught.exception))
        finally:
            connection.set_trace_callback(None)
            connection.close()

    def test_receipt_projects_to_pdf_but_manifest_keeps_identity_hashed(self):
        workbook = build_purchase_documents_workbook(
            [receipt_row()], template_path=GOLDEN,
        )
        try:
            sections = workbook_sections("purchases", workbook)
            self.assertEqual(len(sections), 2)
            receipt = sections[1]
            self.assertEqual(receipt["document_type"], "purchase_receipt")
            self.assertEqual(receipt["title"], "GIẤY BIÊN NHẬN")
            self.assertEqual(len(receipt["rows"]), 1)
            self.assertIn(FAKE_ID, repr(receipt["notes"]))
            with tempfile.TemporaryDirectory(prefix="tdp_receipt_pdf_") as temp_name:
                output = Path(temp_name) / "receipts.pdf"
                manifest = build_pdf_bundle(
                    sections,
                    output,
                    company="CÔNG TY KIỂM THỬ",
                    document_title="Bộ bảng kê và biên nhận kiểm thử",
                    generated_at="2026-09-03T00:00:00",
                    paper="A4",
                )
                self.assertTrue(manifest["ok"])
                self.assertEqual(manifest["section_count"], 2)
                self.assertEqual(manifest["sections"][1]["title"], "GIẤY BIÊN NHẬN")
                self.assertNotIn(FAKE_ID, repr(manifest))
                self.assertTrue(output.is_file())
        finally:
            workbook.close()


if __name__ == "__main__":
    unittest.main()
