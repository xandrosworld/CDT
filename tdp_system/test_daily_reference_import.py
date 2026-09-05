from __future__ import annotations

import io
import json
import tempfile
import unittest
from datetime import date
from pathlib import Path

from openpyxl import Workbook

try:
    from . import daily_reference_import, server
except ImportError:  # pragma: no cover - direct file invocation
    import daily_reference_import
    import server


class DailyReferenceImportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        cls.original_db_path = server.DB_PATH
        cls.original_data_dir = server.DATA_DIR
        cls.original_master_source = server.MASTER_SOURCE
        cls.original_testing = server.app.config.get("TESTING", False)
        server.DATA_DIR = Path(cls.temp.name) / "data"
        server.DATA_DIR.mkdir(parents=True, exist_ok=True)
        server.DB_PATH = server.DATA_DIR / "references.sqlite3"
        server.MASTER_SOURCE = Path(cls.temp.name) / "master-disabled.xlsx"
        server.app.config["TESTING"] = True
        server.init_database()
        cls.client = server.app.test_client()
        with server.db() as conn:
            conn.execute(
                "INSERT INTO contractors(code,name,price_group,pricing_mode) VALUES('C1','C1','C1','group')"
            )
            conn.execute(
                "INSERT INTO kitchens(code,contractor,name,address) VALUES('K1','C1','Kitchen 1','Old address')"
            )
            conn.execute("INSERT INTO suppliers(code,name) VALUES('S1','Supplier 1')")
            conn.execute(
                """INSERT INTO products(
                       code,name,unit,tax,supplier,buy_price,purchase_list,seller,cccd
                   ) VALUES('P1','Product 1','kg','8%','S1',10000,1,'Seller protected','ID protected')"""
            )
            conn.execute(
                "INSERT INTO product_prices(product_code,price_group,price_text,price_value) VALUES('P1','C1','12000',12000)"
            )

    @classmethod
    def tearDownClass(cls):
        with daily_reference_import.DAILY_REFERENCE_IMPORT_LOCK:
            daily_reference_import.PENDING_DAILY_REFERENCE_IMPORTS.clear()
        with server.PENDING_IMPORT_LOCK:
            for pending in server.PENDING_IMPORTS.values():
                pending["path"].unlink(missing_ok=True)
            server.PENDING_IMPORTS.clear()
        server.DB_PATH = cls.original_db_path
        server.DATA_DIR = cls.original_data_dir
        server.MASTER_SOURCE = cls.original_master_source
        server.app.config["TESTING"] = cls.original_testing
        cls.temp.cleanup()

    def setUp(self):
        with daily_reference_import.DAILY_REFERENCE_IMPORT_LOCK:
            daily_reference_import.PENDING_DAILY_REFERENCE_IMPORTS.clear()
        with server.PENDING_IMPORT_LOCK:
            for pending in server.PENDING_IMPORTS.values():
                pending["path"].unlink(missing_ok=True)
            server.PENDING_IMPORTS.clear()
        with server.db() as conn:
            conn.execute("DROP TRIGGER IF EXISTS fail_reference_kitchen")
            conn.execute("DELETE FROM daily_reference_import_receipts")
            conn.execute("DELETE FROM audit_log WHERE event_type='daily_reference.confirm'")
            conn.execute("DELETE FROM order_import_receipts")
            conn.execute("DELETE FROM batches")
            conn.execute("DELETE FROM kitchens WHERE code!='K1'")
            conn.execute("DELETE FROM contractors WHERE code!='C1'")
            conn.execute("DELETE FROM suppliers WHERE code!='S1'")
            conn.execute(
                """UPDATE products SET name='Product 1',unit='kg',tax='8%',supplier='S1',
                   buy_price=10000,purchase_list=1,seller='Seller protected',cccd='ID protected',
                   product_group='' WHERE code='P1'"""
            )
            conn.execute("DELETE FROM outgoing_product_names WHERE product_code='P1'")

    @staticmethod
    def workbook_bytes(secret="CCCD-MUST-STAY-PRIVATE") -> bytes:
        workbook = Workbook()
        identity = workbook.active
        identity.title = "CCCD"
        identity.append(["CCCD", "Người bán"])
        identity.append([secret, "Private Person"])

        crosscheck = workbook.create_sheet("T.chiếu")
        crosscheck.append(["THAM CHIẾU"])
        crosscheck.append([
            "STT", "Nhà thầu", "tên bếp or mã bếp", "Tên bếp ghi trên đơn hàn",
            "Địa chỉ giao hàng",
        ])
        crosscheck.append([1, "C2", "K2", "Kitchen 2", "Business address 2"])

        catalog = workbook.create_sheet("danh mục hàng hóa")
        catalog.append([
            "STT", "MÃ HÀNG", "Nhóm hàng", "TÊN THÀNH ĐẠT PHÁT",
            "TÊN XUẤT HÓA ĐƠN", "ĐVT", "Thuế",
        ])
        catalog.append([1, "P1", "G", "Product 1", "Invoice Product 1", "kg", "8%"])

        price = workbook.create_sheet("BÁO GIÁ")
        price.append(["BẢNG GIÁ"])
        price.append(["STT", "MÃ HÀNG", "TÊN THÀNH ĐẠT PHÁT", "Giá mua", "NCC", "ĐVT", "THUẾ", "C1"])
        price.append([1, "P1", "Product 1", 10_000, "S1", "kg", "8%", 12_000])

        suppliers = workbook.create_sheet("danh mục nhà cc")
        suppliers.append(["NHÀ CUNG CẤP"])
        suppliers.append(["STT", "MÃ HÀNG", "TÊN THÀNH ĐẠT PHÁT", "NCC", "ĐVT"])
        suppliers.append([1, "P1", "Product 1", "S2", "kg"])

        merged = workbook.create_sheet("gộp đơn")
        merged.append(["Mã hàng", "Mã bếp", "Tên hàng", "Số lượng", "NCC"])
        merged.append(["P-FAKE", "K-FAKE", secret, 999, "S-FAKE"])

        day_sheet = workbook.create_sheet("01.09")
        day_sheet.append(["ĐƠN HÀNG NGÀY"])
        day_sheet.append([
            "Ngày", "Nhà thầu", "Mã bếp", "Mã hàng", "Tên hàng", "Số lượng",
            "ĐVT", "Chọn NCC", "Giá mua", "Giá bán", "CCCD",
        ])
        day_sheet.append([
            date(2026, 9, 1), "C1", "K1", "P1", "Product 1", 2,
            "kg", None, 10_000, 12_000, None,
        ])
        stream = io.BytesIO()
        workbook.save(stream)
        workbook.close()
        return stream.getvalue()

    def preview_reference(self, role, payload=None, filename="daily.xlsx"):
        return self.client.post(
            "/api/daily-references/preview",
            data={
                "role": role,
                "file": (io.BytesIO(payload or self.workbook_bytes()), filename),
            },
            content_type="multipart/form-data",
        )

    def confirm_reference(self, preview, state_hash=None):
        return self.client.post(
            "/api/daily-references/confirm",
            json={
                "token": preview["token"],
                "confirmed": True,
                "state_hash": preview["stateHash"] if state_hash is None else state_hash,
            },
        )

    @staticmethod
    def protected_snapshot(conn):
        return {
            "contractors": [tuple(row) for row in conn.execute("SELECT * FROM contractors ORDER BY code")],
            "kitchens": [tuple(row) for row in conn.execute("SELECT * FROM kitchens ORDER BY code")],
            "suppliers": [tuple(row) for row in conn.execute("SELECT * FROM suppliers ORDER BY code")],
            "products": [tuple(row) for row in conn.execute("SELECT * FROM products ORDER BY code")],
            "prices": [tuple(row) for row in conn.execute("SELECT * FROM product_prices ORDER BY product_code,price_group")],
            "people": [tuple(row) for row in conn.execute("SELECT * FROM people ORDER BY name")],
        }

    def test_order_confirmation_does_not_mutate_any_reference_master(self):
        with server.db() as conn:
            before = self.protected_snapshot(conn)
        workbook = self.workbook_bytes()
        analyzed_response = self.client.post(
            "/api/import/analyze",
            data={"file": (io.BytesIO(workbook), "daily.xlsx")},
            content_type="multipart/form-data",
        )
        analyzed = analyzed_response.get_json()
        confirmed = self.client.post(
            "/api/import/confirm",
            json={
                "token": analyzed["token"], "sheets": ["01.09"],
                "work_date": analyzed["detectedWorkDate"], "state_hash": analyzed["stateHash"],
            },
        )
        self.assertEqual(confirmed.status_code, 200, confirmed.get_data(as_text=True))
        with server.db() as conn:
            after = self.protected_snapshot(conn)
        self.assertEqual(after, before)

    def test_policies_lock_identity_price_and_intermediate_roles(self):
        secret = "CCCD-NOT-IN-REFERENCE-API"
        workbook = self.workbook_bytes(secret)
        for role in ("identity_reference", "price_reference", "intermediate_reference"):
            response = self.preview_reference(role, workbook)
            self.assertEqual(response.status_code, 400, response.get_data(as_text=True))
            self.assertEqual(response.get_json()["code"], "reference_role_locked")
            self.assertNotIn(secret, response.get_data(as_text=True))

        analyzed_response = self.client.post(
            "/api/import/analyze",
            data={"file": (io.BytesIO(workbook), "daily.xlsx")},
            content_type="multipart/form-data",
        )
        policies = {item["role"]: item for item in analyzed_response.get_json()["ignoredSheets"]}
        self.assertEqual(policies["identity_reference"]["writePolicy"], "locked")
        self.assertEqual(policies["price_reference"]["writePolicy"], "locked_until_tdp050")
        self.assertEqual(policies["intermediate_reference"]["writePolicy"], "locked")
        self.assertEqual(policies["product_reference"]["previewEndpoint"], "/api/daily-references/preview")

    def test_product_reference_updates_only_confirmed_catalog_fields(self):
        workbook = Workbook()
        catalog = workbook.active
        catalog.title = "danh mục hàng hóa"
        catalog.append([
            "STT", "MÃ HÀNG", "Nhóm hàng", "TÊN THÀNH ĐẠT PHÁT",
            "TÊN XUẤT HÓA ĐƠN", "ĐVT", "Thuế",
        ])
        catalog.append([1, "P1", "NEW", "Product 1 revised", "Invoice revised", "box", "10%"])
        stream = io.BytesIO()
        workbook.save(stream)
        workbook.close()
        with server.db() as conn:
            protected_before = tuple(conn.execute(
                "SELECT supplier,buy_price,purchase_list,seller,cccd FROM products WHERE code='P1'"
            ).fetchone())
        preview_response = self.preview_reference("product_reference", stream.getvalue())
        self.assertEqual(preview_response.status_code, 200, preview_response.get_data(as_text=True))
        preview = preview_response.get_json()
        self.assertEqual(preview["counts"]["updated_products"], 1)
        confirmed = self.confirm_reference(preview)
        self.assertEqual(confirmed.status_code, 200, confirmed.get_data(as_text=True))
        with server.db() as conn:
            product = conn.execute("SELECT * FROM products WHERE code='P1'").fetchone()
            protected_after = tuple(conn.execute(
                "SELECT supplier,buy_price,purchase_list,seller,cccd FROM products WHERE code='P1'"
            ).fetchone())
            invoice_name = conn.execute(
                "SELECT invoice_name FROM outgoing_product_names WHERE product_code='P1'"
            ).fetchone()["invoice_name"]
            price_count = conn.execute(
                "SELECT COUNT(*) FROM product_prices WHERE product_code='P1'"
            ).fetchone()[0]
        self.assertEqual(
            (product["name"], product["unit"], product["tax"], product["product_group"]),
            ("Product 1 revised", "box", "0.1", "NEW"),
        )
        self.assertEqual(invoice_name, "Invoice revised")
        self.assertEqual(protected_after, protected_before)
        self.assertEqual(price_count, 1)

    def test_contractor_kitchen_confirm_is_stale_safe_atomic_and_replay_safe(self):
        workbook = self.workbook_bytes()
        preview_response = self.preview_reference("contractor_kitchen_reference", workbook)
        self.assertEqual(preview_response.status_code, 200, preview_response.get_data(as_text=True))
        preview = preview_response.get_json()
        self.assertEqual(preview["counts"]["new_contractors"], 1)
        self.assertEqual(preview["counts"]["new_kitchens"], 1)
        confirmed = self.confirm_reference(preview)
        self.assertEqual(confirmed.status_code, 200, confirmed.get_data(as_text=True))
        self.assertFalse(confirmed.get_json()["idempotent"])
        with server.db() as conn:
            self.assertTrue(conn.execute("SELECT 1 FROM contractors WHERE code='C2'").fetchone())
            kitchen = conn.execute("SELECT * FROM kitchens WHERE code='K2'").fetchone()
            self.assertEqual((kitchen["contractor"], kitchen["name"], kitchen["address"]),
                             ("C2", "Kitchen 2", "Business address 2"))
            audit_json = conn.execute(
                "SELECT metadata_json FROM audit_log WHERE event_type='daily_reference.confirm'"
            ).fetchone()["metadata_json"]
        self.assertNotIn("Kitchen 2", audit_json)
        self.assertNotIn("Business address 2", audit_json)

        replay_preview = self.preview_reference(
            "contractor_kitchen_reference", workbook, "renamed-same-bytes.xlsx",
        ).get_json()
        replay = self.confirm_reference(replay_preview)
        self.assertEqual(replay.status_code, 200, replay.get_data(as_text=True))
        self.assertTrue(replay.get_json()["idempotent"])
        with server.db() as conn:
            self.assertEqual(conn.execute("SELECT COUNT(*) FROM kitchens WHERE code='K2'").fetchone()[0], 1)
            self.assertEqual(conn.execute("SELECT COUNT(*) FROM daily_reference_import_receipts").fetchone()[0], 1)

    def test_wrong_hash_and_database_change_consume_token_without_partial_write(self):
        preview = self.preview_reference("contractor_kitchen_reference").get_json()
        stale = self.confirm_reference(preview, "0" * 64)
        self.assertEqual(stale.status_code, 409, stale.get_data(as_text=True))
        self.assertEqual(stale.get_json()["code"], "stale_preview")
        self.assertEqual(self.confirm_reference(preview).status_code, 410)

        preview = self.preview_reference("contractor_kitchen_reference").get_json()
        with server.db() as conn:
            conn.execute(
                "INSERT INTO contractors(code,name,price_group,pricing_mode) VALUES('STALE','STALE','STALE','group')"
            )
        stale_database = self.confirm_reference(preview)
        self.assertEqual(stale_database.status_code, 409, stale_database.get_data(as_text=True))
        self.assertEqual(stale_database.get_json()["code"], "stale_database")
        with server.db() as conn:
            self.assertFalse(conn.execute("SELECT 1 FROM kitchens WHERE code='K2'").fetchone())
            self.assertFalse(conn.execute("SELECT 1 FROM daily_reference_import_receipts").fetchone())

    def test_supplier_confirm_changes_only_mapping_and_is_replay_safe(self):
        with server.db() as conn:
            protected_before = tuple(conn.execute(
                """SELECT name,unit,tax,buy_price,purchase_list,seller,cccd
                   FROM products WHERE code='P1'"""
            ).fetchone())
        preview = self.preview_reference("supplier_reference").get_json()
        self.assertEqual(preview["counts"]["updated_mappings"], 1)
        confirmed = self.confirm_reference(preview)
        self.assertEqual(confirmed.status_code, 200, confirmed.get_data(as_text=True))
        with server.db() as conn:
            product = conn.execute("SELECT * FROM products WHERE code='P1'").fetchone()
            protected_after = tuple(conn.execute(
                """SELECT name,unit,tax,buy_price,purchase_list,seller,cccd
                   FROM products WHERE code='P1'"""
            ).fetchone())
            prices = conn.execute("SELECT COUNT(*) FROM product_prices WHERE product_code='P1'").fetchone()[0]
        self.assertEqual(product["supplier"], "S2")
        self.assertEqual(protected_after, protected_before)
        self.assertEqual(prices, 1)
        replay = self.confirm_reference(
            self.preview_reference("supplier_reference", filename="renamed.xlsx").get_json()
        )
        self.assertEqual(replay.status_code, 200, replay.get_data(as_text=True))
        self.assertTrue(replay.get_json()["idempotent"])

    def test_reference_write_failure_rolls_back_every_table(self):
        with server.db() as conn:
            conn.execute(
                """CREATE TRIGGER fail_reference_kitchen BEFORE INSERT ON kitchens
                   WHEN NEW.code='K2' BEGIN SELECT RAISE(ABORT,'forced reference failure'); END"""
            )
        preview = self.preview_reference("contractor_kitchen_reference").get_json()
        failed = self.confirm_reference(preview)
        self.assertEqual(failed.status_code, 500, failed.get_data(as_text=True))
        with server.db() as conn:
            conn.execute("DROP TRIGGER fail_reference_kitchen")
            self.assertFalse(conn.execute("SELECT 1 FROM contractors WHERE code='C2'").fetchone())
            self.assertFalse(conn.execute("SELECT 1 FROM kitchens WHERE code='K2'").fetchone())
            self.assertFalse(conn.execute("SELECT 1 FROM daily_reference_import_receipts").fetchone())
            self.assertFalse(conn.execute(
                "SELECT 1 FROM audit_log WHERE event_type='daily_reference.confirm'"
            ).fetchone())


if __name__ == "__main__":
    unittest.main()
