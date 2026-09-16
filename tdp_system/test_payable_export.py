from __future__ import annotations

import io
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

try:
    from . import server
    from .payable_ledger import sync_payable_ledger
except ImportError:  # pragma: no cover - direct invocation
    import server
    from payable_ledger import sync_payable_ledger


class PayableExportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp_dir = tempfile.TemporaryDirectory()
        cls.original_db_path = server.DB_PATH
        cls.original_data_dir = server.DATA_DIR
        cls.original_export_dir = server.EXPORT_DIR
        cls.original_testing = server.app.config.get("TESTING", False)
        server.DATA_DIR = Path(cls.temp_dir.name) / "data"
        server.EXPORT_DIR = Path(cls.temp_dir.name) / "exports"
        server.DATA_DIR.mkdir(parents=True, exist_ok=True)
        server.EXPORT_DIR.mkdir(parents=True, exist_ok=True)
        server.DB_PATH = server.DATA_DIR / "payable_export.sqlite3"
        server.app.config["TESTING"] = True
        server.init_database()
        cls.client = server.app.test_client()

    @classmethod
    def tearDownClass(cls):
        server.DB_PATH = cls.original_db_path
        server.DATA_DIR = cls.original_data_dir
        server.EXPORT_DIR = cls.original_export_dir
        server.app.config["TESTING"] = cls.original_testing
        cls.temp_dir.cleanup()

    def setUp(self):
        with server.db() as conn:
            conn.execute("DELETE FROM payable_payment_revisions")
            conn.execute("DELETE FROM payable_payment_allocations")
            conn.execute("DELETE FROM payments")
            conn.execute("DELETE FROM payable_ledger_revisions")
            conn.execute("DELETE FROM payable_ledger_lines")
            conn.execute("DELETE FROM purchase_workbook_line_revisions")
            conn.execute("DELETE FROM purchase_workbook_lines")
            conn.execute("DELETE FROM purchase_order_imports")
            conn.execute("DELETE FROM purchase_order_lines")
            conn.execute("DELETE FROM orders")
            conn.execute("DELETE FROM batches")
            conn.execute("DELETE FROM historical_payable_lines")
            conn.execute("DELETE FROM settings WHERE key='historical_payables_through_date'")
            conn.execute("DELETE FROM audit_log")
            conn.execute(
                "INSERT OR REPLACE INTO suppliers(code,name) VALUES('S1','Nhà cung cấp Một')"
            )
            conn.execute(
                "INSERT OR REPLACE INTO suppliers(code,name) VALUES('S2','Nhà cung cấp Hai')"
            )

    @staticmethod
    def _batch(conn, *, status="approved") -> int:
        timestamp = server.now_iso()
        return int(conn.execute(
            """INSERT INTO batches(work_date,source_name,status,created_at,approved_at)
               VALUES('2026-09-01','payable export fixture',?,?,?)""",
            (status, timestamp, timestamp if status == "approved" else None),
        ).lastrowid)

    @staticmethod
    def _purchase(
        conn, batch_id: int, *, row_key: str, supplier: str, base_qty: float,
        damaged: float = 0, added: float = 0, reduced: float = 0,
        missing: float = 0, buy_price: int = 10,
    ) -> None:
        actual_qty = base_qty + added - damaged - reduced - missing
        amount = int(actual_qty * buy_price)
        timestamp = server.now_iso()
        conn.execute(
            """INSERT INTO purchase_workbook_lines(
                   batch_id,row_key,source_sheet,source_row,product_code,kitchen,work_date,
                   product_name,base_qty,unit,supplier,note,buy_price,price_source,
                   damaged_qty,added_qty,reduced_qty,missing_qty,actual_qty,amount,status,
                   source_hash,revision,created_at,updated_at
               ) VALUES(?,?,'đặt hàng',3,?,'K1','2026-09-01',?,?,'kg',?, ?,?,
                        'fixture',?,?,?,?,?,?,'confirmed',?,1,?,?)""",
            (
                batch_id, row_key, f"P-{row_key}", f"Hàng {row_key}", base_qty,
                supplier, f"Ghi chú {row_key}", buy_price, damaged, added, reduced,
                missing, actual_qty, amount, f"HASH-{row_key}", timestamp, timestamp,
            ),
        )

    def _lines(self):
        with server.db() as conn:
            approved = self._batch(conn)
            draft = self._batch(conn, status="draft")
            self._purchase(
                conn, approved, row_key="S1-A", supplier="S1", base_qty=10,
                damaged=1, added=2, missing=1,
            )
            self._purchase(conn, approved, row_key="S1-B", supplier="S1", base_qty=20)
            self._purchase(conn, approved, row_key="S2-A", supplier="S2", base_qty=30)
            self._purchase(conn, draft, row_key="S2-REVERSED", supplier="S2", base_qty=4)
            sync_payable_ledger(conn, timestamp=server.now_iso())
            return {
                row["source_ref"].rsplit(":", 1)[-1]: dict(row)
                for row in conn.execute("SELECT * FROM payable_ledger_lines ORDER BY id")
            }

    def _pay(self, request_id, supplier, amount, allocations):
        response = self.client.post("/api/debts/payables/payments", json={
            "request_id": request_id,
            "party_code": supplier,
            "payment_date": "2026-09-05",
            "amount": amount,
            "method": "Chuyển khoản",
            "reference_code": request_id,
            "note": "Thanh toán kiểm thử",
            "allocations": allocations,
        })
        self.assertEqual(response.status_code, 201, response.get_data(as_text=True))
        return response.get_json()

    def _fixture_with_payments(self):
        lines = self._lines()
        first = self._pay("EXPORT-PAY-001", "S1", 40, [{
            "ledger_line_id": lines["S1-A"]["id"], "amount": 40,
            "expected_revision": lines["S1-A"]["revision"],
        }])
        second = self._pay("EXPORT-PAY-002", "S1", 110, [{
            "ledger_line_id": lines["S1-A"]["id"], "amount": 60,
            "expected_revision": lines["S1-A"]["revision"] + 1,
        }, {
            "ledger_line_id": lines["S1-B"]["id"], "amount": 50,
            "expected_revision": lines["S1-B"]["revision"],
        }])
        to_reverse = self._pay("EXPORT-PAY-003", "S2", 80, [{
            "ledger_line_id": lines["S2-A"]["id"], "amount": 80,
            "expected_revision": lines["S2-A"]["revision"],
        }])
        reversed_response = self.client.post(
            f"/api/debts/payables/payments/{to_reverse['id']}/reverse",
            json={"actor": "Người kiểm thử", "expected_revision": 1, "reason": "Giao dịch ngân hàng bị hủy"},
        )
        self.assertEqual(reversed_response.status_code, 200, reversed_response.get_data(as_text=True))
        return lines, first, second, to_reverse

    def _export(self, **query):
        params = {"from": "2026-09-01", "to": "2026-09-30", **query}
        return self.client.get("/api/debts/payables/export", query_string=params)

    def test_workbook_has_exact_customer_columns_and_static_totals(self):
        self._fixture_with_payments()
        response = self._export()
        self.assertEqual(response.status_code, 200, response.get_json(silent=True))
        self.assertIn(
            "Cong_no_phai_tra_2026-09-01_2026-09-30.xlsx",
            response.headers["Content-Disposition"],
        )
        workbook = load_workbook(io.BytesIO(response.data), data_only=False)
        expected_reference_headers = [
            "Tháng", "Tên bếp", "Ngày, tháng", "Tên hàng", "Số lượng", "ĐVT", "NCC",
            "Giá mua", "Hỏng", "Thêm", "Giảm", "Thiếu", "SL thực tế", "Thành tiền",
        ]
        self.assertEqual(workbook.sheetnames, ["Công nợ phải trả", "Tổng NCC", "NCC S1", "NCC S2"])
        sheet = workbook.active
        self.assertEqual([cell.value for cell in sheet[3]], expected_reference_headers)
        self.assertEqual(sheet.max_column, 14)
        self.assertEqual(sheet.max_row, 7)
        self.assertEqual(
            [sheet.cell(7, column).value for column in range(1, 15)],
            ["TỔNG", *([None] * 12), 600],
        )
        self.assertNotIn("Hàng S2-REVERSED", [sheet.cell(row, 4).value for row in range(4, 8)])

        formulas = [
            cell.coordinate
            for sheet in workbook.worksheets
            for row in sheet.iter_rows()
            for cell in row
            if isinstance(cell.value, str) and cell.value.startswith("=")
        ]
        self.assertEqual(formulas, [])

    def test_workbook_topology_print_setup_and_static_source_adjustments(self):
        self._fixture_with_payments()
        response = self._export(supplier="s1")
        self.assertEqual(response.status_code, 200, response.get_json(silent=True))
        workbook = load_workbook(io.BytesIO(response.data), data_only=False)
        self.assertEqual(workbook.sheetnames, ["Công nợ phải trả"])
        detail = workbook.active
        first = next(row for row in detail.iter_rows(min_row=4, values_only=True) if row[3] == "Hàng S1-A")
        self.assertEqual(first[4:14], (10, "kg", "S1", 10, 1, 2, 0, 1, 10, 100))
        self.assertTrue(str(detail.print_area).endswith("$N$6"))
        self.assertTrue(str(detail.auto_filter.ref).endswith("N5"))
        for sheet in workbook.worksheets:
            self.assertEqual(sheet.freeze_panes, "A4")
            self.assertEqual(sheet.page_setup.orientation, "landscape")
            self.assertEqual(sheet.page_setup.paperSize, 9)
            self.assertEqual(sheet.page_setup.fitToWidth, 1)
            self.assertEqual(sheet.page_setup.fitToHeight, 0)
            self.assertEqual(sheet.print_title_rows, "$3:$3")
            self.assertTrue(str(sheet.print_area).startswith(f"'{sheet.title}'!$A$1:"))
            self.assertTrue(sheet.auto_filter.ref.startswith("A3:"))
            self.assertNotIn("9969", str(sheet.freeze_panes))

    def test_invalid_filter_and_integrity_mismatch_fail_closed(self):
        self.assertEqual(
            self.client.get(
                "/api/debts/payables/export",
                query_string={"from": "2026-09-30", "to": "2026-09-01"},
            ).get_json()["code"],
            "invalid_period",
        )
        missing_supplier = self._export(supplier="KHONG-TON-TAI")
        self.assertEqual(missing_supplier.status_code, 400)
        self.assertEqual(missing_supplier.get_json()["code"], "invalid_supplier")

        lines = self._lines()
        with server.db() as conn:
            conn.execute(
                "UPDATE payable_ledger_lines SET paid_amount=1,status='partially_paid' WHERE id=?",
                (lines["S1-A"]["id"],),
            )
        corrupt = self._export()
        self.assertEqual(corrupt.status_code, 409)
        self.assertEqual(corrupt.get_json()["code"], "payable_export_allocation_mismatch")

    def test_empty_period_still_has_customer_sheet_and_zero_totals(self):
        response = self.client.get(
            "/api/debts/payables/export",
            query_string={"from": "2025-01-01", "to": "2025-01-31"},
        )
        self.assertEqual(response.status_code, 200, response.get_json(silent=True))
        workbook = load_workbook(io.BytesIO(response.data))
        self.assertEqual(workbook.sheetnames, ["Công nợ phải trả", "Tổng NCC"])
        sheet = workbook.active
        self.assertEqual(sheet.max_row, 4)
        self.assertEqual([sheet.cell(4, column).value for column in (1, 5, 13, 14)], ["TỔNG", None, None, 0])

    def test_signed_historical_adjustment_is_preserved_not_misread_as_overpayment(self):
        timestamp = server.now_iso()
        with server.db() as conn:
            conn.execute(
                """INSERT INTO historical_payable_lines(
                       purchase_date,kitchen,item_name,qty,unit,supplier,buy_price,
                       damaged_qty,added_qty,reduced_qty,missing_qty,actual_qty,
                       source_amount,calculated_amount,amount,note,source_file,source_sheet,
                       source_row,source_hash,created_at,updated_at
                   ) VALUES('2026-09-01','K1','Điều chỉnh âm',-10,'kg','S1',10,
                            0,0,0,0,-10,-100,-100,-100,'Giữ dấu từ file khách',
                            'history.xlsx','2026 chuẩn',88,'NEGATIVE-HASH',?,?)""",
                (timestamp, timestamp),
            )
            sync_payable_ledger(conn, timestamp=timestamp)
        response = self._export(supplier="S1")
        self.assertEqual(response.status_code, 200, response.get_json(silent=True))
        workbook = load_workbook(io.BytesIO(response.data), data_only=False)
        detail = workbook.active
        self.assertEqual([detail.cell(4, column).value for column in (5, 13, 14)], [-10, -10, -100])
        self.assertEqual([detail.cell(5, column).value for column in (5, 13, 14)], [None, None, -100])

    def test_reversed_source_is_kept_in_audit_database_but_omitted_from_customer_sheet(self):
        lines = self._lines()
        self._pay("EXPORT-PAY-REV-SOURCE", "S1", 100, [{
            "ledger_line_id": lines["S1-A"]["id"], "amount": 100,
            "expected_revision": lines["S1-A"]["revision"],
        }])
        with server.db() as conn:
            conn.execute(
                "DELETE FROM purchase_workbook_lines WHERE id=?",
                (lines["S1-A"]["source_id"],),
            )
            sync_payable_ledger(conn, timestamp=server.now_iso())
        response = self._export(supplier="S1")
        self.assertEqual(response.status_code, 200, response.get_json(silent=True))
        workbook = load_workbook(io.BytesIO(response.data), data_only=False)
        detail = workbook.active
        names = [detail.cell(row, 4).value for row in range(4, detail.max_row)]
        self.assertNotIn("Hàng S1-A", names)
        self.assertEqual(detail.cell(detail.max_row, 14).value, 200)
        with server.db() as conn:
            reversed_line = conn.execute(
                "SELECT status,paid_amount FROM payable_ledger_lines WHERE id=?",
                (lines["S1-A"]["id"],),
            ).fetchone()
        self.assertEqual(tuple(reversed_line), ("reversed", 100))

    def test_range_totals_exclude_reimport_history_and_sort_dates_before_suppliers(self):
        lines = self._lines()
        with server.db() as conn:
            # Earlier day deliberately belongs to a supplier sorted after S1.
            conn.execute("UPDATE payable_ledger_lines SET work_date='2026-09-02' WHERE supplier_code='S1'")
            conn.execute("UPDATE payable_ledger_lines SET status='reversed' WHERE id=?", (lines['S1-A']['id'],))
        for status, expected in [('all', 500), ('open,partially_paid,paid', 500), ('reversed', 0)]:
            with self.subTest(status=status):
                response = self._export(status=status)
                book = load_workbook(io.BytesIO(response.data))
                sheet = book.active
                self.assertEqual(sheet.cell(sheet.max_row, 14).value, expected)
                self.assertEqual(sum(sheet.cell(r, 14).value for r in range(4, sheet.max_row)), expected)
                dates = [sheet.cell(r, 3).value for r in range(4, sheet.max_row)]
                self.assertEqual(dates, sorted(dates))
                self.assertTrue(all(sheet.cell(sheet.max_row, c).value is None for c in (5, 9, 10, 11, 12, 13)))
                self.assertNotIn('Lượng', sheet['A2'].value)
                self.assertLessEqual(sheet.row_dimensions[sheet.max_row].height, 32)
                if status in ('all', 'reversed'):
                    self.assertIn('Lịch sử đã đảo', book.sheetnames)
                    history = book['Lịch sử đã đảo']
                    self.assertEqual(history.cell(history.max_row, 14).value, 140)
                for day, total in [('2026-09-01', 300), ('2026-09-02', 200)]:
                    if status != 'reversed':
                        daily = self._export(**{'from':day, 'to':day, 'status':status})
                        daily_book = load_workbook(io.BytesIO(daily.data))
                        self.assertEqual(daily_book.active.cell(daily_book.active.max_row, 14).value, total)
                        daily_book.close()
                book.close()


if __name__ == "__main__":
    unittest.main()
