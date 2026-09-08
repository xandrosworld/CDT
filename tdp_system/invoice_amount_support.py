"""Explain monetary holds and refresh one verified source, without posting stock."""
import hashlib
import json
from decimal import Decimal
from io import BytesIO

from flask import jsonify, request, send_file

try:
    from .invoice_output_editing import output_amount_review
    from .minvoice_portal import portal_validation_error, portal_status
    from .invoice_output_sync import upsert_output_invoice, normalize_output_invoice, normalize_output_item
    from .invoice_mapping import apply_saved_mappings, refresh_linked_batches
except ImportError:
    from invoice_output_editing import output_amount_review
    from minvoice_portal import portal_validation_error, portal_status
    from invoice_output_sync import upsert_output_invoice, normalize_output_invoice, normalize_output_item
    from invoice_mapping import apply_saved_mappings, refresh_linked_batches


class AmountReviewError(ValueError):
    pass


def digest(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True, ensure_ascii=False).encode()).hexdigest()


def load_invoice(conn, invoice_id, tenant):
    row = conn.execute("SELECT * FROM outgoing_source_invoices WHERE id=? AND tenant=?",
                       (invoice_id, tenant)).fetchone()
    if row is None or not output_amount_review(row):
        raise AmountReviewError("Hóa đơn không còn ở trạng thái chờ đối chiếu tiền. Đóng bảng và tải lại danh sách.")
    return dict(row)


def fingerprint(conn, row):
    lines = [dict(r) for r in conn.execute(
        "SELECT * FROM outgoing_source_invoice_items WHERE invoice_id=? ORDER BY id", (row['id'],))]
    return digest([row, lines])


def inspect_amounts(row, raw, *, fresh=False):
    """Arithmetic observations only: no invented missing rows or inferred edits."""
    lines = []
    money = lambda value: Decimal(str(value))
    for index, line in enumerate(raw['invoiceDetail'], 1):
        qty, price, amount, tax = (money(line[k]) for k in ('quantity', 'unitPrice', 'amountWithoutVAT', 'vatAmount'))
        note = ''
        extras = ('serviceFeeAmount', 'exciseTaxAmount', 'environmentalProtectionFee', 'deductionAmount', 'ortherTax', 'ortherFee')
        if line.get('property') == 1 and not any(line.get(k) for k in extras):
            expected = qty * price - money(line.get('discountAmount') or 0)
            if abs(expected - amount) > 1:
                note = 'Cần đối chiếu: số lượng × đơn giá − chiết khấu khác tiền dòng.'
        lines.append({'line': index, 'name': str(line.get('productName') or ''),
                      'unit': str(line.get('unitCode') or ''), 'qty': float(qty), 'price': float(price),
                      'amount': float(amount), 'tax': float(tax), 'note': note})
    detail_amount = sum(money(l['amountWithoutVAT']) for l in raw['invoiceDetail'])
    detail_tax = sum(money(l['vatAmount']) for l in raw['invoiceDetail'])
    comparisons = [{'label': label, 'detail': float(detail), 'header': float(money(raw[key])),
                    'difference': float(money(raw[key]) - detail)} for label, detail, key in (
                        ('Tiền hàng chưa thuế', detail_amount, 'totalAmountWithoutVAT'),
                        ('Tiền thuế', detail_tax, 'vatAmount'),
                        ('Tổng gồm thuế', detail_amount + detail_tax, 'totalAmount'))]
    validation = portal_validation_error(raw)
    eligible = not validation and portal_status(raw)[1] == 'issued'
    old = json.loads(row['raw_json'])
    old_lines = old['invoiceDetail']
    changes = []
    if fresh:
        # Explicit source comparison; row numbers here refer to each downloaded document.
        if len(old_lines) != len(lines):
            changes.append(f"Số dòng: {len(old_lines)} → {len(lines)}.")
        for label, key in (('Tiền hàng', 'totalAmountWithoutVAT'), ('Tiền thuế', 'vatAmount'), ('Tổng tiền', 'totalAmount')):
            if old[key] != raw[key]:
                changes.append(f"{label}: {old[key]:,.0f} → {raw[key]:,.0f} đ.")
        for index, line in enumerate(lines):
            before = old_lines[index] if index < len(old_lines) else None
            keys = ('productCode', 'productName', 'unitCode', 'quantity', 'unitPrice', 'amountWithoutVAT', 'vatAmount')
            current = raw['invoiceDetail'][index]
            if before is None:
                line['change'] = 'Dòng mới trong lần tải này.'
            elif any(before.get(k) != current.get(k) for k in keys):
                line['change'] = 'Khác lần tải trước: ' + str(before.get('productName') or '') + f"; lượng {before.get('quantity')}; tiền {before.get('amountWithoutVAT')}."
        for index in range(len(lines), len(old_lines)):
            changes.append(f"Lần tải trước có dòng {index + 1}: {old_lines[index].get('productName', '')}; lần này không còn vị trí này.")
    suspects = sum(bool(l['note']) for l in lines)
    if eligible:
        message = 'Dữ liệu mới đã khớp tổng tiền. Bấm Cập nhật hóa đơn để lưu vào web.'
    elif suspects:
        message = f'{suspects} dòng cần đối chiếu phép tính, đã đánh dấu bên dưới. Chưa đủ căn cứ tự sửa tiền.'
    else:
        message = 'Chưa xác định được dòng thiếu hoặc sai từ dữ liệu hiện có. Không cần tự sửa tiền cho khớp.'
    return {'invoice_id': row['id'], 'series': row['invoice_series'], 'number': row['invoice_number'],
            'date': row['invoice_date'], 'buyer': row['buyer_name'], 'fresh': fresh,
            'lines': lines, 'comparisons': comparisons, 'changes': changes,
            'message': message, 'can_apply': fresh and eligible, 'source_digest': digest(raw),
            'status_note': '' if portal_status(raw)[1] == 'issued' else 'Trạng thái nguồn đã thay đổi; cần kiểm tra hóa đơn trên M-Invoice.'}


def source_document(client_factory, row):
    client = client_factory() if client_factory else None
    if not callable(getattr(client, 'get_outgoing_invoice', None)):
        raise AmountReviewError('Chưa kết nối được chức năng kiểm tra từng hóa đơn. Tải hồ sơ hỗ trợ bên dưới.')
    raw = client.get_outgoing_invoice(remote_id=row['remote_id'], series=row['invoice_series'],
                                      number=row['invoice_number'], invoice_date=row['invoice_date'])
    # Validate types/numeric ranges with the same ingestion contract before offering apply.
    data = normalize_output_invoice(raw, status_map={}, status_fields=(), reference_fields=(), now=row['updated_at'])
    if data['identity_key'] != row['identity_key'] or data['business_key'] != row['business_key']:
        raise AmountReviewError('Định danh hóa đơn nguồn đã thay đổi. Chưa cập nhật dữ liệu.')
    for index, line in enumerate(raw['details'], 1):
        normalize_output_item(line, index)
    return raw


def support_workbook(report):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb = Workbook(); ws = wb.active; ws.title = 'Kiểm tra hóa đơn'
    ws.append([f"HÓA ĐƠN {report['series']} / {report['number']}"])
    ws.append(['Ngày', report['date'], 'Khách hàng', report['buyer']])
    ws.append(['Nguồn: ' + ('vừa kiểm tra trực tiếp từ M-Invoice.' if report['fresh'] else 'dữ liệu hóa đơn đã tải về web.') + ' Chưa xác nhận hóa đơn gốc sai.'])
    ws.append([report['message']])
    ws.append(['Gửi file này cho người hỗ trợ; không cần chụp màn hình.'])
    ws.append([])
    ws.append(['Khoản tiền', 'Cộng chi tiết', 'Tổng hóa đơn', 'Chênh lệch'])
    for c in report['comparisons']:
        ws.append([c['label'], c['detail'], c['header'], c['difference']])
    ws.append([])
    ws.append(['Dòng', 'Tên hàng', 'ĐVT', 'Số lượng', 'Đơn giá', 'Tiền chưa thuế', 'Thuế', 'Cần kiểm tra'])
    for line in report['lines']:
        ws.append([line[k] for k in ('line', 'name', 'unit', 'qty', 'price', 'amount', 'tax', 'note')])
    for row in ws:
        for cell in row:
            # Invoice text is untrusted, including strings starting with =.
            if isinstance(cell.value, str): cell.data_type = 's'
            cell.font = Font(name='Arial', size=13)
            cell.alignment = Alignment(vertical='center', wrap_text=True)
        ws.row_dimensions[row[0].row].height = 42
    for rownum in (1, 7, 12):
        for cell in ws[rownum]:
            cell.fill = PatternFill('solid', fgColor='173750')
            cell.font = Font(name='Arial', size=14, bold=True, color='FFFFFF')
    for r in range(8, 11):
        for c in range(2, 5): ws.cell(r,c).number_format = '#,##0.##'
        if abs(ws.cell(r,4).value) > 1:
            ws.cell(r,4).font = Font(name='Arial', size=14, bold=True, color='B42318')
    for r in range(13, ws.max_row + 1):
        for c in range(4, 8): ws.cell(r,c).number_format = '#,##0.##'
        if ws.cell(r,8).value:
            for cell in ws[r]: cell.fill = PatternFill('solid', fgColor='FFF0ED')
    for r in (1, 3, 4, 5): ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
    ws.merge_cells('D2:H2')
    for c, width in enumerate((24, 42, 20, 23, 20, 24, 20, 54), 1):
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.freeze_panes = 'C13'; ws.auto_filter.ref = f'A12:H{ws.max_row}'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape'; ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.print_title_rows = '1:12'; ws.print_options.horizontalCentered = True
    out = BytesIO(); wb.save(out); out.seek(0); return out


def register_amount_support_routes(app, ctx):
    db = ctx['db']; now_iso = ctx['now_iso']; create_client = ctx.get('create_minvoice_client')
    tenant = lambda conn: str(ctx['setting_get'](conn, 'tenant_code', 'TDP') or 'TDP')

    @app.get('/api/invoice-workbench/output/<int:invoice_id>/amount-review')
    def amount_review(invoice_id):
        try:
            with db() as conn:
                row = load_invoice(conn, invoice_id, tenant(conn)); expected = fingerprint(conn, row)
            fresh = request.args.get('fresh') == '1'
            raw = source_document(create_client, row) if fresh else json.loads(row['raw_json'])
            report = inspect_amounts(row, raw, fresh=fresh)
            if request.args.get('download') == '1':
                return send_file(support_workbook(report), as_attachment=True,
                                 download_name=f"Kiem-tra-hoa-don-{invoice_id}.xlsx",
                                 mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
            return jsonify(ok=True, expected=expected, **report)
        except AmountReviewError as error:
            return jsonify(ok=False, error=str(error)), 409
        except Exception:
            return jsonify(ok=False, error='Chưa kiểm tra được hóa đơn. Thử lại hoặc tải hồ sơ hỗ trợ; dữ liệu vẫn được giữ nguyên.'), 502

    @app.post('/api/invoice-workbench/output/<int:invoice_id>/amount-review/apply')
    def apply_amount_review(invoice_id):
        body = request.get_json(silent=True) or {}
        try:
            with db() as conn:
                row = load_invoice(conn, invoice_id, tenant(conn))
                if body.get('expected') != fingerprint(conn, row):
                    raise AmountReviewError('Dữ liệu hoặc mã hàng đã thay đổi. Bấm Kiểm tra lại M-Invoice trước khi cập nhật.')
            raw = source_document(create_client, row)
            if body.get('source_digest') != digest(raw):
                raise AmountReviewError('M-Invoice vừa thay đổi dữ liệu. Bấm Kiểm tra lại M-Invoice để xem bản mới.')
            if portal_validation_error(raw) or portal_status(raw)[1] != 'issued':
                raise AmountReviewError('Nguồn vẫn cần đối chiếu. Chưa cập nhật hoặc ghi kho.')
            with db() as conn:
                conn.execute('BEGIN IMMEDIATE')
                current = load_invoice(conn, invoice_id, tenant(conn))
                if body.get('expected') != fingerprint(conn, current):
                    raise AmountReviewError('Hóa đơn vừa thay đổi. Kiểm tra lại trước khi cập nhật.')
                now = now_iso()
                iid, created, _, held = upsert_output_invoice(conn, raw, tenant=row['tenant'], now=now,
                    status_map={}, status_fields=(), reference_fields=())
                if iid != invoice_id or created or held:
                    raise AmountReviewError('Chưa cập nhật được đúng hóa đơn. Dữ liệu được giữ nguyên.')
                apply_saved_mappings(conn, 'output', iid)
                refresh_linked_batches(conn, 'output', {iid}, now)
                conn.execute("INSERT INTO audit_log(event_type,entity_type,entity_id,status,message,metadata_json,created_at) VALUES(?, ?, ?, 'ok', ?, ?, ?)",
                    ('invoice_amount_source_refresh', 'outgoing_source_invoice', str(iid),
                     'Cập nhật một hóa đơn từ M-Invoice sau khi người dùng kiểm tra; không ghi kho',
                     json.dumps({'before': digest(json.loads(row['raw_json'])), 'after': digest(raw)}), now))
            return jsonify(ok=True, invoice_id=invoice_id, message='Đã cập nhật hóa đơn. Kiểm tra mã hàng rồi ghi xuất kho khi đủ điều kiện.')
        except AmountReviewError as error:
            return jsonify(ok=False, error=str(error)), 409
        except Exception:
            return jsonify(ok=False, error='Chưa cập nhật được hóa đơn. Dữ liệu được giữ nguyên; hãy thử lại.'), 502
