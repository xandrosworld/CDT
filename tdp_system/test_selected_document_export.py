from __future__ import annotations

import hashlib
import io
import tempfile
import unittest
import zipfile
from pathlib import Path

from openpyxl import load_workbook

try:
    from . import server
except ImportError:  # pragma: no cover - direct file invocation
    import server


class SelectedDocumentExportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory(prefix="tdp_selected_print_")
        cls.original_db_path = server.DB_PATH
        cls.original_data_dir = server.DATA_DIR
        cls.original_testing = server.app.config.get("TESTING", False)
        server.DATA_DIR = Path(cls.temp.name) / "data"
        server.DATA_DIR.mkdir(parents=True, exist_ok=True)
        server.DB_PATH = server.DATA_DIR / "selected-print.sqlite3"
        server.app.config["TESTING"] = True
        server.init_database()
        cls.client = server.app.test_client()

    @classmethod
    def tearDownClass(cls):
        server.DB_PATH = cls.original_db_path
        server.DATA_DIR = cls.original_data_dir
        server.app.config["TESTING"] = cls.original_testing
        cls.temp.cleanup()

    def setUp(self):
        with server.db() as conn:
            conn.execute("DELETE FROM orders")
            conn.execute("DELETE FROM batches")
            conn.execute(
                """INSERT INTO contractors(code,name,price_group,pricing_mode)
                   VALUES('PRINT','Nhà thầu in','PRINT','group')
                   ON CONFLICT(code) DO UPDATE SET name=excluded.name"""
            )
            for code, name in (("BEP-A", "Bếp A"), ("BEP-B", "Bếp B")):
                conn.execute(
                    """INSERT INTO kitchens(code,contractor,name,address)
                       VALUES(?, 'PRINT', ?, 'Hải Phòng')
                       ON CONFLICT(code) DO UPDATE SET contractor=excluded.contractor,
                           name=excluded.name,address=excluded.address""",
                    (code, name),
                )
            timestamp = server.now_iso()
            self.batch_id = int(conn.execute(
                """INSERT INTO batches(work_date,source_name,status,created_at,approved_at)
                   VALUES('2026-09-04','Đơn hàng 04.09.xlsx','approved',?,?)""",
                (timestamp, timestamp),
            ).lastrowid)
            for index, kitchen in enumerate(("BEP-A", "BEP-B"), start=1):
                conn.execute(
                    """INSERT INTO orders(
                           batch_id,work_date,contractor,kitchen,product_code,product_name,
                           qty,actual_received,actual_delivered,unit,supplier,buy_price,
                           sell_price,tax,source_sheet,source_row,updated_at
                       ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                    (
                        self.batch_id, "2026-09-04", "PRINT", kitchen,
                        f"PRINT-{index}", f"Mặt hàng {index}", index,
                        index, index, "Kg", "NCC", 10000, 12000, "8%",
                        "04.09", index + 2, timestamp,
                    ),
                )

    def test_bootstrap_lists_each_delivery_note_for_individual_selection(self):
        response = self.client.get(f"/api/bootstrap?batch_id={self.batch_id}")
        self.assertEqual(200, response.status_code)
        batch = next(item for item in response.get_json()["batches"] if item["id"] == self.batch_id)
        self.assertEqual(
            [("BEP-A", "Bếp A", 1), ("BEP-B", "Bếp B", 1)],
            [(item["code"], item["name"], item["line_count"]) for item in batch["delivery_notes"]],
        )

    def test_selected_delivery_export_contains_only_the_chosen_kitchen(self):
        before = hashlib.sha256(server.DB_PATH.read_bytes()).hexdigest()
        response = self.client.post(
            "/api/export/selected-documents",
            json={
                "batch_ids": [self.batch_id],
                "documents": ["deliveries"],
                "delivery_selections": [
                    {"batch_id": self.batch_id, "kitchen": "BEP-B"},
                ],
            },
        )
        self.assertEqual(200, response.status_code, response.get_json(silent=True))
        self.assertEqual("application/zip", response.mimetype)
        after = hashlib.sha256(server.DB_PATH.read_bytes()).hexdigest()
        self.assertEqual(before, after, "Tải file chọn lọc không được sửa dữ liệu vận hành")

        with zipfile.ZipFile(io.BytesIO(response.data)) as archive:
            names = archive.namelist()
            self.assertEqual(1, len(names))
            workbook = load_workbook(io.BytesIO(archive.read(names[0])), data_only=False)
        try:
            self.assertEqual(["BEP-B"], workbook.sheetnames)
            self.assertEqual("Đơn vị mua hàng: Bếp B", workbook["BEP-B"]["C7"].value)
        finally:
            workbook.close()

    def test_outgoing_statement_keeps_stt_column_compact(self):
        response = self.client.get(f"/api/export/outgoing-statement/{self.batch_id}")
        self.assertEqual(200, response.status_code, response.get_json(silent=True))
        workbook = load_workbook(io.BytesIO(response.data), data_only=False)
        try:
            summary = workbook["Tổng hợp"]
            self.assertLess(summary.column_dimensions["A"].width, 20)
            detail = workbook["PRINT"]
            self.assertLess(detail.column_dimensions["A"].width, detail.column_dimensions["D"].width)
            self.assertEqual(detail["A4"].value, "STT")
        finally:
            workbook.close()

    def test_payroll_total_label_uses_customer_three_column_merge(self):
        created = self.client.post("/api/staff", json={
            "employee_code": "PRINT-NV", "full_name": "Nhân viên in",
            "role_name": "Nhân viên", "kitchen": "BEP-A",
            "base_salary": 8_000_000, "standard_days": 26, "standard_hours": 8,
        })
        self.assertEqual(200, created.status_code, created.get_json(silent=True))
        attendance = self.client.post("/api/attendance", json={
            "employee_code": "PRINT-NV", "work_date": "2026-09-04", "normal_hours": 8,
        })
        self.assertEqual(200, attendance.status_code, attendance.get_json(silent=True))
        response = self.client.get("/api/export/payroll?month=2026-09")
        self.assertEqual(200, response.status_code, response.get_json(silent=True))
        workbook = load_workbook(io.BytesIO(response.data), data_only=False)
        try:
            sheet = workbook["LƯƠNG XƯỞNG"]
            total_row = sheet.max_row
            self.assertEqual("TỔNG", sheet.cell(total_row, 1).value)
            self.assertIn(f"A{total_row}:C{total_row}", {str(item) for item in sheet.merged_cells.ranges})
        finally:
            workbook.close()


if __name__ == "__main__":
    unittest.main()
