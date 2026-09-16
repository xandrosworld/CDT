"""Golden-template purchase receipts grouped by legal seller and day.

The customer's ``biên nhận`` sheet is the authoritative form.  Receipt data
comes from the same confirmed purchase/BK projection as the purchase summary;
sales and receivable values are outside this module.  Identity values are
confined to the official workbook/PDF and never copied into error text or
manifest metadata.
"""

from __future__ import annotations

import math
try:
    from document_totals import quantity_cell
    from document_preview import white_print_style
except ImportError:
    from .document_totals import quantity_cell
    from .document_preview import white_print_style

import re
import unicodedata
from copy import copy
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Any, Iterable, Mapping

from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.pagebreak import RowBreak

try:
    from seller_identity_catalog import name_key, is_excluded_seller
    from purchase_summary_export import (
        EM_THANH_SHA256,
        PurchaseSummaryError,
        aggregate_purchase_summary_rows,
        amount_in_words,
        build_purchase_summary_workbook,
    )
    from template_workbook import assert_workbook_safe, write_literal
except ImportError:  # pragma: no cover - package import path
    from .seller_identity_catalog import name_key, is_excluded_seller
    from .purchase_summary_export import (
        EM_THANH_SHA256,
        PurchaseSummaryError,
        aggregate_purchase_summary_rows,
        amount_in_words,
        build_purchase_summary_workbook,
    )
    from .template_workbook import assert_workbook_safe, write_literal


RECEIPT_TEMPLATE_SHEET = "biên nhận"
RECEIPT_MAX_DAILY_AMOUNT = Decimal("5000000")
ITEM_FIRST_ROW = 15
GOLDEN_ITEM_COUNT = 7
GOLDEN_TOTAL_ROW = 22


class ReceiptExportError(PurchaseSummaryError):
    pass


def _plain(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip())


def _key(value: Any) -> str:
    text = _plain(value).replace("đ", "d").replace("Đ", "D")
    text = unicodedata.normalize("NFKD", text).casefold()
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"[^a-z0-9]+", "", text)


def _row_dict(row: Any) -> dict[str, Any]:
    return dict(row) if not isinstance(row, dict) else row.copy()


def _identity(value: Any) -> str:
    text = _plain(value)
    return text if re.fullmatch(r"(?:\d{9}|\d{12})", text) else ""


def _date_value(value: Any, *, label: str) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _plain(value)
    for pattern in ("%Y-%m-%d", "%d/%m/%Y", "%d.%m.%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            continue
    raise ReceiptExportError(f"{label} không hợp lệ", code="invalid_receipt_date")


def _table_columns(conn: Any, table: str) -> set[str]:
    return {str(row[1]) for row in conn.execute(f"PRAGMA table_info({table})")}


def enrich_receipt_identity_rows(
    conn: Any, rows: Iterable[Mapping[str, Any]],
) -> list[dict[str, Any]]:
    """Use one current seller record for all identity fields; refuse ambiguity."""

    materialized = [_row_dict(row) for row in rows]
    columns = _table_columns(conn, "people")
    required = {"name", "cccd", "issue_date", "issue_place", "address"}
    if not required.issubset(columns):
        raise ReceiptExportError(
            "Danh mục định danh chưa đủ trường để lập biên nhận",
            code="receipt_identity_catalogue_incomplete",
        )
    people: dict[str, list[dict[str, Any]]] = {}
    identities: dict[str, int] = {}
    for source in conn.execute(
        "SELECT name,cccd,issue_date,issue_place,address FROM people ORDER BY name"
    ):
        person = _row_dict(source)
        people.setdefault(name_key(person.get("name")), []).append(person)
        identity = _identity(person.get("cccd"))
        # Excluded historical profiles cannot issue new documents. Their retained
        # identity must not block the customer-confirmed, eligible seller.
        if identity and not is_excluded_seller(person.get("name")):
            identities[identity] = identities.get(identity, 0) + 1

    invalid: list[str] = []
    output: list[dict[str, Any]] = []
    for index, item in enumerate(materialized, start=1):
        source_ref = int(item.get("source_ref") or index)
        if is_excluded_seller(item.get("seller")):
            invalid.append(f"dòng nguồn {source_ref}: người bán đã bị loại khỏi bảng kê/biên nhận")
            continue
        matches = people.get(name_key(item.get("seller")), [])
        issues = []
        if len(matches) != 1:
            issues.append("người bán không duy nhất trong danh mục định danh")
            person: dict[str, Any] = {}
        else:
            person = matches[0]
        if not _identity(person.get("cccd")) or _identity(person.get("cccd")) != _identity(item.get("cccd")):
            issues.append("số định danh không khớp danh mục hiện hành")
        if identities.get(_identity(person.get("cccd")), 0) > 1:
            issues.append("số định danh đang dùng cho nhiều người bán; cần đối chiếu nguồn")
        address = _plain(person.get("address"))
        if not address:
            issues.append("thiếu địa chỉ người bán")
        issue_date = _plain(person.get("issue_date"))
        issue_place = _plain(person.get("issue_place"))
        if not issue_date:
            issues.append("thiếu ngày cấp")
        else:
            try:
                issue_date = _date_value(issue_date, label="Ngày cấp").strftime("%d/%m/%Y")
            except ReceiptExportError:
                issues.append("ngày cấp không hợp lệ")
        if not issue_place:
            issues.append("thiếu nơi cấp")
        if issues:
            invalid.append(f"dòng nguồn {source_ref}: " + ", ".join(issues))
            continue
        enriched = item.copy()
        enriched["address"] = address
        enriched["issue_date"] = issue_date
        enriched["issue_place"] = issue_place
        output.append(enriched)

    if invalid:
        detail = "; ".join(invalid[:12])
        suffix = f"; còn {len(invalid) - 12} dòng" if len(invalid) > 12 else ""
        raise ReceiptExportError(
            "Chưa thể lập biên nhận vì dữ liệu định danh chưa hợp lệ: " + detail + suffix,
            code="invalid_receipt_identity",
        )
    return output


def group_receipt_rows(
    rows: Iterable[Mapping[str, Any]],
    *,
    max_daily_amount: Decimal = RECEIPT_MAX_DAILY_AMOUNT,
) -> list[dict[str, Any]]:
    """Create exactly one receipt for each legal seller on each purchase day."""

    materialized = [_row_dict(row) for row in rows]
    aggregated = aggregate_purchase_summary_rows(materialized)
    identities: dict[tuple[str, str], dict[str, str]] = {}
    for index, item in enumerate(materialized, start=1):
        work_date = _date_value(item.get("work_date"), label=f"Ngày mua dòng {index}").isoformat()
        identity = _identity(item.get("cccd"))
        issue_date = _plain(item.get("issue_date"))
        issue_place = _plain(item.get("issue_place"))
        if not identity or not issue_date or not issue_place or not _plain(item.get("address")):
            raise ReceiptExportError(
                f"Dòng {index} thiếu địa chỉ, định danh, ngày cấp hoặc nơi cấp hợp lệ",
                code="invalid_receipt_identity",
            )
        issue_date = _date_value(issue_date, label=f"Ngày cấp dòng {index}").strftime("%d/%m/%Y")
        signature = {
            "seller": _plain(item.get("seller")),
            "address": _plain(item.get("address")),
            "cccd": identity,
            "issue_date": issue_date,
            "issue_place": issue_place,
        }
        key = (work_date, identity)
        current = identities.get(key)
        if current is not None and current != signature:
            raise ReceiptExportError(
                f"Dòng {index} mâu thuẫn thông tin định danh người bán",
                code="conflicting_receipt_identity",
            )
        identities[key] = signature

    grouped: dict[tuple[str, str], dict[str, Any]] = {}
    for item in aggregated:
        key = (str(item["work_date"]), str(item["cccd"]))
        group = grouped.setdefault(key, {**identities[key], "work_date": key[0], "items": []})
        group["items"].append(item)

    output = []
    for key in sorted(grouped, key=lambda value: (value[0], _key(grouped[value]["seller"]), value[1])):
        group = grouped[key]
        total = sum((Decimal(str(item["amount"])) for item in group["items"]), Decimal("0"))
        if total > max_daily_amount:
            raise ReceiptExportError(
                "Tổng biên nhận của một người bán trong ngày vượt 5.000.000 đồng",
                code="receipt_daily_limit_exceeded",
            )
        group["total_amount"] = int(total.quantize(Decimal("1"), rounding=ROUND_HALF_UP))
        group["total_quantity"] = sum(
            (Decimal(str(item["quantity"])) for item in group["items"]), Decimal("0")
        )
        output.append(group)
    return output


def _row_snapshot(sheet: Any, row: int, *, include_values: bool) -> dict[str, Any]:
    cells = []
    for column in range(1, sheet.max_column + 1):
        cell = sheet.cell(row, column)
        if isinstance(cell, MergedCell):
            cells.append(None)
            continue
        cells.append({
            "value": cell.value if include_values else None,
            "style": copy(cell._style),
            "number_format": cell.number_format,
            "font": copy(cell.font),
            "fill": copy(cell.fill),
            "border": copy(cell.border),
            "alignment": copy(cell.alignment),
            "protection": copy(cell.protection),
        })
    return {"height": sheet.row_dimensions[row].height, "cells": cells}


def _apply_row_snapshot(sheet: Any, row: int, snapshot: Mapping[str, Any]) -> None:
    sheet.row_dimensions[row].height = snapshot.get("height")
    for column, source in enumerate(snapshot["cells"], start=1):
        if source is None:
            continue
        cell = sheet.cell(row, column)
        cell.value = source["value"]
        cell._style = copy(source["style"])
        cell.number_format = source["number_format"]
        cell.font = copy(source["font"])
        cell.fill = copy(source["fill"])
        cell.border = copy(source["border"])
        cell.alignment = copy(source["alignment"])
        cell.protection = copy(source["protection"])


def _set_receipt_font(cell: Any, *, size: float, bold: bool | None = None,
                      italic: bool | None = None) -> None:
    font = copy(cell.font)
    font.name = "Times New Roman"
    font.sz = size
    if bold is not None:
        font.bold = bold
    if italic is not None:
        font.italic = italic
    cell.font = font


def _merge_receipt_range(sheet: Any, reference: str) -> None:
    """Replace any smaller overlapping template merge with one clean range."""

    target = sheet[reference]
    min_column = target[0][0].column
    min_row = target[0][0].row
    max_column = target[-1][-1].column
    max_row = target[-1][-1].row
    for merged in list(sheet.merged_cells.ranges):
        if not (
            merged.max_col < min_column or merged.min_col > max_column
            or merged.max_row < min_row or merged.min_row > max_row
        ):
            sheet.unmerge_cells(str(merged))
    sheet.merge_cells(reference)


def _align_receipt_sheet(sheet: Any, *, total_row: int, signature_row: int) -> None:
    """Lay out the customer's receipt as one balanced, readable A4 form."""

    # Legal heading and identity block.
    for reference in ("C1:G1", "C2:G2", "C3:G3", "C4:G4", "C5:G5",
                      "C6:G6", "C7:G7"):
        _merge_receipt_range(sheet, reference)
    for row_index in range(8, 13):
        _merge_receipt_range(sheet, f"D{row_index}:G{row_index}")

    row_heights = {
        1: 18, 2: 18, 3: 26, 4: 38, 5: 20, 6: 23, 7: 26,
        8: 19, 9: 19, 10: 19, 11: 19, 12: 19, 13: 8, 14: 24,
    }
    for row_index, height in row_heights.items():
        sheet.row_dimensions[row_index].height = height

    for reference, size, bold in (
        ("C1", 12, True), ("C2", 12, True), ("C3", 18, True),
        ("C4", 12, False), ("C5", 12, True), ("C6", 13, False),
        ("C7", 12, False),
    ):
        _set_receipt_font(sheet[reference], size=size, bold=bold)
    for reference in ("C1", "C2", "C3"):
        sheet[reference].alignment = Alignment(horizontal="center", vertical="center")
    for reference in ("C4", "C5", "C6", "C7"):
        sheet[reference].alignment = Alignment(
            horizontal="left", vertical="center", wrap_text=True,
        )
    for row_index in range(8, 13):
        label = sheet.cell(row_index, 3)
        value = sheet.cell(row_index, 4)
        _set_receipt_font(label, size=12, bold=(row_index == 8))
        _set_receipt_font(value, size=12)
        label.alignment = Alignment(horizontal="left", vertical="center")
        value.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # Product table: one font family, enough height to read while checking goods.
    for column in range(3, 8):
        header = sheet.cell(14, column)
        _set_receipt_font(header, size=12, bold=True)
        header.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row_index in range(ITEM_FIRST_ROW, total_row):
        sheet.row_dimensions[row_index].height = max(23, 18 * math.ceil(len(str(sheet.cell(row_index, 3).value or '')) / 30))
        for column in range(3, 8):
            cell = sheet.cell(row_index, column)
            _set_receipt_font(cell, size=12)
            cell.alignment = Alignment(
                horizontal="left" if column == 3 else "center" if column == 4 else "right",
                vertical="center",
                wrap_text=(column == 3),
            )
        sheet.cell(row_index, 5).number_format = "#,##0"
        sheet.cell(row_index, 6).number_format = "#,##0.######"
        sheet.cell(row_index, 7).number_format = "#,##0"

    _merge_receipt_range(sheet, f"C{total_row}:E{total_row}")
    sheet.row_dimensions[total_row].height = 22
    for column in range(3, 8):
        cell = sheet.cell(total_row, column)
        _set_receipt_font(cell, size=12, bold=True)
        cell.alignment = Alignment(
            horizontal="center" if column <= 6 else "right", vertical="center",
        )
    sheet.cell(total_row, 6).number_format = "#,##0.######"
    sheet.cell(total_row, 6).alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
    sheet.row_dimensions[total_row].height = max(24, 18 * math.ceil(len(str(sheet.cell(total_row, 6).value or '')) / 12))
    sheet.cell(total_row, 7).number_format = "#,##0"

    # Notes and signatures use two balanced halves. The date belongs to the
    # seller/signature side instead of floating between both parties.
    _merge_receipt_range(sheet, f"C{total_row + 1}:G{total_row + 1}")
    _merge_receipt_range(sheet, f"C{total_row + 3}:G{total_row + 3}")
    _merge_receipt_range(sheet, f"C{total_row + 4}:G{total_row + 4}")
    _merge_receipt_range(sheet, f"E{total_row + 5}:G{total_row + 5}")
    _merge_receipt_range(sheet, f"C{total_row + 6}:D{total_row + 6}")
    _merge_receipt_range(sheet, f"E{total_row + 6}:G{total_row + 6}")
    _merge_receipt_range(sheet, f"E{signature_row}:G{signature_row}")

    sheet.row_dimensions[total_row + 1].height = 22
    sheet.row_dimensions[total_row + 2].height = 8
    sheet.row_dimensions[total_row + 3].height = 18
    sheet.row_dimensions[total_row + 4].height = 18
    sheet.row_dimensions[total_row + 5].height = 20
    sheet.row_dimensions[total_row + 6].height = 22
    for row_index in range(total_row + 7, signature_row):
        sheet.row_dimensions[row_index].height = 22
    sheet.row_dimensions[signature_row].height = 20

    for row_index in (total_row + 1, total_row + 3, total_row + 4):
        cell = sheet.cell(row_index, 3)
        _set_receipt_font(cell, size=12)
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    date_cell = sheet.cell(total_row + 5, 5)
    _set_receipt_font(date_cell, size=12, italic=True)
    date_cell.alignment = Alignment(horizontal="center", vertical="center")
    for column in (3, 5):
        cell = sheet.cell(total_row + 6, column)
        _set_receipt_font(cell, size=12, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    seller_cell = sheet.cell(signature_row, 5)
    _set_receipt_font(seller_cell, size=12, bold=True)
    seller_cell.alignment = Alignment(horizontal="center", vertical="center")

    sheet.column_dimensions["C"].width = 31
    sheet.column_dimensions["D"].width = 10
    sheet.column_dimensions["E"].width = 14
    sheet.column_dimensions["F"].width = 13
    sheet.column_dimensions["G"].width = 16
    sheet.print_options.horizontalCentered = True
    sheet.print_options.verticalCentered = False


def configure_receipt_paper(sheet: Any, paper: str = "A5") -> None:
    """Lay out the receipt at its physical paper size, before Excel/PDF export.

    A5 needs narrower columns and wrapped prose at a readable font size; merely
    setting PaperSize on the A4 form reduced its 12pt body to less than 8pt.
    This only changes formatting, never the seller identity or purchase values.
    """
    if paper not in {"A4", "A5"}:
        raise ValueError("Biên nhận chỉ hỗ trợ khổ A4 hoặc A5")
    total_row = next((r for r in range(ITEM_FIRST_ROW, sheet.max_row + 1)
                      if sheet.cell(r, 3).value == "TỔNG"), None)
    if total_row is not None:
        signature_row = total_row + 10
        # Excel measures column widths using the workbook's Normal font. The
        # golden template uses Aptos Narrow, which LibreOffice substitutes and
        # Windows Excel does not: identical widths then print at different sizes.
        # Keep the receipt's explicit Times New Roman cells, but use a portable
        # default font for column measurement in both renderers.
        normal_font = Font(name="Arial", size=11)
        sheet.parent._fonts[0] = normal_font
        sheet.parent._named_styles["Normal"].font = normal_font
        _align_receipt_sheet(sheet, total_row=total_row, signature_row=signature_row)
        if paper == "A5":
            for column, width in {"C":24, "D":6.5, "E":11, "F":10, "G":12.5}.items():
                sheet.column_dimensions[column].width = width
            for row in sheet.iter_rows(min_row=1, max_row=signature_row, min_col=3, max_col=7):
                for cell in row:
                    if not isinstance(cell, MergedCell) and cell.value is not None:
                        _set_receipt_font(cell, size=16 if cell.coordinate == "C3" else 11)
            # Heights include the wrapped lines, instead of compressing the
            # entire form to one page as the number of purchased items grows.
            import textwrap
            def height(text, width, minimum=16):
                lines = sum(max(1, len(textwrap.wrap(line, width=width)))
                            for line in str(text or "").split('\n'))
                return max(minimum, lines * 14 + 3)
            for r, h in {1:17, 2:17, 3:25, 13:7, 14:23}.items():
                sheet.row_dimensions[r].height = h
            for r in range(4, 8):
                sheet.row_dimensions[r].height = height(sheet.cell(r, 3).value, 76)
            for r in range(8, 13):
                sheet.row_dimensions[r].height = height(sheet.cell(r, 4).value, 44)
            for r in range(ITEM_FIRST_ROW, total_row):
                sheet.row_dimensions[r].height = max(22, height(sheet.cell(r, 3).value, 26))
            sheet.row_dimensions[total_row].height = height(sheet.cell(total_row, 6).value, 12, 22)
            sheet.row_dimensions[total_row + 1].height = height(sheet.cell(total_row + 1, 3).value, 76)
            sheet.row_dimensions[total_row + 2].height = 7
            for r in (total_row + 3, total_row + 4):
                sheet.row_dimensions[r].height = height(sheet.cell(r, 3).value, 76)
            sheet.row_dimensions[total_row + 5].height = 30
            sheet.cell(total_row + 5, 5).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            sheet.row_dimensions[total_row + 6].height = 22
            for r in range(total_row + 7, signature_row):
                sheet.row_dimensions[r].height = 20
            sheet.row_dimensions[signature_row].height = height(sheet.cell(signature_row, 5).value, 32, 22)
            sheet.cell(signature_row, 5).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            # A short form can spill onto a second sheet because of padding,
            # particularly the signing space. Try a denser form before allowing
            # pagination; retain wrapping, font sizes and room to sign. Long
            # forms keep their normal spacing instead of shrinking indefinitely.
            if sum(sheet.row_dimensions[r].height or 15 for r in range(1, signature_row + 1)) > 620:
                compact = {r: sheet.row_dimensions[r].height for r in range(1, signature_row + 1)}
                def compact_height(text, width, minimum=15):
                    lines = sum(max(1, len(textwrap.wrap(line, width=width)))
                                for line in str(text or '').split('\n'))
                    return max(minimum, lines * 13 + 2)
                compact.update({1:15, 2:15, 3:23, 13:4, 14:20})
                for r in range(4, 8):
                    compact[r] = compact_height(sheet.cell(r, 3).value, 76)
                for r in range(8, 13):
                    compact[r] = compact_height(sheet.cell(r, 4).value, 44)
                for r in range(ITEM_FIRST_ROW, total_row):
                    compact[r] = compact_height(sheet.cell(r, 3).value, 26, 18)
                compact[total_row] = compact_height(sheet.cell(total_row, 6).value, 12, 18)
                for r in (total_row + 1, total_row + 3, total_row + 4):
                    compact[r] = compact_height(sheet.cell(r, 3).value, 76)
                compact.update({total_row + 2:4, total_row + 5:28, total_row + 6:18})
                for r in range(total_row + 7, signature_row):
                    compact[r] = 14
                compact[signature_row] = compact_height(sheet.cell(signature_row, 5).value, 32, 18)
                if sum(compact.values()) <= 620:
                    for r, h in compact.items():
                        sheet.row_dimensions[r].height = h
        # Keep the saved selection reconciliation visible when switching paper
        # sizes or printing an individual receipt without its annex.
        selection_footer = next((r for r in range(signature_row + 1, sheet.max_row + 1)
                                 if str(sheet.cell(r, 3).value or '').startswith('Lựa chọn ngày ')), None)
        if selection_footer:
            signature_row = selection_footer
        sheet.print_area = f"$C$1:$G${signature_row}"
    sheet.row_breaks = RowBreak()
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A5 if paper == "A5" else sheet.PAPERSIZE_A4
    sheet.page_setup.orientation = sheet.ORIENTATION_PORTRAIT
    sheet.page_setup.scale = None
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0 if paper == "A5" else 1
    if paper == "A5" and total_row is not None:
        # A modest fit keeps short receipts and signatures together at about 10pt.
        # Long receipts keep the full-size text and continue on the reverse.
        height_points = sum(sheet.row_dimensions[r].height or 15 for r in range(1, signature_row + 1))
        sheet.page_setup.fitToHeight = 1 if height_points <= 620 else 0
    sheet.print_title_rows = "14:14" if sheet.page_setup.fitToHeight == 0 and total_row is not None else None
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_options.horizontalCentered = True
    sheet.print_options.verticalCentered = False
    sheet.page_margins = PageMargins(left=0.24, right=0.24, top=0.28, bottom=0.28, header=0.1, footer=0.1)


def _populate_receipt_sheet(
    sheet: Any,
    group: Mapping[str, Any],
    *,
    buyer_name: str = "",
    buyer_title: str = "",
    company_name: str = "",
    company_address: str = "",
    location: str = "Hải Phòng",
) -> None:
    items = list(group["items"])
    item_count = len(items)
    if not item_count:
        raise ReceiptExportError("Biên nhận không có dòng hàng", code="empty_receipt")

    item_style = _row_snapshot(sheet, ITEM_FIRST_ROW, include_values=False)
    lower_styles = {
        offset: _row_snapshot(sheet, GOLDEN_TOTAL_ROW + offset, include_values=offset >= 2)
        for offset in range(0, 15)
    }
    for merged in ("C22:E22", "E27:G27", "E32:G32"):
        if merged in {str(value) for value in sheet.merged_cells.ranges}:
            sheet.unmerge_cells(merged)

    delta = item_count - GOLDEN_ITEM_COUNT
    if delta > 0:
        sheet.insert_rows(GOLDEN_TOTAL_ROW, delta)
    elif delta < 0:
        sheet.delete_rows(ITEM_FIRST_ROW + item_count, -delta)

    total_row = ITEM_FIRST_ROW + item_count
    for row_index in range(ITEM_FIRST_ROW, total_row):
        _apply_row_snapshot(sheet, row_index, item_style)
    for offset, snapshot in lower_styles.items():
        _apply_row_snapshot(sheet, total_row + offset, snapshot)

    work_date = _date_value(group["work_date"], label="Ngày biên nhận")
    write_literal(
        sheet,
        "C4",
        "Hôm nay, vào lúc 6 giờ 00 phút, ngày "
        f"{work_date.strftime('%d.%m.%Y')}. Tại {_plain(company_name) or 'đơn vị mua'}, chúng tôi gồm",
    )
    if _plain(company_name):
        write_literal(sheet, "C5", f"I - Tên đơn vị mua: {_plain(company_name)}")
    if _plain(buyer_name):
        buyer_line = f"Người mua hàng: {_plain(buyer_name)}"
        if _plain(buyer_title):
            buyer_line += f" - Chức vụ: {_plain(buyer_title)}"
        write_literal(sheet, "C6", buyer_line)
    elif _plain(buyer_title):
        raise ReceiptExportError(
            "Đã cấu hình chức vụ nhưng thiếu tên người mua hàng",
            code="invalid_receipt_buyer",
        )
    if _plain(company_address):
        write_literal(sheet, "C7", f"Địa chỉ: {_plain(company_address)}")

    write_literal(sheet, "D8", group["seller"])
    write_literal(sheet, "D9", group["address"])
    write_literal(sheet, "D10", group["cccd"])
    sheet["D10"].number_format = "@"
    write_literal(sheet, "D11", group["issue_date"])
    sheet["D11"].number_format = "@"
    write_literal(sheet, "D12", group["issue_place"])

    for row_index, item in enumerate(items, start=ITEM_FIRST_ROW):
        write_literal(sheet, f"C{row_index}", item["product_name"])
        write_literal(sheet, f"D{row_index}", item["unit"])
        write_literal(sheet, f"E{row_index}", item["unit_price"])
        write_literal(sheet, f"F{row_index}", item["quantity"])
        write_literal(sheet, f"G{row_index}", item["amount"])

    sheet.merge_cells(start_row=total_row, start_column=3, end_row=total_row, end_column=5)
    write_literal(sheet, f"C{total_row}", "TỔNG")
    total_quantity = group["total_quantity"]
    write_literal(
        sheet,
        f"F{total_row}",
        quantity_cell(items, "quantity"),
    )
    write_literal(sheet, f"G{total_row}", group["total_amount"])
    write_literal(
        sheet,
        f"C{total_row + 1}",
        "Số tiền bằng chữ: " + amount_in_words(group["total_amount"]) + "./.",
    )
    write_literal(sheet, f"D{total_row + 5}", None)
    write_literal(
        sheet,
        f"E{total_row + 5}",
        f"{_plain(location) or 'Hải Phòng'}, ngày {work_date.day:02d} "
        f"tháng {work_date.month:02d} năm {work_date.year}",
    )
    write_literal(sheet, f"E{total_row + 10}", group["seller"])

    signature_row = total_row + 10
    if sheet.max_row > signature_row:
        sheet.delete_rows(signature_row + 1, sheet.max_row - signature_row)
    configure_receipt_paper(sheet, "A5")


def build_purchase_documents_workbook(
    rows: Iterable[Mapping[str, Any]],
    *,
    template_path: str | Path,
    date_from: Any | None = None,
    date_to: Any | None = None,
    buyer_name: str = "",
    buyer_title: str = "",
    company_name: str = "",
    company_address: str = "",
    location: str = "Hải Phòng",
    expected_sha256: str = EM_THANH_SHA256,
) -> Any:
    materialized = [_row_dict(row) for row in rows]
    receipt_groups = group_receipt_rows(materialized)
    workbook = build_purchase_summary_workbook(
        materialized,
        template_path=template_path,
        date_from=date_from,
        date_to=date_to,
        expected_sha256=expected_sha256,
        additional_template_sheets=(RECEIPT_TEMPLATE_SHEET,),
    )
    try:
        template = workbook[RECEIPT_TEMPLATE_SHEET]
        sheets = [template]
        for _ in receipt_groups[1:]:
            sheets.append(workbook.copy_worksheet(template))
        for index, (sheet, group) in enumerate(zip(sheets, receipt_groups), start=1):
            sheet.title = RECEIPT_TEMPLATE_SHEET if index == 1 else f"biên nhận {index:02d}"
            _populate_receipt_sheet(
                sheet,
                group,
                buyer_name=buyer_name,
                buyer_title=buyer_title,
                company_name=company_name,
                company_address=company_address,
                location=location,
            )
        workbook.active = 0
        workbook.worksheets[0].sheet_view.tabSelected = True
        for sheet in workbook.worksheets[1:]:
            sheet.sheet_view.tabSelected = False
        assert_workbook_safe(workbook)
        white_print_style(workbook)
        return workbook
    except Exception:
        workbook.close()
        raise


__all__ = [
    "RECEIPT_MAX_DAILY_AMOUNT",
    "RECEIPT_TEMPLATE_SHEET",
    "ReceiptExportError",
    "build_purchase_documents_workbook",
    "enrich_receipt_identity_rows",
    "group_receipt_rows",
]
