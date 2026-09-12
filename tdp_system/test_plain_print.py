import io
import unittest
import unicodedata
from copy import copy
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Border,Side
from docx import Document
from docx.oxml.ns import qn
from .document_preview import white_print_style,plain_docx_print_style
from .template_workbook import safe_workbook_bytes


class PlainPrintTests(unittest.TestCase):
    def test_workbook_values_formulas_geometry_survive_plain_print(self):
        wb=Workbook();ws=wb.active
        for value in ['Tiêu đề','Vũ Thị Thụy',unicodedata.normalize('NFD','Vũ Thị Thụy'),42,'=SUM(A4:A4)']:
            ws.append([value])
        ws.merge_cells('A1:C1');ws.print_area='A1:C5';ws.column_dimensions['A'].width=29;ws.row_dimensions[1].height=32
        for cells in ws:
            for cell in cells:
                cell.font=Font(name='Times New Roman',size=13,bold=True,color='FF0000')
                cell.fill=PatternFill('solid',fgColor='007777')
                cell.border=Border(bottom=Side(style='thick',color='007777'))
        original=[(c.coordinate,c.value) for cells in ws for c in cells]
        result=load_workbook(io.BytesIO(safe_workbook_bytes(wb)))
        self.assertEqual([(c.coordinate,c.value) for cells in result.active for c in cells],original)
        self.assertEqual(str(result.active.print_area),str(ws.print_area))
        self.assertEqual(list(result.active.merged_cells.ranges),list(ws.merged_cells.ranges))
        self.assertEqual(result.active.column_dimensions['A'].width,29)
        self.assertEqual(result.active.row_dimensions[1].height,32)
        for cells in result.active:
            for cell in cells:
                if cell.value is None:continue
                self.assertEqual(bool(cell.font.bold),cell.row in (2,3))
                self.assertFalse(cell.fill.fill_type)
                self.assertEqual(cell.font.color.rgb,'00000000')
                self.assertEqual(cell.border.bottom.style,'hair')
        result.close();wb.close()

    def test_word_body_header_and_table_keep_only_named_signer_bold(self):
        doc=Document()
        doc.add_paragraph('Tiêu đề','Title').runs[0].bold=True
        doc.add_paragraph('Vũ Thị Thụy').runs[0].bold=False
        doc.sections[0].header.paragraphs[0].add_run('Tên công ty').bold=True
        doc.add_table(rows=1,cols=1).cell(0,0).text='Tên hàng'
        before=[p.text for p in doc.paragraphs]
        result=plain_docx_print_style(doc)
        self.assertEqual([p.text for p in result.paragraphs],before)
        for root in [doc.element,doc.sections[0].header._element]:
            for run in root.xpath('.//w:r'):
                text=''.join(run.xpath('.//w:t/text()'))
                self.assertEqual(run.rPr.find(qn('w:b')).get(qn('w:val')),'1' if text=='Vũ Thị Thụy' else '0')
                self.assertEqual(run.rPr.find(qn('w:color')).get(qn('w:val')),'000000')
        stream=io.BytesIO();result.save(stream)
        self.assertGreater(len(stream.getvalue()),1000)


if __name__=='__main__':unittest.main()
