"""NXT in the customer's four-group template, with monthly stock valuation."""
from collections import defaultdict
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.views import Selection

try:
    from .invoice_line_tax import tax_rate_label
    from .inventory_report_company import write_company_header
    from .template_workbook import clone_template_workbook, copy_row_layout, assert_workbook_safe
except ImportError:
    from invoice_line_tax import tax_rate_label
    from inventory_report_company import write_company_header
    from template_workbook import clone_template_workbook, copy_row_layout, assert_workbook_safe

TEMPLATE_PATH = Path(__file__).resolve().parent / 'templates' / 'inventory_nxt_template.xlsx'
TEMPLATE_SHA256 = 'DFBBE6AE4730C3A35ADA6D17B20E3C71F6AC0B0A3D967E264C6B1CCA5471F6AB'


def number(value):
    return Decimal(str(value or 0))


def average(value, qty):
    return float((number(value) / number(qty)).quantize(Decimal('0.000001'), rounding=ROUND_HALF_UP)) if qty else None


def literal(ws, row, col, value):
    cell = ws.cell(row, col, value)
    if isinstance(value, str): cell.data_type = 's'
    return cell


def tax_value(item, lines):
    catalog = item.get('tax')
    if catalog not in (None, ''):
        if tax_rate_label(catalog) in {'KCT', 'KKKNT'}:
            return tax_rate_label(catalog)
        try:
            return float(Decimal(str(catalog)))
        except InvalidOperation:
            return str(catalog)
    rates = list(dict.fromkeys(tax_rate_label(r.get('tax_rate')) for r in lines))
    if len(rates) == 1 and rates[0]:
        try:
            return float(Decimal(rates[0].replace('%', '')) / 100)
        except InvalidOperation:
            return rates[0]
    return ' / '.join(r + '%' if r.replace('.', '', 1).isdigit() else r for r in rates if r)


def customer_nxt_workbook(model, sales):
    if model.get('valuation_method') != 'monthly_weighted_average':
        raise ValueError('NXT cần dữ liệu bình quân tháng')
    outgoing = defaultdict(list)
    items = {r['product_code']: dict(r) for r in model['items']}
    for line in sales['lines']:
        mapped = line.get('product_code') and line.get('mapping_status') == 'mapped' and not line.get('identity_warning')
        key = line['product_code'] if mapped else ('source', line['invoice_id'], line.get('id'))
        outgoing[key].append(line)
        if key not in items:
            items[key] = dict(product_code=line.get('product_code') if mapped else '',
                product_name=line.get('product_name') if mapped else line.get('source_item_name'),
                invoice_name=line.get('source_item_name'), unit=line.get('product_unit') if mapped else line.get('source_unit'),
                opening_qty=0, opening_value=0, input_qty=0, input_value=0, output_qty=0,
                closing_qty=0, closing_value=0, average_unit_cost=0, tax='', valuation_status='ok')
    wb = clone_template_workbook(TEMPLATE_PATH, sheet_names=['Ton 7 (2)'], expected_sha256=TEMPLATE_SHA256).workbook
    ws = wb.active; ws.title = 'NXT'
    # The sample's company is obsolete. Only reuse its layout; write the
    # current company identity from settings and normalize oversized rows.
    write_company_header(ws, model)
    for row, height in ((1, 26), (2, 24), (3, 22)):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=18)
        ws.cell(row, 1).alignment = Alignment(horizontal='left', vertical='center', wrap_text=False)
        ws.row_dimensions[row].height = height
    ws.merge_cells('A5:F5')
    ws.row_dimensions[5].height = 25
    ws.row_dimensions[6].height = 32
    ws.row_dimensions[8].height = 25
    ws.row_dimensions[9].height = 28
    year, month = model['date_from'][:7].split('-')
    ws['A5'] = f'THÁNG {int(month)}/{year}'
    ws['A6'] = 'TỔNG'
    summary = sales['output_summary']
    note = 'Tồn cuối: bình quân cả tháng theo từng mã. SL theo sổ kho; tiền Xuất theo hóa đơn M-Invoice.'
    if abs(summary['detail_difference']) > 1:
        note += f" Tổng HĐ lệch cộng chi tiết {summary['detail_difference']:,.0f} đ; xem Đối chiếu."
    ws.merge_cells('A7:R7'); ws['A7'] = note
    ws['A7'].alignment = Alignment(wrap_text=True, vertical='center'); ws.row_dimensions[7].height = 32
    review = wb.create_sheet('Đối chiếu')
    review.append(['NỘI DUNG', 'GIÁ TRỊ'])
    review.append(['Phương pháp giá tồn cuối', 'Bình quân gia quyền cả tháng theo từng mã hàng'])
    review.append(['Công thức đơn giá', '(Giá trị tồn đầu + giá trị nhập thuần trong tháng) / (Lượng tồn đầu + lượng nhập thuần trong tháng)'])
    review.append(['Doanh thu M-Invoice chưa thuế', summary['subtotal']])
    review.append(['Cộng tiền Xuất chi tiết', summary['detail_amount']])
    review.append(['Tổng hóa đơn trừ cộng chi tiết', summary['detail_difference']])
    review.append(['Ghi chú', 'Doanh thu bán không dùng để tính giá trị tồn cuối. Dữ liệu kho và hóa đơn cần được đối chiếu trước khi chốt tháng.'])
    review.append(['Mã hàng', 'Cần đối chiếu'])
    for index, (key, item) in enumerate(items.items(), 1):
        row = 9 + index
        # Row 11 is the blank, sanitized prototype from the supplied workbook.
        copy_row_layout(ws, 11, row, include_values=False)
        ws.row_dimensions[row].height = max(28, 14 * (1 + max(len(str(item.get('product_name') or '')), len(str(item.get('invoice_name') or ''))) // 24))
        revenue = float(sum((number(r.get('amount')) for r in outgoing[key]), Decimal(0)))
        values = [index, item['product_code'], item['product_name'], item['invoice_name'], tax_value(item, outgoing[key]), item['unit'],
                  item['opening_qty'], average(item['opening_value'], item['opening_qty']), item['opening_value'],
                  item['input_qty'], average(item['input_value'], item['input_qty']), item['input_value'],
                  item['output_qty'], average(revenue, item['output_qty']), revenue,
                  item['closing_qty'], item['average_unit_cost'], item['closing_value']]
        for col, value in enumerate(values, 1): literal(ws, row, col, value)
        rate = ws.cell(row, 5).value
        ws.cell(row, 5).number_format = '0%' if not isinstance(rate, (int, float)) or abs(rate * 100 - round(rate * 100)) < 1e-8 else '0.00%'
        for col in (3, 4): ws.cell(row, col).alignment = Alignment(wrap_text=True, vertical='center')
        for col in (7, 10, 13, 16): ws.cell(row, col).number_format = '#,##0.######'
        for col in (8, 9, 11, 12, 14, 15, 17, 18): ws.cell(row, col).number_format = '#,##0.##'
        if item['valuation_status'] != 'ok': review.append([item['product_code'], item['valuation_status']])
        if revenue and not item['output_qty']: review.append([item['product_code'], 'Có tiền bán; chưa có lượng xuất thuần trong sổ kho'])
        if isinstance(key, tuple): review.append([item['product_code'], 'Dòng hóa đơn chưa khớp mã hoặc dòng tài chính'])
    last = max(10, 9 + len(items))
    if ws.max_row > last: ws.delete_rows(last + 1, ws.max_row - last)
    qty = defaultdict(lambda: [Decimal(0)] * 4)
    for item in model['items']:
        for i, field in enumerate(('opening_qty', 'input_qty', 'output_qty', 'closing_qty')):
            qty[item['unit']][i] += number(item[field])
    qty = {unit: values for unit, values in qty.items() if any(values)}
    for i, col in enumerate((7, 10, 13, 16)):
        ws.cell(6, col, float(next(iter(qty.values()))[i]) if len(qty) == 1 else f'{len(qty)} ĐVT · xem Tổng ĐVT' if qty else 0)
    for col, field in ((9, 'opening_value'), (12, 'input_value'), (18, 'closing_value')):
        ws.cell(6, col, float(sum((number(item[field]) for item in model['items']), Decimal(0))))
        ws.cell(6, col).number_format = '#,##0.##'
    ws.cell(6, 15, summary['subtotal']).number_format = '#,##0.##'
    for col in (7, 10, 13, 16):
        ws.cell(6, col).alignment = Alignment(wrap_text=True, horizontal='center', vertical='center')
        ws.cell(6, col).font = Font(name='Times New Roman', size=10, bold=True)
    if len(qty) > 1:
        units = wb.create_sheet('Tổng ĐVT'); units.append(['ĐVT', 'Tồn đầu', 'Nhập', 'Xuất', 'Tồn cuối'])
        for unit, values in sorted(qty.items()): units.append([unit, *map(float, values)])
    for sheet in wb:
        for cells in sheet:
            for cell in cells:
                if isinstance(cell.value, str): cell.data_type = 's'
        if sheet != ws:
            sheet.column_dimensions['A'].width = 42; sheet.column_dimensions['B'].width = 100
            for cells in sheet:
                for cell in cells: cell.alignment = Alignment(wrap_text=True, vertical='top')
    ws.sheet_view.pane = None; ws.sheet_view.selection = [Selection(activeCell='A1', sqref='A1')]
    ws.freeze_panes = 'G10'; ws.auto_filter.ref = f'A9:R{last}'
    ws.print_area = f'A1:R{last}'; ws.print_title_rows = '1:9'
    ws.page_setup.orientation = 'landscape'; ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.scale = None; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.properties.title = f'NXT {month}/{year}'
    wb.properties.description = 'Mẫu khách cung cấp; bình quân cả tháng; tiền Xuất theo M-Invoice.'
    assert_workbook_safe(wb)
    return wb
