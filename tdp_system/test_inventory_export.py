from __future__ import annotations

import io
import sqlite3
import unittest
import zipfile
from pathlib import Path

from flask import Flask
from openpyxl import load_workbook

try:
    from .inventory_export import (
        EXPORT_KINDS,
        InventoryExportError,
        OFFICIAL_LAYOUT_ID,
        OPENING_TEMPLATE_SHA256,
        collect_inventory_export_model,
        inventory_archive_bytes,
        inventory_workbook_bytes,
        register_inventory_export_routes,
    )
    from .test_invoice_input_sync import init_test_database
except ImportError:  # pragma: no cover - direct invocation
    from inventory_export import (
        EXPORT_KINDS,
        InventoryExportError,
        OFFICIAL_LAYOUT_ID,
        OPENING_TEMPLATE_SHA256,
        collect_inventory_export_model,
        inventory_archive_bytes,
        inventory_workbook_bytes,
        register_inventory_export_routes,
    )
    from test_invoice_input_sync import init_test_database


ROOT = Path(__file__).resolve().parent.parent
OPENING_TEMPLATE = ROOT / "_HANDOFF" / "EXTERNAL_INPUTS" / "TĐK T8-2026.xlsx thụy.xlsx"
NOW = "2026-08-31T12:00:00"


class InventoryExportTests(unittest.TestCase):
    def setUp(self):
        self.conn = sqlite3.connect(":memory:")
        init_test_database(self.conn)
        self.conn.execute("INSERT INTO products(code,name,unit) VALUES('P-001','Thịt thử','kg')")
        if self._table("outgoing_product_names"):
            self.conn.execute(
                "INSERT INTO outgoing_product_names(product_code,invoice_name,updated_at) VALUES(?,?,?)",
                ("P-001", "Thịt trên hóa đơn", NOW),
            )
        self.conn.execute(
            """INSERT INTO inventory_transactions(
                   txn_date,product_code,warehouse_codes_json,qty_in,qty_out,unit_cost,source_type,source_id,
                   source_line,status,note,created_at,updated_at
               ) VALUES('2026-08-01','P-001','["KHO-THIT-01"]',10,0,100,'OPENING','2026-08','P-001',
                        'posted','fixture',?,?)""",
            (NOW, NOW),
        )
        self.input_invoice_id, self.input_line_id = self._input_source()
        self.output_invoice_id, self.output_line_id = self._output_source()
        input_revision = self._mapping_revision("input")
        output_revision = self._mapping_revision("output")
        input_confirmation = self._confirmation("input", "msmi_invoices", self.input_invoice_id)
        output_confirmation = self._confirmation(
            "output", "outgoing_source_invoices", self.output_invoice_id,
        )
        self._ledger(
            "IN-EVENT", "input", "POST", "msmi_invoices", self.input_invoice_id,
            self.input_line_id, "2026-08-05", 5, 200, input_revision, input_confirmation,
        )
        self._ledger(
            "OUT-EVENT", "output", "POST", "outgoing_source_invoices",
            self.output_invoice_id, self.output_line_id, "2026-08-10", -3,
            133.333333, output_revision, output_confirmation,
        )
        self.conn.commit()

    def tearDown(self):
        self.conn.close()

    def _table(self, name: str) -> bool:
        return self.conn.execute(
            "SELECT 1 FROM sqlite_master WHERE type='table' AND name=?", (name,),
        ).fetchone() is not None

    def _input_source(self) -> tuple[int, int]:
        cursor = self.conn.execute(
            """INSERT INTO msmi_invoices(
                   remote_id,tenant,invoice_type,seller_tax_code,seller_name,invoice_number,
                   invoice_series,invoice_date,subtotal,tax_amount,total_amount,sync_status,
                   receipt_status,raw_json,error_message,synced_at,created_at,updated_at
               ) VALUES('REMOTE-IN','TDP','INPUT_ELECTRONIC_INVOICE','0200000001',
                        'NCC thử','0001','AA/26E','2026-08-05',1000,80,1080,'synced',
                        'posted','{}','',?,?,?)""",
            (NOW, NOW, NOW),
        )
        invoice_id = int(cursor.lastrowid)
        cursor = self.conn.execute(
            """INSERT INTO msmi_invoice_items(
                   invoice_id,line_index,source_item_code,source_item_name,source_unit,qty,
                   unit_price,amount,tax_rate,source_nature,inventory_eligible,validation_note,
                   product_code,mapping_status
               ) VALUES(?,1,'NCC-THIT','Thịt mua theo HĐ','kg',5,200,1000,'8','1',1,'',
                        'P-001','mapped')""",
            (invoice_id,),
        )
        return invoice_id, int(cursor.lastrowid)

    def _output_source(self) -> tuple[int, int]:
        cursor = self.conn.execute(
            """INSERT INTO outgoing_source_invoices(
                   tenant,source,identity_key,remote_id,business_key,buyer_tax_code,buyer_name,
                   invoice_number,invoice_series,invoice_date,subtotal,tax_amount,total_amount,
                   source_status_raw,source_status_class,source_status_field,relation_reference,
                   sync_status,stock_status,raw_json,error_message,synced_at,created_at,updated_at
               ) VALUES('TDP','msmi','OUT-IDENTITY','REMOTE-OUT','OUT-BUSINESS','0200000002',
                        'Khách thử','0002','BB/26E','2026-08-10',900,72,972,'issued','issued',
                        'fixtureStatus','','synced','posted','{}','',?,?,?)""",
            (NOW, NOW, NOW),
        )
        invoice_id = int(cursor.lastrowid)
        cursor = self.conn.execute(
            """INSERT INTO outgoing_source_invoice_items(
                   invoice_id,line_index,source_item_code,source_item_name,source_unit,qty,
                   unit_price,amount,tax_rate,source_nature,inventory_eligible,validation_note,
                   product_code,mapping_status,conversion_factor,stock_qty,stock_unit_price
               ) VALUES(?,1,'KH-THIT','Thịt bán theo HĐ','kg',3,300,900,'8','1',1,'',
                        'P-001','mapped',1,3,300)""",
            (invoice_id,),
        )
        return invoice_id, int(cursor.lastrowid)

    def _mapping_revision(self, direction: str) -> int:
        invoice_type = (
            "INPUT_ELECTRONIC_INVOICE" if direction == "input" else "OUTPUT_ELECTRONIC_INVOICE"
        )
        cursor = self.conn.execute(
            """INSERT INTO invoice_line_mappings(
                   tenant,source,invoice_type,partner_key,scope_key,source_item_code,
                   source_item_name,source_unit,product_code,target_unit,mapping_status,
                   conversion_factor,effective_from,effective_to,confirmed_at,updated_at
               ) VALUES('TDP','msmi',?,?,?,?,?,?,'P-001','kg','confirmed',1,'','',?,?)""",
            (
                invoice_type,
                direction,
                "scope-" + direction,
                "SRC-" + direction,
                "Thịt " + direction,
                "kg",
                NOW,
                NOW,
            ),
        )
        mapping_id = int(cursor.lastrowid)
        cursor = self.conn.execute(
            """INSERT INTO invoice_mapping_revisions(
                   revision_key,mapping_id,product_code,source_unit,target_unit,
                   conversion_factor,effective_from,effective_to,created_at
               ) VALUES(?,?,'P-001','kg','kg',1,'','',?)""",
            ("revision-" + direction, mapping_id, NOW),
        )
        return int(cursor.lastrowid)

    def _confirmation(
        self, direction: str, table: str, invoice_id: int, *, action: str = "post",
    ) -> int:
        cursor = self.conn.execute(
            """INSERT INTO invoice_inventory_confirmations(
                   confirmation_key,direction,source_invoice_table,source_invoice_id,
                   action,confirmed,note,created_at
               ) VALUES(?,?,?,?,?,1,'fixture',?)""",
            (
                "confirmation-" + direction + "-" + action,
                direction,
                table,
                invoice_id,
                action,
                NOW,
            ),
        )
        return int(cursor.lastrowid)

    def _ledger(
        self, key: str, direction: str, event_type: str, table: str, invoice_id: int,
        line_id: int, txn_date: str, qty: float, cost: float, revision_id: int,
        confirmation_id: int, *, reverses: str = "",
    ) -> None:
        self.conn.execute(
            """INSERT INTO invoice_inventory_ledger(
                   event_key,direction,event_type,source_invoice_table,source_invoice_id,
                   source_line_id,source_line_index,product_code,txn_date,qty_delta,unit_cost,
                   mapping_revision_id,confirmation_id,reverses_event_key,status,created_at
               ) VALUES(?,?,?,?,?,?,1,'P-001',?,?,?,?,?,?,'posted',?)""",
            (
                key, direction, event_type, table, invoice_id, line_id, txn_date, qty, cost,
                revision_id, confirmation_id, reverses, NOW,
            ),
        )

    def _model(self):
        return collect_inventory_export_model(
            self.conn, date_from="2026-08-01", date_to="2026-08-31",
        )

    def test_four_files_share_contract_and_reconcile_to_nxt(self):
        self.assertTrue(OPENING_TEMPLATE.is_file())
        model = self._model()
        self.assertEqual(OFFICIAL_LAYOUT_ID, model["layout_id"])
        self.assertEqual("official_customer_tdk_derived_layout", model["visual_status"])
        self.assertEqual(
            {
                "opening_qty": 10,
                "opening_value": 1000,
                "input_qty": 5,
                "input_value": 1000,
                "output_qty": 3,
                "output_value": 400,
                "closing_qty": 12,
                "closing_value": 1600,
            },
            model["totals"],
        )
        payloads = {
            kind: inventory_workbook_bytes(
                model, kind, template_path=OPENING_TEMPLATE,
            )
            for kind in EXPORT_KINDS
        }
        for kind, payload in payloads.items():
            self.assertTrue(payload.startswith(b"PK"), kind)
            workbook = load_workbook(io.BytesIO(payload), data_only=False, keep_links=False)
            self.assertEqual(model["contract_id"], workbook["_ĐỐI_CHIẾU"]["B1"].value)
            self.assertEqual(kind, workbook["_ĐỐI_CHIẾU"]["B2"].value)
            self.assertEqual("veryHidden", workbook["_ĐỐI_CHIẾU"].sheet_state)
            manifest = {
                row[0].value: row[1].value
                for row in workbook["_ĐỐI_CHIẾU"].iter_rows(min_col=1, max_col=2)
            }
            self.assertEqual(OFFICIAL_LAYOUT_ID, manifest["MẪU_BIỂU"])
            self.assertEqual("CHÍNH THỨC", manifest["TRẠNG_THÁI_BIỂU_MẪU"])
            self.assertEqual("Bình quân gia quyền di động", manifest["PHƯƠNG_PHÁP_GIÁ"])
            self.assertEqual(OPENING_TEMPLATE_SHA256, manifest["NGUỒN_MẪU_TĐK_SHA256"])
            visible = workbook.worksheets[0]
            self.assertEqual("landscape", visible.page_setup.orientation)
            self.assertEqual("9", str(visible.page_setup.paperSize))
            self.assertEqual(1, visible.page_setup.fitToWidth)
            self.assertIsNone(visible.page_setup.scale)
            self.assertTrue(str(visible.print_area))
            self.assertTrue(str(visible.print_title_rows))
            self.assertEqual("Times New Roman", visible["A1"].font.name)
            if kind == "opening":
                self.assertEqual(1, len(visible.sheet_view.selection))
                self.assertEqual("bottomLeft", visible.sheet_view.selection[0].pane)
            self.assertFalse(getattr(workbook, "_external_links", []))
            self.assertFalse(any(
                cell.data_type == "f"
                for sheet in workbook.worksheets
                for row in sheet.iter_rows()
                for cell in row
            ))
            workbook.close()

        opening = load_workbook(io.BytesIO(payloads["opening"]), data_only=True)
        self.assertEqual(["TĐK", "_ĐỐI_CHIẾU"], opening.sheetnames)
        self.assertEqual("TỒN ĐẦU KỲ", opening["TĐK"]["G5"].value)
        self.assertEqual("KHO-THIT-01", opening["TĐK"]["D7"].value)
        self.assertEqual((10, 100, 1000), tuple(opening["TĐK"][cell].value for cell in ("G7", "H7", "I7")))
        opening.close()

        inputs = load_workbook(io.BytesIO(payloads["input"]), data_only=True)
        self.assertEqual(("0001", "NCC thử", 5, 200, 1000), tuple(
            inputs["NHẬP TRONG KỲ"][cell].value for cell in ("D5", "F5", "K5", "L5", "M5")
        ))
        inputs.close()
        outputs = load_workbook(io.BytesIO(payloads["output"]), data_only=True)
        self.assertEqual(("0002", "Khách thử", 3, 133.333333, 400), tuple(
            outputs["XUẤT TRONG KỲ"][cell].value for cell in ("D5", "F5", "K5", "L5", "M5")
        ))
        outputs.close()
        nxt = load_workbook(io.BytesIO(payloads["nxt"]), data_only=True)
        row = nxt["NXT"]
        self.assertEqual(
            ("Mã TĐP", "Tên TĐP", "Tên trên HĐ", "Mã kho", "T/Suất", "ĐVT"),
            tuple(row.cell(4, column).value for column in range(1, 7)),
        )
        self.assertEqual(
            ("TỒN ĐẦU KỲ", "NHẬP TRONG KỲ", "XUẤT TRONG KỲ", "TỒN CUỐI KỲ"),
            tuple(row.cell(4, column).value for column in (7, 10, 13, 16)),
        )
        self.assertEqual(
            (10, 1000, 5, 1000, 3, 400, 12, 1600, "OK"),
            tuple(row[cell].value for cell in ("G6", "I6", "J6", "L6", "M6", "O6", "P6", "R6", "S6")),
        )
        self.assertEqual("KHO-THIT-01", row["D6"].value)
        self.assertIn("bình quân gia quyền di động", row["A3"].value.lower())
        self.assertNotIn("Q-005", row["A2"].value)
        nxt.close()

    def test_mixed_units_are_separate_and_review_is_not_reported_as_matched(self):
        self.conn.execute("INSERT INTO products(code,name,unit) VALUES('P-002','Hàng thứ hai','cái')")
        self.conn.execute("""INSERT INTO inventory_transactions(
            txn_date,product_code,qty_in,qty_out,unit_cost,source_type,source_id,
            source_line,status,note,created_at,updated_at)
            VALUES('2026-08-01','P-002',-3,0,100.5,'OPENING','2026-08','P-002','posted','fixture',?,?)""", (NOW,NOW))
        model = self._model()
        for kind in ('opening','input','output','nxt'):
            payload = inventory_workbook_bytes(model, kind, template_path=OPENING_TEMPLATE)
            wb = load_workbook(io.BytesIO(payload))
            self.assertIn('Tổng ĐVT', wb.sheetnames)
            units = {r[0]: r[1:] for r in wb['Tổng ĐVT'].iter_rows(min_row=2, values_only=True)}
            self.assertEqual(units, {'cái':(-3,0,0,-3),'kg':(10,5,3,12)})
            if kind == 'opening':
                self.assertEqual(wb.active['G4'].value, '2 ĐVT · xem Tổng ĐVT')
                self.assertEqual(wb.active['I4'].number_format, '#,##0')
            if kind == 'nxt':
                last = wb.active.max_row
                self.assertEqual(wb.active.cell(last,7).value, '2 ĐVT · xem Tổng ĐVT')
                self.assertEqual(wb.active.cell(last,19).value, 'CẦN KIỂM TRA')
                self.assertEqual(wb.active['I7'].value, -301.5)
                self.assertEqual(wb.active['I7'].number_format, '#,##0')
            wb.close()

    def test_archive_and_read_only_routes_return_exactly_four_workbooks(self):
        model = self._model()
        archive = inventory_archive_bytes(model, template_path=OPENING_TEMPLATE)
        with zipfile.ZipFile(io.BytesIO(archive)) as bundle:
            self.assertEqual(4, len(bundle.namelist()))
            self.assertEqual(4, sum(name.endswith(".xlsx") for name in bundle.namelist()))

        app = Flask(__name__)
        app.config["TESTING"] = True

        class SharedDb:
            def __enter__(inner):
                return self.conn

            def __exit__(inner, exc_type, exc, tb):
                return False

        register_inventory_export_routes(app, {
            "db": SharedDb,
            "opening_template_path": OPENING_TEMPLATE,
            "opening_template_sha256": OPENING_TEMPLATE_SHA256,
        })
        client = app.test_client()
        response = client.get(
            "/api/invoice-valuation/export?from=2026-08-01&to=2026-08-31"
        )
        self.assertEqual(200, response.status_code)
        self.assertTrue(response.content_type.startswith("application/zip"))
        with zipfile.ZipFile(io.BytesIO(response.data)) as bundle:
            self.assertEqual(4, len(bundle.namelist()))
        one = client.get(
            "/api/invoice-valuation/export/nxt?from=2026-08-01&to=2026-08-31"
        )
        self.assertEqual(200, one.status_code)
        self.assertTrue(one.data.startswith(b"PK"))
        self.assertEqual(
            400,
            client.get(
                "/api/invoice-valuation/export?from=2026-08-31&to=2026-08-01"
            ).status_code,
        )
        self.assertEqual(
            404,
            client.get(
                "/api/invoice-valuation/export/unknown?from=2026-08-01&to=2026-08-31"
            ).status_code,
        )

    def test_missing_invoice_line_trace_is_blocked_not_fabricated(self):
        self.conn.execute("DELETE FROM msmi_invoice_items WHERE id=?", (self.input_line_id,))
        with self.assertRaisesRegex(InventoryExportError, "thiếu dòng hóa đơn nguồn") as raised:
            self._model()
        self.assertEqual("source_trace_missing", raised.exception.code)

    def test_output_reversal_is_negative_in_detail_and_reconciles_nxt(self):
        revision_id = int(self.conn.execute(
            "SELECT mapping_revision_id FROM invoice_inventory_ledger WHERE event_key='OUT-EVENT'"
        ).fetchone()[0])
        confirmation_id = self._confirmation(
            "output", "outgoing_source_invoices", self.output_invoice_id, action="reversal",
        )
        self._ledger(
            "OUT-REVERSAL", "output", "REVERSAL", "outgoing_source_invoices",
            self.output_invoice_id, self.output_line_id, "2026-08-20", 1,
            133.333333, revision_id, confirmation_id, reverses="OUT-EVENT",
        )
        model = self._model()
        self.assertEqual([3, -1], [row["quantity"] for row in model["output_rows"]])
        self.assertEqual([400, -133.33], [row["amount"] for row in model["output_rows"]])
        self.assertEqual(2, model["totals"]["output_qty"])
        self.assertEqual(266.67, model["totals"]["output_value"])
        self.assertEqual(13, model["totals"]["closing_qty"])
        self.assertEqual(1733.33, model["totals"]["closing_value"])

    def test_ui_names_tdk_nxt_and_exposes_period_scoped_four_file_download(self):
        script = (Path(__file__).resolve().parent / "static" / "app.js").read_text(
            encoding="utf-8"
        )
        page = (Path(__file__).resolve().parent / "static" / "index.html").read_text(
            encoding="utf-8"
        )
        for contract in (
            'id="inventoryFrom"', 'id="inventoryTo"',
            "/api/invoice-valuation/export",
            "/api/invoice-valuation/export/opening",
            "/api/invoice-valuation/export/input",
            "/api/invoice-valuation/export/output",
            "/api/invoice-valuation/export/nxt",
            "Tải đủ 4 file ZIP",
            "Xem chi tiết và nhập dữ liệu",
        ):
            self.assertIn(contract, script)
        self.assertNotIn("Q-005", script)
        self.assertIn("Báo cáo vật tư hàng hóa", page)
        self.assertNotIn("Kho hóa đơn", page)


if __name__ == "__main__":
    unittest.main()
