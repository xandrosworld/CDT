import io
import unittest
from openpyxl import load_workbook

from . import test_inventory_period_close as fixtures
from . import test_inventory_export as export_fixtures
from .invoice_monthly_valuation import monthly_average_report
from .invoice_valuation import InvoiceValuationError, moving_average_report
from .inventory_export import monthly_customer_model
from .inventory_customer_report import customer_nxt_workbook, TEMPLATE_PATH
from .invoice_workbench_listing import invoice_range_payload


class MonthlyValuationTests(unittest.TestCase):
    def setUp(self):
        self.fixture = fixtures.InventoryPeriodCloseTests(); self.fixture.setUp()
        self.conn = self.fixture.conn

    def tearDown(self): self.fixture.tearDown()

    def report(self, month='08'):
        return monthly_average_report(self.conn, date_from=f'2026-{month}-01',
            date_to=f'2026-{month}-' + ('31' if month=='08' else '30'), include_zero=True)

    def test_late_purchase_changes_entire_month_average_per_product_without_writes(self):
        self.fixture._event('OUT', 'output', '2026-08-02', -5)
        self.fixture._event('IN', 'input', '2026-08-28', 10, 200)
        self.conn.execute("UPDATE inventory_transactions SET qty_in=4,unit_cost=900 WHERE source_id='2026-08' AND product_code='P2'")
        before = list(self.conn.iterdump())
        rows = {r['product_code']: r for r in self.report()['items']}
        self.assertEqual((150, 2250, 750), tuple(rows['P1'][k] for k in ('average_unit_cost','closing_value','output_value')))
        self.assertEqual((900, 3600), tuple(rows['P2'][k] for k in ('average_unit_cost','closing_value')))
        moving = moving_average_report(self.conn, date_from='2026-08-01', date_to='2026-08-31')
        self.assertEqual(2500, next(r for r in moving['items'] if r['product_code']=='P1')['closing_value'])
        self.assertEqual(before, list(self.conn.iterdump()))

    def test_next_month_uses_previous_month_average_even_before_explicit_close(self):
        self.fixture._event('OUT', 'output', '2026-08-02', -5)
        self.fixture._event('IN', 'input', '2026-08-28', 10, 200)
        september = next(r for r in self.report('09')['items'] if r['product_code']=='P1')
        self.assertEqual((15, 2250, 150), tuple(september[k] for k in ('opening_qty','opening_value','average_unit_cost')))
        result = self.fixture._close()
        self.assertEqual(2250, result['preview']['total_value'])
        stored = self.conn.execute("SELECT unit_cost,qty_in FROM inventory_transactions WHERE source_id='2026-09' AND product_code='P1'").fetchone()
        self.assertEqual((150,15), tuple(stored))
        september = next(r for r in self.report('09')['items'] if r['product_code']=='P1')
        self.assertEqual((15,2250,150), tuple(september[k] for k in ('opening_qty','opening_value','average_unit_cost')))

    def test_month_excludes_next_month_purchase_and_rejects_partial_period(self):
        self.fixture._event('IN-SEP','input','2026-09-01',10,999)
        item = next(r for r in self.report()['items'] if r['product_code']=='P1')
        self.assertEqual(100, item['average_unit_cost'])
        for start,end in [('2026-08-02','2026-08-31'),('2026-08-01','2026-09-30')]:
            with self.assertRaises(InvoiceValuationError):
                monthly_average_report(self.conn,date_from=start,date_to=end)

    def test_zero_stock_month_keeps_finite_average_and_zero_closing_value(self):
        self.fixture._event('OUT','output','2026-08-02',-10)
        item = next(r for r in self.report()['items'] if r['product_code']=='P1')
        self.assertEqual((0,0,100),tuple(item[k] for k in ('closing_qty','closing_value','average_unit_cost')))

    def test_partial_replenishment_of_unpriced_negative_opening_uses_receipt_price(self):
        self.conn.execute("ALTER TABLE products ADD COLUMN tax TEXT DEFAULT ''")
        self.conn.execute("UPDATE products SET tax='KKKNT' WHERE code='P1'")
        self.conn.execute("UPDATE inventory_transactions SET qty_in=0,qty_out=3,unit_cost=0 WHERE source_id='2026-08' AND product_code='P1'")
        self.fixture._event('IN','input','2026-08-12',1,20000)
        before=list(self.conn.iterdump())
        item=next(r for r in self.report()['items'] if r['product_code']=='P1')
        self.assertEqual((-2,-40000,20000),tuple(item[k] for k in ('closing_qty','closing_value','average_unit_cost')))
        self.assertEqual(before,list(self.conn.iterdump()))

    def test_priced_negative_opening_is_not_replaced_by_receipt_price(self):
        self.conn.execute("ALTER TABLE products ADD COLUMN tax TEXT DEFAULT ''")
        self.conn.execute("UPDATE products SET tax='KKKNT' WHERE code='P1'")
        self.conn.execute("UPDATE inventory_transactions SET qty_in=0,qty_out=3,unit_cost=100 WHERE source_id='2026-08' AND product_code='P1'")
        self.fixture._event('IN','input','2026-08-12',1,120)
        item=next(r for r in self.report()['items'] if r['product_code']=='P1')
        self.assertEqual((-2,-180,90),tuple(item[k] for k in ('closing_qty','closing_value','average_unit_cost')))

    def test_replenishment_above_unpriced_shortage_values_only_remaining_stock(self):
        self.conn.execute("ALTER TABLE products ADD COLUMN tax TEXT DEFAULT ''")
        self.conn.execute("UPDATE products SET tax='KKKNT' WHERE code='P1'")
        self.conn.execute("UPDATE inventory_transactions SET qty_in=0,qty_out=3,unit_cost=0 WHERE source_id='2026-08' AND product_code='P1'")
        self.fixture._event('IN','input','2026-08-12',4,20000)
        item=next(r for r in self.report()['items'] if r['product_code']=='P1')
        self.assertEqual((1,20000,20000,60000),tuple(item[k] for k in ('closing_qty','closing_value','average_unit_cost','output_value')))


class CustomerTemplateTests(unittest.TestCase):
    def test_exact_customer_groups_no_extra_tax_columns_and_invoice_revenue(self):
        fixture=export_fixtures.InventoryExportTests();fixture.setUp()
        try:
            model=monthly_customer_model(fixture.conn,fixture._model())
            sales=invoice_range_payload(fixture.conn,tenant='TDP',invoice_type='output',date_from='2026-08-01',date_to='2026-08-31')
            book=customer_nxt_workbook(model,sales);ws=book.active
            template=load_workbook(TEMPLATE_PATH);source=template.active
            for row in (8,9):
                for col in range(1,19):
                    self.assertEqual(source.cell(row,col).value,ws.cell(row,col).value)
                    self.assertEqual(str(source.cell(row,col).font),str(ws.cell(row,col).font))
            self.assertEqual('THÁNG 8/2026',ws['A5'].value)
            self.assertIn('A1:R1', [str(r) for r in ws.merged_cells.ranges])
            self.assertLessEqual(ws.row_dimensions[1].height, 30)
            self.assertLessEqual(ws.row_dimensions[2].height, 30)
            self.assertEqual(900,ws['O6'].value)
            self.assertEqual(900,ws['O10'].value)
            self.assertEqual(1600,ws['R10'].value)
            self.assertEqual(133.333333,ws['Q10'].value)
            self.assertNotIn('Tiền thuế',str(list(ws.values)))
            self.assertEqual('8',str(ws.page_setup.paperSize))
            self.assertEqual('G10',ws.freeze_panes)
            self.assertEqual('A1:R10',str(ws.print_area).split('!')[-1].replace('$',''))
            template.close();book.close()
        finally:fixture.tearDown()
