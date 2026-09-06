import sqlite3
import unittest
from contextlib import contextmanager
from pathlib import Path
from flask import Flask

from . import test_safe_mapping_bulk as fixture
from .automatic_invoice_mapping import apply_automatic_input_mappings
from .contract_modules import register_contract_routes
from .invoice_mapping import save_mapping
from .invoice_workbench_listing import invoice_range_payload, range_workbook
from openpyxl import load_workbook


class AutomaticMappingTests(unittest.TestCase):
    setUp = fixture.SafeBulkMappingTests.setUp
    tearDown = fixture.SafeBulkMappingTests.tearDown

    def apply(self):
        return apply_automatic_input_mappings(self.conn, tenant="TDP", date_from="2026-08-01",
                                             date_to="2026-08-31", now_iso=fixture.now_iso)

    def client(self):
        @contextmanager
        def db():
            with self.conn:
                yield self.conn
        app = Flask(__name__)
        register_contract_routes(app, dict(db=db, now_iso=fixture.now_iso,
            clean_text=lambda v: str(v or "").strip(), number_value=lambda v: float(v or 0),
            tax_factor=lambda v: 1, setting_get=lambda c,k,d="": d, setting_set=lambda *a: None,
            root=Path.cwd(), data_dir=Path.cwd()))
        return app.test_client()

    def test_one_post_without_confirmation_and_retry_has_no_writes(self):
        client = self.client()
        response = client.post('/api/msmi/auto-mappings', json=dict(**{'from':'2026-08-01','to':'2026-08-31'}))
        self.assertEqual(200, response.status_code, response.get_json())
        result = response.get_json()
        self.assertEqual((2,3,1,False), (result['mapped_lines_in_period'], result['mapped_lines_all_periods'],
                                       result['ready_invoices_in_period'], result['stock_changed']))
        before = list(self.conn.iterdump())
        result = client.post('/api/msmi/auto-mappings', json={'from':'2026-08-01','to':'2026-08-31'}).get_json()
        self.assertEqual(0, result['applied_rules'])
        self.assertEqual(before, list(self.conn.iterdump()))
        self.assertEqual(0, self.conn.execute('SELECT COUNT(*) FROM invoice_inventory_ledger').fetchone()[0])

    def test_manual_choice_on_related_invoice_is_preserved(self):
        self.conn.execute("INSERT INTO products(code,name,unit) VALUES('MANUAL','Other rice','kg')")
        item = self.conn.execute("SELECT id FROM msmi_invoice_items WHERE invoice_id=?", (self.invoice_ids[0],)).fetchone()[0]
        save_mapping(self.conn, direction='input', item_id=item, product_code='MANUAL', now_iso=fixture.now_iso)
        # Simulate a historical unlinked line with the same source identity.
        self.conn.execute("UPDATE msmi_invoice_items SET mapping_status='unmapped',product_code='' WHERE id=?", (item,))
        before = list(self.conn.iterdump())
        self.assertEqual(0, self.apply()['applied_rules'])
        self.assertEqual(before, list(self.conn.iterdump()))

    def test_audit_failure_rolls_back_every_mapping(self):
        self.conn.execute("""CREATE TRIGGER reject_auto BEFORE INSERT ON audit_log
            WHEN NEW.event_type='msmi.mapping_auto' BEGIN SELECT RAISE(ABORT,'audit unavailable'); END""")
        before = list(self.conn.iterdump())
        with self.assertRaises(sqlite3.IntegrityError):
            self.apply()
        self.assertEqual(before, list(self.conn.iterdump()))

    def test_attention_then_stock_confirmation_then_completed_same_in_excel(self):
        self.apply()
        payload = invoice_range_payload(self.conn, tenant='TDP', invoice_type='input',
                                        date_from='2026-08-01', date_to='2026-08-31')
        ranks = [r['action_rank'] for r in payload['lines']]
        self.assertEqual(sorted(ranks), ranks)
        self.assertEqual({0,1,2}, set(ranks))
        for row in payload['lines']:
            if row['action_rank'] == 1:
                self.assertTrue(row['needs_confirmation'])
                self.assertFalse(row['issue'])
        wb = load_workbook(range_workbook(payload), data_only=True)
        self.assertEqual([r['source_item_name'] for r in payload['lines']],
                         [wb.active.cell(i+3,5).value for i in range(len(payload['lines']))])
        wb.close()

    def test_invalid_period_cannot_write(self):
        before = list(self.conn.iterdump())
        result = self.client().post('/api/msmi/auto-mappings', json={'from':'2026-09-01','to':'2026-08-01'})
        self.assertEqual(400, result.status_code)
        self.assertEqual(before, list(self.conn.iterdump()))

    def test_download_route_automatically_maps_and_reports_current_batch_counts(self):
        from .invoice_workbench import register_invoice_workbench_routes
        from .test_invoice_input_sync import DateBoundedMsmi
        remote = fixture.remote_invoice(6)
        remote['hdhhdvu'][0].update(ma='NEW-RICE', ten='Gạo', dvtinh='kg')
        @contextmanager
        def db():
            with self.conn:
                yield self.conn
        app = Flask('auto-download')
        register_invoice_workbench_routes(app, dict(db=db, now_iso=fixture.now_iso,
            setting_get=lambda c,k,d='': d, create_msmi_client=lambda: DateBoundedMsmi([remote])))
        client = app.test_client()
        prepared = client.post('/api/invoice-workbench/batches', json=dict(
            source='msmi', invoice_type='input', date_from='2026-08-01', date_to='2026-08-31'))
        self.assertEqual(200, prepared.status_code, prepared.get_json())
        batch = prepared.get_json()['batch']['id']
        response = client.post(f'/api/invoice-workbench/batches/{batch}/sync', json={})
        self.assertEqual(200, response.status_code, response.get_json())
        result = response.get_json()
        self.assertGreater(result['automatic_mapping']['mapped_lines_in_period'], 0)
        self.assertEqual(('ready', 1, 0), (result['status'], result['ready_count'], result['needs_mapping_count']))
        self.assertEqual(0, self.conn.execute('SELECT COUNT(*) FROM invoice_inventory_ledger').fetchone()[0])
