"""Outgoing invoice workbooks backed by the four customer tax goldens."""

from __future__ import annotations

import hashlib
import io
import math
import re
import unicodedata
import zipfile
from collections import defaultdict
from copy import copy
from datetime import datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

try:
    from template_workbook import safe_workbook_bytes
except ImportError:  # pragma: no cover - package invocation
    from .template_workbook import safe_workbook_bytes


INVOICE_HEADERS = (
    "Mã hàng", "Tên hàng", "Đơn vị tính", "Số lượng", "Đơn giá",
    "Cộng tiền hàng", "%CK", "Tiền CK", "Tiền trước thuế", "% VAT",
    "Tiền thuế GTGT", "Tổng tiền", "Tính chất",
)

TEMPLATE_SPECS = {
    "KKKNT": {
        "filename": "thue 0.xlsx",
        "sha256": "82642C993BF26E52B54762E397ED5A1C3F5886F308834582DBCD01C4FBEA0619",
        "sheet": "Sheet2",
    },
    "VAT8": {
        "filename": "thue 8.xlsx",
        "sha256": "892C57A688E28AFADDB58E2070E122945A926FB6EF1DC417714A004885D4C573",
        "sheet": "Sheet 1 (2)",
    },
    "VAT10": {
        "filename": "thue 10.xlsx",
        "sha256": "D4548D308CDD236A854972A629DB63B4CEA2FE1B65816EAFF6AEBE57040D0DE9",
        "sheet": "Sheet 1 (2)",
    },
    "VAT10_PROMOTION": {
        "filename": "thue 10 có khuyến mại.xlsx",
        "sha256": "243EC60B25C3235F447A7C773E66C7948F36B3AA5BE3F3AFD0A4F06A78F0691B",
        "sheet": "Sheet 1 (2)",
    },
}


class InvoiceTaxExportError(ValueError):
    def __init__(self, message: str, *, code: str = "invoice_export_invalid", status: int = 409):
        super().__init__(message)
        self.code = code
        self.status = status


def _number(value: Any, label: str) -> float:
    try:
        result = float(value)
    except (TypeError, ValueError):
        raise InvoiceTaxExportError(f"{label} không phải là số hữu hạn") from None
    if not math.isfinite(result):
        raise InvoiceTaxExportError(f"{label} không phải là số hữu hạn")
    return result


def _vnd(*values: Any) -> int:
    try:
        result = Decimal("1")
        for value in values:
            result *= Decimal(str(value))
        return int(result.quantize(Decimal("1"), rounding=ROUND_HALF_UP))
    except (InvalidOperation, TypeError, ValueError):
        raise InvoiceTaxExportError("Không thể làm tròn số tiền hóa đơn") from None


def _literal(value: Any) -> str:
    text = str(value or "")
    if text.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


def _safe_name(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or "KHAC"))
    text = "".join(char for char in text if not unicodedata.combining(char))
    text = re.sub(r"[^A-Za-z0-9_-]+", "_", text).strip("_")
    return text[:60] or "KHAC"


def _template_kind(vat_percent: float, has_promotion: bool) -> str:
    if abs(vat_percent + 2) <= 1e-9:
        return "KKKNT"
    if abs(vat_percent - 8) <= 1e-9:
        return "VAT8"
    if abs(vat_percent - 10) <= 1e-9:
        return "VAT10_PROMOTION" if has_promotion else "VAT10"
    return "GENERIC"


def _load_golden(template_dir: Path, template_kind: str):
    spec = TEMPLATE_SPECS[template_kind]
    source = template_dir / spec["filename"]
    if not source.is_file():
        raise InvoiceTaxExportError(
            f"Thiếu mẫu hóa đơn chuẩn {spec['filename']}",
            code="invoice_golden_missing",
            status=500,
        )
    payload = source.read_bytes()
    actual_hash = hashlib.sha256(payload).hexdigest().upper()
    if actual_hash != spec["sha256"]:
        raise InvoiceTaxExportError(
            f"Mẫu hóa đơn {spec['filename']} đã khác bản khách chốt; dừng xuất",
            code="invoice_golden_changed",
            status=500,
        )
    workbook = load_workbook(io.BytesIO(payload), data_only=False, keep_links=False)
    if workbook.active.title != spec["sheet"]:
        workbook.close()
        raise InvoiceTaxExportError(
            f"Mẫu hóa đơn {spec['filename']} sai sheet active",
            code="invoice_golden_changed",
            status=500,
        )
    source_sheet = workbook.active
    headers = tuple(source_sheet.cell(1, column).value for column in range(1, 14))
    if headers != INVOICE_HEADERS:
        workbook.close()
        raise InvoiceTaxExportError(
            f"Mẫu hóa đơn {spec['filename']} sai 13 cột chuẩn",
            code="invoice_golden_changed",
            status=500,
        )
    return workbook, source_sheet


def _copy_cell_style(source, target) -> None:
    target.font = copy(source.font)
    target.fill = copy(source.fill)
    target.border = copy(source.border)
    target.number_format = source.number_format
    target.alignment = copy(source.alignment)
    target.protection = copy(source.protection)


def _styled_workbook(template_dir: Path, template_kind: str) -> Workbook:
    if template_kind == "GENERIC":
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Sheet1"
        for column, value in enumerate(INVOICE_HEADERS, start=1):
            sheet.cell(1, column).value = value
        return workbook

    source_book, source_sheet = _load_golden(template_dir, template_kind)
    target_book = Workbook()
    try:
        target_sheet = target_book.active
        target_sheet.title = source_sheet.title
        target_sheet.sheet_format = copy(source_sheet.sheet_format)
        target_sheet.sheet_properties = copy(source_sheet.sheet_properties)
        target_sheet.page_margins = copy(source_sheet.page_margins)
        target_sheet.page_setup = copy(source_sheet.page_setup)
        target_sheet.print_options = copy(source_sheet.print_options)
        for letter in "ABCDEFGHIJKLM":
            target_sheet.column_dimensions[letter].width = source_sheet.column_dimensions[letter].width
            target_sheet.column_dimensions[letter].hidden = source_sheet.column_dimensions[letter].hidden
        target_sheet.row_dimensions[1].height = source_sheet.row_dimensions[1].height
        for column, value in enumerate(INVOICE_HEADERS, start=1):
            target = target_sheet.cell(1, column)
            target.value = value
            _copy_cell_style(source_sheet.cell(1, column), target)
        target_sheet.auto_filter.ref = "A1:M1"
        return target_book
    finally:
        source_book.close()


def build_invoice_workbook(
    rows: list[dict[str, Any]],
    *,
    vat_percent: float,
    template_dir: Path,
) -> tuple[bytes, str]:
    if not rows:
        raise InvoiceTaxExportError("Dự thảo hóa đơn không có dòng hàng")
    natures = {str(row.get("invoice_nature") or "1").strip() for row in rows}
    if not natures <= {"1", "2"}:
        raise InvoiceTaxExportError("Tính chất hóa đơn chỉ nhận 1 hoặc 2")
    has_promotion = "2" in natures
    template_kind = _template_kind(vat_percent, has_promotion)
    workbook = _styled_workbook(template_dir, template_kind)
    try:
        sheet = workbook.active
        style_source_book = style_source_sheet = None
        if template_kind != "GENERIC":
            style_source_book, style_source_sheet = _load_golden(template_dir, template_kind)
        try:
            for output_row, item in enumerate(rows, start=2):
                code = str(item.get("product_code") or "").strip()
                name = str(item.get("product_name") or "").strip()
                unit = str(item.get("unit") or "").strip()
                contractor = str(item.get("contractor") or "").strip()
                if not code or not name or not unit or not contractor:
                    raise InvoiceTaxExportError(
                        f"Dòng dự thảo {item.get('line_id') or item.get('id') or '?'} thiếu mã/tên/ĐVT/nhà thầu"
                    )
                qty = _number(item.get("qty"), f"Số lượng {code}")
                unit_price = _number(item.get("unit_price"), f"Đơn giá {code}")
                if qty <= 0 or unit_price < 0:
                    raise InvoiceTaxExportError(f"Dòng {code} phải có số lượng dương và đơn giá không âm")
                nature = str(item.get("invoice_nature") or "1").strip()
                amount = _vnd(qty, unit_price)
                stored_amount = _number(item.get("amount", amount), f"Thành tiền {code}")
                if abs(stored_amount - amount) > 0.01:
                    raise InvoiceTaxExportError(
                        f"Dòng {code} lệch Số lượng × Đơn giá; cần tạo lại dự thảo",
                        code="invoice_draft_amount_mismatch",
                    )
                if nature == "2" and (abs(unit_price) > 1e-9 or abs(stored_amount) > 0.01):
                    raise InvoiceTaxExportError(
                        f"Dòng khuyến mại {code} phải giữ đơn giá và thành tiền bằng 0"
                    )
                tax_amount = 0 if vat_percent <= 0 else _vnd(amount, vat_percent / 100)
                total = amount + tax_amount
                values = [
                    _literal(code), _literal(name), _literal(unit), qty,
                    None if nature == "2" else unit_price,
                    None if nature == "2" else amount,
                    None, None,
                    None if nature == "2" else amount,
                    vat_percent,
                    None if nature == "2" else tax_amount,
                    None if nature == "2" else total,
                    nature,
                ]
                for column, value in enumerate(values, start=1):
                    target = sheet.cell(output_row, column)
                    target.value = value
                    if style_source_sheet is not None:
                        source_row = 4 if nature == "2" and style_source_sheet.max_row >= 4 else 2
                        _copy_cell_style(style_source_sheet.cell(source_row, column), target)
                sheet.cell(output_row, 4).number_format = "#,##0.######"
                for column in (5, 6, 8, 9, 11, 12):
                    sheet.cell(output_row, column).number_format = "#,##0"
                sheet.cell(output_row, 5).number_format = "#,##0.##########"
            sheet.auto_filter.ref = f"A1:M{sheet.max_row}"
            sheet.freeze_panes = None
            return safe_workbook_bytes(workbook), template_kind
        finally:
            if style_source_book is not None:
                style_source_book.close()
    finally:
        workbook.close()


def _zip_entry(name: str, payload: bytes, work_date: str) -> tuple[zipfile.ZipInfo, bytes]:
    try:
        parsed = datetime.strptime(work_date, "%Y-%m-%d")
        year = max(parsed.year, 1980)
        date_time = (year, parsed.month, parsed.day, 0, 0, 0)
    except ValueError:
        date_time = (1980, 1, 1, 0, 0, 0)
    info = zipfile.ZipInfo(name, date_time=date_time)
    info.compress_type = zipfile.ZIP_DEFLATED
    info.external_attr = 0o600 << 16
    return info, payload


def export_invoice_drafts_zip(
    lines: list[dict[str, Any]],
    *,
    work_date: str,
    template_dir: Path,
) -> io.BytesIO:
    if not lines:
        raise InvoiceTaxExportError(
            "Không có vòng dự thảo chưa phát hành; cần kiểm tra tồn và tạo dự thảo trước khi tải",
            code="invoice_draft_not_found",
        )
    groups: dict[tuple[str, int, float], list[dict[str, Any]]] = defaultdict(list)
    line_ids: set[int] = set()
    for line in lines:
        line_id = int(line.get("line_id") or line.get("id") or 0)
        if not line_id or line_id in line_ids:
            raise InvoiceTaxExportError("Dòng dự thảo trùng hoặc thiếu khóa; dừng xuất")
        line_ids.add(line_id)
        contractor = str(line.get("contractor") or "").strip()
        if not contractor:
            raise InvoiceTaxExportError("Dự thảo thiếu nhà thầu; dừng xuất")
        vat_percent = _number(line.get("vat_percent"), "Thuế suất")
        groups[(contractor, int(line.get("round_no") or 1), vat_percent)].append(line)

    output = io.BytesIO()
    manifest = [
        "FILE TẢI PHẦN MỀM TRUNG GIAN", "",
        "Mỗi file chỉ chứa một nhà thầu, một vòng dự thảo và một nhóm thuế.",
        "Số lượng lấy từ dòng dự thảo đã giữ tồn, không lấy lại toàn bộ đơn khách.", "",
    ]
    with zipfile.ZipFile(output, "w") as archive:
        for (contractor, round_no, vat_percent), items in sorted(groups.items()):
            has_promotion = any(str(item.get("invoice_nature") or "1") == "2" for item in items)
            payload, template_kind = build_invoice_workbook(
                items, vat_percent=vat_percent, template_dir=template_dir,
            )
            label = _template_kind(vat_percent, has_promotion)
            if label == "GENERIC":
                label = f"VAT{vat_percent:g}" if vat_percent >= 0 else f"THUE{vat_percent:g}"
            filename = (
                f"Hoa_don_{_safe_name(contractor)}_{_safe_name(label)}_"
                f"lan_{round_no}_{work_date}.xlsx"
            )
            info, data = _zip_entry(filename, payload, work_date)
            archive.writestr(info, data)
            manifest.append(
                f"- {filename}: {len(items)} dòng · mẫu {template_kind} · vòng {round_no}"
            )
        info, data = _zip_entry("HUONG_DAN.txt", "\r\n".join(manifest).encode("utf-8-sig"), work_date)
        archive.writestr(info, data)
    output.seek(0)
    return output


__all__ = [
    "INVOICE_HEADERS",
    "InvoiceTaxExportError",
    "TEMPLATE_SPECS",
    "build_invoice_workbook",
    "export_invoice_drafts_zip",
]
