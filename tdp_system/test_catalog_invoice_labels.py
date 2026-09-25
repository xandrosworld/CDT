import io
import unittest
import zipfile
from openpyxl import Workbook,load_workbook
from . import server
from . import test_outgoing_readiness as support
from .contract_modules import parse_catalog_workbook
from .outgoing_waiting import refresh_waiting
from .outgoing_names import refresh_editable_names


class CatalogInvoiceLabelsTests(unittest.TestCase):
    setUpClass=classmethod(support.OutgoingReadinessTests.setUpClass.__func__)
    tearDownClass=classmethod(support.OutgoingReadinessTests.tearDownClass.__func__)
    setUp=support.OutgoingReadinessTests.setUp

    def book(self,rows):
        book=Workbook();s=book.active;s.title='danh mục hh'
        s.append(['Mã hàng','Tên Thành Đạt Phát','Tên xuất hóa đơn','ĐVT','Thuế'])
        for row in rows:s.append(row)
        return book

    def preview(self,rows,mode='full'):
        book=self.book(rows);data=io.BytesIO();book.save(data);book.close();data.seek(0)
        response=self.client.post('/api/catalog/import/preview',data={'mode':mode,'file':(data,'Đơn hàng.xlsx')},content_type='multipart/form-data')
        self.assertEqual(200,response.status_code,response.get_json())
        return response.get_json()

    def seed(self):
        with server.db() as conn:
            support.OutgoingReadinessTests.add_opening(conn,20)
            batch,ids=support.OutgoingReadinessTests.add_batch(conn,'2026-09-01',[{'qty':7}])
        return batch,ids

    def test_new_only_recovers_price_sheet_code_without_overwriting_catalog(self):
        book=self.book([])
        price=book.create_sheet('BÁO GIÁ')
        price.append(['Mã hàng','Tên Thành Đạt Phát','ĐVT','Thuế','Giá mua','NCC'])
        price.append([1,2,3,4,5,6])
        price.append(['C000030','Thịt thăn bò CN','Kg','KKKNT',192000,'khánh'])
        price.append(['C000030','Thịt thăn bò cn','Kg','KKKNT',195000,'khác'])
        with server.db() as conn:
            before=list(conn.iterdump())
            result=parse_catalog_workbook(conn,book,mode='new_only')
            self.assertEqual(before,list(conn.iterdump()))
            self.assertTrue(result['can_confirm'],result)
            self.assertEqual(1,result['counts']['new_products'])
            item=result['items'][0]
            self.assertEqual(('C000030','Kg','KKKNT','BÁO GIÁ'),(item['product_code'],item['unit'],item['tax'],item['source_sheet']))
            price.append(['C000030','Thịt khác','Kg','KKKNT'])
            conflict=parse_catalog_workbook(conn,book,mode='new_only')
            self.assertFalse(conflict['can_confirm'])
        book.close()

    def test_cumulative_and_waiting_drafts_use_invoice_name(self):
        batch,_=self.seed()
        with server.db() as conn:
            conn.execute("INSERT OR REPLACE INTO outgoing_product_names VALUES('HH-01','Tên xuất hóa đơn','2026-09-13')")
            refresh_waiting(conn,'2026-09-13')
            self.assertEqual({'Tên xuất hóa đơn'},{r[0] for r in conn.execute("SELECT l.product_name FROM outgoing_invoice_lines l JOIN outgoing_invoice_drafts d ON d.id=l.draft_id WHERE d.status='draft'")})
        result=self.client.post('/api/export/order-invoices',json={'contractor':'NT-A','from':'2026-09-01','to':'2026-09-01'})
        self.assertEqual(200,result.status_code,result.get_json(silent=True))
        with zipfile.ZipFile(io.BytesIO(result.data)) as z:
            for name in z.namelist():
                if name.endswith('.xlsx'):
                    book=load_workbook(io.BytesIO(z.read(name)),data_only=True)
                    self.assertEqual('Tên xuất hóa đơn',book.active['B2'].value);book.close()

    def test_name_update_repairs_editable_draft_without_changing_stock_or_money(self):
        self.seed()
        with server.db() as conn:
            refresh_waiting(conn,'2026-09-13')
            lines=[dict(r) for r in conn.execute('SELECT * FROM outgoing_invoice_lines')]
            stock=[tuple(r) for r in conn.execute('SELECT * FROM inventory_transactions')]
            orders=[tuple(r) for r in conn.execute('SELECT * FROM orders')]
        response=self.client.put('/api/outgoing-product-names/HH-01',json={'invoice_name':'Tên mới đúng hóa đơn'})
        self.assertEqual(200,response.status_code)
        with server.db() as conn:
            after=[dict(r) for r in conn.execute('SELECT * FROM outgoing_invoice_lines')]
            for old,new in zip(lines,after):
                self.assertEqual('Tên mới đúng hóa đơn',new['product_name'])
                self.assertEqual({**old,'product_name':new['product_name']},new)
            self.assertEqual(stock,[tuple(r) for r in conn.execute('SELECT * FROM inventory_transactions')])
            self.assertEqual(orders,[tuple(r) for r in conn.execute('SELECT * FROM orders')])

    def test_signed_saved_and_inflight_drafts_keep_their_snapshots(self):
        self.seed()
        with server.db() as conn:
            refresh_waiting(conn,'2026-09-13')
            did=conn.execute("SELECT id FROM outgoing_invoice_drafts WHERE status='draft'").fetchone()[0]
            conn.execute("INSERT OR REPLACE INTO outgoing_product_names VALUES('HH-01','Tên mới','2026-09-13')")
            for status,remote in [('issued','not_sent'),('cancelled','not_sent'),('draft','saved'),('draft','saving'),('draft','unknown')]:
                conn.execute('UPDATE outgoing_invoice_drafts SET status=?,minvoice_status=? WHERE id=?',(status,remote,did))
                before=conn.serialize();self.assertEqual(0,refresh_editable_names(conn,'2026-09-13'));self.assertEqual(before,conn.serialize())

    def test_full_import_rejects_unit_change_of_used_product(self):
        self.seed()
        p=self.preview([['HH-01','Hàng nội bộ','Tên hóa đơn','Gói','KKKNT']])
        self.assertFalse(p['can_confirm']);self.assertEqual(1,p['counts']['error'])
        self.assertIn('không đổi',p['rows'][0]['errors'][0])

    def test_names_and_new_import_preserves_existing_unit_tax_and_adds_new_code(self):
        self.seed()
        with server.db() as conn:before=dict(conn.execute("SELECT * FROM products WHERE code='HH-01'").fetchone())
        p=self.preview([['HH-01','Tên nội bộ khác','Tên hóa đơn đúng','Gói','KKKNT'],['NEW-CAT','Mã mới','Tên mới xuất','Kg','8%']],mode='names_and_new')
        self.assertTrue(p['can_confirm'],p)
        response=self.client.post('/api/catalog/import/confirm',json={'token':p['token'],'confirmed':True})
        self.assertEqual(200,response.status_code,response.get_json())
        with server.db() as conn:
            after=dict(conn.execute("SELECT * FROM products WHERE code='HH-01'").fetchone())
            for field in ('name','unit','tax','buy_price','supplier'):self.assertEqual(before[field],after[field])
            self.assertEqual('Tên hóa đơn đúng',conn.execute("SELECT invoice_name FROM outgoing_product_names WHERE product_code='HH-01'").fetchone()[0])
            self.assertIsNotNone(conn.execute("SELECT * FROM products WHERE code='NEW-CAT'").fetchone())

    def test_stray_tax_footer_is_reported_but_incomplete_product_still_blocks(self):
        p=self.preview([['NEW-X','Tên mới','Tên xuất','Kg','8%'],['','','','',4.55]])
        self.assertTrue(p['can_confirm']);self.assertEqual([3],p['ignored_non_product_rows'])
        p=self.preview([['','Thiếu mã hàng','Tên xuất','Kg','8%']])
        self.assertFalse(p['can_confirm']);self.assertIn('Thiếu mã hàng',p['rows'][0]['errors'])

    def test_new_only_import_ignores_old_unit_changes_and_replay_keeps_existing_data(self):
        self.seed()
        rows=[['HH-01','Tên cũ trong file','Tên hóa đơn cũ','Gói','KKKNT'],
              ['NEW-ONLY-CAT','Vỏ đậu hũ','Vỏ đậu hũ','Gói','8%']]
        protected=('products','outgoing_product_names','outgoing_product_units','orders',
                   'inventory_transactions','outgoing_invoice_lines')
        with server.db() as conn:
            before={t:[tuple(r) for r in conn.execute('SELECT * FROM '+t)] for t in protected}
        preview=self.preview(rows,mode='new_only')
        self.assertTrue(preview['can_confirm']);self.assertEqual(1,preview['counts']['new_products'])
        self.assertEqual(0,preview['counts']['error'])
        response=self.client.post('/api/catalog/import/confirm',json={'token':preview['token'],'confirmed':True})
        self.assertEqual(200,response.status_code,response.get_json())
        self.assertEqual(1,response.get_json()['inserted_products'])
        with server.db() as conn:
            self.assertEqual(('Vỏ đậu hũ','Gói','0.08'),tuple(conn.execute("SELECT name,unit,tax FROM products WHERE code='NEW-ONLY-CAT'").fetchone()))
            for table,values in before.items():
                after=[tuple(r) for r in conn.execute('SELECT * FROM '+table)]
                self.assertEqual(values,[r for r in after if r[0]!='NEW-ONLY-CAT'],table)
            first={t:[tuple(r) for r in conn.execute('SELECT * FROM '+t)] for t in protected}
        replay=self.preview(rows,mode='new_only')
        self.assertEqual(0,replay['counts']['new_products'])
        response=self.client.post('/api/catalog/import/confirm',json={'token':replay['token'],'confirmed':True})
        self.assertEqual(200,response.status_code);self.assertEqual(0,response.get_json()['processed'])
        with server.db() as conn:
            self.assertEqual(first,{t:[tuple(r) for r in conn.execute('SELECT * FROM '+t)] for t in protected})

    def test_new_only_import_still_blocks_conflicting_new_product_rows(self):
        preview=self.preview([['NEW-ONLY-CONFLICT','Mặt hàng A','A','Gói','8%'],
                              ['NEW-ONLY-CONFLICT','Mặt hàng B','B','Gói','8%']],mode='new_only')
        self.assertFalse(preview['can_confirm']);self.assertEqual(2,preview['counts']['error'])
        response=self.client.post('/api/catalog/import/confirm',json={'token':preview['token'],'confirmed':True})
        self.assertEqual(400,response.status_code)
        with server.db() as conn:
            self.assertIsNone(conn.execute("SELECT 1 FROM products WHERE code='NEW-ONLY-CONFLICT'").fetchone())

    def test_reference_added_after_preview_blocks_unit_change_and_rolls_back(self):
        p=self.preview([['NEW-FIRST','New','New','Kg','8%'],['HH-01','Hàng hóa 01','Tên mới','Gói','0%']])
        self.assertTrue(p['can_confirm'])
        self.seed()
        with server.db() as conn:before=conn.serialize()
        result=self.client.post('/api/catalog/import/confirm',json={'token':p['token'],'confirmed':True})
        self.assertEqual(409,result.status_code)
        with server.db() as conn:self.assertEqual(before,conn.serialize())

    def test_duplicate_rows_in_names_mode_use_source_values(self):
        self.seed()
        row=['HH-01','Tên khác','Tên hóa đơn','Gói','KKKNT']
        p=self.preview([row,row],mode='names_and_new')
        self.assertTrue(p['can_confirm']);self.assertEqual(1,p['counts']['duplicate'])

    def test_new_unit_headers_select_catalog_and_never_price_sheet(self):
        book=self.book([['HH-01','Tên nội bộ','Tên hóa đơn','Gói','KKKNT']])
        sheet=book.active;sheet['D1']='ĐVT \nthành đạt phát';sheet['F1']='ĐVT Xuất HĐ';sheet['F2']='Kg'
        price=book.create_sheet('BÁO GIÁ');price.append(['Mã hàng','Tên hàng','ĐVT','Thuế']);price.append(['BAD','Wrong sheet','Kg','8%'])
        with server.db() as conn:
            full=parse_catalog_workbook(conn,book)
            self.assertEqual('danh mục hh',full['sheet']);self.assertTrue(full['can_confirm'])
            self.assertEqual('Kg',full['rows'][0]['invoice_unit'])
            names=parse_catalog_workbook(conn,book,'names_and_new');self.assertTrue(names['can_confirm'])
            self.assertTrue(any('chờ quy đổi' in w for w in names['rows'][0]['warnings']))
        book.close()

    def test_malformed_explicit_catalog_does_not_fall_back_to_price_sheet(self):
        book=self.book([]);book.active['D1']='Unknown unit header'
        sheet=book.create_sheet('BÁO GIÁ');sheet.append(['Mã hàng','Tên hàng','ĐVT','Thuế']);sheet.append(['BAD','Wrong sheet','Kg','8%'])
        with server.db() as conn:
            with self.assertRaises(ValueError):parse_catalog_workbook(conn,book)
        book.close()

    def test_unissued_template_uses_invoice_label_but_keeps_internal_name_for_reconciliation(self):
        self.seed()
        with server.db() as conn:conn.execute("INSERT OR REPLACE INTO outgoing_product_names VALUES('HH-01','Tên xuất đúng','2026-09-13')")
        payload=self.client.get('/api/outgoing-invoices/unissued?to=2026-09-13&contractor=NT-A').get_json()
        self.assertEqual('Tên xuất đúng',payload['details'][0]['invoice_name'])
        result=self.client.get('/api/outgoing-invoices/unissued-template.zip?to=2026-09-13&contractor=NT-A&portion=all')
        self.assertEqual(200,result.status_code)
        with zipfile.ZipFile(io.BytesIO(result.data)) as z:
            name=next(n for n in z.namelist() if n.startswith('CHUA_XUAT_'))
            book=load_workbook(io.BytesIO(z.read(name)),data_only=True)
            self.assertEqual('Tên xuất đúng',book.active['B2'].value);book.close()


if __name__=='__main__':unittest.main()
