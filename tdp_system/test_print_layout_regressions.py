from __future__ import annotations

import unittest

from openpyxl import Workbook

try:
    from .contract_modules import style_export_sheet
    from .server import autosize, set_title
except ImportError:  # pragma: no cover - direct file invocation
    from contract_modules import style_export_sheet
    from server import autosize, set_title


class PrintLayoutRegressionTests(unittest.TestCase):
    def test_short_summary_is_portrait_centered_but_wide_detail_is_landscape(self):
        short = Workbook().active
        short.append(["Đối tượng", "Đầu kỳ", "Phát sinh", "Điều chỉnh", "Đã thu", "Cuối kỳ"])
        short.append(["A", 1, 2, 0, 1, 2])
        style_export_sheet(short, "CÔNG NỢ PHẢI THU", "Kỳ kiểm tra", [
            "Đối tượng", "Đầu kỳ", "Phát sinh", "Điều chỉnh", "Đã thu", "Cuối kỳ",
        ])
        self.assertEqual("portrait", short.page_setup.orientation)
        self.assertTrue(short.print_options.horizontalCentered)

        wide = Workbook().active
        headers = [f"Cột {index}" for index in range(1, 10)]
        wide.append(headers)
        style_export_sheet(wide, "CHI TIẾT", "Kỳ kiểm tra", headers)
        self.assertEqual("landscape", wide.page_setup.orientation)

    def test_merged_title_does_not_expand_stt_column(self):
        sheet = Workbook().active
        set_title(sheet, "BẢNG KÊ HÀNG HÓA ĐẦU RA", "Một tiêu đề phụ rất dài", 4)
        sheet.append([])
        sheet.append(["STT", "Bếp", "Mã hàng", "Tên hàng xuất hóa đơn"])
        sheet.append([1, "Bếp A", "HH-01", "Cà rốt loại một"])
        autosize(sheet, min_row=4)
        self.assertLess(sheet.column_dimensions["A"].width, sheet.column_dimensions["D"].width)


if __name__ == "__main__":
    unittest.main()
