import io
import unittest
import zipfile
from openpyxl import Workbook, load_workbook
from .outgoing_unissued_export import _contractor_workbook

class TestUnissuedReconciliationNotice(unittest.TestCase):
    row = ['ITEM', 'Test', 'kg', 1, 100, 100, 0, 0, 100, 8, 8, 108, 1]

    def render(self, warnings=None, complete=True):
        source = Workbook()
        sheet = source.active
        sheet.append(['header'] * 13)
        sheet.append(self.row)
        self.assertEqual(sheet['J2'].value, 8)
        buf=io.BytesIO(); source.save(buf); source.close()
        restored=load_workbook(io.BytesIO(buf.getvalue()))
        self.assertEqual(restored.active['J2'].value,8); restored.close()
        archive=io.BytesIO()
        with zipfile.ZipFile(archive,'w') as z: z.writestr('test.xlsx',buf.getvalue())
        detail={'order_id':1,'work_date':'2026-09-01','contractor':'TEST','product_code':'ITEM','product_name':'Test','unit':'kg','unissued_qty':1,'unit_price':100,'tax':8,'needs_conversion':False}
        payload={'from':'2026-09-01','asof':'2026-09-15','warnings':warnings or [],'reconciliation_complete':complete,'details':[detail]}
        with zipfile.ZipFile(io.BytesIO(archive.getvalue())) as z:
            result=_contractor_workbook(z,['test.xlsx'],'TEST',[detail],payload,lambda v:8)
        return load_workbook(io.BytesIO(result))

    def text(self,sheet):
        return '\n'.join(str(c.value or '') for row in sheet for c in row)

    def test_scoped_warning_in_all_three_places(self):
        w=self.render([{'contractor':'TEST','message':'NEED_REVIEW'}],False)
        for name in ['Chua xuat','Tong hop','Canh bao doi chieu']:
            self.assertIn('CHƯA ĐỐI CHIẾU XONG',self.text(w[name]))
        self.assertIn('NEED_REVIEW',self.text(w['Canh bao doi chieu']));w.close()

    def test_global_included_other_customer_excluded(self):
        w=self.render([{'contractor':'','message':'GLOBAL_NOTICE'},{'contractor':'OTHER','message':'PRIVATE_OTHER'}],False)
        text='\n'.join(self.text(s) for s in w)
        self.assertIn('GLOBAL_NOTICE',text);self.assertNotIn('PRIVATE_OTHER',text);w.close()

    def test_no_warning_does_not_claim_money_verified(self):
        w=self.render()
        self.assertIn('chưa xác nhận đối chiếu tiền hóa đơn',self.text(w['Tong hop']));w.close()

    def test_incomplete_without_explanation(self):
        w=self.render([],False)
        self.assertIn('CHƯA ĐỐI CHIẾU XONG',self.text(w['Tong hop']));w.close()

    def test_data_and_money_unchanged(self):
        w=self.render([{'contractor':'TEST','message':'REVIEW'}],False);s=w['Chua xuat']
        self.assertEqual([c.value for c in s[2]],self.row)
        self.assertEqual((s['I4'].value,s['K4'].value,s['L4'].value),(100,8,108))
        self.assertEqual(s.max_column,13);w.close()

    def test_warning_formula_text_is_escaped(self):
        w=self.render([{'contractor':'TEST','code':'=1+1','message':'@FORMULA'}],False)
        s=w['Canh bao doi chieu'];values=[c.value for row in s for c in row]
        self.assertIn("'=1+1",values);self.assertIn("'@FORMULA",values)
        self.assertFalse(any(c.data_type=='f' for row in s for c in row));w.close()

if __name__=='__main__':unittest.main()
