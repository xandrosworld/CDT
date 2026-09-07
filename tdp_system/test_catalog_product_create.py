"""Manual product creation never overwrites catalog or inventory data."""
import sqlite3
import unittest
from contextlib import contextmanager
from pathlib import Path
from unittest.mock import patch

from flask import Flask

from .contract_modules import register_contract_routes, init_contract_schema
from .invoice_workbench import init_invoice_workbench_schema
from .test_invoice_input_sync import now_iso
from .server import SCHEMA
from .test_invoice_workbench_listing import seed_round1


class CatalogProductCreateTests(unittest.TestCase):
    def setUp(self):
        self.conn = sqlite3.connect(':memory:')
        self.conn.row_factory = sqlite3.Row
        self.conn.execute('PRAGMA foreign_keys=ON')
        self.conn.executescript(SCHEMA)
        init_contract_schema(self.conn)
        init_invoice_workbench_schema(self.conn)
        seed_round1(self.conn)
        self.conn.commit()

        @contextmanager
        def db():
            with self.conn:
                yield self.conn

        app = Flask(__name__)
        register_contract_routes(app, dict(db=db, now_iso=now_iso,
            clean_text=lambda value: str(value or '').strip(), number_value=lambda value: float(value or 0),
            tax_factor=lambda value: 1, setting_get=lambda c,k,d='': d, setting_set=lambda *a: None,
            root=Path.cwd(), data_dir=Path.cwd()))
        self.client = app.test_client()
        self.data = {'code': ' h000003 ', 'name': 'Đậu phụ chiên', 'unit': 'Cái', 'tax': '8%'}

    def tearDown(self):
        self.conn.close()

    def post(self, data=None):
        return self.client.post('/api/catalog/products', json=self.data if data is None else data)

    def test_creates_normalized_product_without_changing_existing_records_or_stock(self):
        before = {name: [tuple(row) for row in self.conn.execute('SELECT * FROM ' + name)] for name in
                  ['products', 'product_prices', 'invoice_inventory_ledger', 'msmi_invoice_items']}
        response = self.post()
        self.assertEqual(201, response.status_code, response.json)
        product = response.json['product']
        self.assertEqual({'code': 'H000003', 'name': 'Đậu phụ chiên', 'unit': 'Cái', 'tax': '0.08'}, product)
        saved = dict(self.conn.execute("SELECT * FROM products WHERE code='H000003'").fetchone())
        self.assertEqual((0, '', 0), (saved['buy_price'], saved['supplier'], saved['purchase_list']))
        self.assertTrue(saved['catalog_updated_at'])
        for name, rows in before.items():
            query = "SELECT * FROM products WHERE code!='H000003'" if name == 'products' else 'SELECT * FROM ' + name
            self.assertEqual(rows, [tuple(row) for row in self.conn.execute(query)])

    def test_duplicate_or_retry_never_overwrites_product(self):
        self.assertEqual(201, self.post().status_code)
        before = '\n'.join(self.conn.iterdump())
        self.assertEqual(409, self.post({**self.data, 'name': 'Tên khác', 'unit': 'Kg'}).status_code)
        self.assertEqual(before, '\n'.join(self.conn.iterdump()))
        self.assertEqual(409, self.post({**self.data, 'code': 'r1-kg'}).status_code)

    def test_rejects_missing_invalid_and_oversized_input_without_writes(self):
        bad = [[], {'code': 'NEW'}, *[{**self.data, key: ''} for key in self.data],
               {**self.data, 'tax': 'NaN'}, {**self.data, 'tax': '7%'}, {**self.data, 'tax': True},
               {**self.data, 'code': 'A B'}, {**self.data, 'code': '../A'},
               {**self.data, 'name': ['invalid']}, {**self.data, 'unit': 'x' * 51}]
        before = '\n'.join(self.conn.iterdump())
        for data in bad:
            with self.subTest(data=data):
                self.assertEqual(400, self.post(data).status_code)
        self.assertEqual(before, '\n'.join(self.conn.iterdump()))

    def test_all_supported_tax_choices_are_saved(self):
        for index, tax in enumerate(['KKKNT', 'KCT', '0', '0.05', '0.08', '0.1']):
            response = self.post({**self.data, 'code': 'TAX-' + str(index), 'tax': tax})
            self.assertEqual(201, response.status_code, response.json)
            self.assertEqual(tax, response.json['product']['tax'])

    def edit(self, row, **changes):
        expected = {key: row[key] for key in ('name', 'unit', 'tax', 'invoice_name', 'catalog_updated_at')}
        return self.client.put('/api/catalog/products', json={**row, **changes, 'expected': expected})

    def test_unified_list_edit_preserves_names_and_rejects_stale_edits(self):
        self.assertEqual(201, self.post({**self.data, 'invoice_name': 'Tên hóa đơn riêng'}).status_code)
        row = self.client.get('/api/catalog/products?q=H000003').json['items'][0]
        self.assertEqual('Tên hóa đơn riêng', row['invoice_name'])
        self.assertEqual(200, self.edit(row, name='Đậu chiên mới', invoice_name='Tên đầu ra mới').status_code)
        self.assertEqual(409, self.edit(row, name='Ghi đè dữ liệu cũ').status_code)
        updated = self.client.get('/api/catalog/products?q=H000003').json['items'][0]
        self.assertEqual('Tên đầu ra mới', updated['invoice_name'])
        self.assertEqual(200, self.edit(updated, invoice_name='').status_code)
        self.assertEqual('', self.client.get('/api/catalog/products?q=H000003').json['items'][0]['invoice_name'])

    def test_unit_change_of_used_code_is_blocked_without_partial_name_changes(self):
        self.conn.execute("UPDATE products SET tax='0.08' WHERE code='R1-KG'")
        self.conn.commit()
        row = self.client.get('/api/catalog/products?q=R1-KG').json['items'][0]
        before = '\n'.join(self.conn.iterdump())
        self.assertEqual(409, self.edit(row, name='Không đổi', unit='Cái', invoice_name='Không đổi').status_code)
        self.assertEqual(before, '\n'.join(self.conn.iterdump()))
        self.assertEqual(200, self.edit(row, name='Tên hàng mới', invoice_name='Tên trên hóa đơn mới').status_code)

    def test_excel_second_method_preserves_manual_codes_and_existing_invoice_names(self):
        from io import BytesIO
        from openpyxl import Workbook
        self.post({**self.data, 'invoice_name': 'Tên đã lưu'})
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'BÁO GIÁ'
        for col, label in {3:'Mã SP',4:'Tên hàng',10:'ĐVT',11:'Thuế'}.items():
            sheet.cell(2, col, label)
        for index, code in enumerate(['H000003','EXCEL-NEW'], start=4):
            for col, value in {3:code,4:'Tên hàng Excel',10:'Cái',11:'8%'}.items():
                sheet.cell(index,col,value)
        output = BytesIO(); workbook.save(output); output.seek(0)
        preview = self.client.post('/api/catalog/import/preview', data={'file':(output,'Em Thành.xlsx')}, content_type='multipart/form-data')
        self.assertEqual(200, preview.status_code, preview.json)
        confirmed = self.client.post('/api/catalog/import/confirm', json={'token':preview.json['token'], 'confirmed':True})
        self.assertEqual(200, confirmed.status_code, confirmed.json)
        self.assertEqual('Tên đã lưu', self.client.get('/api/catalog/products?q=H000003').json['items'][0]['invoice_name'])
        self.assertEqual('EXCEL-NEW', self.client.get('/api/catalog/products?q=EXCEL-NEW').json['items'][0]['code'])
        self.assertIsNotNone(self.conn.execute("SELECT code FROM products WHERE code='R1-KG'").fetchone())

    def test_audit_failure_rolls_back_product_creation(self):
        with self.assertLogs(level='ERROR'), patch('tdp_system.contract_modules.audit', side_effect=RuntimeError('test failure')):
            response = self.post()
        self.assertEqual(500, response.status_code)
        self.assertIsNone(self.conn.execute("SELECT code FROM products WHERE code='H000003'").fetchone())

    def test_missing_server_workbook_does_not_report_success(self):
        import tempfile
        from . import server
        with tempfile.TemporaryDirectory() as directory, patch.object(server, 'MASTER_SOURCE', Path(directory) / 'missing.xlsx'), patch.object(server, 'sync_master_if_needed') as sync:
            response = server.app.test_client().post('/api/master/sync')
        self.assertEqual(404, response.status_code)
        sync.assert_not_called()
