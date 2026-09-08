import io
import zipfile
import tempfile
import unittest
from unittest.mock import patch
from concurrent.futures import ThreadPoolExecutor
from threading import Barrier
from pathlib import Path

from openpyxl import load_workbook

try:
    from . import server
    from .outgoing_readiness import canonical_available_stock
except ImportError:  # pragma: no cover - direct invocation
    import server
    from outgoing_readiness import canonical_available_stock


class OutgoingReadinessTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp_dir = tempfile.TemporaryDirectory()
        cls.original_db_path = server.DB_PATH
        cls.original_data_dir = server.DATA_DIR
        cls.original_testing = server.app.config.get("TESTING", False)
        server.DATA_DIR = Path(cls.temp_dir.name) / "data"
        server.DATA_DIR.mkdir(parents=True, exist_ok=True)
        server.DB_PATH = server.DATA_DIR / "outgoing_readiness.sqlite3"
        server.app.config["TESTING"] = True
        server.init_database()
        cls.client = server.app.test_client()

    @classmethod
    def tearDownClass(cls):
        server.DB_PATH = cls.original_db_path
        server.DATA_DIR = cls.original_data_dir
        server.app.config["TESTING"] = cls.original_testing
        cls.temp_dir.cleanup()

    def setUp(self):
        with server.db() as conn:
            conn.execute("DELETE FROM invoice_inventory_ledger")
            conn.execute("DELETE FROM invoice_inventory_confirmations")
            conn.execute("DELETE FROM outgoing_source_invoice_items")
            conn.execute("DELETE FROM outgoing_source_events")
            conn.execute("DELETE FROM outgoing_source_invoices")
            conn.execute("DELETE FROM outgoing_invoice_lines")
            conn.execute("DELETE FROM outgoing_invoice_drafts")
            conn.execute("DELETE FROM inventory_transactions")
            conn.execute("DELETE FROM orders")
            conn.execute("DELETE FROM batches")
            conn.execute("DELETE FROM audit_log")
            conn.execute(
                "INSERT OR REPLACE INTO contractors(code,name,price_group,pricing_mode) "
                "VALUES('NT-A','Nhà thầu A','NT-A','group')"
            )
            conn.execute(
                "INSERT OR REPLACE INTO contractors(code,name,price_group,pricing_mode) "
                "VALUES('NT-B','Nhà thầu B','NT-B','group')"
            )
            conn.execute(
                """INSERT OR REPLACE INTO products(
                       code,name,unit,tax,supplier,buy_price,purchase_list,seller,cccd
                   ) VALUES('HH-01','Hàng hóa 01','kg','0%','NCC-A',10,0,'','')"""
            )

    @staticmethod
    def add_batch(conn, work_date, rows):
        now = server.now_iso()
        batch_id = conn.execute(
            """INSERT INTO batches(work_date,source_name,status,created_at,approved_at)
               VALUES(?,?,'approved',?,?)""",
            (work_date, f"don-{work_date}.xlsx", now, now),
        ).lastrowid
        order_ids = []
        for index, row in enumerate(rows, start=1):
            qty = row.get("qty", 1)
            order_ids.append(conn.execute(
                """INSERT INTO orders(
                       batch_id,work_date,contractor,kitchen,product_code,product_name,qty,
                       actual_received,actual_delivered,unit,supplier,buy_price,sell_price,tax,
                       purchase_list,errors,warnings,updated_at
                   ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,0,'[]','[]',?)""",
                (
                    batch_id, work_date, row.get("contractor", "NT-A"), row.get("kitchen", f"BEP-{index}"),
                    row.get("product_code", "HH-01"), "Hàng hóa 01", qty, qty, qty, "kg", "NCC-A",
                    10, row.get("sell_price", 20), "0%", now,
                ),
            ).lastrowid)
        return batch_id, order_ids

    @staticmethod
    def add_opening(conn, qty, work_date="2026-08-01", product_code="HH-01"):
        now = server.now_iso()
        conn.execute(
            """INSERT INTO inventory_transactions(
                   txn_date,product_code,qty_in,qty_out,unit_cost,source_type,source_id,
                   source_line,status,note,created_at,updated_at
               ) VALUES(?,?,?,0,10,'OPENING',?,?,'posted','opening',?,?)""",
            (work_date, product_code, qty, "OPEN-TEST", product_code, now, now),
        )

    @staticmethod
    def add_canonical_event(
        conn, qty_delta, direction, key, work_date="2026-08-15",
        source_table="test_source", source_id=None, product_code="HH-01",
    ):
        now = server.now_iso()
        safe_source_id = len(key) if source_id is None else source_id
        confirmation_id = conn.execute(
            """INSERT INTO invoice_inventory_confirmations(
                   confirmation_key,direction,source_invoice_table,source_invoice_id,
                   action,confirmed,note,created_at
               ) VALUES(?,?,?,?, 'post',1,'test',?)""",
            (f"confirm-{key}", direction, source_table, safe_source_id, now),
        ).lastrowid
        conn.execute(
            """INSERT INTO invoice_inventory_ledger(
                   event_key,direction,event_type,source_invoice_table,source_invoice_id,
                   source_line_id,source_line_index,product_code,txn_date,qty_delta,unit_cost,
                   mapping_revision_id,confirmation_id,reverses_event_key,status,created_at
               ) VALUES(?,?,'POST',?,?,?,?,?,?,?,10,NULL,?,'','posted',?)""",
            (
                key, direction, source_table, safe_source_id, safe_source_id * 10, 1,
                product_code, work_date, qty_delta, confirmation_id, now,
            ),
        )

    @classmethod
    def add_posted_source(
        cls, conn, *, source, number, series="1C26TDP", invoice_date="2026-08-20",
        product_code="HH-01", qty=4,
    ):
        now = server.now_iso()
        source_id = conn.execute(
            """INSERT INTO outgoing_source_invoices(
                   tenant,source,identity_key,invoice_number,invoice_series,invoice_date,
                   source_status_class,sync_status,stock_status,synced_at,created_at,updated_at
               ) VALUES('default',?,?,?,?,?,'issued','synced','posted',?,?,?)""",
            (
                source, f"{source}-{series}-{number}-{invoice_date}", number, series,
                invoice_date, now, now, now,
            ),
        ).lastrowid
        cls.add_canonical_event(
            conn, -qty, "output", f"{source}-POST-{source_id}", invoice_date,
            source_table="outgoing_source_invoices", source_id=source_id,
            product_code=product_code,
        )
        return source_id

    @staticmethod
    def add_local_issued_draft(
        conn, batch_id, order_id, *, number, series="1C26TDP",
        invoice_date="2026-08-20", product_code="HH-01", qty=4, unit="kg",
    ):
        now = server.now_iso()
        draft_id = conn.execute(
            """INSERT INTO outgoing_invoice_drafts(
                   batch_id,contractor,invoice_date,status,subtotal,tax_amount,total_amount,
                   created_at,issued_at,round_no,issued_invoice_number,
                   issued_invoice_series,issued_invoice_date
               ) VALUES(?,'NT-A',?,'issued',80,0,80,?,?,1,?,?,?)""",
            (batch_id, invoice_date, now, now, number, series, invoice_date),
        ).lastrowid
        conn.execute(
            """INSERT INTO outgoing_invoice_lines(
                   draft_id,order_id,product_code,product_name,qty,unit,unit_price,tax,
                   invoice_nature,amount
               ) VALUES(?,?,?,'Hàng hóa',?,?,20,'0%','1',80)""",
            (draft_id, order_id, product_code, qty, unit),
        )
        return draft_id

    def test_two_rounds_keep_shortage_and_are_idempotent(self):
        with server.db() as conn:
            self.add_opening(conn, 7)
            batch_id, _ = self.add_batch(conn, "2026-08-20", [{"qty": 10}])

        first_readiness = self.client.get(f"/api/outgoing-invoices/readiness/{batch_id}").get_json()
        self.assertEqual((first_readiness["invoiceable_qty"], first_readiness["pending_qty"]), (7, 3))
        first = self.client.post(f"/api/outgoing-invoices/draft/{batch_id}")
        self.assertEqual(first.status_code, 200, first.get_data(as_text=True))
        self.assertEqual((len(first.get_json()["drafts"]), first.get_json()["pending_qty"]), (1, 3))

        with server.db() as conn:
            draft_id = conn.execute(
                "SELECT id FROM outgoing_invoice_drafts WHERE batch_id=?", (batch_id,),
            ).fetchone()["id"]
            conn.execute(
                """UPDATE outgoing_invoice_drafts
                      SET status='issued',issued_invoice_series='1C26TDP',
                          issued_invoice_number='0000001',issued_invoice_date='2026-08-20'
                    WHERE id=?""",
                (draft_id,),
            )
            conn.execute(
                """UPDATE inventory_transactions SET status='posted'
                    WHERE source_type='OUTGOING_DRAFT' AND source_id=?""",
                (str(draft_id),),
            )
            self.add_canonical_event(conn, 3, "input", "INPUT-3")

        between = self.client.get(f"/api/outgoing-invoices/readiness/{batch_id}").get_json()
        self.assertEqual((between["issued_qty"], between["invoiceable_qty"], between["pending_qty"]), (7, 3, 0))
        second = self.client.post(f"/api/outgoing-invoices/draft/{batch_id}")
        self.assertEqual(second.status_code, 200, second.get_data(as_text=True))
        self.assertEqual(second.get_json()["drafts"][0]["round_no"], 2)
        replay = self.client.post(f"/api/outgoing-invoices/draft/{batch_id}")
        self.assertEqual(replay.status_code, 200, replay.get_data(as_text=True))
        with server.db() as conn:
            drafts = conn.execute(
                "SELECT round_no,status FROM outgoing_invoice_drafts WHERE batch_id=? ORDER BY round_no",
                (batch_id,),
            ).fetchall()
            self.assertEqual([(row["round_no"], row["status"]) for row in drafts], [(1, "issued"), (2, "draft")])
            qty = conn.execute(
                """SELECT COALESCE(SUM(qty_out),0) qty FROM inventory_transactions
                    WHERE source_type='OUTGOING_DRAFT' AND status='reserved'"""
            ).fetchone()["qty"]
            self.assertEqual(qty, 3)

    def test_tax_files_have_separate_drafts_and_independent_issue_confirmations(self):
        with server.db() as conn:
            self.add_opening(conn, 20)
            batch_id, order_ids = self.add_batch(conn, '2026-08-20', [{'qty': 4}, {'qty': 6}])
            conn.execute("UPDATE orders SET tax='KKKNT' WHERE id=?", (order_ids[0],))
            conn.execute("UPDATE orders SET tax='8%' WHERE id=?", (order_ids[1],))
        response = self.client.post(f'/api/outgoing-invoices/draft/{batch_id}')
        self.assertEqual(response.status_code, 200, response.json)
        drafts = response.json['drafts']
        self.assertEqual(len(drafts), 2, 'One separately uploaded tax file must have its own invoice number')
        exported = self.client.get(f'/api/export/invoices/{batch_id}')
        self.assertEqual(exported.status_code, 200, exported.get_data()[:200])
        with zipfile.ZipFile(io.BytesIO(exported.data)) as archive:
            files = [n for n in archive.namelist() if n.endswith('.xlsx')]
            self.assertEqual(len(files), len(drafts))
            for draft in drafts:
                workbook = load_workbook(io.BytesIO(archive.read(next(n for n in files if f"lan_{draft['round_no']}_" in n))), data_only=True)
                self.assertEqual(sum(row[8] or 0 for row in list(workbook.active.values)[1:]), draft['subtotal'])
                self.assertEqual(sum(row[10] or 0 for row in list(workbook.active.values)[1:]), draft['tax_amount'])
                workbook.close()
        replay = self.client.post(f'/api/outgoing-invoices/draft/{batch_id}').json['drafts']
        self.assertEqual({d['id'] for d in replay}, {d['id'] for d in drafts})
        with server.db() as conn:
            conn.execute("UPDATE outgoing_invoice_drafts SET buyer_name_snapshot='Buyer',buyer_tax_code_snapshot='0200000001',"
                         "buyer_address_snapshot='Address',company_name_snapshot='TDP',company_tax_code_snapshot='0100000001',"
                         "company_address_snapshot='Address',payment_requester_snapshot='Requester',"
                         "payment_bank_name_snapshot='Bank',payment_bank_account_snapshot='Account' WHERE batch_id=?", (batch_id,))
        for index, draft in enumerate(drafts):
            issued = self.client.post(f"/api/outgoing-invoices/{draft['id']}/confirm-issued", json={
                'confirmed': True, 'invoice_number': str(900+index), 'invoice_series': '1C26TDP', 'invoice_date': '2026-08-20'})
            self.assertEqual(issued.status_code, 200, issued.json)
            repeated = self.client.post(f"/api/outgoing-invoices/{draft['id']}/confirm-issued", json={
                'confirmed': True, 'invoice_number': str(900+index), 'invoice_series': '1C26TDP', 'invoice_date': '2026-08-20'})
            self.assertTrue(repeated.json['idempotent'])
            changed_number = self.client.post(f"/api/outgoing-invoices/{draft['id']}/confirm-issued", json={
                'confirmed': True, 'invoice_number': '999', 'invoice_series': '1C26TDP', 'invoice_date': '2026-08-20'})
            self.assertEqual(changed_number.status_code, 409)
            if index == 0:
                remaining = self.client.get(f'/api/export/invoices/{batch_id}')
                self.assertEqual(remaining.status_code, 200)
                with zipfile.ZipFile(io.BytesIO(remaining.data)) as archive:
                    self.assertEqual(len([n for n in archive.namelist() if n.endswith('.xlsx')]), 1)
        with server.db() as conn:
            self.assertEqual(conn.execute('SELECT count(*) FROM invoice_inventory_ledger').fetchone()[0], 0)
            self.assertEqual(conn.execute("SELECT SUM(qty_out) FROM inventory_transactions WHERE source_type='OUTGOING_DRAFT' AND status='posted'").fetchone()[0], 10)

    def test_old_mixed_tax_draft_must_be_split_before_file_or_local_issue(self):
        with server.db() as conn:
            self.add_opening(conn, 20)
            batch_id, orders = self.add_batch(conn, '2026-08-20', [{'qty': 4}, {'qty': 6}])
            conn.execute("UPDATE orders SET tax='8%' WHERE id=?", (orders[1],))
        drafts = self.client.post(f'/api/outgoing-invoices/draft/{batch_id}').json['drafts']
        first, second = [r['id'] for r in drafts]
        with server.db() as conn:
            conn.execute('UPDATE outgoing_invoice_lines SET draft_id=? WHERE draft_id=?', (first, second))
            conn.execute("UPDATE inventory_transactions SET source_id=? WHERE source_type='OUTGOING_DRAFT' AND source_id=?", (str(first), str(second)))
            conn.execute("UPDATE outgoing_invoice_drafts SET status='cancelled' WHERE id=?", (second,))
        blocked = self.client.get(f'/api/export/invoices/{batch_id}')
        self.assertEqual(blocked.status_code, 409)
        self.assertEqual(blocked.json['code'], 'invoice_tax_split_required')
        blocked = self.client.post(f'/api/outgoing-invoices/{first}/confirm-issued', json={
            'confirmed': True, 'invoice_number': '100', 'invoice_series': '1C26TDP', 'invoice_date': '2026-08-20'})
        self.assertEqual(blocked.json['code'], 'invoice_tax_split_required')
        rebuilt = self.client.post(f'/api/outgoing-invoices/draft/{batch_id}')
        self.assertEqual(rebuilt.status_code, 200, rebuilt.json)
        self.assertEqual(len(rebuilt.json['drafts']), 2)
        self.assertEqual(self.client.get(f'/api/export/invoices/{batch_id}').status_code, 200)
        with server.db() as conn:
            self.assertEqual(conn.execute("SELECT SUM(qty_out) FROM inventory_transactions WHERE source_type='OUTGOING_DRAFT' AND status='reserved'").fetchone()[0], 10)

    def test_selected_batch_invoices_remain_visible_after_200_newer_invoices(self):
        with server.db() as conn:
            self.add_opening(conn, 20)
            old, _ = self.add_batch(conn, '2026-08-20', [{'qty': 4}])
            new, _ = self.add_batch(conn, '2026-08-21', [{'qty': 1}])
        old_id = self.client.post(f'/api/outgoing-invoices/draft/{old}').json['drafts'][0]['id']
        with server.db() as conn:
            for round_no in range(1,202):
                conn.execute("INSERT INTO outgoing_invoice_drafts(batch_id,contractor,invoice_date,status,subtotal,tax_amount,total_amount,created_at,round_no) VALUES(?,'NT-A','2026-08-21','cancelled',0,0,0,?,?)", (new,server.now_iso(),round_no))
        result = self.client.get(f'/api/outgoing-invoices?batch_id={old}')
        self.assertEqual([r['id'] for r in result.json['items']], [old_id])
        self.assertEqual(result.json['items'][0]['tax_label'], '0%')
        self.assertEqual(self.client.get('/api/outgoing-invoices?batch_id=bad').status_code, 400)

    def test_negative_stock_is_explained_before_creating_files(self):
        with server.db() as conn:
            self.add_opening(conn, 0.5)
            self.add_canonical_event(conn,-2,'output','NEGATIVE-OPENING')
            batch_id, _ = self.add_batch(conn,'2026-08-20',[{'qty':4}])
        preview = self.client.get(f'/api/outgoing-invoices/readiness/{batch_id}')
        self.assertEqual(preview.status_code,200)
        self.assertEqual(preview.json['blocking_issues'][0]['product_code'],'HH-01')
        self.assertEqual(preview.json['blocking_issues'][0]['qty'],-1.5)
        self.assertEqual(self.client.post(f'/api/outgoing-invoices/draft/{batch_id}').status_code,409)

    def test_allocation_can_use_later_input_but_actual_backdated_issue_is_blocked(self):
        from .outgoing_readiness import OutgoingReadinessError, validate_issued_draft_stock
        with server.db() as conn:
            self.add_opening(conn, 0)
            self.add_canonical_event(conn, 7, "input", "FUTURE-INPUT", work_date="2026-08-21")
            batch_id, _ = self.add_batch(conn, "2026-08-20", [{"qty": 10}])
        payload = self.client.get(f"/api/outgoing-invoices/readiness/{batch_id}").get_json()
        self.assertEqual((payload["invoiceable_qty"], payload["pending_qty"]), (7, 3))
        response = self.client.post(f"/api/outgoing-invoices/draft/{batch_id}")
        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        draft_id = response.get_json()["drafts"][0]["id"]
        with server.db() as conn:
            with self.assertRaises(OutgoingReadinessError):
                validate_issued_draft_stock(conn, draft_id, "2026-08-20", "1C26TDP", "100")
            validate_issued_draft_stock(conn, draft_id, "2026-08-21", "1C26TDP", "100")

    def test_confirm_rechecks_stock_and_preserves_draft_on_shortage(self):
        with server.db() as conn:
            self.add_opening(conn, 7)
            batch_id, _ = self.add_batch(conn, "2026-08-20", [{"qty": 7}])
        draft_id = self.client.post(f"/api/outgoing-invoices/draft/{batch_id}").get_json()["drafts"][0]["id"]
        with server.db() as conn:
            conn.execute(
                """UPDATE outgoing_invoice_drafts SET buyer_name_snapshot='Buyer',
                   buyer_tax_code_snapshot='0200000001',buyer_address_snapshot='Address',
                   company_name_snapshot='TDP',company_tax_code_snapshot='0100000001',
                   company_address_snapshot='Address',payment_requester_snapshot='Requester',
                   payment_bank_name_snapshot='Bank',payment_bank_account_snapshot='Account'
                   WHERE id=?""", (draft_id,),
            )
            self.add_canonical_event(conn, -5, "output", "STOCK-CHANGED")
        response = self.client.post(f"/api/outgoing-invoices/{draft_id}/confirm-issued", json={
            "confirmed": True, "invoice_number": "100", "invoice_series": "1C26TDP", "invoice_date": "2026-08-20",
        })
        self.assertEqual(response.status_code, 409, response.get_data(as_text=True))
        with server.db() as conn:
            self.assertEqual(conn.execute("SELECT status FROM outgoing_invoice_drafts WHERE id=?", (draft_id,)).fetchone()[0], "draft")
            self.assertEqual(conn.execute("SELECT SUM(qty_out) FROM inventory_transactions WHERE status='reserved'").fetchone()[0], 7)

    def test_failed_second_contractor_rolls_back_first_contractor(self):
        with server.db() as conn:
            self.add_opening(conn, 10)
            batch_id, ids = self.add_batch(conn, "2026-08-20", [{"qty": 3}, {"qty": 3, "contractor": "NT-B"}])
        with patch("tdp_system.contract_modules.invoice_tax_percent", side_effect=[0, ValueError("fixture failure")]):
            response = self.client.post(f"/api/outgoing-invoices/draft/{batch_id}")
        self.assertGreaterEqual(response.status_code, 400)
        with server.db() as conn:
            self.assertEqual(conn.execute("SELECT COUNT(*) FROM outgoing_invoice_drafts").fetchone()[0], 0)
            self.assertEqual(conn.execute("SELECT COUNT(*) FROM inventory_transactions WHERE status='reserved'").fetchone()[0], 0)

    def test_concurrent_batches_cannot_share_the_same_stock(self):
        with server.db() as conn:
            self.add_opening(conn, 7)
            batches = [self.add_batch(conn, "2026-08-20", [{"qty": 7}])[0] for _ in range(2)]
        barrier = Barrier(2)
        def create(batch_id):
            with server.app.test_client() as client:
                barrier.wait(timeout=10)
                response = client.post(f"/api/outgoing-invoices/draft/{batch_id}")
                return response.status_code, response.get_json()
        with ThreadPoolExecutor(max_workers=2) as pool:
            results = list(pool.map(create, batches))
        self.assertEqual([r[0] for r in results], [200, 200])
        self.assertEqual(sum(r[1]["pending_qty"] for r in results), 7)
        with server.db() as conn:
            self.assertEqual(conn.execute("SELECT SUM(qty_out) FROM inventory_transactions WHERE status='reserved'").fetchone()[0], 7)
            self.assertEqual(canonical_available_stock(conn)["HH-01"]["available_qty"], 0)

    def test_export_rechecks_stock_and_missing_hold_without_writing(self):
        with server.db() as conn:
            self.add_opening(conn, 7)
            batch_id, _ = self.add_batch(conn, "2026-08-20", [{"qty": 7}])
        self.client.post(f"/api/outgoing-invoices/draft/{batch_id}")
        with server.db() as conn:
            self.add_canonical_event(conn, -5, "output", "EXPORT-STOCK-CHANGED")
        response = self.client.get(f"/api/export/invoices/{batch_id}")
        self.assertEqual(response.status_code, 409, response.get_data(as_text=True))
        self.assertEqual(response.get_json()["code"], "canonical_stock_overcommitted")
        with server.db() as conn:
            conn.execute("UPDATE inventory_transactions SET status='cancelled' WHERE source_type='OUTGOING_DRAFT'")
        response = self.client.get(f"/api/export/invoices/{batch_id}")
        self.assertEqual(response.status_code, 409)
        self.assertEqual(response.get_json()["code"], "draft_reservation_mismatch")

    def test_confirm_already_synced_invoice_does_not_double_deduct(self):
        from .outgoing_readiness import validate_issued_draft_stock
        with server.db() as conn:
            self.add_opening(conn, 7)
            batch_id, _ = self.add_batch(conn, "2026-08-20", [{"qty": 7}])
        draft_id = self.client.post(f"/api/outgoing-invoices/draft/{batch_id}").get_json()["drafts"][0]["id"]
        with server.db() as conn:
            self.add_posted_source(conn, source="minvoice", number="100", qty=7)
            validate_issued_draft_stock(conn, draft_id, "2026-08-20", "1C26TDP", "100")
            # The validation savepoint is rolled back even on success.
            self.assertEqual(conn.execute("SELECT status FROM outgoing_invoice_drafts WHERE id=?", (draft_id,)).fetchone()[0], "draft")
            self.assertEqual(conn.execute("SELECT COUNT(*) FROM inventory_transactions WHERE status='reserved'").fetchone()[0], 1)

    def test_multiple_contractors_are_separate_and_never_overallocate(self):
        with server.db() as conn:
            self.add_opening(conn, 7)
            batch_id, _ = self.add_batch(conn, "2026-08-20", [
                {"qty": 6, "contractor": "NT-A"},
                {"qty": 6, "contractor": "NT-B"},
            ])
        payload = self.client.get(f"/api/outgoing-invoices/readiness/{batch_id}").get_json()
        self.assertEqual([(row["contractor"], row["invoiceable_qty"], row["pending_qty"])
                          for row in payload["contractors"]], [("NT-A", 6, 0), ("NT-B", 1, 5)])
        response = self.client.post(f"/api/outgoing-invoices/draft/{batch_id}")
        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        self.assertEqual({row["contractor"] for row in response.get_json()["drafts"]}, {"NT-A", "NT-B"})
        with server.db() as conn:
            self.assertEqual(conn.execute(
                "SELECT COUNT(*) n FROM outgoing_invoice_drafts WHERE batch_id=? AND status='draft'",
                (batch_id,),
            ).fetchone()["n"], 2)
            self.assertEqual(conn.execute(
                "SELECT SUM(qty_out) qty FROM inventory_transactions WHERE status='reserved'",
            ).fetchone()["qty"], 7)

    def test_product_code_is_mandatory_and_manual_adjustment_does_not_unlock(self):
        with server.db() as conn:
            self.add_opening(conn, 2)
            batch_id, order_ids = self.add_batch(conn, "2026-08-20", [{"qty": 5}])
            now = server.now_iso()
            conn.execute(
                """INSERT INTO inventory_transactions(
                       txn_date,product_code,qty_in,qty_out,unit_cost,source_type,source_id,
                       source_line,status,note,created_at,updated_at
                   ) VALUES('2026-08-10','HH-01',100,0,10,'ADJUSTMENT','ADJ-1','1',
                            'posted','manual',?,?)""",
                (now, now),
            )
        payload = self.client.get(f"/api/outgoing-invoices/readiness/{batch_id}").get_json()
        self.assertEqual((payload["invoiceable_qty"], payload["pending_qty"]), (2, 3))
        with server.db() as conn:
            conn.execute("UPDATE orders SET product_code='' WHERE id=?", (order_ids[0],))
        blocked = self.client.get(f"/api/outgoing-invoices/readiness/{batch_id}")
        self.assertEqual(blocked.status_code, 409)
        self.assertEqual(blocked.get_json()["code"], "missing_product_code")
        self.assertEqual(self.client.post(f"/api/outgoing-invoices/draft/{batch_id}").status_code, 409)

    def test_canonical_output_reduces_readiness(self):
        with server.db() as conn:
            self.add_opening(conn, 5)
            self.add_canonical_event(conn, -2, "output", "OUTPUT-2")
            batch_id, _ = self.add_batch(conn, "2026-08-20", [{"qty": 4}])
        payload = self.client.get(f"/api/outgoing-invoices/readiness/{batch_id}").get_json()
        self.assertEqual((payload["invoiceable_qty"], payload["pending_qty"]), (3, 1))

    def test_replacing_draft_cannot_restore_more_than_current_canonical_stock(self):
        with server.db() as conn:
            self.add_opening(conn, 7)
            batch_id, _ = self.add_batch(conn, "2026-08-20", [{"qty": 7}])
        first = self.client.post(f"/api/outgoing-invoices/draft/{batch_id}")
        self.assertEqual(first.status_code, 200, first.get_data(as_text=True))
        with server.db() as conn:
            self.add_canonical_event(conn, -5, "output", "LATE-OUTPUT-5")
        recalculated = self.client.post(f"/api/outgoing-invoices/draft/{batch_id}")
        self.assertEqual(recalculated.status_code, 200, recalculated.get_data(as_text=True))
        self.assertEqual(recalculated.get_json()["pending_qty"], 5)
        with server.db() as conn:
            row = conn.execute(
                """SELECT COUNT(*) lines,COALESCE(SUM(qty_out),0) qty
                     FROM inventory_transactions
                    WHERE source_type='OUTGOING_DRAFT' AND status='reserved'"""
            ).fetchone()
            self.assertEqual((row["lines"], row["qty"]), (1, 2))

    def test_local_confirm_keeps_hold_without_creating_canonical_ledger(self):
        with server.db() as conn:
            self.add_opening(conn, 10)
            batch_id, _ = self.add_batch(conn, "2026-08-20", [{"qty": 4}])
        created = self.client.post(f"/api/outgoing-invoices/draft/{batch_id}")
        self.assertEqual(created.status_code, 200, created.get_data(as_text=True))
        draft_id = created.get_json()["drafts"][0]["id"]
        with server.db() as conn:
            conn.execute(
                """UPDATE outgoing_invoice_drafts SET
                       buyer_name_snapshot='Buyer',buyer_tax_code_snapshot='0200000001',
                       buyer_address_snapshot='Address',company_name_snapshot='TDP',
                       company_tax_code_snapshot='0100000001',company_address_snapshot='Address',
                       payment_requester_snapshot='Requester',payment_bank_name_snapshot='Bank',
                       payment_bank_account_snapshot='Account'
                     WHERE id=?""",
                (draft_id,),
            )
        confirmed = self.client.post(
            f"/api/outgoing-invoices/{draft_id}/confirm-issued",
            json={
                "confirmed": True, "invoice_number": "0000041",
                "invoice_series": "1C26TDP", "invoice_date": "2026-08-20",
            },
        )
        self.assertEqual(confirmed.status_code, 200, confirmed.get_data(as_text=True))
        with server.db() as conn:
            self.assertEqual(
                conn.execute("SELECT COUNT(*) FROM invoice_inventory_ledger").fetchone()[0], 0,
            )
            draft = conn.execute(
                "SELECT status FROM outgoing_invoice_drafts WHERE id=?", (draft_id,),
            ).fetchone()
            reservation = conn.execute(
                """SELECT status FROM inventory_transactions
                    WHERE source_type='OUTGOING_DRAFT' AND source_id=?""",
                (str(draft_id),),
            ).fetchone()
            stock = canonical_available_stock(conn)["HH-01"]
            self.assertEqual((draft["status"], reservation["status"]), ("issued", "posted"))
            self.assertEqual(
                (stock["canonical_qty"], stock["reserved_qty"],
                 stock["pending_sync_issued_qty"], stock["available_qty"]),
                (10, 0, 4, 6),
            )

    def test_matching_msmi_identity_does_not_release_local_hold(self):
        with server.db() as conn:
            self.add_opening(conn, 10)
            batch_id, order_ids = self.add_batch(conn, "2026-08-20", [{"qty": 10}])
            self.add_local_issued_draft(
                conn, batch_id, order_ids[0], number="0000042", qty=4,
            )
            self.add_posted_source(conn, source="msmi", number="0000042", qty=4)
            stock = canonical_available_stock(conn)["HH-01"]
            self.assertEqual(
                (stock["canonical_qty"], stock["pending_sync_issued_qty"], stock["available_qty"]),
                (6, 4, 2),
            )

    def test_matching_minvoice_canonical_post_replaces_local_hold_once(self):
        with server.db() as conn:
            self.add_opening(conn, 4)
            batch_id, order_ids = self.add_batch(conn, "2026-08-20", [{"qty": 10}])
            self.add_local_issued_draft(
                conn, batch_id, order_ids[0], number="0000042", qty=4,
            )
            self.add_canonical_event(conn, 6, "input", "MATCH-INPUT-6")
            self.add_posted_source(conn, source="minvoice", number="0000042", qty=4)
            first = canonical_available_stock(conn)["HH-01"]
            second = canonical_available_stock(conn)["HH-01"]
            self.assertEqual(
                (first["canonical_qty"], first["pending_sync_issued_qty"], first["available_qty"]),
                (6, 0, 6),
            )
            self.assertEqual(first, second)
        payload = self.client.get(f"/api/outgoing-invoices/readiness/{batch_id}").get_json()
        self.assertEqual((payload["issued_qty"], payload["invoiceable_qty"], payload["pending_qty"]), (4, 6, 0))

    def test_minvoice_wrong_quantity_does_not_release_local_hold(self):
        with server.db() as conn:
            self.add_opening(conn, 10)
            batch_id, order_ids = self.add_batch(conn, "2026-08-20", [{"qty": 10}])
            self.add_local_issued_draft(
                conn, batch_id, order_ids[0], number="0000043", qty=4,
            )
            self.add_posted_source(conn, source="minvoice", number="0000043", qty=3)
            stock = canonical_available_stock(conn)["HH-01"]
            self.assertEqual(
                (stock["canonical_qty"], stock["pending_sync_issued_qty"], stock["available_qty"]),
                (7, 4, 3),
            )

    def test_minvoice_wrong_product_does_not_release_local_hold(self):
        with server.db() as conn:
            conn.execute(
                """INSERT OR REPLACE INTO products(
                       code,name,unit,tax,supplier,buy_price,purchase_list,seller,cccd
                   ) VALUES('HH-02','Hàng hóa 02','kg','0%','NCC-A',10,0,'','')"""
            )
            self.add_opening(conn, 10, product_code="HH-01")
            self.add_opening(conn, 10, product_code="HH-02")
            batch_id, order_ids = self.add_batch(conn, "2026-08-20", [{"qty": 10}])
            self.add_local_issued_draft(
                conn, batch_id, order_ids[0], number="0000044", qty=4,
            )
            self.add_posted_source(
                conn, source="minvoice", number="0000044", product_code="HH-02", qty=4,
            )
            stock = canonical_available_stock(conn)
            self.assertEqual(
                (stock["HH-01"]["canonical_qty"],
                 stock["HH-01"]["pending_sync_issued_qty"],
                 stock["HH-01"]["available_qty"]),
                (10, 4, 6),
            )
            self.assertEqual(stock["HH-02"]["available_qty"], 6)

    def test_minvoice_fingerprint_uses_six_decimal_stock_precision(self):
        with server.db() as conn:
            self.add_opening(conn, 10)
            batch_id, order_ids = self.add_batch(conn, "2026-08-20", [{"qty": 10}])
            self.add_local_issued_draft(
                conn, batch_id, order_ids[0], number="0000045", qty=4.0000004,
            )
            self.add_posted_source(conn, source="minvoice", number="0000045", qty=4)
            stock = canonical_available_stock(conn)["HH-01"]
            self.assertEqual(
                (stock["canonical_qty"], stock["pending_sync_issued_qty"], stock["available_qty"]),
                (6, 0, 6),
            )

    def test_minvoice_local_unit_must_match_catalog_before_hold_release(self):
        with server.db() as conn:
            self.add_opening(conn, 10)
            batch_id, order_ids = self.add_batch(conn, "2026-08-20", [{"qty": 10}])
            self.add_local_issued_draft(
                conn, batch_id, order_ids[0], number="0000046", qty=4, unit="thùng",
            )
            self.add_posted_source(conn, source="minvoice", number="0000046", qty=4)
            stock = canonical_available_stock(conn)["HH-01"]
            self.assertEqual(
                (stock["canonical_qty"], stock["pending_sync_issued_qty"], stock["available_qty"]),
                (6, 4, 2),
            )

    def test_cancel_releases_global_stock_for_backdated_batch(self):
        with server.db() as conn:
            self.add_opening(conn, 5)
            future_batch, _ = self.add_batch(conn, "2026-09-10", [{"qty": 5}])
            old_batch, _ = self.add_batch(conn, "2026-08-10", [{"qty": 5}])
        self.assertEqual(self.client.post(f"/api/outgoing-invoices/draft/{future_batch}").status_code, 200)
        blocked = self.client.get(f"/api/outgoing-invoices/readiness/{old_batch}").get_json()
        self.assertEqual((blocked["invoiceable_qty"], blocked["pending_qty"]), (0, 5))
        with server.db() as conn:
            draft_id = conn.execute(
                "SELECT id FROM outgoing_invoice_drafts WHERE batch_id=?", (future_batch,),
            ).fetchone()["id"]
        cancelled = self.client.post(
            f"/api/outgoing-invoices/{draft_id}/cancel", json={"confirmed": True},
        )
        self.assertEqual(cancelled.status_code, 200, cancelled.get_data(as_text=True))
        released = self.client.get(f"/api/outgoing-invoices/readiness/{old_batch}").get_json()
        self.assertEqual((released["invoiceable_qty"], released["pending_qty"]), (5, 0))

    def test_period_shortage_filter_and_static_export(self):
        with server.db() as conn:
            self.add_opening(conn, 5)
            first_batch, _ = self.add_batch(conn, "2026-08-01", [
                {"qty": 4, "contractor": "NT-A", "sell_price": 20},
            ])
            self.add_batch(conn, "2026-08-02", [
                {"qty": 4, "contractor": "NT-B", "sell_price": 30},
            ])
        self.assertEqual(self.client.post(f"/api/outgoing-invoices/draft/{first_batch}").status_code, 200)
        response = self.client.get("/api/outgoing-invoices/shortages?from=2026-08-01&to=2026-08-31")
        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        payload = response.get_json()
        self.assertEqual(payload["shortage_count"], 1)
        self.assertEqual(
            (payload["shortages"][0]["contractor"], payload["shortages"][0]["pending_qty"],
             payload["shortages"][0]["pending_value"]),
            ("NT-B", 3, 90),
        )
        filtered = self.client.get(
            "/api/outgoing-invoices/shortages?from=2026-08-01&to=2026-08-31&contractor=NT-A"
        ).get_json()
        self.assertEqual(filtered["contractor"], "NT-A")
        self.assertEqual(filtered["pending_qty"], 0)

        exported = self.client.get(
            "/api/outgoing-invoices/shortages/export?from=2026-08-01&to=2026-08-31"
        )
        self.assertEqual(exported.status_code, 200)
        workbook = load_workbook(io.BytesIO(exported.data), data_only=False)
        try:
            self.assertIn("TỔNG HỢP THIẾU", workbook.sheetnames)
            self.assertIn("NT-B", workbook.sheetnames)
            self.assertEqual(
                "Kỳ 01/08/2026 đến 31/08/2026 · tồn kho theo hóa đơn",
                workbook["TỔNG HỢP THIẾU"]["A2"].value,
            )
            formulas = []
            hyperlinks = []
            for sheet in workbook.worksheets:
                for row in sheet.iter_rows():
                    for cell in row:
                        if cell.data_type == "f":
                            formulas.append(cell.coordinate)
                        if cell.hyperlink:
                            hyperlinks.append(cell.coordinate)
            self.assertEqual(formulas, [])
            self.assertEqual(hyperlinks, [])
        finally:
            workbook.close()
        self.assertEqual(
            self.client.get("/api/outgoing-invoices/shortages?from=2026-08-31&to=2026-08-01").status_code,
            400,
        )

    def test_ui_exposes_round_status_and_period_shortage_export(self):
        script = (Path(server.STATIC_DIR) / "app.js").read_text(encoding="utf-8")
        self.assertIn("Đã dự thảo", script)
        self.assertIn("Đã phát hành", script)
        self.assertIn("Danh sách còn thiếu theo kỳ", script)
        self.assertIn("/api/outgoing-invoices/shortages/export?", script)
        self.assertIn('data-action="load-outgoing-shortages"', script)


if __name__ == "__main__":
    unittest.main()
