from __future__ import annotations

import io
import tempfile
import unittest
import zipfile
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook

try:
    from . import server
    from .invoice_input_sync import sync_input_batch
    from .invoice_inventory import invoice_stock_rows, post_output_invoice
    from .invoice_mapping import save_conversion, save_mapping
    from .invoice_output_sync import sync_output_batch
    from .invoice_receipt import create_input_receipt
    from .invoice_workbench import prepare_sync_batch
    from .test_invoice_input_sync import DateBoundedMsmi
    from .test_invoice_output_sync import OutputFixtureMsmi, STATUS_MAP, output_invoice
    from .test_msmi_sync import remote_invoice
except ImportError:  # pragma: no cover - direct pytest collection
    import server
    from invoice_input_sync import sync_input_batch
    from invoice_inventory import invoice_stock_rows, post_output_invoice
    from invoice_mapping import save_conversion, save_mapping
    from invoice_output_sync import sync_output_batch
    from invoice_receipt import create_input_receipt
    from invoice_workbench import prepare_sync_batch
    from test_invoice_input_sync import DateBoundedMsmi
    from test_invoice_output_sync import OutputFixtureMsmi, STATUS_MAP, output_invoice
    from test_msmi_sync import remote_invoice


class GoalAEndToEndSmoke(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory(prefix="tdp_goal_a_e2e_")
        cls.original_db_path = server.DB_PATH
        cls.original_data_dir = server.DATA_DIR
        cls.original_master_source = server.MASTER_SOURCE
        cls.original_testing = server.app.config.get("TESTING", False)
        server.DATA_DIR = Path(cls.temp.name) / "data"
        server.DATA_DIR.mkdir(parents=True, exist_ok=True)
        server.DB_PATH = server.DATA_DIR / "goal-a-e2e.sqlite3"
        server.MASTER_SOURCE = Path(cls.temp.name) / "master-disabled.xlsx"
        server.app.config["TESTING"] = True
        server.init_database()
        # Keep initialization isolated, then expose the approved golden only
        # to read-only template cloning performed by the export steps.
        server.MASTER_SOURCE = cls.original_master_source
        cls.client = server.app.test_client()

    @classmethod
    def tearDownClass(cls):
        with server.PENDING_IMPORT_LOCK:
            for pending in server.PENDING_IMPORTS.values():
                pending["path"].unlink(missing_ok=True)
            server.PENDING_IMPORTS.clear()
        server.DB_PATH = cls.original_db_path
        server.DATA_DIR = cls.original_data_dir
        server.MASTER_SOURCE = cls.original_master_source
        server.app.config["TESTING"] = cls.original_testing
        cls.temp.cleanup()

    def setUp(self):
        with server.db() as conn:
            conn.execute(
                """INSERT OR REPLACE INTO contractors(code,name,price_group,pricing_mode)
                   VALUES('TOYOTA','Nhà thầu E2E','TOYOTA','daily')"""
            )
            conn.execute(
                "INSERT OR REPLACE INTO kitchens(code,contractor,name) "
                "VALUES('K1','TOYOTA','Bếp E2E')"
            )
            conn.executemany(
                "INSERT OR REPLACE INTO suppliers(code,name) VALUES(?,?)",
                [("S1", "NCC E2E 1"), ("S2", "NCC E2E 2")],
            )
            conn.execute(
                """INSERT OR REPLACE INTO people(name,cccd,issue_date,issue_place,address)
                   VALUES('Người bán E2E','012345678901','01/01/2020',
                          'Nơi cấp E2E','Địa chỉ người bán E2E')"""
            )
            conn.executemany(
                """INSERT OR REPLACE INTO products(
                       code,name,unit,tax,supplier,buy_price,purchase_list,seller,cccd
                   ) VALUES(?,?,?,?,?,?,?,?,?)""",
                [
                    ("A000001", "Hàng E2E A", "kg", "8%", "S1", 9000, 1,
                     "Người bán E2E", "012345678901"),
                    ("B000001", "Hàng E2E B", "kg", "10%", "S2", 8000, 0, "", ""),
                ],
            )
            conn.execute(
                """INSERT OR REPLACE INTO dated_prices(
                       product_code,price_group,period,price_value,updated_at
                   ) VALUES('A000001','HATRAN','2026-08',9000,?)""",
                (server.now_iso(),),
            )
            for key, value in {
                "company": "Công ty E2E",
                "company_tax_code": "0200000002",
                "company_address": "Địa chỉ công ty E2E",
                "payment_requester": "Người lập E2E",
                "payment_bank_name": "Ngân hàng E2E",
                "payment_bank_account": "0000000000",
            }.items():
                server.setting_set(conn, key, value)

        buyer = self.client.put("/api/outgoing-buyers/TOYOTA", json={
            "display_name": "Khách E2E",
            "legal_name": "Công ty khách E2E",
            "tax_code": "0200000003",
            "address": "Địa chỉ khách E2E",
            "email": "",
        })
        self.assertEqual(buyer.status_code, 200, buyer.get_data(as_text=True))

    @staticmethod
    def daily_workbook(*, finalized: bool) -> bytes:
        workbook = Workbook()
        day_sheet = workbook.active
        day_sheet.title = "20.08"
        day_sheet.append(["ĐƠN HÀNG NGÀY"])
        day_sheet.append([
            "Ngày", "Nhà thầu", "Mã bếp", "Mã hàng", "Tên hàng", "Số lượng",
            "Thực giao", "ĐVT", "Chọn NCC", "Giá mua", "Giá bán",
        ])
        rows = [
            ("A000001", "Hàng E2E A", 10, 9 if finalized else 10, "S1", 9000, 13000),
            ("B000001", "Hàng E2E B", 8, 7 if finalized else 8, "S2", 8000, 12000),
        ]
        for code, name, qty, delivered, supplier, buy_price, sell_price in rows:
            day_sheet.append([
                date(2026, 8, 20), "TOYOTA", "K1", code, name, qty, delivered,
                "kg", supplier, buy_price, sell_price,
            ])

        purchase = workbook.create_sheet("đặt hàng")
        purchase.append(["ĐẶT NHÀ CUNG CẤP"])
        purchase.append([
            "Mã hàng", "Mã bếp", "Ngày", "Tên hàng", "Số lượng", "ĐVT", "NCC",
            "ghi chú", "giá mua", "hỏng", "thêm", "Giảm", "thiếu",
            "SL thực tế", "Thành tiền",
        ])
        for code, name, qty, _delivered, supplier, buy_price, _sell_price in rows:
            purchase.append([
                code, "K1", date(2026, 8, 20), name, qty, "kg", supplier,
                "Chốt mua E2E", buy_price, 0, 0, 0, 0, qty, qty * buy_price,
            ])
        stream = io.BytesIO()
        workbook.save(stream)
        workbook.close()
        return stream.getvalue()

    def analyze_daily(self, payload: bytes, filename: str) -> dict:
        response = self.client.post(
            "/api/import/analyze",
            data={"file": (io.BytesIO(payload), filename)},
            content_type="multipart/form-data",
        )
        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        return response.get_json()

    def confirm_daily(self, analysis: dict, sheets: list[str]):
        response = self.client.post("/api/import/confirm", json={
            "token": analysis["token"],
            "sheets": sheets,
            "work_date": "2026-08-20",
            "state_hash": analysis["stateHash"],
        })
        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        return response.get_json()

    @staticmethod
    def input_fixture(number: int, code: str, unit: str, qty: float, unit_price: float) -> dict:
        invoice = remote_invoice(number)
        subtotal = qty * unit_price
        invoice["hdhhdvu"][0].update(
            ma=code, ten=f"Nguồn {code}", dvtinh=unit,
            sluong=qty, dgia=unit_price, thtien=subtotal,
        )
        invoice.update(tgtcthue=subtotal, tgtthue=subtotal * 0.08, tgtttbso=subtotal * 1.08)
        return invoice

    @staticmethod
    def assert_workbook_response(testcase, response):
        testcase.assertEqual(
            response.status_code,
            200,
            f"HTTP {response.status_code}; content-type={response.content_type}",
        )
        testcase.assertGreater(len(response.data), 1000)
        workbook = load_workbook(io.BytesIO(response.data), data_only=False, keep_links=False)
        try:
            testcase.assertGreaterEqual(len(workbook.sheetnames), 1)
        finally:
            workbook.close()

    def test_representative_goal_a_end_to_end(self):
        # 1–3. First-load daily workbook, approve operational delivery/order,
        # then complete every supplier and explicitly reopen one supplier.
        first = self.confirm_daily(
            self.analyze_daily(self.daily_workbook(finalized=False), "e2e-first.xlsx"),
            ["20.08"],
        )
        batch_id = first["batch"]["id"]
        approved = self.client.post(f"/api/batches/{batch_id}/approve")
        self.assertEqual(approved.status_code, 200, approved.get_data(as_text=True))
        supplier_file = self.client.get(f"/api/export/suppliers/{batch_id}")
        self.assert_workbook_response(self, supplier_file)

        checklist = self.client.get(f"/api/supplier-needs/{batch_id}").get_json()
        self.assertEqual({item["supplier_key"] for item in checklist["checklist"]}, {"s1", "s2"})
        for item in checklist["checklist"]:
            marked = self.client.put(
                f"/api/supplier-order-status/{batch_id}/{item['supplier_key']}",
                json={"status": "ordered", "revision": item["revision"]},
            )
            self.assertEqual(marked.status_code, 200, marked.get_data(as_text=True))
        reopened = self.client.put(
            f"/api/supplier-order-status/{batch_id}/s1",
            json={"status": "reopened", "revision": 1},
        )
        self.assertEqual(reopened.status_code, 200, reopened.get_data(as_text=True))
        reordered = self.client.put(
            f"/api/supplier-order-status/{batch_id}/s1",
            json={"status": "ordered", "revision": 2},
        )
        self.assertEqual(reordered.status_code, 200, reordered.get_data(as_text=True))
        checklist_after = self.client.get(f"/api/supplier-needs/{batch_id}").get_json()
        self.assertEqual(checklist_after["checklist_counts"], {
            "pending": 0, "reopened": 0, "ordered": 2,
        })

        # 4–6. Second load commits purchase and customer scopes independently;
        # both ledgers and both exports must contain current, reconciled rows.
        final_bytes = self.daily_workbook(finalized=True)
        purchase_analysis = self.analyze_daily(final_bytes, "e2e-final-purchase.xlsx")
        self.assertEqual(purchase_analysis["phase"], "finalization")
        purchase = self.confirm_daily(purchase_analysis, ["đặt hàng"])
        self.assertEqual(purchase["selectedScopes"], ["purchase_orders"])
        sales = self.confirm_daily(
            self.analyze_daily(final_bytes, "e2e-final-sales.xlsx"), ["20.08"],
        )
        self.assertEqual(sales["selectedScopes"], ["customer_orders"])
        self.assertEqual(sales["finalized"]["lifecycle_status"], "finalized")
        with server.db() as conn:
            self.assertEqual(conn.execute(
                "SELECT COUNT(*) FROM payable_ledger_lines WHERE status!='reversed'"
            ).fetchone()[0], 2)
            self.assertEqual(conn.execute(
                "SELECT COUNT(*) FROM receivable_ledger_lines WHERE status='active'"
            ).fetchone()[0], 2)
        self.assert_workbook_response(self, self.client.get(
            "/api/debts/payables/export?from=2026-08-01&to=2026-08-31"
        ))
        self.assert_workbook_response(self, self.client.get(
            "/api/debts/receivables/export?from=2026-08-01&to=2026-08-31&contractor=TOYOTA"
        ))

        # 7–9. Exact August input/output fixture pull, independent two-way
        # mappings, explicit conversions, posts, replay, and stock equation.
        opening = self.client.post("/api/inventory/opening", json={
            "period": "2026-08",
            "items": [{"product_code": "A000001", "qty": 10, "unit_cost": 9000}],
        })
        self.assertEqual(opening.status_code, 200, opening.get_data(as_text=True))

        input_remote = self.input_fixture(10, "IN-BOX-A", "thùng", 2, 45000)
        output_remote = output_invoice(11, status="FIXTURE_ISSUED", series="E2E-OUT")
        output_remote["hdhhdvu"][0].update(
            ma="OUT-BOX-A", ten="Nguồn xuất A", dvtinh="thùng", sluong=1,
            dgia=78000, thtien=78000,
        )
        output_remote.update(tgtcthue=78000, tgtthue=6240, tgtttbso=84240)

        with server.db() as conn:
            input_batch, _ = prepare_sync_batch(
                conn, tenant="TDP", source="msmi", invoice_type="input",
                date_from="2026-08-01", date_to="2026-08-31", now_iso=server.now_iso,
            )
            output_batch, _ = prepare_sync_batch(
                conn, tenant="TDP", source="msmi", invoice_type="output",
                date_from="2026-08-01", date_to="2026-08-31", now_iso=server.now_iso,
            )
            input_client = DateBoundedMsmi([input_remote])
            output_client = OutputFixtureMsmi([output_remote])
            sync_input_batch(conn, input_client, input_batch["id"], server.now_iso)
            sync_output_batch(
                conn, output_client, output_batch["id"], server.now_iso,
                status_map=STATUS_MAP, status_fields=["fixtureStatus"],
            )
            input_invoice_id = conn.execute(
                "SELECT id FROM msmi_invoices WHERE remote_id=?", (input_remote["_id"],),
            ).fetchone()[0]
            input_item_id = conn.execute(
                "SELECT id FROM msmi_invoice_items WHERE invoice_id=?", (input_invoice_id,),
            ).fetchone()[0]
            output_invoice_id = conn.execute(
                "SELECT id FROM outgoing_source_invoices WHERE remote_id=?", (output_remote["_id"],),
            ).fetchone()[0]
            output_item_id = conn.execute(
                "SELECT id FROM outgoing_source_invoice_items WHERE invoice_id=?", (output_invoice_id,),
            ).fetchone()[0]
            self.assertTrue(save_mapping(
                conn, direction="input", item_id=input_item_id,
                product_code="A000001", now_iso=server.now_iso,
            )["requires_unit_conversion"])
            save_conversion(
                conn, direction="input", item_id=input_item_id,
                conversion_factor=5, now_iso=server.now_iso,
            )
            self.assertFalse(save_mapping(
                conn, direction="output", item_id=output_item_id,
                product_code="A000001", now_iso=server.now_iso,
            )["requires_unit_conversion"])
            # Outgoing invoices preserve source quantity; only incoming lines
            # use a conversion factor under the confirmed customer workflow.
            create_input_receipt(conn, input_invoice_id, server.now_iso)
            post_output_invoice(
                conn, output_invoice_id, confirmed=True, now_iso=server.now_iso,
            )
            event_count = conn.execute(
                "SELECT COUNT(*) FROM invoice_inventory_ledger"
            ).fetchone()[0]

            sync_input_batch(conn, input_client, input_batch["id"], server.now_iso)
            sync_output_batch(
                conn, output_client, output_batch["id"], server.now_iso,
                status_map=STATUS_MAP, status_fields=["fixtureStatus"],
            )
            self.assertTrue(create_input_receipt(
                conn, input_invoice_id, server.now_iso,
            )["idempotent"])
            self.assertTrue(post_output_invoice(
                conn, output_invoice_id, confirmed=True, now_iso=server.now_iso,
            )["idempotent"])
            self.assertEqual(conn.execute(
                "SELECT COUNT(*) FROM invoice_inventory_ledger"
            ).fetchone()[0], event_count)
            a_stock = next(
                item for item in invoice_stock_rows(conn, "2026-08-31", include_zero=True)
                if item["product_code"] == "A000001"
            )
            self.assertEqual(
                (a_stock["opening_qty"], a_stock["input_qty"],
                 a_stock["output_qty"], a_stock["closing_qty"]),
                (10, 10, 1, 19),
            )

        first_readiness = self.client.get(
            f"/api/outgoing-invoices/readiness/{batch_id}"
        ).get_json()
        self.assertEqual(
            (first_readiness["invoiceable_qty"], first_readiness["pending_qty"]),
            (9, 7),
        )
        first_round = self.client.post(f"/api/outgoing-invoices/draft/{batch_id}")
        self.assertEqual(first_round.status_code, 200, first_round.get_data(as_text=True))
        first_draft = first_round.get_json()["drafts"][0]
        first_tax_file = self.client.get(f"/api/export/invoices/{batch_id}")
        self.assertEqual(
            first_tax_file.status_code, 200,
            f"HTTP {first_tax_file.status_code}; content-type={first_tax_file.content_type}",
        )
        with zipfile.ZipFile(io.BytesIO(first_tax_file.data)) as archive:
            self.assertTrue(any(name.endswith(".xlsx") for name in archive.namelist()))
        first_issued = self.client.post(
            f"/api/outgoing-invoices/{first_draft['id']}/confirm-issued",
            json={
                "confirmed": True, "invoice_number": "1001",
                "invoice_series": "E2E26", "invoice_date": "2026-08-20",
            },
        )
        self.assertEqual(first_issued.status_code, 200, first_issued.get_data(as_text=True))

        refill_remote = self.input_fixture(30, "IN-B", "kg", 7, 8000)
        with server.db() as conn:
            refill_batch, _ = prepare_sync_batch(
                conn, tenant="TDP", source="msmi", invoice_type="input",
                date_from="2026-08-30", date_to="2026-08-30", now_iso=server.now_iso,
            )
            sync_input_batch(
                conn, DateBoundedMsmi([refill_remote]), refill_batch["id"], server.now_iso,
            )
            refill_invoice_id = conn.execute(
                "SELECT id FROM msmi_invoices WHERE remote_id=?", (refill_remote["_id"],),
            ).fetchone()[0]
            refill_item_id = conn.execute(
                "SELECT id FROM msmi_invoice_items WHERE invoice_id=?", (refill_invoice_id,),
            ).fetchone()[0]
            mapped = save_mapping(
                conn, direction="input", item_id=refill_item_id,
                product_code="B000001", now_iso=server.now_iso,
            )
            self.assertFalse(mapped["requires_unit_conversion"])
            create_input_receipt(conn, refill_invoice_id, server.now_iso)

        second_readiness = self.client.get(
            f"/api/outgoing-invoices/readiness/{batch_id}"
        ).get_json()
        self.assertEqual(
            (second_readiness["issued_qty"], second_readiness["invoiceable_qty"],
             second_readiness["pending_qty"]),
            (9, 7, 0),
        )
        second_round = self.client.post(f"/api/outgoing-invoices/draft/{batch_id}")
        self.assertEqual(second_round.status_code, 200, second_round.get_data(as_text=True))
        second_draft = second_round.get_json()["drafts"][0]
        self.assertEqual(second_draft["round_no"], 2)
        second_tax_file = self.client.get(f"/api/export/invoices/{batch_id}")
        self.assertEqual(
            second_tax_file.status_code, 200,
            f"HTTP {second_tax_file.status_code}; content-type={second_tax_file.content_type}",
        )
        with zipfile.ZipFile(io.BytesIO(second_tax_file.data)) as archive:
            self.assertTrue(any(name.endswith(".xlsx") for name in archive.namelist()))
        second_issued = self.client.post(
            f"/api/outgoing-invoices/{second_draft['id']}/confirm-issued",
            json={
                "confirmed": True, "invoice_number": "1002",
                "invoice_series": "E2E26", "invoice_date": "2026-08-30",
            },
        )
        self.assertEqual(second_issued.status_code, 200, second_issued.get_data(as_text=True))
        final_readiness = self.client.get(
            f"/api/outgoing-invoices/readiness/{batch_id}"
        ).get_json()
        self.assertEqual((final_readiness["issued_qty"], final_readiness["pending_qty"]), (16, 0))
        with server.db() as conn:
            stocks = invoice_stock_rows(conn, "2026-08-31", include_zero=True)
            self.assertTrue(all(item["closing_qty"] >= 0 for item in stocks))
            for item in stocks:
                self.assertAlmostEqual(
                    item["opening_qty"] + item["input_qty"] - item["output_qty"],
                    item["closing_qty"], places=6,
                )

        # 10–11. Quote, approved golden-derived operational forms, and the
        # invoice-only delivery statement all open and reconcile.
        quote = self.client.get(f"/api/export/quote/TOYOTA?batch_id={batch_id}")
        self.assert_workbook_response(self, quote)
        delivery = self.client.get(f"/api/export/deliveries/{batch_id}")
        report = self.client.get(f"/api/export/report/{batch_id}")
        purchases = self.client.get(f"/api/export/purchases/{batch_id}")
        for label, response in (
            ("delivery", delivery), ("report", report), ("purchases", purchases),
        ):
            with self.subTest(golden_export=label):
                self.assertEqual(
                    response.status_code, 200,
                    f"{label}: HTTP {response.status_code}; payload={response.get_json(silent=True)}",
                )
                self.assert_workbook_response(self, response)
        purchase_book = load_workbook(io.BytesIO(purchases.data), data_only=False, keep_links=False)
        try:
            self.assertIn("bảng kê tổng", purchase_book.sheetnames)
            self.assertTrue(any(name.startswith("biên nhận") for name in purchase_book.sheetnames))
        finally:
            purchase_book.close()

        scope_response = self.client.get(
            "/api/outgoing-invoices/payment-scope/TOYOTA"
            "?from=2026-08-01&to=2026-08-31"
        )
        self.assertEqual(scope_response.status_code, 200, scope_response.get_data(as_text=True))
        scope = scope_response.get_json()
        self.assertEqual(len(scope["invoices"]), 2)
        statement = self.client.get(
            "/api/outgoing-invoices/delivery-statement/TOYOTA"
            f"?from=2026-08-01&to=2026-08-31&scope_id={scope['scope_id']}"
        )
        self.assertEqual(
            statement.status_code, 200,
            f"HTTP {statement.status_code}; content-type={statement.content_type}",
        )
        statement_book = load_workbook(
            io.BytesIO(statement.data), data_only=False, keep_links=False,
        )
        try:
            reconcile = statement_book["Đối chiếu hóa đơn"]
            self.assertTrue(all(
                reconcile.cell(row, 10).value == 0
                for row in range(4, reconcile.max_row + 1)
            ))
        finally:
            statement_book.close()

        # 12. Kitchen/PO and payroll remain writable/readable in the same
        # source candidate after all prior modules have run.
        plan = self.client.post("/api/kitchen/plans", json={
            "work_date": "2026-08-20", "kitchen": "K1", "shift": "Trưa",
            "meal_count": 10, "xcom_code": "XCOM-K1", "menu_count": 1,
            "servings_per_menu": 10, "meal_price": 25000, "other_cost": 0,
            "items": [{
                "dish_name": "Món E2E", "product_code": "A000001",
                "norm_qty": 0.1, "unit": "kg", "supplier": "S1",
            }],
        })
        self.assertEqual(plan.status_code, 200, plan.get_data(as_text=True))
        plan_id = plan.get_json()["id"]
        approved_plan = self.client.post(f"/api/kitchen/plans/{plan_id}/approve")
        self.assertEqual(approved_plan.status_code, 200, approved_plan.get_data(as_text=True))
        self.assert_workbook_response(
            self, self.client.get("/api/kitchen/po?date=2026-08-20"),
        )

        staff = self.client.post("/api/staff", json={
            "employee_code": "NV-E2E", "full_name": "Nhân sự E2E",
            "role_name": "Kiểm thử", "kitchen": "K1", "base_salary": 8000000,
            "standard_days": 26, "standard_hours": 8,
            "bhxh_employee_rate": 0, "bhxh_company_rate": 0,
        })
        self.assertEqual(staff.status_code, 200, staff.get_data(as_text=True))
        attendance = self.client.post("/api/attendance", json={
            "employee_code": "NV-E2E", "work_date": "2026-08-20",
            "normal_hours": 8, "overtime_hours": 2, "sunday_hours": 0,
            "night_hours": 0, "holiday_hours": 0,
        })
        self.assertEqual(attendance.status_code, 200, attendance.get_data(as_text=True))
        payroll = self.client.get("/api/payroll?month=2026-08")
        self.assertEqual(payroll.status_code, 200, payroll.get_data(as_text=True))
        self.assertEqual(len(payroll.get_json()["items"]), 1)
        self.assert_workbook_response(
            self, self.client.get("/api/export/payroll?month=2026-08"),
        )

        with server.db() as conn:
            self.assertEqual(conn.execute("PRAGMA integrity_check").fetchone()[0], "ok")
            self.assertEqual(conn.execute("PRAGMA foreign_key_check").fetchall(), [])
