"""Review worksheets record choices on existing orders; notes track the work period."""
from io import BytesIO
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
import json
import secrets
import uuid
import zipfile

from flask import jsonify, request, send_file
from itsdangerous import URLSafeTimedSerializer, BadSignature
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

try:
    from .outgoing_contractors import save_line_choices, line_choices_payload
    from .outgoing_unissued import unissued_payload
    from .template_workbook import safe_workbook_bytes
except ImportError:
    from outgoing_contractors import save_line_choices, line_choices_payload
    from outgoing_unissued import unissued_payload
    from template_workbook import safe_workbook_bytes

SCHEMA='''CREATE TABLE IF NOT EXISTS outgoing_review_notes (
    date_from TEXT NOT NULL, date_to TEXT NOT NULL, contractor TEXT NOT NULL,
    note TEXT NOT NULL, revision TEXT NOT NULL, updated_at TEXT NOT NULL,
    PRIMARY KEY(date_from,date_to,contractor)
);'''
SHEET='Lua chon mat hang'
HEADERS=['Đưa vào file (1/0)','Dòng đơn','Ngày đơn','Nhà thầu','Bếp','Mã hàng',
         'Tên xuất hóa đơn','ĐVT đơn','SL chưa xuất','Đơn giá','Tiền hàng','Thuế','ĐVT hóa đơn']
MAX_ROWS=10000


def literal(value):
    text=str(value or '')
    return "'"+text if text.startswith(('=','+','-','@')) else text


def scope(values,validate_date):
    start=validate_date(values.get('from'),'Từ ngày')
    end=validate_date(values.get('to'),'Đến ngày')
    if start>end:raise ValueError('Từ ngày phải nhỏ hơn hoặc bằng Đến ngày.')
    party=str(values.get('contractor') or '').strip().upper()
    return {'from':start,'to':end,'contractor':'' if party=='*' else party}


def choices(conn,period):
    payload=unissued_payload(conn,period['to'],period['contractor'],start=period['from'],respect_export_choices=True)
    units={r['code']:r['unit'] for r in conn.execute('SELECT code,unit FROM products')}
    units.update({r['product_code']:r['invoice_unit'] for r in conn.execute('SELECT * FROM outgoing_product_units')})
    tax={r['id']:r['tax'] for r in conn.execute('SELECT id,tax FROM orders WHERE work_date BETWEEN ? AND ?',(period['from'],period['to']))}
    result=[]
    for r in payload['line_choices']:
        amount=int((Decimal(str(r['qty']))*Decimal(str(r['price']))).quantize(Decimal('1'),rounding=ROUND_HALF_UP))
        result.append({**r,'tax':tax[r['order_id']],'invoice_unit':units.get(r['product_code'],r['unit']),'amount':amount})
    return result


def workbook(rows,period,serializer):
    if not rows:raise ValueError('Không có dòng chưa xuất trong khoảng ngày và nhà thầu đã chọn.')
    if len(rows)>MAX_ROWS:raise ValueError('Chọn khoảng ngày nhỏ hơn để bảng lựa chọn dưới 10.000 dòng.')
    wb=Workbook();ws=wb.active;ws.title=SHEET;ws.append(HEADERS)
    manifest=[]
    for r in rows:
        values=[r['order_id'],r['date'],literal(r['contractor']),literal(r['kitchen']),literal(r['product_code']),
                literal(r['invoice_name']),literal(r['unit']),r['qty'],r['price'],r['amount'],literal(r['tax']),literal(r['invoice_unit'])]
        ws.append([int(r['enabled']),*values])
        manifest.append({'order_id':r['order_id'],'token':r['token'],'values':values,'enabled':r['enabled']})
    ws.freeze_panes='H2';ws.auto_filter.ref=ws.dimensions
    validation=DataValidation(type='list',formula1='"1,0"');validation.errorTitle='Chọn 1 hoặc 0';validation.error='1: đưa vào file; 0: bỏ khỏi file.';validation.showErrorMessage=True
    ws.add_data_validation(validation);validation.add(f'A2:A{ws.max_row}')
    for column,width in enumerate([20,12,16,22,22,16,38,14,18,18,20,14,18],1):
        ws.column_dimensions[ws.cell(1,column).column_letter].width=width
    for cell in ws[1]:cell.font=Font(bold=True)
    for row in ws.iter_rows(min_row=2):
        row[0].fill=PatternFill('solid',fgColor='FFF2CC')
        for cell in row[8:11]:cell.number_format='#,##0.######'
    meta=wb.create_sheet('_DoiChieu');meta.sheet_state='veryHidden'
    signed=serializer.dumps({'scope':period,'rows':manifest})
    for offset in range(0,len(signed),30000):meta.append([signed[offset:offset+30000]])
    guide=wb.create_sheet('Huong dan')
    for line in [
        'BẢNG LỰA CHỌN MẶT HÀNG — KIỂM TRA TRƯỚC KHI LẬP HÓA ĐƠN',
        'Cột đầu: 1 là đưa vào file, 0 là bỏ chọn. Có thể xóa dòng khỏi sheet Lua chon mat hang.',
        'Nhập lại file trên web, xem thay đổi rồi bấm Lưu lựa chọn. Các dòng bỏ chọn sẽ không tự trở lại ở lần tải sau; có thể chọn lại trên web.',
        'Giữ nguyên các cột thông tin khác. Bảng này chỉ lưu lựa chọn theo từng dòng đơn; không đổi hàng, số lượng, giá hay đơn vị.',
        'Dòng mới ngoài file không bị bỏ chọn. Nếu đơn hoặc hóa đơn đã ký vừa thay đổi, tải lại bảng mới trước khi lưu.',
        'Sau khi lưu, dùng Tải bảng kê đã chọn để up M-Invoice. Kho ghi theo hóa đơn đã ký được đồng bộ và đối chiếu; không theo bảng lựa chọn.',
        f'Ngày đơn: {period["from"]} → {period["to"]}; nhà thầu: {period["contractor"] or "tất cả đã chọn"}.',
    ]:guide.append([line])
    guide.column_dimensions['A'].width=110
    try:return safe_workbook_bytes(wb)
    finally:wb.close()


def _same(a,b):
    if isinstance(a,(int,float)) and not isinstance(a,bool):
        if isinstance(b,bool):return False
        try:return abs(Decimal(str(a))-Decimal(str(b)))<=Decimal('0.00000001')
        except (InvalidOperation,ValueError):return False
    return a==b


def read_choices(data,serializer):
    if not data or len(data)>10*1024*1024:raise ValueError('Chọn file .xlsx không quá 10 MB.')
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            if sum(r.file_size for r in archive.infolist())>50*1024*1024:raise ValueError('File Excel quá lớn.')
        wb=load_workbook(BytesIO(data),read_only=True,data_only=False)
        try:
            if SHEET not in wb or '_DoiChieu' not in wb:raise ValueError('Dùng đúng bảng lựa chọn đã tải từ web, giữ sheet đối chiếu.')
            signed=''.join(str(r[0] or '') for r in wb['_DoiChieu'].iter_rows(values_only=True))
            manifest=serializer.loads(signed,max_age=7*24*60*60)
            source={r['order_id']:r for r in manifest['rows']}
            if not source or len(source)>MAX_ROWS:raise ValueError('Phạm vi file không hợp lệ.')
            iterator=wb[SHEET].iter_rows(values_only=True)
            if list(next(iterator))!=HEADERS:raise ValueError('Giữ nguyên các cột của bảng lựa chọn.')
            selected={}
            for number,row in enumerate(iterator,2):
                if number>MAX_ROWS+1:raise ValueError('Bảng có quá nhiều dòng.')
                if not any(v is not None for v in row):continue
                flag,oid,*values=row
                if type(oid) is not int or oid not in source or oid in selected:raise ValueError(f'Dòng {number}: mã dòng đơn bị đổi hoặc trùng.')
                if any(not _same(a,b) for a,b in zip(source[oid]['values'],[oid,*values])):
                    raise ValueError(f'Dòng {number}: chỉ sửa cột Đưa vào file hoặc xóa dòng; giữ nguyên thông tin đơn.')
                if str(flag).strip() not in ('0','1','True','False'):raise ValueError(f'Dòng {number}: chọn 1 hoặc 0 ở cột Đưa vào file.')
                selected[oid]=str(flag).strip() in ('1','True')
            return manifest,[{'order_id':oid,'enabled':selected.get(oid,False),'token':r['token']} for oid,r in source.items()]
        finally:wb.close()
    except (BadSignature,KeyError,StopIteration) as exc:
        raise ValueError('File hết hạn hoặc phần đối chiếu đã thay đổi. Tải bảng lựa chọn mới.') from exc
    except ValueError:raise
    except Exception as exc:raise ValueError('Không đọc được bảng lựa chọn. Dùng file .xlsx tải từ màn này.') from exc


def check_current(conn,manifest,items):
    current={r['order_id']:r for r in line_choices_payload(conn,'9999-12-31')}
    changed=[]
    for item in items:
        r=current.get(item['order_id'])
        if not r or r['token']!=item['token']:raise ValueError('Dòng đơn, phần đã ký hoặc lựa chọn đã thay đổi. Tải lại bảng trước khi lưu.')
        if r['enabled']!=item['enabled']:
            if not r['editable']:raise ValueError(r['reason'])
            changed.append({**item,'contractor':r['contractor'],'product_code':r['product_code'],'invoice_name':r['invoice_name'],'date':r['date']})
    return changed


def notes_payload(conn,period):
    current=conn.execute('SELECT * FROM outgoing_review_notes WHERE date_from=? AND date_to=? AND contractor=?',
                         (period['from'],period['to'],period['contractor'])).fetchone()
    recent=[dict(r) for r in conn.execute('SELECT * FROM outgoing_review_notes ORDER BY updated_at DESC LIMIT 10')]
    return {'scope':period,'current':dict(current) if current else {'note':'','revision':'','updated_at':''},'recent':recent}


def register(app,ctx):
    fallback=secrets.token_hex(32)
    def signer():return URLSafeTimedSerializer(app.secret_key or fallback,salt='tdp-order-review-v1')

    @app.get('/api/outgoing-invoices/review.xlsx')
    def review_template():
        try:
            period=scope(request.args,ctx['valid_iso_date'])
            with ctx['db']() as conn:
                conn.execute('PRAGMA query_only=ON');conn.execute('BEGIN')
                data=workbook(choices(conn,period),period,signer())
            return send_file(BytesIO(data),as_attachment=True,download_name=f'LUA_CHON_MAT_HANG_TU_{period["from"]}_DEN_{period["to"]}.xlsx',mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        except ValueError as exc:return jsonify(ok=False,error=str(exc)),400

    @app.post('/api/outgoing-invoices/review/preview')
    def review_preview():
        try:
            upload=request.files.get('file')
            if not upload:raise ValueError('Chọn bảng lựa chọn đã chỉnh sửa.')
            manifest,items=read_choices(upload.read(10*1024*1024+1),signer())
            with ctx['db']() as conn:
                conn.execute('PRAGMA query_only=ON');conn.execute('BEGIN')
                changed=check_current(conn,manifest,items)
            token=signer().dumps({'scope':manifest['scope'],'items':items,'kind':'save_choices'})
            return jsonify(ok=True,scope=manifest['scope'],changes=changed,total=len(items),
                           selected=sum(r['enabled'] for r in items),excluded=sum(not r['enabled'] for r in items),token=token)
        except ValueError as exc:return jsonify(ok=False,error=str(exc)),409

    @app.put('/api/outgoing-invoices/review/choices')
    def review_save():
        try:
            body=request.get_json(silent=True) or {}
            if not isinstance(body,dict) or body.get('confirmed') is not True:raise ValueError('Kiểm tra thay đổi trước khi lưu lựa chọn.')
            if not isinstance(body.get('token'),str) or not body['token'] or len(body['token'])>2*1024*1024:raise ValueError('Bản xem trước không hợp lệ. Nhập lại file.')
            plan=signer().loads(body.get('token',''),max_age=30*60)
            if plan.get('kind')!='save_choices':raise ValueError('Bản xem trước không hợp lệ.')
            with ctx['db']() as conn:
                conn.execute('BEGIN IMMEDIATE')
                changed=check_current(conn,plan,plan['items'])
                result=save_line_choices(conn,{'items':changed},ctx['now_iso']()) if changed else {'changed':[],'released_draft_ids':[]}
            return jsonify(ok=True,scope=plan['scope'],**result)
        except (ValueError,BadSignature) as exc:return jsonify(ok=False,error=str(exc) if isinstance(exc,ValueError) else 'Bản xem trước hết hạn. Nhập lại file.'),409

    @app.route('/api/outgoing-invoices/review-notes',methods=['GET','PUT'])
    def review_notes():
        try:
            body=request.get_json(silent=True) if request.method=='PUT' else request.args
            if body is None or not hasattr(body,'get'):raise ValueError('Ghi chú không hợp lệ.')
            period=scope(body,ctx['valid_iso_date'])
            with ctx['db']() as conn:
                if request.method=='PUT':
                    conn.execute('BEGIN IMMEDIATE')
                    note=body.get('note')
                    if not isinstance(note,str) or len(note)>3000:raise ValueError('Ghi chú tối đa 3.000 ký tự.')
                    current=notes_payload(conn,period)['current']
                    if body.get('revision','')!=current['revision']:raise ValueError('Ghi chú vừa thay đổi. Tải lại để xem nội dung mới trước khi lưu.')
                    conn.execute('''INSERT INTO outgoing_review_notes VALUES(?,?,?,?,?,?)
                        ON CONFLICT(date_from,date_to,contractor) DO UPDATE SET note=excluded.note,revision=excluded.revision,updated_at=excluded.updated_at''',
                        (period['from'],period['to'],period['contractor'],note.strip(),uuid.uuid4().hex,ctx['now_iso']()))
                else:conn.execute('PRAGMA query_only=ON');conn.execute('BEGIN')
                return jsonify(ok=True,**notes_payload(conn,period))
        except ValueError as exc:return jsonify(ok=False,error=str(exc)),409 if request.method=='PUT' else 400
