"""Safe, template-preserving Excel output primitives.

The customer's golden workbook is loaded from immutable bytes for every
generation.  Selected sheets keep their native layout while formulas that
would depend on an external workbook or an omitted sheet are replaced with
their cached values from the same golden file.  Callers then write operational
data as literals and can hand the resulting :class:`Workbook` to the existing
print/PDF pipeline.
"""

from __future__ import annotations

import hashlib
import io
import math
import re
from copy import copy
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.formula import Tokenizer
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter


MAX_TEMPLATE_BYTES = 50 * 1024 * 1024
DANGEROUS_FORMULA_FUNCTION = re.compile(
    r"(?i)(?:^|[^A-Z0-9_.])(?:WEBSERVICE|RTD|HYPERLINK|INDIRECT|DDE|CALL|EXEC|REGISTER\.ID)\s*\("
)
EXTERNAL_URI = re.compile(r"(?i)(?:https?|ftp|file)://|\\\\")
EXTERNAL_BOOK_REFERENCE = re.compile(r"\[[^\]]+\][^!]*!")
INVALID_SHEET_CHARACTER = re.compile(r"[\\/*?:\[\]]")


class TemplateWorkbookError(ValueError):
    def __init__(self, message: str, *, code: str = "invalid_template"):
        super().__init__(message)
        self.code = code


@dataclass
class TemplateClone:
    workbook: Any
    source_path: str
    source_hash: str
    retained_sheets: tuple[str, ...]
    formula_report: dict[str, int]

    def close(self) -> None:
        self.workbook.close()

    def to_bytes(self) -> bytes:
        return safe_workbook_bytes(self.workbook)


def _plain(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip())


def _formula_text(value: Any) -> str:
    text = str(value or "")
    return text if text.startswith("=") else "=" + text


def _sheet_references(formula: str) -> tuple[list[str], bool]:
    references: list[str] = []
    try:
        tokens = Tokenizer(_formula_text(formula)).items
    except Exception:
        return references, False
    for token in tokens:
        if token.type != "OPERAND" or token.subtype != "RANGE" or "!" not in token.value:
            continue
        prefix = token.value.rsplit("!", 1)[0].strip()
        if prefix.startswith("'") and prefix.endswith("'"):
            prefix = prefix[1:-1].replace("''", "'")
        if prefix:
            references.append(prefix)
    return references, True


def unsafe_formula_reason(formula: Any, available_sheets: Iterable[str]) -> str:
    text = _formula_text(formula)
    if "#REF!" in text.upper():
        return "broken_reference"
    if EXTERNAL_BOOK_REFERENCE.search(text) or EXTERNAL_URI.search(text) or "|" in text:
        return "external_reference"
    if DANGEROUS_FORMULA_FUNCTION.search(text):
        return "unsafe_function"
    references, parsed = _sheet_references(text)
    if not parsed:
        return "invalid_formula"
    available = {name.casefold() for name in available_sheets}
    for reference in references:
        if "[" in reference or "]" in reference:
            return "external_reference"
        # 3-D references are not needed by the approved output templates and
        # are unsafe once a subset of sheets is exported independently.
        if ":" in reference:
            return "unsupported_3d_reference"
        if reference.casefold() not in available:
            return "omitted_sheet_reference"
    return ""


def _set_literal(cell: Any, value: Any) -> None:
    if isinstance(value, bool) or value is None:
        cell.value = value
        return
    if isinstance(value, (int, float)):
        if not math.isfinite(float(value)):
            raise TemplateWorkbookError(
                f"Giá trị tại {cell.coordinate} phải là số hữu hạn",
                code="non_finite_value",
            )
        cell.value = value
        return
    text = str(value)
    if text.startswith("="):
        raise TemplateWorkbookError(
            f"Dữ liệu động tại {cell.coordinate} không được là công thức",
            code="formula_injection",
        )
    cell.value = value


def write_literal(sheet: Any, coordinate: str, value: Any) -> None:
    cell = sheet[coordinate]
    if isinstance(cell, MergedCell):
        raise TemplateWorkbookError(
            f"Không thể ghi vào ô phụ của vùng gộp {coordinate}",
            code="merged_cell_write",
        )
    _set_literal(cell, value)


def write_formula(sheet: Any, coordinate: str, formula: str) -> None:
    cell = sheet[coordinate]
    if isinstance(cell, MergedCell):
        raise TemplateWorkbookError(
            f"Không thể ghi vào ô phụ của vùng gộp {coordinate}",
            code="merged_cell_write",
        )
    text = _formula_text(formula)
    reason = unsafe_formula_reason(text, sheet.parent.sheetnames)
    if reason:
        raise TemplateWorkbookError(
            f"Công thức tại {coordinate} không an toàn: {reason}",
            code="unsafe_formula",
        )
    cell.value = text


def _defined_name_reason(value: Any, available_sheets: Sequence[str]) -> str:
    text = _plain(getattr(value, "attr_text", ""))
    if not text:
        return ""
    if EXTERNAL_BOOK_REFERENCE.search(text) or EXTERNAL_URI.search(text) or "|" in text:
        return "external_reference"
    references, parsed = _sheet_references("=" + text)
    if not parsed:
        return "invalid_formula"
    available = {name.casefold() for name in available_sheets}
    for reference in references:
        if "[" in reference or "]" in reference:
            return "external_reference"
        if ":" in reference or reference.casefold() not in available:
            return "omitted_sheet_reference"
    return ""


def _sanitize_clone(workbook: Any, values_book: Any) -> dict[str, int]:
    report = {
        "preserved_formulas": 0,
        "replaced_external_reference": 0,
        "replaced_omitted_sheet_reference": 0,
        "replaced_unsafe_formula": 0,
        "missing_cached_values": 0,
        "external_hyperlinks_removed": 0,
        "defined_names_removed": 0,
    }
    available = tuple(workbook.sheetnames)
    for sheet in workbook.worksheets:
        values_sheet = values_book[sheet.title]
        for cells in sheet.iter_rows():
            for cell in cells:
                if isinstance(cell, MergedCell):
                    continue
                hyperlink = cell.hyperlink
                if hyperlink is not None and getattr(hyperlink, "target", None):
                    target = str(hyperlink.target)
                    if not target.startswith("#"):
                        cell._hyperlink = None
                        report["external_hyperlinks_removed"] += 1
                if cell.data_type != "f" and not (
                    isinstance(cell.value, str) and cell.value.startswith("=")
                ):
                    continue
                reason = unsafe_formula_reason(cell.value, available)
                if not reason:
                    report["preserved_formulas"] += 1
                    continue
                cached = values_sheet[cell.coordinate].value
                cell.value = cached
                if isinstance(cached, str) and cached.startswith("="):
                    # Preserve a cached string literally; do not reactivate it.
                    cell.data_type = "s"
                if cached is None:
                    report["missing_cached_values"] += 1
                if reason == "external_reference":
                    report["replaced_external_reference"] += 1
                elif reason == "omitted_sheet_reference":
                    report["replaced_omitted_sheet_reference"] += 1
                else:
                    report["replaced_unsafe_formula"] += 1

    for name, defined in list(workbook.defined_names.items()):
        if _defined_name_reason(defined, available):
            del workbook.defined_names[name]
            report["defined_names_removed"] += 1
    workbook._external_links = []
    return report


def clone_template_workbook(
    template_path: str | Path, *, sheet_names: Sequence[str], expected_sha256: str = "",
) -> TemplateClone:
    """Clone selected golden sheets without ever saving back to the source."""
    path = Path(template_path).resolve()
    if not path.is_file():
        raise TemplateWorkbookError("Không tìm thấy workbook golden", code="template_not_found")
    size = path.stat().st_size
    if size <= 0 or size > MAX_TEMPLATE_BYTES:
        raise TemplateWorkbookError("Workbook golden có kích thước không hợp lệ", code="template_size")
    payload = path.read_bytes()
    source_hash = hashlib.sha256(payload).hexdigest().upper()
    expected = _plain(expected_sha256).upper()
    if expected and source_hash != expected:
        raise TemplateWorkbookError(
            "Workbook golden không khớp SHA-256 đã khóa",
            code="template_hash_mismatch",
        )
    selected = tuple(str(name) for name in sheet_names)
    if not selected or len(set(selected)) != len(selected):
        raise TemplateWorkbookError("Danh sách sheet golden phải có ít nhất một tên duy nhất")
    for name in selected:
        if not name or len(name) > 31 or INVALID_SHEET_CHARACTER.search(name):
            raise TemplateWorkbookError(f"Tên sheet golden không hợp lệ: {name!r}")

    workbook = values_book = None
    try:
        workbook = load_workbook(
            io.BytesIO(payload), data_only=False, read_only=False, keep_links=False,
        )
        values_book = load_workbook(
            io.BytesIO(payload), data_only=True, read_only=False, keep_links=False,
        )
        missing = [name for name in selected if name not in workbook.sheetnames]
        if missing:
            raise TemplateWorkbookError(
                "Thiếu sheet golden: " + ", ".join(missing), code="template_sheet_missing",
            )
        for sheet in list(workbook.worksheets):
            if sheet.title not in selected:
                workbook.remove(sheet)
        # Removing sheets keeps their original relative order.  Reorder only
        # when the caller explicitly requested a different deterministic order.
        workbook._sheets = [workbook[name] for name in selected]
        report = _sanitize_clone(workbook, values_book)
        assert_workbook_safe(workbook)
        return TemplateClone(
            workbook=workbook,
            source_path=str(path),
            source_hash=source_hash,
            retained_sheets=selected,
            formula_report=report,
        )
    except TemplateWorkbookError:
        if workbook is not None:
            workbook.close()
        raise
    except Exception as error:
        if workbook is not None:
            workbook.close()
        raise TemplateWorkbookError(
            f"Không thể đọc workbook golden: {type(error).__name__}",
            code="template_unreadable",
        ) from error
    finally:
        if values_book is not None:
            values_book.close()


def copy_row_layout(
    sheet: Any, source_row: int, target_row: int, *, include_values: bool = False,
    translate_formulas: bool = True,
) -> None:
    """Copy one prototype row's dimensions and cell styles deterministically."""
    source_row = int(source_row)
    target_row = int(target_row)
    if source_row <= 0 or target_row <= 0:
        raise TemplateWorkbookError("Số dòng mẫu/đích phải lớn hơn 0")
    source_dimension = sheet.row_dimensions[source_row]
    target_dimension = sheet.row_dimensions[target_row]
    for attribute in ("height", "hidden", "outlineLevel", "collapsed", "thickTop", "thickBot"):
        setattr(target_dimension, attribute, getattr(source_dimension, attribute))
    available = sheet.parent.sheetnames
    for column in range(1, sheet.max_column + 1):
        source = sheet.cell(source_row, column)
        target = sheet.cell(target_row, column)
        if isinstance(source, MergedCell) or isinstance(target, MergedCell):
            continue
        target._style = copy(source._style)
        target.protection = copy(source.protection)
        target.alignment = copy(source.alignment)
        target.comment = copy(source.comment) if source.comment is not None else None
        target._hyperlink = None
        if not include_values:
            target.value = None
            continue
        value = source.value
        if source.data_type == "f" or (isinstance(value, str) and value.startswith("=")):
            formula = str(value)
            if translate_formulas and source_row != target_row:
                formula = Translator(
                    formula, origin=f"{get_column_letter(column)}{source_row}",
                ).translate_formula(f"{get_column_letter(column)}{target_row}")
            reason = unsafe_formula_reason(formula, available)
            if reason:
                raise TemplateWorkbookError(
                    f"Công thức dòng mẫu không an toàn tại {source.coordinate}: {reason}",
                    code="unsafe_formula",
                )
            target.value = formula
        else:
            _set_literal(target, value)


def write_table_rows(
    sheet: Any, *, start_row: int, records: Iterable[Mapping[str, Any]],
    columns: Mapping[str, int], prototype_row: int | None = None,
) -> int:
    """Write mapping records as literals while reusing an optional golden row."""
    start_row = int(start_row)
    if start_row <= 0 or not columns:
        raise TemplateWorkbookError("Vùng dữ liệu bảng không hợp lệ")
    normalized_columns = {str(key): int(column) for key, column in columns.items()}
    if any(column <= 0 for column in normalized_columns.values()):
        raise TemplateWorkbookError("Số cột bảng phải lớn hơn 0")
    last_row = start_row - 1
    for offset, record in enumerate(records):
        if not isinstance(record, Mapping):
            raise TemplateWorkbookError("Mỗi dòng bảng phải là một mapping")
        target_row = start_row + offset
        if prototype_row is not None and target_row != int(prototype_row):
            copy_row_layout(sheet, int(prototype_row), target_row, include_values=False)
        for key, column in normalized_columns.items():
            cell = sheet.cell(target_row, column)
            if isinstance(cell, MergedCell):
                raise TemplateWorkbookError(
                    f"Cột dữ liệu {key!r} trỏ vào ô phụ của vùng gộp {cell.coordinate}",
                    code="merged_cell_write",
                )
            _set_literal(cell, record.get(key))
        last_row = target_row
    return last_row


def assert_workbook_safe(workbook: Any) -> None:
    if getattr(workbook, "_external_links", None):
        raise TemplateWorkbookError(
            "Output còn external workbook link", code="external_link_remaining",
        )
    available = workbook.sheetnames
    for sheet in workbook.worksheets:
        for cells in sheet.iter_rows():
            for cell in cells:
                if isinstance(cell, MergedCell):
                    continue
                hyperlink = cell.hyperlink
                if hyperlink is not None and getattr(hyperlink, "target", None):
                    target = str(hyperlink.target)
                    if not target.startswith("#"):
                        raise TemplateWorkbookError(
                            f"Output còn hyperlink ngoài tại {sheet.title}!{cell.coordinate}",
                            code="external_hyperlink_remaining",
                        )
                if cell.data_type == "f" or (
                    isinstance(cell.value, str) and cell.value.startswith("=")
                ):
                    reason = unsafe_formula_reason(cell.value, available)
                    if reason:
                        raise TemplateWorkbookError(
                            f"Output còn công thức không an toàn tại {sheet.title}!{cell.coordinate}: {reason}",
                            code="unsafe_formula",
                        )
    for defined in workbook.defined_names.values():
        reason = _defined_name_reason(defined, available)
        if reason:
            raise TemplateWorkbookError(
                f"Output còn defined name không an toàn: {reason}",
                code="unsafe_defined_name",
            )


def safe_workbook_bytes(workbook: Any) -> bytes:
    try:
        from .document_preview import white_print_style
    except ImportError:
        from document_preview import white_print_style
    white_print_style(workbook)
    assert_workbook_safe(workbook)
    output = io.BytesIO()
    workbook.save(output)
    payload = output.getvalue()
    verification = None
    try:
        verification = load_workbook(
            io.BytesIO(payload), data_only=False, read_only=False, keep_links=False,
        )
        assert_workbook_safe(verification)
    except TemplateWorkbookError:
        raise
    except Exception as error:
        raise TemplateWorkbookError(
            f"Output template không thể mở lại: {type(error).__name__}",
            code="output_unreadable",
        ) from error
    finally:
        if verification is not None:
            verification.close()
    return payload


def workbook_topology_signature(workbook: Any) -> tuple[Any, ...]:
    """Return a stable topology signature without exposing business cell data."""
    sheets = []
    for sheet in workbook.worksheets:
        row_dimensions = tuple(sorted(
            (
                index, dimension.height, bool(dimension.hidden),
                dimension.outlineLevel, bool(dimension.collapsed),
            )
            for index, dimension in sheet.row_dimensions.items()
            if dimension.height is not None or dimension.hidden
            or dimension.outlineLevel or dimension.collapsed
        ))
        column_dimensions = tuple(sorted(
            (
                index, dimension.width, bool(dimension.hidden),
                dimension.outlineLevel, bool(dimension.collapsed),
            )
            for index, dimension in sheet.column_dimensions.items()
            if dimension.width is not None or dimension.hidden
            or dimension.outlineLevel or dimension.collapsed
        ))
        margins = tuple(
            getattr(sheet.page_margins, name)
            for name in ("left", "right", "top", "bottom", "header", "footer")
        )
        sheets.append((
            sheet.title, sheet.sheet_state, sheet.max_row, sheet.max_column,
            tuple(sorted(str(value) for value in sheet.merged_cells.ranges)),
            str(sheet.print_area or ""), str(sheet.print_title_rows or ""),
            str(sheet.print_title_cols or ""), str(sheet.freeze_panes or ""),
            sheet.page_setup.orientation, sheet.page_setup.paperSize,
            sheet.page_setup.scale, sheet.page_setup.fitToWidth,
            sheet.page_setup.fitToHeight, margins, row_dimensions, column_dimensions,
        ))
    return tuple(sheets)


__all__ = [
    "TemplateClone",
    "TemplateWorkbookError",
    "assert_workbook_safe",
    "clone_template_workbook",
    "copy_row_layout",
    "safe_workbook_bytes",
    "unsafe_formula_reason",
    "workbook_topology_signature",
    "write_formula",
    "write_literal",
    "write_table_rows",
]
