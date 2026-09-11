from __future__ import annotations

import io
import json
import tempfile
import unittest
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook

try:
    from . import server
    from .daily_workbook_import import analyze_daily_workbook
except ImportError:  # pragma: no cover - direct file invocation
    import server
    from daily_workbook_import import analyze_daily_workbook


class DailyWorkbookImportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        cls.original_db_path = server.DB_PATH
        cls.original_data_dir = server.DATA_DIR
        cls.original_master_source = server.MASTER_SOURCE
        cls.original_testing = server.app.config.get("TESTING", False)
        server.DATA_DIR = Path(cls.temp.name) / "data"
        server.DATA_DIR.mkdir(parents=True, exist_ok=True)
        server.DB_PATH = server.DATA_DIR / "daily.sqlite3"
        server.MASTER_SOURCE = Path(cls.temp.name) / "master-disabled.xlsx"
        server.app.config["TESTING"] = True
        server.init_database()
        cls.client = server.app.test_client()
        with server.db() as conn:
            conn.execute(
                "INSERT INTO contractors(code,name,price_group,pricing_mode) VALUES('C1','C1','C1','group')"
            )
            conn.execute("INSERT INTO kitchens(code,contractor,name) VALUES('K1','C1','Kitchen 1')")
            conn.execute("INSERT INTO suppliers(code,name) VALUES('S1','Supplier 1')")
            conn.execute(
                """INSERT INTO products(
                       code,name,unit,tax,supplier,buy_price,purchase_list,seller,cccd
                   ) VALUES('P1','Product 1','kg','0%','S1',10000,0,'','')"""
            )

    @classmethod
    def tearDownClass(cls):
        with server.PENDING_IMPORT_LOCK:
            for pending in server.PENDING_IMPORTS.values():
                pending["path"].unlink(missing_ok=True)
            server.PENDING_IMPORTS.clear()
        server.DB_PATH = cls.original_db_path
        server.DATA_DIR = cls.original_data_dir
        server.MASTER_SOURCE = cls.original_master_source
        server.app.config["TESTING"] = cls.original_testing
        cls.temp.cleanup()

    def setUp(self):
        with server.PENDING_IMPORT_LOCK:
            for pending in server.PENDING_IMPORTS.values():
                pending["path"].unlink(missing_ok=True)
            server.PENDING_IMPORTS.clear()
        with server.db() as conn:
            conn.execute("DROP TRIGGER IF EXISTS fail_daily_order_insert")
            conn.execute("DELETE FROM receivable_ledger_revisions")
            conn.execute("DELETE FROM receivable_ledger_lines")
            conn.execute("DELETE FROM order_import_receipts")
            conn.execute("DELETE FROM batches")
            conn.execute("UPDATE products SET buy_price=10000 WHERE code='P1'")

    @staticmethod
    def workbook_bytes(*, quantities=(2, 3), secret="CCCD-NEVER-EXPOSE") -> bytes:
        workbook = Workbook()
        identity = workbook.active
        identity.title = "CCCD"
        identity.append(["Số CCCD", "Người bán"])
        identity.append([secret, "PRIVATE PERSON"])

        day_sheet = workbook.create_sheet("01.09")
        day_sheet.append(["ĐƠN HÀNG NGÀY"])
        day_sheet.append([
            "Ngày", "Nhà thầu", "Mã bếp", "Mã hàng", "Tên hàng", "Số lượng",
            "ĐVT", "Chọn NCC", "Giá mua", "Giá bán", "CCCD",
        ])
        for qty in quantities:
            day_sheet.append([
                date(2026, 9, 1), "C1", "K1", "P1", "Product 1", qty,
                "kg", None, 10_000, 12_000, None,
            ])

        merged = workbook.create_sheet("gộp đơn")
        merged.append([
            "Ngày", "Nhà thầu", "Mã bếp", "Mã hàng", "Tên hàng", "Số lượng", "Giá bán",
        ])
        merged.append([date(2026, 9, 1), "FAKE", "FAKE", "FAKE", secret, 999, 999])

        purchase = workbook.create_sheet("đặt hàng")
        purchase.append(["ĐẶT NHÀ CUNG CẤP"])
        purchase.append([
            "Mã hàng", "Mã bếp", "Tên hàng", "Số lượng", "ĐVT", "NCC", "ghi chú",
            "giá mua", "hỏng", "thêm", "Giảm", "thiếu", "SL thực tế", "Thành tiền",
        ])
        purchase.append([
            "P1", "K1", "Product 1", sum(quantities), "kg", None, None,
            10_000, 0, 0, 0, 0, sum(quantities), sum(quantities) * 10_000,
        ])
        stream = io.BytesIO()
        workbook.save(stream)
        workbook.close()
        return stream.getvalue()

    @staticmethod
    def finalization_workbook_bytes(
        *, actual_delivered=10, sell_price=12000, damaged=0, added=0,
        extra_qty=None,
    ) -> bytes:
        workbook = Workbook()
        day_sheet = workbook.active
        day_sheet.title = "01.09"
        day_sheet.append(["ĐƠN HÀNG NGÀY"])
        day_sheet.append([
            "Ngày", "Nhà thầu", "Mã bếp", "Mã hàng", "Tên hàng", "Số lượng",
            "Thực giao", "ĐVT", "Chọn NCC", "Giá mua", "Giá bán",
        ])
        day_sheet.append([
            date(2026, 9, 1), "C1", "K1", "P1", "Product 1", 10,
            actual_delivered, "kg", "S1", 0, sell_price,
        ])
        if extra_qty is not None:
            day_sheet.append([
                date(2026, 9, 1), "C1", "K1", "P1", "Product 1", extra_qty,
                extra_qty, "kg", "S1", 0, sell_price,
            ])
        purchase = workbook.create_sheet("đặt hàng")
        purchase.append(["ĐẶT NHÀ CUNG CẤP"])
        purchase.append([
            "Mã hàng", "Mã bếp", "Ngày", "Tên hàng", "Số lượng", "ĐVT", "NCC",
            "ghi chú", "giá mua", "hỏng", "thêm", "Giảm", "thiếu",
            "SL thực tế", "Thành tiền",
        ])
        actual = 10 + added - damaged
        purchase.append([
            "P1", "K1", date(2026, 9, 1), "Product 1", 10, "kg", "S1",
            "Chốt mua", 9000, damaged, added, 0, 0, actual, actual * 9000,
        ])
        if extra_qty is not None:
            purchase.append([
                "P1", "K1", date(2026, 9, 1), "Product 1", extra_qty, "kg", "S1",
                "Chốt mua thêm", 9000, 0, 0, 0, 0, extra_qty, extra_qty * 9000,
            ])
        stream = io.BytesIO()
        workbook.save(stream)
        workbook.close()
        return stream.getvalue()

    def test_new_product_keeps_vietnamese_unit_and_reports_missing_catalog(self):
        workbook = load_workbook(io.BytesIO(self.workbook_bytes()))
        sheet = workbook['01.09']
        sheet.cell(3, 4, 'NEW-CAN')
        sheet.cell(3, 5, 'New canned product')
        sheet.cell(3, 7, 'Can')
        path = Path(self.temp.name) / 'new-product.xlsx'
        workbook.save(path)
        workbook.close()
        rows, _ = server.parse_workbook(path, '2026-09-01', selected_sheets=['01.09'])
        row = rows[0]
        self.assertEqual((row['product_code'], row['unit'], row['qty']), ('NEW-CAN', 'Can', 2))
        self.assertIn('Mã hàng chưa có trong danh mục: NEW-CAN', row['errors'])
        with server.db() as conn:
            empty_unit = server.resolve_order(conn, {**row, 'unit': ''}, '2026-09-01', *server.product_lookup(conn))
        self.assertIn('Thiếu đơn vị tính', empty_unit['errors'])

    def analyze_api(self, payload: bytes, filename="orders.xlsx") -> dict:
        response = self.client.post(
            "/api/import/analyze",
            data={"file": (io.BytesIO(payload), filename)},
            content_type="multipart/form-data",
        )
        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        return response.get_json()

    def confirm_api(self, analyzed: dict, *, sheets=None, state_hash=None):
        return self.client.post(
            "/api/import/confirm",
            json={
                "token": analyzed["token"],
                "sheets": ["01.09"] if sheets is None else sheets,
                "work_date": "2026-09-01",
                "state_hash": analyzed.get("stateHash") if state_hash is None else state_hash,
            },
        )

    def counts(self):
        with server.db() as conn:
            return {
                table: int(conn.execute(f"SELECT COUNT(*) n FROM {table}").fetchone()["n"])
                for table in (
                    "batches", "orders", "order_import_receipts", "daily_workdays",
                    "daily_import_versions", "daily_import_scopes", "daily_import_rows",
                )
            }

    def continuous_file(self, *, qty=2, complete=True):
        workbook = load_workbook(io.BytesIO(self.workbook_bytes(quantities=(qty,))))
        sheet = workbook['01.09']
        sheet['H2'] = 'NCC'
        sheet['H3'] = 'S1'
        sheet['J3'] = 12000 if complete else 0
        stream = io.BytesIO(); workbook.save(stream); workbook.close()
        return stream.getvalue()

    def continuous_analyze(self, data):
        response = self.client.post('/api/import/analyze', data={
            'file': (io.BytesIO(data), 'orders-01.09.xlsx'), 'continuous': '1'}, content_type='multipart/form-data')
        self.assertEqual(response.status_code, 200, response.json)
        return response.json

    def test_continuous_day_accepts_incomplete_first_then_three_revisions(self):
        first = self.confirm_api(self.continuous_analyze(self.continuous_file(complete=False)))
        self.assertEqual(first.status_code, 200, first.json)
        self.assertGreater(first.json['summary']['totals']['errors'], 0)
        batch_id = first.json['batch']['id']
        for qty in (4, 7, 9):
            analysis = self.continuous_analyze(self.continuous_file(qty=qty))
            self.assertTrue(analysis['sheets'][0]['confirmAvailable'], analysis)
            result = self.confirm_api(analysis)
            self.assertEqual(result.status_code, 200, result.json)
            self.assertEqual(result.json['batch']['id'], batch_id)
            self.assertEqual(result.json['batch']['status'], 'draft')
            self.assertEqual(len(result.json['orders']), 1)
            self.assertEqual(result.json['orders'][0]['qty'], qty)
            self.assertEqual(result.json['orders'][0]['supplier'], 'S1')
            self.assertNotEqual(result.json['orders'][0]['note'], 'edited on web')
            with server.db() as conn:
                self.assertEqual(conn.execute('SELECT lifecycle_status FROM daily_workdays WHERE batch_id=?', (batch_id,)).fetchone()[0], 'picking')
                conn.execute("UPDATE orders SET note='edited on web',supplier='web supplier' WHERE batch_id=?", (batch_id,))
        with server.db() as conn:
            history = conn.execute('SELECT rows_json FROM order_reimport_history WHERE batch_id=? ORDER BY id', (batch_id,)).fetchall()
            self.assertEqual(len(history), 3)
            self.assertIn('edited on web', history[-1][0])
            self.assertEqual(conn.execute('SELECT COUNT(*) FROM invoice_inventory_ledger').fetchone()[0], 0)

    def test_continuous_invalid_latest_preserves_current_and_duplicate_is_noop(self):
        data = self.continuous_file(qty=3)
        first = self.confirm_api(self.continuous_analyze(data))
        self.assertEqual(first.status_code, 200, first.json)
        repeated = self.confirm_api(self.continuous_analyze(data))
        self.assertEqual(repeated.status_code, 200, repeated.json)
        self.assertTrue(repeated.json['idempotent'])
        invalid = self.continuous_analyze(self.continuous_file(qty=4, complete=False))
        self.assertFalse(invalid['sheets'][0]['confirmAvailable'])
        details=invalid['sheets'][0]['issueDetails']
        self.assertTrue(details)
        self.assertEqual(details[0]['row'],3)
        self.assertTrue(details[0]['errors'])
        self.assertNotIn('cccd',details[0])
        self.assertEqual(self.confirm_api(invalid).status_code, 400)
        with server.db() as conn:
            self.assertEqual(conn.execute('SELECT qty FROM orders WHERE batch_id=?', (first.json['batch']['id'],)).fetchone()[0], 3)

    def test_continuous_new_file_restores_note_even_when_other_values_match(self):
        import zipfile
        data = self.continuous_file(qty=3)
        first = self.confirm_api(self.continuous_analyze(data)).json
        with server.db() as conn:
            conn.execute("UPDATE orders SET note='web only' WHERE batch_id=?", (first['batch']['id'],))
        stream = io.BytesIO(data)
        with zipfile.ZipFile(stream, 'a') as archive:
            archive.comment = b'Later revision with identical business cells'
        result = self.confirm_api(self.continuous_analyze(stream.getvalue()))
        self.assertEqual(result.status_code, 200, result.json)
        self.assertEqual(result.json['orders'][0]['note'], first['orders'][0]['note'])

    def test_structure_first_analyzer_never_serializes_reference_values(self):
        secret = "CCCD-DO-NOT-SERIALIZE"
        path = Path(self.temp.name) / "structure.xlsx"
        path.write_bytes(self.workbook_bytes(secret=secret))
        result = analyze_daily_workbook(path)
        self.assertTrue(result["strictCustomerWorkbook"])
        self.assertEqual(
            [(item["name"], item["rows"], item["confirmAvailable"]) for item in result["daySheets"]],
            [("01.09", 2, True)],
        )
        self.assertEqual(
            [(item["name"], item["rows"], item["confirmAvailable"]) for item in result["purchaseSheets"]],
            [("đặt hàng", 1, False)],
        )
        self.assertTrue(any(
            item["name"] == "gộp đơn" and item["role"] == "intermediate_reference"
            for item in result["ignoredSheets"]
        ))
        self.assertNotIn(secret, json.dumps(result, ensure_ascii=False))

    def test_preview_exposes_only_safe_counts_scopes_and_hash(self):
        secret = "CCCD-NOT-IN-API"
        analyzed = self.analyze_api(self.workbook_bytes(secret=secret))
        self.assertTrue(analyzed["strictDaily"])
        self.assertEqual(analyzed["detectedWorkDate"], "2026-09-01")
        self.assertEqual(len(analyzed["stateHash"]), 64)
        self.assertEqual([item["name"] for item in analyzed["sheets"]], ["01.09", "đặt hàng"])
        daily, purchase = analyzed["sheets"]
        self.assertEqual(
            (daily["rows"], daily["parsedRows"], daily["errorRows"], daily["writeScope"]),
            (2, 2, 0, "customer_orders"),
        )
        self.assertFalse(purchase["confirmAvailable"])
        self.assertEqual(purchase["writeScope"], "purchase_orders_locked_until_round_trip")
        self.assertNotIn(secret, json.dumps(analyzed, ensure_ascii=False))

    def test_dated_filename_scopes_a_multi_day_workbook_to_one_exact_sheet(self):
        workbook = load_workbook(io.BytesIO(self.workbook_bytes()))
        prior = workbook.copy_worksheet(workbook["01.09"])
        prior.title = "31.08"
        for row in range(3, prior.max_row + 1):
            prior.cell(row, 1).value = date(2026, 8, 31)
        stream = io.BytesIO()
        workbook.save(stream)
        workbook.close()

        analyzed = self.analyze_api(stream.getvalue(), filename="Đơn hàng 01.09.2026.xlsx")
        self.assertEqual("2026-09-01", analyzed["detectedWorkDate"])
        self.assertEqual(["01.09", "đặt hàng"], [item["name"] for item in analyzed["sheets"]])
        self.assertTrue(any(
            item["name"] == "31.08" and item["role"] == "historical_daily_sheet"
            for item in analyzed["ignoredSheets"]
        ))
        confirmed = self.confirm_api(analyzed)
        self.assertEqual(200, confirmed.status_code, confirmed.get_data(as_text=True))

    def test_state_hash_and_scope_failures_consume_the_token(self):
        analyzed = self.analyze_api(self.workbook_bytes())
        stale = self.confirm_api(analyzed, state_hash="0" * 64)
        self.assertEqual(stale.status_code, 409, stale.get_data(as_text=True))
        self.assertEqual(stale.get_json()["code"], "stale_preview")
        replay = self.confirm_api(analyzed)
        self.assertEqual(replay.status_code, 400, replay.get_data(as_text=True))

        analyzed = self.analyze_api(self.workbook_bytes())
        wrong_scope = self.confirm_api(analyzed, sheets=["đặt hàng"])
        self.assertEqual(wrong_scope.status_code, 400, wrong_scope.get_data(as_text=True))
        self.assertEqual(wrong_scope.get_json()["code"], "scope_not_confirmable")
        replay = self.confirm_api(analyzed)
        self.assertEqual(replay.status_code, 400, replay.get_data(as_text=True))
        self.assertEqual(self.counts()["batches"], 0)

    def test_confirm_is_atomic_and_filename_independent(self):
        workbook = self.workbook_bytes()
        first_analysis = self.analyze_api(workbook, "first-name.xlsx")
        first = self.confirm_api(first_analysis)
        self.assertEqual(first.status_code, 200, first.get_data(as_text=True))
        first_payload = first.get_json()
        self.assertFalse(first_payload["idempotent"])
        self.assertEqual(len(first_payload["orders"]), 2)
        self.assertEqual(first_payload["dailyImport"]["scopes"][0]["state"], "confirmed")

        renamed_analysis = self.analyze_api(workbook, "renamed-same-bytes.xlsx")
        replay = self.confirm_api(renamed_analysis)
        self.assertEqual(replay.status_code, 200, replay.get_data(as_text=True))
        replay_payload = replay.get_json()
        self.assertTrue(replay_payload["idempotent"])
        self.assertEqual(replay_payload["batch"]["id"], first_payload["batch"]["id"])
        self.assertEqual(replay_payload["importKey"], first_payload["importKey"])
        self.assertEqual(
            self.counts(),
            {
                "batches": 1, "orders": 2, "order_import_receipts": 1,
                "daily_workdays": 1, "daily_import_versions": 1,
                "daily_import_scopes": 1, "daily_import_rows": 2,
            },
        )

    def test_changed_workbook_reuses_workday_batch_and_replaces_scope(self):
        first = self.confirm_api(self.analyze_api(self.workbook_bytes(quantities=(2, 3))))
        changed_analysis = self.analyze_api(self.workbook_bytes(quantities=(7,)))
        self.assertEqual(changed_analysis["phase"], "finalization")
        self.assertEqual(changed_analysis["sheets"][0]["diff"], {
            "added": 0, "updated": 1, "unchanged": 0,
            "removed": 1, "conflicts": 0, "total": 1,
        })
        changed = self.confirm_api(changed_analysis)
        self.assertEqual(first.status_code, 200, first.get_data(as_text=True))
        self.assertEqual(changed.status_code, 200, changed.get_data(as_text=True))
        self.assertEqual(first.get_json()["batch"]["id"], changed.get_json()["batch"]["id"])
        self.assertEqual(len(changed.get_json()["orders"]), 1)
        counts = self.counts()
        self.assertEqual(counts["batches"], 1)
        self.assertEqual(counts["orders"], 1)
        self.assertEqual(counts["daily_import_versions"], 2)
        self.assertEqual(counts["order_import_receipts"], 1)
        with server.db() as conn:
            receivable_rows = [dict(row) for row in conn.execute(
                "SELECT source_key,source_hash,status,reversal_reason "
                "FROM receivable_ledger_lines ORDER BY id"
            )]
            receivable_revisions = conn.execute(
                "SELECT COUNT(*) n FROM receivable_ledger_revisions"
            ).fetchone()["n"]
        self.assertEqual(len(receivable_rows), 2)
        self.assertEqual(len({row["source_key"] for row in receivable_rows}), 2)
        self.assertEqual(
            sorted(row["reversal_reason"] for row in receivable_rows),
            ["batch_not_approved", "source_removed"],
        )
        current_receivable = next(
            row for row in receivable_rows if row["reversal_reason"] == "batch_not_approved"
        )
        self.assertEqual(current_receivable["source_hash"], changed.get_json()["sourceHash"])
        self.assertEqual(receivable_revisions, 5)

    def test_second_load_confirms_purchase_and_sales_as_independent_scopes(self):
        with server.db() as conn:
            conn.execute("UPDATE products SET buy_price=0 WHERE code='P1'")
            price_book_before = [tuple(row) for row in conn.execute(
                "SELECT * FROM product_prices ORDER BY product_code,price_group"
            )]

        first_bytes = self.finalization_workbook_bytes()
        first_analysis = self.analyze_api(first_bytes, "night-before.xlsx")
        self.assertEqual(first_analysis["phase"], "first_load")
        self.assertFalse(first_analysis["scopeSelectionRequired"])
        first = self.confirm_api(first_analysis, sheets=["01.09"])
        self.assertEqual(first.status_code, 200, first.get_data(as_text=True))
        batch_id = first.get_json()["batch"]["id"]
        self.assertEqual(first.get_json()["summary"]["totals"]["revenue"], 120000)
        with server.db() as conn:
            conn.execute(
                "UPDATE batches SET status='approved',approved_at=? WHERE id=?",
                (server.now_iso(), batch_id),
            )
            server.sync_receivable_ledger(conn, timestamp=server.now_iso())

        final_bytes = self.finalization_workbook_bytes(
            actual_delivered=8, sell_price=16000, damaged=1, added=2,
        )
        no_scope_analysis = self.analyze_api(final_bytes, "next-day-no-scope.xlsx")
        no_scope = self.confirm_api(no_scope_analysis, sheets=[])
        self.assertEqual((no_scope.status_code, no_scope.get_json()["code"]), (
            400, "scope_selection_required",
        ))
        analyzed = self.analyze_api(final_bytes, "next-day-final.xlsx")
        self.assertEqual((analyzed["phase"], analyzed["batchId"]), ("finalization", batch_id))
        self.assertTrue(analyzed["scopeSelectionRequired"])
        day_scope, purchase_scope = analyzed["sheets"]
        self.assertTrue(day_scope["confirmAvailable"])
        self.assertTrue(purchase_scope["confirmAvailable"])
        self.assertEqual(day_scope["diff"], {
            "added": 0, "updated": 1, "unchanged": 0,
            "removed": 0, "conflicts": 0, "total": 1,
        })
        self.assertEqual(purchase_scope["diff"], {
            "added": 1, "updated": 0, "unchanged": 0,
            "removed": 0, "conflicts": 0, "total": 1,
        })

        # Confirm only purchase first. Sales/revenue/receivable must remain at
        # the first-load values while payable is created independently.
        purchase_only = self.confirm_api(analyzed, sheets=["đặt hàng"])
        self.assertEqual(purchase_only.status_code, 200, purchase_only.get_data(as_text=True))
        purchase_payload = purchase_only.get_json()
        self.assertEqual(purchase_payload["selectedScopes"], ["purchase_orders"])
        self.assertIsNone(purchase_payload["finalized"])
        self.assertEqual(purchase_payload["summary"]["totals"]["revenue"], 120000)
        debts = self.client.get("/api/debts?from=2026-09-01&to=2026-09-01").get_json()
        self.assertEqual(debts["contractors"]["C1"]["period_charge"], 120000)
        self.assertEqual(debts["suppliers"]["S1"]["period_charge"], 99000)
        with server.db() as conn:
            order = dict(conn.execute(
                "SELECT * FROM orders WHERE batch_id=?", (batch_id,)
            ).fetchone())
            purchase_line = dict(conn.execute(
                "SELECT * FROM purchase_workbook_lines WHERE batch_id=?", (batch_id,)
            ).fetchone())
        self.assertEqual((order["actual_delivered"], order["sell_price"], order["buy_price"]), (10, 12000, 0))
        self.assertEqual((purchase_line["actual_qty"], purchase_line["buy_price"], purchase_line["amount"]), (11, 9000, 99000))

        # A concurrent database change makes the preview stale and consumes
        # its token without partially applying the sales scope.
        stale_analysis = self.analyze_api(final_bytes, "next-day-final-stale.xlsx")
        with server.db() as conn:
            conn.execute("UPDATE orders SET note='concurrent edit' WHERE batch_id=?", (batch_id,))
        stale = self.confirm_api(stale_analysis, sheets=["01.09"])
        self.assertEqual((stale.status_code, stale.get_json()["code"]), (409, "stale_database_state"))
        consumed = self.confirm_api(stale_analysis, sheets=["01.09"])
        self.assertEqual(consumed.status_code, 400)
        with server.db() as conn:
            unchanged = conn.execute(
                "SELECT actual_delivered,sell_price,note FROM orders WHERE batch_id=?",
                (batch_id,),
            ).fetchone()
        self.assertEqual(tuple(unchanged), (10, 12000, "concurrent edit"))

        # Re-preview and confirm only sales. Existing order id is preserved,
        # purchase ledger is untouched, and all scopes now finalize the day.
        sales_analysis = self.analyze_api(final_bytes, "next-day-final-retry.xlsx")
        sales_only = self.confirm_api(sales_analysis, sheets=["01.09"])
        self.assertEqual(sales_only.status_code, 200, sales_only.get_data(as_text=True))
        sales_payload = sales_only.get_json()
        self.assertEqual(sales_payload["selectedScopes"], ["customer_orders"])
        self.assertEqual(sales_payload["finalized"]["lifecycle_status"], "finalized")
        self.assertEqual(sales_payload["summary"]["totals"]["revenue"], 128000)
        debts = self.client.get("/api/debts?from=2026-09-01&to=2026-09-01").get_json()
        self.assertEqual(debts["contractors"]["C1"]["period_charge"], 128000)
        self.assertEqual(debts["suppliers"]["S1"]["period_charge"], 99000)
        with server.db() as conn:
            final_order = dict(conn.execute(
                "SELECT * FROM orders WHERE batch_id=?", (batch_id,)
            ).fetchone())
            final_purchase = dict(conn.execute(
                "SELECT * FROM purchase_workbook_lines WHERE batch_id=?", (batch_id,)
            ).fetchone())
            states = {
                row["scope"]: row["state"] for row in conn.execute(
                    """SELECT scope,state FROM daily_import_scopes
                       WHERE version_id=(SELECT MAX(id) FROM daily_import_versions)"""
                )
            }
            price_book_after = [tuple(row) for row in conn.execute(
                "SELECT * FROM product_prices ORDER BY product_code,price_group"
            )]
        self.assertEqual(final_order["id"], order["id"])
        self.assertEqual((final_order["actual_delivered"], final_order["sell_price"]), (8, 16000))
        self.assertEqual((final_order["supplier"], final_order["buy_price"]), ("S1", 0))
        self.assertEqual((final_purchase["actual_qty"], final_purchase["amount"]), (11, 99000))
        self.assertEqual(states, {"customer_orders": "confirmed", "purchase_orders": "confirmed"})
        self.assertEqual(price_book_after, price_book_before)

        # Exact replay can select both scopes and must not duplicate either ledger.
        replay_analysis = self.analyze_api(final_bytes, "same-final-renamed.xlsx")
        replay = self.confirm_api(replay_analysis, sheets=["01.09", "đặt hàng"])
        self.assertEqual(replay.status_code, 200, replay.get_data(as_text=True))
        self.assertTrue(replay.get_json()["idempotent"])
        with server.db() as conn:
            self.assertEqual(conn.execute(
                "SELECT COUNT(*) FROM orders WHERE batch_id=?", (batch_id,)
            ).fetchone()[0], 1)
            self.assertEqual(conn.execute(
                "SELECT COUNT(*) FROM purchase_workbook_lines WHERE batch_id=?", (batch_id,)
            ).fetchone()[0], 1)
            self.assertEqual(conn.execute(
                "SELECT COUNT(*) FROM purchase_workbook_line_revisions WHERE batch_id=?",
                (batch_id,),
            ).fetchone()[0], 1)

    def test_second_load_reports_conflict_instead_of_deleting_linked_order(self):
        first = self.confirm_api(self.analyze_api(self.workbook_bytes(quantities=(2, 3))))
        self.assertEqual(first.status_code, 200, first.get_data(as_text=True))
        batch_id = first.get_json()["batch"]["id"]
        with server.db() as conn:
            linked_order = conn.execute(
                "SELECT * FROM orders WHERE batch_id=? ORDER BY source_row,id LIMIT 1 OFFSET 1",
                (batch_id,),
            ).fetchone()
            timestamp = server.now_iso()
            conn.execute(
                """INSERT INTO purchase_workbook_lines(
                       batch_id,row_key,order_id,source_sheet,source_row,product_code,kitchen,
                       work_date,product_name,base_qty,unit,supplier,buy_price,actual_qty,
                       amount,source_hash,created_at,updated_at
                   ) VALUES(?, ?,?,'đặt hàng',4,'P1','K1','2026-09-01','Product 1',
                            3,'kg','S1',10000,3,30000,'fixture',?,?)""",
                (batch_id, "A" * 64, int(linked_order["id"]), timestamp, timestamp),
            )
        analyzed = self.analyze_api(self.workbook_bytes(quantities=(7,)))
        day_scope = analyzed["sheets"][0]
        self.assertFalse(day_scope["confirmAvailable"])
        self.assertEqual(day_scope["previewIssue"], "customer_scope_conflict")
        self.assertEqual(day_scope["diff"], {
            "added": 0, "updated": 1, "unchanged": 0,
            "removed": 0, "conflicts": 1, "total": 1,
        })
        blocked = self.confirm_api(analyzed, sheets=["01.09"])
        self.assertEqual((blocked.status_code, blocked.get_json()["code"]), (
            400, "scope_not_confirmable",
        ))
        with server.db() as conn:
            self.assertEqual(conn.execute(
                "SELECT COUNT(*) FROM orders WHERE batch_id=?", (batch_id,)
            ).fetchone()[0], 2)

    def test_two_scope_confirm_links_new_purchase_row_to_new_customer_order(self):
        with server.db() as conn:
            conn.execute("UPDATE products SET buy_price=0 WHERE code='P1'")
        first = self.confirm_api(self.analyze_api(self.finalization_workbook_bytes()))
        self.assertEqual(first.status_code, 200, first.get_data(as_text=True))
        batch_id = first.get_json()["batch"]["id"]

        final_bytes = self.finalization_workbook_bytes(extra_qty=2)
        analyzed = self.analyze_api(final_bytes, "final-with-new-row.xlsx")
        day_scope, purchase_scope = analyzed["sheets"]
        self.assertEqual(day_scope["diff"], {
            "added": 1, "updated": 0, "unchanged": 1,
            "removed": 0, "conflicts": 0, "total": 2,
        })
        self.assertEqual(purchase_scope["diff"], {
            "added": 2, "updated": 0, "unchanged": 0,
            "removed": 0, "conflicts": 0, "total": 2,
        })

        confirmed = self.confirm_api(analyzed, sheets=["01.09", "đặt hàng"])
        self.assertEqual(confirmed.status_code, 200, confirmed.get_data(as_text=True))
        payload = confirmed.get_json()
        self.assertEqual(payload["selectedScopes"], ["customer_orders", "purchase_orders"])
        self.assertEqual(payload["finalized"]["lifecycle_status"], "finalized")
        self.assertEqual(payload["summary"]["totals"]["revenue"], 144000)
        with server.db() as conn:
            order_ids = {
                int(row["id"]) for row in conn.execute(
                    "SELECT id FROM orders WHERE batch_id=?", (batch_id,)
                )
            }
            purchase_order_ids = {
                int(row["order_id"]) for row in conn.execute(
                    "SELECT order_id FROM purchase_workbook_lines WHERE batch_id=?",
                    (batch_id,),
                )
                if row["order_id"] is not None
            }
            self.assertEqual(conn.execute(
                "SELECT COUNT(*) FROM purchase_workbook_lines WHERE batch_id=?",
                (batch_id,),
            ).fetchone()[0], 2)
        self.assertEqual(len(order_ids), 2)
        self.assertEqual(purchase_order_ids, order_ids)

        replay = self.confirm_api(
            self.analyze_api(final_bytes, "same-final-with-new-row.xlsx"),
            sheets=["01.09", "đặt hàng"],
        )
        self.assertEqual(replay.status_code, 200, replay.get_data(as_text=True))
        self.assertTrue(replay.get_json()["idempotent"])
        with server.db() as conn:
            self.assertEqual(conn.execute(
                "SELECT COUNT(*) FROM orders WHERE batch_id=?", (batch_id,)
            ).fetchone()[0], 2)
            self.assertEqual(conn.execute(
                "SELECT COUNT(*) FROM purchase_workbook_lines WHERE batch_id=?",
                (batch_id,),
            ).fetchone()[0], 2)

    def test_sales_scope_new_row_does_not_create_supplier_payable(self):
        first = self.confirm_api(self.analyze_api(self.finalization_workbook_bytes()))
        self.assertEqual(first.status_code, 200, first.get_data(as_text=True))
        batch_id = first.get_json()["batch"]["id"]
        approved = self.client.post(f"/api/batches/{batch_id}/approve")
        self.assertEqual(approved.status_code, 200, approved.get_data(as_text=True))
        before = self.client.get(
            "/api/debts?from=2026-09-01&to=2026-09-01"
        ).get_json()["suppliers"].get("S1", {}).get("period_charge", 0)

        analyzed = self.analyze_api(
            self.finalization_workbook_bytes(extra_qty=2), "sales-only-new-row.xlsx"
        )
        confirmed = self.confirm_api(analyzed, sheets=["01.09"])
        self.assertEqual(confirmed.status_code, 200, confirmed.get_data(as_text=True))
        after = self.client.get(
            "/api/debts?from=2026-09-01&to=2026-09-01"
        ).get_json()["suppliers"].get("S1", {}).get("period_charge", 0)
        self.assertEqual(after, before)
        with server.db() as conn:
            new_order = conn.execute(
                """SELECT actual_received,damaged_qty,supplier_return_qty,supplier,
                          buy_price,purchase_list
                   FROM orders WHERE batch_id=? ORDER BY source_row DESC,id DESC LIMIT 1""",
                (batch_id,),
            ).fetchone()
            self.assertEqual(tuple(new_order), (0, 0, 0, "", 0, 0))
            self.assertEqual(conn.execute(
                "SELECT COUNT(*) FROM purchase_workbook_lines WHERE batch_id=?", (batch_id,)
            ).fetchone()[0], 0)

    def test_two_scope_failure_rolls_back_sales_purchase_and_lifecycle_together(self):
        with server.db() as conn:
            conn.execute("UPDATE products SET buy_price=0 WHERE code='P1'")
        first = self.confirm_api(self.analyze_api(self.finalization_workbook_bytes()))
        self.assertEqual(first.status_code, 200, first.get_data(as_text=True))
        batch_id = first.get_json()["batch"]["id"]
        final_bytes = self.finalization_workbook_bytes(
            actual_delivered=8, sell_price=16000, damaged=1, added=2,
        )
        analyzed = self.analyze_api(final_bytes)
        with server.db() as conn:
            conn.execute(
                """CREATE TRIGGER fail_daily_purchase_scope
                   BEFORE INSERT ON purchase_workbook_lines
                   BEGIN SELECT RAISE(ABORT,'forced daily purchase failure'); END"""
            )
        failed = self.confirm_api(analyzed, sheets=["01.09", "đặt hàng"])
        self.assertEqual(failed.status_code, 500, failed.get_data(as_text=True))
        with server.db() as conn:
            conn.execute("DROP TRIGGER fail_daily_purchase_scope")
            order = conn.execute(
                "SELECT actual_delivered,sell_price FROM orders WHERE batch_id=?",
                (batch_id,),
            ).fetchone()
            self.assertEqual(tuple(order), (10, 12000))
            self.assertEqual(conn.execute(
                "SELECT COUNT(*) FROM purchase_workbook_lines WHERE batch_id=?", (batch_id,)
            ).fetchone()[0], 0)
            self.assertEqual(conn.execute(
                "SELECT COUNT(*) FROM daily_import_versions"
            ).fetchone()[0], 1)
            self.assertEqual(conn.execute(
                "SELECT COUNT(*) FROM daily_import_scopes"
            ).fetchone()[0], 1)

    def test_order_insert_failure_rolls_back_lifecycle_and_receipt(self):
        with server.db() as conn:
            conn.execute(
                """CREATE TRIGGER fail_daily_order_insert BEFORE INSERT ON orders
                   BEGIN SELECT RAISE(ABORT,'forced order failure'); END"""
            )
        analyzed = self.analyze_api(self.workbook_bytes())
        failed = self.confirm_api(analyzed)
        self.assertEqual(failed.status_code, 500, failed.get_data(as_text=True))
        with server.db() as conn:
            conn.execute("DROP TRIGGER fail_daily_order_insert")
        self.assertEqual(
            self.counts(),
            {
                "batches": 0, "orders": 0, "order_import_receipts": 0,
                "daily_workdays": 0, "daily_import_versions": 0,
                "daily_import_scopes": 0, "daily_import_rows": 0,
            },
        )


if __name__ == "__main__":
    unittest.main()
