from __future__ import annotations

import io
import sqlite3
import tempfile
import time
import unittest
import zipfile
from contextlib import contextmanager
from datetime import date
from pathlib import Path

from flask import Flask
from openpyxl import Workbook, load_workbook

try:
    from .contract_modules import (
        MAPPING_IMPORT_TTL_SECONDS,
        PENDING_LEGACY_ATTENDANCE_IMPORTS,
        apply_legacy_attendance_snapshot,
        import_legacy_attendance,
        parse_legacy_attendance_workbooks,
        register_contract_routes,
        validate_legacy_attendance_period,
    )
except ImportError:  # pragma: no cover - direct file invocation
    from contract_modules import (
        MAPPING_IMPORT_TTL_SECONDS,
        PENDING_LEGACY_ATTENDANCE_IMPORTS,
        apply_legacy_attendance_snapshot,
        import_legacy_attendance,
        parse_legacy_attendance_workbooks,
        register_contract_routes,
        validate_legacy_attendance_period,
    )


SYSTEM_DIR = Path(__file__).resolve().parent
PROJECT_DIR = SYSTEM_DIR.parent
REAL_T8_FILE = (
    PROJECT_DIR / "bosung.30.8.26" / "CHẤM CÔNG+ SUẤT ĂN  2026" / "CHẤM CÔNG T8.2026.xlsx"
)


SCHEMA = """
CREATE TABLE audit_log(
    id INTEGER PRIMARY KEY AUTOINCREMENT,event_type TEXT,entity_type TEXT,entity_id TEXT,
    status TEXT,message TEXT,metadata_json TEXT,created_at TEXT
);
CREATE TABLE staff(
    id INTEGER PRIMARY KEY AUTOINCREMENT,employee_code TEXT NOT NULL UNIQUE,
    full_name TEXT NOT NULL,role_name TEXT,kitchen TEXT,base_salary REAL NOT NULL DEFAULT 0,
    standard_days REAL NOT NULL DEFAULT 26,standard_hours REAL NOT NULL DEFAULT 8,
    bhxh_employee_rate REAL NOT NULL DEFAULT 0,bhxh_company_rate REAL NOT NULL DEFAULT 0,
    active INTEGER NOT NULL DEFAULT 1,created_at TEXT NOT NULL,updated_at TEXT NOT NULL
);
CREATE TABLE attendance_entries(
    id INTEGER PRIMARY KEY AUTOINCREMENT,employee_id INTEGER NOT NULL,
    work_date TEXT NOT NULL,normal_hours REAL NOT NULL DEFAULT 0,
    overtime_hours REAL NOT NULL DEFAULT 0,sunday_hours REAL NOT NULL DEFAULT 0,
    night_hours REAL NOT NULL DEFAULT 0,holiday_hours REAL NOT NULL DEFAULT 0,
    note TEXT,source TEXT,updated_at TEXT NOT NULL,UNIQUE(employee_id,work_date)
);
CREATE TABLE payroll_adjustments(
    employee_id INTEGER NOT NULL,month TEXT NOT NULL,allowance REAL NOT NULL DEFAULT 0,
    responsibility REAL NOT NULL DEFAULT 0,advance REAL NOT NULL DEFAULT 0,
    probation_deduction REAL NOT NULL DEFAULT 0,bhxh_employee_amount REAL NOT NULL DEFAULT 0,
    bhxh_company_amount REAL NOT NULL DEFAULT 0,gross_override REAL NOT NULL DEFAULT 0,
    net_override REAL NOT NULL DEFAULT 0,use_override INTEGER NOT NULL DEFAULT 0,
    note TEXT,updated_at TEXT NOT NULL,PRIMARY KEY(employee_id,month)
);
CREATE TABLE kitchen_labor_costs(
    work_date TEXT NOT NULL,kitchen TEXT NOT NULL,amount REAL NOT NULL DEFAULT 0,
    source TEXT,updated_at TEXT NOT NULL,PRIMARY KEY(work_date,kitchen,source)
);
"""


def memory_database():
    connection = sqlite3.connect(":memory:")
    connection.row_factory = sqlite3.Row
    connection.executescript(SCHEMA)
    return connection


def build_legacy_workbook(
    *, period="2026-08", names=("AN",), hours=None, payroll=True, labor=True,
):
    year, month = map(int, period.split("-"))
    workbook = Workbook()
    attendance = workbook.active
    attendance.title = "Xưởng Cơm"
    attendance["C2"] = year
    attendance["C3"] = month
    for day in range(1, 4):
        attendance.cell(4, 5 + day, date(year, month, day))
    hours = hours or {}
    for index, name in enumerate(names):
        row = 6 + index * 2
        attendance.cell(row, 1, index + 1)
        attendance.cell(row, 2, name)
        attendance.cell(row, 3, "Helper")
        for day, value in hours.get(name, {1: 8}).items():
            attendance.cell(row, 5 + day, value)

    if payroll:
        salary = workbook.create_sheet("LƯƠNG XƯỞNG")
        salary["C2"] = 5_000_000
        salary["I1"] = month
        salary["J1"] = year
        for index, name in enumerate(names):
            row = 4 + index * 2
            salary.cell(row, 1, index + 1)
            salary.cell(row, 2, name)
            salary.cell(row, 3, "Helper")
            salary.cell(row + 1, 4, 8)
            salary.cell(row + 1, 11, 1_000_000 + index)
            salary.cell(row + 1, 15, 900_000 + index)

    if labor:
        market = workbook.create_sheet("chấm công chợ")
        market["C3"] = "ATV"
        market["A4"] = date(year, month, 1)
        market["C4"] = 167_000
        market["A5"] = date(year, month, 2)
        market["C5"] = 168_000
    return workbook


def parse_book(connection, workbook, period="2026-08", formula_workbook=None, source="source.xlsx"):
    return parse_legacy_attendance_workbooks(
        connection, workbook, formula_workbook or workbook, source, period,
    )


class LegacyAttendanceParserTests(unittest.TestCase):
    def test_period_is_required_and_sheet_mismatch_is_blocked(self):
        with self.assertRaises(ValueError):
            validate_legacy_attendance_period("")
        workbook = build_legacy_workbook(period="2026-07")
        connection = memory_database()
        try:
            with self.assertRaisesRegex(ValueError, r"Xưởng Cơm ghi kỳ 2026-07.*2026-08"):
                parse_book(connection, workbook, "2026-08")
        finally:
            workbook.close()
            connection.close()

    def test_negative_hours_and_missing_formula_cache_are_blocked(self):
        negative = build_legacy_workbook(hours={"AN": {1: -1}})
        connection = memory_database()
        try:
            with self.assertRaisesRegex(ValueError, "giờ âm"):
                parse_book(connection, negative)
        finally:
            negative.close()
            connection.close()

        values = build_legacy_workbook(hours={"AN": {}})
        formulas = build_legacy_workbook(hours={"AN": {}})
        formulas["Xưởng Cơm"]["F6"] = "=4+4"
        connection = memory_database()
        try:
            with self.assertRaisesRegex(ValueError, r"F6.*công thức chưa có kết quả"):
                parse_book(connection, values, formula_workbook=formulas)
        finally:
            values.close()
            formulas.close()
            connection.close()

    def test_daily_wage_is_reported_and_employee_codes_do_not_collide(self):
        workbook = build_legacy_workbook(
            names=("A-B", "AB"), hours={"A-B": {1: 167_000}, "AB": {1: 8}},
        )
        connection = memory_database()
        try:
            preview = parse_book(connection, workbook)
            codes = [row["employee_code"] for row in preview["snapshot"]["staff"]]
            self.assertEqual(len(codes), len(set(codes)))
            self.assertEqual(preview["counts"]["attendance_entries"], 1)
            self.assertTrue(any("số >24" in warning for warning in preview["warnings"]))
        finally:
            workbook.close()
            connection.close()

    def test_period_snapshot_replaces_old_import_but_preserves_manual_rows(self):
        connection = memory_database()
        initial = build_legacy_workbook(hours={"AN": {1: 8, 2: 8}})
        revised = build_legacy_workbook(hours={"AN": {1: 6}}, labor=False)
        try:
            first = parse_book(connection, initial, source="old-name.xlsx")
            apply_legacy_attendance_snapshot(
                connection, first["snapshot"], "old-name.xlsx", "2026-09-01 01:00:00",
            )
            person = connection.execute("SELECT id FROM staff WHERE full_name='AN'").fetchone()["id"]
            connection.execute(
                """UPDATE attendance_entries SET normal_hours=9,source='manual'
                   WHERE employee_id=? AND work_date='2026-08-01'""",
                (person,),
            )
            connection.execute(
                """UPDATE payroll_adjustments SET gross_override=999,note='Chốt tay'
                   WHERE employee_id=? AND month='2026-08'""",
                (person,),
            )

            second = parse_book(connection, revised, source="renamed.xlsx")
            result = apply_legacy_attendance_snapshot(
                connection, second["snapshot"], "renamed.xlsx", "2026-09-01 02:00:00",
            )
            rows = connection.execute(
                "SELECT work_date,normal_hours,source FROM attendance_entries ORDER BY work_date"
            ).fetchall()
            self.assertEqual([tuple(row) for row in rows], [("2026-08-01", 9.0, "manual")])
            payroll = connection.execute(
                "SELECT gross_override,note FROM payroll_adjustments WHERE employee_id=? AND month='2026-08'",
                (person,),
            ).fetchone()
            self.assertEqual(tuple(payroll), (999.0, "Chốt tay"))
            self.assertEqual(result["manual_attendance_preserved"], 1)
            self.assertEqual(result["manual_payroll_preserved"], 1)
            self.assertEqual(connection.execute(
                "SELECT COUNT(*) FROM kitchen_labor_costs"
            ).fetchone()[0], 0)

            # Repeating the same renamed/revised file is state-idempotent.
            apply_legacy_attendance_snapshot(
                connection, second["snapshot"], "another-name.xlsx", "2026-09-01 03:00:00",
            )
            self.assertEqual(connection.execute(
                "SELECT COUNT(*) FROM attendance_entries"
            ).fetchone()[0], 1)
            self.assertEqual(connection.execute(
                "SELECT COUNT(*) FROM payroll_adjustments"
            ).fetchone()[0], 1)
        finally:
            initial.close()
            revised.close()
            connection.close()

    def test_direct_import_requires_explicit_period(self):
        workbook = build_legacy_workbook()
        connection = memory_database()
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "attendance.xlsx"
            workbook.save(path)
            workbook.close()
            try:
                with self.assertRaises(ValueError):
                    import_legacy_attendance(
                        connection, path, path.name, "2026-09-01 00:00:00", "",
                    )
            finally:
                connection.close()

    def test_real_august_customer_file_matches_selected_period(self):
        self.assertTrue(REAL_T8_FILE.exists(), REAL_T8_FILE)
        values = load_workbook(REAL_T8_FILE, data_only=True, keep_links=False)
        formulas = load_workbook(REAL_T8_FILE, data_only=False, keep_links=False)
        connection = memory_database()
        try:
            preview = parse_legacy_attendance_workbooks(
                connection, values, formulas, REAL_T8_FILE.name, "2026-08",
            )
            self.assertTrue(preview["can_confirm"])
            self.assertGreater(preview["counts"]["staff"], 0)
            self.assertGreater(preview["counts"]["attendance_entries"], 0)
            self.assertGreater(preview["counts"]["labor_cost_entries"], 0)
            self.assertTrue(any("số >24" in warning for warning in preview["warnings"]))
        finally:
            values.close()
            formulas.close()
            connection.close()


class LegacyAttendanceApiTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.database_path = Path(self.temp.name) / "attendance.sqlite3"
        connection = sqlite3.connect(self.database_path)
        connection.executescript(SCHEMA)
        connection.close()

        @contextmanager
        def database():
            conn = sqlite3.connect(self.database_path)
            conn.row_factory = sqlite3.Row
            try:
                yield conn
                conn.commit()
            except Exception:
                conn.rollback()
                raise
            finally:
                conn.close()

        self.app = Flask(__name__)
        self.app.config["TESTING"] = True
        register_contract_routes(self.app, {
            "db": database,
            "now_iso": lambda: "2026-09-01 04:00:00",
            "clean_text": lambda value: str(value or "").strip(),
            "number_value": lambda value, default=0: float(value or default),
            "tax_factor": lambda value: 0,
            "setting_get": lambda conn, key, default="": default,
            "setting_set": lambda conn, key, value: None,
            "root": Path(self.temp.name),
            "data_dir": Path(self.temp.name),
            "require_batch": lambda conn, batch_id: None,
            "export_supplier_orders": lambda *args, **kwargs: None,
            "export_deliveries": lambda *args, **kwargs: None,
            "export_purchase_documents": lambda *args, **kwargs: None,
            "export_report": lambda *args, **kwargs: None,
        })
        self.client = self.app.test_client()

    def tearDown(self):
        PENDING_LEGACY_ATTENDANCE_IMPORTS.clear()
        self.temp.cleanup()

    @staticmethod
    def workbook_bytes():
        workbook = build_legacy_workbook()
        stream = io.BytesIO()
        workbook.save(stream)
        workbook.close()
        stream.seek(0)
        return stream

    def test_preview_confirm_and_direct_route_lock(self):
        locked = self.client.post("/api/attendance/import")
        self.assertEqual(locked.status_code, 410)
        preview = self.client.post(
            "/api/attendance/import/preview",
            data={"period": "2026-08", "file": (self.workbook_bytes(), "attendance.xlsx")},
            content_type="multipart/form-data",
        )
        self.assertEqual(preview.status_code, 200, preview.get_json())
        payload = preview.get_json()
        self.assertTrue(payload["can_confirm"])
        confirmed = self.client.post(
            "/api/attendance/import/confirm",
            json={"token": payload["token"], "confirmed": True},
        )
        self.assertEqual(confirmed.status_code, 200, confirmed.get_json())
        replay = self.client.post(
            "/api/attendance/import/confirm",
            json={"token": payload["token"], "confirmed": True},
        )
        self.assertEqual(replay.status_code, 410)

    def test_corrupt_zip_and_expired_token_are_rejected(self):
        corrupt = self.client.post(
            "/api/attendance/import/preview",
            data={"period": "2026-08", "file": (io.BytesIO(b"not-an-xlsx"), "bad.xlsx")},
            content_type="multipart/form-data",
        )
        self.assertEqual(corrupt.status_code, 400)

        expanded = io.BytesIO()
        with zipfile.ZipFile(expanded, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("xl/worksheets/sheet1.xml", b"0" * (50 * 1024 * 1024 + 1))
        expanded.seek(0)
        oversized_structure = self.client.post(
            "/api/attendance/import/preview",
            data={"period": "2026-08", "file": (expanded, "expanded.xlsx")},
            content_type="multipart/form-data",
        )
        self.assertEqual(oversized_structure.status_code, 413)

        preview = self.client.post(
            "/api/attendance/import/preview",
            data={"period": "2026-08", "file": (self.workbook_bytes(), "attendance.xlsx")},
            content_type="multipart/form-data",
        ).get_json()
        PENDING_LEGACY_ATTENDANCE_IMPORTS[preview["token"]]["created"] = (
            time.time() - MAPPING_IMPORT_TTL_SECONDS - 1
        )
        expired = self.client.post(
            "/api/attendance/import/confirm",
            json={"token": preview["token"], "confirmed": True},
        )
        self.assertEqual(expired.status_code, 410)


if __name__ == "__main__":
    unittest.main()
