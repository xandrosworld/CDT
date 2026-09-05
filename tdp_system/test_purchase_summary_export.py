import hashlib
import io
import sqlite3
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

try:
    from . import server as server_module
    from .pdf_documents import build_pdf_bundle
    from .print_bundle import workbook_sections
    from .purchase_summary_export import (
        PurchaseSummaryError,
        aggregate_purchase_summary_rows,
        build_purchase_summary_workbook,
        collect_purchase_summary_rows,
    )
    from .template_workbook import TemplateWorkbookError, safe_workbook_bytes
except ImportError:  # pragma: no cover - direct file invocation
    import server as server_module
    from pdf_documents import build_pdf_bundle
    from print_bundle import workbook_sections
    from purchase_summary_export import (
        PurchaseSummaryError,
        aggregate_purchase_summary_rows,
        build_purchase_summary_workbook,
        collect_purchase_summary_rows,
    )
    from template_workbook import TemplateWorkbookError, safe_workbook_bytes


ROOT = Path(__file__).resolve().parent.parent
GOLDEN = ROOT / "Em Thành.xlsx"
FAKE_ID = "0" * 12


def summary_row(**overrides):
    row = {
        "work_date": "2026-09-02",
        "seller": "Người bán kiểm thử",
        "address": "Địa chỉ kiểm thử",
        "cccd": FAKE_ID,
        "product_name": "Cà rốt",
        "unit": "Kg",
        "quantity": 2,
        "buy_price": 100,
        "amount": 200,
        "supplier": "NCC A",
        "kitchen": "BEP-A",
        "source_ref": 3,
    }
    row.update(overrides)
    return row


class PurchaseSummaryExportTests(unittest.TestCase):
    def test_same_item_and_legal_seller_merge_across_supplier_and_kitchen(self):
        rows = [
            summary_row(quantity=2, amount=200, supplier="NCC A", kitchen="BEP-A"),
            summary_row(
                unit="kg", quantity=3, buy_price=120, amount=360,
                supplier="NCC B", kitchen="BEP-B", source_ref=4,
            ),
            summary_row(
                seller="Người bán kiểm thử B", cccd="1" * 12,
                address="Địa chỉ kiểm thử B", quantity=1, amount=130,
                supplier="NCC A", kitchen="BEP-A", source_ref=5,
            ),
        ]
        grouped = aggregate_purchase_summary_rows(rows)
        self.assertEqual(len(grouped), 2)
        first = next(item for item in grouped if item["cccd"] == FAKE_ID)
        self.assertEqual(first["quantity"], 5)
        self.assertEqual(first["amount"], 560)
        self.assertEqual(first["unit_price"], 112)
        self.assertEqual(first["suppliers"], ("NCC A", "NCC B"))
        self.assertEqual(first["kitchens"], ("BEP-A", "BEP-B"))
        self.assertEqual(first["source_refs"], (3, 4))

    def test_identity_conflict_and_formula_injection_fail_without_echoing_identity(self):
        private_a = "2" * 12
        private_b = "3" * 12
        rows = [
            summary_row(cccd=private_a),
            summary_row(cccd=private_b, product_name="Khoai tây"),
        ]
        with self.assertRaises(PurchaseSummaryError) as caught:
            aggregate_purchase_summary_rows(rows)
        self.assertEqual(caught.exception.code, "conflicting_purchase_identity")
        self.assertNotIn(private_a, str(caught.exception))
        self.assertNotIn(private_b, str(caught.exception))

        injected = [summary_row(product_name='=WEBSERVICE("https://invalid")')]
        with self.assertRaises(TemplateWorkbookError):
            build_purchase_summary_workbook(injected, template_path=GOLDEN)

    def test_real_golden_clone_is_dynamic_safe_and_removes_all_sample_values(self):
        source_hash = hashlib.sha256(GOLDEN.read_bytes()).hexdigest()
        rows = [
            summary_row(
                product_name=f"Mặt hàng kiểm thử {index:02d}",
                quantity=index,
                buy_price=1_000,
                amount=index * 1_000,
                source_ref=index + 10,
            )
            for index in range(1, 24)
        ]
        source = load_workbook(GOLDEN, data_only=True, keep_links=False)
        try:
            source_ids = {
                str(source["bảng kê tổng"].cell(row, 4).value or "").strip()
                for row in range(11, 32)
                if source["bảng kê tổng"].cell(row, 4).value not in (None, "")
            }
        finally:
            source.close()

        workbook = build_purchase_summary_workbook(rows, template_path=GOLDEN)
        try:
            self.assertEqual(workbook.sheetnames, ["bảng kê tổng"])
            self.assertIs(workbook.active, workbook["bảng kê tổng"])
            sheet = workbook.active
            self.assertEqual(sheet["A2"].value, "Từ ngày 02/09/2026 đến ngày 02/09/2026")
            self.assertTrue(all(sheet.cell(7, column).value is None for column in range(1, 11)))
            self.assertEqual(sheet["A11"].value, "02/09/2026")
            self.assertEqual(sheet["E33"].value, "Mặt hàng kiểm thử 23")
            self.assertEqual(sheet["A34"].value, "TỔNG CỘNG")
            self.assertEqual(sheet["G34"].value, 276)
            self.assertEqual(sheet["I34"].value, 276_000)
            self.assertIn("Hai trăm bảy mươi sáu nghìn đồng", sheet["A35"].value)
            self.assertIn("A34:E34", {str(value) for value in sheet.merged_cells.ranges})
            self.assertIn("A37:C37", {str(value) for value in sheet.merged_cells.ranges})
            self.assertIn("H37:J37", {str(value) for value in sheet.merged_cells.ranges})
            self.assertIn("B38:C38", {str(value) for value in sheet.merged_cells.ranges})
            self.assertIn("A39:C39", {str(value) for value in sheet.merged_cells.ranges})
            self.assertIn("H39:J39", {str(value) for value in sheet.merged_cells.ranges})
            self.assertEqual(str(sheet.print_area), "'bảng kê tổng'!$A$1:$J$44")
            self.assertEqual(sheet.print_title_rows, "$8:$10")
            self.assertEqual(str(sheet.page_setup.paperSize), "9")
            self.assertEqual(sheet.page_setup.orientation, "landscape")
            self.assertEqual(sheet.page_setup.fitToWidth, 1)
            self.assertEqual(sheet.page_setup.fitToHeight, 0)
            self.assertIsNone(sheet["C49"].value)
            self.assertIsNone(sheet["C50"].value)
            output_ids = {
                str(sheet.cell(row, 4).value or "").strip()
                for row in range(11, 34)
                if sheet.cell(row, 4).value not in (None, "")
            }
            self.assertTrue(source_ids.isdisjoint(output_ids))
            self.assertEqual(output_ids, {FAKE_ID})
            self.assertFalse(workbook._external_links)
            self.assertFalse(any(
                isinstance(cell.value, str) and cell.value.startswith("=")
                for row in sheet.iter_rows()
                for cell in row
            ))
            payload = safe_workbook_bytes(workbook)
        finally:
            workbook.close()
        self.assertEqual(hashlib.sha256(GOLDEN.read_bytes()).hexdigest(), source_hash)

        reopened = load_workbook(io.BytesIO(payload), data_only=False, keep_links=False)
        try:
            self.assertEqual(reopened.sheetnames, ["bảng kê tổng"])
            self.assertFalse(reopened._external_links)
            self.assertIsNone(reopened.active["C50"].value)
        finally:
            reopened.close()

    def test_collect_uses_confirmed_purchase_scope_and_never_receivables(self):
        connection = sqlite3.connect(":memory:")
        connection.row_factory = sqlite3.Row
        connection.executescript(
            """
            CREATE TABLE people(
                name TEXT, cccd TEXT, issue_date TEXT, issue_place TEXT, address TEXT
            );
            CREATE TABLE settings(key TEXT PRIMARY KEY, value TEXT);
            CREATE TABLE products(
                code TEXT, purchase_list INTEGER, seller TEXT, cccd TEXT, supplier TEXT
            );
            CREATE TABLE orders(
                id INTEGER, purchase_list INTEGER, seller TEXT, cccd TEXT, kitchen TEXT,
                source_row INTEGER, supplier TEXT, actual_received REAL, damaged_qty REAL,
                supplier_return_qty REAL, product_code TEXT, product_name TEXT, unit TEXT,
                buy_price REAL, work_date TEXT
            );
            CREATE TABLE purchase_workbook_lines(
                id INTEGER, batch_id INTEGER, row_key TEXT, order_id INTEGER,
                source_row INTEGER, product_code TEXT, kitchen TEXT, work_date TEXT,
                product_name TEXT, unit TEXT, supplier TEXT, buy_price REAL,
                actual_qty REAL, amount REAL, status TEXT
            );
            CREATE TABLE receivable_ledger_lines(id INTEGER, amount REAL);
            """
        )
        connection.execute(
            "INSERT INTO people VALUES(?,?,?,?,?)",
            ("Người bán kiểm thử", FAKE_ID, "02/01/2020", "Nơi cấp kiểm thử", "Địa chỉ kiểm thử"),
        )
        connection.execute(
            "INSERT INTO products VALUES(?,?,?,?,?)",
            ("P1", 1, "Người bán kiểm thử", FAKE_ID, "NCC A"),
        )
        orders = [
            {
                "id": 10, "purchase_list": 1, "seller": "Người bán kiểm thử",
                "cccd": FAKE_ID, "kitchen": "BEP-A", "source_row": 3,
            },
            {
                "id": 11, "purchase_list": 1, "seller": "Người bán kiểm thử",
                "cccd": FAKE_ID, "kitchen": "BEP-B", "source_row": 4,
            },
        ]
        connection.executemany(
            "INSERT INTO purchase_workbook_lines VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            [
                (1, 7, "a", 10, 3, "P1", "BEP-A", "2026-09-02", "Cà rốt", "Kg", "NCC A", 100, 2, 200, "confirmed"),
                (2, 7, "b", 11, 4, "P1", "BEP-B", "2026-09-02", "Cà rốt", "kg", "NCC B", 120, 3, 360, "confirmed"),
            ],
        )
        connection.execute("INSERT INTO receivable_ledger_lines VALUES(1,999999999)")
        traced = []
        connection.set_trace_callback(traced.append)
        try:
            projected = collect_purchase_summary_rows(
                connection, {"id": 7, "work_date": "2026-09-02", "status": "approved"}, orders,
            )
            self.assertEqual(len(projected), 2)
            self.assertEqual(sum(item["amount"] for item in projected), 560)
            self.assertFalse(any("receivable" in statement.casefold() for statement in traced))

            previous = server_module.MASTER_SOURCE
            server_module.MASTER_SOURCE = GOLDEN
            try:
                workbook = server_module.export_purchase_documents(
                    connection,
                    {"id": 7, "work_date": "2026-09-02", "status": "approved"},
                    orders,
                )
                try:
                    self.assertEqual(workbook.sheetnames, ["bảng kê tổng", "biên nhận"])
                    self.assertEqual(workbook.active["G11"].value, 5)
                    self.assertEqual(workbook.active["H11"].value, 112)
                    self.assertEqual(workbook.active["I11"].value, 560)
                    self.assertEqual(workbook["biên nhận"]["D10"].value, FAKE_ID)
                finally:
                    workbook.close()
            finally:
                server_module.MASTER_SOURCE = previous
        finally:
            connection.set_trace_callback(None)
            connection.close()

    def test_legacy_bk_row_uses_configured_percentage_only_when_buy_price_is_missing(self):
        connection = sqlite3.connect(":memory:")
        connection.row_factory = sqlite3.Row
        connection.executescript(
            """
            CREATE TABLE settings(key TEXT PRIMARY KEY, value TEXT);
            CREATE TABLE people(name TEXT, cccd TEXT, address TEXT);
            CREATE TABLE products(
                code TEXT, purchase_list INTEGER, seller TEXT, cccd TEXT, supplier TEXT
            );
            INSERT INTO settings VALUES('purchase_rate','0.95');
            """
        )
        connection.execute(
            "INSERT INTO people VALUES(?,?,?)",
            ("Người bán kiểm thử", FAKE_ID, "Địa chỉ kiểm thử"),
        )
        order = {
            "id": 20, "purchase_list": 1, "seller": "Người bán kiểm thử",
            "cccd": FAKE_ID, "kitchen": "BEP-A", "source_row": 8,
            "supplier": "NCC A", "actual_received": 2, "damaged_qty": 0,
            "supplier_return_qty": 0, "product_code": "P1",
            "product_name": "Cà rốt", "unit": "Kg", "buy_price": 0,
            "sell_price": 101, "work_date": "2026-09-02",
        }
        try:
            projected = collect_purchase_summary_rows(
                connection,
                {"id": 8, "work_date": "2026-09-02", "status": "draft"},
                [order],
            )
            self.assertEqual(len(projected), 1)
            self.assertEqual(projected[0]["buy_price"], 96)
            self.assertEqual(projected[0]["amount"], 192)
        finally:
            connection.close()

    def test_optional_bundle_export_skips_only_a_batch_without_bk_rows(self):
        connection = sqlite3.connect(":memory:")
        connection.row_factory = sqlite3.Row
        connection.executescript(
            """
            CREATE TABLE people(name TEXT, cccd TEXT, address TEXT);
            CREATE TABLE products(
                code TEXT, purchase_list INTEGER, seller TEXT, cccd TEXT, supplier TEXT
            );
            """
        )
        try:
            batch = {"id": 9, "work_date": "2026-09-02", "status": "approved"}
            self.assertIsNone(
                server_module.export_optional_purchase_documents(connection, batch, [])
            )
            with self.assertRaises(PurchaseSummaryError) as caught:
                server_module.export_purchase_documents(connection, batch, [])
            self.assertEqual(caught.exception.code, "no_purchase_summary_rows")
        finally:
            connection.close()

    def test_purchase_summary_projects_to_verified_a4_pdf_without_identity_metadata(self):
        workbook = build_purchase_summary_workbook(
            [summary_row(), summary_row(product_name="Khoai tây", amount=300)],
            template_path=GOLDEN,
        )
        try:
            sections = workbook_sections("purchases", workbook)
            self.assertEqual(len(sections), 1)
            self.assertEqual(len(sections[0]["rows"]), 2)
            self.assertEqual(
                [column["label"] for column in sections[0]["columns"]],
                [
                    "Ngày tháng năm mua hàng", "Tên người bán", "Địa chỉ", "Số CCCD",
                    "Tên mặt hàng", "ĐVT", "Số lượng", "Đơn giá",
                    "Tổng giá thanh toán", "Ghi chú",
                ],
            )
            with tempfile.TemporaryDirectory(prefix="tdp_purchase_summary_pdf_") as temp_name:
                output = Path(temp_name) / "summary.pdf"
                manifest = build_pdf_bundle(
                    sections,
                    output,
                    company="CÔNG TY TNHH THỰC PHẨM THÀNH ĐẠT PHÁT",
                    document_title="Bảng kê thu mua ngày 02/09/2026",
                    generated_at="2026-09-02T12:00:00",
                    paper="A4",
                )
                self.assertTrue(manifest["ok"])
                self.assertEqual(manifest["paper"], "A4")
                self.assertEqual(manifest["section_count"], 1)
                self.assertNotIn(FAKE_ID, repr(manifest))
                self.assertTrue(output.is_file())
        finally:
            workbook.close()


if __name__ == "__main__":
    unittest.main()
