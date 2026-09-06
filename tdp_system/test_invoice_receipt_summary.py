"""Paid and zero-value goods must share quantity/value without changing source rows."""
import json
import sqlite3
import subprocess
import unittest
from pathlib import Path

from openpyxl import load_workbook

from .invoice_input_sync import sync_input_batch
from .invoice_mapping import save_mapping
from .invoice_receipt import create_input_receipt
from .invoice_valuation import moving_average_report
from .invoice_workbench_listing import invoice_range_payload, range_workbook, input_receipt_summary
from .test_invoice_input_sync import DateBoundedMsmi, init_test_database, now_iso
from .test_invoice_workbench_listing import prepare
from .test_msmi_sync import remote_invoice


def seed_promotion(conn, mapped=True):
    conn.executemany("INSERT INTO products(code,name,unit) VALUES(?,?,?)",
                     [('QA-OIL', 'Dầu hào kiểm thử', 'Can'), ('QA-CHILI', 'Tương ớt kiểm thử', 'Can')])
    remote = remote_invoice(28)
    remote['_id'] = 'QA-PAID-AND-FREE'
    remote['hdhhdvu'] = [
        dict(stt=i+1, ma=code, ten=name, dvtinh='Can', sluong=qty, dgia=amount/qty,
             thtien=amount, tchat=2 if amount == 0 else 1, tsuat='0%')
        for i, (code, name, qty, amount) in enumerate([
            ('QA-OIL-PAID', 'Dầu hào MISA can 2L', 24, 1288889),
            ('QA-CHILI-PAID', 'Tương ớt cay đặc biệt MISA can 2L', 86, 3264815),
            ('QA-OIL-FREE', 'Dầu hào MISA can 2L (Hàng khuyến mại không thu tiền)', 6, 0),
            ('QA-CHILI-FREE', 'Tương ớt cay đặc biệt MISA can 2L (Hàng khuyến mại không thu tiền)', 21, 0),
        ])]
    remote.update(tgtcthue=4553704, tgtthue=0, tgtttbso=4553704)
    batch = prepare(conn)
    sync_input_batch(conn, DateBoundedMsmi([remote]), batch['id'], now_iso)
    invoice_id = conn.execute("SELECT id FROM msmi_invoices WHERE remote_id=?", (remote['_id'],)).fetchone()[0]
    ids = {}
    for row in conn.execute('SELECT id,source_item_code FROM msmi_invoice_items WHERE invoice_id=?', (invoice_id,)).fetchall():
        ids[row['source_item_code']] = row['id']
        if mapped:
            save_mapping(conn, direction='input', item_id=row['id'],
                         product_code='QA-OIL' if 'OIL' in row['source_item_code'] else 'QA-CHILI', now_iso=now_iso)
    return {'promotion_invoice': invoice_id, 'promotion_lines': ids}


class ReceiptSummaryTests(unittest.TestCase):
    def setUp(self):
        self.conn = sqlite3.connect(':memory:')
        init_test_database(self.conn)

    def tearDown(self):
        self.conn.close()

    def payload(self, **kwargs):
        return invoice_range_payload(self.conn, tenant='TDP', invoice_type='input',
                                     date_from='2026-08-01', date_to='2026-08-31', **kwargs)

    def test_paid_plus_free_post_report_repeat_and_excel_preserve_amount(self):
        fixture = seed_promotion(self.conn)
        before = list(self.conn.execute('SELECT qty,unit_price,amount FROM msmi_invoice_items'))
        changes = self.conn.total_changes
        payload = self.payload()
        summary = payload['items'][0]['receipt_summary']
        oil = next(r for r in summary['items'] if r['product_code'] == 'QA-OIL')
        chili = next(r for r in summary['items'] if r['product_code'] == 'QA-CHILI')
        self.assertEqual((30, 6, 1288889), (oil['qty'], oil['zero_amount_qty'], oil['amount']))
        self.assertAlmostEqual(1288889/30, oil['average_unit_cost'])
        self.assertEqual((107, 21, 3264815), (chili['qty'], chili['zero_amount_qty'], chili['amount']))
        self.assertEqual(0, summary['pending_lines'])
        wb = load_workbook(range_workbook(payload))
        ws = wb['Tong nhap theo ma']
        row = next(r for r in ws.iter_rows(min_row=3) if r[2].value == 'QA-OIL')
        self.assertEqual([30, 6, 1288889], [c.value for c in row[5:8]])
        self.assertAlmostEqual(1288889/30, row[8].value)
        self.assertEqual('#,##0', row[8].number_format)
        self.assertEqual(changes, self.conn.total_changes, 'preview and export are read-only')
        first = create_input_receipt(self.conn, fixture['promotion_invoice'], now_iso)
        self.assertEqual(4, first['new_invoice_ledger_lines'])
        create_input_receipt(self.conn, fixture['promotion_invoice'], now_iso)
        self.assertEqual(4, self.conn.execute('SELECT COUNT(*) FROM invoice_inventory_ledger').fetchone()[0])
        report = moving_average_report(self.conn, date_from='2026-08-01', date_to='2026-08-31')
        row = next(r for r in report['items'] if r['product_code'] == 'QA-OIL')
        self.assertEqual((30, 1288889, 30, 1288889),
                         (row['input_qty'], row['input_value'], row['closing_qty'], row['closing_value']))
        self.assertAlmostEqual(42962.966667, row['average_unit_cost'], places=6)
        self.assertEqual([tuple(r) for r in before], [tuple(r) for r in self.conn.execute('SELECT qty,unit_price,amount FROM msmi_invoice_items')])

    def test_existing_stock_is_included_only_in_warehouse_average(self):
        fixture = seed_promotion(self.conn)
        self.conn.execute("""INSERT INTO inventory_transactions(txn_date,product_code,qty_in,qty_out,unit_cost,
                          source_type,source_id,source_line,status,note,created_at,updated_at)
                          VALUES('2026-08-01','QA-OIL',10,0,50000,'OPENING','2026-08','QA-OIL','posted','QA',?,?)""", (now_iso(),now_iso()))
        create_input_receipt(self.conn, fixture['promotion_invoice'], now_iso)
        row = next(r for r in moving_average_report(self.conn, date_from='2026-08-01', date_to='2026-08-31')['items'] if r['product_code']=='QA-OIL')
        self.assertEqual(40, row['closing_qty'])
        self.assertEqual(1788889, row['closing_value'])
        self.assertAlmostEqual(1788889/40, row['average_unit_cost'])

    def test_filter_does_not_hide_mapped_quantity_from_summary(self):
        fixture = seed_promotion(self.conn, mapped=False)
        paid = fixture['promotion_lines']['QA-OIL-PAID']
        save_mapping(self.conn, direction='input', item_id=paid, product_code='QA-OIL', now_iso=now_iso)
        payload = self.payload(line_filter='unmapped')
        self.assertNotIn(paid, [r['id'] for r in payload['lines']])
        summary = payload['items'][0]['receipt_summary']
        self.assertEqual(3, summary['pending_lines'])
        self.assertEqual(24, summary['items'][0]['qty'])
        self.assertEqual('pending_mapping', self.conn.execute('SELECT receipt_status FROM msmi_invoices').fetchone()[0])

    def test_summary_uses_converted_quantity_and_keeps_units_separate(self):
        rows = [dict(inventory_eligible=1, mapping_status='mapped', product_code='P', product_unit=unit,
                     stock_qty=qty, qty=1, amount=amount) for unit,qty,amount in [('Can', 24, 1288889), ('Can',6,0), ('kg',2,100)]]
        summary = input_receipt_summary({'items': rows})
        self.assertEqual([30,2], [r['qty'] for r in summary['items']])
        rows.append(dict(rows[0], stock_qty=None))
        self.assertEqual(1, input_receipt_summary({'items': rows})['pending_lines'])

    def test_render_separate_source_quantity_unit_and_summary_button(self):
        seed_promotion(self.conn)
        script = """global.window={};require('./tdp_system/static/invoice-workbench.js');
          const html=window.TdpInvoiceWorkbench({invoiceDirection:'input',invoiceListing:JSON.parse(process.argv[1])},
            {esc:x=>String(x??''),num:x=>String(x??0),money:x=>String(x??0),dateVN:x=>x});console.log(html);"""
        html = subprocess.check_output(['node','-e',script,json.dumps(self.payload())],cwd=Path(__file__).resolve().parent.parent).decode('utf-8')
        self.assertIn('class="num-cell invoice-source-qty">24</td><td class="invoice-source-unit">Can</td>',html)
        self.assertIn('<th>Số lượng nguồn</th><th>ĐVT</th><th>Đơn giá</th>',html)
        self.assertIn('data-action="view-invoice-receipt-summary"',html)


if __name__ == '__main__':
    unittest.main()
