"""Excel export for input invoices already fetched into one workbench batch.

The export is intentionally downstream of the read-only mSMI synchronization
and upstream of mapping/inventory posting.  Users can therefore download the
source invoice list immediately after a pull without changing accounting
stock.  Raw connector payloads, remote identifiers and credentials are never
written to the workbook.
"""

from __future__ import annotations

import io
import math
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from typing import Any

from flask import jsonify, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins


INPUT_INVOICE = "INPUT_ELECTRONIC_INVOICE"
MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class InvoiceInputExportError(ValueError):
    def __init__(self, message: str, *, code: str = "invalid_invoice_input_export", status: int = 400):
        super().__init__(message)
        self.code = code
        self.status = status


def _plain(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").replace("\u00a0", " ").strip())


def _excel_text(value: Any) -> str:
    text = _plain(value)
    return "'" + text if text.startswith(("=", "+", "-", "@")) else text


def _vnd(value: Any, label: str = "Số tiền") -> int:
    try:
        number = Decimal(str(value or 0))
    except (InvalidOperation, TypeError, ValueError):
        raise InvoiceInputExportError(
            f"{label} không hợp lệ", code="invalid_invoice_input_money", status=500,
        ) from None
    if not number.is_finite():
        raise InvoiceInputExportError(
            f"{label} không hữu hạn", code="invalid_invoice_input_money", status=500,
        )
    return int(number.quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def _number(value: Any, label: str) -> float:
    if isinstance(value, bool):
        raise InvoiceInputExportError(f"{label} không hợp lệ", status=500)
    try:
        number = float(value or 0)
    except (TypeError, ValueError):
        raise InvoiceInputExportError(f"{label} không hợp lệ", status=500) from None
    if not math.isfinite(number):
        raise InvoiceInputExportError(f"{label} không hữu hạn", status=500)
    return number


def _excel_date(value: Any) -> date | str:
    text = _plain(value)
    try:
        return datetime.strptime(text, "%Y-%m-%d").date()
    except ValueError:
        return _excel_text(text)


def _date_text(value: Any) -> str:
    parsed = _excel_date(value)
    return parsed.strftime("%d/%m/%Y") if isinstance(parsed, date) else str(parsed)


STATUS_LABELS = {
    "synced": "Đã tải",
    "needs_mapping": "Cần ghép mã",
    "pending_mapping": "Chờ ghép mã",
    "unmapped": "Chưa ghép mã",
    "mapped": "Đã ghép mã",
    "ready": "Sẵn sàng ghi kho",
    "posted": "Đã ghi kho",
    "ignored": "Không đưa vào kho",
    "not_inventory": "Không nhập kho",
    "error": "Có lỗi cần kiểm tra",
}


def _status_label(value: Any) -> str:
    text = _plain(value)
    return STATUS_LABELS.get(text.casefold(), text or "Chưa có")


NATURE_LABELS = {
    "1": "Hàng hóa, dịch vụ",
    "2": "Khuyến mại",
    "3": "Chiết khấu",
    "4": "Ghi chú, diễn giải",
}


def _nature_label(value: Any) -> str:
    text = _plain(value)
    if not text:
        return "Không ghi"
    return NATURE_LABELS.get(text.casefold(), f"Khác (mã {text})")


def input_invoice_export_data(conn, batch_id: Any) -> dict[str, Any]:
    try:
        safe_batch_id = int(batch_id)
    except (TypeError, ValueError):
        raise InvoiceInputExportError("Mã phiên tải hóa đơn không hợp lệ") from None
    batch_row = conn.execute(
        "SELECT * FROM invoice_sync_batches WHERE id=?", (safe_batch_id,),
    ).fetchone()
    if batch_row is None:
        raise InvoiceInputExportError(
            "Không tìm thấy phiên tải hóa đơn", code="invoice_input_batch_not_found", status=404,
        )
    batch = dict(batch_row)
    if batch.get("invoice_type") != INPUT_INVOICE:
        raise InvoiceInputExportError(
            "Chỉ phiên hóa đơn đầu vào mới được xuất bằng chức năng này",
            code="invoice_input_batch_required", status=409,
        )

    invoices = [dict(row) for row in conn.execute(
        """SELECT i.id,i.seller_tax_code,i.seller_name,i.invoice_number,i.invoice_series,
                  i.invoice_date,i.subtotal,i.tax_amount,i.total_amount,i.sync_status,
                  i.receipt_status,i.synced_at,i.raw_json
             FROM invoice_sync_batch_invoices bi
             JOIN msmi_invoices i ON i.id=bi.invoice_id
            WHERE bi.batch_id=?
            ORDER BY i.invoice_date,i.invoice_series,i.invoice_number,i.id""",
        (safe_batch_id,),
    )]
    if not invoices:
        raise InvoiceInputExportError(
            "Phiên chưa có hóa đơn đầu vào để xuất; hãy tải dữ liệu trước",
            code="invoice_input_batch_empty", status=404,
        )
    invoice_ids = [int(item["id"]) for item in invoices]
    placeholders = ",".join("?" for _ in invoice_ids)
    lines = [dict(row) for row in conn.execute(
        f"""SELECT l.id,l.invoice_id,l.line_index,l.source_item_code,l.source_item_name,
                   l.source_unit,l.qty,l.unit_price,l.amount,l.tax_rate,l.source_nature,
                   l.inventory_eligible,l.validation_note,l.product_code,l.mapping_status,
                   l.conversion_factor,l.stock_qty,l.stock_unit_price
              FROM msmi_invoice_items l
             WHERE l.invoice_id IN ({placeholders})
             ORDER BY l.invoice_id,l.line_index,l.id""",
        invoice_ids,
    )]
    invoice_by_id = {int(item["id"]): item for item in invoices}
    for invoice in invoices:
        invoice["subtotal"] = _vnd(invoice.get("subtotal"), "Tiền trước thuế")
        invoice["tax_amount"] = _vnd(invoice.get("tax_amount"), "Tiền thuế")
        invoice["total_amount"] = _vnd(invoice.get("total_amount"), "Tổng thanh toán")
    for line in lines:
        if int(line["invoice_id"]) not in invoice_by_id:
            raise InvoiceInputExportError("Dòng hóa đơn nằm ngoài phiên tải", status=500)
        line["qty"] = _number(line.get("qty"), "Số lượng")
        line["unit_price"] = _vnd(line.get("unit_price"), "Đơn giá")
        line["amount"] = _vnd(line.get("amount"), "Thành tiền")
        if line.get("conversion_factor") is not None:
            line["conversion_factor"] = _number(line["conversion_factor"], "Hệ số quy đổi")
        line["stock_qty"] = _number(line.get("stock_qty"), "Số lượng kho")
        line["stock_unit_price"] = _vnd(line.get("stock_unit_price"), "Đơn giá kho")

    try:
        from .invoice_line_tax import annotate_invoice_tax
    except ImportError:
        from invoice_line_tax import annotate_invoice_tax
    for invoice in invoices:
        invoice['items'] = [line for line in lines if line['invoice_id'] == invoice['id']]
        annotate_invoice_tax(invoice, invoice.pop('raw_json', '{}'))
        try:
            from .input_discount import annotate, DiscountError
        except ImportError:
            from input_discount import annotate, DiscountError
        try:
            annotate(conn,invoice)
        except DiscountError:
            pass
        for line in invoice['items']:
            if 'stock_amount' in line and line['stock_qty'] > 0:
                line['stock_unit_price'] = _vnd(line['stock_amount'] / line['stock_qty'], 'Đơn giá kho sau chiết khấu')
    return {
        "batch": batch,
        "invoices": invoices,
        "lines": lines,
        "totals": {
            "invoice_count": len(invoices),
            "line_count": len(lines),
            "subtotal": sum(item["subtotal"] for item in invoices),
            "tax_amount": sum(item["tax_amount"] for item in invoices),
            "total_amount": sum(item["total_amount"] for item in invoices),
        },
    }


THIN = Side(style="thin", color="7F7F7F")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TITLE_FILL = PatternFill(fill_type=None)
HEADER_FILL = PatternFill(fill_type=None)
TOTAL_FILL = PatternFill(fill_type=None)


def _style_table(
    ws, *, title: str, subtitle: str, headers: list[str], widths: list[float],
    money_columns: tuple[int, ...] = (), quantity_columns: tuple[int, ...] = (),
    date_columns: tuple[int, ...] = (), total_row: int | None = None,
) -> None:
    end_column = len(headers)
    ws.insert_rows(1, 2)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_column)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end_column)
    ws.cell(1, 1, title)
    ws.cell(2, 1, subtitle)
    ws.cell(1, 1).font = Font(name="Times New Roman", size=16, bold=True)
    ws.cell(1, 1).fill = TITLE_FILL
    ws.cell(1, 1).alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(2, 1).font = Font(name="Times New Roman", size=11, italic=True)
    ws.cell(2, 1).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 26
    ws.row_dimensions[2].height = 28
    ws.row_dimensions[3].height = 38
    for cell in ws[3]:
        cell.font = Font(name="Times New Roman", size=11, bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    for row in ws.iter_rows(min_row=4, max_col=end_column):
        for cell in row:
            cell.font = Font(name="Times New Roman", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = BORDER
    for column in money_columns:
        for row in range(4, ws.max_row + 1):
            ws.cell(row, column).number_format = "#,##0"
    for column in quantity_columns:
        for row in range(4, ws.max_row + 1):
            ws.cell(row, column).number_format = "#,##0.###"
    for column in date_columns:
        for row in range(4, ws.max_row + 1):
            ws.cell(row, column).number_format = "dd/mm/yyyy"
    if total_row:
        for cell in ws[total_row]:
            cell.font = Font(name="Times New Roman", size=11, bold=True)
            cell.fill = TOTAL_FILL
            cell.border = BORDER
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:{get_column_letter(end_column)}{max(ws.max_row, 3)}"
    ws.print_title_rows = "3:3"
    ws.print_area = f"A1:{get_column_letter(end_column)}{max(ws.max_row, 3)}"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.scale = None
    ws.page_margins = PageMargins(left=0.25, right=0.25, top=0.45, bottom=0.45)
    ws.print_options.horizontalCentered = True
    ws.oddFooter.center.text = "Trang &P / &N"
    ws.sheet_view.showGridLines = False


def input_invoice_workbook(data: dict[str, Any]) -> Workbook:
    batch = data["batch"]
    invoices = data["invoices"]
    lines = data["lines"]
    totals = data["totals"]
    period = (
        f"Từ {_date_text(batch['date_from'])} đến {_date_text(batch['date_to'])} · "
        f"lần tải số {batch['id']} · chỉ xuất dữ liệu đã tải, chưa cộng vào kho"
    )
    workbook = Workbook()
    workbook.properties.title = f"Hóa đơn đầu vào {batch['date_from']} - {batch['date_to']}"
    workbook.properties.subject = "Danh sách và chi tiết hóa đơn đầu vào đã kéo từ nguồn chỉ đọc"
    workbook.properties.creator = "Thành Đạt Phát"
    workbook.calculation.fullCalcOnLoad = False
    workbook.calculation.forceFullCalc = False
    workbook.calculation.calcMode = "manual"

    summary_headers = [
        "STT", "Ngày hóa đơn", "Ký hiệu", "Số hóa đơn", "MST người bán",
        "Tên người bán", "Tiền trước thuế", "Tiền thuế", "Tổng thanh toán",
        "Trạng thái dữ liệu", "Trạng thái nhập kho",
    ]
    ws = workbook.active
    ws.title = "Hóa đơn đầu vào"
    ws.append(summary_headers)
    for index, invoice in enumerate(invoices, start=1):
        ws.append([
            index, _excel_date(invoice.get("invoice_date")), _excel_text(invoice.get("invoice_series")),
            _excel_text(invoice.get("invoice_number")), _excel_text(invoice.get("seller_tax_code")),
            _excel_text(invoice.get("seller_name")), invoice["subtotal"], invoice["tax_amount"],
            invoice["total_amount"], _status_label(invoice.get("sync_status")),
            _status_label(invoice.get("receipt_status")),
        ])
    ws.append([
        "TỔNG", "", "", "", "", "", totals["subtotal"], totals["tax_amount"],
        totals["total_amount"], "", "",
    ])
    total_row = ws.max_row + 2
    _style_table(
        ws, title="DANH SÁCH HÓA ĐƠN ĐẦU VÀO ĐÃ TẢI", subtitle=period,
        headers=summary_headers,
        widths=[7, 14, 15, 16, 18, 34, 18, 16, 19, 18, 19],
        money_columns=(7, 8, 9), date_columns=(2,), total_row=total_row,
    )

    detail_headers = [
        "STT", "Ngày HĐ", "Ký hiệu", "Số HĐ", "MST người bán", "Tên người bán",
        "Dòng", "Mã hàng nguồn", "Tên hàng nguồn", "ĐVT nguồn", "Số lượng",
        "Đơn giá", "Thành tiền", "Thuế suất", "Tính chất", "Có vào kho",
        "Mã TĐP", "Trạng thái ghép mã", "Hệ số quy đổi", "SL kho", "Giá kho",
        "Ghi chú kiểm tra", "Tiền thuế", "Tiền gồm thuế", "Đối chiếu thuế",
    ]
    ws = workbook.create_sheet("Chi tiết hàng hóa")
    ws.append(detail_headers)
    by_id = {int(item["id"]): item for item in invoices}
    for index, line in enumerate(lines, start=1):
        invoice = by_id[int(line["invoice_id"])]
        ws.append([
            index, _excel_date(invoice.get("invoice_date")), _excel_text(invoice.get("invoice_series")),
            _excel_text(invoice.get("invoice_number")), _excel_text(invoice.get("seller_tax_code")),
            _excel_text(invoice.get("seller_name")), int(line.get("line_index") or 0),
            _excel_text(line.get("source_item_code")), _excel_text(line.get("source_item_name")),
            _excel_text(line.get("source_unit")), line["qty"], line["unit_price"], line["amount"],
            _excel_text(line.get("tax_rate")), _nature_label(line.get("source_nature")),
            "Có" if int(line.get("inventory_eligible") or 0) else "Không",
            _excel_text(line.get("product_code")), _status_label(line.get("mapping_status")),
            line.get("conversion_factor"), line["stock_qty"], line["stock_unit_price"],
            _excel_text(line.get("validation_note")), line.get('line_tax_amount'), line.get('amount_with_tax'),
            _excel_text(line.get('tax_note')),
        ])
    _style_table(
        ws, title="CHI TIẾT HÀNG HÓA TRÊN HÓA ĐƠN ĐẦU VÀO", subtitle=period,
        headers=detail_headers,
        widths=[7, 13, 14, 15, 17, 30, 8, 18, 34, 13, 13, 16, 17, 13, 13, 12, 14, 18, 15, 14, 16, 30, 18, 20, 44],
        money_columns=(12, 13, 21, 23, 24), quantity_columns=(11, 19, 20), date_columns=(2,),
    )

    proof = workbook.create_sheet("Thông tin lần tải")
    proof.append(["Mục", "Giá trị"])
    proof_rows = [
        ("Mã lần tải", int(batch["id"])),
        ("Dữ liệu từ", "mSMI - hóa đơn đầu vào" if _plain(batch.get("source")).casefold() == "msmi" else _excel_text(batch.get("source"))),
        ("Loại hóa đơn", "Hóa đơn điện tử đầu vào"),
        ("Từ ngày", _date_text(batch["date_from"])),
        ("Đến ngày", _date_text(batch["date_to"])),
        ("Trạng thái", _status_label(batch.get("status"))),
        ("Số hóa đơn", totals["invoice_count"]),
        ("Số dòng hàng hóa", totals["line_count"]),
        ("Tổng trước thuế", totals["subtotal"]),
        ("Tổng thuế", totals["tax_amount"]),
        ("Tổng thanh toán", totals["total_amount"]),
        ("Tải file Excel có tự cộng vào kho không?", "Không"),
    ]
    for row in proof_rows:
        proof.append(list(row))
    _style_table(
        proof, title="THÔNG TIN LẦN TẢI HÓA ĐƠN ĐẦU VÀO",
        subtitle="File chỉ có dữ liệu để kiểm tra, không chứa mật khẩu hay thông tin đăng nhập",
        headers=["Mục", "Giá trị"], widths=[34, 48], money_columns=(),
    )
    for row in range(12, 15):
        proof.cell(row, 2).number_format = "#,##0"
    return workbook


def register_invoice_input_export_routes(app, ctx) -> None:
    db_factory = ctx["db"]
    now_iso = ctx.get("now_iso")

    @app.get("/api/invoice-workbench/batches/<int:batch_id>/export-input-xlsx")
    def api_export_input_invoice_batch(batch_id: int):
        with db_factory() as conn:
            try:
                data = input_invoice_export_data(conn, batch_id)
                workbook = input_invoice_workbook(data)
            except InvoiceInputExportError as error:
                return jsonify({"ok": False, "error": str(error), "code": error.code}), error.status
            stream = io.BytesIO()
            try:
                workbook.save(stream)
            finally:
                workbook.close()
            if now_iso is not None:
                conn.execute(
                    """INSERT INTO audit_log(
                           event_type,entity_type,entity_id,status,message,metadata_json,created_at
                       ) VALUES('invoice_input.export','invoice_sync_batch',?,'ok','',?,?)""",
                    (
                        str(batch_id),
                        '{"invoice_count":%d,"line_count":%d}' % (
                            data["totals"]["invoice_count"], data["totals"]["line_count"],
                        ),
                        now_iso(),
                    ),
                )
            stream.seek(0)
            response = send_file(
                stream, as_attachment=True,
                download_name=(
                    f"Hoa_don_dau_vao_{data['batch']['date_from']}_den_"
                    f"{data['batch']['date_to']}.xlsx"
                ),
                mimetype=MIME_XLSX,
            )
            response.headers["X-TDP-Invoice-Batch"] = str(batch_id)
            response.headers["X-TDP-Inventory-Effect"] = "none"
            return response


__all__ = [
    "InvoiceInputExportError", "input_invoice_export_data", "input_invoice_workbook",
    "register_invoice_input_export_routes",
]
