"""Shared contractor statement matching the customer's BK-Yjilink sheet."""
from collections import OrderedDict, defaultdict
from copy import copy
from decimal import Decimal
from pathlib import Path
import math
import textwrap

from openpyxl import Workbook, load_workbook


def customer_statement(scope):
    try:
        from .invoice_payment_documents import _strict_date, _excel_text, number_to_vietnamese, InvoicePaymentDocumentError
        from .invoice_line_tax import tax_fields, number
        from .payment_print_layout import readable_payment_layout, keep_payment_footer
        from .invoice_payment_scope import _tax_percent
    except ImportError:
        from invoice_payment_documents import _strict_date, _excel_text, number_to_vietnamese, InvoicePaymentDocumentError
        from invoice_line_tax import tax_fields, number
        from payment_print_layout import readable_payment_layout, keep_payment_footer
        from invoice_payment_scope import _tax_percent

    def numeric(value, label):
        parsed = number(value)
        if parsed is None:
            raise InvoicePaymentDocumentError(f'Bảng kê thiếu {label}. Tải lại chi tiết hóa đơn đầu ra rồi mở lại Đề nghị thanh toán.')
        return parsed

    groups, actual = OrderedDict(), defaultdict(lambda: [Decimal(0), Decimal(0), Decimal(0)])
    local_lines = scope.get('invoice_lines', scope.get('lines', []))
    lines = [(r, True) for r in scope.get('source_lines', [])] + [(r, False) for r in local_lines]
    for row, source in lines:
        name, unit = (row['source_item_name'], row['source_unit']) if source else (row['product_name'],row['unit'])
        # Local invoices store 8% as either "8%", 8 or the ratio 0.08.
        # Use the same interpretation as issued-invoice validation. Provider
        # rates retain their source-specific interpretation and actual VAT.
        local_rate = _tax_percent(row.get('tax')) if not source else None
        rate_value = row.get('tax_rate') if source else (
            'KKKNT' if local_rate == -2 else 'KCT' if local_rate == -1 else f'{local_rate:g}%')
        tax = tax_fields({'tax_rate':rate_value, 'amount':row['amount']})
        if 'line_tax_amount' in row:
            tax['line_tax_amount'] = row['line_tax_amount']
        amount = numeric(row['amount'], 'Thành tiền')
        vat = numeric(tax['line_tax_amount'], 'Tiền thuế')
        quantity, price = numeric(row['qty'], 'Số lượng'), numeric(row['unit_price'], 'Đơn giá')
        identity = ('source', row['invoice_id']) if source else ('local', row['draft_id'])
        for index,value in enumerate((amount,vat,amount+vat)):
            actual[identity][index] += value
        rate = tax['tax_rate']
        numeric_rate = number(rate)
        if numeric_rate is not None and 0 <= numeric_rate <= 100:
            rate = format(numeric_rate.normalize(),'f')+'%'
        key = (str(name).strip(), str(unit).strip(), price, rate)
        if key not in groups:
            groups[key] = [Decimal(0),Decimal(0),Decimal(0)]
        groups[key][0] += quantity; groups[key][1] += amount; groups[key][2] += vat
    # Keep invoice totals authoritative; never invent an allocation to make a table fit.
    for invoice in scope['invoices']:
        identity = ('local', invoice['draft_id']) if invoice.get('draft_id') else ('source', invoice['source_invoice_id'])
        expected = [numeric(invoice[k], k) for k in ('subtotal','tax_amount','total_amount')]
        # Same one-dong rounding tolerance as the verified invoice scope.
        if identity not in actual or any(abs(a-b) > 1 for a,b in zip(actual[identity],expected)):
            raise InvoicePaymentDocumentError(
                f"Hóa đơn {invoice['invoice_series']}/{invoice['invoice_number']}: Thành tiền, Tiền thuế hoặc Thanh toán của các dòng chưa khớp tổng hóa đơn. "
                'Vào Hóa đơn đầu vào + đầu ra, chọn Đầu ra và tải lại chi tiết hóa đơn này rồi mở lại Đề nghị thanh toán.',
                code='payment_statement_line_totals_mismatch')
    if not groups:
        raise InvoicePaymentDocumentError('Chưa có chi tiết mặt hàng để lập bảng kê. Tải lại chi tiết hóa đơn đầu ra rồi mở lại Đề nghị thanh toán.')
    book = Workbook(); ws = book.active; ws.title = 'Bảng tổng hợp giao nhận'
    book.properties.subject = 'Mẫu bảng kê BK-Yjilink khách hàng'
    template_book = load_workbook(Path(__file__).parent/'templates'/'payment_statement_customer_20260930.xlsx')
    template = template_book.active
    count = len(groups); total_row = 11+count
    row_pairs = [(r,r) for r in range(1,11)] + [(11,11+i) for i in range(count)] + [(r,r+count-1) for r in range(12,16)]
    for source_row,target_row in row_pairs:
        ws.row_dimensions[target_row].height = template.row_dimensions[source_row].height
        for column in range(1,10):
            a,b = template.cell(source_row,column),ws.cell(target_row,column)
            for attr in ('font','fill','border','alignment','protection','number_format'):
                setattr(b,attr,copy(getattr(a,attr)))
            b.value = a.value
    for key,dim in template.column_dimensions.items():
        ws.column_dimensions[key].width = dim.width
    for merged in template.merged_cells.ranges:
        offset=count-1 if merged.min_row>=12 else 0
        ws.merge_cells(start_row=merged.min_row+offset,end_row=merged.max_row+offset,start_column=merged.min_col,end_column=merged.max_col)
    ws.page_margins=copy(template.page_margins);template_book.close()
    snapshot=scope['snapshot'];start=_strict_date(scope['date_from'],'Từ ngày');end=_strict_date(scope['date_to'],'Đến ngày')
    def write(cell,value):ws[cell]=_excel_text(value) if isinstance(value,str) else value
    write('A1',snapshot['company_name_snapshot'])
    write('A2','Địa chỉ: '+snapshot['company_address_snapshot'])
    write('A4','MST: '+snapshot['company_tax_code_snapshot'])
    write('A5',f'BẢNG TỔNG HỢP GIAO NHẬN THÁNG {end.month:02d} NĂM {end.year}' if start.strftime('%Y-%m')==end.strftime('%Y-%m') else 'BẢNG TỔNG HỢP GIAO NHẬN')
    write('A6',f"(Từ ngày {start:%d/%m/%Y} – {end:%d/%m/%Y})")
    for cell,label,key in [('A7','Khách hàng: ','buyer_name_snapshot'),('A8','Địa chỉ: ','buyer_address_snapshot'),('A9','MST: ','buyer_tax_code_snapshot')]:
        write(cell,label+snapshot[key])
    for index,((name,unit,price,rate),(quantity,amount,vat)) in enumerate(groups.items(),1):
        row=10+index
        for column,value in enumerate((index,name,unit,float(quantity),float(price),float(amount),rate,float(vat),float(amount+vat)),1):
            write(f'{chr(64+column)}{row}',value)
        for column in (4,5):ws.cell(row,column).number_format='#,##0.######'
        for column in (6,8,9):ws.cell(row,column).number_format='#,##0.##'
        ws.row_dimensions[row].height=max(25,len(textwrap.wrap(name,width=18))*20)
    for cell,key in [('F','subtotal'),('H','tax_amount'),('I','total_amount')]:
        write(f'{cell}{total_row}',scope['totals'][key]);ws[f'{cell}{total_row}'].number_format='#,##0.##'
    write(f'B{total_row+1}',number_to_vietnamese(scope['totals']['total_amount'])+'./.')
    # The original label occupied the wide STT column. Keep it on a full
    # line now that that column is appropriately narrow for an A4 table.
    words = f"Bằng chữ: {ws.cell(total_row+1,2).value}"
    ws.unmerge_cells(f'B{total_row+1}:I{total_row+1}')
    ws.cell(total_row+1,2).value = None
    ws.merge_cells(f'A{total_row+1}:I{total_row+1}')
    write(f'A{total_row+1}', words)
    write(f'E{total_row+2}',f'Hải Phòng, ngày {end.day:02d} tháng {end.month:02d} năm {end.year}')
    for row in (1,2,5,7,8):ws.row_dimensions[row].height=max(25,math.ceil(len(ws.cell(row,1).value or '')/65)*22)
    ws.row_dimensions[total_row+1].height=max(38,math.ceil(len(words)/75)*20)
    ws.row_dimensions[10].height=60
    ws.row_dimensions[total_row].height=28
    ws.row_dimensions[total_row+2].height=40
    ws.row_dimensions[total_row+3].height=38
    ws.sheet_view.showGridLines=False;ws.freeze_panes='A11';ws.print_title_rows='10:10'
    ws.print_area=f'A1:I{total_row+6}'
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.orientation='portrait';ws.page_setup.paperSize=ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
    readable_payment_layout(ws, statement=True)
    # Paginate using the rendered A4 geometry, not unscaled Excel row heights.
    # Explicit estimated breaks leave large empty areas after fitting the width.
    # Avoid Excel's trailing decimal dot for integral amounts, without dropping
    # fractional quantities or provider VAT decimals.
    for row in ws:
        for cell in row:
            if isinstance(cell.value,(int,float)) and float(cell.value).is_integer():
                cell.number_format=cell.number_format.replace('.######','').replace('.##','')
    proof=book.create_sheet('Đối chiếu hóa đơn')
    proof.append(['Ngày hóa đơn','Ký hiệu','Số hóa đơn','Thành tiền','Tiền thuế','Thanh toán','Nguồn xác minh','ID nguồn','ID dự thảo','Chênh lệch thành tiền dòng - HĐ','Chênh lệch thuế dòng - HĐ','Chênh lệch thanh toán dòng - HĐ'])
    for invoice in scope['invoices']:
        identity=('local',invoice['draft_id']) if invoice.get('draft_id') else ('source',invoice['source_invoice_id'])
        proof.append([_excel_text(str(invoice.get(k) or '')) if k in ('invoice_date','invoice_series','invoice_number','verification_source') else invoice.get(k)
                      for k in ('invoice_date','invoice_series','invoice_number','subtotal','tax_amount','total_amount','verification_source','source_invoice_id','draft_id')] +
                     [float(actual[identity][i]-numeric(invoice[k],k)) for i,k in enumerate(('subtotal','tax_amount','total_amount'))])
    proof.sheet_state='hidden'
    return book
