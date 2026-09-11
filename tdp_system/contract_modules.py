from __future__ import annotations

import io
import hashlib
import json
import math
import os
import re
import subprocess
import threading
import time
import unicodedata
import uuid
import warnings
import zipfile
from collections import Counter, defaultdict
from copy import deepcopy
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from statistics import median

from flask import jsonify, request, send_file
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel

try:
    from .purchase_money_adjustments import DEDUCTION_KIND, DEDUCTION_LABEL, approved_phong_source
except ImportError:
    from purchase_money_adjustments import DEDUCTION_KIND, DEDUCTION_LABEL, approved_phong_source


PRINT_TEMPLATE_DIR = Path(__file__).resolve().parent / "templates"
PAYROLL_TEMPLATE_SOURCE = PRINT_TEMPLATE_DIR / "payroll_template.xlsx"

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt
except ImportError:  # Kept optional until a payment-request DOCX is requested.
    Document = None

try:
    from msmi_client import MsmiClient, MsmiConfig, MsmiError
except ImportError:
    from .msmi_client import MsmiClient, MsmiConfig, MsmiError

try:
    from minvoice_client import MinvoiceError, MinvoiceOutcomeUnknown
except ImportError:
    from .minvoice_client import MinvoiceError, MinvoiceOutcomeUnknown

try:
    from outgoing_readiness import (
        invoice_order_issues,
        batch_product_stock_trace,
        OutgoingReadinessError,
        batch_readiness_payload,
        canonical_available_stock,
        validate_issued_draft_stock,
        validate_draft_export_stock,
        period_shortage_payload,
        shortage_workbook_bytes,
        validate_demand_orders,
    )
except ImportError:
    from .outgoing_readiness import (
        invoice_order_issues,
        batch_product_stock_trace,
        OutgoingReadinessError,
        batch_readiness_payload,
        canonical_available_stock,
        validate_issued_draft_stock,
        validate_draft_export_stock,
        period_shortage_payload,
        shortage_workbook_bytes,
        validate_demand_orders,
    )

try:
    from outgoing_substitution import mark_substitution_actions_reversed_for_draft
except ImportError:
    from .outgoing_substitution import mark_substitution_actions_reversed_for_draft

try:
    from invoice_payment_scope import (
        InvoicePaymentScopeError,
        issued_invoice_payment_scope,
        public_payment_scope,
    )
except ImportError:
    from .invoice_payment_scope import (
        InvoicePaymentScopeError,
        issued_invoice_payment_scope,
        public_payment_scope,
    )

try:
    from invoice_delivery_statement import (
        InvoiceDeliveryStatementError,
        invoice_delivery_statement_workbook as _invoice_delivery_statement_workbook,
    )
except ImportError:
    from .invoice_delivery_statement import (
        InvoiceDeliveryStatementError,
        invoice_delivery_statement_workbook as _invoice_delivery_statement_workbook,
    )

try:
    from invoice_payment_documents import (
        InvoicePaymentDocumentError,
        invoice_payment_request_workbook,
    )
except ImportError:
    from .invoice_payment_documents import (
        InvoicePaymentDocumentError,
        invoice_payment_request_workbook,
    )

try:
    from excel_print_renderer import ExcelPrintError, build_excel_pdf_bundle, verify_excel_pdf
    from pdf_documents import write_manifest
except ImportError:
    from .excel_print_renderer import ExcelPrintError, build_excel_pdf_bundle, verify_excel_pdf
    from .pdf_documents import write_manifest

try:
    from xcom_payment_documents import (
        XcomPaymentError,
        assign_payment_scope,
        consume_payment_preview,
        create_payment_preview,
        delete_meal_tariff,
        delete_payment_profile,
        init_xcom_payment_schema,
        list_payment_profiles,
        remove_payment_scope,
        upsert_meal_tariff,
        upsert_payment_profile,
    )
except ImportError:
    from .xcom_payment_documents import (
        XcomPaymentError,
        assign_payment_scope,
        consume_payment_preview,
        create_payment_preview,
        delete_meal_tariff,
        delete_payment_profile,
        init_xcom_payment_schema,
        list_payment_profiles,
        remove_payment_scope,
        upsert_meal_tariff,
        upsert_payment_profile,
    )


ADVANCED_SCHEMA = """
CREATE TABLE IF NOT EXISTS audit_log (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    event_type TEXT NOT NULL,
    entity_type TEXT,
    entity_id TEXT,
    status TEXT NOT NULL,
    message TEXT,
    metadata_json TEXT NOT NULL DEFAULT '{}',
    created_at TEXT NOT NULL
);
CREATE INDEX IF NOT EXISTS idx_audit_created ON audit_log(created_at DESC);

CREATE TABLE IF NOT EXISTS supplier_rules (
    supplier_code TEXT PRIMARY KEY,
    combine_kitchens INTEGER NOT NULL DEFAULT 0,
    updated_at TEXT NOT NULL
);

CREATE TABLE IF NOT EXISTS supplier_order_statuses (
    batch_id INTEGER NOT NULL REFERENCES batches(id) ON DELETE CASCADE,
    supplier_key TEXT NOT NULL,
    supplier_label TEXT NOT NULL DEFAULT '',
    status TEXT NOT NULL DEFAULT 'pending',
    revision INTEGER NOT NULL DEFAULT 1,
    ordered_at TEXT,
    reopened_at TEXT,
    updated_at TEXT NOT NULL,
    PRIMARY KEY(batch_id,supplier_key),
    CHECK(status IN ('pending','ordered','reopened')),
    CHECK(revision >= 1)
);
CREATE INDEX IF NOT EXISTS idx_supplier_order_statuses_batch_status
    ON supplier_order_statuses(batch_id,status,supplier_key);

CREATE TABLE IF NOT EXISTS inventory_transactions (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    txn_date TEXT NOT NULL,
    product_code TEXT NOT NULL,
    warehouse_codes_json TEXT NOT NULL DEFAULT '[]',
    qty_in REAL NOT NULL DEFAULT 0,
    qty_out REAL NOT NULL DEFAULT 0,
    unit_cost REAL NOT NULL DEFAULT 0,
    source_type TEXT NOT NULL,
    source_id TEXT NOT NULL,
    source_line TEXT NOT NULL DEFAULT '',
    kitchen TEXT,
    status TEXT NOT NULL DEFAULT 'posted',
    note TEXT,
    created_at TEXT NOT NULL,
    updated_at TEXT NOT NULL,
    UNIQUE(source_type, source_id, source_line)
);
CREATE INDEX IF NOT EXISTS idx_inventory_product_date ON inventory_transactions(product_code,txn_date);
CREATE INDEX IF NOT EXISTS idx_inventory_source ON inventory_transactions(source_type,source_id);

CREATE TABLE IF NOT EXISTS msmi_invoices (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    remote_id TEXT NOT NULL UNIQUE,
    tenant TEXT NOT NULL DEFAULT 'default',
    invoice_type TEXT NOT NULL,
    seller_tax_code TEXT,
    seller_name TEXT,
    invoice_number TEXT,
    invoice_series TEXT,
    invoice_date TEXT,
    subtotal REAL NOT NULL DEFAULT 0,
    tax_amount REAL NOT NULL DEFAULT 0,
    total_amount REAL NOT NULL DEFAULT 0,
    sync_status TEXT NOT NULL DEFAULT 'synced',
    receipt_status TEXT NOT NULL DEFAULT 'pending_mapping',
    raw_json TEXT NOT NULL,
    error_message TEXT,
    synced_at TEXT NOT NULL,
    created_at TEXT NOT NULL,
    updated_at TEXT NOT NULL
);
CREATE INDEX IF NOT EXISTS idx_msmi_invoice_date ON msmi_invoices(invoice_date DESC);

CREATE TABLE IF NOT EXISTS msmi_invoice_items (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    invoice_id INTEGER NOT NULL REFERENCES msmi_invoices(id) ON DELETE CASCADE,
    line_index INTEGER NOT NULL,
    source_item_code TEXT,
    source_item_name TEXT,
    source_unit TEXT,
    qty REAL NOT NULL DEFAULT 0,
    unit_price REAL NOT NULL DEFAULT 0,
    amount REAL NOT NULL DEFAULT 0,
    tax_rate TEXT,
    source_nature TEXT,
    inventory_eligible INTEGER NOT NULL DEFAULT 1,
    validation_note TEXT NOT NULL DEFAULT '',
    product_code TEXT,
    mapping_status TEXT NOT NULL DEFAULT 'unmapped',
    UNIQUE(invoice_id,line_index)
);
CREATE INDEX IF NOT EXISTS idx_msmi_items_invoice ON msmi_invoice_items(invoice_id);

CREATE TABLE IF NOT EXISTS item_mappings (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    tenant TEXT NOT NULL DEFAULT 'default',
    seller_tax_code TEXT NOT NULL DEFAULT '',
    source_item_code TEXT NOT NULL DEFAULT '',
    source_item_name TEXT NOT NULL DEFAULT '',
    product_code TEXT NOT NULL,
    updated_at TEXT NOT NULL,
    UNIQUE(tenant,seller_tax_code,source_item_code,source_item_name)
);

CREATE TABLE IF NOT EXISTS msmi_sync_state (
    invoice_type TEXT PRIMARY KEY,
    last_remote_id TEXT,
    last_invoice_date TEXT,
    last_synced_at TEXT,
    last_status TEXT,
    last_error TEXT,
    backfill_anchor_id TEXT,
    backfill_anchor_date TEXT,
    backfill_complete INTEGER NOT NULL DEFAULT 0
);

CREATE TABLE IF NOT EXISTS outgoing_invoice_drafts (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    batch_id INTEGER NOT NULL,
    contractor TEXT NOT NULL,
    invoice_date TEXT NOT NULL,
    status TEXT NOT NULL DEFAULT 'draft',
    subtotal REAL NOT NULL DEFAULT 0,
    tax_amount REAL NOT NULL DEFAULT 0,
    total_amount REAL NOT NULL DEFAULT 0,
    created_at TEXT NOT NULL,
    issued_at TEXT,
    round_no INTEGER NOT NULL DEFAULT 1,
    draft_kind TEXT NOT NULL DEFAULT 'standard',
    UNIQUE(batch_id,contractor,round_no)
);
CREATE TABLE IF NOT EXISTS outgoing_invoice_lines (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    draft_id INTEGER NOT NULL REFERENCES outgoing_invoice_drafts(id) ON DELETE CASCADE,
    order_id INTEGER NOT NULL,
    product_code TEXT NOT NULL,
    product_name TEXT,
    qty REAL NOT NULL,
    unit TEXT,
    unit_price REAL NOT NULL,
    tax TEXT,
    invoice_nature TEXT NOT NULL DEFAULT '1',
    amount REAL NOT NULL,
    UNIQUE(draft_id,order_id)
);

CREATE TABLE IF NOT EXISTS outgoing_product_names (
    product_code TEXT PRIMARY KEY,
    invoice_name TEXT NOT NULL,
    updated_at TEXT NOT NULL
);

CREATE TABLE IF NOT EXISTS outgoing_buyer_profiles (
    contractor TEXT PRIMARY KEY,
    display_name TEXT,
    legal_name TEXT,
    tax_code TEXT,
    address TEXT NOT NULL,
    email TEXT,
    updated_at TEXT NOT NULL
);

CREATE TABLE IF NOT EXISTS purchase_order_lines (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    batch_id INTEGER NOT NULL,
    order_id INTEGER NOT NULL UNIQUE,
    demand_qty REAL NOT NULL DEFAULT 0,
    physical_stock_used REAL NOT NULL DEFAULT 0,
    order_qty REAL NOT NULL DEFAULT 0,
    supplier TEXT NOT NULL,
    buy_price REAL NOT NULL DEFAULT 0,
    price_source TEXT NOT NULL DEFAULT '',
    note TEXT NOT NULL DEFAULT '',
    status TEXT NOT NULL DEFAULT 'confirmed',
    source_hash TEXT NOT NULL DEFAULT '',
    created_at TEXT NOT NULL,
    updated_at TEXT NOT NULL
);
CREATE INDEX IF NOT EXISTS idx_purchase_order_lines_batch
    ON purchase_order_lines(batch_id,order_id);

CREATE TABLE IF NOT EXISTS purchase_order_imports (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    batch_id INTEGER NOT NULL,
    source_hash TEXT NOT NULL,
    source_name TEXT NOT NULL,
    line_count INTEGER NOT NULL DEFAULT 0,
    total_amount REAL NOT NULL DEFAULT 0,
    imported_at TEXT NOT NULL,
    UNIQUE(batch_id,source_hash)
);

CREATE TABLE IF NOT EXISTS purchase_workbook_lines (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    batch_id INTEGER NOT NULL REFERENCES batches(id) ON DELETE CASCADE,
    row_key TEXT NOT NULL,
    order_id INTEGER REFERENCES orders(id) ON DELETE SET NULL,
    source_sheet TEXT NOT NULL DEFAULT '',
    source_row INTEGER NOT NULL DEFAULT 0,
    product_code TEXT NOT NULL DEFAULT '',
    kitchen TEXT NOT NULL DEFAULT '',
    work_date TEXT NOT NULL,
    product_name TEXT NOT NULL DEFAULT '',
    base_qty REAL NOT NULL DEFAULT 0,
    unit TEXT NOT NULL DEFAULT '',
    supplier TEXT NOT NULL DEFAULT '',
    note TEXT NOT NULL DEFAULT '',
    buy_price REAL NOT NULL DEFAULT 0,
    price_source TEXT NOT NULL DEFAULT '',
    damaged_qty REAL NOT NULL DEFAULT 0,
    added_qty REAL NOT NULL DEFAULT 0,
    reduced_qty REAL NOT NULL DEFAULT 0,
    missing_qty REAL NOT NULL DEFAULT 0,
    actual_qty REAL NOT NULL DEFAULT 0,
    amount REAL NOT NULL DEFAULT 0,
    status TEXT NOT NULL DEFAULT 'confirmed',
    source_hash TEXT NOT NULL DEFAULT '',
    revision INTEGER NOT NULL DEFAULT 1,
    created_at TEXT NOT NULL,
    updated_at TEXT NOT NULL,
    UNIQUE(batch_id,row_key),
    CHECK(source_row >= 0),
    CHECK(revision >= 1)
);
CREATE INDEX IF NOT EXISTS idx_purchase_workbook_lines_batch
    ON purchase_workbook_lines(batch_id,source_row,id);
CREATE INDEX IF NOT EXISTS idx_purchase_workbook_lines_supplier_date
    ON purchase_workbook_lines(supplier,work_date,batch_id);

CREATE TABLE IF NOT EXISTS purchase_workbook_line_revisions (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    batch_id INTEGER NOT NULL REFERENCES batches(id) ON DELETE CASCADE,
    row_key TEXT NOT NULL,
    order_id INTEGER,
    revision INTEGER NOT NULL,
    change_kind TEXT NOT NULL,
    source_hash TEXT NOT NULL,
    source_row INTEGER NOT NULL DEFAULT 0,
    base_qty REAL NOT NULL DEFAULT 0,
    damaged_qty REAL NOT NULL DEFAULT 0,
    added_qty REAL NOT NULL DEFAULT 0,
    reduced_qty REAL NOT NULL DEFAULT 0,
    missing_qty REAL NOT NULL DEFAULT 0,
    actual_qty REAL NOT NULL DEFAULT 0,
    buy_price REAL NOT NULL DEFAULT 0,
    amount REAL NOT NULL DEFAULT 0,
    created_at TEXT NOT NULL,
    UNIQUE(batch_id,row_key,revision),
    CHECK(revision >= 1),
    CHECK(change_kind IN ('insert','update'))
);
CREATE INDEX IF NOT EXISTS idx_purchase_workbook_line_revisions_batch
    ON purchase_workbook_line_revisions(batch_id,row_key,revision);

CREATE TABLE IF NOT EXISTS debt_adjustments (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    adjustment_date TEXT NOT NULL,
    party_type TEXT NOT NULL,
    party_code TEXT NOT NULL,
    amount REAL NOT NULL,
    note TEXT,
    created_at TEXT NOT NULL
);

CREATE TABLE IF NOT EXISTS historical_payable_lines (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    purchase_date TEXT NOT NULL,
    kitchen TEXT,
    item_name TEXT NOT NULL,
    qty REAL NOT NULL DEFAULT 0,
    unit TEXT,
    supplier TEXT NOT NULL,
    buy_price REAL NOT NULL DEFAULT 0,
    damaged_qty REAL NOT NULL DEFAULT 0,
    added_qty REAL NOT NULL DEFAULT 0,
    reduced_qty REAL NOT NULL DEFAULT 0,
    missing_qty REAL NOT NULL DEFAULT 0,
    actual_qty REAL NOT NULL DEFAULT 0,
    source_amount REAL NOT NULL DEFAULT 0,
    calculated_amount REAL NOT NULL DEFAULT 0,
    amount REAL NOT NULL DEFAULT 0,
    note TEXT,
    source_file TEXT NOT NULL,
    source_sheet TEXT NOT NULL,
    source_row INTEGER NOT NULL,
    source_hash TEXT NOT NULL,
    created_at TEXT NOT NULL,
    updated_at TEXT NOT NULL,
    UNIQUE(source_file,source_sheet,source_row)
);
CREATE INDEX IF NOT EXISTS idx_historical_payable_date_supplier
    ON historical_payable_lines(purchase_date,supplier);

CREATE TABLE IF NOT EXISTS kitchen_units (
    kitchen_code TEXT PRIMARY KEY,
    unit_code TEXT NOT NULL,
    updated_at TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS dated_prices (
    product_code TEXT NOT NULL,
    price_group TEXT NOT NULL,
    period TEXT NOT NULL,
    price_value REAL NOT NULL,
    updated_at TEXT NOT NULL,
    PRIMARY KEY(product_code,price_group,period)
);
CREATE TABLE IF NOT EXISTS meal_plans (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    work_date TEXT NOT NULL,
    kitchen TEXT NOT NULL,
    shift TEXT NOT NULL,
    meal_count REAL NOT NULL,
    unit_code TEXT,
    status TEXT NOT NULL DEFAULT 'draft',
    note TEXT,
    created_at TEXT NOT NULL,
    updated_at TEXT NOT NULL
);
CREATE INDEX IF NOT EXISTS idx_meal_plan_date ON meal_plans(work_date,kitchen);
CREATE TABLE IF NOT EXISTS meal_plan_items (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    plan_id INTEGER NOT NULL REFERENCES meal_plans(id) ON DELETE CASCADE,
    dish_name TEXT,
    product_code TEXT NOT NULL,
    product_name TEXT,
    norm_qty REAL NOT NULL,
    unit TEXT,
    supplier TEXT,
    buy_price REAL NOT NULL DEFAULT 0,
    price_source TEXT,
    UNIQUE(plan_id,dish_name,product_code)
);

CREATE TABLE IF NOT EXISTS meal_attendance (
    work_date TEXT NOT NULL,
    kitchen TEXT NOT NULL,
    shift TEXT NOT NULL,
    actual_count REAL NOT NULL DEFAULT 0,
    ordered_count REAL NOT NULL DEFAULT 0,
    source_type TEXT NOT NULL DEFAULT 'MONTHLY_WORKBOOK',
    source_file TEXT,
    source_sheet TEXT,
    source_column TEXT,
    updated_at TEXT NOT NULL,
    PRIMARY KEY(work_date,kitchen,shift)
);
CREATE INDEX IF NOT EXISTS idx_meal_attendance_month
    ON meal_attendance(work_date,kitchen);

CREATE TABLE IF NOT EXISTS staff (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    employee_code TEXT NOT NULL UNIQUE,
    full_name TEXT NOT NULL,
    role_name TEXT,
    kitchen TEXT,
    base_salary REAL NOT NULL DEFAULT 0,
    standard_days REAL NOT NULL DEFAULT 26,
    standard_hours REAL NOT NULL DEFAULT 8,
    bhxh_employee_rate REAL NOT NULL DEFAULT 0,
    bhxh_company_rate REAL NOT NULL DEFAULT 0,
    active INTEGER NOT NULL DEFAULT 1,
    created_at TEXT NOT NULL,
    updated_at TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS attendance_entries (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    employee_id INTEGER NOT NULL REFERENCES staff(id) ON DELETE CASCADE,
    work_date TEXT NOT NULL,
    normal_hours REAL NOT NULL DEFAULT 0,
    overtime_hours REAL NOT NULL DEFAULT 0,
    sunday_hours REAL NOT NULL DEFAULT 0,
    night_hours REAL NOT NULL DEFAULT 0,
    holiday_hours REAL NOT NULL DEFAULT 0,
    note TEXT,
    source TEXT,
    updated_at TEXT NOT NULL,
    UNIQUE(employee_id,work_date)
);
CREATE TABLE IF NOT EXISTS payroll_adjustments (
    employee_id INTEGER NOT NULL REFERENCES staff(id) ON DELETE CASCADE,
    month TEXT NOT NULL,
    allowance REAL NOT NULL DEFAULT 0,
    responsibility REAL NOT NULL DEFAULT 0,
    advance REAL NOT NULL DEFAULT 0,
    probation_deduction REAL NOT NULL DEFAULT 0,
    bhxh_employee_amount REAL NOT NULL DEFAULT 0,
    bhxh_company_amount REAL NOT NULL DEFAULT 0,
    note TEXT,
    updated_at TEXT NOT NULL,
    PRIMARY KEY(employee_id,month)
);
CREATE TABLE IF NOT EXISTS kitchen_labor_costs (
    work_date TEXT NOT NULL,
    kitchen TEXT NOT NULL,
    amount REAL NOT NULL DEFAULT 0,
    source TEXT,
    updated_at TEXT NOT NULL,
    PRIMARY KEY(work_date,kitchen,source)
);

CREATE TABLE IF NOT EXISTS print_jobs (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    batch_id INTEGER NOT NULL,
    document_type TEXT NOT NULL,
    file_path TEXT NOT NULL,
    status TEXT NOT NULL DEFAULT 'prepared',
    approved_at TEXT,
    printed_at TEXT,
    printer_name TEXT,
    error_message TEXT,
    created_at TEXT NOT NULL,
    UNIQUE(batch_id,document_type)
);
"""


MAPPING_IMPORT_MAX_BYTES = 10 * 1024 * 1024
MAPPING_IMPORT_MAX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAPPING_IMPORT_MAX_ROWS = 10_000
MAPPING_IMPORT_TTL_SECONDS = 30 * 60
PENDING_MAPPING_IMPORTS = {}
MAPPING_IMPORT_LOCK = threading.Lock()
PENDING_LEGACY_INVOICE_MAPPING_IMPORTS = {}
LEGACY_INVOICE_MAPPING_IMPORT_LOCK = threading.Lock()
PENDING_CATALOG_IMPORTS = {}
CATALOG_IMPORT_LOCK = threading.Lock()
PENDING_KITCHEN_IMPORTS = {}
KITCHEN_IMPORT_LOCK = threading.Lock()
PENDING_OPENING_IMPORTS = {}
OPENING_IMPORT_LOCK = threading.Lock()
PENDING_MEAL_ATTENDANCE_IMPORTS = {}
MEAL_ATTENDANCE_IMPORT_LOCK = threading.Lock()
PENDING_LEGACY_ATTENDANCE_IMPORTS = {}
LEGACY_ATTENDANCE_IMPORT_LOCK = threading.Lock()
PENDING_PAYABLE_IMPORTS = {}
PAYABLE_IMPORT_LOCK = threading.Lock()
PENDING_PURCHASE_ORDER_IMPORTS = {}
PURCHASE_ORDER_IMPORT_LOCK = threading.Lock()
PRINT_SUBMISSION_LOCK = threading.Lock()

MAPPING_ALIASES = {
    "invoice_names": {
        "product_code": {
            "mahang", "mahanghoa", "mahh", "mavt", "mavattu", "mavattutdp", "masp",
            "masanpham", "matdp", "masptdp", "mahhtdp",
            "productcode", "itemcode", "code",
        },
        "product_name": {
            "tenhang", "tenhanghoa", "tenhanghoatdp", "tenhh", "tenvattu", "tensp",
            "tensanpham", "tentdp", "tenthanhdatphat", "tentrenphanmem", "tenhientai",
            "tengoc", "productname", "itemname",
        },
        "target_value": {
            "tenxuathoadon", "tenxuathd", "tenxhd", "tenhoadon", "tenhd", "tendaura",
            "tenhangxuathoadon", "tenhanghoaxuathoadon", "tenhangxhd", "tenhanghoadon",
            "tenchuanhoadon", "invoicename",
        },
    },
    "kitchen_units": {
        "source_code": {
            "mabep", "bep", "kitchen", "kitchencode", "code",
        },
        "source_name": {
            "tenbep", "tennhabep", "kitchenname",
        },
        "target_value": {
            "unit", "maunit", "unitcode", "donvi", "donvigop", "nhomunit",
            "xcom", "maxcom", "xuongcom", "maxuongcom", "xcomxuongcom",
        },
    },
}

CATALOG_ALIASES = {
    "product_code": {
        "mahang", "mahanghoa", "mahh", "mavt", "mavattu", "matdp", "masp",
        "masanpham", "mahhtdp", "productcode", "itemcode", "code",
    },
    "product_group": {"nhomhang", "nhomhanghoa", "nhom", "category", "group"},
    "product_name": {
        "tenhang", "tenhanghoa", "tenhanghoatdp", "tentdp", "tenthanhdatphat",
        "productname", "itemname",
    },
    "invoice_name": {
        "tenxuathoadon", "tenxuathd", "tenxhd", "tenhoadon", "tenhd", "tendaura",
        "tenhangxuathoadon", "tenhanghoaxuathoadon", "tenhangxhd", "invoicename",
    },
    "unit": {"dvt", "donvitinh", "unit"},
    "tax": {"thue", "thuesuat", "thuegtgt", "tsuat", "vat", "tax"},
}

CATALOG_CANONICAL_UNITS = {
    "kg": "Kg", "cai": "Cái", "goi": "Gói", "hop": "Hộp", "chai": "Chai",
    "can": "Can", "qua": "Quả", "thung": "Thùng", "lo": "Lọ", "tui": "Túi",
    "doi": "Đôi", "lit": "Lít", "cuon": "Cuộn", "bich": "Bịch", "coc": "Cốc",
    "vi": "Vỉ", "bo": "Bộ", "mo": "Mớ", "la": "Lá", "bao": "Bao",
    "tuyp": "Tuýp", "con": "Con", "to": "Tô", "binh": "Bình",
    "cay": "Cây", "vien": "Viên", "ly": "Ly", "cu": "Củ", "tap": "Tập",
    "lon": "Lon", "ong": "Ống", "le": "Lễ", "tep": "Tệp", "dia": "Đĩa",
    "mieng": "Miếng", "day": "Dây", "suat": "Suất",
}


def ensure_column(conn, table: str, name: str, definition: str):
    columns = {row["name"] for row in conn.execute(f"PRAGMA table_info({table})")}
    if name not in columns:
        conn.execute(f"ALTER TABLE {table} ADD COLUMN {name} {definition}")


def migrate_outgoing_invoice_rounds(conn):
    """Allow several partial invoice rounds for one contractor and batch."""
    columns = {row["name"] for row in conn.execute("PRAGMA table_info(outgoing_invoice_drafts)")}
    if "round_no" in columns:
        return
    conn.execute("ALTER TABLE outgoing_invoice_lines RENAME TO outgoing_invoice_lines_legacy")
    conn.execute("ALTER TABLE outgoing_invoice_drafts RENAME TO outgoing_invoice_drafts_legacy")
    conn.executescript(
        """
        CREATE TABLE outgoing_invoice_drafts (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            batch_id INTEGER NOT NULL,
            contractor TEXT NOT NULL,
            invoice_date TEXT NOT NULL,
            status TEXT NOT NULL DEFAULT 'draft',
            subtotal REAL NOT NULL DEFAULT 0,
            tax_amount REAL NOT NULL DEFAULT 0,
            total_amount REAL NOT NULL DEFAULT 0,
            created_at TEXT NOT NULL,
            issued_at TEXT,
            round_no INTEGER NOT NULL DEFAULT 1,
            draft_kind TEXT NOT NULL DEFAULT 'standard',
            minvoice_status TEXT NOT NULL DEFAULT 'not_sent',
            minvoice_series TEXT,
            minvoice_remote_id TEXT,
            minvoice_saved_at TEXT,
            minvoice_error TEXT,
            minvoice_key_api TEXT,
            minvoice_started_at TEXT,
            minvoice_reconciled_at TEXT,
            external_key_uuid TEXT,
            issued_invoice_number TEXT,
            issued_invoice_series TEXT,
            issued_invoice_date TEXT,
            buyer_name_snapshot TEXT,
            buyer_tax_code_snapshot TEXT,
            buyer_address_snapshot TEXT,
            buyer_email_snapshot TEXT,
            company_name_snapshot TEXT,
            company_tax_code_snapshot TEXT,
            company_address_snapshot TEXT,
            payment_requester_snapshot TEXT,
            payment_bank_name_snapshot TEXT,
            payment_bank_account_snapshot TEXT,
            UNIQUE(batch_id,contractor,round_no)
        );
        CREATE TABLE outgoing_invoice_lines (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            draft_id INTEGER NOT NULL REFERENCES outgoing_invoice_drafts(id) ON DELETE CASCADE,
            order_id INTEGER NOT NULL,
            product_code TEXT NOT NULL,
            product_name TEXT,
            qty REAL NOT NULL,
            unit TEXT,
            unit_price REAL NOT NULL,
            tax TEXT,
            invoice_nature TEXT NOT NULL DEFAULT '1',
            amount REAL NOT NULL,
            UNIQUE(draft_id,order_id)
        );
        """
    )
    new_draft_columns = [row["name"] for row in conn.execute("PRAGMA table_info(outgoing_invoice_drafts)")]
    old_draft_columns = {row["name"] for row in conn.execute("PRAGMA table_info(outgoing_invoice_drafts_legacy)")}
    copied = [name for name in new_draft_columns if name in old_draft_columns]
    column_sql = ",".join(f'"{name}"' for name in copied)
    conn.execute(
        f"INSERT INTO outgoing_invoice_drafts({column_sql}) "
        f"SELECT {column_sql} FROM outgoing_invoice_drafts_legacy"
    )
    new_line_columns = [row["name"] for row in conn.execute("PRAGMA table_info(outgoing_invoice_lines)")]
    old_line_columns = {row["name"] for row in conn.execute("PRAGMA table_info(outgoing_invoice_lines_legacy)")}
    copied_lines = [name for name in new_line_columns if name in old_line_columns]
    line_sql = ",".join(f'"{name}"' for name in copied_lines)
    conn.execute(
        f"INSERT INTO outgoing_invoice_lines({line_sql}) "
        f"SELECT {line_sql} FROM outgoing_invoice_lines_legacy"
    )
    conn.execute("DROP TABLE outgoing_invoice_lines_legacy")
    conn.execute("DROP TABLE outgoing_invoice_drafts_legacy")


def init_contract_schema(conn, opening_template_path=None):
    conn.executescript(ADVANCED_SCHEMA)
    try:
        from .outgoing_consolidation import SCHEMA as CONSOLIDATION_SCHEMA
    except ImportError:
        from outgoing_consolidation import SCHEMA as CONSOLIDATION_SCHEMA
    try:
        from .catalog_products import SCHEMA as CATALOG_SCHEMA
    except ImportError:
        from catalog_products import SCHEMA as CATALOG_SCHEMA
    conn.executescript(CATALOG_SCHEMA)
    migrate_outgoing_invoice_rounds(conn)
    conn.execute(
        "INSERT OR IGNORE INTO settings(key,value) VALUES('installation_uuid',?)",
        (uuid.uuid4().hex.upper(),),
    )
    # Schema creation belongs to application startup/migration only.  Route
    # helpers below never run DDL inside a business transaction.
    init_xcom_payment_schema(conn)
    ensure_column(conn, "products", "product_group", "TEXT")
    ensure_column(conn, "products", "catalog_updated_at", "TEXT")
    # Preserve the customer's distinct `MÃ KHO` values from each opening
    # snapshot.  They are presentation/source identity, not interchangeable
    # with the TĐP product code.
    ensure_column(conn, "inventory_transactions", "warehouse_codes_json", "TEXT NOT NULL DEFAULT '[]'")
    backfill_opening_warehouse_codes(conn, opening_template_path)
    ensure_column(conn, "balances", "as_of_date", "TEXT NOT NULL DEFAULT '1900-01-01'")
    # Older builds only remembered the newest mSMI invoice.  That was not enough
    # to continue an initial history import after max_pages was reached: the next
    # run saw an already-known first page and stopped forever.  The anchor marks
    # the *oldest fully processed position* and is deliberately migrated as
    # incomplete so an existing database safely resumes/backfills idempotently.
    ensure_column(conn, "msmi_sync_state", "backfill_anchor_id", "TEXT")
    ensure_column(conn, "msmi_sync_state", "backfill_anchor_date", "TEXT")
    ensure_column(conn, "msmi_sync_state", "backfill_complete", "INTEGER NOT NULL DEFAULT 0")
    ensure_column(conn, "msmi_sync_state", "reconcile_next_page", "INTEGER NOT NULL DEFAULT 1")
    ensure_column(conn, "msmi_invoice_items", "source_nature", "TEXT")
    ensure_column(conn, "msmi_invoice_items", "inventory_eligible", "INTEGER NOT NULL DEFAULT 1")
    ensure_column(conn, "msmi_invoice_items", "validation_note", "TEXT NOT NULL DEFAULT ''")
    ensure_column(conn, "orders", "damaged_qty", "REAL NOT NULL DEFAULT 0")
    ensure_column(conn, "orders", "supplier_return_qty", "REAL NOT NULL DEFAULT 0")
    ensure_column(conn, "orders", "customer_return_qty", "REAL NOT NULL DEFAULT 0")
    ensure_column(conn, "outgoing_invoice_drafts", "draft_kind", "TEXT NOT NULL DEFAULT 'standard'")
    ensure_column(conn, "outgoing_invoice_drafts", "minvoice_status", "TEXT NOT NULL DEFAULT 'not_sent'")
    ensure_column(conn, "outgoing_invoice_drafts", "minvoice_series", "TEXT")
    ensure_column(conn, "outgoing_invoice_drafts", "minvoice_remote_id", "TEXT")
    ensure_column(conn, "outgoing_invoice_drafts", "minvoice_saved_at", "TEXT")
    ensure_column(conn, "outgoing_invoice_drafts", "minvoice_error", "TEXT")
    ensure_column(conn, "outgoing_invoice_drafts", "minvoice_key_api", "TEXT")
    ensure_column(conn, "outgoing_invoice_drafts", "minvoice_started_at", "TEXT")
    ensure_column(conn, "outgoing_invoice_drafts", "minvoice_reconciled_at", "TEXT")
    ensure_column(conn, "outgoing_invoice_drafts", "external_key_uuid", "TEXT")
    for draft_row in conn.execute(
        "SELECT id FROM outgoing_invoice_drafts WHERE TRIM(COALESCE(external_key_uuid,''))=''"
    ):
        conn.execute(
            "UPDATE outgoing_invoice_drafts SET external_key_uuid=? WHERE id=?",
            (uuid.uuid4().hex.upper(), draft_row["id"]),
        )
    conn.execute(
        """CREATE UNIQUE INDEX IF NOT EXISTS idx_outgoing_external_key_uuid
           ON outgoing_invoice_drafts(external_key_uuid)
           WHERE external_key_uuid IS NOT NULL AND external_key_uuid!=''"""
    )
    ensure_column(conn, "outgoing_invoice_drafts", "issued_invoice_number", "TEXT")
    ensure_column(conn, "outgoing_invoice_drafts", "issued_invoice_series", "TEXT")
    ensure_column(conn, "outgoing_invoice_drafts", "issued_invoice_date", "TEXT")
    ensure_column(conn, "outgoing_invoice_drafts", "buyer_name_snapshot", "TEXT")
    ensure_column(conn, "outgoing_invoice_drafts", "buyer_tax_code_snapshot", "TEXT")
    ensure_column(conn, "outgoing_invoice_drafts", "buyer_address_snapshot", "TEXT")
    ensure_column(conn, "outgoing_invoice_drafts", "buyer_email_snapshot", "TEXT")
    ensure_column(conn, "outgoing_invoice_drafts", "company_name_snapshot", "TEXT")
    ensure_column(conn, "outgoing_invoice_drafts", "company_tax_code_snapshot", "TEXT")
    ensure_column(conn, "outgoing_invoice_drafts", "company_address_snapshot", "TEXT")
    ensure_column(conn, "outgoing_invoice_drafts", "payment_requester_snapshot", "TEXT")
    ensure_column(conn, "outgoing_invoice_drafts", "payment_bank_name_snapshot", "TEXT")
    ensure_column(conn, "outgoing_invoice_drafts", "payment_bank_account_snapshot", "TEXT")
    ensure_column(conn, "outgoing_invoice_lines", "invoice_nature", "TEXT NOT NULL DEFAULT '1'")
    ensure_column(conn, "purchase_order_lines", "price_source", "TEXT NOT NULL DEFAULT ''")
    ensure_column(conn, "purchase_workbook_lines", "price_source", "TEXT NOT NULL DEFAULT ''")
    ensure_column(conn, "purchase_workbook_lines", "line_kind", "TEXT NOT NULL DEFAULT 'goods'")
    ensure_column(conn, "purchase_workbook_line_revisions", "line_kind", "TEXT NOT NULL DEFAULT 'goods'")
    ensure_column(conn, "historical_payable_lines", "source_amount", "REAL NOT NULL DEFAULT 0")
    ensure_column(conn, "historical_payable_lines", "calculated_amount", "REAL NOT NULL DEFAULT 0")
    ensure_column(conn, "historical_payable_lines", "source_hash", "TEXT NOT NULL DEFAULT ''")
    ensure_column(conn, "historical_payable_lines", "source_sheet", "TEXT NOT NULL DEFAULT ''")
    ensure_column(conn, "historical_payable_lines", "source_row", "INTEGER NOT NULL DEFAULT 0")
    conn.execute("DROP INDEX IF EXISTS idx_historical_payable_source_row")
    duplicate_payable_sources = [
        dict(row) for row in conn.execute(
            """SELECT source_hash,source_sheet,source_row,GROUP_CONCAT(id) ids,COUNT(*) count
               FROM historical_payable_lines
               GROUP BY source_hash,source_sheet,source_row HAVING COUNT(*)>1"""
        )
    ]
    if duplicate_payable_sources:
        conn.execute(
            """INSERT INTO settings(key,value)
               VALUES('schema_warning_duplicate_payable_source_ids',?)
               ON CONFLICT(key) DO UPDATE SET value=excluded.value""",
            (json.dumps(duplicate_payable_sources, ensure_ascii=False),),
        )
    else:
        conn.execute(
            "DELETE FROM settings WHERE key='schema_warning_duplicate_payable_source_ids'"
        )
        conn.execute(
            """CREATE UNIQUE INDEX idx_historical_payable_source_row
               ON historical_payable_lines(source_hash,source_sheet,source_row)"""
        )
    ensure_column(conn, "print_jobs", "file_sha256", "TEXT")
    ensure_column(conn, "print_jobs", "input_sha256", "TEXT")
    ensure_column(conn, "print_jobs", "manifest_path", "TEXT")
    ensure_column(conn, "print_jobs", "page_count", "INTEGER NOT NULL DEFAULT 0")
    ensure_column(conn, "print_jobs", "paper", "TEXT NOT NULL DEFAULT 'A4'")
    ensure_column(conn, "print_jobs", "copies", "INTEGER NOT NULL DEFAULT 1")
    ensure_column(conn, "print_jobs", "submitted_at", "TEXT")
    ensure_column(conn, "print_jobs", "claim_id", "TEXT")
    ensure_column(conn, "payroll_adjustments", "gross_override", "REAL NOT NULL DEFAULT 0")
    ensure_column(conn, "payroll_adjustments", "net_override", "REAL NOT NULL DEFAULT 0")
    ensure_column(conn, "payroll_adjustments", "use_override", "INTEGER NOT NULL DEFAULT 0")
    ensure_column(conn, "meal_plans", "import_key", "TEXT")
    ensure_column(conn, "meal_plans", "source_file", "TEXT")
    ensure_column(conn, "meal_plans", "source_sheet", "TEXT")
    ensure_column(conn, "meal_plans", "source_row_start", "INTEGER")
    ensure_column(conn, "meal_plans", "source_row_end", "INTEGER")
    ensure_column(conn, "meal_plans", "menu_count", "INTEGER NOT NULL DEFAULT 1")
    ensure_column(conn, "meal_plans", "servings_per_menu", "REAL NOT NULL DEFAULT 0")
    ensure_column(conn, "meal_plans", "meal_price", "REAL NOT NULL DEFAULT 0")
    ensure_column(conn, "meal_plans", "other_cost", "REAL NOT NULL DEFAULT 0")
    ensure_column(conn, "meal_plans", "source_financials_json", "TEXT NOT NULL DEFAULT '{}'")
    ensure_column(conn, "meal_plan_items", "source_norm_per_1000", "REAL NOT NULL DEFAULT 0")
    ensure_column(conn, "meal_plan_items", "applicable_meal_count", "REAL NOT NULL DEFAULT 0")
    ensure_column(conn, "meal_plan_items", "source_row", "INTEGER")
    ensure_column(conn, "meal_plan_items", "source_amount", "REAL NOT NULL DEFAULT 0")
    conn.execute(
        """CREATE UNIQUE INDEX IF NOT EXISTS idx_meal_plan_import_key
           ON meal_plans(import_key) WHERE import_key IS NOT NULL AND import_key!=''"""
    )
    # Canonical invoice identity is case-insensitive.  Legacy databases may
    # already contain conflicting variants; keep the app startable and surface
    # the exact IDs instead of crashing halfway through migration or deleting a
    # legal document automatically.
    conn.execute("DROP INDEX IF EXISTS idx_outgoing_issued_number")
    duplicate_invoice_ids = [
        dict(row) for row in conn.execute(
            """SELECT UPPER(TRIM(COALESCE(issued_invoice_series,''))) series_key,
                      TRIM(COALESCE(issued_invoice_number,'')) number_key,
                      GROUP_CONCAT(id) ids,COUNT(*) count
               FROM outgoing_invoice_drafts
               WHERE TRIM(COALESCE(issued_invoice_number,''))!=''
               GROUP BY series_key,number_key HAVING COUNT(*)>1"""
        )
    ]
    if duplicate_invoice_ids:
        conn.execute(
            """INSERT INTO settings(key,value) VALUES('schema_warning_duplicate_invoice_ids',?)
               ON CONFLICT(key) DO UPDATE SET value=excluded.value""",
            (json.dumps(duplicate_invoice_ids, ensure_ascii=False),),
        )
    else:
        conn.execute(
            """UPDATE outgoing_invoice_drafts
               SET issued_invoice_series=UPPER(TRIM(COALESCE(issued_invoice_series,''))),
                   issued_invoice_number=TRIM(COALESCE(issued_invoice_number,''))
               WHERE TRIM(COALESCE(issued_invoice_number,''))!=''"""
        )
        conn.execute("DELETE FROM settings WHERE key='schema_warning_duplicate_invoice_ids'")
        conn.execute(
            """CREATE UNIQUE INDEX idx_outgoing_issued_number
               ON outgoing_invoice_drafts(issued_invoice_series COLLATE NOCASE,issued_invoice_number)
               WHERE issued_invoice_number IS NOT NULL AND issued_invoice_number!=''"""
        )
    conn.executescript(CONSOLIDATION_SCHEMA)
    defaults = {
        "tenant_code": "TDP",
        "printer_name": "",
        "print_copies": "1",
        "print_paper": "A4",
        "print_other_paper": "A5",
        # Từ mẫu "Đề nghị Thanh toán TĐP (T04.26).xlsx" khách đã cung cấp.
        "payment_requester": "VŨ THỊ THỤY",
        "payment_bank_name": "Ngân hàng TMCP Ngoại Thương Việt Nam",
        "payment_bank_account": "1052787580",
        "company_tax_code": "0202265016",
        "company_address": "Số nhà 112 ngõ 366, Đường Hùng Vương, Phường Hồng Bàng, Thành phố Hải Phòng, Việt Nam",
    }
    for key, value in defaults.items():
        conn.execute(
            "INSERT INTO settings(key,value) VALUES(?,?) ON CONFLICT(key) DO NOTHING",
            (key, value),
        )
    # Bản demo cũ đã tạo ba khóa này với giá trị rỗng; chỉ điền phần còn trống để
    # không ghi đè cấu hình mà người dùng đã chủ động sửa.
    for key in ("payment_requester", "payment_bank_name", "payment_bank_account"):
        conn.execute(
            "UPDATE settings SET value=? WHERE key=? AND TRIM(COALESCE(value,''))=''",
            (defaults[key], key),
        )


def first_value(data: dict, *keys, default=""):
    for key in keys:
        value = data.get(key)
        if value not in (None, ""):
            return value
    return default


def as_number(value, default=0.0) -> float:
    if value in (None, ""):
        return float(default)
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(" ", "")
    if "," in text:
        text = text.replace(".", "").replace(",", ".")
    try:
        return float(text)
    except (TypeError, ValueError):
        return float(default)


def msmi_number(value, label: str, default=0.0) -> float:
    """Parse an API number without ever converting malformed data to zero."""
    if value in (None, ""):
        value = default
    if isinstance(value, bool):
        raise MsmiError(f"{label} trên hóa đơn mSMI không hợp lệ")
    text = str(value).strip().replace(" ", "")
    if "," in text:
        text = text.replace(".", "").replace(",", ".")
    try:
        number = Decimal(text)
    except (InvalidOperation, TypeError, ValueError):
        raise MsmiError(f"{label} trên hóa đơn mSMI không phải là số") from None
    if not number.is_finite():
        raise MsmiError(f"{label} trên hóa đơn mSMI phải là số hữu hạn")
    return float(number)


def import_cell_number(value, label: str, default=0.0) -> float:
    """Strict financial-cell parser used by authoritative snapshot imports."""
    if value in (None, ""):
        return float(default)
    if isinstance(value, bool):
        raise ValueError(f"{label} không phải là số")
    text = str(value).strip().replace(" ", "")
    if "," in text:
        text = text.replace(".", "").replace(",", ".")
    try:
        number = Decimal(text)
    except (InvalidOperation, TypeError, ValueError):
        raise ValueError(f"{label} không phải là số") from None
    if not number.is_finite():
        raise ValueError(f"{label} phải là số hữu hạn")
    return float(number)


def vnd_round(value):
    """Round VND with the same HALF_UP rule as the M-Invoice client."""
    try:
        return int(Decimal(str(value)).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
    except (InvalidOperation, TypeError, ValueError):
        return 0


def vnd_product(*values):
    try:
        total = Decimal("1")
        for value in values:
            total *= Decimal(str(value))
        return int(total.quantize(Decimal("1"), rounding=ROUND_HALF_UP))
    except (InvalidOperation, TypeError, ValueError):
        return 0


MINVOICE_SAVING_STALE_SECONDS = 120
VIETNAM_TIMEZONE = timezone(timedelta(hours=7))


def minvoice_remote_id(data) -> str:
    """Extract only a remote identifier; never persist the raw invoice body."""
    candidates = [data] if isinstance(data, dict) else data if isinstance(data, list) else []
    for candidate in candidates:
        if not isinstance(candidate, dict):
            continue
        for key in ("inv_invoiceAuth_Id", "invoiceAuthId", "id", "_id"):
            value = candidate.get(key)
            if value not in (None, ""):
                return str(value).strip()
    return ""


def minvoice_saving_is_fresh(started_at, current_time) -> bool:
    if not started_at:
        return False
    try:
        age = (
            datetime.fromisoformat(str(current_time))
            - datetime.fromisoformat(str(started_at))
        ).total_seconds()
    except (TypeError, ValueError):
        return False
    # A future timestamp is treated as active: clock skew must never allow a
    # second POST while the first request may still be running.
    return age < MINVOICE_SAVING_STALE_SECONDS


def as_date(value) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value or "").strip()
    if not text:
        return ""
    # mSMI represents local midnight in Vietnam as 17:00:00Z on the
    # preceding UTC day.  Taking the first ten characters would therefore
    # move every invoice back one business day and corrupt month boundaries.
    # Convert timezone-aware ISO values to the application's Vietnam business
    # timezone before taking the calendar date.  Plain dates and naive local
    # timestamps intentionally keep their existing calendar date.
    iso_text = text[:-1] + "+00:00" if text.upper().endswith("Z") else text
    try:
        parsed_iso = datetime.fromisoformat(iso_text)
    except ValueError:
        parsed_iso = None
    if parsed_iso is not None:
        if parsed_iso.tzinfo is not None:
            parsed_iso = parsed_iso.astimezone(VIETNAM_TIMEZONE)
        return parsed_iso.date().isoformat()
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue
    return ""


def audit(conn, now_iso, event_type: str, status: str, message: str = "", entity_type="", entity_id="", metadata=None):
    conn.execute(
        "INSERT INTO audit_log(event_type,entity_type,entity_id,status,message,metadata_json,created_at) "
        "VALUES(?,?,?,?,?,?,?)",
        (event_type, entity_type, str(entity_id or ""), status, message,
         json.dumps(metadata or {}, ensure_ascii=False), now_iso()),
    )


def mapping_key(value) -> str:
    text = str(value or "").strip().lower().replace("đ", "d")
    text = "".join(
        char for char in unicodedata.normalize("NFKD", text)
        if not unicodedata.combining(char)
    )
    return re.sub(r"[^a-z0-9]+", "", text)


def mapping_cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return re.sub(r"\s+", " ", str(value).strip())


def catalog_unit(value) -> str:
    text = mapping_cell_text(value)
    return CATALOG_CANONICAL_UNITS.get(mapping_key(text), text)


def catalog_tax(value) -> str:
    if value in (None, "") or isinstance(value, bool):
        return ""
    if isinstance(value, (int, float)):
        number = float(value)
    else:
        text = mapping_cell_text(value).upper()
        if mapping_key(text) in {"kkknt", "khongkekhai"}:
            return "KKKNT"
        if mapping_key(text) in {"kct", "khongchiuthue"}:
            return "KCT"
        text = text.replace("%", "").replace(",", ".").strip()
        try:
            number = float(text)
        except ValueError:
            return mapping_cell_text(value)
    if number > 1:
        number /= 100
    return f"{number:.6f}".rstrip("0").rstrip(".")


def invoice_tax_percent(value) -> float:
    raw = mapping_cell_text(value).upper().replace(" ", "")
    if raw in {"KKKNT", "KHÔNGKÊKHAI", "KHONGKEKHAI"}:
        return -2
    if raw in {"KCT", "KHÔNGCHỊUTHUẾ", "KHONGCHIUTHUE"}:
        return -1
    has_percent = raw.endswith("%")
    number = as_number(raw.rstrip("%"), 0)
    if not has_percent and 0 < abs(number) < 1:
        number *= 100
    return round(number, 4)


def catalog_header_fields(row) -> dict:
    found = {}
    for column_index, value in enumerate(row, start=1):
        key = mapping_key(value)
        if not key:
            continue
        for field_name, accepted in CATALOG_ALIASES.items():
            if field_name not in found and key in accepted:
                found[field_name] = column_index
                break
    return found


def find_catalog_sheet(workbook):
    candidates = []
    required = {"product_code", "product_name", "unit", "tax"}
    for sheet_index, worksheet in enumerate(workbook.worksheets):
        for row_index, row in enumerate(
            worksheet.iter_rows(min_row=1, max_row=25, max_col=30, values_only=True),
            start=1,
        ):
            fields = catalog_header_fields(row)
            if required.issubset(fields):
                candidates.append((len(fields), -sheet_index, -row_index, worksheet, row_index, fields))
    if not candidates:
        return None
    _, _, _, worksheet, row_index, fields = max(candidates, key=lambda item: item[:3])
    return worksheet, row_index, fields


def catalog_database_state_hash(conn) -> str:
    digest = hashlib.sha256()
    for row in conn.execute(
        "SELECT code,name,unit,tax,COALESCE(product_group,'') product_group FROM products ORDER BY code"
    ):
        digest.update(json.dumps(list(row), ensure_ascii=False, separators=(",", ":")).encode("utf-8"))
        digest.update(b"\n")
    for row in conn.execute(
        "SELECT product_code,invoice_name FROM outgoing_product_names ORDER BY product_code"
    ):
        digest.update(json.dumps(list(row), ensure_ascii=False, separators=(",", ":")).encode("utf-8"))
        digest.update(b"\n")
    return digest.hexdigest()


def query_database_state_hash(conn, query_specs) -> str:
    """Fingerprint the exact database rows a preview was calculated from.

    Preview/confirm imports can be open in two browser tabs at once.  A token
    therefore protects the uploaded bytes but not the database snapshot.  The
    caller supplies deterministic ``(label, sql, params)`` queries; confirmation
    must compare the fingerprint again inside the same write transaction before
    applying an authoritative snapshot.
    """

    digest = hashlib.sha256()
    for label, sql, params in query_specs:
        digest.update(str(label).encode("utf-8"))
        digest.update(b"\0")
        for row in conn.execute(sql, tuple(params or ())):
            digest.update(
                json.dumps(list(row), ensure_ascii=False, separators=(",", ":"), default=str)
                .encode("utf-8")
            )
            digest.update(b"\n")
    return digest.hexdigest()


def mapping_database_state_hash(conn, mapping_type: str) -> str:
    if mapping_type == "invoice_names":
        return query_database_state_hash(conn, (
            ("products", "SELECT code,name FROM products ORDER BY code", ()),
            (
                "outgoing_product_names",
                "SELECT product_code,invoice_name FROM outgoing_product_names ORDER BY product_code",
                (),
            ),
        ))
    return query_database_state_hash(conn, (
        ("kitchens", "SELECT code,name FROM kitchens ORDER BY code", ()),
        (
            "kitchen_units",
            "SELECT kitchen_code,unit_code FROM kitchen_units ORDER BY kitchen_code",
            (),
        ),
    ))


def kitchen_import_database_state_hash(conn) -> str:
    return query_database_state_hash(conn, (
        (
            "products",
            "SELECT code,name,unit,supplier,buy_price FROM products ORDER BY code",
            (),
        ),
        (
            "dated_prices",
            "SELECT product_code,price_group,period,price_value FROM dated_prices "
            "ORDER BY product_code,price_group,period",
            (),
        ),
        ("kitchens", "SELECT code,name FROM kitchens ORDER BY code", ()),
        (
            "kitchen_units",
            "SELECT kitchen_code,unit_code FROM kitchen_units ORDER BY kitchen_code",
            (),
        ),
        (
            "meal_plans",
            "SELECT id,work_date,kitchen,shift,meal_count,unit_code,status,import_key,"
            "meal_price,other_cost,updated_at FROM meal_plans ORDER BY id",
            (),
        ),
        (
            "meal_plan_items",
            "SELECT id,plan_id,product_code,norm_qty,buy_price,price_source,"
            "applicable_meal_count,source_amount "
            "FROM meal_plan_items ORDER BY id",
            (),
        ),
    ))


def opening_import_database_state_hash(conn, period: str) -> str:
    return query_database_state_hash(conn, (
        (
            "products",
            "SELECT code,name,unit,tax,buy_price FROM products ORDER BY code",
            (),
        ),
        (
            "opening",
            "SELECT product_code,warehouse_codes_json,qty_in,qty_out,unit_cost,source_line,status,note "
            "FROM inventory_transactions WHERE source_type='OPENING' AND source_id=? "
            "ORDER BY source_line",
            (period,),
        ),
    ))


def meal_attendance_database_state_hash(conn, periods) -> str:
    normalized = sorted({str(period) for period in periods})
    if not normalized:
        return query_database_state_hash(conn, ())
    placeholders = ",".join("?" for _ in normalized)
    return query_database_state_hash(conn, ((
        "meal_attendance",
        f"SELECT work_date,kitchen,shift,actual_count,ordered_count,source_type,"
        f"source_file,source_sheet,source_column,updated_at FROM meal_attendance "
        f"WHERE substr(work_date,1,7) IN ({placeholders}) "
        f"ORDER BY work_date,kitchen,shift",
        tuple(normalized),
    ),))


def payables_database_state_hash(conn) -> str:
    return query_database_state_hash(conn, (
        (
            "supplier_master",
            "SELECT code,name FROM suppliers ORDER BY code",
            (),
        ),
        (
            "product_suppliers",
            "SELECT code,supplier FROM products ORDER BY code",
            (),
        ),
        (
            "historical_payable_lines",
            "SELECT id,purchase_date,kitchen,item_name,qty,unit,supplier,buy_price,"
            "damaged_qty,added_qty,reduced_qty,missing_qty,actual_qty,source_amount,"
            "calculated_amount,amount,note,source_file,source_sheet,source_row,source_hash "
            ",updated_at FROM historical_payable_lines ORDER BY id",
            (),
        ),
    ))


def parse_catalog_workbook(conn, workbook) -> dict:
    found = find_catalog_sheet(workbook)
    if not found:
        raise ValueError("Không tìm thấy bảng có Mã hàng, Tên Thành Đạt Phát, ĐVT và Thuế")
    worksheet, header_row, fields = found
    existing_products = {
        row["code"]: dict(row)
        for row in conn.execute(
            "SELECT code,name,unit,tax,COALESCE(product_group,'') product_group FROM products"
        )
    }
    existing_names = {
        row["product_code"]: row["invoice_name"]
        for row in conn.execute("SELECT product_code,invoice_name FROM outgoing_product_names")
    }
    rows = []
    unique_items = {}
    scanned = 0
    max_column = max(fields.values())
    for row_index, row in enumerate(
        worksheet.iter_rows(
            min_row=header_row + 1,
            max_row=min(worksheet.max_row, header_row + MAPPING_IMPORT_MAX_ROWS + 1),
            max_col=max_column,
            values_only=True,
        ),
        start=header_row + 1,
    ):
        code = mapping_cell_text(row[fields["product_code"] - 1]).upper()
        name = mapping_cell_text(row[fields["product_name"] - 1])
        unit = catalog_unit(row[fields["unit"] - 1])
        tax = catalog_tax(row[fields["tax"] - 1])
        group = mapping_cell_text(row[fields["product_group"] - 1]).upper() if "product_group" in fields else ""
        invoice_name = mapping_cell_text(row[fields["invoice_name"] - 1]) if "invoice_name" in fields else ""
        if not any((code, name, unit, tax, group, invoice_name)):
            continue
        scanned += 1
        item = {
            "source_row": row_index,
            "product_code": code,
            "product_group": group,
            "product_name": name,
            "invoice_name": invoice_name,
            "unit": unit,
            "tax": tax,
            "product_status": "error",
            "invoice_status": "none",
            "warnings": [],
            "errors": [],
            "apply": False,
        }
        if not code:
            item["errors"].append("Thiếu mã hàng")
        if not name:
            item["errors"].append("Thiếu tên Thành Đạt Phát")
        if not unit:
            item["errors"].append("Thiếu đơn vị tính")
        if not tax:
            item["errors"].append("Thiếu thuế")
        if "invoice_name" in fields and not invoice_name:
            item["errors"].append("Thiếu tên xuất hóa đơn")

        previous = unique_items.get(code) if code else None
        if previous:
            signature = (mapping_key(name), mapping_key(invoice_name), mapping_key(unit), tax, group)
            previous_signature = (
                mapping_key(previous["product_name"]), mapping_key(previous["invoice_name"]),
                mapping_key(previous["unit"]), previous["tax"], previous["product_group"],
            )
            if signature == previous_signature:
                item["product_status"] = "duplicate"
                item["invoice_status"] = "duplicate"
                item["warnings"].append(
                    f"Trùng hoàn toàn với dòng {previous['source_row']}; chỉ nhập một lần"
                )
            else:
                item["errors"].append(
                    f"Cùng mã nhưng khác dữ liệu với dòng {previous['source_row']}"
                )
                previous["errors"].append(f"Cùng mã nhưng khác dữ liệu với dòng {row_index}")
                previous["product_status"] = "error"
                previous["invoice_status"] = "error"
                previous["apply"] = False
            rows.append(item)
            continue

        if code:
            unique_items[code] = item
        if item["errors"]:
            rows.append(item)
            continue

        current = existing_products.get(code)
        if not current:
            item["product_status"] = "new"
        else:
            base_changed = any((
                mapping_key(current.get("name")) != mapping_key(name),
                mapping_key(current.get("unit")) != mapping_key(unit),
                catalog_tax(current.get("tax")) != tax,
                mapping_key(current.get("product_group")) != mapping_key(group),
            ))
            item["product_status"] = "update" if base_changed else "unchanged"

        if invoice_name:
            current_invoice_name = existing_names.get(code, "")
            if not current_invoice_name:
                item["invoice_status"] = "new"
            elif mapping_key(current_invoice_name) == mapping_key(invoice_name):
                item["invoice_status"] = "unchanged"
            else:
                item["invoice_status"] = "update"
                item["warnings"].append("Tên xuất hóa đơn sẽ được cập nhật sau khi xác nhận")
        item["apply"] = True
        rows.append(item)

    if scanned > MAPPING_IMPORT_MAX_ROWS:
        raise ValueError(f"File vượt quá giới hạn {MAPPING_IMPORT_MAX_ROWS:,} dòng dữ liệu")
    if not rows:
        raise ValueError("Sheet được nhận diện nhưng không có dòng dữ liệu")

    unique_rows = [item for item in rows if item["apply"] and item["product_status"] != "duplicate"]
    incoming_codes = {item["product_code"] for item in unique_rows}
    retained_codes = sorted(set(existing_products) - incoming_codes)
    counts = {
        "total": len(rows),
        "unique_products": len(unique_rows),
        "new_products": sum(item["product_status"] == "new" for item in unique_rows),
        "update_products": sum(item["product_status"] == "update" for item in unique_rows),
        "unchanged_products": sum(item["product_status"] == "unchanged" for item in unique_rows),
        "retained_products": len(retained_codes),
        "new_names": sum(item["invoice_status"] == "new" for item in unique_rows),
        "update_names": sum(item["invoice_status"] == "update" for item in unique_rows),
        "unchanged_names": sum(item["invoice_status"] == "unchanged" for item in unique_rows),
        "duplicate": sum(item["product_status"] == "duplicate" for item in rows),
        "error": sum(bool(item["errors"]) for item in rows),
    }
    return {
        "sheet": worksheet.title,
        "header_row": header_row,
        "rows": rows,
        "items": unique_rows,
        "counts": counts,
        "retained_codes": retained_codes,
        "database_state_hash": catalog_database_state_hash(conn),
        "can_confirm": counts["error"] == 0,
    }


def mapping_header_fields(row, mapping_type: str) -> dict:
    aliases = MAPPING_ALIASES[mapping_type]
    found = {}
    for column_index, value in enumerate(row, start=1):
        key = mapping_key(value)
        if not key:
            continue
        for field_name, accepted in aliases.items():
            if field_name not in found and key in accepted:
                found[field_name] = column_index
                break
    return found


def mapping_header_is_valid(fields: dict, mapping_type: str) -> bool:
    if "target_value" not in fields:
        return False
    if mapping_type == "invoice_names":
        return "product_code" in fields or "product_name" in fields
    return "source_code" in fields or "source_name" in fields


def find_mapping_sheet(workbook, mapping_type: str):
    candidates = []
    for sheet_index, worksheet in enumerate(workbook.worksheets):
        for row_index, row in enumerate(
            worksheet.iter_rows(min_row=1, max_row=25, max_col=30, values_only=True),
            start=1,
        ):
            fields = mapping_header_fields(row, mapping_type)
            if mapping_header_is_valid(fields, mapping_type):
                candidates.append((len(fields), -sheet_index, -row_index, worksheet, row_index, fields))
    if not candidates:
        return None
    _, _, _, worksheet, row_index, fields = max(candidates, key=lambda item: item[:3])
    return worksheet, row_index, fields


def mapping_catalog(conn, mapping_type: str) -> dict:
    if mapping_type == "invoice_names":
        rows = [dict(row) for row in conn.execute("SELECT code,name FROM products ORDER BY code")]
    else:
        rows = [dict(row) for row in conn.execute("SELECT code,name FROM kitchens ORDER BY code")]
    by_code = {mapping_cell_text(row["code"]).upper(): row for row in rows}
    by_name = defaultdict(list)
    for row in rows:
        name_key = mapping_key(row.get("name"))
        if name_key:
            by_name[name_key].append(row)
    return {"by_code": by_code, "by_name": by_name}


def resolve_mapping_source(raw_code: str, raw_name: str, catalog: dict) -> tuple[dict | None, list, list]:
    warnings = []
    errors = []
    code = mapping_cell_text(raw_code).upper()
    name = mapping_cell_text(raw_name)
    resolved = catalog["by_code"].get(code) if code else None
    if resolved:
        if name and mapping_key(name) != mapping_key(resolved.get("name")):
            warnings.append("Tên trong file khác danh mục; hệ thống dùng mã để ghép")
        return resolved, warnings, errors

    lookup_name = name or (raw_code if code else "")
    matches = catalog["by_name"].get(mapping_key(lookup_name), []) if lookup_name else []
    if len(matches) == 1:
        if code:
            warnings.append(f"Mã/giá trị {code} chưa có trong danh mục; hệ thống ghép theo tên")
        return matches[0], warnings, errors
    if len(matches) > 1:
        errors.append("Tên nguồn trùng nhiều mã trong danh mục; cần bổ sung mã")
    elif code:
        errors.append(f"Mã {code} chưa có trong danh mục")
    elif name:
        errors.append("Tên nguồn chưa khớp danh mục")
    else:
        errors.append("Thiếu mã hoặc tên nguồn")
    return None, warnings, errors


def parse_mapping_workbook(conn, workbook, mapping_type: str) -> dict:
    found = find_mapping_sheet(workbook, mapping_type)
    if not found:
        expected = "Mã/Tên hàng và Tên xuất hóa đơn" if mapping_type == "invoice_names" else "Mã/Tên bếp và XCOM (xưởng cơm)"
        raise ValueError(f"Không tìm thấy dòng tiêu đề có {expected}")
    worksheet, header_row, fields = found
    catalog = mapping_catalog(conn, mapping_type)
    existing_table = "outgoing_product_names" if mapping_type == "invoice_names" else "kitchen_units"
    existing_key = "product_code" if mapping_type == "invoice_names" else "kitchen_code"
    existing_value = "invoice_name" if mapping_type == "invoice_names" else "unit_code"
    existing = {
        row[existing_key]: row[existing_value]
        for row in conn.execute(f"SELECT {existing_key},{existing_value} FROM {existing_table}")
    }
    source_code_field = "product_code" if mapping_type == "invoice_names" else "source_code"
    source_name_field = "product_name" if mapping_type == "invoice_names" else "source_name"
    rows = []
    seen = {}
    scanned = 0
    for row_index, row in enumerate(
        worksheet.iter_rows(
            min_row=header_row + 1,
            max_row=header_row + MAPPING_IMPORT_MAX_ROWS + 1,
            max_col=max(fields.values()),
            values_only=True,
        ),
        start=header_row + 1,
    ):
        raw_code = mapping_cell_text(row[fields[source_code_field] - 1]) if source_code_field in fields else ""
        raw_name = mapping_cell_text(row[fields[source_name_field] - 1]) if source_name_field in fields else ""
        target = mapping_cell_text(row[fields["target_value"] - 1])
        if not raw_code and not raw_name and not target:
            continue
        if mapping_key(target) in MAPPING_ALIASES[mapping_type]["target_value"]:
            continue
        scanned += 1
        item = {
            "source_row": row_index,
            "source_code": raw_code,
            "source_name": raw_name,
            "resolved_code": "",
            "resolved_name": "",
            "target_value": target.upper() if mapping_type == "kitchen_units" else target,
            "current_value": "",
            "status": "error",
            "warnings": [],
            "errors": [],
            "apply": False,
        }
        if not target:
            item["errors"].append("Thiếu giá trị cần nhập")
        resolved, warnings, errors = resolve_mapping_source(raw_code, raw_name, catalog)
        item["warnings"].extend(warnings)
        item["errors"].extend(errors)
        if resolved:
            code = mapping_cell_text(resolved["code"]).upper()
            item["resolved_code"] = code
            item["resolved_name"] = mapping_cell_text(resolved.get("name"))
            item["current_value"] = mapping_cell_text(existing.get(code))
        if item["errors"]:
            rows.append(item)
            continue

        code = item["resolved_code"]
        target_compare = mapping_key(item["target_value"])
        previous = seen.get(code)
        if previous:
            if mapping_key(previous["target_value"]) == target_compare:
                item["status"] = "duplicate"
                item["warnings"].append(f"Trùng nội dung với dòng {previous['source_row']}; chỉ nhập một lần")
                previous["warnings"].append(f"Có dòng trùng {row_index}; chỉ nhập một lần")
            else:
                message = f"Cùng mã nhưng khác giá trị với dòng {previous['source_row']}"
                item["errors"].append(message)
                item["status"] = "error"
                previous["errors"].append(f"Cùng mã nhưng khác giá trị với dòng {row_index}")
                previous["status"] = "error"
                previous["apply"] = False
            rows.append(item)
            continue

        seen[code] = item
        if not item["current_value"]:
            item["status"] = "new"
        elif mapping_key(item["current_value"]) == target_compare:
            item["status"] = "unchanged"
        else:
            item["status"] = "update"
            item["warnings"].append("Sẽ thay giá trị đang lưu sau khi xác nhận")
        item["apply"] = True
        rows.append(item)

    if scanned > MAPPING_IMPORT_MAX_ROWS:
        raise ValueError(f"File vượt quá giới hạn {MAPPING_IMPORT_MAX_ROWS:,} dòng dữ liệu")
    if not rows:
        raise ValueError("Sheet được nhận diện nhưng không có dòng dữ liệu")

    counts = {
        "total": len(rows),
        "new": sum(item["status"] == "new" for item in rows),
        "update": sum(item["status"] == "update" for item in rows),
        "unchanged": sum(item["status"] == "unchanged" for item in rows),
        "duplicate": sum(item["status"] == "duplicate" for item in rows),
        "error": sum(bool(item["errors"]) for item in rows),
    }
    return {
        "sheet": worksheet.title,
        "header_row": header_row,
        "rows": rows,
        "items": [
            {"code": item["resolved_code"], "target_value": item["target_value"]}
            for item in rows if item["apply"] and not item["errors"]
        ],
        "counts": counts,
        "can_confirm": counts["error"] == 0,
    }


OPENING_ALIASES = {
    "product_code": {"matdp", "mahangtdp", "mavattutdp", "masptdp", "mahhtdp"},
    "product_name": {"tentdp", "tenhangtdp", "tenhanghoatdp", "tenvattutdp"},
    "invoice_name": {"tentrenhd", "tentrenhoadon", "tenhoadon", "tenxuathoadon"},
    "warehouse_code": {"makho", "mahangkho", "mavattukho"},
    "tax": {"tsuat", "thuesuat", "thuegtgt", "vat"},
    "unit": {"dvt", "donvitinh", "unit"},
    "qty": {"soluong", "toncuoikysoluong", "tondaukysoluong", "toncuoiky", "tondauky"},
    "unit_cost": {"dongia", "giavon", "dongiavon"},
    "amount": {"thanhtien", "giatriton", "giatritonkho"},
}

# The only legacy snapshot eligible for automatic MÃ KHO recovery is the
# customer-approved August 2026 golden. The file is bundled into the portable
# executable and its digest is also used by the official TĐK–NXT exporter.
OPENING_WAREHOUSE_BACKFILL_PERIOD = "2026-08"
OPENING_WAREHOUSE_GOLDEN_SHA256 = (
    "36DF2BA86D13307F96BB5944FCECB19A4A81C093B4AC6A98EA71330D68204DA6"
)
OPENING_WAREHOUSE_BACKFILL_SETTING = "opening_warehouse_codes_backfill_2026_08"


def opening_header_fields(*header_rows) -> dict:
    width = max((len(row) for row in header_rows), default=0)
    fields = {}
    for column in range(width):
        parts = [mapping_key(row[column]) for row in header_rows if column < len(row)]
        keys = {part for part in parts if part}
        combined = "".join(part for part in parts if part)
        if combined:
            keys.add(combined)
        for field, aliases in OPENING_ALIASES.items():
            if field not in fields and keys.intersection(aliases):
                fields[field] = column + 1
    return fields


def find_opening_sheet(workbook):
    candidates = []
    for sheet_index, worksheet in enumerate(workbook.worksheets):
        rows = list(worksheet.iter_rows(
            min_row=1, max_row=min(worksheet.max_row or 0, 26), max_col=30, values_only=True,
        ))
        for index, row in enumerate(rows):
            single = opening_header_fields(row)
            if {"product_code", "qty"}.issubset(single):
                candidates.append((len(single), -sheet_index, -(index + 1), worksheet, index + 1, single))
            if index + 1 < len(rows):
                double = opening_header_fields(row, rows[index + 1])
                if {"product_code", "qty"}.issubset(double):
                    candidates.append((len(double), -sheet_index, -(index + 2), worksheet, index + 2, double))
    if not candidates:
        return None
    _, _, _, worksheet, header_end, fields = max(candidates, key=lambda item: item[:3])
    return worksheet, header_end, fields


def _warehouse_codes_from_golden(workbook) -> dict[str, list[str]]:
    """Read only the product/MÃ KHO identity columns from a trusted workbook."""

    found = find_opening_sheet(workbook)
    if not found:
        raise ValueError("Golden TĐK không có dòng tiêu đề Mã TĐP")
    worksheet, header_end, fields = found
    if "warehouse_code" not in fields:
        raise ValueError("Golden TĐK thiếu cột MÃ KHO")
    code_column = fields["product_code"]
    warehouse_column = fields["warehouse_code"]
    mapping: dict[str, list[str]] = {}
    for row in worksheet.iter_rows(
        min_row=header_end + 1,
        max_row=header_end + MAPPING_IMPORT_MAX_ROWS + 1,
        max_col=max(code_column, warehouse_column),
        values_only=True,
    ):
        code = mapping_cell_text(row[code_column - 1]).upper()
        warehouse_code = mapping_cell_text(row[warehouse_column - 1]).upper()
        if not code and not warehouse_code:
            continue
        if mapping_key(code) in OPENING_ALIASES["product_code"]:
            continue
        if not code or not warehouse_code:
            raise ValueError("Golden TĐK có dòng thiếu Mã TĐP hoặc MÃ KHO")
        codes = mapping.setdefault(code, [])
        if warehouse_code not in codes:
            codes.append(warehouse_code)
    if not mapping:
        raise ValueError("Golden TĐK không có dữ liệu MÃ KHO")
    return mapping


def backfill_opening_warehouse_codes(
    conn,
    template_path=None,
    *,
    expected_sha256: str = OPENING_WAREHOUSE_GOLDEN_SHA256,
    period: str = OPENING_WAREHOUSE_BACKFILL_PERIOD,
) -> dict:
    """Recover MÃ KHO lost by the pre-column August 2026 importer.

    This migration is intentionally narrow and idempotent. It never guesses
    from Mã TĐP: a legacy snapshot is changed only when its complete product
    set matches the hash-pinned customer golden. Source/dev installations that
    do not carry the optional golden keep starting normally; a present but
    altered golden fails closed before any row is updated.
    """

    target_rows = conn.execute(
        """SELECT id,product_code,warehouse_codes_json
             FROM inventory_transactions
            WHERE source_type='OPENING' AND source_id=?
            ORDER BY product_code,id""",
        (period,),
    ).fetchall()
    if not target_rows:
        return {"status": "snapshot_absent", "updated": 0, "period": period}

    parsed_existing: dict[int, list[str]] = {}
    missing = 0
    for row in target_rows:
        try:
            raw_codes = json.loads(str(row["warehouse_codes_json"] or "[]"))
        except (TypeError, ValueError, json.JSONDecodeError):
            raw_codes = []
        if not isinstance(raw_codes, list):
            raw_codes = []
        codes = list(dict.fromkeys(
            mapping_cell_text(value).upper() for value in raw_codes
            if mapping_cell_text(value)
        ))
        parsed_existing[int(row["id"])] = codes
        missing += int(not codes)
    if not missing:
        return {"status": "already_complete", "updated": 0, "period": period}

    if template_path in (None, ""):
        return {"status": "asset_absent", "updated": 0, "period": period}
    path = Path(template_path)
    if not path.is_file():
        return {"status": "asset_absent", "updated": 0, "period": period}

    payload = path.read_bytes()
    actual_sha256 = hashlib.sha256(payload).hexdigest().upper()
    trusted_sha256 = str(expected_sha256 or "").strip().upper()
    if actual_sha256 != trusted_sha256:
        raise ValueError("Golden TĐK không đúng SHA-256 đã khóa; không backfill MÃ KHO")

    # The approved workbook carries a harmless x14 conditional-formatting
    # extension that openpyxl cannot render. We only read two value columns;
    # suppress that known warning so first startup does not look like a failed
    # migration to the operator.
    with warnings.catch_warnings():
        warnings.filterwarnings(
            "ignore",
            message="Conditional Formatting extension is not supported.*",
            category=UserWarning,
            module="openpyxl",
        )
        workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
        try:
            golden = _warehouse_codes_from_golden(workbook)
        finally:
            workbook.close()

    target_codes = {mapping_cell_text(row["product_code"]).upper() for row in target_rows}
    if target_codes != set(golden):
        raise ValueError(
            "Snapshot tồn 08/2026 không khớp trọn bộ Mã TĐP trong golden; "
            "không backfill MÃ KHO"
        )

    updates: list[tuple[str, int]] = []
    for row in target_rows:
        row_id = int(row["id"])
        code = mapping_cell_text(row["product_code"]).upper()
        expected_codes = golden[code]
        existing_codes = parsed_existing[row_id]
        if existing_codes and existing_codes != expected_codes:
            raise ValueError(
                "Snapshot tồn 08/2026 đã có MÃ KHO khác golden; không ghi đè"
            )
        if not existing_codes:
            updates.append((
                json.dumps(expected_codes, ensure_ascii=False, separators=(",", ":")),
                row_id,
            ))

    if updates:
        evidence = {
            "sha256": actual_sha256,
            "period": period,
            "updated": len(updates),
            "products": len(golden),
            "different_from_tdp": sum(
                any(warehouse != code for warehouse in codes)
                for code, codes in golden.items()
            ),
            "multiple_warehouse_codes": sum(len(codes) > 1 for codes in golden.values()),
        }
        savepoint = "opening_warehouse_codes_backfill"
        conn.execute(f"SAVEPOINT {savepoint}")
        try:
            conn.executemany(
                "UPDATE inventory_transactions SET warehouse_codes_json=? WHERE id=?",
                updates,
            )
            conn.execute(
                """INSERT INTO settings(key,value) VALUES(?,?)
                   ON CONFLICT(key) DO UPDATE SET value=excluded.value""",
                (
                    OPENING_WAREHOUSE_BACKFILL_SETTING,
                    json.dumps(evidence, ensure_ascii=False, sort_keys=True, separators=(",", ":")),
                ),
            )
            conn.execute(f"RELEASE SAVEPOINT {savepoint}")
        except Exception:
            conn.execute(f"ROLLBACK TO SAVEPOINT {savepoint}")
            conn.execute(f"RELEASE SAVEPOINT {savepoint}")
            raise
    return {
        "status": "backfilled" if updates else "already_complete",
        "updated": len(updates),
        "period": period,
        "products": len(golden),
    }


def opening_number(value):
    if value in (None, "") or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        number = float(value)
    else:
        text = str(value).strip().replace(" ", "").replace("₫", "").replace("đ", "")
        if not text:
            return None
        if "," in text:
            text = text.replace(".", "").replace(",", ".")
        elif text.count(".") > 1:
            text = text.replace(".", "")
        try:
            number = float(text)
        except (TypeError, ValueError):
            return None
    return number if math.isfinite(number) else None


def first_most_common(values) -> str:
    values = [mapping_cell_text(value) for value in values if mapping_cell_text(value)]
    return Counter(values).most_common(1)[0][0] if values else ""


def parse_opening_workbook(conn, workbook, period: str) -> dict:
    found = find_opening_sheet(workbook)
    if not found:
        raise ValueError("Không tìm thấy dòng tiêu đề có Mã TĐP và Số lượng tồn")
    worksheet, header_end, fields = found
    missing_columns = [
        label for field, label in (
            ("product_name", "Tên TĐP"), ("unit", "ĐVT"),
            ("unit_cost", "Đơn giá"), ("amount", "Thành tiền"),
        ) if field not in fields
    ]
    if missing_columns:
        raise ValueError("File thiếu cột: " + ", ".join(missing_columns))

    source_rows = []
    scanned = 0
    max_column = max(fields.values())
    for row_index, row in enumerate(
        worksheet.iter_rows(
            min_row=header_end + 1,
            max_row=header_end + MAPPING_IMPORT_MAX_ROWS + 1,
            max_col=max_column,
            values_only=True,
        ),
        start=header_end + 1,
    ):
        raw = {field: row[column - 1] for field, column in fields.items()}
        code = mapping_cell_text(raw.get("product_code")).upper()
        if not code and not any(value not in (None, "") for value in raw.values()):
            continue
        if mapping_key(code) in OPENING_ALIASES["product_code"]:
            continue
        scanned += 1
        item = {
            "source_row": row_index,
            "product_code": code,
            "product_name": mapping_cell_text(raw.get("product_name")),
            "invoice_name": mapping_cell_text(raw.get("invoice_name")),
            "warehouse_code": mapping_cell_text(raw.get("warehouse_code")).upper(),
            "tax": mapping_cell_text(raw.get("tax")),
            "unit": mapping_cell_text(raw.get("unit")),
            "qty": opening_number(raw.get("qty")),
            "unit_cost": opening_number(raw.get("unit_cost")),
            "amount": opening_number(raw.get("amount")),
            "warnings": [],
            "errors": [],
        }
        if not code:
            item["errors"].append("Thiếu Mã TĐP")
        if item["qty"] is None:
            item["errors"].append("Số lượng tồn không hợp lệ")
        if item["unit_cost"] is None and item["amount"] is None:
            item["warnings"].append("Thiếu cả đơn giá và thành tiền; giá trị tồn được ghi là 0")
        if item["qty"] is not None:
            if item["amount"] is None and item["unit_cost"] is not None:
                item["amount"] = item["qty"] * item["unit_cost"]
                item["warnings"].append("Thiếu thành tiền; hệ thống tính Số lượng × Đơn giá")
            elif item["amount"] is not None and item["unit_cost"] is None and item["qty"]:
                item["unit_cost"] = abs(item["amount"] / item["qty"])
                item["warnings"].append("Thiếu đơn giá; hệ thống suy ra từ Thành tiền / Số lượng")
            elif item["amount"] is not None and item["unit_cost"] is not None:
                calculated = item["qty"] * item["unit_cost"]
                if abs(item["amount"] - calculated) > max(1, abs(item["amount"]) * 0.000001):
                    item["warnings"].append("Thành tiền lệch Số lượng × Đơn giá; ưu tiên Thành tiền sổ sách")
        item["amount"] = item["amount"] if item["amount"] is not None else 0.0
        item["unit_cost"] = item["unit_cost"] if item["unit_cost"] is not None else 0.0
        source_rows.append(item)

    if scanned > MAPPING_IMPORT_MAX_ROWS:
        raise ValueError(f"File vượt quá giới hạn {MAPPING_IMPORT_MAX_ROWS:,} dòng dữ liệu")
    if not source_rows:
        raise ValueError("Sheet được nhận diện nhưng không có dòng tồn đầu kỳ")

    products = {
        row["code"]: dict(row)
        for row in conn.execute("SELECT code,name,unit,tax,buy_price FROM products")
    }
    existing_openings = {
        row["source_line"]: dict(row)
        for row in conn.execute(
            """SELECT source_line,qty_in,qty_out,unit_cost FROM inventory_transactions
               WHERE source_type='OPENING' AND source_id=?""",
            (period,),
        )
    }
    grouped = defaultdict(list)
    orphan_errors = []
    for item in source_rows:
        if item["product_code"]:
            grouped[item["product_code"]].append(item)
        elif item["errors"]:
            orphan_errors.append(item)

    rows = []
    items = []
    for code, parts in grouped.items():
        names = [part["product_name"] for part in parts if part["product_name"]]
        units = [part["unit"] for part in parts if part["unit"]]
        taxes = [part["tax"] for part in parts if part["tax"]]
        invoice_names = [part["invoice_name"] for part in parts if part["invoice_name"]]
        qty = sum(part["qty"] or 0 for part in parts)
        amount = sum(part["amount"] or 0 for part in parts)
        source_row_numbers = [part["source_row"] for part in parts]
        warehouse_codes = list(dict.fromkeys(
            part["warehouse_code"] for part in parts if part["warehouse_code"]
        ))
        product = products.get(code)
        item = {
            "product_code": code,
            "product_name": first_most_common(names),
            "unit": first_most_common(units),
            "tax": first_most_common(taxes),
            "invoice_name": first_most_common(invoice_names),
            "qty": qty,
            "amount": amount,
            "unit_cost": abs(amount / qty) if abs(qty) > 1e-12 else 0,
            "source_rows": source_row_numbers,
            "warehouse_codes": warehouse_codes,
            "source_row_count": len(parts),
            "create_product": product is None,
            "status": "update" if code in existing_openings else "new",
            "current_qty": (
                existing_openings[code]["qty_in"] - existing_openings[code]["qty_out"]
                if code in existing_openings else 0
            ),
            "warnings": [],
            "errors": [],
        }
        for part in parts:
            item["warnings"].extend(part["warnings"])
            item["errors"].extend(part["errors"])
        if len(parts) > 1:
            item["warnings"].append(
                f"Đã gộp {len(parts)} dòng/mã kho theo cùng Mã TĐP"
            )
        distinct_names = list(dict.fromkeys(mapping_cell_text(value) for value in names if value))
        distinct_units = list(dict.fromkeys(mapping_cell_text(value) for value in units if value))
        if len({mapping_key(value) for value in distinct_names}) > 1:
            item["warnings"].append("Một Mã TĐP có nhiều tên nguồn; tồn được gộp theo mã")
        if len({mapping_key(value) for value in distinct_units}) > 1:
            item["warnings"].append("Một Mã TĐP có nhiều ĐVT nguồn; cần hiểu số lượng là tổng theo mã")
        if abs(qty) <= 1e-12:
            item["errors"].append("Tổng số lượng sau gộp bằng 0")
        if qty < 0:
            item["warnings"].append("Tồn đầu kỳ âm; hệ thống giữ đúng số sổ sách")
        if amount and qty and amount * qty < 0:
            item["errors"].append("Thành tiền và số lượng trái dấu")
        if product:
            if distinct_names and all(
                mapping_key(name) != mapping_key(product["name"]) for name in distinct_names
            ):
                item["warnings"].append(
                    f"Tên trong file khác danh mục ({product['name']}); hệ thống giữ danh mục hiện tại"
                )
            if distinct_units and all(
                mapping_key(unit) != mapping_key(product["unit"]) for unit in distinct_units
            ):
                item["warnings"].append(
                    f"ĐVT trong file khác danh mục ({product['unit']}); hệ thống giữ danh mục hiện tại"
                )
        else:
            if not item["product_name"]:
                item["errors"].append("Mã mới thiếu Tên TĐP")
            if not item["unit"]:
                item["errors"].append("Mã mới thiếu ĐVT")
            item["warnings"].append("Mã mới sẽ được thêm vào danh mục; chưa có NCC mặc định")
        item["warnings"] = list(dict.fromkeys(item["warnings"]))
        item["errors"] = list(dict.fromkeys(item["errors"]))
        rows.append(item)
        if not item["errors"]:
            items.append({key: item[key] for key in (
                "product_code", "product_name", "unit", "tax", "invoice_name", "qty",
                "amount", "unit_cost", "source_rows", "warehouse_codes", "create_product",
            )})

    rows.extend({
        "product_code": "", "product_name": item["product_name"], "unit": item["unit"],
        "tax": item["tax"], "invoice_name": item["invoice_name"], "qty": item["qty"] or 0,
        "amount": item["amount"], "unit_cost": item["unit_cost"],
        "source_rows": [item["source_row"]], "warehouse_codes": [item["warehouse_code"]],
        "source_row_count": 1, "create_product": False, "status": "error", "current_qty": 0,
        "warnings": item["warnings"], "errors": item["errors"],
    } for item in orphan_errors)
    rows.sort(key=lambda item: (not bool(item["errors"]), not bool(item["warnings"]), item["product_code"]))
    counts = {
        "source_rows": len(source_rows),
        "items": len(grouped),
        "new_products": sum(item["create_product"] for item in rows),
        "new": sum(item["status"] == "new" for item in rows),
        "update": sum(item["status"] == "update" for item in rows),
        "negative": sum(item["qty"] < 0 for item in rows),
        "warning": sum(bool(item["warnings"]) for item in rows),
        "error": sum(bool(item["errors"]) for item in rows),
    }
    return {
        "sheet": worksheet.title,
        "header_row": header_end,
        "period": period,
        "rows": rows,
        "items": items,
        "counts": counts,
        "totals": {
            "qty": sum(item["qty"] for item in rows),
            "amount": sum(item["amount"] for item in rows),
        },
        "can_confirm": counts["error"] == 0,
    }


MEAL_ATTENDANCE_GROUPS = (
    ("sangolf", "SANGOLF"),
    ("sangold", "SANGOLF"),
    ("dainam", "DAINAM"),
    ("havico", "HAVICO"),
    ("thaco", "THACO"),
    ("sunby", "SUNBY"),
    ("lucky", "LUCKY"),
    ("vina", "VINA"),
    ("united", "UNI"),
    ("uni", "UNI"),
    ("bot", "BOT"),
    ("pot", "BOT"),
    ("tq", "TQ"),
)


def meal_attendance_group(value) -> str:
    key = mapping_key(value)
    for token, code in MEAL_ATTENDANCE_GROUPS:
        if token in key:
            return code
    return ""


def meal_attendance_shift(value) -> str:
    key = mapping_key(value)
    # "Suất ăn đêm + phụ sáng" is the night shift.  Check night
    # first so the explanatory "phụ sáng" text cannot relabel it as morning.
    if "dem" in key or "toi" in key:
        return "Đêm"
    if "sang" in key:
        return "Sáng"
    if "trua" in key:
        return "Trưa"
    if "chieu" in key:
        return "Chiều"
    if "vaosuat" in key:
        return "Tổng"
    return ""


def meal_attendance_header_kind(value) -> str:
    key = mapping_key(value)
    if not key or any(token in key for token in ("chenhlech", "cl", "tong")):
        return ""
    if "dat" in key:
        return "ordered"
    if "an" in key or "vaosuat" in key or key.startswith("ca"):
        return "actual"
    return ""


def workbook_cell_date(value, workbook, period_override="", force_period=False) -> str:
    parsed = ""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if 20_000 <= float(value) <= 80_000:
            try:
                parsed = from_excel(value, workbook.epoch).date().isoformat()
            except (TypeError, ValueError, OverflowError):
                parsed = ""
        elif period_override and float(value).is_integer() and 1 <= int(value) <= 31:
            parsed = f"{period_override}-{int(value):02d}"
    else:
        parsed = as_date(value)
    if parsed and period_override and force_period:
        try:
            year, month = map(int, period_override.split("-"))
            parsed_date = datetime.strptime(parsed, "%Y-%m-%d").date()
            parsed = date(year, month, parsed_date.day).isoformat()
        except ValueError:
            return ""
    try:
        return datetime.strptime(parsed, "%Y-%m-%d").date().isoformat()
    except (TypeError, ValueError):
        return ""


def find_meal_date_column(worksheet, workbook, period_override="") -> tuple[int, dict[int, str]]:
    candidates = []
    max_row = min(worksheet.max_row or 0, 380)
    for column in range(1, min(worksheet.max_column or 0, 6) + 1):
        dates = {}
        for row_index in range(3, max_row + 1):
            value = workbook_cell_date(worksheet.cell(row_index, column).value, workbook)
            if value:
                dates[row_index] = value
        periods = {value[:7] for value in dates.values()}
        if len(set(dates.values())) >= 5 and len(periods) == 1:
            header_key = mapping_key(
                f"{worksheet.cell(1, column).value} {worksheet.cell(2, column).value}"
            )
            bonus = 100 if any(token in header_key for token in ("ngay", "nt", "date")) else 0
            candidates.append((len(set(dates.values())) + bonus, -column, column, dates))
    if candidates:
        _, _, column, dates = max(candidates)
        return column, dates
    if period_override:
        for column in range(1, min(worksheet.max_column or 0, 6) + 1):
            dates = {}
            for row_index in range(3, max_row + 1):
                value = worksheet.cell(row_index, column).value
                if isinstance(value, (int, float)) and not isinstance(value, bool):
                    if float(value).is_integer() and 1 <= int(value) <= 31:
                        parsed = workbook_cell_date(value, workbook, period_override)
                        if parsed:
                            dates[row_index] = parsed
            days = sorted({int(value[-2:]) for value in dates.values()})
            if len(days) >= 5 and all(right - left == 1 for left, right in zip(days, days[1:])):
                header_key = mapping_key(
                    f"{worksheet.cell(1, column).value} {worksheet.cell(2, column).value}"
                )
                bonus = 100 if any(token in header_key for token in ("ngay", "nt", "date")) else 0
                candidates.append((len(days) + bonus, -column, column, dates))
        if candidates:
            _, _, column, dates = max(candidates)
            return column, dates
    raise ValueError("Không nhận diện được cột ngày; nếu file cũ hỏng ngày hãy chọn rõ kỳ tháng rồi thử lại")


LEGACY_MEAL_SHEET_GROUPS = {
    "uni": "UNI",
    "united": "UNI",
    "tq": "TQ",
    "trungquoc": "TQ",
    "lianxin": "LIANXIN",
    "dainam": "DAINAM",
    "sunby": "SUNBY",
    "sangolf": "SANGOLF",
    "sangold": "SANGOLF",
    "havico": "HAVICO",
    "thaco": "THACO",
    "lucky": "LUCKY",
    "vina": "VINA",
}


def legacy_meal_sheet_group(sheet_name) -> str:
    """Return the business unit for a legacy per-unit meal sheet.

    Old workbooks contain both a delivery sheet and one or more invoice copies
    of the same figures.  Importing an invoice copy would double the meals, so
    only the canonical per-unit sheet names are accepted here.
    """
    key = mapping_key(sheet_name)
    if "hoadon" in key or key.endswith("hd"):
        return ""
    return LEGACY_MEAL_SHEET_GROUPS.get(key, "")


def legacy_meal_invoice_sheet_group(sheet_name) -> str:
    """Identify a paired invoice sheet, used only to recover broken displays."""
    key = mapping_key(sheet_name)
    if "hoadon" in key:
        key = key.replace("hoadon", "")
    elif key.endswith("hd"):
        key = key[:-2]
    else:
        return ""
    return LEGACY_MEAL_SHEET_GROUPS.get(key, "")


def legacy_meal_table_layout(worksheet):
    """Locate STT/day, shift and displayed-total columns in a legacy sheet."""
    max_column = min(worksheet.max_column or 0, 40)
    for row_index in range(1, min(worksheet.max_row or 0, 15) + 1):
        stt_column = 0
        date_column = 0
        total_column = 0
        shift_columns = []
        for column in range(1, max_column + 1):
            value = worksheet.cell(row_index, column).value
            key = mapping_key(value)
            if key.startswith("stt"):
                stt_column = column
            elif key.startswith("ngay"):
                date_column = column
            elif key in {"tong", "tongcong"} or key.startswith("tongcong"):
                total_column = column
            shift = meal_attendance_shift(value)
            if shift and shift != "Tổng":
                shift_columns.append((column, shift))
        if stt_column and shift_columns:
            return {
                "header_row": row_index,
                "stt_column": stt_column,
                "date_column": date_column,
                "total_column": total_column,
                "shift_columns": shift_columns,
            }
    return None


def legacy_meal_numeric_cell(cell):
    """Read a count without silently turning a broken Excel value into zero."""
    value = cell.value
    if value in (None, ""):
        return "blank", 0.0
    if isinstance(value, str):
        text = value.strip()
        if text.startswith("#"):
            return "error", 0.0
        if text in {"", "-", "--"}:
            return "blank", 0.0
        normalized = text.replace(" ", "")
        if "," in normalized:
            normalized = normalized.replace(".", "").replace(",", ".")
        try:
            return "number", float(normalized)
        except ValueError as exc:
            raise ValueError(
                f"Sheet {cell.parent.title}, ô {cell.coordinate}: số suất '{text}' không hợp lệ"
            ) from exc
    if isinstance(value, bool):
        raise ValueError(
            f"Sheet {cell.parent.title}, ô {cell.coordinate}: số suất không được là TRUE/FALSE"
        )
    if isinstance(value, (int, float)):
        return "number", float(value)
    raise ValueError(
        f"Sheet {cell.parent.title}, ô {cell.coordinate}: không đọc được số suất"
    )


def parse_legacy_meal_attendance_sheets(workbook, period_override) -> tuple[list, list]:
    """Parse the pre-summary layout used by the January 2026 workbook.

    The workbook's linked/cached date cells are corrupt (mixed 2025 and
    2027-2036 dates).  With an explicit period, STT 1..31 is the only stable
    source for the day.  Broken shift formulas may use a displayed total, but
    the parser never distributes that total across shifts by guessing.
    """
    try:
        year, month = map(int, period_override.split("-"))
        date(year, month, 1)
    except (TypeError, ValueError) as exc:
        raise ValueError("Kỳ chấm suất phải có dạng YYYY-MM") from exc

    # Invoice-format copies are never imported as a second source.  They are
    # retained only as a row-level visible fallback when the delivery sheet has
    # a cached #REF! (as in TQ, day 15 of the real January workbook).
    invoice_fallbacks = {}
    for candidate in workbook.worksheets:
        group = legacy_meal_invoice_sheet_group(candidate.title)
        layout = legacy_meal_table_layout(candidate) if group else None
        if not group or not layout:
            continue
        day_rows = {}
        for row_index in range(layout["header_row"] + 1, min(candidate.max_row or 0, 100) + 1):
            day_value = as_number(candidate.cell(row_index, layout["stt_column"]).value, -1)
            if float(day_value).is_integer() and 1 <= int(day_value) <= 31:
                day_rows[int(day_value)] = row_index
        invoice_fallbacks[group] = (candidate, layout, day_rows)

    raw_items = []
    source_sheet_names = []
    for worksheet in workbook.worksheets:
        kitchen = legacy_meal_sheet_group(worksheet.title)
        if not kitchen:
            continue
        layout = legacy_meal_table_layout(worksheet)
        if not layout:
            continue
        source_sheet_names.append(worksheet.title)
        shift_columns = layout["shift_columns"]
        total_column = layout["total_column"]
        for row_index in range(layout["header_row"] + 1, min(worksheet.max_row or 0, 100) + 1):
            stt_cell = worksheet.cell(row_index, layout["stt_column"])
            stt_key = mapping_key(stt_cell.value)
            if "tong" in stt_key or stt_key in {"cong", "total"}:
                break
            day_number = as_number(stt_cell.value, -1)
            if not float(day_number).is_integer() or not 1 <= int(day_number) <= 31:
                # Header spacers and signature rows are harmless.  A count on a
                # row without a usable STT is not: its date cannot be inferred.
                populated = []
                for column, _ in shift_columns:
                    state, value = legacy_meal_numeric_cell(worksheet.cell(row_index, column))
                    if state == "error" or abs(value) > 1e-12:
                        populated.append(worksheet.cell(row_index, column).coordinate)
                if total_column:
                    state, value = legacy_meal_numeric_cell(worksheet.cell(row_index, total_column))
                    if state == "error" or abs(value) > 1e-12:
                        populated.append(worksheet.cell(row_index, total_column).coordinate)
                if populated:
                    raise ValueError(
                        f"Sheet {worksheet.title}, dòng {row_index}: có số suất tại "
                        f"{', '.join(populated)} nhưng STT/ngày không hợp lệ"
                    )
                continue
            day_number = int(day_number)
            try:
                work_date = date(year, month, day_number).isoformat()
            except ValueError as exc:
                populated = []
                for column, _ in shift_columns:
                    state, value = legacy_meal_numeric_cell(worksheet.cell(row_index, column))
                    if state == "error" or abs(value) > 1e-12:
                        populated.append(worksheet.cell(row_index, column).coordinate)
                if populated:
                    raise ValueError(
                        f"Sheet {worksheet.title}, dòng {row_index}: STT {day_number} "
                        f"không tồn tại trong kỳ {period_override}"
                    ) from exc
                continue

            shift_values = []
            broken_cells = []
            for column, shift in shift_columns:
                cell = worksheet.cell(row_index, column)
                state, value = legacy_meal_numeric_cell(cell)
                if state == "error":
                    broken_cells.append(cell.coordinate)
                else:
                    if value < 0:
                        raise ValueError(
                            f"Sheet {worksheet.title}, ô {cell.coordinate}: số suất không được âm"
                        )
                    shift_values.append((column, shift, value))

            total_state = "blank"
            displayed_total = 0.0
            total_cell = None
            if total_column:
                total_cell = worksheet.cell(row_index, total_column)
                total_state, displayed_total = legacy_meal_numeric_cell(total_cell)
                if total_state == "number" and displayed_total < 0:
                    raise ValueError(
                        f"Sheet {worksheet.title}, ô {total_cell.coordinate}: tổng số suất không được âm"
                    )

            known_total = sum(value for _, _, value in shift_values)
            if broken_cells:
                fallback_note = ""
                total_source_sheet = worksheet.title
                if total_state != "number" and kitchen in invoice_fallbacks:
                    fallback_sheet, fallback_layout, fallback_rows = invoice_fallbacks[kitchen]
                    fallback_row = fallback_rows.get(day_number)
                    if fallback_row:
                        fallback_values = []
                        fallback_broken = []
                        for column, _ in fallback_layout["shift_columns"]:
                            cell = fallback_sheet.cell(fallback_row, column)
                            state, value = legacy_meal_numeric_cell(cell)
                            if state == "error":
                                fallback_broken.append(cell.coordinate)
                            else:
                                fallback_values.append(value)
                        fallback_total_cell = (
                            fallback_sheet.cell(fallback_row, fallback_layout["total_column"])
                            if fallback_layout["total_column"] else None
                        )
                        fallback_total_state, fallback_total = (
                            legacy_meal_numeric_cell(fallback_total_cell)
                            if fallback_total_cell else ("blank", 0.0)
                        )
                        if not fallback_broken and fallback_total_state == "number":
                            fallback_detail_total = sum(fallback_values)
                            if (
                                fallback_total > 0
                                and fallback_detail_total > 0
                                and abs(fallback_total - fallback_detail_total) > 1e-9
                            ):
                                raise ValueError(
                                    f"Sheet {fallback_sheet.title}, dòng {fallback_row}: tổng hiển thị "
                                    f"{fallback_total:g} khác tổng chi tiết ca {fallback_detail_total:g}"
                                )
                            displayed_total = fallback_total
                            total_state = "number"
                            total_cell = fallback_total_cell
                            total_source_sheet = fallback_sheet.title
                            fallback_note = (
                                f"; đối chiếu tổng hiển thị tại "
                                f"{fallback_sheet.title}!{fallback_total_cell.coordinate}"
                            )
                if total_state != "number":
                    total_ref = total_cell.coordinate if total_cell else "không có cột Tổng"
                    raise ValueError(
                        f"Sheet {worksheet.title}, dòng {row_index}: ô ca "
                        f"{', '.join(broken_cells)} bị lỗi và {total_ref} không có tổng hiển thị hợp lệ"
                    )
                if known_total > displayed_total + 1e-9:
                    raise ValueError(
                        f"Sheet {worksheet.title}, dòng {row_index}: tổng hiển thị "
                        f"{displayed_total:g} tại {total_cell.coordinate} nhỏ hơn chi tiết ca đã biết {known_total:g}"
                    )
                if displayed_total > 0:
                    raw_items.append({
                        "work_date": work_date, "kitchen": kitchen, "shift": "Tổng",
                        "actual_count": displayed_total, "ordered_count": 0,
                        "source_sheet": total_source_sheet,
                        "source_column": total_cell.coordinate,
                        "errors": [],
                        "warnings": [
                            f"Chi tiết ca {', '.join(broken_cells)} bị lỗi; dùng tổng hiển thị {total_cell.coordinate}{fallback_note}, không tự phân bổ theo ca"
                        ],
                    })
                continue

            if total_state == "number" and displayed_total > 0:
                if known_total <= 1e-12:
                    raw_items.append({
                        "work_date": work_date, "kitchen": kitchen, "shift": "Tổng",
                        "actual_count": displayed_total, "ordered_count": 0,
                        "source_sheet": worksheet.title,
                        "source_column": total_cell.coordinate,
                        "errors": [],
                        "warnings": [
                            f"Chi tiết ca trống; dùng tổng hiển thị {total_cell.coordinate}, không tự phân bổ theo ca"
                        ],
                    })
                    continue
                if abs(displayed_total - known_total) > 1e-9:
                    source_cells = ", ".join(
                        worksheet.cell(row_index, column).coordinate
                        for column, _, value in shift_values if abs(value) > 1e-12
                    )
                    raise ValueError(
                        f"Sheet {worksheet.title}, dòng {row_index}: tổng hiển thị "
                        f"{displayed_total:g} tại {total_cell.coordinate} khác tổng chi tiết ca "
                        f"{known_total:g} tại {source_cells}"
                    )

            common_warnings = []
            if total_state == "error":
                common_warnings.append(
                    f"Ô tổng {total_cell.coordinate} bị lỗi; dùng các số ca hiển thị"
                )
            for column, shift, value in shift_values:
                if value <= 1e-12:
                    continue
                raw_items.append({
                    "work_date": work_date, "kitchen": kitchen, "shift": shift,
                    "actual_count": value, "ordered_count": 0,
                    "source_sheet": worksheet.title,
                    "source_column": worksheet.cell(row_index, column).coordinate,
                    "errors": [], "warnings": list(common_warnings),
                })
    return raw_items, source_sheet_names


def parse_meal_attendance_workbook(conn, workbook, period_override="") -> dict:
    if period_override and not re.fullmatch(r"\d{4}-\d{2}", period_override):
        raise ValueError("Kỳ chấm suất phải có dạng YYYY-MM")
    worksheet = workbook["SUẤT ĂN"] if "SUẤT ĂN" in workbook.sheetnames else None
    if worksheet is None and not period_override:
        raise ValueError(
            "File cũ thiếu sheet SUẤT ĂN và kỳ nguồn mâu thuẫn; cần chọn rõ kỳ tháng trước khi nhập"
        )
    raw_items = []
    source_sheet_names = []
    if worksheet is not None:
        source_sheet_names.append(worksheet.title)
        columns = {}
        current_group = ""
        for column in range(3, min(worksheet.max_column or 0, 80) + 1):
            group_label = worksheet.cell(1, column).value
            if group_label not in (None, ""):
                current_group = meal_attendance_group(group_label)
            shift = meal_attendance_shift(worksheet.cell(2, column).value)
            kind = meal_attendance_header_kind(worksheet.cell(2, column).value)
            if current_group and shift and kind:
                columns.setdefault((current_group, shift), {})[kind] = column
        if not columns:
            raise ValueError("Không nhận diện được cột bếp/ca trong sheet SUẤT ĂN")
        _, row_dates = find_meal_date_column(worksheet, workbook, period_override)
        source_value_columns = sorted({column for group in columns.values() for column in group.values()})
        medians = {}
        for column in source_value_columns:
            values = [
                abs(as_number(worksheet.cell(row_index, column).value))
                for row_index in row_dates
                if abs(as_number(worksheet.cell(row_index, column).value)) > 0
            ]
            medians[column] = median(values) if values else 0
        for row_index, work_date in sorted(row_dates.items()):
            for (kitchen, shift), source_columns in columns.items():
                actual_column = source_columns.get("actual")
                ordered_column = source_columns.get("ordered")
                actual = as_number(worksheet.cell(row_index, actual_column).value) if actual_column else 0
                ordered = as_number(worksheet.cell(row_index, ordered_column).value) if ordered_column else 0
                if abs(actual) <= 1e-12 and abs(ordered) <= 1e-12:
                    continue
                errors = ["Số suất không được âm"] if actual < 0 or ordered < 0 else []
                for label, value, source_column in (
                    ("Số ăn thực tế", actual, actual_column), ("Số đặt", ordered, ordered_column),
                ):
                    if source_column and value > max(1000, medians.get(source_column, 0) * 10):
                        errors.append(f"{label} {value:g} là ngoại lệ quá lớn; cần kiểm tra ô {get_column_letter(source_column)}{row_index}")
                source_column = "/".join(
                    get_column_letter(column) for column in sorted(set(source_columns.values()))
                )
                raw_items.append({
                    "work_date": work_date, "kitchen": kitchen, "shift": shift,
                    "actual_count": actual, "ordered_count": ordered,
                    "source_sheet": worksheet.title, "source_column": source_column,
                    "errors": errors, "warnings": [],
                })

    if worksheet is None:
        legacy_items, legacy_sheet_names = parse_legacy_meal_attendance_sheets(
            workbook, period_override,
        )
        raw_items.extend(legacy_items)
        source_sheet_names.extend(legacy_sheet_names)

    # TTS dùng ma trận ngày theo cột, nhóm người ăn theo dòng; tổng từng ngày là số suất thực tế.
    if "TTS" in workbook.sheetnames:
        tts = workbook["TTS"]
        source_sheet_names.append(tts.title)
        force_period = worksheet is None and bool(period_override)
        for column in range(3, min(tts.max_column or 0, 80) + 1):
            if mapping_key(tts.cell(4, column).value) == "tong":
                break
            work_date = workbook_cell_date(
                tts.cell(4, column).value, workbook, period_override, force_period=force_period,
            )
            if not work_date:
                continue
            actual = 0.0
            for row_index in range(6, min(tts.max_row or 0, 300) + 1):
                label = mapping_key(tts.cell(row_index, 1).value)
                if label == "tong":
                    break
                if label:
                    actual += max(as_number(tts.cell(row_index, column).value), 0)
            if actual > 0:
                raw_items.append({
                    "work_date": work_date, "kitchen": "TTS", "shift": "Tổng",
                    "actual_count": actual, "ordered_count": 0,
                    "source_sheet": tts.title, "source_column": tts.cell(4, column).column_letter,
                    "errors": [], "warnings": [],
                })

    if not raw_items:
        raise ValueError("File không có số suất ăn theo ngày để nhập")
    grouped = {}
    for item in raw_items:
        key = (item["work_date"], item["kitchen"], item["shift"])
        if key in grouped:
            grouped[key]["actual_count"] += item["actual_count"]
            grouped[key]["ordered_count"] += item["ordered_count"]
            grouped[key]["source_column"] += "+" + item["source_column"]
            grouped[key]["warnings"].append("Đã gộp nhiều vùng nguồn cùng ngày/bếp/ca")
            grouped[key]["errors"].extend(item["errors"])
        else:
            grouped[key] = item

    existing = {
        (row["work_date"], row["kitchen"], row["shift"]): dict(row)
        for row in conn.execute("SELECT * FROM meal_attendance")
    }
    rows = []
    items = []
    for key in sorted(grouped):
        item = grouped[key]
        for label, value in (
            ("Số ăn thực tế", item["actual_count"]),
            ("Số đặt", item["ordered_count"]),
        ):
            if not math.isfinite(float(value)):
                item["errors"].append(f"{label} phải là số hữu hạn")
            elif abs(float(value) - round(float(value))) > 1e-9:
                item["errors"].append(f"{label} phải là số suất nguyên")
        previous = existing.get(key)
        if previous is None:
            status = "new"
        elif (
            abs(previous["actual_count"] - item["actual_count"]) <= 1e-9
            and abs(previous["ordered_count"] - item["ordered_count"]) <= 1e-9
        ):
            status = "unchanged"
        else:
            status = "update"
        item["status"] = status
        if item["ordered_count"] > 0 and abs(item["ordered_count"] - item["actual_count"]) > 1e-9:
            item["warnings"].append("Số đặt và số ăn thực tế chênh nhau")
        item["warnings"] = list(dict.fromkeys(item["warnings"]))
        item["errors"] = list(dict.fromkeys(item["errors"]))
        rows.append(item)
        if not item["errors"]:
            items.append({field: item[field] for field in (
                "work_date", "kitchen", "shift", "actual_count", "ordered_count",
                "source_sheet", "source_column",
            )})
    periods = sorted({item["work_date"][:7] for item in rows})
    if len(periods) > 1:
        for item in rows:
            item["errors"].append("File chứa nhiều kỳ tháng; cần tách file hoặc chọn đúng kỳ nguồn")
        items = []
    counts = {
        "items": len(rows),
        "dates": len({item["work_date"] for item in rows}),
        "kitchens": len({item["kitchen"] for item in rows}),
        "new": sum(item["status"] == "new" for item in rows),
        "update": sum(item["status"] == "update" for item in rows),
        "unchanged": sum(item["status"] == "unchanged" for item in rows),
        "warning": sum(bool(item["warnings"]) for item in rows),
        "error": sum(bool(item["errors"]) for item in rows),
    }
    return {
        "sheet": ", ".join(dict.fromkeys(source_sheet_names)),
        "periods": periods,
        "rows": rows,
        "items": items,
        "counts": counts,
        "totals": {
            "actual": sum(item["actual_count"] for item in rows),
            "ordered": sum(item["ordered_count"] for item in rows),
        },
        "can_confirm": counts["error"] == 0,
    }


def kitchen_shift_label(values, previous="Ca sáng") -> str:
    for value in values:
        key = mapping_key(value)
        if "cachieu" in key:
            return "Ca chiều"
        if "catrua" in key:
            return "Ca trưa"
        if "cadem" in key:
            return "Ca đêm"
        if "casang" in key:
            return "Ca sáng"
    return previous or "Ca sáng"


def kitchen_block_name(segment, xcom_code: str) -> str:
    ignored = ("sosuat", "casang", "cachieu", "catrua", "cadem", "mahang", "mabep", "nhathau")
    candidates = []
    for row in segment:
        value = mapping_cell_text(row[4])
        key = mapping_key(value)
        if value and not any(token in key for token in ignored):
            candidates.append(value)
    name = candidates[-1] if candidates else xcom_code
    return re.sub(r"^(BẾP|BEP)\s+", "", name, flags=re.IGNORECASE).strip().upper()


def kitchen_financials(block_rows) -> dict:
    result = {
        "meal_price": 0,
        "source_revenue": 0,
        "source_food_cost": 0,
        "source_total_cost": 0,
        "source_profit": 0,
        "other_cost": 0,
        "components": [],
        "warnings": [],
    }
    for row in block_rows:
        label = mapping_cell_text(row[10])
        if not label:
            continue
        key = mapping_key(label)
        raw_value = row[11]
        value = as_number(raw_value)
        result["components"].append({"label": label, "value": value if raw_value not in (None, "") else None})
        if "tientinhcost" in key or key == "suatan":
            result["meal_price"] = value
        elif "doanhthu" in key or "tienthukhachhang" in key:
            result["source_revenue"] = value
        elif "tienchiphithucpham" in key:
            result["source_food_cost"] = value
        elif "tongchi" in key:
            result["source_total_cost"] = value
        elif "loinhuan" in key or "costca" in key:
            result["source_profit"] = value
        elif any(token in key for token in ("giavi", "matbang", "nhancong", "vanchuyen")) or key.startswith("gachi"):
            if raw_value in (None, ""):
                result["warnings"].append(f"{label}: chưa có số tiền")
            else:
                result["other_cost"] += value
    return result


KITCHEN_MISSING_MEAL_COUNT_ERROR = "Không tìm thấy số suất của nhóm"


def kitchen_sheet_day_offset(sheet_title: str):
    """Return the T2..CN offset from Monday without treating T20 as T2."""
    key = mapping_key(sheet_title)
    for prefix, offset in (("t2", 0), ("t3", 1), ("t4", 2), ("t5", 3), ("t6", 4), ("t7", 5)):
        if key.startswith(prefix) and (len(key) == len(prefix) or not key[len(prefix)].isdigit()):
            return offset
    if key.startswith("cn") or key.startswith("chunhat"):
        return 6
    return None


def apply_kitchen_meal_count_overrides(plans: list, overrides: dict) -> list[str]:
    """Apply explicit user-entered totals to plans whose workbook has no meal count."""
    errors = []
    overrides = overrides if isinstance(overrides, dict) else {}
    for plan in plans:
        if not plan.get("needs_meal_count"):
            continue
        plan_key = plan.get("plan_key", "")
        raw_value = overrides.get(plan_key)
        meal_count = as_number(raw_value)
        if meal_count <= 0 or abs(meal_count - round(meal_count)) > 1e-9:
            errors.append(
                f"{plan.get('sheet', '')} · {plan.get('kitchen', '')} · "
                f"{plan.get('shift', '')}: cần nhập số suất nguyên lớn hơn 0"
            )
            continue
        meal_count = float(round(meal_count))
        plan["meal_count"] = meal_count
        for item in plan.get("items", []):
            required_qty = as_number(item.get("file_required_qty"))
            if required_qty <= 0:
                required_qty = as_number(item.get("norm_per_1000")) * meal_count / 1000
            item["required_qty"] = required_qty
            item["norm_qty"] = required_qty / meal_count
            norm_per_1000 = as_number(item.get("norm_per_1000"))
            item["applicable_meal_count"] = (
                required_qty * 1000 / norm_per_1000
                if norm_per_1000 > 0 and required_qty > 0
                else meal_count
            )
            item["calculated_amount"] = required_qty * as_number(item.get("buy_price"))

        financials = plan.get("source_financials") or {}
        food_cost = sum(as_number(item.get("calculated_amount")) for item in plan.get("items", []))
        revenue = meal_count * as_number(plan.get("meal_price"))
        total_cost = food_cost + as_number(plan.get("other_cost"))
        financials.update({
            "calculated_revenue": revenue,
            "calculated_food_cost": food_cost,
            "calculated_total_cost": total_cost,
            "calculated_profit": revenue - total_cost,
        })
        plan["source_financials"] = financials
        plan["errors"] = [
            message for message in plan.get("errors", [])
            if message != KITCHEN_MISSING_MEAL_COUNT_ERROR
        ]
        plan["warnings"] = list(dict.fromkeys([
            *plan.get("warnings", []),
            f"Số suất {int(meal_count)} được người dùng nhập khi xác nhận vì file để trống",
        ]))
        plan["needs_meal_count"] = False
        plan["override_only"] = False
    return errors


def parse_kitchen_workbook(conn, workbook, work_date: str) -> dict:
    anchor_date = date.fromisoformat(work_date)
    products = {
        row["code"]: dict(row)
        for row in conn.execute("SELECT code,name,unit,supplier FROM products")
    }
    product_codes_by_name = defaultdict(list)
    for product in products.values():
        product_codes_by_name[mapping_key(product["name"])].append(product["code"])
    known_kitchens = {row["code"] for row in conn.execute("SELECT code FROM kitchens")}
    dated_hatran_prices = {
        (row["product_code"], row["period"]): as_number(row["price_value"])
        for row in conn.execute(
            "SELECT product_code,period,price_value FROM dated_prices WHERE price_group='HATRAN'"
        )
    }
    plans = []
    plan_occurrences = Counter()
    sheet_offsets = {
        worksheet.title: kitchen_sheet_day_offset(worksheet.title)
        for worksheet in workbook.worksheets
    }
    weekly_mode = len({offset for offset in sheet_offsets.values() if offset is not None}) >= 2
    if weekly_mode and anchor_date.weekday() != 0:
        raise ValueError("File tuần T2–CN yêu cầu chọn ngày Thứ Hai làm ngày bắt đầu tuần")
    for worksheet in workbook.worksheets:
        sheet_offset = sheet_offsets[worksheet.title]
        if weekly_mode and sheet_offset is None:
            continue
        plan_work_date = (
            anchor_date + timedelta(days=sheet_offset)
            if sheet_offset is not None
            else anchor_date
        ).isoformat()
        values = [
            list(row)
            for row in worksheet.iter_rows(
                min_row=1,
                max_row=min(worksheet.max_row or 0, 2_000),
                max_col=13,
                values_only=True,
            )
        ]
        header_row = 0
        for index, row in enumerate(values, start=1):
            if mapping_key(row[0]) == "nhathau" and mapping_key(row[1]) in {"mahang", "mavt", "masp"} and mapping_key(row[2]) == "mabep":
                header_row = index
                break
        if not header_row:
            continue

        markers = [header_row]
        for row_number in range(header_row + 1, len(values) + 1):
            row = values[row_number - 1]
            code = mapping_cell_text(row[1]).upper()
            label = mapping_cell_text(row[4])
            if (not code or code == "-") and label and as_number(row[5]) > 0:
                markers.append(row_number)

        previous_shift = "Ca sáng"
        for run_index, marker_row in enumerate(markers, start=1):
            block_end = (markers[run_index] - 1) if run_index < len(markers) else len(values)
            run = []
            for row_number in range(marker_row + 1, block_end + 1):
                row = values[row_number - 1]
                code = mapping_cell_text(row[1]).upper()
                xcom = mapping_cell_text(row[2]).upper()
                if code and code != "-" and xcom and as_number(row[5]) > 0:
                    run.append(row_number)
            if not run:
                continue
            start_row, end_row = run[0], run[-1]
            marker = values[marker_row - 1]
            marker_label = mapping_cell_text(marker[4])
            previous_shift = kitchen_shift_label([marker_label], previous_shift)
            servings_per_menu = as_number(marker[5])
            marker_total = as_number(marker[6])
            meal_count = marker_total if marker_total >= servings_per_menu > 0 else servings_per_menu
            ratio = meal_count / servings_per_menu if servings_per_menu > 0 else 1
            menu_count = max(int(round(ratio)), 1) if abs(ratio - round(ratio)) <= 0.01 else 1
            xcom_counts = Counter(
                mapping_cell_text(values[row_number - 1][2]).upper() for row_number in run
            )
            xcom_code = xcom_counts.most_common(1)[0][0]
            marker_key = mapping_key(marker_label)
            if xcom_code and xcom_code != "XCOM":
                kitchen = xcom_code
            elif marker_label and not any(token in marker_key for token in (
                "sosuat", "casang", "cachieu", "catrua", "cadem", "mahang", "mabep", "nhathau",
            )):
                kitchen = re.sub(r"^(BẾP|BEP)\s+", "", marker_label, flags=re.IGNORECASE).strip().upper()
            else:
                segment = values[max(0, marker_row - 2):marker_row - 1]
                kitchen = kitchen_block_name(segment, xcom_code)
            plan_errors = []
            plan_warnings = []
            if meal_count <= 0:
                plan_errors.append(KITCHEN_MISSING_MEAL_COUNT_ERROR)
            if servings_per_menu > 0 and marker_total > 0 and abs(ratio - round(ratio)) > 0.01:
                plan_warnings.append("Tổng suất không chia hết cho số suất/thực đơn; cần kiểm tra lại")
            if kitchen not in known_kitchens:
                plan_warnings.append(f"{kitchen}: chưa có trong danh mục bếp; sẽ ghi nhận từ file xưởng cơm")
            if len(xcom_counts) > 1:
                plan_errors.append("Một nhóm có nhiều mã XCOM khác nhau")

            items = []
            current_dish = ""
            seen_item_names = defaultdict(list)
            for row_number in run:
                row = values[row_number - 1]
                code = mapping_cell_text(row[1]).upper()
                dish = mapping_cell_text(row[3])
                if dish:
                    current_dish = dish
                norm_raw = as_number(row[5])
                file_required = as_number(row[6])
                formula_norm_qty = norm_raw / 1000
                formula_required = formula_norm_qty * meal_count
                required_qty = file_required if file_required > 0 else formula_required
                norm_qty = required_qty / meal_count if meal_count > 0 else formula_norm_qty
                applicable_meal_count = (
                    required_qty * 1000 / norm_raw if norm_raw > 0 and required_qty > 0 else meal_count
                )
                file_price = as_number(row[8])
                price_group = mapping_cell_text(row[0]).upper()
                price_period = plan_work_date[:7]
                dated_price = dated_hatran_prices.get((code, price_period), 0)
                if dated_price > 0:
                    price = dated_price
                    price_source = f"HATRAN {price_period} · giá kỳ đã khóa"
                elif file_price > 0:
                    # The workbook itself is the current-period HATRAN source.
                    # Confirmation below locks this exact value to the period;
                    # never borrow the timeless/previous catalogue price.
                    price = file_price
                    price_source = f"HATRAN {price_period} · giá file chờ xác nhận"
                else:
                    price = 0
                    price_source = f"HATRAN {price_period} · chưa có giá"
                source_amount_raw = row[9]
                source_amount = as_number(source_amount_raw)
                calculated_amount = required_qty * price
                item_errors = []
                item_warnings = []
                product = products.get(code)
                if not product:
                    item_errors.append(f"Mã {code} chưa có trong danh mục")
                item_key = (mapping_key(current_dish), code)
                source_name = mapping_cell_text(row[4])
                source_name_key = mapping_key(source_name)
                previous_names = seen_item_names[item_key]
                if previous_names:
                    if all(previous["key"] == source_name_key for previous in previous_names):
                        item_warnings.append(
                            f"Mã {code} và tên {source_name or '(trống)'} lặp trong cùng món; giữ nguyên từng dòng"
                        )
                    else:
                        prior_names = ", ".join(dict.fromkeys(
                            previous["name"] or "(trống)" for previous in previous_names
                        ))
                        message = (
                            f"Mã {code} dùng cho nhiều tên trong cùng món {current_dish or '(chưa có tên)'}: "
                            f"{prior_names} / {source_name or '(trống)'}"
                        )
                        exact_codes = [
                            exact_code for exact_code in product_codes_by_name.get(source_name_key, [])
                            if exact_code != code
                        ]
                        if exact_codes:
                            message += f"; danh mục có tên khớp chính xác ở mã {', '.join(sorted(exact_codes))}"
                        item_errors.append(message)
                seen_item_names[item_key].append({"key": source_name_key, "name": source_name})
                if price_group and price_group != "HATRAN":
                    item_warnings.append(f"Nguồn file ghi {price_group}; hệ thống vẫn bắt buộc đối chiếu giá HATRAN {price_period}")
                if price <= 0:
                    item_errors.append(f"Mã {code} chưa có giá HATRAN đúng kỳ {price_period}")
                elif file_price > 0 and abs(file_price - price) > 1:
                    item_warnings.append(
                        f"Giá trong file {file_price:,.0f} lệch giá HATRAN {price_period} {price:,.0f}; hệ thống dùng giá HATRAN"
                    )
                if file_required > 0 and abs(file_required - formula_required) > max(0.01, formula_required * 0.01):
                    item_warnings.append("Số lượng áp dụng khác tổng suất (có thể do chia thực đơn); hệ thống giữ đúng số trong file")
                if calculated_amount > 0 and source_amount_raw in (None, ""):
                    item_warnings.append("Cột Thành tiền đang trống; hệ thống tự tính Số lượng × Đơn giá")
                elif source_amount_raw not in (None, "") and abs(source_amount - calculated_amount) > max(1, calculated_amount * 0.001):
                    item_warnings.append("Thành tiền trong file lệch Số lượng × Đơn giá; hệ thống dùng số tính lại")
                item = {
                    "source_row": row_number,
                    "product_code": code,
                    "source_name": source_name,
                    "product_name": product["name"] if product else source_name,
                    "dish_name": current_dish,
                    "norm_per_1000": norm_raw,
                    "norm_qty": norm_qty,
                    "applicable_meal_count": applicable_meal_count,
                    "required_qty": required_qty,
                    "file_required_qty": file_required,
                    "source_amount": source_amount,
                    "calculated_amount": calculated_amount,
                    "unit": mapping_cell_text(row[7]) or (product["unit"] if product else ""),
                    "supplier": product["supplier"] if product else "",
                    "buy_price": price,
                    "price_source": price_source,
                    "errors": item_errors,
                    "warnings": item_warnings,
                }
                items.append(item)
                plan_errors.extend(item_errors)
                plan_warnings.extend(item_warnings)

            financials = kitchen_financials(values[marker_row:block_end])
            plan_warnings.extend(financials["warnings"])
            food_cost = sum(item["calculated_amount"] for item in items)
            calculated_revenue = meal_count * financials["meal_price"]
            calculated_total_cost = food_cost + financials["other_cost"]
            calculated_profit = calculated_revenue - calculated_total_cost
            if financials["meal_price"] <= 0:
                plan_warnings.append("Chưa tìm thấy đơn giá suất ăn trong bảng tính cost")
            if financials["source_revenue"] > 0 and abs(financials["source_revenue"] - calculated_revenue) > 1:
                plan_warnings.append("Doanh thu trong file lệch Tổng suất × Đơn giá suất")
            if financials["source_food_cost"] > 0 and abs(financials["source_food_cost"] - food_cost) > 1:
                plan_warnings.append("Chi phí thực phẩm tổng trong file bị thiếu/cũ; hệ thống đã tính lại từng nguyên liệu")
            if financials["source_total_cost"] > 0 and abs(financials["source_total_cost"] - calculated_total_cost) > 1:
                plan_warnings.append("Tổng chi trong file bị thiếu/cũ; hệ thống đã tính lại từ chi tiết")
            financials.update({
                "calculated_revenue": calculated_revenue,
                "calculated_food_cost": food_cost,
                "calculated_total_cost": calculated_total_cost,
                "calculated_profit": calculated_profit,
            })
            financials["warnings"] = list(dict.fromkeys(
                warning for warning in plan_warnings
                if "chưa có trong danh mục bếp" not in warning
            ))
            occurrence_key = (plan_work_date, kitchen, previous_shift)
            plan_occurrences[occurrence_key] += 1
            key_source = "|".join([
                plan_work_date, kitchen, previous_shift, str(plan_occurrences[occurrence_key]),
            ])
            import_key = hashlib.sha256(key_source.encode("utf-8")).hexdigest()
            existing = conn.execute("SELECT id FROM meal_plans WHERE import_key=?", (import_key,)).fetchone()
            if not existing:
                existing = conn.execute(
                    """SELECT id FROM meal_plans WHERE work_date=? AND kitchen=? AND shift=?
                       AND COALESCE(source_file,'')!='' ORDER BY id DESC LIMIT 1""",
                    (plan_work_date, kitchen, previous_shift),
                ).fetchone()
            dishes = list(dict.fromkeys(item["dish_name"] for item in items if item["dish_name"]))
            plans.append({
                "import_key": import_key,
                "plan_key": import_key,
                "existing_id": existing["id"] if existing else None,
                "status": "update" if existing else "new",
                "sheet": worksheet.title,
                "work_date": plan_work_date,
                "source_row_start": start_row,
                "source_row_end": end_row,
                "kitchen": kitchen,
                "xcom_code": xcom_code,
                "shift": previous_shift,
                "meal_count": meal_count,
                "menu_count": menu_count,
                "servings_per_menu": servings_per_menu,
                "meal_price": financials["meal_price"],
                "other_cost": financials["other_cost"],
                "source_financials": financials,
                "menu_name": " · ".join(dishes),
                "items": items,
                "errors": list(dict.fromkeys(plan_errors)),
                "warnings": list(dict.fromkeys(plan_warnings)),
            })
    if not plans:
        raise ValueError("Không tìm thấy bảng xưởng cơm có Nhà thầu, Mã hàng, Mã bếp và định lượng")
    xcom_by_kitchen = defaultdict(set)
    for plan in plans:
        xcom_by_kitchen[plan["kitchen"]].add(plan["xcom_code"])
    for kitchen, xcom_codes in xcom_by_kitchen.items():
        if len(xcom_codes) > 1:
            message = f"Bếp {kitchen} có nhiều mã XCOM trong cùng file: {', '.join(sorted(xcom_codes))}"
            for plan in plans:
                if plan["kitchen"] == kitchen:
                    plan["errors"].append(message)
    for plan in plans:
        plan["errors"] = list(dict.fromkeys(plan["errors"]))
        plan["warnings"] = list(dict.fromkeys(plan["warnings"]))
        plan["needs_meal_count"] = plan["meal_count"] <= 0
        plan["override_only"] = (
            plan["needs_meal_count"]
            and bool(plan["errors"])
            and all(message == KITCHEN_MISSING_MEAL_COUNT_ERROR for message in plan["errors"])
        )
    counts = {
        "plans": len(plans),
        "items": sum(len(plan["items"]) for plan in plans),
        "new": sum(plan["status"] == "new" for plan in plans),
        "update": sum(plan["status"] == "update" for plan in plans),
        "errors": sum(bool(plan["errors"]) for plan in plans),
        "warnings": sum(bool(plan["warnings"]) for plan in plans),
    }
    work_dates = sorted({plan["work_date"] for plan in plans})
    return {
        "plans": plans,
        "counts": counts,
        "can_confirm": counts["errors"] == 0,
        "can_confirm_with_overrides": all(
            not plan["errors"] or plan["override_only"] for plan in plans
        ),
        "weekly": weekly_mode,
        "work_dates": work_dates,
    }


PAYABLE_ALIASES = {
    "purchase_date": {"ngaythang", "ngay", "ngaymua", "ngaynhap"},
    "kitchen": {"tenbep", "bep", "mabep"},
    "item_name": {"tenhang", "tenhanghoa", "tenvattu"},
    "qty": {"soluong", "sldat", "slnhan"},
    "unit": {"dvt", "donvitinh"},
    "supplier": {"ncc", "nhacungcap", "nguoinhan"},
    "buy_price": {"giamua", "dongiamua", "dongia"},
    "damaged_qty": {"hong", "hanghong"},
    "added_qty": {"them", "phatsinhthem"},
    "reduced_qty": {"giam", "tralai"},
    "missing_qty": {"thieu", "giaothieu"},
    "actual_qty": {"slthucte", "soluongthucte", "thucte"},
    "amount": {"thanhtien", "tongtien", "sotien"},
    "note": {"ghichu", "diengiai"},
}


def payable_header_fields(row) -> dict:
    found = {}
    for column_index, value in enumerate(row, start=1):
        key = mapping_key(value)
        for field_name, accepted in PAYABLE_ALIASES.items():
            if field_name not in found and key in accepted:
                found[field_name] = column_index
                break
    return found


def find_payable_sheet(workbook):
    required = {"purchase_date", "kitchen", "item_name", "qty", "supplier", "buy_price"}
    candidates = []
    for sheet_index, worksheet in enumerate(workbook.worksheets):
        for row_index, row in enumerate(
            worksheet.iter_rows(min_row=1, max_row=25, max_col=30, values_only=True), start=1,
        ):
            fields = payable_header_fields(row)
            if required.issubset(fields):
                candidates.append((len(fields), -sheet_index, -row_index, worksheet, row_index, fields))
    if not candidates:
        return None
    distinct_locations = {
        (candidate[3].title, candidate[4]) for candidate in candidates
    }
    if len(distinct_locations) > 1:
        locations = ", ".join(
            f"{sheet}!{row}" for sheet, row in sorted(distinct_locations)[:5]
        )
        raise ValueError(
            "File có nhiều bảng công nợ hợp lệ "
            f"({locations}). Hãy chỉ giữ một bảng tổng hợp duy nhất để tránh "
            "bỏ sót sheet hoặc thay nhầm toàn bộ lịch sử."
        )
    _, _, _, worksheet, row_index, fields = max(candidates, key=lambda item: item[:3])
    return worksheet, row_index, fields


def parse_historical_payables(conn, workbook) -> dict:
    found = find_payable_sheet(workbook)
    if not found:
        raise ValueError("Không tìm thấy bảng công nợ có Ngày, Tên bếp, Tên hàng, Số lượng, NCC và Giá mua")
    worksheet, header_row, fields = found
    supplier_names = {}
    for row in conn.execute(
        "SELECT code supplier FROM suppliers WHERE TRIM(COALESCE(code,''))!='' "
        "UNION SELECT supplier FROM products WHERE TRIM(COALESCE(supplier,''))!='' "
        "UNION SELECT supplier FROM historical_payable_lines"
    ):
        value = mapping_cell_text(row["supplier"])
        supplier_names.setdefault(mapping_key(value), value)

    rows = []
    items = []
    scanned = 0
    for row_number, row in enumerate(
        worksheet.iter_rows(
            min_row=header_row + 1,
            max_row=min(worksheet.max_row, header_row + MAPPING_IMPORT_MAX_ROWS),
            max_col=max(max(fields.values()), fields.get("actual_qty", 0) + 1),
            values_only=True,
        ),
        start=header_row + 1,
    ):
        def cell(name, default=None):
            column = fields.get(name)
            return row[column - 1] if column else default

        raw_date = cell("purchase_date")
        kitchen = mapping_cell_text(cell("kitchen")).upper()
        item_name = mapping_cell_text(cell("item_name"))
        raw_supplier = mapping_cell_text(cell("supplier"))
        errors = []

        def financial_number(field, label):
            try:
                return import_cell_number(cell(field), label)
            except ValueError as exc:
                errors.append(str(exc))
                return 0.0

        qty = financial_number("qty", "Số lượng")
        buy_price = financial_number("buy_price", "Giá mua")
        damaged = financial_number("damaged_qty", "Số lượng hỏng")
        added = financial_number("added_qty", "Số lượng thêm")
        reduced = financial_number("reduced_qty", "Số lượng giảm/trả")
        missing = financial_number("missing_qty", "Số lượng thiếu")
        cached_actual = financial_number("actual_qty", "Số thực tế")
        amount_column = fields.get("amount") or (fields.get("actual_qty", 0) + 1)
        raw_source_amount = row[amount_column - 1] if 0 < amount_column <= len(row) else None
        try:
            source_amount = import_cell_number(raw_source_amount, "Thành tiền")
        except ValueError as exc:
            errors.append(str(exc))
            source_amount = 0.0
        if not any((raw_date, kitchen, item_name, raw_supplier, qty, buy_price, damaged, added, reduced, missing)):
            continue
        scanned += 1
        purchase_date = as_date(raw_date)
        actual_qty = qty + added - damaged - reduced - missing
        calculated_amount = actual_qty * buy_price
        amount = source_amount if raw_source_amount not in (None, "") else calculated_amount
        supplier = supplier_names.get(mapping_key(raw_supplier), raw_supplier.upper())
        warnings = []
        apply = True
        zero_value_line = abs(actual_qty) > 1e-9 and buy_price <= 0 and abs(amount) <= 1e-9
        if actual_qty < -1e-9:
            warnings.append("Số lượng âm được giữ làm dòng ghi giảm/trả lại")
        if abs(amount) > 1e-9 and not purchase_date:
            errors.append("Thiếu hoặc sai ngày mua")
        elif not purchase_date:
            apply = False
            warnings.append("Dòng không phát sinh phải trả và thiếu ngày nên được bỏ qua")
        if zero_value_line:
            warnings.append("Có số lượng nhưng Giá mua và Thành tiền đều bằng 0; giữ dòng lịch sử với giá trị phải trả 0 đồng")
        if apply and abs(actual_qty) > 1e-9 and not item_name:
            errors.append("Thiếu tên hàng")
        if apply and abs(actual_qty) > 1e-9 and not supplier:
            errors.append("Thiếu nhà cung cấp")
        if apply and abs(actual_qty) > 1e-9 and buy_price <= 0 and not zero_value_line:
            errors.append("Thiếu giá mua")
        if "actual_qty" in fields and abs(cached_actual - actual_qty) > 0.001:
            warnings.append("Số thực tế lưu trong Excel lệch công thức; hệ thống tính lại từ hỏng/thêm/giảm/thiếu")
        if raw_source_amount not in (None, "") and abs(source_amount - calculated_amount) > 1:
            warnings.append("Thành tiền trong file lệch SL thực tế × Giá mua; hệ thống giữ số tiền khách đã chốt")
        status = "error" if errors else ("skip" if not apply else "ready")
        item = {
            "source_row": row_number,
            "purchase_date": purchase_date,
            "kitchen": kitchen,
            "item_name": item_name,
            "qty": qty,
            "unit": mapping_cell_text(cell("unit")),
            "supplier": supplier,
            "buy_price": buy_price,
            "damaged_qty": damaged,
            "added_qty": added,
            "reduced_qty": reduced,
            "missing_qty": missing,
            "actual_qty": actual_qty,
            "source_amount": source_amount,
            "calculated_amount": calculated_amount,
            "amount": amount,
            "note": mapping_cell_text(cell("note")),
            "status": status,
            "errors": errors,
            "warnings": warnings,
        }
        rows.append(item)
        if apply and not errors:
            items.append(item)

    if scanned >= MAPPING_IMPORT_MAX_ROWS and worksheet.max_row > header_row + MAPPING_IMPORT_MAX_ROWS:
        raise ValueError(f"File vượt quá giới hạn {MAPPING_IMPORT_MAX_ROWS:,} dòng dữ liệu")
    if not rows:
        raise ValueError("Bảng công nợ không có dòng dữ liệu")
    counts = {
        "total": len(rows),
        "ready": len(items),
        "skipped": sum(item["status"] == "skip" for item in rows),
        "errors": sum(bool(item["errors"]) for item in rows),
        "warnings": sum(bool(item["warnings"]) for item in rows),
        "suppliers": len({item["supplier"] for item in items}),
    }
    return {
        "sheet": worksheet.title,
        "header_row": header_row,
        "rows": rows[:200],
        "issues": [item for item in rows if item["errors"] or item["warnings"]][:1000],
        "items": items,
        "counts": counts,
        "totals": {
            "actual_qty": sum(item["actual_qty"] for item in items),
            "amount": sum(item["amount"] for item in items),
        },
        "preview_truncated": len(rows) > 200,
        "can_confirm": counts["errors"] == 0 and counts["ready"] > 0,
    }


def net_received(order) -> float:
    return max(
        as_number(order["actual_received"])
        - as_number(order.get("damaged_qty") if isinstance(order, dict) else order["damaged_qty"])
        - as_number(order.get("supplier_return_qty") if isinstance(order, dict) else order["supplier_return_qty"]),
        0,
    )


def net_delivered(order) -> float:
    return max(
        as_number(order["actual_delivered"])
        - as_number(order.get("customer_return_qty") if isinstance(order, dict) else order["customer_return_qty"]),
        0,
    )


def inventory_rows(conn, as_of="", include_zero=False):
    if as_of:
        # Posted movements are chronological, but an active reservation is a
        # global claim on stock regardless of its requested delivery date.  If
        # future reservations were hidden here, a later back-dated draft could
        # reserve the same units a second time.
        query = """
            SELECT p.code product_code,p.name product_name,p.unit,
                   COALESCE(SUM(CASE WHEN t.status='posted' AND t.txn_date<=:as_of
                                     THEN t.qty_in-t.qty_out ELSE 0 END),0) accounting_qty,
                   COALESCE(SUM(CASE WHEN (t.status='posted' AND t.txn_date<=:as_of)
                                          OR t.status='reserved'
                                     THEN t.qty_in-t.qty_out ELSE 0 END),0) available_qty,
                   COALESCE(SUM(CASE WHEN t.status='posted' AND t.txn_date<=:as_of
                                     THEN t.qty_in*t.unit_cost ELSE 0 END),0) input_value,
                   COALESCE(SUM(CASE WHEN t.status='posted' AND t.txn_date<=:as_of
                                     THEN t.qty_out*t.unit_cost ELSE 0 END),0) output_value
            FROM products p LEFT JOIN inventory_transactions t
                 ON t.product_code=p.code AND t.source_type!='BK_INPUT'
            GROUP BY p.code,p.name,p.unit
            ORDER BY p.name
        """
        rows = [dict(row) for row in conn.execute(query, {"as_of": as_of})]
    else:
        query = """
        SELECT p.code product_code,p.name product_name,p.unit,
               COALESCE(SUM(CASE WHEN t.status='posted' THEN t.qty_in-t.qty_out ELSE 0 END),0) accounting_qty,
               COALESCE(SUM(CASE WHEN t.status IN ('posted','reserved') THEN t.qty_in-t.qty_out ELSE 0 END),0) available_qty,
               COALESCE(SUM(CASE WHEN t.status='posted' THEN t.qty_in*t.unit_cost ELSE 0 END),0) input_value,
               COALESCE(SUM(CASE WHEN t.status='posted' THEN t.qty_out*t.unit_cost ELSE 0 END),0) output_value
        FROM products p LEFT JOIN inventory_transactions t
             ON t.product_code=p.code AND t.source_type!='BK_INPUT'
        GROUP BY p.code,p.name,p.unit
        ORDER BY p.name
        """
        rows = [dict(row) for row in conn.execute(query)]
    if not include_zero:
        rows = [row for row in rows if abs(row["accounting_qty"]) > 1e-9 or abs(row["available_qty"]) > 1e-9]
    return rows


def inventory_lookup(conn, as_of=""):
    return {row["product_code"]: row for row in inventory_rows(conn, as_of, include_zero=True)}


SUPPLIER_MERGE_ALL = {"dung", "thu", "tan", "phuong", "ky", "kho"}
SUPPLIER_POLICY_MANAGED = SUPPLIER_MERGE_ALL | {"hoai"}
SUPPLIER_ORDER_IMAGE_COLUMNS = (
    {"key": "kitchen", "label": "Mã bếp"},
    {"key": "work_date", "label": "Ngày"},
    {"key": "product_name", "label": "Tên hàng"},
    {"key": "order_qty", "label": "Số lượng"},
    {"key": "unit", "label": "ĐVT"},
    {"key": "supplier", "label": "NCC"},
    {"key": "note", "label": "Ghi chú"},
)
SUPPLIER_ORDER_IMAGE_FORBIDDEN_TEXT = (
    "Tồn tủ", "Giá mua", "Thành tiền", "Tổng cần mua sau trừ tồn",
    "Mã dòng hệ thống", "Số lượng thực tế", "Hỏng", "Thêm", "Giảm", "Thiếu",
)


def supplier_merge_key(supplier: str) -> str:
    """Normalize a supplier identity for the customer's confirmed merge rules."""
    supplier_key = mapping_key(supplier)
    if supplier_key.startswith("nha") and len(supplier_key) > 3:
        supplier_key = supplier_key[3:]
    return supplier_key


def supplier_merge_policy(supplier: str, product_name: str) -> tuple[bool, bool]:
    """Return (combine kitchens, merge identical lines) for confirmed NCC rules."""
    supplier_key = supplier_merge_key(supplier)
    if supplier_key in SUPPLIER_MERGE_ALL:
        return True, True
    if supplier_key == "hoai":
        is_carrot = mapping_key(product_name) == "carot"
        return is_carrot, is_carrot
    return False, False


def is_internal_stock_supplier(supplier: str) -> bool:
    return supplier_merge_key(supplier) == "kho"


def supplier_order_status_map(conn, batch_id: int) -> dict[str, dict]:
    return {
        row["supplier_key"]: dict(row)
        for row in conn.execute(
            """SELECT supplier_key,supplier_label,status,revision,ordered_at,reopened_at,updated_at
               FROM supplier_order_statuses WHERE batch_id=? ORDER BY supplier_key""",
            (batch_id,),
        )
    }


def purchase_order_payload(conn, batch_id: int) -> dict:
    """Return the editable NCC-order plan without touching invoice inventory.

    Physical leftovers in the freezer/cabinet are an operational purchasing
    decision entered by the user.  They are intentionally separate from the
    accounting/invoice inventory used to gate outgoing invoices.
    """
    batch = conn.execute("SELECT * FROM batches WHERE id=?", (batch_id,)).fetchone()
    if not batch:
        raise ValueError("Không tìm thấy phiên đơn")
    rows = []
    ordered_total = 0.0
    required_total = 0.0
    physical_total = 0.0
    amount_total = 0
    canonical_sources = [dict(row) for row in conn.execute(
        "SELECT * FROM purchase_workbook_lines WHERE batch_id=? ORDER BY source_row,id",
        (batch_id,),
    )]
    if canonical_sources:
        for item in canonical_sources:
            base_qty = max(as_number(item["base_qty"]), 0)
            actual_qty = max(as_number(item["actual_qty"]), 0)
            row = {
                "row_key": item["row_key"], "order_id": item["order_id"],
                "source_row": item["source_row"],
                "line_kind": item.get("line_kind", "goods"),
                "product_code": item["product_code"], "product_name": item["product_name"],
                "kitchen": item["kitchen"], "work_date": item["work_date"],
                "unit": item["unit"], "demand_qty": base_qty, "customer_qty": base_qty,
                "physical_stock_used": 0, "order_qty": actual_qty, "required_qty": actual_qty,
                "supplier": item["supplier"], "buy_price": max(as_number(item["buy_price"]), 0),
                "price_source": item["price_source"] or "Sheet đặt hàng chuẩn",
                "amount": as_number(item["amount"]),
                "note": item["note"], "damaged_qty": as_number(item["damaged_qty"]),
                "added_qty": as_number(item["added_qty"]),
                "reduced_qty": as_number(item["reduced_qty"]),
                "missing_qty": as_number(item["missing_qty"]),
                "confirmed": item["status"] == "confirmed",
            }
            rows.append(row)
            ordered_total += base_qty
            required_total += actual_qty
            amount_total += row["amount"]
    else:
        for source in conn.execute(
            """SELECT o.*,p.physical_stock_used,p.order_qty,p.supplier plan_supplier,
                      p.buy_price plan_buy_price,p.price_source plan_price_source,
                      p.note plan_note,p.status plan_status
               FROM orders o LEFT JOIN purchase_order_lines p ON p.order_id=o.id
               WHERE o.batch_id=? ORDER BY o.id""",
            (batch_id,),
        ):
            item = dict(source)
            demand_qty = max(as_number(item["qty"]), 0)
            has_plan = item["order_qty"] is not None
            order_qty = max(as_number(item["order_qty"] if has_plan else demand_qty), 0)
            physical_stock_used = max(as_number(item["physical_stock_used"] if has_plan else 0), 0)
            supplier = mapping_cell_text(item["plan_supplier"] if has_plan else item["supplier"])
            buy_price = max(as_number(item["plan_buy_price"] if has_plan else item["buy_price"]), 0)
            note = mapping_cell_text(item["plan_note"] if has_plan else item["note"])
            row = {
                "order_id": item["id"], "product_code": item["product_code"],
                "product_name": item["product_name"], "kitchen": item["kitchen"],
                "work_date": item["work_date"], "unit": item["unit"],
                "demand_qty": demand_qty, "customer_qty": demand_qty,
                "physical_stock_used": physical_stock_used,
                "order_qty": order_qty, "required_qty": order_qty,
                "supplier": supplier, "buy_price": buy_price,
                "price_source": mapping_cell_text(item["plan_price_source"]) if has_plan else "Bảng báo giá",
                "amount": vnd_product(order_qty, buy_price), "note": note,
                "confirmed": has_plan and item["plan_status"] == "confirmed",
            }
            rows.append(row)
            ordered_total += demand_qty
            required_total += order_qty
            physical_total += physical_stock_used
            amount_total += row["amount"]

    manual_rules = {}
    for rule in conn.execute(
        "SELECT supplier_code,combine_kitchens FROM supplier_rules ORDER BY updated_at,supplier_code"
    ):
        manual_rules[supplier_merge_key(rule["supplier_code"])] = bool(rule["combine_kitchens"])

    order_statuses = supplier_order_status_map(conn, batch_id)
    groups = {}
    for item in rows:
        if item["order_qty"] <= 1e-9:
            continue
        supplier_key = supplier_merge_key(item["supplier"])
        forced_combined, merge_identical = supplier_merge_policy(
            item["supplier"], item["product_name"]
        )
        policy_managed = supplier_key in SUPPLIER_POLICY_MANAGED
        manual_combined = bool(manual_rules.get(supplier_key)) if not policy_managed else False
        combined = forced_combined or manual_combined
        key = f"{supplier_key}|{'ALL' if combined else item['kitchen']}"
        if supplier_key == "hoai":
            policy_label = "Tự dồn Cà rốt" if forced_combined else "Giữ riêng theo rule Hoài"
        elif supplier_key in SUPPLIER_MERGE_ALL:
            policy_label = "Tự dồn tên hàng giống nhau"
        else:
            policy_label = ""
        group = groups.setdefault(key, {
            "supplier": item["supplier"],
            "kitchen": "GỘP NHIỀU BẾP" if combined else item["kitchen"],
            "combine_kitchens": combined, "items": [], "total_qty": 0,
            "total_amount": 0, "raw_line_count": 0,
            "policy_managed": policy_managed,
            "automatic_merge": forced_combined,
            "manual_combine": manual_combined,
            "policy_label": policy_label,
        })
        group["raw_line_count"] += 1
        merge_key = (
            mapping_key(item["product_name"]), mapping_key(item["unit"])
        ) if merge_identical else ("order", item.get("row_key") or item["order_id"])
        merged = next(
            (line for line in group["items"] if line.get("_merge_key") == merge_key), None
        )
        if merged is None:
            merged = {
                **item, "_merge_key": merge_key,
                "order_ids": [item["order_id"]] if item.get("order_id") is not None else [],
                "source_refs": [{
                    "row_key": item.get("row_key"), "order_id": item.get("order_id"),
                    "source_row": item.get("source_row"),
                }],
                "_kitchens": [item["kitchen"]],
            }
            group["items"].append(merged)
        else:
            if abs(merged["buy_price"] - item["buy_price"]) > 1e-9:
                merged["mixed_buy_prices"] = True
            for field in (
                "customer_qty", "demand_qty", "physical_stock_used",
                "order_qty", "required_qty", "amount",
            ):
                merged[field] += item[field]
            if item.get("order_id") is not None:
                merged["order_ids"].append(item["order_id"])
            merged["source_refs"].append({
                "row_key": item.get("row_key"), "order_id": item.get("order_id"),
                "source_row": item.get("source_row"),
            })
            if item["kitchen"] not in merged["_kitchens"]:
                merged["_kitchens"].append(item["kitchen"])
            if item["note"] and item["note"] not in merged["note"]:
                merged["note"] = " | ".join(filter(None, (merged["note"], item["note"])))
        group["total_qty"] += item["order_qty"]
        group["total_amount"] += item["amount"]
    for group in groups.values():
        group["presented_line_count"] = len(group["items"])
        status_row = order_statuses.get(supplier_merge_key(group["supplier"]))
        group["supplier_key"] = supplier_merge_key(group["supplier"])
        group["order_status"] = status_row["status"] if status_row else "pending"
        group["order_status_revision"] = int(status_row["revision"]) if status_row else 0
        group["order_status_updated_at"] = status_row["updated_at"] if status_row else ""
        group["ordered_at"] = status_row["ordered_at"] if status_row else None
        group["reopened_at"] = status_row["reopened_at"] if status_row else None
        for item in group["items"]:
            item.pop("_merge_key", None)
            item["kitchen"] = ", ".join(item.pop("_kitchens"))
    group_list = list(groups.values())
    status_priority = {"reopened": 0, "pending": 1, "ordered": 2}
    group_list.sort(key=lambda group: (
        status_priority[group["order_status"]],
        supplier_merge_key(group["supplier"]),
        mapping_key(group["kitchen"]),
    ))
    checklist_by_supplier = {}
    for group in group_list:
        item = checklist_by_supplier.setdefault(group["supplier_key"], {
            "supplier_key": group["supplier_key"],
            "supplier": group["supplier"],
            "status": group["order_status"],
            "revision": group["order_status_revision"],
            "updated_at": group["order_status_updated_at"],
            "ordered_at": group["ordered_at"],
            "reopened_at": group["reopened_at"],
            "group_count": 0, "raw_line_count": 0, "presented_line_count": 0,
            "total_qty": 0,
        })
        item["group_count"] += 1
        item["raw_line_count"] += group["raw_line_count"]
        item["presented_line_count"] += group["presented_line_count"]
        item["total_qty"] += group["total_qty"]
    checklist = list(checklist_by_supplier.values())
    checklist.sort(key=lambda item: (
        status_priority[item["status"]], mapping_key(item["supplier"]),
    ))
    checklist_counts = {
        status: sum(item["status"] == status for item in checklist)
        for status in ("pending", "reopened", "ordered")
    }
    return {
        "batch_id": batch_id, "work_date": batch["work_date"],
        "format": "customer_canonical" if canonical_sources else "legacy_or_customer_orders",
        "formula": "Số đặt NCC do người dùng chốt sau khi trừ tồn tủ thực tế",
        "ordered_qty": ordered_total, "required_qty": required_total,
        "physical_stock_used": physical_total, "total_amount": amount_total,
        "confirmed_lines": sum(1 for item in rows if item["confirmed"]),
        "image_contract": {
            "columns": list(SUPPLIER_ORDER_IMAGE_COLUMNS),
            "forbidden_text": list(SUPPLIER_ORDER_IMAGE_FORBIDDEN_TEXT),
        },
        "raw_line_count": len(rows),
        "presented_line_count": sum(len(group["items"]) for group in group_list),
        "checklist": checklist, "checklist_counts": checklist_counts,
        "rows": rows, "groups": group_list,
        "money_adjustments": [row for row in rows if row.get("line_kind") == DEDUCTION_KIND],
        "plan_hash": purchase_order_database_state_hash(conn, batch_id),
    }


def purchase_order_database_state_hash(conn, batch_id: int) -> str:
    return query_database_state_hash(conn, (
        (
            "orders",
            "SELECT id,batch_id,work_date,kitchen,product_code,product_name,qty,unit,supplier,buy_price,note "
            "FROM orders WHERE batch_id=? ORDER BY id",
            (batch_id,),
        ),
        (
            "purchase_order_lines",
            "SELECT order_id,demand_qty,physical_stock_used,order_qty,supplier,buy_price,price_source,note,status "
            "FROM purchase_order_lines WHERE batch_id=? ORDER BY order_id",
            (batch_id,),
        ),
        (
            "purchase_workbook_lines",
            "SELECT row_key,order_id,product_code,kitchen,work_date,product_name,base_qty,unit,"
            "supplier,note,buy_price,price_source,damaged_qty,added_qty,reduced_qty,missing_qty,actual_qty,"
            "amount,status,revision,line_kind FROM purchase_workbook_lines WHERE batch_id=? ORDER BY row_key",
            (batch_id,),
        ),
    ))


PURCHASE_CANONICAL_ALIASES = {
    "product_code": {"mahang", "mahanghoa", "mavattu"},
    "kitchen": {"mabep", "bep"},
    "work_date": {"ngay", "ngaythang", "ngaydat"},
    "product_name": {"tenhang", "tenhanghoa"},
    "base_qty": {"soluong"},
    "unit": {"dvt", "donvitinh"},
    "supplier": {"ncc", "nhacungcap"},
    "note": {"ghichu", "ghichudathang"},
    "buy_price": {"giamua", "dongiamua"},
    "damaged_qty": {"hong", "hanghong"},
    "added_qty": {"them", "hangthem"},
    "reduced_qty": {"giam", "hanggiam"},
    "missing_qty": {"thieu", "hangthieu"},
    "actual_qty": {"slthucte", "soluongthucte", "slthucte"},
    "amount": {"thanhtien"},
    "row_key": {"madonghethong", "rowkey", "khoadong"},
    "line_kind": {"loaidong"},
}


def purchase_canonical_header_fields(row) -> dict:
    try:
        from .customer_purchase_layout import customer_purchase_fields
    except ImportError:
        from customer_purchase_layout import customer_purchase_fields
    found = {}
    normalized_aliases = {
        field: {mapping_key(alias) for alias in aliases}
        for field, aliases in PURCHASE_CANONICAL_ALIASES.items()
    }
    for column_index, value in enumerate(row, 1):
        key = mapping_key(value)
        for field, aliases in normalized_aliases.items():
            if field not in found and key in aliases:
                found[field] = column_index
                break
    # The customer's canonical sheet intentionally leaves the visible date
    # header blank. It is the column between Mã bếp and Tên hàng.
    if "work_date" not in found and {
        "product_code", "kitchen", "product_name",
    }.issubset(found):
        inferred = found["kitchen"] + 1
        if inferred + 1 == found["product_name"]:
            found["work_date"] = inferred
    found.update(customer_purchase_fields(row))
    return found


def purchase_work_date(value, fallback: str) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = mapping_cell_text(value)
    if not text:
        return fallback
    for pattern in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, pattern).date().isoformat()
        except ValueError:
            pass
    raise ValueError("Ngày trong sheet đặt hàng không hợp lệ")


def purchase_business_row_key(
    work_date: str, product_code: str, kitchen: str, product_name: str,
    unit: str, occurrence: int,
) -> str:
    identity = "\0".join((
        "purchase_orders", work_date, mapping_key(product_code), mapping_key(kitchen),
        mapping_key(product_name), mapping_key(unit), str(int(occurrence)),
    ))
    return hashlib.sha256(identity.encode("utf-8")).hexdigest().upper()


def purchase_scope_hash(items) -> str:
    digest = hashlib.sha256()
    text_fields = (
        "row_key", "product_code", "kitchen", "work_date", "product_name",
        "unit", "supplier", "note", "price_source",
    )
    number_fields = (
        "base_qty", "buy_price", "damaged_qty", "added_qty", "reduced_qty",
        "missing_qty", "actual_qty", "amount",
    )
    for item in sorted(items, key=lambda row: row["row_key"]):
        normalized = {field: mapping_cell_text(item.get(field)) for field in text_fields}
        if item.get("line_kind", "goods") != "goods":
            normalized["line_kind"] = item["line_kind"]
        normalized.update({field: as_number(item.get(field)) for field in number_fields})
        encoded = json.dumps(
            normalized,
            ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str,
        ).encode("utf-8")
        digest.update(encoded)
        digest.update(b"\n")
    return digest.hexdigest().upper()


def purchase_formula_key(value) -> str:
    if not isinstance(value, str) or not value.lstrip().startswith("="):
        return ""
    return re.sub(r"[\s$]+", "", value).upper()


def expected_purchase_formulas(fields: dict, source_row: int) -> tuple[str, set[str]]:
    ref = lambda field: f"{get_column_letter(fields[field])}{source_row}"
    actual_ref = ref("actual_qty")
    price_ref = ref("buy_price")
    actual_formula = (
        f"={ref('base_qty')}+{ref('added_qty')}-{ref('damaged_qty')}"
        f"-{ref('reduced_qty')}-{ref('missing_qty')}"
    )
    amount_product = f"{actual_ref}*{price_ref}"
    return purchase_formula_key(actual_formula), {
        purchase_formula_key(f"={amount_product}"),
        purchase_formula_key(f"={price_ref}*{actual_ref}"),
        purchase_formula_key(f"=IFERROR({amount_product},0)"),
        purchase_formula_key(f"=IFERROR({price_ref}*{actual_ref},0)"),
    }


def purchase_line_changed(previous: dict | None, item: dict) -> bool:
    if previous is None:
        return True
    if previous.get("line_kind", "goods") != item.get("line_kind", "goods"):
        return True
    text_fields = (
        "product_code", "kitchen", "work_date", "product_name", "unit",
        "supplier", "note", "price_source",
    )
    number_fields = (
        "base_qty", "buy_price", "damaged_qty", "added_qty", "reduced_qty",
        "missing_qty", "actual_qty", "amount",
    )
    if any(mapping_cell_text(previous.get(field)) != mapping_cell_text(item.get(field))
           for field in text_fields):
        return True
    if any(abs(as_number(previous.get(field)) - as_number(item.get(field))) > 1e-9
           for field in number_fields):
        return True
    return int(previous.get("order_id") or 0) != int(item.get("order_id") or 0)


def purchase_component_totals(items) -> dict:
    fields = (
        "base_qty", "damaged_qty", "added_qty", "reduced_qty", "missing_qty",
        "actual_qty", "amount",
    )
    return {
        field: vnd_round(sum(as_number(item.get(field)) for item in items))
        if field == "amount" else sum(as_number(item.get(field)) for item in items)
        for field in fields
    }


class PurchaseOrderApplyError(ValueError):
    def __init__(self, message: str, *, code: str = "purchase_apply_failed", status: int = 409):
        super().__init__(message)
        self.code = code
        self.status = status


def refresh_payable_ledger(conn, timestamp: str) -> dict:
    """Refresh the small dedicated payable projection inside this transaction."""
    try:
        from payable_ledger import sync_payable_ledger
    except ImportError:  # pragma: no cover - package invocation
        from .payable_ledger import sync_payable_ledger
    return sync_payable_ledger(conn, timestamp=timestamp)


def apply_purchase_order_preview(
    conn, *, batch_id: int, items: list[dict], source_hash: str,
    source_name: str, format_name: str, now_iso,
) -> dict:
    """Apply one validated purchase preview inside the caller's transaction.

    This is shared by the dedicated NCC round-trip and the second daily
    finalization flow so both entry points keep identical revision/payable
    semantics without ever updating customer orders.
    """
    try:
        from .batch_bk_approval import mutation_blocker
    except ImportError:
        from batch_bk_approval import mutation_blocker
    blocked = mutation_blocker(conn, batch_id)
    if blocked:
        raise PurchaseOrderApplyError(blocked, code='batch_bk_posted')
    already = conn.execute(
        "SELECT 1 FROM purchase_order_imports WHERE batch_id=? AND source_hash=?",
        (batch_id, source_hash),
    ).fetchone()
    if already:
        current_items = [dict(row) for row in conn.execute(
            "SELECT * FROM purchase_workbook_lines WHERE batch_id=?", (batch_id,),
        )]
        current_hash = purchase_scope_hash(current_items) if current_items else ""
        if current_hash != source_hash:
            raise PurchaseOrderApplyError(
                "Bản đặt hàng này đã bị phiên bản mới hơn thay thế; hãy tải file mới nhất",
                code="superseded_purchase_version",
            )
        current_by_key = {row['row_key']: row for row in current_items}
        # Prices/quantities can be identical while a newly imported sales row
        # supplies a previously missing link. Persist that link once.
        already = not any(purchase_line_changed(current_by_key.get(item['row_key']), item) for item in items)
    if already:
        result = purchase_order_payload(conn, batch_id)
        audit(
            conn, now_iso, "purchase_orders.import", "unchanged",
            entity_type="batch", entity_id=batch_id,
            metadata={"rows": len(items), "source_hash": source_hash[:16]},
        )
        payable_ledger = refresh_payable_ledger(conn, now_iso())
        return {
            "processed": 0, "count": len(items), "idempotent": True,
            "purchase_order": result, "payable_ledger": payable_ledger,
        }

    timestamp = now_iso()
    incoming_keys = {item["row_key"] for item in items}
    existing_by_key = {
        row["row_key"]: dict(row) for row in conn.execute(
            "SELECT * FROM purchase_workbook_lines WHERE batch_id=?", (batch_id,),
        )
    }
    changed_count = 0
    inserted_count = 0
    updated_count = 0
    for item in items:
        if item.get("errors"):
            raise PurchaseOrderApplyError("Dòng mua còn lỗi; chưa được ghi")
        previous = existing_by_key.get(item["row_key"])
        if not purchase_line_changed(previous, item):
            continue
        revision = int(previous["revision"] or 0) + 1 if previous else 1
        change_kind = "update" if previous else "insert"
        conn.execute(
            """INSERT INTO purchase_workbook_lines(
                   batch_id,row_key,order_id,source_sheet,source_row,product_code,kitchen,
                   work_date,product_name,base_qty,unit,supplier,note,buy_price,price_source,damaged_qty,
                   added_qty,reduced_qty,missing_qty,actual_qty,amount,status,source_hash,
                   revision,created_at,updated_at,line_kind
               ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,'confirmed',?,?,?,?,?)
               ON CONFLICT(batch_id,row_key) DO UPDATE SET
                   order_id=excluded.order_id,source_sheet=excluded.source_sheet,
                   source_row=excluded.source_row,product_code=excluded.product_code,
                   kitchen=excluded.kitchen,work_date=excluded.work_date,
                   product_name=excluded.product_name,base_qty=excluded.base_qty,
                   unit=excluded.unit,supplier=excluded.supplier,note=excluded.note,
                   buy_price=excluded.buy_price,price_source=excluded.price_source,
                   damaged_qty=excluded.damaged_qty,
                   added_qty=excluded.added_qty,reduced_qty=excluded.reduced_qty,
                   missing_qty=excluded.missing_qty,actual_qty=excluded.actual_qty,
                   amount=excluded.amount,status='confirmed',source_hash=excluded.source_hash,
                   revision=excluded.revision,updated_at=excluded.updated_at,line_kind=excluded.line_kind""",
            (
                batch_id, item["row_key"], item.get("order_id"), item["source_sheet"],
                item["source_row"], item["product_code"], item["kitchen"], item["work_date"],
                item["product_name"], item["base_qty"], item["unit"], item["supplier"],
                item["note"], item["buy_price"], item["price_source"], item["damaged_qty"],
                item["added_qty"], item["reduced_qty"], item["missing_qty"],
                item["actual_qty"], item["amount"], source_hash, revision,
                previous["created_at"] if previous else timestamp, timestamp,
                item.get("line_kind", "goods"),
            ),
        )
        conn.execute(
            """INSERT INTO purchase_workbook_line_revisions(
                   batch_id,row_key,order_id,revision,change_kind,source_hash,source_row,
                   base_qty,damaged_qty,added_qty,reduced_qty,missing_qty,actual_qty,
                   buy_price,amount,created_at,line_kind
               ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (
                batch_id, item["row_key"], item.get("order_id"), revision,
                change_kind, source_hash, item["source_row"], item["base_qty"],
                item["damaged_qty"], item["added_qty"], item["reduced_qty"],
                item["missing_qty"], item["actual_qty"], item["buy_price"],
                item["amount"], timestamp,
                item.get("line_kind", "goods"),
            ),
        )
        changed_count += 1
        inserted_count += int(previous is None)
        updated_count += int(previous is not None)

    stale_keys = [row["row_key"] for row in conn.execute(
        "SELECT row_key FROM purchase_workbook_lines WHERE batch_id=?", (batch_id,),
    ) if row["row_key"] not in incoming_keys]
    if stale_keys:
        placeholders = ",".join("?" for _ in stale_keys)
        conn.execute(
            f"DELETE FROM purchase_workbook_lines WHERE batch_id=? AND row_key IN ({placeholders})",
            (batch_id, *stale_keys),
        )
    conn.execute("DELETE FROM purchase_order_lines WHERE batch_id=?", (batch_id,))
    if format_name == "legacy_13":
        for item in items:
            conn.execute(
                """INSERT INTO purchase_order_lines(
                       batch_id,order_id,demand_qty,physical_stock_used,order_qty,
                       supplier,buy_price,price_source,note,status,source_hash,created_at,updated_at
                   ) VALUES(?,?,?,?,?,?,?,?,?,'confirmed',?,?,?)""",
                (
                    batch_id, item["order_id"], item["demand_qty"],
                    item["physical_stock_used"], item["order_qty"], item["supplier"],
                    item["buy_price"], item["price_source"], item["note"],
                    source_hash, timestamp, timestamp,
                ),
            )
    total_amount = sum(item["amount"] for item in items)
    conn.execute(
        """INSERT INTO purchase_order_imports(
               batch_id,source_hash,source_name,line_count,total_amount,imported_at
           ) VALUES(?,?,?,?,?,?) ON CONFLICT(batch_id,source_hash) DO NOTHING""",
        (batch_id, source_hash, source_name, len(items), total_amount, timestamp),
    )
    audit(
        conn, now_iso, "purchase_orders.import", "ok",
        entity_type="batch", entity_id=batch_id,
        metadata={
            "rows": len(items), "changed": changed_count,
            "inserted": inserted_count, "updated": updated_count,
            "source_hash": source_hash[:16],
            "components": purchase_component_totals(items),
        },
    )
    payable_ledger = refresh_payable_ledger(conn, timestamp)
    return {
        "processed": changed_count, "count": len(items), "idempotent": False,
        "purchase_order": purchase_order_payload(conn, batch_id),
        "payable_ledger": payable_ledger,
    }


def parse_canonical_purchase_workbook(
    conn, workbook, batch_id: int, expected: dict, formula_workbook=None, *, pricing_orders=None,
) -> dict | None:
    batch = conn.execute("SELECT work_date FROM batches WHERE id=?", (batch_id,)).fetchone()
    if not batch:
        raise ValueError("Không tìm thấy phiên đơn")
    required = {
        "product_code", "kitchen", "product_name", "base_qty", "unit", "supplier",
        "note", "buy_price", "damaged_qty", "added_qty", "reduced_qty",
        "missing_qty", "actual_qty", "amount",
    }
    target = None
    for worksheet in workbook.worksheets:
        for row_index, values in enumerate(
            worksheet.iter_rows(min_row=1, max_row=min(20, worksheet.max_row), values_only=True), 1
        ):
            fields = purchase_canonical_header_fields(values)
            if required.issubset(fields):
                if target is not None:
                    raise ValueError("File có nhiều bảng đặt hàng chuẩn; chỉ được có một sheet đặt hàng")
                target = (worksheet, row_index, fields)
                break
    if target is None:
        return None

    worksheet, header_row, fields = target
    formula_worksheet = None
    if formula_workbook is not None and worksheet.title in formula_workbook.sheetnames:
        formula_worksheet = formula_workbook[worksheet.title]
    existing_by_key = {
        row["row_key"]: dict(row) for row in conn.execute(
            "SELECT * FROM purchase_workbook_lines WHERE batch_id=?", (batch_id,)
        )
    }
    orders_by_identity = defaultdict(list)
    for order in expected.values():
        identity = (mapping_key(order["product_code"]), mapping_key(order["kitchen"]))
        orders_by_identity[identity].append(order)
    for matches in orders_by_identity.values():
        matches.sort(key=lambda row: int(row["id"]))

    # A combined daily confirmation applies sales first. Resolve links against
    # those updated rows, but keep the pricing inputs used by the validated
    # purchase preview: writing sales must not silently reprice purchasing.
    pricing_by_identity = orders_by_identity
    if pricing_orders is not None:
        pricing_by_identity = defaultdict(list)
        for order in pricing_orders:
            identity = (mapping_key(order['product_code']), mapping_key(order['kitchen']))
            pricing_by_identity[identity].append(order)
        for matches in pricing_by_identity.values():
            matches.sort(key=lambda row: int(row['id']))

    parsed = []
    occurrences = Counter()
    seen_row_keys = set()
    seen_deductions = set()
    for source_row, values in enumerate(
        worksheet.iter_rows(min_row=header_row + 1, values_only=True), header_row + 1
    ):
        def cell(field):
            column = fields.get(field)
            return values[column - 1] if column and column <= len(values) else None

        product_code = mapping_cell_text(cell("product_code"))
        product_name = mapping_cell_text(cell("product_name"))
        if not product_code and not product_name:
            continue
        kitchen = mapping_cell_text(cell("kitchen"))
        unit = mapping_cell_text(cell("unit"))
        supplier = mapping_cell_text(cell("supplier"))
        note = mapping_cell_text(cell("note"))
        errors = []
        warnings = []
        try:
            work_date = purchase_work_date(cell("work_date"), batch["work_date"])
        except ValueError as exc:
            work_date = batch["work_date"]
            errors.append(str(exc))
        if work_date != batch["work_date"]:
            errors.append("Ngày dòng đặt hàng không khớp phiên đang chọn")
        numbers = {}
        for field, label in (
            ("base_qty", "Số lượng"), ("buy_price", "Giá mua"),
            ("damaged_qty", "Hỏng"), ("added_qty", "Thêm"),
            ("reduced_qty", "Giảm"), ("missing_qty", "Thiếu"),
        ):
            try:
                numbers[field] = import_cell_number(cell(field), label)
            except ValueError as exc:
                errors.append(str(exc))
                numbers[field] = 0
        explicit_kind = mapping_key(cell("line_kind"))
        explicit_deduction = explicit_kind in {mapping_key(DEDUCTION_KIND), mapping_key(DEDUCTION_LABEL)}
        if explicit_kind and explicit_kind not in {"goods", "hanghoa"} and not explicit_deduction:
            errors.append("Loại dòng không hợp lệ")
        approved_deduction = approved_phong_source(
            work_date, supplier, kitchen, product_code, product_name, unit, numbers,
        ) and not explicit_kind
        deduction = explicit_deduction or approved_deduction
        if deduction:
            deduction_identity = (mapping_key(supplier), mapping_key(kitchen), mapping_key(product_name))
            if deduction_identity in seen_deductions:
                errors.append("Khoản trừ tiền mua hộ bị lặp trong file")
            seen_deductions.add(deduction_identity)
        if explicit_deduction and any(numbers.values()):
            errors.append("Dòng trừ tiền mua hộ chỉ nhập thành tiền âm; số lượng, giá mua và các cột điều chỉnh phải bằng 0")
        identity = (mapping_key(product_code), mapping_key(kitchen))
        occurrences[identity] += 1
        occurrence = occurrences[identity]
        computed_key = purchase_business_row_key(
            work_date, product_code, kitchen, product_name, unit, occurrence,
        )
        supplied_key = mapping_cell_text(cell("row_key")).upper()
        if supplied_key and not re.fullmatch(r"[0-9A-F]{64}", supplied_key):
            errors.append("Mã dòng hệ thống không hợp lệ")
            supplied_key = ""
        row_key = supplied_key or computed_key
        if row_key in seen_row_keys:
            errors.append("Mã dòng hệ thống bị lặp trong file")
        seen_row_keys.add(row_key)
        previous = existing_by_key.get(row_key)
        if supplied_key and previous:
            immutable_now = tuple(mapping_key(value) for value in (
                product_code, kitchen, product_name, unit,
            ))
            immutable_before = tuple(mapping_key(previous[field]) for field in (
                "product_code", "kitchen", "product_name", "unit",
            ))
            if immutable_now != immutable_before:
                errors.append("Mã hàng/bếp/tên/ĐVT đã bị sửa so với file xuất")
        elif supplied_key and supplied_key != computed_key:
            errors.append("Mã dòng không thuộc sheet đặt hàng đang chọn")
        matches = orders_by_identity.get(identity) or []
        matched_order = matches[occurrence - 1] if occurrence <= len(matches) else None
        order_id = int(matched_order["id"]) if matched_order and not deduction else None
        sheet_buy_price = numbers["buy_price"]
        quoted_buy_price = 0
        price_matches = pricing_by_identity.get(identity) or []
        pricing_order = price_matches[occurrence - 1] if occurrence <= len(price_matches) else None
        if previous and previous.get('price_source') == 'Bảng báo giá' and not deduction:
            # A confirmed purchase owns its quoted cost. Reimporting sales must
            # not rewrite that quote or make an identical purchase file drift.
            quoted_buy_price = max(as_number(previous.get('buy_price')), 0)
        elif pricing_order and not deduction and not previous:
            quoted_buy_price = max(as_number(pricing_order.get("buy_price")), 0)
        if quoted_buy_price > 0:
            numbers["buy_price"] = quoted_buy_price
            price_source = "Bảng báo giá"
        else:
            price_source = "Sheet đặt hàng chuẩn"
        if not kitchen:
            errors.append("Dòng đặt hàng thiếu mã bếp")
        if not product_code and not deduction:
            errors.append("Dòng đặt hàng thiếu mã hàng")
        if not product_name:
            errors.append("Dòng đặt hàng thiếu tên hàng")
        if not deduction and (numbers["base_qty"] < 0 or any(
            numbers[field] < 0 for field in ("damaged_qty", "added_qty", "reduced_qty", "missing_qty")
        )):
            errors.append("Số lượng và các cột điều chỉnh không được âm")
        computed_actual = (
            numbers["base_qty"] + numbers["added_qty"] - numbers["damaged_qty"]
            - numbers["reduced_qty"] - numbers["missing_qty"]
        )
        expected_actual_formula, expected_amount_formulas = expected_purchase_formulas(
            fields, source_row,
        )
        formula_actual = (
            formula_worksheet.cell(source_row, fields["actual_qty"]).value
            if formula_worksheet is not None else cell("actual_qty")
        )
        actual_formula_key = purchase_formula_key(formula_actual)
        # Subtracting missing before reduced is the same verified formula used
        # by the customer's workbook. Compare signed cell terms, never eval Excel.
        def signed_terms(formula):
            if not re.fullmatch(r'=[A-Z]+[0-9]+(?:[+-][A-Z]+[0-9]+)*', formula):
                return formula
            return sorted(re.findall(r'[+-][A-Z]+[0-9]+', '+' + formula[1:]))
        if actual_formula_key and signed_terms(actual_formula_key) != signed_terms(expected_actual_formula):
            errors.append(
                "Công thức Số lượng thực tế phải là Số lượng + Thêm - Hỏng - Giảm - Thiếu"
            )
        try:
            cached_actual = import_cell_number(cell("actual_qty"), "Số lượng thực tế", computed_actual)
        except ValueError as exc:
            errors.append(str(exc))
            cached_actual = computed_actual
        if cell("actual_qty") not in (None, "") and abs(cached_actual - computed_actual) > 0.001:
            errors.append("Số lượng thực tế lệch công thức chuẩn")
        actual_qty = computed_actual
        formula_amount = (
            formula_worksheet.cell(source_row, fields["amount"]).value
            if formula_worksheet is not None else cell("amount")
        )
        amount_formula_key = purchase_formula_key(formula_amount)
        if explicit_deduction and amount_formula_key:
            errors.append("Thành tiền mua hộ phải nhập số tiền trực tiếp")
        elif amount_formula_key and amount_formula_key not in expected_amount_formulas:
            errors.append("Công thức Thành tiền phải là Số lượng thực tế × Giá mua")
        source_formula_amount = vnd_product(actual_qty, sheet_buy_price)
        try:
            cached_amount = import_cell_number(
                cell("amount"), "Thành tiền", source_formula_amount,
            )
        except ValueError as exc:
            errors.append(str(exc))
            cached_amount = source_formula_amount
        if not explicit_deduction and cell("amount") not in (None, "") and vnd_round(cached_amount) != source_formula_amount:
            errors.append("Thành tiền lệch Số lượng thực tế × Giá mua")
        amount = vnd_product(actual_qty, numbers["buy_price"])
        if actual_qty < -1e-9 and not deduction:
            errors.append("Số lượng thực tế không được âm")
        if numbers["buy_price"] < 0:
            errors.append("Giá mua không được âm")
        if actual_qty > 1e-9 and not supplier:
            errors.append("Dòng có đặt hàng phải có NCC")
        if actual_qty > 1e-9 and numbers["buy_price"] <= 0 and not is_internal_stock_supplier(supplier):
            errors.append("Dòng mua ngoài có số lượng dương phải có giá mua lớn hơn 0")
        if deduction:
            if not supplier or is_internal_stock_supplier(supplier):
                errors.append("Khoản trừ tiền mua hộ phải có nhà cung cấp bên ngoài")
            amount = vnd_round(cached_amount)
            if cell("amount") in (None, "") or amount >= 0:
                errors.append("Khoản trừ tiền mua hộ phải có thành tiền âm")
            # This line has a monetary effect only. Preserve its source row and
            # business identity, without fabricating a stock code or order link.
            numbers = {field: 0 for field in numbers}
            actual_qty = 0
            price_source = DEDUCTION_LABEL
            if not note.startswith(DEDUCTION_LABEL):
                note = DEDUCTION_LABEL + (" · " + note if note else "")
        parsed.append({
            "source_sheet": worksheet.title, "source_row": source_row,
            "format": "customer_canonical", "row_key": row_key, "order_id": order_id,
            "_issue_columns": fields,
            "product_code": product_code, "kitchen": kitchen, "work_date": work_date,
            "product_name": product_name, "base_qty": numbers["base_qty"], "unit": unit,
            "supplier": supplier, "note": note, "buy_price": numbers["buy_price"],
            "price_source": price_source,
            "damaged_qty": numbers["damaged_qty"], "added_qty": numbers["added_qty"],
            "reduced_qty": numbers["reduced_qty"], "missing_qty": numbers["missing_qty"],
            "actual_qty": actual_qty, "amount": amount, "errors": errors, "warnings": warnings,
            "line_kind": DEDUCTION_KIND if deduction else "goods",
        })
    if not parsed:
        raise ValueError("Sheet đặt hàng chuẩn không có dòng dữ liệu")
    missing_existing = set(existing_by_key) - seen_row_keys
    if missing_existing:
        raise ValueError(
            f"File thiếu {len(missing_existing)} dòng mua đã có; hãy tải lại sheet đặt hàng mới nhất"
        )
    error_rows = sum(bool(item["errors"]) for item in parsed)
    return {
        "format": "customer_canonical", "sheet": worksheet.title,
        "items": parsed, "rows": parsed[:200], "count": len(parsed),
        "issues": [item for item in parsed if item["errors"] or item["warnings"]][:200],
        "preview_truncated": len(parsed) > 200,
        "error_rows": error_rows, "can_confirm": error_rows == 0,
        "warning_rows": sum(bool(item["warnings"]) for item in parsed),
        "total_qty": sum(item["actual_qty"] for item in parsed if not item["errors"]),
        "total_amount": sum(item["amount"] for item in parsed if not item["errors"]),
        "content_hash": purchase_scope_hash(parsed),
    }


PURCHASE_ORDER_ALIASES = {
    "order_id": {"madonghethong", "madong", "orderid", "dong"},
    "product_code": {"mahang", "mahanghoa", "mavattu"},
    "kitchen": {"mabep", "bep"},
    "work_date": {"ngay", "ngaythang", "ngaydat"},
    "product_name": {"tenhang", "tenhanghoa"},
    "demand_qty": {"nhucautudonkhach", "nhucau", "soluongkhachdat"},
    "physical_stock_used": {"tontudatr u", "tontudatru", "tontu", "tonthuctedatru"},
    "order_qty": {"soluongdatncc", "soluongdat", "soluongthucdat"},
    "unit": {"dvt", "donvitinh"},
    "supplier": {"ncc", "nhacungcap"},
    "buy_price": {"giamua", "dongiamua"},
    "note": {"ghichu", "ghichudathang"},
}


def purchase_order_header_fields(row) -> dict:
    found = {}
    for column_index, value in enumerate(row, 1):
        key = mapping_key(value)
        for field, aliases in PURCHASE_ORDER_ALIASES.items():
            if field not in found and key in {mapping_key(alias) for alias in aliases}:
                found[field] = column_index
                break
    return found


def parse_purchase_order_workbook(conn, workbook, batch_id: int, formula_workbook=None, *, pricing_orders=None) -> dict:
    expected = {
        row["id"]: dict(row) for row in conn.execute(
            """SELECT o.*,p.price_source plan_price_source
               FROM orders o LEFT JOIN purchase_order_lines p ON p.order_id=o.id
               WHERE o.batch_id=? ORDER BY o.id""", (batch_id,)
        )
    }
    if not expected:
        raise ValueError("Phiên đơn không có dòng đặt hàng")
    canonical = parse_canonical_purchase_workbook(
        conn, workbook, batch_id, expected, formula_workbook=formula_workbook,
        pricing_orders=pricing_orders,
    )
    if canonical is not None:
        return canonical
    parsed = []
    seen = set()
    legacy_occurrences = Counter()
    for worksheet in workbook.worksheets:
        header_row = None
        fields = {}
        for row_index, values in enumerate(
            worksheet.iter_rows(min_row=1, max_row=min(20, worksheet.max_row), values_only=True), 1
        ):
            candidate = purchase_order_header_fields(values)
            if {"order_id", "order_qty", "buy_price", "supplier"}.issubset(candidate):
                header_row, fields = row_index, candidate
                break
        if not header_row:
            continue
        for source_row, values in enumerate(
            worksheet.iter_rows(min_row=header_row + 1, values_only=True), header_row + 1
        ):
            def cell(field):
                column = fields.get(field)
                return values[column - 1] if column and column <= len(values) else None

            raw_id = cell("order_id")
            if raw_id in (None, ""):
                continue
            errors = []
            try:
                order_id = int(import_cell_number(raw_id, "Mã dòng hệ thống"))
            except ValueError as exc:
                parsed.append({"source_sheet": worksheet.title, "source_row": source_row,
                               "order_id": 0, "errors": [str(exc)]})
                continue
            order = expected.get(order_id)
            if not order:
                errors.append("Mã dòng không thuộc phiên đơn đang chọn")
            if order_id in seen:
                errors.append("Mã dòng bị lặp trong file")
            seen.add(order_id)
            try:
                order_qty = import_cell_number(cell("order_qty"), "Số lượng đặt NCC")
            except ValueError as exc:
                errors.append(str(exc)); order_qty = 0
            try:
                buy_price = import_cell_number(cell("buy_price"), "Giá mua")
            except ValueError as exc:
                errors.append(str(exc)); buy_price = 0
            try:
                physical = import_cell_number(cell("physical_stock_used"), "Tồn tủ đã trừ")
            except ValueError as exc:
                errors.append(str(exc)); physical = 0
            supplier = mapping_cell_text(cell("supplier"))
            prior_source = mapping_cell_text(order["plan_price_source"]) if order else ""
            quoted_buy_price = (
                max(as_number(order["buy_price"]), 0)
                if order and prior_source != "File đặt hàng lần hai" else 0
            )
            effective_buy_price = quoted_buy_price if quoted_buy_price > 0 else max(buy_price, 0)
            if order_qty < 0:
                errors.append("Số lượng đặt NCC không được âm")
            if physical < 0:
                errors.append("Tồn tủ đã trừ không được âm")
            if order_qty > 0 and not supplier:
                errors.append("Dòng có đặt hàng phải có NCC")
            if (
                order_qty > 0
                and effective_buy_price <= 0
                and not is_internal_stock_supplier(supplier)
            ):
                errors.append("Dòng có đặt hàng phải có giá mua lớn hơn 0")
            if order:
                file_code = mapping_cell_text(cell("product_code")).upper()
                file_kitchen = mapping_cell_text(cell("kitchen")).upper()
                if file_code and file_code != mapping_cell_text(order["product_code"]).upper():
                    errors.append("Mã hàng đã bị sửa so với file xuất")
                if file_kitchen and file_kitchen != mapping_cell_text(order["kitchen"]).upper():
                    errors.append("Mã bếp đã bị sửa so với file xuất")
                demand_qty = max(as_number(order["qty"]), 0)
                # Quantity ordered is authoritative.  Derive the amount taken
                # from the physical cabinet, while still allowing whole-carton
                # purchases larger than demand.
                physical = max(demand_qty - order_qty, 0)
            else:
                demand_qty = 0
            identity = (
                mapping_key(order["product_code"] if order else ""),
                mapping_key(order["kitchen"] if order else ""),
            )
            legacy_occurrences[identity] += 1
            row_key = purchase_business_row_key(
                order["work_date"] if order else "1900-01-01",
                order["product_code"] if order else "",
                order["kitchen"] if order else "",
                order["product_name"] if order else "",
                order["unit"] if order else "",
                legacy_occurrences[identity],
            )
            parsed.append({
                "source_sheet": worksheet.title, "source_row": source_row,
                "format": "legacy_13", "row_key": row_key,
                "_issue_columns": fields,
                "order_id": order_id, "demand_qty": demand_qty,
                "physical_stock_used": physical, "order_qty": max(order_qty, 0),
                "supplier": supplier, "buy_price": effective_buy_price,
                "price_source": "Bảng báo giá" if quoted_buy_price > 0 else "File đặt hàng lần hai",
                "note": mapping_cell_text(cell("note")),
                "product_code": mapping_cell_text(order["product_code"] if order else ""),
                "kitchen": mapping_cell_text(order["kitchen"] if order else ""),
                "work_date": mapping_cell_text(order["work_date"] if order else ""),
                "product_name": mapping_cell_text(order["product_name"] if order else ""),
                "base_qty": max(order_qty, 0),
                "unit": mapping_cell_text(order["unit"] if order else ""),
                "damaged_qty": 0, "added_qty": 0, "reduced_qty": 0, "missing_qty": 0,
                "actual_qty": max(order_qty, 0),
                "amount": vnd_product(max(order_qty, 0), effective_buy_price),
                "errors": errors,
            })
    missing = sorted(set(expected) - seen)
    if missing:
        raise ValueError(
            f"File thiếu {len(missing)} dòng của phiên đơn; hãy tải lại file đặt NCC mới nhất rồi chỉnh"
        )
    if not parsed:
        raise ValueError("Không tìm thấy bảng đặt NCC đúng mẫu hệ thống")
    error_rows = sum(bool(item["errors"]) for item in parsed)
    return {
        "format": "legacy_13",
        "items": parsed, "rows": parsed[:200], "count": len(parsed),
        "error_rows": error_rows, "can_confirm": error_rows == 0,
        "warning_rows": 0,
        "total_qty": sum(item["order_qty"] for item in parsed if not item["errors"]),
        "total_amount": sum(
            vnd_product(item["order_qty"], item["buy_price"])
            for item in parsed if not item["errors"]
        ),
        "content_hash": purchase_scope_hash(parsed),
    }


def outgoing_invoice_readiness_payload(conn, batch_id: int) -> dict:
    return batch_readiness_payload(conn, batch_id)


def create_partial_outgoing_drafts(conn, batch_id: int, now_iso, *, contractor_filter: str = "") -> dict:
    """Reserve every quantity currently invoiceable and retain the remainder."""
    batch = conn.execute("SELECT * FROM batches WHERE id=?", (batch_id,)).fetchone()
    if not batch:
        raise ValueError("Không tìm thấy phiên đơn")
    if batch["status"] != "approved":
        raise ValueError("Phải duyệt phiên đơn trước khi lập hóa đơn đầu ra")
    selected_orders = [dict(r) for r in conn.execute(
        'SELECT * FROM orders WHERE batch_id=? AND (? = \'\' OR contractor=?) ORDER BY contractor,id',
        (batch_id, contractor_filter, contractor_filter))]
    issues = invoice_order_issues(selected_orders)
    if issues:
        details = '; '.join(f"Dòng {r['order_id']} · {r['product_code']}: {', '.join(r['messages'])}" for r in issues[:5])
        raise OutgoingReadinessError(f"Còn {len(issues)} dòng đơn cần sửa trước khi tạo file. {details}", code='invalid_invoice_orders')
    orders = [row for row in selected_orders if net_delivered(row) > 1e-9]
    if not orders:
        raise ValueError("Phiên không còn lượng thực giao dương để lập hóa đơn")
    validate_demand_orders(conn, orders)

    replaceable = [dict(row) for row in conn.execute(
        """SELECT * FROM outgoing_invoice_drafts
           WHERE batch_id=? AND status='draft' AND (? = '' OR contractor=?)
             AND COALESCE(draft_kind,'standard')='standard'
             AND COALESCE(minvoice_status,'not_sent') NOT IN ('saved','saving','unknown')
           ORDER BY id""",
        (batch_id, contractor_filter, contractor_filter),
    )]
    replaceable_ids = {row["id"] for row in replaceable}
    # Only canonical invoice inventory may unlock an outgoing draft.  Active
    # reservations and locally issued invoices awaiting source sync are already
    # deducted by this projection.
    stock = canonical_available_stock(conn)
    raw_available = {code: as_number(row["raw_available_qty"]) for code, row in stock.items()}
    released = defaultdict(float)
    if replaceable_ids:
        placeholders = ",".join("?" for _ in replaceable_ids)
        for reserved in conn.execute(
            f"""SELECT product_code,SUM(qty_out) qty FROM inventory_transactions
                 WHERE source_type='OUTGOING_DRAFT' AND status='reserved'
                   AND source_id IN ({placeholders}) GROUP BY product_code""",
            tuple(str(value) for value in sorted(replaceable_ids)),
        ):
            released[reserved["product_code"]] += as_number(reserved["qty"])
    available = {
        code: max(raw_available.get(code, 0) + released.get(code, 0), 0)
        for code in set(raw_available) | set(released)
    }
    try:
        from .stock_tax_policy import exempt_order_codes
    except ImportError:
        from stock_tax_policy import exempt_order_codes
    exempt = exempt_order_codes(conn, orders)
    demand_codes = {item["product_code"] for item in orders} - exempt
    unresolved_holds = {
        code: -(raw_available.get(code, 0) + released.get(code, 0))
        for code in demand_codes
        if raw_available.get(code, 0) + released.get(code, 0) < -1e-9
    }
    if unresolved_holds:
        details = ", ".join(
            f"{code}: {qty:g}" for code, qty in sorted(unresolved_holds.items())[:10]
        )
        raise OutgoingReadinessError(
            "Tồn hóa đơn chuẩn đang thấp hơn phần đã khóa/đã phát hành chưa đồng bộ; "
            f"cần đối chiếu trước khi lập tiếp ({details})",
            code="canonical_stock_overcommitted",
        )

    allocated_locked = {
        row["order_id"]: as_number(row["qty"])
        for row in conn.execute(
            """SELECT l.order_id,SUM(l.qty) qty
               FROM outgoing_order_allocations l JOIN outgoing_invoice_drafts d ON d.id=l.draft_id
               JOIN orders o ON o.id=l.order_id
               WHERE o.batch_id=? AND d.status='draft'
                 AND (
                     COALESCE(d.draft_kind,'standard')!='standard'
                     OR COALESCE(d.minvoice_status,'not_sent') IN ('saved','saving','unknown')
                 )
               GROUP BY l.order_id""",
            (batch_id,),
        )
    }
    try:
        from .outgoing_unissued import issued_allocations
    except ImportError:
        from outgoing_unissued import issued_allocations
    issued, source_warnings = issued_allocations(conn)
    selected_parties={item['contractor'] for item in orders}
    source_warnings=[w for w in source_warnings if not w['contractor'] or w['contractor'] in selected_parties]
    if source_warnings:
        raise OutgoingReadinessError(
            'Cần đối chiếu hóa đơn đã phát hành trước khi lập tiếp để tránh xuất trùng. '+source_warnings[0]['message'],
            code='issued_source_unresolved',
        )
    for item in orders:
        allocated_locked[item['id']] = allocated_locked.get(item['id'], 0) + issued.get(item['id'], 0)
    allocations = defaultdict(list)
    pending_qty = 0.0
    for item in orders:
        demand = net_delivered(item)
        locked = allocated_locked.get(item["id"], 0)
        if locked > demand + 1e-9:
            raise OutgoingReadinessError(
                f"Dòng đơn {item['id']} đã phân bổ vượt số thực giao; cần đối chiếu trước khi lập tiếp",
                code="allocated_over_demand",
            )
        remaining = max(demand - locked, 0)
        have = available.get(item["product_code"], 0)
        qty = remaining if item['product_code'] in exempt else min(remaining, have)
        available[item["product_code"]] = max(have - qty, 0)
        if qty > 1e-9:
            # Each tax workbook is imported and issued separately. Keep its
            # number, buyer snapshot and stock hold on a separate local draft.
            allocations[(item["contractor"], invoice_tax_percent(item["tax"]))].append((item, qty))
        pending_qty += max(remaining - qty, 0)

    reusable_by_tax = {}
    for draft in replaceable:
        taxes = {invoice_tax_percent(row['tax']) for row in conn.execute(
            'SELECT tax FROM outgoing_invoice_lines WHERE draft_id=?', (draft['id'],))}
        if len(taxes) != 1:
            # Old mixed-tax drafts are cancelled below, never silently issued
            # against one invoice number for several upload files.
            continue
        key = (draft['contractor'], next(iter(taxes)))
        current = reusable_by_tax.get(key)
        if current is None or draft["id"] > current["id"]:
            reusable_by_tax[key] = draft
    used_ids = set()
    created = []
    for (contractor, tax_group), lines in allocations.items():
        subtotal = tax_amount = 0
        calculated = []
        for item, qty in lines:
            unit_price = vnd_round(as_number(item["sell_price"]))
            amount = vnd_product(qty, unit_price)
            vat_percent = invoice_tax_percent(item["tax"])
            line_tax = 0 if vat_percent <= 0 else vnd_product(amount, vat_percent / 100)
            calculated.append((item, qty, unit_price, amount))
            subtotal += amount
            tax_amount += line_tax
        total = subtotal + tax_amount
        draft = reusable_by_tax.get((contractor, tax_group))
        if draft:
            draft_id = draft["id"]
            round_no = draft.get("round_no") or 1
            used_ids.add(draft_id)
            conn.execute(
                """UPDATE outgoing_invoice_drafts SET invoice_date=?,status='draft',
                       subtotal=?,tax_amount=?,total_amount=?,minvoice_status='not_sent',
                       minvoice_series=NULL,minvoice_remote_id=NULL,minvoice_saved_at=NULL,
                       minvoice_error=NULL,minvoice_key_api=NULL,minvoice_started_at=NULL,
                       minvoice_reconciled_at=NULL
                   WHERE id=?""",
                (batch["work_date"], subtotal, tax_amount, total, draft_id),
            )
        else:
            round_no = conn.execute(
                "SELECT COALESCE(MAX(round_no),0)+1 n FROM outgoing_invoice_drafts WHERE batch_id=? AND contractor=?",
                (batch_id, contractor),
            ).fetchone()["n"]
            cursor = conn.execute(
                """INSERT INTO outgoing_invoice_drafts(
                       batch_id,contractor,invoice_date,status,subtotal,tax_amount,total_amount,
                       created_at,external_key_uuid,round_no
                   ) VALUES(?,?,?,'draft',?,?,?,?,?,?)""",
                (
                    batch_id, contractor, batch["work_date"], subtotal, tax_amount, total,
                    now_iso(), uuid.uuid4().hex.upper(), round_no,
                ),
            )
            draft_id = cursor.lastrowid
            used_ids.add(draft_id)
        conn.execute("DELETE FROM outgoing_invoice_lines WHERE draft_id=?", (draft_id,))
        conn.execute(
            """UPDATE inventory_transactions SET status='cancelled',updated_at=?
               WHERE source_type='OUTGOING_DRAFT' AND source_id=? AND status='reserved'""",
            (now_iso(), str(draft_id)),
        )
        for item, qty, unit_price, amount in calculated:
            invoice_name_row = conn.execute(
                "SELECT invoice_name FROM outgoing_product_names WHERE product_code=?",
                (item["product_code"],),
            ).fetchone()
            invoice_name = invoice_name_row["invoice_name"] if invoice_name_row else item["product_name"]
            conn.execute(
                """INSERT INTO outgoing_invoice_lines(
                       draft_id,order_id,product_code,product_name,qty,unit,unit_price,tax,
                       invoice_nature,amount
                   ) VALUES(?,?,?,?,?,?,?,?,?,?)""",
                (
                    draft_id, item["id"], item["product_code"], invoice_name, qty,
                    item["unit"], unit_price, item["tax"], item.get("invoice_nature") or "1", amount,
                ),
            )
            conn.execute(
                """INSERT INTO inventory_transactions(
                       txn_date,product_code,qty_in,qty_out,unit_cost,source_type,source_id,source_line,
                       kitchen,status,note,created_at,updated_at
                   ) VALUES(?,?,0,?,?, 'OUTGOING_DRAFT',?,?,?,'reserved',?,?,?)
                   ON CONFLICT(source_type,source_id,source_line) DO UPDATE SET
                       txn_date=excluded.txn_date,product_code=excluded.product_code,
                       qty_out=excluded.qty_out,unit_cost=excluded.unit_cost,
                       kitchen=excluded.kitchen,status='reserved',note=excluded.note,
                       updated_at=excluded.updated_at""",
                (
                    batch["work_date"], item["product_code"], qty, item["buy_price"],
                    str(draft_id), str(item["id"]), item["kitchen"],
                    f"Dự thảo hóa đơn {contractor} lần {round_no}", now_iso(), now_iso(),
                ),
            )
        created.append(dict(conn.execute(
            "SELECT * FROM outgoing_invoice_drafts WHERE id=?", (draft_id,)
        ).fetchone()))

    for draft in replaceable:
        if draft["id"] in used_ids:
            continue
        conn.execute("UPDATE outgoing_invoice_drafts SET status='cancelled' WHERE id=?", (draft["id"],))
        conn.execute(
            """UPDATE inventory_transactions SET status='cancelled',updated_at=?
               WHERE source_type='OUTGOING_DRAFT' AND source_id=? AND status='reserved'""",
            (now_iso(), str(draft["id"])),
        )
    audit(
        conn, now_iso, "outgoing.draft", "ok", entity_type="batch", entity_id=batch_id,
        metadata={"drafts": len(created), "pending_qty": pending_qty, "partial": pending_qty > 1e-9},
    )
    return {
        "drafts": created, "pending_qty": pending_qty,
        "partial": pending_qty > 1e-9,
        "requires_user_sign_and_issue": True,
    }


def post_purchase_list_inventory(conn, batch_id: int, now_iso):
    """Purge the retired automatic BK projection; never create stock from a flag.

    Column ``bk`` only identifies a candidate purchase-list row. Approving a
    sales batch must not turn that marker directly into stock: the official BK
    workbook still has to pass preview and explicit confirmation through
    ``bk_import``. Keeping this compatibility function lets old approval paths
    remove a stale automatic projection atomically. Historical ``BK_INPUT``
    rows are also excluded by :func:`inventory_rows`.
    """
    removed = conn.execute(
        "DELETE FROM inventory_transactions WHERE source_type='BK_INPUT' AND source_id=?",
        (str(batch_id),),
    ).rowcount
    if removed:
        audit(
            conn, now_iso, "inventory.bk_legacy_purged", "ok",
            entity_type="batch", entity_id=batch_id,
            metadata={"lines": removed, "replacement_flow": "bk_import_confirm"},
        )
    return 0


def invoice_details(remote: dict) -> list[dict]:
    details = first_value(remote, "hdhhdvu", "invoiceItems", "details", default=[])
    return details if isinstance(details, list) else []


def normalize_invoice(remote: dict, invoice_type: str, now: str) -> dict:
    remote_id = str(first_value(remote, "_id", "id", "invoiceId", default="")).strip()
    if not remote_id:
        raise MsmiError("Hóa đơn mSMI thiếu khóa _id")
    subtotal = msmi_number(
        first_value(remote, "tgtcthue", "subtotal", "totalBeforeTax"), "Tiền trước thuế",
    )
    tax_amount = msmi_number(
        first_value(remote, "tgtthue", "taxAmount", "totalTax"), "Tiền thuế",
    )
    total = msmi_number(
        first_value(remote, "tgtttbso", "tgtttbchu", "totalAmount", "total"), "Tổng tiền",
    )
    if subtotal < 0 or tax_amount < 0 or total < 0:
        raise MsmiError("Hóa đơn mSMI có tổng tiền âm; cần đối chiếu thủ công")
    if total <= 0:
        total = subtotal + tax_amount
    invoice_date = as_date(first_value(remote, "tdlap", "nlap", "invoiceDate", "signedDate"))
    if not invoice_date:
        raise MsmiError("Hóa đơn mSMI thiếu ngày lập hợp lệ")
    try:
        raw_json = json.dumps(remote, ensure_ascii=False, separators=(",", ":"), allow_nan=False)
    except (TypeError, ValueError):
        raise MsmiError("Hóa đơn mSMI chứa dữ liệu JSON không hợp lệ") from None
    return {
        "remote_id": remote_id,
        "invoice_type": invoice_type,
        "seller_tax_code": str(first_value(remote, "mstNban", "nbmst", "sellerTaxCode", "sellerTaxId")),
        "seller_name": str(first_value(remote, "tenNban", "nbten", "sellerName")),
        "invoice_number": str(first_value(remote, "shdon", "soHoaDon", "invoiceNumber")),
        "invoice_series": str(first_value(remote, "khhdon", "khmshdon", "series", "invoiceSeries")),
        "invoice_date": invoice_date,
        "subtotal": subtotal,
        "tax_amount": tax_amount,
        "total_amount": total,
        "raw_json": raw_json,
        "synced_at": now,
    }


def normalize_invoice_item(remote_item: dict, line_index: int) -> dict:
    qty = msmi_number(first_value(remote_item, "sluong", "quantity", "qty"), f"Số lượng dòng {line_index}")
    unit_price = msmi_number(first_value(remote_item, "dgia", "unitPrice", "price"), f"Đơn giá dòng {line_index}")
    amount = msmi_number(first_value(remote_item, "thtien", "amount", "total"), f"Thành tiền dòng {line_index}")
    nature = str(first_value(remote_item, "tchat", "nature", "itemNature", "type")).strip()
    # mSMI represents invoice-level discounts/financial adjustments as a
    # zero-quantity, zero-unit-price line with a negative amount.  That is a
    # legitimate accounting line, not a stock movement, so it must not block
    # the positive goods lines on the same invoice.  Negative quantity/unit
    # price, or a negative amount attached to a physical quantity, remains
    # unsafe and is quarantined for manual review.
    financial_adjustment = amount < 0 and qty == 0 and unit_price == 0
    if qty < 0 or unit_price < 0 or (amount < 0 and not financial_adjustment):
        raise MsmiError(
            f"Dòng {line_index} mSMI có số lượng, đơn giá hoặc thành tiền âm; cần đối chiếu thủ công"
        )
    try:
        from .invoice_input_integrity import is_goods_line
    except ImportError:
        from invoice_input_integrity import is_goods_line
    inventory_eligible = not financial_adjustment and is_goods_line(qty, unit_price, amount, nature)
    if financial_adjustment:
        validation_note = (
            "Không ghi kho: dòng chiết khấu/điều chỉnh tài chính âm, được giữ nguyên để đối chiếu tổng hóa đơn"
        )
    elif inventory_eligible:
        validation_note = ""
    else:
        validation_note = (
            "Không ghi kho: dòng nguồn không có số lượng/đơn giá dương; vẫn giữ nguyên để đối chiếu hóa đơn"
        )
    return {
        "line_index": line_index,
        "source_item_code": str(first_value(remote_item, "ma", "mhhhoa", "itemCode", "code")),
        "source_item_name": str(first_value(remote_item, "ten", "tenhh", "itemName", "name")),
        "source_unit": str(first_value(remote_item, "dvtinh", "dvt", "unit")),
        "qty": qty,
        "unit_price": unit_price,
        "amount": amount,
        "tax_rate": str(first_value(remote_item, "tsuat", "taxRate", "tax")),
        "source_nature": nature,
        "inventory_eligible": 1 if inventory_eligible else 0,
        "validation_note": validation_note,
    }


def quarantine_msmi_invoice(conn, remote: dict, invoice_type: str, tenant: str, now: str, error: Exception):
    """Persist an identifiable bad source invoice for review without inventing accounting values."""
    remote_id = str(first_value(remote, "_id", "id", "invoiceId", default="")).strip()
    if not remote_id:
        raise MsmiError("Hóa đơn mSMI lỗi dữ liệu và thiếu khóa _id nên không thể cách ly an toàn")
    try:
        raw_json = json.dumps(remote, ensure_ascii=False, separators=(",", ":"), allow_nan=False)
    except (TypeError, ValueError):
        raise MsmiError("Hóa đơn mSMI lỗi dữ liệu chứa JSON không hợp lệ nên không thể cách ly an toàn") from None

    def safe_nonnegative(*keys):
        try:
            value = msmi_number(first_value(remote, *keys), "Giá trị hóa đơn lỗi")
            return value if value >= 0 else 0
        except MsmiError:
            return 0

    existing = conn.execute(
        "SELECT id,receipt_status FROM msmi_invoices WHERE remote_id=?", (remote_id,)
    ).fetchone()
    message = f"Dữ liệu nguồn cần đối chiếu thủ công: {str(error)[:220]}"
    if existing and existing["receipt_status"] == "posted":
        conn.execute(
            """UPDATE msmi_invoices SET sync_status='review_required',error_message=?,
                      synced_at=?,updated_at=? WHERE id=?""",
            (message, now, now, existing["id"]),
        )
        return existing["id"], False, {
            "remote_id": remote_id,
            "invoice_date": "",
        }

    invoice_date = as_date(first_value(remote, "tdlap", "nlap", "invoiceDate", "signedDate"))
    conn.execute(
        """INSERT INTO msmi_invoices(
               remote_id,tenant,invoice_type,seller_tax_code,seller_name,invoice_number,invoice_series,
               invoice_date,subtotal,tax_amount,total_amount,sync_status,receipt_status,raw_json,
               error_message,synced_at,created_at,updated_at
           ) VALUES(?,?,?,?,?,?,?,?,?,?,?,'review_required','blocked',?,?,?,?,?)
           ON CONFLICT(remote_id) DO UPDATE SET
               tenant=excluded.tenant,invoice_type=excluded.invoice_type,
               seller_tax_code=excluded.seller_tax_code,seller_name=excluded.seller_name,
               invoice_number=excluded.invoice_number,invoice_series=excluded.invoice_series,
               invoice_date=excluded.invoice_date,subtotal=excluded.subtotal,
               tax_amount=excluded.tax_amount,total_amount=excluded.total_amount,
               sync_status='review_required',receipt_status='blocked',raw_json=excluded.raw_json,
               error_message=excluded.error_message,synced_at=excluded.synced_at,updated_at=excluded.updated_at""",
        (
            remote_id, tenant, invoice_type,
            str(first_value(remote, "mstNban", "nbmst", "sellerTaxCode", "sellerTaxId")),
            str(first_value(remote, "tenNban", "nbten", "sellerName")),
            str(first_value(remote, "shdon", "soHoaDon", "invoiceNumber")),
            str(first_value(remote, "khhdon", "khmshdon", "series", "invoiceSeries")),
            invoice_date,
            safe_nonnegative("tgtcthue", "subtotal", "totalBeforeTax"),
            safe_nonnegative("tgtthue", "taxAmount", "totalTax"),
            safe_nonnegative("tgtttbso", "tgtttbchu", "totalAmount", "total"),
            raw_json, message, now, now, now,
        ),
    )
    invoice_id = conn.execute(
        "SELECT id FROM msmi_invoices WHERE remote_id=?", (remote_id,)
    ).fetchone()["id"]
    conn.execute("DELETE FROM msmi_invoice_items WHERE invoice_id=?", (invoice_id,))
    return invoice_id, existing is None, {
        "remote_id": remote_id,
        "invoice_date": invoice_date,
    }


def saved_mapping(conn, tenant: str, seller_tax_code: str, item: dict):
    return conn.execute(
        """SELECT product_code FROM item_mappings
           WHERE tenant=? AND seller_tax_code=? AND source_item_code=? AND source_item_name=?""",
        (tenant, seller_tax_code, item["source_item_code"], item["source_item_name"]),
    ).fetchone()


def product_name_candidates(conn) -> dict[str, list[dict]]:
    """Index canonical and approved invoice names without hiding ambiguous aliases."""
    grouped = defaultdict(list)
    for row in conn.execute(
        """SELECT p.code,p.name,COALESCE(o.invoice_name,'') invoice_name
           FROM products p
           LEFT JOIN outgoing_product_names o ON o.product_code=p.code
           WHERE p.code IS NOT NULL AND p.name IS NOT NULL
           ORDER BY p.code"""
    ):
        candidate = {
            "code": row["code"],
            "name": row["name"],
            "invoice_name": row["invoice_name"],
        }
        for value in (row["name"], row["invoice_name"]):
            key = mapping_key(value)
            if key and all(item["code"] != row["code"] for item in grouped[key]):
                grouped[key].append(candidate)
    return grouped


def unique_product_name_suggestions(conn) -> dict[str, dict]:
    """Return only deterministic name matches; ambiguous names are never suggested."""
    grouped = product_name_candidates(conn)
    return {key: rows[0] for key, rows in grouped.items() if len(rows) == 1}


def safe_input_mapping_suggestion_plan(conn, *, tenant: str, date_from: str, date_to: str) -> dict:
    """Preview exact-name, same-unit mappings without changing invoice or stock data."""
    try:
        from invoice_mapping import mapping_scope_key, mapping_units_match
        from invoice_workbench import validate_date_range
    except ImportError:  # pragma: no cover - package invocation
        from .invoice_mapping import mapping_scope_key, mapping_units_match
        from .invoice_workbench import validate_date_range

    start, end = validate_date_range(date_from, date_to)
    safe_tenant = str(tenant or "TDP").strip() or "TDP"
    suggestions = unique_product_name_suggestions(conn)
    products = {
        str(row["code"]): dict(row)
        for row in conn.execute("SELECT code,name,COALESCE(unit,'') unit FROM products")
    }
    rows = conn.execute(
        """SELECT li.id,li.invoice_id,li.source_item_code,li.source_item_name,li.source_unit,
                  i.seller_tax_code
             FROM msmi_invoice_items li
             JOIN msmi_invoices i ON i.id=li.invoice_id
            WHERE i.tenant=? AND i.invoice_type='INPUT_ELECTRONIC_INVOICE'
              AND i.invoice_date>=? AND i.invoice_date<=? AND i.sync_status='synced'
              AND i.receipt_status NOT IN ('posted','blocked')
              AND li.inventory_eligible=1 AND li.mapping_status!='mapped'
            ORDER BY li.id""",
        (safe_tenant, start, end),
    ).fetchall()

    representatives = {}
    safe_line_ids = set()
    exact_name_unit_review = 0
    no_unique_name_match = 0
    for row in rows:
        suggestion = suggestions.get(mapping_key(row["source_item_name"]))
        product = products.get(str(suggestion["code"])) if suggestion else None
        if not product:
            no_unique_name_match += 1
            continue
        if not mapping_units_match(row["source_unit"], product["unit"]):
            exact_name_unit_review += 1
            continue
        scope = mapping_scope_key(row["source_item_code"], row["source_item_name"], row["source_unit"])
        key = (str(row["seller_tax_code"] or ""), scope)
        representatives.setdefault(key, {"item_id": int(row["id"]), "product_code": str(product["code"])})
        safe_line_ids.add(int(row["id"]))

    # save_mapping deliberately remembers one confirmed rule for all still-editable
    # invoices of the same supplier/source identity. Show that full impact before POST.
    affected_line_ids = set()
    if representatives:
        for row in conn.execute(
            """SELECT li.id,li.source_item_code,li.source_item_name,li.source_unit,i.seller_tax_code
                 FROM msmi_invoice_items li
                 JOIN msmi_invoices i ON i.id=li.invoice_id
                WHERE i.tenant=? AND i.invoice_type='INPUT_ELECTRONIC_INVOICE'
                  AND i.sync_status='synced' AND i.receipt_status NOT IN ('posted','blocked')
                  AND li.inventory_eligible=1 AND li.mapping_status!='mapped'""",
            (safe_tenant,),
        ):
            scope = mapping_scope_key(row["source_item_code"], row["source_item_name"], row["source_unit"])
            if (str(row["seller_tax_code"] or ""), scope) in representatives:
                affected_line_ids.add(int(row["id"]))

    remaining_by_invoice = defaultdict(set)
    for row in rows:
        remaining_by_invoice[int(row["invoice_id"])].add(int(row["id"]))
    ready_after = sum(bool(ids) and ids.issubset(safe_line_ids) for ids in remaining_by_invoice.values())
    snapshot_data = {
        "tenant": safe_tenant,
        "date_from": start,
        "date_to": end,
        "rules": sorted((partner, scope, item["product_code"]) for (partner, scope), item in representatives.items()),
        "selected_line_ids": sorted(safe_line_ids),
        "affected_line_ids": sorted(affected_line_ids),
    }
    snapshot = hashlib.sha256(
        json.dumps(snapshot_data, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    ).hexdigest()
    return {
        "tenant": safe_tenant,
        "date_from": start,
        "date_to": end,
        "safe_lines_in_period": len(safe_line_ids),
        "safe_rules": len(representatives),
        "invoices_with_safe_lines": len({int(row["invoice_id"]) for row in rows if int(row["id"]) in safe_line_ids}),
        "invoices_ready_after": ready_after,
        "affected_lines_all_periods": len(affected_line_ids),
        "exact_name_needing_unit_review": exact_name_unit_review,
        "lines_without_unique_name": no_unique_name_match,
        "snapshot": snapshot,
        "_representatives": list(representatives.values()),
    }


def upsert_msmi_invoice(conn, remote: dict, invoice_type: str, tenant: str, now: str) -> tuple[int, bool]:
    data = normalize_invoice(remote, invoice_type, now)
    existing = conn.execute("SELECT * FROM msmi_invoices WHERE remote_id=?", (data["remote_id"],)).fetchone()
    detail_rows = invoice_details(remote)
    normalized_items = [
        normalize_invoice_item(raw_item if isinstance(raw_item, dict) else {}, index)
        for index, raw_item in enumerate(detail_rows, start=1)
    ]

    def business_signature(header, items):
        header_fields = {
            "tenant": str(header.get("tenant") or tenant),
            "invoice_type": str(header.get("invoice_type") or invoice_type),
            "seller_tax_code": str(header.get("seller_tax_code") or ""),
            "seller_name": str(header.get("seller_name") or ""),
            "invoice_number": str(header.get("invoice_number") or ""),
            "invoice_series": str(header.get("invoice_series") or ""),
            "invoice_date": str(header.get("invoice_date") or ""),
            "subtotal": round(as_number(header.get("subtotal")), 6),
            "tax_amount": round(as_number(header.get("tax_amount")), 6),
            "total_amount": round(as_number(header.get("total_amount")), 6),
        }
        item_fields = [{
            "line_index": int(item.get("line_index") or 0),
            "source_item_code": str(item.get("source_item_code") or ""),
            "source_item_name": str(item.get("source_item_name") or ""),
            "source_unit": str(item.get("source_unit") or ""),
            "qty": round(as_number(item.get("qty")), 6),
            "unit_price": round(as_number(item.get("unit_price")), 6),
            "amount": round(as_number(item.get("amount")), 6),
            "tax_rate": str(item.get("tax_rate") or ""),
            "source_nature": str(item.get("source_nature") or ""),
        } for item in items]
        return json.dumps({"header": header_fields, "items": item_fields}, ensure_ascii=False, sort_keys=True)

    if existing and existing["receipt_status"] == "posted":
        stored_items = [dict(row) for row in conn.execute(
            "SELECT * FROM msmi_invoice_items WHERE invoice_id=? ORDER BY line_index", (existing["id"],),
        )]
        changed = business_signature(dict(existing), stored_items) != business_signature(data, normalized_items)
        if changed:
            message = "Hóa đơn mSMI đã thay đổi sau khi tạo phiếu nhập; kho được giữ nguyên và cần đối chiếu thủ công"
            conn.execute(
                """UPDATE msmi_invoices SET sync_status='review_required',error_message=?,
                   synced_at=?,updated_at=? WHERE id=?""",
                (message, now, now, existing["id"]),
            )
            conn.execute(
                """INSERT INTO audit_log(event_type,entity_type,entity_id,status,message,metadata_json,created_at)
                   VALUES('msmi.remote_change_after_receipt','msmi_invoice',?,'warning',?,'{}',?)""",
                (data["remote_id"], message, now),
            )
        else:
            conn.execute(
                """UPDATE msmi_invoices SET sync_status='synced',error_message=NULL,
                   synced_at=?,updated_at=? WHERE id=?""",
                (now, now, existing["id"]),
            )
        return existing["id"], False

    conn.execute(
        """INSERT INTO msmi_invoices(
            remote_id,tenant,invoice_type,seller_tax_code,seller_name,invoice_number,invoice_series,
            invoice_date,subtotal,tax_amount,total_amount,sync_status,receipt_status,raw_json,
            synced_at,created_at,updated_at
        ) VALUES(?,?,?,?,?,?,?,?,?,?,?,'synced','pending_mapping',?,?,?,?)
        ON CONFLICT(remote_id) DO UPDATE SET
            seller_tax_code=excluded.seller_tax_code,seller_name=excluded.seller_name,
            invoice_number=excluded.invoice_number,invoice_series=excluded.invoice_series,
            invoice_date=excluded.invoice_date,subtotal=excluded.subtotal,tax_amount=excluded.tax_amount,
            total_amount=excluded.total_amount,raw_json=excluded.raw_json,synced_at=excluded.synced_at,
            sync_status='synced',error_message=NULL,updated_at=excluded.updated_at""",
        (data["remote_id"], tenant, data["invoice_type"], data["seller_tax_code"], data["seller_name"],
         data["invoice_number"], data["invoice_series"], data["invoice_date"], data["subtotal"],
         data["tax_amount"], data["total_amount"], data["raw_json"], now, now, now),
    )
    invoice_id = conn.execute("SELECT id FROM msmi_invoices WHERE remote_id=?", (data["remote_id"],)).fetchone()["id"]
    # Unposted invoices are safe to rebuild from the remote source. Mappings are
    # recovered only by the exact seller + source code + source name key.
    conn.execute("DELETE FROM msmi_invoice_items WHERE invoice_id=?", (invoice_id,))
    mapped = 0
    inventory_items = 0
    for item in normalized_items:
        eligible = bool(item["inventory_eligible"])
        inventory_items += eligible
        mapping = saved_mapping(conn, tenant, data["seller_tax_code"], item) if eligible else None
        product_code = mapping["product_code"] if mapping else ""
        conn.execute(
            """INSERT INTO msmi_invoice_items(
                invoice_id,line_index,source_item_code,source_item_name,source_unit,qty,unit_price,
                amount,tax_rate,source_nature,inventory_eligible,validation_note,product_code,mapping_status
            ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (invoice_id, item["line_index"], item["source_item_code"], item["source_item_name"], item["source_unit"],
             item["qty"], item["unit_price"], item["amount"], item["tax_rate"], item["source_nature"],
             item["inventory_eligible"], item["validation_note"], product_code,
             "mapped" if product_code else ("unmapped" if eligible else "not_inventory")),
        )
        mapped += bool(product_code)
    if not inventory_items:
        receipt_status = "not_inventory"
    else:
        receipt_status = "ready" if mapped == inventory_items else "pending_mapping"
    conn.execute("UPDATE msmi_invoices SET receipt_status=? WHERE id=?", (receipt_status, invoice_id))
    try:
        from .invoice_expenses import restore_expenses
        from .invoice_mapping import _refresh_input_invoice
    except ImportError:
        from invoice_expenses import restore_expenses
        from invoice_mapping import _refresh_input_invoice
    restore_expenses(conn,invoice_id)
    _refresh_input_invoice(conn,invoice_id)
    return invoice_id, existing is None


def sync_msmi(
    conn,
    client,
    now_iso,
    tenant="default",
    max_pages=5,
    page_size=199,
    invoice_type="INPUT_ELECTRONIC_INVOICE",
    from_date="",
    to_date="",
):
    if invoice_type not in {"INPUT_ELECTRONIC_INVOICE", "OUTPUT_ELECTRONIC_INVOICE"}:
        raise MsmiError("Loại hóa đơn mSMI không hợp lệ")
    if bool(from_date) != bool(to_date):
        raise MsmiError("Cần truyền đủ từ ngày và đến ngày khi đồng bộ mSMI")
    new_count = 0
    seen_count = 0
    item_count = 0
    review_count = 0
    pages = 0
    progress_pages = 0
    state = conn.execute(
        "SELECT * FROM msmi_sync_state WHERE invoice_type=?", (invoice_type,)
    ).fetchone()
    newest_id = str(state["last_remote_id"] or "") if state else ""
    newest_date = str(state["last_invoice_date"] or "") if state else ""
    newest_captured = False
    backfill_complete = bool(state["backfill_complete"]) if state else False
    backfill_anchor_id = str(state["backfill_anchor_id"] or "") if state else ""
    backfill_anchor_date = str(state["backfill_anchor_date"] or "") if state else ""
    reconcile_next_page = max(int(state["reconcile_next_page"] or 1), 1) if state else 1
    page_limit = max(2, min(int(max_pages), 50))

    def consume(remote):
        nonlocal new_count, seen_count, item_count, review_count, newest_id, newest_date, newest_captured
        timestamp = now_iso()
        try:
            invoice_id, created = upsert_msmi_invoice(conn, remote, invoice_type, tenant, timestamp)
            normalized = normalize_invoice(remote, invoice_type, timestamp)
        except MsmiError as error:
            invoice_id, created, normalized = quarantine_msmi_invoice(
                conn, remote, invoice_type, tenant, timestamp, error,
            )
            review_count += 1
        if created:
            new_count += 1
        else:
            seen_count += 1
        item_count += conn.execute(
            "SELECT COUNT(*) n FROM msmi_invoice_items WHERE invoice_id=?", (invoice_id,)
        ).fetchone()["n"]
        if not newest_captured:
            newest_id = normalized["remote_id"]
            newest_date = normalized["invoice_date"]
            newest_captured = True
        return normalized, created

    # One sync is atomic: if a later page fails, invoices from earlier pages
    # are rolled back before the durable error state is recorded.
    savepoint = "msmi_sync_atomic"
    conn.execute(f"SAVEPOINT {savepoint}")
    try:
        if backfill_complete:
            # Always scan the newest page, then use the remaining page budget
            # as a rolling reconciliation cursor through older history.  A
            # back-dated/late invoice can therefore never hide forever behind
            # one page containing only already-known IDs.
            page_sequence = [0] + list(
                range(reconcile_next_page, reconcile_next_page + page_limit - 1)
            )
            next_reconcile_page = reconcile_next_page
            for page in page_sequence:
                request_args = {"invoice_type": invoice_type, "page": page, "size": page_size}
                if from_date:
                    request_args.update({"from_date": from_date, "to_date": to_date})
                result = client.list_invoices(**request_args)
                pages += 1
                remote_items = result["items"]
                if not remote_items:
                    if page > 0:
                        next_reconcile_page = 1
                    break
                for remote in remote_items:
                    if not isinstance(remote, dict):
                        continue
                    consume(remote)
                if page > 0:
                    next_reconcile_page = page + 1
                if not result["has_more"]:
                    next_reconcile_page = 1
                    break
            reconcile_next_page = next_reconcile_page
        else:
            # Initial history import is resumable.  A page number alone is not a
            # safe cursor because newly arriving invoices shift every later
            # page.  Instead, locate the exact oldest processed invoice again,
            # then spend max_pages only on records *after* that anchor.  Pages
            # used to relocate the anchor are intentionally not counted as
            # progress; this prevents both skipped history and permanent stalls.
            locating_anchor = bool(backfill_anchor_id)
            page = 0
            anchor_scan_limit = 5000
            while True:
                if page >= anchor_scan_limit:
                    raise MsmiError("Không tìm thấy mốc tiếp tục mSMI trong giới hạn an toàn")
                request_args = {"invoice_type": invoice_type, "page": page, "size": page_size}
                if from_date:
                    request_args.update({"from_date": from_date, "to_date": to_date})
                result = client.list_invoices(**request_args)
                pages += 1
                remote_items = result["items"]
                if not remote_items:
                    backfill_complete = True
                    break

                page_made_progress = False
                for remote in remote_items:
                    if not isinstance(remote, dict):
                        continue
                    normalized, _ = consume(remote)
                    if locating_anchor:
                        if normalized["remote_id"] == backfill_anchor_id:
                            locating_anchor = False
                        continue
                    # The anchor itself was already persisted.  Only records
                    # following it advance the backfill cursor.
                    backfill_anchor_id = normalized["remote_id"]
                    backfill_anchor_date = normalized["invoice_date"]
                    page_made_progress = True

                if locating_anchor:
                    # If the remote invoice was deleted/reordered, scan to the
                    # end and upsert everything encountered.  This may cost more
                    # reads once, but it is the only safe behavior: never guess a
                    # page and silently skip accounting history.
                    if not result["has_more"]:
                        backfill_complete = True
                        break
                    page += 1
                    continue

                if page_made_progress:
                    progress_pages += 1
                if not result["has_more"]:
                    backfill_complete = True
                    break
                if progress_pages >= page_limit:
                    break
                page += 1

        conn.execute(
            """INSERT INTO msmi_sync_state(
                   invoice_type,last_remote_id,last_invoice_date,last_synced_at,last_status,last_error,
                   backfill_anchor_id,backfill_anchor_date,backfill_complete,reconcile_next_page
               ) VALUES(?,?,?,?,?,'',?,?,?,?) ON CONFLICT(invoice_type) DO UPDATE SET
               last_remote_id=excluded.last_remote_id,last_invoice_date=excluded.last_invoice_date,
               last_synced_at=excluded.last_synced_at,last_status='ok',last_error='',
               backfill_anchor_id=excluded.backfill_anchor_id,
               backfill_anchor_date=excluded.backfill_anchor_date,
               backfill_complete=excluded.backfill_complete,
               reconcile_next_page=excluded.reconcile_next_page""",
            (invoice_type, newest_id, newest_date, now_iso(), "ok",
             backfill_anchor_id, backfill_anchor_date, 1 if backfill_complete else 0,
             reconcile_next_page),
        )
        audit(conn, now_iso, "msmi.sync", "ok", metadata={
            "new_invoices": new_count, "known_invoices": seen_count,
            "invoice_items": item_count, "pages": pages,
            "review_required": review_count,
            "backfill_progress_pages": progress_pages,
            "backfill_complete": backfill_complete,
            "reconcile_next_page": reconcile_next_page,
        })
        conn.execute(f"RELEASE SAVEPOINT {savepoint}")
        return {
            "new_invoices": new_count, "known_invoices": seen_count,
            "items": item_count, "pages": pages,
            "review_required": review_count,
            "backfill_complete": backfill_complete,
            "more_history": not backfill_complete,
            "reconcile_next_page": reconcile_next_page,
        }
    except Exception as error:
        conn.execute(f"ROLLBACK TO SAVEPOINT {savepoint}")
        conn.execute(f"RELEASE SAVEPOINT {savepoint}")
        message = str(error)[:300]
        conn.execute(
            """INSERT INTO msmi_sync_state(invoice_type,last_synced_at,last_status,last_error)
               VALUES(?,?,'error',?) ON CONFLICT(invoice_type) DO UPDATE SET
               last_synced_at=excluded.last_synced_at,last_status='error',last_error=excluded.last_error""",
            (invoice_type, now_iso(), message),
        )
        audit(conn, now_iso, "msmi.sync", "error", message=message)
        raise


def payroll_rows(conn, month: str):
    start = month + "-01"
    try:
        start_date = datetime.strptime(start, "%Y-%m-%d").date()
    except ValueError:
        raise ValueError("Tháng tính lương phải có dạng YYYY-MM") from None
    if start_date.month == 12:
        end_date = date(start_date.year + 1, 1, 1)
    else:
        end_date = date(start_date.year, start_date.month + 1, 1)
    output = []
    for staff_row in conn.execute("SELECT * FROM staff WHERE active=1 ORDER BY full_name"):
        person = dict(staff_row)
        hours = conn.execute(
            """SELECT COALESCE(SUM(normal_hours),0) normal,COALESCE(SUM(overtime_hours),0) overtime,
                      COALESCE(SUM(sunday_hours),0) sunday,COALESCE(SUM(night_hours),0) night,
                      COALESCE(SUM(holiday_hours),0) holiday
               FROM attendance_entries WHERE employee_id=? AND work_date>=? AND work_date<?""",
            (person["id"], start_date.isoformat(), end_date.isoformat()),
        ).fetchone()
        adjustment = conn.execute(
            "SELECT * FROM payroll_adjustments WHERE employee_id=? AND month=?",
            (person["id"], month),
        ).fetchone()
        adj = dict(adjustment) if adjustment else defaultdict(float)
        hourly = person["base_salary"] / max(person["standard_days"] * person["standard_hours"], 1)
        normal_pay = hourly * hours["normal"]
        overtime_pay = hourly * hours["overtime"] * 1.5
        sunday_pay = hourly * hours["sunday"] * 2
        night_pay = hourly * hours["night"] * 1.3
        holiday_pay = hourly * hours["holiday"] * 3
        gross = (normal_pay + overtime_pay + sunday_pay + night_pay + holiday_pay
                 + as_number(adj["allowance"]) + as_number(adj["responsibility"]))
        employee_bhxh = as_number(adj["bhxh_employee_amount"])
        if employee_bhxh <= 0:
            employee_bhxh = person["base_salary"] * person["bhxh_employee_rate"]
        company_bhxh = as_number(adj["bhxh_company_amount"])
        if company_bhxh <= 0:
            company_bhxh = person["base_salary"] * person["bhxh_company_rate"]
        net = gross - as_number(adj["advance"]) - as_number(adj["probation_deduction"]) - employee_bhxh
        if int(as_number(adj["use_override"])):
            gross = as_number(adj["gross_override"])
            net = as_number(adj["net_override"])
        output.append({
            **person, "month": month, "normal_hours": hours["normal"], "overtime_hours": hours["overtime"],
            "sunday_hours": hours["sunday"], "night_hours": hours["night"], "holiday_hours": hours["holiday"],
            "normal_pay": normal_pay, "overtime_pay": overtime_pay, "sunday_pay": sunday_pay,
            "night_pay": night_pay, "holiday_pay": holiday_pay, "allowance": as_number(adj["allowance"]),
            "responsibility": as_number(adj["responsibility"]), "advance": as_number(adj["advance"]),
            "probation_deduction": as_number(adj["probation_deduction"]), "bhxh_employee": employee_bhxh,
            "bhxh_company": company_bhxh, "gross_salary": gross, "net_salary": net,
        })
    return output


VIET_DIGITS = ["không", "một", "hai", "ba", "bốn", "năm", "sáu", "bảy", "tám", "chín"]


def number_to_vietnamese(value: float) -> str:
    number = int(round(value))
    if number == 0:
        return "Không đồng"

    def three_digits(group: int, full=False) -> str:
        hundred, remainder = divmod(group, 100)
        ten, unit = divmod(remainder, 10)
        words = []
        if hundred or full:
            words += [VIET_DIGITS[hundred], "trăm"]
        if ten > 1:
            words += [VIET_DIGITS[ten], "mươi"]
            if unit == 1:
                words.append("mốt")
            elif unit == 5:
                words.append("lăm")
            elif unit:
                words.append(VIET_DIGITS[unit])
        elif ten == 1:
            words.append("mười")
            if unit == 5:
                words.append("lăm")
            elif unit:
                words.append(VIET_DIGITS[unit])
        elif unit:
            if hundred or full:
                words.append("lẻ")
            words.append(VIET_DIGITS[unit])
        return " ".join(words)

    groups = []
    while number:
        groups.append(number % 1000)
        number //= 1000
    units = ["", "nghìn", "triệu", "tỷ", "nghìn tỷ", "triệu tỷ"]
    words = []
    for index in range(len(groups) - 1, -1, -1):
        group = groups[index]
        if group:
            words.append(three_digits(group, full=bool(words) and group < 100))
            if index < len(units) and units[index]:
                words.append(units[index])
    text = " ".join(words).strip()
    return text[:1].upper() + text[1:] + " đồng"


def windows_printer_state() -> dict:
    """Return installed/default printers without changing Windows settings."""

    if os.name != "nt":
        return {"supported": False, "default": "", "installed": [], "error": "Chỉ hỗ trợ Windows"}
    try:
        import win32print

        flags = win32print.PRINTER_ENUM_LOCAL | win32print.PRINTER_ENUM_CONNECTIONS
        installed = sorted({
            str(item[2]).strip()
            for item in win32print.EnumPrinters(flags)
            if len(item) > 2 and str(item[2]).strip()
        })
        try:
            default = str(win32print.GetDefaultPrinter() or "").strip()
        except Exception:
            default = ""
        return {"supported": True, "default": default, "installed": installed, "error": ""}
    except Exception:
        return {
            "supported": True,
            "default": "",
            "installed": [],
            "error": "Không đọc được danh sách máy in Windows",
        }


def safe_print_path(data_dir: Path, batch_id: int, value: str) -> Path:
    root = (data_dir / "print_jobs" / str(batch_id)).resolve()
    path = Path(value).resolve()
    if not path.is_relative_to(root):
        raise ValueError("Đường dẫn file in nằm ngoài thư mục phiên được phép")
    return path


def claim_approved_print_jobs(conn, batch_id: int, job_ids: list[int]) -> bool:
    """Atomically claim the exact approved jobs before touching Windows.

    The caller must commit this transition before starting the PDF print tool.
    A savepoint keeps a future multi-document bundle all-or-nothing if one row
    was claimed, invalidated or submitted by another request in the meantime.
    """

    ids = sorted({int(job_id) for job_id in job_ids})
    if not ids:
        return False
    placeholders = ",".join("?" for _ in ids)
    conn.execute("SAVEPOINT claim_print_jobs")
    try:
        changed = conn.execute(
            f"""UPDATE print_jobs
                SET status='submitting',submitted_at=NULL,printed_at=NULL,
                    printer_name=NULL,error_message=NULL
                WHERE batch_id=? AND status='approved' AND id IN ({placeholders})""",
            (batch_id, *ids),
        ).rowcount
        if changed != len(ids):
            conn.execute("ROLLBACK TO claim_print_jobs")
            conn.execute("RELEASE claim_print_jobs")
            return False
        conn.execute("RELEASE claim_print_jobs")
        return True
    except Exception:
        conn.execute("ROLLBACK TO claim_print_jobs")
        conn.execute("RELEASE claim_print_jobs")
        raise


def finish_claimed_print_jobs(
    conn,
    job_ids: list[int],
    *,
    status: str,
    submitted_at=None,
    printer_name="",
    copies=1,
    error_message=None,
) -> bool:
    """CAS a claimed bundle to its terminal, non-retryable state."""

    if status not in {"submitted", "submission_unknown"}:
        raise ValueError("Trạng thái kết thúc lệnh in không hợp lệ")
    ids = sorted({int(job_id) for job_id in job_ids})
    if not ids:
        return False
    placeholders = ",".join("?" for _ in ids)
    conn.execute("SAVEPOINT finish_print_jobs")
    try:
        changed = conn.execute(
            f"""UPDATE print_jobs
                SET status=?,submitted_at=?,printed_at=NULL,printer_name=?,copies=?,error_message=?
                WHERE status='submitting' AND id IN ({placeholders})""",
            (
                status,
                submitted_at if status == "submitted" else None,
                printer_name,
                copies,
                error_message,
                *ids,
            ),
        ).rowcount
        if changed != len(ids):
            conn.execute("ROLLBACK TO finish_print_jobs")
            conn.execute("RELEASE finish_print_jobs")
            return False
        conn.execute("RELEASE finish_print_jobs")
        return True
    except Exception:
        conn.execute("ROLLBACK TO finish_print_jobs")
        conn.execute("RELEASE finish_print_jobs")
        raise


def windows_print_job_ids(win32print_module, printer_name: str) -> set[int] | None:
    """Return live spool job ids, or ``None`` when Windows cannot query them."""

    handle = None
    try:
        handle = win32print_module.OpenPrinter(printer_name)
        return {
            int(job["JobId"])
            for job in win32print_module.EnumJobs(handle, 0, 999, 1)
        }
    except Exception:
        return None
    finally:
        if handle is not None:
            try:
                win32print_module.ClosePrinter(handle)
            except Exception:
                pass


WINDOWS_PAPER_CODES = {"A4": 9, "A5": 11}
WINDOWS_DM_PAPERSIZE = 0x00000002
WINDOWS_DM_DEFAULTSOURCE = 0x00000200
WINDOWS_DM_IN_BUFFER = 0x00000008
WINDOWS_DM_OUT_BUFFER = 0x00000002
SUMATRA_EXE_NAME = "SumatraPDF-3.6.1-64.exe"
SUMATRA_SHA256 = "719F689B34F47BE8CA105CE8484948474DAFDE0E106BAB599E4A89326070C3D0"


def windows_printer_paper_code(win32print_module, printer_name: str) -> int | None:
    """Read the printer driver's current default paper code."""

    handle = None
    try:
        handle = win32print_module.OpenPrinter(printer_name)
        info = win32print_module.GetPrinter(handle, 2)
        code = int(getattr(info.get("pDevMode"), "PaperSize", 0) or 0)
        return code if code > 0 else None
    except Exception:
        return None
    finally:
        if handle is not None:
            try:
                win32print_module.ClosePrinter(handle)
            except Exception:
                pass


def windows_printer_media_state(
    win32print_module, printer_name: str
) -> tuple[int, int] | None:
    """Read the driver's paper size and input-bin codes."""

    handle = None
    try:
        handle = win32print_module.OpenPrinter(printer_name)
        info = win32print_module.GetPrinter(handle, 2)
        devmode = info.get("pDevMode")
        paper_code = int(getattr(devmode, "PaperSize", 0) or 0)
        bin_code = int(getattr(devmode, "DefaultSource", 0) or 0)
        if paper_code <= 0 or bin_code <= 0:
            return None
        return paper_code, bin_code
    except Exception:
        return None
    finally:
        if handle is not None:
            try:
                win32print_module.ClosePrinter(handle)
            except Exception:
                pass


def windows_printer_bins(win32print_module, printer_name: str) -> list[dict]:
    """Return the driver's input-bin codes and exact command-line names."""

    handle = None
    try:
        handle = win32print_module.OpenPrinter(printer_name)
        info = win32print_module.GetPrinter(handle, 2)
        port_name = str(info.get("pPortName") or "").strip()
        codes = tuple(win32print_module.DeviceCapabilities(printer_name, port_name, 6) or ())
        names = tuple(win32print_module.DeviceCapabilities(printer_name, port_name, 12) or ())
        return [
            {"code": int(code), "name": str(name).strip()}
            for code, name in zip(codes, names)
            if int(code) > 0 and str(name).strip()
        ]
    except Exception:
        return []
    finally:
        if handle is not None:
            try:
                win32print_module.ClosePrinter(handle)
            except Exception:
                pass


def select_windows_printer_bin(bins: list[dict], paper: str) -> dict | None:
    """Route A4 to Drawer 1 and A5 to the registered multi-purpose tray."""

    paper = str(paper).strip().upper()
    if paper not in WINDOWS_PAPER_CODES:
        return None
    target_code = 1 if paper == "A4" else 4
    target_names = (
        {"drawer 1", "cassette 1"}
        if paper == "A4"
        else {"multi-purpose tray", "multipurpose tray", "manual feed"}
    )
    return next(
        (
            item for item in bins
            if item["code"] == target_code
            or item["name"].casefold() in target_names
        ),
        None,
    )


def set_windows_printer_media(
    win32print_module, printer_name: str, paper, bin_code: int | None = None
) -> tuple[int, int]:
    """Set and verify paper size plus the physical input tray."""

    code = WINDOWS_PAPER_CODES.get(str(paper).strip().upper())
    if code is None:
        try:
            code = int(paper)
        except (TypeError, ValueError) as exc:
            raise ValueError("Khổ giấy Windows không hợp lệ") from exc
    if code <= 0:
        raise ValueError("Khổ giấy Windows không hợp lệ")
    handle = None
    try:
        access = getattr(win32print_module, "PRINTER_ALL_ACCESS", 0x000F000C)
        handle = win32print_module.OpenPrinter(
            printer_name, {"DesiredAccess": access}
        )
        info = win32print_module.GetPrinter(handle, 2)
        devmode = info.get("pDevMode")
        if devmode is None:
            raise RuntimeError("printer driver has no DEVMODE")
        devmode.Fields = int(getattr(devmode, "Fields", 0)) | WINDOWS_DM_PAPERSIZE
        devmode.PaperSize = code
        if bin_code is not None:
            bin_code = int(bin_code)
            if bin_code <= 0:
                raise ValueError("Khay giấy Windows không hợp lệ")
            devmode.Fields |= WINDOWS_DM_DEFAULTSOURCE
            devmode.DefaultSource = bin_code
        # Canon Generic Plus stores the effective paper/feed combination in
        # its private DEVMODE area as well as the public Windows fields.  Ask
        # the driver to merge and validate both halves before applying it;
        # otherwise A5 can feed correctly but render as a blank page.
        result = win32print_module.DocumentProperties(
            0,
            handle,
            printer_name,
            devmode,
            devmode,
            WINDOWS_DM_IN_BUFFER | WINDOWS_DM_OUT_BUFFER,
        )
        if int(result) < 0:
            raise RuntimeError("printer driver rejected requested paper/bin")
        info["pDevMode"] = devmode
        win32print_module.SetPrinter(handle, 2, info, 0)
        actual_mode = win32print_module.GetPrinter(handle, 2).get("pDevMode")
        actual_paper = int(getattr(actual_mode, "PaperSize", 0) or 0)
        actual_bin = int(getattr(actual_mode, "DefaultSource", 0) or 0)
        if actual_paper != code or (bin_code is not None and actual_bin != bin_code):
            raise RuntimeError("printer driver did not retain requested paper/bin")
        return actual_paper, actual_bin
    finally:
        if handle is not None:
            try:
                win32print_module.ClosePrinter(handle)
            except Exception:
                pass


def set_windows_printer_paper(win32print_module, printer_name: str, paper) -> int:
    """Backward-compatible paper-only wrapper used by diagnostics."""

    return set_windows_printer_media(win32print_module, printer_name, paper)[0]


def pdf_print_tool_path() -> Path | None:
    """Locate and integrity-check the bundled standalone PDF print tool."""

    override = str(os.environ.get("TDP_PDF_PRINT_TOOL", "")).strip()
    candidates = [Path(override)] if override else []
    candidates.append(Path(__file__).resolve().parent / "print_tools" / SUMATRA_EXE_NAME)
    for candidate in candidates:
        try:
            if (
                candidate.is_file()
                and hashlib.sha256(candidate.read_bytes()).hexdigest().upper() == SUMATRA_SHA256
            ):
                return candidate.resolve()
        except OSError:
            continue
    return None


def sumatra_print_settings(paper: str, bin_item: dict) -> str:
    """Build the explicit paper/tray/sides contract for one print job."""

    paper = str(paper).strip().upper()
    if paper not in WINDOWS_PAPER_CODES:
        raise ValueError("Khổ giấy in không hợp lệ")
    sides = "duplexlong" if paper == "A4" else "simplex"
    return ",".join((
        f"paperkind={WINDOWS_PAPER_CODES[paper]}",
        f"bin={int(bin_item['code'])}",
        "fit", "center", sides, "monochrome", "ignore-pdf-print-settings",
    ))


def submit_pdf_via_sumatra(
    tool_path: Path,
    pdf_path: Path,
    printer_name: str,
    paper: str,
    bin_item: dict,
    win32print_module,
    known_job_ids: set[int],
    *,
    appdata_dir: Path,
    timeout_seconds: float = 60.0,
) -> set[int]:
    """Print without the unstable Windows/WPS PDF shell association."""

    appdata_dir.mkdir(parents=True, exist_ok=True)
    settings = sumatra_print_settings(paper, bin_item)
    process = subprocess.Popen(
        [
            str(tool_path), "-silent", "-appdata", str(appdata_dir),
            "-print-to", printer_name, "-print-settings", settings, str(pdf_path),
        ],
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )
    observed: set[int] = set()
    deadline = time.monotonic() + max(1.0, timeout_seconds)
    exit_seen_at = None
    while time.monotonic() < deadline:
        current = windows_print_job_ids(win32print_module, printer_name)
        if current is not None:
            observed.update(current - known_job_ids)
        return_code = process.poll()
        if return_code is not None:
            if return_code != 0:
                raise RuntimeError(f"PDF print tool failed with exit code {return_code}")
            if observed:
                return observed
            if exit_seen_at is None:
                exit_seen_at = time.monotonic()
            elif time.monotonic() - exit_seen_at >= 3.0:
                break
        time.sleep(0.05)
    if process.poll() is None:
        process.terminate()
        try:
            process.wait(timeout=3)
        except subprocess.TimeoutExpired:
            process.kill()
    return observed


def wait_for_spool_submission(
    win32print_module,
    printer_name: str,
    known_job_ids: set[int],
    *,
    timeout_seconds: float = 20.0,
    poll_seconds: float = 0.1,
) -> set[int]:
    """Wait for the PDF handler to create a job on the selected printer.

    Shell handlers such as WPS return from ``os.startfile(..., 'print')``
    before they resolve the default printer. Restoring the previous default
    immediately can silently redirect the PDF to another queue.
    """

    deadline = time.monotonic() + max(0.1, timeout_seconds)
    while time.monotonic() < deadline:
        current = windows_print_job_ids(win32print_module, printer_name)
        if current is None:
            return set()
        new_ids = current - known_job_ids
        if new_ids:
            return new_ids
        time.sleep(max(0.02, poll_seconds))
    return set()


def register_contract_routes(app, ctx):
    db_factory = ctx["db"]
    now_iso = ctx["now_iso"]
    clean_text = ctx["clean_text"]
    number_value = ctx["number_value"]
    tax_factor = ctx["tax_factor"]
    setting_get = ctx["setting_get"]
    setting_set = ctx["setting_set"]
    create_minvoice_client = ctx.get("create_minvoice_client")
    configured_msmi_client_factory = ctx.get("create_msmi_client")
    root = ctx["root"]
    data_dir = ctx["data_dir"]

    def xcom_payment_error_response(exc):
        status = 404 if exc.code == "not_found" else 409 if exc.code in {
            "missing_config", "missing_scope", "missing_tariff", "no_actual_attendance",
            "preview_required", "preview_not_found", "preview_used", "preview_expired",
            "preview_mismatch", "stale_preview",
        } else 400
        return jsonify({
            "ok": False,
            "error": str(exc),
            "code": exc.code,
            "details": exc.details,
        }), status

    def create_msmi_client():
        if configured_msmi_client_factory is not None:
            return configured_msmi_client_factory()
        return MsmiClient(MsmiConfig.from_env_files([root / ".env", data_dir.parent / ".env"]))

    invoice_snapshot_columns = (
        "buyer_name_snapshot", "buyer_tax_code_snapshot", "buyer_address_snapshot",
        "buyer_email_snapshot", "company_name_snapshot", "company_tax_code_snapshot",
        "company_address_snapshot", "payment_requester_snapshot",
        "payment_bank_name_snapshot", "payment_bank_account_snapshot",
    )

    def current_invoice_snapshot(conn, draft):
        buyer = conn.execute(
            "SELECT * FROM outgoing_buyer_profiles WHERE contractor=?", (draft["contractor"],),
        ).fetchone()
        if not buyer or any(not clean_text(buyer[key]) for key in ("legal_name", "tax_code", "address")):
            raise ValueError(
                "Cần lưu đủ tên pháp lý, mã số thuế và địa chỉ người mua trước khi khóa hóa đơn"
            )
        snapshot = {
            "buyer_name_snapshot": clean_text(buyer["legal_name"]),
            "buyer_tax_code_snapshot": clean_text(buyer["tax_code"]),
            "buyer_address_snapshot": clean_text(buyer["address"]),
            "buyer_email_snapshot": clean_text(buyer["email"]),
            "company_name_snapshot": clean_text(setting_get(conn, "company", "")),
            "company_tax_code_snapshot": clean_text(setting_get(conn, "company_tax_code", "")),
            "company_address_snapshot": clean_text(setting_get(conn, "company_address", "")),
            "payment_requester_snapshot": clean_text(setting_get(conn, "payment_requester", "")),
            "payment_bank_name_snapshot": clean_text(setting_get(conn, "payment_bank_name", "")),
            "payment_bank_account_snapshot": clean_text(setting_get(conn, "payment_bank_account", "")),
        }
        required = tuple(column for column in invoice_snapshot_columns if column != "buyer_email_snapshot")
        if any(not snapshot[column] for column in required):
            raise ValueError("Thiếu hồ sơ công ty hoặc tài khoản thanh toán để khóa cùng hóa đơn")
        return snapshot, buyer

    def stored_invoice_snapshot(draft):
        snapshot = {column: clean_text(draft[column]) for column in invoice_snapshot_columns}
        required = tuple(column for column in invoice_snapshot_columns if column != "buyer_email_snapshot")
        return snapshot if all(snapshot[column] for column in required) else None

    def snapshot_update_values(snapshot):
        return tuple(snapshot[column] for column in invoice_snapshot_columns)

    def finite_number(value, default=0.0, label="Giá trị"):
        if value in (None, ""):
            result = float(default)
        else:
            result = number_value(value, math.nan)
        try:
            result = float(result)
        except (TypeError, ValueError, OverflowError):
            result = math.nan
        if not math.isfinite(result):
            raise ValueError(f"{label} phải là số hữu hạn")
        return result

    def finite_integer(value, default=0, label="Giá trị"):
        result = finite_number(value, default, label)
        if abs(result - round(result)) > 1e-9:
            raise ValueError(f"{label} phải là số nguyên")
        return int(round(result))

    def minvoice_reconciliation_mismatches(remote_data, draft, series, local_lines):
        data = remote_data
        if isinstance(data, list):
            data = data[0] if data else {}
        if isinstance(data, dict) and isinstance(data.get("data"), dict):
            data = data["data"]
        if not isinstance(data, dict):
            return ["dữ liệu đối soát không đúng cấu trúc"]

        def present(*keys):
            for key in keys:
                if key in data and data[key] not in (None, ""):
                    return data[key]
            return None

        mismatches = []
        remote_series = present("inv_invoiceSeries", "khhdon", "series", "invoiceSeries")
        if remote_series is None:
            mismatches.append("thiếu ký hiệu để đối soát")
        elif clean_text(remote_series).upper() != clean_text(series).upper():
            mismatches.append("ký hiệu")
        remote_date = present("inv_invoiceIssuedDate", "tdlap", "invoiceDate", "issuedDate")
        if remote_date is None:
            mismatches.append("thiếu ngày để đối soát")
        elif as_date(remote_date) != clean_text(draft["invoice_date"]):
            mismatches.append("ngày hóa đơn")
        remote_tax_code = present("inv_buyerTaxCode", "nmmst", "buyerTaxCode")
        local_tax_code = clean_text(draft["buyer_tax_code_snapshot"])
        if local_tax_code:
            if remote_tax_code is None:
                mismatches.append("thiếu mã số thuế người mua để đối soát")
            elif clean_text(remote_tax_code) != local_tax_code:
                mismatches.append("mã số thuế người mua")
        for label, aliases, expected in (
            ("tiền trước thuế", ("inv_TotalAmountWithoutVat", "tgtcthue", "subtotal"), draft["subtotal"]),
            ("tiền thuế", ("inv_vatAmount", "tgtthue", "taxAmount"), draft["tax_amount"]),
            ("tổng tiền", ("inv_TotalAmount", "tgtttbso", "totalAmount", "total"), draft["total_amount"]),
        ):
            remote_value = present(*aliases)
            if remote_value is None:
                mismatches.append(f"thiếu {label} để đối soát")
                continue
            parsed_value = number_value(remote_value, math.nan)
            if not math.isfinite(parsed_value) or abs(parsed_value - float(expected)) > 1:
                mismatches.append(label)

        remote_lines = present("details", "hdhhdvu", "invoiceItems")
        if isinstance(remote_lines, list) and remote_lines and all(
            isinstance(wrapper, dict) and isinstance(wrapper.get("data"), list)
            for wrapper in remote_lines
        ):
            remote_lines = [line for wrapper in remote_lines for line in wrapper["data"]]
        if not isinstance(remote_lines, list) or not remote_lines:
            mismatches.append("thiếu chi tiết hàng hóa để đối soát")
            return mismatches
        if len(remote_lines) != len(local_lines):
            mismatches.append("số dòng hàng")
            return mismatches
        for index, (remote_line, local_line) in enumerate(zip(remote_lines, local_lines), 1):
            if not isinstance(remote_line, dict):
                mismatches.append(f"cấu trúc dòng {index}")
                continue
            remote_code = first_value(
                remote_line, "inv_itemCode", "ma", "itemCode", "product_code", default=None,
            )
            remote_qty = first_value(
                remote_line, "inv_quantity", "sluong", "quantity", "qty", default=None,
            )
            remote_price = first_value(
                remote_line, "inv_unitPrice", "dgia", "unitPrice", "unit_price", default=None,
            )
            remote_tax = first_value(remote_line, "ma_thue", "tax", "tax_rate", default=None)
            remote_nature = first_value(remote_line, "tchat", "invoice_nature", default=None)
            if None in (remote_code, remote_qty, remote_price, remote_tax, remote_nature):
                mismatches.append(f"dòng {index} thiếu trường đối soát")
                continue
            parsed_qty = number_value(remote_qty, math.nan)
            parsed_price = number_value(remote_price, math.nan)
            if (
                clean_text(remote_code).upper() != clean_text(local_line["product_code"]).upper()
                or not math.isfinite(parsed_qty)
                or abs(parsed_qty - float(local_line["qty"])) > 1e-6
                or not math.isfinite(parsed_price)
                or abs(parsed_price - float(local_line["unit_price"])) > 1
                or clean_text(remote_tax) != clean_text(local_line["tax"])
                or clean_text(remote_nature) != clean_text(local_line["invoice_nature"])
            ):
                mismatches.append(f"chi tiết dòng {index}")
        return mismatches

    def invoice_payload(conn, limit=100, invoice_id=None):
        params = []
        where = ""
        if invoice_id:
            where = " WHERE i.id=?"
            params.append(invoice_id)
        query = f"""SELECT i.id,i.remote_id,i.invoice_type,i.seller_tax_code,i.seller_name,
                           i.invoice_number,i.invoice_series,i.invoice_date,i.subtotal,i.tax_amount,
                           i.total_amount,i.sync_status,i.receipt_status,i.error_message,i.synced_at,
                           COUNT(li.id) item_count,
                           SUM(CASE WHEN li.inventory_eligible=1 THEN 1 ELSE 0 END) inventory_item_count,
                           SUM(CASE WHEN li.inventory_eligible=1 AND li.mapping_status='mapped' THEN 1 ELSE 0 END) mapped_count
                    FROM msmi_invoices i LEFT JOIN msmi_invoice_items li ON li.invoice_id=i.id
                    {where} GROUP BY i.id ORDER BY i.invoice_date DESC,i.id DESC LIMIT ?"""
        params.append(max(1, min(int(limit), 500)))
        invoices = [dict(row) for row in conn.execute(query, params)]
        name_candidates = product_name_candidates(conn)
        name_suggestions = {
            key: rows[0] for key, rows in name_candidates.items() if len(rows) == 1
        }
        for item in invoices:
            item["mapped_count"] = item["mapped_count"] or 0
            item["inventory_item_count"] = item["inventory_item_count"] or 0
            item["items"] = [dict(row) for row in conn.execute(
                """SELECT li.id,li.line_index,li.source_item_code,li.source_item_name,li.source_unit,
                          li.qty,li.unit_price,li.amount,li.tax_rate,li.source_nature,
                          li.inventory_eligible,li.validation_note,li.product_code,li.mapping_status,
                          li.conversion_factor,li.stock_qty,li.stock_unit_price,
                          p.name product_name,p.unit product_unit
                   FROM msmi_invoice_items li LEFT JOIN products p ON p.code=li.product_code
                   WHERE li.invoice_id=? ORDER BY li.line_index""",
                (item["id"],),
            )]
            for line in item["items"]:
                line["suggested_product_code"] = ""
                line["suggested_product_name"] = ""
                line["candidate_products"] = []
                if line["inventory_eligible"] and line["mapping_status"] != "mapped":
                    source_name_key = mapping_key(line["source_item_name"])
                    line["candidate_products"] = name_candidates.get(source_name_key, [])[:20]
                    suggestion = name_suggestions.get(source_name_key)
                    if suggestion:
                        line["suggested_product_code"] = suggestion["code"]
                        line["suggested_product_name"] = suggestion["name"]
        return invoices

    try:
        from .catalog_products import CatalogError, save_product, register_worksheet_routes
    except ImportError:
        from catalog_products import CatalogError, save_product, register_worksheet_routes
    catalog_context = {'db': db_factory, 'now_iso': now_iso, 'clean_text': clean_text,
                       'catalog_unit': catalog_unit, 'catalog_tax': catalog_tax,
                       'audit': lambda *args, **kwargs: audit(*args, **kwargs)}
    register_worksheet_routes(app, catalog_context)

    @app.get("/api/catalog/products")
    def api_catalog_products():
        term = request.args.get('q', '').strip()[:255]
        offset = max(0, request.args.get('offset', 0, type=int))
        join = 'FROM products p LEFT JOIN outgoing_product_names n ON n.product_code=p.code '
        with db_factory() as conn:
            conn.execute('BEGIN')
            items = [dict(row) for row in conn.execute(
                "SELECT p.code,p.name,p.unit,p.tax,p.catalog_updated_at,COALESCE(n.invoice_name,'') invoice_name " + join + ' ORDER BY p.code')]
        catalog_total = len(items)
        def folded(value):
            text = unicodedata.normalize('NFD', str(value or '').casefold().replace('đ', 'd'))
            return ' '.join(''.join(c for c in text if unicodedata.category(c) != 'Mn').split())
        needle = folded(term)
        if needle:
            items = [r for r in items if any(needle in folded(r[k]) for k in ('code','name','invoice_name'))]
        total = len(items)
        offset = min(offset, ((total - 1) // 50) * 50) if total else 0
        if request.args.get('all') == '1':
            return jsonify(ok=True, items=items, total=total, catalog_total=catalog_total, offset=0, limit=total)
        return jsonify(ok=True, items=items[offset:offset+50], total=total, catalog_total=catalog_total, offset=offset, limit=50)

    @app.route("/api/catalog/products", methods=['POST', 'PUT'])
    def api_catalog_create_product():
        try:
            with db_factory() as conn:
                conn.execute('BEGIN IMMEDIATE')
                product = save_product(conn, request.get_json(silent=True), editing=request.method == 'PUT', ctx=catalog_context)
            return jsonify(ok=True, product=product), 200 if request.method == 'PUT' else 201
        except CatalogError as error:
            return jsonify(ok=False, error=str(error)), error.status

    @app.post("/api/catalog/import/preview")
    def api_catalog_import_preview():
        upload = request.files.get("file")
        if not upload or not upload.filename:
            return jsonify({"ok": False, "error": "Chưa chọn file danh mục Excel"}), 400
        filename = Path(upload.filename).name
        if Path(filename).suffix.lower() not in {".xlsx", ".xlsm"}:
            return jsonify({"ok": False, "error": "Chỉ nhận file .xlsx hoặc .xlsm"}), 400
        payload = upload.read(MAPPING_IMPORT_MAX_BYTES + 1)
        if len(payload) > MAPPING_IMPORT_MAX_BYTES:
            return jsonify({"ok": False, "error": "File Excel vượt quá giới hạn 10 MB"}), 413
        try:
            with zipfile.ZipFile(io.BytesIO(payload)) as archive:
                entries = archive.infolist()
                expanded_size = sum(item.file_size for item in entries)
                if len(entries) > 2_000 or expanded_size > MAPPING_IMPORT_MAX_UNCOMPRESSED_BYTES:
                    return jsonify({"ok": False, "error": "File Excel có cấu trúc quá lớn để đọc an toàn"}), 413
        except zipfile.BadZipFile:
            return jsonify({"ok": False, "error": "File Excel bị hỏng hoặc không đúng định dạng"}), 400

        cutoff = time.time() - MAPPING_IMPORT_TTL_SECONDS
        with CATALOG_IMPORT_LOCK:
            for old_token, item in list(PENDING_CATALOG_IMPORTS.items()):
                if item["created"] < cutoff:
                    PENDING_CATALOG_IMPORTS.pop(old_token, None)
        workbook = None
        try:
            workbook = load_workbook(
                io.BytesIO(payload), read_only=True, data_only=True, keep_links=False,
            )
            with db_factory() as conn:
                preview = parse_catalog_workbook(conn, workbook)
        except ValueError as exc:
            return jsonify({"ok": False, "error": str(exc)}), 400
        except Exception:
            return jsonify({"ok": False, "error": "Không đọc được file danh mục Excel"}), 400
        finally:
            if workbook is not None:
                workbook.close()

        token = uuid.uuid4().hex
        source_hash = hashlib.sha256(payload).hexdigest().upper()
        pending = {
            "created": time.time(),
            "filename": filename,
            "source_hash": source_hash,
            "sheet": preview["sheet"],
            "header_row": preview["header_row"],
            "items": preview.pop("items"),
            "counts": preview["counts"],
            "retained_codes": preview["retained_codes"],
            "database_state_hash": preview.pop("database_state_hash"),
            "has_errors": not preview["can_confirm"],
        }
        with CATALOG_IMPORT_LOCK:
            PENDING_CATALOG_IMPORTS[token] = pending
        return jsonify({
            "ok": True,
            "token": token,
            "filename": filename,
            "source_hash": source_hash,
            "expires_in_minutes": MAPPING_IMPORT_TTL_SECONDS // 60,
            **preview,
        })

    @app.post("/api/catalog/import/confirm")
    def api_catalog_import_confirm():
        body = request.get_json(force=True) or {}
        if body.get("confirmed") is not True:
            return jsonify({"ok": False, "error": "Cần xác nhận trước khi ghi danh mục"}), 400
        token = clean_text(body.get("token"))
        with CATALOG_IMPORT_LOCK:
            pending = PENDING_CATALOG_IMPORTS.pop(token, None)
        if not pending or time.time() - pending["created"] > MAPPING_IMPORT_TTL_SECONDS:
            return jsonify({"ok": False, "error": "Phiên xem trước đã hết hạn; vui lòng chọn lại file"}), 410
        if pending["has_errors"]:
            return jsonify({"ok": False, "error": "File còn dòng lỗi nên chưa thể nhập"}), 400

        inserted_products = updated_products = unchanged_products = 0
        inserted_names = updated_names = unchanged_names = 0
        try:
            with db_factory() as conn:
                conn.execute("BEGIN IMMEDIATE")
                if catalog_database_state_hash(conn) != pending["database_state_hash"]:
                    raise ValueError(
                        "Danh mục đã thay đổi sau khi xem trước; dữ liệu chưa được ghi, vui lòng chọn lại file"
                    )
                timestamp = now_iso()
                for item in pending["items"]:
                    code = item["product_code"]
                    current_product = conn.execute(
                        "SELECT code FROM products WHERE code=?", (code,),
                    ).fetchone()
                    if current_product:
                        if item["product_status"] == "update":
                            updated_products += 1
                        else:
                            unchanged_products += 1
                    else:
                        inserted_products += 1
                    conn.execute(
                        """INSERT INTO products(
                               code,name,unit,tax,supplier,buy_price,purchase_list,
                               product_group,catalog_updated_at
                           ) VALUES(?,?,?,?, '',0,0,?,?)
                           ON CONFLICT(code) DO UPDATE SET
                               name=excluded.name,
                               unit=excluded.unit,
                               tax=excluded.tax,
                               product_group=excluded.product_group,
                               catalog_updated_at=excluded.catalog_updated_at""",
                        (
                            code, item["product_name"], item["unit"], item["tax"],
                            item["product_group"], timestamp,
                        ),
                    )

                    invoice_name = item["invoice_name"]
                    if invoice_name:
                        current_name = conn.execute(
                            "SELECT invoice_name FROM outgoing_product_names WHERE product_code=?",
                            (code,),
                        ).fetchone()
                        if not current_name:
                            inserted_names += 1
                        elif mapping_key(current_name["invoice_name"]) == mapping_key(invoice_name):
                            unchanged_names += 1
                        else:
                            updated_names += 1
                        conn.execute(
                            """INSERT INTO outgoing_product_names(product_code,invoice_name,updated_at)
                               VALUES(?,?,?) ON CONFLICT(product_code) DO UPDATE SET
                               invoice_name=excluded.invoice_name,updated_at=excluded.updated_at""",
                            (code, invoice_name, timestamp),
                        )
                result_counts = {
                    "inserted_products": inserted_products,
                    "updated_products": updated_products,
                    "unchanged_products": unchanged_products,
                    "retained_products": len(pending["retained_codes"]),
                    "inserted_names": inserted_names,
                    "updated_names": updated_names,
                    "unchanged_names": unchanged_names,
                    "processed": len(pending["items"]),
                }
                audit(
                    conn, now_iso, "catalog.bulk_import", "ok",
                    entity_type="catalog", entity_id=pending["source_hash"][:16],
                    metadata={
                        "filename": pending["filename"], "sheet": pending["sheet"],
                        "header_row": pending["header_row"], "source_hash": pending["source_hash"],
                        **result_counts,
                    },
                )
        except ValueError as exc:
            return jsonify({"ok": False, "error": str(exc)}), 409
        return jsonify({"ok": True, "source_hash": pending["source_hash"], **result_counts})

    @app.post("/api/mappings/import/preview")
    def api_mapping_import_preview():
        mapping_type = clean_text(request.form.get("mapping_type"))
        if mapping_type not in MAPPING_ALIASES:
            return jsonify({"ok": False, "error": "Loại dữ liệu mapping không hợp lệ"}), 400
        upload = request.files.get("file")
        if not upload or not upload.filename:
            return jsonify({"ok": False, "error": "Chưa chọn file Excel"}), 400
        filename = Path(upload.filename).name
        if Path(filename).suffix.lower() not in {".xlsx", ".xlsm"}:
            return jsonify({"ok": False, "error": "Chỉ nhận file .xlsx hoặc .xlsm"}), 400
        payload = upload.read(MAPPING_IMPORT_MAX_BYTES + 1)
        if len(payload) > MAPPING_IMPORT_MAX_BYTES:
            return jsonify({"ok": False, "error": "File Excel vượt quá giới hạn 10 MB"}), 413
        try:
            with zipfile.ZipFile(io.BytesIO(payload)) as archive:
                entries = archive.infolist()
                expanded_size = sum(item.file_size for item in entries)
                if len(entries) > 2_000 or expanded_size > MAPPING_IMPORT_MAX_UNCOMPRESSED_BYTES:
                    return jsonify({"ok": False, "error": "File Excel có cấu trúc quá lớn để đọc an toàn"}), 413
        except zipfile.BadZipFile:
            return jsonify({"ok": False, "error": "File Excel bị hỏng hoặc không đúng định dạng"}), 400

        cutoff = time.time() - MAPPING_IMPORT_TTL_SECONDS
        with MAPPING_IMPORT_LOCK:
            for old_token, item in list(PENDING_MAPPING_IMPORTS.items()):
                if item["created"] < cutoff:
                    PENDING_MAPPING_IMPORTS.pop(old_token, None)
        workbook = None
        try:
            workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True, keep_links=False)
            with db_factory() as conn:
                preview = parse_mapping_workbook(conn, workbook, mapping_type)
                database_state_hash = mapping_database_state_hash(conn, mapping_type)
        except ValueError as exc:
            return jsonify({"ok": False, "error": str(exc)}), 400
        except Exception:
            return jsonify({"ok": False, "error": "Không đọc được file Excel; vui lòng kiểm tra lại file"}), 400
        finally:
            if workbook is not None:
                workbook.close()

        token = uuid.uuid4().hex
        pending = {
            "created": time.time(),
            "mapping_type": mapping_type,
            "filename": filename,
            "sheet": preview["sheet"],
            "header_row": preview["header_row"],
            "items": preview.pop("items"),
            "counts": preview["counts"],
            "database_state_hash": database_state_hash,
            "has_errors": not preview["can_confirm"],
        }
        with MAPPING_IMPORT_LOCK:
            PENDING_MAPPING_IMPORTS[token] = pending
        return jsonify({
            "ok": True,
            "token": token,
            "mapping_type": mapping_type,
            "filename": filename,
            "expires_in_minutes": MAPPING_IMPORT_TTL_SECONDS // 60,
            **preview,
        })

    @app.post("/api/mappings/import/confirm")
    def api_mapping_import_confirm():
        body = request.get_json(force=True) or {}
        if body.get("confirmed") is not True:
            return jsonify({"ok": False, "error": "Cần xác nhận trước khi ghi dữ liệu"}), 400
        token = clean_text(body.get("token"))
        with MAPPING_IMPORT_LOCK:
            pending = PENDING_MAPPING_IMPORTS.pop(token, None)
        if not pending or time.time() - pending["created"] > MAPPING_IMPORT_TTL_SECONDS:
            return jsonify({"ok": False, "error": "Phiên xem trước đã hết hạn; vui lòng chọn lại file"}), 410
        if pending["has_errors"]:
            return jsonify({"ok": False, "error": "File còn dòng lỗi nên chưa thể nhập"}), 400

        mapping_type = pending["mapping_type"]
        is_invoice = mapping_type == "invoice_names"
        catalog_table = "products" if is_invoice else "kitchens"
        mapping_table = "outgoing_product_names" if is_invoice else "kitchen_units"
        mapping_key_column = "product_code" if is_invoice else "kitchen_code"
        mapping_value_column = "invoice_name" if is_invoice else "unit_code"
        inserted = updated = unchanged = 0
        try:
            with db_factory() as conn:
                conn.execute("BEGIN IMMEDIATE")
                if mapping_database_state_hash(conn, mapping_type) != pending["database_state_hash"]:
                    raise ValueError(
                        "Dữ liệu ánh xạ đã thay đổi sau khi xem trước; dữ liệu chưa được ghi, vui lòng chọn lại file"
                    )
                for item in pending["items"]:
                    code = item["code"]
                    target_value = item["target_value"]
                    if not conn.execute(f"SELECT 1 FROM {catalog_table} WHERE code=?", (code,)).fetchone():
                        raise ValueError(f"Mã {code} không còn trong danh mục; dữ liệu chưa được ghi")
                    current = conn.execute(
                        f"SELECT {mapping_value_column} value FROM {mapping_table} WHERE {mapping_key_column}=?",
                        (code,),
                    ).fetchone()
                    if current and mapping_key(current["value"]) == mapping_key(target_value):
                        unchanged += 1
                        continue
                    if current:
                        updated += 1
                    else:
                        inserted += 1
                    conn.execute(
                        f"""INSERT INTO {mapping_table}({mapping_key_column},{mapping_value_column},updated_at)
                            VALUES(?,?,?) ON CONFLICT({mapping_key_column}) DO UPDATE SET
                            {mapping_value_column}=excluded.{mapping_value_column},updated_at=excluded.updated_at""",
                        (code, target_value, now_iso()),
                    )
                result_counts = {
                    "inserted": inserted,
                    "updated": updated,
                    "unchanged": unchanged,
                    "processed": len(pending["items"]),
                }
                audit(
                    conn,
                    now_iso,
                    "mapping.bulk_import",
                    "ok",
                    entity_type="mapping",
                    entity_id=mapping_type,
                    metadata={
                        "filename": pending["filename"],
                        "sheet": pending["sheet"],
                        "header_row": pending["header_row"],
                        **result_counts,
                    },
                )
        except ValueError as exc:
            return jsonify({"ok": False, "error": str(exc)}), 409
        return jsonify({"ok": True, "mapping_type": mapping_type, **result_counts})

    @app.post("/api/kitchen/import/preview")
    def api_kitchen_import_preview():
        work_date = clean_text(request.form.get("work_date"))
        if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", work_date):
            return jsonify({"ok": False, "error": "Ngày xưởng cơm phải có dạng YYYY-MM-DD"}), 400
        try:
            datetime.strptime(work_date, "%Y-%m-%d")
        except ValueError:
            return jsonify({"ok": False, "error": "Ngày xưởng cơm không hợp lệ"}), 400
        upload = request.files.get("file")
        if not upload or not upload.filename:
            return jsonify({"ok": False, "error": "Chưa chọn file xưởng cơm"}), 400
        filename = Path(upload.filename).name
        if Path(filename).suffix.lower() not in {".xlsx", ".xlsm"}:
            return jsonify({"ok": False, "error": "Chỉ nhận file .xlsx hoặc .xlsm"}), 400
        payload = upload.read(MAPPING_IMPORT_MAX_BYTES + 1)
        if len(payload) > MAPPING_IMPORT_MAX_BYTES:
            return jsonify({"ok": False, "error": "File Excel vượt quá giới hạn 10 MB"}), 413
        try:
            with zipfile.ZipFile(io.BytesIO(payload)) as archive:
                entries = archive.infolist()
                expanded_size = sum(item.file_size for item in entries)
                if len(entries) > 2_000 or expanded_size > MAPPING_IMPORT_MAX_UNCOMPRESSED_BYTES:
                    return jsonify({"ok": False, "error": "File Excel có cấu trúc quá lớn để đọc an toàn"}), 413
        except zipfile.BadZipFile:
            return jsonify({"ok": False, "error": "File Excel bị hỏng hoặc không đúng định dạng"}), 400

        cutoff = time.time() - MAPPING_IMPORT_TTL_SECONDS
        with KITCHEN_IMPORT_LOCK:
            for old_token, item in list(PENDING_KITCHEN_IMPORTS.items()):
                if item["created"] < cutoff:
                    PENDING_KITCHEN_IMPORTS.pop(old_token, None)
        workbook = None
        try:
            workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True, keep_links=False)
            with db_factory() as conn:
                preview = parse_kitchen_workbook(conn, workbook, work_date)
                database_state_hash = kitchen_import_database_state_hash(conn)
        except ValueError as exc:
            return jsonify({"ok": False, "error": str(exc)}), 400
        except Exception:
            app.logger.exception("Kitchen workbook preview failed")
            return jsonify({"ok": False, "error": "Không đọc được file xưởng cơm; vui lòng kiểm tra lại file"}), 400
        finally:
            if workbook is not None:
                workbook.close()

        token = uuid.uuid4().hex
        with KITCHEN_IMPORT_LOCK:
            PENDING_KITCHEN_IMPORTS[token] = {
                "created": time.time(),
                "filename": filename,
                "work_date": work_date,
                "plans": preview["plans"],
                "counts": preview["counts"],
                "database_state_hash": database_state_hash,
                "has_errors": not preview["can_confirm"],
            }
        return jsonify({
            "ok": True,
            "token": token,
            "filename": filename,
            "work_date": work_date,
            "expires_in_minutes": MAPPING_IMPORT_TTL_SECONDS // 60,
            **preview,
        })

    @app.post("/api/kitchen/import/confirm")
    def api_kitchen_import_confirm():
        body = request.get_json(force=True) or {}
        if body.get("confirmed") is not True:
            return jsonify({"ok": False, "error": "Cần xác nhận trước khi ghi dữ liệu"}), 400
        token = clean_text(body.get("token"))
        with KITCHEN_IMPORT_LOCK:
            pending_source = PENDING_KITCHEN_IMPORTS.get(token)
        if not pending_source or time.time() - pending_source["created"] > MAPPING_IMPORT_TTL_SECONDS:
            with KITCHEN_IMPORT_LOCK:
                PENDING_KITCHEN_IMPORTS.pop(token, None)
            return jsonify({"ok": False, "error": "Phiên xem trước đã hết hạn; vui lòng chọn lại file"}), 410
        pending = deepcopy(pending_source)
        override_errors = apply_kitchen_meal_count_overrides(
            pending["plans"], body.get("meal_count_overrides") or {}
        )
        if override_errors:
            return jsonify({"ok": False, "error": override_errors[0]}), 400
        remaining_errors = [
            f"{plan['sheet']} · {plan['kitchen']} · {plan['shift']}: {plan['errors'][0]}"
            for plan in pending["plans"] if plan.get("errors")
        ]
        if remaining_errors:
            return jsonify({"ok": False, "error": remaining_errors[0]}), 400
        with KITCHEN_IMPORT_LOCK:
            if PENDING_KITCHEN_IMPORTS.get(token) is not pending_source:
                return jsonify({"ok": False, "error": "Phiên xem trước đã thay đổi; vui lòng chọn lại file"}), 409
            PENDING_KITCHEN_IMPORTS.pop(token, None)

        inserted = updated = saved_items = saved_mappings = 0
        try:
            with db_factory() as conn:
                conn.execute("BEGIN IMMEDIATE")
                if kitchen_import_database_state_hash(conn) != pending["database_state_hash"]:
                    raise ValueError(
                        "Dữ liệu xưởng cơm hoặc bảng giá đã thay đổi sau khi xem trước; "
                        "dữ liệu chưa được ghi, vui lòng chọn lại file"
                    )
                for plan in pending["plans"]:
                    plan_work_date = plan["work_date"]
                    conn.execute(
                        "INSERT INTO kitchens(code,name) VALUES(?,?) ON CONFLICT(code) DO NOTHING",
                        (plan["kitchen"], plan["kitchen"]),
                    )
                    conn.execute(
                        """INSERT INTO kitchen_units(kitchen_code,unit_code,updated_at) VALUES(?,?,?)
                           ON CONFLICT(kitchen_code) DO UPDATE SET
                           unit_code=excluded.unit_code,updated_at=excluded.updated_at""",
                        (plan["kitchen"], plan["xcom_code"], now_iso()),
                    )
                    saved_mappings += 1
                    existing = conn.execute(
                        "SELECT id FROM meal_plans WHERE import_key=?", (plan["import_key"],)
                    ).fetchone()
                    if not existing and plan.get("existing_id"):
                        existing = conn.execute(
                            "SELECT id FROM meal_plans WHERE id=?", (plan["existing_id"],)
                        ).fetchone()
                    if existing:
                        plan_id = existing["id"]
                        conn.execute(
                            """UPDATE meal_plans SET work_date=?,kitchen=?,shift=?,meal_count=?,unit_code=?,
                               status='draft',note=?,source_file=?,source_sheet=?,source_row_start=?,
                               source_row_end=?,menu_count=?,servings_per_menu=?,meal_price=?,other_cost=?,
                               source_financials_json=?,import_key=?,updated_at=? WHERE id=?""",
                            (plan_work_date, plan["kitchen"], plan["shift"], plan["meal_count"],
                             plan["xcom_code"], plan["menu_name"], pending["filename"], plan["sheet"],
                             plan["source_row_start"], plan["source_row_end"], plan["menu_count"],
                             plan["servings_per_menu"], plan["meal_price"], plan["other_cost"],
                             json.dumps(plan["source_financials"], ensure_ascii=False), plan["import_key"],
                             now_iso(), plan_id),
                        )
                        conn.execute("DELETE FROM meal_plan_items WHERE plan_id=?", (plan_id,))
                        updated += 1
                    else:
                        cur = conn.execute(
                            """INSERT INTO meal_plans(
                               work_date,kitchen,shift,meal_count,unit_code,status,note,created_at,updated_at,
                               import_key,source_file,source_sheet,source_row_start,source_row_end,menu_count,
                               servings_per_menu,meal_price,other_cost,source_financials_json
                               ) VALUES(?,?,?,?,?,'draft',?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                            (plan_work_date, plan["kitchen"], plan["shift"], plan["meal_count"],
                             plan["xcom_code"], plan["menu_name"], now_iso(), now_iso(), plan["import_key"],
                             pending["filename"], plan["sheet"], plan["source_row_start"], plan["source_row_end"],
                             plan["menu_count"], plan["servings_per_menu"], plan["meal_price"], plan["other_cost"],
                             json.dumps(plan["source_financials"], ensure_ascii=False)),
                        )
                        plan_id = cur.lastrowid
                        inserted += 1
                    for item in plan["items"]:
                        if not conn.execute(
                            "SELECT 1 FROM products WHERE code=?", (item["product_code"],)
                        ).fetchone():
                            raise ValueError(
                                f"Mã {item['product_code']} không còn trong danh mục; dữ liệu chưa được ghi"
                            )
                        # Confirmation locks the exact current-file/current-period
                        # HATRAN value.  Later catalogue changes cannot silently
                        # rewrite historical cost.
                        conn.execute(
                            """INSERT INTO dated_prices(product_code,price_group,period,price_value,updated_at)
                               VALUES(?,'HATRAN',?,?,?)
                               ON CONFLICT(product_code,price_group,period) DO NOTHING""",
                            (item["product_code"], plan_work_date[:7], item["buy_price"], now_iso()),
                        )
                        # Store one canonical post-confirmation provenance.  On
                        # the first import the price comes from this workbook;
                        # on a replay it is read from the just-locked period.
                        # Those paths must produce byte-for-byte equal state.
                        confirmed_price_source = (
                            f"HATRAN {plan_work_date[:7]} · giá kỳ đã khóa"
                        )
                        conn.execute(
                            """INSERT INTO meal_plan_items(
                               plan_id,dish_name,product_code,product_name,norm_qty,unit,supplier,buy_price,
                               price_source,source_norm_per_1000,applicable_meal_count,source_row,source_amount
                               ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                            (plan_id, item["dish_name"], item["product_code"], item["product_name"],
                             item["norm_qty"], item["unit"], item["supplier"], item["buy_price"],
                             confirmed_price_source, item["norm_per_1000"], item["applicable_meal_count"],
                             item["source_row"], item["source_amount"]),
                        )
                        saved_items += 1
                result_counts = {
                    "inserted": inserted,
                    "updated": updated,
                    "items": saved_items,
                    "mappings": saved_mappings,
                }
                audit(
                    conn,
                    now_iso,
                    "kitchen.bulk_import",
                    "ok",
                    entity_type="meal_plan",
                    entity_id=pending["work_date"],
                    metadata={
                        "filename": pending["filename"],
                        "work_dates": sorted({plan["work_date"] for plan in pending["plans"]}),
                        **result_counts,
                    },
                )
        except ValueError as exc:
            return jsonify({"ok": False, "error": str(exc)}), 409
        except Exception:
            return jsonify({"ok": False, "error": "Không ghi được file xưởng cơm; dữ liệu chưa được thay đổi"}), 409
        return jsonify({
            "ok": True,
            "work_date": pending["work_date"],
            "work_dates": sorted({plan["work_date"] for plan in pending["plans"]}),
            **result_counts,
        })

    @app.get("/api/mappings/<mapping_type>")
    def api_mapping_list(mapping_type):
        with db_factory() as conn:
            if mapping_type == "invoice_names":
                rows = [dict(row) for row in conn.execute(
                    """SELECT m.product_code code,p.name source_name,m.invoice_name target_value,m.updated_at
                       FROM outgoing_product_names m JOIN products p ON p.code=m.product_code
                       ORDER BY m.product_code"""
                )]
            elif mapping_type == "kitchen_units":
                rows = [dict(row) for row in conn.execute(
                    """SELECT m.kitchen_code code,k.name source_name,m.unit_code target_value,m.updated_at
                       FROM kitchen_units m LEFT JOIN kitchens k ON k.code=m.kitchen_code
                       ORDER BY m.kitchen_code"""
                )]
            else:
                return jsonify({"ok": False, "error": "Loại dữ liệu mapping không hợp lệ"}), 404
            return jsonify({"ok": True, "mapping_type": mapping_type, "items": rows})

    @app.get("/api/operations/bootstrap")
    def api_operations_bootstrap():
        as_of = request.args.get("as_of") or date.today().isoformat()
        month = request.args.get("month") or as_of[:7]
        active_only = request.args.get("active_only") == "1"
        with db_factory() as conn:
            sync_state = [dict(row) for row in conn.execute("SELECT * FROM msmi_sync_state")]
            inventory = inventory_rows(conn, as_of)
            labor = [] if active_only else [dict(row) for row in conn.execute(
                "SELECT work_date,kitchen,amount,source FROM kitchen_labor_costs WHERE substr(work_date,1,7)=? ORDER BY work_date,kitchen",
                (month,),
            )]
            meal_attendance = [] if active_only else [dict(row) for row in conn.execute(
                """SELECT work_date,kitchen,shift,actual_count,ordered_count,source_file,source_sheet
                   FROM meal_attendance WHERE substr(work_date,1,7)=?
                   ORDER BY work_date,kitchen,shift""",
                (month,),
            )]
            printer_state = windows_printer_state()
            return jsonify({
                "ok": True,
                "inventory": inventory,
                "inventory_as_of": as_of,
                "inventory_totals": {
                    "products": len(inventory),
                    "accounting_qty": sum(row["accounting_qty"] for row in inventory),
                    "available_qty": sum(row["available_qty"] for row in inventory),
                },
                "msmi": {"state": sync_state, "invoices": invoice_payload(conn, 30)},
                "meal_plans": [] if active_only else meal_plan_payload(conn, request.args.get("date", "")),
                "meal_attendance": meal_attendance,
                "meal_attendance_totals": {
                    "actual": sum(row["actual_count"] for row in meal_attendance),
                    "ordered": sum(row["ordered_count"] for row in meal_attendance),
                    "rows": len(meal_attendance),
                    "kitchens": len({row["kitchen"] for row in meal_attendance}),
                },
                "kitchen_units": [] if active_only else [dict(row) for row in conn.execute("SELECT * FROM kitchen_units ORDER BY kitchen_code")],
                "xcom_payment_profiles": [] if active_only else list_payment_profiles(conn),
                "payroll": [] if active_only else payroll_rows(conn, month),
                "labor_costs": labor,
                "staff": [] if active_only else [dict(row) for row in conn.execute("SELECT * FROM staff ORDER BY full_name")],
                "print_jobs": [dict(row) for row in conn.execute("SELECT * FROM print_jobs ORDER BY id DESC LIMIT 50")],
                "printer": {
                    "name": setting_get(conn, "printer_name", ""),
                    "copies": int(as_number(setting_get(conn, "print_copies", "1"), 1)),
                    "paper": setting_get(conn, "print_other_paper", "A5"),
                    "delivery_paper": "A4",
                    "other_paper": setting_get(conn, "print_other_paper", "A5"),
                    "default": printer_state["default"],
                    "installed": printer_state["installed"],
                    "supported": printer_state["supported"],
                    "error": printer_state["error"],
                },
                "document_settings": {
                    "requester": setting_get(conn, "payment_requester", ""),
                    "bank_name": setting_get(conn, "payment_bank_name", ""),
                    "bank_account": setting_get(conn, "payment_bank_account", ""),
                },
            })

    @app.get("/api/inventory")
    def api_inventory():
        with db_factory() as conn:
            as_of = request.args.get("as_of") or date.today().isoformat()
            return jsonify({"ok": True, "as_of": as_of, "items": inventory_rows(conn, as_of)})

    @app.post("/api/inventory/opening/import/preview")
    def api_inventory_opening_import_preview():
        period = clean_text(request.form.get("period"))
        if not re.fullmatch(r"\d{4}-\d{2}", period):
            return jsonify({"ok": False, "error": "Kỳ tồn đầu phải có dạng YYYY-MM"}), 400
        try:
            datetime.strptime(period + "-01", "%Y-%m-%d")
        except ValueError:
            return jsonify({"ok": False, "error": "Kỳ tồn đầu không hợp lệ"}), 400
        upload = request.files.get("file")
        if not upload or not upload.filename:
            return jsonify({"ok": False, "error": "Chưa chọn file tồn đầu kỳ"}), 400
        filename = Path(upload.filename).name
        if Path(filename).suffix.lower() not in {".xlsx", ".xlsm"}:
            return jsonify({"ok": False, "error": "Chỉ nhận file .xlsx hoặc .xlsm"}), 400
        payload = upload.read(MAPPING_IMPORT_MAX_BYTES + 1)
        if len(payload) > MAPPING_IMPORT_MAX_BYTES:
            return jsonify({"ok": False, "error": "File Excel vượt quá giới hạn 10 MB"}), 413
        try:
            with zipfile.ZipFile(io.BytesIO(payload)) as archive:
                entries = archive.infolist()
                expanded_size = sum(item.file_size for item in entries)
                if len(entries) > 2_000 or expanded_size > MAPPING_IMPORT_MAX_UNCOMPRESSED_BYTES:
                    return jsonify({"ok": False, "error": "File Excel có cấu trúc quá lớn để đọc an toàn"}), 413
        except zipfile.BadZipFile:
            return jsonify({"ok": False, "error": "File Excel bị hỏng hoặc không đúng định dạng"}), 400

        cutoff = time.time() - MAPPING_IMPORT_TTL_SECONDS
        with OPENING_IMPORT_LOCK:
            for old_token, item in list(PENDING_OPENING_IMPORTS.items()):
                if item["created"] < cutoff:
                    PENDING_OPENING_IMPORTS.pop(old_token, None)
        workbook = None
        try:
            workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True, keep_links=False)
            with db_factory() as conn:
                preview = parse_opening_workbook(conn, workbook, period)
                database_state_hash = opening_import_database_state_hash(conn, period)
        except ValueError as exc:
            return jsonify({"ok": False, "error": str(exc)}), 400
        except Exception:
            return jsonify({"ok": False, "error": "Không đọc được file tồn đầu kỳ; vui lòng kiểm tra lại file"}), 400
        finally:
            if workbook is not None:
                workbook.close()

        token = uuid.uuid4().hex
        pending_items = preview.pop("items")
        with OPENING_IMPORT_LOCK:
            PENDING_OPENING_IMPORTS[token] = {
                "created": time.time(),
                "filename": filename,
                "source_hash": hashlib.sha256(payload).hexdigest().upper(),
                "period": period,
                "sheet": preview["sheet"],
                "header_row": preview["header_row"],
                "items": pending_items,
                "counts": preview["counts"],
                "totals": preview["totals"],
                "database_state_hash": database_state_hash,
                "has_errors": not preview["can_confirm"],
            }
        return jsonify({
            "ok": True,
            "token": token,
            "filename": filename,
            "expires_in_minutes": MAPPING_IMPORT_TTL_SECONDS // 60,
            **preview,
        })

    @app.post("/api/inventory/opening/import/confirm")
    def api_inventory_opening_import_confirm():
        body = request.get_json(force=True) or {}
        if body.get("confirmed") is not True:
            return jsonify({"ok": False, "error": "Cần xác nhận trước khi ghi tồn đầu kỳ"}), 400
        token = clean_text(body.get("token"))
        with OPENING_IMPORT_LOCK:
            pending = PENDING_OPENING_IMPORTS.pop(token, None)
        if not pending or time.time() - pending["created"] > MAPPING_IMPORT_TTL_SECONDS:
            return jsonify({"ok": False, "error": "Phiên xem trước đã hết hạn; vui lòng chọn lại file"}), 410
        if pending["has_errors"]:
            return jsonify({"ok": False, "error": "File còn dòng lỗi nên chưa thể nhập"}), 400

        inserted_products = inserted = updated = 0
        try:
            with db_factory() as conn:
                conn.execute("BEGIN IMMEDIATE")
                if opening_import_database_state_hash(conn, pending["period"]) != pending["database_state_hash"]:
                    raise ValueError(
                        "Danh mục hoặc tồn đầu kỳ đã thay đổi sau khi xem trước; "
                        "dữ liệu chưa được ghi, vui lòng chọn lại file"
                    )
                incoming_codes = {item["product_code"] for item in pending["items"]}
                for item in pending["items"]:
                    code = item["product_code"]
                    product = conn.execute("SELECT code FROM products WHERE code=?", (code,)).fetchone()
                    if not product and not item["create_product"]:
                        raise ValueError(f"Mã {code} không còn trong danh mục; dữ liệu chưa được ghi")
                    if not product:
                        conn.execute(
                            """INSERT INTO products(
                               code,name,unit,tax,supplier,buy_price,purchase_list,seller,cccd
                               ) VALUES(?,?,?,?,?, ?,0,'','')""",
                            (code, item["product_name"], item["unit"], item["tax"], "",
                             max(number_value(item["unit_cost"]), 0)),
                        )
                        inserted_products += 1
                    existing = conn.execute(
                        """SELECT id FROM inventory_transactions
                           WHERE source_type='OPENING' AND source_id=? AND source_line=?""",
                        (pending["period"], code),
                    ).fetchone()
                    if existing:
                        updated += 1
                    else:
                        inserted += 1
                    qty = number_value(item["qty"])
                    note = (
                        f"Tồn đầu kỳ {pending['period']} từ {pending['filename']}; "
                        f"sheet {pending['sheet']}; dòng {','.join(map(str, item['source_rows']))}"
                    )
                    conn.execute(
                        """INSERT INTO inventory_transactions(
                           txn_date,product_code,warehouse_codes_json,qty_in,qty_out,unit_cost,source_type,source_id,source_line,
                           status,note,created_at,updated_at
                           ) VALUES(?,?,?,?,?,?,'OPENING',?,?,'posted',?,?,?)
                           ON CONFLICT(source_type,source_id,source_line) DO UPDATE SET
                           txn_date=excluded.txn_date,product_code=excluded.product_code,
                           warehouse_codes_json=excluded.warehouse_codes_json,
                           qty_in=excluded.qty_in,qty_out=excluded.qty_out,unit_cost=excluded.unit_cost,
                           status='posted',note=excluded.note,updated_at=excluded.updated_at""",
                        (pending["period"] + "-01", code,
                         json.dumps(item["warehouse_codes"], ensure_ascii=False, separators=(",", ":")),
                         max(qty, 0), max(-qty, 0),
                         max(number_value(item["unit_cost"]), 0), pending["period"], code,
                         note, now_iso(), now_iso()),
                    )
                stale_rows = conn.execute(
                    "SELECT source_line FROM inventory_transactions WHERE source_type='OPENING' AND source_id=?",
                    (pending["period"],),
                ).fetchall()
                stale_codes = [row["source_line"] for row in stale_rows if row["source_line"] not in incoming_codes]
                if stale_codes:
                    placeholders = ",".join("?" for _ in stale_codes)
                    conn.execute(
                        f"DELETE FROM inventory_transactions WHERE source_type='OPENING' AND source_id=? "
                        f"AND source_line IN ({placeholders})",
                        (pending["period"], *stale_codes),
                    )
                result_counts = {
                    "inserted_products": inserted_products,
                    "inserted": inserted,
                    "updated": updated,
                    "deleted_stale": len(stale_codes),
                    "processed": len(pending["items"]),
                }
                audit(
                    conn, now_iso, "inventory.opening.bulk_import", "ok",
                    entity_type="period", entity_id=pending["period"],
                    metadata={
                        "filename": pending["filename"], "source_hash": pending["source_hash"],
                        "sheet": pending["sheet"], "header_row": pending["header_row"],
                        **result_counts,
                    },
                )
        except ValueError as exc:
            return jsonify({"ok": False, "error": str(exc)}), 409
        except Exception:
            return jsonify({"ok": False, "error": "Không ghi được tồn đầu kỳ; dữ liệu chưa được thay đổi"}), 409
        year, month = map(int, pending["period"].split("-"))
        next_month = date(year + (month == 12), 1 if month == 12 else month + 1, 1)
        as_of = (next_month - timedelta(days=1)).isoformat()
        with db_factory() as conn:
            inventory = inventory_rows(conn, as_of)
        return jsonify({
            "ok": True, "period": pending["period"], **result_counts,
            "totals": pending["totals"], "items": inventory,
        })

    @app.post("/api/kitchen/attendance/import/preview")
    def api_meal_attendance_import_preview():
        period_override = clean_text(request.form.get("period"))
        if period_override and not re.fullmatch(r"\d{4}-\d{2}", period_override):
            return jsonify({"ok": False, "error": "Kỳ chấm suất phải có dạng YYYY-MM"}), 400
        upload = request.files.get("file")
        if not upload or not upload.filename:
            return jsonify({"ok": False, "error": "Chưa chọn file chấm suất ăn"}), 400
        filename = Path(upload.filename).name
        if Path(filename).suffix.lower() not in {".xlsx", ".xlsm"}:
            return jsonify({"ok": False, "error": "Chỉ nhận file .xlsx hoặc .xlsm"}), 400
        payload = upload.read(MAPPING_IMPORT_MAX_BYTES + 1)
        if len(payload) > MAPPING_IMPORT_MAX_BYTES:
            return jsonify({"ok": False, "error": "File Excel vượt quá giới hạn 10 MB"}), 413
        try:
            with zipfile.ZipFile(io.BytesIO(payload)) as archive:
                entries = archive.infolist()
                expanded_size = sum(item.file_size for item in entries)
                if len(entries) > 2_000 or expanded_size > MAPPING_IMPORT_MAX_UNCOMPRESSED_BYTES:
                    return jsonify({"ok": False, "error": "File Excel có cấu trúc quá lớn để đọc an toàn"}), 413
        except zipfile.BadZipFile:
            return jsonify({"ok": False, "error": "File Excel bị hỏng hoặc không đúng định dạng"}), 400

        cutoff = time.time() - MAPPING_IMPORT_TTL_SECONDS
        with MEAL_ATTENDANCE_IMPORT_LOCK:
            for old_token, item in list(PENDING_MEAL_ATTENDANCE_IMPORTS.items()):
                if item["created"] < cutoff:
                    PENDING_MEAL_ATTENDANCE_IMPORTS.pop(old_token, None)
        workbook = None
        try:
            workbook = load_workbook(io.BytesIO(payload), read_only=False, data_only=True, keep_links=False)
            with db_factory() as conn:
                preview = parse_meal_attendance_workbook(conn, workbook, period_override)
                database_state_hash = meal_attendance_database_state_hash(conn, preview["periods"])
        except ValueError as exc:
            return jsonify({"ok": False, "error": str(exc)}), 400
        except Exception:
            return jsonify({"ok": False, "error": "Không đọc được file chấm suất ăn; vui lòng kiểm tra lại file"}), 400
        finally:
            if workbook is not None:
                workbook.close()

        token = uuid.uuid4().hex
        pending_items = preview.pop("items")
        with MEAL_ATTENDANCE_IMPORT_LOCK:
            PENDING_MEAL_ATTENDANCE_IMPORTS[token] = {
                "created": time.time(),
                "filename": filename,
                "source_hash": hashlib.sha256(payload).hexdigest().upper(),
                "periods": preview["periods"],
                "items": pending_items,
                "counts": preview["counts"],
                "totals": preview["totals"],
                "database_state_hash": database_state_hash,
                "has_errors": not preview["can_confirm"],
            }
        return jsonify({
            "ok": True, "token": token, "filename": filename,
            "expires_in_minutes": MAPPING_IMPORT_TTL_SECONDS // 60,
            **preview,
        })

    @app.post("/api/kitchen/attendance/import/confirm")
    def api_meal_attendance_import_confirm():
        body = request.get_json(force=True) or {}
        if body.get("confirmed") is not True:
            return jsonify({"ok": False, "error": "Cần xác nhận trước khi ghi chấm suất ăn"}), 400
        token = clean_text(body.get("token"))
        with MEAL_ATTENDANCE_IMPORT_LOCK:
            pending = PENDING_MEAL_ATTENDANCE_IMPORTS.pop(token, None)
        if not pending or time.time() - pending["created"] > MAPPING_IMPORT_TTL_SECONDS:
            return jsonify({"ok": False, "error": "Phiên xem trước đã hết hạn; vui lòng chọn lại file"}), 410
        if pending["has_errors"]:
            return jsonify({"ok": False, "error": "File còn dòng lỗi nên chưa thể nhập"}), 400

        inserted = updated = unchanged = deleted = 0
        try:
            with db_factory() as conn:
                conn.execute("BEGIN IMMEDIATE")
                if meal_attendance_database_state_hash(conn, pending["periods"]) != pending["database_state_hash"]:
                    raise ValueError(
                        "Chấm suất ăn trong kỳ đã thay đổi sau khi xem trước; "
                        "dữ liệu chưa được ghi, vui lòng chọn lại file"
                    )
                existing = {
                    (row["work_date"], row["kitchen"], row["shift"]): dict(row)
                    for period in pending["periods"]
                    for row in conn.execute(
                        "SELECT * FROM meal_attendance WHERE source_type='MONTHLY_WORKBOOK' AND substr(work_date,1,7)=?",
                        (period,),
                    )
                }
                incoming_keys = set()
                for item in pending["items"]:
                    key = (item["work_date"], item["kitchen"], item["shift"])
                    incoming_keys.add(key)
                    previous = existing.get(key)
                    if previous is None:
                        inserted += 1
                    elif (
                        abs(previous["actual_count"] - item["actual_count"]) <= 1e-9
                        and abs(previous["ordered_count"] - item["ordered_count"]) <= 1e-9
                    ):
                        unchanged += 1
                    else:
                        updated += 1
                    conn.execute(
                        """INSERT INTO meal_attendance(
                           work_date,kitchen,shift,actual_count,ordered_count,source_type,
                           source_file,source_sheet,source_column,updated_at
                           ) VALUES(?,?,?,?,?,'MONTHLY_WORKBOOK',?,?,?,?)
                           ON CONFLICT(work_date,kitchen,shift) DO UPDATE SET
                           actual_count=excluded.actual_count,ordered_count=excluded.ordered_count,
                           source_type='MONTHLY_WORKBOOK',source_file=excluded.source_file,
                           source_sheet=excluded.source_sheet,source_column=excluded.source_column,
                           updated_at=excluded.updated_at""",
                        (item["work_date"], item["kitchen"], item["shift"], item["actual_count"],
                         item["ordered_count"], pending["filename"], item["source_sheet"],
                         item["source_column"], now_iso()),
                    )
                for key in set(existing).difference(incoming_keys):
                    conn.execute(
                        "DELETE FROM meal_attendance WHERE work_date=? AND kitchen=? AND shift=? AND source_type='MONTHLY_WORKBOOK'",
                        key,
                    )
                    deleted += 1
                result_counts = {
                    "inserted": inserted, "updated": updated, "unchanged": unchanged,
                    "deleted": deleted, "processed": len(pending["items"]),
                }
                audit(
                    conn, now_iso, "kitchen.meal_attendance.bulk_import", "ok",
                    entity_type="period", entity_id=",".join(pending["periods"]),
                    metadata={
                        "filename": pending["filename"], "source_hash": pending["source_hash"],
                        **result_counts,
                    },
                )
        except ValueError as exc:
            return jsonify({"ok": False, "error": str(exc)}), 409
        except Exception:
            return jsonify({"ok": False, "error": "Không ghi được chấm suất ăn; dữ liệu chưa được thay đổi"}), 409
        return jsonify({
            "ok": True, "periods": pending["periods"], **result_counts,
            "totals": pending["totals"],
        })

    @app.post("/api/inventory/opening")
    def api_inventory_opening():
        body = request.get_json(force=True) or {}
        period = clean_text(body.get("period"))
        items = body.get("items") or []
        if not re.fullmatch(r"\d{4}-\d{2}", period) or not isinstance(items, list):
            return jsonify({"ok": False, "error": "Kỳ tồn đầu phải có dạng YYYY-MM và có danh sách hàng"}), 400
        try:
            datetime.strptime(period + "-01", "%Y-%m-%d")
        except ValueError:
            return jsonify({"ok": False, "error": "Kỳ tồn đầu không hợp lệ"}), 400
        txn_date = period + "-01"
        try:
            if any(not isinstance(item, dict) for item in items):
                raise ValueError("Danh sách tồn đầu có dòng không hợp lệ")
            normalized_items = [{
                **item,
                "_qty": finite_number(item.get("qty"), 0, "Số lượng tồn đầu"),
                "_unit_cost": finite_number(item.get("unit_cost"), 0, "Đơn giá tồn đầu"),
            } for item in items]
        except ValueError as exc:
            return jsonify({"ok": False, "error": str(exc)}), 400
        with db_factory() as conn:
            conn.execute('BEGIN IMMEDIATE')
            opening_change = None
            if 'expected_opening' in body:
                expected = body['expected_opening']
                if not isinstance(expected, dict) or len(normalized_items) != 1:
                    return jsonify(ok=False, error='Cần tải lại dòng tồn đầu trước khi sửa.'), 400
                current = conn.execute(
                    "SELECT id,txn_date,product_code,source_id,source_line,qty_in,qty_out,unit_cost,note,updated_at "
                    "FROM inventory_transactions WHERE id=? AND source_type='OPENING' AND status='posted' "
                    "AND source_id=? AND source_line=? AND product_code=? AND txn_date=?",
                    (expected.get('id'), period, clean_text(normalized_items[0].get('product_code')).upper(),
                     clean_text(normalized_items[0].get('product_code')).upper(), txn_date),
                ).fetchone()
                newer_opening = conn.execute(
                    "SELECT 1 FROM inventory_transactions WHERE source_type='OPENING' AND status='posted' "
                    "AND (txn_date>? OR (txn_date=? AND source_id<>?)) LIMIT 1", (txn_date, txn_date, period),
                ).fetchone()
                if not current or dict(current) != expected or newer_opening:
                    return jsonify(ok=False, error='Tồn đầu đã thay đổi. Đóng cửa sổ rồi mở lại để kiểm tra.'), 409
                normalized_items[0]['_unit_cost'] = current['unit_cost']
                normalized_items[0]['note'] = (current['note'] or '') + ' · Đã sửa số tồn trên web'
                opening_change = {'product_code': current['product_code'],
                                  'before_qty': current['qty_in'] - current['qty_out'],
                                  'after_qty': normalized_items[0]['_qty'], 'source_note': current['note']}
            known_codes = {
                row["code"] for row in conn.execute("SELECT code FROM products")
            }
            missing_codes = sorted({
                clean_text(item.get("product_code")).upper()
                for item in normalized_items
                if clean_text(item.get("product_code")).upper() not in known_codes
            })
            if missing_codes:
                return jsonify({
                    "ok": False,
                    "error": "Tồn đầu có mã hàng chưa tồn tại: " + ", ".join(missing_codes[:10]),
                }), 400
            saved = 0
            for item in normalized_items:
                code = clean_text(item.get("product_code")).upper()
                qty = item["_qty"]
                conn.execute(
                    """INSERT INTO inventory_transactions(
                        txn_date,product_code,qty_in,qty_out,unit_cost,source_type,source_id,source_line,
                        status,note,created_at,updated_at
                    ) VALUES(?,?,?,?,?,'OPENING',?,?,'posted',?,?,?)
                    ON CONFLICT(source_type,source_id,source_line) DO UPDATE SET
                        qty_in=excluded.qty_in,qty_out=excluded.qty_out,unit_cost=excluded.unit_cost,
                        note=excluded.note,updated_at=excluded.updated_at""",
                    (txn_date, code, max(qty, 0), max(-qty, 0),
                     max(item["_unit_cost"], 0), period, code,
                     clean_text(item.get("note")), now_iso(), now_iso()),
                )
                saved += 1
            audit(conn, now_iso, "inventory.opening", "ok", entity_type="period", entity_id=period,
                  metadata={"items": saved, **({'change': opening_change} if opening_change else {})})
            return jsonify({"ok": True, "saved": saved, "items": inventory_rows(conn, date.today().isoformat())})

    @app.post("/api/inventory/adjustments")
    def api_inventory_adjustment():
        body = request.get_json(force=True) or {}
        code = clean_text(body.get("product_code")).upper()
        try:
            qty = finite_number(body.get("qty"), 0, "Số lượng điều chỉnh")
            unit_cost = finite_number(body.get("unit_cost"), 0, "Đơn giá điều chỉnh")
        except ValueError as exc:
            return jsonify({"ok": False, "error": str(exc)}), 400
        if not code or qty == 0:
            return jsonify({"ok": False, "error": "Cần mã hàng và số lượng điều chỉnh khác 0"}), 400
        txn_date = clean_text(body.get("txn_date")) or date.today().isoformat()
        if as_date(txn_date) != txn_date:
            return jsonify({"ok": False, "error": "Ngày điều chỉnh phải hợp lệ dạng YYYY-MM-DD"}), 400
        with db_factory() as conn:
            if not conn.execute("SELECT 1 FROM products WHERE code=?", (code,)).fetchone():
                return jsonify({"ok": False, "error": "Mã hàng chưa có trong danh mục"}), 400
            source_id = clean_text(body.get("reference")) or f"ADJ-{datetime.now():%Y%m%d%H%M%S%f}"
            conn.execute(
                """INSERT INTO inventory_transactions(
                    txn_date,product_code,qty_in,qty_out,unit_cost,source_type,source_id,source_line,
                    status,note,created_at,updated_at
                ) VALUES(?,?,?,?,?,'ADJUSTMENT',?,'','posted',?,?,?)""",
                (txn_date, code, max(qty, 0), max(-qty, 0),
                 max(unit_cost, 0), source_id, clean_text(body.get("note")),
                 now_iso(), now_iso()),
            )
            audit(conn, now_iso, "inventory.adjust", "ok", entity_type="product", entity_id=code,
                  metadata={"qty": qty, "reference": source_id})
            return jsonify({"ok": True})

    @app.get("/api/supplier-needs/<int:batch_id>")
    def api_supplier_needs(batch_id):
        with db_factory() as conn:
            try:
                payload = purchase_order_payload(conn, batch_id)
            except ValueError as exc:
                return jsonify({"ok": False, "error": str(exc)}), 404
            return jsonify({"ok": True, **payload})

    @app.put("/api/supplier-rules/<supplier_code>")
    def api_supplier_rule(supplier_code):
        body = request.get_json(force=True) or {}
        code = supplier_merge_key(clean_text(supplier_code))
        if not code:
            return jsonify({"ok": False, "error": "Nhà cung cấp không hợp lệ"}), 400
        if code in SUPPLIER_POLICY_MANAGED:
            return jsonify({
                "ok": False,
                "error": "NCC này dùng quy tắc dồn cố định đã chốt với khách",
            }), 409
        with db_factory() as conn:
            conn.execute(
                """INSERT INTO supplier_rules(supplier_code,combine_kitchens,updated_at) VALUES(?,?,?)
                   ON CONFLICT(supplier_code) DO UPDATE SET combine_kitchens=excluded.combine_kitchens,updated_at=excluded.updated_at""",
                (code, 1 if body.get("combine_kitchens") else 0, now_iso()),
            )
            return jsonify({"ok": True})

    @app.put("/api/supplier-order-status/<int:batch_id>/<supplier_code>")
    def api_supplier_order_status(batch_id, supplier_code):
        body = request.get_json(force=True) or {}
        target_status = clean_text(body.get("status")).lower()
        if target_status not in {"ordered", "reopened"}:
            return jsonify({
                "ok": False,
                "error": "Trạng thái NCC chỉ được chuyển sang đã đặt hoặc mở lại",
                "code": "invalid_supplier_order_status",
            }), 400
        try:
            expected_revision = int(body.get("revision"))
        except (TypeError, ValueError):
            return jsonify({
                "ok": False,
                "error": "Thiếu revision trạng thái NCC; hãy tải lại màn hình",
                "code": "supplier_status_revision_required",
            }), 400
        if expected_revision < 0:
            return jsonify({
                "ok": False,
                "error": "Revision trạng thái NCC không hợp lệ",
                "code": "invalid_supplier_status_revision",
            }), 400
        supplier_key = supplier_merge_key(supplier_code)
        if not supplier_key:
            return jsonify({"ok": False, "error": "Nhà cung cấp không hợp lệ"}), 400

        with db_factory() as conn:
            conn.execute("BEGIN IMMEDIATE")
            if not conn.execute("SELECT 1 FROM batches WHERE id=?", (batch_id,)).fetchone():
                return jsonify({"ok": False, "error": "Không tìm thấy phiên đơn"}), 404
            current_payload = purchase_order_payload(conn, batch_id)
            if body.get("plan_hash") and body["plan_hash"] != current_payload["plan_hash"]:
                return jsonify(ok=False, error="Nội dung đặt NCC đã đổi trong lúc tạo ảnh; tải lại và sao chép ảnh mới", code="stale_supplier_plan"), 409
            checklist_item = next(
                (item for item in current_payload["checklist"] if item["supplier_key"] == supplier_key),
                None,
            )
            if checklist_item is None:
                return jsonify({
                    "ok": False,
                    "error": "NCC không còn dòng đặt hàng trong phiên này",
                    "code": "supplier_not_in_batch",
                }), 404
            current = conn.execute(
                """SELECT supplier_label,status,revision,ordered_at,reopened_at,updated_at
                   FROM supplier_order_statuses WHERE batch_id=? AND supplier_key=?""",
                (batch_id, supplier_key),
            ).fetchone()
            current_status = current["status"] if current else "pending"
            current_revision = int(current["revision"]) if current else 0
            if expected_revision != current_revision:
                return jsonify({
                    "ok": False,
                    "error": "Trạng thái NCC đã được thay đổi; hãy tải lại trước khi thao tác",
                    "code": "stale_supplier_order_status",
                    "supplier_key": supplier_key,
                    "status": current_status,
                    "revision": current_revision,
                }), 409
            allowed = (
                target_status == "ordered" and current_status in {"pending", "reopened"}
            ) or (
                target_status == "reopened" and current_status == "ordered"
            )
            if not allowed:
                return jsonify({
                    "ok": False,
                    "error": "Chuyển trạng thái NCC không hợp lệ; hãy tải lại màn hình",
                    "code": "invalid_supplier_order_transition",
                    "supplier_key": supplier_key,
                    "status": current_status,
                    "revision": current_revision,
                }), 409

            timestamp = now_iso()
            new_revision = current_revision + 1
            ordered_at = timestamp if target_status == "ordered" else (
                current["ordered_at"] if current else None
            )
            reopened_at = timestamp if target_status == "reopened" else (
                current["reopened_at"] if current else None
            )
            supplier_label = checklist_item["supplier"]
            if current:
                cursor = conn.execute(
                    """UPDATE supplier_order_statuses
                       SET supplier_label=?,status=?,revision=?,ordered_at=?,reopened_at=?,updated_at=?
                       WHERE batch_id=? AND supplier_key=? AND revision=?""",
                    (
                        supplier_label, target_status, new_revision, ordered_at, reopened_at,
                        timestamp, batch_id, supplier_key, current_revision,
                    ),
                )
                if cursor.rowcount != 1:
                    raise RuntimeError("optimistic supplier-order status update lost inside transaction")
            else:
                conn.execute(
                    """INSERT INTO supplier_order_statuses(
                           batch_id,supplier_key,supplier_label,status,revision,
                           ordered_at,reopened_at,updated_at
                       ) VALUES(?,?,?,?,?,?,?,?)""",
                    (
                        batch_id, supplier_key, supplier_label, target_status, new_revision,
                        ordered_at, reopened_at, timestamp,
                    ),
                )
            audit(
                conn, now_iso, "supplier_order.status", "ok",
                entity_type="supplier_order",
                entity_id=f"{batch_id}:{supplier_key}",
                metadata={
                    "batch_id": batch_id, "supplier_key": supplier_key,
                    "supplier": supplier_label, "previous_status": current_status,
                    "status": target_status, "expected_revision": expected_revision,
                    "revision": new_revision,
                },
            )
            return jsonify({
                "ok": True, "batch_id": batch_id, "supplier_key": supplier_key,
                "supplier": supplier_label, "previous_status": current_status,
                "status": target_status, "revision": new_revision,
                "ordered_at": ordered_at, "reopened_at": reopened_at,
                "updated_at": timestamp,
            })

    @app.post("/api/purchase-orders/import/preview")
    def api_purchase_order_import_preview():
        upload = request.files.get("file")
        if not upload or not upload.filename:
            return jsonify({"ok": False, "error": "Chưa chọn file đặt NCC đã chỉnh"}), 400
        try:
            batch_id = int(request.form.get("batch_id") or 0)
        except (TypeError, ValueError):
            batch_id = 0
        if batch_id <= 0:
            return jsonify({"ok": False, "error": "Chưa chọn phiên đơn cần nạp lại"}), 400
        filename = Path(upload.filename).name
        if Path(filename).suffix.lower() not in {".xlsx", ".xlsm"}:
            return jsonify({"ok": False, "error": "Chỉ nhận file .xlsx hoặc .xlsm"}), 400
        payload = upload.read(MAPPING_IMPORT_MAX_BYTES + 1)
        if len(payload) > MAPPING_IMPORT_MAX_BYTES:
            return jsonify({"ok": False, "error": "File Excel vượt quá giới hạn 10 MB"}), 413
        try:
            with zipfile.ZipFile(io.BytesIO(payload)) as archive:
                entries = archive.infolist()
                if (
                    len(entries) > 2_000
                    or sum(item.file_size for item in entries) > MAPPING_IMPORT_MAX_UNCOMPRESSED_BYTES
                ):
                    return jsonify({"ok": False, "error": "File Excel có cấu trúc quá lớn"}), 413
        except zipfile.BadZipFile:
            return jsonify({"ok": False, "error": "File Excel bị hỏng hoặc sai định dạng"}), 400

        cutoff = time.time() - MAPPING_IMPORT_TTL_SECONDS
        with PURCHASE_ORDER_IMPORT_LOCK:
            for old_token, item in list(PENDING_PURCHASE_ORDER_IMPORTS.items()):
                if item["created"] < cutoff:
                    PENDING_PURCHASE_ORDER_IMPORTS.pop(old_token, None)
        workbook = None
        formula_workbook = None
        try:
            workbook = load_workbook(
                io.BytesIO(payload), read_only=True, data_only=True, keep_links=False
            )
            formula_workbook = load_workbook(
                io.BytesIO(payload), read_only=True, data_only=False, keep_links=False
            )
            with db_factory() as conn:
                preview = parse_purchase_order_workbook(
                    conn, workbook, batch_id, formula_workbook=formula_workbook,
                )
                database_state_hash = purchase_order_database_state_hash(conn, batch_id)
        except ValueError as exc:
            return jsonify({"ok": False, "error": str(exc)}), 400
        except Exception:
            return jsonify({"ok": False, "error": "Không đọc được file đặt NCC"}), 400
        finally:
            if workbook is not None:
                workbook.close()
            if formula_workbook is not None:
                formula_workbook.close()

        token = uuid.uuid4().hex
        items = preview.pop("items")
        source_hash = preview.pop("content_hash")
        file_hash = hashlib.sha256(payload).hexdigest().upper()
        with PURCHASE_ORDER_IMPORT_LOCK:
            PENDING_PURCHASE_ORDER_IMPORTS[token] = {
                "created": time.time(), "filename": filename, "batch_id": batch_id,
                "source_hash": source_hash, "items": items,
                "file_hash": file_hash, "format": preview["format"],
                "database_state_hash": database_state_hash,
                "has_errors": not preview["can_confirm"],
            }
        return jsonify({
            "ok": True, "token": token, "filename": filename,
            "source_hash": source_hash,
            "expires_in_minutes": MAPPING_IMPORT_TTL_SECONDS // 60,
            **preview,
        })

    @app.post("/api/purchase-orders/import/confirm")
    def api_purchase_order_import_confirm():
        body = request.get_json(force=True) or {}
        if body.get("confirmed") is not True:
            return jsonify({"ok": False, "error": "Cần xác nhận trước khi ghi file đặt NCC"}), 400
        token = clean_text(body.get("token"))
        with PURCHASE_ORDER_IMPORT_LOCK:
            pending = PENDING_PURCHASE_ORDER_IMPORTS.pop(token, None)
        if not pending or time.time() - pending["created"] > MAPPING_IMPORT_TTL_SECONDS:
            return jsonify({"ok": False, "error": "Phiên xem trước đã hết hạn; hãy chọn lại file"}), 410
        if pending["has_errors"]:
            return jsonify({"ok": False, "error": "File còn dòng lỗi nên chưa thể nhập"}), 400
        try:
            with db_factory() as conn:
                conn.execute("BEGIN IMMEDIATE")
                batch_id = pending["batch_id"]
                if purchase_order_database_state_hash(conn, batch_id) != pending["database_state_hash"]:
                    return jsonify({
                        "ok": False,
                        "error": "Dữ liệu đơn đã thay đổi sau khi xem trước; hãy nạp lại file",
                    }), 409
                result = apply_purchase_order_preview(
                    conn,
                    batch_id=batch_id,
                    items=pending["items"],
                    source_hash=pending["source_hash"],
                    source_name=pending["filename"],
                    format_name=pending["format"],
                    now_iso=now_iso,
                )
            return jsonify({"ok": True, **result})
        except PurchaseOrderApplyError as error:
            return jsonify({
                "ok": False, "error": str(error), "code": error.code,
            }), error.status

    @app.get("/api/msmi/status")
    def api_msmi_status():
        try:
            return jsonify({"ok": True, **create_msmi_client().status()})
        except MsmiError as error:
            return jsonify({"ok": False, "connected": False, "read_only": True, "error": str(error)}), 502

    @app.post("/api/msmi/sync")
    def api_msmi_sync():
        # The former global sync had no date scope and shared one cursor across
        # unrelated periods. Keep the URL as an explicit compatibility guard;
        # all ingestion now starts from a date-bounded workbench batch.
        return jsonify({
            "ok": False,
            "error": "Hãy tạo phiên có Từ ngày–Đến ngày tại bàn làm việc hóa đơn rồi đồng bộ phiên đó",
            "code": "date_bounded_batch_required",
            "read_only": True,
        }), 409

    @app.get("/api/msmi/invoices")
    def api_msmi_invoices():
        with db_factory() as conn:
            return jsonify({"ok": True, "items": invoice_payload(
                conn, request.args.get("limit", 100, type=int), request.args.get("id", type=int)
            )})

    @app.put("/api/msmi/items/<int:item_id>/mapping")
    def api_msmi_mapping(item_id):
        body = request.get_json(force=True) or {}
        try:
            from invoice_mapping import InvoiceMappingError, save_mapping
        except ImportError:  # pragma: no cover - package invocation
            from .invoice_mapping import InvoiceMappingError, save_mapping
        with db_factory() as conn:
            conn.execute("BEGIN IMMEDIATE")
            try:
                result = save_mapping(
                    conn,
                    direction="input",
                    item_id=item_id,
                    product_code=body.get("product_code"),
                    now_iso=now_iso,
                )
            except InvoiceMappingError as error:
                return jsonify({"ok": False, "error": str(error), "code": error.code}), error.status
            item = conn.execute(
                "SELECT invoice_id FROM msmi_invoice_items WHERE id=?", (item_id,)
            ).fetchone()
            return jsonify({
                "ok": True,
                **result,
                "invoice": invoice_payload(conn, 1, item["invoice_id"])[0],
            })

    @app.post("/api/msmi/auto-mappings")
    def api_msmi_auto_mappings():
        try:
            from automatic_invoice_mapping import apply_automatic_input_mappings
            from invoice_workbench import validate_date_range
        except ImportError:
            from .automatic_invoice_mapping import apply_automatic_input_mappings
            from .invoice_workbench import validate_date_range
        workbook = None
        try:
            body = request.form if request.mimetype == "multipart/form-data" else (request.get_json(silent=True) or {})
            start, end = validate_date_range(body.get("from"), body.get("to"))
            if request.mimetype == "multipart/form-data":
                upload = request.files.get("file")
                if not upload or Path(upload.filename or "").suffix.lower() not in {".xlsx", ".xlsm"}:
                    return jsonify({"ok": False, "error": "Hãy chọn bảng kê nhập .xlsx hoặc .xlsm"}), 400
                payload = upload.read(MAPPING_IMPORT_MAX_BYTES + 1)
                if len(payload) > MAPPING_IMPORT_MAX_BYTES:
                    return jsonify({"ok": False, "error": "File Excel vượt quá giới hạn 10 MB"}), 413
                with zipfile.ZipFile(io.BytesIO(payload)) as archive:
                    entries = archive.infolist()
                    if len(entries) > 2000 or sum(item.file_size for item in entries) > MAPPING_IMPORT_MAX_UNCOMPRESSED_BYTES:
                        return jsonify({"ok": False, "error": "File Excel có cấu trúc quá lớn"}), 413
                workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True, keep_links=False)
            with db_factory() as conn:
                conn.execute("BEGIN IMMEDIATE")
                result = apply_automatic_input_mappings(
                    conn, tenant=setting_get(conn, "tenant_code", "TDP"),
                    date_from=start, date_to=end, now_iso=now_iso, workbook=workbook,
                )
            return jsonify({"ok": True, **result})
        except (ValueError, zipfile.BadZipFile) as error:
            return jsonify({"ok": False, "error": str(error), "code": "automatic_mapping_invalid"}), 400
        except Exception:
            return jsonify({"ok": False, "error": "Chưa tự ghép được mã; dữ liệu lượt này đã hoàn tác", "code": "automatic_mapping_failed"}), 409
        finally:
            if workbook is not None:
                workbook.close()

    @app.get("/api/msmi/suggested-mappings/preview")
    def api_msmi_safe_suggestions_preview():
        with db_factory() as conn:
            try:
                plan = safe_input_mapping_suggestion_plan(
                    conn,
                    tenant=setting_get(conn, "tenant_code", "TDP"),
                    date_from=request.args.get("from"),
                    date_to=request.args.get("to"),
                )
            except ValueError as error:
                return jsonify({"ok": False, "error": str(error), "code": "invalid_period"}), 400
            return jsonify({"ok": True, **{key: value for key, value in plan.items() if not key.startswith("_")}})

    @app.post("/api/msmi/suggested-mappings")
    def api_msmi_safe_suggestions_apply():
        body = request.get_json(force=True) or {}
        if body.get("confirmed") is not True:
            return jsonify({
                "ok": False,
                "error": "Cần xác nhận rõ trước khi ghi nhớ mã hàng khớp chắc chắn",
                "code": "confirmation_required",
            }), 400
        try:
            from invoice_mapping import InvoiceMappingError, save_mapping
        except ImportError:  # pragma: no cover - package invocation
            from .invoice_mapping import InvoiceMappingError, save_mapping
        with db_factory() as conn:
            conn.execute("BEGIN IMMEDIATE")
            try:
                plan = safe_input_mapping_suggestion_plan(
                    conn,
                    tenant=setting_get(conn, "tenant_code", "TDP"),
                    date_from=body.get("from"),
                    date_to=body.get("to"),
                )
            except ValueError as error:
                return jsonify({"ok": False, "error": str(error), "code": "invalid_period"}), 400
            if not plan["safe_rules"]:
                return jsonify({
                    "ok": False,
                    "error": "Không còn dòng nào khớp duy nhất cả tên hàng và đơn vị trong kỳ đã chọn",
                    "code": "no_safe_suggestions",
                }), 400
            if str(body.get("snapshot") or "") != plan["snapshot"]:
                return jsonify({
                    "ok": False,
                    "error": "Dữ liệu ghép mã đã thay đổi sau lúc kiểm tra; hãy kiểm tra lại trước khi xác nhận",
                    "code": "stale_suggestion_preview",
                }), 409

            conn.execute("SAVEPOINT safe_suggestion_bulk")
            applied_rules = 0
            applied_lines_all_periods = 0
            try:
                for item in plan["_representatives"]:
                    result = save_mapping(
                        conn,
                        direction="input",
                        item_id=item["item_id"],
                        product_code=item["product_code"],
                        now_iso=now_iso,
                    )
                    if result["requires_unit_conversion"]:
                        raise InvoiceMappingError(
                            "Đơn vị nguồn đã thay đổi; hãy kiểm tra lại trước khi ghép mã hàng loạt",
                            code="stale_suggestion_preview",
                            status=409,
                        )
                    applied_rules += 1
                    applied_lines_all_periods += int(result["applied_lines"])
            except Exception as error:
                conn.execute("ROLLBACK TO safe_suggestion_bulk")
                conn.execute("RELEASE safe_suggestion_bulk")
                status = error.status if isinstance(error, InvoiceMappingError) else 409
                code = error.code if isinstance(error, InvoiceMappingError) else "bulk_mapping_failed"
                return jsonify({"ok": False, "error": str(error), "code": code}), status
            conn.execute("RELEASE safe_suggestion_bulk")
            status_counts = conn.execute(
                """SELECT receipt_status,COUNT(*) count FROM msmi_invoices
                    WHERE tenant=? AND invoice_type='INPUT_ELECTRONIC_INVOICE'
                      AND invoice_date>=? AND invoice_date<=?
                    GROUP BY receipt_status""",
                (plan["tenant"], plan["date_from"], plan["date_to"]),
            ).fetchall()
            counts = {str(row["receipt_status"]): int(row["count"]) for row in status_counts}
            audit(
                conn, now_iso, "msmi.mapping_suggestions_bulk", "ok",
                entity_type="invoice_period", entity_id=f'{plan["date_from"]}:{plan["date_to"]}',
                metadata={
                    "safe_rules": applied_rules,
                    "safe_lines_in_period": plan["safe_lines_in_period"],
                    "applied_lines_all_periods": applied_lines_all_periods,
                    "ready_invoices_in_period": counts.get("ready", 0),
                },
            )
            return jsonify({
                "ok": True,
                "applied_rules": applied_rules,
                "mapped_lines_in_period": plan["safe_lines_in_period"],
                "mapped_lines_all_periods": applied_lines_all_periods,
                "ready_invoices_in_period": counts.get("ready", 0),
                "pending_invoices_in_period": counts.get("pending_mapping", 0),
                "stock_changed": False,
            })

    @app.post("/api/msmi/legacy-mappings/preview")
    def api_msmi_legacy_mappings_preview():
        upload = request.files.get("file")
        if not upload or not upload.filename:
            return jsonify({"ok": False, "error": "Chưa chọn file bảng kê nhập của phần mềm cũ"}), 400
        filename = Path(upload.filename).name
        if Path(filename).suffix.lower() not in {".xlsx", ".xlsm"}:
            return jsonify({"ok": False, "error": "Chỉ nhận file .xlsx hoặc .xlsm"}), 400
        payload = upload.read(MAPPING_IMPORT_MAX_BYTES + 1)
        if len(payload) > MAPPING_IMPORT_MAX_BYTES:
            return jsonify({"ok": False, "error": "File Excel vượt quá giới hạn 10 MB"}), 413
        try:
            with zipfile.ZipFile(io.BytesIO(payload)) as archive:
                entries = archive.infolist()
                expanded_size = sum(item.file_size for item in entries)
                if len(entries) > 2_000 or expanded_size > MAPPING_IMPORT_MAX_UNCOMPRESSED_BYTES:
                    return jsonify({"ok": False, "error": "File Excel có cấu trúc quá lớn để đọc an toàn"}), 413
        except zipfile.BadZipFile:
            return jsonify({"ok": False, "error": "File Excel bị hỏng hoặc không đúng định dạng"}), 400
        try:
            try:
                from legacy_invoice_mapping import legacy_input_mapping_plan
            except ImportError:  # pragma: no cover - package invocation
                from .legacy_invoice_mapping import legacy_input_mapping_plan
            workbook = load_workbook(
                io.BytesIO(payload), read_only=True, data_only=True, keep_links=False,
            )
            with db_factory() as conn:
                plan = legacy_input_mapping_plan(
                    conn,
                    workbook,
                    tenant=setting_get(conn, "tenant_code", "TDP"),
                    date_from=request.form.get("from"),
                    date_to=request.form.get("to"),
                )
        except ValueError as error:
            return jsonify({"ok": False, "error": str(error), "code": "invalid_legacy_mapping_file"}), 400
        except Exception:
            return jsonify({"ok": False, "error": "Không đọc được bảng kê nhập; hãy kiểm tra lại file"}), 400
        finally:
            if "workbook" in locals():
                workbook.close()

        cutoff = time.time() - MAPPING_IMPORT_TTL_SECONDS
        with LEGACY_INVOICE_MAPPING_IMPORT_LOCK:
            for old_token, item in list(PENDING_LEGACY_INVOICE_MAPPING_IMPORTS.items()):
                if item["created"] < cutoff:
                    PENDING_LEGACY_INVOICE_MAPPING_IMPORTS.pop(old_token, None)
            token = uuid.uuid4().hex
            PENDING_LEGACY_INVOICE_MAPPING_IMPORTS[token] = {
                "created": time.time(),
                "filename": filename,
                "payload": payload,
                "date_from": plan["date_from"],
                "date_to": plan["date_to"],
                "snapshot": plan["snapshot"],
            }
        return jsonify({
            "ok": True,
            "token": token,
            "filename": filename,
            "expires_in_minutes": MAPPING_IMPORT_TTL_SECONDS // 60,
            **{key: value for key, value in plan.items() if not key.startswith("_")},
        })

    @app.post("/api/msmi/legacy-mappings/confirm")
    def api_msmi_legacy_mappings_confirm():
        body = request.get_json(force=True) or {}
        if body.get("confirmed") is not True:
            return jsonify({"ok": False, "error": "Cần xác nhận trước khi ghi nhớ mã từ bảng kê cũ"}), 400
        token = clean_text(body.get("token"))
        with LEGACY_INVOICE_MAPPING_IMPORT_LOCK:
            pending = PENDING_LEGACY_INVOICE_MAPPING_IMPORTS.pop(token, None)
        if not pending or time.time() - pending["created"] > MAPPING_IMPORT_TTL_SECONDS:
            return jsonify({"ok": False, "error": "Bản xem trước đã hết hạn; hãy chọn lại file"}), 410
        try:
            try:
                from legacy_invoice_mapping import legacy_input_mapping_plan
                from invoice_mapping import InvoiceMappingError, save_mapping
            except ImportError:  # pragma: no cover - package invocation
                from .legacy_invoice_mapping import legacy_input_mapping_plan
                from .invoice_mapping import InvoiceMappingError, save_mapping
            workbook = load_workbook(
                io.BytesIO(pending["payload"]), read_only=True, data_only=True, keep_links=False,
            )
            with db_factory() as conn:
                conn.execute("BEGIN IMMEDIATE")
                plan = legacy_input_mapping_plan(
                    conn,
                    workbook,
                    tenant=setting_get(conn, "tenant_code", "TDP"),
                    date_from=pending["date_from"],
                    date_to=pending["date_to"],
                )
                if plan["snapshot"] != pending["snapshot"]:
                    raise ValueError(
                        "Dữ liệu hóa đơn hoặc danh mục đã thay đổi sau lúc xem trước; hãy chọn lại file"
                    )
                if not plan["can_confirm"]:
                    raise ValueError("Không có quy tắc khớp chắc chắn nào để xác nhận")
                mapped_before_period = conn.execute(
                    """SELECT COUNT(*) FROM msmi_invoice_items li
                        JOIN msmi_invoices i ON i.id=li.invoice_id
                        WHERE i.tenant=? AND i.invoice_type='INPUT_ELECTRONIC_INVOICE'
                          AND i.invoice_date>=? AND i.invoice_date<=?
                          AND li.inventory_eligible=1 AND li.mapping_status='mapped'""",
                    (plan["tenant"], plan["date_from"], plan["date_to"]),
                ).fetchone()[0]
                conn.execute("SAVEPOINT legacy_mapping_bulk")
                applied_rules = 0
                applied_lines_all_periods = 0
                try:
                    for item in plan["_representatives"]:
                        result = save_mapping(
                            conn,
                            direction="input",
                            item_id=item["item_id"],
                            product_code=item["product_code"],
                            now_iso=now_iso,
                        )
                        if result["requires_unit_conversion"]:
                            raise InvoiceMappingError(
                                "Đơn vị đã thay đổi sau lúc xem trước",
                                code="stale_legacy_mapping_preview",
                                status=409,
                            )
                        applied_rules += 1
                        applied_lines_all_periods += int(result["applied_lines"])
                except Exception:
                    conn.execute("ROLLBACK TO legacy_mapping_bulk")
                    conn.execute("RELEASE legacy_mapping_bulk")
                    raise
                conn.execute("RELEASE legacy_mapping_bulk")
                status_counts = conn.execute(
                    """SELECT receipt_status,COUNT(*) count FROM msmi_invoices
                        WHERE tenant=? AND invoice_type='INPUT_ELECTRONIC_INVOICE'
                          AND invoice_date>=? AND invoice_date<=?
                        GROUP BY receipt_status""",
                    (plan["tenant"], plan["date_from"], plan["date_to"]),
                ).fetchall()
                counts = {str(row["receipt_status"]): int(row["count"]) for row in status_counts}
                mapped_after_period = conn.execute(
                    """SELECT COUNT(*) FROM msmi_invoice_items li
                        JOIN msmi_invoices i ON i.id=li.invoice_id
                        WHERE i.tenant=? AND i.invoice_type='INPUT_ELECTRONIC_INVOICE'
                          AND i.invoice_date>=? AND i.invoice_date<=?
                          AND li.inventory_eligible=1 AND li.mapping_status='mapped'""",
                    (plan["tenant"], plan["date_from"], plan["date_to"]),
                ).fetchone()[0]
                audit(
                    conn, now_iso, "msmi.legacy_mapping_import", "ok",
                    entity_type="invoice_period", entity_id=f'{plan["date_from"]}:{plan["date_to"]}',
                    metadata={
                        "filename": pending["filename"],
                        "safe_rules": applied_rules,
                        "safe_lines_in_period": plan["safe_lines_in_period"],
                        "applied_lines_all_periods": applied_lines_all_periods,
                    },
                )
        except InvoiceMappingError as error:
            return jsonify({"ok": False, "error": str(error), "code": error.code}), error.status
        except ValueError as error:
            return jsonify({"ok": False, "error": str(error), "code": "stale_legacy_mapping_preview"}), 409
        except Exception:
            return jsonify({
                "ok": False,
                "error": "Không thể ghi nhớ mã từ bảng kê cũ; dữ liệu chưa được thay đổi",
                "code": "legacy_mapping_failed",
            }), 409
        finally:
            if "workbook" in locals():
                workbook.close()
        return jsonify({
            "ok": True,
            "applied_rules": applied_rules,
            "mapped_lines_in_period": int(mapped_after_period) - int(mapped_before_period),
            "mapped_lines_all_periods": applied_lines_all_periods,
            "ready_invoices_in_period": counts.get("ready", 0),
            "pending_invoices_in_period": counts.get("pending_mapping", 0),
            "stock_changed": False,
        })

    @app.post("/api/msmi/invoices/<int:invoice_id>/suggested-mappings")
    def api_msmi_suggested_mappings(invoice_id):
        try:
            from invoice_mapping import InvoiceMappingError, save_mapping
        except ImportError:  # pragma: no cover - package invocation
            from .invoice_mapping import InvoiceMappingError, save_mapping
        with db_factory() as conn:
            conn.execute("BEGIN IMMEDIATE")
            invoice = conn.execute("SELECT * FROM msmi_invoices WHERE id=?", (invoice_id,)).fetchone()
            if not invoice:
                return jsonify({"ok": False, "error": "Không tìm thấy hóa đơn đầu vào"}), 404
            if invoice["receipt_status"] == "posted":
                return jsonify({
                    "ok": False,
                    "error": "Phiếu nhập đã ghi kho; không được đổi ghép mã",
                }), 409
            suggestions = unique_product_name_suggestions(conn)
            items = conn.execute(
                """SELECT * FROM msmi_invoice_items
                   WHERE invoice_id=? AND inventory_eligible=1 AND mapping_status!='mapped'
                   ORDER BY line_index""",
                (invoice_id,),
            ).fetchall()
            applied = 0
            for item in items:
                suggestion = suggestions.get(mapping_key(item["source_item_name"]))
                if not suggestion:
                    continue
                product_code = suggestion["code"]
                try:
                    save_mapping(
                        conn,
                        direction="input",
                        item_id=item["id"],
                        product_code=product_code,
                        now_iso=now_iso,
                    )
                    applied += 1
                except InvoiceMappingError:
                    # A suggestion must never bypass unit/frozen/scope gates.
                    continue
            if not applied:
                return jsonify({
                    "ok": False,
                    "error": "Không có mã nào khớp duy nhất theo tên để xác nhận tự động",
                }), 400
            remaining = conn.execute(
                """SELECT COUNT(*) n FROM msmi_invoice_items
                   WHERE invoice_id=? AND inventory_eligible=1 AND mapping_status!='mapped'""",
                (invoice_id,),
            ).fetchone()["n"]
            conn.execute(
                "UPDATE msmi_invoices SET receipt_status=?,updated_at=? WHERE id=?",
                ("ready" if remaining == 0 else "pending_mapping", now_iso(), invoice_id),
            )
            audit(conn, now_iso, "msmi.mapping_suggestions", "ok",
                  entity_type="msmi_invoice", entity_id=invoice_id,
                  metadata={"applied": applied, "remaining": remaining})
            return jsonify({
                "ok": True,
                "applied": applied,
                "remaining": remaining,
                "invoice": invoice_payload(conn, 1, invoice_id)[0],
            })

    @app.post("/api/msmi/invoices/<int:invoice_id>/receipt")
    def api_msmi_receipt(invoice_id):
        try:
            from invoice_receipt import InvoiceReceiptError, create_input_receipt
        except ImportError:  # pragma: no cover - package invocation
            from .invoice_receipt import InvoiceReceiptError, create_input_receipt
        with db_factory() as conn:
            conn.execute("BEGIN IMMEDIATE")
            try:
                return jsonify({"ok": True, **create_input_receipt(conn, invoice_id, now_iso)})
            except InvoiceReceiptError as error:
                return jsonify({"ok": False, "error": str(error), "code": error.code}), error.status

    @app.get("/api/outgoing-invoices/readiness/<int:batch_id>")
    def api_outgoing_invoice_readiness(batch_id):
        with db_factory() as conn:
            try:
                payload = outgoing_invoice_readiness_payload(conn, batch_id)
            except OutgoingReadinessError as exc:
                return jsonify({"ok": False, "error": str(exc), "code": exc.code}), exc.status
            return jsonify({"ok": True, **payload})

    @app.get("/api/outgoing-invoices/readiness/<int:batch_id>/stock/<path:product_code>")
    def api_outgoing_product_stock_trace(batch_id, product_code):
        with db_factory() as conn:
            conn.execute('BEGIN')
            try:
                return jsonify(ok=True, **batch_product_stock_trace(conn, batch_id, product_code))
            except OutgoingReadinessError as exc:
                return jsonify(ok=False, error=str(exc), code=exc.code), exc.status

    @app.get("/api/outgoing-invoices/shortages")
    def api_outgoing_invoice_shortages():
        try:
            with db_factory() as conn:
                payload = period_shortage_payload(
                    conn,
                    request.args.get("from"),
                    request.args.get("to"),
                    request.args.get("contractor", ""),
                )
        except OutgoingReadinessError as exc:
            return jsonify({"ok": False, "error": str(exc), "code": exc.code}), exc.status
        return jsonify({"ok": True, **payload})

    @app.get("/api/outgoing-invoices/shortages/export")
    def api_export_outgoing_invoice_shortages():
        try:
            with db_factory() as conn:
                payload = period_shortage_payload(
                    conn,
                    request.args.get("from"),
                    request.args.get("to"),
                    request.args.get("contractor", ""),
                )
            workbook = shortage_workbook_bytes(payload)
        except OutgoingReadinessError as exc:
            return jsonify({"ok": False, "error": str(exc), "code": exc.code}), exc.status
        suffix = re.sub(r"[^A-Za-z0-9_-]+", "_", payload["contractor"]).strip("_")
        filename = f"DANH_SACH_THIEU_HDDV_{payload['date_from']}_{payload['date_to']}"
        if suffix:
            filename += f"_{suffix[:40]}"
        return send_file(
            io.BytesIO(workbook),
            as_attachment=True,
            download_name=filename + ".xlsx",
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    @app.post("/api/outgoing-invoices/draft/<int:batch_id>")
    def api_create_outgoing_drafts(batch_id):
        with db_factory() as conn:
            # Serialize the stock check and the reservation writes.  A deferred
            # transaction lets two batches both observe the same stock (or an
            # order mutate after we read it) before either draft exists.
            conn.execute("BEGIN IMMEDIATE")
            try:
                result = create_partial_outgoing_drafts(conn, batch_id, now_iso)
            except OutgoingReadinessError as exc:
                conn.rollback()
                return jsonify({"ok": False, "error": str(exc), "code": exc.code}), exc.status
            except ValueError as exc:
                conn.rollback()
                return jsonify({"ok": False, "error": str(exc)}), 400
            return jsonify({"ok": True, **result})
            batch = conn.execute("SELECT * FROM batches WHERE id=?", (batch_id,)).fetchone()
            if not batch:
                return jsonify({"ok": False, "error": "Không tìm thấy phiên đơn"}), 404
            if batch["status"] != "approved":
                return jsonify({"ok": False, "error": "Phải duyệt phiên đơn trước khi lập hóa đơn đầu ra"}), 400
            orders = [dict(row) for row in conn.execute(
                "SELECT * FROM orders WHERE batch_id=? ORDER BY contractor,id", (batch_id,)
            )]
            orders = [item for item in orders if net_delivered(item) > 1e-9]
            if not orders:
                return jsonify({
                    "ok": False,
                    "error": "Phiên không còn lượng thực giao dương để lập hóa đơn đầu ra",
                }), 409
            issued_contractors = {
                row["contractor"] for row in conn.execute(
                    "SELECT contractor FROM outgoing_invoice_drafts WHERE batch_id=? AND status='issued'",
                    (batch_id,),
                )
            }
            demand_orders = [
                item for item in orders if item["contractor"] not in issued_contractors
            ]
            stock = inventory_lookup(conn, batch["work_date"])
            available = {code: max(row["available_qty"], 0) for code, row in stock.items()}
            # Khi người dùng bấm tạo lại, tồn khả dụng đã trừ phần chính dự thảo
            # của batch này đang giữ. Cộng phần giữ đó lại trước khi kiểm tra để
            # thao tác idempotent, rồi phía dưới thay thế đúng các dòng reserve.
            for reserved in conn.execute(
                """SELECT t.product_code,SUM(t.qty_out) qty
                   FROM inventory_transactions t
                   JOIN outgoing_invoice_drafts d ON CAST(d.id AS TEXT)=t.source_id
                   WHERE t.source_type='OUTGOING_DRAFT' AND t.status='reserved'
                     AND d.batch_id=? AND d.status!='issued'
                   GROUP BY t.product_code""",
                (batch_id,),
            ):
                available[reserved["product_code"]] = (
                    available.get(reserved["product_code"], 0) + reserved["qty"]
                )
            shortages = defaultdict(lambda: {"required": 0, "available": 0, "product_name": ""})
            for item in demand_orders:
                qty = net_delivered(item)
                code = item["product_code"]
                have = available.get(code, 0)
                if qty > have + 1e-9:
                    shortages[code]["required"] += qty - have
                    shortages[code]["available"] = have
                    shortages[code]["product_name"] = item["product_name"]
                available[code] = max(have - qty, 0)
            if shortages:
                return jsonify({
                    "ok": False,
                    "error": "Không đủ tồn hóa đơn để lập toàn bộ hóa đơn đầu ra",
                    "shortages": [{"product_code": code, **values} for code, values in shortages.items()],
                }), 409

            created = []
            for contractor, group in _group_rows(orders, "contractor").items():
                existing = conn.execute(
                    "SELECT * FROM outgoing_invoice_drafts WHERE batch_id=? AND contractor=?",
                    (batch_id, contractor),
                ).fetchone()
                if existing and existing["status"] == "issued":
                    created.append(dict(existing))
                    continue
                if existing and existing["minvoice_status"] in {"saved", "saving", "unknown"}:
                    created.append(dict(existing))
                    continue
                calculated_lines = []
                for item in group:
                    qty = net_delivered(item)
                    unit_price = vnd_round(number_value(item["sell_price"]))
                    amount = vnd_product(qty, unit_price)
                    vat_percent = invoice_tax_percent(item["tax"])
                    line_tax = 0 if vat_percent <= 0 else vnd_product(amount, vat_percent / 100)
                    calculated_lines.append((item, qty, unit_price, amount, line_tax))
                subtotal = sum(line[3] for line in calculated_lines)
                tax_amount = sum(line[4] for line in calculated_lines)
                total = subtotal + tax_amount
                conn.execute(
                    """INSERT INTO outgoing_invoice_drafts(
                        batch_id,contractor,invoice_date,status,subtotal,tax_amount,total_amount,
                        created_at,external_key_uuid
                    ) VALUES(?,?,?,'draft',?,?,?,?,?)
                    ON CONFLICT(batch_id,contractor) DO UPDATE SET
                        invoice_date=excluded.invoice_date,status='draft',subtotal=excluded.subtotal,
                        tax_amount=excluded.tax_amount,total_amount=excluded.total_amount,
                        external_key_uuid=CASE
                          WHEN TRIM(COALESCE(outgoing_invoice_drafts.external_key_uuid,''))=''
                          THEN excluded.external_key_uuid
                          ELSE outgoing_invoice_drafts.external_key_uuid END""",
                    (
                        batch_id, contractor, batch["work_date"], subtotal, tax_amount, total,
                        now_iso(), uuid.uuid4().hex.upper(),
                    ),
                )
                draft = conn.execute(
                    "SELECT * FROM outgoing_invoice_drafts WHERE batch_id=? AND contractor=?",
                    (batch_id, contractor),
                ).fetchone()
                draft_id = draft["id"]
                conn.execute("DELETE FROM outgoing_invoice_lines WHERE draft_id=?", (draft_id,))
                conn.execute(
                    "UPDATE inventory_transactions SET status='cancelled',updated_at=? "
                    "WHERE source_type='OUTGOING_DRAFT' AND source_id=? AND status='reserved'",
                    (now_iso(), str(draft_id)),
                )
                for item, qty, unit_price, amount, _ in calculated_lines:
                    invoice_name_row = conn.execute(
                        "SELECT invoice_name FROM outgoing_product_names WHERE product_code=?",
                        (item["product_code"],),
                    ).fetchone()
                    invoice_name = invoice_name_row["invoice_name"] if invoice_name_row else item["product_name"]
                    conn.execute(
                        """INSERT INTO outgoing_invoice_lines(
                            draft_id,order_id,product_code,product_name,qty,unit,unit_price,tax,
                            invoice_nature,amount
                        ) VALUES(?,?,?,?,?,?,?,?,?,?)""",
                        (draft_id, item["id"], item["product_code"], invoice_name, qty,
                         item["unit"], unit_price, item["tax"], item.get("invoice_nature") or "1", amount),
                    )
                    conn.execute(
                        """INSERT INTO inventory_transactions(
                            txn_date,product_code,qty_in,qty_out,unit_cost,source_type,source_id,source_line,
                            kitchen,status,note,created_at,updated_at
                        ) VALUES(?,?,0,?,?, 'OUTGOING_DRAFT',?,?,?,'reserved',?,?,?)
                        ON CONFLICT(source_type,source_id,source_line) DO UPDATE SET
                            txn_date=excluded.txn_date,product_code=excluded.product_code,qty_out=excluded.qty_out,
                            unit_cost=excluded.unit_cost,kitchen=excluded.kitchen,status='reserved',updated_at=excluded.updated_at""",
                        (batch["work_date"], item["product_code"], qty, item["buy_price"], str(draft_id),
                         str(item["id"]), item["kitchen"], f"Dự thảo hóa đơn {contractor}", now_iso(), now_iso()),
                    )
                created.append(dict(draft))
            audit(conn, now_iso, "outgoing.draft", "ok", entity_type="batch", entity_id=batch_id,
                  metadata={"drafts": len(created)})
            return jsonify({"ok": True, "drafts": created, "requires_user_sign_and_issue": True})

    @app.put("/api/outgoing-product-names/<product_code>")
    def api_outgoing_product_name(product_code):
        body = request.get_json(force=True) or {}
        code = clean_text(product_code).upper()
        invoice_name = clean_text(body.get("invoice_name"))
        if not code or not invoice_name:
            return jsonify({"ok": False, "error": "Cần mã hàng và tên xuất hóa đơn"}), 400
        with db_factory() as conn:
            product = conn.execute("SELECT code FROM products WHERE code=?", (code,)).fetchone()
            if not product:
                return jsonify({"ok": False, "error": "Mã hàng chưa có trong danh mục"}), 404
            conn.execute(
                """INSERT INTO outgoing_product_names(product_code,invoice_name,updated_at) VALUES(?,?,?)
                   ON CONFLICT(product_code) DO UPDATE SET invoice_name=excluded.invoice_name,updated_at=excluded.updated_at""",
                (code, invoice_name, now_iso()),
            )
            audit(conn, now_iso, "outgoing.name_mapping", "ok", entity_type="product", entity_id=code)
            return jsonify({"ok": True, "product_code": code, "invoice_name": invoice_name})

    @app.post("/api/outgoing-invoices/<int:draft_id>/confirm-issued")
    def api_confirm_outgoing_issued(draft_id):
        body = request.get_json(silent=True) or {}
        if not body.get("confirmed"):
            return jsonify({"ok": False, "error": "Cần xác nhận người dùng đã ký và phát hành hóa đơn"}), 400
        with db_factory() as conn:
            conn.execute("BEGIN IMMEDIATE")
            draft = conn.execute("SELECT * FROM outgoing_invoice_drafts WHERE id=?", (draft_id,)).fetchone()
            if not draft:
                return jsonify({"ok": False, "error": "Không tìm thấy dự thảo hóa đơn"}), 404
            if draft["status"] == "cancelled":
                return jsonify({"ok": False, "error": "Dự thảo đã hủy; cần tạo lại trước khi xác nhận phát hành"}), 409
            if draft["status"] == "issued":
                if (clean_text(body.get('invoice_number')) != clean_text(draft['issued_invoice_number'])
                        or clean_text(body.get('invoice_series')).upper() != clean_text(draft['issued_invoice_series']).upper()
                        or as_date(body.get('invoice_date')) != (draft['issued_invoice_date'] or draft['invoice_date'])):
                    return jsonify(ok=False, error='Dự thảo này đã ghi nhận một số hóa đơn khác. Tải lại để kiểm tra.'), 409
                return jsonify({
                    "ok": True,
                    "idempotent": True,
                    "invoice_number": draft["issued_invoice_number"] or "",
                    "invoice_series": draft["issued_invoice_series"] or "",
                    "invoice_date": draft["issued_invoice_date"] or draft["invoice_date"],
                })
            invoice_number = unicodedata.normalize(
                "NFKC", clean_text(body.get("invoice_number"))
            )
            invoice_series = unicodedata.normalize(
                "NFKC", clean_text(body.get("invoice_series") or draft["minvoice_series"])
            ).upper()
            invoice_date = as_date(body.get("invoice_date"))
            if not re.fullmatch(r"\d{1,20}", invoice_number):
                return jsonify({"ok": False, "error": "Số hóa đơn đã phát hành chỉ được gồm 1–20 chữ số"}), 400
            if not invoice_series or len(invoice_series) > 50 or any(ord(ch) < 32 for ch in invoice_series):
                return jsonify({"ok": False, "error": "Cần ký hiệu hóa đơn hợp lệ, tối đa 50 ký tự"}), 400
            if not invoice_date:
                return jsonify({"ok": False, "error": "Ngày hóa đơn phải hợp lệ dạng YYYY-MM-DD"}), 400
            tax_groups = {invoice_tax_percent(r['tax']) for r in conn.execute(
                'SELECT tax FROM outgoing_invoice_lines WHERE draft_id=?', (draft_id,))}
            if len(tax_groups) > 1 and clean_text(draft['minvoice_status']) not in {'saved', 'saving', 'unknown'}:
                return jsonify(ok=False, error='Dự thảo cũ có nhiều nhóm thuế. Bấm Tính lại dự thảo trước khi tạo file.',
                               code='invoice_tax_split_required'), 409
            try:
                validate_issued_draft_stock(conn, draft_id, invoice_date, invoice_series, invoice_number)
            except OutgoingReadinessError as exc:
                return jsonify({"ok": False, "error": str(exc), "code": exc.code}), exc.status
            snapshot = stored_invoice_snapshot(draft)
            if not snapshot:
                if clean_text(draft["minvoice_status"]) in {"saved", "saving", "unknown"}:
                    return jsonify({
                        "ok": False,
                        "error": "Bản nháp M-Invoice cũ chưa có hồ sơ bất biến. Cần đối chiếu hóa đơn rồi xác nhận khóa hồ sơ cũ trước khi ghi phát hành",
                    }), 409
                try:
                    snapshot, _ = current_invoice_snapshot(conn, draft)
                except ValueError as exc:
                    return jsonify({"ok": False, "error": str(exc)}), 409
            duplicate = conn.execute(
                """SELECT id FROM outgoing_invoice_drafts
                   WHERE id!=? AND UPPER(TRIM(COALESCE(issued_invoice_series,'')))=?
                     AND TRIM(COALESCE(issued_invoice_number,''))=?""",
                (draft_id, invoice_series, invoice_number),
            ).fetchone()
            if duplicate:
                return jsonify({"ok": False, "error": "Ký hiệu và số hóa đơn này đã được ghi nhận"}), 409
            if draft["status"] != "issued":
                conn.execute(
                    """UPDATE outgoing_invoice_drafts
                       SET status='issued',issued_at=?,issued_invoice_number=?,
                           issued_invoice_series=?,issued_invoice_date=?,buyer_name_snapshot=?,
                           buyer_tax_code_snapshot=?,buyer_address_snapshot=?,buyer_email_snapshot=?,
                           company_name_snapshot=?,company_tax_code_snapshot=?,company_address_snapshot=?,
                           payment_requester_snapshot=?,payment_bank_name_snapshot=?,
                           payment_bank_account_snapshot=? WHERE id=? AND status='draft'""",
                    (
                        now_iso(), invoice_number, invoice_series, invoice_date,
                        *snapshot_update_values(snapshot),
                        draft_id,
                    ),
                )
                if conn.execute("SELECT changes() n").fetchone()["n"] != 1:
                    return jsonify({"ok": False, "error": "Trạng thái hóa đơn vừa thay đổi; vui lòng tải lại"}), 409
                conn.execute(
                    "UPDATE inventory_transactions SET status='posted',updated_at=? "
                    "WHERE source_type='OUTGOING_DRAFT' AND source_id=? AND status='reserved'",
                    (now_iso(), str(draft_id)),
                )
                audit(conn, now_iso, "outgoing.confirm_issued", "ok", entity_type="outgoing_invoice", entity_id=draft_id,
                      metadata={"invoice_number": invoice_number, "invoice_series": invoice_series,
                                "invoice_date": invoice_date})
            return jsonify({"ok": True, "idempotent": False, "invoice_number": invoice_number,
                            "invoice_series": invoice_series, "invoice_date": invoice_date})

    @app.post("/api/outgoing-invoices/<int:draft_id>/capture-legacy-snapshot")
    def api_capture_legacy_invoice_snapshot(draft_id):
        body = request.get_json(silent=True) or {}
        if body.get("confirmed_profile_matches_invoice") is not True:
            return jsonify({
                "ok": False,
                "error": "Cần đối chiếu hồ sơ hiện tại với bản nháp/hóa đơn cũ rồi xác nhận rõ",
            }), 400
        with db_factory() as conn:
            conn.execute("BEGIN IMMEDIATE")
            draft = conn.execute(
                "SELECT * FROM outgoing_invoice_drafts WHERE id=?", (draft_id,),
            ).fetchone()
            if not draft:
                return jsonify({"ok": False, "error": "Không tìm thấy hóa đơn"}), 404
            existing = stored_invoice_snapshot(draft)
            if existing:
                return jsonify({"ok": True, "idempotent": True})
            if draft["status"] != "issued" and clean_text(draft["minvoice_status"]) not in {
                "saved", "saving", "unknown",
            }:
                return jsonify({
                    "ok": False,
                    "error": "Chỉ khóa hồi tố cho hóa đơn đã phát hành hoặc bản nháp M-Invoice cũ",
                }), 409
            try:
                snapshot, _ = current_invoice_snapshot(conn, draft)
            except ValueError as exc:
                return jsonify({"ok": False, "error": str(exc)}), 409
            conn.execute(
                """UPDATE outgoing_invoice_drafts SET
                       buyer_name_snapshot=?,buyer_tax_code_snapshot=?,buyer_address_snapshot=?,
                       buyer_email_snapshot=?,company_name_snapshot=?,company_tax_code_snapshot=?,
                       company_address_snapshot=?,payment_requester_snapshot=?,
                       payment_bank_name_snapshot=?,payment_bank_account_snapshot=?
                   WHERE id=?""",
                (*snapshot_update_values(snapshot), draft_id),
            )
            if conn.execute("SELECT changes() n").fetchone()["n"] != 1:
                return jsonify({
                    "ok": False,
                    "error": "Hồ sơ hóa đơn vừa thay đổi; vui lòng tải lại",
                }), 409
            audit(
                conn, now_iso, "outgoing.capture_legacy_snapshot", "ok",
                entity_type="outgoing_invoice", entity_id=draft_id,
                metadata={"explicit_historical_confirmation": True},
            )
            return jsonify({"ok": True, "idempotent": False})

    @app.post("/api/outgoing-invoices/<int:draft_id>/cancel")
    def api_cancel_outgoing_draft(draft_id):
        body = request.get_json(silent=True) or {}
        if not body.get("confirmed"):
            return jsonify({"ok": False, "error": "Cần xác nhận hủy dự thảo hóa đơn"}), 400
        with db_factory() as conn:
            conn.execute("BEGIN IMMEDIATE")
            draft = conn.execute("SELECT * FROM outgoing_invoice_drafts WHERE id=?", (draft_id,)).fetchone()
            if not draft:
                return jsonify({"ok": False, "error": "Không tìm thấy dự thảo hóa đơn"}), 404
            if draft["status"] == "issued":
                return jsonify({"ok": False, "error": "Hóa đơn đã phát hành nên không thể hủy dự thảo"}), 409
            if draft["minvoice_status"] in {"saved", "saving", "unknown"}:
                return jsonify({
                    "ok": False,
                    "error": "Dự thảo đã lưu hoặc chưa đối soát xong với M-Invoice; không được hủy cục bộ",
                }), 409
            if draft["status"] != "cancelled":
                timestamp = now_iso()
                conn.execute(
                    "UPDATE outgoing_invoice_drafts SET status='cancelled' WHERE id=?", (draft_id,)
                )
                conn.execute(
                    """UPDATE inventory_transactions SET status='cancelled',updated_at=?
                       WHERE source_type='OUTGOING_DRAFT' AND source_id=? AND status='reserved'""",
                    (timestamp, str(draft_id)),
                )
                mark_substitution_actions_reversed_for_draft(conn, draft_id, timestamp)
                audit(conn, now_iso, "outgoing.cancel", "ok", entity_type="outgoing_invoice", entity_id=draft_id)
            return jsonify({"ok": True, "idempotent": draft["status"] == "cancelled"})

    @app.get("/api/outgoing-invoices")
    def api_outgoing_invoices():
        batch_id = request.args.get('batch_id', type=int)
        if 'batch_id' in request.args and (batch_id is None or batch_id <= 0):
            return jsonify(ok=False, error='Đơn hàng không hợp lệ.'), 400
        with db_factory() as conn:
            rows = [dict(row) for row in conn.execute(
                'SELECT * FROM outgoing_invoice_drafts ' +
                ('WHERE id IN (SELECT l.draft_id FROM outgoing_order_allocations l JOIN orders o ON o.id=l.order_id WHERE o.batch_id=? UNION SELECT draft_id FROM outgoing_consolidated_days WHERE batch_id=?) ORDER BY invoice_date DESC,id DESC' if batch_id else
                 'ORDER BY invoice_date DESC,id DESC LIMIT 200'), (batch_id,batch_id) if batch_id else ()
            )]
            taxes = defaultdict(set)
            for line in conn.execute('SELECT draft_id,tax FROM outgoing_invoice_lines WHERE draft_id IN '
                                     '(SELECT value FROM json_each(?))', (json.dumps([r['id'] for r in rows]),)):
                taxes[line['draft_id']].add(invoice_tax_percent(line['tax']))
            profiles = {
                row["contractor"]: dict(row)
                for row in conn.execute("SELECT * FROM outgoing_buyer_profiles")
            }
            for row in rows:
                row['source_batch_ids'] = [r[0] for r in conn.execute('SELECT DISTINCT o.batch_id FROM outgoing_order_allocations l JOIN orders o ON o.id=l.order_id WHERE l.draft_id=? UNION SELECT batch_id FROM outgoing_consolidated_days WHERE draft_id=?',(row['id'],row['id']))]
                row["buyer"] = profiles.get(row["contractor"])
                row['tax_label'] = ', '.join('KKKNT' if tax == -2 else 'KCT' if tax == -1 else f'{tax:g}%'
                                            for tax in sorted(taxes[row['id']]))
            return jsonify({"ok": True, "items": rows, "buyer_profiles": profiles})

    @app.put("/api/outgoing-buyers/<contractor>")
    def api_outgoing_buyer(contractor):
        body = request.get_json(force=True) or {}
        code = clean_text(contractor).upper()
        legal_name = clean_text(body.get("legal_name"))
        display_name = clean_text(body.get("display_name"))
        tax_code = clean_text(body.get("tax_code"))
        address = clean_text(body.get("address"))
        email = clean_text(body.get("email"))
        if not code or not address:
            return jsonify({"ok": False, "error": "Cần nhà thầu và địa chỉ người mua"}), 400
        if bool(legal_name) != bool(tax_code):
            return jsonify({"ok": False, "error": "Người mua là công ty phải có đủ tên pháp lý và mã số thuế"}), 400
        if tax_code and not re.fullmatch(r"\d{10}(?:-\d{3})?", tax_code):
            return jsonify({"ok": False, "error": "Mã số thuế người mua không hợp lệ"}), 400
        if not legal_name and not display_name:
            return jsonify({"ok": False, "error": "Cần tên người mua"}), 400
        with db_factory() as conn:
            conn.execute(
                """INSERT INTO outgoing_buyer_profiles(
                       contractor,display_name,legal_name,tax_code,address,email,updated_at
                   ) VALUES(?,?,?,?,?,?,?) ON CONFLICT(contractor) DO UPDATE SET
                       display_name=excluded.display_name,legal_name=excluded.legal_name,
                       tax_code=excluded.tax_code,address=excluded.address,email=excluded.email,
                       updated_at=excluded.updated_at""",
                (code, display_name, legal_name, tax_code, address, email, now_iso()),
            )
            audit(conn, now_iso, "outgoing.buyer_profile", "ok", entity_type="contractor", entity_id=code)
            return jsonify({"ok": True, "contractor": code})

    @app.post("/api/minvoice/drafts/<int:draft_id>")
    def api_save_minvoice_draft(draft_id):
        client_factory = app.config.get("MINVOICE_CLIENT_FACTORY") or create_minvoice_client
        if client_factory is None:
            return jsonify({"ok": False, "error": "Chưa cấu hình M-Invoice client"}), 500
        body = request.get_json(force=True) or {}
        dry_run = body.get("dry_run", True) is not False
        confirmed = body.get("confirm_remote_write") is True
        if not dry_run and not confirmed:
            return jsonify({"ok": False, "error": "Cần xác nhận rõ trước khi lưu dự thảo lên M-Invoice"}), 400
        try:
            minvoice_client = client_factory()
        except MinvoiceError as exc:
            return jsonify({"ok": False, "error": str(exc)}), 502
        if getattr(minvoice_client, "supports_remote_drafts", True) is False:
            return jsonify({"ok": False, "remote_write": False,
                            "error": "Portal mới đang hỗ trợ tải hóa đơn. Để lập hóa đơn, tải file M-Invoice rồi nhập trên portal."}), 409
        requested_series = clean_text(body.get("series")).upper()
        timestamp = now_iso()
        reconcile_only = False
        payload = None
        with db_factory() as conn:
            if not dry_run:
                conn.execute("BEGIN IMMEDIATE")
            draft = conn.execute(
                "SELECT * FROM outgoing_invoice_drafts WHERE id=?", (draft_id,),
            ).fetchone()
            if not draft:
                return jsonify({"ok": False, "error": "Không tìm thấy dự thảo hóa đơn"}), 404
            if draft["status"] != "draft":
                return jsonify({"ok": False, "error": "Chỉ dự thảo chưa phát hành mới được lưu lên M-Invoice"}), 409
            if not dry_run and draft["minvoice_status"] == "saved":
                return jsonify({
                    "ok": True, "dry_run": False, "remote_write": False, "idempotent": True,
                    "remote_id": draft["minvoice_remote_id"],
                    "message": "Dự thảo này đã được lưu lên M-Invoice",
                    "requires_user_sign_and_issue": True,
                })
            persisted_series = clean_text(draft["minvoice_series"]).upper()
            minvoice_status = clean_text(draft["minvoice_status"]) or "not_sent"
            if (
                not dry_run
                and minvoice_status in {"saving", "unknown"}
                and persisted_series
                and requested_series
                and requested_series != persisted_series
            ):
                return jsonify({
                    "ok": False,
                    "error": "Không được đổi ký hiệu khi lần lưu M-Invoice trước đang chờ đối soát",
                    "minvoice_status": minvoice_status,
                }), 409
            # Reconciliation must retain the exact series used by the possibly
            # successful POST.  A newly confirmed attempt may still choose a
            # different series after a definitive error/not-found result.
            series = (
                persisted_series
                if not dry_run and minvoice_status in {"saving", "unknown"}
                else requested_series or persisted_series
            )
            if not series:
                return jsonify({"ok": False, "error": "Cần chọn ký hiệu hóa đơn M-Invoice"}), 400
            installation_uuid = re.sub(
                r"[^A-Z0-9]", "", clean_text(setting_get(conn, "installation_uuid", "")).upper()
            )
            if not installation_uuid:
                return jsonify({
                    "ok": False,
                    "error": "Thiếu mã định danh cài đặt; không thể tạo khóa chống trùng M-Invoice an toàn",
                }), 500
            draft_external_uuid = re.sub(
                r"[^A-Z0-9]", "", clean_text(draft["external_key_uuid"]).upper()
            )
            if not draft_external_uuid:
                draft_external_uuid = uuid.uuid4().hex.upper()
                if not dry_run:
                    conn.execute(
                        "UPDATE outgoing_invoice_drafts SET external_key_uuid=? WHERE id=?",
                        (draft_external_uuid, draft_id),
                    )
            key_api = clean_text(draft["minvoice_key_api"]) or (
                f"TDP-{installation_uuid[:12]}-{draft_external_uuid[:32]}"
            )

            if not dry_run and minvoice_status == "saving":
                if minvoice_saving_is_fresh(draft["minvoice_started_at"], timestamp):
                    return jsonify({
                        "ok": False,
                        "error": "Dự thảo đang được gửi lên M-Invoice; không được gửi đồng thời",
                        "minvoice_status": "saving",
                    }), 409
                conn.execute(
                    """UPDATE outgoing_invoice_drafts
                       SET minvoice_status='unknown',minvoice_error=?
                       WHERE id=? AND minvoice_status='saving'""",
                    ("Lần lưu trước bị gián đoạn; cần đối soát key_api", draft_id),
                )
                reconcile_only = True
            elif not dry_run and minvoice_status == "unknown":
                reconcile_only = True

            snapshot = None
            if not reconcile_only:
                try:
                    validate_draft_export_stock(conn, draft_id, draft["invoice_date"])
                except OutgoingReadinessError as exc:
                    return jsonify({"ok": False, "error": str(exc), "code": exc.code}), exc.status
                if not dry_run:
                    try:
                        snapshot, buyer = current_invoice_snapshot(conn, draft)
                    except ValueError as exc:
                        return jsonify({"ok": False, "error": str(exc)}), 409
                else:
                    buyer = conn.execute(
                        "SELECT * FROM outgoing_buyer_profiles WHERE contractor=?", (draft["contractor"],),
                    ).fetchone()
                    if not buyer:
                        return jsonify({
                            "ok": False,
                            "error": f"Chưa lưu thông tin người mua cho {draft['contractor']}",
                        }), 400
                lines = [dict(row) for row in conn.execute(
                    "SELECT * FROM outgoing_invoice_lines WHERE draft_id=? ORDER BY id", (draft_id,),
                )]
                if not lines:
                    return jsonify({"ok": False, "error": "Dự thảo chưa có dòng hàng"}), 400
                payload = {
                    "invoice_date": draft["invoice_date"],
                    "series": series,
                    "currency": "VND",
                    "payment_method": clean_text(body.get("payment_method")) or "TM/CK",
                    "order_number": f"TDP-{draft['batch_id']}-{draft['contractor']}",
                    "key_api": key_api,
                    "buyer": {
                        "display_name": buyer["display_name"],
                        "legal_name": snapshot["buyer_name_snapshot"] if snapshot else buyer["legal_name"],
                        "tax_code": snapshot["buyer_tax_code_snapshot"] if snapshot else buyer["tax_code"],
                        "address": snapshot["buyer_address_snapshot"] if snapshot else buyer["address"],
                        "email": snapshot["buyer_email_snapshot"] if snapshot else buyer["email"],
                    },
                    "lines": [{
                        "code": line["product_code"],
                        "name": line["product_name"],
                        "unit": line["unit"],
                        "quantity": line["qty"],
                        "unit_price": line["unit_price"],
                        "tax": line["tax"],
                        "tchat": 2 if clean_text(line.get("invoice_nature")) == "2" else 1,
                    } for line in lines],
                }

            if not dry_run and not reconcile_only:
                conn.execute(
                    """UPDATE outgoing_invoice_drafts
                       SET minvoice_status='saving',minvoice_series=?,minvoice_key_api=?,
                           minvoice_started_at=?,minvoice_error=NULL,minvoice_remote_id=NULL,
                           buyer_name_snapshot=?,buyer_tax_code_snapshot=?,buyer_address_snapshot=?,
                           buyer_email_snapshot=?,company_name_snapshot=?,company_tax_code_snapshot=?,
                           company_address_snapshot=?,payment_requester_snapshot=?,
                           payment_bank_name_snapshot=?,payment_bank_account_snapshot=?
                       WHERE id=? AND minvoice_status NOT IN ('saved','saving','unknown')""",
                    (series, key_api, timestamp, *snapshot_update_values(snapshot), draft_id),
                )
                if conn.execute("SELECT changes() n").fetchone()["n"] != 1:
                    return jsonify({
                        "ok": False,
                        "error": "Trạng thái M-Invoice vừa thay đổi; tải lại trước khi thao tác",
                    }), 409

        if not dry_run:
            try:
                remote = minvoice_client.get_invoice_info(key_api=key_api)
            except MinvoiceError as exc:
                with db_factory() as conn:
                    if not reconcile_only:
                        conn.execute(
                            """UPDATE outgoing_invoice_drafts
                               SET minvoice_status='error',minvoice_error=? WHERE id=?""",
                            (str(exc), draft_id),
                        )
                    audit(conn, now_iso, "minvoice.reconcile", "error", str(exc),
                          entity_type="outgoing_invoice", entity_id=draft_id,
                          metadata={"key_api": key_api, "before_save": not reconcile_only})
                return jsonify({"ok": False, "error": str(exc)}), 502

            if remote["found"]:
                remote_id = minvoice_remote_id(remote.get("data"))
                with db_factory() as validation_conn:
                    validation_draft = validation_conn.execute(
                        "SELECT * FROM outgoing_invoice_drafts WHERE id=?", (draft_id,),
                    ).fetchone()
                    validation_lines = [dict(row) for row in validation_conn.execute(
                        "SELECT * FROM outgoing_invoice_lines WHERE draft_id=? ORDER BY id", (draft_id,),
                    )]
                mismatches = minvoice_reconciliation_mismatches(
                    remote.get("data"), validation_draft, series, validation_lines,
                )
                if mismatches:
                    message = "Bản nháp M-Invoice trùng key nhưng lệch " + ", ".join(mismatches)
                    with db_factory() as conn:
                        conn.execute(
                            """UPDATE outgoing_invoice_drafts
                               SET minvoice_status='unknown',minvoice_error=?,minvoice_reconciled_at=?
                               WHERE id=?""",
                            (message, now_iso(), draft_id),
                        )
                        audit(
                            conn, now_iso, "minvoice.reconcile", "mismatch", message,
                            entity_type="outgoing_invoice", entity_id=draft_id,
                            metadata={"fields": mismatches, "remote_id_received": bool(remote_id)},
                        )
                    return jsonify({
                        "ok": False,
                        "error": message + "; không được tự liên kết, cần đối chiếu thủ công",
                        "reconcile_required": True,
                    }), 409
                with db_factory() as conn:
                    conn.execute(
                        """UPDATE outgoing_invoice_drafts
                           SET minvoice_status='saved',minvoice_series=?,minvoice_key_api=?,
                               minvoice_remote_id=?,minvoice_saved_at=?,minvoice_reconciled_at=?,
                               minvoice_error=NULL WHERE id=?""",
                        (series, key_api, remote_id, now_iso(), now_iso(), draft_id),
                    )
                    audit(conn, now_iso, "minvoice.reconcile", "found",
                          entity_type="outgoing_invoice", entity_id=draft_id,
                          metadata={"series": series, "key_api": key_api,
                                    "remote_id_received": bool(remote_id)})
                return jsonify({
                    "ok": True, "dry_run": False, "remote_write": False,
                    "idempotent": True, "reconciled": True, "remote_id": remote_id,
                    "message": "Đã tìm thấy bản nháp M-Invoice theo key_api; không gửi lại",
                    "requires_user_sign_and_issue": True,
                    "draft_id": draft_id, "contractor": draft["contractor"],
                })

            if reconcile_only:
                with db_factory() as conn:
                    conn.execute(
                        """UPDATE outgoing_invoice_drafts
                           SET minvoice_status='not_sent',minvoice_error=NULL,
                               minvoice_started_at=NULL,minvoice_reconciled_at=?
                           WHERE id=? AND minvoice_status IN ('saving','unknown')""",
                        (now_iso(), draft_id),
                    )
                    audit(conn, now_iso, "minvoice.reconcile", "not_found",
                          entity_type="outgoing_invoice", entity_id=draft_id,
                          metadata={"key_api": key_api})
                return jsonify({
                    "ok": False,
                    "error": "Đối soát không thấy bản nháp trên M-Invoice; hệ thống chưa tự gửi lại. Hãy kiểm tra rồi bấm Lưu nháp lần nữa.",
                    "reconciled": True, "remote_write": False,
                    "retry_requires_new_confirmation": True,
                }), 409

        try:
            result = minvoice_client.create_draft(
                payload, dry_run=dry_run, confirm_remote_write=confirmed,
            )
        except MinvoiceOutcomeUnknown as exc:
            with db_factory() as conn:
                conn.execute(
                    """UPDATE outgoing_invoice_drafts
                       SET minvoice_status='unknown',minvoice_error=? WHERE id=?""",
                    (str(exc), draft_id),
                )
                audit(conn, now_iso, "minvoice.save_draft", "unknown", str(exc),
                      entity_type="outgoing_invoice", entity_id=draft_id,
                      metadata={"key_api": key_api})
            return jsonify({
                "ok": False, "error": str(exc), "minvoice_status": "unknown",
                "reconcile_required": True, "remote_write": False,
            }), 409
        except MinvoiceError as exc:
            if not dry_run:
                with db_factory() as conn:
                    conn.execute(
                        "UPDATE outgoing_invoice_drafts SET minvoice_status='error',minvoice_error=? WHERE id=?",
                        (str(exc), draft_id),
                    )
                    audit(conn, now_iso, "minvoice.save_draft", "error", str(exc),
                          entity_type="outgoing_invoice", entity_id=draft_id)
            return jsonify({"ok": False, "error": str(exc)}), 400
        if result.get("remote_write"):
            remote_id = minvoice_remote_id(result.get("data"))
            with db_factory() as conn:
                conn.execute(
                    """UPDATE outgoing_invoice_drafts
                       SET minvoice_status='saved',minvoice_series=?,minvoice_key_api=?,
                           minvoice_remote_id=?,minvoice_saved_at=?,minvoice_reconciled_at=?,
                           minvoice_error=NULL WHERE id=?""",
                    (series, key_api, remote_id, now_iso(), now_iso(), draft_id),
                )
                audit(conn, now_iso, "minvoice.save_draft", "ok", entity_type="outgoing_invoice",
                      entity_id=draft_id, metadata={"series": series, "key_api": key_api,
                                                    "remote_id_received": bool(remote_id)})
            result = {key: value for key, value in result.items() if key != "data"}
            result["remote_id"] = remote_id
        return jsonify({**result, "draft_id": draft_id, "contractor": draft["contractor"]})

    @app.post("/api/debts/payables/import/preview")
    def api_payables_import_preview():
        upload = request.files.get("file")
        if not upload or not upload.filename:
            return jsonify({"ok": False, "error": "Chưa chọn file công nợ phải trả"}), 400
        filename = Path(upload.filename).name
        if Path(filename).suffix.lower() not in {".xlsx", ".xlsm"}:
            return jsonify({"ok": False, "error": "Chỉ nhận file .xlsx hoặc .xlsm"}), 400
        payload = upload.read(MAPPING_IMPORT_MAX_BYTES + 1)
        if len(payload) > MAPPING_IMPORT_MAX_BYTES:
            return jsonify({"ok": False, "error": "File Excel vượt quá giới hạn 10 MB"}), 413
        try:
            with zipfile.ZipFile(io.BytesIO(payload)) as archive:
                entries = archive.infolist()
                if len(entries) > 2_000 or sum(item.file_size for item in entries) > MAPPING_IMPORT_MAX_UNCOMPRESSED_BYTES:
                    return jsonify({"ok": False, "error": "File Excel có cấu trúc quá lớn để đọc an toàn"}), 413
        except zipfile.BadZipFile:
            return jsonify({"ok": False, "error": "File Excel bị hỏng hoặc không đúng định dạng"}), 400

        cutoff = time.time() - MAPPING_IMPORT_TTL_SECONDS
        with PAYABLE_IMPORT_LOCK:
            for old_token, item in list(PENDING_PAYABLE_IMPORTS.items()):
                if item["created"] < cutoff:
                    PENDING_PAYABLE_IMPORTS.pop(old_token, None)
        workbook = None
        try:
            workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True, keep_links=False)
            with db_factory() as conn:
                preview = parse_historical_payables(conn, workbook)
                database_state_hash = payables_database_state_hash(conn)
        except ValueError as exc:
            return jsonify({"ok": False, "error": str(exc)}), 400
        except Exception:
            return jsonify({"ok": False, "error": "Không đọc được file công nợ phải trả"}), 400
        finally:
            if workbook is not None:
                workbook.close()

        token = uuid.uuid4().hex
        pending_items = preview.pop("items")
        source_hash = hashlib.sha256(payload).hexdigest().upper()
        with PAYABLE_IMPORT_LOCK:
            PENDING_PAYABLE_IMPORTS[token] = {
                "created": time.time(),
                "filename": filename,
                "source_hash": source_hash,
                "sheet": preview["sheet"],
                "items": pending_items,
                "database_state_hash": database_state_hash,
                "has_errors": not preview["can_confirm"],
            }
        return jsonify({
            "ok": True, "token": token, "filename": filename, "source_hash": source_hash,
            "expires_in_minutes": MAPPING_IMPORT_TTL_SECONDS // 60, **preview,
        })

    @app.post("/api/debts/payables/import/confirm")
    def api_payables_import_confirm():
        body = request.get_json(force=True) or {}
        if body.get("confirmed") is not True:
            return jsonify({"ok": False, "error": "Cần xác nhận trước khi ghi công nợ phải trả"}), 400
        token = clean_text(body.get("token"))
        with PAYABLE_IMPORT_LOCK:
            pending = PENDING_PAYABLE_IMPORTS.pop(token, None)
        if not pending or time.time() - pending["created"] > MAPPING_IMPORT_TTL_SECONDS:
            return jsonify({"ok": False, "error": "Phiên xem trước đã hết hạn; vui lòng chọn lại file"}), 410
        if pending["has_errors"]:
            return jsonify({"ok": False, "error": "File còn dòng lỗi nên chưa thể nhập"}), 400
        with db_factory() as conn:
            conn.execute("BEGIN IMMEDIATE")
            if payables_database_state_hash(conn) != pending["database_state_hash"]:
                return jsonify({
                    "ok": False,
                    "error": "Công nợ phải trả đã thay đổi sau khi xem trước; dữ liệu chưa được ghi, vui lòng chọn lại file",
                }), 409
            historical_through = max(
                (item["purchase_date"] for item in pending["items"]), default="",
            )
            if historical_through:
                # The imported workbook is an authoritative historical snapshot.
                # System purchase orders on/before this date are already covered
                # by that snapshot and must not be charged a second time.
                setting_set(conn, "historical_payables_through_date", historical_through)
            unchanged = conn.execute(
                "SELECT COUNT(*) n FROM historical_payable_lines WHERE source_hash=? AND source_sheet=?",
                (pending["source_hash"], pending["sheet"]),
            ).fetchone()["n"]
            if unchanged:
                stale = conn.execute(
                    "SELECT COUNT(*) n FROM historical_payable_lines WHERE source_hash!=?",
                    (pending["source_hash"],),
                ).fetchone()["n"]
                if stale:
                    conn.execute(
                        "DELETE FROM historical_payable_lines WHERE source_hash!=?",
                        (pending["source_hash"],),
                    )
                audit(
                    conn, now_iso, "debts.payables_import", "unchanged", entity_type="source_file",
                    entity_id=pending["source_hash"][:16], metadata={
                        "filename": pending["filename"], "sheet": pending["sheet"], "rows": unchanged,
                        "removed_stale_rows": stale,
                    },
                )
                payable_ledger = refresh_payable_ledger(conn, now_iso())
                return jsonify({
                    "ok": True, "inserted": 0, "replaced": stale, "unchanged": unchanged,
                    "idempotent": stale == 0, "source_hash": pending["source_hash"],
                    "payable_ledger": payable_ledger,
                })
            # This workbook is an authoritative historical-payables snapshot,
            # not an append-only journal.  Replace the prior snapshot even when
            # the customer renamed the revised file; otherwise one filename
            # change would double every payable line.
            replaced = conn.execute(
                "SELECT COUNT(*) n FROM historical_payable_lines"
            ).fetchone()["n"]
            conn.execute("DELETE FROM historical_payable_lines")
            timestamp = now_iso()
            for item in pending["items"]:
                conn.execute(
                    """INSERT INTO historical_payable_lines(
                           purchase_date,kitchen,item_name,qty,unit,supplier,buy_price,
                           damaged_qty,added_qty,reduced_qty,missing_qty,actual_qty,
                           source_amount,calculated_amount,amount,note,
                           source_file,source_sheet,source_row,source_hash,created_at,updated_at
                       ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                    (
                        item["purchase_date"], item["kitchen"], item["item_name"], item["qty"],
                        item["unit"], item["supplier"], item["buy_price"], item["damaged_qty"],
                        item["added_qty"], item["reduced_qty"], item["missing_qty"],
                        item["actual_qty"], item["source_amount"], item["calculated_amount"],
                        item["amount"], item["note"], pending["filename"],
                        pending["sheet"], item["source_row"], pending["source_hash"], timestamp, timestamp,
                    ),
                )
            audit(
                conn, now_iso, "debts.payables_import", "ok", entity_type="source_file",
                entity_id=pending["source_hash"][:16], metadata={
                    "filename": pending["filename"], "sheet": pending["sheet"],
                    "rows": len(pending["items"]), "replaced_rows": replaced,
                    "snapshot_scope": "all_historical_payables",
                },
            )
            payable_ledger = refresh_payable_ledger(conn, timestamp)
        return jsonify({
            "ok": True, "inserted": len(pending["items"]), "replaced": replaced, "unchanged": 0,
            "idempotent": False,
            "source_hash": pending["source_hash"], "payable_ledger": payable_ledger,
        })

    @app.post("/api/debt-adjustments")
    def api_debt_adjustment():
        body = request.get_json(force=True) or {}
        if body.get("party_type") not in {"contractor", "supplier"}:
            return jsonify({"ok": False, "error": "Nhóm công nợ không hợp lệ"}), 400
        with db_factory() as conn:
            cur = conn.execute(
                """INSERT INTO debt_adjustments(
                    adjustment_date,party_type,party_code,amount,note,created_at
                ) VALUES(?,?,?,?,?,?)""",
                (body.get("adjustment_date") or date.today().isoformat(), body["party_type"],
                 clean_text(body.get("party_code")), number_value(body.get("amount")),
                 clean_text(body.get("note")), now_iso()),
            )
            return jsonify({"ok": True, "id": cur.lastrowid})

    @app.get("/api/debts")
    def api_debts():
        period_from = request.args.get("from") or date.today().replace(day=1).isoformat()
        period_to = request.args.get("to") or date.today().isoformat()
        try:
            validate_date_range(period_from, period_to)
        except ValueError as exc:
            return jsonify({"ok": False, "error": str(exc)}), 400
        with db_factory() as conn:
            conn.execute("BEGIN")
            payload = debt_period_payload(conn, period_from, period_to, tax_factor)
            payload["receipts"] = [dict(row) for row in conn.execute(
                """SELECT id,payment_date,party_code,amount,note,status,revision,created_by,reversed_by,reversal_reason
                   FROM payments WHERE kind='receipt' AND party_type='contractor'
                   AND payment_date>=? AND payment_date<=? ORDER BY payment_date,id""",
                (period_from, period_to),
            )]
            return jsonify({"ok": True, **payload})

    @app.get("/api/export/debts")
    def api_export_debts():
        period_from = request.args.get("from") or date.today().replace(day=1).isoformat()
        period_to = request.args.get("to") or date.today().isoformat()
        try:
            validate_date_range(period_from, period_to)
        except ValueError as exc:
            return jsonify({"ok": False, "error": str(exc)}), 400
        with db_factory() as conn:
            conn.execute("BEGIN")
            payload = debt_period_payload(conn, period_from, period_to, tax_factor)
            workbook = debt_period_workbook(conn, payload)
            return send_workbook(
                workbook,
                f"Cong_no_{period_from}_{period_to}.xlsx",
            )

    @app.put("/api/kitchen-units/<kitchen_code>")
    def api_kitchen_unit(kitchen_code):
        body = request.get_json(force=True) or {}
        unit_code = clean_text(body.get("xcom_code") or body.get("unit_code")).upper()
        if not unit_code:
            return jsonify({"ok": False, "error": "Cần mã XCOM (xưởng cơm)"}), 400
        with db_factory() as conn:
            conn.execute(
                """INSERT INTO kitchen_units(kitchen_code,unit_code,updated_at) VALUES(?,?,?)
                   ON CONFLICT(kitchen_code) DO UPDATE SET unit_code=excluded.unit_code,updated_at=excluded.updated_at""",
                (clean_text(kitchen_code).upper(), unit_code, now_iso()),
            )
            return jsonify({"ok": True})

    @app.get("/api/kitchen/payment-profiles")
    def api_xcom_payment_profiles():
        with db_factory() as conn:
            return jsonify({"ok": True, "items": list_payment_profiles(conn)})

    @app.put("/api/kitchen/payment-profiles/<profile_code>")
    def api_xcom_payment_profile_upsert(profile_code):
        body = request.get_json(force=True) or {}
        body["profile_code"] = clean_text(profile_code).upper()
        try:
            with db_factory() as conn:
                result = upsert_payment_profile(conn, body, now_iso=now_iso)
                audit(
                    conn, now_iso, "kitchen.payment_profile.upsert", "ok",
                    entity_type="xcom_payment_profile", entity_id=result["profile_code"],
                    metadata={
                        "document_type": clean_text(body.get("document_type") or "MEAL_SIMPLE").upper(),
                        "created": result["created"], "changed": result["changed"],
                    },
                )
                profile = next(
                    item for item in list_payment_profiles(conn)
                    if item["profile_code"] == result["profile_code"]
                )
                return jsonify({"ok": True, **result, "profile": profile})
        except XcomPaymentError as exc:
            return xcom_payment_error_response(exc)

    @app.delete("/api/kitchen/payment-profiles/<profile_code>")
    def api_xcom_payment_profile_delete(profile_code):
        body = request.get_json(silent=True) or {}
        if body.get("confirmed") is not True:
            return jsonify({"ok": False, "error": "Cần xác nhận xóa hồ sơ thanh toán"}), 400
        try:
            with db_factory() as conn:
                deleted = delete_payment_profile(conn, profile_code)
                if not deleted:
                    return jsonify({"ok": False, "error": "Hồ sơ thanh toán không tồn tại"}), 404
                audit(
                    conn, now_iso, "kitchen.payment_profile.delete", "ok",
                    entity_type="xcom_payment_profile", entity_id=clean_text(profile_code).upper(),
                )
                return jsonify({"ok": True, "deleted": True})
        except XcomPaymentError as exc:
            return xcom_payment_error_response(exc)

    @app.put("/api/kitchen/payment-profiles/<profile_code>/scopes/<scope_type>/<scope_code>")
    def api_xcom_payment_scope_upsert(profile_code, scope_type, scope_code):
        try:
            with db_factory() as conn:
                result = assign_payment_scope(
                    conn, profile_code, scope_type, scope_code, now_iso=now_iso
                )
                audit(
                    conn, now_iso, "kitchen.payment_scope.upsert", "ok",
                    entity_type="xcom_payment_profile", entity_id=result["profile_code"],
                    metadata={
                        "scope_type": result["scope_type"], "scope_code": result["scope_code"],
                        "changed": result["changed"],
                    },
                )
                return jsonify({"ok": True, **result})
        except XcomPaymentError as exc:
            return xcom_payment_error_response(exc)

    @app.delete("/api/kitchen/payment-scopes/<scope_type>/<scope_code>")
    def api_xcom_payment_scope_delete(scope_type, scope_code):
        body = request.get_json(silent=True) or {}
        if body.get("confirmed") is not True:
            return jsonify({"ok": False, "error": "Cần xác nhận bỏ phạm vi hồ sơ thanh toán"}), 400
        try:
            with db_factory() as conn:
                deleted = remove_payment_scope(conn, scope_type, scope_code)
                if not deleted:
                    return jsonify({"ok": False, "error": "Phạm vi không tồn tại"}), 404
                audit(
                    conn, now_iso, "kitchen.payment_scope.delete", "ok",
                    entity_type="xcom_payment_scope",
                    entity_id=f"{clean_text(scope_type).upper()}:{clean_text(scope_code).upper()}",
                )
                return jsonify({"ok": True, "deleted": True})
        except XcomPaymentError as exc:
            return xcom_payment_error_response(exc)

    @app.put("/api/kitchen/payment-profiles/<profile_code>/tariffs/<period>/<shift>")
    def api_xcom_meal_tariff_upsert(profile_code, period, shift):
        body = request.get_json(force=True) or {}
        try:
            with db_factory() as conn:
                result = upsert_meal_tariff(
                    conn, profile_code, period, shift, body.get("unit_price"), now_iso=now_iso
                )
                audit(
                    conn, now_iso, "kitchen.meal_tariff.upsert", "ok",
                    entity_type="xcom_payment_profile", entity_id=result["profile_code"],
                    metadata={
                        "period": result["period"], "shift": result["shift"],
                        "unit_price": result["unit_price"], "changed": result["changed"],
                    },
                )
                return jsonify({"ok": True, **result})
        except XcomPaymentError as exc:
            return xcom_payment_error_response(exc)

    @app.delete("/api/kitchen/payment-profiles/<profile_code>/tariffs/<period>/<shift>")
    def api_xcom_meal_tariff_delete(profile_code, period, shift):
        body = request.get_json(silent=True) or {}
        if body.get("confirmed") is not True:
            return jsonify({"ok": False, "error": "Cần xác nhận xóa đơn giá suất ăn"}), 400
        try:
            with db_factory() as conn:
                deleted = delete_meal_tariff(conn, profile_code, period, shift)
                if not deleted:
                    return jsonify({"ok": False, "error": "Đơn giá kỳ/ca không tồn tại"}), 404
                audit(
                    conn, now_iso, "kitchen.meal_tariff.delete", "ok",
                    entity_type="xcom_payment_profile", entity_id=clean_text(profile_code).upper(),
                    metadata={"period": clean_text(period), "shift": clean_text(shift)},
                )
                return jsonify({"ok": True, "deleted": True})
        except XcomPaymentError as exc:
            return xcom_payment_error_response(exc)

    @app.post("/api/kitchen/payment-documents/preview")
    def api_xcom_payment_document_preview():
        body = request.get_json(force=True) or {}
        try:
            with db_factory() as conn:
                # One writer snapshot prevents profile/scope/tariff/attendance
                # edits from being mixed across the preview's multiple reads.
                conn.execute("BEGIN IMMEDIATE")
                result = create_payment_preview(
                    conn,
                    body.get("profile_code"),
                    body.get("date_from"),
                    body.get("date_to"),
                    body.get("issue_date"),
                    now_iso=now_iso,
                )
                audit(
                    conn, now_iso, "kitchen.payment_document.preview", "ok",
                    entity_type="xcom_payment_profile",
                    entity_id=result["summary"]["profile_code"],
                    metadata={
                        "from": result["summary"]["date_from"],
                        "to": result["summary"]["date_to"],
                        "issue_date": result["issue_date"],
                        "input_sha256": result["summary"]["input_sha256"],
                        "expires_at": result["expires_at"],
                    },
                )
                return jsonify({"ok": True, **result})
        except XcomPaymentError as exc:
            return xcom_payment_error_response(exc)

    @app.post("/api/kitchen/payment-documents/export")
    def api_xcom_payment_document_export():
        body = request.get_json(force=True) or {}
        try:
            with db_factory() as conn:
                # Token verification, source re-hash and single-use CAS must be
                # one atomic snapshot.  Export is POST so browser/link prefetch
                # cannot consume a one-time approval token.
                conn.execute("BEGIN IMMEDIATE")
                result = consume_payment_preview(
                    conn,
                    body.get("preview_token"),
                    body.get("profile_code"),
                    body.get("date_from"),
                    body.get("date_to"),
                    body.get("issue_date"),
                    now_iso=now_iso,
                )
                audit(
                    conn, now_iso, "kitchen.payment_document.export", "ok",
                    entity_type="xcom_payment_profile", entity_id=result["summary"]["profile_code"],
                    metadata={
                        "from": result["summary"]["date_from"],
                        "to": result["summary"]["date_to"],
                        "preview_created_at": result["preview_created_at"],
                        "input_sha256": result["input_sha256"],
                        "file_sha256": result["file_sha256"],
                        "document_type": result["summary"]["document_type"],
                    },
                )
                return send_file(
                    io.BytesIO(result["payload"]),
                    as_attachment=True,
                    download_name=result["filename"],
                    mimetype=result["mimetype"],
                )
        except XcomPaymentError as exc:
            return xcom_payment_error_response(exc)

    @app.put("/api/dated-prices/<product_code>")
    def api_dated_price(product_code):
        body = request.get_json(force=True) or {}
        period = clean_text(body.get("period"))
        if not re.fullmatch(r"\d{4}-\d{2}", period):
            return jsonify({"ok": False, "error": "Kỳ giá phải có dạng YYYY-MM"}), 400
        try:
            datetime.strptime(period + "-01", "%Y-%m-%d")
        except ValueError:
            return jsonify({"ok": False, "error": "Kỳ giá không hợp lệ"}), 400
        code = clean_text(product_code).upper()
        price_group = clean_text(body.get("price_group") or "HATRAN").upper()
        try:
            price_value = finite_number(body.get("price_value"), 0, "Giá theo kỳ")
        except ValueError as exc:
            return jsonify({"ok": False, "error": str(exc)}), 400
        if price_value <= 0:
            return jsonify({"ok": False, "error": "Giá theo kỳ phải lớn hơn 0"}), 400
        with db_factory() as conn:
            if not conn.execute("SELECT 1 FROM products WHERE code=?", (code,)).fetchone():
                return jsonify({"ok": False, "error": "Mã hàng chưa có trong danh mục"}), 404
            conn.execute(
                """INSERT INTO dated_prices(product_code,price_group,period,price_value,updated_at)
                   VALUES(?,?,?,?,?) ON CONFLICT(product_code,price_group,period)
                   DO UPDATE SET price_value=excluded.price_value,updated_at=excluded.updated_at""",
                (code, price_group, period, price_value, now_iso()),
            )
            audit(conn, now_iso, "dated_price.upsert", "ok", entity_type="product", entity_id=code,
                  metadata={"price_group": price_group, "period": period, "price_value": price_value})
            return jsonify({"ok": True})

    @app.get("/api/kitchen/plans")
    def api_meal_plans():
        with db_factory() as conn:
            return jsonify({"ok": True, "items": meal_plan_payload(conn, request.args.get("date", ""))})

    @app.post("/api/kitchen/plans")
    def api_save_meal_plan():
        body = request.get_json(force=True) or {}
        items = body.get("items") or []
        if not body.get("work_date") or not body.get("kitchen") or not body.get("shift"):
            return jsonify({"ok": False, "error": "Cần ngày, bếp và ca"}), 400
        if as_date(body.get("work_date")) != clean_text(body.get("work_date")):
            return jsonify({"ok": False, "error": "Ngày kế hoạch phải hợp lệ dạng YYYY-MM-DD"}), 400
        try:
            if not isinstance(items, list) or any(not isinstance(raw, dict) for raw in items):
                raise ValueError("Danh sách định lượng không hợp lệ")
            meal_count = finite_number(body.get("meal_count"), 0, "Số suất")
            menu_count = finite_integer(body.get("menu_count"), 1, "Số thực đơn")
            servings_per_menu = finite_number(
                body.get("servings_per_menu"), 0, "Số suất mỗi thực đơn",
            ) or (meal_count / max(menu_count, 1))
            meal_price = finite_number(body.get("meal_price"), 0, "Đơn giá suất ăn")
            other_cost = finite_number(body.get("other_cost"), 0, "Chi phí khác")
            plan_id = finite_integer(body.get("id"), 0, "Mã kế hoạch")
            normalized_items = [{
                **raw,
                "_norm_qty": finite_number(raw.get("norm_qty"), 0, "Định lượng nguyên liệu"),
                "_source_norm": finite_number(
                    raw.get("source_norm_per_1000"), 0, "Định lượng nguồn",
                ),
                "_meal_count": finite_number(
                    raw.get("applicable_meal_count"), 0, "Số suất áp dụng",
                ),
            } for raw in items]
        except ValueError as exc:
            return jsonify({"ok": False, "error": str(exc)}), 400
        if meal_count <= 0:
            return jsonify({"ok": False, "error": "Số suất phải lớn hơn 0"}), 400
        if menu_count <= 0 or servings_per_menu <= 0 or meal_price < 0 or other_cost < 0:
            return jsonify({"ok": False, "error": "Số thực đơn/suất phải dương và chi phí không được âm"}), 400
        with db_factory() as conn:
            if not normalized_items:
                return jsonify({
                    "ok": False,
                    "error": "Cần ít nhất một dòng nguyên liệu có mã hàng và định lượng",
                }), 400
            item_errors = []
            for item_number, raw in enumerate(normalized_items, start=1):
                code = clean_text(raw.get("product_code")).upper()
                if not code:
                    item_errors.append(f"Dòng {item_number}: thiếu mã hàng")
                    continue
                if raw["_norm_qty"] <= 0:
                    item_errors.append(f"Dòng {item_number}: định lượng phải lớn hơn 0")
                    continue
                if not conn.execute("SELECT 1 FROM products WHERE code=?", (code,)).fetchone():
                    item_errors.append(f"Dòng {item_number}: mã hàng {code} chưa có trong danh mục")
            if item_errors:
                return jsonify({"ok": False, "error": "; ".join(item_errors)}), 400
            kitchen = clean_text(body["kitchen"]).upper()
            unit = conn.execute("SELECT unit_code FROM kitchen_units WHERE kitchen_code=?", (kitchen,)).fetchone()
            unit_code = clean_text(body.get("xcom_code") or body.get("unit_code")) or (unit["unit_code"] if unit else "")
            if plan_id:
                conn.execute(
                    """UPDATE meal_plans SET work_date=?,kitchen=?,shift=?,meal_count=?,unit_code=?,
                       menu_count=?,servings_per_menu=?,meal_price=?,other_cost=?,status='draft',note=?,
                       updated_at=? WHERE id=?""",
                    (body["work_date"], kitchen, clean_text(body["shift"]), meal_count, unit_code,
                     menu_count, servings_per_menu, meal_price, other_cost, clean_text(body.get("note")),
                     now_iso(), plan_id),
                )
                conn.execute("DELETE FROM meal_plan_items WHERE plan_id=?", (plan_id,))
            else:
                cur = conn.execute(
                    """INSERT INTO meal_plans(
                        work_date,kitchen,shift,meal_count,unit_code,status,note,created_at,updated_at,
                        menu_count,servings_per_menu,meal_price,other_cost
                    ) VALUES(?,?,?,?,?,'draft',?,?,?,?,?,?,?)""",
                    (body["work_date"], kitchen, clean_text(body["shift"]), meal_count,
                     unit_code, clean_text(body.get("note")), now_iso(), now_iso(), menu_count,
                     servings_per_menu, meal_price, other_cost),
                )
                plan_id = cur.lastrowid
            period = body["work_date"][:7]
            saved = 0
            warnings = []
            for raw in normalized_items:
                code = clean_text(raw.get("product_code")).upper()
                product = conn.execute("SELECT * FROM products WHERE code=?", (code,)).fetchone()
                dated = conn.execute(
                    "SELECT price_value FROM dated_prices WHERE product_code=? AND price_group='HATRAN' AND period=?",
                    (code, period),
                ).fetchone()
                if dated and number_value(dated["price_value"]) > 0:
                    price = number_value(dated["price_value"])
                    source = f"HATRAN {period} · giá kỳ đã khóa"
                else:
                    price = 0
                    source = f"HATRAN {period} · chưa có giá"
                    warnings.append(f"{code}: chưa có giá HATRAN đúng kỳ {period}")
                conn.execute(
                    """INSERT INTO meal_plan_items(
                        plan_id,dish_name,product_code,product_name,norm_qty,unit,supplier,buy_price,price_source,
                        source_norm_per_1000,applicable_meal_count,source_amount
                    ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)""",
                    (plan_id, clean_text(raw.get("dish_name")), code, product["name"],
                     raw["_norm_qty"], clean_text(raw.get("unit")) or product["unit"],
                     clean_text(raw.get("supplier")) or product["supplier"], price, source,
                     raw["_source_norm"] or raw["_norm_qty"] * 1000,
                     raw["_meal_count"] or meal_count, 0),
                )
                saved += 1
            audit(conn, now_iso, "kitchen.save_plan", "ok", entity_type="meal_plan", entity_id=plan_id,
                  metadata={"items": saved, "warnings": len(warnings)})
            return jsonify({"ok": True, "id": plan_id, "warnings": warnings,
                            "items": meal_plan_payload(conn, body["work_date"])})

    @app.post("/api/kitchen/plans/<int:plan_id>/approve")
    def api_approve_meal_plan(plan_id):
        with db_factory() as conn:
            plan = conn.execute(
                """SELECT mp.*,
                          COALESCE(NULLIF(mp.unit_code,''),ku.unit_code,'') mapped_unit
                   FROM meal_plans mp
                   LEFT JOIN kitchen_units ku ON ku.kitchen_code=mp.kitchen
                   WHERE mp.id=?""",
                (plan_id,),
            ).fetchone()
            if not plan:
                return jsonify({"ok": False, "error": "Không tìm thấy kế hoạch xưởng cơm"}), 404
            if not clean_text(plan["mapped_unit"]):
                return jsonify({"ok": False, "error": "Phải ghép bếp vào XCOM trước khi duyệt"}), 400
            if number_value(plan["meal_price"]) <= 0:
                return jsonify({"ok": False, "error": "Phải có đơn giá suất ăn trước khi duyệt"}), 400
            count = conn.execute("SELECT COUNT(*) n FROM meal_plan_items WHERE plan_id=?", (plan_id,)).fetchone()["n"]
            if count == 0:
                return jsonify({"ok": False, "error": "Kế hoạch chưa có định lượng nguyên liệu"}), 400
            missing_price = conn.execute(
                "SELECT COUNT(*) n FROM meal_plan_items WHERE plan_id=? AND buy_price<=0", (plan_id,)
            ).fetchone()["n"]
            if missing_price:
                return jsonify({"ok": False, "error": f"Còn {missing_price} nguyên liệu thiếu giá"}), 400
            expected_prefix = f"HATRAN {plan['work_date'][:7]}"
            wrong_period_price = conn.execute(
                """SELECT COUNT(*) n
                   FROM meal_plan_items i
                   LEFT JOIN dated_prices d
                     ON d.product_code=i.product_code AND d.price_group='HATRAN' AND d.period=?
                   WHERE i.plan_id=? AND (
                     d.price_value IS NULL OR d.price_value<=0 OR ABS(i.buy_price-d.price_value)>0.01
                   )""",
                (plan["work_date"][:7], plan_id),
            ).fetchone()["n"]
            if wrong_period_price:
                return jsonify({
                    "ok": False,
                    "error": f"Còn {wrong_period_price} nguyên liệu chưa dùng giá {expected_prefix}",
                }), 400
            conn.execute("UPDATE meal_plans SET status='approved',updated_at=? WHERE id=?", (now_iso(), plan_id))
            return jsonify({"ok": True})

    @app.get("/api/kitchen/po")
    def api_kitchen_po():
        work_date = request.args.get("date") or date.today().isoformat()
        with db_factory() as conn:
            plans = meal_plan_payload(conn, work_date)
            if not plans:
                return jsonify({"ok": False, "error": "Ngày này chưa có kế hoạch xưởng cơm"}), 404
            wb = meal_po_workbook(plans, work_date)
            approval_label = "DA_DUYET" if all(plan["status"] == "approved" for plan in plans) else "NHAP"
            stream = io.BytesIO()
            wb.save(stream)
            stream.seek(0)
            return send_file(stream, as_attachment=True,
                             download_name=f"PO_xuong_com_{approval_label}_{work_date}.xlsx",
                             mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    @app.post("/api/staff")
    def api_save_staff():
        body = request.get_json(force=True) or {}
        code = clean_text(body.get("employee_code")).upper()
        name = clean_text(body.get("full_name"))
        if not code or not name:
            return jsonify({"ok": False, "error": "Cần mã và họ tên nhân sự"}), 400
        try:
            staff_numbers = {
                "base_salary": finite_number(body.get("base_salary"), 0, "Lương cơ bản"),
                "standard_days": finite_number(body.get("standard_days"), 26, "Ngày công chuẩn"),
                "standard_hours": finite_number(body.get("standard_hours"), 8, "Giờ công chuẩn"),
                "bhxh_employee_rate": finite_number(body.get("bhxh_employee_rate"), 0, "Tỷ lệ BHXH nhân viên"),
                "bhxh_company_rate": finite_number(body.get("bhxh_company_rate"), 0, "Tỷ lệ BHXH công ty"),
            }
        except ValueError as exc:
            return jsonify({"ok": False, "error": str(exc)}), 400
        if any(value < 0 for value in staff_numbers.values()) or staff_numbers["standard_days"] <= 0 or staff_numbers["standard_hours"] <= 0:
            return jsonify({"ok": False, "error": "Lương/tỷ lệ không được âm; ngày và giờ chuẩn phải dương"}), 400
        with db_factory() as conn:
            conn.execute(
                """INSERT INTO staff(
                    employee_code,full_name,role_name,kitchen,base_salary,standard_days,standard_hours,
                    bhxh_employee_rate,bhxh_company_rate,active,created_at,updated_at
                ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)
                ON CONFLICT(employee_code) DO UPDATE SET
                    full_name=excluded.full_name,role_name=excluded.role_name,kitchen=excluded.kitchen,
                    base_salary=excluded.base_salary,standard_days=excluded.standard_days,
                    standard_hours=excluded.standard_hours,bhxh_employee_rate=excluded.bhxh_employee_rate,
                    bhxh_company_rate=excluded.bhxh_company_rate,active=excluded.active,updated_at=excluded.updated_at""",
                (code, name, clean_text(body.get("role_name")), clean_text(body.get("kitchen")).upper(),
                 staff_numbers["base_salary"], staff_numbers["standard_days"],
                 staff_numbers["standard_hours"], staff_numbers["bhxh_employee_rate"],
                 staff_numbers["bhxh_company_rate"], 0 if body.get("active") is False else 1,
                 now_iso(), now_iso()),
            )
            return jsonify({"ok": True})

    @app.post("/api/attendance")
    def api_save_attendance():
        body = request.get_json(force=True) or {}
        work_date = clean_text(body.get("work_date"))
        try:
            attendance_numbers = [
                finite_number(body.get(field), 0, label) for field, label in (
                    ("normal_hours", "Giờ thường"), ("overtime_hours", "Giờ tăng ca"),
                    ("sunday_hours", "Giờ Chủ nhật"), ("night_hours", "Giờ đêm"),
                    ("holiday_hours", "Giờ ngày lễ"),
                )
            ]
        except ValueError as exc:
            return jsonify({"ok": False, "error": str(exc)}), 400
        if as_date(work_date) != work_date or any(value < 0 for value in attendance_numbers):
            return jsonify({"ok": False, "error": "Ngày phải hợp lệ và số giờ không được âm"}), 400
        with db_factory() as conn:
            staff_row = conn.execute(
                "SELECT id FROM staff WHERE employee_code=?", (clean_text(body.get("employee_code")).upper(),)
            ).fetchone()
            if not staff_row:
                return jsonify({"ok": False, "error": "Nhân sự hoặc ngày chấm công không hợp lệ"}), 400
            conn.execute(
                """INSERT INTO attendance_entries(
                    employee_id,work_date,normal_hours,overtime_hours,sunday_hours,night_hours,
                    holiday_hours,note,source,updated_at
                ) VALUES(?,?,?,?,?,?,?,?,?,?) ON CONFLICT(employee_id,work_date) DO UPDATE SET
                    normal_hours=excluded.normal_hours,overtime_hours=excluded.overtime_hours,
                    sunday_hours=excluded.sunday_hours,night_hours=excluded.night_hours,
                    holiday_hours=excluded.holiday_hours,note=excluded.note,source=excluded.source,
                    updated_at=excluded.updated_at""",
                (staff_row["id"], work_date, *attendance_numbers,
                 clean_text(body.get("note")), "manual", now_iso()),
            )
            return jsonify({"ok": True})

    @app.post("/api/attendance/import")
    def api_import_attendance():
        return jsonify({
            "ok": False,
            "error": "Luồng nạp trực tiếp đã khóa; hãy xem trước file rồi xác nhận",
        }), 410

    @app.post("/api/attendance/import/preview")
    def api_attendance_import_preview():
        period = clean_text(request.form.get("period"))
        try:
            validate_legacy_attendance_period(period)
        except ValueError as exc:
            return jsonify({"ok": False, "error": str(exc)}), 400
        upload = request.files.get("file")
        if not upload or not upload.filename:
            return jsonify({"ok": False, "error": "Chưa chọn file chấm công"}), 400
        filename = Path(upload.filename).name
        if Path(filename).suffix.lower() not in {".xlsx", ".xlsm"}:
            return jsonify({"ok": False, "error": "Chỉ nhận file chấm công .xlsx/.xlsm"}), 400
        payload = upload.read(MAPPING_IMPORT_MAX_BYTES + 1)
        if len(payload) > MAPPING_IMPORT_MAX_BYTES:
            return jsonify({"ok": False, "error": "File Excel vượt quá giới hạn 10 MB"}), 413
        try:
            with zipfile.ZipFile(io.BytesIO(payload)) as archive:
                entries = archive.infolist()
                expanded_size = sum(item.file_size for item in entries)
                if len(entries) > 2_000 or expanded_size > MAPPING_IMPORT_MAX_UNCOMPRESSED_BYTES:
                    return jsonify({
                        "ok": False,
                        "error": "File Excel có cấu trúc quá lớn để đọc an toàn",
                    }), 413
        except zipfile.BadZipFile:
            return jsonify({"ok": False, "error": "File Excel bị hỏng hoặc không đúng định dạng"}), 400

        cutoff = time.time() - MAPPING_IMPORT_TTL_SECONDS
        with LEGACY_ATTENDANCE_IMPORT_LOCK:
            for old_token, item in list(PENDING_LEGACY_ATTENDANCE_IMPORTS.items()):
                if item["created"] < cutoff:
                    PENDING_LEGACY_ATTENDANCE_IMPORTS.pop(old_token, None)

        values_book = formulas_book = None
        try:
            values_book = load_workbook(
                io.BytesIO(payload), read_only=False, data_only=True, keep_links=False,
            )
            formulas_book = load_workbook(
                io.BytesIO(payload), read_only=False, data_only=False, keep_links=False,
            )
            with db_factory() as conn:
                preview = parse_legacy_attendance_workbooks(
                    conn, values_book, formulas_book, filename, period,
                )
                database_state_hash = legacy_attendance_database_state_hash(conn, period)
        except ValueError as exc:
            return jsonify({"ok": False, "error": str(exc)}), 400
        except Exception:
            return jsonify({
                "ok": False,
                "error": "Không đọc được file chấm công; vui lòng kiểm tra lại file",
            }), 400
        finally:
            if values_book is not None:
                values_book.close()
            if formulas_book is not None:
                formulas_book.close()

        token = uuid.uuid4().hex
        snapshot = preview.pop("snapshot")
        source_hash = hashlib.sha256(payload).hexdigest().upper()
        with LEGACY_ATTENDANCE_IMPORT_LOCK:
            PENDING_LEGACY_ATTENDANCE_IMPORTS[token] = {
                "created": time.time(),
                "filename": filename,
                "period": period,
                "source_hash": source_hash,
                "database_state_hash": database_state_hash,
                "snapshot": snapshot,
                "counts": preview["counts"],
                "warnings": preview["warnings"],
            }
        return jsonify({
            "ok": True,
            "token": token,
            "filename": filename,
            "source_hash": source_hash,
            "expires_in_minutes": MAPPING_IMPORT_TTL_SECONDS // 60,
            **preview,
        })

    @app.post("/api/attendance/import/confirm")
    def api_attendance_import_confirm():
        body = request.get_json(force=True) or {}
        if body.get("confirmed") is not True:
            return jsonify({"ok": False, "error": "Cần xác nhận trước khi ghi chấm công"}), 400
        token = clean_text(body.get("token"))
        with LEGACY_ATTENDANCE_IMPORT_LOCK:
            pending = PENDING_LEGACY_ATTENDANCE_IMPORTS.pop(token, None)
        if not pending or time.time() - pending["created"] > MAPPING_IMPORT_TTL_SECONDS:
            return jsonify({
                "ok": False,
                "error": "Phiên xem trước đã hết hạn; vui lòng chọn lại file",
            }), 410
        try:
            with db_factory() as conn:
                conn.execute("BEGIN IMMEDIATE")
                if legacy_attendance_database_state_hash(conn, pending["period"]) != pending["database_state_hash"]:
                    raise ValueError(
                        "Dữ liệu chấm công/lương đã thay đổi sau khi xem trước; "
                        "chưa ghi file, vui lòng xem trước lại"
                    )
                result = apply_legacy_attendance_snapshot(
                    conn, pending["snapshot"], pending["filename"], now_iso(),
                )
                audit(
                    conn, now_iso, "attendance.import", "ok",
                    entity_type="period", entity_id=pending["period"],
                    metadata={
                        "filename": pending["filename"],
                        "source_hash": pending["source_hash"],
                        **result,
                    },
                )
        except ValueError as exc:
            return jsonify({"ok": False, "error": str(exc)}), 409
        except Exception:
            return jsonify({
                "ok": False,
                "error": "Không ghi được chấm công; dữ liệu chưa được thay đổi",
            }), 409
        return jsonify({"ok": True, **result})

    @app.put("/api/payroll-adjustments/<employee_code>/<month>")
    def api_payroll_adjustment(employee_code, month):
        body = request.get_json(force=True) or {}
        if not re.fullmatch(r"\d{4}-\d{2}", month):
            return jsonify({"ok": False, "error": "Tháng phải có dạng YYYY-MM"}), 400
        numeric_fields = (
            "allowance", "responsibility", "advance", "probation_deduction",
            "bhxh_employee_amount", "bhxh_company_amount", "gross_override", "net_override",
        )
        try:
            values = {
                field: finite_number(body.get(field), 0, field)
                for field in numeric_fields
            }
        except ValueError as exc:
            return jsonify({"ok": False, "error": str(exc)}), 400
        with db_factory() as conn:
            person = conn.execute("SELECT id FROM staff WHERE employee_code=?", (employee_code.upper(),)).fetchone()
            if not person:
                return jsonify({"ok": False, "error": "Không tìm thấy nhân sự"}), 404
            conn.execute(
                """INSERT INTO payroll_adjustments(
                    employee_id,month,allowance,responsibility,advance,probation_deduction,
                    bhxh_employee_amount,bhxh_company_amount,gross_override,net_override,use_override,note,updated_at
                ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?) ON CONFLICT(employee_id,month) DO UPDATE SET
                    allowance=excluded.allowance,responsibility=excluded.responsibility,
                    advance=excluded.advance,probation_deduction=excluded.probation_deduction,
                    bhxh_employee_amount=excluded.bhxh_employee_amount,
                    bhxh_company_amount=excluded.bhxh_company_amount,
                    gross_override=excluded.gross_override,net_override=excluded.net_override,
                    use_override=excluded.use_override,note=excluded.note,updated_at=excluded.updated_at""",
                (person["id"], month, *(values[field] for field in numeric_fields),
                 1 if body.get("use_override") else 0,
                 clean_text(body.get("note")), now_iso()),
            )
            return jsonify({"ok": True, "payroll": payroll_rows(conn, month)})

    @app.get("/api/payroll")
    def api_payroll():
        month = request.args.get("month") or date.today().strftime("%Y-%m")
        if not re.fullmatch(r"\d{4}-\d{2}", month):
            return jsonify({"ok": False, "error": "Tháng tính lương phải có dạng YYYY-MM"}), 400
        with db_factory() as conn:
            try:
                rows = payroll_rows(conn, month)
            except ValueError as exc:
                return jsonify({"ok": False, "error": str(exc)}), 400
            labor = [dict(row) for row in conn.execute(
                """SELECT kitchen,SUM(amount) amount FROM kitchen_labor_costs
                   WHERE substr(work_date,1,7)=? GROUP BY kitchen ORDER BY kitchen""", (month,)
            )]
            return jsonify({"ok": True, "month": month, "items": rows, "labor_by_kitchen": labor,
                            "total_net": sum(row["net_salary"] for row in rows)})

    @app.get("/api/export/payroll")
    def api_export_payroll():
        month = request.args.get("month") or date.today().strftime("%Y-%m")
        if not re.fullmatch(r"\d{4}-\d{2}", month):
            return jsonify({"ok": False, "error": "Tháng tính lương phải có dạng YYYY-MM"}), 400
        with db_factory() as conn:
            try:
                workbook = payroll_workbook(conn, month)
            except ValueError as exc:
                return jsonify({"ok": False, "error": str(exc)}), 400
            return send_workbook(workbook, f"Bang_luong_{month}.xlsx")

    @app.put("/api/print/settings")
    def api_print_settings():
        body = request.get_json(force=True) or {}
        other_paper = clean_text(
            body.get("other_paper") or body.get("paper") or "A5"
        ).upper()
        if other_paper not in {"A4", "A5"}:
            return jsonify({"ok": False, "error": "Chứng từ khác chỉ hỗ trợ khổ A4 hoặc A5"}), 400
        printer_name = clean_text(body.get("printer_name"))
        printer_state = windows_printer_state()
        if printer_name and printer_name not in printer_state["installed"]:
            return jsonify({"ok": False, "error": "Tên máy in không có trong danh sách Windows"}), 400
        try:
            copies = finite_integer(body.get("copies"), 1, "Số bản in")
        except ValueError as exc:
            return jsonify({"ok": False, "error": str(exc)}), 400
        if not 1 <= copies <= 10:
            return jsonify({"ok": False, "error": "Số bản in phải từ 1 đến 10"}), 400
        with db_factory() as conn:
            setting_set(conn, "printer_name", printer_name)
            setting_set(conn, "print_copies", copies)
            setting_set(conn, "print_other_paper", other_paper)
            invalidated = conn.execute(
                """UPDATE print_jobs SET status='stale',approved_at=NULL,error_message=?
                   WHERE status IN ('preparing','prepared','approved')""",
                ("Cấu hình máy in/số bản đã thay đổi; cần chuẩn bị và duyệt lại",),
            ).rowcount
            if invalidated:
                audit(
                    conn, now_iso, "print.settings_invalidate", "ok",
                    entity_type="print_settings", entity_id="default",
                    metadata={
                        "jobs": invalidated, "copies": copies, "printer": printer_name,
                        "delivery_paper": "A4", "other_paper": other_paper,
                    },
                )
            return jsonify({
                "ok": True,
                "printer": printer_name or printer_state["default"],
                "copies": copies,
                "paper": other_paper,
                "delivery_paper": "A4",
                "other_paper": other_paper,
                "invalidated_jobs": invalidated,
                "hardware_ready": bool(
                    printer_state["supported"]
                    and (printer_name or printer_state["default"]) in printer_state["installed"]
                ),
            })

    @app.post("/api/print/prepare/<int:batch_id>")
    def api_prepare_print(batch_id):
        export_dir = data_dir / "print_jobs" / str(batch_id)
        export_dir.mkdir(parents=True, exist_ok=True)
        claim_id = uuid.uuid4().hex
        attempt_dir = export_dir / "attempts" / claim_id
        attempt_dir.mkdir(parents=True, exist_ok=False)
        with db_factory() as conn:
            conn.execute("BEGIN IMMEDIATE")
            batch, orders = ctx["require_batch"](conn, batch_id)
            if batch["status"] != "approved":
                return jsonify({"ok": False, "error": "Phải duyệt phiên đơn trước khi chuẩn bị in"}), 400
            locked_job = conn.execute(
                """SELECT status FROM print_jobs
                   WHERE batch_id=?
                     AND status IN ('preparing','approved','submitting','submitted','submission_unknown','printed')
                   ORDER BY id DESC LIMIT 1""",
                (batch_id,),
            ).fetchone()
            if locked_job:
                message = (
                    "Bộ chứng từ đang/đã gửi sang máy in nên không được ghi đè"
                    if locked_job["status"] in {"submitting", "submitted", "submission_unknown", "printed"}
                    else "Bộ PDF đã duyệt; hãy hủy bộ in cũ trước khi chuẩn bị lại"
                )
                return jsonify({"ok": False, "error": message}), 409
            copies = max(1, min(int(as_number(setting_get(conn, "print_copies", "1"), 1)), 10))
            other_paper = clean_text(setting_get(conn, "print_other_paper", "A5")).upper()
            if other_paper not in {"A4", "A5"}:
                other_paper = "A5"
            printer_state = windows_printer_state()
            prepared_printer = clean_text(setting_get(conn, "printer_name", "")) or printer_state["default"]
            job_specs = (
                ("delivery_pdf", "A4"),
                ("other_pdf", other_paper),
            )
            for document_type, paper in job_specs:
                conn.execute(
                    """INSERT INTO print_jobs(
                           batch_id,document_type,file_path,status,created_at,paper,copies,printer_name,
                           error_message,claim_id
                       ) VALUES(?,?,'','preparing',?,?,?,?,NULL,?)
                       ON CONFLICT(batch_id,document_type) DO UPDATE SET
                           file_path='',status='preparing',created_at=excluded.created_at,
                           paper=excluded.paper,copies=excluded.copies,printer_name=excluded.printer_name,
                           file_sha256=NULL,input_sha256=NULL,manifest_path=NULL,page_count=0,
                           approved_at=NULL,printed_at=NULL,submitted_at=NULL,error_message=NULL,
                           claim_id=excluded.claim_id""",
                    (
                        batch_id, document_type, now_iso(), paper, copies,
                        prepared_printer, claim_id,
                    ),
                )
            # Make the preparing claim visible before the comparatively slow
            # XLSX/PDF render.  Order mutation sees this row and is rejected.
            conn.commit()
            documents = {}
            bundles = []
            try:
                documents["deliveries"] = (
                    ctx["export_deliveries"](conn, batch, orders),
                    f"Phieu_giao_{batch['work_date']}.xlsx",
                )
                purchase_exporter = ctx.get(
                    "export_optional_purchase_documents",
                    ctx["export_purchase_documents"],
                )
                purchase_workbook = purchase_exporter(conn, batch, orders)
                if purchase_workbook is not None:
                    documents["purchases"] = (
                        purchase_workbook,
                        f"Bang_ke_{batch['work_date']}.xlsx",
                    )
                documents["report"] = (
                    ctx["export_report"](conn, batch, orders),
                    f"Bao_cao_{batch['work_date']}.xlsx",
                )
                for _, (workbook, filename) in documents.items():
                    workbook.save(attempt_dir / filename)
                bundle_inputs = (
                    (
                        "delivery_pdf", "A4",
                        ("deliveries",),
                        f"Bo_phieu_giao_A4_{batch['work_date']}.pdf",
                    ),
                    (
                        "other_pdf", other_paper,
                        tuple(
                            source_type
                            for source_type in ("purchases", "report")
                            if source_type in documents
                        ),
                        f"Bo_chung_tu_khac_{other_paper}_{batch['work_date']}.pdf",
                    ),
                )
                print_titles = {
                    "deliveries": "Phiếu giao hàng",
                    "purchases": "Bảng kê và giấy biên nhận",
                    "report": "Báo cáo tổng hợp",
                }
                for document_type, paper, source_types, filename in bundle_inputs:
                    pdf_path = attempt_dir / filename
                    manifest = build_excel_pdf_bundle(
                        [
                            {
                                "path": attempt_dir / documents[source_type][1],
                                "document_type": source_type,
                                "title": print_titles[source_type],
                            }
                            for source_type in source_types
                        ],
                        pdf_path,
                        generated_at=now_iso(),
                        paper=paper,
                        duplex=paper == 'A4',
                    )
                    manifest_path = write_manifest(
                        manifest, attempt_dir / f"{document_type}.manifest.json"
                    )
                    bundles.append({
                        "document_type": document_type,
                        "paper": paper,
                        "pdf_path": pdf_path,
                        "manifest": manifest,
                        "manifest_path": manifest_path,
                    })
            except Exception as exc:
                app.logger.exception("Print bundle preparation failed for batch %s", batch_id)
                conn.execute(
                    """UPDATE print_jobs SET status='stale',error_message=?
                       WHERE batch_id=? AND status='preparing' AND claim_id=?""",
                    (f"Chuẩn bị PDF thất bại: {str(exc)[:180]}", batch_id, claim_id),
                )
                return jsonify({"ok": False, "error": f"Không tạo được PDF chuẩn in: {str(exc)[:240]}"}), 500
            finally:
                for workbook, _ in documents.values():
                    workbook.close()

            # Refresh after rendering.  A concurrent explicit invalidation may
            # release the batch; never resurrect that obsolete render.
            conn.rollback()
            conn.execute("BEGIN IMMEDIATE")
            current_batch = conn.execute("SELECT status FROM batches WHERE id=?", (batch_id,)).fetchone()
            preparing_count = conn.execute(
                """SELECT COUNT(*) n FROM print_jobs WHERE batch_id=?
                   AND document_type IN ('delivery_pdf','other_pdf')
                   AND status='preparing' AND claim_id=?""",
                (batch_id, claim_id),
            ).fetchone()["n"]
            if (
                not current_batch or current_batch["status"] != "approved"
                or preparing_count != len(job_specs)
            ):
                return jsonify({
                    "ok": False,
                    "error": "Phiên đơn hoặc yêu cầu chuẩn bị in đã thay đổi; PDF vừa tạo không được duyệt",
                }), 409
            conn.execute(
                """DELETE FROM print_jobs WHERE batch_id=?
                   AND document_type NOT IN ('delivery_pdf','other_pdf')
                   AND status IN ('stale','cancelled','error','prepared')""",
                (batch_id,),
            )
            for bundle in bundles:
                manifest = bundle["manifest"]
                changed = conn.execute(
                    """UPDATE print_jobs SET
                           file_path=?,status='prepared',created_at=?,file_sha256=?,input_sha256=?,
                           manifest_path=?,page_count=?,paper=?,copies=?,printer_name=?,
                           approved_at=NULL,printed_at=NULL,submitted_at=NULL,error_message=NULL
                       WHERE batch_id=? AND document_type=? AND status='preparing'
                         AND claim_id=?""",
                    (
                        str(bundle["pdf_path"].resolve()), now_iso(), manifest["sha256"],
                        manifest["input_sha256"], str(bundle["manifest_path"].resolve()),
                        manifest["pages"], bundle["paper"], copies, prepared_printer,
                        batch_id, bundle["document_type"], claim_id,
                    ),
                ).rowcount
                if changed != 1:
                    return jsonify({
                        "ok": False,
                        "error": "Lượt chuẩn bị PDF đã bị thay thế; kết quả cũ không được công bố",
                    }), 409
            total_pages = sum(bundle["manifest"]["pages"] for bundle in bundles)
            total_sections = sum(bundle["manifest"]["section_count"] for bundle in bundles)
            audit(
                conn, now_iso, "print.prepare", "ok", entity_type="batch", entity_id=batch_id,
                metadata={
                    "documents": [
                        {
                            "type": bundle["document_type"], "paper": bundle["paper"],
                            "pages": bundle["manifest"]["pages"],
                            "sha256": bundle["manifest"]["sha256"],
                        }
                        for bundle in bundles
                    ],
                    "pages": total_pages, "sections": total_sections,
                },
            )
            return jsonify({
                "ok": True,
                "prepared": [bundle["document_type"] for bundle in bundles],
                "requires_approval": True,
                "pages": total_pages,
                "sections": total_sections,
                "documents": [
                    {
                        "document_type": bundle["document_type"],
                        "paper": bundle["paper"],
                        "pages": bundle["manifest"]["pages"],
                        "sha256": bundle["manifest"]["sha256"],
                        "pdf_url": f"/api/print/pdf/{batch_id}/{bundle['document_type']}",
                    }
                    for bundle in bundles
                ],
            })

    @app.post("/api/print/approve/<int:batch_id>")
    def api_approve_print(batch_id):
        with db_factory() as conn:
            conn.execute("BEGIN IMMEDIATE")
            batch = conn.execute("SELECT status FROM batches WHERE id=?", (batch_id,)).fetchone()
            if not batch or batch["status"] != "approved":
                return jsonify({"ok": False, "error": "Phiên đơn không còn ở trạng thái đã duyệt"}), 409
            count = conn.execute(
                "SELECT COUNT(*) n FROM print_jobs WHERE batch_id=? AND status='prepared'", (batch_id,)
            ).fetchone()["n"]
            if not count:
                return jsonify({"ok": False, "error": "Không có PDF mới ở trạng thái chờ duyệt"}), 400
            conn.execute(
                """UPDATE print_jobs SET status='approved',approved_at=?,error_message=NULL
                   WHERE batch_id=? AND status='prepared'""",
                (now_iso(), batch_id),
            )
            audit(conn, now_iso, "print.approve", "ok", entity_type="batch", entity_id=batch_id,
                  metadata={"jobs": count})
            return jsonify({"ok": True, "approved": count})

    @app.post("/api/print/invalidate/<int:batch_id>")
    def api_invalidate_print(batch_id):
        body = request.get_json(silent=True) or {}
        if body.get("confirmed") is not True:
            return jsonify({"ok": False, "error": "Cần xác nhận hủy bộ PDF in cũ"}), 400
        with db_factory() as conn:
            submitted = conn.execute(
                """SELECT COUNT(*) n FROM print_jobs
                   WHERE batch_id=? AND status IN ('submitting','submitted','submission_unknown','printed')""",
                (batch_id,),
            ).fetchone()["n"]
            if submitted:
                return jsonify({
                    "ok": False,
                    "error": "Bộ chứng từ đang/đã gửi sang máy in nên không được hủy khỏi lịch sử",
                }), 409
            changed = conn.execute(
                """UPDATE print_jobs SET status='stale',approved_at=NULL,error_message=?
                   WHERE batch_id=? AND status NOT IN ('stale','cancelled')""",
                ("Bộ PDF đã bị hủy vì cần sửa dữ liệu nguồn", batch_id),
            ).rowcount
            if changed:
                audit(conn, now_iso, "print.invalidate", "ok", entity_type="batch", entity_id=batch_id,
                      metadata={"jobs": changed})
            return jsonify({"ok": True, "invalidated": changed, "idempotent": changed == 0})

    @app.post("/api/print/run/<int:batch_id>")
    def api_run_print(batch_id):
        body = request.get_json(silent=True) or {}
        dry_run = body.get("dry_run") is True
        with db_factory() as conn:
            batch = conn.execute("SELECT status FROM batches WHERE id=?", (batch_id,)).fetchone()
            if not batch or batch["status"] != "approved":
                return jsonify({"ok": False, "error": "Phiên đơn không còn ở trạng thái đã duyệt"}), 409
            jobs = [dict(row) for row in conn.execute(
                "SELECT * FROM print_jobs WHERE batch_id=? AND status='approved' ORDER BY id", (batch_id,)
            )]
            if not jobs:
                locked = conn.execute(
                    """SELECT status FROM print_jobs
                       WHERE batch_id=?
                         AND status IN ('submitting','submitted','submission_unknown','printed')
                       ORDER BY id DESC LIMIT 1""",
                    (batch_id,),
                ).fetchone()
                if locked:
                    return jsonify({
                        "ok": False,
                        "status": locked["status"],
                        "error": "Bộ chứng từ đang/đã được Windows tiếp nhận; không gửi lại để tránh in trùng",
                    }), 409
                return jsonify({"ok": False, "error": "Không có chứng từ đã duyệt để in"}), 400
            verified_jobs = []
            try:
                for job in jobs:
                    path = safe_print_path(data_dir, batch_id, job["file_path"])
                    if not path.is_file():
                        raise ValueError("PDF in bị thiếu")
                    if hashlib.sha256(path.read_bytes()).hexdigest() != clean_text(job.get("file_sha256")):
                        raise ValueError("PDF đã thay đổi sau bước chuẩn bị")
                    paper = clean_text(job.get("paper")).upper()
                    verify_excel_pdf(
                        path,
                        minimum_pages=max(1, int(job.get("page_count") or 1)),
                        paper=paper,
                    )
                    verified_jobs.append((job, path))
            except (OSError, ValueError, ExcelPrintError) as exc:
                return jsonify({"ok": False, "error": f"Bộ PDF không còn hợp lệ: {str(exc)[:220]}"}), 409

            job_types = {clean_text(job.get("document_type")) for job in jobs}
            papers = {
                clean_text(job.get("document_type")): clean_text(job.get("paper")).upper()
                for job in jobs
            }
            if any(paper not in {"A4", "A5"} for paper in papers.values()):
                return jsonify({"ok": False, "error": "Khổ giấy đã duyệt chỉ được là A4 hoặc A5"}), 409
            split_types = {"delivery_pdf", "other_pdf"}
            if job_types & split_types and job_types & split_types != split_types:
                return jsonify({"ok": False, "error": "Bộ in A4/A5 chưa đủ hai nhóm chứng từ"}), 409
            if papers.get("delivery_pdf") not in {None, "A4"}:
                return jsonify({"ok": False, "error": "Phiếu giao hàng bắt buộc khổ A4"}), 409
            approved_copies = {max(1, min(int(as_number(job.get("copies"), 1)), 10)) for job in jobs}
            approved_printers = {clean_text(job.get("printer_name")) for job in jobs}
            if len(approved_copies) != 1 or len(approved_printers) != 1:
                return jsonify({"ok": False, "error": "Cấu hình các chứng từ đã duyệt không đồng nhất"}), 409
            copies = approved_copies.pop()
            printer_name = approved_printers.pop()
            printer_state = windows_printer_state()
            hardware_ready = bool(
                printer_state["supported"] and printer_name and printer_name in printer_state["installed"]
            )
            if dry_run:
                audit(
                    conn, now_iso, "print.run", "dry_run",
                    entity_type="batch", entity_id=batch_id,
                    metadata={"jobs": len(jobs), "copies": copies, "hardware_ready": hardware_ready},
                )
                return jsonify({
                    "ok": True,
                    "dry_run": True,
                    "status": "verified",
                    "jobs": len(jobs),
                    "copies": copies,
                    "printer": printer_name or "Chưa có máy in mặc định Windows",
                    "hardware_ready": hardware_ready,
                    "physical_confirmation_required": False,
                })

            print_tool = pdf_print_tool_path()
            if not hardware_ready or not print_tool:
                return jsonify({
                    "ok": False,
                    "error": "Chưa có máy in Windows hoặc bộ in PDF độc lập hợp lệ; chưa gửi lệnh in",
                }), 409
            try:
                import win32print

                # Snapshot all global driver state before claiming the
                # immutable job. Sumatra targets the printer directly, while
                # the driver still needs a paper/bin pair per physical tray.
                restore_printer = str(win32print.GetDefaultPrinter() or "").strip()
                if not restore_printer or restore_printer not in printer_state["installed"]:
                    return jsonify({
                        "ok": False,
                        "error": "Windows chưa có máy in mặc định hợp lệ để khôi phục sau lệnh in",
                    }), 409
                restore_media_state = windows_printer_media_state(win32print, printer_name)
                printer_bins = windows_printer_bins(win32print, printer_name)
                bin_by_paper = {
                    paper: select_windows_printer_bin(printer_bins, paper)
                    for paper in set(papers.values())
                }
                if not restore_media_state or any(not item for item in bin_by_paper.values()):
                    return jsonify({
                        "ok": False,
                        "error": "Không xác định được đúng khay A4/A5 của driver máy in; chưa gửi lệnh in",
                    }), 409
            except Exception:
                return jsonify({
                    "ok": False,
                    "error": "Không đọc được máy in mặc định Windows; chưa gửi lệnh in",
                }), 409

            job_ids = [int(job["id"]) for job in jobs]
            if not claim_approved_print_jobs(conn, batch_id, job_ids):
                return jsonify({
                    "ok": False,
                    "error": "Bộ chứng từ vừa được một lệnh khác tiếp nhận; không gửi lại để tránh in trùng",
                }), 409
            audit(
                conn, now_iso, "print.run", "submitting",
                entity_type="batch", entity_id=batch_id,
                metadata={"jobs": len(jobs), "copies": copies, "hardware_ready": hardware_ready},
            )
            # The claim must be durable before the first irreversible shell
            # call.  The context manager's final commit alone would be too late.
            conn.commit()

        submission_failed = False
        default_restore_failed = False
        paper_restore_failed = False
        commands_started = 0
        spool_job_ids: set[int] = set()
        with PRINT_SUBMISSION_LOCK:
            live_restore_printer = restore_printer
            live_restore_media_state = restore_media_state
            try:
                current_default = str(win32print.GetDefaultPrinter() or "").strip()
                if not current_default or current_default not in printer_state["installed"]:
                    raise RuntimeError("missing live default printer")
                live_restore_printer = current_default
                baseline_job_ids = windows_print_job_ids(win32print, printer_name)
                if baseline_job_ids is None:
                    raise RuntimeError("cannot read target print queue")
                known_job_ids = set(baseline_job_ids)
                current_media_state = windows_printer_media_state(win32print, printer_name)
                if not current_media_state:
                    raise RuntimeError("cannot read target printer media")
                live_restore_media_state = current_media_state
                for job, path in verified_jobs:
                    paper = clean_text(job.get("paper")).upper()
                    bin_item = bin_by_paper[paper]
                    set_windows_printer_media(
                        win32print, printer_name, paper, bin_item["code"]
                    )
                    for _ in range(copies):
                        commands_started += 1
                        new_job_ids = submit_pdf_via_sumatra(
                            print_tool, path, printer_name, paper, bin_item,
                            win32print, known_job_ids,
                            appdata_dir=data_dir / "sumatra_appdata",
                        )
                        if not new_job_ids:
                            raise RuntimeError("PDF print tool did not create a spool job")
                        known_job_ids.update(new_job_ids)
                        spool_job_ids.update(new_job_ids)
            except Exception:
                submission_failed = True
            finally:
                # Restore both the paper form and its physical input tray only
                # after each job has materialized in the target spooler.
                try:
                    set_windows_printer_media(
                        win32print, printer_name,
                        live_restore_media_state[0], live_restore_media_state[1],
                    )
                except Exception:
                    paper_restore_failed = True
                # Always compare the live default.  SetDefaultPrinter can fail
                # after partially changing Windows state, so a boolean flag is
                # not a reliable indication that restoration is unnecessary.
                try:
                    if str(win32print.GetDefaultPrinter() or "").strip() != live_restore_printer:
                        win32print.SetDefaultPrinter(live_restore_printer)
                except Exception:
                    default_restore_failed = True

        restore_failed = default_restore_failed or paper_restore_failed

        finished_at = now_iso()
        if submission_failed:
            error_message = "Không xác định Windows đã nhận bao nhiêu bản; không được bấm in lại"
            try:
                with db_factory() as conn:
                    finished = finish_claimed_print_jobs(
                        conn,
                        job_ids,
                        status="submission_unknown",
                        printer_name=printer_name,
                        copies=copies,
                        error_message=error_message,
                    )
                    if not finished:
                        raise RuntimeError("print claim changed unexpectedly")
                    audit(
                        conn, now_iso, "print.run", "submission_unknown",
                        entity_type="batch", entity_id=batch_id,
                        metadata={
                            "jobs": len(jobs), "copies": copies,
                            "commands_started": commands_started,
                            "spool_jobs_observed": len(spool_job_ids),
                            "default_restored": not restore_failed,
                            "printer_paper_restored": not paper_restore_failed,
                        },
                    )
            except Exception:
                return jsonify({
                    "ok": False,
                    "status": "submitting",
                    "error": "Lệnh in có thể đã vào hàng đợi nhưng chưa ghi được kết quả; tuyệt đối không bấm lại",
                }), 500
            return jsonify({
                "ok": False,
                "status": "submission_unknown",
                "error": "Windows không gửi được trọn bộ PDF; có thể một phần đã vào hàng đợi, không bấm lại",
                "default_printer_restored": not restore_failed,
                "printer_paper_restored": not paper_restore_failed,
            }), 500

        restore_warning = (
            "Đã gửi PDF nhưng Windows không khôi phục được máy in mặc định; cần đặt lại thủ công"
            if restore_failed else None
        )
        try:
            with db_factory() as conn:
                finished = finish_claimed_print_jobs(
                    conn,
                    job_ids,
                    status="submitted",
                    submitted_at=finished_at,
                    printer_name=printer_name,
                    copies=copies,
                    error_message=restore_warning,
                )
                if not finished:
                    raise RuntimeError("print claim changed unexpectedly")
                audit(
                    conn, now_iso, "print.run", "submitted",
                    entity_type="batch", entity_id=batch_id,
                    metadata={
                        "jobs": len(jobs), "copies": copies,
                        "commands_started": commands_started,
                        "spool_jobs_observed": len(spool_job_ids),
                        "default_restored": not restore_failed,
                        "printer_paper_restored": not paper_restore_failed,
                    },
                )
        except Exception:
            return jsonify({
                "ok": False,
                "status": "submitting",
                "error": "PDF đã được gửi sang Windows nhưng chưa ghi được kết quả; tuyệt đối không bấm lại",
            }), 500
        response = {
            "ok": True,
            "dry_run": False,
            "status": "submitted",
            "jobs": len(jobs),
            "copies": copies,
            "printer": printer_name,
            "hardware_ready": True,
            "physical_confirmation_required": True,
            "spool_confirmed": True,
            "spool_jobs_observed": len(spool_job_ids),
            "default_printer_restored": not restore_failed,
            "printer_paper_restored": not paper_restore_failed,
            "print_engine": "SumatraPDF 3.6.1",
        }
        if restore_warning:
            response["warning"] = restore_warning
        return jsonify(response)

    @app.get("/api/print/pdf/<int:batch_id>")
    @app.get("/api/print/pdf/<int:batch_id>/<document_type>")
    def api_print_pdf(batch_id, document_type=None):
        if document_type and document_type not in {"delivery_pdf", "other_pdf", "pdf_bundle"}:
            return jsonify({"ok": False, "error": "Loại PDF in không hợp lệ"}), 400
        with db_factory() as conn:
            if document_type:
                job = conn.execute(
                    """SELECT * FROM print_jobs WHERE batch_id=? AND document_type=?
                       ORDER BY id DESC LIMIT 1""",
                    (batch_id, document_type),
                ).fetchone()
            else:
                job = conn.execute(
                    """SELECT * FROM print_jobs WHERE batch_id=?
                       AND document_type IN ('pdf_bundle','delivery_pdf')
                       ORDER BY CASE document_type WHEN 'pdf_bundle' THEN 0 ELSE 1 END,id DESC
                       LIMIT 1""",
                    (batch_id,),
                ).fetchone()
            if not job:
                return jsonify({"ok": False, "error": "Chưa chuẩn bị PDF cho phiên này"}), 404
            data = dict(job)
            if data.get("status") in {"stale", "cancelled"}:
                return jsonify({"ok": False, "error": "Bộ PDF này đã bị hủy; cần chuẩn bị lại"}), 409
            try:
                path = safe_print_path(data_dir, batch_id, data["file_path"])
                if not path.is_file() or hashlib.sha256(path.read_bytes()).hexdigest() != clean_text(data["file_sha256"]):
                    raise ValueError("PDF bị thiếu hoặc đã thay đổi")
                verify_excel_pdf(
                    path,
                    minimum_pages=max(1, int(data.get("page_count") or 1)),
                    paper=clean_text(data.get("paper")).upper(),
                )
            except (OSError, ValueError, ExcelPrintError) as exc:
                return jsonify({"ok": False, "error": str(exc)[:220]}), 409
            return send_file(
                path,
                as_attachment=True,
                download_name=path.name,
                mimetype="application/pdf",
            )

    @app.get("/api/print/jobs/<int:batch_id>")
    def api_print_jobs(batch_id):
        with db_factory() as conn:
            return jsonify({"ok": True, "items": [dict(row) for row in conn.execute(
                "SELECT * FROM print_jobs WHERE batch_id=? ORDER BY id", (batch_id,)
            )]})

    @app.put("/api/document-settings")
    def api_document_settings():
        body = request.get_json(force=True) or {}
        allowed = {"payment_requester", "payment_bank_name", "payment_bank_account"}
        with db_factory() as conn:
            for key in allowed:
                if key in body:
                    setting_set(conn, key, clean_text(body[key]))
            return jsonify({"ok": True})

    @app.get("/api/export/payment-request/<contractor>")
    def api_payment_request(contractor):
        # Backward-compatible URL, but the source of truth is now restricted
        # to invoices the user explicitly confirmed as issued.  This prevents
        # a cancelled/unissued draft from producing an official-looking request.
        return api_invoice_payment_bundle(contractor)

    @app.get("/api/outgoing-invoices/payment-scope/<contractor>")
    def api_invoice_payment_scope(contractor):
        try:
            with db_factory() as conn:
                conn.execute("PRAGMA query_only=ON")
                scope = issued_invoice_payment_scope(
                    conn,
                    contractor,
                    request.args.get("from"),
                    request.args.get("to"),
                )
        except InvoicePaymentScopeError as exc:
            return jsonify({"ok": False, "error": str(exc), "code": exc.code}), exc.status
        return jsonify({"ok": True, **public_payment_scope(scope)})

    @app.get("/api/outgoing-invoices/delivery-statement/<contractor>")
    def api_invoice_delivery_statement(contractor):
        period_from = request.args.get("from")
        period_to = request.args.get("to")
        with db_factory() as conn:
            try:
                scope = issued_invoice_payment_scope(conn, contractor, period_from, period_to)
            except InvoicePaymentScopeError as exc:
                return jsonify({"ok": False, "error": str(exc), "code": exc.code}), exc.status
            expected_scope = clean_text(request.args.get("scope_id")).upper()
            if expected_scope and expected_scope != scope["scope_id"]:
                return jsonify({
                    "ok": False,
                    "error": "Phạm vi hóa đơn đã thay đổi; cần xem lại danh sách trước khi tải",
                    "code": "stale_invoice_payment_scope",
                }), 409
            try:
                statement = invoice_delivery_statement_scope_workbook(scope)
            except InvoiceDeliveryStatementError as exc:
                return jsonify({"ok": False, "error": str(exc), "code": exc.code}), exc.status
            stream = io.BytesIO()
            statement.save(stream)
            statement.close()
            stream.seek(0)
            safe_code = re.sub(r"[^A-Z0-9_-]+", "_", scope["contractor"])[:40] or "KHACH_HANG"
            audit(
                conn, now_iso, "outgoing.invoice_delivery_statement", "ok",
                entity_type="contractor", entity_id=scope["contractor"],
                metadata={
                    "from": scope["date_from"], "to": scope["date_to"],
                    "invoices": len(scope["invoices"]),
                    "amount": scope["totals"]["total_amount"],
                    "scope_id": scope["scope_id"],
                },
            )
            statement_prefix = 'Bang_ke_hoa_don_VAT' if scope.get('statement_kind') == 'invoices' else 'Bang_tong_hop_giao_nhan'
            response = send_file(
                stream, as_attachment=True,
                download_name=(
                    f"{statement_prefix}_{safe_code}_"
                    f"{scope['date_from']}_{scope['date_to']}.xlsx"
                ),
                mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
            response.headers["X-TDP-Invoice-Scope"] = scope["scope_id"]
            return response

    @app.get("/api/export/invoice-payment-bundle/<contractor>")
    def api_invoice_payment_bundle(contractor):
        period_from = request.args.get("from")
        period_to = request.args.get("to")
        with db_factory() as conn:
            try:
                scope = issued_invoice_payment_scope(
                    conn, contractor, period_from, period_to,
                )
            except InvoicePaymentScopeError as exc:
                return jsonify({"ok": False, "error": str(exc), "code": exc.code}), exc.status
            expected_scope = clean_text(request.args.get("scope_id")).upper()
            if expected_scope and expected_scope != scope["scope_id"]:
                return jsonify({
                    "ok": False,
                    "error": "Phạm vi hóa đơn đã thay đổi; cần xem lại danh sách trước khi tải",
                    "code": "stale_invoice_payment_scope",
                }), 409
            contractor = scope["contractor"]
            period_from = scope["date_from"]
            period_to = scope["date_to"]
            drafts = scope["drafts"]
            lines = scope["lines"]
            details = scope["invoices"]
            snapshot = scope["snapshot"]
            payment_total = scope["totals"]["total_amount"]
            try:
                statement = invoice_delivery_statement_scope_workbook(scope)
                payment_request = invoice_payment_request_workbook(
                    scope,
                    issue_date=request.args.get("issue_date") or period_to,
                    request_number=request.args.get("request_number") or "……/CV/ĐNTT",
                    contract_no=request.args.get("contract_no") or "",
                    contract_date=request.args.get("contract_date") or "",
                )
            except (InvoiceDeliveryStatementError, InvoicePaymentDocumentError) as exc:
                return jsonify({"ok": False, "error": str(exc), "code": exc.code}), exc.status
            payment_stream = io.BytesIO()
            payment_request.save(payment_stream)
            payment_request.close()
            xlsx_stream = io.BytesIO()
            statement.save(xlsx_stream)
            statement.close()
            safe_code = re.sub(r"[^A-Z0-9_-]+", "_", contractor)[:40] or "KHACH_HANG"
            bundle = io.BytesIO()
            statement_prefix = 'Bang_ke_hoa_don_VAT' if scope.get('statement_kind') == 'invoices' else 'Bang_tong_hop_giao_nhan'
            with zipfile.ZipFile(bundle, "w", zipfile.ZIP_DEFLATED) as archive:
                archive.writestr(
                    f"De_nghi_thanh_toan_{safe_code}_{period_from}_{period_to}.xlsx",
                    payment_stream.getvalue(),
                )
                archive.writestr(
                    f"{statement_prefix}_{safe_code}_{period_from}_{period_to}.xlsx",
                    xlsx_stream.getvalue(),
                )
                invoice_sources = "\n".join(
                    f"- {item['invoice_series']} / {item['invoice_number']} / {item['invoice_date']}: "
                    f"{item['verification_source']}"
                    for item in details
                )
                archive.writestr(
                    "THONG_TIN_DOI_CHIEU.txt",
                    ("HỒ SƠ HÓA ĐƠN ĐÃ ĐỐI CHIẾU\n"
                     "Đề nghị thanh toán dùng biểu mẫu chính thức theo mẫu khách hàng.\n"
                     f"{scope['warning']}\n"
                     f"Mã phạm vi: {scope['scope_id']}\n"
                     f"Nhà thầu: {contractor}\nTừ ngày: {period_from}\nĐến ngày: {period_to}\n"
                     f"Số hóa đơn: {len(details)}\nTổng đề nghị: {payment_total:,.0f} VNĐ\n"
                     "Nguồn: chỉ hóa đơn đỏ đã xác nhận phát hành và chi tiết đã đối chiếu.\n"
                     f"{invoice_sources}\n"),
                )
            bundle.seek(0)
            audit(conn, now_iso, "outgoing.invoice_payment_bundle", "ok", entity_type="contractor",
                  entity_id=contractor, metadata={"from": period_from, "to": period_to,
                                                   "invoices": len(details), "amount": payment_total,
                                                   "scope_id": scope["scope_id"],
                                                   "template_status": "official_customer_xlsx"})
            response = send_file(
                bundle, as_attachment=True,
                download_name=f"Ho_so_de_nghi_thanh_toan_{safe_code}_{period_from}_{period_to}.zip",
                mimetype="application/zip",
            )
            response.headers["X-TDP-Template-Status"] = "official-customer-xlsx"
            response.headers["X-TDP-Invoice-Scope"] = scope["scope_id"]
            return response


def _group_rows(rows, key):
    groups = defaultdict(list)
    for row in rows:
        groups[row.get(key) or "CHƯA XÁC ĐỊNH"].append(row)
    return groups


def debt_period_payload(conn, period_from: str, period_to: str, tax_factor):
    if period_from > period_to:
        raise ValueError("Ngày bắt đầu phải trước ngày kết thúc")
    parties = {
        "contractor": defaultdict(lambda: {
            "opening": 0, "period_charge": 0, "period_paid": 0,
            "period_adjustment": 0, "closing": 0,
        }),
        "supplier": defaultdict(lambda: {
            "opening": 0, "period_charge": 0, "period_paid": 0,
            "period_adjustment": 0, "closing": 0,
        }),
    }
    party_maps = {"contractor": {}, "supplier": {}}

    def remember_party(party_type, value):
        code = mapping_cell_text(value)
        if code:
            party_maps[party_type].setdefault(code.casefold(), code)

    # Master tables win; later case-only variants from imports/orders are
    # folded into the same ledger account instead of creating phantom debts.
    for row in conn.execute("SELECT code FROM contractors ORDER BY code"):
        remember_party("contractor", row["code"])
    for row in conn.execute("SELECT code FROM suppliers ORDER BY code"):
        remember_party("supplier", row["code"])
    for row in conn.execute("SELECT DISTINCT contractor FROM orders ORDER BY contractor"):
        remember_party("contractor", row["contractor"])
    for row in conn.execute(
        "SELECT DISTINCT contractor_code FROM receivable_ledger_lines ORDER BY contractor_code"
    ):
        remember_party("contractor", row["contractor_code"])
    for query in (
        "SELECT DISTINCT supplier FROM products ORDER BY supplier",
        "SELECT DISTINCT supplier FROM orders ORDER BY supplier",
        "SELECT DISTINCT supplier FROM purchase_order_lines ORDER BY supplier",
        "SELECT DISTINCT supplier FROM purchase_workbook_lines ORDER BY supplier",
        "SELECT DISTINCT supplier FROM historical_payable_lines ORDER BY supplier",
    ):
        for row in conn.execute(query):
            remember_party("supplier", row["supplier"])
    for table in ("balances", "payments", "debt_adjustments"):
        for row in conn.execute(
            f"SELECT DISTINCT party_type,party_code FROM {table} ORDER BY party_type,party_code"
        ):
            if row["party_type"] in party_maps:
                remember_party(row["party_type"], row["party_code"])

    def party_key(party_type, value):
        code = mapping_cell_text(value)
        return party_maps[party_type].get(code.casefold(), code)
    snapshot_cutoffs = {
        "contractor": defaultdict(lambda: "0001-01-01"),
        "supplier": defaultdict(lambda: "0001-01-01"),
    }
    snapshots_used = {"contractor": {}, "supplier": {}}
    cutoff_setting = conn.execute(
        "SELECT value FROM settings WHERE key='historical_payables_through_date'"
    ).fetchone()
    historical_through = mapping_cell_text(
        cutoff_setting["value"] if cutoff_setting else ""
    )
    if not historical_through:
        historical_row = conn.execute(
            "SELECT MAX(purchase_date) through_date FROM historical_payable_lines"
        ).fetchone()
        historical_through = mapping_cell_text(
            historical_row["through_date"] if historical_row else ""
        )
    for row in conn.execute("SELECT * FROM balances"):
        code = party_key(row["party_type"], row["party_code"]) if row["party_type"] in parties else ""
        as_of_date = row["as_of_date"] or "1900-01-01"
        if row["party_type"] in parties and as_of_date <= period_from:
            parties[row["party_type"]][code]["opening"] += row["opening"]
            snapshot_cutoffs[row["party_type"]][code] = as_of_date
            snapshots_used[row["party_type"]][code] = as_of_date
    def add_operational_charge(party_type, raw_code, amount, work_date):
        code = party_key(party_type, raw_code)
        if not code:
            return
        if work_date < snapshot_cutoffs[party_type][code]:
            return
        target = parties[party_type][code]
        if work_date < period_from:
            target["opening"] += amount
        else:
            target["period_charge"] += amount

    # Contractor charges consume the dedicated operational receivable ledger.
    # That ledger is sourced from approved net deliveries and transaction sell
    # prices; issued VAT invoices are deliberately not an input to this balance.
    for row in conn.execute(
        """SELECT work_date,contractor_code,amount
             FROM receivable_ledger_lines
            WHERE status='active' AND work_date<=?
            ORDER BY work_date,id""",
        (period_to,),
    ):
        add_operational_charge(
            "contractor", row["contractor_code"], vnd_round(row["amount"]), row["work_date"],
        )

    # Supplier charges use exactly the same source-line projection as the
    # detailed payable ledger.  This keeps current canonical purchases,
    # legacy orders and the authoritative history snapshot on one identity
    # and cutoff policy, including supplier-code references after a rename.
    try:
        from payable_ledger import discover_payable_sources
    except ImportError:  # pragma: no cover - package invocation
        from .payable_ledger import discover_payable_sources
    historical_summary = {"rows": 0, "amount": 0, "period_rows": 0, "period_amount": 0}
    payable_ledger_existing = {
        row["source_key"]: dict(row)
        for row in conn.execute("SELECT * FROM payable_ledger_lines ORDER BY id")
    }
    payable_sources = sorted(
        discover_payable_sources(conn, payable_ledger_existing),
        key=lambda item: (item["work_date"], item["source_type"], item["source_key"]),
    )
    for row in payable_sources:
        if row["work_date"] > period_to or row["reversal_reason"]:
            continue
        supplier = party_key("supplier", row["supplier_code"])
        if not supplier:
            continue
        if row["work_date"] < snapshot_cutoffs["supplier"][supplier]:
            continue
        amount = vnd_round(row["amount"])
        target = parties["supplier"][supplier]
        if row["source_type"] == "historical_import":
            historical_summary["rows"] += 1
            historical_summary["amount"] += amount
        if row["work_date"] < period_from:
            target["opening"] += amount
        else:
            target["period_charge"] += amount
            if row["source_type"] == "historical_import":
                historical_summary["period_rows"] += 1
                historical_summary["period_amount"] += amount
    for row in conn.execute(
        """SELECT * FROM payments
            WHERE payment_date<=? AND COALESCE(status,'posted')='posted'
              AND ((party_type='contractor' AND kind='receipt')
                OR (party_type='supplier' AND kind='payment'))""",
        (period_to,),
    ):
        party_type = row["party_type"]
        if party_type not in parties:
            continue
        code = party_key(party_type, row["party_code"])
        if row["payment_date"] < snapshot_cutoffs[party_type][code]:
            continue
        target = parties[party_type][code]
        if row["payment_date"] < period_from:
            target["opening"] -= vnd_round(row["amount"])
        else:
            target["period_paid"] += vnd_round(row["amount"])
    for row in conn.execute("SELECT * FROM debt_adjustments WHERE adjustment_date<=?", (period_to,)):
        if row["party_type"] not in parties:
            continue
        code = party_key(row["party_type"], row["party_code"])
        if row["adjustment_date"] < snapshot_cutoffs[row["party_type"]][code]:
            continue
        target = parties[row["party_type"]][code]
        if row["adjustment_date"] < period_from:
            target["opening"] += vnd_round(row["amount"])
        else:
            target["period_adjustment"] += vnd_round(row["amount"])
    for group in parties.values():
        for values in group.values():
            values["closing"] = (
                values["opening"] + values["period_charge"]
                + values["period_adjustment"] - values["period_paid"]
            )
    return {
        "period_from": period_from,
        "period_to": period_to,
        "contractors": dict(parties["contractor"]),
        "suppliers": dict(parties["supplier"]),
        "receivable_source": "receivable_ledger_lines",
        "historical_payables": historical_summary,
        "historical_payables_through_date": historical_through or None,
        "snapshots_used": snapshots_used,
    }


def validate_date_range(period_from: str, period_to: str):
    try:
        start = datetime.strptime(period_from, "%Y-%m-%d").date()
        end = datetime.strptime(period_to, "%Y-%m-%d").date()
    except ValueError:
        raise ValueError("Kỳ công nợ phải dùng ngày hợp lệ dạng YYYY-MM-DD") from None
    if start > end:
        raise ValueError("Ngày bắt đầu phải trước hoặc bằng ngày kết thúc")


def display_date_vn(value: str) -> str:
    try:
        return datetime.strptime(str(value), "%Y-%m-%d").strftime("%d/%m/%Y")
    except (TypeError, ValueError):
        return str(value or "")


def display_period_vn_in_text(value) -> str:
    return re.sub(
        r"(?<!\d)(20\d{2})-(0[1-9]|1[0-2])(?!\d)",
        lambda match: f"{match.group(2)}/{match.group(1)}",
        mapping_cell_text(value),
    )


def display_datetime_vn(value: str) -> str:
    text = str(value or "").strip()
    for pattern in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d"):
        try:
            parsed = datetime.strptime(text, pattern)
            return parsed.strftime("%d/%m/%Y %H:%M:%S" if "%H" in pattern else "%d/%m/%Y")
        except ValueError:
            pass
    return text


def style_export_sheet(ws, title: str, subtitle: str, headers: list[str], money_columns=()):
    end_col = max(len(headers), 1)
    ws.insert_rows(1, 2)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_col)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end_col)
    ws.cell(1, 1, title)
    ws.cell(2, 1, subtitle)
    # These are operational Excel sheets, not dashboard cards.  Keep their
    # appearance close to the customer's plain workbooks: black Times New
    # Roman text, no decorative colour bands, and a title that remains fully
    # visible even on narrow two-column summaries.
    ws.cell(1, 1).font = Font(name="Times New Roman", size=14, bold=True)
    ws.cell(1, 1).alignment = Alignment(
        horizontal="center", vertical="center", wrap_text=True,
        shrink_to_fit=True,
    )
    ws.cell(2, 1).font = Font(name="Times New Roman", size=11, italic=True)
    ws.cell(2, 1).alignment = Alignment(
        horizontal="center", vertical="center", wrap_text=True,
        shrink_to_fit=True,
    )
    ws.row_dimensions[1].height = 34 if end_col <= 2 and len(title) > 24 else 24
    ws.row_dimensions[2].height = 28 if end_col <= 2 and len(subtitle) > 28 else 20
    table_border = Border(
        left=Side(style="thin", color="000000"),
        right=Side(style="thin", color="000000"),
        top=Side(style="thin", color="000000"),
        bottom=Side(style="thin", color="000000"),
    )
    for cell in ws[3]:
        cell.font = Font(name="Times New Roman", size=11, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = table_border
    for row in ws.iter_rows(min_row=4):
        for cell in row:
            cell.font = Font(name="Times New Roman", size=11)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = table_border
    for column in money_columns:
        for row in range(4, ws.max_row + 1):
            ws.cell(row, column).number_format = "#,##0"
    ws.freeze_panes = "A4"
    if ws.max_row >= 3:
        ws.auto_filter.ref = f"A3:{get_column_letter(end_col)}{ws.max_row}"
    for index in range(1, end_col + 1):
        values = [str(ws.cell(row, index).value or "") for row in range(3, ws.max_row + 1)]
        width = min(max([len(value) for value in values] + [10]) + 2, 42)
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    # Short operational summaries are materially easier to read on portrait
    # A4.  Wide detail tables remain landscape so no business column is lost.
    ws.page_setup.orientation = (
        ws.ORIENTATION_PORTRAIT if end_col <= 8 else ws.ORIENTATION_LANDSCAPE
    )
    ws.page_setup.scale = None
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "$3:$3"
    ws.print_area = f"A1:{get_column_letter(end_col)}{max(ws.max_row, 3)}"
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.45
    ws.page_margins.bottom = 0.45
    ws.page_margins.header = 0.15
    ws.page_margins.footer = 0.2
    ws.print_options.horizontalCentered = True
    ws.oddFooter.center.text = "Trang &P / &N"
    ws.sheet_view.showGridLines = False


def send_workbook(workbook: Workbook, filename: str):
    stream = io.BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return send_file(
        stream,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def debt_period_workbook(conn, payload: dict):
    period_from = payload["period_from"]
    period_to = payload["period_to"]
    workbook = Workbook()
    for index, (sheet_name, title, items) in enumerate((
        ("Phải thu", "CÔNG NỢ PHẢI THU", payload["contractors"]),
        ("Phải trả", "CÔNG NỢ PHẢI TRẢ", payload["suppliers"]),
    )):
        ws = workbook.active if index == 0 else workbook.create_sheet()
        ws.title = sheet_name
        headers = ["Đối tượng", "Số dư đầu kỳ", "Phát sinh", "Điều chỉnh", "Đã thu/trả", "Số dư cuối kỳ"]
        ws.append(headers)
        for code, values in sorted(items.items()):
            ws.append([
                code,
                values["opening"],
                values["period_charge"],
                values["period_adjustment"],
                values["period_paid"],
                values["closing"],
            ])
        style_export_sheet(
            ws,
            title,
            f"Từ {display_date_vn(period_from)} đến {display_date_vn(period_to)} · đầu kỳ + phát sinh + điều chỉnh − đã thu/trả",
            headers,
            money_columns=(2, 3, 4, 5, 6),
        )

    ws = workbook.create_sheet("Thu chi")
    payment_headers = ["Ngày", "Loại", "Nhóm", "Đối tượng", "Số tiền", "Nội dung", "Ngày tạo"]
    ws.append(payment_headers)
    for row in conn.execute(
        "SELECT * FROM payments WHERE payment_date>=? AND payment_date<=? "
        "AND COALESCE(status,'posted')='posted' ORDER BY payment_date,id",
        (period_from, period_to),
    ):
        ws.append([
            display_date_vn(row["payment_date"]),
            "Thu khách hàng" if row["kind"] == "receipt" else "Trả nhà cung cấp",
            "Phải thu" if row["party_type"] == "contractor" else "Phải trả",
            row["party_code"],
            row["amount"],
            row["note"],
            display_datetime_vn(row["created_at"]),
        ])
    style_export_sheet(ws, "LỊCH SỬ THU – CHI", f"Từ {display_date_vn(period_from)} đến {display_date_vn(period_to)}", payment_headers, (5,))

    ws = workbook.create_sheet("Điều chỉnh")
    adjustment_headers = ["Ngày", "Nhóm", "Đối tượng", "Số điều chỉnh", "Lý do", "Ngày tạo"]
    ws.append(adjustment_headers)
    for row in conn.execute(
        "SELECT * FROM debt_adjustments WHERE adjustment_date>=? AND adjustment_date<=? ORDER BY adjustment_date,id",
        (period_from, period_to),
    ):
        ws.append([
            display_date_vn(row["adjustment_date"]),
            "Phải thu" if row["party_type"] == "contractor" else "Phải trả",
            row["party_code"],
            row["amount"],
            row["note"],
            display_datetime_vn(row["created_at"]),
        ])
    style_export_sheet(ws, "ĐIỀU CHỈNH CÔNG NỢ", f"Từ {display_date_vn(period_from)} đến {display_date_vn(period_to)}", adjustment_headers, (4,))

    ws = workbook.create_sheet("Chi tiết phải trả cũ")
    payable_headers = [
        "Ngày", "Bếp", "Tên hàng", "Số lượng đặt", "ĐVT", "NCC", "Giá mua",
        "Hỏng", "Thêm", "Giảm", "Thiếu", "SL thực tế", "Tiền file nguồn",
        "Tiền theo công thức", "Chênh lệch", "Thành tiền công nợ", "Ghi chú", "File nguồn", "Dòng",
    ]
    ws.append(payable_headers)
    for row in conn.execute(
        """SELECT * FROM historical_payable_lines
           WHERE purchase_date>=? AND purchase_date<=? ORDER BY purchase_date,id""",
        (period_from, period_to),
    ):
        ws.append([
            display_date_vn(row["purchase_date"]), row["kitchen"], row["item_name"], row["qty"], row["unit"],
            row["supplier"], row["buy_price"], row["damaged_qty"], row["added_qty"],
            row["reduced_qty"], row["missing_qty"], row["actual_qty"], row["source_amount"],
            row["calculated_amount"], row["source_amount"] - row["calculated_amount"], row["amount"],
            row["note"], row["source_file"], row["source_row"],
        ])
    style_export_sheet(
        ws, "CHI TIẾT PHẢI TRẢ TỪ FILE KHÁCH",
        f"Từ {display_date_vn(period_from)} đến {display_date_vn(period_to)} · giữ tiền khách đã chốt và hiển thị chênh lệch công thức",
        payable_headers, (7, 13, 14, 15, 16),
    )
    return workbook


def payroll_workbook(conn, month: str):
    rows = payroll_rows(conn, month)
    if not PAYROLL_TEMPLATE_SOURCE.is_file():
        raise FileNotFoundError(
            "Thiếu mẫu LƯƠNG XƯỞNG đã chốt; không tạo bảng lương bằng mẫu tự đoán"
        )
    workbook = load_workbook(PAYROLL_TEMPLATE_SOURCE, data_only=False, read_only=False)
    ws = workbook["LƯƠNG XƯỞNG"]
    headers = [
        "STT", "HỌ VÀ TÊN", "CHỨC VỤ", "Công HC", "Công TC", "Công CN",
        "Công Đêm", "Công NL", "Phụ cấp", "Trách nhiệm", "Tổng Lương",
        "Tạm ứng", "Thử việc", "BHXH người lao động đóng", "TỔNG",
        "BHXH Công ty đóng",
    ]
    data_styles = [deepcopy(ws.cell(4, column)._style) for column in range(1, 17)]
    total_styles = [deepcopy(ws.cell(5, column)._style) for column in range(1, 17)]
    for column, label in enumerate(headers, 1):
        ws.cell(3, column, label)
    year, month_number = (int(value) for value in month.split("-"))
    ws["A1"] = "LƯƠNG THÁNG"
    ws["I1"] = month_number
    ws["J1"] = year
    ws["A2"] = "Lương cơ bản:"
    # The customer's row 2 lists the base-salary rates actually used in the
    # month (for example 5,000,000 and 8,000,000), rather than a prose note.
    # Reproduce that behavior from the configured staff records.
    for column in range(3, 17):
        ws.cell(2, column, None)
    base_salaries = sorted({
        round(as_number(item.get("base_salary")))
        for item in rows if as_number(item.get("base_salary")) > 0
    })
    for offset, amount in enumerate(base_salaries[:14], start=3):
        ws.cell(2, offset, amount)
        ws.cell(2, offset).number_format = "#,##0"
    for row_number in range(4, max(ws.max_row, len(rows) + 4) + 1):
        for column in range(1, 17):
            ws.cell(row_number, column, None)
    for index, item in enumerate(rows, 1):
        row_number = index + 3
        values = [
            index, item["full_name"], item["role_name"], item["normal_hours"],
            item["overtime_hours"], item["sunday_hours"], item["night_hours"],
            item["holiday_hours"], item["allowance"], item["responsibility"],
            item["gross_salary"], item["advance"], item["probation_deduction"],
            item["bhxh_employee"], item["net_salary"], item["bhxh_company"],
        ]
        for column, value in enumerate(values, 1):
            cell = ws.cell(row_number, column, value)
            cell._style = deepcopy(data_styles[column - 1])
            if column >= 9:
                cell.number_format = "#,##0"
    total_row = len(rows) + 4
    # The customer workbook merges the first three cells for the total label.
    # Preserve that merge so ``TỔNG`` is not clipped inside the very narrow
    # STT column when printed.
    for merged_range in list(ws.merged_cells.ranges):
        if merged_range.min_row == total_row and merged_range.max_row == total_row:
            ws.unmerge_cells(str(merged_range))
    ws.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=3)
    for column in range(1, 17):
        ws.cell(total_row, column)._style = deepcopy(total_styles[column - 1])
        ws.cell(total_row, column, None)
    ws.cell(total_row, 1, "TỔNG")
    ws.cell(total_row, 1).alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(total_row, 15, f"=SUM(O4:O{total_row - 1})" if rows else 0)
    ws.cell(total_row, 15).number_format = "#,##0"
    # The customer workbook keeps the ``Tạm ứng`` column unusually narrow.
    # Real six/seven-digit advances otherwise print as ``#####`` even though
    # the underlying value is valid. Keep the golden layout and widen only
    # this numeric column enough for an ordinary VND amount.
    ws.column_dimensions["L"].width = max(ws.column_dimensions["L"].width or 0, 10)
    ws.freeze_panes = "A4"
    ws.print_title_rows = "$1:$3"
    ws.print_area = f"A1:P{total_row}"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.scale = None
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_view.showGridLines = False

    ws = workbook.create_sheet("Chi phí theo bếp")
    labor_headers = ["Bếp", "Chi phí lao động"]
    ws.append(labor_headers)
    for row in conn.execute(
        """SELECT kitchen,SUM(amount) amount FROM kitchen_labor_costs
           WHERE substr(work_date,1,7)=? GROUP BY kitchen ORDER BY kitchen""",
        (month,),
    ):
        ws.append([row["kitchen"], row["amount"]])
    year_text, month_text = month.split("-", 1)
    style_export_sheet(
        ws,
        f"CHI PHÍ LAO ĐỘNG THÁNG {month_text}/{year_text}",
        "Tổng hợp theo bếp",
        labor_headers,
        (2,),
    )
    return workbook


def meal_plan_payload(conn, work_date=""):
    params = []
    where = ""
    if work_date:
        where = "WHERE mp.work_date=?"
        params.append(work_date)
    plans = [dict(row) for row in conn.execute(
        f"""SELECT mp.*,COALESCE(NULLIF(mp.unit_code,''),ku.unit_code,'') mapped_unit
             FROM meal_plans mp LEFT JOIN kitchen_units ku ON ku.kitchen_code=mp.kitchen
             {where} ORDER BY mp.work_date DESC,mp.kitchen,mp.shift,mp.id""", params
    )]
    meal_totals_by_kitchen_day = defaultdict(float)
    for plan in plans:
        meal_totals_by_kitchen_day[(plan["work_date"], plan["kitchen"])] += max(
            as_number(plan.get("meal_count")), 0
        )
    labor_by_kitchen_day = {
        (row["work_date"], row["kitchen"]): as_number(row["amount"])
        for row in conn.execute(
            """SELECT work_date,kitchen,SUM(amount) amount
               FROM kitchen_labor_costs GROUP BY work_date,kitchen"""
        )
    }
    for plan in plans:
        items = [dict(row) for row in conn.execute(
            "SELECT * FROM meal_plan_items WHERE plan_id=? ORDER BY dish_name,product_name", (plan["id"],)
        )]
        for item in items:
            item["required_qty"] = plan["meal_count"] * item["norm_qty"]
            item["cost"] = item["required_qty"] * item["buy_price"]
            if not item.get("applicable_meal_count"):
                item["applicable_meal_count"] = plan["meal_count"]
            if not item.get("source_norm_per_1000"):
                item["source_norm_per_1000"] = item["norm_qty"] * 1000
        plan["items"] = items
        plan["food_cost"] = sum(item["cost"] for item in items)
        plan["revenue"] = plan["meal_count"] * plan.get("meal_price", 0)
        kitchen_day_key = (plan["work_date"], plan["kitchen"])
        kitchen_day_meals = meal_totals_by_kitchen_day.get(kitchen_day_key, 0)
        kitchen_day_labor = labor_by_kitchen_day.get(kitchen_day_key, 0)
        plan["labor_cost"] = (
            kitchen_day_labor * max(as_number(plan["meal_count"]), 0) / kitchen_day_meals
            if kitchen_day_meals > 0 else 0
        )
        plan["total_cost"] = (
            plan["food_cost"] + plan.get("other_cost", 0) + plan["labor_cost"]
        )
        plan["profit"] = plan["revenue"] - plan["total_cost"]
        plan["profit_margin"] = plan["profit"] / plan["revenue"] if plan["revenue"] else 0
        plan["cost_per_meal"] = plan["total_cost"] / plan["meal_count"] if plan["meal_count"] else 0
        try:
            plan["source_financials"] = json.loads(plan.get("source_financials_json") or "{}")
        except (TypeError, ValueError):
            plan["source_financials"] = {}
        plan["warnings"] = []
        if not plan["mapped_unit"]:
            plan["warnings"].append("Bếp chưa được ghép XCOM (xưởng cơm)")
        if any(item["buy_price"] <= 0 for item in items):
            plan["warnings"].append("Có nguyên liệu thiếu giá")
        expected_price_source = f"HATRAN {plan['work_date'][:7]}"
        wrong_period = sum(
            not mapping_cell_text(item.get("price_source")).startswith(expected_price_source)
            for item in items
        )
        if wrong_period:
            plan["warnings"].append(
                f"Có {wrong_period} nguyên liệu chưa dùng giá {expected_price_source}"
            )
        if plan.get("meal_price", 0) <= 0:
            plan["warnings"].append("Chưa có đơn giá suất ăn")
        plan["warnings"].extend(plan["source_financials"].get("warnings") or [])
    return plans


def meal_po_workbook(plans, work_date: str):
    groups = defaultdict(list)
    for plan in plans:
        for item in plan["items"]:
            groups[plan["mapped_unit"] or "CHƯA CÓ XCOM"].append((plan, item))
    wb = Workbook()
    wb.remove(wb.active)
    for unit_code, rows in sorted(groups.items()):
        unique_plans = []
        seen_plans = set()
        for plan, _ in rows:
            if plan["id"] not in seen_plans:
                seen_plans.add(plan["id"])
                unique_plans.append(plan)
        approval_label = "ĐÃ DUYỆT" if all(
            plan["status"] == "approved" for plan in unique_plans
        ) else "BẢN NHÁP"
        title = re.sub(r"[\\/*?:\[\]]", "-", unit_code)[:31]
        ws = wb.create_sheet(title or "PO")
        ws.append([f"PO XƯỞNG CƠM – {approval_label}"])
        ws.merge_cells("A1:M1")
        ws["A1"].font = Font(name="Times New Roman", size=14, bold=True)
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        price_sources = sorted({
            display_period_vn_in_text(item["price_source"])
            for _, item in rows if item.get("price_source")
        })
        ws.append([f"Ngày {display_date_vn(work_date)}", "", "", f"XCOM {unit_code}", "", "", "",
                   f"Nguồn giá: {', '.join(price_sources)}"])
        for cell in ws[2]:
            cell.font = Font(name="Times New Roman", size=11, bold=True)
        ws.append(["STT", "Bếp / ca", "Mã hàng", "Tên nguyên liệu", "ĐVT", "Suất áp dụng",
                   "Định lượng/1.000 suất", "Số lượng cần", "Đơn giá", "Thành tiền", "NCC", "Nguồn giá"])
        for index, (plan, item) in enumerate(rows, 1):
            ws.append([
                index, f"{plan['kitchen']} / {plan['shift']}", item["product_code"], item["product_name"],
                item["unit"], item["applicable_meal_count"], item["source_norm_per_1000"],
                item["required_qty"], item["buy_price"], item["cost"], item["supplier"],
                display_period_vn_in_text(item["price_source"]),
            ])
        for cell in ws[3]:
            cell.font = Font(name="Times New Roman", size=11, bold=True)
            cell.border = Border(
                left=Side(style="thin", color="000000"),
                right=Side(style="thin", color="000000"),
                top=Side(style="thin", color="000000"),
                bottom=Side(style="thin", color="000000"),
            )
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        detail_end = ws.max_row
        ws.append([])
        ws.append(["TỔNG HỢP SUẤT ĂN / COST"])
        ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=12)
        ws.cell(ws.max_row, 1).font = Font(name="Times New Roman", size=11, bold=True)
        ws.cell(ws.max_row, 1).alignment = Alignment(horizontal="center")
        ws.append(["Bếp / ca", "Số thực đơn", "Suất/thực đơn", "Tổng suất", "Đơn giá suất",
                   "Doanh thu", "Chi phí thực phẩm", "Chi phí khác", "Chi phí lao động",
                   "Tổng chi", "Lợi nhuận", "Biên lợi nhuận", "Thực đơn"])
        summary_header = ws.max_row
        for plan in unique_plans:
            ws.append([
                f"{plan['kitchen']} / {plan['shift']}", plan.get("menu_count", 1),
                plan.get("servings_per_menu") or plan["meal_count"], plan["meal_count"],
                plan.get("meal_price", 0), plan["revenue"], plan["food_cost"], plan.get("other_cost", 0),
                plan.get("labor_cost", 0), plan["total_cost"], plan["profit"], plan["profit_margin"],
                plan.get("note", ""),
            ])
        for cell in ws[summary_header]:
            cell.font = Font(name="Times New Roman", size=11, bold=True)
            cell.border = Border(
                left=Side(style="thin", color="000000"),
                right=Side(style="thin", color="000000"),
                top=Side(style="thin", color="000000"),
                bottom=Side(style="thin", color="000000"),
            )
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.freeze_panes = "A4"
        ws.auto_filter.ref = f"A3:L{detail_end}"
        widths = [7, 20, 14, 30, 10, 14, 20, 15, 16, 16, 16, 16, 32]
        for index, width in enumerate(widths, 1):
            ws.column_dimensions[chr(64 + index)].width = width
        for row in range(4, detail_end + 1):
            for col in (9, 10):
                ws.cell(row, col).number_format = "#,##0"
        for row in range(summary_header + 1, ws.max_row + 1):
            for col in range(5, 12):
                ws.cell(row, col).number_format = "#,##0"
            ws.cell(row, 12).number_format = "0.0%"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
        ws.page_setup.scale = None
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.print_title_rows = "$3:$3"
        ws.print_area = f"A1:M{ws.max_row}"
        ws.page_margins.left = 0.2
        ws.page_margins.right = 0.2
        ws.page_margins.top = 0.4
        ws.page_margins.bottom = 0.4
        ws.page_margins.header = 0.15
        ws.page_margins.footer = 0.2
        ws.oddFooter.center.text = "Trang &P / &N"
        ws.sheet_view.showGridLines = False
    if not wb.sheetnames:
        wb.create_sheet("PO")
    return wb


def employee_code(name: str) -> str:
    text = unicodedata.normalize("NFD", name.upper())
    text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn")
    return re.sub(r"[^A-Z0-9]+", "", text)[:30] or "NV"


LEGACY_ATTENDANCE_SOURCE_PREFIX = "LEGACY_ATTENDANCE:"
LEGACY_PAYROLL_NOTE_PREFIX = "[IMPORT_CHAM_CONG:"


def validate_legacy_attendance_period(period: str) -> tuple[int, int]:
    text = str(period or "").strip()
    if not re.fullmatch(r"\d{4}-\d{2}", text):
        raise ValueError("Kỳ chấm công phải có dạng YYYY-MM")
    year, month_num = map(int, text.split("-"))
    try:
        date(year, month_num, 1)
    except ValueError:
        raise ValueError("Kỳ chấm công không hợp lệ") from None
    return year, month_num


def legacy_name_identity(name: str) -> str:
    return re.sub(r"\s+", " ", unicodedata.normalize("NFKC", str(name or "")).strip()).casefold()


def allocate_legacy_employee_codes(conn, names: list[str]) -> dict[str, str]:
    existing = [dict(row) for row in conn.execute("SELECT employee_code,full_name FROM staff")]
    by_identity = {legacy_name_identity(row["full_name"]): row["employee_code"] for row in existing}
    code_owner = {row["employee_code"].upper(): legacy_name_identity(row["full_name"]) for row in existing}
    result = {}
    for name in names:
        identity = legacy_name_identity(name)
        if not identity or identity in result:
            continue
        if identity in by_identity:
            result[identity] = by_identity[identity]
            continue
        base = employee_code(name)
        if base not in code_owner:
            code = base
        else:
            digest = hashlib.sha256(identity.encode("utf-8")).hexdigest().upper()
            code = ""
            for suffix_len in (6, 8, 10, 12):
                candidate = f"{base[:29 - suffix_len]}-{digest[:suffix_len]}"
                if candidate not in code_owner or code_owner[candidate] == identity:
                    code = candidate
                    break
            if not code:
                raise ValueError(f"Không tạo được mã nhân sự không trùng cho {name}")
        result[identity] = code
        code_owner[code] = identity
    return result


def legacy_formula_missing(formula_sheet, values_sheet, row: int, column: int) -> bool:
    if formula_sheet is None:
        return False
    formula_cell = formula_sheet.cell(row, column)
    return formula_cell.data_type == "f" and values_sheet.cell(row, column).value is None


def legacy_number(values_sheet, formula_sheet, row: int, column: int, label: str,
                  errors: list[str], warnings: list[str] | None = None) -> float:
    cell = values_sheet.cell(row, column)
    if legacy_formula_missing(formula_sheet, values_sheet, row, column):
        errors.append(f"{label} ({cell.coordinate}) là công thức chưa có kết quả lưu trong file")
        return 0.0
    value = cell.value
    if value in (None, ""):
        return 0.0
    parsed = as_number(value, float("nan"))
    if not math.isfinite(parsed):
        message = f"{label} ({cell.coordinate}) ghi '{value}', được hiểu là 0"
        if warnings is not None:
            warnings.append(message)
        else:
            errors.append(f"{label} ({cell.coordinate}) không phải số hợp lệ: {value}")
        return 0.0
    return float(parsed)


def legacy_money_number(values_sheet, formula_sheet, row: int, column: int, label: str,
                        errors: list[str], warnings: list[str]) -> float:
    cell = values_sheet.cell(row, column)
    if legacy_formula_missing(formula_sheet, values_sheet, row, column):
        errors.append(f"{label} ({cell.coordinate}) là công thức chưa có kết quả lưu trong file")
        return 0.0
    value = cell.value
    if value in (None, ""):
        return 0.0
    parsed = as_number(value, float("nan"))
    if math.isfinite(parsed):
        return float(parsed)
    tokens = [token.strip() for token in re.findall(r"[-+]?\d[\d\s.,]*", str(value))]
    numbers = [as_number(token, float("nan")) for token in tokens]
    numbers = [number for number in numbers if math.isfinite(number)]
    if len(numbers) == 1:
        warnings.append(
            f"{label} ({cell.coordinate}) ghi kèm chữ '{value}'; hệ thống lấy số {numbers[0]:g}"
        )
        return float(numbers[0])
    if not numbers:
        warnings.append(f"{label} ({cell.coordinate}) ghi '{value}', được hiểu là 0")
        return 0.0
    errors.append(f"{label} ({cell.coordinate}) có nhiều số, không xác định được số tiền: {value}")
    return 0.0


def legacy_cell_date(cell) -> str:
    value = cell.value
    parsed = as_date(value)
    if parsed:
        return parsed
    if isinstance(value, (int, float)) and 20_000 <= float(value) <= 80_000:
        try:
            converted = from_excel(value)
            return converted.date().isoformat() if isinstance(converted, datetime) else converted.isoformat()
        except (TypeError, ValueError, OverflowError):
            return ""
    return ""


def legacy_attendance_database_state_hash(conn, period: str) -> str:
    payload = {
        "staff": [tuple(row) for row in conn.execute(
            "SELECT employee_code,full_name,role_name,base_salary FROM staff ORDER BY employee_code"
        )],
        "attendance": [tuple(row) for row in conn.execute(
            """SELECT s.employee_code,a.work_date,a.normal_hours,a.overtime_hours,a.sunday_hours,
                      a.night_hours,a.holiday_hours,a.note,a.source
               FROM attendance_entries a JOIN staff s ON s.id=a.employee_id
               WHERE substr(a.work_date,1,7)=? ORDER BY s.employee_code,a.work_date""",
            (period,),
        )],
        "payroll": [tuple(row) for row in conn.execute(
            """SELECT s.employee_code,p.allowance,p.responsibility,p.advance,p.probation_deduction,
                      p.bhxh_employee_amount,p.bhxh_company_amount,p.gross_override,p.net_override,
                      p.use_override,p.note
               FROM payroll_adjustments p JOIN staff s ON s.id=p.employee_id
               WHERE p.month=? ORDER BY s.employee_code""",
            (period,),
        )],
        "labor": [tuple(row) for row in conn.execute(
            """SELECT work_date,kitchen,amount,source FROM kitchen_labor_costs
               WHERE substr(work_date,1,7)=? ORDER BY work_date,kitchen,source""",
            (period,),
        )],
    }
    encoded = json.dumps(payload, ensure_ascii=False, separators=(",", ":"), default=str)
    return hashlib.sha256(encoded.encode("utf-8")).hexdigest().upper()


def parse_legacy_attendance_workbooks(conn, values_book, formulas_book, source_name: str,
                                      period: str) -> dict:
    year, month_num = validate_legacy_attendance_period(period)
    errors = []
    warnings = []
    attendance_sheet = next(
        (ws for ws in values_book.worksheets if "xưởng cơm" in ws.title.lower()), None
    )
    if attendance_sheet is None:
        raise ValueError("Không tìm thấy sheet Xưởng Cơm trong file chấm công")
    attendance_formula_sheet = (
        formulas_book[attendance_sheet.title] if formulas_book and attendance_sheet.title in formulas_book.sheetnames else None
    )
    sheet_year = legacy_number(
        attendance_sheet, attendance_formula_sheet, 2, 3, "Năm trên sheet Xưởng Cơm", errors,
    )
    sheet_month = legacy_number(
        attendance_sheet, attendance_formula_sheet, 3, 3, "Tháng trên sheet Xưởng Cơm", errors,
    )
    if not sheet_year.is_integer() or not sheet_month.is_integer():
        errors.append("Năm/tháng trên sheet Xưởng Cơm phải là số nguyên")
    elif f"{int(sheet_year):04d}-{int(sheet_month):02d}" != period:
        errors.append(
            f"Sheet Xưởng Cơm ghi kỳ {int(sheet_year):04d}-{int(sheet_month):02d}, "
            f"khác kỳ đã chọn {period}"
        )

    payroll_sheet = next(
        (ws for ws in values_book.worksheets if "lương xưởng" in ws.title.lower()), None
    )
    payroll_formula_sheet = (
        formulas_book[payroll_sheet.title] if payroll_sheet and formulas_book
        and payroll_sheet.title in formulas_book.sheetnames else None
    )
    base_salary = 0.0
    if payroll_sheet is not None:
        base_salary = legacy_number(
            payroll_sheet, payroll_formula_sheet, 2, 3, "Lương cơ bản", errors,
        )
        payroll_month = as_number(payroll_sheet["I1"].value, float("nan"))
        payroll_year = as_number(payroll_sheet["J1"].value, float("nan"))
        if math.isfinite(payroll_month) or math.isfinite(payroll_year):
            if not (math.isfinite(payroll_month) and math.isfinite(payroll_year)
                    and payroll_month.is_integer() and payroll_year.is_integer()):
                errors.append("Kỳ trên sheet LƯƠNG XƯỞNG chưa đủ tháng và năm")
            elif f"{int(payroll_year):04d}-{int(payroll_month):02d}" != period:
                errors.append(
                    f"Sheet LƯƠNG XƯỞNG ghi kỳ {int(payroll_year):04d}-{int(payroll_month):02d}, "
                    f"khác kỳ đã chọn {period}"
                )

    staff_by_identity = {}

    def remember_staff(name, role=""):
        identity = legacy_name_identity(name)
        if not identity:
            return identity
        record = staff_by_identity.setdefault(identity, {
            "name": re.sub(r"\s+", " ", str(name).strip()),
            "role": "",
            "base_salary": base_salary,
        })
        if role and not record["role"]:
            record["role"] = re.sub(r"\s+", " ", str(role).strip())
        return identity

    attendance_records = []
    for row in range(6, attendance_sheet.max_row + 1):
        name = re.sub(r"\s+", " ", str(attendance_sheet.cell(row, 2).value or "").strip())
        if (not name or name == "0" or mapping_key(name) in {
                "hoten", "hovaten", "tong", "tongluongthang",
        }):
            continue
        identity = remember_staff(name, attendance_sheet.cell(row, 3).value or "")
        overtime_row = (
            row + 1 if row + 1 <= attendance_sheet.max_row
            and not attendance_sheet.cell(row + 1, 2).value else None
        )
        wage_cells = []
        for day_index, column in enumerate(range(6, 37), start=1):
            try:
                work_day = date(year, month_num, day_index)
            except ValueError:
                break
            header_cell = attendance_sheet.cell(4, column)
            if legacy_formula_missing(attendance_formula_sheet, attendance_sheet, 4, column):
                errors.append(
                    f"Ngày ở {attendance_sheet.title}!{header_cell.coordinate} là công thức chưa có kết quả lưu"
                )
            else:
                header_date = legacy_cell_date(header_cell)
                if header_date and header_date != work_day.isoformat():
                    errors.append(
                        f"{attendance_sheet.title}!{header_cell.coordinate} ghi {header_date}, "
                        f"khác ngày phải có {work_day.isoformat()}"
                    )
            normal = legacy_number(
                attendance_sheet, attendance_formula_sheet, row, column,
                f"Giờ thường của {name}", errors, warnings,
            )
            overtime = 0.0
            if overtime_row:
                overtime = legacy_number(
                    attendance_sheet, attendance_formula_sheet, overtime_row, column,
                    f"Giờ tăng ca của {name}", errors, warnings,
                )
            if normal < 0 or overtime < 0:
                cells = [attendance_sheet.cell(row, column).coordinate]
                if overtime_row:
                    cells.append(attendance_sheet.cell(overtime_row, column).coordinate)
                errors.append(
                    f"{attendance_sheet.title}!{'/'.join(cells)} có giờ âm; không được nạp"
                )
                continue
            if normal > 24 or overtime > 24:
                wage_cells.append(attendance_sheet.cell(row, column).coordinate)
                continue
            if normal == 0 and overtime == 0:
                continue
            sunday = work_day.weekday() == 6
            attendance_records.append({
                "identity": identity,
                "work_date": work_day.isoformat(),
                "normal_hours": 0.0 if sunday else normal,
                "overtime_hours": 0.0 if sunday else overtime,
                "sunday_hours": normal + overtime if sunday else 0.0,
                "night_hours": 0.0,
                "holiday_hours": 0.0,
                "note": "",
            })
        if wage_cells:
            visible = ", ".join(wage_cells[:5])
            extra = f" và {len(wage_cells) - 5} ô khác" if len(wage_cells) > 5 else ""
            warnings.append(
                f"{name}: {visible}{extra} có số >24, được coi là tiền công/ngày nên không nhập thành giờ"
            )

    payroll_records_by_identity = {}
    if payroll_sheet is not None:
        for row in range(4, payroll_sheet.max_row + 1):
            name = re.sub(r"\s+", " ", str(payroll_sheet.cell(row, 2).value or "").strip())
            if not name or name == "0" or mapping_key(name).startswith("tong"):
                continue
            identity = remember_staff(name, payroll_sheet.cell(row, 3).value or "")
            next_row = row + 1
            detail_row = row
            if (next_row <= payroll_sheet.max_row
                    and not payroll_sheet.cell(next_row, 1).value
                    and not payroll_sheet.cell(next_row, 2).value
                    and any(as_number(payroll_sheet.cell(next_row, col).value) for col in range(4, 17))):
                detail_row = next_row
            values = {
                "allowance": legacy_number(payroll_sheet, payroll_formula_sheet, detail_row, 9, f"Phụ cấp {name}", errors, warnings),
                "responsibility": legacy_number(payroll_sheet, payroll_formula_sheet, detail_row, 10, f"Trách nhiệm {name}", errors, warnings),
                "advance": legacy_number(payroll_sheet, payroll_formula_sheet, detail_row, 12, f"Tạm ứng {name}", errors, warnings),
                "probation": legacy_number(payroll_sheet, payroll_formula_sheet, detail_row, 13, f"Thử việc {name}", errors, warnings),
                "bhxh_employee": legacy_number(payroll_sheet, payroll_formula_sheet, detail_row, 14, f"BHXH NLĐ {name}", errors, warnings),
                "bhxh_company": legacy_number(payroll_sheet, payroll_formula_sheet, detail_row, 16, f"BHXH công ty {name}", errors, warnings),
                "gross": legacy_number(payroll_sheet, payroll_formula_sheet, detail_row, 11, f"Tổng lương {name}", errors, warnings),
                "net": legacy_number(payroll_sheet, payroll_formula_sheet, detail_row, 15, f"Thực lĩnh {name}", errors, warnings),
            }
            if values["gross"] == 0 and values["net"] != 0:
                values["gross"] = values["net"]
            payroll_records_by_identity[identity] = {"identity": identity, **values}

        for row in range(4, payroll_sheet.max_row + 1):
            label = str(payroll_sheet.cell(row, 16).value or "").strip()
            if not label.lower().startswith("cho "):
                continue
            supplement = legacy_number(
                payroll_sheet, payroll_formula_sheet, row, 15,
                f"Khoản bổ sung {label}", errors, warnings,
            )
            if supplement == 0:
                continue
            label_code = employee_code(
                re.sub(r"^cho\s+(chị|anh|em)?\s*", "", label, flags=re.IGNORECASE)
            )
            match = next((
                identity for identity, record in payroll_records_by_identity.items()
                if employee_code(staff_by_identity[identity]["name"]) == label_code
                or employee_code(staff_by_identity[identity]["name"]) in label_code
                or label_code in employee_code(staff_by_identity[identity]["name"])
            ), None)
            if match:
                payroll_records_by_identity[match]["gross"] += supplement
                payroll_records_by_identity[match]["net"] += supplement
            else:
                warnings.append(f"Không ghép được khoản bổ sung '{label}' ở dòng {row}")

    labor_records = []
    labor_sheet = next(
        (ws for ws in values_book.worksheets if "chấm công chợ" in ws.title.lower()), None
    )
    labor_formula_sheet = (
        formulas_book[labor_sheet.title] if labor_sheet and formulas_book
        and labor_sheet.title in formulas_book.sheetnames else None
    )
    if labor_sheet is not None:
        column_scores = []
        for column in range(1, min(labor_sheet.max_column, 6) + 1):
            dated_rows = []
            for row in range(4, min(labor_sheet.max_row, 45) + 1):
                parsed = legacy_cell_date(labor_sheet.cell(row, column))
                if parsed:
                    dated_rows.append((row, parsed))
            column_scores.append((len(dated_rows), column, dated_rows))
        score, date_column, dated_rows = max(column_scores, default=(0, 0, []))
        if score < 2:
            errors.append(f"Sheet {labor_sheet.title}: không xác định được cột ngày của chấm công chợ")
        else:
            mismatches = [(row, value) for row, value in dated_rows if value[:7] != period]
            if mismatches:
                row, value = mismatches[0]
                errors.append(
                    f"Sheet {labor_sheet.title}!{labor_sheet.cell(row, date_column).coordinate} "
                    f"ghi kỳ {value[:7]}, khác kỳ đã chọn {period}"
                )
            first_data_row = min(row for row, _ in dated_rows)
            header_candidates = []
            for header_row in range(1, first_data_row):
                text_count = sum(
                    1 for column in range(date_column + 1, min(labor_sheet.max_column, 24) + 1)
                    if isinstance(labor_sheet.cell(header_row, column).value, str)
                    and labor_sheet.cell(header_row, column).value.strip()
                )
                header_candidates.append((text_count, header_row))
            _, header_row = max(header_candidates, default=(0, max(1, first_data_row - 1)))
            kitchens = {}
            for column in range(date_column + 1, min(labor_sheet.max_column, 24) + 1):
                kitchen = re.sub(
                    r"\s+", " ", str(labor_sheet.cell(header_row, column).value or "").strip()
                ).upper()
                if kitchen:
                    kitchens[column] = kitchen
            if not kitchens:
                errors.append(f"Sheet {labor_sheet.title}: không tìm thấy tên bếp ở hàng tiêu đề")
            seen_labor = set()
            for row, work_date in dated_rows:
                if work_date[:7] != period:
                    continue
                for column, kitchen in kitchens.items():
                    amount = legacy_money_number(
                        labor_sheet, labor_formula_sheet, row, column,
                        f"Chi phí {kitchen} ngày {work_date}", errors, warnings,
                    )
                    if amount < 0:
                        errors.append(
                            f"{labor_sheet.title}!{labor_sheet.cell(row, column).coordinate} có chi phí âm"
                        )
                        continue
                    if amount == 0:
                        continue
                    key = (work_date, kitchen)
                    if key in seen_labor:
                        errors.append(f"Sheet {labor_sheet.title} trùng chi phí {kitchen} ngày {work_date}")
                        continue
                    seen_labor.add(key)
                    labor_records.append({
                        "work_date": work_date, "kitchen": kitchen, "amount": amount,
                    })

    if errors:
        message = "File chấm công chưa thể nạp: " + " · ".join(errors[:12])
        if len(errors) > 12:
            message += f" · và {len(errors) - 12} lỗi khác"
        raise ValueError(message)

    names = [record["name"] for record in staff_by_identity.values()]
    codes = allocate_legacy_employee_codes(conn, names)
    staff_records = []
    for identity, record in staff_by_identity.items():
        staff_records.append({"identity": identity, "employee_code": codes[identity], **record})
    for item in attendance_records:
        item["employee_code"] = codes[item.pop("identity")]
    payroll_records = []
    for item in payroll_records_by_identity.values():
        item["employee_code"] = codes[item.pop("identity")]
        payroll_records.append(item)

    manual_attendance = 0
    for item in attendance_records:
        row = conn.execute(
            """SELECT a.source FROM attendance_entries a JOIN staff s ON s.id=a.employee_id
               WHERE s.employee_code=? AND a.work_date=?""",
            (item["employee_code"], item["work_date"]),
        ).fetchone()
        if row and str(row["source"] or "").lower() == "manual":
            manual_attendance += 1
    manual_payroll = 0
    for item in payroll_records:
        row = conn.execute(
            """SELECT p.note FROM payroll_adjustments p JOIN staff s ON s.id=p.employee_id
               WHERE s.employee_code=? AND p.month=?""",
            (item["employee_code"], period),
        ).fetchone()
        if row and not legacy_payroll_note_is_imported(row["note"], period):
            manual_payroll += 1
    if manual_attendance:
        warnings.append(
            f"Có {manual_attendance} dòng đã sửa tay; hệ thống sẽ giữ nguyên, không ghi đè từ file"
        )
    if manual_payroll:
        warnings.append(
            f"Có {manual_payroll} khoản lương đã sửa tay; hệ thống sẽ giữ nguyên, không ghi đè từ file"
        )

    snapshot = {
        "period": period,
        "staff": staff_records,
        "attendance": attendance_records,
        "payroll": payroll_records,
        "labor": labor_records,
    }
    counts = {
        "staff": len(staff_records),
        "attendance_entries": len(attendance_records),
        "payroll_overrides": len(payroll_records),
        "labor_cost_entries": len(labor_records),
        "warnings": len(warnings),
        "manual_attendance_preserved": manual_attendance,
        "manual_payroll_preserved": manual_payroll,
    }
    return {
        "month": period,
        "can_confirm": True,
        "counts": counts,
        "warnings": warnings[:100],
        "snapshot": snapshot,
    }


def legacy_payroll_note_is_imported(note, period: str) -> bool:
    text = str(note or "").strip()
    return (
        text.startswith(f"{LEGACY_PAYROLL_NOTE_PREFIX}{period}]")
        or (text.lower().startswith("nạp từ ") and ".xls" in text.lower())
    )


def apply_legacy_attendance_snapshot(conn, snapshot: dict, source_name: str, timestamp: str) -> dict:
    period = snapshot["period"]
    validate_legacy_attendance_period(period)
    marker = f"{LEGACY_ATTENDANCE_SOURCE_PREFIX}{period}"
    payroll_note = f"{LEGACY_PAYROLL_NOTE_PREFIX}{period}] {Path(source_name).name}"
    for record in snapshot["staff"]:
        conn.execute(
            """INSERT INTO staff(
                   employee_code,full_name,role_name,base_salary,standard_days,standard_hours,
                   created_at,updated_at
               ) VALUES(?,?,?,?,26,8,?,?)
               ON CONFLICT(employee_code) DO UPDATE SET
                   full_name=excluded.full_name,
                   role_name=CASE WHEN TRIM(COALESCE(staff.role_name,''))=''
                                  THEN excluded.role_name ELSE staff.role_name END,
                   base_salary=CASE WHEN staff.base_salary<=0 THEN excluded.base_salary ELSE staff.base_salary END,
                   updated_at=excluded.updated_at""",
            (record["employee_code"], record["name"], record["role"], record["base_salary"],
             timestamp, timestamp),
        )

    imported_attendance_before = conn.execute(
        """SELECT COUNT(*) qty FROM attendance_entries
           WHERE substr(work_date,1,7)=? AND (
               source=? OR (LOWER(COALESCE(source,''))!='manual' AND LOWER(COALESCE(source,'')) LIKE '%.xls%')
           )""",
        (period, marker),
    ).fetchone()["qty"]
    conn.execute(
        """DELETE FROM attendance_entries
           WHERE substr(work_date,1,7)=? AND (
               source=? OR (LOWER(COALESCE(source,''))!='manual' AND LOWER(COALESCE(source,'')) LIKE '%.xls%')
           )""",
        (period, marker),
    )
    attendance_inserted = attendance_manual_preserved = 0
    for item in snapshot["attendance"]:
        staff_id = conn.execute(
            "SELECT id FROM staff WHERE employee_code=?", (item["employee_code"],),
        ).fetchone()["id"]
        existing = conn.execute(
            "SELECT source FROM attendance_entries WHERE employee_id=? AND work_date=?",
            (staff_id, item["work_date"]),
        ).fetchone()
        if existing:
            attendance_manual_preserved += 1
            continue
        conn.execute(
            """INSERT INTO attendance_entries(
                   employee_id,work_date,normal_hours,overtime_hours,sunday_hours,night_hours,
                   holiday_hours,note,source,updated_at
               ) VALUES(?,?,?,?,?,?,?,?,?,?)""",
            (staff_id, item["work_date"], item["normal_hours"], item["overtime_hours"],
             item["sunday_hours"], item["night_hours"], item["holiday_hours"], item["note"],
             marker, timestamp),
        )
        attendance_inserted += 1

    imported_payroll_ids = [
        row["employee_id"] for row in conn.execute(
            "SELECT employee_id,note FROM payroll_adjustments WHERE month=?", (period,),
        ) if legacy_payroll_note_is_imported(row["note"], period)
    ]
    if imported_payroll_ids:
        conn.executemany(
            "DELETE FROM payroll_adjustments WHERE employee_id=? AND month=?",
            [(employee_id, period) for employee_id in imported_payroll_ids],
        )
    payroll_inserted = payroll_manual_preserved = 0
    for item in snapshot["payroll"]:
        staff_id = conn.execute(
            "SELECT id FROM staff WHERE employee_code=?", (item["employee_code"],),
        ).fetchone()["id"]
        existing = conn.execute(
            "SELECT 1 FROM payroll_adjustments WHERE employee_id=? AND month=?",
            (staff_id, period),
        ).fetchone()
        if existing:
            payroll_manual_preserved += 1
            continue
        conn.execute(
            """INSERT INTO payroll_adjustments(
                   employee_id,month,allowance,responsibility,advance,probation_deduction,
                   bhxh_employee_amount,bhxh_company_amount,gross_override,net_override,
                   use_override,note,updated_at
               ) VALUES(?,?,?,?,?,?,?,?,?,?,1,?,?)""",
            (staff_id, period, item["allowance"], item["responsibility"], item["advance"],
             item["probation"], item["bhxh_employee"], item["bhxh_company"], item["gross"],
             item["net"], payroll_note, timestamp),
        )
        payroll_inserted += 1

    imported_labor_before = conn.execute(
        """SELECT COUNT(*) qty FROM kitchen_labor_costs
           WHERE substr(work_date,1,7)=? AND (source=? OR LOWER(COALESCE(source,'')) LIKE '%.xls%')""",
        (period, marker),
    ).fetchone()["qty"]
    conn.execute(
        """DELETE FROM kitchen_labor_costs
           WHERE substr(work_date,1,7)=? AND (source=? OR LOWER(COALESCE(source,'')) LIKE '%.xls%')""",
        (period, marker),
    )
    for item in snapshot["labor"]:
        conn.execute(
            """INSERT INTO kitchen_labor_costs(work_date,kitchen,amount,source,updated_at)
               VALUES(?,?,?,?,?)""",
            (item["work_date"], item["kitchen"], item["amount"], marker, timestamp),
        )

    return {
        "month": period,
        "staff": len(snapshot["staff"]),
        "attendance_entries": attendance_inserted,
        "labor_cost_entries": len(snapshot["labor"]),
        "payroll_overrides": payroll_inserted,
        "manual_attendance_preserved": attendance_manual_preserved,
        "manual_payroll_preserved": payroll_manual_preserved,
        "replaced_attendance_entries": imported_attendance_before,
        "replaced_labor_cost_entries": imported_labor_before,
        "replaced_payroll_overrides": len(imported_payroll_ids),
        "source": Path(source_name).name,
    }


def import_legacy_attendance(conn, path: Path, source_name: str, now_iso, period: str):
    validate_legacy_attendance_period(period)
    values_book = formulas_book = None
    try:
        values_book = load_workbook(path, data_only=True, read_only=False, keep_links=False)
        formulas_book = load_workbook(path, data_only=False, read_only=False, keep_links=False)
        preview = parse_legacy_attendance_workbooks(
            conn, values_book, formulas_book, source_name, period,
        )
        result = apply_legacy_attendance_snapshot(
            conn, preview["snapshot"], source_name, now_iso,
        )
        result["warnings"] = preview["warnings"]
        return result
    finally:
        if values_book is not None:
            values_book.close()
        if formulas_book is not None:
            formulas_book.close()


def _legacy_invoice_payment_request_document(company: str, company_tax_code: str, company_address: str,
                                             recipient: str, recipient_tax_code: str, recipient_address: str,
                                             requester: str, bank_name: str, bank_account: str,
                                             period_from: str, period_to: str, details: list[dict]):
    """Retained only for old-file compatibility; current routes use the approved XLSX form."""
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)

    control_notice = document.add_paragraph()
    control_notice.alignment = WD_ALIGN_PARAGRAPH.CENTER
    control_run = control_notice.add_run("BIỂU MẪU DOCX CŨ – KHÔNG DÙNG CHO BÀN GIAO")
    control_run.bold = True
    control_run.font.size = Pt(11)

    header = document.add_table(rows=1, cols=2)
    header.autofit = False
    header.columns[0].width = Cm(8.3)
    header.columns[1].width = Cm(8.3)
    left = header.cell(0, 0).paragraphs[0]
    left.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = left.add_run(company.upper())
    run.bold = True
    if company_tax_code:
        left.add_run(f"\nMST: {company_tax_code}")
    right = header.cell(0, 1).paragraphs[0]
    right.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = right.add_run("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\nĐộc lập - Tự do - Hạnh phúc")
    run.bold = True

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(12)
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run("ĐỀ NGHỊ THANH TOÁN")
    run.bold = True
    run.font.size = Pt(16)
    subtitle = document.add_paragraph(f"Từ ngày {period_from} đến ngày {period_to}")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    labels = [
        ("Kính gửi: ", recipient.upper()),
        ("Mã số thuế bên mua: ", recipient_tax_code),
        ("Địa chỉ bên mua: ", recipient_address),
        ("Đơn vị đề nghị: ", company),
        ("Địa chỉ đơn vị đề nghị: ", company_address),
        ("Người đề nghị: ", requester),
        ("Nội dung: ", "Đề nghị thanh toán các hóa đơn hàng hóa/dịch vụ đã giao trong kỳ"),
    ]
    for label, value in labels:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(2)
        label_run = paragraph.add_run(label)
        label_run.bold = True
        paragraph.add_run(str(value or ""))

    table = document.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    table.autofit = False
    widths = (Cm(1), Cm(2.2), Cm(3.3), Cm(3.2), Cm(2.7), Cm(3.2))
    headers = ("STT", "Ngày HĐ", "Ký hiệu / Số HĐ", "Tiền trước thuế", "Tiền thuế", "Tổng thanh toán")
    for index, cell in enumerate(table.rows[0].cells):
        cell.width = widths[index]
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(headers[index])
        run.bold = True
    header_properties = table.rows[0]._tr.get_or_add_trPr()
    repeat_header = OxmlElement("w:tblHeader")
    repeat_header.set(qn("w:val"), "true")
    header_properties.append(repeat_header)

    total_subtotal = total_tax = total_amount = 0
    for index, item in enumerate(details, start=1):
        subtotal = vnd_round(item.get("subtotal"))
        tax_amount = vnd_round(item.get("tax_amount"))
        amount = vnd_round(item.get("total_amount"))
        total_subtotal += subtotal
        total_tax += tax_amount
        total_amount += amount
        invoice_ref = " / ".join(filter(None, [
            str(item.get("invoice_series") or "").strip(),
            str(item.get("invoice_number") or "").strip(),
        ]))
        invoice_date = as_date(item.get("invoice_date"))
        if invoice_date:
            invoice_date = datetime.strptime(invoice_date, "%Y-%m-%d").strftime("%d/%m/%Y")
        values = (
            index, invoice_date, invoice_ref,
            f"{subtotal:,.0f}".replace(",", "."),
            f"{tax_amount:,.0f}".replace(",", "."),
            f"{amount:,.0f}".replace(",", "."),
        )
        row = table.add_row().cells
        for column, value in enumerate(values):
            row[column].width = widths[column]
            paragraph = row[column].paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT if column >= 3 else WD_ALIGN_PARAGRAPH.CENTER
            paragraph.add_run(str(value))
    totals = table.add_row().cells
    totals[0].merge(totals[2])
    label = totals[0].paragraphs[0]
    label.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = label.add_run("TỔNG CỘNG")
    run.bold = True
    for column, value in enumerate((total_subtotal, total_tax, total_amount), start=3):
        paragraph = totals[column].paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = paragraph.add_run(f"{value:,.0f}".replace(",", "."))
        run.bold = True

    payment_lines = [
        ("Số tiền bằng chữ: ", number_to_vietnamese(total_amount) + "./."),
        ("Hình thức thanh toán: ", "Chuyển khoản"),
        ("Đơn vị thụ hưởng: ", company),
        ("Số tài khoản: ", bank_account),
        ("Tại ngân hàng: ", bank_name),
    ]
    for label, value in payment_lines:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(2)
        run = paragraph.add_run(label)
        run.bold = True
        paragraph.add_run(str(value or ""))
    signatures = document.add_table(rows=1, cols=2)
    for cell, text in zip(signatures.rows[0].cells, (
        "NGƯỜI LẬP\n(Ký, ghi rõ họ tên)",
        "ĐẠI DIỆN CÔNG TY\n(Ký, ghi rõ họ tên, đóng dấu)",
    )):
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(text)
        run.bold = True
    return document


def invoice_delivery_statement_workbook(
    lines: list[dict], drafts: list[dict], tax_factor=None, *, contractor="", scope_id="",
    period_from="", period_to="",
):
    """Backward-compatible entrypoint for the dedicated issued-invoice exporter."""
    return _invoice_delivery_statement_workbook(
        lines, drafts, tax_factor, contractor=contractor, scope_id=scope_id,
        period_from=period_from, period_to=period_to,
    )


def invoice_delivery_statement_scope_workbook(scope: dict):
    """Render a verified scope while retaining each invoice's provenance."""
    if scope.get('statement_kind') == 'invoices':
        try:
            from .invoice_payment_documents import synced_invoice_statement_workbook
        except ImportError:
            from invoice_payment_documents import synced_invoice_statement_workbook
        return synced_invoice_statement_workbook(scope)
    provenance = {int(item["draft_id"]): item for item in scope["invoices"]}
    drafts = []
    for source_draft in scope["drafts"]:
        draft = dict(source_draft)
        proof = provenance[int(draft["id"])]
        draft["verification_source"] = proof["verification_source"]
        draft["source_invoice_id"] = proof.get("source_invoice_id")
        drafts.append(draft)
    return invoice_delivery_statement_workbook(
        scope["lines"], drafts,
        contractor=scope["contractor"], scope_id=scope["scope_id"],
        period_from=scope["date_from"], period_to=scope["date_to"],
    )


def payment_request_document(company: str, recipient: str, requester: str, bank_name: str,
                             bank_account: str, amount: float, period_from: str, period_to: str,
                             details=None):
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(13)
    header = document.add_table(rows=1, cols=2)
    header.autofit = False
    header.columns[0].width = Cm(8)
    header.columns[1].width = Cm(9)
    left = header.cell(0, 0).paragraphs[0]
    left.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = left.add_run(company.upper())
    run.bold = True
    right = header.cell(0, 1).paragraphs[0]
    right.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = right.add_run("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\nĐộc lập - Tự do - Hạnh phúc")
    run.bold = True
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(18)
    title.paragraph_format.space_after = Pt(8)
    run = title.add_run("ĐỀ NGHỊ THANH TOÁN")
    run.bold = True
    run.font.size = Pt(16)
    today = date.today()
    date_line = document.add_paragraph(f"Ngày {today.day:02d} tháng {today.month:02d} năm {today.year}")
    date_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lines = [
        ("Kính gửi: ", recipient.upper()),
        ("Đơn vị đề nghị: ", company),
        ("Người đại diện: ", requester),
        ("Nội dung: ", f"Đề nghị thanh toán tiền hàng/dịch vụ từ {period_from} đến {period_to}"),
    ]
    for label, value in lines:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(4)
        label_run = paragraph.add_run(label)
        label_run.bold = True
        paragraph.add_run(value)

    detail_rows = details or [{"work_date": f"{period_from} – {period_to}", "amount": amount}]
    detail_table = document.add_table(rows=1, cols=4)
    detail_table.style = "Table Grid"
    detail_table.autofit = False
    widths = (Cm(1.4), Cm(3.2), Cm(8.4), Cm(3.8))
    headers = ("STT", "Ngày", "Nội dung", "Thành tiền (VNĐ)")
    for index, cell in enumerate(detail_table.rows[0].cells):
        cell.width = widths[index]
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(headers[index])
        run.bold = True
    # Word otherwise starts a continuation page with bare data rows.  Mark the
    # header row so monthly requests remain self-explanatory on every page.
    header_properties = detail_table.rows[0]._tr.get_or_add_trPr()
    repeat_header = OxmlElement("w:tblHeader")
    repeat_header.set(qn("w:val"), "true")
    header_properties.append(repeat_header)
    for index, item in enumerate(detail_rows, start=1):
        row_cells = detail_table.add_row().cells
        work_date = str(item.get("work_date") or "").strip()
        if re.fullmatch(r"\d{4}-\d{2}-\d{2}", work_date):
            work_date = datetime.strptime(work_date, "%Y-%m-%d").strftime("%d/%m/%Y")
        values = (
            str(index),
            work_date,
            "Tiền hàng/dịch vụ đã giao",
            f"{as_number(item.get('amount')):,.0f}".replace(",", "."),
        )
        for column, value in enumerate(values):
            row_cells[column].width = widths[column]
            paragraph = row_cells[column].paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT if column == 3 else WD_ALIGN_PARAGRAPH.CENTER
            paragraph.add_run(value)
    total_cells = detail_table.add_row().cells
    total_cells[0].merge(total_cells[2])
    total_label = total_cells[0].paragraphs[0]
    total_label.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = total_label.add_run("TỔNG CỘNG")
    run.bold = True
    total_value = total_cells[3].paragraphs[0]
    total_value.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = total_value.add_run(f"{amount:,.0f}".replace(",", "."))
    run.bold = True

    payment_lines = [
        ("Số tiền bằng chữ: ", number_to_vietnamese(amount) + "./."),
        ("Hình thức thanh toán: ", "Chuyển khoản"),
        ("Đơn vị thụ hưởng: ", company),
        ("Số tài khoản: ", bank_account),
        ("Tại ngân hàng: ", bank_name),
    ]
    for label, value in payment_lines:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(4)
        label_run = paragraph.add_run(label)
        label_run.bold = True
        paragraph.add_run(value)

    closing = document.add_paragraph("Rất mong Quý đơn vị xem xét và thanh toán.\nTrân trọng cảm ơn!")
    closing.paragraph_format.space_after = Pt(8)
    signatures = document.add_table(rows=1, cols=2)
    signatures.autofit = False
    signatures.columns[0].width = Cm(8)
    signatures.columns[1].width = Cm(8)
    for cell, text in zip(signatures.rows[0].cells, (
        "NGƯỜI LẬP\n(Ký, ghi rõ họ tên)",
        "ĐẠI DIỆN CÔNG TY\n(Ký, ghi rõ họ tên, đóng dấu)",
    )):
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(text)
        run.bold = True
    document.add_paragraph("\n\n\n")
    return document
