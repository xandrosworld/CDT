"""Supplementary paperwork for saved purchase remainders; no stock/debt writes."""
import io
import json
from datetime import date, timedelta
from decimal import Decimal
from html import escape

from flask import Response, jsonify, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

try:
    from . import purchase_document_selection as selection
except ImportError:
    import purchase_document_selection as selection


def queue(conn, start, end):
    try:
        a, b = date.fromisoformat(start), date.fromisoformat(end)
        if a > b or (b - a).days > 31:
            raise ValueError()
    except (TypeError, ValueError):
        raise selection.SelectionError('Chọn Từ ngày và Đến ngày hợp lệ, tối đa 32 ngày.') from None
    entries, cursor = [], (a - timedelta(days=1)).isoformat()
    while True:
        page = selection.pending_queue(conn, end, cursor)
        entries.extend(page['entries'])
        cursor = page['next_after']
        if not cursor:
            break
    for entry in entries:
        if entry['status'] == 'saved':
            plan = selection.saved_plan(conn, entry['date'])
            entry['source_hash'] = plan['source_hash']
    history = []
    for record in conn.execute('SELECT * FROM purchase_document_supplements ORDER BY id DESC'):
        dates = json.loads(record['days_json'])
        if not any(start <= day <= end for day in dates):
            continue
        rows = json.loads(record['rows_json'])
        history.append({k: record[k] for k in ('id', 'reference', 'actor', 'created_at', 'status', 'cancel_reason')})
        history[-1].update(dates=sorted(dates), amount=float(sum(Decimal(str(r['amount'])) for r in rows)))
    state_hash = selection.digest({'entries': entries, 'history': history})
    # Seller identity is needed internally for grouping, not for this public list.
    for entry in entries:
        for group in entry.get('groups', []):
            group.pop('identity', None)
            group['rows'] = [{k: r[k] for k in ('work_date', 'seller', 'product_name', 'unit', 'quantity', 'buy_price', 'amount', 'kitchen', 'source_ref')}
                             for r in group['rows']]
    return {'from': start, 'to': end, 'entries': entries, 'history': history, 'state_hash': state_hash}


def create(conn, body, timestamp, audit):
    actor, reference = str(body.get('actor') or '').strip(), str(body.get('reference') or '').strip()
    request_id = str(body.get('request_id') or '').strip()
    if not actor or len(actor) > 120:
        raise selection.SelectionError('Điền Người lập, tối đa 120 ký tự.')
    if not reference or len(reference) > 100:
        raise selection.SelectionError('Điền Số bảng kê bổ sung, tối đa 100 ký tự.')
    if not request_id or len(request_id) > 100 or body.get('confirmed') is not True:
        raise selection.SelectionError('Tích xác nhận đã kiểm tra phần chờ trước khi lập bảng kê bổ sung.')
    prior = conn.execute('SELECT * FROM purchase_document_supplements WHERE request_id=?', (request_id,)).fetchone()
    if prior:
        if prior['actor'] != actor or prior['reference'] != reference or sorted(json.loads(prior['days_json'])) != sorted(body.get('dates') or []):
            raise selection.SelectionError('Lần lưu này đã dùng cho bảng kê khác. Cập nhật danh sách trước khi lập tiếp.')
        return {'id': prior['id'], 'unchanged': True}
    if conn.execute("SELECT 1 FROM purchase_document_supplements WHERE reference=? AND status='active'", (reference,)).fetchone():
        raise selection.SelectionError('Số bảng kê bổ sung đã dùng. Mở Đã lập để tải lại hoặc nhập số khác.')
    current = queue(conn, body.get('from'), body.get('to'))
    if current['state_hash'] != body.get('state_hash'):
        raise selection.SelectionError('Phần chờ vừa thay đổi. Bấm Xem phần chờ rồi kiểm tra lại; Người lập và Số bảng kê bổ sung vẫn được giữ.')
    dates = body.get('dates')
    if not isinstance(dates, list) or not dates or any(not isinstance(d, str) for d in dates) or len(set(dates)) != len(dates):
        raise selection.SelectionError('Chọn ít nhất một ngày còn chờ để lập bảng kê bổ sung.')
    available = {e['date']: e for e in current['entries'] if e['status'] == 'saved'}
    if set(dates) - set(available):
        raise selection.SelectionError('Có ngày chưa lưu lựa chọn hoặc cần đối chiếu. Bấm Mở lựa chọn của ngày đó để xử lý trước.')
    rows, days = [], {}
    for day in sorted(dates):
        source, source_hash = selection.day_source(conn, day)
        plan = selection.saved_plan(conn, day)
        _, pending = selection.split_rows(source, json.loads(plan['quantities_json']))
        pending = selection.remaining_supplement_rows(conn, day, source_hash, pending)
        if not pending:
            raise selection.SelectionError('Ngày ' + day + ' không còn phần chờ. Bấm Xem phần chờ để cập nhật.')
        rows.extend(pending)
        days[day] = {'source_hash': source_hash, 'revision': plan['revision']}
    row_id = conn.execute('''INSERT INTO purchase_document_supplements
        (request_id,reference,actor,created_at,days_json,rows_json) VALUES(?,?,?,?,?,?)''',
        (request_id, reference, actor, timestamp, json.dumps(days, sort_keys=True), json.dumps(rows, ensure_ascii=False))).lastrowid
    audit(conn, 'purchase_document.supplement', entity_type='purchase_document_supplement', entity_id=row_id,
          metadata={'actor': actor, 'reference': reference, 'dates': sorted(days), 'row_count': len(rows)})
    return {'id': row_id, 'unchanged': False}


def cancel(conn, record_id, body, timestamp, audit):
    actor, reason = str(body.get('actor') or '').strip(), str(body.get('reason') or '').strip()
    if not actor or len(actor) > 120 or not reason or len(reason) > 500 or body.get('confirmed') is not True:
        raise selection.SelectionError('Điền Người lập và Lý do hủy, rồi tích xác nhận hủy bảng kê bổ sung.')
    record = get_record(conn, record_id)
    if record['status'] == 'cancelled':
        return
    conn.execute("UPDATE purchase_document_supplements SET status='cancelled',cancelled_by=?,cancel_reason=?,cancelled_at=? WHERE id=?", (actor, reason, timestamp, record_id))
    audit(conn, 'purchase_document.supplement_cancel', entity_type='purchase_document_supplement', entity_id=record_id,
          metadata={'actor': actor, 'reason': reason})


def get_record(conn, record_id):
    record = conn.execute('SELECT * FROM purchase_document_supplements WHERE id=?', (record_id,)).fetchone()
    if not record:
        raise selection.SelectionError('Không tìm thấy bảng kê bổ sung. Bấm Xem phần chờ để cập nhật.')
    return dict(record)


HEADERS = ['Ngày mua', 'Người bán', 'Tên hàng', 'ĐVT', 'Số lượng bổ sung', 'Giá mua BK', 'Thành tiền', 'Bếp', 'Dòng nguồn']


def document_rows(record):
    rows = sorted(json.loads(record['rows_json']), key=lambda r: (r['work_date'], r['seller'], r['source_ref']))
    return [[r[k] for k in ('work_date', 'seller', 'product_name', 'unit', 'quantity', 'buy_price', 'amount', 'kitchen', 'source_ref')] for r in rows]


def workbook(record):
    book = Workbook(); sheet = book.active; sheet.title = 'Bảng kê bổ sung'
    sheet.append(['BẢNG KÊ BỔ SUNG — PHẦN CHỜ ĐÃ CHỌN'])
    sheet.append(['Số bảng kê: ' + record['reference'] + ' · Người lập: ' + record['actor']])
    sheet.append(['Trạng thái: ' + ('Đã hủy' if record['status'] == 'cancelled' else 'Đã lập') + ' · Giữ nguyên ngày mua; không ghi thêm kho hoặc công nợ.'])
    sheet.append(HEADERS)
    for row in document_rows(record):
        sheet.append(row)
    sheet.append(['TỔNG', '', '', '', '', '', sum(Decimal(str(r[6])) for r in document_rows(record))])
    for row in sheet:
        for cell in row:
            if isinstance(cell.value, str): cell.data_type = 's'
            cell.font = Font(name='Times New Roman', size=11, bold=cell.row <= 4 or cell.row == sheet.max_row)
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if cell.column in (5, 6, 7) and cell.row > 4: cell.number_format = '#,##0.######'
    for number in (1, 2, 3): sheet.merge_cells(start_row=number, start_column=1, end_row=number, end_column=9)
    for cell in sheet[4]: cell.fill = PatternFill('solid', fgColor='E9F2F1')
    for col, width in zip('ABCDEFGHI', [15, 25, 32, 10, 16, 16, 20, 18, 20]): sheet.column_dimensions[col].width = width
    sheet.freeze_panes = 'C5'; sheet.auto_filter.ref = f'A4:I{sheet.max_row - 1}'
    sheet.print_title_rows = '1:4'; sheet.print_area = f'A1:I{sheet.max_row}'
    sheet.page_setup.orientation = 'landscape'; sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1; sheet.page_setup.fitToHeight = 0; sheet.sheet_properties.pageSetUpPr.fitToPage = True
    return book


def register_routes(app, ctx):
    @app.route('/api/purchase-document-supplements', methods=['GET', 'POST'])
    def supplements():
        try:
            with ctx['db']() as conn:
                conn.execute('BEGIN IMMEDIATE' if request.method == 'POST' else 'BEGIN')
                if request.method == 'GET':
                    return jsonify(ok=True, **queue(conn, request.args.get('from'), request.args.get('to')))
                body = request.get_json(silent=True)
                if not isinstance(body, dict): raise selection.SelectionError('Nội dung lập bảng kê bổ sung không hợp lệ.')
                return jsonify(ok=True, **create(conn, body, ctx['now_iso'](), ctx['audit_event']))
        except (ValueError, TypeError) as exc:
            return jsonify(ok=False, error=str(exc)), 409

    @app.post('/api/purchase-document-supplements/<int:record_id>/cancel')
    def cancel_supplement(record_id):
        try:
            with ctx['db']() as conn:
                conn.execute('BEGIN IMMEDIATE')
                cancel(conn, record_id, request.get_json(silent=True) or {}, ctx['now_iso'](), ctx['audit_event'])
                return jsonify(ok=True)
        except (ValueError, TypeError) as exc:
            return jsonify(ok=False, error=str(exc)), 409

    @app.get('/api/purchase-document-supplements/<int:record_id>/<kind>')
    def supplement_document(record_id, kind):
        try:
            with ctx['db']() as conn: record = get_record(conn, record_id)
            if kind == 'excel':
                book = workbook(record); output = io.BytesIO(); book.save(output); book.close(); output.seek(0)
                return send_file(output, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', as_attachment=True, download_name=f'Bang_ke_bo_sung_{record_id}.xlsx')
            if kind != 'print': return jsonify(ok=False, error='Không tìm thấy bản in.'), 404
            rows = document_rows(record)
            table = '<table><thead><tr>' + ''.join('<th>'+escape(h)+'</th>' for h in HEADERS) + '</tr></thead><tbody>'
            for row in rows:
                table += '<tr>' + ''.join('<td>'+escape(str(value))+'</td>' for value in row) + '</tr>'
            table += '<tr><th colspan="6">TỔNG</th><th>'+escape(str(sum(Decimal(str(r[6])) for r in rows)))+'</th><td colspan="2"></td></tr></tbody></table>'
            return Response('<!doctype html><html lang="vi"><meta charset="utf-8"><title>Bảng kê bổ sung</title><style>body{font-family:serif}table{border-collapse:collapse;width:100%}td,th{border:1px solid #777;padding:6px}thead{display:table-header-group}tr{break-inside:avoid}@page{size:A4 landscape;margin:12mm}@media print{button{display:none}}</style><button onclick="window.print()">In / Lưu PDF</button><h1>Bảng kê bổ sung'+(' — ĐÃ HỦY' if record['status']=='cancelled' else '')+'</h1><p>Số bảng kê: '+escape(record['reference'])+' · Người lập: '+escape(record['actor'])+'</p><p>Ngày mua giữ nguyên theo nguồn. Bổ sung chứng từ cho phần chờ; không ghi thêm kho hoặc công nợ.</p>'+table+'</html>', mimetype='text/html')
        except (ValueError, TypeError) as exc:
            return jsonify(ok=False, error=str(exc)), 409
