from __future__ import annotations

import hashlib
import io
import tempfile
import unittest
import zipfile
from pathlib import Path

from openpyxl import load_workbook

try:
    from . import server
    from .receivable_ledger import sync_receivable_ledger
except ImportError:  # pragma: no cover - direct invocation
    import server
    from receivable_ledger import sync_receivable_ledger


class ReceivableExportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        cls.original_db_path = server.DB_PATH
        cls.original_data_dir = server.DATA_DIR
        cls.original_master_source = server.MASTER_SOURCE
        cls.original_testing = server.app.config.get("TESTING", False)
        server.DATA_DIR = Path(cls.temp.name) / "data"
        server.DATA_DIR.mkdir(parents=True, exist_ok=True)
        server.DB_PATH = server.DATA_DIR / "receivable-export.sqlite3"
        server.MASTER_SOURCE = Path(cls.temp.name) / "master-disabled.xlsx"
        server.app.config["TESTING"] = True
        server.init_database()
        cls.client = server.app.test_client()
        with server.db() as conn:
            conn.execute(
                "INSERT INTO contractors(code,name,price_group,pricing_mode) "
                "VALUES('C1','Nhà thầu Một','C1','group')"
            )
            conn.execute(
                "INSERT INTO contractors(code,name,price_group,pricing_mode) "
                "VALUES('C2','Nhà thầu Hai','C2','group')"
            )
            conn.execute(
                "INSERT INTO suppliers(code,name) VALUES('S1','Nhà cung cấp Một')"
            )

    @classmethod
    def tearDownClass(cls):
        server.DB_PATH = cls.original_db_path
        server.DATA_DIR = cls.original_data_dir
        server.MASTER_SOURCE = cls.original_master_source
        server.app.config["TESTING"] = cls.original_testing
        cls.temp.cleanup()

    def setUp(self):
        with server.db() as conn:
            conn.execute("DELETE FROM receivable_ledger_revisions")
            conn.execute("DELETE FROM receivable_ledger_lines")
            conn.execute("DELETE FROM outgoing_invoice_lines")
            conn.execute("DELETE FROM outgoing_invoice_drafts")
            conn.execute("DELETE FROM order_import_receipts")
            conn.execute("DELETE FROM orders")
            conn.execute("DELETE FROM batches")
            conn.execute("DELETE FROM debt_adjustments")
            conn.execute("DELETE FROM payments")
            conn.execute("DELETE FROM balances")
            conn.execute("DELETE FROM kitchens")
            conn.execute("DELETE FROM audit_log")

    @staticmethod
    def _batch(conn, work_date: str) -> int:
        timestamp = server.now_iso()
        return int(conn.execute(
            """INSERT INTO batches(work_date,source_name,status,created_at,approved_at)
               VALUES(?,?,'approved',?,?)""",
            (work_date, "receivable-export-fixture", timestamp, timestamp),
        ).lastrowid)

    @staticmethod
    def _order(
        conn, batch_id: int, work_date: str, *, contractor: str, kitchen: str,
        product_code: str, product_name: str, ordered: float, delivered: float,
        returned: float, sell_price: float, tax: str,
    ) -> int:
        return int(conn.execute(
            """INSERT INTO orders(
                   batch_id,work_date,contractor,kitchen,product_code,product_name,qty,
                   actual_received,actual_delivered,unit,supplier,buy_price,sell_price,tax,
                   purchase_list,source_sheet,source_row,errors,warnings,updated_at,
                   customer_return_qty
               ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (
                batch_id, work_date, contractor, kitchen, product_code, product_name,
                ordered, ordered, delivered, "kg", "S1", 5_000, sell_price, tax, 0,
                "đơn hàng", 3, "[]", "[]", server.now_iso(), returned,
            ),
        ).lastrowid)

    def _fixture(self) -> tuple[str, str]:
        # Both labels sanitize to the same 31-character prefix.  The exporter
        # must create legal, distinct worksheet names deterministically.
        kitchen_one = "BEP/" + "A" * 36 + "1"
        kitchen_two = "BEP:" + "A" * 36 + "2"
        with server.db() as conn:
            conn.execute(
                "INSERT INTO kitchens(code,contractor,name) VALUES(?,?,?)",
                (kitchen_one, "C1", "Bếp tên rất dài / ca sáng [A]"),
            )
            conn.execute(
                "INSERT INTO kitchens(code,contractor,name) VALUES(?,?,?)",
                (kitchen_two, "C1", "Bếp tên rất dài : ca chiều *A*"),
            )
            conn.execute(
                "INSERT INTO kitchens(code,contractor,name) VALUES('K3','C2','Bếp Ba')"
            )
            first = self._batch(conn, "2026-09-01")
            second = self._batch(conn, "2026-09-15")
            third = self._batch(conn, "2026-09-20")
            self._order(
                conn, first, "2026-09-01", contractor="C1", kitchen=kitchen_one,
                product_code="P1", product_name="Cà rốt", ordered=10, delivered=8,
                returned=2, sell_price=10_000, tax="10%",
            )
            self._order(
                conn, second, "2026-09-15", contractor="C1", kitchen=kitchen_two,
                product_code="P2", product_name="Bí ngòi", ordered=5, delivered=5,
                returned=0, sell_price=10_000, tax="0%",
            )
            self._order(
                conn, third, "2026-09-20", contractor="C2", kitchen="K3",
                product_code="P3", product_name="Khoai tây", ordered=2, delivered=2,
                returned=0, sell_price=20_000, tax="0%",
            )
            sync_receivable_ledger(conn, timestamp=server.now_iso())
            conn.execute(
                """INSERT INTO balances(party_type,party_code,opening,as_of_date)
                   VALUES('contractor','C1',100000,'2026-08-31')""",
            )
            conn.execute(
                """INSERT INTO payments(
                       payment_date,kind,party_type,party_code,amount,note,created_at,updated_at
                   ) VALUES('2026-09-10','receipt','contractor','C1',30000,'Khách trả',?,?)""",
                (server.now_iso(), server.now_iso()),
            )
            conn.execute(
                """INSERT INTO payments(
                       payment_date,kind,party_type,party_code,amount,note,created_at,updated_at
                   ) VALUES('2026-09-11','payment','contractor','C1',999000,
                            'Sai loại, không được tính',?,?)""",
                (server.now_iso(), server.now_iso()),
            )
            conn.execute(
                """INSERT INTO debt_adjustments(
                       adjustment_date,party_type,party_code,amount,note,created_at
                   ) VALUES('2026-09-18','contractor','C1',-5000,'Giảm trừ',?)""",
                (server.now_iso(),),
            )
        return kitchen_one, kitchen_two

    def _export(self, **query):
        return self.client.get(
            "/api/debts/receivables/export",
            query_string={"from": "2026-09-01", "to": "2026-09-30", **query},
        )

    def test_single_contractor_workbook_reconciles_account_summary_and_kitchen_sheets(self):
        self._fixture()
        response = self._export(contractor="Nhà thầu Một")
        self.assertEqual(response.status_code, 200, response.get_json(silent=True))
        self.assertEqual(
            response.mimetype,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn(
            "Cong_no_phai_thu_C1_2026-09-01_2026-09-30.xlsx",
            response.headers["Content-Disposition"],
        )
        workbook = load_workbook(io.BytesIO(response.data), data_only=False, keep_links=False)
        self.assertNotIn("Tổng nhà thầu", workbook.sheetnames)
        self.assertEqual(len(workbook.sheetnames), 2)
        self.assertEqual(len({name.casefold() for name in workbook.sheetnames}), 2)
        for name in workbook.sheetnames:
            self.assertLessEqual(len(name), 31)
            self.assertFalse(any(character in name for character in "\\/*?:[]"))

        child_totals = []
        for sheet in workbook:
            self.assertEqual([cell.value for cell in sheet[3]], [
                "Ngày", "Tên bếp", "Tên hàng", "SL đặt", "ĐVT", "Giá bán giao dịch",
                "Thuế suất (%)", "Tiền trước thuế", "Tiền thuế", "Phát sinh phải thu",
            ])
            child_totals.append(sheet.cell(sheet.max_row, 10).value)
            self.assertEqual(sheet.freeze_panes, "A4")
            self.assertEqual(sheet.page_setup.orientation, "landscape")
            self.assertEqual(sheet.page_setup.paperSize, 9)
            self.assertEqual(sheet.page_setup.fitToWidth, 1)
            self.assertEqual(sheet.print_title_rows, "$3:$3")
            self.assertTrue(str(sheet.print_area).endswith(f"$J${sheet.max_row}"))
            dates = [sheet.cell(r, 1).value for r in range(4, sheet.max_row)]
            self.assertEqual(dates, sorted(dates))
        # 8 delivered - 2 returned = 6, at 10,000 + 10% VAT: hiding
        # quantity columns must not change net-delivery accounting.
        self.assertEqual(sorted(child_totals), [50_000, 66_000])
        self.assertEqual(sum(child_totals), 116_000)

        formulas = [
            cell.coordinate
            for sheet in workbook.worksheets
            for row in sheet.iter_rows()
            for cell in row
            if isinstance(cell.value, str) and cell.value.startswith("=")
        ]
        self.assertEqual(formulas, [])
        self.assertEqual(workbook._external_links, [])
        workbook.close()

    def test_all_contractors_is_zip_with_one_independent_workbook_per_contractor(self):
        self._fixture()
        before = hashlib.sha256(server.DB_PATH.read_bytes()).hexdigest()
        response = self._export()
        after = hashlib.sha256(server.DB_PATH.read_bytes()).hexdigest()
        self.assertEqual(response.status_code, 200, response.get_json(silent=True))
        self.assertEqual(response.mimetype, "application/zip")
        self.assertEqual(before, after, "Kết xuất chỉ đọc không được sửa database")
        with zipfile.ZipFile(io.BytesIO(response.data)) as archive:
            names = archive.namelist()
            self.assertEqual(len(names), 2)
            self.assertTrue(all(name.endswith(".xlsx") for name in names))
            self.assertEqual(len({name.casefold() for name in names}), 2)
            seen = {}
            for name in names:
                workbook = load_workbook(io.BytesIO(archive.read(name)), data_only=False)
                self.assertNotIn("Tổng nhà thầu", workbook.sheetnames)
                contractor_code = workbook.active["A2"].value.split(" · ")[0]
                child_total = sum(
                    sheet.cell(sheet.max_row, 10).value
                    for sheet in workbook.worksheets
                )
                seen[contractor_code] = child_total
                workbook.close()
        self.assertEqual(seen, {"C1": 116_000, "C2": 40_000})

    def test_custom_period_empty_data_and_invalid_filters_are_stable(self):
        self._fixture()
        empty = self.client.get(
            "/api/debts/receivables/export",
            query_string={
                "from": "2025-01-01", "to": "2025-01-31", "contractor": "c1",
            },
        )
        self.assertEqual(empty.status_code, 200, empty.get_json(silent=True))
        workbook = load_workbook(io.BytesIO(empty.data), data_only=False)
        self.assertEqual(workbook.sheetnames, ["Không phát sinh"])
        self.assertIn("Không có dòng phải thu", workbook.active["A2"].value)
        workbook.close()

        empty_bundle = self.client.get(
            "/api/debts/receivables/export",
            query_string={"from": "2025-01-01", "to": "2025-01-31"},
        )
        self.assertEqual(empty_bundle.status_code, 200)
        with zipfile.ZipFile(io.BytesIO(empty_bundle.data)) as archive:
            self.assertEqual(archive.namelist(), ["KHONG_PHAT_SINH.txt"])
            self.assertIn("không dùng hóa đơn đỏ", archive.read(archive.namelist()[0]).decode("utf-8"))

        invalid_cases = [
            ({"from": "2026-09-31", "to": "2026-10-01", "contractor": "C1"}, "invalid_period"),
            ({"from": "2026-10-01", "to": "2026-09-01", "contractor": "C1"}, "invalid_period"),
            ({"from": "2026-09-01", "to": "2026-09-30", "contractor": "KHONG-CO"}, "invalid_contractor"),
        ]
        for query, code in invalid_cases:
            with self.subTest(query=query):
                response = self.client.get("/api/debts/receivables/export", query_string=query)
                self.assertEqual(response.status_code, 400, response.get_data(as_text=True))
                self.assertEqual(response.get_json()["code"], code)


if __name__ == "__main__":
    unittest.main()
