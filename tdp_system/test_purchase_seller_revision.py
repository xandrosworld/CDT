import io
import json
import sqlite3
import unittest

from openpyxl import Workbook
from flask import Flask

from . import purchase_seller_revision as revision
from . import purchase_document_selection as selection
from .purchase_summary_export import collect_purchase_summary_rows


class SellerRevisionTests(unittest.TestCase):
    def setUp(self):
        self.c = sqlite3.connect(':memory:')
        self.c.row_factory = sqlite3.Row
        self.addCleanup(self.c.close)
        self.c.executescript('''
            CREATE TABLE batches(id INTEGER, work_date TEXT, status TEXT);
            INSERT INTO batches VALUES(1,'2026-09-15','approved');
            CREATE TABLE people(name TEXT,cccd TEXT,address TEXT,issue_date TEXT,issue_place TEXT);
            CREATE TABLE orders(id INTEGER,batch_id INTEGER,contractor TEXT,product_code TEXT,
                kitchen TEXT,unit TEXT,sell_price REAL,tax REAL,qty REAL,purchase_list INTEGER,
                seller TEXT,cccd TEXT);
            CREATE TABLE purchase_workbook_lines(id INTEGER,batch_id INTEGER,order_id INTEGER,
                source_row INTEGER,work_date TEXT,product_code TEXT,product_name TEXT,unit TEXT,
                supplier TEXT,buy_price REAL,actual_qty REAL,status TEXT);
            CREATE TABLE batch_bk_approvals(batch_id INTEGER,document_id INTEGER);
            CREATE TABLE bk_import_documents(id INTEGER,status TEXT);
            CREATE TABLE bk_import_lines(document_id INTEGER,source_line INTEGER,qty REAL,unit_cost REAL,amount REAL);
            CREATE TABLE receivable_ledger_lines(amount REAL);
            CREATE TABLE payable_ledger_lines(amount REAL);
            CREATE TABLE invoice_inventory_ledger(qty REAL);
            INSERT INTO receivable_ledger_lines VALUES(9500000);
            INSERT INTO payable_ledger_lines VALUES(7200000);
            INSERT INTO invoice_inventory_ledger VALUES(36);
            INSERT INTO batch_bk_approvals VALUES(1,10);
            INSERT INTO bk_import_documents VALUES(10,'posted');
            INSERT INTO bk_import_lines VALUES(10,10,36,247000,8892000);
            INSERT INTO bk_import_lines VALUES(10,11,2,100000,200000);
        ''')
        for i, name in enumerate(('Seller A', 'Seller B', 'Seller C'), 1):
            self.c.execute('INSERT INTO people VALUES(?,?,?,?,?)', (name, str(i)*12, 'Test address', '01/01/2020', 'Test place'))
        self.c.executemany('INSERT INTO orders VALUES(?,?,?,?,?,?,?,?,?,?,?,?)', [
            (1,1,'ATV','P1','K1','Kg',260000,-2,36,1,'Seller A','1'*12),
            (2,1,'ATV','P2','K2','Kg',120000,-2,2,1,'Seller C','3'*12)])
        self.c.executemany('INSERT INTO purchase_workbook_lines VALUES(?,?,?,?,?,?,?,?,?,?,?,?)', [
            (10,1,1,39,'2026-09-15','P1','Test shrimp','Kg','NCC',230000,36,'confirmed'),
            (11,1,2,40,'2026-09-15','P2','Other goods','Kg','NCC',90000,2,'confirmed')])
        selection.init_schema(self.c)
        self.blob = self.workbook()

    def workbook(self, parts=None, change=None):
        parts = parts if parts is not None else [(18, 'Seller A', '1'*12), (18, 'Seller B', '2'*12)]
        book = Workbook(); ws = book.active
        ws.cell(2,5,'Mã hàng'); ws.cell(2,20,'CCCD')
        rows = [('P1','K1',260000,q,n,c) for q,n,c in parts] + [('P2','K2',120000,2,'Seller C','3'*12)]
        for index,(code,kitchen,price,qty,name,identity) in enumerate(rows,3):
            row = ['ATV',None,None,None,code,kitchen,'2026-09-15',None,qty,'Kg',None,price,None,-2,None,None,None,'bk',name,identity]
            for col,value in enumerate(row,1): ws.cell(index,col,value)
        if change: change(ws)
        output=io.BytesIO();book.save(output);book.close();return output.getvalue()

    def preview(self, blob=None):
        return revision.build_preview(self.c,blob or self.blob,'2026-09-15','corrected.xlsx')

    def save(self, preview=None, blob=None):
        return revision.save_preview(self.c,preview or self.preview(blob),blob or self.blob,'Reviewer','Actual seller corrected','2026-09-17 18:00:00')

    def business_snapshot(self):
        tables=('orders','purchase_workbook_lines','batch_bk_approvals','bk_import_documents','bk_import_lines',
                'receivable_ledger_lines','payable_ledger_lines','invoice_inventory_ledger','batches')
        return {t:[tuple(r) for r in self.c.execute('SELECT * FROM '+t)] for t in tables}

    def test_posted_order_correction_preserves_all_financial_sources_and_exact_amount(self):
        before=self.business_snapshot();p=self.preview()
        self.assertEqual(1,len(p['changes']))
        self.assertEqual([4446000,4446000],[r['amount'] for r in p['changes'][0]['after']])
        self.save(p)
        self.assertEqual(before,self.business_snapshot())
        rows,_=selection.day_source(self.c,'2026-09-15')
        self.assertEqual(9092000,sum(r['amount'] for r in rows))
        self.assertEqual(36,sum(r['quantity'] for r in rows if r['product_name']=='Test shrimp'))
        self.assertFalse(selection.scope(self.c,'2026-09-15','2026-09-15')['days'][0]['stale'])
        self.assertTrue(self.save(p)['unchanged'])
        self.assertEqual(1,self.c.execute('SELECT COUNT(*) FROM purchase_seller_revisions').fetchone()[0])

    def test_business_changes_and_invalid_identity_rejected(self):
        for column,value in ((9,19),(12,270000),(5,'NEW'),(6,'K9'),(14,8),(20,'9'*12),(19,'Unknown')):
            with self.subTest(column=column),self.assertRaises(ValueError):
                self.preview(self.workbook(change=lambda ws:ws.cell(3,column,value)))
        # The file must not silently change even an otherwise unchanged identity.
        with self.assertRaisesRegex(ValueError,'CCCD'):
            self.preview(self.workbook(change=lambda ws:ws.cell(5,20,'9'*12)))
        with self.assertRaisesRegex(ValueError,'Excel'):
            self.preview(b'not an xlsx')

    def test_concurrent_selection_or_catalogue_change_requires_new_preview(self):
        p=self.preview()
        state=selection.scope(self.c,'2026-09-15','2026-09-15')
        selection.save(self.c,{**state,'actor':'Other reviewer','days':[{'date':'2026-09-15','quantities':{}}]},'now',lambda *a,**k:None)
        with self.assertRaisesRegex(ValueError,'vừa thay đổi'):self.save(p)
        p=self.preview();self.c.execute("UPDATE people SET address='New address' WHERE name='Seller B'")
        with self.assertRaisesRegex(ValueError,'vừa thay đổi'):self.save(p)

    def test_existing_selection_retained_only_for_unchanged_rows_and_history_preserved(self):
        state=selection.scope(self.c,'2026-09-15','2026-09-15')
        selection.save(self.c,{**state,'actor':'First reviewer','days':[{'date':'2026-09-15','quantities':{'1:purchase_workbook_lines:10':'10','1:purchase_workbook_lines:11':'2'}}]},'now',lambda *a,**k:None)
        old_history=tuple(self.c.execute('SELECT * FROM purchase_document_selection_history').fetchone())
        self.save()
        plan=selection.saved_plan(self.c,'2026-09-15')
        self.assertEqual({'1:purchase_workbook_lines:11':'2'},json.loads(plan['quantities_json']))
        self.assertEqual(old_history,tuple(self.c.execute('SELECT * FROM purchase_document_selection_history WHERE revision=1').fetchone()))
        self.assertEqual(2,plan['revision'])

    def test_reverting_to_original_seller_keeps_daily_limit_and_audit_history(self):
        self.save()
        original=self.workbook([(36,'Seller A','1'*12)])
        p=self.preview(original);self.assertEqual(1,len(p['changes']))
        self.assertTrue(next(g for g in p['groups'] if g['seller']=='Seller A')['over_limit'])
        self.save(p,original)
        state=selection.scope(self.c,'2026-09-15','2026-09-15')
        group=next(g for g in state['days'][0]['groups'] if g['seller']=='Seller A')
        self.assertEqual(0,group['selected'])
        self.assertEqual(8892000,group['pending'])
        self.assertEqual(2,self.c.execute('SELECT COUNT(*) FROM purchase_seller_revisions').fetchone()[0])

    def test_changed_source_cannot_export_saved_identity_allocation(self):
        self.save()
        self.c.execute('UPDATE bk_import_lines SET amount=amount+1,unit_cost=unit_cost+0.0277777778 WHERE source_line=10')
        with self.assertRaisesRegex(ValueError,'Nguồn mua đã đổi'):
            selection.day_source(self.c,'2026-09-15')

    def test_second_revision_keeps_keys_for_unchanged_previous_split(self):
        self.save(); first,_=selection.day_source(self.c,'2026-09-15')
        blob=self.workbook(change=lambda ws:(ws.cell(5,19,'Seller B'),ws.cell(5,20,'2'*12)))
        self.save(blob=blob);second,_=selection.day_source(self.c,'2026-09-15')
        self.assertEqual([r for r in first if r['product_name']=='Test shrimp'],[r for r in second if r['product_name']=='Test shrimp'])

    def test_route_requires_confirmation_and_records_audit_only_after_save(self):
        self.c.commit()
        app=Flask(__name__);events=[]
        revision.register_routes(app,{'db':lambda:self.c,'now_iso':lambda:'now',
            'audit_event':lambda *a,**kw:events.append(kw)})
        client=app.test_client()
        response=client.post('/api/purchase-sellers/preview',data={
            'date':'2026-09-15','file':(io.BytesIO(self.blob),'corrected.xlsx')})
        self.assertEqual(200,response.status_code,response.get_json())
        body={'token':response.get_json()['token'],'actor':'Reviewer','reason':'Verified actual seller'}
        before=self.business_snapshot()
        self.assertEqual(409,client.post('/api/purchase-sellers/confirm',json=body).status_code)
        self.assertFalse(events)
        body['confirmed']=True
        result=client.post('/api/purchase-sellers/confirm',json=body)
        self.assertEqual(200,result.status_code,result.get_json())
        self.assertEqual(before,self.business_snapshot())
        self.assertEqual(1,len(events))
        self.assertTrue(client.post('/api/purchase-sellers/confirm',json=body).get_json()['unchanged'])
        self.assertEqual(409,client.post('/api/purchase-sellers/confirm',json=['invalid']).status_code)

    def test_stale_saved_selection_cannot_be_silently_rebased(self):
        state=selection.scope(self.c,'2026-09-15','2026-09-15')
        selection.save(self.c,{**state,'actor':'Reviewer','days':[{'date':'2026-09-15','quantities':{}}]},'now',lambda *a,**k:None)
        self.c.execute("UPDATE people SET issue_place='Changed' WHERE name='Seller C'")
        with self.assertRaisesRegex(ValueError,'đã cũ'):self.preview()


if __name__ == '__main__':
    unittest.main()
