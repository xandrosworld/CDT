import io
import unittest
from openpyxl import load_workbook
from . import server
from . import test_bk_import as fixture


class BkDraftTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls): fixture.BKImportTests.setUpClass.__func__(cls)
    @classmethod
    def tearDownClass(cls): fixture.BKImportTests.tearDownClass.__func__(cls)
    def setUp(self): fixture.BKImportTests.setUp(self)
    preview = fixture.BKImportTests.preview
    confirm = fixture.BKImportTests.confirm
    counts = staticmethod(fixture.BKImportTests.counts)

    def body(self, complete=False):
        return {'from':'2026-08-01','to':'2026-08-31',
                'document_date':'2026-08-31' if complete else '',
                'reference':'BK-AUG-TEST' if complete else '',
                'rows':[{'product_code':'BK-P1','qty':2,'unit_cost':100 if complete else '',
                         'source_party':'BK-S1' if complete else '', 'note':'Hàng mua bổ sung'}]}

    def test_incomplete_draft_downloads_but_cannot_print_final_receipts_or_post(self):
        before = self.counts()
        response = self.client.post('/api/bk-import/draft/excel',json=self.body())
        self.assertEqual(response.status_code,200,response.get_json())
        wb=load_workbook(io.BytesIO(response.data))
        self.assertEqual(wb['BK_IMPORT']['E4'].value,'BK-P1')
        self.assertIsNone(wb['BK_IMPORT']['A4'].value)
        self.assertIsNone(wb['BK_IMPORT']['I4'].value)
        self.assertEqual(wb['BK_IMPORT']['H4'].value,2)
        wb.close()
        snapshot=self.client.post('/api/bk-import/draft/preview',json=self.body())
        self.assertEqual(snapshot.status_code,400,snapshot.get_json())
        self.assertEqual(self.counts(),before)
        preview=self.preview(response.data).get_json()
        self.assertFalse(preview['canConfirm'])
        self.assertGreater(preview['counts']['errorRows'],0)
        self.assertEqual(self.confirm(preview).status_code,400)
        self.assertEqual(self.counts(),before)

    def test_complete_draft_roundtrip_requires_confirmation_and_is_idempotent(self):
        before=self.counts()
        response=self.client.post('/api/bk-import/draft/excel',json=self.body(True))
        self.assertEqual(response.status_code,200,response.get_json())
        preview=self.preview(response.data).get_json()
        self.assertTrue(preview['canConfirm'],preview)
        self.assertEqual(self.counts(),before)
        result=self.confirm(preview)
        self.assertEqual(result.status_code,200,result.get_json())
        self.assertEqual(result.get_json()['newInventoryLines'],1)
        after=self.counts()
        repeated=self.preview(response.data).get_json()
        result=self.confirm(repeated)
        self.assertEqual(result.status_code,200,result.get_json())
        self.assertTrue(result.get_json()['idempotent'])
        self.assertEqual(self.counts(),after)

    def test_real_cutoff_and_tax_filter_do_not_use_later_stock(self):
        with server.db() as conn:
            conn.execute("UPDATE products SET tax='KKKNT' WHERE code='BK-P1'")
            cid=conn.execute("""INSERT INTO invoice_inventory_confirmations(
                confirmation_key,direction,source_invoice_table,source_invoice_id,action,confirmed,note,created_at)
                VALUES('DRAFT-TEST','output','outgoing_source_invoices',999,'post',1,'test',?)""",(fixture.NOW,)).lastrowid
            for index,(code,day,qty) in enumerate([('BK-P1','2026-08-15',-2),('BK-P2','2026-08-16',-3),('BK-P1','2026-09-01',-10)]):
                conn.execute("""INSERT INTO invoice_inventory_ledger(
                    event_key,direction,event_type,source_invoice_table,source_invoice_id,source_line_id,
                    source_line_index,product_code,txn_date,qty_delta,unit_cost,confirmation_id,status,created_at)
                    VALUES(?,'output','POST','outgoing_source_invoices',999,?,?, ?,?,?,100,?,'posted',?)""",
                    (f'DRAFT-OUT-{index}',index+1,index+1,code,day,qty,cid,fixture.NOW))
        before=self.counts()
        result=self.client.get('/api/bk-import/shortages?from=2026-08-01&to=2026-08-31').get_json()
        self.assertTrue(result['ok'],result)
        self.assertEqual([(r['product_code'],r['closing_qty'],r['suggested_qty']) for r in result['items']],[('BK-P1',-2,2)])
        all_rows=self.client.get('/api/bk-import/shortages?from=2026-08-01&to=2026-08-31&tax=all').get_json()['items']
        self.assertEqual(len(all_rows),2)
        self.assertEqual(self.counts(),before)

    def test_reject_bad_dates_numbers_codes_without_writes(self):
        before=self.counts()
        for change in [{'from':'2026-09-01'},{'document_date':'2026-02-30'},{'rows':[]},
                       {'rows':[{'product_code':'MISSING','qty':1}]},
                       *[{'rows':[{'product_code':'BK-P1','qty':v,'unit_cost':v}]} for v in [-1,0,True,'NaN','Infinity','1e1000','1e15']]]:
            with self.subTest(change=change):
                body=self.body();body.update(change)
                result=self.client.post('/api/bk-import/draft/excel',json=body)
                self.assertEqual(result.status_code,400,result.get_json())
        self.assertEqual(self.counts(),before)

    def test_text_is_literal_canonical_units_and_tax_unchanged(self):
        body=self.body(True)
        body['rows'][0].update(note='<b>Chưa có hóa đơn</b>',unit='fake',product_name='fake')
        response=self.client.post('/api/bk-import/draft/excel',json=body)
        self.assertEqual(response.status_code,200,response.get_json())
        wb=load_workbook(io.BytesIO(response.data))
        self.assertEqual(wb['BK_IMPORT']['F4'].value,'Hàng BK')
        self.assertEqual(wb['BK_IMPORT']['G4'].value,'kg')
        self.assertEqual(wb['BK_IMPORT']['L4'].data_type,'s')
        wb.close()
        body['rows'][0]['note']='=HYPERLINK("https://example.invalid")'
        self.assertEqual(self.client.post('/api/bk-import/draft/excel',json=body).status_code,400)
        with server.db() as conn:
            self.assertEqual(conn.execute("SELECT tax FROM products WHERE code='BK-P1'").fetchone()[0],'8%')

    def sale(self, day, price, unit='Kg', status='approved', returned=0):
        with server.db() as conn:
            bid=conn.execute('INSERT INTO batches(work_date,status,created_at) VALUES(?,?,?)', (day,status,fixture.NOW)).lastrowid
            return conn.execute('''INSERT INTO orders(batch_id,work_date,contractor,kitchen,product_code,
                unit,qty,actual_delivered,customer_return_qty,sell_price,updated_at)
                VALUES(?,?,'C1','K1','BK-P1',?,2,2,?,?,?)''',
                (bid,day,unit,returned,price,fixture.NOW)).lastrowid

    def price(self, end='2026-08-31', code='BK-P1'):
        r=self.client.get('/api/bk-import/suggested-price',query_string={'to':end,'product_code':code})
        self.assertEqual(r.status_code,200,r.get_json())
        return r.get_json()

    def test_price_95_uses_latest_approved_same_unit_and_cutoff(self):
        self.sale('2026-08-05',10000)
        chosen=self.sale('2026-08-12',175001)
        self.sale('2026-08-13',999999,unit='Gói')
        self.sale('2026-08-14',999999,status='draft')
        self.sale('2026-08-15',999999,returned=2)
        self.sale('2026-09-01',999999)
        before=self.counts()
        p=self.price()
        self.assertEqual(p['unit_cost'],166251)
        self.assertEqual(p['reference_sell_price'],175001)
        self.assertEqual(p['price_order_id'],chosen)
        self.assertIn('12/08/2026',p['price_source'])
        self.assertEqual(self.counts(),before)

    def test_missing_price_stays_editable_and_manual_cost_is_preserved(self):
        self.sale('2026-08-15',10000,unit='Gói')
        self.assertEqual(self.price()['unit_cost'],'')
        self.assertEqual(self.price(code='BK-P2')['unit_cost'],'')
        self.sale('2026-08-16',10000)
        body=self.body(True)
        body['rows'][0]['unit_cost']=1234
        response=self.client.post('/api/bk-import/draft/excel',json=body)
        wb=load_workbook(io.BytesIO(response.data),data_only=True)
        self.assertEqual(wb['BK_IMPORT']['I4'].value,1234)
        self.assertEqual(wb['BK_IMPORT']['J4'].value,2468)
        wb.close()
        self.assertEqual(self.client.get('/api/bk-import/suggested-price?to=bad&product_code=BK-P1').status_code,400)
        self.assertEqual(self.client.get('/api/bk-import/suggested-price?to=2026-08-31&product_code=UNKNOWN').status_code,400)

    def test_older_shortage_can_use_signed_invoice_without_guessing_unit(self):
        try:
            with server.db() as conn:
                for index,(day,unit,status,stock,price) in enumerate([
                    ('2026-08-01','Kg','issued','posted',35000),
                    ('2026-08-02','Gói','issued','posted',90000),
                    ('2026-08-03','Kg','draft','blocked',90000),
                    ('2026-08-04','Kg','issued','ready',90000),
                    ('2026-09-01','Kg','issued','posted',90000)]):
                    iid=conn.execute('''INSERT INTO outgoing_source_invoices(tenant,source,identity_key,
                        invoice_number,invoice_date,source_status_class,stock_status,synced_at,created_at,updated_at)
                        VALUES('test','bk-price-test',?,?,?,?,?,?,?,?)''',
                        (str(index),str(index+700),day,status,stock,fixture.NOW,fixture.NOW,fixture.NOW)).lastrowid
                    conn.execute('''INSERT INTO outgoing_source_invoice_items(invoice_id,line_index,product_code,
                        source_unit,qty,unit_price,mapping_status,conversion_factor)
                        VALUES(?,1,'BK-P1',?,1,?,'mapped',1)''',(iid,unit,price))
            before=self.counts();p=self.price()
            self.assertEqual(p['unit_cost'],33250)
            self.assertIn('HĐ 700',p['price_source'])
            self.assertEqual(p['price_date'],'2026-08-01')
            self.assertEqual(self.counts(),before)
            self.sale('2026-08-15',40000)
            self.assertEqual(self.price()['unit_cost'],38000)
        finally:
            with server.db() as conn:conn.execute("DELETE FROM outgoing_source_invoices WHERE source='bk-price-test'")


if __name__=='__main__': unittest.main()
