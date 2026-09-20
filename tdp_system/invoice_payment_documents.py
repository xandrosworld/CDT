"""Official contractor payment-request workbook based on the approved sample.

The customer workbook establishes the visible six-column invoice table and
the legal/payment wording.  It also contains stale sheets, external VnTools
formulas and fixed blank rows, so this renderer writes reconciled literal
values, spells the amount natively and paginates dynamically.
"""

from __future__ import annotations

import math
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from typing import Any

from openpyxl import Workbook, load_workbook
from copy import copy
from pathlib import Path

try:
    from .invoice_line_tax import tax_rate_label
except ImportError:
    from invoice_line_tax import tax_rate_label
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins


def synced_invoice_statement_workbook(scope):
    """VAT statement without fabricated kitchens, deliveries or inventory entries."""
    book = Workbook()
    sheet = book.active
    sheet.title = 'Bảng kê hóa đơn'
    sheet.append(['BẢNG KÊ HÓA ĐƠN VAT ĐÃ PHÁT HÀNH'])
    sheet.merge_cells('A1:G1')
    sheet.append([scope['warning']]); sheet.merge_cells('A2:G2')
    sheet.append(['STT', 'Ngày hóa đơn', 'Ký hiệu', 'Số hóa đơn', 'Trước thuế', 'Thuế', 'Thanh toán'])
    for i, invoice in enumerate(scope['invoices'], 1):
        sheet.append([i, _strict_date(invoice['invoice_date'], 'Ngày hóa đơn'),
                      _excel_text(invoice['invoice_series']), _excel_text(invoice['invoice_number']),
                      invoice['subtotal'], invoice['tax_amount'], invoice['total_amount']])
    sheet.append(['TỔNG', None, None, None, scope['totals']['subtotal'], scope['totals']['tax_amount'], scope['totals']['total_amount']])
    details = book.create_sheet('Chi tiết hóa đơn')
    details.append(['CHI TIẾT THEO HÓA ĐƠN VAT']); details.merge_cells('A1:I1')
    details.append(['Ngày dưới đây là ngày hóa đơn; không suy ra bếp hoặc ngày giao hàng.']); details.merge_cells('A2:I2')
    details.append(['Ngày hóa đơn', 'Ký hiệu / số', 'Tên hàng', 'ĐVT', 'Số lượng', 'Đơn giá', 'Thành tiền', 'Thuế suất', 'Ghi chú'])
    for item in scope.get('source_lines', []):
        details.append([_strict_date(item['invoice_date'], 'Ngày hóa đơn'),
                        _excel_text(item['invoice_series'] + ' / ' + item['invoice_number']),
                        _excel_text(item['source_item_name']), _excel_text(item['source_unit']),
                        item['qty'], item['unit_price'], item['amount'], _excel_text(tax_rate_label(item['tax_rate'])),
                        _excel_text(item['validation_note'])])
    for item in scope.get('lines', []):
        details.append([_strict_date(item['issued_invoice_date'], 'Ngày hóa đơn'),
                        _excel_text(item['issued_invoice_series'] + ' / ' + item['issued_invoice_number']),
                        _excel_text(item['product_name']), _excel_text(item['unit']), item['qty'],
                        item['unit_price'], item['amount'], _excel_text(tax_rate_label(item['tax'])), ''])
    from collections import defaultdict
    try:
        from .document_totals import quantity_text
    except ImportError:
        from document_totals import quantity_text
    quantities = defaultdict(float)
    for row in details.iter_rows(min_row=4, values_only=True):
        quantities[str(row[3] or '')] += float(row[4] or 0)
    details.append(['TỔNG', None, None, None,
                    quantity_text([{'unit': u, 'quantity': q} for u, q in quantities.items()]),
                    None, scope['totals']['subtotal']])
    for ws in book:
        ws.freeze_panes = 'A4'
        ws.sheet_view.showGridLines = False
        ws.row_dimensions[1].height = 30
        ws.row_dimensions[2].height = 66
        ws.row_dimensions[3].height = 32
        for row in ws:
            for cell in row:
                cell.font = Font(name='Times New Roman', size=12, bold=cell.row <= 3 or cell.row == ws.max_row)
                cell.alignment = Alignment(vertical='center', wrap_text=True)
                if cell.row >= 3: cell.border = TABLE_BORDER
                if isinstance(cell.value, (date, datetime)): cell.number_format = 'dd/mm/yyyy'
                elif isinstance(cell.value, (int, float)):
                    cell.number_format = '#,##0.######' if ws == details and cell.column == 5 else '#,##0'
                    cell.alignment = Alignment(horizontal='right', vertical='center')
        for col, width in enumerate((12, 24, 36, 16, 22, 22, 24, 14, 34), 1):
            ws.column_dimensions[get_column_letter(col)].width = width
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = 'landscape'
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
        ws.print_title_rows = '1:3'
        ws.print_area = f'A1:{get_column_letter(ws.max_column)}{ws.max_row}'
    return book


class InvoicePaymentDocumentError(ValueError):
    def __init__(self, message: str, *, code: str = "invalid_invoice_payment_document", status: int = 409):
        super().__init__(message)
        self.code = code
        self.status = status


def _plain(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").replace("\u00a0", " ").strip())


def _excel_text(value: Any) -> str:
    text = _plain(value)
    return "'" + text if text.startswith(("=", "+", "-", "@")) else text


def _strict_date(value: Any, label: str) -> date:
    text = _plain(value)
    try:
        return datetime.strptime(text, "%Y-%m-%d").date()
    except ValueError:
        raise InvoicePaymentDocumentError(
            f"{label} phải là ngày YYYY-MM-DD hợp lệ", code="invalid_payment_document_date",
        ) from None


def _vnd(value: Any, label: str = "Số tiền") -> int:
    try:
        number = Decimal(str(value or 0))
    except (InvalidOperation, TypeError, ValueError):
        raise InvoicePaymentDocumentError(f"{label} không hợp lệ") from None
    if not number.is_finite():
        raise InvoicePaymentDocumentError(f"{label} không hữu hạn")
    return int(number.quantize(Decimal("1"), rounding=ROUND_HALF_UP))


VIET_DIGITS = ("không", "một", "hai", "ba", "bốn", "năm", "sáu", "bảy", "tám", "chín")


def number_to_vietnamese(value: Any) -> str:
    number = _vnd(value)
    if number < 0:
        raise InvoicePaymentDocumentError("Tổng đề nghị thanh toán không được âm")
    if number == 0:
        return "Không đồng"

    def three_digits(group: int, full: bool = False) -> str:
        hundred, remainder = divmod(group, 100)
        ten, unit = divmod(remainder, 10)
        words: list[str] = []
        if hundred or full:
            words.extend((VIET_DIGITS[hundred], "trăm"))
        if ten > 1:
            words.extend((VIET_DIGITS[ten], "mươi"))
            if unit == 1:
                words.append("mốt")
            elif unit == 5:
                words.append("lăm")
            elif unit:
                words.append(VIET_DIGITS[unit])
        elif ten == 1:
            words.append("mười")
            if unit == 5:
                words.append("lăm")
            elif unit:
                words.append(VIET_DIGITS[unit])
        elif unit:
            if hundred or full:
                words.append("lẻ")
            words.append(VIET_DIGITS[unit])
        return " ".join(words)

    groups: list[int] = []
    while number:
        groups.append(number % 1000)
        number //= 1000
    units = ("", "nghìn", "triệu", "tỷ", "nghìn tỷ", "triệu tỷ")
    if len(groups) > len(units):
        raise InvoicePaymentDocumentError("Tổng đề nghị thanh toán vượt giới hạn đọc bằng chữ")
    words: list[str] = []
    for index in range(len(groups) - 1, -1, -1):
        group = groups[index]
        if not group:
            continue
        words.append(three_digits(group, full=bool(words) and group < 100))
        if units[index]:
            words.append(units[index])
    result = " ".join(words).strip()
    return result[:1].upper() + result[1:] + " đồng"


THIN = Side(style="thin", color="000000")
TABLE_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _font(*, size: float = 12, bold: bool = False, italic: bool = False, color: str = "000000") -> Font:
    return Font(name="Times New Roman", size=size, bold=bold, italic=italic, color=color)


def _merge_write(
    ws, cell_range: str, value: Any, *, size: float = 12, bold: bool = False,
    italic: bool = False, color: str = "000000", horizontal: str = "left",
    vertical: str = "center", wrap: bool = True,
) -> None:
    ws.merge_cells(cell_range)
    cell = ws[cell_range.split(":", 1)[0]]
    text = _plain(value)
    # openpyxl only marks leading '=' as a formula.  Keep ordinary legal
    # punctuation such as dashed separators and '- Như trên' visually clean.
    cell.value = "'" + text if text.startswith("=") else text
    cell.font = _font(size=size, bold=bold, italic=italic, color=color)
    cell.alignment = Alignment(horizontal=horizontal, vertical=vertical, wrap_text=wrap)


def _validated_scope(scope: dict[str, Any]) -> tuple[list[dict[str, Any]], dict[str, Any], dict[str, Any]]:
    invoices = list(scope.get("invoices") or [])
    snapshot = dict(scope.get("snapshot") or {})
    totals = dict(scope.get("totals") or {})
    if not invoices:
        raise InvoicePaymentDocumentError("Không có hóa đơn đã phát hành để lập đề nghị")
    required_snapshot = (
        "buyer_name_snapshot", "buyer_tax_code_snapshot", "buyer_address_snapshot",
        "company_name_snapshot", "company_tax_code_snapshot", "company_address_snapshot",
        "payment_bank_name_snapshot", "payment_bank_account_snapshot",
    )
    missing = [name for name in required_snapshot if not _plain(snapshot.get(name))]
    if missing:
        raise InvoicePaymentDocumentError("Thiếu hồ sơ pháp lý hoặc thanh toán đã khóa cùng hóa đơn")
    computed = {
        "subtotal": sum(_vnd(item.get("subtotal")) for item in invoices),
        "tax_amount": sum(_vnd(item.get("tax_amount")) for item in invoices),
        "total_amount": sum(_vnd(item.get("total_amount")) for item in invoices),
    }
    expected = {key: _vnd(totals.get(key)) for key in computed}
    if computed != expected or expected["total_amount"] <= 0:
        raise InvoicePaymentDocumentError(
            "Tổng đề nghị thanh toán không khớp phạm vi hóa đơn",
            code="payment_document_total_mismatch",
        )
    identities: set[tuple[str, str, str]] = set()
    for item in invoices:
        invoice_date = _strict_date(item.get("invoice_date"), "Ngày hóa đơn").isoformat()
        number = _plain(item.get("invoice_number"))
        series = _plain(item.get("invoice_series")).upper()
        if not number:
            raise InvoicePaymentDocumentError("Hóa đơn thiếu số hóa đơn")
        identity = (invoice_date, series, number)
        if identity in identities:
            raise InvoicePaymentDocumentError(
                "Phạm vi có hóa đơn bị lặp", code="payment_document_duplicate_invoice",
            )
        identities.add(identity)
    return invoices, snapshot, expected


def invoice_payment_request_workbook(
    scope: dict[str, Any], *, issue_date: Any | None = None,
    request_number: Any = "……/CV/ĐNTT", contract_no: Any = "", contract_date: Any = "",
) -> Workbook:
    """Create one official, print-ready payment request for one contractor."""
    invoices, snapshot, totals = _validated_scope(scope)
    period_from = _strict_date(scope.get("date_from"), "Từ ngày")
    period_to = _strict_date(scope.get("date_to"), "Đến ngày")
    if period_from > period_to:
        raise InvoicePaymentDocumentError("Từ ngày không được lớn hơn Đến ngày")
    issued_on = _strict_date(issue_date or period_to.isoformat(), "Ngày lập")
    safe_contract_date = None
    if _plain(contract_date):
        safe_contract_date = _strict_date(contract_date, "Ngày hợp đồng")
    safe_contract_no = _plain(contract_no)

    workbook = Workbook()
    workbook.properties.title = (
        f"Đề nghị thanh toán {_plain(scope.get('contractor'))} "
        f"{period_from.isoformat()} - {period_to.isoformat()}"
    )
    workbook.properties.subject = "Đề nghị thanh toán từ hóa đơn đỏ đã phát hành"
    workbook.properties.creator = "Thành Đạt Phát"
    workbook.calculation.fullCalcOnLoad = False
    workbook.calculation.forceFullCalc = False
    workbook.calculation.calcMode = "manual"
    # Use the customer's print sheet artwork, with only current scoped values.
    template_book = load_workbook(Path(__file__).parent / 'templates' / 'invoice_payment_customer_20260920.xlsx')
    template = template_book.active
    ws = workbook.active
    ws.title = "Đề nghị thanh toán"
    ws.sheet_view.showGridLines = False
    for letter, dimension in template.column_dimensions.items():
        ws.column_dimensions[letter].width = dimension.width
    count = len(invoices)
    row_pairs = [(r, r) for r in range(1, 11)]
    row_pairs += [(11, 11 + i) for i in range(count)]
    row_pairs += [(r, r + count - 1) for r in range(12, 23)]
    for source_row, target_row in row_pairs:
        ws.row_dimensions[target_row].height = template.row_dimensions[source_row].height
        for column in range(1, 9):
            source, target = template.cell(source_row, column), ws.cell(target_row, column)
            for attribute in ('font', 'fill', 'border', 'alignment', 'protection', 'number_format'):
                setattr(target, attribute, copy(getattr(source, attribute)))
            font = copy(target.font)
            font.color = '000000'  # Keep the customer's established black print output.
            target.font = font
            target.value = source.value
    for merged in template.merged_cells.ranges:
        offset = count - 1 if merged.min_row >= 12 else 0
        ws.merge_cells(start_row=merged.min_row + offset, end_row=merged.max_row + offset,
                       start_column=merged.min_col, end_column=merged.max_col)
    ws.page_margins = copy(template.page_margins)
    template_book.close()

    def write(cell, value):
        # Literal data, never formula expressions from names or request fields.
        ws[cell] = "'" + value if isinstance(value, str) and value.startswith(('=', '+', '-', '@')) else value

    write('A1', snapshot['company_name_snapshot'].upper())
    write('A2', _plain(request_number))
    write('E3', f"Hải Phòng, ngày {issued_on.day:02d} tháng {issued_on.month:02d} năm {issued_on.year}")
    write('A6', f"Kính gửi: {snapshot['buyer_name_snapshot'].upper()}")
    if safe_contract_no:
        contract_text = f"Căn cứ vào hợp đồng mua bán số: {safe_contract_no}"
        if safe_contract_date:
            contract_text += f" ngày {safe_contract_date.day:02d} tháng {safe_contract_date.month:02d} năm {safe_contract_date.year}"
        contract_text += f" giữa {snapshot['buyer_name_snapshot']} và {snapshot['company_name_snapshot']}."
    else:
        contract_text = f"Căn cứ hàng hóa đã cung cấp giữa {snapshot['buyer_name_snapshot']} và {snapshot['company_name_snapshot']}."
    write('A7', '          ' + contract_text)
    write('A8', f"         Thời gian từ ngày {period_from.strftime('%d/%m/%Y')} – {period_to.strftime('%d/%m/%Y')} "
          f"chúng tôi đã cung cấp hàng hoá cho {snapshot['buyer_name_snapshot']}, dựa theo số lượng "
          "bàn giao chúng tôi đã xuất hóa đơn như sau:")
    for index, item in enumerate(invoices, 1):
        row = 10 + index
        values = (index, _strict_date(item['invoice_date'], 'Ngày hóa đơn'),
                  _excel_text(item['invoice_number']), _vnd(item['subtotal']),
                  _vnd(item['tax_amount']), _vnd(item['total_amount']))
        for column, value in enumerate(values, 2):
            ws.cell(row, column, value)
        ws.cell(row, 3).number_format = 'dd/mm/yyyy'
        for column in (5, 6, 7):
            ws.cell(row, column).number_format = '#,##0'
    total_row = 11 + count
    for column, key in zip((5, 6, 7), ('subtotal', 'tax_amount', 'total_amount')):
        ws.cell(total_row, column, totals[key]).number_format = '#,##0'
    write(f'C{total_row + 1}', number_to_vietnamese(totals['total_amount']) + './.')
    write(f'A{total_row + 3}', f"        Tên tài khoản: {snapshot['company_name_snapshot'].upper()}")
    write(f'A{total_row + 4}', f"        Số tài khoản: {snapshot['payment_bank_account_snapshot']} Tại {snapshot['payment_bank_name_snapshot']}")
    # Grow narrative rows for long customer names; keep the invoice table compact.
    for row in (7, 8):
        ws.row_dimensions[row].height = max(ws.row_dimensions[row].height or 0,
                                            math.ceil(len(ws.cell(row, 1).value or '') / 110) * 15)
    ws.print_area = f'A1:H{total_row + 14}'
    ws.print_title_rows = '10:10'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'portrait'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_options.horizontalCentered = True

    proof = workbook.create_sheet("Đối chiếu hóa đơn")
    proof.append([
        "Ngày HĐ", "Ký hiệu", "Số HĐ", "Trước thuế", "Thuế", "Tổng",
        "Nguồn xác minh", "ID nguồn", "ID draft",
    ])
    for item in invoices:
        proof.append([
            _strict_date(item.get("invoice_date"), "Ngày hóa đơn"),
            _excel_text(item.get("invoice_series")), _excel_text(item.get("invoice_number")),
            _vnd(item.get("subtotal")), _vnd(item.get("tax_amount")),
            _vnd(item.get("total_amount")), _excel_text(item.get("verification_source")),
            item.get("source_invoice_id") or "", item.get("draft_id") or "",
        ])
    proof.append(["TỔNG", "", "", totals["subtotal"], totals["tax_amount"], totals["total_amount"], "", "", ""])
    for cell in proof[1]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17324D")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in proof.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=9, bold=cell.row == proof.max_row)
            cell.border = TABLE_BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    for column in (4, 5, 6):
        for row in range(2, proof.max_row + 1):
            proof.cell(row, column).number_format = "#,##0"
    for column, width in enumerate((14, 15, 15, 17, 15, 18, 24, 14, 12), start=1):
        proof.column_dimensions[get_column_letter(column)].width = width
    proof.freeze_panes = "A2"
    proof.auto_filter.ref = f"A1:I{proof.max_row - 1}"
    proof.sheet_state = "hidden"
    return workbook


__all__ = [
    "InvoicePaymentDocumentError", "invoice_payment_request_workbook",
    "number_to_vietnamese",
]
