"""Explain catalog versus source-invoice tax without changing either source."""
from decimal import Decimal, InvalidOperation
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

try:
    from .invoice_line_tax import tax_rate_label
    from .stock_tax_policy import is_kkknt
except ImportError:
    from invoice_line_tax import tax_rate_label
    from stock_tax_policy import is_kkknt


def display_tax(value):
    if is_kkknt(value):
        return 'KKKNT'
    raw = tax_rate_label(value).strip()
    if not raw:
        return 'Chưa rõ'
    if raw.upper() in {'KCT', 'KHÔNG CHỊU THUẾ', 'KHONG CHIU THUE'}:
        return 'KCT'
    try:
        rate = Decimal(raw.rstrip('%').replace(',', '.'))
        if not rate.is_finite():
            return 'Chưa rõ'
        if '%' not in raw and 0 < rate < 1:
            rate *= 100
        return format(rate.normalize(), 'f') + '%'
    except InvalidOperation:
        return raw


def compare_tax(item, lines):
    catalog = display_tax(item.get('tax'))
    rates = sorted({display_tax(r.get('tax_rate')) for r in lines})
    differs = catalog != 'Chưa rõ' and any(r != 'Chưa rõ' and r != catalog for r in rates)
    note = ('Thuế danh mục khác thuế HĐ; kiểm tra danh mục và mã ghép, giữ nguyên HĐ gốc.' if differs
            else 'Danh mục chưa có thuế.' if catalog == 'Chưa rõ'
            else 'Có dòng HĐ chưa rõ thuế.' if 'Chưa rõ' in rates
            else 'Chưa có dòng HĐ để đối chiếu.' if not rates else 'Khớp')
    return dict(catalog=catalog, source=' / '.join(rates) or 'Không có dòng HĐ',
                differs=differs, note=note)


def format_review_sheet(sheet, header_row, widths):
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = 's'
            cell.alignment = Alignment(wrap_text=True, vertical='center')
            if cell.row == header_row:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='17354A')
        sheet.row_dimensions[row[0].row].height = 34
    sheet.freeze_panes = f'C{header_row+1}'
    sheet.auto_filter.ref = f'A{header_row}:{get_column_letter(len(widths))}{max(header_row,sheet.max_row)}'


def add_tax_review(book, main_sheet, items, outgoing, first_row, invoices=()):
    reviews = []
    headers = {invoice['id']:invoice for invoice in invoices}
    for index, item in enumerate(items, first_row):
        lines = outgoing.get(item['product_code'], [])
        review = compare_tax(item, lines)
        if review['differs']:
            main_sheet.cell(index, 5).comment = Comment(
                f"Thuế danh mục: {review['catalog']}. Thuế HĐ đầu ra trong kỳ: {review['source']}. "
                'Xem sheet Đối chiếu thuế. Chưa tự thay đổi nguồn thuế.', 'TDP')
            references = set()
            for line in lines:
                header = headers.get(line.get('invoice_id'), line)
                if header.get('invoice_number') and display_tax(line.get('tax_rate')) not in ('Chưa rõ',review['catalog']):
                    references.add(str(header.get('invoice_series') or '') + ' / ' + str(header['invoice_number']))
            reviews.append([item['product_code'], item['product_name'], review['catalog'],
                            review['source'], len(lines), review['note'], ', '.join(sorted(references))])
    if reviews:
        sheet = book.create_sheet('Đối chiếu thuế')
        sheet.append(['T/Suất trên báo cáo tồn ưu tiên danh mục; thuế hóa đơn giữ theo nguồn. Khác nhau cần đối chiếu.'])
        sheet.merge_cells('A1:G1')
        sheet.append(['Mã nội bộ', 'Tên hàng', 'Thuế danh mục', 'Thuế HĐ đầu ra trong kỳ', 'Số dòng HĐ', 'Đối chiếu', 'Hóa đơn có thuế khác danh mục'])
        for row in reviews:
            sheet.append(row)
        format_review_sheet(sheet, 2, (18, 40, 20, 28, 16, 65, 55))
    return len(reviews)
