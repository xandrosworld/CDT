"""Build the plain-language, screenshot-led TDP user guide as a DOCX.

The guide is intentionally written for an office user with no technical
background.  Screenshots are captured from the exact release executable on an
isolated test database; no connector credentials are read by this script.
"""

from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SCREENS = ROOT / "tmp" / "user_guide_screens_20260904_cap_nhat"
OUTPUT_DIR = ROOT / "HUONG_DAN_TDP_20260904_CAP_NHAT"
OUTPUT = OUTPUT_DIR / "HUONG_DAN_SU_DUNG_THANH_DAT_PHAT_CAP_NHAT.docx"

VERSION = "2026.09.04.1"
ISSUE_DATE = "04/09/2026"

NAVY = "17324D"
TEAL = "078778"
TEAL_DARK = "006B61"
TEXT = "203040"
MUTED = "607284"
WHITE = "FFFFFF"
LINE = "D6E2E8"
PALE_TEAL = "EAF6F3"
PALE_BLUE = "EEF5F9"
PALE_YELLOW = "FFF6DF"
PALE_RED = "FDEEEE"
PALE_GRAY = "F4F7F9"
RED = "B53030"
AMBER = "A76A00"


def set_cell_shading(cell, fill: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    shade = props.find(qn("w:shd"))
    if shade is None:
        shade = OxmlElement("w:shd")
        props.append(shade)
    shade.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = LINE, size: str = "5") -> None:
    props = cell._tc.get_or_add_tcPr()
    borders = props.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        props.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_cell_margins(cell, top: int = 120, start: int = 150,
                     bottom: int = 120, end: int = 150) -> None:
    props = cell._tc.get_or_add_tcPr()
    margins = props.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        props.append(margins)
    for name, value in (("top", top), ("start", start),
                        ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_run(run, *, size: float | None = None, bold: bool | None = None,
            color: str | None = None, italic: bool | None = None) -> None:
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Arial")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if italic is not None:
        run.italic = italic


def add_field(run, instruction: str) -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, text, separate, end))


def set_repeat_table_header(row) -> None:
    props = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    props.append(marker)


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(1.55)
section.bottom_margin = Cm(1.45)
section.left_margin = Cm(1.65)
section.right_margin = Cm(1.65)
section.header_distance = Cm(0.55)
section.footer_distance = Cm(0.55)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Arial")
normal.font.size = Pt(12)
normal.font.color.rgb = RGBColor.from_string(TEXT)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.13

for name, size, color, before, after in (
    ("Title", 29, NAVY, 0, 12),
    ("Heading 1", 20, NAVY, 0, 8),
    ("Heading 2", 15.5, TEAL_DARK, 7, 5),
    ("Heading 3", 13, NAVY, 5, 3),
):
    style = styles[name]
    style.font.name = "Arial"
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Arial")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

settings = doc.settings._element
compat = settings.find(qn("w:compat"))
if compat is None:
    compat = OxmlElement("w:compat")
    settings.append(compat)
compatibility_mode = next(
    (
        item
        for item in compat.iter(qn("w:compatSetting"))
        if item.get(qn("w:name")) == "compatibilityMode"
    ),
    None,
)
if compatibility_mode is None:
    compatibility_mode = OxmlElement("w:compatSetting")
    compatibility_mode.set(qn("w:name"), "compatibilityMode")
    compatibility_mode.set(qn("w:uri"), "http://schemas.microsoft.com/office/word")
    compat.append(compatibility_mode)
compatibility_mode.set(qn("w:val"), "15")

update_fields = settings.find(qn("w:updateFields"))
if update_fields is None:
    update_fields = OxmlElement("w:updateFields")
    settings.append(update_fields)
update_fields.set(qn("w:val"), "true")

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.LEFT
run = header.add_run("THÀNH ĐẠT PHÁT  •  HƯỚNG DẪN SỬ DỤNG")
set_run(run, size=8.5, bold=True, color=TEAL_DARK)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = footer.add_run("Trang ")
set_run(run, size=9, color=MUTED)
add_field(run, " PAGE ")


def paragraph(text: str = "", *, bold: bool = False, color: str | None = None,
              size: float | None = None, align=None, italic: bool = False,
              keep: bool = False):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.keep_together = keep
    r = p.add_run(text)
    set_run(r, size=size, bold=bold, color=color, italic=italic)
    return p


def page_break() -> None:
    # Put the break at the end of the preceding body paragraph.  A separate
    # empty paragraph can be pushed onto the next page when a screenshot page
    # is already full, which would create an unwanted blank page.
    if doc.paragraphs:
        doc.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)
    else:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def callout(title: str, text: str, *, fill: str = PALE_TEAL,
            accent: str = TEAL_DARK) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    left, body = table.rows[0].cells
    left.width = Cm(0.25)
    body.width = Cm(16.9)
    set_cell_shading(left, accent)
    set_cell_shading(body, fill)
    for cell in (left, body):
        set_cell_border(cell, fill, "0")
        set_cell_margins(cell, 130, 160, 130, 160)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = body.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title)
    set_run(r, size=11.5, bold=True, color=accent)
    p = body.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    set_run(r, size=11)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)


def steps(items: list[str]) -> None:
    for index, item in enumerate(items, start=1):
        table = doc.add_table(rows=1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        number, body = table.rows[0].cells
        number.width = Cm(0.9)
        body.width = Cm(16.25)
        set_cell_shading(number, TEAL)
        set_cell_shading(body, PALE_GRAY if index % 2 else WHITE)
        for cell in (number, body):
            set_cell_border(cell, WHITE, "0")
            set_cell_margins(cell, 85, 120, 85, 120)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = number.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(str(index))
        set_run(r, size=12, bold=True, color=WHITE)
        p = body.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(item)
        set_run(r, size=11.3)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)


def bullets(items: list[str], *, checkbox: bool = False) -> None:
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.55)
        p.paragraph_format.first_line_indent = Cm(-0.38)
        p.paragraph_format.space_after = Pt(3)
        marker = "☐" if checkbox else "•"
        r = p.add_run(f"{marker}  ")
        set_run(r, size=11.5, bold=True, color=TEAL_DARK)
        r = p.add_run(item)
        set_run(r, size=11.3)


def simple_table(headers: list[str], rows: list[tuple[str, ...]],
                 widths: list[float]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, text in enumerate(headers):
        cell = header.cells[i]
        cell.width = Cm(widths[i])
        set_cell_shading(cell, NAVY)
        set_cell_border(cell, WHITE, "4")
        set_cell_margins(cell)
        r = cell.paragraphs[0].add_run(text)
        set_run(r, size=10.5, bold=True, color=WHITE)
    for row_index, row in enumerate(rows, start=1):
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cell = cells[i]
            cell.width = Cm(widths[i])
            set_cell_shading(cell, WHITE if row_index % 2 else PALE_GRAY)
            set_cell_border(cell)
            set_cell_margins(cell)
            r = cell.paragraphs[0].add_run(text)
            set_run(r, size=10.5)
    doc.add_paragraph()


def add_screen(filename: str, caption: str) -> None:
    path = SCREENS / filename
    if not path.is_file():
        raise FileNotFoundError(path)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(path), width=Cm(17.1))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(caption)
    set_run(r, size=9.2, color=MUTED, italic=True)


def screen_page(title: str, filename: str, caption: str,
                intro: str, actions: list[str], *, note: tuple[str, str] | None = None,
                heading_level: int = 2) -> None:
    page_break()
    doc.add_heading(title, level=heading_level)
    paragraph(intro, size=11.5, color=MUTED)
    add_screen(filename, caption)
    steps(actions)
    if note:
        callout(note[0], note[1], fill=PALE_YELLOW, accent=AMBER)


# Cover
paragraph("THÀNH ĐẠT PHÁT", size=16, bold=True, color=TEAL_DARK,
          align=WD_ALIGN_PARAGRAPH.CENTER)
paragraph("HƯỚNG DẪN SỬ DỤNG\nHỆ THỐNG VẬN HÀNH", size=29, bold=True,
          color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER)
paragraph("Dành cho người mới sử dụng – không cần biết kỹ thuật",
          size=14, bold=True, color=TEAL, align=WD_ALIGN_PARAGRAPH.CENTER)
paragraph("", size=6)

cover = doc.add_table(rows=3, cols=1)
cover.alignment = WD_TABLE_ALIGNMENT.CENTER
cover.autofit = False
cover_items = (
    ("MỖI NGÀY", "Chọn ngày → nạp đơn → kiểm tra → đặt hàng → tải giấy tờ → sao lưu", PALE_TEAL),
    ("KHI CÓ LỖI", "Dừng lại, đọc dòng báo lỗi và sửa đúng chỗ. Không bấm liên tiếp nhiều lần.", PALE_YELLOW),
    ("KHI CHƯA CHẮC", "Chụp toàn màn hình và hỏi người phụ trách trước khi xác nhận hoặc ghi kho.", PALE_BLUE),
)
for row, (title, body, fill) in zip(cover.rows, cover_items):
    cell = row.cells[0]
    cell.width = Cm(15.8)
    set_cell_shading(cell, fill)
    set_cell_border(cell, WHITE, "8")
    set_cell_margins(cell, 190, 250, 190, 250)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    set_run(r, size=12, bold=True, color=NAVY)
    p = cell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(body)
    set_run(r, size=11.5)

paragraph("", size=7)
paragraph(f"Phiên bản phần mềm: {VERSION}  •  Ngày cập nhật: {ISSUE_DATE}",
          size=10.5, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER)
paragraph("Ảnh trong tài liệu được chụp trực tiếp từ bản phần mềm bàn giao. "
          "Số liệu trên máy thực tế có thể khác ảnh minh họa.",
          size=9.5, color=MUTED, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER)


# Quick start
page_break()
doc.add_heading("LÀM VIỆC HẰNG NGÀY – CHỈ CẦN NHỚ 7 VIỆC", level=1)
steps([
    "Mở file Thanh_Dat_Phat.exe và chờ trình duyệt hiện ra. Chỉ mở một lần.",
    "Chọn đúng ngày hoặc đúng đơn ở ô phía trên bên phải.",
    "Vào Công việc hằng ngày, chọn Từ ngày – Đến ngày rồi nạp file Excel.",
    "Vào Nhập & sửa đơn, sửa hết những dòng hệ thống báo thiếu hoặc sai.",
    "Vào Đặt hàng nhà cung cấp để gửi đơn; vào Phiếu giao để tải phiếu giao hàng.",
    "Khi cần giấy tờ, vào In giấy tờ, chọn đúng loại và đúng ngày rồi tải file.",
    "Cuối ngày vào Danh mục & sao lưu, bấm Tải bản sao lưu dữ liệu.",
])
callout(
    "Dấu hiệu đã lưu",
    "Góc dưới bên trái có chấm xanh và dòng “Dữ liệu đang được tự động lưu”. "
    "Sau mỗi nút Lưu/Xác nhận, chờ thông báo thành công rồi mới chuyển sang việc khác.",
)
doc.add_heading("Ba màu cần nhớ", level=2)
simple_table(
    ["Màu nhìn thấy", "Nghĩa là", "Cần làm"],
    [
        ("Xanh", "Đã sẵn sàng hoặc đã làm xong", "Có thể sang bước tiếp theo"),
        ("Vàng", "Cần đọc và kiểm tra", "Đối chiếu file/số liệu trước khi tiếp tục"),
        ("Đỏ", "Đang thiếu hoặc sai", "Phải sửa; chưa duyệt, chưa in, chưa ghi kho"),
    ],
    [3.0, 5.1, 9.0],
)
callout(
    "Quy tắc an toàn",
    "Không xóa thư mục data được tạo cạnh file EXE. Không gửi file dữ liệu, bản sao lưu "
    "hoặc thông tin đăng nhập cho người không có trách nhiệm.",
    fill=PALE_RED,
    accent=RED,
)


# Contents
page_break()
doc.add_heading("MỤC LỤC", level=1)
p = doc.add_paragraph()
r = p.add_run()
add_field(r, ' TOC \\o "1-2" \\h \\z \\u ')
set_run(r, size=11)
paragraph("Nếu số trang chưa tự cập nhật: bấm chuột phải vào mục lục → Update Field → "
          "Update entire table.", size=9.5, color=MUTED, italic=True)


# Starting the app
page_break()
doc.add_heading("1. MỞ VÀ ĐÓNG HỆ THỐNG", level=1)
doc.add_heading("1.1. Lần đầu nhận file", level=2)
steps([
    "Tạo một thư mục riêng, ví dụ THANH_DAT_PHAT, rồi đặt Thanh_Dat_Phat.exe vào trong đó.",
    "Không đặt lẫn với nhiều file tải về khác. Hệ thống sẽ tự tạo thư mục data cạnh file EXE để lưu số liệu.",
    "Nhấp đúp Thanh_Dat_Phat.exe. Lần mở đầu có thể chờ khoảng 30–60 giây.",
    "Khi trình duyệt mở trang Thành Đạt Phát, bắt đầu làm việc. Giữ cửa sổ màu đen đang chạy ở phía sau.",
])
callout(
    "Nếu Windows hỏi xác nhận",
    "Chỉ khi đúng file do đơn vị triển khai gửi: chọn More info → Run anyway. "
    "Nếu tên file hoặc nguồn gửi không đúng, dừng lại và hỏi người phụ trách.",
    fill=PALE_YELLOW,
    accent=AMBER,
)
doc.add_heading("1.2. Đóng đúng cách", level=2)
steps([
    "Lưu xong công việc và tải bản sao lưu dữ liệu.",
    "Đóng các file Excel/PDF đang mở.",
    "Đóng tab trình duyệt của hệ thống.",
    "Đóng cửa sổ màu đen của Thanh Đạt Phát. Không tắt máy giữa lúc hệ thống đang lưu.",
])
callout(
    "Không mở hai lần",
    "Nếu đã có một cửa sổ Thanh Đạt Phát đang chạy, không nhấp đúp EXE lần nữa. "
    "Chỉ quay lại tab trình duyệt đang mở.",
    fill=PALE_RED,
    accent=RED,
)


# Screen 1
screen_page(
    "2. CÔNG VIỆC HẰNG NGÀY",
    "01_cong_viec_hang_ngay.png",
    "Màn hình tổng quan: bốn việc chính trong ngày và khoảng ngày đang chọn.",
    "Đây là trang nên mở đầu tiên mỗi ngày. Bốn ô lớn dẫn thẳng tới bốn việc chính.",
    [
        "Chọn Từ ngày và Đến ngày theo ngày cần làm.",
        "Tạo phiếu đặt hàng: bấm Chọn file Excel để nạp file đơn mới hoặc file hoàn chỉnh cuối ngày.",
        "In đơn hàng: bấm Chọn đơn để in khi cần lấy phiếu giao.",
        "Bảng kê và biên nhận: bấm Chọn bảng kê để tải khi cần.",
        "Duyệt đơn: nếu ô viền đỏ, bấm Sửa các dòng đang thiếu rồi sửa hết lỗi trước.",
    ],
    note=(
        "File cuối cùng là file đang dùng",
        "Nếu cùng một ngày được nạp lại bằng file hợp lệ mới hơn, hệ thống dùng file mới nhất làm bản chuẩn. "
        "Vì vậy phải chọn đúng file khách đã chốt cuối cùng.",
    ),
    heading_level=1,
)


screen_page(
    "3.1. NHẬP VÀ SỬA ĐƠN",
    "02_nhap_va_sua_don.png",
    "Màn hình kiểm tra từng dòng đơn hàng.",
    "Dùng màn hình này để xem kỹ, sửa một dòng hoặc sửa nhanh cả bảng.",
    [
        "Nhìn ô nguồn để chắc chắn đang mở đúng file và đúng ngày.",
        "Muốn đổi file: bấm Chọn file đơn hàng khác.",
        "Muốn thêm nhiều dòng từ Excel: bấm Dán nhiều dòng.",
        "Muốn sửa nhiều ô: bấm Sửa nhanh cả bảng, sửa xong bấm Lưu một lần cho cả bảng.",
        "Dòng nào hệ thống báo thiếu nhà cung cấp, thiếu giá hoặc sai số liệu thì sửa dòng đó trước.",
        "Chỉ duyệt/chốt sau khi không còn lỗi đỏ.",
    ],
    note=(
        "Không sửa mù",
        "Nếu không biết giá, mã hàng hoặc nhà cung cấp, đối chiếu file gốc hoặc hỏi người phụ trách. "
        "Không điền tạm để cho hết màu đỏ.",
    ),
)


screen_page(
    "3.2. THÊM NHANH MỘT MẶT HÀNG",
    "03_them_nhanh_mat_hang.png",
    "Cửa sổ thêm một mặt hàng vào đúng bếp/ngày đang xem.",
    "Dùng khi đơn bị thiếu đúng một mặt hàng. Hệ thống tự giữ bếp và ngày đang chọn.",
    [
        "Tại đúng bếp cần thêm, bấm nút thêm hàng.",
        "Nhập Tên hàng.",
        "Nhập Số lượng.",
        "Nhập Đơn vị tính, ví dụ Kg, Cái, Túi.",
        "Bấm Thêm vào đơn này và kiểm tra dòng mới trong bảng.",
    ],
    note=(
        "Chỉ có ba ô cần nhập",
        "Không cần nhập lại bếp, ngày hoặc nhóm đang làm. Nếu mở nhầm bếp, bấm Hủy và mở lại ở đúng bếp.",
    ),
)


screen_page(
    "4. ĐẶT HÀNG NHÀ CUNG CẤP",
    "04_dat_hang_nha_cung_cap.png",
    "Mỗi nhà cung cấp là một thẻ gọn; chi tiết chỉ mở khi cần xem.",
    "Màn hình này dùng để gửi đơn cho từng nhà cung cấp và theo dõi đã đặt/chưa đặt.",
    [
        "Nhìn tên nhà cung cấp và tổng số dòng trên từng thẻ.",
        "Nếu cần xem hàng cụ thể, bấm mở chi tiết của thẻ.",
        "Bấm Sao chép ảnh để lấy ảnh gửi nhà cung cấp. Tạo ảnh thành công sẽ tự chuyển sang Đã đặt.",
        "Bấm Tải ảnh nếu muốn lưu ảnh thành file trước khi gửi.",
        "Khi cần gộp hoặc tách theo bếp, dùng nút trên đúng thẻ nhà cung cấp.",
        "Nếu đã đặt nhưng cần sửa/gửi lại, bấm Mở lại rồi thực hiện lại.",
    ],
    note=(
        "Ảnh gửi nhà cung cấp",
        "Ảnh chỉ có thông tin cần để đặt hàng; không hiện giá, tồn kho hoặc số liệu kiểm soát nội bộ.",
    ),
    heading_level=1,
)


screen_page(
    "5.1. PHIẾU GIAO – XEM TỔNG QUAN",
    "05_phieu_giao_tong_quan.png",
    "Màn hình chỉ giữ ba thông tin: ngày, tổng số phiếu và nút tải.",
    "Dùng khi cần biết trong ngày có bao nhiêu phiếu hoặc tải nhanh tất cả phiếu giao.",
    [
        "Chọn đúng đơn/ngày ở góc trên bên phải.",
        "Đọc tổng số phiếu giao trong ngày.",
        "Bấm Tải toàn bộ phiếu giao hàng để lấy tất cả phiếu của ngày.",
        "Bấm Xem chi tiết khi muốn kiểm tra riêng từng bếp.",
    ],
)


screen_page(
    "5.2. PHIẾU GIAO – XEM TỪNG BẾP",
    "06_phieu_giao_chi_tiet.png",
    "Chi tiết được đặt bên trong để màn hình ngoài vẫn gọn.",
    "Danh sách chi tiết giúp kiểm tra tên hàng và số lượng của từng bếp trước khi tải/in.",
    [
        "Mở Xem chi tiết từ màn hình tổng quan.",
        "Tìm đúng tên bếp cần kiểm tra.",
        "Đối chiếu tên hàng, số lượng và đơn vị tính.",
        "Dùng nút tải của bếp nếu chỉ cần một phiếu; dùng Tải toàn bộ nếu cần cả ngày.",
        "Mở file Excel vừa tải để kiểm tra rồi nhấn Ctrl+P khi cần in.",
    ],
    note=(
        "Mẫu in đã chốt",
        "Ngày nằm ở góc phải. Tên hàng và số lượng đặt gần nhau, chữ lớn để công nhân soạn hàng dễ đọc. "
        "Phần ký tên là chữ/ô Excel thật, không phải ảnh chèn.",
    ),
)


screen_page(
    "6.1. BÁO GIÁ",
    "07_bao_gia.png",
    "Hai lựa chọn rõ ràng: Báo giá tổng và Báo giá chi tiết theo nhà thầu.",
    "Dùng để nạp bảng giá tháng mới, tải toàn bộ hoặc lấy riêng cho một nhà thầu.",
    [
        "Chọn đúng Tháng báo giá.",
        "Khi có file giá mới, bấm Nạp báo giá tháng mới, xem trước rồi mới xác nhận.",
        "Báo giá tổng: tải bản mới nhất của tháng; mỗi nhà thầu là một file Excel riêng.",
        "Báo giá chi tiết: chọn đúng Nhà thầu rồi bấm Xem bảng chi tiết hoặc tải file khi nút đã sáng.",
        "Bấm Xem các lần báo giá đã lưu khi cần kiểm tra lịch sử.",
    ],
    note=(
        "Nếu nút đang mờ",
        "Tháng đó chưa có bảng giá đã xác nhận hoặc chưa đủ giá để xuất. Nạp/xác nhận bảng giá đúng tháng trước.",
    ),
)


screen_page(
    "6.2. BÁO CÁO TỔNG HỢP",
    "08_bao_cao_tong_hop.png",
    "Màn hình tải báo cáo theo tháng, đúng biểu mẫu khách hàng đã gửi.",
    "Dùng khi cần doanh số, giá vốn và lợi nhuận theo nhà thầu/bếp trong tháng.",
    [
        "Chọn Tháng cần xem.",
        "Đọc số ngày dữ liệu và số phiếu theo bếp đang có.",
        "Bấm Tải báo cáo tổng hợp.",
        "Mở file Excel và kiểm tra tháng, nhà thầu, bếp và dòng TỔNG THÁNG trước khi gửi.",
    ],
)


screen_page(
    "7.1. CÔNG NỢ – CHỌN ĐÚNG LOẠI",
    "09_cong_no_ba_lua_chon.png",
    "Màn hình ngoài chỉ có ba loại công nợ; chi tiết nằm bên trong.",
    "Chọn theo câu hỏi mình đang cần trả lời, không trộn tiền khách trả với tiền trả nhà cung cấp.",
    [
        "Công nợ phải thu (bếp): xem khách còn nợ chi tiết theo từng bếp và mặt hàng.",
        "Công nợ phải thu (tổng): xem tổng theo nhà thầu và ghi nhận tiền khách đã trả.",
        "Công nợ phải trả: xem tổng số lượng, tổng tiền, ghi nhận thanh toán và tải Excel đúng 14 cột đã chốt.",
    ],
    note=(
        "Nhớ cách phân biệt",
        "Phải thu = khách hàng trả tiền cho mình. Phải trả = mình trả tiền cho nhà cung cấp.",
    ),
    heading_level=1,
)


screen_page(
    "7.2. CÔNG NỢ PHẢI THU THEO BẾP",
    "10_cong_no_phai_thu_theo_bep.png",
    "Bộ lọc theo ngày, nhà thầu, bếp và trạng thái; chi tiết nằm ở bảng dưới.",
    "Dùng để biết một bếp phát sinh bao nhiêu, đã thu bao nhiêu và còn phải thu bao nhiêu.",
    [
        "Chọn Từ ngày – Đến ngày.",
        "Chọn Nhà thầu và Bếp; nếu cần xem hết thì để Tất cả.",
        "Bấm Xem công nợ hoặc Lọc danh sách.",
        "Đọc các ô tổng phía trên, sau đó xem từng dòng ở bảng chi tiết.",
        "Khi cần gửi đối chiếu, tải Excel của nhà thầu; nút ZIP sẽ tải riêng từng file cho mọi nhà thầu.",
    ],
)


screen_page(
    "7.3. CÔNG NỢ PHẢI THU TỔNG",
    "11_cong_no_phai_thu_tong.png",
    "Màn tổng theo nhà thầu, có phần ghi nhận tiền khách đã thanh toán.",
    "Dùng khi khách trả tiền hoặc cần xem số dư tổng, không cần mở từng mặt hàng.",
    [
        "Chọn khoảng ngày và đúng nhà thầu.",
        "Đọc Số dư đầu kỳ, Phát sinh, Đã thu và Còn phải thu.",
        "Khi nhận tiền, nhập đúng ngày, số tiền, hình thức và nội dung/mã tham chiếu.",
        "Bấm Ghi nhận đã thu một lần và kiểm tra giao dịch mới trong lịch sử.",
        "Nếu ghi nhầm, dùng Hoàn tác và nhập lý do; không xóa lịch sử.",
    ],
    note=(
        "Không bấm lặp",
        "Sau khi bấm ghi nhận, chờ thông báo thành công và kiểm tra lịch sử trước khi bấm lại.",
    ),
)


screen_page(
    "7.4. CÔNG NỢ PHẢI TRẢ NHÀ CUNG CẤP",
    "12_cong_no_phai_tra.png",
    "Màn hình chọn từng khoản nợ và phân bổ số tiền trả.",
    "Dùng khi thanh toán cho nhà cung cấp. Mỗi khoản trả phải chỉ rõ trả cho những dòng nào.",
    [
        "Chọn khoảng ngày, nhà cung cấp và trạng thái rồi bấm Xem công nợ.",
        "Tích Chọn ở đúng các dòng cần trả.",
        "Nhập Phân bổ lần này cho từng dòng; có thể trả một phần hoặc trả đủ.",
        "Kiểm tra Tổng phân bổ phải bằng đúng số tiền thanh toán.",
        "Nhập ngày, nội dung và mã tham chiếu rồi xác nhận một lần.",
        "Kiểm tra lịch sử; nếu sai, dùng Hoàn tác có lý do.",
    ],
    note=(
        "Hệ thống không tự chia tiền",
        "Người dùng phải tự chọn dòng và nhập số phân bổ. Hệ thống sẽ chặn nếu trả vượt số còn phải trả.",
    ),
)


screen_page(
    "8.1. BẢNG KÊ, BIÊN NHẬN VÀ HÓA ĐƠN",
    "13_bang_ke_va_hoa_don.png",
    "Màn tổng quan phần có thể lập hóa đơn, phần còn thiếu và các file cần tải.",
    "Dùng để kiểm tra theo nhà thầu trước khi làm file hóa đơn, bảng kê hoặc hồ sơ thanh toán.",
    [
        "Bấm Tính lại để lấy số mới nhất.",
        "Đọc năm ô tổng: Tổng cần lập, Đã dự thảo, Đã phát hành, Có thể lập và Còn chờ hóa đơn đầu vào.",
        "Màu xanh là phần có thể làm ngay; màu vàng là phần phải chờ thêm hóa đơn đầu vào.",
        "Mở Xem chi tiết từng mặt hàng chỉ khi cần kiểm tra mã hoặc số lượng.",
        "Ở File tải hóa đơn, chọn đúng nhà thầu/nhóm thuế rồi tải file 13 cột.",
        "Ở Bảng kê từ hóa đơn đỏ, chỉ chọn hóa đơn đã phát hành đúng kỳ trước khi tải.",
    ],
    note=(
        "Đề nghị thanh toán",
        "Chỉ lập từ hóa đơn đỏ đã phát hành của đúng một nhà thầu. Xem danh sách hóa đơn trước, "
        "sau đó mới tải Đề nghị thanh toán + bảng kê.",
    ),
    heading_level=1,
)


screen_page(
    "8.2. XỬ LÝ CHI TIẾT HÓA ĐƠN",
    "14_xu_ly_chi_tiet_hoa_don.png",
    "Vùng chọn mặt hàng thay thế và tải danh sách phần còn thiếu.",
    "Đây là thao tác ít dùng. Chỉ làm khi khách đã đồng ý cho thay mặt hàng và có người xác nhận.",
    [
        "Chọn đúng Dòng đang thiếu.",
        "Nhập mã TĐP thay thế, số lượng, người xác nhận và lý do.",
        "Nếu dùng giá riêng, nhập giá, lý do và tích xác nhận đã được người có quyền duyệt.",
        "Bấm Xem trước, chưa ghi; đọc kỹ mã gốc → mã thay thế và tồn còn lại.",
        "Chỉ xác nhận cuối khi thông tin đúng và đã có chấp thuận.",
        "Ở Danh sách còn thiếu theo kỳ, chọn ngày/nhà thầu để tải Excel phần chưa làm được.",
    ],
    note=(
        "Không tự thay mã",
        "Hệ thống không tự quyết định mặt hàng thay thế. Nếu chưa có xác nhận của khách/người có quyền, dừng tại bước xem trước.",
    ),
)


screen_page(
    "9.1. BÁO CÁO VẬT TƯ HÀNG HÓA",
    "15_bao_cao_vat_tu.png",
    "Màn ngoài chỉ có khoảng ngày và năm nút tải báo cáo.",
    "Dùng để tải tồn đầu, nhập, xuất và nhập–xuất–tồn theo đúng khoảng ngày.",
    [
        "Chọn Từ ngày – Đến ngày.",
        "Bấm Tải đủ 4 file ZIP khi cần cả bộ báo cáo.",
        "Nếu chỉ cần một file, bấm Tồn đầu kỳ, Nhập, Xuất hoặc Nhập – xuất – tồn.",
        "Khi tháng đã kết thúc, kiểm tra tổng tồn cuối rồi bấm Chốt tháng và chuyển sang tháng sau. Hệ thống thay tồn đầu tháng sau, không cộng chồng.",
        "Nếu cần làm lại tháng đã chốt, bấm Mở lại tháng; sửa xong phải chốt lại trước khi dùng số tháng sau.",
        "Bấm Xem chi tiết và nhập dữ liệu khi cần nạp tồn đầu hoặc xem từng mã hàng.",
    ],
    note=(
        "Bốn file đi cùng một kỳ",
        "Không đổi ngày giữa lúc tải từng file. Nếu cần gửi cả bộ, dùng nút ZIP để tránh lấy nhầm kỳ.",
    ),
    heading_level=1,
)


screen_page(
    "9.2. CHỐT THÁNG VÀ CHUYỂN TỒN SANG THÁNG SAU",
    "15b_chot_kho_theo_thang.png",
    "Khu vực chốt kho: tồn cuối tháng này được chuyển thành tồn đầu tháng sau.",
    "Dùng sau khi đã làm xong toàn bộ nhập, xuất và điều chỉnh của tháng. Không cần cộng hoặc nhập lại tồn đầu tháng sau bằng tay.",
    [
        "Chọn Từ ngày là ngày đầu tháng và Đến ngày là ngày cuối tháng cần chốt.",
        "Đối chiếu Số mặt hàng, Tổng lượng tồn cuối và Tổng giá trị tồn cuối.",
        "Nếu hệ thống báo còn vướng, quay lại sửa đúng dữ liệu được nêu; chưa bấm chốt.",
        "Khi số liệu đúng, bấm Chốt tháng và chuyển sang tháng sau, đọc lại thông báo rồi xác nhận một lần.",
        "Mở tháng sau và kiểm tra tồn đầu đã bằng đúng tồn cuối của tháng vừa chốt.",
        "Nếu buộc phải sửa lại tháng cũ, bấm Mở lại tháng. Sửa xong phải chốt lại trước khi tiếp tục dùng số tháng sau.",
    ],
    note=(
        "Không cộng chồng tồn",
        "Mỗi lần chốt lại, hệ thống thay đúng tồn đầu do lần chốt trước tạo ra; không cộng thêm lần thứ hai. "
        "Không tự nhập một bộ tồn đầu khác cho cùng tháng sau khi đã chốt.",
    ),
)


screen_page(
    "9.3. VẬT TƯ – XEM CHI TIẾT VÀ NHẬP DỮ LIỆU",
    "16_vat_tu_chi_tiet.png",
    "Chi tiết tồn đầu, nhập, xuất và tồn cuối theo từng mã hàng.",
    "Phần này dành cho người phụ trách kho/sổ vật tư; không cần mở trong công việc thường ngày.",
    [
        "Chọn đúng kỳ trước khi nạp tồn đầu.",
        "Có thể nhập một mã hoặc bấm Nạp Excel tồn đầu kỳ để nạp cả file.",
        "Luôn xem trước số mã, tổng giá trị, dòng âm và cảnh báo rồi mới xác nhận.",
        "Bảng phía dưới dùng để đối chiếu: Tồn đầu + Nhập − Xuất = Tồn cuối.",
        "Điều chỉnh nội bộ chỉ dùng khi có lý do rõ ràng; nhập ngày, mã, số tăng/giảm và lý do.",
    ],
    note=(
        "Không sửa trực tiếp để làm đẹp số",
        "Nếu cần sửa sai, dùng chức năng điều chỉnh có lý do để hệ thống giữ lịch sử. Không sửa/xóa dữ liệu cũ bên ngoài hệ thống.",
    ),
)


screen_page(
    "10.1. HÓA ĐƠN ĐẦU VÀO",
    "17_hoa_don_dau_vao.png",
    "Hóa đơn đầu vào được lấy từ mSMI theo khoảng ngày.",
    "Phần mềm chỉ tải dữ liệu về để kiểm tra; không ký, sửa hoặc phát hành hóa đơn.",
    [
        "Chọn tab Hóa đơn đầu vào và khoảng Từ ngày – Đến ngày.",
        "Bấm Tải/tiếp tục đầu vào. Nếu dữ liệu nhiều, bấm tiếp tục cho đến khi hoàn tất.",
        "Tại lần tải vừa tạo, bấm Xuất Excel nếu chỉ cần danh sách hóa đơn đầu vào.",
        "Muốn ghi kho: mở chi tiết, ghép đúng mã hàng TĐP và hệ số đổi đơn vị cho mọi dòng.",
        "Khi tất cả dòng đã Sẵn sàng, kiểm tra lần cuối rồi mới bấm Tạo phiếu nhập/ghi kho.",
    ],
    note=(
        "Xuất Excel không làm thay đổi kho",
        "Có thể tải Excel ngay sau khi kéo hóa đơn; chưa cần ghép mã. Chỉ nút xác nhận ghi kho mới làm tăng sổ vật tư.",
    ),
    heading_level=1,
)


screen_page(
    "10.2. HÓA ĐƠN ĐẦU RA",
    "18_hoa_don_dau_ra.png",
    "Hóa đơn đầu ra được lấy từ M-Invoice theo khoảng ngày.",
    "Chỉ hóa đơn đã phát hành hợp lệ mới được dùng để ghi giảm kho.",
    [
        "Chọn tab Hóa đơn đầu ra và khoảng Từ ngày – Đến ngày.",
        "Bấm Tải/tiếp tục đầu ra. Nếu hệ thống báo còn dữ liệu, bấm tiếp cho đến khi tải hết đúng khoảng ngày.",
        "Mở lần tải, ghép đúng mã hàng TĐP và hệ số đổi đơn vị cho mọi dòng.",
        "Kiểm tra trạng thái: hóa đơn nháp, hủy, thay thế hoặc điều chỉnh sẽ không được ghi kho như hóa đơn bình thường.",
        "Khi dòng đã Sẵn sàng, bấm Xác nhận xuất kho hóa đơn.",
        "Nếu hóa đơn đã ghi kho sau đó bị hủy/thay thế/điều chỉnh, dùng Hoàn tác xuất kho và ghi rõ lý do.",
    ],
    note=(
        "Phần mềm không phát hành hóa đơn",
        "Việc ký và phát hành vẫn thực hiện trên M-Invoice. Màn hình này chỉ tải về, đối chiếu mã và ghi sổ kho sau khi hóa đơn đã phát hành.",
    ),
)


screen_page(
    "11. SUẤT ĂN VÀ ĐẶT HÀNG BẾP",
    "19_xuong_com_va_po.png",
    "Màn hình nạp định mức/PO, chấm suất và ghép bếp với xưởng cơm.",
    "Cách làm thường ngày đã được ghi ngay trên dải xanh ở đầu màn hình.",
    [
        "Chọn đúng Ngày hoặc Thứ Hai đầu tuần.",
        "Bấm Nạp file định mức và đặt hàng, kiểm tra từng bếp, ca, số suất và nguyên liệu rồi xác nhận.",
        "Bấm Tải đơn đặt bếp nháp để xem; chỉ duyệt kế hoạch khi đủ bếp/XCOM và giá.",
        "Chấm suất thực tế theo tháng: chọn Tháng chấm suất → Nạp chấm suất tháng → xem trước → xác nhận.",
        "Ghép bếp vào XCOM chỉ làm khi thêm/sửa bếp; có thể nạp danh sách bếp → XCOM bằng Excel.",
        "Nhập tay kế hoạch bếp/ca chỉ dùng khi không có file nguồn.",
    ],
    note=(
        "Giá theo ngày/kỳ",
        "Giá nguyên liệu phải đúng ngày hoặc đúng kỳ đang làm. Nếu hệ thống báo thiếu giá, bổ sung đúng bảng giá trước khi duyệt.",
    ),
    heading_level=1,
)


screen_page(
    "12. CHẤM CÔNG VÀ TÍNH LƯƠNG",
    "20_cham_cong_va_luong.png",
    "Màn hình nhân sự, nạp chấm công, sửa một ngày và tải bảng lương.",
    "Dùng theo tháng. Nên nạp file chấm công trước; chỉ sửa tay khi có phát sinh riêng lẻ.",
    [
        "Chọn đúng tháng ở ô mm/yyyy.",
        "Bấm Nạp file chấm công, xem trước dữ liệu rồi xác nhận.",
        "Nhân sự mới: nhập mã, họ tên, chức vụ, bếp, lương cơ bản và thông số cần thiết rồi bấm Lưu nhân sự.",
        "Phát sinh một ngày: nhập mã nhân sự, ngày, giờ thường/tăng ca/Chủ nhật/ca đêm/ngày lễ rồi bấm Lưu chấm công.",
        "Nhập các khoản phụ cấp, bảo hiểm, tạm ứng hoặc điều chỉnh ở phần phía dưới của màn hình.",
        "Kiểm tra tổng theo người và theo bếp, sau đó bấm Tải bảng lương.",
    ],
    note=(
        "Không nạp lặp khi chưa kiểm tra",
        "Nếu vừa nạp xong, kiểm tra tháng và số nhân sự trước. Chỉ nạp lại khi có file đã sửa/chốt mới.",
    ),
    heading_level=1,
)


screen_page(
    "13. IN GIẤY TỜ",
    "21_in_giay_to.png",
    "Chọn loại giấy tờ, chọn ngày/bếp rồi tải file để xem trước và in.",
    "Đây là cách in chính, dễ kiểm tra nhất. Không cần tìm giấy tờ ở nhiều menu khác nhau.",
    [
        "Chọn một loại: Đơn hàng đi giao, Đơn đặt nhà cung cấp, Bảng kê và biên nhận hoặc Báo cáo tổng hợp.",
        "Chọn Từ ngày – Đến ngày.",
        "Với phiếu giao, bấm Chọn tất cả hoặc bỏ tích những bếp không cần in.",
        "Nếu còn nhãn đỏ “Còn … dòng lỗi”, quay lại Nhập & sửa đơn; chưa tải/in.",
        "Bấm Tải file đã chọn.",
        "Mở file vừa tải, kiểm tra trang rồi nhấn Ctrl+P để in.",
    ],
    note=(
        "Khổ giấy",
        "Phiếu giao dùng A4. Các giấy tờ khác chọn A5 hoặc A4 theo nhu cầu. Luôn kiểm tra giấy ra thực tế trước khi in nhiều bản.",
    ),
    heading_level=1,
)


screen_page(
    "14.1. DANH MỤC VÀ CẤU HÌNH",
    "22_danh_muc_va_sao_luu.png",
    "Nơi quản lý mã hàng, bếp, nhà cung cấp, nhóm nhà thầu và thông tin dùng trên giấy tờ.",
    "Phần lớn nội dung đã được cài sẵn. Chỉ thay đổi khi có danh mục đã được khách hàng/người phụ trách chốt.",
    [
        "Bấm Kiểm tra ngay khi cần thử kết nối M-Invoice.",
        "Đồng bộ lại từ Em Thành.xlsx khi có file chuẩn mới đã được xác nhận.",
        "Nạp danh mục khách chốt khi có file danh mục chính thức mới; xem trước rồi mới xác nhận.",
        "Thông tin thanh toán mặc định phải đủ người đề nghị, số tài khoản và ngân hàng trước khi lập Đề nghị thanh toán.",
        "Tên xuất hóa đơn chỉ sửa/nạp khi có danh sách đã được xác nhận.",
    ],
    note=(
        "Không tự sửa hàng loạt",
        "Nếu cùng một mã có hai thông tin khác nhau, hệ thống sẽ chặn. Sửa file nguồn cho thống nhất rồi nạp lại.",
    ),
    heading_level=1,
)


screen_page(
    "14.2. SAO LƯU DỮ LIỆU",
    "23_sao_luu_du_lieu.png",
    "Ô Sao lưu dữ liệu nằm trong menu Danh mục & sao lưu.",
    "Nên sao lưu cuối mỗi ngày làm việc và luôn sao lưu trước khi chuyển máy hoặc cập nhật phần mềm.",
    [
        "Mở Danh mục & sao lưu.",
        "Tìm ô Sao lưu dữ liệu.",
        "Bấm Tải bản sao lưu dữ liệu.",
        "Chờ tải xong và kiểm tra có file mới trong thư mục Tải xuống.",
        "Lưu file vào nơi nội bộ an toàn; không gửi qua nhóm chat công cộng.",
    ],
    note=(
        "Đây là dữ liệu quan trọng",
        "Bản sao lưu có đơn hàng, công nợ và lịch sử vận hành. Chỉ người có trách nhiệm được giữ và sử dụng.",
    ),
)


# Troubleshooting
page_break()
doc.add_heading("15. KHI GẶP SỰ CỐ", level=1)
simple_table(
    ["Hiện tượng", "Làm lần lượt"],
    [
        ("Nhấp EXE nhưng chưa thấy trang", "Chờ 30–60 giây → kiểm tra có cửa sổ màu đen → không nhấp lại nhiều lần."),
        ("Trang bị đứng hoặc chưa đổi", "Nhấn F5 một lần → chờ tải xong → chọn lại đúng ngày/đơn ở góc trên bên phải."),
        ("Có dòng màu đỏ", "Đọc nội dung → vào Nhập & sửa đơn → sửa đúng dòng → lưu → quay lại kiểm tra."),
        ("Nút tải/xác nhận bị mờ", "Phần trước chưa đủ điều kiện. Đọc dòng giải thích cạnh nút và hoàn thành dữ liệu còn thiếu."),
        ("Tải file nhưng không thấy", "Mở thư mục Downloads/Tải xuống → sắp xếp theo thời gian mới nhất → mở đúng file vừa tải."),
        ("Bấm in nhưng chưa ra giấy", "Không bấm lặp → kiểm tra màn hình máy in, khay giấy và hàng đợi in → thử một bản trước."),
        ("Sai ngày hoặc sai đơn", "Chọn lại ô đơn ở góc trên bên phải. Không nạp lại file chỉ để xem dữ liệu cũ."),
        ("Không biết có nên xác nhận", "Dừng lại, chụp toàn màn hình và hỏi người phụ trách. Không điền số tạm."),
    ],
    [5.5, 11.6],
)
callout(
    "Khi nhờ hỗ trợ",
    "Gửi ảnh chụp toàn màn hình, tên menu đang mở, ngày/đơn đang chọn và nút vừa bấm. "
    "Không gửi mật khẩu, file cấu hình hoặc bản sao lưu dữ liệu.",
    fill=PALE_YELLOW,
    accent=AMBER,
)


# Final checklist
page_break()
doc.add_heading("16. CHECKLIST CUỐI NGÀY", level=1)
paragraph("Có thể in riêng trang này và để cạnh máy làm việc.", size=11.5,
          color=MUTED, italic=True)
doc.add_heading("Đơn hàng", level=2)
bullets([
    "Đang chọn đúng ngày và đúng file cuối cùng khách đã chốt.",
    "Không còn dòng lỗi đỏ.",
    "Tên hàng, số lượng, bếp, nhà thầu, nhà cung cấp và giá đã được kiểm tra.",
    "Các nhà cung cấp cần đặt đã chuyển sang trạng thái Đã đặt.",
], checkbox=True)
doc.add_heading("Giấy tờ", level=2)
bullets([
    "Đã kiểm tra tổng số phiếu giao trong ngày.",
    "Các file Excel/PDF đã mở kiểm tra trước khi gửi hoặc in.",
    "Ngày, tên đơn vị, số lượng, đơn giá và tổng tiền trên file đúng.",
    "Máy in đã ra đủ giấy, đúng khổ và không còn lệnh in bị treo.",
], checkbox=True)
doc.add_heading("Tiền và dữ liệu", level=2)
bullets([
    "Khoản đã thu/đã trả trong ngày đã xuất hiện đúng trong lịch sử.",
    "Hóa đơn nháp và hóa đơn đã phát hành được phân biệt đúng.",
    "Nếu là cuối tháng, đã đối chiếu tồn cuối và chốt chuyển sang tồn đầu tháng sau.",
    "Đã tải bản sao lưu dữ liệu.",
    "Đã đóng hệ thống đúng cách sau khi lưu xong.",
], checkbox=True)
callout(
    "Ba điều cuối cùng",
    "Đúng ngày – Đúng số liệu – Xem trước rồi mới xác nhận. Nếu chưa chắc, dừng lại và hỏi; không bấm thử.",
)


# Back cover
page_break()
paragraph("THÀNH ĐẠT PHÁT", size=16, bold=True, color=TEAL_DARK,
          align=WD_ALIGN_PARAGRAPH.CENTER)
paragraph("HẾT TÀI LIỆU HƯỚNG DẪN", size=24, bold=True, color=NAVY,
          align=WD_ALIGN_PARAGRAPH.CENTER)
paragraph("Tài liệu này dùng cùng phiên bản phần mềm đã bàn giao. Khi giao diện được cập nhật, "
          "cần cập nhật lại ảnh và hướng dẫn để tránh làm theo bản cũ.",
          size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
paragraph(f"Phiên bản {VERSION}  •  {ISSUE_DATE}", size=10.5, color=MUTED,
          align=WD_ALIGN_PARAGRAPH.CENTER)


props = doc.core_properties
props.title = "Hướng dẫn sử dụng hệ thống Thành Đạt Phát"
props.subject = "Hướng dẫn có ảnh, dành cho người mới không biết kỹ thuật"
props.author = "Thành Đạt Phát"
props.keywords = "Thành Đạt Phát, hướng dẫn sử dụng, vận hành, đơn hàng, công nợ"
props.comments = f"Chụp từ bản EXE {VERSION}; cập nhật {ISSUE_DATE}"

OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)

with ZipFile(OUTPUT) as archive:
    media = sorted(name for name in archive.namelist() if name.startswith("word/media/"))
if len(media) != 24:
    raise RuntimeError(f"Expected 24 embedded screenshots, found {len(media)}")

print(f"output={OUTPUT}")
print(f"embedded_screenshots={len(media)}")
print(f"bytes={OUTPUT.stat().st_size}")
