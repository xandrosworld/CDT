import io
import json
import unittest
from datetime import date

from openpyxl import load_workbook

from . import server
from .test_outgoing_readiness import OutgoingReadinessTests as Seed
from .output_stock_remap import export_workbook, preview_workbook, confirm_preview, RemapError
from .invoice_monthly_valuation import monthly_average_report
from .invoice_inventory import reverse_output_invoice
from .inventory_period_close import inventory_period_close_preview, close_inventory_period
from .outgoing_readiness import canonical_available_stock


class OutputStockRemapTests(unittest.TestCase):
    setUpClass = classmethod(Seed.setUpClass.__func__)
    tearDownClass = classmethod(Seed.tearDownClass.__func__)

    def setUp(self):
        with server.db() as conn:
            conn.execute('DELETE FROM output_stock_remaps')
            conn.execute('DELETE FROM output_stock_excel_sessions')
            conn.execute('DELETE FROM inventory_period_closures')
            conn.execute('DELETE FROM invoice_inventory_ledger')
            conn.execute('DELETE FROM invoice_mapping_revisions')
            conn.execute('DELETE FROM invoice_line_mappings')
            conn.execute('DELETE FROM bk_import_lines')
            conn.execute('DELETE FROM bk_import_documents')
        Seed.setUp(self)
        with server.db() as conn:
            conn.execute("INSERT OR REPLACE INTO products(code,name,unit,tax) VALUES('REMAP-B','Hàng nhận','kg','8%')")
            conn.execute("INSERT OR REPLACE INTO products(code,name,unit,tax) VALUES('REMAP-C','Hoa cúng','kg','KKKNT')")
            Seed.add_opening(conn, 0)
            Seed.add_opening(conn, 10, product_code='REMAP-B')
            Seed.add_opening(conn, -3, product_code='REMAP-C')
            conn.execute("UPDATE inventory_transactions SET unit_cost=30 WHERE product_code='REMAP-B'")
            conn.execute("UPDATE inventory_transactions SET source_id='2026-08'")
            self.source_id = Seed.add_posted_source(conn, source='minvoice', number='REMAP-1')
            conn.execute('UPDATE outgoing_source_invoices SET raw_json=?,subtotal=80,tax_amount=6,total_amount=86 WHERE id=?',
                         (json.dumps({'_tdp_source_contract':'minvoice_portal_v1','details':[{'taxAmount':6}]}),self.source_id))
            self.line_id = conn.execute("""INSERT INTO outgoing_source_invoice_items(invoice_id,line_index,source_item_code,
                source_item_name,source_unit,qty,unit_price,amount,tax_rate,product_code,mapping_status,stock_qty)
                VALUES(?,1,'SOURCE-A','Tên trên hóa đơn','kg',4,20,80,'8','HH-01','mapped',4)""",(self.source_id,)).lastrowid
            conn.execute('UPDATE invoice_inventory_ledger SET source_line_id=? WHERE source_invoice_id=?',(self.line_id,self.source_id))
            self.file = export_workbook(conn, '2026-08-01', '2026-08-31')

    def edited(self, code='REMAP-B', name='Hàng nhận', mutate=None):
        wb = load_workbook(io.BytesIO(self.file)); ws=wb['Doi ma xuat kho']
        ws['Q2']=code; ws['R2']=name
        if mutate: mutate(ws)
        data=io.BytesIO(); wb.save(data); wb.close(); return data.getvalue()

    def test_round_trip_preserves_source_and_ledger_and_closes_with_negative_kkknt(self):
        with server.db() as conn:
            invoice = dict(conn.execute('SELECT * FROM outgoing_source_invoices WHERE id=?',(self.source_id,)).fetchone())
            line = dict(conn.execute('SELECT * FROM outgoing_source_invoice_items WHERE id=?',(self.line_id,)).fetchone())
            ledger = [dict(r) for r in conn.execute('SELECT * FROM invoice_inventory_ledger')]
            preview = preview_workbook(conn, self.edited())
            self.assertTrue(preview['can_confirm'],preview)
            self.assertEqual(0,conn.execute('SELECT COUNT(*) FROM output_stock_remaps').fetchone()[0])
            self.assertEqual(6,preview['changes'][0]['new_closing_after'])
            result = confirm_preview(conn,preview['token'],'Người kiểm thử',server.now_iso())
            self.assertEqual(1,result['changed_lines'])
            self.assertTrue(confirm_preview(conn,preview['token'],'Người kiểm thử',server.now_iso())['idempotent'])
            self.assertEqual(invoice,dict(conn.execute('SELECT * FROM outgoing_source_invoices WHERE id=?',(self.source_id,)).fetchone()))
            after=dict(conn.execute('SELECT * FROM outgoing_source_invoice_items WHERE id=?',(self.line_id,)).fetchone())
            self.assertEqual({**line,'product_code':'REMAP-B'},after)
            self.assertEqual(ledger,[dict(r) for r in conn.execute('SELECT * FROM invoice_inventory_ledger')])
            stock=canonical_available_stock(conn)
            self.assertEqual((0,6,-3),tuple(stock[c]['canonical_qty'] for c in ['HH-01','REMAP-B','REMAP-C']))
            close=inventory_period_close_preview(conn,'2026-08',today=date(2026,9,10))
            self.assertTrue(close['can_close'],close['issues'])
            self.assertEqual(1,close['kkknt_negative_count'])
            close_inventory_period(conn,'2026-08',expected_source_hash=close['source_hash'],expected_target_hash=close['target_hash'],timestamp=server.now_iso(),today=date(2026,9,10))
            report=monthly_average_report(conn,date_from='2026-09-01',date_to='2026-09-30',include_zero=True)
            flower=next(r for r in report['items'] if r['product_code']=='REMAP-C')
            self.assertEqual((-3,-30), (flower['opening_qty'],flower['opening_value']))

    def test_only_new_identity_columns_are_editable(self):
        for column in ['A','B','C','D','E','F','G','H','I','J','K','L','M','N','O','P']:
            with self.subTest(column=column),server.db() as conn:
                with self.assertRaises(RemapError): preview_workbook(conn,self.edited(mutate=lambda ws: ws.__setitem__(column+'2','Changed')))
                self.assertEqual(0,conn.execute('SELECT COUNT(*) FROM output_stock_remaps').fetchone()[0])

    def test_bad_code_name_unit_formula_and_missing_row(self):
        with server.db() as conn:
            self.assertFalse(preview_workbook(conn,self.edited(code='NO-SUCH-CODE'))['can_confirm'])
            self.assertFalse(preview_workbook(conn,self.edited(name='Tên tự bịa'))['can_confirm'])
            with self.assertRaises(RemapError): preview_workbook(conn,self.edited(code='=1+1'))
            with self.assertRaises(RemapError): preview_workbook(conn,self.edited(mutate=lambda ws:ws.delete_rows(2)))
            conn.execute("UPDATE products SET unit='chai' WHERE code='REMAP-B'")
            self.file=export_workbook(conn,'2026-08-01','2026-08-31')
            p=preview_workbook(conn,self.edited()); self.assertFalse(p['can_confirm']); self.assertIn('đơn vị',p['errors'][0])

    def test_target_shortage_stale_file_and_stale_preview_are_rejected(self):
        with server.db() as conn:
            p=preview_workbook(conn,self.edited()); self.assertTrue(p['can_confirm'])
            conn.execute("UPDATE inventory_transactions SET qty_in=1 WHERE product_code='REMAP-B'")
            with self.assertRaises(RemapError): confirm_preview(conn,p['token'],'Test',server.now_iso())
            with self.assertRaises(RemapError): preview_workbook(conn,self.edited())
            self.file=export_workbook(conn,'2026-08-01','2026-08-31')
            p=preview_workbook(conn,self.edited()); self.assertFalse(p['can_confirm']); self.assertIn('sẽ âm',p['errors'][0])

    def test_reversal_returns_stock_to_remapped_product(self):
        with server.db() as conn:
            mapping = conn.execute("""INSERT INTO invoice_line_mappings(tenant,source,invoice_type,scope_key,product_code,mapping_status,conversion_factor,confirmed_at,updated_at)
                VALUES('test','minvoice','OUTPUT_ELECTRONIC_INVOICE',?,'HH-01','confirmed',1,'test','test')""",(str(self.line_id),)).lastrowid
            revision = conn.execute("""INSERT INTO invoice_mapping_revisions(revision_key,mapping_id,product_code,source_unit,target_unit,conversion_factor,created_at)
                VALUES(?,?,'HH-01','kg','kg',1,'test')""",(str(self.line_id),mapping)).lastrowid
            conn.execute('UPDATE invoice_inventory_ledger SET mapping_revision_id=? WHERE source_line_id=?',(revision,self.line_id))
            self.file=export_workbook(conn,'2026-08-01','2026-08-31')
            p=preview_workbook(conn,self.edited()); confirm_preview(conn,p['token'],'Test',server.now_iso())
            self.assertEqual(30,conn.execute("SELECT unit_cost FROM invoice_inventory_effective_ledger WHERE source_line_id=?",(self.line_id,)).fetchone()[0])
            close=inventory_period_close_preview(conn,'2026-08',today=date(2026,9,10))
            close_inventory_period(conn,'2026-08',expected_source_hash=close['source_hash'],expected_target_hash=close['target_hash'],timestamp=server.now_iso(),today=date(2026,9,10))
            conn.execute("UPDATE outgoing_source_invoices SET stock_status='reversal_required',source_status_class='cancelled' WHERE id=?",(self.source_id,))
            reverse_output_invoice(conn,self.source_id,confirmed=True,note='fixture',now_iso=lambda:'2026-09-25T12:00:00')
            stocks=canonical_available_stock(conn)
            self.assertEqual((0,10),tuple(stocks[c]['canonical_qty'] for c in ['HH-01','REMAP-B']))
            report=monthly_average_report(conn,date_from='2026-09-01',date_to='2026-09-30',include_zero=True)
            target=next(r for r in report['items'] if r['product_code']=='REMAP-B')
            self.assertEqual((10,300),(target['closing_qty'],target['closing_value']))

    def test_remapping_does_not_restore_local_issued_hold(self):
        with server.db() as conn:
            batch,orders=Seed.add_batch(conn,'2026-08-20',[{'qty':4}])
            Seed.add_local_issued_draft(conn,batch,orders[0],number='REMAP-1')
            self.file=export_workbook(conn,'2026-08-01','2026-08-31')
            p=preview_workbook(conn,self.edited()); confirm_preview(conn,p['token'],'Test',server.now_iso())
            self.assertEqual(0,canonical_available_stock(conn)['HH-01']['pending_sync_issued_qty'])

    def test_write_failure_rolls_back_every_change(self):
        with server.db() as conn:
            p=preview_workbook(conn,self.edited())
            conn.execute("CREATE TEMP TRIGGER fail_remap BEFORE UPDATE ON outgoing_source_invoice_items BEGIN SELECT RAISE(ABORT,'fixture'); END")
            with self.assertRaises(Exception): confirm_preview(conn,p['token'],'Test',server.now_iso())
            self.assertEqual(0,conn.execute('SELECT COUNT(*) FROM output_stock_remaps').fetchone()[0])
            self.assertEqual('HH-01',conn.execute('SELECT product_code FROM outgoing_source_invoice_items WHERE id=?',(self.line_id,)).fetchone()[0])

    def test_negative_kkknt_can_allocate_export_issue_and_recalculate(self):
        with server.db() as conn:
            batch,orders=Seed.add_batch(conn,'2026-08-20',[{'product_code':'REMAP-C','qty':100}])
            conn.execute("UPDATE orders SET tax='KKKNT' WHERE batch_id=?",(batch,))
        ready=self.client.get(f'/api/outgoing-invoices/readiness/{batch}').get_json()
        self.assertEqual([],ready['blocking_issues']); self.assertEqual(100,ready['invoiceable_qty'])
        self.assertEqual(1,len(ready['negative_stock_warnings']))
        for _ in range(2):
            created=self.client.post(f'/api/outgoing-invoices/draft/{batch}')
            self.assertEqual(200,created.status_code,created.get_json())
            self.assertEqual(0,created.get_json()['pending_qty'])
        exported=self.client.get(f'/api/export/invoices/{batch}')
        self.assertEqual(200,exported.status_code,exported.get_json(silent=True))
        draft=created.get_json()['drafts'][0]['id']
        with server.db() as conn:
            conn.execute("UPDATE outgoing_invoice_drafts SET buyer_name_snapshot='Buyer',buyer_tax_code_snapshot='0200000001',"
                         "buyer_address_snapshot='Address',company_name_snapshot='TDP',company_tax_code_snapshot='0100000001',"
                         "company_address_snapshot='Address',payment_requester_snapshot='Requester',"
                         "payment_bank_name_snapshot='Bank',payment_bank_account_snapshot='Account' WHERE batch_id=?", (batch,))
        issued=self.client.post(f'/api/outgoing-invoices/{draft}/confirm-issued',json={
            'confirmed':True,'invoice_number':'999001','invoice_series':'1C26TDP','invoice_date':'2026-08-20'})
        self.assertEqual(200,issued.status_code,issued.get_json())
        ready=self.client.get(f'/api/outgoing-invoices/readiness/{batch}').get_json()
        self.assertEqual((100,0),(ready['issued_qty'],ready['invoiceable_qty']))

    def test_zero_kct_unknown_and_mixed_tax_do_not_get_negative_exception(self):
        for tax in ['0%','KCT','INVALID','8%']:
            with self.subTest(tax=tax),server.db() as conn:
                conn.execute('UPDATE products SET tax=? WHERE code=?',(tax,'REMAP-C'))
                batch,_=Seed.add_batch(conn,'2026-08-20',[{'product_code':'REMAP-C','qty':1}])
                conn.execute('UPDATE orders SET tax=? WHERE batch_id=?',(tax,batch))
            ready=self.client.get(f'/api/outgoing-invoices/readiness/{batch}').get_json()
            self.assertEqual(1,len(ready['blocking_issues']))
            self.assertEqual(0,ready['invoiceable_qty'])
            self.assertNotEqual(200,self.client.post(f'/api/outgoing-invoices/draft/{batch}').status_code)
        with server.db() as conn:
            conn.execute("UPDATE products SET tax='KKKNT' WHERE code='REMAP-C'")
            conn.execute("UPDATE orders SET tax='0%' WHERE batch_id=?",(batch,))
        self.assertEqual(1,len(self.client.get(f'/api/outgoing-invoices/readiness/{batch}').get_json()['blocking_issues']))

    def test_kkknt_supplementary_purchase_register_does_not_require_nonnegative_stock(self):
        with server.db() as conn:
            conn.execute("INSERT OR REPLACE INTO suppliers(code,name) VALUES('BK-S1','Nhà cung cấp')")
            conn.execute("UPDATE products SET purchase_list=1,supplier='BK-S1',buy_price=10 WHERE code='REMAP-C'")
        template=self.client.get('/api/bk-import/template')
        wb=load_workbook(io.BytesIO(template.data)); ws=wb['BK_IMPORT']
        ws.append(['2026-08-31','','BK-FLOWER',1,'REMAP-C','Hoa cúng','kg',1,10,10,'BK-S1','Bổ sung mua vào'])
        data=io.BytesIO(); wb.save(data); wb.close()
        response=self.client.post('/api/bk-import/preview',data={'file':(io.BytesIO(data.getvalue()),'bk.xlsx')})
        self.assertEqual(200,response.status_code,response.get_json())
        preview=response.get_json()
        posted=self.client.post('/api/bk-import/confirm',json={'confirmed':True,'token':preview['token'],'previewId':preview['previewId']})
        self.assertEqual(200,posted.status_code,(posted.get_json(),preview))
        with server.db() as conn:
            self.assertEqual(-2,canonical_available_stock(conn)['REMAP-C']['canonical_qty'])
            report=monthly_average_report(conn,date_from='2026-08-01',date_to='2026-08-31',include_zero=True)
            self.assertEqual(-2,next(r for r in report['items'] if r['product_code']=='REMAP-C')['closing_qty'])


if __name__ == '__main__': unittest.main()
