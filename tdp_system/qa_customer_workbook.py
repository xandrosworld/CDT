"""Verify the supplied workbook on a fresh SQLite backup, never the source DB."""
import argparse
from collections import Counter
from contextlib import closing
import hashlib
import json
import os
from pathlib import Path
import sqlite3


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--database', type=Path, required=True)
    parser.add_argument('--workbook', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    output = args.output.resolve(); output.mkdir(parents=True, exist_ok=False)
    source = args.database.resolve(); target = output / 'copy.sqlite3'
    before = hashlib.sha256(source.read_bytes()).hexdigest()
    with closing(sqlite3.connect(source.as_uri() + '?mode=ro', uri=True)) as read, closing(sqlite3.connect(target)) as write:
        read.backup(write)
    os.environ.update(TDP_DB_PATH=str(target), TDP_DATA_DIR=str(output), TDP_EXPORT_DIR=str(output / 'exports'),
                      MSMI_API_BASE_URL='http://127.0.0.1:9', MSMI_API_TOKEN='offline',
                      MINVOICE_API_BASE_URL='http://127.0.0.1:9', MINVOICE_USERNAME='offline', MINVOICE_PASSWORD='offline')
    from . import server
    from .daily_workbook_import import analyze_daily_workbook
    from .contract_modules import parse_purchase_order_workbook, purchase_order_payload
    from openpyxl import load_workbook
    server.init_database(sync_master=False)
    analysis = analyze_daily_workbook(args.workbook)
    assert len(analysis['purchaseSheets']) == 1
    orders, skipped = server.parse_workbook(args.workbook, '2026-09-03', ['03.09'])
    chosen = next(r for r in orders if r['source_row'] == 342)
    assert chosen['supplier'] == 'kho' and chosen['qty'] == .54
    with server.db() as conn:
        batch = conn.execute("INSERT INTO batches(work_date,source_name,created_at) VALUES('2026-09-03','acceptance-copy','now')").lastrowid
        server.save_imported_orders(conn, batch, orders)
        persisted = dict(conn.execute('SELECT * FROM orders WHERE batch_id=? AND source_row=342', (batch,)).fetchone())
        assert persisted['supplier'] == 'kho' and persisted['qty'] == .54
        plan = purchase_order_payload(conn, batch)
        plan_row = next(r for r in plan['rows'] if r['order_id'] == persisted['id'])
        assert plan_row['supplier'] == 'kho'
        values = load_workbook(args.workbook, data_only=True)
        formulas = load_workbook(args.workbook, data_only=False)
        purchase = parse_purchase_order_workbook(conn, values, batch, formulas)
        purchase_row = next(r for r in purchase['items'] if r['source_row'] == 103)
        assert purchase_row['supplier'] == 'kho' and purchase_row['actual_qty'] == .54 and purchase_row['amount'] == 14040
        result = {
            'orders': len(orders), 'day_sheet': analysis['daySheets'], 'purchase_sheets': analysis['purchaseSheets'],
            'order_342': {k: persisted[k] for k in ('supplier', 'qty', 'product_code', 'kitchen')},
            'purchase_103': {k: purchase_row[k] for k in ('supplier', 'actual_qty', 'amount', 'errors')},
            'purchase_count': purchase['count'], 'purchase_error_rows': purchase['error_rows'],
            'purchase_errors': dict(Counter(error for r in purchase['items'] for error in r['errors'])),
            'order_error_rows': sum(bool(r['errors']) for r in orders),
            'order_errors': dict(Counter(error for r in orders for error in r['errors'])),
            'source_db_unchanged': hashlib.sha256(source.read_bytes()).hexdigest() == before,
        }
        values.close(); formulas.close()
    (output / 'result.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(result, ensure_ascii=True))


if __name__ == '__main__': main()
