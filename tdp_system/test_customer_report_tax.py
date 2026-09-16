import io
import json
import unittest
import zipfile
from contextlib import contextmanager
from unittest.mock import patch

from flask import Flask
from openpyxl import load_workbook

from .invoice_line_tax import annotate_invoice_tax, grouped_tax, tax_fields
from .inventory_export import register_inventory_export_routes
from . import test_inventory_export as inventory_fixtures
from . import test_invoice_output_register as sales_fixtures
from .invoice_output_register import output_sales_workbook


class LineTaxTests(unittest.TestCase):
    def test_msmi_fractional_rate_matches_pdf_without_changing_header_or_source(self):
        invoice = dict(tax_amount=369334, total_amount=4986000, items=[
            dict(line_index=1, amount=983333, tax_rate='0.08'),
            dict(line_index=2, amount=3633333, tax_rate='0.08'),
        ])
        raw = json.dumps({'hdhhdvu': [{'tsuat':0.08,'tthue':None}, {'tsuat':0.08,'tthue':None}]})
        annotate_invoice_tax(invoice, raw)
        self.assertEqual(['8%', '8%'], [r['tax_rate'] for r in invoice['items']])
        self.assertEqual([78667, 290667], [r['line_tax_amount'] for r in invoice['items']])
        self.assertEqual(['0.08', '0.08'], [r['source_tax_rate'] for r in invoice['items']])
        self.assertEqual(0, invoice['detail_tax_difference'])
        self.assertEqual(4986000, invoice['total_amount'])
        self.assertEqual(123, tax_fields(dict(amount=983333,tax_rate='0.08'), {'tsuat':0.08,'tthue':123})['line_tax_amount'])
        for raw_line in ({}, {'taxRate':0.08}, {'tsuat':'0.08%'}):
            self.assertEqual(80, tax_fields(dict(amount=100000,tax_rate='0.08'),raw_line)['line_tax_amount'])
        self.assertEqual(80, tax_fields(dict(amount=100000,tax_rate='0.08%'), {'tsuat':0.08})['line_tax_amount'])

    def test_minvoice_category_codes_display_as_labels_without_negative_tax(self):
        for code, label in [('-2', 'KKKNT'), (-2, 'KKKNT'), ('-2.0', 'KKKNT'), ('-1', 'KCT'), ('-1%', 'KCT')]:
            line = dict(amount=100000, tax_rate=code)
            self.assertEqual((label, 0, 100000), tuple(tax_fields(line)[k] for k in ('tax_rate', 'line_tax_amount', 'amount_with_tax')))
            self.assertEqual(code, line['tax_rate'])
            self.assertEqual(123, tax_fields(line, {'vatAmount': 123})['line_tax_amount'])
        self.assertEqual(('0', 0), tuple(tax_fields(dict(amount=100, tax_rate=0))[k] for k in ('tax_rate', 'line_tax_amount')))

    def test_explicit_zero_and_signed_source_tax_override_rate_arithmetic(self):
        for tax in (0, -8001.25, 8123):
            fields = tax_fields(dict(amount=100000, tax_rate='8'), {'vatAmount': tax})
            self.assertEqual(tax, fields['line_tax_amount'])
            self.assertEqual(100000 + tax, fields['amount_with_tax'])
            self.assertEqual('Thuế theo dòng hóa đơn', fields['tax_note'])

    def test_minvoice_api_line_tax_alias_keeps_exact_zero_decimal_and_signed_tax(self):
        for tax in (0, 1.6, -1.6):
            fields = tax_fields(dict(amount=20,tax_rate='8'),{'inv_vatAmount':tax})
            self.assertEqual(tax,fields['line_tax_amount'])
            self.assertEqual(20+tax,fields['amount_with_tax'])
            self.assertEqual('Thuế theo dòng hóa đơn',fields['tax_note'])
        self.assertIsNone(tax_fields(dict(amount=20,tax_rate='8'),{'inv_vatAmount':'bad'})['line_tax_amount'])

    def test_missing_or_invalid_tax_is_not_silently_zero(self):
        for raw in ({}, {'tthue': 'invalid'}, {'tthue': float('nan')}):
            fields = tax_fields(dict(amount=100, tax_rate='KHAC'), raw)
            self.assertIsNone(fields['line_tax_amount'])
            self.assertIsNone(fields['amount_with_tax'])
        # A malformed explicit tax cannot be masked by a plausible tax rate.
        self.assertIsNone(tax_fields(dict(amount=100, tax_rate='8'), {'tthue': 'bad'})['line_tax_amount'])

    def test_computed_tax_is_labelled_and_groups_sum_member_tax(self):
        first = tax_fields(dict(amount=101, tax_rate='8%'))
        second = tax_fields(dict(amount=100, tax_rate='0'))
        self.assertEqual(8, first['line_tax_amount'])
        self.assertIn('Tính từ', first['tax_note'])
        merged = grouped_tax([first, second])
        self.assertEqual(8, merged['line_tax_amount'])
        self.assertEqual(209, merged['amount_with_tax'])
        self.assertEqual('8% / 0', merged['tax_rate'])
        self.assertIsNone(grouped_tax([first, tax_fields({'amount': 50})])['line_tax_amount'])

    def test_invoice_tax_mismatch_is_preserved_without_allocating_to_goods(self):
        invoice = dict(tax_amount=99, items=[dict(id=1, line_index=1, amount=1000, tax_rate='8')])
        annotate_invoice_tax(invoice, json.dumps({'hdhhdvu': [{'tthue': 81}]}))
        self.assertEqual(81, invoice['items'][0]['line_tax_amount'])
        self.assertEqual(18, invoice['detail_tax_difference'])


class CustomerReportTests(unittest.TestCase):
    def setUp(self):
        self.fixture = inventory_fixtures.InventoryExportTests(); self.fixture.setUp()
        self.conn = self.fixture.conn
        self.conn.execute('UPDATE msmi_invoices SET raw_json=? WHERE id=?',
            (json.dumps({'hdhhdvu': [{'tthue': 80}]}), self.fixture.input_invoice_id))
        self.conn.execute('UPDATE outgoing_source_invoices SET raw_json=? WHERE id=?',
            (json.dumps({'details': [{'vatAmount': 72}]}), self.fixture.output_invoice_id))
        @contextmanager
        def db(): yield self.conn
        app = Flask(__name__); app.config['TESTING'] = True
        register_inventory_export_routes(app, dict(db=db, opening_template_path=inventory_fixtures.OPENING_TEMPLATE))
        self.client = app.test_client()
        self.query = '?from=2026-08-01&to=2026-08-31'

    def tearDown(self): self.fixture.tearDown()

    def test_sales_uses_900_source_revenue_even_when_valuation_fails(self):
        before = list(self.conn.iterdump())
        with patch('tdp_system.inventory_export.moving_average_report', side_effect=AssertionError('sales must not read cost')):
            response = self.client.get('/api/invoice-valuation/export/output' + self.query)
            self.assertEqual(200, response.status_code)
            wb = load_workbook(io.BytesIO(response.data)); ws = wb.active
            self.assertEqual((900, 72, 972), tuple(ws[f'I{row}'].value for row in (3, 4, 5)))
            self.assertEqual((900, '8', 72, 972), tuple(ws[c].value for c in ('I8', 'K8', 'L8', 'M8')))
            self.assertNotIn('Giá vốn', str(list(ws.values)))
            self.assertEqual(900, ws.cell(ws.max_row, 9).value)
            wb.close()
            self.assertEqual(200, self.client.get('/api/invoice-valuation/preview/output' + self.query).status_code)
        self.assertEqual(before, list(self.conn.iterdump()))

    def test_zip_nxt_input_and_preview_share_source_tax_without_database_writes(self):
        before = list(self.conn.iterdump())
        response = self.client.get('/api/invoice-valuation/export' + self.query)
        self.assertEqual(200, response.status_code)
        with zipfile.ZipFile(io.BytesIO(response.data)) as archive:
            self.assertEqual(5, len(archive.namelist()))
            self.assertFalse(any('gia_von' in name for name in archive.namelist()))
            for kind, prefix in [('input', 'Nhap_'), ('nxt', 'NXT_'), ('closing', 'Ton_trong_ky_')]:
                name = next(n for n in archive.namelist() if n.startswith(prefix))
                zipped = load_workbook(io.BytesIO(archive.read(name)))
                individual = load_workbook(io.BytesIO(self.client.get('/api/invoice-valuation/export/' + kind + self.query).data))
                self.assertEqual(list(zipped.active.values), list(individual.active.values))
                ws = zipped.active
                if kind == 'input':
                    self.assertEqual(('8', 80, 1080), tuple(ws[c].value for c in ('O5', 'P5', 'Q5')))
                elif kind == 'closing':
                    self.assertEqual((12, 133.333333, 1600), tuple(ws.cell(7, c).value for c in (7, 8, 9)))
                    self.assertEqual(1600, ws.cell(ws.max_row, 9).value)
                else:
                    self.assertEqual((10, 5, 3, 12), tuple(ws.cell(10, c).value for c in (7, 10, 13, 16)))
                    self.assertEqual((100, 1000, 200, 1000, 300, 900, 133.333333, 1600), tuple(ws.cell(10, c).value for c in (8, 9, 11, 12, 14, 15, 17, 18)))
                    self.assertNotIn('Tiền thuế', str(list(ws.values)))
                    self.assertNotIn('Giá vốn', str(list(ws.values)))
                zipped.close(); individual.close()
        self.assertEqual(before, list(self.conn.iterdump()))

    def test_unmapped_sales_are_visible_in_nxt_and_do_not_create_stock_quantity(self):
        self.conn.execute("UPDATE outgoing_source_invoice_items SET product_code='',mapping_status='unmapped'")
        response = self.client.get('/api/invoice-valuation/export/nxt' + self.query)
        wb = load_workbook(io.BytesIO(response.data)); ws = wb.active
        self.assertEqual((0, 0, 0, 0, 900), tuple(ws.cell(11, c).value for c in (7, 10, 13, 16, 15)))
        self.assertIn('chưa khớp mã', str(list(wb['Đối chiếu'].values)))
        wb.close()

    def test_invalid_dates_and_removed_cost_route(self):
        for kind in ('input', 'output', 'nxt', 'closing'):
            self.assertEqual(400, self.client.get('/api/invoice-valuation/export/' + kind + '?from=2026-08-31&to=2026-08-01').status_code)
        self.assertEqual(404, self.client.get('/api/invoice-valuation/export/output_cost' + self.query).status_code)

    def test_closing_requires_complete_month(self):
        for route in ('export', 'preview'):
            response = self.client.get('/api/invoice-valuation/' + route + '/closing?from=2026-08-02&to=2026-08-31')
            self.assertEqual(400, response.status_code)
            self.assertEqual('full_calendar_month_required', response.json['code'])

    def test_special_tax_labels_reach_download_and_preview_without_source_edits(self):
        self.conn.execute("UPDATE outgoing_source_invoice_items SET tax_rate='-2'")
        before = list(self.conn.iterdump())
        response = self.client.get('/api/invoice-valuation/export/output' + self.query)
        self.assertEqual(200, response.status_code)
        book = load_workbook(io.BytesIO(response.data))
        self.assertEqual('KKKNT', book.active['K8'].value)
        self.assertEqual(72, book.active['L8'].value)
        book.close()
        response = self.client.get('/api/invoice-valuation/preview/output' + self.query)
        self.assertIn('KKKNT', json.dumps(response.json))
        self.assertEqual(before, list(self.conn.iterdump()))

    def test_month_transfer_does_not_request_stock_posting_for_financial_only_invoices(self):
        from .inventory_period_close import init_inventory_period_close_schema, inventory_period_close_preview
        from datetime import date
        init_inventory_period_close_schema(self.conn)
        self.conn.execute("UPDATE outgoing_source_invoices SET stock_status='blocked'")
        self.conn.execute('UPDATE outgoing_source_invoice_items SET inventory_eligible=0')
        before = list(self.conn.iterdump())
        preview = inventory_period_close_preview(self.conn, '2026-08', today=date(2026,9,9))
        self.assertEqual(0, preview['unposted_output_count'])
        self.assertEqual(before, list(self.conn.iterdump()))
        self.conn.execute('UPDATE outgoing_source_invoice_items SET inventory_eligible=1')
        preview = inventory_period_close_preview(self.conn, '2026-08', today=date(2026,9,9))
        self.assertEqual(1, preview['unposted_output_count'])


class SignedSalesTaxTests(unittest.TestCase):
    def test_negative_source_tax_and_mismatched_header_survive_excel_export(self):
        fixture = sales_fixtures.OutputRegisterTests(); fixture.setUp()
        try:
            payload = fixture.payload()
            wb = load_workbook(output_sales_workbook(payload)); ws = wb.active
            self.assertEqual([8000, 4000, 8000, -8000], [ws.cell(r, 12).value for r in range(8, 12)])
            self.assertEqual((170000, 13600, 183600), tuple(ws.cell(ws.max_row, c).value for c in (9, 12, 13)))
            self.assertEqual(1600, ws['L13'].value)
            wb.close()
        finally:
            fixture.tearDown()
