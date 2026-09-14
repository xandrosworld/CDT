"""Customer flow: edit downloaded lines, sign externally, reconcile exactly once."""
import io
import unittest
import zipfile

from openpyxl import load_workbook

from . import server, test_outgoing_readiness as fixture
from .outgoing_waiting import refresh_waiting
from .outgoing_unissued import unissued_payload
from .outgoing_source_refresh import refresh_sources
from .receivable_ledger import sync_receivable_ledger
from .invoice_payment_scope import issued_invoice_payment_scope
from .test_invoice_output_sync import documented_minvoice_invoice, OutputFixtureMsmi


class InvoiceFinalScopeTests(unittest.TestCase):
    setUpClass = classmethod(fixture.OutgoingReadinessTests.setUpClass.__func__)
    tearDownClass = classmethod(fixture.OutgoingReadinessTests.tearDownClass.__func__)

    def setUp(self):
        with server.db() as conn:
            for table in ('invoice_inventory_ledger', 'invoice_sync_batch_output_invoices', 'invoice_mapping_revisions',
                          'invoice_line_mappings', 'outgoing_source_order_scopes',
                          'outgoing_buyer_profiles', 'outgoing_product_names', 'outgoing_product_units',
                          'receivable_ledger_revisions', 'receivable_ledger_lines'):
                conn.execute('DELETE FROM '+table)
        fixture.OutgoingReadinessTests.setUp(self)

    def business(self, conn):
        names = [r[0] for r in conn.execute("SELECT name FROM sqlite_master WHERE type='table'")
                 if r[0] in ('orders', 'batches') or r[0].startswith(('receivable_', 'payable_'))]
        return {name: [tuple(r) for r in conn.execute('SELECT * FROM '+name+' ORDER BY rowid')]
                for name in names}

    def test_external_signed_subset_only_reduces_actual_stock_and_payment_once(self):
        f=fixture.OutgoingReadinessTests
        with server.db() as conn:
            conn.execute("INSERT OR REPLACE INTO products(code,name,unit,tax) VALUES('HH-02','Hàng 02','kg','0%')")
            f.add_opening(conn,20)
            f.add_opening(conn,20,product_code='HH-02')
            _, ids=f.add_batch(conn,'2026-09-01',[{'qty':5},{'qty':3,'product_code':'HH-02'}])
            conn.execute("INSERT INTO outgoing_buyer_profiles(contractor,legal_name,tax_code,address,updated_at) VALUES('NT-A','Khách hàng kiểm thử','0209999999','Địa chỉ mua','now')")
            for key,value in {'company':'TĐP','company_tax_code':'0202265016','company_address':'Hải Phòng',
                              'payment_requester':'Thụy','payment_bank_name':'VCB','payment_bank_account':'123'}.items():
                server.setting_set(conn,key,value)
            sync_receivable_ledger(conn,timestamp=server.now_iso())
            refresh_waiting(conn,server.now_iso())
            before=self.business(conn)
            report=unissued_payload(conn,'2026-09-01','NT-A')
            self.assertEqual(sum(r['unissued_qty'] for r in report['details']),8)
            self.assertEqual(sum(r['ready_qty'] for r in report['details']),8)
            self.assertEqual(conn.execute("SELECT COUNT(*) FROM invoice_inventory_ledger WHERE direction='output'").fetchone()[0],0)

        # The downloaded sheet contained both codes; the customer deleted HH-02
        # and signed only 2 kg of HH-01 on M-Invoice, after the order cutoff.
        raw=documented_minvoice_invoice(3)
        raw.update(inv_invoiceIssuedDate='2026-09-03',tgtcthue=40,tgtthue=0,tgtttbso=40,
                   inv_buyerAddressLine='Địa chỉ mua')
        raw['details'][0].update(inv_itemCode='HH-01',inv_itemName='Hàng hóa 01',inv_unitCode='kg',
                                 inv_quantity=2,inv_unitPrice=20,inv_TotalAmountWithoutVat=40,
                                 inv_vatAmount=0,inv_TotalAmount=40,ma_thue='0')
        remote=OutputFixtureMsmi([raw])
        for attempt in range(2):
            refresh_sources(server.db,lambda:remote,server.now_iso,'2026-09-01','2026-09-03')
            with server.db() as conn:
                self.assertEqual(self.business(conn),before)
                report=unissued_payload(conn,'2026-09-01','NT-A')
                self.assertEqual(report['warnings'],[])
                self.assertEqual({r['order_id']:r['unissued_qty'] for r in report['details']},
                                 {ids[0]:3,ids[1]:3})
                outputs=[tuple(r) for r in conn.execute("SELECT product_code,qty_delta FROM invoice_inventory_ledger WHERE direction='output' AND status='posted'")]
                self.assertEqual(outputs,[('HH-01',-2)])
                scope=issued_invoice_payment_scope(conn,'NT-A','2026-09-01','2026-09-30')
                self.assertEqual(scope['totals']['total_amount'],40)
                self.assertEqual(scope['statement_kind'],'invoices')
                self.assertEqual(len(scope['invoices']),1)
                self.assertEqual(len(scope['source_lines']),1)
                self.assertEqual(scope['source_lines'][0]['qty'],2)

    def test_full_excel_keeps_ready_unsigned_and_red_conversion_with_original_values(self):
        f=fixture.OutgoingReadinessTests
        with server.db() as conn:
            f.add_opening(conn,20)
            _,ids=f.add_batch(conn,'2026-09-01',[{'qty':5,'sell_price':20},{'qty':14,'sell_price':8000}])
            conn.execute("UPDATE orders SET unit='Gói',tax='KKKNT' WHERE id=?",(ids[1],))
            conn.execute("INSERT INTO outgoing_product_names(product_code,invoice_name,updated_at) VALUES('HH-01','Tên xuất hóa đơn','now')")
            refresh_waiting(conn,server.now_iso())
            before=conn.serialize()
        url='/api/outgoing-invoices/unissued'
        scope={'to':'2026-09-01','contractor':'NT-A'}
        data=self.client.get(url,query_string=scope).json
        self.assertEqual(len(data['details']),2)
        self.assertEqual([r['order_id'] for r in data['details'] if r['needs_conversion']],[ids[1]])
        self.assertEqual(data['details'][0]['ready_qty'],5)
        response=self.client.get(url+'-template.zip',query_string=scope)
        self.assertEqual(response.status_code,200,response.get_json(silent=True))
        with zipfile.ZipFile(io.BytesIO(response.data)) as archive:
            files=[n for n in archive.namelist() if n.startswith('CHUA_XUAT_')]
            self.assertEqual(len(files),2)
            for name in files:
                book=load_workbook(io.BytesIO(archive.read(name)))
                sheet=book.active
                self.assertEqual(sheet.max_column,13)
                self.assertEqual(sheet['B2'].value,'Tên xuất hóa đơn')
                if 'KKKNT' in name:
                    self.assertEqual([sheet.cell(2,c).value for c in (3,4,5,6,10)],['Gói',14,8000,112000,-2])
                    self.assertEqual(sheet['B2'].font.color.rgb,'00B42318')
                    self.assertIn('Cần quy đổi',sheet['C2'].comment.text)
                else:
                    self.assertEqual(sheet['D2'].value,5)
                    self.assertIsNone(sheet['C2'].comment)
                book.close()
        with server.db() as conn:
            self.assertEqual(conn.serialize(),before)


if __name__=='__main__':unittest.main()
