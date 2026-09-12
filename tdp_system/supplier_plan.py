"""Supplier instructions from the customer's purchase sheet, without posting debt.

The initial workbook is a plan and may not have purchase prices yet. Its rows
must never be reconstructed from sales or treated as confirmed purchases.
"""
import hashlib
import json

from openpyxl import load_workbook


def _modules():
    try:
        from . import contract_modules
    except ImportError:
        import contract_modules
    return contract_modules


def parse_supplier_plan(conn, workbook, batch_id, formula_workbook=None):
    cm = _modules()
    preview = cm.parse_canonical_purchase_workbook(
        conn, workbook, batch_id, {}, formula_workbook, plan_only=True,
    )
    if preview is None:
        raise ValueError('Không tìm thấy đúng sheet đặt hàng. Hãy chọn file Excel gốc có sheet này.')
    for item in preview['items']:
        if item['actual_qty'] > 0 and not item['unit']:
            item['errors'].append('Dòng đặt hàng thiếu đơn vị tính')
    preview['error_rows'] = sum(bool(r['errors']) for r in preview['items'])
    preview['can_confirm'] = not preview['error_rows']
    preview['issues'] = [r for r in preview['items'] if r['errors'] or r['warnings']][:200]
    preview['plan_only'] = True
    return preview


def _image_rows(payload):
    cm = _modules()
    groups = {}
    for row in payload['rows']:
        if row['order_qty'] <= 0:
            continue
        key = cm.supplier_merge_key(row['supplier'])
        groups.setdefault(key, []).append(tuple(row.get(k) for k in (
            'work_date', 'kitchen', 'product_code', 'product_name', 'order_qty', 'unit', 'note',
        )))
    return {k: sorted(rows, key=repr) for k, rows in groups.items()}


def reopen_changed_suppliers(conn, batch_id, before, now_iso):
    cm = _modules()
    old = _image_rows(before)
    new = _image_rows(cm.purchase_order_payload(conn, batch_id, for_sending=True))
    for key in old.keys() | new.keys():
        if old.get(key) != new.get(key):
            conn.execute("""UPDATE supplier_order_statuses SET status='reopened',
                revision=revision+1,reopened_at=?,updated_at=?
                WHERE batch_id=? AND supplier_key=?""", (now_iso(), now_iso(), batch_id, key))


def save_supplier_plan(conn, *, batch_id, preview, source_hash, source_name, now_iso):
    """Caller owns the transaction. No order, payable or inventory writes."""
    cm = _modules()
    items = [{k: v for k, v in item.items() if k not in ('_issue_columns', 'errors', 'warnings')}
             for item in preview.get('items', [])]
    issues = preview.get('source_issues', []) or [
        f"Dòng {r['source_row']}: {error}"
        for r in preview.get('items', []) for error in r.get('errors', [])
    ]
    encoded = json.dumps(items, ensure_ascii=False, sort_keys=True)
    encoded_issues = json.dumps(issues, ensure_ascii=False)
    content_hash = hashlib.sha256((encoded + encoded_issues).encode()).hexdigest().upper()
    previous = conn.execute('SELECT * FROM supplier_plan_sources WHERE batch_id=?', (batch_id,)).fetchone()
    if previous and previous['content_hash'] == content_hash and previous['source_hash'] == source_hash:
        return {'processed': 0, 'count': len(items), 'idempotent': True, 'plan_only': True}
    if conn.execute('SELECT 1 FROM supplier_plan_history WHERE batch_id=? AND source_hash=? AND content_hash=?',
                    (batch_id, source_hash, content_hash)).fetchone():
        raise cm.PurchaseOrderApplyError('File đặt hàng này đã được thay bằng bản mới hơn. Hãy chọn file hiện hành.',
                                         code='superseded_supplier_plan')
    before = cm.purchase_order_payload(conn, batch_id, for_sending=True)
    revision = (previous['revision'] if previous else 0) + 1
    timestamp = now_iso()
    conn.execute('''INSERT INTO supplier_plan_sources
        (batch_id,source_hash,source_name,source_sheet,content_hash,items_json,issues_json,revision,updated_at)
        VALUES(?,?,?,?,?,?,?,?,?) ON CONFLICT(batch_id) DO UPDATE SET
        source_hash=excluded.source_hash,source_name=excluded.source_name,source_sheet=excluded.source_sheet,
        content_hash=excluded.content_hash,items_json=excluded.items_json,issues_json=excluded.issues_json,
        revision=excluded.revision,updated_at=excluded.updated_at''',
        (batch_id, source_hash, source_name, preview.get('sheet', 'đặt hàng'), content_hash,
         encoded, encoded_issues, revision, timestamp))
    conn.execute('''INSERT INTO supplier_plan_history
        (batch_id,source_hash,content_hash,source_name,items_json,issues_json,revision,created_at)
        VALUES(?,?,?,?,?,?,?,?)''',
        (batch_id, source_hash, content_hash, source_name, encoded, encoded_issues, revision, timestamp))
    reopen_changed_suppliers(conn, batch_id, before, now_iso)
    cm.audit(conn, now_iso, 'supplier_plan.import', 'ok', entity_type='batch', entity_id=batch_id,
             metadata={'rows': len(items), 'source_hash': source_hash[:16], 'issues': len(issues)})
    return {'processed': len(items), 'count': len(items), 'idempotent': False, 'plan_only': True}


def capture_supplier_plan(conn, *, batch_id, path, source_hash, source_name, now_iso, only_missing=False):
    if only_missing and (conn.execute('SELECT 1 FROM supplier_plan_sources WHERE batch_id=?', (batch_id,)).fetchone()
                         or conn.execute('SELECT 1 FROM purchase_workbook_lines WHERE batch_id=?', (batch_id,)).fetchone()):
        return None
    workbook = formula_workbook = None
    try:
        workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        formula_workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        preview = parse_supplier_plan(conn, workbook, batch_id, formula_workbook)
    except ValueError as error:
        preview = {'items': [], 'source_issues': [str(error)]}
    finally:
        if workbook is not None:
            workbook.close()
        if formula_workbook is not None:
            formula_workbook.close()
    return save_supplier_plan(conn, batch_id=batch_id, preview=preview, source_hash=source_hash,
                              source_name=source_name, now_iso=now_iso)
