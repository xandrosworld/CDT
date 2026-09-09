import copy
import json
import sqlite3
import subprocess
import unittest
from contextlib import contextmanager
from io import BytesIO
from pathlib import Path

from flask import Flask
from openpyxl import load_workbook

from .test_invoice_input_sync import init_test_database, now_iso
from .test_minvoice_portal import document
from .minvoice_portal import normalize_portal_document
from .invoice_output_sync import upsert_output_invoice
from .invoice_mapping import save_mapping
from .invoice_workbench import register_invoice_workbench_routes
from .invoice_workbench_listing import invoice_range_payload
from .invoice_output_register import output_sales_workbook


class OutputRegisterTests(unittest.TestCase):
    def setUp(self):
        self.conn = sqlite3.connect(':memory:'); init_test_database(self.conn)
        self.conn.execute('ALTER TABLE products ADD COLUMN purchase_list INTEGER DEFAULT 0')
        self.conn.execute("INSERT INTO products(code,name,unit,purchase_list) VALUES('A','Goods','Chai',1)")
        self.ids = []
        for index in range(3):
            raw = document(); raw['id'] = 'register-' + str(index); raw['invoiceNumber'] = 12 + index
            raw['invoiceDetail'][0]['unitCode'] = 'xách'
            if index == 0:
                second = copy.deepcopy(raw['invoiceDetail'][0])
                second.update(quantity=1, amount=50000, amountWithoutVAT=50000, vatAmount=4000)
                raw['invoiceDetail'].append(second)
                raw.update(totalAmountWithoutVAT=150000, vatAmount=12000, totalAmount=162000)
            if index == 1:
                # The source header and details disagree. Preserve both.
                raw.update(totalAmountWithoutVAT=120000, vatAmount=9600, totalAmount=129600)
            if index == 2:
                raw.update(invoiceStatus=2, totalAmountWithoutVAT=-100000, vatAmount=-8000, totalAmount=-108000)
                raw['invoiceDetail'][0].update(quantity=-2, amount=-100000, amountWithoutVAT=-100000, vatAmount=-8000)
            iid = upsert_output_invoice(self.conn, normalize_portal_document(raw), tenant='TDP', now=now_iso(),
                                        status_map={}, status_fields=(), reference_fields=())[0]
            self.ids.append(iid)
            for line in self.conn.execute('SELECT id,inventory_eligible FROM outgoing_source_invoice_items WHERE invoice_id=?', (iid,)).fetchall():
                if line['inventory_eligible']:
                    save_mapping(self.conn, direction='output', item_id=line['id'], product_code='A', now_iso=now_iso)
        # A previously posted invoice and BK with negative opening must both
        # remain in this source report; no inventory operation is requested.
        self.conn.execute("UPDATE outgoing_source_invoices SET stock_status='posted' WHERE id=?", (self.ids[0],))
        self.conn.execute("""INSERT INTO inventory_transactions(txn_date,product_code,qty_in,qty_out,unit_cost,
            source_type,source_id,source_line,status,created_at,updated_at)
            VALUES('2026-08-01','A',0,10,100,'OPENING','2026-08','A','posted',?,?)""", (now_iso(), now_iso()))

    def tearDown(self): self.conn.close()

    def payload(self, **kwargs):
        return invoice_range_payload(self.conn, tenant='TDP', invoice_type='output',
                                     date_from='2026-08-01', date_to='2026-08-31', **kwargs)

    def test_period_totals_survive_line_and_stock_filters_without_writing(self):
        before = list(self.conn.iterdump())
        full = self.payload(); summary = full['output_summary']
        self.assertEqual(3, summary['invoice_count']); self.assertEqual(4, summary['line_count'])
        self.assertEqual(170000, summary['subtotal']); self.assertEqual(13600, summary['tax_amount'])
        self.assertEqual(183600, summary['total_amount'])
        self.assertEqual(150000, summary['detail_amount']); self.assertEqual(20000, summary['detail_difference'])
        self.assertEqual(3, summary['mapped_line_count'])
        for filters in ({'line_filter':'needs_attention'}, {'status':'posted'}, {'scope':'pending'}):
            result = self.payload(**filters)
            self.assertEqual(summary, result['output_summary'])
        self.assertNotIn(self.ids[0], [h['id'] for h in self.payload(scope='pending')['items']])
        self.assertEqual(before, list(self.conn.iterdump()))

    def test_export_preserves_all_source_lines_signed_adjustments_and_header_amounts(self):
        payload = self.payload(); before = list(self.conn.iterdump())
        wb = load_workbook(output_sales_workbook(payload)); ws = wb.active
        self.assertEqual(['Dau ra M-Invoice'], wb.sheetnames)
        self.assertEqual((170000, 13600, 183600), (ws['I3'].value, ws['I4'].value, ws['I5'].value))
        rows = list(ws.iter_rows(min_row=8, max_row=11, values_only=True))
        self.assertEqual(4, len(rows)); self.assertEqual({'xách'}, {r[5] for r in rows})
        self.assertEqual([2, 1, 2, -2], [r[6] for r in rows])
        self.assertEqual([100000, 50000, 100000, -100000], [r[8] for r in rows])
        self.assertEqual(150000, ws['I12'].value); self.assertEqual(20000, ws['I13'].value)
        self.assertNotIn('Giá vốn', str(list(ws.values)))
        self.assertEqual(before, list(self.conn.iterdump()))
        wb.close()

    def test_export_route_ignores_old_filters_and_is_tenant_scoped(self):
        @contextmanager
        def db(): yield self.conn
        app = Flask(__name__)
        register_invoice_workbench_routes(app, {'db':db, 'now_iso':now_iso, 'setting_get':lambda c,k,d:d})
        client = app.test_client()
        url = '/api/invoice-workbench/output-register/export?from=2026-08-01&to=2026-08-31&scope=pending&status=posted&line_filter=unmapped&invoice_type=input'
        response = client.get(url); self.assertEqual(200, response.status_code)
        wb = load_workbook(BytesIO(response.data)); self.assertEqual(170000, wb.active['I3'].value); wb.close()
        other = invoice_range_payload(self.conn, tenant='OTHER', invoice_type='output',date_from='2026-08-01',date_to='2026-08-31')
        self.assertEqual(0, other['output_summary']['invoice_count'])
        self.assertEqual(400, client.get(url.replace('from=2026-08-01','from=2026-09-01')).status_code)

    def test_source_strings_are_not_excel_formulas(self):
        self.conn.execute("UPDATE outgoing_source_invoice_items SET source_item_name='=1+1' WHERE invoice_id=?", (self.ids[0],))
        wb = load_workbook(output_sales_workbook(self.payload()))
        self.assertEqual('=1+1', wb.active['E8'].value); self.assertEqual('s', wb.active['E8'].data_type); wb.close()

    def test_renderer_exposes_full_export_and_keeps_stock_actions_available(self):
        payload = self.payload(line_filter='needs_attention')
        script = """
        const assert=require('node:assert/strict');global.window={};
        require('./tdp_system/static/invoice-workbench.js');
        const state={invoiceDirection:'output',invoiceFrom:'2026-08-01',invoiceTo:'2026-08-31',invoiceStatus:'all',invoiceLineFilter:'needs_attention',invoiceListing:JSON.parse(process.argv[1]),invoiceWorkbench:{batches:[]}};
        const f={esc:x=>String(x??''),num:x=>String(x??0),money:x=>String(x??0),dateVN:x=>x};
        const html=window.TdpInvoiceWorkbench(state,f);
        for(const text of ['Tiền hàng M-Invoice (chưa thuế): 170000','Tải bảng đầu ra','Xem toàn bộ đầu ra','Ghi xuất kho','Tạo file đưa lên M-Invoice','BK âm vẫn khớp mã và tải bảng được.'])assert(html.includes(text),text);
        const link=html.match(/href="([^"]+)"[^>]*>Tải bảng đầu ra/)[1];
        assert(link.includes('/output-register/export?from=2026-08-01&to=2026-08-31'));
        assert(!link.includes('line_filter')&&!link.includes('scope=pending'));
        assert(!html.includes('invoice-draft-factor'));
        """
        result = subprocess.run(['node','-e',script,json.dumps(payload)],cwd=Path(__file__).resolve().parent.parent,
                                capture_output=True,text=True,encoding='utf-8')
        self.assertEqual(0,result.returncode,result.stderr)


if __name__ == '__main__': unittest.main()
