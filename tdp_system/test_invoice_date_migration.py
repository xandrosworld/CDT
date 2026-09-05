"""Offline regression for upgrading persisted UTC dates, not merely parsing new ones."""
import json
import sqlite3
import tempfile
import unittest
from contextlib import closing
from pathlib import Path

from openpyxl import load_workbook

from .invoice_date_migration import repair_legacy_input_dates, input_date_repair_report
from .contract_modules import upsert_msmi_invoice
from .invoice_workbench_listing import invoice_range_payload, range_workbook
from .build_release_inputs import snapshot_database
from .test_invoice_input_sync import init_test_database, now_iso
from .test_msmi_sync import remote_invoice


class LegacyInvoiceDateTests(unittest.TestCase):
    def setUp(self):
        self.conn = sqlite3.connect(":memory:")
        init_test_database(self.conn)

    def tearDown(self):
        self.conn.close()

    def invoice(self, number=1, source="2026-07-31T17:00:00Z", stored="2026-07-31"):
        raw = remote_invoice(number)
        raw["tdlap"] = source
        invoice_id, _ = upsert_msmi_invoice(self.conn, raw, "INPUT_ELECTRONIC_INVOICE", "TDP", now_iso())
        self.conn.execute("UPDATE msmi_invoices SET invoice_date=? WHERE id=?", (stored, invoice_id))
        return invoice_id

    def listing(self):
        return invoice_range_payload(self.conn, tenant="TDP", invoice_type="input",
                                     date_from="2026-08-01", date_to="2026-08-31")

    def report(self, tenant="TDP"):
        return input_date_repair_report(self.conn, tenant=tenant, date_from="2026-08-01", date_to="2026-08-31")

    def test_persisted_257_becomes_exact_266_without_reimport(self):
        for n in range(1, 267):
            self.invoice(n, source="2026-07-31T17:00:00Z" if n <= 9 else "2026-08-15T17:00:00Z",
                         stored="2026-07-31" if n <= 9 else "2026-08-15")
        before = self.listing()
        line_rows = [tuple(r) for r in self.conn.execute("SELECT * FROM msmi_invoice_items ORDER BY id")]
        source_rows = [tuple(r) for r in self.conn.execute("SELECT id,remote_id,raw_json,subtotal,tax_amount,total_amount,receipt_status FROM msmi_invoices ORDER BY id")]
        self.assertEqual(257, before["counts"]["all"])
        self.assertEqual({"corrected": 266, "blocked": 0}, repair_legacy_input_dates(self.conn))
        after = self.listing()
        self.assertEqual(266, after["counts"]["all"])
        self.assertEqual(266, after["totals"]["invoice_count"])
        self.assertEqual(266 * 10800, after["totals"]["invoice_amount"])
        self.assertEqual(9, len({r["id"] for r in after["items"]} - {r["id"] for r in before["items"]}))
        self.assertEqual(line_rows, [tuple(r) for r in self.conn.execute("SELECT * FROM msmi_invoice_items ORDER BY id")])
        self.assertEqual(source_rows, [tuple(r) for r in self.conn.execute("SELECT id,remote_id,raw_json,subtotal,tax_amount,total_amount,receipt_status FROM msmi_invoices ORDER BY id")])
        self.assertEqual(0, self.conn.execute("SELECT COUNT(*) FROM invoice_inventory_ledger").fetchone()[0])
        workbook = load_workbook(range_workbook(after), data_only=True)
        self.assertEqual(266, sum(isinstance(r[0], str) and r[0].startswith("2026-08-") for r in workbook.active.values))
        workbook.close()

    def test_restart_is_idempotent_and_audit_retains_old_and_new_dates(self):
        invoice_id = self.invoice()
        repair_legacy_input_dates(self.conn, timestamp="first")
        changes = self.conn.total_changes
        self.assertEqual({"corrected": 0, "blocked": 0}, repair_legacy_input_dates(self.conn, timestamp="second"))
        self.assertEqual(changes, self.conn.total_changes)
        self.assertEqual((invoice_id, "2026-07-31", "2026-08-01", "corrected", "first"), tuple(self.conn.execute(
            "SELECT invoice_id,old_date,new_date,status,created_at FROM invoice_date_repairs").fetchone()))

    def test_september_first_is_removed_from_august_not_double_counted(self):
        self.invoice(source="2026-08-31T17:00:00Z", stored="2026-08-31")
        repair_legacy_input_dates(self.conn)
        self.assertEqual(0, self.listing()["counts"]["all"])
        self.assertEqual("2026-09-01", self.conn.execute("SELECT invoice_date FROM msmi_invoices").fetchone()[0])

    def test_posted_invoice_is_visible_in_review_even_outside_filtered_month(self):
        invoice_id = self.invoice()
        self.conn.execute("UPDATE msmi_invoices SET receipt_status='posted' WHERE id=?", (invoice_id,))
        self.assertEqual({"corrected": 0, "blocked": 1}, repair_legacy_input_dates(self.conn))
        self.assertEqual(("2026-07-31", "posted", "review_required"), tuple(self.conn.execute(
            "SELECT invoice_date,receipt_status,sync_status FROM msmi_invoices").fetchone()))
        self.assertEqual(0, self.listing()["counts"]["all"])
        self.assertEqual(1, self.listing()["date_repair"]["blocked_count"])
        self.assertEqual(invoice_id, self.report()["items"][0]["invoice_id"])
        self.assertEqual(0, self.report("OTHER")["blocked_count"])
        before = self.conn.total_changes
        repair_legacy_input_dates(self.conn)
        self.assertEqual(before, self.conn.total_changes)

    def test_existing_legacy_stock_footprint_blocks_even_if_header_status_wrong(self):
        self.invoice()
        self.conn.execute("INSERT INTO products(code,name,unit) VALUES('P','P','kg')")
        self.conn.execute("""INSERT INTO inventory_transactions
            (txn_date,product_code,qty_in,qty_out,unit_cost,source_type,source_id,source_line,status,note,created_at,updated_at)
            VALUES('2026-07-31','P',1,0,10000,'MSMI_INPUT','MSMI-001','1','posted','','a','a')""")
        before = [tuple(r) for r in self.conn.execute("SELECT * FROM inventory_transactions")]
        self.assertEqual(1, repair_legacy_input_dates(self.conn)["blocked"])
        self.assertEqual(before, [tuple(r) for r in self.conn.execute("SELECT * FROM inventory_transactions")])

    def test_canonical_confirmation_alone_prevents_date_rewrite(self):
        invoice_id = self.invoice()
        self.conn.execute("""INSERT INTO invoice_inventory_confirmations
            (confirmation_key,direction,source_invoice_table,source_invoice_id,action,confirmed,note,created_at)
            VALUES('guard','input','msmi_invoices',?,'post',1,'','a')""", (invoice_id,))
        self.assertEqual(1, repair_legacy_input_dates(self.conn)["blocked"])

    def test_ambiguous_manual_date_is_blocked_not_shifted(self):
        self.invoice(stored="2026-07-20")
        self.assertEqual(1, repair_legacy_input_dates(self.conn)["blocked"])
        self.assertEqual("2026-07-20", self.conn.execute("SELECT invoice_date FROM msmi_invoices").fetchone()[0])

    def test_date_only_local_timestamps_and_malformed_payloads_are_not_guessed(self):
        for n, source in enumerate(["2026-08-01", "2026-08-01T00:00:00", "2026-08-01T00:00:00+07:00"], 1):
            self.invoice(n, source=source, stored="2026-08-01")
        self.invoice(4)
        self.conn.execute("UPDATE msmi_invoices SET raw_json='broken' WHERE invoice_number='4'")
        self.assertEqual({"corrected": 0, "blocked": 0}, repair_legacy_input_dates(self.conn))

    def test_source_changes_do_not_leave_stale_review_notices(self):
        invoice_id = self.invoice()
        self.conn.execute("UPDATE msmi_invoices SET receipt_status='posted' WHERE id=?", (invoice_id,))
        repair_legacy_input_dates(self.conn)
        self.conn.execute("UPDATE msmi_invoices SET invoice_date='2026-08-01' WHERE id=?", (invoice_id,))
        self.assertEqual(0, self.report()["blocked_count"])

    def test_failure_rolls_back_all_date_repairs_and_history(self):
        self.invoice(1)
        self.invoice(2)
        self.conn.execute("""CREATE TRIGGER fail_repair BEFORE UPDATE OF invoice_date ON msmi_invoices
            WHEN NEW.invoice_number='2' BEGIN SELECT RAISE(ABORT,'simulated failure'); END""")
        with self.assertRaises(sqlite3.IntegrityError):
            repair_legacy_input_dates(self.conn)
        self.assertEqual(["2026-07-31", "2026-07-31"], [r[0] for r in self.conn.execute("SELECT invoice_date FROM msmi_invoices")])
        self.assertIsNone(self.conn.execute("SELECT name FROM sqlite_master WHERE name='invoice_date_repairs'").fetchone())

    def test_caller_transaction_can_roll_back_repair(self):
        self.invoice()
        self.conn.commit()
        self.conn.execute("BEGIN")
        repair_legacy_input_dates(self.conn)
        self.conn.rollback()
        self.assertEqual("2026-07-31", self.conn.execute("SELECT invoice_date FROM msmi_invoices").fetchone()[0])

    def test_release_snapshot_is_corrected_without_touching_source(self):
        self.invoice()
        self.conn.commit()
        with tempfile.TemporaryDirectory() as temp:
            source, dest = Path(temp) / "source.sqlite3", Path(temp) / "seed.sqlite3"
            with closing(sqlite3.connect(source)) as disk:
                self.conn.backup(disk)
            snapshot_database(source, dest)
            a, b = sqlite3.connect(source), sqlite3.connect(dest)
            try:
                self.assertEqual("2026-07-31", a.execute("SELECT invoice_date FROM msmi_invoices").fetchone()[0])
                self.assertEqual("2026-08-01", b.execute("SELECT invoice_date FROM msmi_invoices").fetchone()[0])
            finally:
                a.close(); b.close()

    def test_build_refuses_unresolved_seed_and_preserves_previous_artifact(self):
        self.invoice()
        self.conn.execute("UPDATE msmi_invoices SET receipt_status='posted'")
        self.conn.commit()
        with tempfile.TemporaryDirectory() as temp:
            source, dest = Path(temp) / "source.sqlite3", Path(temp) / "seed.sqlite3"
            with closing(sqlite3.connect(source)) as disk:
                self.conn.backup(disk)
            dest.write_bytes(b"previous-seed")
            with self.assertRaises(sqlite3.DatabaseError):
                snapshot_database(source, dest)
            self.assertEqual(b"previous-seed", dest.read_bytes())

    def test_review_export_escapes_formula_and_does_not_add_to_invoice_total(self):
        self.invoice()
        self.conn.execute("UPDATE msmi_invoices SET receipt_status='posted',seller_name='=1+1'")
        repair_legacy_input_dates(self.conn)
        payload = self.listing()
        before = self.conn.total_changes
        workbook = load_workbook(range_workbook(payload), data_only=False)
        self.assertEqual("s", workbook['Ngay can doi chieu']['C3'].data_type)
        self.assertEqual("=1+1", workbook['Ngay can doi chieu']['C3'].value)
        self.assertEqual(0, payload['totals']['invoice_amount'])
        self.assertEqual(before, self.conn.total_changes)
        workbook.close()


if __name__ == '__main__':
    unittest.main()
