import io
import unittest
import zipfile
from openpyxl import load_workbook
from . import server
from . import test_outgoing_readiness as support
from .outgoing_weights import invoice_rows, confirmed_weights
from .outgoing_readiness import validate_draft_export_stock, canonical_available_stock
from .outgoing_consolidation import consolidate
from .contract_modules import invoice_tax_percent, create_partial_outgoing_drafts
from .invoice_mapping import validated_output_stock_snapshot, InvoiceMappingError
from .invoice_inventory import post_output_invoice


class ActualWeightTests(unittest.TestCase):
    setUpClass = classmethod(support.OutgoingReadinessTests.setUpClass.__func__)
    tearDownClass = classmethod(support.OutgoingReadinessTests.tearDownClass.__func__)

    def setUp(self):
        with server.db() as conn:
            conn.execute('DELETE FROM invoice_inventory_ledger')
            conn.execute('DELETE FROM outgoing_weight_exports')
            conn.execute("DELETE FROM invoice_mapping_revisions WHERE mapping_id IN (SELECT id FROM invoice_line_mappings WHERE source='tdp_actual_weight')")
            conn.execute("DELETE FROM invoice_line_mappings WHERE source='tdp_actual_weight'")
            conn.execute('DELETE FROM outgoing_product_units')
        support.OutgoingReadinessTests.setUp(self)
        with server.db() as conn:
            conn.execute("UPDATE products SET unit='Gói' WHERE code='HH-01'")
            conn.execute("INSERT INTO outgoing_product_units VALUES('HH-01','Kg','now')")
            conn.execute("INSERT OR REPLACE INTO outgoing_product_names VALUES('HH-01','Nấm kim châm','now')")
            self.bid, self.oids = support.OutgoingReadinessTests.add_batch(conn, '2026-09-01', [{'qty':14,'sell_price':8000}])
            conn.execute("UPDATE orders SET unit='Gói',tax='KKKNT'")
            support.OutgoingReadinessTests.add_opening(conn, 20)

    def rows(self):
        r = self.client.get('/api/outgoing-invoices/actual-weights?to=2026-09-13')
        self.assertEqual(200, r.status_code, r.json)
        return r.json['rows']

    def save(self, kg=2.8, row=None, **changes):
        row = row or self.rows()[0]
        return self.client.put('/api/outgoing-invoices/actual-weights/'+str(row['order_id']), json={
            'token':row['token'], 'actual_kg':kg, 'note':'Đã cân thực tế',
            'confirmed_actual_weight':True, **changes})

    def draft(self):
        with server.db() as conn:
            create_partial_outgoing_drafts(conn,self.bid,server.now_iso)
            return conn.execute("SELECT id FROM outgoing_invoice_drafts WHERE status='draft' ORDER BY id DESC").fetchone()[0]

    def test_separate_real_rows_alias_and_no_implicit_weight(self):
        row = self.rows()[0]
        self.assertEqual(('Nấm kim châm','Gói','Kg',14,112000,None), tuple(row[k] for k in
            ('invoice_name','unit','invoice_unit','remaining_qty','amount','actual_kg')))
        self.assertTrue(row['editable'])
        self.assertEqual(row['review_kind'],'actual_kg')
        with server.db() as conn:
            self.assertEqual({},confirmed_weights(conn))

    def test_weight_editor_respects_both_dates_and_contractor(self):
        with server.db() as conn:
            _, middle = support.OutgoingReadinessTests.add_batch(conn, '2026-09-05', [
                {'qty':2}, {'qty':3,'contractor':'NT-B'}])
            support.OutgoingReadinessTests.add_batch(conn, '2026-09-10', [{'qty':4}])
            conn.execute("UPDATE orders SET unit='Gói'")
        with server.db() as conn:
            before = conn.serialize()
        response = self.client.get('/api/outgoing-invoices/actual-weights', query_string={
            'from':'2026-09-05','to':'2026-09-07','contractor':'NT-A'})
        self.assertEqual(response.status_code, 200, response.json)
        self.assertEqual(response.json['from'], '2026-09-05')
        self.assertEqual([r['order_id'] for r in response.json['rows']], [middle[0]])
        self.assertTrue(response.json['rows'][0]['batch_id'])
        with server.db() as conn:
            self.assertEqual(conn.serialize(), before)
        invalid = self.client.get('/api/outgoing-invoices/actual-weights?from=2026-09-08&to=2026-09-07')
        self.assertEqual(invalid.status_code, 400)

    def test_weight_editor_does_not_restore_excluded_lines_or_contractors(self):
        with server.db() as conn:
            conn.execute("INSERT INTO outgoing_order_choices VALUES(?,0,'test','now')", (self.oids[0],))
        self.assertEqual(self.rows(), [])
        with server.db() as conn:
            conn.execute('DELETE FROM outgoing_order_choices')
            conn.execute("INSERT INTO outgoing_contractor_choices VALUES('NT-A',0,'test','now')")
        self.assertEqual(self.rows(), [])

    def test_save_changes_neither_orders_stock_nor_financial_ledgers(self):
        tables=['orders','inventory_transactions','outgoing_invoice_lines','payable_ledger_lines','receivable_ledger_lines']
        with server.db() as conn:
            before={t:[tuple(r) for r in conn.execute('SELECT * FROM '+t)] for t in tables}
        r=self.save();self.assertEqual(200,r.status_code,r.json)
        self.assertEqual((2.8,40000,112000),tuple(r.json[k] for k in ('actual_kg','price_per_kg','amount')))
        with server.db() as conn:
            for t in tables:self.assertEqual(before[t],[tuple(r) for r in conn.execute('SELECT * FROM '+t)])
        self.assertTrue(self.rows()[0]['confirmed'])

    def test_invalid_unconfirmed_and_stale_save_rejected(self):
        row=self.rows()[0]
        for kg in (0,-1,'NaN','Infinity',True,'abc',1000000,0.1234567):
            r=self.save(kg,row);self.assertEqual(409,r.status_code,(kg,r.json))
        self.assertEqual(409,self.save(row=row,note='').status_code)
        self.assertEqual(409,self.save(row=row,confirmed_actual_weight=False).status_code)
        self.assertEqual(200,self.save(row=row).status_code)
        self.assertEqual(409,self.save(3,row).status_code)
        self.assertEqual(2.8,self.rows()[0]['actual_kg'])

    def test_order_stock_mismatch_never_unlocked_by_kg(self):
        with server.db() as conn:conn.execute("UPDATE products SET unit='Hộp' WHERE code='HH-01'")
        self.assertFalse(self.rows()[0]['editable'])
        self.assertEqual(self.rows()[0]['review_kind'],'unit_mismatch')
        self.assertEqual(409,self.save().status_code)

    def test_package_label_mismatch_is_not_an_actual_kg_request(self):
        with server.db() as conn:
            conn.execute("UPDATE products SET unit='Hộp' WHERE code='HH-01'")
            conn.execute("UPDATE orders SET unit='Hộp'")
            conn.execute("UPDATE outgoing_product_units SET invoice_unit='Chai' WHERE product_code='HH-01'")
        row=self.rows()[0]
        self.assertEqual(row['review_kind'],'unit_mismatch')
        self.assertFalse(row['editable'])
        self.assertIn('Hộp',row['reason'])
        self.assertIn('Chai',row['reason'])
        self.assertEqual(409,self.save().status_code)

    def test_changed_order_invalidates_confirmation(self):
        self.assertEqual(200,self.save().status_code)
        with server.db() as conn:
            conn.execute('UPDATE orders SET actual_delivered=15')
            self.assertEqual({},confirmed_weights(conn))
        self.assertFalse(self.rows()[0]['confirmed'])

    def test_kg_file_has_alias_price_amount_and_original_stock_hold(self):
        self.save();did=self.draft()
        with server.db() as conn:
            validate_draft_export_stock(conn,did)
            batch=conn.execute('SELECT * FROM batches WHERE id=?',(self.bid,)).fetchone()
            output=server.export_invoices_zip(conn,batch,[])
            stored=conn.execute('SELECT * FROM outgoing_invoice_lines WHERE draft_id=?',(did,)).fetchone()
            self.assertEqual((14,'Gói',8000,112000),tuple(stored[k] for k in ('qty','unit','unit_price','amount')))
            self.assertEqual(14,conn.execute("SELECT SUM(qty_out) FROM inventory_transactions WHERE status='reserved'").fetchone()[0])
        with zipfile.ZipFile(output) as z:
            file=next(n for n in z.namelist() if n.endswith('.xlsx'))
            w=load_workbook(io.BytesIO(z.read(file)),data_only=True)
            values=[list(r) for r in w.active.values]
            r=next(r for r in values if 'HH-01' in r)
            for v in ('Nấm kim châm','Kg',2.8,40000,112000):self.assertIn(v,r)
            w.close()
        self.assertFalse(self.rows()[0]['editable'])
        self.assertEqual(409,self.save(3).status_code)

    def test_kg_does_not_remove_stock_shortage(self):
        with server.db() as conn:conn.execute('UPDATE inventory_transactions SET qty_in=0')
        self.save()
        p=self.client.get('/api/outgoing-invoices/unissued?to=2026-09-13').json
        self.assertEqual(0,sum(r['ready_qty'] for r in p['details']))
        self.assertEqual(14,sum(r['waiting_qty'] for r in p['details']))

    def test_partial_stock_converts_only_selected_packages(self):
        with server.db() as conn:conn.execute('UPDATE inventory_transactions SET qty_in=7')
        self.save();did=self.draft()
        with server.db() as conn:
            lines=invoice_rows(conn,conn.execute('SELECT * FROM outgoing_invoice_lines WHERE draft_id=?',(did,)))
            self.assertEqual((1.4,40000,56000,7),tuple(lines[0][k] for k in ('qty','unit_price','amount','stock_qty')))

    def test_two_orders_with_different_weights_keep_separate_prices(self):
        self.save()
        with server.db() as conn:
            bid,oids=support.OutgoingReadinessTests.add_batch(conn,'2026-09-02',[{'qty':2,'sell_price':8000}])
            conn.execute("UPDATE orders SET unit='Gói',tax='KKKNT' WHERE id=?",(oids[0],))
        self.assertEqual(200,self.save(1,next(r for r in self.rows() if r['order_id']==oids[0])).status_code)
        with server.db() as conn:
            for batch in (self.bid,bid):create_partial_outgoing_drafts(conn,batch,server.now_iso)
            ids=consolidate(conn,[self.bid,bid],'NT-A',invoice_tax_percent,server.now_iso())
            lines=invoice_rows(conn,conn.execute('SELECT * FROM outgoing_invoice_lines WHERE draft_id=?',(ids[0],)),freeze=True,timestamp=server.now_iso())
            self.assertEqual(2,len(lines))
            self.assertEqual([16000,40000],sorted(r['unit_price'] for r in lines))
            self.assertEqual(128000,sum(r['amount'] for r in lines))

    def test_main_download_is_idempotent_and_freezes_actual_kg(self):
        self.save()
        def download():
            r=self.client.post('/api/export/order-invoices',json={'contractor':'NT-A','from':'2026-09-01','to':'2026-09-13'})
            self.assertEqual(200,r.status_code,r.get_json(silent=True))
            with zipfile.ZipFile(io.BytesIO(r.data)) as z:
                w=load_workbook(io.BytesIO(z.read(next(n for n in z.namelist() if n.endswith('.xlsx')))),data_only=True)
                row=next(list(r) for r in w.active.values if 'HH-01' in r);w.close()
                self.assertEqual(['HH-01','Nấm kim châm','Kg',2.8,40000],row[:5])
        download()
        with server.db() as conn:before=conn.serialize()
        download()
        with server.db() as conn:self.assertEqual(before,conn.serialize())

    def test_minvoice_payload_uses_same_measured_weight_and_rounded_price(self):
        from .test_minvoice_client import RecordingMinvoiceClient
        self.save(3);did=self.draft()
        with server.db() as conn:
            conn.execute("""INSERT OR REPLACE INTO outgoing_buyer_profiles(contractor,display_name,legal_name,tax_code,address,email,updated_at)
                VALUES('NT-A','','CONG TY A','0201234567','Hai Phong','','now')""")
        original=server.app.config.get('MINVOICE_CLIENT_FACTORY')
        client=RecordingMinvoiceClient()
        server.app.config['MINVOICE_CLIENT_FACTORY']=lambda:client
        try:
            response=self.client.post(f'/api/minvoice/drafts/{did}',json={'series':'1C26TDP','dry_run':True})
            self.assertEqual(200,response.status_code,response.json)
            line=response.json['payload']['data'][0]['details'][0]['data'][0]
            self.assertEqual('Kg',line['inv_unitCode']);self.assertEqual(3,line['inv_quantity'])
            self.assertEqual(37333.333333,line['inv_unitPrice'])
            self.assertEqual(112000,response.json['payload']['data'][0]['inv_TotalAmountWithoutVat'])
            self.assertEqual([],client.calls)
        finally:
            if original is None:server.app.config.pop('MINVOICE_CLIENT_FACTORY',None)
            else:server.app.config['MINVOICE_CLIENT_FACTORY']=original

    def signed(self,kg=2.8):
        self.save(kg);did=self.draft()
        with server.db() as conn:
            invoice_rows(conn,conn.execute('SELECT * FROM outgoing_invoice_lines WHERE draft_id=?',(did,)),freeze=True,timestamp=server.now_iso())
            conn.execute("UPDATE outgoing_invoice_drafts SET status='issued',issued_invoice_series='1C26TDP',issued_invoice_number='001',issued_invoice_date='2026-09-13',buyer_tax_code_snapshot='MST-A' WHERE id=?",(did,))
            conn.execute("UPDATE inventory_transactions SET status='posted' WHERE source_type='OUTGOING_DRAFT'")
            sid=conn.execute("""INSERT INTO outgoing_source_invoices(tenant,source,identity_key,invoice_number,invoice_series,invoice_date,buyer_tax_code,
                source_status_class,sync_status,stock_status,synced_at,created_at,updated_at)
                VALUES('default','minvoice','weight-test','001','1C26TDP','2026-09-13','MST-A','issued','synced','ready','now','now','now')""").lastrowid
            item=conn.execute("""INSERT INTO outgoing_source_invoice_items(invoice_id,line_index,source_item_code,source_item_name,source_unit,qty,unit_price,
                amount,tax_rate,product_code,mapping_status,stock_qty,stock_unit_price,conversion_factor,inventory_eligible)
                VALUES(?,1,'HH-01','Nấm kim châm','Kg',?,?,112000,'KKKNT','HH-01','mapped',?,?,1,1)""",(sid,kg,112000/kg,kg,112000/kg)).lastrowid
        return sid,item,did

    def test_measured_one_to_one_does_not_teach_other_buyers_package_weight(self):
        from .invoice_output_mapping import _confirmed_names
        sid,item,did=self.signed(14)
        with server.db() as conn:
            post_output_invoice(conn,sid,confirmed=True,now_iso=server.now_iso)
            evidence=_confirmed_names(conn,'default',{('nấm kim châm','kg')})
            self.assertEqual([None],evidence[('nấm kim châm','kg')])

    def test_signed_invoice_posts_original_packages_once(self):
        sid,item,did=self.signed()
        with server.db() as conn:
            snap=validated_output_stock_snapshot(conn,item)
            self.assertEqual(14,snap['stock_qty']);self.assertEqual(5,snap['conversion_factor'])
            r=post_output_invoice(conn,sid,confirmed=True,now_iso=server.now_iso)
            self.assertEqual(1,r['new_inventory_lines'])
            self.assertEqual(-14,conn.execute("SELECT SUM(qty_delta) FROM invoice_inventory_ledger WHERE direction='output'").fetchone()[0])
            self.assertEqual((2.8,14,5),tuple(conn.execute('SELECT qty,stock_qty,conversion_factor FROM outgoing_source_invoice_items WHERE id=?',(item,)).fetchone()))
            self.assertEqual(6,canonical_available_stock(conn)['HH-01']['available_qty'])
            self.assertTrue(post_output_invoice(conn,sid,confirmed=True,now_iso=server.now_iso)['idempotent'])

    def test_signed_invoice_without_proven_matching_weight_cannot_use_factor_one(self):
        sid,item,did=self.signed()
        for column,value in [('qty',3),('amount',99999),('source_unit','Gói'),('source_item_name','Mặt hàng khác')]:
            with server.db() as conn:
                conn.execute('SAVEPOINT change_source')
                conn.execute('UPDATE outgoing_source_invoice_items SET '+column+'=? WHERE id=?',(value,item))
                # A source in the original stock unit follows the existing policy;
                # all mismatched Kg exports must block instead of subtracting Kg as packages.
                with self.assertRaises(InvoiceMappingError):validated_output_stock_snapshot(conn,item)
                conn.execute('ROLLBACK TO change_source');conn.execute('RELEASE change_source')
        with server.db() as conn:
            conn.execute("UPDATE outgoing_invoice_drafts SET issued_invoice_number='002' WHERE id=?",(did,))
            with self.assertRaises(InvoiceMappingError):validated_output_stock_snapshot(conn,item)


if __name__=='__main__':unittest.main()
