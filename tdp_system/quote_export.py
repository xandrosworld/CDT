"""Golden-shaped quotation workbook builders.

Toyota remains the visual golden, but the builder is contractor-neutral.
Prices are supplied by the caller from exactly one confirmed period version or
one selected daily batch; this module never looks up or mixes price groups.
"""

from __future__ import annotations

import math
import re
from collections import OrderedDict
from copy import copy
from datetime import datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


COMPANY_NAME = "CÔNG TY TNHH THỰC PHẨM THÀNH ĐẠT PHÁT"
COMPANY_ADDRESS = "Địa chỉ: Số 112, ngõ 366, đường Hùng Vương, Sở Dầu, Hồng Bàng, Hải Phòng"
COMPANY_EMAIL = "Email: thanhdatphatnct@gmail.com"
COMPANY_PHONE = "ĐT: 0904.495.655"
TOYOTA_RECIPIENT = "CÔNG TY TNHH TOYOTA NANKAI HẢI PHÒNG"

# Source-backed display names.  These are deliberately separate from the
# contractor codes used for price selection.  A deployment can override one
# recipient through ``settings.quote_recipient_<CODE>`` without changing the
# price group or the historical quotation version.
QUOTE_RECIPIENTS = {
    "ATV": "CÔNG TY CỔ PHẦN SUẤT ĂN CÔNG NGHIỆP ATV",
    "BIADAUVOI": "NHÀ HÀNG BIA ĐẦU VÒI",
    "GIANHAPTAY": "GIANHAPTAY",
    "HATRAN": "HÀ TRÂN",
    "NGUYENGIA": "BẾP NGUYỄN GIA",
    "NHUAHAIPHONG": "CÔNG TY CỔ PHẦN NHỰA VÀ CƠ KHÍ HẢI PHÒNG",
    "SUPPY": "CÔNG TY TNNN SUPPLY",
    "TOYOTA": TOYOTA_RECIPIENT,
    "YLKHAN": "BẾP YLKHAN",
}

QUOTE_GROUPS = OrderedDict((
    ("A", "THỊT HEO"),
    ("B", "GIA CẦM"),
    ("C", "BÒ, BÊ, DÊ"),
    ("D", "THỦY SẢN"),
    ("E", "HẢI SẢN"),
    ("F", "GIÒ CHẢ"),
    ("G", "BÚN-BÁNH"),
    ("H", "TRỨNG - ĐẬU"),
    ("I", "RAU CỦ"),
    ("J", "HOA QUẢ"),
    ("K", "ĐÔNG LẠNH"),
    ("L", "GẠO"),
    ("M", "GIA VỊ-ĐỒ KHÔ"),
    ("N", "ĐỒ LỄ-BÁNH SỮA-NƯỚC NGỌT"),
    ("O", "CÔNG CỤ DỤNG CỤ"),
    ("P", "CHẤT TẨY RỬA-GIẤY CÁC LOẠI"),
))

PRICE_NUMBER_FORMAT = '#,##0'
YELLOW = "FFFF00"
GREEN = "92D050"
BLACK = "000000"


class QuoteExportError(ValueError):
    def __init__(self, message: str, *, code: str = "invalid_quote_export", status: int = 409):
        super().__init__(message)
        self.code = code
        self.status = status


def _period_parts(period: Any) -> tuple[str, int, int]:
    text = str(period or "").strip()
    if not re.fullmatch(r"\d{4}-\d{2}", text):
        raise QuoteExportError("Kỳ báo giá phải có dạng YYYY-MM", code="invalid_period", status=400)
    try:
        value = datetime.strptime(text + "-01", "%Y-%m-%d")
    except ValueError:
        raise QuoteExportError("Kỳ báo giá không hợp lệ", code="invalid_period", status=400) from None
    return text, value.month, value.year


def _contractor_code(value: Any) -> str:
    code = _literal(value, "mã nhà thầu", required=True).upper()
    if not re.fullmatch(r"[A-Z0-9_-]+", code):
        raise QuoteExportError(
            "Mã nhà thầu chỉ được chứa chữ Latin, số, gạch ngang hoặc gạch dưới",
            code="invalid_contractor",
            status=400,
        )
    return code


def quote_recipient(
    contractor: Any, *, configured: Any = "", fallback_name: Any = "",
) -> str:
    """Resolve a display recipient without coupling it to price selection."""
    code = _contractor_code(contractor)
    return (
        _literal(configured, "tên đơn vị nhận báo giá")
        or QUOTE_RECIPIENTS.get(code)
        or _literal(fallback_name, "tên nhà thầu")
        or code
    )


def _daily_source_parts(source: dict[str, Any]) -> tuple[dict[str, Any], int, int, str]:
    if not isinstance(source, dict):
        raise QuoteExportError(
            "Nguồn báo giá theo ngày không hợp lệ",
            code="invalid_daily_source",
            status=400,
        )
    try:
        batch_id = int(source["batch_id"])
    except (KeyError, TypeError, ValueError):
        raise QuoteExportError(
            "Thiếu phiên đơn nguồn cho báo giá theo ngày",
            code="daily_batch_required",
        ) from None
    if batch_id <= 0:
        raise QuoteExportError(
            "Phiên đơn nguồn cho báo giá theo ngày không hợp lệ",
            code="daily_batch_required",
        )
    work_date = _literal(source.get("work_date"), "ngày của phiên đơn", required=True)
    try:
        parsed = datetime.strptime(work_date, "%Y-%m-%d")
    except ValueError:
        raise QuoteExportError(
            "Ngày của phiên đơn phải có dạng YYYY-MM-DD",
            code="invalid_daily_source",
            status=400,
        ) from None
    normalized = {
        "batch_id": batch_id,
        "work_date": parsed.strftime("%Y-%m-%d"),
        "source_name": _literal(source.get("source_name"), "tên nguồn phiên đơn"),
    }
    return normalized, parsed.month, parsed.year, parsed.strftime("%d/%m/%Y")


def _literal(value: Any, field: str, *, required: bool = False) -> str:
    text = re.sub(r"\s+", " ", str(value or "").strip())
    if required and not text:
        raise QuoteExportError(f"Thiếu {field} trong dòng báo giá")
    if text.startswith("="):
        raise QuoteExportError(f"{field} không được là công thức", code="unsafe_formula")
    return text


def _price(value: Any) -> int:
    if isinstance(value, bool):
        raise QuoteExportError("Giá bán không hợp lệ")
    try:
        number = Decimal(str(value))
    except (InvalidOperation, ValueError):
        raise QuoteExportError("Giá bán không hợp lệ") from None
    if not number.is_finite() or number < 0:
        raise QuoteExportError("Giá bán phải là số VND hữu hạn không âm")
    return int(number.quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def _tax_value(value: Any) -> tuple[Any, str]:
    if isinstance(value, bool):
        return _literal(value, "thuế"), "General"
    if isinstance(value, (int, float)) and math.isfinite(float(value)):
        number = float(value)
        if 0 <= number <= 1:
            return number, "0%"
    text = _literal(value, "thuế")
    if not text:
        return "", "General"
    candidate = text.replace(" ", "").replace(",", ".")
    try:
        if candidate.endswith("%"):
            number = float(candidate[:-1]) / 100
        else:
            number = float(candidate)
    except ValueError:
        return text, "General"
    if math.isfinite(number) and 0 <= number <= 1:
        return number, "0%"
    return text, "General"


def _normalized_rows(rows: Iterable[dict[str, Any]]) -> list[dict[str, Any]]:
    output: list[dict[str, Any]] = []
    seen: set[str] = set()
    for item in rows:
        if item.get("exportable") is not True:
            continue
        state = _literal(item.get("price_state"), "trạng thái giá").lower()
        if state not in {"numeric", "zero"}:
            continue
        code = _literal(item.get("product_code"), "mã hàng", required=True).upper()
        if code in seen:
            raise QuoteExportError(
                f"Mã {code} còn xuất hiện nhiều hơn một dòng",
                code="quote_duplicate_conflict",
            )
        seen.add(code)
        price = _price(item.get("sell_price"))
        if state == "zero" and price != 0:
            raise QuoteExportError(f"Trạng thái giá 0 của mã {code} không khớp giá trị")
        group_code = code[:1] if code else ""
        output.append({
            "product_code": code,
            "product_name": _literal(item.get("product_name"), "tên hàng", required=True),
            "unit": _literal(item.get("unit"), "ĐVT"),
            "tax": _tax_value(item.get("tax")),
            "sell_price": price,
            "group_code": group_code,
        })
    rank = {code: index for index, code in enumerate(QUOTE_GROUPS)}
    output.sort(key=lambda item: (
        rank.get(item["group_code"], len(rank)), item["group_code"],
        item["product_code"].casefold(), item["product_name"].casefold(),
    ))
    return output


def _thin_border() -> Border:
    side = Side(style="thin", color=BLACK)
    return Border(left=side, right=side, top=side, bottom=side)


def _style_header(sheet, row: int) -> None:
    fill = PatternFill(fill_type=None)
    border = _thin_border()
    alignments = (None, "center", "left", "center", "center", "center")
    for column, alignment in enumerate(alignments, 1):
        cell = sheet.cell(row, column)
        cell.font = Font(name="Times New Roman", size=10, bold=True)
        cell.fill = fill
        cell.border = border
        cell.alignment = Alignment(
            horizontal=alignment, vertical="center", shrink_to_fit=True,
        )


def _style_group(sheet, row: int) -> None:
    fill = PatternFill(fill_type=None)
    border = _thin_border()
    for column in range(1, 7):
        cell = sheet.cell(row, column)
        cell.font = Font(name="Times New Roman", size=14, bold=(column == 3))
        cell.fill = fill
        cell.border = border
        cell.alignment = Alignment(
            horizontal="center" if column != 3 else "left", vertical="center",
            shrink_to_fit=(column != 3),
        )
    sheet.cell(row, 5).number_format = PRICE_NUMBER_FORMAT
    sheet.cell(row, 6).number_format = "0%"


def _style_item(sheet, row: int) -> None:
    border = _thin_border()
    for column in range(1, 7):
        cell = sheet.cell(row, column)
        cell.font = Font(name="Times New Roman", size=14 if column == 1 else 13)
        cell.border = border
        cell.alignment = Alignment(
            horizontal=("center" if column in {1, 4, 6} else "right" if column == 5 else "left"),
            vertical="center", shrink_to_fit=(column in {1,5}),
        )
    sheet.cell(row, 5).number_format = PRICE_NUMBER_FORMAT


def build_contractor_quote_workbook(
    rows: Iterable[dict[str, Any]], *, contractor: str, recipient: str,
    period: str = "", version: dict[str, Any] | None = None,
    daily_source: dict[str, Any] | None = None,
) -> Workbook:
    """Build one isolated contractor quotation in the Toyota golden shape."""
    contractor = _contractor_code(contractor)
    recipient = _literal(recipient, "tên đơn vị nhận báo giá", required=True)
    if daily_source is not None:
        if version:
            raise QuoteExportError(
                "Báo giá không thể đồng thời lấy từ kỳ và phiên đơn theo ngày",
                code="ambiguous_quote_source",
                status=400,
            )
        daily_source, month, year, display_date = _daily_source_parts(daily_source)
        period = daily_source["work_date"][:7]
        version_no = None
        source_hash = ""
        title_text = f"BẢNG BÁO GIÁ NGÀY {display_date}"
        source_text = f"{contractor} · ngày {display_date} · phiên đơn {daily_source['batch_id']}"
        subject_text = source_text
        keyword_text = f"{contractor},{daily_source['work_date']},batch-{daily_source['batch_id']}"
    else:
        period, month, year = _period_parts(period)
        if not version:
            raise QuoteExportError(
                f"Kỳ báo giá {contractor} chưa có phiên bản đã xác nhận",
                code="quote_period_not_confirmed",
            )
        try:
            version_no = int(version["version_no"])
        except (KeyError, TypeError, ValueError):
            raise QuoteExportError("Thiếu phiên bản nguồn báo giá đã xác nhận") from None
        if version_no <= 0:
            raise QuoteExportError("Phiên bản nguồn báo giá không hợp lệ")
        source_hash = _literal(version.get("source_hash"), "hash nguồn", required=True).upper()
        title_text = f"BẢNG BÁO GIÁ THÁNG {month:02d} - NĂM {year}"
        source_text = f"{contractor} · kỳ {period} · phiên bản {version_no} · {source_hash[:16]}"
        subject_text = f"{contractor} · kỳ {period} · phiên bản {version_no}"
        keyword_text = f"{contractor},{period},version-{version_no},{source_hash}"
    items = _normalized_rows(rows)

    workbook = Workbook()
    normal_font = Font(name="Times New Roman", size=12, charset=134)
    workbook._fonts[0] = normal_font
    workbook._named_styles["Normal"].font = normal_font
    workbook.properties.creator = COMPANY_NAME
    workbook.properties.title = f"Báo giá {contractor} tháng {month:02d}/{year}"
    workbook.properties.subject = subject_text
    workbook.properties.keywords = keyword_text
    sheet = workbook.active
    sheet.title = "all"
    sheet.sheet_view.zoomScale = 115
    sheet.sheet_format.defaultColWidth = 6.375
    sheet.sheet_format.defaultRowHeight = 16.9
    sheet.sheet_format.customHeight = True

    widths = {"A": 5.125, "B": 9.625, "C": 42.825, "D": 9.125, "E": 11.625, "F": 9.75, "G": 6.375}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    sheet.merge_cells("A1:F1")
    sheet["A1"] = COMPANY_NAME
    sheet["A1"].font = Font(name="Times New Roman", size=12, bold=True)
    sheet["A1"].alignment = Alignment(horizontal="left")
    sheet["A2"] = COMPANY_ADDRESS
    sheet["A2"].font = Font(name="Times New Roman", size=12, bold=True)
    sheet.merge_cells("A3:C3")
    sheet["A3"] = COMPANY_EMAIL
    sheet["A3"].font = Font(name="Times New Roman", size=14)
    sheet["A3"].alignment = Alignment(horizontal="left", vertical="center")
    sheet.merge_cells("A4:F4")
    sheet["A4"] = COMPANY_PHONE
    sheet["A4"].font = Font(name="Times New Roman", size=14, bold=True)
    sheet.merge_cells("A5:F5")
    sheet["A5"] = title_text
    sheet["A5"].font = Font(name="Times New Roman", size=16, bold=True)
    sheet["A5"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[5].height = 25
    sheet.merge_cells("A6:F6")
    sheet["A6"] = f"KÍNH GỬI: {recipient}"
    sheet["A6"].font = Font(name="Times New Roman", size=11, bold=True, italic=True)
    sheet["A6"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[7].height = 19.2

    headings = ("STT", "MÃ", "TÊN THÀNH ĐẠT PHÁT", "ĐVT", "Giá chưa VAT", "Thuế")
    for column, value in enumerate(headings, 1):
        sheet.cell(8, column, value)
    _style_header(sheet, 8)

    current_group = None
    row_number = 9
    serial = 0
    for item in items:
        group_code = item["group_code"]
        if group_code != current_group:
            current_group = group_code
            sheet.cell(row_number, 3, QUOTE_GROUPS.get(group_code, f"NHÓM {group_code or 'KHÁC'}"))
            _style_group(sheet, row_number)
            row_number += 1
        serial += 1
        tax_value, tax_format = item["tax"]
        values = (
            serial, item["product_code"], item["product_name"], item["unit"],
            item["sell_price"], tax_value,
        )
        for column, value in enumerate(values, 1):
            sheet.cell(row_number, column, value)
        _style_item(sheet, row_number)
        sheet.cell(row_number, 6).number_format = tax_format
        row_number += 1

    separator_row = row_number
    for body_row in range(8, separator_row):
        sheet.row_dimensions[body_row].height = max(22, 18 * math.ceil(len(str(sheet.cell(body_row, 3).value or '')) / 38))
        alignment = copy(sheet.cell(body_row, 3).alignment)
        alignment.wrap_text = True
        sheet.cell(body_row, 3).alignment = alignment
    sheet.row_dimensions[separator_row].height = 9
    sheet.cell(separator_row, 1).border = Border(left=Side(style="thin", color=BLACK))
    note_row = separator_row + 1
    sheet.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=3)
    sheet.cell(note_row, 1, "Báo giá trên chưa bao gồm VAT!")
    sheet.cell(note_row, 1).font = Font(name="Times New Roman", size=13, italic=True)
    sheet.cell(note_row, 1).alignment = Alignment(horizontal="left", shrink_to_fit=True)
    sheet.cell(note_row, 1).border = Border(left=Side(style="thin", color=BLACK))
    # Provenance belongs in workbook properties (subject/keywords above), not
    # on the customer's quotation artwork. The golden keeps D:F of this row
    # blank, so no technical hash/version text is allowed to appear on paper.
    signature_row = note_row + 2
    sheet.cell(signature_row, 3, "XÁC NHẬN CỦA BÊN BÁN")
    sheet.cell(signature_row, 3).font = Font(name="Times New Roman", size=13, bold=True)
    sheet.cell(signature_row, 3).alignment = Alignment(horizontal="left")
    final_row = signature_row + 1
    for column in range(1, 7):
        sheet.cell(final_row, column).font = Font(name="Times New Roman", size=12)

    sheet.print_title_rows = "8:8"
    sheet.auto_filter.ref = f"A8:F{signature_row}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = "9"
    sheet.page_setup.scale = 87
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins.left = 0.2
    sheet.page_margins.right = 0.2
    sheet.page_margins.top = 0.23
    sheet.page_margins.bottom = 0.24
    sheet.page_margins.header = 0.3
    sheet.page_margins.footer = 0.3
    return workbook


def build_toyota_quote_workbook(
    rows: Iterable[dict[str, Any]], *, period: str, version: dict[str, Any],
) -> Workbook:
    """Compatibility wrapper retaining the verified Toyota API."""
    return build_contractor_quote_workbook(
        rows,
        contractor="TOYOTA",
        recipient=TOYOTA_RECIPIENT,
        period=period,
        version=version,
    )


def contractor_quote_filename(
    contractor: str, *, period: str = "", version_no: int | None = None,
    daily_source: dict[str, Any] | None = None,
) -> str:
    contractor = _contractor_code(contractor)
    if daily_source is not None:
        source, _month, _year, _display_date = _daily_source_parts(daily_source)
        return f"BAO_GIA_{contractor}_{source['work_date']}_PHIEN_{source['batch_id']}.xlsx"
    _, month, year = _period_parts(period)
    try:
        version_no = int(version_no)
    except (TypeError, ValueError):
        raise QuoteExportError("Thiếu phiên bản nguồn báo giá đã xác nhận") from None
    if version_no <= 0:
        raise QuoteExportError("Phiên bản nguồn báo giá không hợp lệ")
    return f"BAO_GIA_{contractor}_T{month:02d}-{year}_V{version_no}.xlsx"


def toyota_quote_filename(period: str, version_no: int) -> str:
    return contractor_quote_filename("TOYOTA", period=period, version_no=version_no)
