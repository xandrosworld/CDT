import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from . import server
from .daily_workbook_import import _header_map, analyze_daily_workbook
from .contract_modules import purchase_canonical_header_fields, parse_canonical_purchase_workbook
from .customer_purchase_layout import customer_purchase_fields


class CustomerPurchaseLayoutTests(unittest.TestCase):
    def test_explicit_supplier_wins_over_warning_column(self):
        book = Workbook()
        sheet = book.active
        sheet.append(['Chọn NCC', 'NCC', 'Mã hàng', 'Mã bếp', 'Tên hàng', 'Số lượng', 'Nhà cung cấp'])
        sheet.append(['Chọn NCC', 'kho', 'L000004', 'LSVINA', 'Gạo nếp', .54, 'ánh'])
        self.assertEqual(server.detect_header(sheet)[1]['supplier'], 2)
        self.assertEqual(_header_map(sheet)[1]['supplier'], 2)
        # A legacy workbook using only Chọn NCC still remains readable.
        sheet.cell(1, 2, 'Khác'); sheet.cell(1, 7, 'Khác')
        self.assertEqual(server.detect_header(sheet)[1]['supplier'], 1)
        book.close()

    @staticmethod
    def workbook():
        book = Workbook()
        sheet = book.active
        sheet.title = 'đặt hàng'
        sheet.append(['ĐẶT HÀNG'])
        sheet.append(['Mã tham chiếu', 'Mã hàng', 'Mã bếp', 'Ngày', 'Tên hàng', 'Số lượng',
                      'ĐVT', 'NCC', 'Ghi chú đặt hàng', 'Đơn giá', 'hỏng', 'thêm', 'giảm',
                      'thiếu', 'SL \nthực té', 14040])
        sheet.append(['L000004kho', 'L000004', 'LSVINA', '03.09.2026', 'Gạo nếp', .54,
                      'Kg', 'kho', '', 26000, 0, 0, 0, 0, .54, 14040])
        return book

    def test_verified_layout_reaches_purchase_parser_and_preserves_quantity(self):
        with tempfile.TemporaryDirectory() as root:
            root = Path(root)
            book = self.workbook()
            path = root / 'customer.xlsx'; book.save(path)
            analysis = analyze_daily_workbook(path)
            self.assertEqual(len(analysis['purchaseSheets']), 1)
            fields = purchase_canonical_header_fields([c.value for c in book.active[2]])
            self.assertEqual((fields['buy_price'], fields['amount']), (10, 16))
            with patch.object(server, 'DB_PATH', root / 'test.sqlite3'), \
                    patch.object(server, 'DATA_DIR', root), patch.object(server, 'auto_backup'):
                server.init_database(sync_master=False)
                with server.db() as conn:
                    batch = conn.execute("INSERT INTO batches(source_name,work_date,created_at,status) VALUES('test','2026-09-03','now','draft')").lastrowid
                    formulas = self.workbook()
                    formulas.active['O3'] = '=F3+L3-K3-N3-M3'
                    formulas.active['P3'] = '=IFERROR(J3*O3,0)'
                    result = parse_canonical_purchase_workbook(conn, book, batch, {}, formulas)
                    row = result['items'][0]
                    self.assertEqual(row['errors'], [])
                    self.assertEqual((row['supplier'], row['actual_qty'], row['amount']), ('kho', .54, 14040))
                    formulas.active['O3'] = '=F3+L3+K3-N3-M3'
                    self.assertTrue(parse_canonical_purchase_workbook(conn, book, batch, {}, formulas)['items'][0]['errors'])
                    formulas.close()
            book.close()

    def test_unrelated_numeric_column_is_not_guessed(self):
        book = self.workbook()
        values = [c.value for c in book.active[2]]
        values[13] = 'Khác'
        self.assertEqual(customer_purchase_fields(values), {})
        book.close()


if __name__ == '__main__': unittest.main()
