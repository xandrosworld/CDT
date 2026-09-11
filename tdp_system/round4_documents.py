"""Selected documents and issued-invoice paperwork, without business writes."""
from __future__ import annotations

import hashlib
import io
import re
import zipfile
from datetime import datetime

from flask import jsonify, request, send_file
from openpyxl import load_workbook

try:
    from .document_preview import create_snapshot, snapshot_files, PDF_LOCK
    from .excel_print_renderer import build_excel_pdf_bundle, ExcelPrintError, is_receipt_sheet
    from .invoice_payment_scope import issued_invoice_payment_scope
    from .invoice_payment_documents import invoice_payment_request_workbook
    from .contract_modules import invoice_delivery_statement_scope_workbook
except ImportError:
    from document_preview import create_snapshot, snapshot_files, PDF_LOCK
    from excel_print_renderer import build_excel_pdf_bundle, ExcelPrintError, is_receipt_sheet
    from invoice_payment_scope import issued_invoice_payment_scope
    from invoice_payment_documents import invoice_payment_request_workbook
    from contract_modules import invoice_delivery_statement_scope_workbook


def positive_id(value):
    if isinstance(value, bool) or not re.fullmatch(r'[1-9][0-9]*', str(value)):
        raise ValueError('Mã phiếu không hợp lệ')
    return int(value)


def selected_workbooks(conn, body, ctx):
    """Same exporters as Excel download; explicit selection, fail all on bad rows."""
    kind = body.get('kind')
    result = []
    try:
        if kind == 'quote':
            contractor = str(body.get('contractor') or '').strip().upper()
            period = ctx['quote_period_arg'](body.get('period'))
            version_id = ctx['quote_version_arg'](body.get('version_id'))
            batch_id = positive_id(body['batch_id']) if body.get('batch_id') else None
            mode, rows, meta = ctx['get_quote_rows'](conn, contractor, batch_id, period, version_id=version_id)
            if meta['conflicts']:
                raise ValueError('Báo giá còn dòng trùng xung đột; cần sửa trước khi xem bản gửi khách')
            if (mode == 'daily' and not meta['daily_source']) or (mode != 'daily' and not meta['version']):
                raise ValueError('Chưa có báo giá đã xác nhận cho kỳ được chọn')
            recipient = ctx['quote_recipient'](contractor,
                configured=ctx['setting_get'](conn, f'quote_recipient_{contractor}', ''),
                fallback_name=meta['contractor_name'])
            result.append(('Bao_gia.xlsx', ctx['build_contractor_quote_workbook'](rows,
                contractor=contractor, recipient=recipient, period=meta['period'],
                version=meta['version'], daily_source=meta['daily_source'] if mode == 'daily' else None)))
        elif kind == 'quotes':
            period = ctx['quote_period_arg'](body.get('period'))
            if not period: raise ValueError('Hãy chọn kỳ báo giá')
            version_id = ctx['quote_version_arg'](body.get('version_id'))
            version = conn.execute("""SELECT id FROM quote_versions WHERE effective_period=? AND status='confirmed'
                AND (? IS NULL OR id=?) ORDER BY version_no DESC LIMIT 1""",(period,version_id,version_id)).fetchone()
            if not version: raise ValueError('Kỳ này chưa có báo giá đã xác nhận')
            batch_id = positive_id(body['batch_id']) if body.get('batch_id') else None
            if batch_id:
                batch = conn.execute('SELECT work_date FROM batches WHERE id=?',(batch_id,)).fetchone()
                if not batch or batch['work_date'][:7] != period:
                    raise ValueError('Đơn theo ngày không thuộc kỳ báo giá đang xem')
            for contractor in conn.execute("SELECT code,pricing_mode FROM contractors ORDER BY code"):
                daily = contractor['pricing_mode']=='daily'
                if daily and not batch_id: continue
                mode,rows,meta=ctx['get_quote_rows'](conn,contractor['code'],batch_id if daily else None,period,
                    version_id=None if daily else int(version['id']))
                if meta['conflicts']: raise ValueError('Báo giá '+contractor['code']+' còn dòng trùng xung đột')
                if not any(row.get('exportable') for row in rows): continue
                name,payload=ctx['quote_workbook_payload'](conn,contractor['code'],rows,meta,mode)
                result.append((name,load_workbook(io.BytesIO(payload))))
        elif kind == 'payment':
            scope = issued_invoice_payment_scope(conn, body.get('contractor'), body.get('from'), body.get('to'))
            if body.get('scope_id') and body['scope_id'] != scope['scope_id']:
                raise ValueError('Phạm vi hóa đơn đã đổi; hãy kiểm tra lại danh sách hóa đơn')
            statement_name = 'Bang_ke_hoa_don_VAT.xlsx' if scope.get('statement_kind') == 'invoices' else 'Bang_tong_hop_giao_nhan.xlsx'
            result.append((statement_name, invoice_delivery_statement_scope_workbook(scope)))
            result.append(('De_nghi_thanh_toan.xlsx', invoice_payment_request_workbook(scope,
                issue_date=body.get('issue_date') or scope['date_to'],
                request_number=body.get('request_number') or '……/CV/ĐNTT',
                contract_no=body.get('contract_no') or '', contract_date=body.get('contract_date') or '')))
        else:
            builders = {'deliveries': 'export_deliveries', 'suppliers': 'export_supplier_orders',
                        'purchases': 'export_purchase_documents', 'report': 'export_report',
                        'outgoing-statement': 'export_outgoing_statement'}
            if kind not in builders:
                raise ValueError('Loại chứng từ không hợp lệ')
            selections = body.get('selections')
            if not isinstance(selections, list) or not selections or len(selections) > 300:
                raise ValueError('Hãy chọn từ 1 đến 300 phiếu')
            grouped = {}
            for selected in selections:
                if not isinstance(selected, dict):
                    raise ValueError('Danh sách phiếu không hợp lệ')
                batch_id = positive_id(selected.get('batch_id'))
                grouped.setdefault(batch_id, set())
                kitchen = str(selected.get('kitchen') or '').strip().upper()
                if kind == 'deliveries' and not kitchen:
                    raise ValueError('Cần chọn rõ bếp của từng phiếu giao')
                if kitchen:
                    grouped[batch_id].add(kitchen)
            if len(grouped) > 100:
                raise ValueError('Mỗi lần chọn tối đa 100 ngày đơn')
            periods = set()
            for batch_id, kitchens in grouped.items():
                batch, orders = ctx['require_batch'](conn, batch_id)
                if kind == 'deliveries':
                    available = {str(o['kitchen']).strip().upper() for o in orders if ctx['net_delivered'](o) > 0}
                    if not kitchens <= available:
                        raise ValueError('Phiếu giao đã thay đổi hoặc không còn hàng; hãy tải lại danh sách')
                    workbook = ctx[builders[kind]](conn, batch, orders, kitchens)
                elif kind == 'report':
                    period = batch['work_date'][:7]
                    if period in periods: continue
                    periods.add(period)
                    workbook = ctx[builders[kind]](conn, {'id':-1, 'status':'approved', 'work_date':period+'-01'}, [])
                else:
                    workbook = ctx[builders[kind]](conn, batch, orders)
                result.append((f'{kind}_{batch["work_date"]}_{batch_id}.xlsx', workbook))
        return result
    except Exception:
        for _, workbook in result: workbook.close()
        raise


def snapshot_selection(root, token, raw):
    directory, files = snapshot_files(root, token)
    sheets, books = [], []
    try:
        for file_index, (name, path) in enumerate(files):
            workbook = load_workbook(path)
            books.append((name, workbook))
            sheets.extend((file_index, ws.title) for ws in workbook if ws.sheet_state == 'visible')
        if raw is None:
            indices = set(range(len(sheets)))
        elif not re.fullmatch(r'\d+(?:,\d+)*', raw):
            raise ValueError('Hãy chọn ít nhất một phiếu')
        else:
            indices = {int(value) for value in raw.split(',')}
        if not indices or max(indices) >= len(sheets):
            raise ValueError('Có phiếu không thuộc bản xem đang mở')
        selected = {sheets[i] for i in indices}
        result = []
        for file_index, (name, workbook) in enumerate(books):
            included = any(owner == file_index for owner, _ in selected)
            for ws in list(workbook):
                # Preserve hidden reconciliation/evidence sheets in Excel.
                # They stay hidden and the PDF renderer never prints them.
                if ws.sheet_state == 'visible' and (file_index, ws.title) not in selected:
                    workbook.remove(ws)
            if included:
                result.append((name, workbook))
            else:
                workbook.close()
        return directory, result, ','.join(map(str, sorted(indices)))
    except Exception:
        for _, workbook in books: workbook.close()
        raise


def register_document_routes(app, context_factory):
    def root(): return context_factory()['DATA_DIR'] / 'document_previews'

    @app.get('/api/documents/list')
    def document_list():
        try:
            start, end = request.args.get('from',''), request.args.get('to','')
            for value in (start,end):
                if not re.fullmatch(r'\d{4}-\d{2}-\d{2}',value): raise ValueError('Hãy chọn đủ từ ngày và đến ngày')
                datetime.strptime(value,'%Y-%m-%d')
            if start > end: raise ValueError('Từ ngày phải nhỏ hơn hoặc bằng đến ngày')
            kind = request.args.get('kind','deliveries')
            if kind not in {'deliveries','suppliers','purchases','report'}: raise ValueError('Loại giấy tờ không hợp lệ')
            customer = request.args.get('customer','').strip().upper()
            rows=[]
            with context_factory()['db']() as conn:
                conn.execute('PRAGMA query_only=ON'); conn.execute('BEGIN')
                batches=conn.execute("""SELECT b.*,
                    (SELECT count(*) FROM orders o WHERE o.batch_id=b.id) line_count
                    FROM batches b WHERE b.work_date BETWEEN ? AND ? ORDER BY b.work_date,b.id""",(start,end)).fetchall()
                for source in batches:
                    batch=dict(source)
                    public={k:batch.get(k) for k in ('id','work_date','status','source_name','line_count')}
                    if kind=='deliveries':
                        notes=conn.execute("""SELECT o.kitchen code,COALESCE(k.name,o.kitchen) name,count(*) line_count
                            FROM orders o LEFT JOIN kitchens k ON k.code=o.kitchen
                            WHERE o.batch_id=? AND COALESCE(o.actual_delivered,0)-COALESCE(o.customer_return_qty,0)>0
                            AND (?='' OR upper(o.kitchen)=?) GROUP BY o.kitchen ORDER BY o.kitchen""",(batch['id'],customer,customer))
                        for note in notes:
                            rows.append({'key':f'delivery:{batch["id"]}:{note["code"]}', 'batch_id':batch['id'],
                                'kitchen':note['code'],'title':note['name'],'subtitle':batch['work_date'],
                                'line_count':note['line_count'],'batch':public})
                    elif batch['line_count']:
                        rows.append({'key':f'batch:{batch["id"]}', 'batch_id':batch['id'],
                            'title':batch['work_date'],'subtitle':batch['source_name'],
                            'line_count':batch['line_count'],'batch':public})
            return jsonify(ok=True,rows=rows,total=len(rows))
        except ValueError as exc:
            return jsonify(ok=False,error=str(exc)),400

    @app.post('/api/documents/preview')
    def document_preview():
        books = []
        try:
            body = request.get_json(silent=True)
            if not isinstance(body, dict): raise ValueError('Yêu cầu xem chứng từ không hợp lệ')
            ctx = context_factory()
            with ctx['db']() as conn:
                conn.execute('PRAGMA query_only=ON')
                conn.execute('BEGIN')
                books = selected_workbooks(conn, body, ctx)
            return jsonify(create_snapshot(root(), books))
        except (ValueError, RuntimeError) as exc:
            return jsonify(ok=False, error=str(exc), code=getattr(exc, 'code', 'document_preview_invalid')), getattr(exc, 'status', 422)
        finally:
            for _, workbook in books: workbook.close()

    @app.get('/api/documents/<token>/<output>')
    def document_snapshot_download(token, output):
        books = []
        try:
            if output not in {'excel', 'pdf'}:
                return jsonify(ok=False, error='Định dạng không hợp lệ'), 404
            directory, books, selection = snapshot_selection(root(), token, request.args.get('sheets'))
            paper = request.args.get('paper', 'A4').upper()
            if paper not in {'A4', 'A5'}:
                raise ValueError('Chọn khổ giấy A4 hoặc A5')
            if paper == 'A5':
                if any(not is_receipt_sheet(sheet.title) for _, book in books for sheet in book if sheet.sheet_state == 'visible'):
                    raise ValueError('Chọn riêng các biên nhận để in A5. Bảng kê tổng in A4.')
                for _, book in books:
                    for sheet in book:
                        sheet.page_setup.paperSize = sheet.PAPERSIZE_A5
                        sheet.page_setup.fitToWidth = 1
                        sheet.page_setup.fitToHeight = 1
                        sheet.page_margins.left = sheet.page_margins.right = 0.2
                        sheet.page_margins.top = sheet.page_margins.bottom = 0.2
                        sheet.page_margins.header = sheet.page_margins.footer = 0.1
            if output == 'excel':
                if len(books) == 1:
                    stream = io.BytesIO()
                    books[0][1].save(stream)
                    stream.seek(0)
                    return send_file(stream, as_attachment=True, download_name=books[0][0],
                        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
                stream = io.BytesIO()
                with zipfile.ZipFile(stream, 'w', zipfile.ZIP_DEFLATED) as archive:
                    for name, workbook in books:
                        payload = io.BytesIO()
                        workbook.save(payload)
                        archive.writestr(name, payload.getvalue())
                stream.seek(0)
                return send_file(stream, as_attachment=True, download_name='Chung_tu_da_chon.zip', mimetype='application/zip')
            sides = request.args.get('sides', 'duplex')
            if sides not in {'simplex', 'duplex'}:
                raise ValueError('Chọn cách in một mặt hoặc hai mặt')
            digest = hashlib.sha256(('sheet-scope-v4:' + paper + ':' + sides + ':' + selection).encode()).hexdigest()[:20]
            pdf = directory / f'{digest}.pdf'
            with PDF_LOCK:
                if not pdf.exists():
                    sources = []
                    for index, (_, workbook) in enumerate(books):
                        path = directory / f'print_{digest}_{index}.xlsx'
                        workbook.save(path)
                        sources.append({'path':path, 'document_type':'selected', 'title':'Chứng từ đã chọn'})
                    build_excel_pdf_bundle(sources, pdf, paper=paper, duplex=sides == 'duplex')
            return send_file(pdf, mimetype='application/pdf', as_attachment=False, download_name='Chung_tu_da_chon.pdf')
        except ExcelPrintError as exc:
            if getattr(exc, 'code', '') == 'receipt_requires_one_page':
                return jsonify(ok=False, error=str(exc), code=exc.code), 422
            return jsonify(ok=False, error='Chưa tạo được bản in đúng mẫu. Hãy thử lại hoặc tải Excel để in; nếu lỗi tiếp diễn, kiểm tra bộ chuyển PDF trên máy chủ.'), 503
        except (ValueError, OSError) as exc:
            return jsonify(ok=False, error=str(exc)), 422
        finally:
            for _, workbook in books: workbook.close()
