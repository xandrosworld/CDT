"""Catch-up must consume approved orders once, never create another sale."""
import io
import unittest
import zipfile
from unittest.mock import patch

from openpyxl import load_workbook

from . import server, test_outgoing_readiness as fixtures
from .invoice_tax_export import INVOICE_HEADERS
from .outgoing_readiness import canonical_available_stock
from .receivable_ledger import sync_receivable_ledger


class CatchUpExportTests(unittest.TestCase):
    setUpClass = classmethod(fixtures.OutgoingReadinessTests.setUpClass.__func__)
    tearDownClass = classmethod(fixtures.OutgoingReadinessTests.tearDownClass.__func__)
    add_batch = staticmethod(fixtures.OutgoingReadinessTests.add_batch)
    add_opening = staticmethod(fixtures.OutgoingReadinessTests.add_opening)

    def setUp(self):
        fixtures.OutgoingReadinessTests.setUp(self)
        with server.db() as conn:
            conn.execute('DELETE FROM receivable_ledger_revisions')
            conn.execute('DELETE FROM receivable_ledger_lines')
            conn.execute('DELETE FROM outgoing_buyer_profiles')
            conn.execute("DELETE FROM settings WHERE key='minvoice_active_connection'")

    def catch_up(self, **extra):
        return self.client.post('/api/export/catch-up-invoices', json={
            'contractor':'NT-A', 'to':'2026-09-03', **extra})

    def report(self, **extra):
        result = self.client.get('/api/outgoing-invoices/unissued', query_string={
            'contractor':'NT-A', 'to':'2026-09-03', **extra})
        self.assertEqual(result.status_code, 200, result.get_json())
        return result.get_json()

    def template(self, **extra):
        return self.client.get('/api/outgoing-invoices/unissued-template.zip', query_string={
            'contractor':'NT-A', 'to':'2026-09-03', **extra})

    def files(self, response):
        self.assertEqual(response.status_code, 200, response.get_json(silent=True))
        books = {}
        with zipfile.ZipFile(io.BytesIO(response.data)) as archive:
            for name in archive.namelist():
                if not name.endswith('.xlsx') or name == 'DOI_CHIEU_CHI_TIET.xlsx':
                    continue
                book = load_workbook(io.BytesIO(archive.read(name)), data_only=True)
                self.assertEqual(tuple(next(book.active.values)), INVOICE_HEADERS)
                self.assertEqual(book.active.max_column, 13)
                books[name] = list(book.active.values)[1:]
                book.close()
        return books

    def rows(self, response):
        return [row for rows in self.files(response).values() for row in rows]

    def financial_state(self):
        with server.db() as conn:
            return {table: [tuple(row) for row in conn.execute('SELECT * FROM '+table+' ORDER BY 1')]
                    for table in ('orders', 'batches', 'payments', 'receivable_ledger_lines',
                                  'receivable_ledger_revisions', 'payable_ledger_lines')}

    def test_missing_input_template_is_read_only_exact_quantity_and_same_thirteen_columns(self):
        with server.db() as conn:
            self.add_batch(conn, '2026-09-01', [{'qty':27.6, 'sell_price':11000}])
            conn.execute("UPDATE orders SET unit='Cái',product_name='Bánh bao nhân (120g/cái)',tax='8%'")
            conn.execute("UPDATE products SET unit='Cái',tax='8%' WHERE code='HH-01'")
            before = conn.serialize()
        response = self.template()
        row = self.rows(response)[0]
        self.assertEqual((row[2], row[3], row[4], row[8], row[9]), ('Cái', 27.6, 11000, 303600, 8))
        with zipfile.ZipFile(io.BytesIO(response.data)) as archive:
            self.assertIn('Không dùng bộ đối chiếu này để nhập M-Invoice', archive.read('HUONG_DAN.txt').decode('utf-8-sig'))
        self.assertEqual(self.catch_up().status_code, 409)
        with server.db() as conn:
            self.assertEqual(conn.serialize(), before)

    def test_preserves_source_prices_tax_promotion_and_fraction_without_averaging(self):
        with server.db() as conn:
            self.add_batch(conn, '2026-09-01', [{'qty':.04, 'sell_price':10000}])
            self.add_batch(conn, '2026-09-02', [{'qty':.06, 'sell_price':10000}, {'qty':1.25, 'sell_price':20000}])
            _, ids = self.add_batch(conn, '2026-09-03', [{'qty':2, 'sell_price':0}, {'qty':3, 'sell_price':30000}])
            conn.execute("UPDATE orders SET tax='KKKNT'")
            conn.execute("UPDATE orders SET tax='10%',invoice_nature='2' WHERE id=?", (ids[0],))
            conn.execute("UPDATE orders SET tax='8%' WHERE id=?", (ids[1],))
        response = self.template()
        self.assertEqual(response.headers['X-Unissued-Files'], '3')
        rows = self.rows(response)
        self.assertEqual(sorted((r[3],r[4],r[8],r[9],str(r[12])) for r in rows),
                         [(.1,10000,1000,-2,'1'), (1.25,20000,25000,-2,'1'),
                          (2,None,None,10,'2'), (3,30000,90000,8,'1')])

    def test_later_signed_invoice_deducted_once_in_both_exports(self):
        with server.db() as conn:
            self.add_opening(conn, 20)
            batch, ids = self.add_batch(conn, '2026-09-01', [{'qty':10}])
            conn.execute("INSERT INTO outgoing_buyer_profiles(contractor,legal_name,tax_code,address,updated_at) VALUES('NT-A','Test','0101234567','Test','2026-09-01')")
            source = fixtures.OutgoingReadinessTests.add_posted_source(conn, source='minvoice', qty=3, invoice_date='2026-09-10', number='790')
            conn.execute("UPDATE outgoing_source_invoices SET buyer_tax_code='0101234567' WHERE id=?", (source,))
            fixtures.OutgoingReadinessTests.add_local_issued_draft(conn, batch, ids[0], number='790', invoice_date='2026-09-10', qty=3)
        self.assertEqual(self.rows(self.template())[0][3], 7)
        self.assertEqual(self.rows(self.catch_up())[0][3], 7)
        self.assertEqual(self.report()['rows'][0]['unissued_qty'], 7)

    def test_catch_up_always_uses_old_approved_orders_excludes_draft_and_after_cutoff(self):
        with server.db() as conn:
            self.add_opening(conn, 100)
            self.add_batch(conn, '2026-08-31', [{'qty':2}])
            self.add_batch(conn, '2026-09-03', [{'qty':3}])
            draft, _ = self.add_batch(conn, '2026-09-02', [{'qty':10}])
            conn.execute("UPDATE batches SET status='draft' WHERE id=?", (draft,))
            self.add_batch(conn, '2026-09-04', [{'qty':20}])
        before = self.financial_state()
        response = self.catch_up(**{'scope':'range', 'from':'2026-09-03'})
        self.assertEqual(response.headers['X-Order-Scope'], 'unissued')
        self.assertIn('XUAT_BU_UP_M_INVOICE', response.headers['Content-Disposition'])
        self.assertEqual(self.rows(response)[0][3], 5)
        self.assertEqual(self.rows(self.template())[0][3], 5)
        self.assertEqual(before, self.financial_state())

    def test_shortage_then_two_signed_catch_up_rounds_never_duplicate_revenue_or_debt(self):
        with server.db() as conn:
            self.add_opening(conn, 4)
            self.add_batch(conn, '2026-09-01', [{'qty':10, 'sell_price':20000}])
            sync_receivable_ledger(conn, timestamp=server.now_iso())
            self.assertEqual(conn.execute("SELECT SUM(subtotal) FROM receivable_ledger_lines WHERE status='active'").fetchone()[0], 200000)
        before = self.financial_state()
        buyer = self.client.put('/api/outgoing-buyers/NT-A', json={
            'legal_name':'Công ty thử', 'tax_code':'0100000001', 'address':'Địa chỉ thử'})
        self.assertEqual(buyer.status_code, 200, buyer.get_json())
        for number, qty in [('901', 4), ('902', 6)]:
            rows = self.rows(self.catch_up())
            self.assertEqual(sum(r[3] for r in rows), qty)
            with server.db() as conn:
                exported = conn.serialize()
                did = conn.execute("SELECT id FROM outgoing_invoice_drafts WHERE status='draft' AND draft_kind!='waiting'").fetchone()[0]
            self.assertEqual(self.rows(self.catch_up()), rows)
            with server.db() as conn:
                self.assertEqual(conn.serialize(), exported)
            confirmation = {'confirmed':True, 'invoice_number':number, 'invoice_series':'C26TEST', 'invoice_date':'2026-09-10'}
            for _ in range(2):
                confirmed = self.client.post(f'/api/outgoing-invoices/{did}/confirm-issued', json=confirmation)
                self.assertEqual(confirmed.status_code, 200, confirmed.get_json())
            with server.db() as conn:
                refreshed = sync_receivable_ledger(conn, timestamp=server.now_iso())
                self.assertEqual(refreshed['inserted'], 0)
                self.assertEqual(refreshed['updated'], 0)
                self.assertAlmostEqual(canonical_available_stock(conn)['HH-01']['raw_available_qty'], 0)
            self.assertEqual(before, self.financial_state())
            self.assertEqual(self.catch_up().status_code, 409)
            if number == '901':
                self.assertEqual(self.rows(self.template())[0][3], 6)
                with server.db() as conn:
                    fixtures.OutgoingReadinessTests.add_canonical_event(conn, 6, 'input', 'new-real-input', work_date='2026-09-09')
        self.assertEqual(self.report()['rows'], [])
        self.assertEqual(self.template().status_code, 400)
        self.assertEqual(before, self.financial_state())

    def test_unknown_weight_is_visible_in_template_but_not_converted_or_invoiceable(self):
        with server.db() as conn:
            self.add_opening(conn, 50)
            self.add_batch(conn, '2026-09-01', [{'qty':2}])
            conn.execute("UPDATE orders SET unit='Túi',tax='KKKNT'")
        self.assertEqual(self.rows(self.template())[0][2:4], ('Túi', 2))
        self.assertEqual(self.catch_up().status_code, 409)
        self.assertTrue(self.report()['held_line_issues'])

    def test_unresolved_signed_invoice_blocks_catch_up_and_labels_reconciliation(self):
        with server.db() as conn:
            self.add_opening(conn, 20)
            self.add_batch(conn, '2026-09-01', [{'qty':10}])
            fixtures.OutgoingReadinessTests.add_posted_source(conn, source='minvoice', qty=3, invoice_date='2026-09-03', number='792')
            before = conn.serialize()
        with zipfile.ZipFile(io.BytesIO(self.template().data)) as archive:
            self.assertIn('CÒN HÓA ĐƠN CẦN ĐỐI CHIẾU', archive.read('HUONG_DAN.txt').decode('utf-8-sig'))
        response = self.catch_up()
        self.assertEqual(response.status_code, 409)
        self.assertEqual(response.get_json()['code'], 'issued_source_unresolved')
        with server.db() as conn:
            self.assertEqual(before, conn.serialize())

    def test_connected_catch_up_syncs_through_today_before_allocating_and_failure_leaves_no_draft(self):
        with server.db() as conn:
            self.add_opening(conn, 10)
            self.add_batch(conn, '2026-08-31', [{'qty':5}])
            server.setting_set(conn, 'minvoice_active_connection', 'test')
            before = conn.serialize()
        with patch('tdp_system.outgoing_source_refresh.refresh_sources', side_effect=ValueError('connection failed')) as sync:
            response = self.catch_up()
            self.assertEqual(response.status_code, 409)
            self.assertEqual(response.get_json()['code'], 'issued_sync_required')
            self.assertEqual(sync.call_args.args[3], '2026-08-31')
            self.assertGreaterEqual(sync.call_args.args[4], '2026-09-03')
        with server.db() as conn:
            self.assertEqual(before, conn.serialize())

    def test_invalid_filters_rejected_without_changes(self):
        for args in [{'to':'bad'}, {'contractor':'MISSING'}]:
            self.assertEqual(self.template(**args).status_code, 400)
            self.assertEqual(self.catch_up(**args).status_code, 409)
        self.assertEqual(self.client.post('/api/export/catch-up-invoices', json=['bad']).status_code, 409)

    def test_shared_stock_released_by_later_signed_buyer_is_usable_on_first_download(self):
        with server.db() as conn:
            self.add_opening(conn, 20)
            self.add_batch(conn, '2026-09-01', [{'qty':10}, {'qty':10,'contractor':'NT-B'}])
        self.assertEqual(sum(r[3] for r in self.rows(self.catch_up(contractor=''))),20)
        with server.db() as conn:
            conn.execute("INSERT INTO outgoing_buyer_profiles(contractor,legal_name,tax_code,address,updated_at) VALUES('NT-B','Test','0101234567','Test','2026-09-01')")
            source=fixtures.OutgoingReadinessTests.add_posted_source(conn,source='minvoice',qty=10,invoice_date='2026-09-10',number='801')
            conn.execute("UPDATE outgoing_source_invoices SET buyer_tax_code='0101234567' WHERE id=?",(source,))
        first=self.rows(self.catch_up(contractor=''))
        self.assertEqual(sum(r[3] for r in first),10)
        with server.db() as conn:
            before=conn.serialize()
            audit_before=[tuple(r) for r in conn.execute('SELECT * FROM audit_log')]
        self.assertEqual(first,self.rows(self.catch_up(contractor='')))
        with server.db() as conn:
            self.assertEqual(audit_before,[tuple(r) for r in conn.execute('SELECT * FROM audit_log')])
            self.assertEqual(before,conn.serialize())

    def test_rounding_shared_stock_is_completed_before_first_download(self):
        with server.db() as conn:
            self.add_opening(conn,1.05)
            self.add_batch(conn,'2026-09-01',[{'qty':.16},{'qty':.96,'contractor':'NT-B'}])
            # Fractional waiting holds already exist before export (as when a
            # partially signed external invoice releases part of a reservation).
            from .outgoing_waiting import refresh_waiting
            refresh_waiting(conn,server.now_iso())
        first=self.rows(self.catch_up(contractor=''))
        self.assertEqual(sorted(r[3] for r in first),[.1,.9])
        with server.db() as conn:before=conn.serialize()
        self.assertEqual(first,self.rows(self.catch_up(contractor='')))
        with server.db() as conn:self.assertEqual(before,conn.serialize())


if __name__ == '__main__':
    unittest.main()
