import io
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

try:
    from . import contract_modules, print_bundle, server
except ImportError:  # pragma: no cover - direct file invocation
    import contract_modules
    import print_bundle
    import server


class PurchaseOrderRoundtripTests(unittest.TestCase):
    """Regression coverage for the customer's real two-workbook workflow.

    The customer order is the receivable source of truth.  The downloaded NCC
    workbook is a separate purchasing decision: the operator subtracts the
    physical cabinet/freezer stock, enters the actual order quantity and buy
    price, then uploads it again.  Invoice inventory belongs only to the
    outgoing-invoice workflow and must never reduce the NCC order.
    """

    @classmethod
    def setUpClass(cls):
        cls.temp_dir = tempfile.TemporaryDirectory()
        cls.original_db_path = server.DB_PATH
        cls.original_data_dir = server.DATA_DIR
        cls.original_testing = server.app.config.get("TESTING", False)
        server.DATA_DIR = Path(cls.temp_dir.name) / "data"
        server.DATA_DIR.mkdir(parents=True, exist_ok=True)
        server.DB_PATH = server.DATA_DIR / "purchase_roundtrip.sqlite3"
        server.app.config["TESTING"] = True
        server.init_database()
        cls.client = server.app.test_client()

    @classmethod
    def tearDownClass(cls):
        server.DB_PATH = cls.original_db_path
        server.DATA_DIR = cls.original_data_dir
        server.app.config["TESTING"] = cls.original_testing
        cls.temp_dir.cleanup()

    def setUp(self):
        pending = getattr(contract_modules, "PENDING_PURCHASE_ORDER_IMPORTS", None)
        if pending is not None:
            pending.clear()
        with server.db() as conn:
            conn.execute("DELETE FROM invoice_inventory_ledger")
            conn.execute("DELETE FROM invoice_inventory_confirmations")
            conn.execute("DELETE FROM receivable_ledger_revisions")
            conn.execute("DELETE FROM receivable_ledger_lines")
            conn.execute("DELETE FROM outgoing_invoice_lines")
            conn.execute("DELETE FROM outgoing_invoice_drafts")
            conn.execute("DELETE FROM inventory_transactions")
            conn.execute("DELETE FROM purchase_order_imports")
            conn.execute("DELETE FROM purchase_workbook_lines")
            conn.execute("DELETE FROM purchase_order_lines")
            conn.execute("DELETE FROM order_import_receipts")
            conn.execute("DELETE FROM orders")
            conn.execute("DELETE FROM batches")
            conn.execute("DELETE FROM historical_payable_lines")
            conn.execute("DELETE FROM settings WHERE key='historical_payables_through_date'")
            conn.execute("DELETE FROM payments")
            conn.execute("DELETE FROM balances")
            conn.execute("DELETE FROM debt_adjustments")
            conn.execute("DELETE FROM supplier_rules")
            conn.execute("DELETE FROM audit_log")
            conn.execute(
                "INSERT OR REPLACE INTO contractors(code,name,price_group,pricing_mode) "
                "VALUES('HATRAN','Hà Trân','HATRAN','group')"
            )
            conn.execute(
                "INSERT OR REPLACE INTO kitchens(code,contractor,name,address,show_price) "
                "VALUES('POT','HATRAN','Bếp POT','',0)"
            )
            conn.execute(
                "INSERT OR REPLACE INTO suppliers(code,name) VALUES('DUNG','Nhà Dung')"
            )
            conn.execute(
                """INSERT OR REPLACE INTO products(
                       code,name,unit,tax,supplier,buy_price,purchase_list,seller,cccd
                   ) VALUES('PO-P1','Sườn non','kg','0%','DUNG',12000,0,'','')"""
            )

    @staticmethod
    def _insert_approved_order(
        conn, *, qty=10, product_code="PO-P1", product_name="Sườn non",
        original_buy_price=0,
    ):
        timestamp = server.now_iso()
        batch_id = conn.execute(
            """INSERT INTO batches(work_date,source_name,status,created_at,approved_at)
               VALUES('2026-09-01','customer-order.xlsx','approved',?,?)""",
            (timestamp, timestamp),
        ).lastrowid
        order_id = conn.execute(
            """INSERT INTO orders(
                   batch_id,work_date,contractor,kitchen,product_code,product_name,qty,
                   actual_received,actual_delivered,unit,supplier,buy_price,sell_price,tax,
                   purchase_list,note,errors,warnings,updated_at
               ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,? ,0,'','[]','[]',?)""",
            (
                batch_id, "2026-09-01", "HATRAN", "POT", product_code,
                product_name, qty, qty, qty, "kg", "DUNG", original_buy_price, 30000,
                "0%", timestamp,
            ),
        ).lastrowid
        server.sync_receivable_ledger(conn, timestamp=timestamp)
        return batch_id, order_id

    @staticmethod
    def _add_canonical_invoice_stock(conn, qty, key, work_date="2026-09-01"):
        timestamp = server.now_iso()
        source_id = sum((index + 1) * ord(char) for index, char in enumerate(key))
        confirmation_id = conn.execute(
            """INSERT INTO invoice_inventory_confirmations(
                   confirmation_key,direction,source_invoice_table,source_invoice_id,
                   action,confirmed,note,created_at
               ) VALUES(?,'input','test_source',?,'post',1,'regression',?)""",
            (f"confirm-{key}", source_id, timestamp),
        ).lastrowid
        conn.execute(
            """INSERT INTO invoice_inventory_ledger(
                   event_key,direction,event_type,source_invoice_table,source_invoice_id,
                   source_line_id,source_line_index,product_code,txn_date,qty_delta,unit_cost,
                   mapping_revision_id,confirmation_id,reverses_event_key,status,created_at
               ) VALUES(?,'input','POST','test_source',?,?,1,'PO-P1',?,?,12000,
                        NULL,?,'','posted',?)""",
            (key, source_id, source_id, work_date, qty, confirmation_id, timestamp),
        )

    @staticmethod
    def _edited_export(
        response, *, order_qty=7, buy_price=50000, physical_stock=3,
        supplier="DUNG",
    ):
        workbook = load_workbook(io.BytesIO(response.data))
        try:
            if workbook.sheetnames != ["đặt hàng"]:
                raise AssertionError(f"Expected canonical sheet, got {workbook.sheetnames}")
            sheet = workbook["đặt hàng"]
            headers = [sheet.cell(2, column).value for column in range(1, 18)]
            expected = [
                "Mã hàngNCC", "Mã hàng", "Mã bếp", None, "Tên hàng ", "Số lượng",
                "ĐVT", "NCC", "ghi chú", "giá mua", "hỏng", "thêm", "Giảm",
                "thiếu", "SL \nthực té", "Thành tiền", "Mã dòng hệ thống",
            ]
            if headers != expected:
                raise AssertionError(f"Unexpected NCC workbook headers: {headers!r}")

            # The customer checks physical stock and records the quantity she
            # will really buy directly in the canonical quantity/adjustment columns.
            sheet.cell(3, 6, order_qty)
            sheet.cell(3, 8, supplier)
            sheet.cell(3, 9, "Đã kiểm tra tủ đông")
            sheet.cell(3, 10, buy_price)
            stream = io.BytesIO()
            workbook.save(stream)
            stream.seek(0)
            return stream
        finally:
            workbook.close()

    def test_internal_kho_rows_allow_zero_buy_price(self):
        with server.db() as conn:
            batch_id, _ = self._insert_approved_order(conn)
            conn.execute(
                "UPDATE orders SET supplier='KHO',buy_price=0 WHERE batch_id=?",
                (batch_id,),
            )
        exported = self.client.get(f"/api/export/suppliers/{batch_id}")
        edited = self._edited_export(
            exported, order_qty=7, physical_stock=3, buy_price=0, supplier="KHO",
        )
        response = self.client.post(
            "/api/purchase-orders/import/preview",
            data={
                "batch_id": str(batch_id),
                "file": (edited, "Dat_Kho_da_chinh.xlsx"),
            },
            content_type="multipart/form-data",
        )
        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        preview = response.get_json()
        self.assertTrue(preview["can_confirm"])
        self.assertEqual(preview["error_rows"], 0)

    def test_all_purchase_adjustments_round_vnd_and_keep_component_history(self):
        with server.db() as conn:
            batch_id, _ = self._insert_approved_order(conn)
            timestamp = server.now_iso()
            for index in range(2, 6):
                conn.execute(
                    """INSERT INTO orders(
                           batch_id,work_date,contractor,kitchen,product_code,product_name,qty,
                           actual_received,actual_delivered,unit,supplier,buy_price,sell_price,tax,
                           purchase_list,note,errors,warnings,updated_at
                       ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,0,'','[]','[]',?)""",
                    (
                        batch_id, "2026-09-01", "HATRAN", "POT", "PO-P1",
                        f"Sườn non {index}", 10, 10, 10, "kg", "DUNG", 0,
                        30000, "0%", timestamp,
                    ),
                )
            sales_before = [tuple(row) for row in conn.execute(
                "SELECT id,qty,buy_price,sell_price FROM orders WHERE batch_id=? ORDER BY id",
                (batch_id,),
            )]
        exported = self.client.get(f"/api/export/suppliers/{batch_id}")
        workbook = load_workbook(io.BytesIO(exported.data), data_only=False)
        try:
            sheet = workbook["đặt hàng"]
            adjustments = [
                (1, 0, 0, 0),   # hỏng
                (0, 2, 0, 0),   # thêm
                (0, 0, 3, 0),   # giảm
                (0, 0, 0, 4),   # thiếu
                (1, 5, 2, 3),   # tổ hợp
            ]
            for row_index, values in enumerate(adjustments, 3):
                sheet.cell(row_index, 10, 10.5)
                for column, value in zip((11, 12, 13, 14), values):
                    sheet.cell(row_index, column, value)
            stream = io.BytesIO()
            workbook.save(stream)
            payload_bytes = stream.getvalue()
        finally:
            workbook.close()
        preview_response = self.client.post(
            "/api/purchase-orders/import/preview",
            data={"batch_id": str(batch_id), "file": (io.BytesIO(payload_bytes), "adjustments.xlsx")},
            content_type="multipart/form-data",
        )
        self.assertEqual(preview_response.status_code, 200, preview_response.get_data(as_text=True))
        preview = preview_response.get_json()
        self.assertEqual((preview["error_rows"], preview["total_qty"], preview["total_amount"]),
                         (0, 43, 453))
        confirmed = self.client.post(
            "/api/purchase-orders/import/confirm",
            json={"token": preview["token"], "confirmed": True},
        )
        self.assertEqual(confirmed.status_code, 200, confirmed.get_data(as_text=True))
        self.assertEqual(confirmed.get_json()["processed"], 5)
        with server.db() as conn:
            rows = [dict(row) for row in conn.execute(
                """SELECT damaged_qty,added_qty,reduced_qty,missing_qty,actual_qty,amount,revision
                   FROM purchase_workbook_lines WHERE batch_id=? ORDER BY source_row""",
                (batch_id,),
            )]
            history = [dict(row) for row in conn.execute(
                """SELECT damaged_qty,added_qty,reduced_qty,missing_qty,actual_qty,amount,revision
                   FROM purchase_workbook_line_revisions WHERE batch_id=? ORDER BY source_row""",
                (batch_id,),
            )]
            metadata = json.loads(conn.execute(
                """SELECT metadata_json FROM audit_log
                   WHERE event_type='purchase_orders.import' AND status='ok'
                   ORDER BY id DESC LIMIT 1"""
            ).fetchone()[0])
            sales_after = [tuple(row) for row in conn.execute(
                "SELECT id,qty,buy_price,sell_price FROM orders WHERE batch_id=? ORDER BY id",
                (batch_id,),
            )]
        self.assertEqual(
            [(row["damaged_qty"], row["added_qty"], row["reduced_qty"],
              row["missing_qty"], row["actual_qty"], row["amount"])
             for row in rows],
            [(1, 0, 0, 0, 9, 95), (0, 2, 0, 0, 12, 126),
             (0, 0, 3, 0, 7, 74), (0, 0, 0, 4, 6, 63),
             (1, 5, 2, 3, 9, 95)],
        )
        self.assertEqual(history, rows)
        self.assertEqual(metadata["components"], {
            "base_qty": 50.0, "damaged_qty": 2.0, "added_qty": 7.0,
            "reduced_qty": 5.0, "missing_qty": 7.0, "actual_qty": 43.0,
            "amount": 453,
        })
        self.assertEqual(sales_after, sales_before)

        replay_preview = self.client.post(
            "/api/purchase-orders/import/preview",
            data={"batch_id": str(batch_id), "file": (io.BytesIO(payload_bytes), "renamed.xlsx")},
            content_type="multipart/form-data",
        ).get_json()
        replay = self.client.post(
            "/api/purchase-orders/import/confirm",
            json={"token": replay_preview["token"], "confirmed": True},
        )
        self.assertEqual((replay.status_code, replay.get_json()["idempotent"]), (200, True))
        with server.db() as conn:
            self.assertEqual(conn.execute(
                "SELECT COUNT(*) FROM purchase_workbook_line_revisions WHERE batch_id=?",
                (batch_id,),
            ).fetchone()[0], 5)

    def test_invalid_adjustments_prices_and_formula_mismatches_are_blocked(self):
        with server.db() as conn:
            batch_id, _ = self._insert_approved_order(conn)
        exported = self.client.get(f"/api/export/suppliers/{batch_id}").data
        variants = (
            ("negative component", 11, -1, "không được âm"),
            ("nan component", 12, "NaN", "phải là số hữu hạn"),
            ("negative actual", 11, 11, "Số lượng thực tế không được âm"),
            ("zero external price", 10, 0, "phải có giá mua lớn hơn 0"),
            ("negative price", 10, -1, "Giá mua không được âm"),
            ("wrong actual formula", 15, "=F3+L3+K3-M3-N3", "Công thức Số lượng thực tế"),
            ("wrong amount formula", 16, "=O3+J3", "Công thức Thành tiền"),
            ("stale actual cache", 15, 9, "Số lượng thực tế lệch công thức"),
            ("stale amount cache", 16, 999, "Thành tiền lệch"),
        )
        for label, column, value, expected_error in variants:
            with self.subTest(label=label):
                workbook = load_workbook(io.BytesIO(exported), data_only=False)
                try:
                    sheet = workbook["đặt hàng"]
                    sheet.cell(3, 10, 10.5)
                    sheet.cell(3, column, value)
                    stream = io.BytesIO()
                    workbook.save(stream)
                    stream.seek(0)
                finally:
                    workbook.close()
                response = self.client.post(
                    "/api/purchase-orders/import/preview",
                    data={"batch_id": str(batch_id), "file": (stream, f"{label}.xlsx")},
                    content_type="multipart/form-data",
                )
                self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
                payload = response.get_json()
                self.assertFalse(payload["can_confirm"])
                self.assertEqual(payload["error_rows"], 1)
                self.assertTrue(any(
                    expected_error in message for message in payload["rows"][0]["errors"]
                ), payload["rows"][0]["errors"])
        with server.db() as conn:
            self.assertEqual(conn.execute(
                "SELECT COUNT(*) FROM purchase_workbook_lines WHERE batch_id=?", (batch_id,),
            ).fetchone()[0], 0)

    def _preview_and_confirm(self, batch_id, workbook_stream):
        preview_response = self.client.post(
            "/api/purchase-orders/import/preview",
            data={
                "batch_id": str(batch_id),
                "file": (workbook_stream, "Dat_NCC_da_chinh.xlsx"),
            },
            content_type="multipart/form-data",
        )
        self.assertEqual(
            preview_response.status_code, 200, preview_response.get_data(as_text=True),
        )
        preview = preview_response.get_json()
        self.assertTrue(preview["ok"])
        self.assertTrue(preview["can_confirm"])
        self.assertEqual(preview["count"], 1)
        self.assertEqual(preview["error_rows"], 0)
        self.assertEqual(preview["total_qty"], 7)
        self.assertEqual(preview["total_amount"], 350000)

        confirm_response = self.client.post(
            "/api/purchase-orders/import/confirm",
            json={"token": preview["token"], "confirmed": True},
        )
        self.assertEqual(
            confirm_response.status_code, 200, confirm_response.get_data(as_text=True),
        )
        self.assertTrue(confirm_response.get_json()["ok"])
        return preview, confirm_response.get_json()

    def test_ncc_file_roundtrip_is_separate_from_customer_order_and_invoice_stock(self):
        with server.db() as conn:
            batch_id, order_id = self._insert_approved_order(conn)
            timestamp = server.now_iso()
            # Deliberately abundant accounting/invoice stock: the old bug used
            # this to reduce the NCC need to zero.
            conn.execute(
                """INSERT INTO inventory_transactions(
                       txn_date,product_code,qty_in,qty_out,unit_cost,source_type,
                       source_id,source_line,status,note,created_at,updated_at
                   ) VALUES('2026-09-01','PO-P1',100,0,12000,'MSMI','INV-1','1',
                            'posted','invoice stock',?,?)""",
                (timestamp, timestamp),
            )

        needs_response = self.client.get(f"/api/supplier-needs/{batch_id}")
        self.assertEqual(needs_response.status_code, 200, needs_response.get_data(as_text=True))
        needs = needs_response.get_json()
        self.assertEqual(needs["ordered_qty"], 10)
        self.assertEqual(needs["required_qty"], 10)

        export_response = self.client.get(f"/api/export/suppliers/{batch_id}")
        self.assertEqual(
            export_response.status_code, 200,
            f"Unexpected export status/content-type: {export_response.status_code} {export_response.content_type}",
        )
        edited = self._edited_export(export_response)
        edited_bytes = edited.getvalue()
        _, first_confirm = self._preview_and_confirm(batch_id, edited)
        self.assertFalse(first_confirm["idempotent"])

        with server.db() as conn:
            order = dict(conn.execute("SELECT * FROM orders WHERE id=?", (order_id,)).fetchone())
            plan = dict(conn.execute(
                "SELECT * FROM purchase_workbook_lines WHERE order_id=?", (order_id,),
            ).fetchone())
        # Receivable source data is untouched by the purchasing roundtrip.
        self.assertEqual(order["qty"], 10)
        self.assertEqual(order["actual_delivered"], 10)
        self.assertEqual(order["sell_price"], 30000)
        self.assertEqual(order["buy_price"], 0)
        self.assertEqual(plan["base_qty"], 7)
        self.assertEqual(plan["actual_qty"], 7)
        self.assertEqual(plan["buy_price"], 50000)
        self.assertEqual(plan["price_source"], "Sheet đặt hàng chuẩn")
        first_plan_updated_at = plan["updated_at"]

        needs_response = self.client.get(f"/api/supplier-needs/{batch_id}")
        self.assertEqual(needs_response.status_code, 200, needs_response.get_data(as_text=True))
        needs = needs_response.get_json()
        self.assertEqual(needs["ordered_qty"], 7)
        self.assertEqual(needs["required_qty"], 7)
        self.assertEqual(needs["physical_stock_used"], 0)

        debts_response = self.client.get("/api/debts?from=2026-09-01&to=2026-09-01")
        self.assertEqual(debts_response.status_code, 200, debts_response.get_data(as_text=True))
        debts = debts_response.get_json()
        self.assertEqual(debts["contractors"]["HATRAN"]["period_charge"], 300000)
        self.assertEqual(debts["suppliers"]["DUNG"]["period_charge"], 350000)

        # Re-importing the exact same business values is idempotent: one plan
        # line and one source receipt, never duplicated supplier payable.
        _, second_confirm = self._preview_and_confirm(batch_id, io.BytesIO(edited_bytes))
        self.assertTrue(second_confirm["idempotent"])
        with server.db() as conn:
            self.assertEqual(conn.execute(
                "SELECT COUNT(*) n FROM purchase_workbook_lines WHERE batch_id=?", (batch_id,),
            ).fetchone()["n"], 1)
            self.assertEqual(conn.execute(
                "SELECT COUNT(*) n FROM purchase_order_imports WHERE batch_id=?", (batch_id,),
            ).fetchone()["n"], 1)
            plan_after_repeat = dict(conn.execute(
                "SELECT * FROM purchase_workbook_lines WHERE order_id=?", (order_id,),
            ).fetchone())
            self.assertEqual(plan_after_repeat["price_source"], "Sheet đặt hàng chuẩn")
            self.assertEqual(plan_after_repeat["updated_at"], first_plan_updated_at)
        debts = self.client.get("/api/debts?from=2026-09-01&to=2026-09-01").get_json()
        self.assertEqual(debts["suppliers"]["DUNG"]["period_charge"], 350000)

    def test_actual_customer_sheet_259_rows_roundtrips_without_touching_sales(self):
        manifest = json.loads(
            (Path(__file__).with_name("golden_manifest.json")).read_text(encoding="utf-8")
        )
        source_path = Path(next(
            item["path"] for item in manifest["artifacts"]
            if item["id"] == "daily_order_2026_09_01"
        ))
        orders, skipped = server.parse_workbook(source_path, "2026-09-01", ["01.09"])
        self.assertEqual((len(orders), skipped), (323, []))
        with server.db() as conn:
            batch_id = int(conn.execute(
                """INSERT INTO batches(work_date,source_name,status,created_at,approved_at)
                   VALUES('2026-09-01','golden-01.09','approved',?,?)""",
                (server.now_iso(), server.now_iso()),
            ).lastrowid)
            server.save_imported_orders(conn, batch_id, orders)
            sales_before = [tuple(row) for row in conn.execute(
                """SELECT id,qty,actual_received,actual_delivered,buy_price,sell_price,
                          damaged_qty,supplier_return_qty,customer_return_qty
                   FROM orders WHERE batch_id=? ORDER BY id""",
                (batch_id,),
            )]

        preview_response = self.client.post(
            "/api/purchase-orders/import/preview",
            data={
                "batch_id": str(batch_id),
                "file": (io.BytesIO(source_path.read_bytes()), source_path.name),
            },
            content_type="multipart/form-data",
        )
        self.assertEqual(preview_response.status_code, 200, preview_response.get_data(as_text=True))
        preview = preview_response.get_json()
        self.assertEqual((preview["format"], preview["count"], preview["error_rows"]),
                         ("customer_canonical", 259, 1))
        self.assertEqual(preview["warning_rows"], 0)
        self.assertFalse(preview["can_confirm"])
        blocked = self.client.post(
            "/api/purchase-orders/import/confirm",
            json={"token": preview["token"], "confirmed": True},
        )
        self.assertEqual(blocked.status_code, 400, blocked.get_data(as_text=True))
        with server.db() as conn:
            self.assertEqual(conn.execute(
                "SELECT COUNT(*) FROM purchase_workbook_lines WHERE batch_id=?", (batch_id,),
            ).fetchone()[0], 0)

        # The real customer file contains exactly one external positive-quantity
        # line without a buy price. It must be corrected before confirmation.
        fixed_workbook = load_workbook(source_path, data_only=False)
        fixed_cache = load_workbook(source_path, data_only=True)
        try:
            fixed_sheet = fixed_workbook["đặt hàng"]
            cache_sheet = fixed_cache["đặt hàng"]
            fixed_rows = []
            for row_index in range(3, fixed_sheet.max_row + 1):
                # openpyxl does not preserve cached formula results when saving.
                # Materialize only the formula cells from the customer's cache
                # so the corrected fixture remains a valid static equivalent.
                for column in (6, 10, 11, 12, 13, 14, 15, 16):
                    value = fixed_sheet.cell(row_index, column).value
                    if isinstance(value, str) and value.startswith("="):
                        fixed_sheet.cell(row_index, column, cache_sheet.cell(row_index, column).value)
                base = float(cache_sheet.cell(row_index, 6).value or 0)
                damaged = float(cache_sheet.cell(row_index, 11).value or 0)
                added = float(cache_sheet.cell(row_index, 12).value or 0)
                reduced = float(cache_sheet.cell(row_index, 13).value or 0)
                missing = float(cache_sheet.cell(row_index, 14).value or 0)
                actual = base + added - damaged - reduced - missing
                supplier = str(cache_sheet.cell(row_index, 8).value or "")
                price = float(cache_sheet.cell(row_index, 10).value or 0)
                if actual > 0 and price <= 0 and not contract_modules.is_internal_stock_supplier(supplier):
                    fixed_sheet.cell(row_index, 10, 1000)
                    fixed_sheet.cell(row_index, 16, server.vnd_product(actual, 1000))
                    fixed_rows.append(row_index)
            self.assertEqual(len(fixed_rows), 1)
            fixed_stream = io.BytesIO()
            fixed_workbook.save(fixed_stream)
            fixed_stream.seek(0)
        finally:
            fixed_workbook.close()
            fixed_cache.close()
        fixed_preview_response = self.client.post(
            "/api/purchase-orders/import/preview",
            data={
                "batch_id": str(batch_id),
                "file": (fixed_stream, "daily-order-purchase-priced.xlsx"),
            },
            content_type="multipart/form-data",
        )
        self.assertEqual(
            fixed_preview_response.status_code, 200,
            fixed_preview_response.get_data(as_text=True),
        )
        preview = fixed_preview_response.get_json()
        self.assertEqual(
            (preview["count"], preview["error_rows"]), (259, 0),
            [(item["source_row"], item["errors"]) for item in preview["issues"]],
        )
        self.assertTrue(preview["can_confirm"])
        source_total = preview["total_amount"]
        confirm = self.client.post(
            "/api/purchase-orders/import/confirm",
            json={"token": preview["token"], "confirmed": True},
        )
        self.assertEqual(confirm.status_code, 200, confirm.get_data(as_text=True))
        with server.db() as conn:
            self.assertEqual(conn.execute(
                "SELECT COUNT(*) FROM purchase_workbook_lines WHERE batch_id=?", (batch_id,),
            ).fetchone()[0], 259)
            self.assertEqual(conn.execute(
                "SELECT COUNT(*) FROM purchase_workbook_lines WHERE batch_id=? AND order_id IS NULL",
                (batch_id,),
            ).fetchone()[0], 6)
            self.assertEqual(conn.execute(
                "SELECT COUNT(*) FROM purchase_workbook_line_revisions WHERE batch_id=?",
                (batch_id,),
            ).fetchone()[0], 259)
            sales_after = [tuple(row) for row in conn.execute(
                """SELECT id,qty,actual_received,actual_delivered,buy_price,sell_price,
                          damaged_qty,supplier_return_qty,customer_return_qty
                   FROM orders WHERE batch_id=? ORDER BY id""",
                (batch_id,),
            )]
        self.assertEqual(sales_after, sales_before)

        exported = self.client.get(f"/api/export/suppliers/{batch_id}")
        self.assertEqual(exported.status_code, 200)
        workbook = load_workbook(io.BytesIO(exported.data), data_only=False)
        try:
            sheet = workbook["đặt hàng"]
            self.assertEqual(sheet.max_row - 2, 259)
            self.assertTrue(sheet.column_dimensions["A"].hidden)
            self.assertTrue(sheet.column_dimensions["B"].hidden)
            self.assertTrue(sheet.column_dimensions["Q"].hidden)
            self.assertEqual(sheet["O3"].value, "=F3+L3-K3-M3-N3")
            self.assertEqual(sheet["P3"].value, "=IFERROR(O3*J3,0)")
            sections = print_bundle.workbook_sections("supplier_orders", workbook)
            self.assertEqual((len(sections), sections[0]["title"]), (1, "đặt hàng"))
            printed = {
                column["label"]: sections[0]["rows"][0][column["key"]]
                for column in sections[0]["columns"]
            }
            expected_actual = (
                float(sheet["F3"].value or 0) + float(sheet["L3"].value or 0)
                - float(sheet["K3"].value or 0) - float(sheet["M3"].value or 0)
                - float(sheet["N3"].value or 0)
            )
            self.assertEqual(printed["SL \nthực té"], expected_actual)
            self.assertEqual(printed["Thành tiền"], expected_actual * float(sheet["J3"].value or 0))
            self.assertFalse(any(
                isinstance(value, str) and value.startswith("=")
                for value in printed.values()
            ))
        finally:
            workbook.close()
        roundtrip_preview = self.client.post(
            "/api/purchase-orders/import/preview",
            data={"batch_id": str(batch_id), "file": (io.BytesIO(exported.data), "dat-hang.xlsx")},
            content_type="multipart/form-data",
        )
        self.assertEqual(roundtrip_preview.status_code, 200, roundtrip_preview.get_data(as_text=True))
        roundtrip = roundtrip_preview.get_json()
        self.assertEqual(roundtrip["count"], 259)
        self.assertAlmostEqual(roundtrip["total_amount"], source_total, places=5)
        replay = self.client.post(
            "/api/purchase-orders/import/confirm",
            json={"token": roundtrip["token"], "confirmed": True},
        )
        self.assertEqual(replay.status_code, 200, replay.get_data(as_text=True))
        self.assertTrue(replay.get_json()["idempotent"])

        debts_before = self.client.get("/api/debts?from=2026-09-01&to=2026-09-01").get_json()
        payable_before = sum(item["period_charge"] for item in debts_before["suppliers"].values())
        with server.db() as conn:
            editable = next(
                dict(row) for row in conn.execute(
                """SELECT * FROM purchase_workbook_lines
                   WHERE batch_id=? AND order_id IS NULL AND actual_qty>0 ORDER BY source_row""",
                (batch_id,),
                ) if not contract_modules.is_internal_stock_supplier(row["supplier"])
            )
        workbook = load_workbook(io.BytesIO(exported.data))
        try:
            sheet = workbook["đặt hàng"]
            export_row = int(editable["source_row"])
            new_price = editable["buy_price"] + 1000
            sheet.cell(export_row, 10, new_price)
            changed = io.BytesIO()
            workbook.save(changed)
            changed.seek(0)
        finally:
            workbook.close()
        changed_preview = self.client.post(
            "/api/purchase-orders/import/preview",
            data={"batch_id": str(batch_id), "file": (changed, "dat-hang-sua-gia.xlsx")},
            content_type="multipart/form-data",
        )
        self.assertEqual(changed_preview.status_code, 200, changed_preview.get_data(as_text=True))
        changed_payload = changed_preview.get_json()
        changed_confirm = self.client.post(
            "/api/purchase-orders/import/confirm",
            json={"token": changed_payload["token"], "confirmed": True},
        )
        self.assertEqual(changed_confirm.status_code, 200, changed_confirm.get_data(as_text=True))
        self.assertEqual(changed_confirm.get_json()["processed"], 1)
        debts_after = self.client.get("/api/debts?from=2026-09-01&to=2026-09-01").get_json()
        payable_after = sum(item["period_charge"] for item in debts_after["suppliers"].values())
        expected_delta = (
            server.vnd_round(editable["actual_qty"] * new_price)
            - server.vnd_round(editable["amount"])
        )
        self.assertEqual(payable_after - payable_before, expected_delta)
        with server.db() as conn:
            sales_final = [tuple(row) for row in conn.execute(
                """SELECT id,qty,actual_received,actual_delivered,buy_price,sell_price,
                          damaged_qty,supplier_return_qty,customer_return_qty
                   FROM orders WHERE batch_id=? ORDER BY id""",
                (batch_id,),
            )]
            revisions = [dict(row) for row in conn.execute(
                """SELECT row_key,revision,damaged_qty,added_qty,reduced_qty,missing_qty,
                          actual_qty,buy_price,amount
                   FROM purchase_workbook_line_revisions WHERE batch_id=?
                   ORDER BY id""",
                (batch_id,),
            )]
            audit_metadata = json.loads(conn.execute(
                """SELECT metadata_json FROM audit_log
                   WHERE event_type='purchase_orders.import' AND status='ok'
                   ORDER BY id DESC LIMIT 1"""
            ).fetchone()[0])
        self.assertEqual(sales_final, sales_before)
        self.assertEqual((len(revisions), revisions[-1]["revision"]), (260, 2))
        self.assertEqual(
            set(audit_metadata["components"]),
            {"base_qty", "damaged_qty", "added_qty", "reduced_qty", "missing_qty",
             "actual_qty", "amount"},
        )
        self.assertEqual(
            (audit_metadata["changed"], audit_metadata["inserted"], audit_metadata["updated"]),
            (1, 0, 1),
        )

    def test_legacy_13_column_file_migrates_into_canonical_purchase_scope(self):
        with server.db() as conn:
            batch_id, order_id = self._insert_approved_order(conn)
        workbook = load_workbook(io.BytesIO(
            self.client.get(f"/api/export/suppliers/{batch_id}").data
        ))
        workbook.remove(workbook.active)
        sheet = workbook.create_sheet("DUNG-POT")
        sheet.append([
            "Mã dòng hệ thống", "Mã hàng", "Mã bếp", "Ngày", "Tên hàng",
            "Nhu cầu từ đơn khách", "Tồn tủ đã trừ", "Số lượng đặt NCC",
            "ĐVT", "NCC", "Giá mua", "Thành tiền", "Ghi chú",
        ])
        sheet.append([
            order_id, "PO-P1", "POT", "2026-09-01", "Sườn non", 10, 3, 7,
            "kg", "DUNG", 50000, 350000, "legacy",
        ])
        legacy = io.BytesIO()
        workbook.save(legacy)
        workbook.close()
        legacy.seek(0)
        preview_response = self.client.post(
            "/api/purchase-orders/import/preview",
            data={"batch_id": str(batch_id), "file": (legacy, "legacy-13.xlsx")},
            content_type="multipart/form-data",
        )
        self.assertEqual(preview_response.status_code, 200, preview_response.get_data(as_text=True))
        preview = preview_response.get_json()
        self.assertEqual((preview["format"], preview["count"], preview["total_amount"]),
                         ("legacy_13", 1, 350000))
        confirmed = self.client.post(
            "/api/purchase-orders/import/confirm",
            json={"token": preview["token"], "confirmed": True},
        )
        self.assertEqual(confirmed.status_code, 200, confirmed.get_data(as_text=True))
        with server.db() as conn:
            canonical = dict(conn.execute(
                "SELECT * FROM purchase_workbook_lines WHERE batch_id=?", (batch_id,),
            ).fetchone())
            legacy_plan = dict(conn.execute(
                "SELECT * FROM purchase_order_lines WHERE batch_id=?", (batch_id,),
            ).fetchone())
            order = dict(conn.execute("SELECT * FROM orders WHERE id=?", (order_id,)).fetchone())
        self.assertEqual((canonical["base_qty"], canonical["actual_qty"], canonical["buy_price"]),
                         (7, 7, 50000))
        self.assertEqual((legacy_plan["order_qty"], legacy_plan["physical_stock_used"]), (7, 3))
        self.assertEqual(order["buy_price"], 0)

    def test_canonical_confirm_failure_rolls_back_all_purchase_rows_and_receipt(self):
        with server.db() as conn:
            batch_id, _ = self._insert_approved_order(conn)
            timestamp = server.now_iso()
            conn.execute(
                """INSERT INTO orders(
                       batch_id,work_date,contractor,kitchen,product_code,product_name,qty,
                       actual_received,actual_delivered,unit,supplier,buy_price,sell_price,tax,
                       purchase_list,note,errors,warnings,updated_at
                   ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,0,'','[]','[]',?)""",
                (
                    batch_id, "2026-09-01", "HATRAN", "POT", "PO-P2", "Thịt vai",
                    2, 2, 2, "kg", "DUNG", 20000, 30000, "0%", timestamp,
                ),
            )
        exported = self.client.get(f"/api/export/suppliers/{batch_id}")
        edited = self._edited_export(exported, order_qty=10, buy_price=50000)
        preview_response = self.client.post(
            "/api/purchase-orders/import/preview",
            data={"batch_id": str(batch_id), "file": (edited, "canonical.xlsx")},
            content_type="multipart/form-data",
        )
        self.assertEqual(preview_response.status_code, 200, preview_response.get_data(as_text=True))
        preview = preview_response.get_json()
        self.assertEqual(preview["count"], 2)
        with server.db() as conn:
            conn.execute(
                """CREATE TRIGGER fail_second_canonical_purchase
                   BEFORE INSERT ON purchase_workbook_lines WHEN NEW.source_row=4
                   BEGIN SELECT RAISE(ABORT,'forced canonical purchase failure'); END"""
            )
        failed = self.client.post(
            "/api/purchase-orders/import/confirm",
            json={"token": preview["token"], "confirmed": True},
        )
        self.assertEqual(failed.status_code, 500, failed.get_data(as_text=True))
        with server.db() as conn:
            conn.execute("DROP TRIGGER fail_second_canonical_purchase")
            self.assertEqual(conn.execute(
                "SELECT COUNT(*) FROM purchase_workbook_lines WHERE batch_id=?", (batch_id,),
            ).fetchone()[0], 0)
            self.assertEqual(conn.execute(
                "SELECT COUNT(*) FROM purchase_workbook_line_revisions WHERE batch_id=?",
                (batch_id,),
            ).fetchone()[0], 0)
            self.assertEqual(conn.execute(
                "SELECT COUNT(*) FROM purchase_order_imports WHERE batch_id=?", (batch_id,),
            ).fetchone()[0], 0)
            self.assertEqual(conn.execute(
                "SELECT COUNT(*) FROM audit_log WHERE event_type='purchase_orders.import'",
            ).fetchone()[0], 0)

    def test_outgoing_readiness_reports_invoiceable_and_pending_quantities(self):
        with server.db() as conn:
            batch_id, _ = self._insert_approved_order(conn)
            self._add_canonical_invoice_stock(conn, 4, "INV-PART")

        response = self.client.get(f"/api/outgoing-invoices/readiness/{batch_id}")
        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        payload = response.get_json()
        self.assertTrue(payload["ok"])
        self.assertEqual(payload["invoiceable_qty"], 4)
        self.assertEqual(payload["pending_qty"], 6)
        self.assertEqual(len(payload["rows"]), 1)
        self.assertEqual(payload["rows"][0]["demand_qty"], 10)
        self.assertEqual(payload["rows"][0]["invoiceable_qty"], 4)
        self.assertEqual(payload["rows"][0]["pending_qty"], 6)

    def test_outgoing_invoice_can_be_issued_in_two_inventory_rounds(self):
        with server.db() as conn:
            batch_id, order_id = self._insert_approved_order(conn)
            self._add_canonical_invoice_stock(conn, 4, "INV-ROUND-1")

        first = self.client.post(f"/api/outgoing-invoices/draft/{batch_id}")
        self.assertEqual(first.status_code, 200, first.get_data(as_text=True))
        with server.db() as conn:
            first_draft = dict(conn.execute(
                "SELECT * FROM outgoing_invoice_drafts WHERE batch_id=?", (batch_id,),
            ).fetchone())
            first_line = dict(conn.execute(
                "SELECT * FROM outgoing_invoice_lines WHERE draft_id=?", (first_draft["id"],),
            ).fetchone())
            self.assertEqual(first_draft["round_no"], 1)
            self.assertEqual(first_line["order_id"], order_id)
            self.assertEqual(first_line["qty"], 4)

            # The legal confirmation route is covered elsewhere.  Mark the
            # first round as issued here so this test can focus on allocation.
            conn.execute(
                "UPDATE outgoing_invoice_drafts SET status='issued',issued_at=? WHERE id=?",
                (server.now_iso(), first_draft["id"]),
            )
            conn.execute(
                """UPDATE inventory_transactions SET status='posted',updated_at=?
                   WHERE source_type='OUTGOING_DRAFT' AND source_id=? AND status='reserved'""",
                (server.now_iso(), str(first_draft["id"])),
            )
            self._add_canonical_invoice_stock(conn, 6, "INV-ROUND-2", "2026-09-06")

        second = self.client.post(f"/api/outgoing-invoices/draft/{batch_id}")
        self.assertEqual(second.status_code, 200, second.get_data(as_text=True))
        with server.db() as conn:
            drafts = [dict(row) for row in conn.execute(
                "SELECT * FROM outgoing_invoice_drafts WHERE batch_id=? ORDER BY round_no",
                (batch_id,),
            )]
            self.assertEqual([row["round_no"] for row in drafts], [1, 2])
            quantities = [conn.execute(
                "SELECT SUM(qty) qty FROM outgoing_invoice_lines WHERE draft_id=?", (row["id"],),
            ).fetchone()["qty"] for row in drafts]
            self.assertEqual(quantities, [4, 6])
            self.assertEqual(sum(quantities), 10)

        readiness = self.client.get(
            f"/api/outgoing-invoices/readiness/{batch_id}",
        ).get_json()
        self.assertEqual(readiness["allocated_qty"], 10)
        self.assertEqual(readiness["invoiceable_qty"], 0)
        self.assertEqual(readiness["pending_qty"], 0)

    def test_purchase_sheet_price_wins_over_sales_quotation(self):
        with server.db() as conn:
            batch_id, _ = self._insert_approved_order(conn, original_buy_price=12000)
        exported = self.client.get(f"/api/export/suppliers/{batch_id}")
        self.assertEqual(exported.status_code, 200)
        # The customer's purchase sheet owns cost even when sales has a quote.
        edited = self._edited_export(exported, order_qty=7, buy_price=50000)
        workbook = load_workbook(edited, read_only=True, data_only=True, keep_links=False)
        try:
            with server.db() as conn:
                preview = contract_modules.parse_purchase_order_workbook(
                    conn, workbook, batch_id,
                )
        finally:
            workbook.close()
        self.assertTrue(preview["can_confirm"])
        self.assertEqual(preview["rows"][0]["buy_price"], 50000)
        self.assertEqual(preview["rows"][0]["price_source"], "Sheet đặt hàng chuẩn")
        self.assertEqual(preview["total_amount"], 350000)

    def test_purchase_confirmation_rejects_a_stale_preview_atomically(self):
        with server.db() as conn:
            batch_id, order_id = self._insert_approved_order(conn)
        exported = self.client.get(f"/api/export/suppliers/{batch_id}")
        edited = self._edited_export(exported)
        preview_response = self.client.post(
            "/api/purchase-orders/import/preview",
            data={
                "batch_id": str(batch_id),
                "file": (edited, "Dat_NCC_da_chinh.xlsx"),
            },
            content_type="multipart/form-data",
        )
        self.assertEqual(preview_response.status_code, 200, preview_response.get_data(as_text=True))
        preview = preview_response.get_json()

        # Another operator changes the customer order between preview and
        # confirmation.  The stale file must not overwrite that newer state.
        with server.db() as conn:
            conn.execute("UPDATE orders SET qty=11 WHERE id=?", (order_id,))
        confirm = self.client.post(
            "/api/purchase-orders/import/confirm",
            json={"token": preview["token"], "confirmed": True},
        )
        self.assertEqual(confirm.status_code, 409, confirm.get_data(as_text=True))
        with server.db() as conn:
            self.assertEqual(conn.execute(
                "SELECT COUNT(*) n FROM purchase_order_lines WHERE batch_id=?", (batch_id,),
            ).fetchone()["n"], 0)
            self.assertEqual(conn.execute(
                "SELECT qty FROM orders WHERE id=?", (order_id,),
            ).fetchone()["qty"], 11)

    def test_confirmed_supplier_merge_policy_is_exact(self):
        for supplier in (
            "Dung", "NHÀ DŨNG", "Nhà Thu", "TÂN", "Nhà Phượng", "Kỳ", "Kho",
        ):
            with self.subTest(supplier=supplier):
                self.assertEqual(
                    contract_modules.supplier_merge_policy(supplier, "Hành lá"),
                    (True, True),
                )

        self.assertEqual(
            contract_modules.supplier_merge_policy("NHÀ HOÀI", "  CÀ-rốt "),
            (True, True),
        )
        self.assertEqual(
            contract_modules.supplier_merge_policy("Nhà Hoài", "Rau muống"),
            (False, False),
        )
        # The older note mentioned Hương, but the latest note and both calls do
        # not.  Do not silently broaden the confirmed customer rule.
        self.assertEqual(
            contract_modules.supplier_merge_policy("Nhà Hương", "Hành lá"),
            (False, False),
        )
        self.assertEqual(
            contract_modules.supplier_merge_policy("Nhà Thuận", "Hành lá"),
            (False, False),
        )

    def test_policy_managed_suppliers_cannot_be_overridden(self):
        for supplier in ("DUNG", "Nhà Hoài", "kho"):
            with self.subTest(supplier=supplier):
                response = self.client.put(
                    f"/api/supplier-rules/{supplier}",
                    json={"combine_kitchens": False},
                )
                self.assertEqual(response.status_code, 409, response.get_data(as_text=True))

        manual = self.client.put(
            "/api/supplier-rules/Nhà Hương",
            json={"combine_kitchens": True},
        )
        self.assertEqual(manual.status_code, 200, manual.get_data(as_text=True))
        with server.db() as conn:
            stored = conn.execute(
                "SELECT supplier_code,combine_kitchens FROM supplier_rules"
            ).fetchall()
        self.assertEqual([tuple(row) for row in stored], [("huong", 1)])

    def test_supplier_groups_merge_only_identical_names_under_confirmed_rules(self):
        with server.db() as conn:
            conn.execute(
                "INSERT OR REPLACE INTO kitchens(code,contractor,name,address,show_price) "
                "VALUES('BIA','HATRAN','Bếp BIA','',0)"
            )
            batch_id, _ = self._insert_approved_order(
                conn, qty=10, product_name="Hành lá",
            )
            timestamp = server.now_iso()
            conn.execute(
                """INSERT INTO orders(
                       batch_id,work_date,contractor,kitchen,product_code,product_name,qty,
                       actual_received,actual_delivered,unit,supplier,buy_price,sell_price,tax,
                       purchase_list,note,errors,warnings,updated_at
                   ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,0,'','[]','[]',?)""",
                (
                    batch_id, "2026-09-01", "HATRAN", "BIA", "PO-P1", "Hành lá",
                    5, 5, 5, "kg", "DUNG", 0, 30000, "0%", timestamp,
                ),
            )
            dung = contract_modules.purchase_order_payload(conn, batch_id)
            self.assertEqual(len(dung["groups"]), 1)
            self.assertEqual(len(dung["groups"][0]["items"]), 1)
            self.assertEqual(dung["groups"][0]["items"][0]["demand_qty"], 15)
            self.assertEqual(len(dung["groups"][0]["items"][0]["order_ids"]), 2)

            conn.execute(
                "UPDATE orders SET supplier='HOAI',product_name='Rau muống' WHERE batch_id=?",
                (batch_id,),
            )
            hoai_other = contract_modules.purchase_order_payload(conn, batch_id)
            self.assertEqual(len(hoai_other["groups"]), 2)

            conn.execute(
                "UPDATE orders SET product_name='Cà rốt' WHERE batch_id=?", (batch_id,),
            )
            hoai_carrot = contract_modules.purchase_order_payload(conn, batch_id)
            self.assertEqual(len(hoai_carrot["groups"]), 1)
            self.assertEqual(len(hoai_carrot["groups"][0]["items"]), 1)
            self.assertEqual(hoai_carrot["groups"][0]["items"][0]["demand_qty"], 15)

    def test_canonical_merge_is_presentation_only_and_unicode_stable(self):
        with server.db() as conn:
            batch_id, _ = self._insert_approved_order(conn)
            timestamp = server.now_iso()
            fixtures = [
                ("Nhà Dũng", "POT", "Cà rốt", 2),
                ("DUNG", "BIA", "cà rốt", 3),
                ("Dung", "BIA", "Cà rốt loại 1", 4),
                ("NHÀ HOÀI", "POT", "CÀ-RỐT", 5),
                ("hoài", "BIA", "cà rốt", 6),
                ("Nhà Hoài", "POT", "Rau muống", 7),
                ("HOÀI", "POT", "rau MUỐNG", 8),
                ("Hoài", "BIA", "Rau muống", 9),
                ("HƯƠNG", "POT", "Hành lá", 10),
                ("Nhà Hương", "POT", "hành LÁ", 11),
                ("Hương", "BIA", "Hành lá", 12),
            ]
            for index, (supplier, kitchen, product_name, qty) in enumerate(fixtures, start=1):
                conn.execute(
                    """INSERT INTO purchase_workbook_lines(
                           batch_id,row_key,source_sheet,source_row,product_code,kitchen,
                           work_date,product_name,base_qty,unit,supplier,buy_price,
                           actual_qty,amount,source_hash,created_at,updated_at
                       ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                    (
                        batch_id, f"raw-{index}", "đặt hàng", index + 2,
                        f"P-MERGE-{index}", kitchen, "2026-09-01", product_name,
                        qty, "kg", supplier, 100, qty, qty * 100,
                        "merge-source", timestamp, timestamp,
                    ),
                )
            raw_before = [tuple(row) for row in conn.execute(
                """SELECT row_key,source_row,kitchen,product_name,base_qty,supplier,actual_qty
                   FROM purchase_workbook_lines WHERE batch_id=? ORDER BY source_row""",
                (batch_id,),
            )]
            payload = contract_modules.purchase_order_payload(conn, batch_id)
            raw_after = [tuple(row) for row in conn.execute(
                """SELECT row_key,source_row,kitchen,product_name,base_qty,supplier,actual_qty
                   FROM purchase_workbook_lines WHERE batch_id=? ORDER BY source_row""",
                (batch_id,),
            )]

        self.assertEqual(raw_after, raw_before)
        self.assertEqual((payload["raw_line_count"], payload["presented_line_count"]), (11, 9))

        dung = next(
            group for group in payload["groups"]
            if contract_modules.supplier_merge_key(group["supplier"]) == "dung"
        )
        self.assertTrue(dung["policy_managed"])
        self.assertTrue(dung["automatic_merge"])
        self.assertEqual((dung["raw_line_count"], dung["presented_line_count"]), (3, 2))
        dung_carrot = next(
            item for item in dung["items"]
            if contract_modules.mapping_key(item["product_name"]) == "carot"
        )
        self.assertEqual(dung_carrot["order_qty"], 5)
        self.assertEqual(
            {ref["row_key"] for ref in dung_carrot["source_refs"]},
            {"raw-1", "raw-2"},
        )

        hoai_groups = [
            group for group in payload["groups"]
            if contract_modules.supplier_merge_key(group["supplier"]) == "hoai"
        ]
        self.assertEqual(len(hoai_groups), 3)
        hoai_carrot = next(group for group in hoai_groups if group["automatic_merge"])
        self.assertEqual((hoai_carrot["raw_line_count"], len(hoai_carrot["items"])), (2, 1))
        self.assertEqual(hoai_carrot["items"][0]["order_qty"], 11)
        hoai_other = [group for group in hoai_groups if not group["automatic_merge"]]
        self.assertEqual(sum(group["raw_line_count"] for group in hoai_other), 3)
        self.assertEqual(sum(len(group["items"]) for group in hoai_other), 3)

        huong_groups = [
            group for group in payload["groups"]
            if contract_modules.supplier_merge_key(group["supplier"]) == "huong"
        ]
        self.assertEqual(len(huong_groups), 2)
        self.assertTrue(all(not group["policy_managed"] for group in huong_groups))
        self.assertEqual(sum(group["raw_line_count"] for group in huong_groups), 3)
        self.assertEqual(sum(len(group["items"]) for group in huong_groups), 3)


if __name__ == "__main__":
    unittest.main()
