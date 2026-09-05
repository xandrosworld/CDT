import io
import math
import tempfile
import unittest
from datetime import date
from pathlib import Path

from openpyxl import Workbook

try:
    from . import server
except ImportError:  # pragma: no cover - direct file invocation
    import server


class ServerFinancialGuardTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp_dir = tempfile.TemporaryDirectory()
        cls.original_db_path = server.DB_PATH
        cls.original_data_dir = server.DATA_DIR
        cls.original_testing = server.app.config.get("TESTING", False)
        server.DATA_DIR = Path(cls.temp_dir.name) / "data"
        server.DATA_DIR.mkdir(parents=True, exist_ok=True)
        server.DB_PATH = server.DATA_DIR / "guard_test.sqlite3"
        server.app.config["TESTING"] = True
        server.init_database()
        cls.client = server.app.test_client()

    @classmethod
    def tearDownClass(cls):
        server.DB_PATH = cls.original_db_path
        server.DATA_DIR = cls.original_data_dir
        server.app.config["TESTING"] = cls.original_testing
        cls.temp_dir.cleanup()

    def setUp(self):
        with server.db() as conn:
            conn.execute("DELETE FROM outgoing_invoice_lines")
            conn.execute("DELETE FROM outgoing_invoice_drafts")
            conn.execute("DELETE FROM inventory_transactions")
            conn.execute("DELETE FROM order_import_receipts")
            conn.execute("DELETE FROM orders")
            conn.execute("DELETE FROM batches")
            conn.execute("DELETE FROM historical_payable_lines")
            conn.execute("DELETE FROM settings WHERE key='historical_payables_through_date'")
            conn.execute("DELETE FROM attendance_entries")
            conn.execute("DELETE FROM payroll_adjustments")
            conn.execute("DELETE FROM staff")
            conn.execute("DELETE FROM payments")
            conn.execute("DELETE FROM balances")
            conn.execute("DELETE FROM debt_adjustments")
            conn.execute("DELETE FROM audit_log")
            conn.execute(
                "INSERT OR REPLACE INTO contractors(code,name,price_group,pricing_mode) "
                "VALUES('HATRAN','Hà Trân','HATRAN','group')"
            )
            conn.execute(
                "INSERT OR REPLACE INTO suppliers(code,name) VALUES('NCC-A','Nhà cung cấp A')"
            )
            conn.execute(
                """INSERT OR REPLACE INTO products(
                       code,name,unit,tax,supplier,buy_price,purchase_list,seller,cccd
                   ) VALUES('RES-P1','Hàng kiểm thử giữ tồn','kg','0%','NCC-A',10,0,'','')"""
            )
        response = self.client.post("/api/staff", json={
            "employee_code": "NVTEST",
            "full_name": "Nhân viên test",
            "role_name": "Bếp",
            "kitchen": "POT",
            "base_salary": 5_200_000,
        })
        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))

    @staticmethod
    def order_row():
        return {
            "contractor": "HATRAN",
            "kitchen": "POT",
            "supplier": "NCC-A",
            "actual_received": 2,
            "actual_delivered": 2,
            "damaged_qty": 0,
            "supplier_return_qty": 0,
            "customer_return_qty": 0,
            "buy_price": 60_000,
            "sell_price": 100_000,
            "tax": 0.1,
            "errors": [],
            "warnings": [],
        }

    def audit_count(self, event_type):
        with server.db() as conn:
            return conn.execute(
                "SELECT COUNT(*) n FROM audit_log WHERE event_type=?", (event_type,)
            ).fetchone()["n"]

    @staticmethod
    def insert_approved_order(conn, work_date, qty, *, buy_price=10, sell_price=20):
        timestamp = server.now_iso()
        batch_id = conn.execute(
            """INSERT INTO batches(work_date,source_name,status,created_at,approved_at)
               VALUES(?,?,'approved',?,?)""",
            (work_date, f"regression-{work_date}.xlsx", timestamp, timestamp),
        ).lastrowid
        conn.execute(
            """INSERT INTO orders(
                   batch_id,work_date,contractor,kitchen,product_code,product_name,qty,
                   actual_received,actual_delivered,unit,supplier,buy_price,sell_price,tax,
                   purchase_list,errors,warnings,updated_at
               ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,0,'[]','[]',?)""",
            (
                batch_id, work_date, "HATRAN", "POT", "RES-P1",
                "Hàng kiểm thử giữ tồn", qty, qty, qty, "kg", "NCC-A",
                buy_price, sell_price, "0%", timestamp,
            ),
        )
        return batch_id

    def test_future_reservation_limits_a_later_backdated_draft_to_available_stock(self):
        """A future reservation stays global while the available part may be drafted."""
        with server.db() as conn:
            timestamp = server.now_iso()
            conn.execute(
                """INSERT INTO inventory_transactions(
                       txn_date,product_code,qty_in,qty_out,unit_cost,source_type,source_id,
                       source_line,status,note,created_at,updated_at
                   ) VALUES('2026-08-01','RES-P1',10,0,10,'OPENING','2026-08',
                            'RES-P1','posted','opening',?,?)""",
                (timestamp, timestamp),
            )
            future_batch = self.insert_approved_order(conn, "2026-09-10", 7)
            backdated_batch = self.insert_approved_order(conn, "2026-08-10", 4)

        future = self.client.post(f"/api/outgoing-invoices/draft/{future_batch}")
        self.assertEqual(future.status_code, 200, future.get_data(as_text=True))

        with server.db() as conn:
            stock = server.inventory_lookup(conn, "2026-08-10")["RES-P1"]
            self.assertEqual(stock["accounting_qty"], 10)
            self.assertEqual(stock["available_qty"], 3)

        backdated = self.client.post(f"/api/outgoing-invoices/draft/{backdated_batch}")
        self.assertEqual(backdated.status_code, 200, backdated.get_data(as_text=True))
        self.assertTrue(backdated.get_json()["partial"])
        self.assertEqual(backdated.get_json()["pending_qty"], 1)
        with server.db() as conn:
            draft = conn.execute(
                "SELECT id FROM outgoing_invoice_drafts WHERE batch_id=?", (backdated_batch,),
            ).fetchone()
            self.assertIsNotNone(draft)
            self.assertEqual(conn.execute(
                "SELECT SUM(qty) qty FROM outgoing_invoice_lines WHERE draft_id=?", (draft["id"],),
            ).fetchone()["qty"], 3)
            reservation = conn.execute(
                """SELECT COUNT(*) lines,COALESCE(SUM(qty_out),0) qty
                   FROM inventory_transactions
                   WHERE source_type='OUTGOING_DRAFT' AND status='reserved'"""
            ).fetchone()
            self.assertEqual((reservation["lines"], reservation["qty"]), (2, 10))

    @staticmethod
    def payable_workbook_with_two_tables():
        workbook = Workbook()
        headers = ["Ngày", "Tên bếp", "Tên hàng", "Số lượng", "ĐVT", "NCC", "Giá mua", "Thành tiền"]
        for index, title in enumerate(("Tổng hợp A", "Tổng hợp B")):
            sheet = workbook.active if index == 0 else workbook.create_sheet()
            sheet.title = title
            sheet.append(headers)
            sheet.append([date(2026, 8, 15), "POT", f"Hàng {index + 1}", 1, "kg", "NCC-A", 100, 100])
        stream = io.BytesIO()
        workbook.save(stream)
        workbook.close()
        stream.seek(0)
        return stream

    def test_payables_preview_rejects_multiple_valid_sheets(self):
        response = self.client.post(
            "/api/debts/payables/import/preview",
            data={"file": (self.payable_workbook_with_two_tables(), "payables.xlsx")},
            content_type="multipart/form-data",
        )
        self.assertEqual(response.status_code, 400, response.get_data(as_text=True))
        self.assertIn("nhiều bảng công nợ", response.get_json()["error"])
        with server.db() as conn:
            self.assertEqual(
                conn.execute("SELECT COUNT(*) n FROM historical_payable_lines").fetchone()["n"], 0,
            )

    def test_historical_payables_cutoff_prevents_order_double_count(self):
        with server.db() as conn:
            timestamp = server.now_iso()
            conn.execute(
                """INSERT INTO historical_payable_lines(
                       purchase_date,kitchen,item_name,qty,unit,supplier,buy_price,
                       actual_qty,source_amount,calculated_amount,amount,note,source_file,
                       source_sheet,source_row,source_hash,created_at,updated_at
                   ) VALUES('2026-08-31','POT','Ảnh chụp công nợ',1,'kg','NCC-A',100,
                            1,100,100,100,'snapshot','payables.xlsx','Tổng hợp',2,
                            'HASH-SNAPSHOT',?,?)""",
                (timestamp, timestamp),
            )
            # Both approved orders through 31/08 are already represented by the
            # authoritative history snapshot and must be skipped for suppliers.
            self.insert_approved_order(conn, "2026-08-15", 1, buy_price=100)
            self.insert_approved_order(conn, "2026-08-31", 1, buy_price=50)
            self.insert_approved_order(conn, "2026-09-01", 1, buy_price=200)

        for explicit_setting in (False, True):
            with self.subTest(explicit_setting=explicit_setting), server.db() as conn:
                if explicit_setting:
                    server.setting_set(conn, "historical_payables_through_date", "2026-08-31")
                else:
                    conn.execute("DELETE FROM settings WHERE key='historical_payables_through_date'")
            response = self.client.get("/api/debts?from=2026-08-01&to=2026-09-30")
            self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
            payload = response.get_json()
            self.assertEqual(payload["historical_payables_through_date"], "2026-08-31")
            self.assertEqual(payload["historical_payables"], {
                "rows": 1, "amount": 100, "period_rows": 1, "period_amount": 100,
            })
            # 100 historical + 200 after cutoff. The 100 and 50 approved-order
            # charges on/before cutoff are intentionally not added again.
            self.assertEqual(payload["suppliers"]["NCC-A"]["period_charge"], 300)
            self.assertEqual(payload["suppliers"]["NCC-A"]["closing"], 300)

    def test_batch_summary_never_subtracts_global_ledger_entries(self):
        with server.db() as conn:
            conn.execute(
                "INSERT INTO balances(party_type,party_code,opening,as_of_date) VALUES('contractor','HATRAN',999999,'2026-01-01')"
            )
            conn.execute(
                """INSERT INTO payments(payment_date,kind,party_type,party_code,amount,note,created_at)
                   VALUES('2026-08-01','receipt','contractor','HATRAN',888888,'old',?)""",
                (server.now_iso(),),
            )
            summary = server.calculate_summary(conn, [self.order_row()])
            self.assertEqual(summary["scope"], "batch_only")
            self.assertEqual(summary["contractors"]["HATRAN"]["opening"], 0)
            self.assertEqual(summary["contractors"]["HATRAN"]["paid"], 0)
            self.assertAlmostEqual(summary["contractors"]["HATRAN"]["balance"], 220_000)
            workbook = server.export_report(
                conn, {"work_date": "2026-08-30"}, [self.order_row()]
            )
        try:
            self.assertEqual(workbook.sheetnames, ["báo cáo tổng hợp"])
            sheet = workbook["báo cáo tổng hợp"]
            self.assertEqual(sheet["B2"].value, "Khách hàng")
            pot = next(row for row in sheet.iter_rows(values_only=True) if row[2] == "POT")
            self.assertEqual(pot[3:7], (200_000, 120_000, 80_000, 220_000))
            self.assertEqual(sheet.cell(sheet.max_row, 1).value, "TỔNG THÁNG")
        finally:
            workbook.close()

    def test_payment_validation_and_create_reverse_audit(self):
        invalid_payloads = [
            {"payment_date": "2026-02-30", "kind": "receipt", "party_type": "contractor", "party_code": "HATRAN", "amount": 1},
            {"payment_date": "2026-08-30", "kind": "refund", "party_type": "contractor", "party_code": "HATRAN", "amount": 1},
            {"payment_date": "2026-08-30", "kind": "receipt", "party_type": "supplier", "party_code": "HATRAN", "amount": 1},
            {"payment_date": "2026-08-30", "kind": "receipt", "party_type": "contractor", "party_code": "", "amount": 1},
            {"payment_date": "2026-08-30", "kind": "receipt", "party_type": "contractor", "party_code": "HATRAN", "amount": 0},
            {"payment_date": "2026-08-30", "kind": "receipt", "party_type": "contractor", "party_code": "HATRAN", "amount": "not-a-number"},
        ]
        for payload in invalid_payloads:
            with self.subTest(payload=payload):
                self.assertEqual(self.client.post("/api/payments", json=payload).status_code, 400)

        created = self.client.post("/api/payments", json={
            "payment_date": "2026-08-30", "kind": "receipt",
            "party_type": "contractor", "party_code": "hatran",
            "amount": 125_000, "note": "Thu chuyển khoản",
            "actor": "Người thử", "request_id": "GUARD-RECEIPT-0001",
        })
        self.assertEqual(created.status_code, 201, created.get_data(as_text=True))
        payment_id = created.get_json()["id"]
        with server.db() as conn:
            self.assertEqual(
                conn.execute("SELECT party_code FROM payments WHERE id=?", (payment_id,)).fetchone()["party_code"],
                "HATRAN",
            )
        self.assertEqual(self.client.post("/api/payments", json={
            "payment_date": "2026-08-30", "kind": "receipt",
            "party_type": "contractor", "party_code": "HATRANN", "amount": 1,
        }).status_code, 400)
        self.assertEqual(self.audit_count("payment.create"), 1)
        self.assertEqual(self.client.delete("/api/payments/999999").status_code, 404)
        self.assertEqual(self.client.delete(f"/api/payments/{payment_id}").status_code, 409)
        self.assertEqual(self.client.post(f"/api/debts/receipts/{payment_id}/reverse", json={
            "actor": "Người thử", "reason": "Nhập nhầm", "expected_revision": 1,
        }).status_code, 200)
        self.assertEqual(self.audit_count("payment.delete"), 0)
        self.assertEqual(self.audit_count("payment.reverse"), 1)

    def test_balance_and_debt_adjustment_validation_and_audit(self):
        self.assertEqual(self.client.post("/api/balances", json={
            "party_type": "invalid", "party_code": "A", "opening": 1,
            "as_of_date": "2026-08-01",
        }).status_code, 400)
        self.assertEqual(self.client.post("/api/balances", json={
            "party_type": "supplier", "party_code": "A", "opening": "bad",
            "as_of_date": "2026-08-01",
        }).status_code, 400)
        saved = self.client.post("/api/balances", json={
            "party_type": "supplier", "party_code": "NCC-A", "opening": -50_000,
            "as_of_date": "2026-08-01",
        })
        self.assertEqual(saved.status_code, 200, saved.get_data(as_text=True))
        self.assertEqual(self.audit_count("balance.upsert"), 1)

        invalid_adjustments = [
            {"adjustment_date": "bad", "party_type": "supplier", "party_code": "NCC-A", "amount": 1, "note": "x"},
            {"adjustment_date": "2026-08-30", "party_type": "supplier", "party_code": "", "amount": 1, "note": "x"},
            {"adjustment_date": "2026-08-30", "party_type": "supplier", "party_code": "NCC-A", "amount": 0, "note": "x"},
            {"adjustment_date": "2026-08-30", "party_type": "supplier", "party_code": "NCC-A", "amount": 1, "note": ""},
        ]
        for payload in invalid_adjustments:
            with self.subTest(payload=payload):
                self.assertEqual(self.client.post("/api/debt-adjustments", json=payload).status_code, 400)
        adjusted = self.client.post("/api/debt-adjustments", json={
            "adjustment_date": "2026-08-30", "party_type": "supplier",
            "party_code": "NCC-A", "amount": -12_500, "note": "Giảm đối chiếu",
        })
        self.assertEqual(adjusted.status_code, 200, adjusted.get_data(as_text=True))
        self.assertEqual(self.audit_count("debt.adjustment.create"), 1)

    def test_attendance_validation_and_audit(self):
        base = {"employee_code": "NVTEST", "work_date": "2026-08-30"}
        bad_rows = [
            {**base, "work_date": "2026-02-30", "normal_hours": 8},
            {**base, "normal_hours": -1},
            {**base, "overtime_hours": 25},
            {**base, "normal_hours": 20, "overtime_hours": 5},
            {**base, "normal_hours": "bad"},
        ]
        for payload in bad_rows:
            with self.subTest(payload=payload):
                self.assertEqual(self.client.post("/api/attendance", json=payload).status_code, 400)
        self.assertEqual(self.client.post("/api/attendance", json={
            **base, "normal_hours": 8, "overtime_hours": 2, "night_hours": 2,
        }).status_code, 200)
        self.assertEqual(self.audit_count("attendance.upsert"), 1)

    def test_health_checks_schema_and_backup_names_are_unique(self):
        healthy = self.client.get("/health")
        self.assertEqual(healthy.status_code, 200, healthy.get_data(as_text=True))
        self.assertTrue(healthy.get_json()["database_ready"])

        first = self.client.get("/api/backup")
        second = self.client.get("/api/backup")
        try:
            self.assertEqual(first.status_code, 200)
            self.assertEqual(second.status_code, 200)
            self.assertNotEqual(
                first.headers["Content-Disposition"], second.headers["Content-Disposition"]
            )
        finally:
            first.close()
            second.close()
        self.assertEqual(
            self.client.get("/api/backup", environ_base={"REMOTE_ADDR": "192.168.1.25"}).status_code,
            403,
        )
        self.assertEqual(self.client.post(
            "/api/payments",
            headers={"Origin": "https://evil.invalid"},
            json={
                "payment_date": "2026-08-30", "kind": "receipt",
                "party_type": "contractor", "party_code": "HATRAN", "amount": 1,
            },
        ).status_code, 403)

        original = server.DB_PATH
        server.DB_PATH = server.DATA_DIR / "empty.sqlite3"
        server.DB_PATH.touch()
        try:
            unhealthy = self.client.get("/health")
            self.assertEqual(unhealthy.status_code, 503)
            self.assertFalse(unhealthy.get_json()["schema_ready"])
        finally:
            server.DB_PATH = original

    def test_missing_purchase_list_cccd_warns_without_blocking_import(self):
        with server.db() as conn:
            by_code, by_name = server.product_lookup(conn)
            resolved = server.resolve_order(conn, {
                "work_date": "2026-08-29",
                "contractor": "HATRAN",
                "kitchen": "POT",
                "product_code": "RES-P1",
                "product_name": "Hàng bảng kê chưa có CCCD",
                "qty": 1,
                "supplier": "NCC-A",
                "buy_price": 10_000,
                "sell_price": 12_000,
                "tax": "0%",
                "purchase_list": "BK",
                "seller": "Người bán chưa bổ sung CCCD",
                "cccd": "",
            }, "2026-08-29", by_code, by_name)

        self.assertFalse(any("CCCD" in error for error in resolved["errors"]))
        self.assertTrue(any("bổ sung sau" in warning for warning in resolved["warnings"]))
        self.assertEqual(resolved["cccd"], "")

    def test_order_quantities_and_tax_codes_are_strict(self):
        self.assertEqual(server.normalize_tax(-2), "KKKNT")
        self.assertEqual(server.normalize_tax(-1), "KCT")
        self.assertEqual(server.normalize_tax("8%"), 0.08)
        self.assertEqual(server.normalize_tax("15%"), "INVALID")
        self.assertEqual(server.normalize_tax("rác"), "INVALID")
        with server.db() as conn:
            by_code, by_name = server.product_lookup(conn)
            raw = {
                "work_date": "2026-08-30", "kitchen": "POT",
                "product_code": "I000060", "qty": 2,
                "actual_received": -1, "actual_delivered": -2,
                "damaged_qty": -0.1, "supplier_return_qty": 0,
                "customer_return_qty": 0, "supplier": "NCC-A",
                "buy_price": 10_000, "sell_price": 12_000, "tax": "15%",
            }
            resolved = server.resolve_order(
                conn, raw, "2026-08-30", by_code, by_name
            )
        self.assertIn("Số lượng thực nhận không được âm", resolved["errors"])
        self.assertIn("Số lượng thực giao không được âm", resolved["errors"])
        self.assertIn("Số lượng hỏng/trả lại không được âm", resolved["errors"])
        self.assertTrue(any("Thuế suất" in error for error in resolved["errors"]))
        for invalid in (float("nan"), float("inf"), float("-inf")):
            with self.subTest(nonfinite=repr(invalid)), server.db() as conn:
                by_code, by_name = server.product_lookup(conn)
                bad = server.resolve_order(conn, {
                    **raw,
                    "actual_received": invalid,
                    "actual_delivered": invalid,
                    "qty": invalid,
                    "buy_price": invalid,
                    "sell_price": invalid,
                    "tax": "8%",
                }, "2026-08-30", by_code, by_name)
                self.assertTrue(any("hữu hạn" in error for error in bad["errors"]))
                for key in ("qty", "actual_received", "actual_delivered", "buy_price", "sell_price"):
                    self.assertTrue(math.isfinite(bad[key]), key)


if __name__ == "__main__":
    unittest.main()
