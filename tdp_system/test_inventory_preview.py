"""The browser preview must show the actual export's cells without database writes."""
import io
import re
import unittest
from datetime import date, datetime
from contextlib import contextmanager

from flask import Flask
from openpyxl import Workbook, load_workbook

from . import test_inventory_export as fixtures
from .inventory_export import register_inventory_export_routes
from .inventory_preview import workbook_preview


class InventoryPreviewTests(unittest.TestCase):
    def setUp(self):
        self.fixture = fixtures.InventoryExportTests()
        self.fixture.setUp()
        self.conn = self.fixture.conn
        @contextmanager
        def db():
            yield self.conn
        app = Flask(__name__)
        app.config['TESTING'] = True
        register_inventory_export_routes(app, {'db': db, 'opening_template_path': fixtures.OPENING_TEMPLATE})
        self.client = app.test_client()

    def tearDown(self):
        self.fixture.tearDown()

    def test_all_reports_match_download_cells_merges_and_formats_without_writes(self):
        before = self.conn.total_changes
        for kind in ('opening', 'input', 'output', 'nxt'):
            with self.subTest(kind=kind):
                query = '?from=2026-08-01&to=2026-08-31'
                preview = self.client.get('/api/invoice-valuation/preview/' + kind + query)
                self.assertEqual(200, preview.status_code, preview.json)
                self.assertTrue(preview.json['read_only'])
                self.assertEqual('2026-08-01', preview.json['date_from'])
                exported = self.client.get('/api/invoice-valuation/export/' + kind + query)
                workbook = load_workbook(io.BytesIO(exported.data))
                sheets = preview.json['workbook']['sheets']
                visible = [sheet for sheet in workbook if sheet.sheet_state == 'visible']
                self.assertEqual([sheet.title for sheet in visible], [sheets[key]['name'] for key in preview.json['workbook']['sheetOrder']])
                for original, key in zip(visible, preview.json['workbook']['sheetOrder']):
                    saved = sheets[key]
                    for row in original.iter_rows(max_row=saved['rowCount'], max_col=saved['columnCount']):
                        for cell in row:
                            if cell.value is None:
                                continue
                            actual = saved['cellData'][str(cell.row - 1)][str(cell.column - 1)]
                            expected = cell.value.strftime('%d/%m/%Y') if isinstance(cell.value, (date,datetime)) else cell.value
                            self.assertEqual(expected, actual['v'], (kind, original.title, cell.coordinate))
                            if isinstance(cell.value, (float, int)):
                                self.assertEqual(2, actual['t'])
                                pattern = re.sub(r'\.#+', '', cell.number_format) if float(cell.value).is_integer() else cell.number_format
                                self.assertEqual(pattern, actual['s']['n']['pattern'])
                    self.assertEqual(len([m for m in original.merged_cells.ranges if m.max_row<=saved['rowCount'] and m.max_col<=saved['columnCount']]), len(saved['mergeData']))
                workbook.close()
        self.assertEqual(before, self.conn.total_changes)

    def test_invalid_kind_dates_empty_period_and_no_internal_sheet(self):
        self.assertEqual(404,self.client.get('/api/invoice-valuation/preview/unknown?from=2026-08-01&to=2026-08-31').status_code)
        self.assertEqual(400,self.client.get('/api/invoice-valuation/preview/nxt?from=2026-08-31&to=2026-08-01').status_code)
        for kind in ('opening','input','output','nxt'):
            response=self.client.get('/api/invoice-valuation/preview/'+kind+'?from=2020-01-01&to=2020-01-31')
            self.assertEqual(200,response.status_code,response.json)
            self.assertNotIn('_ĐỐI_CHIẾU',[s['name'] for s in response.json['workbook']['sheets'].values()])

    def test_literal_codes_and_untrusted_text_are_not_formulas(self):
        book=Workbook();sheet=book.active
        for i,value in enumerate(['00123','=1+1','<img src=x onerror=alert(1)>'],1):
            sheet.cell(i,1,value).data_type='s'
        sheet.cell(4,1,0.855).number_format='#,##0.######'
        result=workbook_preview(book,'Literal')['sheets']['report-0']['cellData']
        self.assertEqual('00123',result[0][0]['v'])
        self.assertEqual('=1+1',result[1][0]['v']);self.assertEqual(1,result[1][0]['t'])
        self.assertEqual(0.855,result[3][0]['v'])
        sheet.cell(5,1,'=SUM(A1:A4)')
        with self.assertRaises(ValueError):workbook_preview(book,'Formula')


if __name__ == '__main__':
    unittest.main()
