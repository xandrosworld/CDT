import io
import json
import tempfile
import unittest
from datetime import date, datetime
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

from openpyxl import load_workbook
from . import server
from .document_totals import quantity_cell
from .receivable_ledger import sync_receivable_ledger
from .payable_ledger import sync_payable_ledger


def seed_round3(conn, count=2):
    stamp = server.now_iso()
    conn.execute("INSERT INTO contractors(code,name,price_group) VALUES('C1','Nhà thầu thử','C1')")
    conn.execute("INSERT INTO contractors(code,name,price_group) VALUES('C2','Nhà thầu Hai','C2')")
    conn.execute("INSERT INTO kitchens(code,name,contractor) VALUES('K1','Bếp Một','C1')")
    conn.execute("INSERT INTO kitchens(code,name,contractor) VALUES('K2','Bếp mới chưa có đơn','C1')")
    for code in ('S1', 'S2'):
        conn.execute("INSERT INTO suppliers(code,name) VALUES(?,?)", (code, 'NCC '+code))
    batch = conn.execute("INSERT INTO batches(work_date,source_name,status,created_at,approved_at) VALUES('2026-09-05','Round 3','approved',?,?)", (stamp, stamp)).lastrowid
    for index in range(count):
        unit = 'kg' if index % 2 == 0 else 'cái'
        qty = .855 if index % 2 == 0 else 2
        conn.execute("""INSERT INTO orders(batch_id,work_date,contractor,kitchen,product_code,product_name,
                        qty,actual_received,actual_delivered,unit,supplier,buy_price,sell_price,tax,
                        purchase_list,source_sheet,source_row,errors,warnings,updated_at)
                        VALUES(?,'2026-09-05','C1','K1',?,?,?,?,?,?,?,1000,152000.994706,'0',0,'Đơn hàng',?,'[]','[]',?)""",
                     (batch, f'P{index}', 'Tên hàng dài để kiểm tra căn dòng '+str(index), qty, qty, qty, unit,
                      'S1' if index % 2 == 0 else 'S2', index+1, stamp))
    # Current payable accounting requires a confirmed purchase-sheet source.
    conn.execute("""INSERT INTO purchase_workbook_lines(
        batch_id,row_key,work_date,product_code,kitchen,product_name,base_qty,actual_qty,
        unit,supplier,buy_price,amount,source_sheet,source_row,source_hash,created_at,updated_at)
        SELECT batch_id,'fixture:'||id,work_date,product_code,kitchen,product_name,qty,actual_received,
               unit,supplier,buy_price,ROUND(actual_received*buy_price),'đặt hàng',source_row,'fixture:'||id,?,?
        FROM orders WHERE batch_id=?""", (stamp, stamp, batch))
    sync_receivable_ledger(conn, timestamp=stamp)
    sync_payable_ledger(conn, timestamp=stamp)
    return batch


class Round3DocumentsTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.original = server.DB_PATH, server.DATA_DIR, server.MASTER_SOURCE
        server.DATA_DIR = Path(self.temp.name)
        server.DB_PATH = Path(self.temp.name) / 'round3.sqlite3'
        server.MASTER_SOURCE = Path(self.temp.name) / 'no-auto-import.xlsx'
        server.init_database()
        server.MASTER_SOURCE = Path(__file__).resolve().parent.parent / 'Em Thành.xlsx'
        with server.db() as conn:
            self.batch = seed_round3(conn)
        self.client = server.app.test_client()

    def tearDown(self):
        server.DB_PATH, server.DATA_DIR, server.MASTER_SOURCE = self.original
        self.temp.cleanup()

    def receipt(self, **changes):
        return {'kind':'receipt','party_type':'contractor','party_code':'C1','amount':100,
                'payment_date':'2026-09-05','note':'Thu tiền thử','actor':'Người kiểm thử',
                'request_id':'ROUND3-RECEIPT-0001', **changes}

    def export(self, url):
        r = self.client.get(url)
        self.assertEqual(r.status_code, 200, r.json)
        return load_workbook(io.BytesIO(r.data), data_only=False)

    def test_report_catalog_zero_kitchen_approved_only_and_excel_match(self):
        with server.db() as conn:
            conn.execute("INSERT INTO batches(work_date,source_name,status,created_at) VALUES('2026-09-10','Chưa duyệt','draft',?)", (server.now_iso(),))
        data = self.client.get('/api/reports/monthly?period=2026-09').json
        self.assertTrue(data['ok'], data)
        self.assertEqual(data['draft_count'], 1)
        self.assertEqual(next(r for r in data['rows'] if r[2] == 'K2')[3:], [0,0,0,0])
        book = self.export('/api/reports/monthly/export?period=2026-09')
        self.assertEqual([list(r) for r in book.active.iter_rows(min_row=3,values_only=True)], data['rows'])
        self.assertEqual(data['rows'][-1][0], 'TỔNG THÁNG')
        book.close()
        empty = self.client.get('/api/reports/monthly?period=2026-08').json
        self.assertEqual(empty['rows'][-1][3:], [0,0,0,0])
        for bad in ('2026-13', 'no', '2026-00'):
            self.assertEqual(self.client.get('/api/reports/monthly?period='+bad).status_code,400)

    def test_quantity_units_and_filtered_receivable_file(self):
        base = '?from=2026-09-01&to=2026-09-30&contractor=C1&kitchen=K1&status=active'
        data = self.client.get('/api/debts/receivables/ledger'+base).json
        summary = data['summary']
        self.assertEqual({i['unit']:i['quantity'] for i in summary['quantities_by_unit']}, {'kg':.855,'cái':2})
        book = self.export('/api/debts/receivables/lines/export'+base)
        self.assertEqual(book.active.cell(book.active.max_row,10).value, summary['filtered_amount'])
        self.assertIsNone(book.active.cell(book.active.max_row,5).value)
        self.assertEqual(book.active['E4'].value,.855)
        book.close()
        book = self.export('/api/debts/receivables/export?from=2026-09-01&to=2026-09-30&contractor=C1')
        self.assertNotIn('Tổng nhà thầu', book.sheetnames)
        forbidden = {'Mã bếp', 'Mã hàng', 'Thực giao', 'Khách trả', 'Giao ròng', 'Mã dòng', 'Lần cập nhật', 'Nguồn', 'Trạng thái'}
        for sheet in book:
            self.assertFalse(forbidden.intersection(cell.value for cell in sheet[3]))
        book.close()

    def assert_preview_matches_export(self, root, query):
        response = self.client.get(root + '/preview' + query)
        self.assertEqual(response.status_code, 200, response.json)
        self.assertNotIn('Content-Disposition', response.headers)
        preview = response.json
        self.assertTrue(preview['read_only'])
        book = self.export(root + '/export' + query)
        try:
            sheets = preview['workbook']['sheets']
            self.assertEqual([sheets[key]['name'] for key in preview['workbook']['sheetOrder']], book.sheetnames)
            for key, expected in zip(preview['workbook']['sheetOrder'], book):
                sheet = sheets[key]
                self.assertEqual(sheet['rowCount'], expected.max_row)
                self.assertEqual(sheet['columnCount'], expected.max_column)
                for row in expected:
                    for cell in row:
                        value = cell.value
                        if isinstance(value, (date, datetime)):
                            value = value.strftime('%d/%m/%Y')
                        actual = sheet['cellData'].get(str(cell.row-1), {}).get(str(cell.column-1), {}).get('v', '')
                        self.assertEqual(actual, value if value is not None else '', (expected.title, cell.coordinate))
            self.assertEqual(preview['line_count'], book.active.max_row - 4)
            self.assertEqual(sheets[preview['workbook']['sheetOrder'][0]]['freeze']['ySplit'], 3)
        finally:
            book.close()
        return preview

    def test_debt_previews_match_excel_filters_empty_periods_and_are_read_only(self):
        with server.db() as conn:
            before = '\n'.join(conn.iterdump())
        period = '?from=2026-09-01&to=2026-09-30'
        for query in ('', '&contractor=C1&kitchen=K1', '&contractor=C2',
                      '&contractor=C1&kitchen=K2', '&status=reversed', '&status=all',
                      '&limit=1&offset=1'):
            with self.subTest(receivable=query):
                self.assert_preview_matches_export('/api/debts/receivables/lines', period + query)
        for query in ('', '&supplier=S1', '&supplier=S2', '&status=paid', '&status=all',
                      '&limit=1&offset=1'):
            with self.subTest(payable=query):
                self.assert_preview_matches_export('/api/debts/payables', period + query)
        for root in ('/api/debts/receivables/lines', '/api/debts/payables'):
            empty = self.assert_preview_matches_export(root, '?from=2026-09-01&to=2026-09-04')
            self.assertEqual(empty['line_count'], 0)
        with server.db() as conn:
            self.assertEqual('\n'.join(conn.iterdump()), before)

    def test_debt_preview_errors_are_visible_and_do_not_return_a_partial_sheet(self):
        from unittest.mock import patch
        for root, module in (('/api/debts/receivables/lines', 'tdp_system.round3_documents'),
                             ('/api/debts/payables', 'tdp_system.payable_export')):
            for query in ('?from=bad&to=2026-09-30', '?from=2026-09-30&to=2026-09-01',
                          '?from=2026-09-01&to=2026-09-30&status=invalid'):
                response = self.client.get(root + '/preview' + query)
                self.assertEqual(response.status_code, 400, response.json)
                self.assertFalse(response.json['ok'])
                self.assertNotIn('workbook', response.json)
            with patch(module + '.workbook_preview', side_effect=ValueError('Thu hẹp khoảng ngày')):
                response = self.client.get(root + '/preview?from=2026-09-01&to=2026-09-30')
                self.assertEqual(response.status_code, 422)
                self.assertEqual(response.json['error'], 'Thu hẹp khoảng ngày')
                self.assertNotIn('workbook', response.json)

    def test_payable_paid_filter_and_all_supplier_sheets(self):
        url = '/api/debts/payables/ledger?from=2026-09-01&to=2026-09-30'
        row = self.client.get(url).json['rows'][0]
        payload = dict(request_id='ROUND3-PAY-00001', party_code=row['supplier']['code'],
                       payment_date='2026-09-05',amount=row['amount'],actor='Người thử',
                       allocations=[dict(ledger_line_id=row['id'],amount=row['amount'],expected_revision=row['ledger_revision'])])
        paid = self.client.post('/api/debts/payables/payments',json=payload)
        self.assertEqual(paid.status_code,201,paid.json)
        self.assertEqual(self.client.get(url).json['pagination']['total'],1)
        for status in ('open,partially_paid', 'paid', 'all'):
            self.assert_preview_matches_export('/api/debts/payables', '?from=2026-09-01&to=2026-09-30&status='+status)
            view = self.client.get(url+'&status='+status).json
            book = self.export('/api/debts/payables/export?from=2026-09-01&to=2026-09-30&status='+status)
            ws = book.active
            self.assertEqual(ws.max_column,14)
            self.assertEqual(ws.max_row-4,view['pagination']['total'])
            self.assertEqual(ws.cell(ws.max_row,14).value,view['summary']['filtered_amount'])
            self.assertIn('Tổng NCC',book.sheetnames)
            for detail in book.worksheets[2:]:
                self.assertEqual(detail.max_column,14)
            book.close()
        undo_url = f"/api/debts/payables/payments/{paid.json['id']}/reverse"
        self.assertEqual(self.client.post(undo_url,json={'expected_revision':1,'reason':'Sai'}).status_code,400)
        undo = self.client.post(undo_url,json={'expected_revision':1,'reason':'Sai','actor':'Chị thử'})
        self.assertEqual(undo.status_code,200,undo.json)
        self.assertEqual(self.client.get(url).json['pagination']['total'],2)

    def test_receipt_retry_conflict_reverse_audit_and_balances(self):
        old = self.client.get('/api/debts?from=2026-09-01&to=2026-09-30').json['contractors']['C1']['closing']
        created = self.client.post('/api/payments',json=self.receipt())
        self.assertEqual(created.status_code,201,created.json)
        pid = created.json['id']
        self.assertEqual(self.client.post('/api/payments',json=self.receipt()).json['id'],pid)
        self.assertEqual(self.client.post('/api/payments',json=self.receipt(amount=200)).status_code,400)
        self.assertEqual(self.client.delete('/api/payments/'+str(pid)).status_code,409)
        self.assertEqual(self.client.get('/api/debts?from=2026-09-01&to=2026-09-30').json['contractors']['C1']['closing'],old-100)
        url = f'/api/debts/receipts/{pid}/reverse'
        undo = dict(reason='Nhập nhầm',actor='Kế toán thử',expected_revision=1)
        self.assertEqual(self.client.post(url,json={**undo,'actor':''}).status_code,400)
        self.assertEqual(self.client.post(url,json={**undo,'expected_revision':8}).status_code,409)
        self.assertEqual(self.client.post(url,json=undo).status_code,200)
        self.assertTrue(self.client.post(url,json=undo).json['idempotent'])
        self.assertEqual(self.client.get('/api/debts?from=2026-09-01&to=2026-09-30').json['contractors']['C1']['closing'],old)
        history = self.client.get('/api/debts/receipts?from=2026-09-01&to=2026-09-30').json['rows']
        self.assertEqual(len(history),1); self.assertEqual(history[0]['status'],'reversed')
        self.assertEqual(history[0]['reversed_by'],'Kế toán thử')
        with server.db() as conn:
            self.assertEqual(conn.execute("SELECT COUNT(*) FROM audit_log WHERE event_type='payment.reverse'").fetchone()[0],1)
            self.assertEqual(conn.execute("SELECT COUNT(*) FROM payments").fetchone()[0],1)
            self.assertEqual(conn.execute("SELECT COUNT(*) FROM invoice_inventory_ledger").fetchone()[0],0)

    def test_receipt_concurrent_retry_and_validation(self):
        def submit(_):
            with server.app.test_client() as client:
                return client.post('/api/payments',json=self.receipt()).status_code
        with ThreadPoolExecutor(max_workers=2) as pool:
            self.assertEqual(sorted(pool.map(submit,range(2))),[200,201])
        for changes in ({'actor':''},{'amount':1.5},{'amount':'NaN'},{'amount':True}, {'request_id':''}, {'amount':-1}):
            self.assertEqual(self.client.post('/api/payments',json=self.receipt(**changes)).status_code,400)

    def test_receipt_audit_failure_rolls_back(self):
        with server.db() as conn:
            conn.execute("CREATE TRIGGER fail_round3_audit BEFORE INSERT ON audit_log BEGIN SELECT RAISE(ABORT,'forced audit failure'); END")
        self.assertEqual(self.client.post('/api/payments',json=self.receipt()).status_code,500)
        with server.db() as conn:
            self.assertEqual(conn.execute('SELECT COUNT(*) FROM payments').fetchone()[0],0)

    def test_read_views_never_write_and_invalid_filters(self):
        with server.db() as conn:
            before = conn.total_changes
            from .receivable_ledger import receivable_ledger_payload
            from .payable_ledger import payable_ledger_payload
            for fn in (receivable_ledger_payload,payable_ledger_payload):
                fn(conn,date_from='2026-09-01',date_to='2026-09-30',limit=1,offset=1)
            self.assertEqual(before,conn.total_changes)
        for path in ('/api/debts/receivables/lines/export','/api/debts/payables/export'):
            self.assertEqual(self.client.get(path+'?from=2026-09-30&to=2026-09-01').status_code,400)
            self.assertEqual(self.client.get(path+'?from=2026-09-01&to=2026-09-30&status=bogus').status_code,400)
        self.assertEqual(quantity_cell([{'unit':'kg','qty':.1},{'unit':'KG','qty':.2}], 'qty'),.3)

    def test_formula_like_names_stay_text_in_all_debt_exports(self):
        with server.db() as conn:
            conn.execute("UPDATE orders SET product_name='=1+1'")
            conn.execute("UPDATE purchase_workbook_lines SET product_name='=1+1'")
            sync_payable_ledger(conn,timestamp=server.now_iso())
            sync_receivable_ledger(conn,timestamp=server.now_iso())
        for url in (
            '/api/debts/payables/export?from=2026-09-01&to=2026-09-30',
            '/api/debts/receivables/export?from=2026-09-01&to=2026-09-30&contractor=C1',
            '/api/debts/receivables/lines/export?from=2026-09-01&to=2026-09-30',
        ):
            book = self.export(url)
            found = [cell for sheet in book for row in sheet for cell in row if cell.value == '=1+1']
            self.assertTrue(found)
            self.assertTrue(all(cell.data_type == 's' for cell in found))
            book.close()

    def test_payment_schema_upgrade_preserves_legacy_receipt(self):
        import sqlite3
        from .payable_payments import init_payable_payment_schema
        conn = sqlite3.connect(':memory:')
        conn.row_factory = sqlite3.Row
        conn.executescript("""CREATE TABLE payments(id INTEGER PRIMARY KEY,payment_date TEXT,kind TEXT,
                         party_type TEXT,party_code TEXT,amount REAL,note TEXT,created_at TEXT);
                         INSERT INTO payments VALUES(7,'2026-08-01','receipt','contractor','C1',123,'Cũ','2026-08-01');""")
        init_payable_payment_schema(conn)
        init_payable_payment_schema(conn)
        row = dict(conn.execute('SELECT * FROM payments').fetchone())
        self.assertEqual((row['id'],row['amount'],row['note'],row['status'],row['created_by']),(7,123,'Cũ','posted',''))
        conn.close()


if __name__ == '__main__':
    unittest.main()
