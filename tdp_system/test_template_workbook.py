from __future__ import annotations

import hashlib
import io
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

try:
    from .pdf_documents import build_pdf_bundle
    from .print_bundle import workbook_sections
    from .template_workbook import (
        TemplateWorkbookError,
        assert_workbook_safe,
        clone_template_workbook,
        safe_workbook_bytes,
        unsafe_formula_reason,
        workbook_topology_signature,
        write_formula,
        write_literal,
        write_table_rows,
    )
except ImportError:  # pragma: no cover - direct file invocation
    from pdf_documents import build_pdf_bundle
    from print_bundle import workbook_sections
    from template_workbook import (
        TemplateWorkbookError,
        assert_workbook_safe,
        clone_template_workbook,
        safe_workbook_bytes,
        unsafe_formula_reason,
        workbook_topology_signature,
        write_formula,
        write_literal,
        write_table_rows,
    )


ROOT = Path(__file__).resolve().parent.parent
EM_THANH_HASH = "66808CD910F63F4D543A72CC51DDAEBA1C60B147E0457FEAFC5BAD65EAD26AD3"


def create_template(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Template"
    sheet.merge_cells("A1:E1")
    sheet["A1"] = "PHIẾU GIAO"
    sheet["A1"].font = Font(name="Times New Roman", size=16, bold=True)
    sheet["A2"] = "Nguồn golden"
    headings = ["STT", "Tên hàng", "Số lượng", "ĐVT", "Thành tiền"]
    for column, value in enumerate(headings, 1):
        cell = sheet.cell(4, column, value)
        cell.font = Font(name="Times New Roman", bold=True)
        cell.fill = PatternFill("solid", fgColor="FFFF00")
    for row in (5, 6):
        for column in range(1, 6):
            sheet.cell(row, column).font = Font(name="Times New Roman", size=12)
    sheet.row_dimensions[5].height = 21
    sheet.column_dimensions["B"].width = 32
    sheet["E8"] = "=SUM(E5:E6)"
    sheet["C9"] = "='Lookup'!A1"
    sheet["D9"] = "='[1]Outside'!A1"
    sheet["A10"] = "NGƯỜI GIAO"
    sheet["D10"] = "NGƯỜI NHẬN"
    sheet["B2"].hyperlink = "https://example.invalid/template"
    sheet.print_area = "A1:E10"
    sheet.print_title_rows = "4:4"
    sheet.page_setup.orientation = "portrait"
    sheet.page_setup.paperSize = "9"
    sheet.freeze_panes = "A5"

    lookup = workbook.create_sheet("Lookup")
    lookup["A1"] = 123
    workbook.save(path)
    workbook.close()


class TemplateWorkbookTests(unittest.TestCase):
    def test_clone_preserves_topology_and_source_but_removes_outside_dependencies(self):
        with tempfile.TemporaryDirectory(prefix="tdp_template_clone_") as temp_name:
            template = Path(temp_name) / "golden.xlsx"
            create_template(template)
            source_hash = hashlib.sha256(template.read_bytes()).hexdigest().upper()
            source_book = load_workbook(template, data_only=False, keep_links=False)
            try:
                source_sheet_signature = workbook_topology_signature(source_book)[0]
            finally:
                source_book.close()

            clone = clone_template_workbook(
                template, sheet_names=["Template"], expected_sha256=source_hash,
            )
            try:
                sheet = clone.workbook["Template"]
                self.assertEqual(workbook_topology_signature(clone.workbook)[0], source_sheet_signature)
                self.assertEqual(sheet["A1"].value, "PHIẾU GIAO")
                self.assertEqual(sheet["A10"].value, "NGƯỜI GIAO")
                self.assertEqual(sheet["D10"].value, "NGƯỜI NHẬN")
                self.assertEqual(sheet["E8"].value, "=SUM(E5:E6)")
                self.assertIsNone(sheet["C9"].value)
                self.assertIsNone(sheet["D9"].value)
                self.assertIsNone(sheet["B2"].hyperlink)
                self.assertEqual(clone.formula_report["preserved_formulas"], 1)
                self.assertEqual(clone.formula_report["replaced_omitted_sheet_reference"], 1)
                self.assertEqual(clone.formula_report["replaced_external_reference"], 1)
                self.assertEqual(clone.formula_report["external_hyperlinks_removed"], 1)
                assert_workbook_safe(clone.workbook)
                reopened = load_workbook(io.BytesIO(clone.to_bytes()), data_only=False, keep_links=False)
                try:
                    self.assertFalse(reopened._external_links)
                    self.assertEqual(reopened["Template"]["E8"].value, "=SUM(E5:E6)")
                finally:
                    reopened.close()
            finally:
                clone.close()
            self.assertEqual(hashlib.sha256(template.read_bytes()).hexdigest().upper(), source_hash)

    def test_two_generations_with_same_data_have_same_topology_totals_and_pdf_contract(self):
        with tempfile.TemporaryDirectory(prefix="tdp_template_repeat_") as temp_name:
            directory = Path(temp_name)
            template = directory / "golden.xlsx"
            create_template(template)
            records = [
                {"stt": 1, "name": "Gạo", "qty": 2, "unit": "Kg", "amount": 40_000},
                {"stt": 2, "name": "Rau", "qty": 3, "unit": "Kg", "amount": 30_000},
            ]
            clones = []
            try:
                for _ in range(2):
                    clone = clone_template_workbook(template, sheet_names=["Template"])
                    clones.append(clone)
                    sheet = clone.workbook["Template"]
                    write_table_rows(
                        sheet,
                        start_row=5,
                        records=records,
                        columns={"stt": 1, "name": 2, "qty": 3, "unit": 4, "amount": 5},
                        prototype_row=5,
                    )
                    write_formula(sheet, "E8", "=SUM(E5:E6)")
                    assert_workbook_safe(clone.workbook)
                self.assertEqual(
                    workbook_topology_signature(clones[0].workbook),
                    workbook_topology_signature(clones[1].workbook),
                )
                for clone in clones:
                    sheet = clone.workbook["Template"]
                    self.assertEqual(sum(sheet.cell(row, 5).value for row in (5, 6)), 70_000)
                    self.assertEqual(sheet["E8"].value, "=SUM(E5:E6)")
                    self.assertEqual(sheet["A10"].value, "NGƯỜI GIAO")
                    self.assertEqual(sheet["D10"].value, "NGƯỜI NHẬN")

                # The native Excel output may retain a safe local formula.  The
                # print/PDF projection receives the already reconciled literal
                # total so it never attempts to execute spreadsheet formulas;
                # signatures are supplied separately by print_bundle.
                write_literal(clones[0].workbook["Template"], "E8", 70_000)
                write_literal(clones[0].workbook["Template"], "A10", None)
                write_literal(clones[0].workbook["Template"], "D10", None)
                sections = workbook_sections("deliveries", clones[0].workbook)
                self.assertEqual((len(sections), sections[0]["document_type"]), (1, "deliveries"))
                for paper in ("A4", "A5"):
                    target = directory / f"template-{paper}.pdf"
                    manifest = build_pdf_bundle(
                        sections, target, paper=paper,
                        generated_at="2026-09-02 23:00:00",
                    )
                    self.assertEqual(manifest["paper"], paper)
                    self.assertTrue(manifest["verification"]["ok"])
                    self.assertTrue(target.is_file())
            finally:
                for clone in clones:
                    clone.close()

    def test_dynamic_formula_injection_nonfinite_and_external_formula_fail_closed(self):
        workbook = Workbook()
        sheet = workbook.active
        try:
            with self.assertRaisesRegex(TemplateWorkbookError, "không được là công thức"):
                write_literal(sheet, "A1", "=HYPERLINK(\"bad\")")
            with self.assertRaisesRegex(TemplateWorkbookError, "số hữu hạn"):
                write_literal(sheet, "A1", float("nan"))
            with self.assertRaisesRegex(TemplateWorkbookError, "không an toàn"):
                write_formula(sheet, "A1", "='[1]Outside'!A1")
            self.assertEqual(
                unsafe_formula_reason("=SUM(A1:A2)", sheet.parent.sheetnames), "",
            )
            self.assertEqual(
                unsafe_formula_reason("=WEBSERVICE(\"https://bad\")", sheet.parent.sheetnames),
                "external_reference",
            )
        finally:
            workbook.close()

    def test_real_em_thanh_receipt_clone_keeps_golden_layout_and_removes_external_formulas(self):
        source = next(path for path in ROOT.glob("*.xlsx") if path.name.startswith("Em "))
        before = hashlib.sha256(source.read_bytes()).hexdigest().upper()
        self.assertEqual(before, EM_THANH_HASH)
        source_book = load_workbook(source, data_only=False, keep_links=False)
        try:
            source_signature = next(
                item for item in workbook_topology_signature(source_book)
                if item[0] == "biên nhận"
            )
        finally:
            source_book.close()
        clone = clone_template_workbook(
            source, sheet_names=["biên nhận"], expected_sha256=EM_THANH_HASH,
        )
        try:
            self.assertEqual(workbook_topology_signature(clone.workbook)[0], source_signature)
            self.assertEqual(len(clone.workbook["biên nhận"].merged_cells.ranges), 11)
            self.assertGreaterEqual(clone.formula_report["replaced_external_reference"], 1)
            self.assertEqual(clone.workbook["biên nhận"]["F22"].value, "=SUM(F15:F21)")
            self.assertEqual(clone.workbook["biên nhận"]["G22"].value, "=SUM(G15:G21)")
            payload = safe_workbook_bytes(clone.workbook)
            reopened = load_workbook(io.BytesIO(payload), data_only=False, keep_links=False)
            try:
                self.assertFalse(reopened._external_links)
                assert_workbook_safe(reopened)
            finally:
                reopened.close()
        finally:
            clone.close()
        self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest().upper(), before)


if __name__ == "__main__":
    unittest.main()
