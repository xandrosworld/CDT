"""Golden-template export for the internal purchase summary.

``Em Thành.xlsx`` is authoritative for the document layout.  Operational
values come only from confirmed purchase/BK rows; receivables and sales
amounts are deliberately outside this module.  Citizen identity values are
written only into the official workbook and are never included in errors or
diagnostic metadata.
"""

from __future__ import annotations

import math
try:
    from seller_identity_catalog import is_excluded_seller
except ImportError:
    from .seller_identity_catalog import is_excluded_seller
try:
    from document_totals import quantity_cell
    from document_preview import white_print_style
except ImportError:
    from .document_totals import quantity_cell
    from .document_preview import white_print_style

import re
import unicodedata
from copy import copy
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any, Iterable, Mapping

from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment

try:
    from template_workbook import (
        TemplateWorkbookError,
        assert_workbook_safe,
        clone_template_workbook,
        write_literal,
    )
except ImportError:  # pragma: no cover - package import path
    from .template_workbook import (
        TemplateWorkbookError,
        assert_workbook_safe,
        clone_template_workbook,
        write_literal,
    )


PURCHASE_SUMMARY_TEMPLATE_SHEET = "bảng kê tổng"
EM_THANH_SHA256 = "66808CD910F63F4D543A72CC51DDAEBA1C60B147E0457FEAFC5BAD65EAD26AD3"

TABLE_NUMBER_ROW = 10
TABLE_FIRST_ROW = 11
GOLDEN_LAST_ITEM_ROW = 31
GOLDEN_TOTAL_ROW = 32
GOLDEN_SIGNATURE_FIRST_ROW = 35
GOLDEN_SIGNATURE_LAST_ROW = 42
GOLDEN_CLEAR_THROUGH_ROW = 50


class PurchaseSummaryError(ValueError):
    def __init__(
        self, message: str, *, code: str = "invalid_purchase_summary", status: int = 422,
    ):
        super().__init__(message)
        self.code = code
        self.status = status


def _plain(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip())


def _key(value: Any) -> str:
    text = _plain(value).replace("đ", "d").replace("Đ", "D")
    text = unicodedata.normalize("NFKD", text).casefold()
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"[^a-z0-9]+", "", text)


def _row_dict(row: Any) -> dict[str, Any]:
    return dict(row) if not isinstance(row, dict) else row.copy()


def _decimal(value: Any, label: str) -> Decimal:
    if isinstance(value, bool):
        raise PurchaseSummaryError(f"{label} phải là số hữu hạn")
    try:
        number = Decimal(str(value))
    except (InvalidOperation, TypeError, ValueError) as error:
        raise PurchaseSummaryError(f"{label} phải là số hữu hạn") from error
    if not number.is_finite():
        raise PurchaseSummaryError(f"{label} phải là số hữu hạn")
    return number


def _vnd_product(quantity: Any, price: Any) -> int:
    return int((_decimal(quantity, "Số lượng") * _decimal(price, "Giá mua")).quantize(
        Decimal("1"), rounding=ROUND_HALF_UP,
    ))


def _iso_date(value: Any, label: str = "Ngày mua") -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = _plain(value)
    for pattern in ("%Y-%m-%d", "%d/%m/%Y", "%d.%m.%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, pattern).date().isoformat()
        except ValueError:
            continue
    raise PurchaseSummaryError(f"{label} không hợp lệ", code="invalid_purchase_date")


def _display_date(value: Any) -> str:
    return datetime.strptime(_iso_date(value), "%Y-%m-%d").strftime("%d/%m/%Y")


def _identity_number(value: Any) -> str:
    text = _plain(value)
    return text if re.fullmatch(r"(?:\d{9}|\d{12})", text) else ""


def _excel_number(value: Decimal) -> int | float:
    if value == value.to_integral_value():
        return int(value)
    return float(value.quantize(Decimal("0.000001"), rounding=ROUND_HALF_UP))


def _table_exists(conn: Any, name: str) -> bool:
    return conn.execute(
        "SELECT 1 FROM sqlite_master WHERE type='table' AND name=?", (name,),
    ).fetchone() is not None


def _people_by_name(conn: Any) -> dict[str, list[dict[str, Any]]]:
    people: dict[str, list[dict[str, Any]]] = {}
    if not _table_exists(conn, "people"):
        return people
    for source in conn.execute("SELECT name,cccd,address FROM people ORDER BY name"):
        item = _row_dict(source)
        people.setdefault(_key(item.get("name")), []).append(item)
    return people


def _products_by_code(conn: Any) -> dict[str, dict[str, Any]]:
    if not _table_exists(conn, "products"):
        return {}
    return {
        _plain(row["code"]).upper(): _row_dict(row)
        for row in conn.execute(
            "SELECT code,purchase_list,seller,cccd,supplier FROM products ORDER BY code"
        )
    }


def _resolved_identity(
    order: Mapping[str, Any] | None,
    product: Mapping[str, Any] | None,
    people: Mapping[str, list[dict[str, Any]]],
) -> tuple[str, str, str, list[str]]:
    order = order or {}
    product = product or {}
    seller = _plain(order.get("seller")) or _plain(product.get("seller"))
    issues: list[str] = []
    matches = people.get(_key(seller), []) if seller else []
    if len(matches) > 1:
        issues.append("người bán không duy nhất trong danh mục định danh")
    person = matches[0] if len(matches) == 1 else {}

    # The internal identity catalogue is the current authoritative value.  A
    # daily order may legitimately retain an older snapshot after a corrected
    # CCCD/CMND is confirmed in that catalogue.
    identity = (
        _identity_number(person.get("cccd"))
        or _identity_number(order.get("cccd"))
        or _identity_number(product.get("cccd"))
    )
    address = _plain(person.get("address"))
    if not seller:
        issues.append("thiếu người bán")
    if not identity:
        issues.append("thiếu hoặc sai định dạng CCCD/CMND")
    if not address:
        issues.append("thiếu địa chỉ người bán trong danh mục định danh")
    return seller, identity, address, issues


def collect_purchase_summary_rows(
    conn: Any, batch: Mapping[str, Any], orders: Iterable[Mapping[str, Any]],
    *, excluded_rows: list | None = None, apply_seller_updates: bool = True,
) -> list[dict[str, Any]]:
    """Project confirmed BK purchase rows without reading a receivable table.

    Canonical ``purchase_workbook_lines`` win whenever that scope exists for
    the batch.  Legacy order rows are used only for old batches that have no
    canonical purchase scope, matching the payable-ledger source boundary.
    """

    batch_id = int(batch.get("id") or 0)
    order_map = {
        int(item.get("id") or 0): _row_dict(item)
        for item in (_row_dict(row) for row in orders)
        if int(item.get("id") or 0) > 0
    }
    products = _products_by_code(conn)
    people = _people_by_name(conn)
    purchase_rate = Decimal("0.95")
    if _table_exists(conn, "settings"):
        rate_row = conn.execute(
            "SELECT value FROM settings WHERE key='purchase_rate'",
        ).fetchone()
        if rate_row is not None:
            purchase_rate = _decimal(rate_row[0], "Tỷ lệ giá bảng kê")
            if purchase_rate <= 0:
                raise PurchaseSummaryError("Tỷ lệ giá bảng kê phải lớn hơn 0")
    candidates: list[dict[str, Any]] = []
    invalid: list[str] = []

    canonical: list[dict[str, Any]] = []
    if batch_id and _table_exists(conn, "purchase_workbook_lines"):
        canonical = [
            _row_dict(row) for row in conn.execute(
                "SELECT * FROM purchase_workbook_lines WHERE batch_id=? ORDER BY source_row,id",
                (batch_id,),
            )
        ]

    if canonical:
        sources = []
        for line in canonical:
            order = order_map.get(int(line.get("order_id") or 0))
            product = products.get(_plain(line.get("product_code")).upper(), {})
            purchase_flag = int((
                (order or {}).get("purchase_list")
                if order is not None
                else product.get("purchase_list")
            ) or 0)
            quantity = _decimal(line.get("actual_qty"), "Số lượng thực tế")
            supplier = _plain(line.get("supplier"))
            if not purchase_flag or quantity <= 0 or _key(supplier) == "kho":
                continue
            source_ref = int(line.get("source_row") or line.get("id") or 0)
            if _plain(line.get("status")).casefold() != "confirmed":
                invalid.append(f"dòng nguồn {source_ref}: phần mua chưa được chốt")
                continue
            sources.append((line, order, product, quantity, supplier, source_ref))
    else:
        sources = []
        for order in order_map.values():
            if not int(order.get("purchase_list") or 0):
                continue
            supplier = _plain(order.get("supplier"))
            if _key(supplier) == "kho":
                continue
            quantity = max(
                _decimal(order.get("actual_received", 0), "Số lượng thực nhận")
                - _decimal(order.get("damaged_qty", 0), "Số lượng hỏng")
                - _decimal(order.get("supplier_return_qty", 0), "Số lượng trả NCC"),
                Decimal("0"),
            )
            if quantity <= 0:
                continue
            product = products.get(_plain(order.get("product_code")).upper(), {})
            source_ref = int(order.get("source_row") or order.get("id") or 0)
            sources.append((order, order, product, quantity, supplier, source_ref))

    excluded_count = 0
    approved_prices = {}
    if _table_exists(conn, 'batch_bk_approvals'):
        approved_prices = {int(r['source_line']): dict(r) for r in conn.execute(
            """SELECT l.source_line,l.qty,l.unit_cost,l.amount FROM batch_bk_approvals a
               JOIN bk_import_documents d ON d.id=a.document_id AND d.status='posted'
               JOIN bk_import_lines l ON l.document_id=d.id WHERE a.batch_id=?""", (batch_id,))}
    for line, order, product, quantity, supplier, source_ref in sources:
        seller, identity, address, identity_issues = _resolved_identity(order, product, people)
        if is_excluded_seller(seller):
            excluded_count += 1
            if excluded_rows is not None:
                excluded_rows.append({"seller": seller, "source_ref": source_ref})
            continue
        buy_price = _decimal(line.get("buy_price"), "Giá mua")
        approved_price = approved_prices.get(int(line.get('id') or 0))
        if approved_price:
            quantity = Decimal(str(approved_price['qty']))
            buy_price = Decimal(str(approved_price['unit_cost']))
        if not canonical and buy_price == 0:
            sell_price = _decimal(line.get("sell_price", 0), "Giá bán")
            if sell_price > 0:
                buy_price = (sell_price * purchase_rate).quantize(
                    Decimal("1"), rounding=ROUND_HALF_UP,
                )
        amount = (
            _decimal(line.get("amount"), "Thành tiền")
            if line.get("amount") not in (None, "")
            else Decimal(_vnd_product(quantity, buy_price))
        )
        expected = Decimal(_vnd_product(quantity, buy_price))
        if approved_price:
            amount = Decimal(str(approved_price['amount']))
        issues = list(identity_issues)
        if buy_price <= 0:
            issues.append("giá mua phải lớn hơn 0")
        if amount < 0 or abs(amount - expected) > Decimal("0.5"):
            issues.append("thành tiền không khớp số lượng thực tế × giá mua")
        work_date = line.get("work_date") or batch.get("work_date")
        try:
            normalized_date = _iso_date(work_date)
        except PurchaseSummaryError:
            normalized_date = ""
            issues.append("ngày mua không hợp lệ")
        if not _plain(line.get("product_name")):
            issues.append("thiếu tên mặt hàng")
        if not _plain(line.get("unit")):
            issues.append("thiếu đơn vị tính")
        if issues:
            invalid.append(f"dòng nguồn {source_ref}: " + ", ".join(issues))
            continue
        candidates.append({
            "work_date": normalized_date,
            "seller": seller,
            "address": address,
            "cccd": identity,
            "product_name": _plain(line.get("product_name")),
            "unit": _plain(line.get("unit")),
            "quantity": _excel_number(quantity),
            "buy_price": _excel_number(buy_price),
            "amount": _excel_number(amount),
            "supplier": supplier,
            "kitchen": _plain(line.get("kitchen") or (order or {}).get("kitchen")),
            "source_ref": source_ref,
            "batch_id": batch_id,
            "selection_key": f"{batch_id}:{'purchase_workbook_lines' if canonical else 'orders'}:{line['id']}",
        })

    if invalid:
        detail = "; ".join(invalid[:12])
        suffix = f"; còn {len(invalid) - 12} dòng" if len(invalid) > 12 else ""
        raise PurchaseSummaryError(
            "Chưa thể lập bảng kê vì dữ liệu BK chưa hợp lệ: " + detail + suffix,
            code="invalid_purchase_identity",
        )
    if not candidates:
        if excluded_count:
            raise PurchaseSummaryError(
                "Không lập chứng từ: các dòng mua thuộc người bán đã bị loại khỏi bảng kê/biên nhận. "
                "Dữ liệu gốc, kho và công nợ vẫn được giữ nguyên.",
                code="excluded_sellers_no_rows",
            )
        raise PurchaseSummaryError(
            "Không có dòng mua/BK đã chốt, có số lượng dương và người bán hợp lệ",
            code="no_purchase_summary_rows",
        )
    if apply_seller_updates:
        try:
            from .purchase_seller_revision import apply_revision
        except ImportError:
            from purchase_seller_revision import apply_revision
        return apply_revision(conn,batch,candidates)
    return candidates


def aggregate_purchase_summary_rows(
    rows: Iterable[Mapping[str, Any]],
) -> list[dict[str, Any]]:
    """Merge the same item for the same legal seller and purchase date.

    Supplier and kitchen dimensions remain in internal source metadata.  They
    do not split a legal seller's line; this is the explicit rule recorded in
    the customer's golden sheet.  A weighted-average unit price preserves the
    exact confirmed purchase total when source prices differ.
    """

    groups: dict[tuple[str, str, str, str], dict[str, Any]] = {}
    identities: dict[str, tuple[str, str]] = {}
    for index, source in enumerate(rows, start=1):
        item = _row_dict(source)
        work_date = _iso_date(item.get("work_date"), f"Ngày mua dòng {index}")
        seller = _plain(item.get("seller"))
        if is_excluded_seller(seller):
            raise PurchaseSummaryError("Người bán đã bị loại khỏi bảng kê/biên nhận", code="excluded_seller")
        address = _plain(item.get("address"))
        identity = _identity_number(item.get("cccd"))
        product_name = _plain(item.get("product_name"))
        unit = _plain(item.get("unit"))
        quantity = _decimal(item.get("quantity"), f"Số lượng dòng {index}")
        amount = _decimal(item.get("amount"), f"Thành tiền dòng {index}")
        if not seller or not address or not identity:
            raise PurchaseSummaryError(
                f"Dòng {index} thiếu người bán, địa chỉ hoặc CCCD/CMND hợp lệ",
                code="invalid_purchase_identity",
            )
        if not product_name or not unit or quantity <= 0 or amount < 0:
            raise PurchaseSummaryError(f"Dòng {index} có dữ liệu hàng mua không hợp lệ")

        seller_key = _key(seller)
        identity_signature = (identity, _key(address))
        if seller_key in identities and identities[seller_key] != identity_signature:
            raise PurchaseSummaryError(
                f"Dòng {index} mâu thuẫn định danh của cùng người bán",
                code="conflicting_purchase_identity",
            )
        identities[seller_key] = identity_signature
        group_key = (work_date, identity, _key(product_name), _key(unit))
        group = groups.setdefault(group_key, {
            "work_date": work_date,
            "seller": seller,
            "address": address,
            "cccd": identity,
            "product_name": product_name,
            "unit": unit,
            "quantity_decimal": Decimal("0"),
            "amount_decimal": Decimal("0"),
            "suppliers": set(),
            "kitchens": set(),
            "source_refs": [],
            "references": set(),
        })
        if _key(group["seller"]) != seller_key or _key(group["address"]) != _key(address):
            raise PurchaseSummaryError(
                f"Dòng {index} mâu thuẫn thông tin của cùng số định danh",
                code="conflicting_purchase_identity",
            )
        group["quantity_decimal"] += quantity
        group["amount_decimal"] += amount
        if _plain(item.get("supplier")):
            group["suppliers"].add(_plain(item.get("supplier")))
        if _plain(item.get("kitchen")):
            group["kitchens"].add(_plain(item.get("kitchen")))
        if item.get("source_ref") not in (None, ""):
            group["source_refs"].append(item.get("source_ref"))
        if _plain(item.get("reference")):
            group["references"].add(_plain(item.get("reference")))

    if not groups:
        raise PurchaseSummaryError(
            "Không có dòng mua/BK để lập bảng kê", code="no_purchase_summary_rows",
        )

    output = []
    for group in groups.values():
        quantity = group.pop("quantity_decimal")
        amount = group.pop("amount_decimal")
        average = amount / quantity
        group["quantity"] = _excel_number(quantity)
        group["unit_price"] = _excel_number(average)
        group["amount"] = _excel_number(amount)
        group["suppliers"] = tuple(sorted(group["suppliers"], key=_key))
        group["kitchens"] = tuple(sorted(group["kitchens"], key=_key))
        group["source_refs"] = tuple(group["source_refs"])
        group["reference"] = ", ".join(sorted(group.pop("references"), key=_key))
        output.append(group)
    output.sort(key=lambda item: (
        item["work_date"], _key(item["seller"]), _key(item["product_name"]), _key(item["unit"]),
    ))
    return output


def _row_snapshot(sheet: Any, row: int, *, include_values: bool) -> dict[str, Any]:
    cells = []
    for column in range(1, 11):
        cell = sheet.cell(row, column)
        cells.append({
            "style": copy(cell._style),
            "alignment": copy(cell.alignment),
            "protection": copy(cell.protection),
            "value": None if isinstance(cell, MergedCell) or not include_values else cell.value,
        })
    dimension = sheet.row_dimensions[row]
    return {
        "cells": cells,
        "height": dimension.height,
        "hidden": dimension.hidden,
        "outlineLevel": dimension.outlineLevel,
        "collapsed": dimension.collapsed,
        "thickTop": dimension.thickTop,
        "thickBot": dimension.thickBot,
    }


def _apply_row_snapshot(sheet: Any, row: int, snapshot: Mapping[str, Any]) -> None:
    dimension = sheet.row_dimensions[row]
    for attribute in ("height", "hidden", "outlineLevel", "collapsed", "thickTop", "thickBot"):
        setattr(dimension, attribute, snapshot[attribute])
    for column, source in enumerate(snapshot["cells"], start=1):
        cell = sheet.cell(row, column)
        cell._style = copy(source["style"])
        cell.alignment = copy(source["alignment"])
        cell.protection = copy(source["protection"])
        cell.comment = None
        cell._hyperlink = None
        write_literal(sheet, cell.coordinate, source["value"])


VIET_DIGITS = ("không", "một", "hai", "ba", "bốn", "năm", "sáu", "bảy", "tám", "chín")


def amount_in_words(value: int) -> str:
    number = int(value)
    if number == 0:
        return "Không đồng"

    def three_digits(group: int, full: bool = False) -> str:
        hundred, remainder = divmod(group, 100)
        ten, unit = divmod(remainder, 10)
        words: list[str] = []
        if hundred or full:
            words.extend((VIET_DIGITS[hundred], "trăm"))
        if ten > 1:
            words.extend((VIET_DIGITS[ten], "mươi"))
            words.append("mốt" if unit == 1 else "lăm" if unit == 5 else VIET_DIGITS[unit] if unit else "")
        elif ten == 1:
            words.append("mười")
            words.append("lăm" if unit == 5 else VIET_DIGITS[unit] if unit else "")
        elif unit:
            if hundred or full:
                words.append("lẻ")
            words.append(VIET_DIGITS[unit])
        return " ".join(word for word in words if word)

    groups = []
    while number:
        groups.append(number % 1000)
        number //= 1000
    units = ("", "nghìn", "triệu", "tỷ", "nghìn tỷ", "triệu tỷ")
    words: list[str] = []
    for index in range(len(groups) - 1, -1, -1):
        group = groups[index]
        if not group:
            continue
        words.append(three_digits(group, full=bool(words) and group < 100))
        if index < len(units) and units[index]:
            words.append(units[index])
    text = " ".join(words)
    return text[:1].upper() + text[1:] + " đồng"


def build_purchase_summary_workbook(
    rows: Iterable[Mapping[str, Any]],
    *,
    template_path: str | Path,
    date_from: Any | None = None,
    date_to: Any | None = None,
    expected_sha256: str = EM_THANH_SHA256,
    additional_template_sheets: Iterable[str] = (),
) -> Any:
    """Build one safe, dynamic ``bảng kê tổng`` sheet from the golden."""

    grouped = aggregate_purchase_summary_rows(rows)
    row_dates = [item["work_date"] for item in grouped]
    first_date = _iso_date(date_from) if date_from is not None else min(row_dates)
    last_date = _iso_date(date_to) if date_to is not None else max(row_dates)
    if first_date > last_date or any(not first_date <= value <= last_date for value in row_dates):
        raise PurchaseSummaryError(
            "Khoảng ngày bảng kê không bao phủ đủ các dòng mua", code="purchase_date_range_mismatch",
        )
    try:
        cloned = clone_template_workbook(
            template_path,
            sheet_names=[PURCHASE_SUMMARY_TEMPLATE_SHEET, *additional_template_sheets],
            expected_sha256=expected_sha256,
        )
    except TemplateWorkbookError as error:
        raise PurchaseSummaryError(
            f"Không thể dùng mẫu bảng kê đã khóa: {error}",
            code="purchase_summary_template_error",
            status=500,
        ) from error

    workbook = cloned.workbook
    sheet = workbook[PURCHASE_SUMMARY_TEMPLATE_SHEET]
    # Keep dates and references readable in both the web preview and A4 export.
    for column, width in {'A':16, 'B':24, 'C':36, 'D':18, 'E':28,
                          'F':7, 'G':12, 'H':16, 'I':19, 'J':20}.items():
        sheet.column_dimensions[column].width=width
    item_style = _row_snapshot(sheet, TABLE_FIRST_ROW, include_values=False)
    total_style = _row_snapshot(sheet, GOLDEN_TOTAL_ROW, include_values=False)
    amount_words_style = _row_snapshot(sheet, GOLDEN_TOTAL_ROW + 1, include_values=False)
    spacer_style = _row_snapshot(sheet, GOLDEN_TOTAL_ROW + 2, include_values=False)
    signature_rows = [
        _row_snapshot(sheet, row, include_values=True)
        for row in range(GOLDEN_SIGNATURE_FIRST_ROW, GOLDEN_SIGNATURE_LAST_ROW + 1)
    ]
    signature_merges = [
        (
            merged.min_row - GOLDEN_SIGNATURE_FIRST_ROW,
            merged.min_col,
            merged.max_row - GOLDEN_SIGNATURE_FIRST_ROW,
            merged.max_col,
        )
        for merged in sheet.merged_cells.ranges
        if merged.min_row >= GOLDEN_SIGNATURE_FIRST_ROW
        and merged.max_row <= GOLDEN_SIGNATURE_LAST_ROW
    ]

    total_row = TABLE_FIRST_ROW + len(grouped)
    amount_words_row = total_row + 1
    signature_first = total_row + 3
    signature_last = signature_first + len(signature_rows) - 1
    try:
        for merged in list(sheet.merged_cells.ranges):
            if merged.min_row >= TABLE_FIRST_ROW:
                sheet.unmerge_cells(str(merged))
        clear_through = max(GOLDEN_CLEAR_THROUGH_ROW, signature_last)
        for row in range(TABLE_FIRST_ROW, clear_through + 1):
            for column in range(1, 11):
                cell = sheet.cell(row, column)
                cell.value = None
                cell.comment = None
                cell._hyperlink = None

        write_literal(
            sheet,
            "A2",
            f"Từ ngày {_display_date(first_date)} đến ngày {_display_date(last_date)}",
        )
        # The golden carries two stray calculation remnants immediately above
        # the table.  They are sample data, not part of the approved wording.
        for column in range(1, 11):
            write_literal(sheet, sheet.cell(7, column).coordinate, None)
        total_quantity = Decimal("0")
        total_amount = Decimal("0")
        for offset, item in enumerate(grouped):
            row = TABLE_FIRST_ROW + offset
            _apply_row_snapshot(sheet, row, item_style)
            values = (
                _display_date(item["work_date"]),
                item["seller"],
                item["address"],
                item["cccd"],
                item["product_name"],
                item["unit"],
                item["quantity"],
                item["unit_price"],
                item["amount"],
                item["reference"],
            )
            for column, value in enumerate(values, start=1):
                write_literal(sheet, sheet.cell(row, column).coordinate, value)
            sheet.cell(row, 1).number_format = "@"
            sheet.cell(row, 4).number_format = "@"
            sheet.cell(row, 7).number_format = "#,##0.######"
            sheet.cell(row, 8).number_format = "#,##0"
            sheet.cell(row, 9).number_format = "#,##0"
            sheet.row_dimensions[row].height = max(sheet.row_dimensions[row].height or 22,
                17 * max(math.ceil(len(str(item['product_name'])) / 26), math.ceil(len(str(item['address'])) / 30),
                         math.ceil(len(str(item['seller'])) / 22), math.ceil(len(str(item['reference'])) / 18)))
            for column in (2, 3, 5, 10):
                alignment = copy(sheet.cell(row, column).alignment)
                alignment.wrap_text = True
                alignment.vertical = 'center'
                sheet.cell(row, column).alignment = alignment
            for column in (1,4,6):
                sheet.cell(row,column).alignment=Alignment(horizontal='center',vertical='center',wrap_text=False)
            for column in (7,8,9):
                sheet.cell(row,column).alignment=Alignment(horizontal='right',vertical='center')
            total_quantity += _decimal(item["quantity"], "Tổng số lượng")
            total_amount += _decimal(item["amount"], "Tổng tiền")

        _apply_row_snapshot(sheet, total_row, total_style)
        sheet.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=5)
        write_literal(sheet, f"A{total_row}", "TỔNG CỘNG")
        write_literal(sheet, f"G{total_row}", quantity_cell(grouped, "quantity"))
        if isinstance(sheet[f"G{total_row}"].value, str):
            sheet.merge_cells(start_row=total_row, start_column=7, end_row=total_row, end_column=8)
        write_literal(sheet, f"I{total_row}", _excel_number(total_amount))
        sheet[f"G{total_row}"].number_format = "#,##0.######"
        alignment = copy(sheet[f"G{total_row}"].alignment)
        alignment.wrap_text = True
        sheet[f"G{total_row}"].alignment = alignment
        sheet.row_dimensions[total_row].height = max(28, 21 * math.ceil(len(str(sheet[f"G{total_row}"].value)) / 10))
        sheet[f"I{total_row}"].number_format = "#,##0"

        _apply_row_snapshot(sheet, amount_words_row, amount_words_style)
        sheet.merge_cells(start_row=amount_words_row,start_column=1,end_row=amount_words_row,end_column=10)
        write_literal(
            sheet,
            f"A{amount_words_row}",
            "Số tiền bằng chữ: " + amount_in_words(int(total_amount.quantize(Decimal("1"), rounding=ROUND_HALF_UP))) + "./.",
        )
        _apply_row_snapshot(sheet, total_row + 2, spacer_style)
        sheet.cell(amount_words_row,1).alignment=Alignment(horizontal='left',vertical='center',wrap_text=True)
        sheet.row_dimensions[amount_words_row].height=34
        sheet.row_dimensions[total_row+2].height=14

        for offset, snapshot in enumerate(signature_rows):
            _apply_row_snapshot(sheet, signature_first + offset, snapshot)
        for start_row, start_col, end_row, end_col in signature_merges:
            sheet.merge_cells(
                start_row=signature_first + start_row,
                start_column=start_col,
                end_row=signature_first + end_row,
                end_column=end_col,
            )
        write_literal(
            sheet,
            f"H{signature_first}",
            f"Hải Phòng, ngày {_display_date(last_date)}",
        )

        # Leave twice the template's signing space for the company seal.
        for row in range(signature_first + 3, signature_last):
            height = sheet.row_dimensions[row].height or sheet.sheet_format.defaultRowHeight
            sheet.row_dimensions[row].height = height * 2
        signer = sheet.cell(signature_last, 9)
        signer_name = signer.value
        signer_font = copy(signer.font)
        signer_font.sz = 16
        signer_font.bold = True
        signer_font.color = '000000'
        write_literal(sheet, signer.coordinate, None)
        sheet.merge_cells(start_row=signature_last, start_column=8,
                          end_row=signature_last, end_column=10)
        write_literal(sheet, f"H{signature_last}", signer_name)
        sheet[f"H{signature_last}"].font = signer_font
        sheet[f"H{signature_last}"].alignment = Alignment(horizontal='center', vertical='center')
        sheet.row_dimensions[signature_last].height = max(24, sheet.row_dimensions[signature_last].height or 0)

        for row in range(8, total_row + 1):
            ratio = 1.0
            for cell in sheet[row]:
                if isinstance(cell, MergedCell):continue
                font = copy(cell.font)
                old_size = font.sz or 13
                font.sz = old_size + 1
                cell.font = font
                ratio = max(ratio, font.sz / old_size)
            height = sheet.row_dimensions[row].height or sheet.sheet_format.defaultRowHeight
            sheet.row_dimensions[row].height = height * ratio

        sheet.print_area = f"A1:J{signature_last}"
        sheet.print_title_rows = "$8:$10"
        sheet.page_setup.paperSize = "9"
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.scale = None
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.page_margins.top = 0.2
        sheet.page_margins.header = 0.1
        sheet.page_margins.bottom = 0.5
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.freeze_panes = None
        workbook.active = 0
        sheet.sheet_view.tabSelected = True
        assert_workbook_safe(workbook)
        white_print_style(workbook)
        return workbook
    except Exception:
        workbook.close()
        raise


__all__ = [
    "EM_THANH_SHA256",
    "PURCHASE_SUMMARY_TEMPLATE_SHEET",
    "PurchaseSummaryError",
    "aggregate_purchase_summary_rows",
    "amount_in_words",
    "build_purchase_summary_workbook",
    "collect_purchase_summary_rows",
]
