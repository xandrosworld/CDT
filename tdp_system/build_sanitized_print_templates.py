"""Build data-free runtime templates from the customer's approved documents.

This utility is intentionally separate from normal application startup.  It is
run when the approved customer artwork changes; the resulting small templates
are safe to bundle because all transaction/customer example values are removed.
"""

from __future__ import annotations

import copy
import re
from pathlib import Path

from docx import Document
from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import MergedCell


ROOT = Path(__file__).resolve().parent.parent
OUTPUT_DIR = Path(__file__).resolve().parent / "templates"


def _replace_paragraph(paragraph, text: str) -> None:
    source = paragraph.runs[0] if paragraph.runs else None
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)
    target = paragraph.add_run(text)
    if source is not None:
        target._r.get_or_add_rPr()
        if source._r.rPr is not None:
            target._r.remove(target._r.rPr)
            target._r.insert(0, copy.deepcopy(source._r.rPr))


def _replace_cell(cell, text: str) -> None:
    first = cell.paragraphs[0]
    _replace_paragraph(first, text)
    for paragraph in list(cell.paragraphs[1:]):
        cell._tc.remove(paragraph._p)


def build_daily_order_template() -> Path:
    candidates = (
        ROOT / "Đơn hàng 01.09.2026.xlsx",
        ROOT.parent / "Đơn hàng 01.09.2026.xlsx",
        Path.home() / "Downloads" / "Đơn hàng 01.09.2026.xlsx",
    )
    source = next((item for item in candidates if item.is_file()), None)
    if source is None:
        raise FileNotFoundError("Thiếu Đơn hàng 01.09.2026.xlsx")
    original = load_workbook(source, data_only=False, read_only=False)
    template_sheet = next(
        sheet for sheet in original.worksheets
        if str(sheet.title).strip().casefold() == "đặt hàng".casefold()
    )

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "đặt hàng"
    for row in range(1, 4):
        worksheet.row_dimensions[row].height = template_sheet.row_dimensions[row].height
        for column in range(1, 18):
            source_cell = template_sheet.cell(row, column)
            target_cell = worksheet.cell(row, column)
            target_cell.font = copy.copy(source_cell.font)
            target_cell.fill = copy.copy(source_cell.fill)
            target_cell.border = copy.copy(source_cell.border)
            target_cell.alignment = copy.copy(source_cell.alignment)
            target_cell.protection = copy.copy(source_cell.protection)
            target_cell.number_format = source_cell.number_format
    for column in range(1, 18):
        letter = worksheet.cell(1, column).column_letter
        source_dimension = template_sheet.column_dimensions[letter]
        target_dimension = worksheet.column_dimensions[letter]
        target_dimension.width = source_dimension.width
        target_dimension.hidden = source_dimension.hidden

    headers = (
        "Mã hàngNCC", "Mã hàng", "Mã bếp", "", "Tên hàng ", "Số lượng",
        "ĐVT", "NCC", "ghi chú", "giá mua", "hỏng", "thêm", "Giảm",
        "thiếu", "SL \nthực té", "Thành tiền", "Mã dòng hệ thống",
    )
    for column, label in enumerate(headers, 1):
        worksheet.cell(2, column, label)
        worksheet.cell(1, column, None)
        worksheet.cell(3, column, None)
    worksheet.freeze_panes = "C3"
    worksheet.sheet_view.showGridLines = template_sheet.sheet_view.showGridLines
    worksheet.page_setup = copy.copy(template_sheet.page_setup)
    worksheet.page_margins = copy.copy(template_sheet.page_margins)
    worksheet.sheet_properties.pageSetUpPr = copy.copy(template_sheet.sheet_properties.pageSetUpPr)
    worksheet.column_dimensions["A"].hidden = True
    worksheet.column_dimensions["B"].hidden = True
    worksheet.column_dimensions["Q"].hidden = True
    worksheet.auto_filter.ref = "A2:Q3"
    worksheet.print_area = "A1:Q3"
    worksheet.print_title_rows = "$2:$2"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    output = OUTPUT_DIR / "daily_order_template.xlsx"
    workbook.save(output)
    return output


def build_simple_payment_template() -> Path:
    candidates = sorted((ROOT / "bosung.30.8.26").rglob("*.docx"))
    source = next(
        (
            item for item in candidates
            if "sunby" in item.name.casefold()
        ),
        None,
    )
    if source is None:
        raise FileNotFoundError("Thiếu mẫu đề nghị thanh toán suất ăn SUNBY")
    document = Document(source)
    if len(document.paragraphs) < 11 or len(document.tables) != 2:
        raise ValueError("Mẫu đề nghị thanh toán suất ăn đã đổi cấu trúc")

    replacements = {
        0: "GIẤY ĐỀ NGHỊ THANH TOÁN",
        1: "{{ISSUE_DATE}}",
        2: "Kính gửi: {{RECIPIENT_NAME}}",
        3: "Họ và tên người đề nghị thanh toán: {{REQUESTER}}",
        4: "Bộ phận (hoặc địa chỉ): {{REQUESTER_DEPARTMENT}}",
        5: "Nội dung thanh toán: {{PAYMENT_CONTENT}}",
        6: "Số tiền: {{TOTAL_AMOUNT}} VNĐ",
        7: "(Bằng chữ: {{TOTAL_WORDS}}./.)",
        8: "Thanh toán bằng chuyển khoản",
        9: "Đơn vị thụ hưởng: {{BENEFICIARY_NAME}}",
        10: "Số tài khoản: {{BANK_ACCOUNT}} Tại Ngân hàng {{BANK_NAME}}.",
    }
    for index, paragraph in enumerate(document.paragraphs):
        _replace_paragraph(paragraph, replacements.get(index, ""))

    header = document.tables[0]
    _replace_cell(
        header.cell(0, 0),
        "{{ISSUER_NAME}}\n{{ISSUER_ADDRESS}}",
    )
    _replace_cell(
        header.cell(0, 1),
        "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\nĐộc lập – Tự do – Hạnh phúc",
    )
    signatures = document.tables[1]
    _replace_cell(
        signatures.cell(0, 0),
        "Người đề nghị thanh toán\n(Ký, họ tên)\n\n\n{{REQUESTER}}",
    )
    _replace_cell(signatures.cell(0, 1), "")
    _replace_cell(
        signatures.cell(0, 2),
        "Người duyệt\n(Ký, họ tên)",
    )
    document.core_properties.title = "Mẫu giấy đề nghị thanh toán suất ăn"
    document.core_properties.subject = "Mẫu trắng theo chứng từ khách hàng"
    document.core_properties.author = "Thành Đạt Phát"
    document.core_properties.last_modified_by = "Thành Đạt Phát"
    output = OUTPUT_DIR / "simple_payment_request_template.docx"
    document.save(output)
    return output


def build_bot_payment_template() -> Path:
    source = next(
        (
            item for item in (ROOT / "bosung.30.8.26").rglob("*.xlsx")
            if "BOT" in item.name.upper()
            and set(load_workbook(item, read_only=True).sheetnames) >= {"BBĐC", "ĐNTT", "suất ăn"}
        ),
        None,
    )
    if source is None:
        raise FileNotFoundError("Thiếu mẫu BB đối chiếu + ĐNTT BOT")
    workbook = load_workbook(source, data_only=False, read_only=False)
    for sheet in list(workbook.worksheets):
        if sheet.title not in {"BBĐC", "ĐNTT", "suất ăn"}:
            workbook.remove(sheet)
            continue
        for row in sheet.iter_rows():
            for cell in row:
                if not isinstance(cell, MergedCell):
                    cell.value = None
        sheet.sheet_view.showGridLines = False
    workbook["BBĐC"].print_area = "A1:F78"
    workbook["ĐNTT"].print_area = "A1:H29"
    workbook["suất ăn"].print_area = "A1:F49"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    output = OUTPUT_DIR / "bot_payment_template.xlsx"
    workbook.save(output)
    workbook.close()
    return output


def build_payroll_template() -> Path:
    key = "LƯƠNG XƯỞNG"
    source_book = None
    source_sheet = None
    for candidate in (ROOT / "bosung.30.8.26").rglob("*.xlsx"):
        if "T8.2026" not in candidate.name:
            continue
        try:
            workbook = load_workbook(candidate, data_only=False, read_only=False)
        except Exception:
            continue
        matching = next((sheet for sheet in workbook.worksheets if sheet.title == key), None)
        if matching is not None:
            source_book = workbook
            source_sheet = matching
            break
        workbook.close()
    if source_sheet is None or source_book is None:
        raise FileNotFoundError("Thiếu sheet LƯƠNG XƯỞNG trong file chấm công khách hàng")

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = key
    row_sources = {1: 1, 2: 2, 3: 3, 4: 21, 5: 26}
    for target_row, source_row in row_sources.items():
        worksheet.row_dimensions[target_row].height = source_sheet.row_dimensions[source_row].height
        for column in range(1, 17):
            source_cell = source_sheet.cell(source_row, column)
            target_cell = worksheet.cell(target_row, column)
            target_cell.font = copy.copy(source_cell.font)
            target_cell.fill = copy.copy(source_cell.fill)
            target_cell.border = copy.copy(source_cell.border)
            target_cell.alignment = copy.copy(source_cell.alignment)
            target_cell.protection = copy.copy(source_cell.protection)
            target_cell.number_format = source_cell.number_format
    for column in range(1, 17):
        letter = worksheet.cell(1, column).column_letter
        source_dimension = source_sheet.column_dimensions[letter]
        target_dimension = worksheet.column_dimensions[letter]
        target_dimension.width = source_dimension.width
        target_dimension.hidden = source_dimension.hidden
    worksheet.merge_cells("A1:H1")
    worksheet.merge_cells("K1:P1")
    worksheet.merge_cells("A2:B2")
    worksheet["A1"] = "LƯƠNG THÁNG"
    worksheet["I1"] = "{{MONTH}}"
    worksheet["J1"] = "{{YEAR}}"
    worksheet["A2"] = "Lương cơ bản:"
    worksheet["C2"] = None
    headers = (
        "STT", "HỌ VÀ TÊN", "CHỨC VỤ", "Công HC", "Công TC", "Công CN",
        "Công Đêm", "Công NL", "Phụ cấp", "Trách nhiệm", "Tổng Lương",
        "Tạm ứng", "Thử việc", "BHXH người lao động đóng", "TỔNG",
        "BHXH Công ty đóng",
    )
    for column, label in enumerate(headers, 1):
        worksheet.cell(3, column, label)
        worksheet.cell(4, column, None)
        worksheet.cell(5, column, None)
    worksheet["A5"] = "TỔNG"
    worksheet.page_setup = copy.copy(source_sheet.page_setup)
    worksheet.page_margins = copy.copy(source_sheet.page_margins)
    worksheet.sheet_properties.pageSetUpPr = copy.copy(source_sheet.sheet_properties.pageSetUpPr)
    worksheet.sheet_view.showGridLines = source_sheet.sheet_view.showGridLines
    worksheet.freeze_panes = "A4"
    worksheet.print_title_rows = "$1:$3"
    worksheet.print_area = "A1:P5"
    source_book.close()
    output = OUTPUT_DIR / "payroll_template.xlsx"
    workbook.save(output)
    workbook.close()
    return output


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    for output in (
        build_daily_order_template(),
        build_simple_payment_template(),
        build_bot_payment_template(),
        build_payroll_template(),
    ):
        print(output)


if __name__ == "__main__":
    main()
