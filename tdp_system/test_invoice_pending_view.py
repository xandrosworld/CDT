import io
import unittest
import zipfile

from openpyxl import load_workbook

from . import server
from .test_outgoing_readiness import OutgoingReadinessTests as Fixture
from .outgoing_waiting import refresh_waiting


class InvoicePendingViewTests(unittest.TestCase):
    setUpClass = classmethod(Fixture.setUpClass.__func__)
    tearDownClass = classmethod(Fixture.tearDownClass.__func__)
    setUp = Fixture.setUp
    add_batch = staticmethod(Fixture.add_batch)
    add_opening = staticmethod(Fixture.add_opening)

    def get(self, path='', **extra):
        r = self.client.get('/api/outgoing-invoices/unissued'+path, query_string={
            'to':'2026-09-03', 'contractor':'NT-A', **extra})
        self.assertEqual(r.status_code, 200, r.get_json(silent=True))
        return r

    def seed(self, stock=4, qty=10):
        with server.db() as conn:
            self.add_opening(conn, stock)
            batch, ids = self.add_batch(conn, '2026-09-01', [{'qty':qty}])
            refresh_waiting(conn, server.now_iso())
        return batch, ids

    def test_partial_stock_separates_remainder_in_ui_and_download_without_writes(self):
        self.seed()
        with server.db() as conn:
            before = conn.serialize()
        data = self.get().get_json()
        self.assertEqual((data['rows'][0]['unissued_qty'],data['rows'][0]['ready_qty'],data['pending_rows'][0]['waiting_qty']), (10,4,6))
        self.assertIn('Chưa đủ tồn',data['pending_rows'][0]['pending_reason'])
        response = self.get('-template.zip',portion='waiting')
        with zipfile.ZipFile(io.BytesIO(response.data)) as z:
            names = [n for n in z.namelist() if n.startswith('CON_CHO_')]
            self.assertEqual(len(names),1)
            book=load_workbook(io.BytesIO(z.read(names[0])),data_only=True)
            self.assertEqual(book.active.max_column,13)
            self.assertEqual(book.active['D2'].value,6)
            book.close()
            detail=load_workbook(io.BytesIO(z.read('DOI_CHIEU_CHI_TIET.xlsx')),data_only=True)
            self.assertEqual(detail.active.title,'Hang con cho')
            self.assertEqual(detail.active['G2'].value,6)
            self.assertIn('Chưa đủ tồn',detail.active['H2'].value)
            detail.close()
            self.assertIn('Tải bảng kê để up M-Invoice',z.read('HUONG_DAN.txt').decode('utf-8-sig'))
        all_response=self.get('-template.zip')
        with zipfile.ZipFile(io.BytesIO(all_response.data)) as z:
            name=next(n for n in z.namelist() if n.startswith('CHUA_XUAT_'))
            book=load_workbook(io.BytesIO(z.read(name)),data_only=True)
            self.assertEqual(book.active['D2'].value,10);book.close()
        self.get('.xlsx',portion='waiting')
        with server.db() as conn:self.assertEqual(conn.serialize(),before)

    def test_replenishment_clears_waiting_but_unsigned_quantity_remains(self):
        self.seed()
        with server.db() as conn:
            Fixture.add_canonical_event(conn,6,'input','new-input',work_date='2026-09-03')
            refresh_waiting(conn,server.now_iso())
        d=self.get().get_json()
        self.assertEqual(d['pending_rows'],[])
        self.assertEqual((d['rows'][0]['unissued_qty'],d['rows'][0]['ready_qty']),(10,10))
        self.assertEqual(self.client.get('/api/outgoing-invoices/unissued-template.zip?to=2026-09-03&contractor=NT-A&portion=waiting').status_code,400)

    def test_available_stock_without_reservation_is_not_reported_as_missing_input(self):
        with server.db() as conn:
            self.add_opening(conn,20)
            self.add_batch(conn,'2026-09-01',[{'qty':5}])
        d=self.get().get_json()
        reason=d['pending_rows'][0]['pending_reason']
        self.assertIn('Chờ cập nhật phân bổ',reason)
        self.assertNotIn('Chưa đủ tồn',reason)

    def test_rounding_uses_cumulative_same_price_and_retains_fraction_only(self):
        with server.db() as conn:
            self.add_opening(conn,20)
            self.add_batch(conn,'2026-09-01',[{'qty':0.06}])
            self.add_batch(conn,'2026-09-02',[{'qty':0.06}])
            refresh_waiting(conn,server.now_iso())
        d=self.get().get_json()
        self.assertAlmostEqual(sum(r['waiting_qty'] for r in d['pending_rows']),0.02)
        self.assertIn('phần lẻ',d['pending_rows'][0]['pending_reason'])
        self.assertNotIn('Chưa đủ tồn',d['pending_rows'][0]['pending_reason'])

    def test_unit_error_bk_and_kkknt_remain_distinct(self):
        with server.db() as conn:
            self.add_batch(conn,'2026-09-01',[{'qty':2}])
            conn.execute("UPDATE orders SET unit='Túi',tax='KKKNT'")
            refresh_waiting(conn,server.now_iso())
        self.assertIn('chờ xác nhận đơn vị',self.get().get_json()['pending_rows'][0]['pending_reason'])
        with server.db() as conn:
            conn.execute("UPDATE orders SET unit='Kg'")
            refresh_waiting(conn,server.now_iso())
        self.assertIn('Chưa đủ tồn',self.get().get_json()['pending_rows'][0]['pending_reason'])
        with server.db() as conn:
            conn.execute('UPDATE orders SET purchase_list=1')
            refresh_waiting(conn,server.now_iso())
        self.assertEqual(self.get().get_json()['pending_rows'],[])

    def test_source_warning_is_explained_instead_of_claiming_missing_input(self):
        self.seed(stock=20)
        with server.db() as conn:
            Fixture.add_posted_source(conn,source='minvoice',qty=3,invoice_date='2026-09-03',number='unknown')
        d=self.get().get_json()
        self.assertFalse(d['reconciliation_complete'])
        self.assertTrue(d['pending_rows'])
        self.assertIn(d['warnings'][0]['message'],d['pending_rows'][0]['pending_reason'])

    def test_shared_scope_excludes_other_contractors_and_later_order_dates(self):
        with server.db() as conn:
            self.add_batch(conn,'2026-09-01',[{'qty':2}])
            self.add_batch(conn,'2026-09-04',[{'qty':4}])
            self.add_batch(conn,'2026-09-02',[{'qty':7,'contractor':'NT-B'}])
        d=self.get().get_json()
        self.assertEqual(sum(r['waiting_qty'] for r in d['pending_rows']),2)
        self.assertEqual({r['contractor'] for r in d['pending_rows']},{'NT-A'})
        self.assertEqual(self.client.get('/api/outgoing-invoices/unissued?portion=invalid').status_code,400)


if __name__=='__main__':unittest.main()
