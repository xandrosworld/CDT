import hashlib
import io
import sqlite3
import unittest
from pathlib import Path

from openpyxl import load_workbook

try:
    from .print_bundle import workbook_sections
    from .report_export import (
        ReportExportError,
        aggregate_monthly_report_rows,
        build_monthly_report_workbook,
        collect_monthly_report_rows,
    )
    from .template_workbook import safe_workbook_bytes
except ImportError:  # pragma: no cover - direct file invocation
    from print_bundle import workbook_sections
    from report_export import (
        ReportExportError,
        aggregate_monthly_report_rows,
        build_monthly_report_workbook,
        collect_monthly_report_rows,
    )
    from template_workbook import safe_workbook_bytes


ROOT = Path(__file__).resolve().parent.parent
GOLDEN = ROOT / "Em Thành.xlsx"
GOLDEN_HASH = "66808CD910F63F4D543A72CC51DDAEBA1C60B147E0457FEAFC5BAD65EAD26AD3"


def report_row(**overrides):
    row = {
        "contractor": "ATV",
        "kitchen": "LSVINA",
        "revenue": 100,
        "cost": 60,
        "profit": 40,
        "total": 108,
        "source_ref": 1,
    }
    row.update(overrides)
    return row


class MonthlyReportExportTests(unittest.TestCase):
    def test_dynamic_groups_new_kitchens_and_totals_follow_golden(self):
        before = hashlib.sha256(GOLDEN.read_bytes()).hexdigest().upper()
        self.assertEqual(before, GOLDEN_HASH)
        rows = [
            report_row(),
            report_row(revenue=50, cost=20, profit=30, total=54, source_ref=2),
            report_row(kitchen="NEWA", revenue=200, cost=100, profit=100, total=216, source_ref=3),
            report_row(
                contractor="HATRAN", kitchen="POT", revenue=300, cost=200,
                profit=100, total=324, source_ref=4,
            ),
            report_row(
                contractor="NEWC", kitchen="NEWB", revenue=400, cost=250,
                profit=150, total=432, source_ref=5,
            ),
        ]
        workbook = build_monthly_report_workbook(
            rows,
            period="2026-09",
            template_path=GOLDEN,
            configured_groups={"NEWA": "CHỊ TÚ"},
        )
        try:
            self.assertEqual(workbook.sheetnames, ["báo cáo tổng hợp"])
            sheet = workbook.active
            self.assertEqual(
                [sheet.cell(2, column).value for column in range(2, 8)],
                [
                    "Khách hàng", "Mã khách hàng", "Doanh số bán", "Giá vốn",
                    "Lợi nhuận gộp", "TỔNG THANH TOÁN",
                ],
            )
            self.assertEqual(sheet["C3"].value, "LSVINA")
            self.assertEqual((sheet["D3"].value, sheet["E3"].value, sheet["F3"].value, sheet["G3"].value), (150, 80, 70, 162))
            self.assertEqual((sheet["A4"].value, sheet["B4"].value, sheet["C4"].value), ("CHỊ TÚ", "ATV", "NEWA"))
            self.assertEqual((sheet["D5"].value, sheet["E5"].value, sheet["F5"].value, sheet["G5"].value), (350, 180, 170, 378))
            self.assertEqual(sheet["C6"].value, "POT")
            self.assertEqual(sheet["C7"].value, "NEWB")
            self.assertEqual(sheet["A8"].value, "TỔNG THÁNG")
            self.assertEqual((sheet["D8"].value, sheet["E8"].value, sheet["F8"].value, sheet["G8"].value), (1050, 630, 420, 1134))
            self.assertEqual(sheet["A5"].fill.fgColor.rgb, "FF92D050")
            self.assertEqual(sheet["A8"].fill.fgColor.rgb, "FF92D050")
            self.assertEqual(str(sheet.print_area), "'báo cáo tổng hợp'!$A$2:$G$8")
            self.assertEqual(sheet.print_title_rows, "$2:$2")
            self.assertEqual(str(sheet.page_setup.paperSize), "9")
            self.assertEqual(sheet.page_setup.orientation, "landscape")
            self.assertFalse(any(
                isinstance(cell.value, str) and cell.value.startswith("=")
                for row in sheet.iter_rows()
                for cell in row
            ))
            payload = safe_workbook_bytes(workbook)
        finally:
            workbook.close()
        self.assertEqual(hashlib.sha256(GOLDEN.read_bytes()).hexdigest().upper(), before)
        reopened = load_workbook(io.BytesIO(payload), data_only=False, keep_links=False)
        try:
            self.assertEqual(reopened.active["A8"].value, "TỔNG THÁNG")
            self.assertFalse(reopened._external_links)
        finally:
            reopened.close()

    def test_month_source_uses_approved_batches_plus_selected_draft_only(self):
        connection = sqlite3.connect(":memory:")
        connection.row_factory = sqlite3.Row
        connection.executescript(
            """
            CREATE TABLE batches(id INTEGER, work_date TEXT, status TEXT);
            CREATE TABLE orders(
                id INTEGER, batch_id INTEGER, contractor TEXT, kitchen TEXT,
                revenue REAL, cost REAL, profit REAL, total REAL
            );
            CREATE TABLE payments(id INTEGER, amount REAL);
            CREATE TABLE receivable_ledger_lines(id INTEGER, amount REAL);
            INSERT INTO batches VALUES(1,'2026-09-01','approved');
            INSERT INTO batches VALUES(2,'2026-09-15','draft');
            INSERT INTO batches VALUES(3,'2026-08-31','approved');
            INSERT INTO batches VALUES(4,'2026-09-20','approved');
            INSERT INTO orders VALUES(10,1,'ATV','LSVINA',100,60,40,108);
            INSERT INTO orders VALUES(20,2,'HATRAN','POT',200,100,100,216);
            INSERT INTO orders VALUES(30,3,'OLD','OLD',999,1,998,999);
            INSERT INTO orders VALUES(40,4,'SUPPY','SUPPY',300,200,100,324);
            INSERT INTO payments VALUES(1,999999999);
            INSERT INTO receivable_ledger_lines VALUES(1,999999999);
            """
        )
        traced = []
        connection.set_trace_callback(traced.append)

        def totals(item):
            return item["revenue"], item["cost"], item["profit"], item["total"]

        try:
            selected = dict(connection.execute("SELECT * FROM orders WHERE id=20").fetchone())
            period, rows = collect_monthly_report_rows(
                connection,
                {"id": 2, "work_date": "2026-09-15", "status": "draft"},
                [selected],
                totals_fn=totals,
            )
            self.assertEqual(period, "2026-09")
            self.assertEqual({item["kitchen"] for item in rows}, {"LSVINA", "POT", "SUPPY"})
            self.assertFalse(any("payments" in statement.casefold() for statement in traced))
            self.assertFalse(any("receivable" in statement.casefold() for statement in traced))

            _, approved_rows = collect_monthly_report_rows(
                connection,
                {"id": 1, "work_date": "2026-09-01", "status": "approved"},
                [dict(connection.execute("SELECT * FROM orders WHERE id=10").fetchone())],
                totals_fn=totals,
            )
            self.assertEqual(len([item for item in approved_rows if item["kitchen"] == "LSVINA"]), 1)
        finally:
            connection.set_trace_callback(None)
            connection.close()

    def test_conflicting_contractor_and_bad_profit_fail_closed(self):
        with self.assertRaises(ReportExportError) as conflict:
            aggregate_monthly_report_rows([
                report_row(),
                report_row(contractor="OTHER"),
            ])
        self.assertEqual(conflict.exception.code, "conflicting_report_contractor")
        with self.assertRaises(ReportExportError) as mismatch:
            aggregate_monthly_report_rows([report_row(profit=41)])
        self.assertEqual(mismatch.exception.code, "report_profit_mismatch")

    def test_golden_report_projects_to_pdf_section_without_blank_header(self):
        workbook = build_monthly_report_workbook(
            [report_row()], period="2026-09", template_path=GOLDEN,
        )
        try:
            sections = workbook_sections("report", workbook)
            self.assertEqual(len(sections), 1)
            self.assertEqual(sections[0]["title"], "BÁO CÁO TỔNG HỢP")
            self.assertEqual(
                [column["label"] for column in sections[0]["columns"]],
                [
                    "Nhóm khách hàng", "Khách hàng", "Mã khách hàng", "Doanh số bán",
                    "Giá vốn", "Lợi nhuận gộp", "TỔNG THANH TOÁN",
                ],
            )
            self.assertEqual(len(sections[0]["rows"]), 1)
            self.assertEqual(sections[0]["summary"][-1]["value"], 108)
        finally:
            workbook.close()


if __name__ == "__main__":
    unittest.main()
