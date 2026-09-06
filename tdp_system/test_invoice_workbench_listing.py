"""Offline regression for the complete, filtered invoice table and export."""
import json
import sqlite3
import subprocess
import unittest
from contextlib import contextmanager
from copy import deepcopy
from pathlib import Path

from flask import Flask
from openpyxl import load_workbook

from .invoice_workbench import prepare_sync_batch, register_invoice_workbench_routes, list_sync_batches
from .invoice_workbench_listing import invoice_range_payload, range_workbook, STATUS_LABELS, LINE_LABELS
from .invoice_input_sync import sync_input_batch
from .invoice_output_sync import sync_output_batch
from .invoice_mapping import save_mapping, save_conversion
from .invoice_receipt import create_input_receipt
from .invoice_inventory import invoice_inventory_trace, post_output_invoice
from .test_invoice_input_sync import init_test_database, DateBoundedMsmi, now_iso
from .test_invoice_output_sync import OutputFixtureMsmi, output_invoice
from .test_msmi_sync import remote_invoice


def prepare(conn, direction="input", start="2026-08-01", end="2026-08-31", tenant="TDP"):
    return prepare_sync_batch(conn, tenant=tenant, source="msmi" if direction == "input" else "minvoice",
                              invoice_type=direction, date_from=start, date_to=end, now_iso=now_iso)[0]


def seed_round1(conn, extra_lines=0):
    """Small mixed-state fixture shared with the real-browser integration test."""
    conn.executemany("INSERT INTO products(code,name,unit) VALUES(?,?,?)",
                     [("R1-KG", "Hàng kiểm thử kg", "kg"), ("R1-CAI", "Hàng kiểm thử cái", "cái")])
    a, b, c = [remote_invoice(i) for i in (1, 2, 3)]
    a["tdlap"] = "2026-07-31T17:00:00Z"
    c["tdlap"] = "2026-08-30T17:00:00Z"
    for index in range(extra_lines):
        c["hdhhdvu"].append({**deepcopy(c["hdhhdvu"][0]), "ma":f"FILLER-{index}", "ten":f"Dòng kiểm thử cuộn bảng {index}", "stt":index+2})
    c.update(tgtcthue=10000 * len(c["hdhhdvu"]), tgtthue=800 * len(c["hdhhdvu"]), tgtttbso=10800 * len(c["hdhhdvu"]))
    b["hdhhdvu"].append({**deepcopy(b["hdhhdvu"][0]), "ma":"SECOND", "ten":"Tên hàng dài cần xuống dòng " * 5, "dvtinh":"Thùng", "sluong":2})
    b.update(tgtcthue=20000, tgtthue=1600, tgtttbso=21600)
    batch = prepare(conn)
    sync_input_batch(conn, DateBoundedMsmi([a,b,c]), batch["id"], now_iso)
    records = conn.execute("SELECT id,invoice_number FROM msmi_invoices ORDER BY invoice_number").fetchall()
    invoice_ids = {str(row["invoice_number"]): row["id"] for row in records}
    line_a = conn.execute("SELECT id FROM msmi_invoice_items WHERE invoice_id=?", (invoice_ids["1"],)).fetchone()[0]
    line_b = conn.execute("SELECT id FROM msmi_invoice_items WHERE invoice_id=? AND line_index=1", (invoice_ids["2"],)).fetchone()[0]
    line_box = conn.execute("SELECT id FROM msmi_invoice_items WHERE invoice_id=? AND line_index=2", (invoice_ids["2"],)).fetchone()[0]
    save_mapping(conn, direction="input", item_id=line_a, product_code="R1-KG", now_iso=now_iso)
    create_input_receipt(conn, invoice_ids["1"], now_iso)
    save_mapping(conn, direction="input", item_id=line_box, product_code="R1-CAI", now_iso=now_iso)
    overlap = prepare(conn, start="2026-08-31", end="2026-08-31")
    sync_input_batch(conn, DateBoundedMsmi([c]), overlap["id"], now_iso)
    out = output_invoice(3)
    out["tdlap"] = "2026-08-30T17:00:00Z"
    output_batch = prepare(conn, "output")
    sync_output_batch(conn, OutputFixtureMsmi([out]), output_batch["id"], now_iso)
    output_id = conn.execute("SELECT id FROM outgoing_source_invoices").fetchone()[0]
    output_line = conn.execute("SELECT id FROM outgoing_source_invoice_items").fetchone()[0]
    return {"input_ids":invoice_ids, "line_b":line_b, "line_box":line_box, "output_id":output_id, "output_line":output_line}


def rendered_invoice_html():
    """Execute the actual renderer; avoid assertions tied to JS concatenation syntax."""
    script = """
    global.window = {};
    require('./tdp_system/static/invoice-workbench.js');
    const fmt = {esc:x=>String(x == null ? '' : x),num:x=>String(x||0),money:x=>String(x||0),dateVN:x=>x};
    for (const direction of ['input','output']) {
      const items = ['ready','posted','error'].map((status,i)=>({id:i+1,invoice_series:'TST',invoice_number:String(i+1),
        workbench_status:status,receipt_status:direction==='input'?status:null,stock_status:direction==='output'?(status==='error'?'reversal_required':status):null,
        sync_status:'synced',source_status_class:'issued',invoice_date:'2026-08-31',items:[]}));
      const state = {invoiceDirection:direction,invoiceFrom:'2026-08-01',invoiceTo:'2026-08-31',invoiceStatus:'all',invoiceLineFilter:'all',invoiceWorkbench:{batches:[]},
        invoiceListing:{items,lines:items.map(i=>({id:i.id,invoice_id:i.id,mapping_status:'mapped',inventory_eligible:1,qty:1,amount:2,product_code:'P',issue:''})),totals:{},status_labels:JSON.parse(process.argv[1]),line_labels:JSON.parse(process.argv[2])}};
      console.log(window.TdpInvoiceWorkbench(state,fmt));
    }
    """
    return subprocess.check_output(["node", "-e", script, json.dumps(STATUS_LABELS), json.dumps(LINE_LABELS)],
                                   cwd=Path(__file__).resolve().parent.parent, encoding="utf-8")


class InvoiceRangeTests(unittest.TestCase):
    def setUp(self):
        self.conn = sqlite3.connect(":memory:")
        init_test_database(self.conn)
        self.fixture = seed_round1(self.conn)

    def tearDown(self):
        self.conn.close()

    def payload(self, **kwargs):
        return invoice_range_payload(self.conn, **{ "tenant":"TDP", "invoice_type":"input", "date_from":"2026-08-01", "date_to":"2026-08-31", **kwargs})

    def test_month_union_deduplicates_overlaps_and_uses_dates_not_latest_batch(self):
        payload = self.payload()
        self.assertEqual((3,4), (payload["totals"]["invoice_count"], payload["totals"]["line_count"]))
        self.assertEqual({"1","2","3"}, {x["invoice_number"] for x in payload["items"]})
        self.assertEqual(3, self.conn.execute("SELECT COUNT(*) FROM msmi_invoices").fetchone()[0])
        self.assertEqual(1, self.payload(date_from="2026-08-01", date_to="2026-08-01")["totals"]["invoice_count"])
        self.assertEqual(1, self.payload(date_from="2026-08-31", date_to="2026-08-31")["totals"]["invoice_count"])
        self.assertEqual(0, self.payload(date_from="2026-09-01", date_to="2026-09-30")["totals"]["invoice_count"])
        batches = list_sync_batches(self.conn, tenant="TDP", invoice_type="input", date_from="2026-08-02", date_to="2026-08-02")
        self.assertEqual(1, len(batches["batches"]))

    def test_all_unmapped_rows_precede_mapped_rows_and_filters_match_totals(self):
        payload = self.payload()
        self.assertEqual([True,True,True,False], [bool(x["issue"]) for x in payload["lines"]])
        self.assertEqual({"Thùng":2, "kg":3}, payload["totals"]["qty_by_unit"])
        self.assertEqual(40000, payload["totals"]["line_amount"])
        self.assertEqual(43200, payload["totals"]["invoice_amount"])
        self.assertEqual(1, self.payload(status="posted")["totals"]["invoice_count"])
        self.assertEqual(2, self.payload(status="needs_mapping")["totals"]["invoice_count"])
        self.assertEqual(2, self.payload(line_filter="unmapped")["totals"]["line_count"])
        filtered = self.payload(line_filter="unit_review")
        self.assertEqual(1, filtered["totals"]["line_count"])
        self.assertEqual((10000,21600), (filtered["totals"]["line_amount"], filtered["totals"]["invoice_amount"]))
        self.assertEqual(0, self.payload(status="ready")["totals"]["line_count"])

    def test_read_only_tenant_scope_and_sanitized_payload(self):
        before = self.conn.total_changes
        self.assertEqual(0, self.payload(tenant="OTHER")["counts"]["all"])
        for direction in ("input", "output"):
            data = self.payload(invoice_type=direction)
            self.assertTrue(data["read_only"])
            self.assertFalse(any(key in json.dumps(data) for key in ['raw_json','remote_id','identity_key']))
        self.assertEqual(before, self.conn.total_changes)

    def test_legacy_unbatched_imports_remain_visible_only_to_local_company(self):
        self.conn.execute("DELETE FROM invoice_sync_batch_invoices")
        self.conn.execute("DELETE FROM invoice_sync_batch_output_invoices")
        self.assertEqual(3, self.payload()["counts"]["all"])
        self.assertEqual(1, self.payload(invoice_type="output")["counts"]["all"])
        self.assertEqual(0, self.payload(tenant="OTHER")["counts"]["all"])
        self.conn.execute("UPDATE msmi_invoices SET invoice_type='OUTPUT_ELECTRONIC_INVOICE' WHERE invoice_number='3'")
        self.assertEqual(2, self.payload()["counts"]["all"])

    def test_output_range_overlaps_and_excludes_next_month(self):
        overlap = prepare(self.conn, "output", "2026-08-31", "2026-08-31")
        self.conn.execute("INSERT INTO invoice_sync_batch_output_invoices(batch_id,invoice_id,linked_at) VALUES(?,?,?)", (overlap["id"],self.fixture["output_id"],now_iso()))
        self.assertEqual(1, self.payload(invoice_type="output")["counts"]["all"])
        self.assertEqual(0, self.payload(invoice_type="output",date_to="2026-08-30")["counts"]["all"])
        self.assertEqual(0, self.payload(invoice_type="output",date_from="2026-09-01",date_to="2026-09-30")["counts"]["all"])

    def test_trace_keeps_original_mapped_unit_when_catalog_is_edited(self):
        self.conn.execute("UPDATE products SET unit='thùng' WHERE code='R1-KG'")
        trace = invoice_inventory_trace(self.conn,"input",self.fixture["input_ids"]["1"])
        self.assertEqual("kg",trace["events"][0]["product_unit"])

    def test_frontend_stale_responses_and_decimal_display(self):
        script = r"""
        const fs=require('node:fs'),vm=require('node:vm'),assert=require('node:assert/strict');
        const source=fs.readFileSync('tdp_system/static/app.js','utf8');
        const state={invoiceWorkbenchRequestSerial:0,invoiceDirection:'input',invoiceFrom:'2026-08-01',invoiceTo:'2026-08-31',invoiceStatus:'all',invoiceLineFilter:'all'};
        const requests=[];
        const ctx={state,content:{innerHTML:''},render:()=>{},api:url=>new Promise((resolve,reject)=>requests.push({url,resolve,reject}))};
        vm.createContext(ctx);
        vm.runInContext(source.slice(source.indexOf('  async function loadInvoiceWorkbench('),source.indexOf('  async function prepareInvoiceSyncBatch(')),ctx);
        vm.runInContext(source.slice(source.indexOf('  function n(value)'),source.indexOf('  function dateVN(')),ctx);
        (async()=>{
          const old=ctx.loadInvoiceWorkbench(true);
          state.invoiceDirection='output';state.invoiceFrom='2026-09-01';state.invoiceTo='2026-09-30';
          const current=ctx.loadInvoiceWorkbench(true);
          requests[2].resolve({batches:[]});requests[3].resolve({items:[{id:900}],lines:[]});await current;
          requests[0].resolve({batches:[]});requests[1].resolve({items:[{id:1}],lines:[]});await old;
          assert.equal(state.invoiceOutputRows[0].id,900);assert.equal(state.invoiceInputRows.length,0);
          const fail=ctx.loadInvoiceWorkbench(true);requests[4].reject(Error('offline'));requests[5].resolve({items:[{id:999}]});await fail;
          assert.equal(state.invoiceListing.error,'offline');assert.equal(state.invoiceOutputRows.length,0);
          assert.equal(ctx.stockQty(0.855),'0,855');assert.equal(ctx.stockQty(0.000001),'0,000001');
          // Customer round 4 sample requires comma thousands for VND, not quantity formatting.
          assert.equal(ctx.stockMoney(26465.5),'26,466 đ');assert.equal(ctx.stockMoney(-26465.5),'-26,466 đ');
          assert.equal(ctx.stockQuantitySummary([{unit:'Kg',qty:1},{unit:'kg',qty:2},{unit:'cái',qty:0}],'qty'),'3 kg');
          console.log('OK');
        })().catch(e=>{console.error(e);process.exitCode=1;});
        """
        result = subprocess.run(["node","-e",script], cwd=Path(__file__).resolve().parent.parent,
                                capture_output=True, text=True, encoding="utf-8")
        self.assertEqual(0,result.returncode,result.stderr)

    def test_mapping_conversion_post_trace_and_no_double_entry(self):
        ids = self.fixture
        save_mapping(self.conn, direction="input", item_id=ids["line_b"], product_code="R1-KG", now_iso=now_iso)
        save_conversion(self.conn, direction="input", item_id=ids["line_box"], conversion_factor=12, now_iso=now_iso)
        self.assertEqual(1, self.payload(status="ready")["totals"]["invoice_count"])
        create_input_receipt(self.conn, ids["input_ids"]["2"], now_iso)
        create_input_receipt(self.conn, ids["input_ids"]["2"], now_iso)
        trace = invoice_inventory_trace(self.conn, "input", ids["input_ids"]["2"])
        self.assertEqual(2, len(trace["events"]))
        self.assertEqual("2", trace["invoice"]["invoice_number"])
        self.assertEqual({"R1-KG":1, "R1-CAI":24}, {x["product_code"]:x["qty_delta"] for x in trace["events"]})
        self.assertTrue(all(x["product_name"] for x in trace["events"]))
        save_mapping(self.conn, direction="output", item_id=ids["output_line"], product_code="R1-KG", now_iso=now_iso)
        self.assertEqual(1, self.payload(invoice_type="output", status="ready")["counts"]["ready"])
        post_output_invoice(self.conn, ids["output_id"], confirmed=True, now_iso=now_iso)
        post_output_invoice(self.conn, ids["output_id"], confirmed=True, now_iso=now_iso)
        self.assertEqual(1, self.payload(invoice_type="output", status="posted")["totals"]["invoice_count"])

    def test_changed_posted_source_and_empty_invoice_still_visible_as_error(self):
        self.conn.execute("UPDATE msmi_invoices SET sync_status='review_required',error_message='Nguồn đã thay đổi' WHERE id=?", (self.fixture["input_ids"]["1"],))
        payload = self.payload(status="error")
        self.assertEqual(1, payload["totals"]["invoice_count"])
        self.assertIn("Nguồn", payload["lines"][0]["issue"])
        invoice_id = self.fixture["input_ids"]["3"]
        self.conn.execute("DELETE FROM msmi_invoice_items WHERE invoice_id=?", (invoice_id,))
        self.conn.execute("UPDATE msmi_invoices SET sync_status='review_required' WHERE id=?", (invoice_id,))
        self.assertEqual(2, self.payload(line_filter="error")["totals"]["invoice_count"])

    def test_excel_exactly_matches_filtered_rows_and_never_runs_source_formulas(self):
        self.conn.execute("UPDATE msmi_invoice_items SET source_item_name='=1+1',unit_price=5432.5 WHERE id=?", (self.fixture["line_box"],))
        payload = self.payload(line_filter="unit_review")
        wb = load_workbook(range_workbook(payload))
        ws = wb.active
        self.assertEqual("=1+1", ws["E3"].value)
        self.assertEqual("s", ws["E3"].data_type)
        self.assertEqual("Đơn giá", ws["H2"].value)
        self.assertEqual(5432.5, payload["lines"][0]["unit_price"])
        self.assertEqual(5432.5, ws["H3"].value)  # Source price, not amount / qty.
        self.assertEqual("#,##0", ws["H3"].number_format)
        self.assertIsNone(ws["H4"].value)  # Unit prices are not summed.
        self.assertEqual(10000, ws["I4"].value)
        self.assertEqual(21600, ws.cell(ws.max_row,9).value)
        self.assertEqual(2, ws["G3"].value)
        self.assertEqual("E3", ws.freeze_panes)
        self.assertEqual("$1:$2", ws.print_title_rows)
        wb.close()

    def test_routes_validate_dates_filters_and_return_same_excel(self):
        @contextmanager
        def db():
            yield self.conn
        app = Flask(__name__)
        register_invoice_workbench_routes(app, {"db":db,"now_iso":now_iso,"setting_get":lambda c,k,d:d})
        client = app.test_client()
        url = "/api/invoice-workbench/invoices?from=2026-08-01&to=2026-08-31&invoice_type=input"
        self.assertEqual(3, client.get(url).json["counts"]["all"])
        self.assertEqual(400, client.get(url + "&status=bogus").status_code)
        self.assertEqual(400, client.get("/api/invoice-workbench/invoices?from=2026-08-31&to=2026-08-01").status_code)
        self.assertEqual(200, client.get(url.replace("invoices?", "invoices/export?")).status_code)

    def test_266_fixture_invoices_boundary_days_and_repeat_sync_are_complete(self):
        batch = prepare(self.conn, start="2026-07-01", end="2026-07-31")
        source = []
        for i in range(266):
            item = remote_invoice(i + 1000)
            item["tdlap"] = f"2026-07-{i % 31 + 1:02d}T00:00:00+07:00"
            source.append(item)
        sync_input_batch(self.conn, DateBoundedMsmi(source), batch["id"], now_iso, page_size=50, max_pages=10)
        sync_input_batch(self.conn, DateBoundedMsmi(source), batch["id"], now_iso, page_size=50, max_pages=10)
        payload = self.payload(date_from="2026-07-01",date_to="2026-07-31")
        self.assertEqual(266, payload["totals"]["invoice_count"])
        self.assertEqual(2660000, payload["totals"]["line_amount"])


if __name__ == "__main__":
    unittest.main()
