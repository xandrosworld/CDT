"""Exact, preview-first recovery of input mappings from an old-system report."""

import io
import sqlite3
import unittest
from contextlib import contextmanager
from copy import deepcopy
from pathlib import Path

from flask import Flask
from openpyxl import Workbook

from .contract_modules import init_contract_schema, register_contract_routes, upsert_msmi_invoice
from .invoice_workbench import init_invoice_workbench_schema
from .legacy_invoice_mapping import legacy_input_mapping_plan
from .test_msmi_sync import remote_invoice


NOW = "2026-09-06T07:00:00+00:00"


def workbook_bytes():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "weekend"
    sheet.append(["Bảng kê phiếu nhập"])
    sheet.append(["TT", "Diễn giải", "Mã Vt", "Đvt", "Số lượng", "Thành tiền", "Mã đối tượng", "Số HĐ"])
    sheet.append([0, "Gạo", "P-GAO", "kg", 1, 10000, "0200000001", "1"])
    sheet.append([0, "Hàng khác", "P-KHAC", "kg", 1, 10000, "0200000001", "1"])
    sheet.append([0, "Gạo", "P-GAO", "kg", 1, 10000, "0200000001", "2"])
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


class LegacyInvoiceMappingTests(unittest.TestCase):
    def setUp(self):
        self.conn = sqlite3.connect(":memory:")
        self.conn.row_factory = sqlite3.Row
        self.conn.execute("PRAGMA foreign_keys=ON")
        self.conn.executescript(
            """
            CREATE TABLE settings(key TEXT PRIMARY KEY,value TEXT NOT NULL);
            CREATE TABLE products(code TEXT PRIMARY KEY,name TEXT,unit TEXT);
            CREATE TABLE balances(id INTEGER PRIMARY KEY);
            CREATE TABLE orders(id INTEGER PRIMARY KEY);
            INSERT INTO settings(key,value) VALUES('tenant_code','TDP');
            INSERT INTO products(code,name,unit) VALUES('P-GAO','Gạo kho','kg');
            """
        )
        init_contract_schema(self.conn)
        init_invoice_workbench_schema(self.conn)
        first = remote_invoice(1)
        first["hdhhdvu"][0].update(ma="GAO", ten="Gạo", dvtinh="kg")
        first["hdhhdvu"].append({
            **deepcopy(first["hdhhdvu"][0]), "ma": "KHAC", "ten": "Hàng khác", "stt": 2,
        })
        second = remote_invoice(2)
        second["hdhhdvu"][0].update(ma="GAO", ten="Gạo", dvtinh="kg")
        later = remote_invoice(4)
        later["tdlap"] = "2026-09-01"
        later["hdhhdvu"][0].update(ma="GAO", ten="Gạo", dvtinh="kg")
        unsafe = remote_invoice(5)
        unsafe["hdhhdvu"][0].update(ma="GAO", ten="Gạo", dvtinh="kg")
        ids = []
        for item in (first, second, later, unsafe):
            invoice_id, _created = upsert_msmi_invoice(
                self.conn, item, "INPUT_ELECTRONIC_INVOICE", "TDP", NOW,
            )
            ids.append(invoice_id)
        self.conn.execute("UPDATE msmi_invoices SET sync_status='review_required' WHERE id=?", (ids[-1],))
        self.conn.commit()

    def tearDown(self):
        self.conn.close()

    def test_plan_requires_exact_invoice_line_code_and_unit(self):
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(workbook_bytes()), read_only=True, data_only=True)
        try:
            plan = legacy_input_mapping_plan(
                self.conn, workbook, tenant="TDP", date_from="2026-08-01", date_to="2026-08-31",
            )
        finally:
            workbook.close()
        self.assertEqual((2, 1, 2, 1, 3), (
            plan["safe_lines_in_period"],
            plan["safe_rules"],
            plan["invoices_with_safe_lines"],
            plan["invoices_ready_after"],
            plan["affected_lines_all_periods"],
        ))
        self.assertEqual(1, plan["skipped"]["product_code_missing"])
        self.assertEqual("P-GAO", plan["rows"][0]["product_code"])

    def test_upload_preview_and_explicit_confirm_never_posts_stock(self):
        @contextmanager
        def db_factory():
            yield self.conn
            self.conn.commit()

        def setting_get(conn, key, default=""):
            row = conn.execute("SELECT value FROM settings WHERE key=?", (key,)).fetchone()
            return row[0] if row else default

        app = Flask("legacy-input-mapping")
        register_contract_routes(app, {
            "db": db_factory,
            "now_iso": lambda: NOW,
            "clean_text": lambda value: str(value or "").strip(),
            "number_value": lambda value: float(value or 0),
            "tax_factor": lambda _value: 1.0,
            "setting_get": setting_get,
            "setting_set": lambda conn, key, value: conn.execute(
                "INSERT OR REPLACE INTO settings(key,value) VALUES(?,?)", (key, value)
            ),
            "root": Path.cwd(),
            "data_dir": Path.cwd(),
        })
        client = app.test_client()
        response = client.post(
            "/api/msmi/legacy-mappings/preview",
            data={
                "from": "2026-08-01",
                "to": "2026-08-31",
                "file": (io.BytesIO(workbook_bytes()), "nhap-cu.xlsx"),
            },
            content_type="multipart/form-data",
        )
        self.assertEqual(200, response.status_code, response.get_data(as_text=True))
        preview = response.get_json()
        self.assertNotIn("_representatives", preview)
        unconfirmed = client.post(
            "/api/msmi/legacy-mappings/confirm", json={"token": preview["token"]},
        )
        self.assertEqual(400, unconfirmed.status_code)

        # A missing confirmation does not consume the preview token.
        confirmed = client.post(
            "/api/msmi/legacy-mappings/confirm",
            json={"token": preview["token"], "confirmed": True},
        )
        self.assertEqual(200, confirmed.status_code, confirmed.get_data(as_text=True))
        result = confirmed.get_json()
        self.assertEqual((2, 3, 1, False), (
            result["mapped_lines_in_period"],
            result["mapped_lines_all_periods"],
            result["ready_invoices_in_period"],
            result["stock_changed"],
        ))
        self.assertEqual(0, self.conn.execute(
            "SELECT COUNT(*) FROM invoice_inventory_ledger"
        ).fetchone()[0])


if __name__ == "__main__":
    unittest.main()
