"""Structure-first recognition for the customer's daily workbook.

Only safe sheet metadata and aggregate counts leave this module. In
particular, values from the CCCD sheet are never read into a response, log or
audit payload.
"""

from __future__ import annotations

import hashlib
import io
import json
import math
import re
import unicodedata
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

try:
    from .customer_purchase_layout import customer_purchase_fields
except ImportError:
    from customer_purchase_layout import customer_purchase_fields


REFERENCE_ROLES = {
    "cccd": "identity_reference",
    "t.chieu": "contractor_kitchen_reference",
    "danh muc hang hoa": "product_reference",
    "bao gia": "price_reference",
    "danh muc nha cc": "supplier_reference",
    "gop don": "intermediate_reference",
}

REFERENCE_POLICIES = {
    "identity_reference": {
        "purpose": "documents_only",
        "writePolicy": "locked",
    },
    "contractor_kitchen_reference": {
        "purpose": "contractor_kitchen_validation",
        "writePolicy": "separate_preview_confirm",
        "previewEndpoint": "/api/daily-references/preview",
    },
    "product_reference": {
        "purpose": "product_catalog_validation",
        "writePolicy": "separate_preview_confirm",
        "previewEndpoint": "/api/daily-references/preview",
    },
    "price_reference": {
        "purpose": "price_book_reference",
        "writePolicy": "locked_until_tdp050",
    },
    "supplier_reference": {
        "purpose": "supplier_mapping_validation",
        "writePolicy": "separate_preview_confirm",
        "previewEndpoint": "/api/daily-references/preview",
    },
    "intermediate_reference": {
        "purpose": "intermediate_only",
        "writePolicy": "locked",
    },
}


class DailyWorkbookError(ValueError):
    def __init__(self, message: str, *, code: str = "invalid"):
        super().__init__(message)
        self.code = code


def _slug(value: Any) -> str:
    text = str(value or "").replace("đ", "d").replace("Đ", "D")
    text = unicodedata.normalize("NFKD", text).casefold()
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"[^a-z0-9]+", "", text)


def _reference_role(title: str) -> str:
    raw = str(title or "").replace("đ", "d").replace("Đ", "D")
    normalized = unicodedata.normalize("NFKD", raw).casefold()
    normalized = "".join(char for char in normalized if not unicodedata.combining(char))
    normalized = re.sub(r"\s+", " ", normalized).strip()
    return REFERENCE_ROLES.get(normalized, "")


def _header_map(ws) -> tuple[int, dict[str, int]]:
    aliases = {
        "contractor": {"nhathau", "nhomkhachhang"},
        "product_code": {"mahang", "mavt", "mavattu"},
        "kitchen": {"mabep", "tenbepormabep"},
        "product_name": {"tenhang", "tenvt", "tenhanghoa", "tenthanhdatphat"},
        "qty": {"khoiluong", "soluong", "sldat", "sl"},
        "unit": {"dvt", "donvitinh"},
        "supplier": {"ncc", "nhacungcap", "chonncc"},
        "buy_price": {"giamua", "dongiamua"},
        "sell_price": {"giaban", "dongia", "dongiaban"},
        "date": {"ngaythang", "ngay", "ngaygiao"},
        "actual_qty": {"slthucte", "soluongthucte", "thucnhan"},
        "amount": {"thanhtien"},
        "damaged_qty": {"hong", "hanghong", "soluonghong"},
        "added_qty": {"them", "soluongthem"},
        "reduced_qty": {"giam", "soluonggiam"},
        "missing_qty": {"thieu", "soluongthieu"},
        "note": {"ghichu", "ghichudathang"},
    }
    reverse = {alias: field for field, values in aliases.items() for alias in values}
    best: tuple[int, int, dict[str, int]] | None = None
    for row_index in range(1, min(ws.max_row, 15) + 1):
        mapping: dict[str, int] = {}
        for column in range(1, min(ws.max_column, 50) + 1):
            field = reverse.get(_slug(ws.cell(row_index, column).value))
            if field and (field not in mapping or (
                field == 'supplier' and _slug(ws.cell(row_index, mapping[field]).value) == 'chonncc'
                and _slug(ws.cell(row_index, column).value) in {'ncc', 'nhacungcap'}
            )):
                mapping[field] = column
        if _slug(ws.title) == 'dathang':
            purchase = customer_purchase_fields([ws.cell(row_index, c).value for c in range(1, 17)])
            if purchase:
                mapping.update({{'base_qty': 'qty', 'work_date': 'date'}.get(k, k): v
                                for k, v in purchase.items()})
                mapping.pop('sell_price', None)
        score = len(mapping)
        candidate = (score, row_index, mapping)
        if best is None or candidate[0] > best[0]:
            best = candidate
    if not best:
        return 0, {}
    return best[1], best[2]


def _positive_or_malformed(value: Any) -> bool:
    if value in (None, ""):
        return False
    try:
        numeric = float(str(value).replace(",", ""))
    except (TypeError, ValueError):
        return True
    return (math.isfinite(numeric) and numeric > 0) or not math.isfinite(numeric)


def _business_row_count(ws, header_row: int, mapping: dict[str, int], *, purchase: bool) -> int:
    count = 0
    empty_run = 0
    for row in range(header_row + 1, ws.max_row + 1):
        identity = [
            ws.cell(row, mapping[field]).value
            for field in ("product_code", "product_name", "kitchen") if field in mapping
        ]
        if not any(value not in (None, "") for value in identity):
            empty_run += 1
            if empty_run > 200:
                break
            continue
        empty_run = 0
        if purchase or _positive_or_malformed(ws.cell(row, mapping["qty"]).value):
            count += 1
    return count


def _as_date(value: Any) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value or "").strip()
    for fmt in ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue
    return ""


def _sheet_dates(ws, header_row: int, mapping: dict[str, int]) -> list[str]:
    if "date" not in mapping:
        return []
    values = set()
    for row in range(header_row + 1, min(ws.max_row, header_row + 2000) + 1):
        parsed = _as_date(ws.cell(row, mapping["date"]).value)
        if parsed:
            values.add(parsed)
    return sorted(values)


def _source_hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest().upper()


def analysis_state_hash(analysis: dict[str, Any]) -> str:
    """Bind confirm to safe preview metadata without retaining workbook values."""
    safe = {
        "sourceHash": analysis.get("sourceHash", ""),
        "detectedWorkDate": analysis.get("detectedWorkDate", ""),
        "phase": analysis.get("phase", "first_load"),
        "batchId": analysis.get("batchId"),
        "databaseStateHash": analysis.get("databaseStateHash", ""),
        "daySheets": [
            {
                key: item.get(key)
                for key in (
                    "name", "role", "scope", "rows", "headerRow", "workDates",
                    "titleDateMatches", "confirmAvailable", "errorRows", "warningRows",
                    "writeScope", "diff", "previewIssue",
                )
            }
            for item in analysis.get("daySheets", [])
        ],
        "purchaseSheets": [
            {
                key: item.get(key)
                for key in (
                    "name", "role", "scope", "rows", "headerRow", "confirmAvailable",
                    "errorRows", "warningRows", "writeScope", "diff", "previewIssue",
                )
            }
            for item in analysis.get("purchaseSheets", [])
        ],
        "ignoredSheets": [
            {"name": item.get("name"), "role": item.get("role")}
            for item in analysis.get("ignoredSheets", [])
        ],
    }
    encoded = json.dumps(
        safe, ensure_ascii=False, sort_keys=True, separators=(",", ":")
    ).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest().upper()


def analyze_daily_workbook(path: Path) -> dict[str, Any]:
    source_hash = _source_hash(path)
    workbook = load_workbook(io.BytesIO(path.read_bytes()), data_only=True, read_only=False)
    day_sheets = []
    purchase_sheets = []
    ignored_sheets = []
    try:
        for ws in workbook.worksheets:
            role = _reference_role(ws.title)
            if role:
                # Never inspect sheet values for reference roles, especially CCCD.
                ignored_sheets.append({
                    "name": ws.title,
                    "role": role,
                    **REFERENCE_POLICIES[role],
                })
                continue
            header_row, mapping = _header_map(ws)
            day_required = (
                {"contractor", "kitchen", "qty", "date"}.issubset(mapping)
                and ("product_code" in mapping or "product_name" in mapping)
                and "sell_price" in mapping
            )
            purchase_required = (
                {"product_code", "kitchen", "product_name", "qty", "unit", "supplier",
                 "buy_price", "actual_qty", "amount"}.issubset(mapping)
                and _slug(ws.title) == "dathang"
            )
            if day_required:
                dates = _sheet_dates(ws, header_row, mapping)
                rows = _business_row_count(ws, header_row, mapping, purchase=False)
                title_match = re.search(r"(?:^|\D)(\d{1,2})[.\-_/](\d{1,2})(?:\D|$)", ws.title)
                title_date_matches = True
                if title_match and dates:
                    day, month = int(title_match.group(1)), int(title_match.group(2))
                    title_date_matches = all(
                        int(value[8:10]) == day and int(value[5:7]) == month for value in dates
                    )
                day_sheets.append({
                    "name": ws.title,
                    "role": "daily_orders",
                    "scope": "customer_orders",
                    "rows": rows,
                    "headerRow": header_row,
                    "fields": sorted(mapping),
                    "workDates": dates,
                    "titleDateMatches": title_date_matches,
                    "confirmAvailable": bool(rows and len(dates) == 1 and title_date_matches),
                })
            elif purchase_required:
                purchase_sheets.append({
                    "name": ws.title,
                    "role": "purchase_orders",
                    "scope": "purchase_orders",
                    "rows": _business_row_count(ws, header_row, mapping, purchase=True),
                    "headerRow": header_row,
                    "fields": sorted(mapping),
                    # Recognition is complete, but writing this scope belongs to
                    # the purchase round-trip/finalization tasks.
                    "confirmAvailable": False,
                })
            else:
                ignored_sheets.append({"name": ws.title, "role": "unrecognized_reference"})
    finally:
        workbook.close()
    strict_customer_workbook = bool(day_sheets) and bool(purchase_sheets or any(
        item["role"] in REFERENCE_ROLES.values() for item in ignored_sheets
    ))
    unique_dates = sorted({date_value for item in day_sheets for date_value in item["workDates"]})
    return {
        "sourceHash": source_hash,
        "strictCustomerWorkbook": strict_customer_workbook,
        "daySheets": day_sheets,
        "purchaseSheets": purchase_sheets,
        "ignoredSheets": ignored_sheets,
        "detectedWorkDate": unique_dates[0] if len(unique_dates) == 1 else "",
    }
