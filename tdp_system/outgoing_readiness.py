"""Canonical readiness and shortage projection for outgoing invoices.

The accounting stock used here intentionally comes only from the selected
opening snapshot and ``invoice_inventory_ledger``.  Compatibility movements
such as manual adjustments or the retired BK projection must never unlock an
outgoing invoice.  Local drafts and locally confirmed invoices are overlaid as
holds until the latter can be matched to a posted canonical source invoice.
"""

from __future__ import annotations

import re
import math
import json
from collections import defaultdict
from datetime import datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.page import PageMargins

try:
    from .stock_tax_policy import exempt_order_codes
    from .outgoing_line_policy import unit_issues, draft_policy_rows
except ImportError:
    from stock_tax_policy import exempt_order_codes
    from outgoing_line_policy import unit_issues, draft_policy_rows

try:
    from invoice_inventory import invoice_stock_rows, _minimum_balance_from, selected_opening_snapshot
    from template_workbook import safe_workbook_bytes
except ImportError:  # pragma: no cover - package invocation
    from .invoice_inventory import invoice_stock_rows, _minimum_balance_from, selected_opening_snapshot
    from .template_workbook import safe_workbook_bytes


EPSILON = 1e-9
QTY_QUANTUM = Decimal("0.000001")


class OutgoingReadinessError(ValueError):
    def __init__(self, message: str, *, code: str = "invalid_readiness", status: int = 409):
        super().__init__(message)
        self.code = code
        self.status = status


def _number(value: Any) -> float:
    try:
        result = float(value or 0)
    except (TypeError, ValueError):
        return 0.0
    return result if result == result and abs(result) != float("inf") else 0.0


def _vnd(*values: Any) -> int:
    try:
        result = Decimal("1")
        for value in values:
            result *= Decimal(str(value or 0))
        return int(result.quantize(Decimal("1"), rounding=ROUND_HALF_UP))
    except (InvalidOperation, TypeError, ValueError):
        return 0


def _net_delivered(order: dict[str, Any]) -> float:
    return max(_number(order.get("actual_delivered")) - _number(order.get("customer_return_qty")), 0)


def _strict_date(value: Any, label: str) -> str:
    text = str(value or "").strip()
    try:
        return datetime.strptime(text, "%Y-%m-%d").date().isoformat()
    except ValueError:
        raise OutgoingReadinessError(
            f"{label} phải là ngày YYYY-MM-DD hợp lệ",
            code="invalid_date_range",
            status=400,
        ) from None


def _display_date(value: Any) -> str:
    text = str(value or "").strip()
    try:
        return datetime.strptime(text, "%Y-%m-%d").strftime("%d/%m/%Y")
    except ValueError:
        return text


def _table_exists(conn, table: str) -> bool:
    return bool(conn.execute(
        "SELECT 1 FROM sqlite_master WHERE type='table' AND name=?", (table,),
    ).fetchone())


def _normalized_identity_part(value: Any, *, uppercase: bool = False) -> str:
    text = str(value or "").strip()
    return text.upper() if uppercase else text


def _normalized_unit(value: Any) -> str:
    return " ".join(str(value or "").strip().split()).casefold()


def _fingerprint_qty(value: Any, *, invert: bool = False) -> Decimal | None:
    try:
        qty = Decimal(str(value))
        if invert:
            qty = -qty
        if not qty.is_finite() or qty <= 0:
            return None
        normalized = qty.quantize(QTY_QUANTUM, rounding=ROUND_HALF_UP)
    except (InvalidOperation, TypeError, ValueError):
        return None
    return normalized if normalized > 0 else None


def _canonicalized_local_draft_ids(conn) -> set[int]:
    """Return local issued drafts replaced by the same posted M-Invoice quantity.

    A series/number/date collision alone is not enough to release a local hold.
    The canonical POST must come from M-Invoice and its stock quantities must
    equal the local draft after grouping by the canonical product code.  Local
    units are also required to match the catalog unit because ledger quantities
    are expressed in that unit.
    """
    if not (
        _table_exists(conn, "outgoing_source_invoices")
        and _table_exists(conn, "invoice_inventory_ledger")
    ):
        return set()

    local: dict[int, dict[str, Any]] = {}
    for row in conn.execute(
        """SELECT d.id draft_id,d.issued_invoice_series,d.issued_invoice_number,
                  COALESCE(d.issued_invoice_date,d.invoice_date) issued_invoice_date,
                  l.id line_id,l.product_code,l.qty,l.unit,p.unit catalog_unit
             FROM outgoing_invoice_drafts d
             LEFT JOIN outgoing_invoice_lines l ON l.draft_id=d.id
             LEFT JOIN products p ON p.code=l.product_code
            WHERE d.status='issued'
            ORDER BY d.id,l.id"""
    ):
        draft_id = int(row["draft_id"])
        record = local.setdefault(draft_id, {
            "identity": (
                _normalized_identity_part(row["issued_invoice_series"], uppercase=True),
                _normalized_identity_part(row["issued_invoice_number"]),
                _normalized_identity_part(row["issued_invoice_date"]),
            ),
            "quantities": defaultdict(Decimal),
            "valid": True,
            "has_line": False,
        })
        if row["line_id"] is None:
            record["valid"] = False
            continue
        record["has_line"] = True
        product_code = str(row["product_code"] or "").strip()
        qty = _fingerprint_qty(row["qty"])
        source_unit = _normalized_unit(row["unit"])
        catalog_unit = _normalized_unit(row["catalog_unit"])
        if (
            not product_code
            or qty is None
            or not source_unit
            or not catalog_unit
            or source_unit != catalog_unit
        ):
            record["valid"] = False
            continue
        record["quantities"][product_code] += qty

    candidates: dict[tuple[str, str, str], list[dict[str, Any]]] = defaultdict(list)
    source_records: dict[int, dict[str, Any]] = {}
    for row in conn.execute(
        """SELECT s.id source_id,s.invoice_series,s.invoice_number,s.invoice_date,
                  il.id ledger_id,il.event_type,il.product_code,il.qty_delta
             FROM outgoing_source_invoices s
             LEFT JOIN invoice_inventory_ledger il
               ON il.direction='output' AND il.status='posted'
              AND il.source_invoice_table='outgoing_source_invoices'
              AND il.source_invoice_id=s.id
            WHERE LOWER(TRIM(COALESCE(s.source,'')))='minvoice'
              AND s.source_status_class='issued'
              AND s.sync_status='synced'
              AND s.stock_status='posted'
              AND EXISTS (
                  SELECT 1 FROM outgoing_invoice_drafts d
                   WHERE d.status='issued'
                     AND TRIM(COALESCE(d.issued_invoice_number,''))!=''
                     AND UPPER(TRIM(COALESCE(d.issued_invoice_series,'')))=
                         UPPER(TRIM(COALESCE(s.invoice_series,'')))
                     AND TRIM(COALESCE(d.issued_invoice_number,''))=
                         TRIM(COALESCE(s.invoice_number,''))
                     AND TRIM(COALESCE(d.issued_invoice_date,d.invoice_date,''))=
                         TRIM(COALESCE(s.invoice_date,''))
              )
            ORDER BY s.id,il.id"""
    ):
        source_id = int(row["source_id"])
        record = source_records.setdefault(source_id, {
            "identity": (
                _normalized_identity_part(row["invoice_series"], uppercase=True),
                _normalized_identity_part(row["invoice_number"]),
                _normalized_identity_part(row["invoice_date"]),
            ),
            "quantities": defaultdict(Decimal),
            "valid": True,
            "has_line": False,
        })
        if row["ledger_id"] is None or row["event_type"] != "POST":
            record["valid"] = False
            continue
        record["has_line"] = True
        product_code = str(row["product_code"] or "").strip()
        qty = _fingerprint_qty(row["qty_delta"], invert=True)
        if not product_code or qty is None:
            record["valid"] = False
            continue
        record["quantities"][product_code] += qty

    for record in source_records.values():
        candidates[record["identity"]].append(record)

    matched: set[int] = set()
    for draft_id, record in local.items():
        if (
            not record["valid"]
            or not record["has_line"]
            or not all(record["identity"])
        ):
            continue
        local_fingerprint = tuple(sorted(
            (code, qty.quantize(QTY_QUANTUM, rounding=ROUND_HALF_UP))
            for code, qty in record["quantities"].items()
        ))
        if not local_fingerprint:
            continue
        for candidate in candidates.get(record["identity"], []):
            if not candidate["valid"] or not candidate["has_line"]:
                continue
            source_fingerprint = tuple(sorted(
                (code, qty.quantize(QTY_QUANTUM, rounding=ROUND_HALF_UP))
                for code, qty in candidate["quantities"].items()
            ))
            if source_fingerprint == local_fingerprint:
                matched.add(draft_id)
                break
    return matched


def canonical_available_stock(conn, as_of: str = "") -> dict[str, dict[str, Any]]:
    """Return canonical stock less every active local outgoing hold."""
    try:
        canonical_rows = invoice_stock_rows(conn, include_zero=True)
    except ValueError as exc:
        raise OutgoingReadinessError(
            f"Sổ vật tư hàng hóa chưa sẵn sàng: {exc}",
            code="canonical_stock_conflict",
        ) from None
    stock = {
        str(row["product_code"] or "").strip(): {
            **row,
            "canonical_qty": _number(row.get("closing_qty")),
            "reserved_qty": 0.0,
            "pending_sync_issued_qty": 0.0,
        }
        for row in canonical_rows
    }
    if as_of:
        safe_date = _strict_date(as_of, "Ngày hóa đơn")
        if conn.execute(
            "SELECT 1 FROM inventory_transactions WHERE source_type='OPENING' "
            "AND status='posted' AND txn_date>? LIMIT 1", (safe_date,),
        ).fetchone():
            raise OutgoingReadinessError(
                "Ngày hóa đơn nằm trước kỳ tồn đầu đã chốt; cần kiểm tra và mở lại kỳ trước khi lập tiếp",
                code="backdated_before_opening_snapshot",
            )
        for code, item in stock.items():
            item["canonical_qty"] = _minimum_balance_from(conn, code, safe_date)

    for row in conn.execute(
        """SELECT product_code,SUM(qty_out) qty
             FROM inventory_transactions
            WHERE source_type='OUTGOING_DRAFT' AND status='reserved'
            GROUP BY product_code"""
    ):
        code = str(row["product_code"] or "").strip()
        item = stock.setdefault(code, {
            "product_code": code, "product_name": "", "unit": "", "canonical_qty": 0.0,
            "reserved_qty": 0.0, "pending_sync_issued_qty": 0.0,
        })
        item["reserved_qty"] += max(_number(row["qty"]), 0)

    canonicalized_drafts = _canonicalized_local_draft_ids(conn)
    for row in conn.execute(
        """SELECT d.id draft_id,l.product_code,SUM(l.qty) qty
             FROM outgoing_invoice_drafts d
             JOIN outgoing_invoice_lines l ON l.draft_id=d.id
            WHERE d.status='issued'
            GROUP BY d.id,l.product_code"""
    ):
        if int(row["draft_id"]) in canonicalized_drafts:
            continue
        code = str(row["product_code"] or "").strip()
        item = stock.setdefault(code, {
            "product_code": code, "product_name": "", "unit": "", "canonical_qty": 0.0,
            "reserved_qty": 0.0, "pending_sync_issued_qty": 0.0,
        })
        item["pending_sync_issued_qty"] += max(_number(row["qty"]), 0)

    for item in stock.values():
        item["raw_available_qty"] = (
            item["canonical_qty"] - item["reserved_qty"] - item["pending_sync_issued_qty"]
        )
        item["available_qty"] = max(item["raw_available_qty"], 0)
    return stock


def invoice_tax_error(value):
    raw = str(value if value is not None else '').strip().upper().replace(' ', '')
    if raw in {'KKKNT', 'KCT', 'KHÔNGKÊKHAI', 'KHONGKEKHAI', 'KHÔNGCHỊUTHUẾ', 'KHONGCHIUTHUE'}:
        return ''
    try:
        number = Decimal(raw.rstrip('%'))
        if not number.is_finite():
            raise InvalidOperation
        if not raw.endswith('%') and 0 < abs(number) < 1:
            number *= 100
        if number in {Decimal('-2'), Decimal('-1'), Decimal('0'), Decimal('5'), Decimal('8'), Decimal('10')}:
            return ''
    except (InvalidOperation, ValueError):
        pass
    return 'Thuế chưa hợp lệ; chọn 0%, 5%, 8%, 10%, KCT hoặc KKKNT'


def invoice_order_issues(orders):
    issues = []
    for order in orders:
        saved = order.get('errors') or []
        if isinstance(saved, str):
            try:
                saved = json.loads(saved)
            except (ValueError, TypeError):
                saved = ['Dòng đơn cần kiểm tra lại']
        messages = list(saved) if isinstance(saved, list) else ['Dòng đơn cần kiểm tra lại']
        tax_error = invoice_tax_error(order.get('tax'))
        if tax_error:
            messages.append(tax_error)
        if messages:
            issues.append({'order_id': order['id'], 'product_code': order.get('product_code', ''),
                           'product_name': order.get('product_name', ''), 'kitchen': order.get('kitchen', ''),
                           'source_row': order.get('source_row'), 'messages': list(dict.fromkeys(messages))})
    return issues


def validate_draft_export_stock(conn, draft_id, invoice_date=""):
    """Read-only recheck before handing off a file or saving a remote draft."""
    required = defaultdict(float)
    lines = conn.execute("SELECT product_code,product_name,qty,tax FROM outgoing_invoice_lines WHERE draft_id=?", (draft_id,)).fetchall()
    policy_rows=draft_policy_rows(conn,draft_id)
    # A legacy download or remote-draft action must not bypass reconciliation.
    try:
        from .outgoing_unissued import issued_allocations
    except ImportError:
        from outgoing_unissued import issued_allocations
    external={}
    issued,warnings=issued_allocations(conn,external_quantities=external)
    parties={r['contractor'] for r in policy_rows}
    relevant=[w for w in warnings if not w['contractor'] or w['contractor'] in parties]
    if relevant:
        raise OutgoingReadinessError(relevant[0]['message'],code='issued_source_unresolved')
    held_by_order={r['order_id']:r['qty'] for r in conn.execute("""SELECT a.order_id,SUM(a.qty) qty
        FROM outgoing_order_allocations a JOIN outgoing_invoice_drafts d ON d.id=a.draft_id
        WHERE d.status='draft' GROUP BY a.order_id""")}
    for r in policy_rows:
        settlement=conn.execute('SELECT external_issued_qty FROM outgoing_waiting_settlements WHERE order_id=?',(r['order_id'],)).fetchone()
        if external.get(r['order_id'],0)>(settlement['external_issued_qty'] if settlement else 0)+EPSILON:
            raise OutgoingReadinessError(f"{r['product_code']}: có hóa đơn vừa ký chưa đối trừ vào phần giữ chờ; cập nhật bảng kê trước khi tải.",code='issued_quantity_in_draft')
        order=conn.execute('SELECT actual_delivered,customer_return_qty FROM orders WHERE id=?',(r['order_id'],)).fetchone()
        if held_by_order.get(r['order_id'],0)+issued.get(r['order_id'],0)>order['actual_delivered']-order['customer_return_qty']+EPSILON:
            raise OutgoingReadinessError(f"{r['product_code']}: phần chờ còn chứa lượng đã ký; cập nhật bảng kê trước khi tải.",code='issued_quantity_in_draft')
    units=unit_issues(conn,policy_rows)
    if units:
        raise OutgoingReadinessError(next(iter(units.values()))['message'],code='order_unit_mismatch')
    exempt = exempt_order_codes(conn, policy_rows)
    for line in lines:
        tax_error = invoice_tax_error(line['tax'])
        if tax_error:
            raise OutgoingReadinessError(f"Mã {line['product_code']}: {tax_error}", code='invalid_invoice_tax')
        qty = float(line["qty"])
        if not math.isfinite(qty) or qty <= 0:
            raise OutgoingReadinessError("Số lượng dự thảo không hợp lệ", code="invalid_draft_quantity")
        required[line["product_code"]] += qty
    held = {row["product_code"]: float(row["qty"]) for row in conn.execute(
        "SELECT product_code,SUM(qty_out) qty FROM inventory_transactions "
        "WHERE source_type='OUTGOING_DRAFT' AND source_id=? AND status='reserved' GROUP BY product_code",
        (str(draft_id),),
    )}
    if not required or set(required) != set(held) or any(abs(qty - held[code]) > EPSILON for code, qty in required.items()):
        raise OutgoingReadinessError("Phần giữ tồn không khớp dự thảo; cần tính lại trước khi tải/gửi hóa đơn", code="draft_reservation_mismatch")
    stock = canonical_available_stock(conn, invoice_date)
    for code, qty in required.items():
        if code in exempt:
            continue
        available = max(stock.get(code, {}).get("raw_available_qty", 0) + held[code], 0)
        if qty > available + EPSILON:
            raise OutgoingReadinessError(
                f"Mã {code}: cần {qty:g}, có thể xuất {available:g}, thiếu {qty - available:g}; cần tính lại dự thảo",
                code="canonical_stock_overcommitted",
            )


def validate_issued_draft_stock(conn, draft_id, invoice_date, invoice_series, invoice_number):
    """Preview the confirmed state atomically, without leaving any mutation.

    Matching an already posted M-Invoice must release only this draft's hold,
    not deduct the same physical invoice twice. Other drafts keep their holds.
    """
    lines = conn.execute("SELECT product_code,product_name,qty,tax FROM outgoing_invoice_lines WHERE draft_id=?", (draft_id,)).fetchall()
    policy_rows=draft_policy_rows(conn,draft_id)
    units=unit_issues(conn,policy_rows)
    if units:
        raise OutgoingReadinessError(next(iter(units.values()))['message'],code='order_unit_mismatch')
    exempt = exempt_order_codes(conn, policy_rows)
    if not lines or any(not math.isfinite(float(r["qty"])) or float(r["qty"]) <= 0 for r in lines):
        raise OutgoingReadinessError("Dự thảo thiếu dòng hàng hoặc có số lượng không hợp lệ", code="invalid_draft_quantity")
    conn.execute("SAVEPOINT validate_issued_stock")
    try:
        conn.execute(
            "UPDATE outgoing_invoice_drafts SET status='issued',issued_invoice_date=?,"
            "issued_invoice_series=?,issued_invoice_number=? WHERE id=?",
            (invoice_date, invoice_series, invoice_number, draft_id),
        )
        conn.execute(
            "UPDATE inventory_transactions SET status='posted' WHERE source_type='OUTGOING_DRAFT' "
            "AND source_id=? AND status='reserved'", (str(draft_id),),
        )
        stock = canonical_available_stock(conn, invoice_date)
        required = defaultdict(float)
        for line in lines:
            required[line["product_code"]] += float(line["qty"])
        shortages = []
        for code, quantity in required.items():
            if code in exempt:
                continue
            remaining = stock.get(code, {}).get("raw_available_qty", -quantity)
            if remaining < -EPSILON:
                available = max(quantity + remaining, 0)
                shortages.append(f"{code}: cần {quantity:g}, có thể xuất {available:g}, thiếu {quantity - available:g}")
        if shortages:
            raise OutgoingReadinessError(
                "Không đủ tồn hóa đơn để xác nhận; cần đối chiếu lại: " + "; ".join(shortages),
                code="canonical_stock_overcommitted",
            )
    finally:
        conn.execute("ROLLBACK TO SAVEPOINT validate_issued_stock")
        conn.execute("RELEASE SAVEPOINT validate_issued_stock")


def validate_demand_orders(conn, orders: list[dict[str, Any]]) -> None:
    known_codes = {
        str(row["code"] or "").strip()
        for row in conn.execute("SELECT code FROM products")
    }
    missing_code = [str(row.get("id") or "?") for row in orders if not str(row.get("product_code") or "").strip()]
    if missing_code:
        raise OutgoingReadinessError(
            "Không thể tính hóa đơn: dòng đơn thiếu mã hàng TĐP (ID " + ", ".join(missing_code[:10]) + ")",
            code="missing_product_code",
        )
    unknown_code = [
        str(row.get("product_code") or "").strip() for row in orders
        if str(row.get("product_code") or "").strip() not in known_codes
    ]
    if unknown_code:
        raise OutgoingReadinessError(
            "Không thể tính hóa đơn: mã hàng không còn trong danh mục: " + ", ".join(sorted(set(unknown_code))[:10]),
            code="unknown_product_code",
        )
    missing_contractor = [
        str(row.get("id") or "?") for row in orders if not str(row.get("contractor") or "").strip()
    ]
    if missing_contractor:
        raise OutgoingReadinessError(
            "Không thể tách hóa đơn: dòng đơn thiếu nhà thầu (ID " + ", ".join(missing_contractor[:10]) + ")",
            code="missing_contractor",
        )


def allocation_by_order(conn, batch_ids: list[int]) -> dict[int, dict[str, float]]:
    if not batch_ids:
        return {}
    placeholders = ",".join("?" for _ in batch_ids)
    rows = conn.execute(
        f"""SELECT l.order_id,
                    SUM(CASE WHEN d.status='draft' THEN l.qty ELSE 0 END) drafted_qty,
                    SUM(CASE WHEN d.status='issued' THEN l.qty ELSE 0 END) issued_qty
               FROM outgoing_order_allocations l
               JOIN outgoing_invoice_drafts d ON d.id=l.draft_id
               JOIN orders o ON o.id=l.order_id
              WHERE d.status!='cancelled' AND o.batch_id IN ({placeholders})
              GROUP BY l.order_id""",
        tuple(batch_ids),
    )
    result = {
        int(row["order_id"]): {
            "drafted_qty": max(_number(row["drafted_qty"]), 0),
            "issued_qty": max(_number(row["issued_qty"]), 0),
        }
        for row in rows
    }
    try:
        from .outgoing_unissued import issued_allocations
    except ImportError:
        from outgoing_unissued import issued_allocations
    issued, _ = issued_allocations(conn)
    selected = {r['id'] for r in conn.execute(f'SELECT id FROM orders WHERE batch_id IN ({placeholders})',tuple(batch_ids))}
    for oid in selected:
        if oid in issued or oid in result:
            result.setdefault(oid, {'drafted_qty': 0.0, 'issued_qty': 0.0})['issued_qty'] = issued.get(oid, 0.0)
    return result


def _project_rows(
    conn,
    orders: list[dict[str, Any]],
    batch_ids: list[int],
) -> list[dict[str, Any]]:
    validate_demand_orders(conn, orders)
    stock = canonical_available_stock(conn)
    available = {code: item["available_qty"] for code, item in stock.items()}
    allocated = allocation_by_order(conn, batch_ids)
    exempt = exempt_order_codes(conn, orders)
    units=unit_issues(conn,orders)
    rows: list[dict[str, Any]] = []
    for item in orders:
        demand = _net_delivered(item)
        if demand <= EPSILON:
            continue
        held = allocated.get(int(item["id"]), {"drafted_qty": 0.0, "issued_qty": 0.0})
        drafted = held["drafted_qty"]
        issued = held["issued_qty"]
        if drafted + issued > demand + EPSILON:
            raise OutgoingReadinessError(
                f"Dòng đơn {item['id']} đã phân bổ vượt số thực giao; cần đối chiếu trước khi lập tiếp",
                code="allocated_over_demand",
            )
        remaining = max(demand - drafted - issued, 0)
        code = str(item["product_code"]).strip()
        have = max(available.get(code, 0), 0)
        invoiceable = 0 if item['id'] in units else remaining if code in exempt else min(remaining, have)
        pending = max(remaining - invoiceable, 0)
        available[code] = max(have - invoiceable, 0)
        unit_price = _vnd(item.get("sell_price"))
        rows.append({
            "order_id": int(item["id"]),
            "batch_id": int(item["batch_id"]),
            "work_date": str(item.get("work_date") or ""),
            "contractor": str(item.get("contractor") or "").strip(),
            "kitchen": str(item.get("kitchen") or "").strip(),
            "product_code": code,
            "product_name": str(item.get("product_name") or "").strip(),
            "unit": str(item.get("unit") or "").strip(),
            "unit_price": unit_price,
            "demand_qty": demand,
            "drafted_qty": drafted,
            "issued_qty": issued,
            "allocated_qty": drafted + issued,
            "invoiceable_qty": invoiceable,
            "pending_qty": pending,
            "pending_reason": units[item['id']]['message'] if item['id'] in units else (
                f"Thiếu {pending:g} {item.get('unit') or ''} tồn hóa đơn khả dụng "
                "sau khi trừ phần đã giữ/đã phát hành. Phần thiếu được giữ lại để xử lý tiếp."
                if pending > EPSILON else ""
            ),
            "available_before": have,
            "negative_stock_allowed": code in exempt,
            "demand_value": _vnd(demand, unit_price),
            "drafted_value": _vnd(drafted, unit_price),
            "issued_value": _vnd(issued, unit_price),
            "invoiceable_value": _vnd(invoiceable, unit_price),
            "pending_value": _vnd(pending, unit_price),
        })
    return rows


_TOTAL_FIELDS = (
    "demand_qty", "drafted_qty", "issued_qty", "allocated_qty", "invoiceable_qty", "pending_qty",
    "demand_value", "drafted_value", "issued_value", "invoiceable_value", "pending_value",
)


def _totals(rows: list[dict[str, Any]]) -> dict[str, Any]:
    result: dict[str, Any] = {}
    for field in _TOTAL_FIELDS:
        value = sum(row[field] for row in rows)
        result[field] = int(value) if field.endswith("_value") else value
    return result


def _contractor_summaries(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        grouped[row["contractor"]].append(row)
    return [
        {"contractor": contractor, **_totals(items)}
        for contractor, items in sorted(grouped.items())
    ]


def _replaceable_batch_holds(conn, batch_id: int) -> dict[str, float]:
    return {r['product_code']: float(r['qty']) for r in conn.execute(
        "SELECT t.product_code,SUM(t.qty_out) qty FROM inventory_transactions t "
        "JOIN outgoing_invoice_drafts d ON CAST(d.id AS TEXT)=t.source_id "
        "WHERE t.source_type='OUTGOING_DRAFT' AND t.status='reserved' AND d.batch_id=? AND d.status='draft' "
        "AND COALESCE(d.draft_kind,'standard')='standard' "
        "AND COALESCE(d.minvoice_status,'not_sent') NOT IN ('saved','saving','unknown') GROUP BY t.product_code", (batch_id,))}


def batch_product_stock_trace(conn, batch_id: int, product_code: str) -> dict[str, Any]:
    """Explain the exact stock basis used by the selected batch, without writes."""
    if not conn.execute('SELECT 1 FROM batches WHERE id=?', (batch_id,)).fetchone():
        raise OutgoingReadinessError('Không tìm thấy đơn hàng.', status=404)
    if not conn.execute('SELECT 1 FROM orders WHERE batch_id=? AND product_code=?',
                        (batch_id, product_code)).fetchone():
        raise OutgoingReadinessError('Mặt hàng không có trong đơn đang chọn.', status=404)
    stock = canonical_available_stock(conn).get(product_code)
    if not stock:
        raise OutgoingReadinessError('Không tìm thấy mặt hàng trong kho.', status=404)
    opening = selected_opening_snapshot(conn, '9999-12-31')
    start = opening[1] if opening else '0001-01-01'
    events = []
    opening_editor = None
    if opening:
        opening_rows = conn.execute(
            "SELECT id,txn_date,product_code,source_id,source_line,qty_in,qty_out,unit_cost,note,updated_at FROM inventory_transactions "
            "WHERE source_type='OPENING' AND status='posted' AND source_id=? AND product_code=? ORDER BY id",
            (opening[0], product_code),
        ).fetchall()
        for row in opening_rows:
            events.append({'date': row['txn_date'], 'label': 'Tồn đầu kỳ',
                           'reference': row['note'] or 'Tồn đầu kỳ đã nhập',
                           'qty_delta': float(row['qty_in']) - float(row['qty_out'])})
        if len(opening_rows) == 1:
            row = dict(opening_rows[0])
            if (row['source_id'] == row['txn_date'][:7] and row['txn_date'].endswith('-01')
                    and row['source_line'] == product_code):
                opening_editor = {'period': row['source_id'], 'qty': events[0]['qty_delta'], 'expected': row}
    movements = conn.execute(
        """SELECT l.txn_date,l.direction,l.event_type,l.qty_delta,l.source_line_index,
                  l.source_invoice_table,l.source_invoice_id,l.source_line_id,
                  COALESCE(i.invoice_series,o.invoice_series,'') invoice_series,
                  COALESCE(i.invoice_number,o.invoice_number,'') invoice_number,
                  COALESCE(i.invoice_date,o.invoice_date,l.txn_date) invoice_date,
                  COALESCE(i.id,o.id) linked_invoice_id,
                  COALESCE(c.note,'') note
           FROM invoice_inventory_effective_ledger l
           LEFT JOIN msmi_invoices i ON l.source_invoice_table='msmi_invoices' AND i.id=l.source_invoice_id
           LEFT JOIN outgoing_source_invoices o ON l.source_invoice_table='outgoing_source_invoices' AND o.id=l.source_invoice_id
           LEFT JOIN invoice_inventory_confirmations c ON c.id=l.confirmation_id
           WHERE l.status='posted' AND l.product_code=? AND l.txn_date>=? ORDER BY l.txn_date,l.id""",
        (product_code, start),
    ).fetchall()
    for row in movements:
        label = 'Nhập kho' if row['direction'] == 'input' else 'Xuất kho'
        if row['event_type'] == 'REVERSAL':
            label = 'Hoàn tác nhập' if row['direction'] == 'input' else 'Hoàn tác xuất'
        reference = (f"{row['invoice_series']} / {row['invoice_number']} · Dòng {row['source_line_index']}"
                     if row['invoice_number'] else row['note'])
        events.append({'date': row['txn_date'], 'label': label, 'reference': reference,
                       'qty_delta': float(row['qty_delta']),
                       'direction': row['direction'], 'invoice_id': row['source_invoice_id'],
                       'invoice_date': row['invoice_date'],
                       'line_id': row['source_line_id'],
                       'can_open_invoice': row['linked_invoice_id'] is not None})
    balance = 0.0
    for event in events:
        balance += event['qty_delta']
        event['balance_qty'] = balance
    released = _replaceable_batch_holds(conn, batch_id).get(product_code, 0)
    reserved = stock['reserved_qty'] - released
    available = stock['raw_available_qty'] + released
    return {'batch_id': batch_id, 'product_code': product_code, 'product_name': stock['product_name'],
            'unit': stock['unit'], 'opening_date': opening[1] if opening else '',
            'opening_qty': stock.get('opening_qty', 0), 'input_qty': stock.get('input_qty', 0),
            'output_qty': stock.get('output_qty', 0), 'reversal_qty': stock.get('reversal_qty', 0),
            'closing_qty': stock['canonical_qty'], 'reserved_qty': reserved,
            'pending_sync_issued_qty': stock['pending_sync_issued_qty'], 'available_qty': available,
            'negative_opening_only': stock.get('opening_qty', 0) < -EPSILON and not movements,
            'events': events, 'opening_editor': opening_editor, 'read_only': True}


def batch_readiness_payload(conn, batch_id: int) -> dict[str, Any]:
    batch = conn.execute("SELECT * FROM batches WHERE id=?", (batch_id,)).fetchone()
    if not batch:
        raise OutgoingReadinessError("Không tìm thấy phiên đơn", code="batch_not_found", status=404)
    all_orders = [dict(row) for row in conn.execute('SELECT * FROM orders WHERE batch_id=? ORDER BY contractor,id', (batch_id,))]
    orders = [
        dict(row) for row in conn.execute(
            "SELECT * FROM orders WHERE batch_id=? ORDER BY contractor,id", (batch_id,),
        ) if _net_delivered(dict(row)) > EPSILON
    ]
    rows = _project_rows(conn, orders, [batch_id])
    stock = canonical_available_stock(conn)
    released = _replaceable_batch_holds(conn, batch_id)
    blocking_issues = []
    exempt = exempt_order_codes(conn, orders)
    negative_warnings = []
    for code in sorted({o['product_code'] for o in orders}):
        remaining = stock.get(code,{}).get('raw_available_qty',0) + released.get(code,0)
        if remaining < -EPSILON:
            (negative_warnings if code in exempt else blocking_issues).append({'product_code':code,'qty':remaining,
                                    'product_name': stock.get(code, {}).get('product_name', ''),
                                    'unit': stock.get(code, {}).get('unit', ''),
                                    'message':f'Tồn {code} đang âm {abs(remaining):g}. Cần đối chiếu kho trước khi tạo file.'})
    return {
        "batch_id": batch_id,
        "work_date": batch["work_date"],
        **_totals(rows),
        "contractors": _contractor_summaries(rows),
        "rows": rows,
        "blocking_issues": blocking_issues,
        "negative_stock_warnings": negative_warnings,
        "order_issues": invoice_order_issues(all_orders),
        "stock_basis": "OPENING + invoice_inventory_ledger − active local holds",
    }


def period_shortage_payload(
    conn,
    date_from: Any,
    date_to: Any,
    contractor: Any = "",
) -> dict[str, Any]:
    safe_from = _strict_date(date_from, "Từ ngày")
    safe_to = _strict_date(date_to, "Đến ngày")
    if safe_from > safe_to:
        raise OutgoingReadinessError(
            "Từ ngày không được sau đến ngày", code="invalid_date_range", status=400,
        )
    contractor_text = str(contractor or "").strip()
    params: list[Any] = [safe_from, safe_to]
    contractor_clause = ""
    if contractor_text:
        contractor_clause = " AND TRIM(COALESCE(o.contractor,''))=?"
        params.append(contractor_text)
    orders = [dict(row) for row in conn.execute(
        """SELECT o.* FROM orders o JOIN batches b ON b.id=o.batch_id
            WHERE b.status='approved' AND b.work_date BETWEEN ? AND ?"""
        + contractor_clause
        + " ORDER BY b.work_date,o.batch_id,o.contractor,o.id",
        tuple(params),
    )]
    orders = [row for row in orders if _net_delivered(row) > EPSILON]
    batch_ids = sorted({int(row["batch_id"]) for row in orders})
    detail_rows = _project_rows(conn, orders, batch_ids)

    grouped: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for row in detail_rows:
        grouped[(row["contractor"], row["product_code"])].append(row)
    rows: list[dict[str, Any]] = []
    for (party, code), items in sorted(grouped.items()):
        summary = _totals(items)
        rows.append({
            "contractor": party,
            "product_code": code,
            "product_name": items[0]["product_name"],
            "unit": items[0]["unit"],
            "unit_price": items[0]["unit_price"],
            "work_dates": ", ".join(sorted({item["work_date"] for item in items})),
            "batch_ids": ", ".join(str(value) for value in sorted({item["batch_id"] for item in items})),
            "kitchens": ", ".join(sorted({item["kitchen"] for item in items if item["kitchen"]})),
            **summary,
        })
    shortages = [row for row in rows if row["pending_qty"] > EPSILON]
    return {
        "date_from": safe_from,
        "date_to": safe_to,
        "contractor": contractor_text,
        "rows": rows,
        "shortages": shortages,
        "shortage_count": len(shortages),
        "detail_count": len(detail_rows),
        **_totals(detail_rows),
        "contractors": _contractor_summaries(detail_rows),
        "stock_basis": "OPENING + invoice_inventory_ledger − active local holds",
    }


def _sheet_name(value: str, used: set[str]) -> str:
    base = re.sub(r"[\\/*?:\[\]\x00-\x1f]", "-", value).strip(" '") or "KHONG_MA"
    base = base[:31]
    candidate = base
    index = 2
    while candidate.casefold() in used:
        suffix = f"-{index}"
        candidate = base[: 31 - len(suffix)] + suffix
        index += 1
    used.add(candidate.casefold())
    return candidate


def _excel_text(value: Any) -> str:
    text = str(value or "")
    if text.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


def _write_shortage_sheet(sheet, payload: dict[str, Any], rows: list[dict[str, Any]], title: str) -> None:
    headers = (
        "Nhà thầu", "Mã hàng TĐP", "Tên hàng", "ĐVT", "Ngày nguồn", "Bếp", "Phiên",
        "Tổng cần", "Đã dự thảo", "Đã phát hành", "Có thể lập", "Còn thiếu",
        "Đơn giá bán", "Giá trị còn thiếu",
    )
    sheet.append([title])
    sheet.append([
        f"Kỳ {_display_date(payload['date_from'])} đến {_display_date(payload['date_to'])} "
        "· tồn kho theo hóa đơn"
    ])
    sheet.append(list(headers))
    for row in rows:
        sheet.append([
            _excel_text(row["contractor"]), _excel_text(row["product_code"]),
            _excel_text(row["product_name"]), _excel_text(row["unit"]),
            _excel_text(row["work_dates"]), _excel_text(row["kitchens"]),
            _excel_text(row["batch_ids"]), row["demand_qty"],
            row["drafted_qty"], row["issued_qty"], row["invoiceable_qty"], row["pending_qty"],
            row["unit_price"], row["pending_value"],
        ])
    sheet.freeze_panes = "A4"
    sheet.auto_filter.ref = f"A3:N{max(sheet.max_row, 3)}"
    sheet.merge_cells("A1:N1")
    sheet.merge_cells("A2:N2")
    sheet["A1"].font = Font(name="Times New Roman", size=14, bold=True)
    sheet["A1"].alignment = Alignment(horizontal="center")
    sheet["A2"].font = Font(name="Times New Roman", size=11, italic=True)
    sheet["A2"].alignment = Alignment(horizontal="center")
    for cell in sheet[3]:
        cell.font = Font(name="Times New Roman", size=11, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=4, max_col=len(headers)):
        for cell in row:
            cell.font = Font(name="Times New Roman", size=11)
    widths = (16, 17, 28, 9, 24, 18, 12, 13, 13, 13, 13, 13, 15, 19)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    for row in sheet.iter_rows(min_row=4, min_col=8, max_col=14):
        for cell in row:
            cell.number_format = '#,##0.###' if cell.column <= 12 else '#,##0'
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.orientation = sheet.ORIENTATION_LANDSCAPE
    sheet.page_setup.scale = None
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "$1:$3"
    sheet.print_area = f"A1:N{max(sheet.max_row, 3)}"
    sheet.page_margins = PageMargins(
        left=0.2, right=0.2, top=0.4, bottom=0.4, header=0.15, footer=0.2,
    )
    sheet.oddFooter.center.text = "Trang &P / &N"
    sheet.sheet_view.showGridLines = False


def shortage_workbook_bytes(payload: dict[str, Any]) -> bytes:
    workbook = Workbook()
    try:
        used: set[str] = set()
        summary = workbook.active
        summary.title = _sheet_name("TỔNG HỢP THIẾU", used)
        shortage_rows = payload["shortages"]
        _write_shortage_sheet(summary, payload, shortage_rows, "DANH SÁCH CÒN THIẾU HÓA ĐƠN ĐẦU VÀO")
        contractors = sorted({row["contractor"] for row in shortage_rows})
        for contractor in contractors:
            sheet = workbook.create_sheet(_sheet_name(contractor, used))
            rows = [row for row in shortage_rows if row["contractor"] == contractor]
            _write_shortage_sheet(sheet, payload, rows, f"NHÀ THẦU {contractor} · PHẦN CÒN THIẾU")
        return safe_workbook_bytes(workbook)
    finally:
        workbook.close()


__all__ = [
    "OutgoingReadinessError",
    "allocation_by_order",
    "batch_readiness_payload",
    "canonical_available_stock",
    "period_shortage_payload",
    "shortage_workbook_bytes",
    "validate_demand_orders",
    "validate_draft_export_stock",
    "validate_issued_draft_stock",
]
