import hashlib
import io
import shutil
import tempfile
import unittest
import zipfile
from pathlib import Path

from openpyxl import load_workbook

try:
    from . import server
    from .invoice_tax_export import (
        INVOICE_HEADERS,
        InvoiceTaxExportError,
        TEMPLATE_SPECS,
        build_invoice_workbook,
        export_invoice_drafts_zip,
    )
except ImportError:  # pragma: no cover - direct invocation
    import server
    from invoice_tax_export import (
        INVOICE_HEADERS,
        InvoiceTaxExportError,
        TEMPLATE_SPECS,
        build_invoice_workbook,
        export_invoice_drafts_zip,
    )


PROJECT_ROOT = Path(__file__).resolve().parent.parent
TEMPLATE_DIR = PROJECT_ROOT / "bosung.30.8.26"


class InvoiceTaxGoldenTests(unittest.TestCase):
    def test_all_26_golden_rows_regenerate_with_exact_tax_semantics(self):
        audited = 0
        for template_kind, spec in TEMPLATE_SPECS.items():
            source_path = TEMPLATE_DIR / spec["filename"]
            self.assertEqual(hashlib.sha256(source_path.read_bytes()).hexdigest().upper(), spec["sha256"])
            source = load_workbook(source_path, data_only=True, keep_links=False)
            source_style = load_workbook(source_path, data_only=False, keep_links=False)
            try:
                sheet = source.active
                style_sheet = source_style.active
                source_rows = [
                    list(values) for values in sheet.iter_rows(min_row=2, max_col=13, values_only=True)
                    if any(value not in (None, "") for value in values)
                ]
                rows = []
                for index, values in enumerate(source_rows, start=1):
                    rows.append({
                        "line_id": audited + index,
                        "contractor": "NT-GOLDEN",
                        "product_code": f"GOLDEN-{audited + index:03d}",
                        "product_name": values[1],
                        "unit": values[2],
                        "qty": values[3],
                        "unit_price": values[4] or 0,
                        "amount": values[5] or 0,
                        "invoice_nature": str(values[12]),
                    })
                vat_percent = float(source_rows[0][9])
                payload, actual_kind = build_invoice_workbook(
                    rows, vat_percent=vat_percent, template_dir=TEMPLATE_DIR,
                )
                self.assertEqual(actual_kind, template_kind)
                output = load_workbook(io.BytesIO(payload), data_only=False, keep_links=False)
                values_output = load_workbook(io.BytesIO(payload), data_only=True, keep_links=False)
                try:
                    self.assertEqual(output.sheetnames, [spec["sheet"]])
                    self.assertEqual(tuple(output.active.cell(1, col).value for col in range(1, 14)), INVOICE_HEADERS)
                    for column in range(1, 14):
                        expected_cell = style_sheet.cell(1, column)
                        actual_cell = output.active.cell(1, column)
                        # Customer's updated print style: keep template geometry,
                        # use regular black text, white fill and fine borders.
                        self.assertEqual(actual_cell.font.name, expected_cell.font.name)
                        self.assertEqual(actual_cell.font.sz, expected_cell.font.sz)
                        self.assertFalse(actual_cell.font.bold)
                        self.assertEqual(actual_cell.font.color.rgb, '00000000')
                        self.assertIsNone(actual_cell.fill.patternType)
                        for side in ('left','right','top','bottom'):
                            self.assertEqual(getattr(actual_cell.border,side).style,
                                             'hair' if getattr(expected_cell.border,side).style else None)
                        self.assertEqual(str(actual_cell.alignment), str(expected_cell.alignment))
                        self.assertEqual(actual_cell.number_format, expected_cell.number_format)
                    self.assertEqual(output.active.max_row - 1, len(source_rows))
                    for row_number, expected in enumerate(source_rows, start=2):
                        actual = [values_output.active.cell(row_number, col).value for col in range(1, 14)]
                        self.assertEqual(actual[1:4], expected[1:4])
                        self.assertEqual(actual[9], expected[9])
                        self.assertEqual(str(actual[12]), str(expected[12]))
                        if str(expected[12]) == "2":
                            self.assertTrue(all(actual[col] is None for col in (4, 5, 6, 7, 8, 10, 11)))
                        else:
                            self.assertEqual(actual[5], round(float(actual[3]) * float(actual[4])))
                    for worksheet in output.worksheets:
                        for row in worksheet.iter_rows():
                            for cell in row:
                                self.assertNotEqual(cell.data_type, "f")
                                self.assertIsNone(cell.hyperlink)
                finally:
                    output.close()
                    values_output.close()
                audited += len(source_rows)
            finally:
                source.close()
                source_style.close()
        self.assertEqual(audited, 26)

    def test_zip_never_mixes_contractor_round_or_tax(self):
        lines = [
            {"line_id": 1, "contractor": "NT-A", "round_no": 1, "product_code": "A-K",
             "product_name": "KKKNT", "unit": "kg", "qty": 1, "unit_price": 100,
             "amount": 100, "tax": "KKKNT", "vat_percent": -2, "invoice_nature": "1"},
            {"line_id": 2, "contractor": "NT-A", "round_no": 1, "product_code": "A-8",
             "product_name": "VAT 8", "unit": "kg", "qty": 2, "unit_price": 100,
             "amount": 200, "tax": "8%", "vat_percent": 8, "invoice_nature": "1"},
            {"line_id": 3, "contractor": "NT-B", "round_no": 2, "product_code": "B-10",
             "product_name": "VAT 10", "unit": "thùng", "qty": 3, "unit_price": 100,
             "amount": 300, "tax": "10%", "vat_percent": 10, "invoice_nature": "1"},
            {"line_id": 4, "contractor": "NT-B", "round_no": 2, "product_code": "B-KM",
             "product_name": "Khuyến mại", "unit": "chai", "qty": 4, "unit_price": 0,
             "amount": 0, "tax": "10%", "vat_percent": 10, "invoice_nature": "2"},
        ]
        payload = export_invoice_drafts_zip(
            lines, work_date="2026-09-01", template_dir=TEMPLATE_DIR,
        )
        with zipfile.ZipFile(payload) as archive:
            files = sorted(name for name in archive.namelist() if name.endswith(".xlsx"))
            self.assertEqual(len(files), 3)
            self.assertTrue(any("NT-A_KKKNT_lan_1" in name for name in files))
            self.assertTrue(any("NT-A_VAT8_lan_1" in name for name in files))
            self.assertTrue(any("NT-B_VAT10_PROMOTION_lan_2" in name for name in files))
            seen_codes = set()
            for name in files:
                workbook = load_workbook(io.BytesIO(archive.read(name)), data_only=True, keep_links=False)
                try:
                    codes = {
                        workbook.active.cell(row, 1).value
                        for row in range(2, workbook.active.max_row + 1)
                    }
                    seen_codes.update(codes)
                    if "NT-A" in name:
                        self.assertTrue(all(str(code).startswith("A-") for code in codes))
                    if "KKKNT" in name:
                        self.assertEqual({workbook.active.cell(row, 10).value for row in range(2, workbook.active.max_row + 1)}, {-2})
                    if "VAT8" in name:
                        self.assertEqual({workbook.active.cell(row, 10).value for row in range(2, workbook.active.max_row + 1)}, {8})
                    if "PROMOTION" in name:
                        self.assertEqual({workbook.active.cell(row, 10).value for row in range(2, workbook.active.max_row + 1)}, {10})
                finally:
                    workbook.close()
            self.assertEqual(seen_codes, {"A-K", "A-8", "B-10", "B-KM"})

    def test_bad_amount_is_blocked_and_formula_like_text_is_literal(self):
        base = {
            "line_id": 1, "contractor": "NT", "product_code": "P1",
            "product_name": "=HYPERLINK(\"bad\")", "unit": "kg", "qty": 2,
            "unit_price": 100, "amount": 200, "invoice_nature": "1",
        }
        payload, _ = build_invoice_workbook([base], vat_percent=8, template_dir=TEMPLATE_DIR)
        workbook = load_workbook(io.BytesIO(payload), data_only=False, keep_links=False)
        try:
            self.assertNotEqual(workbook.active["B2"].data_type, "f")
            self.assertTrue(str(workbook.active["B2"].value).startswith("'="))
        finally:
            workbook.close()
        with self.assertRaises(InvoiceTaxExportError) as caught:
            build_invoice_workbook([{**base, "amount": 201}], vat_percent=8, template_dir=TEMPLATE_DIR)
        self.assertEqual(caught.exception.code, "invoice_draft_amount_mismatch")

        with tempfile.TemporaryDirectory() as temp_name:
            target = Path(temp_name) / "thue 8.xlsx"
            shutil.copy2(TEMPLATE_DIR / "thue 8.xlsx", target)
            target.write_bytes(target.read_bytes() + b"changed")
            with self.assertRaises(InvoiceTaxExportError) as changed:
                build_invoice_workbook([base], vat_percent=8, template_dir=Path(temp_name))
            self.assertEqual(changed.exception.code, "invoice_golden_changed")

    def test_portable_build_includes_four_runtime_goldens(self):
        script = (PROJECT_ROOT / "BUILD_PORTABLE.ps1").read_text(encoding="utf-8")
        self.assertIn("invoice_tax_export.py", script)
        self.assertIn("--hidden-import invoice_tax_export", script)
        self.assertEqual(script.count(";tax_templates"), 4)


class InvoiceTaxMultiRoundApiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp_dir = tempfile.TemporaryDirectory()
        cls.original_db_path = server.DB_PATH
        cls.original_data_dir = server.DATA_DIR
        cls.original_testing = server.app.config.get("TESTING", False)
        server.DATA_DIR = Path(cls.temp_dir.name) / "data"
        server.DATA_DIR.mkdir(parents=True, exist_ok=True)
        server.DB_PATH = server.DATA_DIR / "invoice_tax_export.sqlite3"
        server.app.config["TESTING"] = True
        server.init_database()
        cls.client = server.app.test_client()

    @classmethod
    def tearDownClass(cls):
        server.DB_PATH = cls.original_db_path
        server.DATA_DIR = cls.original_data_dir
        server.app.config["TESTING"] = cls.original_testing
        cls.temp_dir.cleanup()

    def setUp(self):
        with server.db() as conn:
            conn.execute("DELETE FROM invoice_inventory_ledger")
            conn.execute("DELETE FROM invoice_inventory_confirmations")
            conn.execute("DELETE FROM outgoing_invoice_lines")
            conn.execute("DELETE FROM outgoing_invoice_drafts")
            conn.execute("DELETE FROM inventory_transactions")
            conn.execute("DELETE FROM orders")
            conn.execute("DELETE FROM batches")
            conn.execute("DELETE FROM audit_log")
            conn.execute(
                "INSERT OR REPLACE INTO contractors(code,name,price_group,pricing_mode) "
                "VALUES('NT-A','Nhà thầu A','NT-A','group')"
            )
            conn.execute(
                """INSERT OR REPLACE INTO products(
                       code,name,unit,tax,supplier,buy_price,purchase_list,seller,cccd
                   ) VALUES('P-VAT8','Hàng VAT 8','kg','8%','NCC',10,0,'','')"""
            )

    @staticmethod
    def add_input_ledger(conn, qty, key):
        now = server.now_iso()
        source_id = sum(ord(char) for char in key)
        confirmation_id = conn.execute(
            """INSERT INTO invoice_inventory_confirmations(
                   confirmation_key,direction,source_invoice_table,source_invoice_id,
                   action,confirmed,note,created_at
               ) VALUES(?,'input','test_source',?,'post',1,'test',?)""",
            ("confirm-" + key, source_id, now),
        ).lastrowid
        conn.execute(
            """INSERT INTO invoice_inventory_ledger(
                   event_key,direction,event_type,source_invoice_table,source_invoice_id,
                   source_line_id,source_line_index,product_code,txn_date,qty_delta,unit_cost,
                   mapping_revision_id,confirmation_id,reverses_event_key,status,created_at
               ) VALUES(?,'input','POST','test_source',?,?,1,'P-VAT8','2026-09-02',?,10,
                        NULL,?,'','posted',?)""",
            (key, source_id, source_id, qty, confirmation_id, now),
        )

    def test_each_export_contains_only_the_current_reserved_round(self):
        with server.db() as conn:
            now = server.now_iso()
            batch_id = conn.execute(
                """INSERT INTO batches(work_date,source_name,status,created_at,approved_at)
                   VALUES('2026-09-01','round.xlsx','approved',?,?)""",
                (now, now),
            ).lastrowid
            conn.execute(
                """INSERT INTO orders(
                       batch_id,work_date,contractor,kitchen,product_code,product_name,qty,
                       actual_received,actual_delivered,unit,supplier,buy_price,sell_price,tax,
                       purchase_list,errors,warnings,updated_at
                   ) VALUES(?,'2026-09-01','NT-A','BEP','P-VAT8','Hàng VAT 8',10,
                            10,10,'kg','NCC',10,100,'8%',0,'[]','[]',?)""",
                (batch_id, now),
            )
            conn.execute(
                """INSERT INTO inventory_transactions(
                       txn_date,product_code,qty_in,qty_out,unit_cost,source_type,source_id,
                       source_line,status,note,created_at,updated_at
                   ) VALUES('2026-08-01','P-VAT8',7,0,10,'OPENING','OPEN-ROUND',
                            'P-VAT8','posted','opening',?,?)""",
                (now, now),
            )

        first = self.client.post(f"/api/outgoing-invoices/draft/{batch_id}")
        self.assertEqual(first.status_code, 200, first.get_data(as_text=True))
        first_export = self.client.get(f"/api/export/invoices/{batch_id}")
        self.assertEqual(first_export.status_code, 200)
        with zipfile.ZipFile(io.BytesIO(first_export.data)) as archive:
            files = [name for name in archive.namelist() if name.endswith(".xlsx")]
            self.assertEqual(len(files), 1)
            self.assertIn("lan_1", files[0])
            workbook = load_workbook(io.BytesIO(archive.read(files[0])), data_only=True)
            try:
                self.assertEqual(workbook.active["D2"].value, 7)
            finally:
                workbook.close()

        with server.db() as conn:
            first_id = conn.execute(
                "SELECT id FROM outgoing_invoice_drafts WHERE batch_id=?", (batch_id,),
            ).fetchone()["id"]
            conn.execute(
                "UPDATE outgoing_invoice_drafts SET status='issued',issued_at=? WHERE id=?",
                (server.now_iso(), first_id),
            )
            conn.execute(
                """UPDATE inventory_transactions SET status='posted'
                    WHERE source_type='OUTGOING_DRAFT' AND source_id=?""",
                (str(first_id),),
            )
            self.add_input_ledger(conn, 3, "INPUT-ROUND-2")

        second = self.client.post(f"/api/outgoing-invoices/draft/{batch_id}")
        self.assertEqual(second.status_code, 200, second.get_data(as_text=True))
        self.assertEqual(second.get_json()["drafts"][0]["round_no"], 2)
        second_export = self.client.get(f"/api/export/invoices/{batch_id}")
        self.assertEqual(second_export.status_code, 200)
        with zipfile.ZipFile(io.BytesIO(second_export.data)) as archive:
            files = [name for name in archive.namelist() if name.endswith(".xlsx")]
            self.assertEqual(len(files), 1)
            self.assertIn("lan_2", files[0])
            workbook = load_workbook(io.BytesIO(archive.read(files[0])), data_only=True)
            try:
                self.assertEqual(workbook.active["D2"].value, 3)
                self.assertEqual(workbook.active["J2"].value, 8)
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
