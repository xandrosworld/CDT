"""Three-layer, privacy-safe QC for TDP Excel outputs.

The report records locations and categories of differences, never the compared
cell values. Visual comparison uses a trusted spreadsheet renderer when one is
available; on this Windows workstation that renderer is Microsoft Excel COM.
"""

from __future__ import annotations

import argparse
from copy import copy
import json
import math
import os
import shutil
import sys
import tempfile
from dataclasses import dataclass, field
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from PIL import Image, ImageChops


MAX_RECORDED_ISSUES = 300
SENSITIVE_SHEET_NAMES = {"cccd", "cmt", "cmnd"}


class WorkbookQCError(RuntimeError):
    """Raised when a requested QC operation cannot be performed safely."""


@dataclass
class WorkbookQCReport:
    golden: str
    candidate: str
    layers: dict[str, dict[str, Any]] = field(
        default_factory=lambda: {
            "data": {"ok": True, "issue_count": 0},
            "topology": {"ok": True, "issue_count": 0},
            "visual": {"ok": False, "issue_count": 0, "status": "not_run"},
        }
    )
    issues: list[dict[str, str]] = field(default_factory=list)

    def add(self, layer: str, kind: str, *, sheet: str = "", location: str = "") -> None:
        summary = self.layers[layer]
        summary["ok"] = False
        summary["issue_count"] = int(summary["issue_count"]) + 1
        if len(self.issues) < MAX_RECORDED_ISSUES:
            issue = {"layer": layer, "kind": kind}
            if sheet:
                issue["sheet"] = sheet
            if location:
                issue["location"] = location
            self.issues.append(issue)

    @property
    def ok(self) -> bool:
        return all(bool(self.layers[name]["ok"]) for name in ("data", "topology", "visual"))

    def as_dict(self) -> dict[str, Any]:
        return {
            "ok": self.ok,
            "golden": self.golden,
            "candidate": self.candidate,
            "layers": self.layers,
            "issues": self.issues,
            "privacy": "locations_and_categories_only_no_cell_values",
        }


def _normalized_decimal(value: int | float | Decimal) -> str:
    try:
        decimal = Decimal(str(value))
    except InvalidOperation:
        return str(value)
    if not decimal.is_finite():
        return str(decimal)
    normalized = decimal.normalize()
    return format(normalized, "f")


def normalize_value(value: Any) -> tuple[str, str]:
    if value is None:
        return ("empty", "")
    if isinstance(value, bool):
        return ("boolean", "1" if value else "0")
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        return ("number", _normalized_decimal(value))
    if isinstance(value, (datetime, date, time)):
        return ("datetime", value.isoformat())
    if isinstance(value, str):
        text = value.replace("\r\n", "\n").replace("\r", "\n").strip()
        return ("formula" if text.startswith("=") else "text", text)
    return (type(value).__name__, str(value))


def _color_signature(color: Any) -> tuple[Any, ...]:
    if color is None:
        return ()
    return (
        getattr(color, "type", None),
        getattr(color, "rgb", None),
        getattr(color, "indexed", None),
        getattr(color, "theme", None),
        getattr(color, "tint", None),
    )


def _side_signature(side: Any) -> tuple[Any, ...]:
    return (getattr(side, "style", None), _color_signature(getattr(side, "color", None)))


def style_signature(cell: Any) -> tuple[Any, ...]:
    font = cell.font
    fill = cell.fill
    border = cell.border
    alignment = cell.alignment
    return (
        font.name,
        font.sz,
        font.bold,
        font.italic,
        font.underline,
        _color_signature(font.color),
        fill.fill_type,
        _color_signature(fill.fgColor),
        _color_signature(fill.bgColor),
        _side_signature(border.left),
        _side_signature(border.right),
        _side_signature(border.top),
        _side_signature(border.bottom),
        alignment.horizontal,
        alignment.vertical,
        alignment.wrap_text,
        alignment.text_rotation,
        cell.number_format,
    )


def _worksheet_topology(sheet: Any) -> dict[str, Any]:
    return {
        "state": sheet.sheet_state,
        "merged_ranges": tuple(sorted(str(item) for item in sheet.merged_cells.ranges)),
        "print_area": str(sheet.print_area or ""),
        "print_title_rows": str(sheet.print_title_rows or ""),
        "print_title_cols": str(sheet.print_title_cols or ""),
        "orientation": sheet.page_setup.orientation,
        "paper_size": sheet.page_setup.paperSize,
        "fit_to_width": sheet.page_setup.fitToWidth,
        "fit_to_height": sheet.page_setup.fitToHeight,
        "freeze_panes": str(sheet.freeze_panes or ""),
        "show_grid_lines": sheet.sheet_view.showGridLines,
        "row_dimensions": tuple(
            sorted(
                (
                    index,
                    dimension.height,
                    bool(dimension.hidden),
                    dimension.outlineLevel,
                )
                for index, dimension in sheet.row_dimensions.items()
                if dimension.height is not None or dimension.hidden or dimension.outlineLevel
            )
        ),
        "column_dimensions": tuple(
            sorted(
                (
                    index,
                    dimension.width,
                    bool(dimension.hidden),
                    dimension.outlineLevel,
                )
                for index, dimension in sheet.column_dimensions.items()
                if dimension.width is not None or dimension.hidden or dimension.outlineLevel
            )
        ),
    }


def compare_workbooks(golden_path: Path, candidate_path: Path) -> WorkbookQCReport:
    golden_path = Path(golden_path).resolve()
    candidate_path = Path(candidate_path).resolve()
    report = WorkbookQCReport(str(golden_path), str(candidate_path))

    golden_book = load_workbook(golden_path, data_only=False, keep_links=False)
    candidate_book = load_workbook(candidate_path, data_only=False, keep_links=False)
    try:
        if golden_book.sheetnames != candidate_book.sheetnames:
            report.add("topology", "sheet_order_or_names")

        common_sheets = [name for name in golden_book.sheetnames if name in candidate_book.sheetnames]
        for sheet_name in common_sheets:
            golden_sheet = golden_book[sheet_name]
            candidate_sheet = candidate_book[sheet_name]

            golden_topology = _worksheet_topology(golden_sheet)
            candidate_topology = _worksheet_topology(candidate_sheet)
            for key in golden_topology:
                if golden_topology[key] != candidate_topology[key]:
                    report.add("topology", key, sheet=sheet_name)

            coordinates = set(golden_sheet._cells) | set(candidate_sheet._cells)
            for row, column in sorted(coordinates):
                golden_cell = golden_sheet.cell(row, column)
                candidate_cell = candidate_sheet.cell(row, column)
                location = golden_cell.coordinate
                if normalize_value(golden_cell.value) != normalize_value(candidate_cell.value):
                    report.add("data", "cell_value", sheet=sheet_name, location=location)
                if style_signature(golden_cell) != style_signature(candidate_cell):
                    report.add("topology", "cell_style", sheet=sheet_name, location=location)
    finally:
        golden_book.close()
        candidate_book.close()
    return report


def detect_renderer() -> dict[str, Any]:
    if os.name == "nt":
        excel_candidates = (
            Path(r"C:\Program Files\Microsoft Office\root\Office16\EXCEL.EXE"),
            Path(r"C:\Program Files (x86)\Microsoft Office\root\Office16\EXCEL.EXE"),
        )
        try:
            import win32com.client  # noqa: F401
        except ImportError:
            pass
        else:
            executable = next((path for path in excel_candidates if path.is_file()), None)
            if executable is not None:
                return {
                    "available": True,
                    "trusted": True,
                    "name": "Microsoft Excel COM",
                    "executable": str(executable),
                }

    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice:
        return {
            "available": True,
            "trusted": False,
            "name": "LibreOffice (not validated for this workbook set)",
            "executable": soffice,
        }
    return {"available": False, "trusted": False, "name": "none", "executable": ""}


def _assert_sheet_safe(sheet_name: str) -> None:
    normalized = "".join(character for character in sheet_name.casefold() if character.isalnum())
    if normalized in SENSITIVE_SHEET_NAMES:
        raise WorkbookQCError(f"Không render sheet nhạy cảm: {sheet_name}")


def render_sheet_with_excel(workbook_path: Path, sheet_name: str, pdf_path: Path) -> None:
    _assert_sheet_safe(sheet_name)
    renderer = detect_renderer()
    if not renderer["available"] or not renderer["trusted"] or renderer["name"] != "Microsoft Excel COM":
        raise WorkbookQCError("Không có renderer Excel đã kiểm chứng trên máy")

    import win32com.client

    workbook_path = Path(workbook_path).resolve()
    pdf_path = Path(pdf_path).resolve()
    pdf_path.parent.mkdir(parents=True, exist_ok=True)
    excel = None
    workbook = None
    try:
        excel = win32com.client.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        excel.AskToUpdateLinks = False
        excel.EnableEvents = False
        excel.AutomationSecurity = 3
        workbook = excel.Workbooks.Open(
            Filename=str(workbook_path),
            UpdateLinks=0,
            ReadOnly=True,
            AddToMru=False,
            IgnoreReadOnlyRecommended=True,
            Notify=False,
        )
        worksheet = workbook.Worksheets(sheet_name)
        worksheet.ExportAsFixedFormat(
            Type=0,
            Filename=str(pdf_path),
            Quality=0,
            IncludeDocProperties=False,
            IgnorePrintAreas=False,
            OpenAfterPublish=False,
        )
    except Exception as error:
        raise WorkbookQCError(f"Excel không render được sheet {sheet_name}: {type(error).__name__}") from error
    finally:
        if workbook is not None:
            workbook.Close(SaveChanges=False)
        if excel is not None:
            excel.Quit()


def pdf_to_pngs(pdf_path: Path, output_dir: Path, prefix: str) -> list[Path]:
    try:
        import fitz
    except ImportError as error:  # pragma: no cover - environment-specific
        raise WorkbookQCError("Thiếu PyMuPDF để raster hóa PDF") from error

    output_dir.mkdir(parents=True, exist_ok=True)
    pages: list[Path] = []
    document = fitz.open(pdf_path)
    try:
        matrix = fitz.Matrix(2, 2)
        for index, page in enumerate(document):
            target = output_dir / f"{prefix}-{index + 1:03d}.png"
            page.get_pixmap(matrix=matrix, alpha=False).save(target)
            pages.append(target)
    finally:
        document.close()
    return pages


def image_difference_ratio(first_path: Path, second_path: Path, noise_threshold: int = 2) -> float:
    with Image.open(first_path) as first_source, Image.open(second_path) as second_source:
        first = first_source.convert("RGB")
        second = second_source.convert("RGB")
        if first.size != second.size:
            return 1.0
        difference = ImageChops.difference(first, second).convert("L")
        pixels = (
            difference.get_flattened_data()
            if hasattr(difference, "get_flattened_data")
            else difference.getdata()
        )
        changed = sum(1 for pixel in pixels if pixel > noise_threshold)
        return changed / float(first.width * first.height)


def render_and_compare(
    report: WorkbookQCReport,
    *,
    sheet_name: str,
    output_dir: Path,
    tolerance: float = 0.0,
) -> None:
    renderer = detect_renderer()
    report.layers["visual"].update({"renderer": renderer["name"], "trusted": renderer["trusted"]})
    if not renderer["available"] or not renderer["trusted"]:
        report.layers["visual"].update({"ok": False, "status": "unavailable"})
        return

    output_dir = Path(output_dir).resolve()
    golden_pdf = output_dir / "golden.pdf"
    candidate_pdf = output_dir / "candidate.pdf"
    render_sheet_with_excel(Path(report.golden), sheet_name, golden_pdf)
    render_sheet_with_excel(Path(report.candidate), sheet_name, candidate_pdf)
    golden_pages = pdf_to_pngs(golden_pdf, output_dir / "golden", "page")
    candidate_pages = pdf_to_pngs(candidate_pdf, output_dir / "candidate", "page")

    report.layers["visual"].update({
        "ok": True,
        "status": "compared",
        "golden_pages": len(golden_pages),
        "candidate_pages": len(candidate_pages),
        "tolerance": tolerance,
    })
    if len(golden_pages) != len(candidate_pages):
        report.add("visual", "page_count", sheet=sheet_name)

    ratios: list[float] = []
    for index, (golden_page, candidate_page) in enumerate(zip(golden_pages, candidate_pages), start=1):
        ratio = image_difference_ratio(golden_page, candidate_page)
        ratios.append(ratio)
        if not math.isfinite(ratio) or ratio > tolerance:
            report.add("visual", "pixel_difference", sheet=sheet_name, location=f"page:{index}")
    report.layers["visual"]["page_difference_ratios"] = ratios


def _write_self_test_workbooks(directory: Path) -> tuple[Path, Path]:
    golden = directory / "golden.xlsx"
    candidate = directory / "candidate.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "QC"
    sheet.append(["Mã", "Số lượng"])
    sheet.append(["A", 1])
    sheet.append(["TỔNG", 1])
    sheet.print_area = "A1:B3"
    sheet.page_setup.paperSize = "9"
    sheet.page_setup.orientation = "portrait"
    workbook.save(golden)
    sheet["B3"] = 2
    changed_font = copy(sheet["A1"].font)
    changed_font.bold = True
    sheet["A1"].font = changed_font
    workbook.save(candidate)
    workbook.close()
    return golden, candidate


def renderer_self_test(output_dir: Path) -> dict[str, Any]:
    output_dir = Path(output_dir).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    golden, candidate = _write_self_test_workbooks(output_dir)

    identical_report = compare_workbooks(golden, golden)
    render_and_compare(
        identical_report,
        sheet_name="QC",
        output_dir=output_dir / "render-identical",
        tolerance=0.0,
    )
    if not identical_report.ok:
        raise WorkbookQCError("Renderer self-test tạo false positive với hai workbook giống nhau")

    report = compare_workbooks(golden, candidate)
    render_and_compare(report, sheet_name="QC", output_dir=output_dir / "render", tolerance=0.0)
    result = report.as_dict()
    if result["layers"]["visual"]["status"] != "compared":
        raise WorkbookQCError("Renderer self-test chưa thực hiện được visual comparison")
    if result["layers"]["visual"]["issue_count"] < 1:
        raise WorkbookQCError("Renderer self-test không bắt được fixture cố ý sai")
    result["renderer_self_test_identical_passed"] = True
    return result


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="TDP golden workbook three-layer QC")
    parser.add_argument("--golden", type=Path)
    parser.add_argument("--candidate", type=Path)
    parser.add_argument("--sheet")
    parser.add_argument("--output-dir", type=Path)
    parser.add_argument("--self-test-renderer", action="store_true")
    arguments = parser.parse_args(argv)

    try:
        if arguments.self_test_renderer:
            target = arguments.output_dir or Path(tempfile.mkdtemp(prefix="tdp_renderer_selftest_"))
            result = renderer_self_test(target)
            self_test_passed = True
        else:
            self_test_passed = False
            if arguments.golden is None or arguments.candidate is None:
                parser.error("--golden và --candidate là bắt buộc")
            result_report = compare_workbooks(arguments.golden, arguments.candidate)
            if arguments.sheet:
                target = arguments.output_dir or Path(tempfile.mkdtemp(prefix="tdp_golden_qc_"))
                render_and_compare(result_report, sheet_name=arguments.sheet, output_dir=target)
            result = result_report.as_dict()
        print(json.dumps(result, ensure_ascii=False, indent=2))
        if self_test_passed:
            return 0
        return 0 if result["layers"]["data"]["ok"] and result["layers"]["topology"]["ok"] else 2
    except (OSError, WorkbookQCError, ValueError) as error:
        print(f"Golden workbook QC FAILED: {error}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
