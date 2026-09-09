"""Separate month-end stock export in the customer's 'Tồn Trong kỳ' layout."""
from collections import defaultdict
from decimal import Decimal
from pathlib import Path

from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.views import Selection

try:
    from .inventory_customer_report import literal, number, tax_value
    from .template_workbook import clone_template_workbook, copy_row_layout, assert_workbook_safe
except ImportError:
    from inventory_customer_report import literal, number, tax_value
    from template_workbook import clone_template_workbook, copy_row_layout, assert_workbook_safe

TEMPLATE_PATH = Path(__file__).resolve().parent / 'templates' / 'inventory_closing_template.xlsx'
TEMPLATE_SHA256 = '3BDD4FE66A7B98CD2CA94CF2CB9197D2CC021C40F8AE98488451E06E6CD990D8'


def customer_closing_workbook(model, sales):
    if model.get('valuation_method') != 'monthly_weighted_average':
        raise ValueError('Tồn trong kỳ cần dữ liệu bình quân tháng')
    outgoing = defaultdict(list)
    for line in sales['lines']:
        if line.get('product_code') and line.get('mapping_status') == 'mapped' and not line.get('identity_warning'):
            outgoing[line['product_code']].append(line)
    wb = clone_template_workbook(TEMPLATE_PATH, sheet_names=['Ton 7 (2)'], expected_sha256=TEMPLATE_SHA256).workbook
    ws = wb.active
    ws.title = 'Tồn trong kỳ'
    for row, height in ((1, 26), (2, 24), (3, 22)):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
        ws.cell(row, 1).alignment = Alignment(horizontal='left', vertical='center')
        ws.row_dimensions[row].height = height
    year, month = model['date_from'][:7].split('-')
    ws.merge_cells('A6:I6')
    ws['A6'] = f'THÁNG {int(month)}/{year} · Tồn cuối ngày {model["date_to"][8:10]}/{month}/{year} · Đơn giá bình quân tháng theo từng mặt hàng'
    ws['A6'].alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[6].height = 30
    ws.row_dimensions[4].height = 25
    ws.row_dimensions[5].height = 28
    quantities = defaultdict(Decimal)
    total = Decimal(0)
    issues = []
    for index, item in enumerate(model['items'], 1):
        row = index + 6
        copy_row_layout(ws, 7, row, include_values=False)
        ws.row_dimensions[row].height = max(28, 14 * (1 + max(len(str(item.get('product_name') or '')), len(str(item.get('invoice_name') or ''))) // 24))
        values = [index, item['product_code'], item['product_name'], item['invoice_name'],
                  tax_value(item, outgoing[item['product_code']]), item['unit'],
                  item['closing_qty'], item['average_unit_cost'], item['closing_value']]
        for col, value in enumerate(values, 1):
            literal(ws, row, col, value)
        rate = ws.cell(row, 5).value
        ws.cell(row, 5).number_format = '0%' if not isinstance(rate, (int, float)) or abs(rate * 100 - round(rate * 100)) < 1e-8 else '0.00%'
        for col in (3, 4):
            ws.cell(row, col).alignment = Alignment(wrap_text=True, vertical='center')
        ws.cell(row, 7).number_format = '#,##0.######'
        for col in (8, 9):
            ws.cell(row, col).number_format = '#,##0.##'
        quantities[item['unit']] += number(item['closing_qty'])
        total += number(item['closing_value'])
        if item['valuation_status'] != 'ok':
            issues.append([item['product_code'], item['valuation_status']])
    last = max(7, 6 + len(model['items']))
    footer = last + 1
    copy_row_layout(ws, 7, footer, include_values=False)
    ws.merge_cells(start_row=footer, start_column=1, end_row=footer, end_column=6)
    ws.cell(footer, 1, 'TỔNG')
    quantities = {unit: qty for unit, qty in quantities.items() if qty}
    ws.cell(footer, 7, float(next(iter(quantities.values()))) if len(quantities) == 1 else f'{len(quantities)} ĐVT · xem Tổng ĐVT' if quantities else 0)
    ws.cell(footer, 9, float(total)).number_format = '#,##0.##'
    ws.row_dimensions[footer].height = 32
    for cell in ws[footer]:
        cell.font = Font(name='Times New Roman', size=11, bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    if len(quantities) > 1:
        units = wb.create_sheet('Tổng ĐVT')
        units.append(['ĐVT', 'Tồn cuối kỳ'])
        for unit, qty in sorted(quantities.items()):
            units.append([unit, float(qty)])
    if issues:
        review = wb.create_sheet('Đối chiếu')
        review.append(['Mã hàng', 'Cần đối chiếu'])
        for issue in issues:
            review.append(issue)
    for sheet in wb:
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
        if sheet != ws:
            sheet.column_dimensions['A'].width = 24
            sheet.column_dimensions['B'].width = 55
    ws.sheet_view.pane = None
    ws.sheet_view.selection = [Selection(activeCell='A1', sqref='A1')]
    ws.freeze_panes = 'G7'
    ws.auto_filter.ref = f'A5:I{last}'
    ws.print_area = f'A1:I{footer}'
    ws.print_title_rows = '1:6'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.scale = None
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.properties.title = f'Tồn trong kỳ {month}/{year}'
    wb.properties.description = 'Tồn cuối kỳ và đơn giá bình quân tháng lấy cùng dữ liệu báo cáo NXT.'
    assert_workbook_safe(wb)
    return wb
