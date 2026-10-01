import io
import json
import unittest
from openpyxl import load_workbook
from . import server
from . import test_bk_supplement as fixture


class SavedHistoryTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):fixture.SupplementTests.setUpClass.__func__(cls)
    @classmethod
    def tearDownClass(cls):fixture.SupplementTests.tearDownClass.__func__(cls)
    body=fixture.SupplementTests.body
    preview=fixture.SupplementTests.preview
    confirm=fixture.SupplementTests.confirm

    def setUp(self):
        fixture.SupplementTests.setUp(self)
        with server.db() as conn:
            from .purchase_seller_revision import init_schema
            init_schema(conn)
            conn.execute('DELETE FROM purchase_seller_revisions')
            conn.execute('DELETE FROM batch_bk_approvals')
        result=self.confirm(self.preview())
        self.assertEqual(200,result.status_code,result.get_json())

    def history(self,start='2026-08-01',end='2026-08-31'):
        return self.client.get('/api/bk-import/history',query_string={'from':start,'to':end})

    def print_body(self,**extra):
        return dict(kind='saved-purchases',document_ids=[self.history().get_json()['items'][0]['id']],
                    **{'from':'2026-08-01','to':'2026-08-31'},**extra)

    def database(self):
        with server.db() as c:return '\n'.join(c.iterdump())

    def test_reprint_saved_august_without_negative_stock_or_reposting(self):
        listing=self.history().get_json()
        self.assertEqual(1,len(listing['items']))
        self.assertEqual(200,listing['items'][0]['amount_total'])
        with server.db() as c:
            # Existing later-period opening must not block a historical print.
            c.execute("INSERT INTO inventory_transactions(txn_date,product_code,qty_in,qty_out,unit_cost,source_type,source_id,source_line,status,created_at,updated_at) VALUES('2026-09-01','BK-P1',2,0,100,'OPENING','2026-09','BK-P1','posted','test','test')")
        before=self.database()
        for receipts in (False,True):
            result=self.client.post('/api/documents/preview',json=self.print_body(receipts=receipts))
            self.assertEqual(200,result.status_code,result.get_json())
            data=result.get_json();self.assertEqual(2 if receipts else 1,len(data['sheets']))
            self.assertEqual('saved',data['purchase_source']['source'])
            self.assertIn('đã ghi kho',data['purchase_source']['title'])
            output=self.client.get('/api/documents/'+data['token']+'/excel')
            self.assertEqual(200,output.status_code)
            book=load_workbook(io.BytesIO(output.data))
            self.assertIn('bảng kê tổng',book.sheetnames)
            text=str(list(book['bảng kê tổng'].values))
            self.assertIn('Từ ngày 01/08/2026 đến ngày 01/08/2026',text)
            self.assertNotIn('31/08/2026',text)
            book.close()
        self.assertEqual(before,self.database())

    def test_period_and_reversed_document_are_not_silently_reprinted(self):
        body=self.print_body()
        self.assertEqual([],self.history('2026-09-01','2026-09-30').get_json()['items'])
        self.assertEqual(400,self.history('bad','2026-08-31').status_code)
        with server.db() as c:c.execute("UPDATE bk_import_documents SET status='reversed'")
        self.assertEqual([],self.history().get_json()['items'])
        result=self.client.post('/api/documents/preview',json=body)
        self.assertEqual(422,result.status_code);self.assertIn('hoàn tác',result.get_json()['error'])

    def test_print_uses_saved_product_name_not_changed_catalog(self):
        with server.db() as c:
            c.execute("UPDATE products SET name='Changed catalogue',unit='other',buy_price=999")
        result=self.client.post('/api/documents/preview',json=self.print_body())
        self.assertEqual(200,result.status_code,result.get_json())
        self.assertIn('Hàng BK',str(result.get_json()['sheets']))
        self.assertNotIn('Changed catalogue',str(result.get_json()['sheets']))

    def test_receipt_limit_names_seller_and_day_without_blocking_saved_summary(self):
        with server.db() as c:
            c.execute('UPDATE bk_import_lines SET amount=6000000,unit_cost=3000000')
        before=self.database()
        summary=self.client.post('/api/documents/preview',json=self.print_body())
        self.assertEqual(200,summary.status_code,summary.get_json())
        receipt=self.client.post('/api/documents/preview',json=self.print_body(receipts=True))
        self.assertEqual(422,receipt.status_code)
        self.assertIn('Seller Test, ngày 01/08/2026',receipt.get_json()['error'])
        self.assertIn('Bảng kê tổng vẫn in được',receipt.get_json()['error'])
        self.assertEqual(before,self.database())

    def test_saved_seller_revision_keeps_amount_and_is_printed(self):
        from .bk_history import workbook
        body=self.print_body();doc=body['document_ids'][0]
        revised=[dict(work_date='2026-08-01',seller='Seller Revised',cccd='222222222222',address='Revised address',
                      product_name='Hàng BK',unit='kg',quantity=2,buy_price=100,amount=200)]
        with server.db() as c:
            c.execute("INSERT INTO batch_bk_approvals VALUES(999,?,'test','test')",(doc,))
            c.execute("INSERT INTO purchase_seller_revisions VALUES(999,1,'test','[]',?,'test','test','test','test','test')",(json.dumps(revised),))
            book=workbook(c,body,vars(server))
            values=str([r for r in book.active.values]);book.close()
            self.assertIn('Seller Revised',values)
            self.assertNotIn('Seller Test',values)
            revised[0]['amount']=201
            c.execute('UPDATE purchase_seller_revisions SET rows_json=?',(json.dumps(revised),))
            with self.assertRaisesRegex(ValueError,'không khớp'):workbook(c,body,vars(server))
