import io
import unittest
import zipfile
import xml.etree.ElementTree as ET
from openpyxl import Workbook,load_workbook
from .invoice_tax_export import INVOICE_HEADERS
from .minvoice_import_workbook import build_import_package,compatible_archive,TEMPLATE


class MinvoiceImportPackageTests(unittest.TestCase):
    def rows(self):
        return [INVOICE_HEADERS,['HH-01','Rau & quả <tươi>','Kg',1.25,8000,10000,None,None,10000,8,800,10800,'1']]

    def values(self,data):
        wb=load_workbook(io.BytesIO(data),data_only=True)
        try:return list(wb.active.values)
        finally:wb.close()

    def test_accepted_package_preserves_values_styles_relationships_and_shared_strings(self):
        data=build_import_package(self.rows())
        self.assertEqual([tuple(r) for r in self.rows()],self.values(data))
        with zipfile.ZipFile(TEMPLATE) as source,zipfile.ZipFile(io.BytesIO(data)) as generated:
            self.assertEqual(source.namelist(),generated.namelist())
            for name in source.namelist():
                if name not in ('xl/sharedStrings.xml','xl/worksheets/sheet1.xml'):
                    self.assertEqual(source.read(name),generated.read(name),name)
            root=ET.fromstring(generated.read('xl/sharedStrings.xml'))
            self.assertEqual(len(root),int(root.attrib['uniqueCount']))
            self.assertNotIn(b'inlineStr',generated.read('xl/worksheets/sheet1.xml'))
        self.assertEqual(1,len(self.values(TEMPLATE.read_bytes())))

    def test_historical_archive_repairs_only_invoice_and_does_not_change_values(self):
        wb=Workbook()
        for row in self.rows():wb.active.append(row)
        original=io.BytesIO();wb.save(original);wb.close()
        archive=io.BytesIO()
        with zipfile.ZipFile(archive,'w') as z:
            z.writestr('Bang_ke_ATV_VAT8_949.xlsx',original.getvalue())
            z.writestr('HUONG_DAN.txt','keep me')
        repaired=compatible_archive(archive.getvalue())
        with zipfile.ZipFile(io.BytesIO(repaired)) as z:
            self.assertEqual(self.values(original.getvalue()),self.values(z.read('Bang_ke_ATV_VAT8_949.xlsx')))
            self.assertEqual(b'keep me',z.read('HUONG_DAN.txt'))
        self.assertEqual(repaired,compatible_archive(repaired))

    def test_empty_discounts_and_promotion_stay_blank(self):
        rows=self.rows()+[['KM','Hàng tặng','Gói',2,None,None,None,None,None,10,None,None,'2']]
        self.assertEqual([tuple(r) for r in rows],self.values(build_import_package(rows)))

    def test_sheet_name_escaped_and_no_extra_columns(self):
        data=build_import_package(self.rows(),sheet_name='VAT & 8')
        wb=load_workbook(io.BytesIO(data))
        self.assertEqual(['VAT & 8'],wb.sheetnames);wb.close()
        with self.assertRaises(ValueError):build_import_package([['bad']])
