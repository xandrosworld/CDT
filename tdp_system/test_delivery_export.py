import hashlib
import io
import sqlite3
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

try:
    from . import server as server_module
    from .delivery_export import (
        DeliveryExportError,
        PRICE_VISIBLE_CONTRACTOR,
        PRICE_VISIBLE_KITCHEN,
        build_delivery_workbook,
        delivery_prices_visible,
    )
    from .pdf_documents import build_pdf_bundle
    from .print_bundle import workbook_sections
    from .template_workbook import TemplateWorkbookError, safe_workbook_bytes
except ImportError:  # pragma: no cover - direct file invocation
    import server as server_module
    from delivery_export import (
        DeliveryExportError,
        PRICE_VISIBLE_CONTRACTOR,
        PRICE_VISIBLE_KITCHEN,
        build_delivery_workbook,
        delivery_prices_visible,
    )
    from pdf_documents import build_pdf_bundle
    from print_bundle import workbook_sections
    from template_workbook import TemplateWorkbookError, safe_workbook_bytes


ROOT = Path(__file__).resolve().parent.parent
GOLDEN = ROOT / "Em Thành.xlsx"


def delivery_payload():
    return [
        {
            "kitchen": PRICE_VISIBLE_KITCHEN,
            "contractor": PRICE_VISIBLE_CONTRACTOR,
            "recipient": "CÔNG TY CỔ PHẦN NHỰA VÀ CƠ KHÍ HẢI PHÒNG",
            "address": "Km 104 + 200 QL5, Phường Đông Hải, TP Hải Phòng, Việt Nam",
            "items": [
                {
                    "product_code": "RAU-01",
                    "product_name": "Bí xanh sơ chế (gọt vỏ)",
                    "quantity": 2,
                    "unit": "Kg",
                    "sell_price": 14_000,
                    "note": "",
                },
                {
                    "product_code": "RAU-02",
                    "product_name": "Cà rốt",
                    "quantity": 1.5,
                    "unit": "Kg",
                    "sell_price": 13_000,
                    "note": "Đã chọn",
                },
            ],
        },
        {
            "kitchen": "POT",
            "contractor": "HATRAN",
            "recipient": "BẾP POT",
            "address": "476 Phạm Văn Đồng, Anh Dũng, Dương Kinh, Hải Phòng",
            "items": [
                {
                    "product_code": "RAU-01",
                    "product_name": "Bí xanh",
                    "quantity": 3,
                    "unit": "Kg",
                    # This sentinel must never reach the non-price workbook.
                    "sell_price": 987_654_321,
                    "note": "",
                },
                {
                    "product_code": "RAU-02",
                    "product_name": "Cà rốt",
                    "quantity": 1,
                    "unit": "Kg",
                    "sell_price": 123_456_789,
                    "note": "Giao tại cổng",
                },
            ],
        },
    ]


class DeliveryExportTests(unittest.TestCase):
    def test_only_exact_source_backed_nhua_pair_can_show_prices(self):
        self.assertTrue(delivery_prices_visible("NHUAHP", "NHUAHAIPHONG"))
        self.assertFalse(delivery_prices_visible("BẾP NHỰA", "NHUAHAIPHONG"))
        self.assertFalse(delivery_prices_visible("NHUAHP", "ATV"))
        self.assertFalse(delivery_prices_visible("NHUAHP-2", "NHUAHAIPHONG"))

    def test_price_and_non_price_sheets_preserve_approved_form_without_leakage(self):
        source_hash = hashlib.sha256(GOLDEN.read_bytes()).hexdigest()
        workbook = build_delivery_workbook(
            delivery_payload(), work_date="2026-09-02", template_path=GOLDEN,
        )
        sections = workbook_sections("deliveries", workbook)
        try:
            self.assertEqual(workbook.sheetnames, ["NHUAHP", "POT"])
            priced = workbook["NHUAHP"]
            hidden = workbook["POT"]

            self.assertEqual(priced["C4"].value, "PHIẾU GIAO HÀNG")
            self.assertEqual(priced["C5"].value, "(Kiêm phiếu xuất kho)")
            self.assertEqual(
                priced["C7"].value,
                "Đơn vị mua hàng: CÔNG TY CỔ PHẦN NHỰA VÀ CƠ KHÍ HẢI PHÒNG",
            )
            self.assertEqual(priced["C6"].value, "Ngày 02 tháng 09 năm 2026")
            self.assertIn("C6:J6", {str(value) for value in priced.merged_cells.ranges})
            self.assertEqual(priced["C6"].alignment.horizontal, "center")
            self.assertGreaterEqual(priced["C6"].font.sz or 0, 14)
            for address in ('C6', 'C1', 'C2', 'C3'):
                self.assertFalse(priced[address].font.bold)
            self.assertEqual(priced["C9"].value, "Hình thức thanh toán: TM/CK")
            self.assertEqual(priced["H10"].value, "Đơn giá")
            self.assertEqual(priced["I10"].value, "Thành tiền")
            self.assertFalse(bool(priced.column_dimensions["H"].hidden))
            self.assertFalse(bool(priced.column_dimensions["I"].hidden))
            self.assertEqual(priced["H11"].value, 14_000)
            self.assertEqual(priced["I11"].value, 28_000)
            self.assertEqual(priced["I12"].value, 19_500)
            self.assertEqual(priced["C13"].value, "TỔNG CỘNG")
            self.assertEqual(priced["E13"].value, "3.5 Kg")
            self.assertEqual(priced["I13"].value, 47_500)
            self.assertEqual(str(priced.print_area), "'NHUAHP'!$C$1:$J$21")
            self.assertIn("C7:J7", {str(value) for value in priced.merged_cells.ranges})
            self.assertIn("C8:J8", {str(value) for value in priced.merged_cells.ranges})
            self.assertIn("C9:J9", {str(value) for value in priced.merged_cells.ranges})
            self.assertIn("C1:J1", {str(value) for value in priced.merged_cells.ranges})
            self.assertIn("C2:J2", {str(value) for value in priced.merged_cells.ranges})
            self.assertIn("C3:J3", {str(value) for value in priced.merged_cells.ranges})
            self.assertTrue(priced["C7"].alignment.shrinkToFit)

            self.assertTrue(hidden.column_dimensions["H"].hidden)
            self.assertTrue(hidden.column_dimensions["I"].hidden)
            self.assertEqual(hidden["C6"].value, "Ngày 02 tháng 09 năm 2026")
            self.assertGreaterEqual(hidden["C6"].font.sz or 0, 14)
            self.assertFalse(hidden["C6"].font.bold)
            self.assertEqual(hidden["C6"].alignment.horizontal, "center")
            self.assertIn("C6:J6", {str(value) for value in hidden.merged_cells.ranges})
            self.assertIn("C7:J7", {str(value) for value in hidden.merged_cells.ranges})
            self.assertEqual(str(hidden.print_area), "'POT'!$C$1:$J$21")
            for row in range(10, 40):
                self.assertIsNone(hidden[f"G{row}"].value)
                self.assertIsNone(hidden[f"H{row}"].value)
                self.assertIsNone(hidden[f"I{row}"].value)
            self.assertNotIn("987654321", repr([[cell.value for cell in row] for row in hidden.iter_rows()]))

            for sheet in (priced, hidden):
                self.assertEqual(str(sheet.page_setup.paperSize), "9")
                self.assertEqual(sheet.page_setup.orientation, "portrait")
                self.assertIsNone(sheet.page_setup.scale)
                self.assertEqual(sheet.page_setup.fitToWidth, 1)
                self.assertEqual(sheet.page_setup.fitToHeight, 0)
                self.assertEqual(len(sheet._images), 0)
                self.assertIsNone(sheet["D37"].value)
                self.assertTrue(sheet.column_dimensions["G"].hidden)
                self.assertEqual(sheet.column_dimensions["D"].width, 43)
                self.assertEqual(sheet.column_dimensions["J"].width, 24)
                self.assertTrue(sheet.print_options.horizontalCentered)
                self.assertEqual(sheet.print_title_rows, "$10:$10")
                self.assertEqual(len(sheet.conditional_formatting), 0)
                for row in range(11, 13):
                    for column in "CDEFGHIJ":
                        self.assertEqual(sheet[f"{column}{row}"].alignment.vertical, "center")
                    self.assertGreaterEqual(sheet[f"D{row}"].font.sz or 0, 16)
                    self.assertGreaterEqual(sheet[f"E{row}"].font.sz or 0, 16)
                    self.assertGreaterEqual(sheet[f"F{row}"].font.sz or 0, 15)
                    self.assertFalse(sheet[f"D{row}"].font.bold)
                    self.assertFalse(sheet[f"E{row}"].font.bold)
                    self.assertFalse(sheet[f"F{row}"].font.bold)
                    self.assertTrue(sheet[f"D{row}"].alignment.wrap_text)
                    self.assertEqual(sheet[f"E{row}"].alignment.horizontal, "center")
                    self.assertEqual(sheet[f"F{row}"].alignment.horizontal, "center")
                    self.assertGreaterEqual(sheet.row_dimensions[row].height or 0, 23)
                    self.assertGreaterEqual(sheet[f"J{row}"].font.sz or 0, 13)
            self.assertEqual(priced["C14"].value, "Ngày ..... tháng ..... năm ........")
            self.assertIn("Trực ban", priced["C15"].value)
            self.assertIn("Người nhận hàng", priced["C15"].value)
            self.assertEqual(priced["C16"].value.count("Ký và ghi rõ họ tên"), 4)
            self.assertEqual(hidden["C13"].value, "TỔNG CỘNG")
            self.assertEqual(hidden["E13"].value, "4 Kg")
            self.assertEqual(hidden["C14"].value, "Ngày ..... tháng ..... năm ........")
            self.assertIn("Người giao hàng", hidden["C15"].value)
            self.assertEqual(hidden["C16"].value.count("Ký và ghi rõ họ tên"), 4)
            self.assertEqual(priced["C14"].border.top.style, "hair")
            self.assertEqual(hidden["C14"].border.top.style, "hair")

            self.assertEqual(
                [column["label"] for column in sections[0]["columns"]],
                ["STT", "Tên hàng", "SL", "ĐVT", "Đơn giá", "Thành tiền", "GC"],
            )
            self.assertEqual(sections[0]["summary"][0]["value"], 47_500)
            self.assertEqual(
                [column["label"] for column in sections[1]["columns"]],
                ["STT", "Tên hàng", "SL", "ĐVT", "GC"],
            )
            self.assertNotIn("summary", sections[1])
            for section in sections:
                self.assertEqual(
                    [item["title"] for item in section["signatures"]],
                    ["TRỰC BAN", "BẢO VỆ", "NGƯỜI GIAO HÀNG", "NGƯỜI NHẬN HÀNG"],
                )

            candidate = safe_workbook_bytes(workbook)
        finally:
            workbook.close()
        self.assertEqual(hashlib.sha256(GOLDEN.read_bytes()).hexdigest(), source_hash)

        reopened = load_workbook(io.BytesIO(candidate), data_only=False, keep_links=False)
        try:
            self.assertEqual(reopened.sheetnames, ["NHUAHP", "POT"])
            self.assertEqual([len(sheet._images) for sheet in reopened.worksheets], [0, 0])
            self.assertIn("Người nhận hàng", reopened["NHUAHP"]["C15"].value)
            self.assertIn("Người nhận hàng", reopened["POT"]["C15"].value)
            self.assertFalse(reopened._external_links)
        finally:
            reopened.close()

    def test_more_than_golden_rows_moves_total_and_signature_without_blank_table_rows(self):
        rows = [
            {
                "product_name": f"Mặt hàng {index}",
                "quantity": index,
                "unit": "Kg",
                "sell_price": 1_000,
                "note": "",
            }
            for index in range(1, 21)
        ]
        rows[0]["product_name"] = "X" * 43
        workbook = build_delivery_workbook(
            [{
                "kitchen": "NHUAHP",
                "contractor": "NHUAHAIPHONG",
                "recipient": "Nhựa Hải Phòng",
                "address": "Hải Phòng",
                "items": rows,
            }],
            work_date="02/09/2026",
            template_path=GOLDEN,
        )
        try:
            sheet = workbook.active
            self.assertEqual(sheet["D30"].value, "Mặt hàng 20")
            self.assertEqual(sheet["C31"].value, "TỔNG CỘNG")
            self.assertEqual(sheet["I31"].value, 210_000)
            self.assertEqual(str(sheet.print_area), "'NHUAHP'!$C$1:$J$39")
            self.assertEqual(len(sheet._images), 0)
            self.assertEqual(sheet["C32"].value, "Ngày ..... tháng ..... năm ........")
            self.assertIn("Người nhận hàng", sheet["C33"].value)
            self.assertGreaterEqual(sheet.row_dimensions[11].height or 0, 48)
            self.assertEqual(
                [sheet[f"C{row}"].value for row in range(11, 31)], list(range(1, 21)),
            )
        finally:
            workbook.close()

    def test_long_delivery_balances_pages_and_repeats_the_table_header(self):
        rows = [
            {
                "product_name": f"Mặt hàng giao số {index}",
                "quantity": index,
                "unit": "Kg",
                "sell_price": 0,
                "note": "",
            }
            for index in range(1, 33)
        ]
        workbook = build_delivery_workbook(
            [{
                "kitchen": "YLK",
                "contractor": "ATV",
                "recipient": "BẾP YJLINK",
                "address": "Hải Phòng",
                "items": rows,
            }],
            work_date="04/09/2026",
            template_path=GOLDEN,
        )
        try:
            sheet = workbook.active
            self.assertEqual([item.id for item in sheet.row_breaks.brk], [26])
            self.assertEqual(sheet.print_title_rows, "$10:$10")
            self.assertTrue(sheet.print_options.horizontalCentered)
            self.assertAlmostEqual(sheet.page_margins.left, 0.25)
            self.assertAlmostEqual(sheet.page_margins.right, 0.25)
            self.assertEqual(len(sheet.conditional_formatting), 0)
            self.assertEqual(sheet.column_dimensions["D"].width, 43)
            self.assertEqual(sheet.column_dimensions["J"].width, 24)
            self.assertEqual(sheet["C26"].value, 16)
            self.assertEqual(sheet["C27"].value, 17)
        finally:
            workbook.close()

    def test_formula_injection_and_empty_deliveries_fail_closed(self):
        with self.assertRaises(DeliveryExportError) as empty:
            build_delivery_workbook([], work_date="2026-09-02", template_path=GOLDEN)
        self.assertEqual(empty.exception.code, "no_delivery_rows")
        payload = delivery_payload()[:1]
        payload[0]["items"][0]["product_name"] = "=WEBSERVICE(\"https://invalid\")"
        with self.assertRaises(TemplateWorkbookError):
            build_delivery_workbook(payload, work_date="2026-09-02", template_path=GOLDEN)

    def test_approved_delivery_sections_build_verified_a4_pdf(self):
        workbook = build_delivery_workbook(
            delivery_payload(), work_date="2026-09-02", template_path=GOLDEN,
        )
        try:
            sections = workbook_sections("deliveries", workbook)
            with tempfile.TemporaryDirectory(prefix="tdp_delivery_pdf_") as temp_name:
                output = Path(temp_name) / "delivery.pdf"
                manifest = build_pdf_bundle(
                    sections,
                    output,
                    company="CÔNG TY TNHH THỰC PHẨM THÀNH ĐẠT PHÁT",
                    document_title="Phiếu giao hàng ngày 02/09/2026",
                    generated_at="2026-09-02T12:00:00",
                    paper="A4",
                )
                self.assertTrue(manifest["ok"])
                self.assertEqual(manifest["paper"], "A4")
                self.assertEqual(manifest["section_count"], 2)
                self.assertTrue(output.is_file())
        finally:
            workbook.close()

    def test_server_groups_net_delivery_and_ignores_stale_show_price_flag(self):
        connection = sqlite3.connect(":memory:")
        connection.row_factory = sqlite3.Row
        connection.executescript(
            """
            CREATE TABLE settings(key TEXT PRIMARY KEY, value TEXT);
            CREATE TABLE kitchens(
                code TEXT PRIMARY KEY, contractor TEXT, name TEXT, address TEXT,
                show_price INTEGER NOT NULL DEFAULT 0
            );
            INSERT INTO kitchens VALUES(
                'NHUAHP','NHUAHAIPHONG','Nhựa Hải Phòng','Địa chỉ Nhựa',1
            );
            INSERT INTO kitchens VALUES(
                'POT','HATRAN','Bếp POT','Địa chỉ POT',1
            );
            INSERT INTO settings VALUES('delivery_recipient_POT','ĐƠN VỊ POT ĐÃ CẤU HÌNH');
            """
        )
        orders = [
            {
                "kitchen": "NHUAHP", "contractor": "NHUAHAIPHONG",
                "product_code": "A", "product_name": "Hàng A", "actual_delivered": 3,
                "customer_return_qty": 1, "unit": "Kg", "sell_price": 10_000, "note": "",
            },
            {
                "kitchen": "POT", "contractor": "HATRAN",
                "product_code": "B", "product_name": "Hàng B", "actual_delivered": 4,
                "customer_return_qty": 0, "unit": "Kg", "sell_price": 777_777, "note": "",
            },
            {
                "kitchen": "POT", "contractor": "HATRAN",
                "product_code": "C", "product_name": "Không giao", "actual_delivered": 1,
                "customer_return_qty": 1, "unit": "Kg", "sell_price": 888_888, "note": "",
            },
        ]
        workbook = server_module.export_deliveries(
            connection, {"work_date": "2026-09-02"}, orders,
        )
        try:
            self.assertEqual(workbook["NHUAHP"]["E11"].value, 2)
            self.assertEqual(workbook["NHUAHP"]["H11"].value, 10_000)
            self.assertEqual(workbook["POT"]["C7"].value, "Đơn vị mua hàng: ĐƠN VỊ POT ĐÃ CẤU HÌNH")
            self.assertTrue(workbook["POT"].column_dimensions["H"].hidden)
            self.assertIsNone(workbook["POT"]["H11"].value)
            self.assertIsNone(workbook["POT"]["I11"].value)
            self.assertEqual(workbook["POT"]["D11"].value, "Hàng B")
            self.assertIsNone(workbook["POT"]["D12"].value)
        finally:
            workbook.close()
            connection.close()


if __name__ == "__main__":
    unittest.main()
