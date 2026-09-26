"""Same 13-column artwork as M-Invoice, with exact unissued order quantities.

This is a read-only reconciliation export. Invoiceable catch-up files must use
the ordinary stock-checked draft/allocation pipeline instead.
"""
import io
import zipfile
from collections import defaultdict
from decimal import Decimal, ROUND_HALF_UP
from copy import copy
from openpyxl import load_workbook
from openpyxl.comments import Comment

try:
    from .invoice_tax_export import build_invoice_workbook, _safe_name
    from .outgoing_unissued import unissued_workbook
    from .template_workbook import safe_workbook_bytes
except ImportError:
    from invoice_tax_export import build_invoice_workbook, _safe_name
    from outgoing_unissued import unissued_workbook
    from template_workbook import safe_workbook_bytes


def _period(payload):
    return (f'TU_{payload["from"]}_' if payload.get('from') else '')+f'DEN_{payload["asof"]}'


def _tax_split_zip(payload, template_dir, tax_percent):
    waiting=payload.get('portion')=='waiting'
    groups = defaultdict(dict)
    for row in payload['details']:
        qty = Decimal(str(row['waiting_qty'] if waiting else row['unissued_qty'])).quantize(Decimal('.000001'), rounding=ROUND_HALF_UP)
        if qty <= 0:
            continue
        vat = tax_percent(row['tax'])
        price = Decimal(str(row['unit_price']))
        nature = str(row.get('invoice_nature') or '1')
        key = (row['product_code'], row['unit'].strip().casefold(), price, nature)
        lines = groups[(row['contractor'], vat)]
        item = lines.setdefault(key, {'contractor':row['contractor'], 'product_code':row['product_code'],
            'product_name':row.get('invoice_name') or row['product_name'], 'unit':row['unit'], 'unit_price':price,
            'invoice_nature':nature, 'qty':Decimal(0), 'conversion_reasons':set()})
        item['qty'] += qty
        if row.get('needs_conversion'):
            item['conversion_reasons'].add(row['conversion_reason'])
    if not groups:
        raise ValueError('Không còn hàng chờ trong phạm vi đã chọn.' if waiting else 'Không còn hàng chưa xuất hóa đơn trong phạm vi đã chọn.')
    output = io.BytesIO()
    count = 0
    with zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as archive:
        filenames = set()
        for (party, vat), group in sorted(groups.items()):
            lines = []
            for key, row in sorted(group.items()):
                amount = (row['qty'] * row['unit_price']).quantize(Decimal('1'), rounding=ROUND_HALF_UP)
                lines.append({**row, 'qty':float(row['qty']), 'unit_price':float(row['unit_price']), 'amount':float(amount)})
            book, _ = build_invoice_workbook(lines, vat_percent=vat, template_dir=template_dir)
            if any(row['conversion_reasons'] for row in lines):
                workbook = load_workbook(io.BytesIO(book))
                try:
                    for number, row in enumerate(lines, start=2):
                        if not row['conversion_reasons']:
                            continue
                        for cell in workbook.active[number]:
                            font = copy(cell.font)
                            font.color = 'B42318'
                            cell.font = font
                        workbook.active.cell(number, 3).comment = Comment(
                            'Cần quy đổi / đối chiếu ĐVT. Số lượng và đơn giá đang theo đơn gốc.\n'
                            + '\n'.join(sorted(row['conversion_reasons'])), 'TĐP')
                    # The base invoice template has already received print styling.
                    # Keep the customer's red conversion indicators on this report.
                    book = safe_workbook_bytes(workbook, apply_print_style=False)
                finally:
                    workbook.close()
            label = 'KKKNT' if vat == -2 else 'KCT' if vat == -1 else f'VAT{vat:g}'
            prefix='CON_CHO' if waiting else 'CHUA_XUAT'
            name = f"{prefix}_{_safe_name(party)}_{label}_{_period(payload)}.xlsx"
            if name.casefold() in filenames:
                raise ValueError('Tên file nhà thầu bị trùng; cần kiểm tra lại mã nhà thầu.')
            filenames.add(name.casefold()); archive.writestr(name, book); count += 1
        archive.writestr('DOI_CHIEU_CHI_TIET.xlsx', unissued_workbook(payload).getvalue())
        guide = [
            'BẢNG HÀNG CHƯA XUẤT HÓA ĐƠN – CÙNG MẪU 13 CỘT',
            f"Đơn đã duyệt đến hết {payload['asof']}. Nhà thầu: {payload['contractor'] or 'Tất cả'}.",
            'Các file CHUA_XUAT dùng để đối chiếu; gồm cả hàng thiếu đầu vào và phần lẻ chưa đủ xuất. Không dùng bộ đối chiếu này để nhập M-Invoice.',
            'Muốn xuất hóa đơn, kể cả xuất bù: bấm “Tải bảng kê để up M-Invoice” ở đầu trang. Web kiểm tra lại hóa đơn đã ký, đơn cũ và tồn kho rồi tạo bộ file đưa lên M-Invoice.',
            'Lượng chưa xuất = lượng thực giao của đơn đã duyệt sau trả hàng − lượng hóa đơn đã phát hành đã đối chiếu. Giữ nguyên phần lẻ trong bản đối chiếu.',
            'Cùng mã, đơn vị, giá bán, thuế và tính chất được cộng lượng. Khác giá giữ dòng riêng; không đổi giá bán trên đơn.',
            'Dòng chữ đỏ cần quy đổi hoặc đối chiếu ĐVT; xem ghi chú ở ô ĐVT. Chưa tự đổi gói/hộp/túi thành kg. Số lượng và đơn giá trong bảng này theo đơn gốc.',
            'Dòng xóa khỏi Excel trước khi ký vẫn còn chưa xuất trên web. Kho hóa đơn chỉ ghi theo hóa đơn đã ký trên M-Invoice được đồng bộ, đối chiếu đúng mã và ĐVT; đồng bộ lại không trừ lần hai.',
            'Công nợ phải thu theo đơn đã duyệt. Đề nghị thanh toán và bảng kê hóa đơn lấy theo hóa đơn đã ký được đồng bộ từ M-Invoice.',
            'Tải bảng hoặc xuất bù không duyệt lại đơn, không ghi thêm doanh thu hay công nợ. Tải file chưa tính là đã phát hành hóa đơn.',
            'Ngày chọn ở web là mốc đơn hàng. Ngày hóa đơn thực tế chọn khi phát hành trên M-Invoice.',
        ]
        if waiting:
            guide[:1]=['PHẦN CÒN CHỜ – CHƯA DÙNG ĐƯA LÊN M-INVOICE',
                       'File CON_CHO chỉ lấy lượng còn chờ, không lấy phần đủ điều kiện đang giữ để xuất. Xem lý do từng mặt hàng trong DOI_CHIEU_CHI_TIET.xlsx.']
        if payload['warnings']:
            guide += ['CÒN HÓA ĐƠN CẦN ĐỐI CHIẾU – số chưa xuất có thể thay đổi sau khi xử lý:']
            guide += [w['message'] for w in payload['warnings']]
        archive.writestr('HUONG_DAN.txt', '\n'.join(guide).encode('utf-8-sig'))
    output.seek(0)
    return output, count


def _literal(value):
    text=str(value or '')
    return "'"+text if text.startswith(('=','+','-','@')) else text


def _contractor_workbook(archive, names, party, details, payload, tax_percent):
    """Combine validated 13-column sheets without merging different VAT rates."""
    workbook=None
    totals=[]
    waiting=payload.get('portion')=='waiting'
    try:
        for name in names:
            part=load_workbook(io.BytesIO(archive.read(name)))
            subtotal=tax_amount=total=0
            for row in part.active.iter_rows(min_row=2):
                subtotal+=row[8].value or 0
                tax_amount+=row[10].value or 0
                total+=row[11].value or 0
            vat=part.active['J2'].value
            label='KKKNT' if vat==-2 else 'KCT' if vat==-1 else f'{vat:g}%'
            totals.append([label,part.active.max_row-1,subtotal,tax_amount,total])
            if workbook is None:
                workbook=part
                sheet=workbook.active
                sheet.title='Con cho' if waiting else 'Chua xuat'
            else:
                try:
                    for source_row in part.active.iter_rows(min_row=2):
                        number=sheet.max_row+1
                        for source in source_row:
                            cell=sheet.cell(number,source.column,source.value)
                            for attr in ('font','fill','border','alignment','protection','comment'):
                                setattr(cell,attr,copy(getattr(source,attr)))
                            cell.number_format=source.number_format
                finally:part.close()
        end=sheet.max_row
        sheet.freeze_panes='A2';sheet.auto_filter.ref=f'A1:M{end}'
        grand=[sum(row[column] for row in totals) for column in (1,2,3,4)]
        footer=end+2
        sheet.cell(footer,2,'TỔNG CỘNG')
        for column,value in ((6,grand[1]),(9,grand[1]),(11,grand[2]),(12,grand[3])):
            sheet.cell(footer,column,value).number_format='#,##0'
        dates=(payload['from']+' → ' if payload.get('from') else 'Đến ')+payload['asof']
        sheet.cell(footer+1,2,_literal(f'{party} · Ngày đơn {dates} · Tiền theo dữ liệu lúc tải'))

        summary=workbook.create_sheet('Tong hop')
        summary.append(['Nhà thầu',_literal(party)])
        summary.append(['Khoảng ngày đơn',dates])
        summary.append(['Phạm vi','Chỉ phần còn chờ' if waiting else 'Toàn bộ phần chưa ký được đối chiếu'])
        summary.append(['Thuế suất','Số dòng sau gộp','Tiền trước thuế','Tiền thuế','Tổng tiền'])
        for row in totals:summary.append(row)
        summary.append(['TỔNG CỘNG',*grand])
        summary.append(['Cách tính','Cùng mã, ĐVT, giá, thuế và tính chất cộng lượng; khác giá/thuế giữ dòng riêng.'])
        summary.append(['Hàng đã bỏ chọn','Vẫn nằm trong bảng chưa xuất; xem lý do ở sheet Chi tiet don. Không tự đưa lại vào bản nháp và không làm giảm doanh thu/công nợ.'])
        summary.append(['Đơn vị','Số lượng và đơn giá theo đơn gốc. Dòng đỏ cần kiểm tra ĐVT; chưa tự thay số lượng.'])
        summary.append(['Đối chiếu tiền','Tiền trong bảng = lượng chưa xuất × giá trên đơn. Nếu giá trên hóa đơn đã ký khác giá đơn, cần đối chiếu khoản chênh riêng; không tự sửa giá hoặc lượng để bù tiền.'])
        summary.append(['Cập nhật hóa đơn đã ký',payload.get('source_checked_at') or 'Bản dữ liệu đã lưu; chưa cập nhật M-Invoice trong lần tải này.'])
        warnings=[w['message'] for w in payload.get('warnings',[]) if not w.get('contractor') or w['contractor']==party]
        if warnings:
            sheet.cell(footer+2,2,'CHƯA ĐỐI CHIẾU XONG — tổng tiền có thể thay đổi. Xem sheet Tong hop.')
            for warning in warnings:summary.append(['CẦN ĐỐI CHIẾU',_literal(warning)])

        detail=workbook.create_sheet('Chi tiet don')
        detail.append(['Dòng đơn','Ngày đơn','Mã hàng','Tên xuất hóa đơn','ĐVT đơn',
                       'Lượng còn chờ' if waiting else 'Lượng chưa xuất','Đơn giá','Thuế',
                       'Tiền hàng','Lý do còn chờ','ĐVT cần đối chiếu'])
        for row in details:
            qty=row['waiting_qty'] if waiting else row['unissued_qty']
            amount=int((Decimal(str(qty))*Decimal(str(row['unit_price']))).quantize(Decimal('1'),rounding=ROUND_HALF_UP))
            vat=tax_percent(row['tax'])
            tax_label='KKKNT' if vat==-2 else 'KCT' if vat==-1 else f'{vat:g}%'
            detail.append([row['order_id'],row['work_date'],_literal(row['product_code']),
                           _literal(row.get('invoice_name') or row['product_name']),_literal(row['unit']),
                           qty,row['unit_price'],tax_label,amount,_literal(row.get('pending_reason') or ('Đã chuẩn bị — chờ gửi/ký, không phải bị chặn vì thiếu tồn.' if row.get('ready_qty',0)>0 else 'Chưa chuẩn bị bản nháp.')),
                           _literal(row.get('conversion_reason',''))])
            if row.get('needs_conversion'):
                for cell in detail[detail.max_row]:
                    font=copy(cell.font);font.color='B42318';cell.font=font
        detail.freeze_panes='A2';detail.auto_filter.ref=detail.dimensions
        for tab in (summary,detail):
            for col in tab.columns:
                tab.column_dimensions[col[0].column_letter].width=min(65,max(16,max(len(str(c.value or '')) for c in col)+2))
            for row in tab.iter_rows(min_row=2):
                for cell in row:
                    if isinstance(cell.value,(int,float)):
                        cell.number_format='#,##0' if float(cell.value).is_integer() else '#,##0.######'
        all_warnings = payload.get('warnings') or []
        scoped_warnings = [w for w in all_warnings if not w.get('contractor') or w.get('contractor') == party]
        incomplete = bool(scoped_warnings) or payload.get('reconciliation_complete') is False
        status_text = ('CHƯA ĐỐI CHIẾU XONG – không dùng tổng này để xuất hóa đơn bổ sung' if incomplete
                      else 'Đã đối chiếu lượng theo dữ liệu đồng bộ; chưa xác nhận đối chiếu tiền hóa đơn')
        policy_text = 'Hóa đơn đã ký có thể đang chờ ghép với đơn. Đây là báo cáo đối chiếu, không phải file hóa đơn đã ký.'
        sheet.merge_cells(start_row=footer+2,start_column=2,end_row=footer+2,end_column=13)
        notice = sheet.cell(footer+2,2,_literal(status_text))
        notice.alignment = copy(notice.alignment)
        notice.alignment = notice.alignment.copy(wrap_text=True)
        sheet.row_dimensions[footer+2].height = 32
        summary.append(['Trạng thái đối chiếu',_literal(status_text)])
        summary.append(['Ghi chú',_literal(policy_text)])
        for row in summary.iter_rows(min_row=summary.max_row-1):
            for cell in row:
                cell.alignment = cell.alignment.copy(wrap_text=True)
            summary.row_dimensions[row[0].row].height = 42
        warning_sheet = workbook.create_sheet('Canh bao doi chieu')
        warning_sheet.append(['Nội dung','Chi tiết'])
        warning_sheet.append(['Nhà thầu',_literal(party)])
        warning_sheet.append(['Ngày đơn',_literal(dates)])
        warning_sheet.append(['Trạng thái',_literal(status_text)])
        warning_sheet.append(['Lưu ý',_literal(policy_text)])
        for warning in scoped_warnings:
            warning_sheet.append([_literal(warning.get('code') or 'Cảnh báo'),_literal(warning.get('message') or '')])
        if incomplete and not scoped_warnings:
            warning_sheet.append(['Cần kiểm tra','Báo cáo nguồn chưa xác nhận hoàn tất đối chiếu. Kiểm tra cảnh báo trên hệ thống trước khi xuất bổ sung.'])
        warning_sheet.column_dimensions['A'].width=24
        warning_sheet.column_dimensions['B'].width=65
        warning_sheet.freeze_panes='A2'
        for row in warning_sheet.iter_rows():
            for cell in row:
                cell.alignment=cell.alignment.copy(wrap_text=True,vertical='top')
            warning_sheet.row_dimensions[row[0].row].height=45

        return safe_workbook_bytes(workbook,apply_print_style=False)
    finally:
        if workbook is not None:workbook.close()


def unissued_template_zip(payload, template_dir, tax_percent):
    # Keep the ordinary M-Invoice export path and its tax-specific templates
    # unchanged. Only the read-only unissued report is combined by contractor.
    if payload.get('portion')!='waiting':
        payload={**payload,'details':payload['details']+payload.get('skipped_details',[])}
    split,_=_tax_split_zip(payload,template_dir,tax_percent)
    waiting=payload.get('portion')=='waiting'
    prefix='CON_CHO' if waiting else 'CHUA_XUAT'
    parties=defaultdict(list)
    for row in payload['details']:
        qty=Decimal(str(row['waiting_qty' if waiting else 'unissued_qty'])).quantize(Decimal('.000001'),rounding=ROUND_HALF_UP)
        if qty>0:parties[row['contractor']].append(row)
    output=io.BytesIO()
    with zipfile.ZipFile(split) as source,zipfile.ZipFile(output,'w',zipfile.ZIP_DEFLATED) as target:
        names=set()
        for party,details in sorted(parties.items()):
            parts=[]
            for vat in sorted({tax_percent(row['tax']) for row in details}):
                label='KKKNT' if vat==-2 else 'KCT' if vat==-1 else f'VAT{vat:g}'
                parts.append(f'{prefix}_{_safe_name(party)}_{label}_{_period(payload)}.xlsx')
            name=f'{prefix}_{_safe_name(party)}_{_period(payload)}.xlsx'
            if name.casefold() in names:raise ValueError('Tên file nhà thầu bị trùng; cần kiểm tra lại mã nhà thầu.')
            names.add(name.casefold())
            target.writestr(name,_contractor_workbook(source,parts,party,details,payload,tax_percent))
        guide=[
            'MỖI NHÀ THẦU MỘT FILE — BẢNG CHƯA XUẤT / CÒN CHỜ',
            f'Ngày đơn: {payload.get("from") or "từ đầu"} → {payload["asof"]}.',
            'Mọi thuế suất của một nhà thầu nằm chung một bảng 13 cột. Cuối bảng có tổng tiền; sheet Tong hop có tổng theo thuế suất và tổng chung. Không cần cộng nhiều file.',
            'Sheet Chi tiet don ghi từng dòng đơn và lý do còn chờ. Khác giá hoặc khác thuế giữ dòng riêng.',
            'Dòng đỏ cần kiểm tra ĐVT; không phải mọi dòng đều cần nhập kg. Số lượng và đơn giá giữ theo đơn gốc.',
            'Lượng chưa xuất = đơn đã duyệt sau trả hàng trừ phần hóa đơn đã ký được đối chiếu; tải file chưa làm giảm lượng chưa xuất.',
            'File này để đối chiếu, gồm cả dòng chưa đủ điều kiện. Muốn xuất: dùng Tải bảng kê để up M-Invoice trên web để kiểm tra tồn và hóa đơn đã ký.',
            'Web cập nhật hóa đơn đã ký trước khi tải bảng; hóa đơn mới được đối chiếu và ghi kho một lần. Không ghi thêm doanh thu/công nợ. Dòng xóa riêng trong Excel trước khi ký vẫn chưa xuất trên web.',
        ]
        if waiting:guide.append('File CON_CHO chỉ lấy lượng còn chờ; không lấy phần đã đủ điều kiện đang giữ để xuất.')
        guide.extend('Cần đối chiếu: '+w['message'] for w in payload['warnings'])
        target.writestr('HUONG_DAN.txt','\n'.join(guide).encode('utf-8-sig'))
    output.seek(0)
    return output,len(parties)
