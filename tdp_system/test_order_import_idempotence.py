from __future__ import annotations

import io
import json
import tempfile
import threading
import unittest
from concurrent.futures import ThreadPoolExecutor
from datetime import date
from pathlib import Path

from openpyxl import Workbook

try:
    from . import server
except ImportError:  # pragma: no cover - direct file invocation
    import server


class OrderImportIdempotenceTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        cls.original_db_path = server.DB_PATH
        cls.original_data_dir = server.DATA_DIR
        cls.original_master_source = server.MASTER_SOURCE
        cls.original_testing = server.app.config.get("TESTING", False)
        server.DATA_DIR = Path(cls.temp.name) / "data"
        server.DATA_DIR.mkdir(parents=True, exist_ok=True)
        server.DB_PATH = server.DATA_DIR / "orders.sqlite3"
        server.MASTER_SOURCE = Path(cls.temp.name) / "master-disabled.xlsx"
        server.app.config["TESTING"] = True
        server.init_database()
        cls.client = server.app.test_client()
        with server.db() as conn:
            conn.execute(
                "INSERT INTO contractors(code,name,price_group,pricing_mode) VALUES('C1','C1','C1','group')"
            )
            conn.execute(
                "INSERT INTO kitchens(code,contractor,name) VALUES('K1','C1','Kitchen 1')"
            )
            conn.execute("INSERT INTO suppliers(code,name) VALUES('S1','Supplier 1')")
            conn.execute(
                """INSERT INTO products(
                       code,name,unit,tax,supplier,buy_price,purchase_list,seller,cccd
                   ) VALUES('P1','Product 1','kg','0%','S1',10000,0,'','')"""
            )

    @classmethod
    def tearDownClass(cls):
        with server.PENDING_IMPORT_LOCK:
            for pending in server.PENDING_IMPORTS.values():
                pending["path"].unlink(missing_ok=True)
            server.PENDING_IMPORTS.clear()
        server.DB_PATH = cls.original_db_path
        server.DATA_DIR = cls.original_data_dir
        server.MASTER_SOURCE = cls.original_master_source
        server.app.config["TESTING"] = cls.original_testing
        cls.temp.cleanup()

    def setUp(self):
        with server.PENDING_IMPORT_LOCK:
            for pending in server.PENDING_IMPORTS.values():
                pending["path"].unlink(missing_ok=True)
            server.PENDING_IMPORTS.clear()
        with server.db() as conn:
            conn.execute("DELETE FROM order_import_receipts")
            conn.execute("DELETE FROM orders")
            conn.execute("DELETE FROM batches")

    @staticmethod
    def workbook_bytes(*, qty=2, include_prior_sheet=False, current_sheet="29.08") -> bytes:
        workbook = Workbook()
        current = workbook.active
        current.title = current_sheet
        headers = [
            "Ngày", "Nhà thầu", "Mã bếp", "Mã hàng", "Tên hàng", "Số lượng",
            "ĐVT", "NCC", "Giá mua", "Giá bán", "Thuế",
        ]
        current.append(headers)
        current.append([
            date(2026, 8, 29), "C1", "K1", "P1", "Product 1", qty,
            "kg", "S1", 10_000, 12_000, "0%",
        ])
        if include_prior_sheet:
            prior = workbook.create_sheet("28.08")
            prior.append(headers)
            prior.append([
                date(2026, 8, 28), "C1", "K1", "P1", "Product 1", qty,
                "kg", "S1", 10_000, 12_000, "0%",
            ])
        stream = io.BytesIO()
        workbook.save(stream)
        workbook.close()
        return stream.getvalue()

    def analyze(self, payload: bytes, filename="orders.xlsx") -> dict:
        response = self.client.post(
            "/api/import/analyze",
            data={"file": (io.BytesIO(payload), filename)},
            content_type="multipart/form-data",
        )
        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        return response.get_json()

    @staticmethod
    def confirm_with_client(client, token: str, sheets=None, work_date="2026-08-29"):
        return client.post(
            "/api/import/confirm",
            json={"token": token, "sheets": sheets or ["29.08"], "work_date": work_date},
        )

    def database_counts(self):
        with server.db() as conn:
            return tuple(
                conn.execute(f"SELECT COUNT(*) n FROM {table}").fetchone()["n"]
                for table in ("batches", "orders", "order_import_receipts")
            )

    def test_sequential_replay_returns_the_original_batch(self):
        payload = self.workbook_bytes()
        first_token = self.analyze(payload)["token"]
        repeat_token = self.analyze(payload, "renamed.xlsx")["token"]

        first = self.confirm_with_client(self.client, first_token)
        repeat = self.confirm_with_client(self.client, repeat_token)
        self.assertEqual(first.status_code, 200, first.get_data(as_text=True))
        self.assertEqual(repeat.status_code, 200, repeat.get_data(as_text=True))
        first_json = first.get_json()
        repeat_json = repeat.get_json()
        self.assertFalse(first_json["idempotent"])
        self.assertTrue(repeat_json["idempotent"])
        self.assertEqual(first_json["batch"]["id"], repeat_json["batch"]["id"])
        self.assertEqual(first_json["importKey"], repeat_json["importKey"])
        self.assertEqual(first_json["sourceHash"], repeat_json["sourceHash"])
        self.assertEqual(self.database_counts(), (1, 1, 1))

    def test_two_confirmations_racing_on_distinct_tokens_cannot_duplicate(self):
        payload = self.workbook_bytes()
        tokens = [self.analyze(payload, f"orders-{index}.xlsx")["token"] for index in range(2)]
        barrier = threading.Barrier(2)

        def worker(token):
            barrier.wait(timeout=5)
            with server.app.test_client() as client:
                response = self.confirm_with_client(client, token)
                return response.status_code, response.get_json()

        with ThreadPoolExecutor(max_workers=2) as executor:
            results = list(executor.map(worker, tokens))
        self.assertEqual([status for status, _ in results], [200, 200], results)
        payloads = [item for _, item in results]
        self.assertEqual({item["batch"]["id"] for item in payloads}, {payloads[0]["batch"]["id"]})
        self.assertEqual(sorted(item["idempotent"] for item in payloads), [False, True])
        self.assertEqual(self.database_counts(), (1, 1, 1))

    def test_sheet_set_work_date_and_bytes_are_part_of_the_key(self):
        payload = self.workbook_bytes(include_prior_sheet=True)
        first = self.confirm_with_client(self.client, self.analyze(payload)["token"], ["29.08"])
        other_sheet = self.confirm_with_client(
            self.client, self.analyze(payload)["token"], ["28.08"],
        )
        other_date = self.confirm_with_client(
            self.client, self.analyze(payload)["token"], ["29.08"], "2026-08-30",
        )
        other_bytes = self.confirm_with_client(
            self.client, self.analyze(self.workbook_bytes(qty=3))["token"], ["29.08"],
        )
        for response in (first, other_sheet, other_date, other_bytes):
            self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
            self.assertFalse(response.get_json()["idempotent"])
        self.assertEqual(len({
            response.get_json()["importKey"]
            for response in (first, other_sheet, other_date, other_bytes)
        }), 4)
        # New bytes in the same date/sheet scope replace the prior import.
        self.assertEqual(self.database_counts(), (3, 3, 3))
        self.assertEqual(first.get_json()["batch"]["id"], other_bytes.get_json()["batch"]["id"])
        self.assertEqual(other_bytes.get_json()["replacement"]["replaced_rows"], 1)
        with server.db() as conn:
            self.assertEqual(conn.execute("SELECT COUNT(*) FROM order_reimport_history").fetchone()[0], 1)

    def test_deleted_batch_receipt_cascades_and_allows_deliberate_reimport(self):
        payload = self.workbook_bytes()
        first = self.confirm_with_client(self.client, self.analyze(payload)["token"])
        batch_id = first.get_json()["batch"]["id"]
        with server.db() as conn:
            conn.execute("DELETE FROM batches WHERE id=?", (batch_id,))
        self.assertEqual(self.database_counts(), (0, 0, 0))

        replay = self.confirm_with_client(self.client, self.analyze(payload)["token"])
        self.assertEqual(replay.status_code, 200, replay.get_data(as_text=True))
        self.assertFalse(replay.get_json()["idempotent"])
        self.assertNotEqual(replay.get_json()["batch"]["id"], batch_id)
        self.assertEqual(self.database_counts(), (1, 1, 1))

    def test_exact_sheet_title_with_trailing_space_is_preserved(self):
        sheet_title = "đơn hàng27.08 "
        payload = self.workbook_bytes(current_sheet=sheet_title)
        first = self.confirm_with_client(
            self.client, self.analyze(payload)["token"], [sheet_title],
        )
        repeat = self.confirm_with_client(
            self.client, self.analyze(payload)["token"], [sheet_title],
        )
        self.assertEqual(first.status_code, 200, first.get_data(as_text=True))
        self.assertEqual(repeat.status_code, 200, repeat.get_data(as_text=True))
        self.assertEqual(first.get_json()["selectedSheets"], [sheet_title])
        self.assertEqual(repeat.get_json()["selectedSheets"], [sheet_title])
        self.assertFalse(first.get_json()["idempotent"])
        self.assertTrue(repeat.get_json()["idempotent"])
        self.assertEqual(self.database_counts(), (1, 1, 1))

    def test_sheet_selection_order_does_not_change_the_receipt(self):
        payload = self.workbook_bytes(include_prior_sheet=True)
        first = self.confirm_with_client(
            self.client, self.analyze(payload)["token"], ["29.08", "28.08"],
        )
        repeat = self.confirm_with_client(
            self.client, self.analyze(payload)["token"], ["28.08", "29.08"],
        )
        self.assertEqual(first.status_code, 200, first.get_data(as_text=True))
        self.assertEqual(repeat.status_code, 200, repeat.get_data(as_text=True))
        self.assertEqual(first.get_json()["importKey"], repeat.get_json()["importKey"])
        self.assertEqual(first.get_json()["batch"]["id"], repeat.get_json()["batch"]["id"])
        self.assertEqual(first.get_json()["selectedSheets"], ["28.08", "29.08"])
        self.assertFalse(first.get_json()["idempotent"])
        self.assertTrue(repeat.get_json()["idempotent"])
        self.assertEqual(self.database_counts(), (1, 2, 1))

    def test_malformed_nonblank_quantity_is_imported_as_blocked_error_row(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "29.08"
        sheet.append([
            "Ngày", "Nhà thầu", "Mã bếp", "Mã hàng", "Tên hàng", "Số lượng",
            "ĐVT", "NCC", "Giá mua", "Giá bán", "Thuế",
        ])
        sheet.append([
            date(2026, 8, 29), "C1", "K1", "P1", "Product 1", 2,
            "kg", "S1", 10_000, 12_000, "0%",
        ])
        sheet.append([
            date(2026, 8, 29), "C1", "K1", "P1", "Product 1", "garbage",
            "kg", "S1", 10_000, 12_000, "0%",
        ])
        stream = io.BytesIO()
        workbook.save(stream)
        workbook.close()

        analyzed = self.analyze(stream.getvalue())
        self.assertEqual(analyzed["sheets"][0]["rows"], 2)
        confirmed = self.confirm_with_client(self.client, analyzed["token"])
        self.assertEqual(confirmed.status_code, 200, confirmed.get_data(as_text=True))
        batch_id = confirmed.get_json()["batch"]["id"]
        with server.db() as conn:
            rows = conn.execute(
                "SELECT qty,errors FROM orders WHERE batch_id=? ORDER BY source_row", (batch_id,)
            ).fetchall()
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1]["qty"], 0)
        self.assertTrue(json.loads(rows[1]["errors"]))
        approval = self.client.post(f"/api/batches/{batch_id}/approve")
        self.assertEqual(approval.status_code, 400, approval.get_data(as_text=True))


if __name__ == "__main__":
    unittest.main()
