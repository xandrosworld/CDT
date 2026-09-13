import io
import zipfile
import unittest
from openpyxl import load_workbook
from . import server
from .test_supplier_plan_source import SupplierPlanSourceTests as Fixture


class SupplierNotesTests(unittest.TestCase):
    setUpClass=classmethod(Fixture.setUpClass.__func__)
    tearDownClass=classmethod(Fixture.tearDownClass.__func__)
    setUp=Fixture.setUp
    file=Fixture.file
    daily=Fixture.daily
    needs=Fixture.needs
    protected=Fixture.protected
    preview=Fixture.preview
    confirm=Fixture.confirm

    def save(self,bid,plan,note,**overrides):
        body={'plan_hash':plan['plan_hash'],'rows':[{'note_key':plan['rows'][0]['note_key'],'note':note}]}
        body.update(overrides)
        return self.client.put(f'/api/supplier-needs/{bid}/notes',json=body)

    def test_persistent_notes_export_reopen_and_do_not_change_money(self):
        bid=self.daily(self.file())
        p=self.needs(bid)
        self.assertEqual(self.client.put(f'/api/supplier-order-status/{bid}/S1',json={
            'status':'ordered','revision':0,'plan_hash':p['plan_hash']}).status_code,200)
        before=self.protected()
        note='Giao riêng bếp K1\nGọi trước khi giao <đủ hàng>'
        saved=self.save(bid,p,note)
        self.assertEqual(saved.status_code,200,saved.get_json())
        p2=self.needs(bid)
        self.assertEqual(p2['rows'][0]['note'],note)
        self.assertEqual(p2['groups'][0]['items'][0]['note'],note)
        self.assertEqual(p2['checklist'][0]['status'],'reopened')
        self.assertNotEqual(p2['plan_hash'],p['plan_hash'])
        self.assertEqual(before,self.protected())
        exported=self.client.get(f'/api/export/suppliers/{bid}')
        self.assertEqual(exported.status_code,200)
        payloads=[exported.data]
        if exported.headers.get('Content-Type','').startswith('application/zip'):
            with zipfile.ZipFile(io.BytesIO(exported.data)) as z:payloads=[z.read(n) for n in z.namelist() if n.endswith('.xlsx')]
        found=False
        for payload in payloads:
            wb=load_workbook(io.BytesIO(payload),data_only=True)
            found=found or any(c==note for s in wb for row in s.iter_rows(values_only=True) for c in row)
            wb.close()
        self.assertTrue(found,'Saved instruction must appear in supplier Excel')
        self.assertEqual(self.save(bid,p2,note).get_json()['changed'],0)
        self.assertEqual(self.needs(bid)['plan_hash'],p2['plan_hash'])
        self.assertEqual(self.save(bid,p2,'').status_code,200)
        self.assertEqual(self.needs(bid)['rows'][0]['note'],'')

    def test_stale_or_invalid_edits_are_atomic(self):
        bid=self.daily(self.file());p=self.needs(bid)
        self.assertEqual(self.save(bid,p,'Đã chốt lời dặn').status_code,200)
        self.assertEqual(self.save(bid,p,'Ghi đè từ tab cũ').status_code,409)
        p=self.needs(bid)
        for edits in [[],[{'note_key':'bad','note':'x'}],
                      [{'note_key':p['rows'][0]['note_key'],'note':'=SUM(A1)'}],
                      [{'note_key':p['rows'][0]['note_key'],'note':'x'*1001}],
                      [{'note_key':p['rows'][0]['note_key'],'note':'Mới'},{'note_key':'bad','note':'x'}]]:
            with self.subTest(edits=edits):
                self.assertEqual(self.save(bid,p,'',rows=edits).status_code,400)
                self.assertEqual(self.needs(bid)['rows'][0]['note'],'Đã chốt lời dặn')

    def test_new_source_note_takes_precedence_and_changed_qty_keeps_note(self):
        bid=self.daily(self.file());p=self.needs(bid)
        self.assertEqual(self.save(bid,p,'Lời dặn trên web').status_code,200)
        self.confirm(self.preview(bid,self.file(8)))
        self.assertEqual(self.needs(bid)['rows'][0]['note'],'Lời dặn trên web')
        wb=load_workbook(io.BytesIO(self.file(9)))
        wb['đặt hàng'].cell(3,7,'Lời dặn mới từ Excel')
        out=io.BytesIO();wb.save(out);wb.close()
        self.confirm(self.preview(bid,out.getvalue()))
        self.assertEqual(self.needs(bid)['rows'][0]['note'],'Lời dặn mới từ Excel')


if __name__=='__main__':unittest.main()
