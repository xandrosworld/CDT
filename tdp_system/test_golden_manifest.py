from __future__ import annotations

import io
import tempfile
import unittest
import zipfile
from pathlib import Path

from openpyxl import Workbook

try:
    from .golden_manifest import (
        ArtifactSpec,
        GoldenManifestError,
        build_manifest,
        inspect_docx,
        manifest_json,
    )
except ImportError:  # pragma: no cover - direct file invocation
    from golden_manifest import (  # type: ignore
        ArtifactSpec,
        GoldenManifestError,
        build_manifest,
        inspect_docx,
        manifest_json,
    )


def workbook_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Golden"
    sheet["A1"] = "PII-SHOULD-NOT-LEAK"
    sheet["B2"] = "=1+1"
    sheet.merge_cells("A3:B3")
    sheet.print_area = "A1:B5"
    sheet.print_title_rows = "1:2"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = "9"
    hidden = workbook.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def minimal_docx_bytes() -> bytes:
    document = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:body><w:p><w:r><w:t>SECRET-TEXT</w:t></w:r></w:p><w:tbl/></w:body></w:document>'
    )
    relationships = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Target="https://example.invalid" TargetMode="External" Type="link"/>'
        '</Relationships>'
    )
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w") as archive:
        archive.writestr("word/document.xml", document)
        archive.writestr("word/_rels/document.xml.rels", relationships)
        archive.writestr("word/media/image1.png", b"not-a-real-image")
    return stream.getvalue()


class GoldenManifestTests(unittest.TestCase):
    def test_manifest_is_deterministic_and_never_copies_cell_values(self):
        with tempfile.TemporaryDirectory(prefix="tdp_golden_manifest_") as temp_name:
            source = Path(temp_name) / "golden.xlsx"
            source.write_bytes(workbook_bytes())
            specs = (ArtifactSpec("fixture", "output_golden", (source,)),)

            first = manifest_json(build_manifest(specs))
            second = manifest_json(build_manifest(specs))

            self.assertEqual(first, second)
            self.assertNotIn("PII-SHOULD-NOT-LEAK", first)
            self.assertIn('"formula_count": 1', first)
            self.assertIn('"merged_ranges"', first)
            self.assertIn('"print_area"', first)
            self.assertIn('"state": "hidden"', first)

    def test_docx_metadata_does_not_copy_document_text(self):
        rendered = manifest_json({"topology": inspect_docx(minimal_docx_bytes())})
        self.assertNotIn("SECRET-TEXT", rendered)
        self.assertIn('"embedded_media_count": 1', rendered)
        self.assertIn('"external_relationship_count": 1', rendered)

    def test_missing_required_artifact_is_rejected(self):
        with tempfile.TemporaryDirectory(prefix="tdp_golden_manifest_") as temp_name:
            missing = Path(temp_name) / "missing.xlsx"
            with self.assertRaises(GoldenManifestError):
                build_manifest((ArtifactSpec("missing", "input_contract", (missing,)),))


if __name__ == "__main__":
    unittest.main()
