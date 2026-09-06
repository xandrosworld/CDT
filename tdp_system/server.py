from __future__ import annotations

import copy
import io
import hashlib
import ipaddress
import json
import math
import os
import re
import shutil
import socket
import sqlite3
import sys
import threading
import time
import unicodedata
import uuid
import webbrowser
import zipfile
from collections import Counter, defaultdict
from contextlib import contextmanager
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from difflib import SequenceMatcher
from pathlib import Path
from urllib.error import URLError
from urllib.parse import urlsplit
from urllib.request import urlopen

from flask import Flask, Response, jsonify, request, send_file, send_from_directory, session
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from waitress import serve
from werkzeug.exceptions import HTTPException
from werkzeug.utils import secure_filename

try:
    import order_worksheet
except ImportError:
    from . import order_worksheet

try:
    from automatic_backup import automatic_backup, backup_status, sqlite_snapshot, start_backup_worker
except ImportError:
    from .automatic_backup import automatic_backup, backup_status, sqlite_snapshot, start_backup_worker

try:
    from invoice_date_migration import repair_legacy_input_dates
except ImportError:
    from .invoice_date_migration import repair_legacy_input_dates

try:
    from seller_identity_catalog import CATALOG_FILENAME, sync_catalog, EXCLUDED_SELLERS, is_excluded_seller
except ImportError:
    from .seller_identity_catalog import CATALOG_FILENAME, sync_catalog, EXCLUDED_SELLERS, is_excluded_seller

try:
    from minvoice_client import MinvoiceClient, MinvoiceConfig, MinvoiceError
except ImportError:  # Allows importing as tdp_system.server in tests/tools.
    from .minvoice_client import MinvoiceClient, MinvoiceConfig, MinvoiceError

try:
    from msmi_client import MsmiClient, MsmiConfig
except ImportError:  # Allows importing as tdp_system.server in tests/tools.
    from .msmi_client import MsmiClient, MsmiConfig

try:
    from contract_modules import (
        PurchaseOrderApplyError,
        apply_purchase_order_preview,
        debt_period_payload,
        init_contract_schema,
        inventory_lookup,
        net_delivered,
        net_received,
        parse_purchase_order_workbook,
        purchase_business_row_key,
        purchase_line_changed,
        post_purchase_list_inventory,
        purchase_order_payload,
        register_contract_routes,
    )
except ImportError:
    from .contract_modules import (
        PurchaseOrderApplyError,
        apply_purchase_order_preview,
        debt_period_payload,
        init_contract_schema,
        inventory_lookup,
        net_delivered,
        net_received,
        parse_purchase_order_workbook,
        purchase_business_row_key,
        purchase_line_changed,
        post_purchase_list_inventory,
        purchase_order_payload,
        register_contract_routes,
    )

try:
    from invoice_workbench import init_invoice_workbench_schema, register_invoice_workbench_routes
except ImportError:
    from .invoice_workbench import init_invoice_workbench_schema, register_invoice_workbench_routes

try:
    from invoice_input_export import register_invoice_input_export_routes
except ImportError:
    from .invoice_input_export import register_invoice_input_export_routes

try:
    from invoice_mapping import register_invoice_mapping_routes
except ImportError:
    from .invoice_mapping import register_invoice_mapping_routes

try:
    from invoice_inventory import register_invoice_inventory_routes
except ImportError:
    from .invoice_inventory import register_invoice_inventory_routes

try:
    from invoice_valuation import register_invoice_valuation_routes
except ImportError:
    from .invoice_valuation import register_invoice_valuation_routes

try:
    from inventory_export import register_inventory_export_routes
except ImportError:
    from .inventory_export import register_inventory_export_routes

try:
    from inventory_period_close import (
        init_inventory_period_close_schema,
        register_inventory_period_close_routes,
    )
except ImportError:
    from .inventory_period_close import (
        init_inventory_period_close_schema,
        register_inventory_period_close_routes,
    )

try:
    from invoice_tax_export import InvoiceTaxExportError, export_invoice_drafts_zip
except ImportError:
    from .invoice_tax_export import InvoiceTaxExportError, export_invoice_drafts_zip

try:
    from bk_import import init_bk_import_schema, register_bk_import_routes
except ImportError:
    from .bk_import import init_bk_import_schema, register_bk_import_routes

try:
    from payable_ledger import (
        init_payable_ledger_schema,
        register_payable_ledger_routes,
        sync_payable_ledger,
    )
except ImportError:
    from .payable_ledger import (
        init_payable_ledger_schema,
        register_payable_ledger_routes,
        sync_payable_ledger,
    )

try:
    from payable_payments import (
        PayablePaymentError,
        create_payable_payment,
        init_payable_payment_schema,
        register_payable_payment_routes,
    )
except ImportError:
    from .payable_payments import (
        PayablePaymentError,
        create_payable_payment,
        init_payable_payment_schema,
        register_payable_payment_routes,
    )

try:
    from payable_export import register_payable_export_routes
except ImportError:
    from .payable_export import register_payable_export_routes

try:
    from receivable_ledger import (
        init_receivable_ledger_schema,
        register_receivable_ledger_routes,
        sync_receivable_ledger,
    )
except ImportError:
    from .receivable_ledger import (
        init_receivable_ledger_schema,
        register_receivable_ledger_routes,
        sync_receivable_ledger,
    )

try:
    from receivable_export import register_receivable_export_routes
except ImportError:
    from .receivable_export import register_receivable_export_routes

try:
    from daily_import_lifecycle import (
        DailyImportError,
        canonical_day_sheet,
        confirm_daily_import_scope,
        daily_payload_hash,
        daily_row_key,
        finalize_daily_workday,
        init_daily_import_schema,
        prepare_daily_import_version,
    )
except ImportError:
    from .daily_import_lifecycle import (
        DailyImportError,
        canonical_day_sheet,
        confirm_daily_import_scope,
        daily_payload_hash,
        daily_row_key,
        finalize_daily_workday,
        init_daily_import_schema,
        prepare_daily_import_version,
    )

try:
    from daily_workbook_import import analysis_state_hash, analyze_daily_workbook
except ImportError:
    from .daily_workbook_import import analysis_state_hash, analyze_daily_workbook

try:
    from daily_reference_import import init_daily_reference_schema, register_daily_reference_routes
except ImportError:
    from .daily_reference_import import init_daily_reference_schema, register_daily_reference_routes

try:
    from order_price_override import (
        direct_price_change_ids,
        init_order_price_override_schema,
        register_order_price_override_routes,
    )
except ImportError:
    from .order_price_override import (
        direct_price_change_ids,
        init_order_price_override_schema,
        register_order_price_override_routes,
    )

try:
    from quote_import import (
        init_quote_import_schema,
        quote_buy_price,
        quote_rows_for_contractor,
        quote_sell_price,
        register_quote_import_routes,
    )
except ImportError:
    from .quote_import import (
        init_quote_import_schema,
        quote_buy_price,
        quote_rows_for_contractor,
        quote_sell_price,
        register_quote_import_routes,
    )

try:
    from outgoing_substitution import (
        init_outgoing_substitution_schema,
        register_outgoing_substitution_routes,
    )
except ImportError:
    from .outgoing_substitution import (
        init_outgoing_substitution_schema,
        register_outgoing_substitution_routes,
    )

try:
    from quote_export import (
        QuoteExportError,
        build_contractor_quote_workbook,
        contractor_quote_filename,
        quote_recipient,
    )
except ImportError:
    from .quote_export import (
        QuoteExportError,
        build_contractor_quote_workbook,
        contractor_quote_filename,
        quote_recipient,
    )

try:
    from delivery_export import (
        DeliveryExportError,
        build_delivery_workbook,
        delivery_prices_visible,
    )
except ImportError:
    from .delivery_export import (
        DeliveryExportError,
        build_delivery_workbook,
        delivery_prices_visible,
    )

try:
    from purchase_summary_export import (
        PurchaseSummaryError,
        collect_purchase_summary_rows,
    )
except ImportError:
    from .purchase_summary_export import (
        PurchaseSummaryError,
        collect_purchase_summary_rows,
    )

try:
    from receipt_export import (
        build_purchase_documents_workbook,
        enrich_receipt_identity_rows,
    )
except ImportError:
    from .receipt_export import (
        build_purchase_documents_workbook,
        enrich_receipt_identity_rows,
    )

try:
    from report_export import (
        ReportExportError,
        build_monthly_report_workbook,
        collect_monthly_report_rows,
    )
except ImportError:
    from .report_export import (
        ReportExportError,
        build_monthly_report_workbook,
        collect_monthly_report_rows,
    )


FROZEN = bool(getattr(sys, "frozen", False))
BUNDLE_DIR = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent))
if FROZEN:
    APP_DIR = Path(sys.executable).resolve().parent
    ROOT = APP_DIR
    STATIC_DIR = BUNDLE_DIR / "static"
    SHARED_DIR = BUNDLE_DIR / "shared"
    MASTER_SOURCE = BUNDLE_DIR / "Em Thành.xlsx"
    OPENING_TEMPLATE_SOURCE = BUNDLE_DIR / "TĐK T8-2026.xlsx thụy.xlsx"
    TAX_TEMPLATE_DIR = BUNDLE_DIR / "tax_templates"
    PRINT_TEMPLATE_DIR = BUNDLE_DIR / "templates"
    BUNDLED_CONNECTOR_CONFIG = BUNDLE_DIR / "config" / "connector.env"
    SEED_DATABASE_SOURCE = BUNDLE_DIR / "seed" / "tdp_seed.sqlite3"
else:
    APP_DIR = Path(__file__).resolve().parent
    ROOT = APP_DIR.parent
    STATIC_DIR = APP_DIR / "static"
    SHARED_DIR = ROOT / "demo_tdp"
    MASTER_SOURCE = ROOT / "Em Thành.xlsx"
    OPENING_TEMPLATE_SOURCE = ROOT / "_HANDOFF" / "EXTERNAL_INPUTS" / "TĐK T8-2026.xlsx thụy.xlsx"
    TAX_TEMPLATE_DIR = ROOT / "bosung.30.8.26"
    PRINT_TEMPLATE_DIR = APP_DIR / "templates"
    BUNDLED_CONNECTOR_CONFIG = APP_DIR / "connector.env"
    SEED_DATABASE_SOURCE = APP_DIR / "tdp_seed.sqlite3"
DAILY_ORDER_TEMPLATE_SOURCE = PRINT_TEMPLATE_DIR / "daily_order_template.xlsx"
DATA_DIR = Path(os.environ.get("TDP_DATA_DIR", str(APP_DIR / "data"))).resolve()
EXPORT_DIR = Path(os.environ.get("TDP_EXPORT_DIR", str(APP_DIR / "exports"))).resolve()
DB_PATH = Path(os.environ.get("TDP_DB_PATH", str(DATA_DIR / "tdp.sqlite3")))

def ensure_writable_directory(path: Path, label: str):
    try:
        path.mkdir(parents=True, exist_ok=True)
        probe = path / f".tdp-write-{uuid.uuid4().hex}.tmp"
        with probe.open("x", encoding="utf-8") as handle:
            handle.write("ok")
        probe.unlink()
    except OSError as exc:
        raise RuntimeError(
            f"{label} không có quyền ghi: {path}. Hãy chuyển bộ cài sang thư mục người dùng "
            "hoặc cấu hình TDP_DATA_DIR/TDP_EXPORT_DIR."
        ) from exc


ensure_writable_directory(DATA_DIR, "Thư mục dữ liệu")
ensure_writable_directory(EXPORT_DIR, "Thư mục xuất file")


def install_seed_database_if_missing(
    seed_path: Path = SEED_DATABASE_SOURCE,
    target_path: Path = DB_PATH,
) -> bool:
    """Install the bundled customer snapshot once without overwriting user data."""

    if target_path.exists() or not seed_path.is_file():
        return False
    target_path.parent.mkdir(parents=True, exist_ok=True)
    temporary = target_path.with_name(
        f".{target_path.name}.installing-{uuid.uuid4().hex}.tmp"
    )
    connection = None
    try:
        shutil.copy2(seed_path, temporary)
        connection = sqlite3.connect(temporary)
        integrity = str(connection.execute("PRAGMA quick_check(1)").fetchone()[0])
        connection.close()
        connection = None
        if integrity.lower() != "ok":
            raise sqlite3.DatabaseError("Dữ liệu khởi tạo không vượt qua kiểm tra toàn vẹn")
        # An existing customer database always wins, including one created while
        # the bundled snapshot was being checked. Both paths are on the same disk.
        try:
            os.link(temporary, target_path)
        except FileExistsError:
            return False
    except Exception:
        if connection is not None:
            connection.close()
        temporary.unlink(missing_ok=True)
        raise
    finally:
        temporary.unlink(missing_ok=True)
    return True


def connector_config_paths() -> list[Path]:
    """Use the bundled configuration as fallback and an adjacent file as override."""

    paths = [BUNDLED_CONNECTOR_CONFIG] if FROZEN else []
    for candidate in (ROOT / ".env", APP_DIR / ".env"):
        if candidate not in paths:
            paths.append(candidate)
    return paths

NAVY = "17324D"
TEAL = "087F73"
PALE = "F5F8FB"
WHITE = "FFFFFF"
GRAY = "5E7083"
RED = "C2413A"

INVOICE_HEADERS = [
    "Mã hàng", "Tên hàng", "Đơn vị tính", "Số lượng", "Đơn giá",
    "Cộng tiền hàng", "%CK", "Tiền CK", "Tiền trước thuế", "% VAT",
    "Tiền thuế GTGT", "Tổng tiền", "Tính chất",
]
MASTER_FORMAT_VERSION = "2"

app = Flask(__name__, static_folder=str(STATIC_DIR), static_url_path="/static")
app.config["MAX_CONTENT_LENGTH"] = 40 * 1024 * 1024
PENDING_IMPORTS = {}
PENDING_IMPORT_LOCK = threading.Lock()
ORDER_IMPORT_MAX_BYTES = 20 * 1024 * 1024
ORDER_IMPORT_MAX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
ORDER_IMPORT_MAX_ENTRIES = 5_000
ORDER_IMPORT_MAX_ROWS = 100_000


@app.after_request
def prevent_stale_application_assets(response):
    """A replaced EXE/source must never leave the browser running old UI code."""
    if request.path == "/" or request.path.startswith("/static/"):
        response.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
        response.headers["Pragma"] = "no-cache"
        response.headers["Expires"] = "0"
    return response


@app.before_request
def reject_cross_origin_mutations():
    """Reject DNS-rebinding hosts and browser cross-site writes."""
    host_name = (urlsplit("//" + request.host).hostname or "").casefold()
    trusted_hosts = {"localhost", "127.0.0.1", "::1"}
    configured_hosts = clean_text(os.environ.get("TDP_TRUSTED_HOSTS", ""))
    trusted_hosts.update(
        item.strip().casefold() for item in configured_hosts.split(",") if item.strip()
    )
    allow_lan = clean_text(os.environ.get("TDP_ALLOW_LAN", "")).lower() in {
        "1", "true", "yes", "co", "có",
    }
    if allow_lan:
        trusted_hosts.update({
            clean_text(local_ip()).casefold(),
            clean_text(socket.gethostname()).casefold(),
            clean_text(socket.getfqdn()).casefold(),
        })
    if host_name not in trusted_hosts:
        return jsonify({"ok": False, "error": "Tên máy truy cập không nằm trong danh sách tin cậy"}), 403
    if request.method not in {"POST", "PUT", "PATCH", "DELETE"}:
        return None
    origin = clean_text(request.headers.get("Origin")).rstrip("/")
    if origin and origin != request.host_url.rstrip("/"):
        return jsonify({"ok": False, "error": "Yêu cầu ghi dữ liệu khác nguồn đã bị chặn"}), 403
    return None


def now_iso() -> str:
    return datetime.now().replace(microsecond=0).isoformat(sep=" ")


def clean_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y")
    return str(value).strip()


def slug(value) -> str:
    text = clean_text(value).lower()
    text = unicodedata.normalize("NFD", text)
    text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn")
    return re.sub(r"[^a-z0-9]+", "", text)


def canonical_party_code(conn, party_type: str, value) -> str:
    """Resolve free-text ledger input to one canonical debt key.

    Debt aggregation groups by the stored key exactly.  Accepting ``hatran``
    beside ``HATRAN`` would create a phantom second account, so every ledger
    write resolves against master data and existing charge sources first.
    """

    raw = clean_text(value)
    if not raw:
        raise ValueError("Thiếu mã nhà thầu/nhà cung cấp")
    candidates = {}
    case_keys = {}

    def remember(code, name=None):
        canonical = clean_text(code)
        if not canonical:
            return
        # Case-only variants are the same ledger key.  Prefer the first value
        # encountered (master tables are read before transactional sources)
        # and retain every spelling as an alias.
        preferred = case_keys.setdefault(canonical.casefold(), canonical)
        candidates.setdefault(preferred, set()).update(
            item for item in (canonical, clean_text(name)) if item
        )

    if party_type == "contractor":
        for row in conn.execute("SELECT code,name FROM contractors"):
            remember(row["code"], row["name"])
        for row in conn.execute(
            "SELECT DISTINCT contractor FROM orders WHERE TRIM(COALESCE(contractor,''))!=''"
        ):
            remember(row["contractor"])
    elif party_type == "supplier":
        for row in conn.execute("SELECT code,name FROM suppliers"):
            remember(row["code"], row["name"])
        for sql in (
            "SELECT DISTINCT supplier FROM products WHERE TRIM(COALESCE(supplier,''))!=''",
            "SELECT DISTINCT supplier FROM orders WHERE TRIM(COALESCE(supplier,''))!=''",
            "SELECT DISTINCT supplier FROM historical_payable_lines WHERE TRIM(COALESCE(supplier,''))!=''",
        ):
            for row in conn.execute(sql):
                remember(row["supplier"])
    else:
        raise ValueError("Nhóm công nợ không hợp lệ")

    exact = [code for code in candidates if code.casefold() == raw.casefold()]
    if len(exact) == 1:
        return exact[0]
    wanted = slug(raw)
    matches = [
        code for code, aliases in candidates.items()
        if wanted and any(slug(alias) == wanted for alias in aliases)
    ]
    matches = sorted(set(matches))
    if len(matches) == 1:
        return matches[0]
    label = "nhà thầu" if party_type == "contractor" else "nhà cung cấp"
    if not matches:
        raise ValueError(f"Không tìm thấy {label} {raw!r}; hãy chọn đúng mã đã có trong hệ thống")
    raise ValueError(f"{label.capitalize()} {raw!r} khớp nhiều mã; hãy nhập đúng mã duy nhất")


def number_value(value, default=0.0):
    if value in (None, ""):
        return default
    if isinstance(value, (int, float)):
        return float(value)
    text = clean_text(value).replace("₫", "").replace("đ", "").replace(" ", "")
    if not text:
        return default
    # Vietnamese input commonly uses dot for thousands and comma for decimals.
    if "," in text:
        text = text.replace(".", "").replace(",", ".")
    elif text.count(".") > 1:
        text = text.replace(".", "")
    try:
        return float(text)
    except ValueError:
        return default


def valid_iso_date(value, label="Ngày") -> str:
    """Return a strict calendar date or raise a user-safe validation error."""
    text = clean_text(value)
    try:
        return datetime.strptime(text, "%Y-%m-%d").date().isoformat()
    except (TypeError, ValueError):
        raise ValueError(f"{label} phải có dạng YYYY-MM-DD và là ngày hợp lệ") from None


def finite_number(value, label: str, *, default=None) -> float:
    """Parse a numeric request field without silently turning bad input into zero."""
    if value in (None, ""):
        if default is not None:
            return float(default)
        raise ValueError(f"Thiếu {label}")
    parsed = number_value(value, float("nan"))
    if not math.isfinite(parsed):
        raise ValueError(f"{label} phải là số hợp lệ")
    return float(parsed)


def audit_event(conn, event_type: str, *, entity_type="", entity_id="", metadata=None):
    """Write a compact business audit record without request bodies or credentials."""
    conn.execute(
        "INSERT INTO audit_log(event_type,entity_type,entity_id,status,message,metadata_json,created_at) "
        "VALUES(?,?,?,?,?,?,?)",
        (
            event_type,
            clean_text(entity_type),
            clean_text(entity_id),
            "ok",
            "",
            json.dumps(metadata or {}, ensure_ascii=False, sort_keys=True, separators=(",", ":")),
            now_iso(),
        ),
    )


def normalize_tax(value):
    text = clean_text(value).upper()
    if text in {"KKKNT", "KHÔNG KÊ KHAI", "KHONG KE KHAI", "K"}:
        return "KKKNT"
    if text in {"KCT", "KHÔNG CHỊU THUẾ", "KHONG CHIU THUE"}:
        return "KCT"
    if text in {"-2", "-2.0"}:
        return "KKKNT"
    if text in {"-1", "-1.0"}:
        return "KCT"
    if not text:
        return ""
    if text.endswith("%"):
        num = number_value(text[:-1], float("nan")) / 100
    else:
        num = number_value(value, float("nan"))
    if num > 1:
        num /= 100
    if not math.isfinite(num):
        return "INVALID"
    num = round(num, 4)
    if num not in {0.0, 0.05, 0.08, 0.1}:
        return "INVALID"
    return num


def tax_factor(value) -> float:
    return 1.0 if clean_text(value).upper() in {"KKKNT", "KCT"} else 1.0 + number_value(value)


def safe_sheet_name(value: str) -> str:
    return re.sub(r"[\\/*?:\[\]]", "-", clean_text(value) or "Sheet")[:31]


try:
    from physical_inventory import init_physical_inventory_schema, register_physical_inventory_routes, physical_order_revision
except ImportError:
    from .physical_inventory import init_physical_inventory_schema, register_physical_inventory_routes, physical_order_revision


@contextmanager
def db():
    conn = sqlite3.connect(DB_PATH, timeout=20)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    conn.execute("PRAGMA journal_mode = WAL")
    try:
        yield conn
        conn.commit()
    finally:
        conn.close()


SCHEMA = """
CREATE TABLE IF NOT EXISTS settings (
    key TEXT PRIMARY KEY,
    value TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS contractors (
    code TEXT PRIMARY KEY,
    name TEXT NOT NULL,
    price_group TEXT,
    pricing_mode TEXT NOT NULL DEFAULT 'group'
);
CREATE TABLE IF NOT EXISTS kitchens (
    code TEXT PRIMARY KEY,
    contractor TEXT,
    name TEXT,
    address TEXT,
    show_price INTEGER NOT NULL DEFAULT 0
);
CREATE TABLE IF NOT EXISTS suppliers (
    code TEXT PRIMARY KEY,
    name TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS people (
    name TEXT PRIMARY KEY,
    cccd TEXT,
    issue_date TEXT,
    issue_place TEXT,
    address TEXT
);
CREATE TABLE IF NOT EXISTS products (
    code TEXT PRIMARY KEY,
    name TEXT NOT NULL,
    unit TEXT,
    tax TEXT,
    supplier TEXT,
    buy_price REAL NOT NULL DEFAULT 0,
    purchase_list INTEGER NOT NULL DEFAULT 0,
    seller TEXT,
    cccd TEXT
);
CREATE INDEX IF NOT EXISTS idx_products_name ON products(name);
CREATE TABLE IF NOT EXISTS product_prices (
    product_code TEXT NOT NULL,
    price_group TEXT NOT NULL,
    price_text TEXT,
    price_value REAL,
    PRIMARY KEY(product_code, price_group)
);
CREATE TABLE IF NOT EXISTS batches (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    work_date TEXT NOT NULL,
    source_name TEXT,
    status TEXT NOT NULL DEFAULT 'draft',
    created_at TEXT NOT NULL,
    approved_at TEXT
);
CREATE TABLE IF NOT EXISTS order_import_receipts (
    import_key TEXT PRIMARY KEY,
    source_hash TEXT NOT NULL,
    source_name TEXT NOT NULL,
    work_date TEXT NOT NULL,
    selected_sheets_json TEXT NOT NULL,
    batch_id INTEGER NOT NULL REFERENCES batches(id) ON DELETE CASCADE,
    created_at TEXT NOT NULL
);
CREATE INDEX IF NOT EXISTS idx_order_import_receipts_batch
    ON order_import_receipts(batch_id);
CREATE TABLE IF NOT EXISTS orders (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    batch_id INTEGER NOT NULL REFERENCES batches(id) ON DELETE CASCADE,
    work_date TEXT NOT NULL,
    contractor TEXT,
    kitchen TEXT,
    product_code TEXT,
    product_name TEXT,
    qty REAL NOT NULL DEFAULT 0,
    actual_received REAL NOT NULL DEFAULT 0,
    actual_delivered REAL NOT NULL DEFAULT 0,
    unit TEXT,
    supplier TEXT,
    buy_price REAL NOT NULL DEFAULT 0,
    sell_price REAL NOT NULL DEFAULT 0,
    tax TEXT,
    invoice_nature TEXT NOT NULL DEFAULT '1',
    purchase_list INTEGER NOT NULL DEFAULT 0,
    seller TEXT,
    cccd TEXT,
    note TEXT,
    source_sheet TEXT,
    source_row INTEGER,
    errors TEXT NOT NULL DEFAULT '[]',
    warnings TEXT NOT NULL DEFAULT '[]',
    updated_at TEXT NOT NULL
);
CREATE INDEX IF NOT EXISTS idx_orders_batch ON orders(batch_id);
CREATE TABLE IF NOT EXISTS balances (
    party_type TEXT NOT NULL,
    party_code TEXT NOT NULL,
    opening REAL NOT NULL DEFAULT 0,
    PRIMARY KEY(party_type, party_code)
);
CREATE TABLE IF NOT EXISTS payments (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    payment_date TEXT NOT NULL,
    kind TEXT NOT NULL,
    party_type TEXT NOT NULL,
    party_code TEXT NOT NULL,
    amount REAL NOT NULL,
    note TEXT,
    created_at TEXT NOT NULL,
    method TEXT NOT NULL DEFAULT '',
    reference_code TEXT NOT NULL DEFAULT '',
    status TEXT NOT NULL DEFAULT 'posted',
    request_key TEXT NOT NULL DEFAULT '',
    request_hash TEXT NOT NULL DEFAULT '',
    revision INTEGER NOT NULL DEFAULT 1,
    updated_at TEXT NOT NULL DEFAULT '',
    reversed_at TEXT,
    reversal_reason TEXT NOT NULL DEFAULT ''
);
"""


def setting_get(conn, key: str, default=None):
    row = conn.execute("SELECT value FROM settings WHERE key=?", (key,)).fetchone()
    return row["value"] if row else default


def setting_set(conn, key: str, value):
    conn.execute(
        "INSERT INTO settings(key,value) VALUES(?,?) "
        "ON CONFLICT(key) DO UPDATE SET value=excluded.value",
        (key, str(value)),
    )


def pre_migration_backup():
    """Create a verified SQLite backup before any startup DDL can commit."""

    if not DB_PATH.is_file() or DB_PATH.stat().st_size == 0:
        return None
    backup_dir = DATA_DIR / "migration_backups"
    backup_dir.mkdir(parents=True, exist_ok=True)
    target = backup_dir / (
        f"tdp_pre_migration_{datetime.now():%Y%m%d_%H%M%S}_{uuid.uuid4().hex[:10]}.sqlite3"
    )
    try:
        sqlite_snapshot(DB_PATH, target)
    except (OSError, sqlite3.Error, TimeoutError) as exc:
        raise RuntimeError("Không tạo được bản sao an toàn trước khi nâng cấp dữ liệu") from exc
    backups = sorted(
        (path for path in backup_dir.glob("tdp_pre_migration_*.sqlite3")
         if re.fullmatch(r"tdp_pre_migration_\d{8}_\d{6}_[a-f0-9]{10}\.sqlite3", path.name)
         and not path.is_symlink() and path.resolve().parent == backup_dir.resolve()),
        key=lambda path: path.stat().st_mtime,
        reverse=True,
    )
    for old in backups[5:]:
        if not old.is_symlink() and old.resolve().parent == backup_dir.resolve():
            try:
                old.unlink(missing_ok=True)
            except OSError:
                pass  # A verified new backup exists; failed pruning is not data loss.
    return target


def init_database(*, sync_master=True):
    pre_migration_backup()
    with db() as conn:
        conn.executescript(SCHEMA)
        init_contract_schema(conn, opening_template_path=OPENING_TEMPLATE_SOURCE)
        init_daily_import_schema(conn)
        order_worksheet.init_schema(conn)
        init_daily_reference_schema(conn)
        init_order_price_override_schema(conn)
        init_quote_import_schema(conn)
        init_outgoing_substitution_schema(conn)
        init_invoice_workbench_schema(conn)
        init_inventory_period_close_schema(conn)
        init_bk_import_schema(conn)
        init_payable_ledger_schema(conn)
        init_payable_payment_schema(conn)
        init_receivable_ledger_schema(conn)
        order_columns = {row["name"] for row in conn.execute("PRAGMA table_info(orders)")}
        if "warnings" not in order_columns:
            conn.execute("ALTER TABLE orders ADD COLUMN warnings TEXT NOT NULL DEFAULT '[]'")
        if "invoice_nature" not in order_columns:
            conn.execute("ALTER TABLE orders ADD COLUMN invoice_nature TEXT NOT NULL DEFAULT '1'")
        init_physical_inventory_schema(conn)
        # The pre-migration backup above precedes all historical date repairs.
        # Posted inventory remains unchanged and is listed for reconciliation.
        repair_legacy_input_dates(conn, timestamp=now_iso())
        defaults = {
            "company": "CÔNG TY TNHH THỰC PHẨM THÀNH ĐẠT PHÁT",
            "purchase_rate": "0.95",
            "master_version": "0",
        }
        for key, value in defaults.items():
            if setting_get(conn, key) is None:
                setting_set(conn, key, value)
        sync_payable_ledger(conn, timestamp=now_iso())
        sync_receivable_ledger(conn, timestamp=now_iso())
    if sync_master:
        sync_master_if_needed()
    auto_backup()


def auto_backup():
    return automatic_backup(DB_PATH, DATA_DIR)


def sync_master_if_needed(force=False):
    identity_catalog = PRINT_TEMPLATE_DIR / CATALOG_FILENAME
    if identity_catalog.exists():
        # Independent of master_version: existing installations also need this correction.
        with db() as conn:
            sync_catalog(conn, identity_catalog)
    if not MASTER_SOURCE.exists():
        return
    version = f"{MASTER_FORMAT_VERSION}-{int(MASTER_SOURCE.stat().st_mtime)}"
    with db() as conn:
        if not force and setting_get(conn, "master_version") == version:
            return

    wb = load_workbook(MASTER_SOURCE, data_only=True, read_only=False)
    with db() as conn:
        # People / CCCD
        if "CCCD" in wb.sheetnames and not identity_catalog.exists():
            ws = wb["CCCD"]
            for row in range(2, ws.max_row + 1):
                name = clean_text(ws.cell(row, 2).value)
                if not name:
                    continue
                issue = ws.cell(row, 4).value
                issue_date = issue.strftime("%d/%m/%Y") if isinstance(issue, datetime) else clean_text(issue)
                conn.execute(
                    "INSERT INTO people(name,cccd,issue_date,issue_place,address) VALUES(?,?,?,?,?) "
                    "ON CONFLICT(name) DO UPDATE SET cccd=excluded.cccd,issue_date=excluded.issue_date,"
                    "issue_place=excluded.issue_place",
                    (name, clean_text(ws.cell(row, 3).value), issue_date,
                     clean_text(ws.cell(row, 5).value), ""),
                )

        # Contractor and kitchen mapping.
        if "T.chiếu" in wb.sheetnames:
            ws = wb["T.chiếu"]
            for row in range(3, ws.max_row + 1):
                contractor = clean_text(ws.cell(row, 2).value).upper()
                kitchen = clean_text(ws.cell(row, 3).value).upper()
                if not kitchen:
                    continue
                if contractor:
                    mode = "daily" if contractor in {"GIANHAPTAY", "YLKHAN"} else "group"
                    conn.execute(
                        "INSERT INTO contractors(code,name,price_group,pricing_mode) VALUES(?,?,?,?) "
                        "ON CONFLICT(code) DO UPDATE SET name=excluded.name",
                        (contractor, contractor, contractor, mode),
                    )
                conn.execute(
                    "INSERT INTO kitchens(code,contractor,name,address,show_price) VALUES(?,?,?,?,?) "
                    "ON CONFLICT(code) DO UPDATE SET contractor=excluded.contractor,name=excluded.name,"
                    "address=excluded.address",
                    (kitchen, contractor, clean_text(ws.cell(row, 4).value) or kitchen,
                     clean_text(ws.cell(row, 5).value), 1 if kitchen == "NHUAHP" else 0),
                )

        # Product catalog, supplier and all price groups.
        if "BÁO GIÁ" in wb.sheetnames:
            ws = wb["BÁO GIÁ"]
            price_cols = {}
            # Only columns L:R are the seven real selling-price groups.
            # Columns after R repeat group names for margin calculations.
            for col in range(12, min(ws.max_column, 18) + 1):
                header = clean_text(ws.cell(2, col).value).upper()
                if header and not header.startswith("=") and slug(header) not in {"them"}:
                    price_cols[col] = header
                    mode = "daily" if header in {"GIANHAPTAY", "YLKHAN"} else "group"
                    conn.execute(
                        "INSERT INTO contractors(code,name,price_group,pricing_mode) VALUES(?,?,?,?) "
                        "ON CONFLICT(code) DO NOTHING",
                        (header, header, header, mode),
                    )
            for row in range(4, ws.max_row + 1):
                code = clean_text(ws.cell(row, 3).value).upper()
                name = clean_text(ws.cell(row, 4).value)
                if not code or not name:
                    continue
                supplier = clean_text(ws.cell(row, 6).value)
                seller = clean_text(ws.cell(row, 9).value)
                cccd_row = conn.execute("SELECT cccd FROM people WHERE name=?", (seller,)).fetchone()
                cccd = cccd_row["cccd"] if cccd_row else ""
                if supplier:
                    conn.execute(
                        "INSERT INTO suppliers(code,name) VALUES(?,?) ON CONFLICT(code) DO NOTHING",
                        (supplier, supplier),
                    )
                conn.execute(
                    "INSERT INTO products(code,name,unit,tax,supplier,buy_price,purchase_list,seller,cccd) "
                    "VALUES(?,?,?,?,?,?,?,?,?) ON CONFLICT(code) DO UPDATE SET "
                    "name=excluded.name,unit=excluded.unit,tax=excluded.tax,"
                    "supplier=COALESCE(NULLIF(products.supplier,''),excluded.supplier),"
                    "buy_price=CASE WHEN excluded.buy_price>0 THEN excluded.buy_price ELSE products.buy_price END,"
                    "purchase_list=excluded.purchase_list,seller=excluded.seller,cccd=excluded.cccd",
                    (code, name, clean_text(ws.cell(row, 10).value),
                     clean_text(normalize_tax(ws.cell(row, 11).value)), supplier,
                     number_value(ws.cell(row, 5).value), 1 if slug(ws.cell(row, 8).value) == "bk" else 0,
                     seller, cccd),
                )
                for col, group in price_cols.items():
                    raw = ws.cell(row, col).value
                    text = clean_text(raw)
                    value = number_value(raw, None) if isinstance(raw, (int, float)) else None
                    conn.execute(
                        "INSERT INTO product_prices(product_code,price_group,price_text,price_value) VALUES(?,?,?,?) "
                        "ON CONFLICT(product_code,price_group) DO UPDATE SET "
                        "price_text=excluded.price_text,price_value=excluded.price_value",
                        (code, group, text, value),
                    )
        setting_set(conn, "master_version", version)
        setting_set(conn, "master_synced_at", now_iso())


def rows_dict(cursor):
    return [dict(row) for row in cursor.fetchall()]


def product_lookup(conn):
    by_code = {}
    by_name = defaultdict(list)
    for row in conn.execute("SELECT * FROM products"):
        item = dict(row)
        by_code[item["code"].upper()] = item
        by_name[slug(item["name"])].append(item)
    return by_code, by_name


def get_sell_price(conn, product_code: str, contractor: str, work_date: str = ""):
    contractor_row = conn.execute(
        "SELECT price_group,pricing_mode FROM contractors WHERE code=?", (contractor,)
    ).fetchone()
    if not contractor_row or contractor_row["pricing_mode"] == "daily":
        return 0.0, "Giá theo ngày: cần nhập giá bán"
    quote = quote_sell_price(conn, product_code, contractor, work_date)
    if quote["has_version"]:
        if quote["value"] is not None:
            return float(quote["value"]), ""
        return 0.0, quote["message"]
    row = conn.execute(
        "SELECT price_text,price_value FROM product_prices WHERE product_code=? AND price_group=?",
        (product_code, contractor_row["price_group"] or contractor),
    ).fetchone()
    if not row:
        return 0.0, "Chưa có giá cho nhóm nhà thầu"
    if row["price_value"] is not None and row["price_value"] > 0:
        return float(row["price_value"]), ""
    text = clean_text(row["price_text"])
    return 0.0, text or "Chưa có giá cho nhóm nhà thầu"


HEADER_ALIASES = {
    "contractor": {"nhathau", "nhomkhachhang"},
    "supplier": {"ncc", "nhacungcap", "chon ncc", "chonncc"},
    "product_code": {"mahang", "mavt", "mavattu"},
    "kitchen": {"mabep", "tenbepormabep"},
    "product_name": {"tenhang", "tenvt", "tenhanghoa", "tenthanhdatphat"},
    "qty": {"khoiluong", "soluong", "sldat", "sl"},
    "unit": {"dvt", "donvitinh"},
    "buy_price": {"giamua", "dongiamua"},
    "sell_price": {"giaban", "dongia", "dongiaban"},
    "tax": {"thue", "thuegtgt", "thuesuat"},
    "invoice_nature": {"tinhchat", "tinhchathanghoa", "loaibong", "nature"},
    "cccd": {"cccd", "socmtnhandan", "socccd"},
    "date": {"ngaythang", "ngay", "ngaygiao"},
    "note": {"ghichu", "ghichudathang"},
    "seller": {"tenlambangke", "nguoiban"},
    "purchase_list": {"bangke", "bk"},
    "actual_received": {"thucnhan", "soluongthucnhan"},
    "actual_delivered": {"thucgiao", "soluongthucgiao"},
    "damaged_qty": {"hanghong", "soluonghong", "slhong"},
    "supplier_return_qty": {"trancc", "tranhacungcap", "soluongtrancc"},
    "customer_return_qty": {"khachtra", "khachhangtra", "soluongkhachtra"},
}


def canonical_header(value):
    key = slug(value)
    for field, aliases in HEADER_ALIASES.items():
        if key in {slug(x) for x in aliases}:
            return field
    text = clean_text(value).lower()
    if "tên hàng" in text:
        return "product_name"
    if "khối" in text and "lượng" in text:
        return "qty"
    if "giá bán" in text:
        return "sell_price"
    if "giá mua" in text:
        return "buy_price"
    if "thuế" in text:
        return "tax"
    return None


def detect_header(ws):
    best = None
    for row in range(1, min(ws.max_row, 15) + 1):
        mapping = {}
        for col in range(1, min(ws.max_column, 40) + 1):
            field = canonical_header(ws.cell(row, col).value)
            if field and (field not in mapping or (
                field == 'supplier' and slug(ws.cell(row, mapping[field]).value) == 'chonncc'
                and slug(ws.cell(row, col).value) in {'ncc', 'nhacungcap'}
            )):
                mapping[field] = col
        score = sum(key in mapping for key in ("product_name", "qty", "kitchen", "product_code"))
        score += min(len(mapping), 8) / 10
        if best is None or score > best[0]:
            best = (score, row, mapping)
    if not best or best[0] < 2:
        return None, {}
    mapping = best[2]
    required = "qty" in mapping and "kitchen" in mapping and (
        "product_name" in mapping or "product_code" in mapping
    )
    if not required:
        return None, {}
    return best[1], mapping


def display_date(value, fallback: str):
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    text = clean_text(value)
    for fmt in ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, fmt).strftime("%Y-%m-%d")
        except ValueError:
            pass
    return fallback


def display_date_vn(value) -> str:
    """Format an operational ISO date for people-facing exports."""

    text = clean_text(value)
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, fmt).strftime("%d/%m/%Y")
        except ValueError:
            pass
    return text


def resolve_order(conn, raw: dict, fallback_date: str, by_code, by_name):
    errors = []
    warnings = []

    try:
        batch_date = valid_iso_date(fallback_date, "Ngày phiên đơn")
    except ValueError:
        batch_date = date.today().isoformat()
        errors.append("Ngày phiên đơn không hợp lệ")
    supplied_work_date = raw.get("work_date")
    if supplied_work_date not in (None, ""):
        parsed_order_date = display_date(supplied_work_date, "")
        if not parsed_order_date:
            errors.append("Ngày dòng đơn không hợp lệ")
        elif parsed_order_date != batch_date:
            errors.append("Ngày dòng đơn phải trùng ngày của phiên đơn")

    def order_number(value, label, default=0.0):
        parsed = (
            number_value(value, math.nan)
            if value not in (None, "")
            else float(default)
        )
        if not math.isfinite(parsed):
            errors.append(f"{label} phải là số hợp lệ và hữu hạn")
            return float(default) if math.isfinite(float(default)) else 0.0
        return parsed

    code = clean_text(raw.get("product_code")).upper()
    name = clean_text(raw.get("product_name"))
    product = by_code.get(code)
    if not product and name:
        matches = by_name.get(slug(name), [])
        if len(matches) == 1:
            product = matches[0]
            code = product["code"]
    if product and not name:
        name = product["name"]

    kitchen = clean_text(raw.get("kitchen")).upper()
    contractor = clean_text(raw.get("contractor")).upper()
    kitchen_row = conn.execute("SELECT * FROM kitchens WHERE code=?", (kitchen,)).fetchone()
    if not contractor and kitchen_row:
        contractor = clean_text(kitchen_row["contractor"]).upper()

    supplier = clean_text(raw.get("supplier"))
    if supplier.lower() in {"-", "kiểm tra", "chon ncc", "chọn ncc"}:
        supplier = ""
    if not supplier and product:
        supplier = clean_text(product["supplier"])

    qty = order_number(raw.get("qty"), "Số lượng")
    actual_received = order_number(raw.get("actual_received"), "Số lượng thực nhận", qty)
    actual_delivered = order_number(raw.get("actual_delivered"), "Số lượng thực giao", qty)
    damaged_qty = order_number(raw.get("damaged_qty"), "Số lượng hỏng")
    supplier_return_qty = order_number(raw.get("supplier_return_qty"), "Số lượng trả NCC")
    customer_return_qty = order_number(raw.get("customer_return_qty"), "Số lượng khách trả")
    unit = clean_text(raw.get("unit")) or (clean_text(product["unit"]) if product else "")
    requested_buy_price = order_number(raw.get("buy_price"), "Giá mua")
    period_buy_price = quote_buy_price(conn, code, batch_date, contractor) if code else {
        "has_version": False, "value": None, "allow_actual_fallback": True,
    }
    if period_buy_price["has_version"]:
        if period_buy_price["value"] is not None:
            # A non-blank price in the confirmed quotation is authoritative for
            # this period, even when the daily workbook carries a stale value.
            buy_price = float(period_buy_price["value"])
        elif period_buy_price["allow_actual_fallback"]:
            # A blank quotation price may be supplemented only by the actual
            # purchase value supplied on the operational row.  Do not revive a
            # previous-period catalogue value here.
            buy_price = requested_buy_price
        else:
            buy_price = 0.0
    else:
        buy_price = requested_buy_price
        if buy_price <= 0 and product:
            buy_price = order_number(product["buy_price"], "Giá mua trong danh mục")
    raw_invoice_nature = clean_text(raw.get("invoice_nature")) or "1"
    requested_sell_price = order_number(raw.get("sell_price"), "Giá bán")
    promotion_marker = "khuyenmai" in slug(
        f"{name} {clean_text(raw.get('note'))}"
    )
    is_promotion = raw_invoice_nature == "2" or (promotion_marker and requested_sell_price <= 0)
    invoice_nature = "2" if is_promotion else "1"
    price_message = ""
    period_sell_price = (
        quote_sell_price(conn, code, contractor, batch_date)
        if code and contractor and not is_promotion
        else {"has_version": False, "applicable": False, "value": None, "message": ""}
    )
    if period_sell_price["has_version"] and period_sell_price["applicable"]:
        # Once a quotation exists for the order's period, group-priced orders
        # must use that version rather than a possibly different embedded value.
        sell_price = float(period_sell_price["value"] or 0)
        price_message = period_sell_price["message"]
    else:
        sell_price = requested_sell_price
    if sell_price <= 0 and code and contractor and not is_promotion and not period_sell_price["has_version"]:
        sell_price, price_message = get_sell_price(conn, code, contractor, batch_date)
        if not math.isfinite(sell_price):
            errors.append("Giá bán trong bảng giá phải là số hữu hạn")
            sell_price = 0
    raw_tax = raw.get("tax")
    if raw_tax in (None, "") and product:
        tax = clean_text(normalize_tax(product["tax"]))
    else:
        tax = clean_text(normalize_tax(raw_tax))
    if not tax:
        tax = "KKKNT"

    purchase_list_raw = clean_text(raw.get("purchase_list")).lower()
    purchase_list = 1 if purchase_list_raw in {"bk", "x", "1", "true", "có", "co"} else 0
    if product and raw.get("purchase_list") in (None, ""):
        purchase_list = int(product["purchase_list"] or 0)
    seller = clean_text(raw.get("seller")) or (clean_text(product["seller"]) if product else "")
    cccd = clean_text(raw.get("cccd")) or (clean_text(product["cccd"]) if product else "")

    if raw_invoice_nature not in {"", "1", "2"}:
        errors.append("Tính chất hóa đơn chỉ nhận 1 (hàng hóa) hoặc 2 (khuyến mại)")
    if is_promotion and requested_sell_price > 0:
        errors.append("Dòng khuyến mại phải để giá bán bằng 0")
    if buy_price <= 0 and sell_price > 0 and (supplier.lower() == "bk" or purchase_list):
        rate = order_number(setting_get(conn, "purchase_rate", "0.95"), "Tỷ lệ giá mua", 0.95)
        buy_price = round(sell_price * rate)
        warnings.append(f"Giá mua tạm tính {rate:.0%} giá bán – cần xác nhận")
    if not kitchen:
        errors.append("Thiếu mã bếp")
    elif not kitchen_row:
        errors.append("Mã bếp chưa có trong danh mục")
    if not contractor:
        errors.append("Chưa xác định nhà thầu")
    if not code:
        errors.append("Chưa tìm thấy mã hàng")
    if not name:
        errors.append("Thiếu tên hàng")
    if qty <= 0:
        errors.append("Số lượng phải lớn hơn 0")
    if actual_received < 0:
        errors.append("Số lượng thực nhận không được âm")
    if actual_delivered < 0:
        errors.append("Số lượng thực giao không được âm")
    if damaged_qty < 0 or supplier_return_qty < 0 or customer_return_qty < 0:
        errors.append("Số lượng hỏng/trả lại không được âm")
    if tax == "INVALID":
        errors.append("Thuế suất chỉ nhận 0%, 5%, 8%, 10%, KCT hoặc KKKNT")
    if not supplier:
        errors.append("Thiếu nhà cung cấp")
    if buy_price <= 0:
        errors.append("Thiếu giá mua")
    if sell_price <= 0 and not is_promotion:
        errors.append(price_message or "Thiếu giá bán")
    if sell_price > 0 and buy_price > sell_price and not is_promotion:
        warnings.append("Giá bán thấp hơn giá mua – cần xác nhận bán lỗ")
    if purchase_list and not cccd:
        warnings.append("Hàng bảng kê chưa có CCCD – người dùng bổ sung sau khi cần lập bảng kê")
    if damaged_qty + supplier_return_qty > actual_received:
        errors.append("Hàng hỏng + trả NCC không được vượt số thực nhận")
    if customer_return_qty > actual_delivered:
        errors.append("Khách trả không được vượt số thực giao")

    return {
        "work_date": batch_date,
        "contractor": contractor,
        "kitchen": kitchen,
        "product_code": code,
        "product_name": name,
        "qty": qty,
        "actual_received": actual_received,
        "actual_delivered": actual_delivered,
        "damaged_qty": damaged_qty,
        "supplier_return_qty": supplier_return_qty,
        "customer_return_qty": customer_return_qty,
        "unit": unit,
        "supplier": supplier,
        "buy_price": buy_price,
        "sell_price": sell_price,
        "tax": tax,
        "invoice_nature": invoice_nature,
        "purchase_list": purchase_list,
        "seller": seller,
        "cccd": cccd,
        "note": clean_text(raw.get("note")),
        "errors": errors,
        "warnings": warnings,
    }


def parse_workbook(path: Path, fallback_date: str, selected_sheets=None):
    wb = load_workbook(io.BytesIO(path.read_bytes()), data_only=True, read_only=False)
    parsed = []
    skipped_sheets = []
    selected = set(selected_sheets) if selected_sheets else None
    with db() as conn:
        by_code, by_name = product_lookup(conn)
        for ws in wb.worksheets:
            if selected is not None and ws.title not in selected:
                continue
            if ws.max_row > ORDER_IMPORT_MAX_ROWS or ws.max_column > 200:
                raise ValueError(
                    f"Sheet {ws.title} vượt giới hạn {ORDER_IMPORT_MAX_ROWS:,} dòng/200 cột"
                )
            header_row, mapping = detect_header(ws)
            if not header_row:
                skipped_sheets.append(ws.title)
                continue
            sheet_count = 0
            empty_run = 0
            for row in range(header_row + 1, ws.max_row + 1):
                raw = {field: ws.cell(row, col).value for field, col in mapping.items()}
                identity = [
                    raw.get("product_name"), raw.get("product_code"),
                    raw.get("kitchen"), raw.get("qty"),
                ]
                if not any(value not in (None, "") for value in identity):
                    empty_run += 1
                    if empty_run > 200:
                        break
                    continue
                empty_run = 0
                # Formula-filled template rows and genuine zero-quantity lines are
                # not orders. A non-blank malformed/non-finite quantity must still
                # reach resolve_order so the draft shows an error and cannot be
                # approved; silently dropping such a business row loses data.
                raw_qty = raw.get("qty")
                parsed_qty = number_value(raw_qty, math.nan)
                if raw_qty in (None, "") or (math.isfinite(parsed_qty) and parsed_qty <= 0):
                    continue
                order = resolve_order(conn, raw, fallback_date, by_code, by_name)
                order["source_sheet"] = ws.title
                order["source_row"] = row
                parsed.append(order)
                sheet_count += 1
            if sheet_count == 0:
                skipped_sheets.append(ws.title)
    wb.close()
    return parsed, skipped_sheets


def analyze_workbook(path: Path):
    wb = load_workbook(io.BytesIO(path.read_bytes()), data_only=True, read_only=False)
    candidates = []
    for ws in wb.worksheets:
        if ws.max_row > ORDER_IMPORT_MAX_ROWS or ws.max_column > 200:
            wb.close()
            raise ValueError(
                f"Sheet {ws.title} vượt giới hạn {ORDER_IMPORT_MAX_ROWS:,} dòng/200 cột"
            )
        header_row, mapping = detect_header(ws)
        if not header_row:
            continue
        count = 0
        empty_run = 0
        for row in range(header_row + 1, ws.max_row + 1):
            values = [
                ws.cell(row, mapping[key]).value
                for key in ("product_name", "product_code", "kitchen", "qty")
                if key in mapping
            ]
            if not any(value not in (None, "") for value in values):
                empty_run += 1
                if empty_run > 200:
                    break
                continue
            empty_run = 0
            raw_qty = ws.cell(row, mapping["qty"]).value
            parsed_qty = number_value(raw_qty, math.nan)
            # Include malformed non-blank quantities in the sheet preview; the
            # confirmation step will import them as blocked/error rows.
            if (math.isfinite(parsed_qty) and parsed_qty > 0) or (
                raw_qty not in (None, "") and not math.isfinite(parsed_qty)
            ):
                count += 1
        if count:
            candidates.append({
                "name": ws.title,
                "rows": count,
                "headerRow": header_row,
                "fields": sorted(mapping),
            })
    wb.close()
    return candidates


def save_imported_orders(conn, batch_id, orders):
    """Insert parsed rows into an existing batch owned by the caller transaction."""
    for item in orders:
        conn.execute(
            """INSERT INTO orders(
                batch_id,work_date,contractor,kitchen,product_code,product_name,qty,
                actual_received,actual_delivered,damaged_qty,supplier_return_qty,customer_return_qty,
                unit,supplier,buy_price,sell_price,tax,invoice_nature,
                purchase_list,seller,cccd,note,source_sheet,source_row,errors,warnings,updated_at
            ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (
                batch_id, item["work_date"], item["contractor"], item["kitchen"],
                item["product_code"], item["product_name"], item["qty"],
                item["actual_received"], item["actual_delivered"], item["damaged_qty"],
                item["supplier_return_qty"], item["customer_return_qty"], item["unit"],
                item["supplier"], item["buy_price"], item["sell_price"], item["tax"],
                item["invoice_nature"],
                item["purchase_list"], item["seller"], item["cccd"], item["note"],
                item["source_sheet"], item["source_row"],
                json.dumps(item["errors"], ensure_ascii=False),
                json.dumps(item["warnings"], ensure_ascii=False), now_iso(),
            ),
        )


def save_imported_batch(conn, orders, work_date, source_name):
    cur = conn.execute(
        "INSERT INTO batches(work_date,source_name,status,created_at) VALUES(?,?,?,?)",
        (work_date, source_name, "draft", now_iso()),
    )
    batch_id = cur.lastrowid
    save_imported_orders(conn, batch_id, orders)
    return batch_id


def order_import_fingerprint(path: Path, work_date: str, selected_sheets) -> tuple[str, str, list[str]]:
    """Bind one import receipt to the exact bytes, date and explicit sheet set."""
    # Worksheet titles are identifiers, not display text.  Excel permits
    # leading/trailing spaces in a title and the customer's workbook contains
    # one (``đơn hàng27.08 ``).  Trimming here would fingerprint one title but
    # ask openpyxl to parse a different, non-existent worksheet.
    canonical_sheets = sorted({
        str(value) for value in selected_sheets if value is not None and str(value) != ""
    })
    source_bytes = path.read_bytes()
    source_hash = hashlib.sha256(source_bytes).hexdigest().upper()
    digest = hashlib.sha256()
    digest.update(source_bytes)
    digest.update(b"\x00work_date\x00")
    digest.update(work_date.encode("utf-8"))
    for sheet in canonical_sheets:
        digest.update(b"\x00sheet\x00")
        digest.update(sheet.encode("utf-8"))
    return digest.hexdigest().upper(), source_hash, canonical_sheets


def valid_order_import_batch(conn, import_key: str):
    """Return a still-usable imported batch; an empty/deleted batch is not reusable."""
    return conn.execute(
        """SELECT b.id
           FROM order_import_receipts r
           JOIN batches b ON b.id=r.batch_id
           WHERE r.import_key=?
             AND EXISTS(SELECT 1 FROM orders o WHERE o.batch_id=b.id)
           LIMIT 1""",
        (import_key,),
    ).fetchone()


def workbook_date_token(value) -> tuple[int, int] | None:
    match = re.search(r"(?:^|\D)(\d{1,2})[.\-_/](\d{1,2})(?:\D|$)", str(value or ""))
    if not match:
        return None
    day, month = int(match.group(1)), int(match.group(2))
    try:
        date(2000, month, day)
    except ValueError:
        return None
    return day, month


def strict_daily_preview(path: Path, source_name: str = "", *, continuous=False) -> dict | None:
    """Return a safe structure-first preview, or None for legacy workbooks."""
    analysis = analyze_daily_workbook(path)
    if not analysis["strictCustomerWorkbook"]:
        return None
    # Customer workbooks retain older daily sheets.  When the uploaded
    # filename names one exact day, scope the strict preview to that matching
    # sheet just as the operator UI does; otherwise a valid workbook with many
    # historical sheets has no single detected date and can never be confirmed.
    # Non-matching day sheets remain visible only as safe name-level references
    # and are never parsed or written in this upload.
    source_token = workbook_date_token(source_name)
    if source_token:
        matching_days = [
            sheet for sheet in analysis["daySheets"]
            if workbook_date_token(sheet.get("name")) == source_token
        ]
        if len(matching_days) == 1:
            selected_name = matching_days[0]["name"]
            analysis["ignoredSheets"].extend({
                "name": sheet["name"], "role": "historical_daily_sheet",
            } for sheet in analysis["daySheets"] if sheet["name"] != selected_name)
            analysis["daySheets"] = matching_days
            selected_dates = matching_days[0].get("workDates") or []
            analysis["detectedWorkDate"] = selected_dates[0] if len(selected_dates) == 1 else ""
    customer_orders_by_sheet = {}
    for sheet in analysis["daySheets"]:
        dates = sheet.get("workDates") or []
        fallback_date = dates[0] if len(dates) == 1 else date.today().isoformat()
        orders, skipped = parse_workbook(path, fallback_date, [sheet["name"]])
        customer_orders_by_sheet[sheet["name"]] = orders
        sheet["parsedRows"] = len(orders)
        sheet["errorRows"] = sum(bool(item.get("errors")) for item in orders)
        sheet["warningRows"] = sum(bool(item.get("warnings")) for item in orders)
        sheet["writeScope"] = "customer_orders"
        if skipped or len(orders) != int(sheet.get("rows") or 0):
            sheet["confirmAvailable"] = False
            sheet["previewIssue"] = "row_count_mismatch"
    for sheet in analysis["purchaseSheets"]:
        sheet["errorRows"] = 0
        sheet["warningRows"] = 0
        sheet["writeScope"] = "purchase_orders_locked_until_round_trip"
        sheet["previewIssue"] = "customer_scope_must_be_loaded_first"
    analysis["phase"] = "first_load"
    analysis["continuous"] = bool(continuous)
    analysis["batchId"] = None
    analysis["scopeSelectionRequired"] = False
    analysis["_customerOrders"] = customer_orders_by_sheet

    confirmable_days = [
        sheet for sheet in analysis["daySheets"] if sheet.get("confirmAvailable")
    ]
    detected_date = analysis.get("detectedWorkDate") or ""
    if len(confirmable_days) == 1 and detected_date:
        day_sheet = confirmable_days[0]
        with db() as conn:
            workday = conn.execute(
                """SELECT * FROM daily_workdays
                   WHERE work_date=? AND day_sheet_key=?""",
                (detected_date, canonical_day_sheet(day_sheet["name"])),
            ).fetchone()
            if workday:
                batch_id = int(workday["batch_id"])
                latest = conn.execute(
                    """SELECT source_hash,phase FROM daily_import_versions
                       WHERE daily_workday_id=? ORDER BY version_no DESC LIMIT 1""",
                    (int(workday["id"]),),
                ).fetchone()
                first_load_replay = bool(
                    latest and latest["phase"] == "first_load"
                    and latest["source_hash"] == analysis["sourceHash"]
                )
                if not first_load_replay:
                    analysis["phase"] = "finalization"
                    analysis["batchId"] = batch_id
                    analysis["scopeSelectionRequired"] = True
                    customer_diff = strict_customer_scope_diff(
                        conn, batch_id, customer_orders_by_sheet[day_sheet["name"]], continuous=continuous,
                    )
                    day_sheet["diff"] = customer_diff
                    if continuous and day_sheet.get("errorRows"):
                        day_sheet["confirmAvailable"] = False
                        day_sheet["previewIssue"] = "customer_scope_has_errors"
                    if customer_diff["conflicts"]:
                        day_sheet["confirmAvailable"] = False
                        day_sheet["previewIssue"] = "customer_scope_conflict"
                    purchase_preview = None
                    if len(analysis["purchaseSheets"]) == 1:
                        purchase_sheet = analysis["purchaseSheets"][0]
                        try:
                            purchase_preview = strict_purchase_preview_from_path(
                                conn, path, batch_id,
                            )
                            purchase_diff = strict_purchase_scope_diff(
                                conn, batch_id, purchase_preview["items"],
                                purchase_preview["error_rows"],
                            )
                            purchase_sheet.update(
                                parsedRows=purchase_preview["count"],
                                errorRows=purchase_preview["error_rows"],
                                warningRows=purchase_preview["warning_rows"],
                                writeScope="purchase_orders",
                                confirmAvailable=bool(purchase_preview["can_confirm"]),
                                diff=purchase_diff,
                            )
                            purchase_sheet.pop("previewIssue", None)
                            if not purchase_preview["can_confirm"]:
                                purchase_sheet["previewIssue"] = "purchase_scope_has_errors"
                        except ValueError:
                            purchase_sheet.update(
                                confirmAvailable=False,
                                errorRows=max(int(purchase_sheet.get("rows") or 0), 1),
                                writeScope="purchase_orders",
                                previewIssue="purchase_scope_parse_failed",
                                diff={
                                    "added": 0, "updated": 0, "unchanged": 0,
                                    "removed": 0, "conflicts": 1,
                                    "total": int(purchase_sheet.get("rows") or 0),
                                },
                            )
                    analysis["_purchasePreview"] = purchase_preview
                    if workday["lifecycle_status"] == "finalized":
                        for scope_sheet in analysis["daySheets"] + analysis["purchaseSheets"]:
                            diff = scope_sheet.get("diff") or {}
                            has_changes = any(int(diff.get(key) or 0) for key in (
                                "added", "updated", "removed", "conflicts",
                            ))
                            if (
                                not latest or latest["source_hash"] != analysis["sourceHash"]
                                or has_changes
                            ):
                                scope_sheet["confirmAvailable"] = False
                                scope_sheet["previewIssue"] = "workday_finalized"
                analysis["databaseStateHash"] = strict_daily_database_state_hash(
                    conn, batch_id, detected_date, day_sheet["name"],
                )
            else:
                analysis["databaseStateHash"] = strict_daily_database_state_hash(
                    conn, 0, detected_date, day_sheet["name"],
                )
    analysis["stateHash"] = analysis_state_hash(analysis)
    return analysis


def strict_daily_contract_rows(orders):
    """Build hash-only lifecycle descriptors; omit note/seller/CCCD from hashes."""
    occurrences = defaultdict(int)
    rows = []
    payload_fields = (
        "work_date", "contractor", "kitchen", "product_code", "product_name", "qty",
        "actual_received", "actual_delivered", "damaged_qty", "supplier_return_qty",
        "customer_return_qty", "unit", "supplier", "buy_price", "sell_price", "tax",
        "invoice_nature", "purchase_list",
    )
    for order in sorted(orders, key=lambda item: int(item.get("source_row") or 0)):
        identity = {
            "contractor": clean_text(order.get("contractor")),
            "kitchen": clean_text(order.get("kitchen")) or "__MISSING_KITCHEN__",
            "product_code": clean_text(order.get("product_code")),
            "product_name": clean_text(order.get("product_name")),
            "unit": clean_text(order.get("unit")),
        }
        if not identity["product_code"] and not identity["product_name"]:
            identity["product_name"] = "__MISSING_PRODUCT__"
        occurrence_key = tuple(identity[field].casefold() for field in (
            "contractor", "kitchen", "product_code", "product_name", "unit",
        ))
        occurrences[occurrence_key] += 1
        identity["occurrence"] = occurrences[occurrence_key]
        rows.append({
            "identity": identity,
            "payload": {field: order.get(field) for field in payload_fields},
            "source_row": int(order.get("source_row") or 1),
        })
    return rows


def strict_final_sales_contract_rows(orders):
    """Describe only fields owned by the sales/delivery finalization scope."""
    allowed = (
        "work_date", "contractor", "kitchen", "product_code", "product_name",
        "qty", "actual_delivered", "customer_return_qty", "unit", "sell_price",
        "tax", "invoice_nature", "note",
    )
    rows = strict_daily_contract_rows(orders)
    for row in rows:
        row["payload"] = {field: row["payload"].get(field) for field in allowed}
    return rows


def strict_purchase_contract_rows(items):
    """Build lifecycle descriptors from validated canonical purchase rows."""
    fields = (
        "product_code", "kitchen", "work_date", "product_name", "base_qty",
        "unit", "supplier", "buy_price", "price_source", "damaged_qty",
        "added_qty", "reduced_qty", "missing_qty", "actual_qty", "amount",
    )
    return [{
        "row_key": item["row_key"],
        "payload_hash": daily_payload_hash({field: item.get(field) for field in fields}),
        "source_row": int(item.get("source_row") or 1),
    } for item in items]


def strict_continuous_contract_rows(orders):
    """Include every imported value when comparing a later daily revision."""
    ordered = sorted(orders, key=lambda item: int(item.get('source_row') or 0))
    rows = strict_daily_contract_rows(ordered)
    for row, item in zip(rows, ordered):
        row['payload'].update({field: item.get(field) for field in ('note', 'seller', 'cccd')})
    return rows


def strict_order_index(orders, *, finalization=False, continuous=False):
    ordered = sorted(orders, key=lambda item: int(item.get("source_row") or 0))
    descriptors = (
        strict_final_sales_contract_rows(ordered)
        if finalization else strict_continuous_contract_rows(ordered) if continuous else strict_daily_contract_rows(ordered)
    )
    output = {}
    for item, descriptor in zip(ordered, descriptors):
        key = daily_row_key(
            work_date=item["work_date"],
            scope="customer_orders",
            identity=descriptor["identity"],
        )
        output[key] = {
            "order": item,
            "payload_hash": daily_payload_hash(descriptor["payload"]),
            "source_row": descriptor["source_row"],
        }
    return output


def strict_order_removal_conflicts(conn, order_ids):
    conflicts = set()
    for order_id in order_ids:
        if conn.execute(
            "SELECT 1 FROM purchase_order_lines WHERE order_id=? LIMIT 1", (order_id,),
        ).fetchone() or conn.execute(
            "SELECT 1 FROM purchase_workbook_lines WHERE order_id=? LIMIT 1", (order_id,),
        ).fetchone() or conn.execute(
            "SELECT 1 FROM outgoing_invoice_lines WHERE order_id=? LIMIT 1", (order_id,),
        ).fetchone():
            conflicts.add(int(order_id))
    return conflicts


def strict_customer_scope_diff(conn, batch_id: int, incoming_orders, *, continuous=False):
    current_orders = rows_dict(conn.execute(
        "SELECT * FROM orders WHERE batch_id=? ORDER BY source_row,id", (batch_id,),
    )) if batch_id else []
    incoming = strict_order_index(incoming_orders, finalization=not continuous, continuous=continuous)
    current = strict_order_index(current_orders, finalization=not continuous, continuous=continuous)
    incoming_keys = set(incoming)
    current_keys = set(current)
    removed_keys = current_keys - incoming_keys
    conflict_ids = strict_order_removal_conflicts(
        conn, [current[key]["order"]["id"] for key in removed_keys],
    ) if removed_keys else set()
    return {
        "added": len(incoming_keys - current_keys),
        "updated": sum(
            incoming[key]["payload_hash"] != current[key]["payload_hash"]
            for key in incoming_keys & current_keys
        ),
        "unchanged": sum(
            incoming[key]["payload_hash"] == current[key]["payload_hash"]
            for key in incoming_keys & current_keys
        ),
        "removed": len(removed_keys - {
            key for key in removed_keys if int(current[key]["order"]["id"]) in conflict_ids
        }),
        "conflicts": len(conflict_ids),
        "total": len(incoming),
    }


def strict_purchase_scope_diff(conn, batch_id: int, items, error_rows=0):
    previous = {
        row["row_key"]: dict(row) for row in conn.execute(
            "SELECT * FROM purchase_workbook_lines WHERE batch_id=?", (batch_id,),
        )
    }
    incoming = {item["row_key"]: item for item in items}
    shared = set(previous) & set(incoming)
    removed = set(previous) - set(incoming)
    return {
        "added": len(set(incoming) - set(previous)),
        "updated": sum(purchase_line_changed(previous[key], incoming[key]) for key in shared),
        "unchanged": sum(not purchase_line_changed(previous[key], incoming[key]) for key in shared),
        "removed": len(removed),
        "conflicts": int(error_rows or 0),
        "total": len(incoming),
    }


def strict_daily_database_state_hash(conn, batch_id: int, work_date: str, day_sheet: str):
    """Hash all data that can change parsing/applying either finalization scope."""
    queries = [
        ("workday", "SELECT work_date,day_sheet_key,batch_id,lifecycle_status,revision "
         "FROM daily_workdays WHERE work_date=? AND day_sheet_key=?", (work_date, canonical_day_sheet(day_sheet))),
        ("contractors", "SELECT code,name,price_group,pricing_mode FROM contractors ORDER BY code", ()),
        ("kitchens", "SELECT code,contractor,name,address,show_price FROM kitchens ORDER BY code", ()),
        ("suppliers", "SELECT code,name FROM suppliers ORDER BY code", ()),
        ("products", "SELECT code,name,unit,tax,supplier,buy_price,purchase_list FROM products ORDER BY code", ()),
        ("prices", "SELECT product_code,price_group,price_text,price_value FROM product_prices "
         "ORDER BY product_code,price_group", ()),
        ("settings", "SELECT key,value FROM settings WHERE key IN ('purchase_rate') ORDER BY key", ()),
    ]
    if batch_id:
        queries.extend([
            ("batch", "SELECT id,work_date,source_name,status,created_at,approved_at "
             "FROM batches WHERE id=?", (batch_id,)),
            ("orders", "SELECT * FROM orders WHERE batch_id=? ORDER BY id", (batch_id,)),
            ("purchase_lines", "SELECT * FROM purchase_workbook_lines WHERE batch_id=? "
             "ORDER BY row_key", (batch_id,)),
            ("legacy_purchase", "SELECT * FROM purchase_order_lines WHERE batch_id=? "
             "ORDER BY order_id", (batch_id,)),
        ])
    digest = hashlib.sha256()
    for label, sql, params in queries:
        digest.update(label.encode("ascii"))
        digest.update(b"\0")
        rows = [tuple(row) for row in conn.execute(sql, params)]
        digest.update(json.dumps(rows, ensure_ascii=False, default=str).encode("utf-8"))
        digest.update(b"\n")
    return digest.hexdigest().upper()


def strict_purchase_preview_from_path(conn, path: Path, batch_id: int):
    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    formula_workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        return parse_purchase_order_workbook(
            conn, workbook, batch_id, formula_workbook=formula_workbook,
        )
    finally:
        workbook.close()
        formula_workbook.close()


def apply_strict_customer_scope(conn, batch_id: int, incoming_orders, *, continuous=False):
    diff = strict_customer_scope_diff(conn, batch_id, incoming_orders, continuous=continuous)
    if diff["conflicts"]:
        raise DailyImportError(
            "Có dòng cũ đã đi vào đặt NCC hoặc hóa đơn; không thể tự xóa khi chốt lại",
            code="customer_scope_conflict",
        )
    current_rows = rows_dict(conn.execute(
        "SELECT * FROM orders WHERE batch_id=? ORDER BY source_row,id", (batch_id,),
    ))
    current = strict_order_index(current_rows, finalization=not continuous, continuous=continuous)
    incoming = strict_order_index(incoming_orders, finalization=not continuous, continuous=continuous)
    removed_keys = set(current) - set(incoming)
    for key in removed_keys:
        conn.execute("DELETE FROM orders WHERE id=?", (int(current[key]["order"]["id"]),))
    if removed_keys:
        # Persist removal tombstones before inserting any new rows. SQLite may
        # reuse a deleted orders.id inside this same transaction; the later
        # sync must then mint a new source generation, not overwrite history.
        sync_receivable_ledger(conn, timestamp=now_iso())
    update_fields = (
        "work_date", "contractor", "kitchen", "product_code", "product_name", "qty",
        "actual_delivered", "customer_return_qty", "unit", "sell_price", "tax",
        "invoice_nature", "note", "source_sheet", "source_row",
    )
    if continuous:
        update_fields = tuple(sorted(ORDER_FIELDS | {"work_date", "source_sheet", "source_row"}))
    for key, value in incoming.items():
        item = value["order"]
        existing = current.get(key)
        if existing is None:
            # A row introduced by the sales/delivery capability must not
            # manufacture a purchase or supplier payable.  The separate
            # purchase capability owns these fields and can link its row after
            # this order exists.
            sales_item = dict(item)
            if not continuous:
                sales_item.update(
                actual_received=0,
                damaged_qty=0,
                supplier_return_qty=0,
                supplier="",
                buy_price=0,
                purchase_list=0,
                )
            save_imported_orders(conn, batch_id, [sales_item])
            continue
        if value["payload_hash"] == existing["payload_hash"]:
            continue
        assignments = ",".join(f"{field}=?" for field in update_fields)
        conn.execute(
            f"UPDATE orders SET {assignments},errors=?,warnings=?,updated_at=? WHERE id=?",
            (
                *(item[field] for field in update_fields),
                json.dumps(item["errors"], ensure_ascii=False),
                json.dumps(item["warnings"], ensure_ascii=False),
                now_iso(), int(existing["order"]["id"]),
            ),
        )
        if continuous:
            conn.execute("UPDATE orders SET sell_price_revision=sell_price_revision+1,"
                         "sell_price_source='import_or_standard',sell_price_override_at=NULL WHERE id=?",
                         (int(existing['order']['id']),))
    sync_receivable_ledger(conn, timestamp=now_iso())
    # The explicit second sales/delivery confirmation settles the existing
    # order claim; it does not post another physical issue.
    if not continuous:
        conn.execute("UPDATE orders SET physical_stage='delivered' WHERE batch_id=?", (batch_id,))
    return diff


def confirm_strict_daily_finalization(pending, body, sheets):
    analysis = pending["strict_analysis"]
    detected_date = analysis.get("detectedWorkDate") or ""
    work_date = valid_iso_date(body.get("work_date") or detected_date, "Ngày phiên đơn")
    if not detected_date or work_date != detected_date:
        raise DailyImportError(
            "Ngày phiên đơn phải trùng ngày đã nhận diện trong workbook",
            code="work_date_mismatch", status=400,
        )
    day_by_name = {item["name"]: item for item in analysis["daySheets"]}
    purchase_by_name = {item["name"]: item for item in analysis["purchaseSheets"]}
    if not sheets:
        raise DailyImportError(
            "Cần chọn rõ ít nhất một phạm vi bán/giao hoặc mua/phải trả",
            code="scope_selection_required", status=400,
        )
    selected_scopes = []
    for sheet_name in sheets:
        sheet = day_by_name.get(sheet_name) or purchase_by_name.get(sheet_name)
        if not sheet or not sheet.get("confirmAvailable"):
            raise DailyImportError(
                "Chỉ được xác nhận phạm vi đã qua preview và không có xung đột",
                code="scope_not_confirmable", status=400,
            )
        scope = sheet["scope"]
        if scope not in selected_scopes:
            selected_scopes.append(scope)
    selected_scopes.sort(key=lambda scope: 0 if scope == "customer_orders" else 1)
    if len(day_by_name) != 1:
        raise DailyImportError(
            "Chốt lần hai cần đúng một sheet ngày đã nhận diện",
            code="ambiguous_daily_sheet", status=400,
        )
    day_sheet_name = next(iter(day_by_name))
    source_hash = hashlib.sha256(pending["path"].read_bytes()).hexdigest().upper()
    if source_hash != analysis["sourceHash"]:
        raise DailyImportError(
            "Nội dung workbook đã thay đổi sau preview", code="source_changed",
        )
    batch_id = int(analysis.get("batchId") or 0)
    customer_orders = analysis["_customerOrders"][day_sheet_name]
    purchase_preview = analysis.get("_purchasePreview")

    with db() as conn:
        conn.execute("BEGIN IMMEDIATE")
        current_hash = strict_daily_database_state_hash(
            conn, batch_id, work_date, day_sheet_name,
        )
        if current_hash != analysis.get("databaseStateHash"):
            raise DailyImportError(
                "Dữ liệu đã đổi sau preview; hãy tải lại file trước khi chốt",
                code="stale_database_state",
            )
        blocked = batch_mutation_blocker(conn, batch_id)
        if blocked:
            raise DailyImportError(blocked, code="batch_locked")
        old_orders = rows_dict(conn.execute("SELECT * FROM orders WHERE batch_id=? ORDER BY id", (batch_id,)))
        previous_keys = [r['import_key'] for r in conn.execute('SELECT import_key FROM order_import_receipts WHERE batch_id=?', (batch_id,))]
        previous_batch = conn.execute('SELECT source_name FROM batches WHERE id=?', (batch_id,)).fetchone()
        conn.execute('''INSERT INTO order_reimport_history
            (batch_id,previous_source,next_source,rows_json,previous_keys_json,created_at)
            VALUES(?,?,?,?,?,?)''', (batch_id, previous_batch['source_name'], pending['name'],
            json.dumps(old_orders, ensure_ascii=False), json.dumps(previous_keys), now_iso()))
        audit_event(conn, 'order.workbook.replace', entity_type='batch', entity_id=batch_id,
                    metadata={'source': pending['name'], 'selected_scopes': selected_scopes,
                              'purchase_before': rows_dict(conn.execute('SELECT * FROM purchase_workbook_lines WHERE batch_id=?', (batch_id,)))})
        rows_by_scope = {
            "customer_orders": (strict_continuous_contract_rows(customer_orders) if analysis.get('continuous')
                                else strict_final_sales_contract_rows(customer_orders)),
        }
        if analysis["purchaseSheets"]:
            if purchase_preview is not None:
                rows_by_scope["purchase_orders"] = strict_purchase_contract_rows(
                    purchase_preview["items"],
                )
            else:
                fallback_key = hashlib.sha256(
                    f"{source_hash}\0purchase_orders".encode("utf-8")
                ).hexdigest().upper()
                rows_by_scope["purchase_orders"] = [{
                    "row_key": fallback_key,
                    "payload_hash": source_hash,
                    "source_row": int(analysis["purchaseSheets"][0].get("headerRow") or 1) + 1,
                }]
        contract = prepare_daily_import_version(
            conn,
            source_hash=source_hash,
            work_date=work_date,
            day_sheet=day_sheet_name,
            phase="finalization",
            rows_by_scope=rows_by_scope,
            now_iso=now_iso,
        )
        if int(contract["batch_id"]) != batch_id:
            raise DailyImportError(
                "Preview không còn trỏ tới đúng phiên ngày", code="stale_daily_batch",
            )
        latest_version = int(conn.execute(
            "SELECT MAX(version_no) n FROM daily_import_versions WHERE daily_workday_id=?",
            (int(contract["daily_workday_id"]),),
        ).fetchone()["n"])
        if contract["idempotent"] and int(contract["version_no"]) != latest_version:
            raise DailyImportError(
                "Phiên bản này đã bị bản chốt mới hơn thay thế",
                code="superseded_version",
            )
        scopes = {item["scope"]: item for item in contract["scopes"]}
        scope_results = {}
        for scope in selected_scopes:
            grant = scopes[scope]
            if scope == "customer_orders":
                diff = apply_strict_customer_scope(conn, batch_id, customer_orders, continuous=analysis.get('continuous', False))
                applied = {
                    "processed": diff["added"] + diff["updated"] + diff["removed"],
                    "diff": diff,
                }
            else:
                if purchase_preview is None or not purchase_preview["can_confirm"]:
                    raise DailyImportError(
                        "Phạm vi mua/phải trả còn lỗi; chưa được ghi",
                        code="purchase_scope_has_errors", status=400,
                    )
                if "customer_orders" in selected_scopes:
                    refreshed_purchase = strict_purchase_preview_from_path(
                        conn, pending["path"], batch_id,
                    )
                    if (
                        not refreshed_purchase["can_confirm"]
                        or refreshed_purchase["content_hash"] != purchase_preview["content_hash"]
                    ):
                        raise DailyImportError(
                            "Phạm vi mua thay đổi sau khi áp dụng bán/giao; hãy xem trước lại",
                            code="purchase_scope_changed", status=409,
                        )
                    purchase_preview = refreshed_purchase
                try:
                    applied = apply_purchase_order_preview(
                        conn,
                        batch_id=batch_id,
                        items=purchase_preview["items"],
                        source_hash=purchase_preview["content_hash"],
                        source_name=pending["name"],
                        format_name=purchase_preview["format"],
                        now_iso=now_iso,
                    )
                except PurchaseOrderApplyError as error:
                    raise DailyImportError(
                        str(error), code=error.code, status=error.status,
                    ) from error
            confirmed = confirm_daily_import_scope(
                conn,
                version_id=int(contract["version_id"]),
                scope=scope,
                expected_state_hash=grant["state_hash"],
                now_iso=now_iso,
            )
            scope_results[scope] = {**applied, "confirmation": confirmed}

        remaining = int(conn.execute(
            """SELECT COUNT(*) n FROM daily_import_scopes
               WHERE version_id=? AND state!='confirmed'""",
            (int(contract["version_id"]),),
        ).fetchone()["n"])
        finalized = None
        if remaining == 0 and not analysis.get('continuous'):
            finalized = finalize_daily_workday(
                conn, version_id=int(contract["version_id"]), now_iso=now_iso,
            )
        import_key, _, _ = order_import_fingerprint(pending["path"], work_date, sheets)
        conn.execute(
            "DELETE FROM order_import_receipts WHERE batch_id=? AND import_key<>?",
            (batch_id, import_key),
        )
        conn.execute(
            """INSERT INTO order_import_receipts(
                   import_key,source_hash,source_name,work_date,selected_sheets_json,batch_id,created_at
               ) VALUES(?,?,?,?,?,?,?) ON CONFLICT(import_key) DO NOTHING""",
            (
                import_key, source_hash, pending["name"], work_date,
                json.dumps(sheets, ensure_ascii=False), batch_id, now_iso(),
            ),
        )
        sync_receivable_ledger(conn, timestamp=now_iso())
        conn.execute("UPDATE batches SET source_name=? WHERE id=?", (pending['name'], batch_id))
        if analysis.get('continuous'):
            conn.execute("UPDATE batches SET status='draft',approved_at=NULL WHERE id=?", (batch_id,))
            sync_payable_ledger(conn, timestamp=now_iso())
            sync_receivable_ledger(conn, timestamp=now_iso())
        daily_import = strict_daily_version_payload(
            conn, batch_id, source_hash, day_sheet_name,
        )
        payload = batch_payload(conn, batch_id)
        payload.update(
            ok=True,
            selectedSheets=sheets,
            selectedScopes=selected_scopes,
            sourceHash=source_hash,
            importKey=import_key,
            idempotent=all(
                result["confirmation"]["idempotent"]
                and (scope != "purchase_orders" or result.get("idempotent", False))
                for scope, result in scope_results.items()
            ),
            scopeResults=scope_results,
            dailyImport=daily_import,
            finalized=finalized,
        )
        return jsonify(payload)


def strict_daily_version_payload(conn, batch_id: int, source_hash: str, sheet_name: str):
    version = conn.execute(
        """SELECT v.id,v.version_no,v.version_key,v.source_hash,v.phase,v.exact_day_sheet,
                  v.confirmed_at,w.id daily_workday_id,w.work_date,w.lifecycle_status,w.revision
           FROM daily_import_versions v
           JOIN daily_workdays w ON w.id=v.daily_workday_id
           WHERE w.batch_id=? AND v.source_hash=? AND v.exact_day_sheet=?
           ORDER BY v.version_no DESC LIMIT 1""",
        (int(batch_id), source_hash, sheet_name),
    ).fetchone()
    if not version:
        return None
    scopes = rows_dict(conn.execute(
        """SELECT scope,write_capability,state,state_hash,row_count,confirmed_at
           FROM daily_import_scopes WHERE version_id=? ORDER BY scope""",
        (int(version["id"]),),
    ))
    return {
        "daily_workday_id": int(version["daily_workday_id"]),
        "batch_id": int(batch_id),
        "work_date": version["work_date"],
        "lifecycle_status": version["lifecycle_status"],
        "lifecycle_revision": int(version["revision"]),
        "version_id": int(version["id"]),
        "version_no": int(version["version_no"]),
        "version_key": version["version_key"],
        "source_hash": version["source_hash"],
        "phase": version["phase"],
        "day_sheet": version["exact_day_sheet"],
        "scopes": scopes,
        "idempotent": True,
    }


def validate_existing_order(conn, item: dict):
    raw = dict(item)
    resolved = resolve_order(conn, raw, item.get("work_date") or date.today().isoformat(),
                             *product_lookup(conn))
    return resolved["errors"]


ORDER_FIELDS = {
    "contractor", "kitchen", "product_code", "product_name", "qty",
    "actual_received", "actual_delivered", "unit", "supplier", "buy_price",
    "sell_price", "tax", "invoice_nature", "purchase_list", "seller", "cccd", "note",
    "damaged_qty", "supplier_return_qty", "customer_return_qty",
}


def order_totals(order):
    revenue = vnd_product(net_delivered(order), number_value(order["sell_price"]))
    cost = vnd_product(net_received(order), number_value(order["buy_price"]))
    vat_percent = invoice_vat_percent(order["tax"])
    tax_amount = 0 if vat_percent <= 0 else vnd_product(revenue, vat_percent / 100)
    total = revenue + tax_amount
    return revenue, cost, revenue - cost, total


def batch_payload(conn, batch_id=None):
    if batch_id is None:
        batch = conn.execute("SELECT * FROM batches ORDER BY id DESC LIMIT 1").fetchone()
    else:
        batch = conn.execute("SELECT * FROM batches WHERE id=?", (batch_id,)).fetchone()
    if not batch:
        return {"batch": None, "orders": [], "summary": calculate_summary(conn, [])}
    orders = rows_dict(conn.execute("SELECT * FROM orders WHERE batch_id=? ORDER BY id", (batch["id"],)))
    for item in orders:
        item["worksheet_revision"] = order_worksheet.row_revision(item)
        item["errors"] = json.loads(item["errors"] or "[]")
        item["warnings"] = json.loads(item["warnings"] or "[]")
        revenue, cost, profit, total = order_totals(item)
        item.update(revenue=revenue, cost=cost, profit=profit, total=total)
        item["physical_revision"] = physical_order_revision(item)
    return {"batch": dict(batch), "orders": orders, "summary": calculate_summary(conn, orders)}


def calculate_summary(conn, orders):
    """Summarize only the selected order batch.

    Payments and opening balances are deliberately excluded because they are
    cumulative party-level ledger entries and are not allocated to one batch.
    Multi-period debt is calculated by ``/api/debts`` instead.
    """
    contractors = defaultdict(lambda: {
        "revenue": 0, "cost": 0, "profit": 0, "total": 0,
        "opening": 0, "paid": 0, "balance": 0,
    })
    kitchens = defaultdict(lambda: {
        "contractor": "", "revenue": 0, "total": 0, "paid": 0, "balance": 0,
    })
    suppliers = defaultdict(lambda: {
        "cost": 0, "opening": 0, "paid": 0, "balance": 0, "lines": 0,
    })
    totals = {
        "ordered_qty": 0, "received_qty": 0, "delivered_qty": 0,
        "revenue": 0, "cost": 0, "profit": 0, "total": 0,
        "errors": 0, "warnings": 0,
    }
    for item in orders:
        revenue, cost, profit, total = order_totals(item)
        totals["ordered_qty"] += number_value(
            item.get("qty", 0) if isinstance(item, dict) else item["qty"]
        )
        totals["received_qty"] += net_received(item)
        totals["delivered_qty"] += net_delivered(item)
        totals["revenue"] += revenue
        totals["cost"] += cost
        totals["profit"] += profit
        totals["total"] += total
        totals["errors"] += 1 if item.get("errors") else 0
        totals["warnings"] += 1 if item.get("warnings") else 0
        c = contractors[item["contractor"] or "CHƯA XÁC ĐỊNH"]
        c["revenue"] += revenue
        c["cost"] += cost
        c["profit"] += profit
        c["total"] += total
        k = kitchens[item["kitchen"] or "CHƯA XÁC ĐỊNH"]
        k["contractor"] = item["contractor"]
        k["revenue"] += revenue
        k["total"] += total
        s = suppliers[item["supplier"] or "CHƯA XÁC ĐỊNH"]
        s["cost"] += cost
        s["lines"] += 1

    for values in contractors.values():
        values["balance"] = values["total"]
    for values in kitchens.values():
        values["balance"] = values["total"]
    for values in suppliers.values():
        values["balance"] = values["cost"]
    return {
        "scope": "batch_only",
        "totals": totals,
        "contractors": dict(contractors),
        "kitchens": dict(kitchens),
        "suppliers": dict(suppliers),
    }


@app.errorhandler(Exception)
def handle_error(error):
    if isinstance(error, HTTPException):
        return jsonify({"ok": False, "error": error.description}), error.code
    reference = uuid.uuid4().hex[:10].upper()
    app.logger.exception("Unhandled error [%s]", reference)
    return jsonify({
        "ok": False,
        "error": f"Lỗi nội bộ. Vui lòng thử lại hoặc báo mã {reference} cho bên hỗ trợ.",
        "reference": reference,
    }), 500


@app.get("/")
def index():
    return send_from_directory(STATIC_DIR, "index.html")


@app.get("/favicon.ico")
def favicon():
    return Response(status=204)


@app.get("/shared/<path:name>")
def shared_asset(name):
    return send_from_directory(SHARED_DIR, name)


@app.get("/health")
def health():
    required_tables = {
        "settings", "batches", "orders", "order_import_receipts",
        "audit_log", "inventory_transactions",
    }
    daily_tables = {
        "daily_workdays", "daily_import_versions", "daily_import_scopes", "daily_import_rows",
    }
    payable_tables = {"payable_ledger_lines", "payable_ledger_revisions"}
    receivable_tables = {"receivable_ledger_lines", "receivable_ledger_revisions"}
    payable_payment_tables = {
        "payable_payment_allocations", "payable_payment_revisions",
    }
    if not DB_PATH.is_file():
        return jsonify({
            "ok": False, "time": now_iso(), "database_ready": False,
            "schema_ready": False, "integrity": "not_checked",
        }), 503
    conn = None
    try:
        database_uri = DB_PATH.resolve().as_uri() + "?mode=ro"
        conn = sqlite3.connect(database_uri, uri=True, timeout=3)
        conn.execute("SELECT 1").fetchone()
        integrity = str(conn.execute("PRAGMA quick_check(1)").fetchone()[0])
        tables = {
            row[0] for row in conn.execute(
                "SELECT name FROM sqlite_master WHERE type='table'"
            )
        }
        duplicate_invoice_warning = False
        duplicate_payable_warning = False
        if "settings" in tables:
            duplicate_invoice_warning = conn.execute(
                "SELECT 1 FROM settings WHERE key='schema_warning_duplicate_invoice_ids' "
                "AND TRIM(value) NOT IN ('','[]')"
            ).fetchone() is not None
            duplicate_payable_warning = conn.execute(
                "SELECT 1 FROM settings WHERE key='schema_warning_duplicate_payable_source_ids' "
                "AND TRIM(value) NOT IN ('','[]')"
            ).fetchone() is not None
        daily_schema_ready = not (tables & daily_tables) or daily_tables.issubset(tables)
        payable_schema_ready = not (tables & payable_tables) or payable_tables.issubset(tables)
        receivable_schema_ready = (
            not (tables & receivable_tables) or receivable_tables.issubset(tables)
        )
        payable_payment_schema_ready = (
            not (tables & payable_payment_tables)
            or payable_payment_tables.issubset(tables)
        )
        schema_ready = (
            required_tables.issubset(tables) and daily_schema_ready and payable_schema_ready
            and payable_payment_schema_ready and receivable_schema_ready
            and not duplicate_invoice_warning
            and not duplicate_payable_warning
        )
        ready = integrity.lower() == "ok" and schema_ready
        return jsonify({
            "ok": ready,
            "time": now_iso(),
            "database_ready": ready,
            "schema_ready": schema_ready,
            "integrity": "ok" if integrity.lower() == "ok" else "failed",
            "duplicate_invoice_identity_warning": duplicate_invoice_warning,
            "duplicate_payable_source_warning": duplicate_payable_warning,
        }), 200 if ready else 503
    except (OSError, sqlite3.Error, ValueError):
        return jsonify({
            "ok": False, "time": now_iso(), "database_ready": False,
            "schema_ready": False, "integrity": "failed",
        }), 503
    finally:
        if conn is not None:
            conn.close()


def create_minvoice_client():
    test_factory = app.config.get("MINVOICE_CLIENT_FACTORY")
    if test_factory is not None:
        return test_factory()
    config = MinvoiceConfig.from_env_files(connector_config_paths())
    return MinvoiceClient(config)


def create_msmi_client():
    test_factory = app.config.get("MSMI_CLIENT_FACTORY")
    if test_factory is not None:
        return test_factory()
    config = MsmiConfig.from_env_files(connector_config_paths())
    return MsmiClient(config)


@app.get("/api/minvoice/status")
def api_minvoice_status():
    """Verify M-Invoice without creating, signing or changing any remote invoice."""
    try:
        client = create_minvoice_client()
        account = client.profile_status()
        outgoing = client.outgoing_summary()
        return jsonify({
            "ok": True,
            "connected": True,
            "account": account,
            "outgoing": outgoing,
            "incoming": {
                "available": False,
                "requires": "mSMI OpenAPI",
                "reason": "Tài khoản/cổng phát hành M-Invoice hiện tại không cung cấp hóa đơn đầu vào từ Tổng cục Thuế.",
            },
            "read_only": True,
            "status_check_read_only": True,
            "draft_save_available": account.get('draft_save_available', not account.get('test_environment', False)),
            "test_environment": account.get('test_environment', False),
            "test_environment_allowed": account.get('test_environment_allowed', False),
            "warning": (('Đang dùng tài khoản M-Invoice hiện tại theo lựa chọn của chủ dự án. '
                         'Máy chủ này là môi trường kiểm thử của nhà cung cấp.'
                         if account.get('test_environment_allowed', False) else
                         'Đang kết nối máy chủ kiểm thử M-Invoice; cần cấu hình tài khoản được chủ dự án chọn.')
                        if account.get('test_environment', False) else ''),
            "automatic_sign_or_issue": False,
            "official_api": True,
        })
    except MinvoiceError as error:
        return jsonify({"ok": False, "connected": False, "error": str(error), "read_only": True}), 502


@app.get("/api/minvoice/series")
def api_minvoice_series():
    """Return only invoice-series metadata; this endpoint never changes M-Invoice."""
    try:
        year = request.args.get("year", date.today().year, type=int)
        target = year % 100
        series = []
        for item in create_minvoice_client().get_invoice_series():
            value = clean_text(item.get("value") or item.get("khhdon")).upper()
            invoice_year = int(number_value(item.get("invoiceYear"), 0))
            if value.startswith("1") and invoice_year == target:
                series.append({"value": value, "invoice_year": invoice_year})
        unique = {item["value"]: item for item in series if item["value"]}
        return jsonify({"ok": True, "year": year, "items": list(unique.values()), "read_only": True})
    except MinvoiceError as error:
        return jsonify({"ok": False, "error": str(error), "read_only": True}), 502


@app.get("/api/bootstrap")
def api_bootstrap():
    with db() as conn:
        payload = batch_payload(conn, request.args.get("batch_id", type=int))
        payload["hosted"] = bool(app.config.get("TDP_CLOUD_ADMIN_USER"))
        payload["batches"] = rows_dict(conn.execute(
            """SELECT b.*,
                      (SELECT COUNT(*) FROM orders o WHERE o.batch_id=b.id) AS line_count,
                      (SELECT COUNT(DISTINCT UPPER(TRIM(o.kitchen)))
                         FROM orders o WHERE o.batch_id=b.id AND TRIM(o.kitchen)!='') AS order_count,
                      (SELECT COUNT(*) FROM orders o
                         WHERE o.batch_id=b.id AND COALESCE(o.errors,'[]')!='[]') AS error_count
                 FROM batches b ORDER BY b.id DESC LIMIT 100"""
        ))
        batch_kitchens = defaultdict(list)
        batch_ids = [int(item["id"]) for item in payload["batches"]]
        if batch_ids:
            placeholders = ",".join("?" for _ in batch_ids)
            for row in conn.execute(
                f"""SELECT o.batch_id,UPPER(TRIM(o.kitchen)) code,
                           COALESCE(NULLIF(TRIM(k.name),''),UPPER(TRIM(o.kitchen))) name,
                           COUNT(*) line_count
                      FROM orders o
                      LEFT JOIN kitchens k ON UPPER(TRIM(k.code))=UPPER(TRIM(o.kitchen))
                     WHERE o.batch_id IN ({placeholders})
                       AND TRIM(o.kitchen)!=''
                       AND MAX(COALESCE(o.actual_delivered,0)-COALESCE(o.customer_return_qty,0),0)>0
                     GROUP BY o.batch_id,UPPER(TRIM(o.kitchen)),
                              COALESCE(NULLIF(TRIM(k.name),''),UPPER(TRIM(o.kitchen)))
                     ORDER BY o.batch_id,UPPER(TRIM(o.kitchen))""",
                batch_ids,
            ):
                batch_kitchens[int(row["batch_id"])].append({
                    "code": row["code"],
                    "name": row["name"],
                    "line_count": int(row["line_count"] or 0),
                })
        for item in payload["batches"]:
            item["delivery_notes"] = batch_kitchens[int(item["id"])]
        payload["master"] = {
            "eligible_sellers": [row["name"] for row in conn.execute("SELECT name FROM people ORDER BY name")
                                 if not is_excluded_seller(row["name"])],
            "excluded_sellers": list(EXCLUDED_SELLERS),
            "contractors": rows_dict(conn.execute("SELECT * FROM contractors ORDER BY code")),
            "kitchens": rows_dict(conn.execute("SELECT * FROM kitchens ORDER BY code")),
            "suppliers": rows_dict(conn.execute("SELECT * FROM suppliers ORDER BY code")),
            "product_count": conn.execute("SELECT COUNT(*) n FROM products").fetchone()["n"],
            "outgoing_name_count": conn.execute("SELECT COUNT(*) n FROM outgoing_product_names").fetchone()["n"],
            "outgoing_names": rows_dict(conn.execute(
                """SELECT m.product_code,p.name source_name,m.invoice_name,m.updated_at
                   FROM outgoing_product_names m JOIN products p ON p.code=m.product_code
                   ORDER BY m.updated_at DESC,m.product_code LIMIT 100"""
            )),
            "settings": {row["key"]: row["value"] for row in conn.execute("SELECT * FROM settings")},
        }
        payload["payments"] = rows_dict(conn.execute(
            "SELECT * FROM payments WHERE COALESCE(status,'posted')='posted' "
            "ORDER BY payment_date DESC,id DESC LIMIT 100"
        ))
        payload["ok"] = True
        return jsonify(payload)


@app.post("/api/import")
def api_import():
    return jsonify({
        "ok": False,
        "error": "Luồng nhập trực tiếp đã khóa để tránh nhập nhầm toàn bộ lịch sử; hãy phân tích file và chọn rõ sheet ngày",
        "required_flow": ["/api/import/analyze", "/api/import/confirm"],
    }), 410


@app.post("/api/import/analyze")
def api_import_analyze():
    upload = request.files.get("file")
    if not upload or not upload.filename:
        return jsonify({"ok": False, "error": "Chưa chọn file Excel"}), 400
    if Path(upload.filename).suffix.lower() not in {".xlsx", ".xlsm"}:
        return jsonify({"ok": False, "error": "Chỉ nhận file .xlsx hoặc .xlsm"}), 400
    token = uuid.uuid4().hex
    suffix = Path(upload.filename).suffix.lower()
    temp = DATA_DIR / f"pending_{token}{suffix}"
    # Register before creating/analysing the file. Holding one lock across the
    # orphan sweep prevents a concurrent request from deleting another request's
    # just-created pending workbook.
    with PENDING_IMPORT_LOCK:
        cutoff = time.time() - 3600
        for old_token, item in list(PENDING_IMPORTS.items()):
            if item["created"] < cutoff:
                try:
                    item["path"].unlink()
                except OSError:
                    pass
                PENDING_IMPORTS.pop(old_token, None)
        active_pending_paths = {
            item["path"].resolve() for item in PENDING_IMPORTS.values()
            if item.get("path")
        }
        for pattern in ("pending_*.xlsx", "pending_*.xlsm"):
            for orphan in DATA_DIR.glob(pattern):
                try:
                    if orphan.is_file() and orphan.resolve() not in active_pending_paths:
                        orphan.unlink()
                except OSError:
                    pass
        PENDING_IMPORTS[token] = {
            "path": temp, "name": upload.filename, "created": time.time(), "state": "analyzing",
        }
    try:
        upload.save(temp)
        if temp.stat().st_size > ORDER_IMPORT_MAX_BYTES:
            raise ValueError("File Excel vượt quá giới hạn 20 MB")
        with zipfile.ZipFile(temp) as archive:
            entries = archive.infolist()
            expanded_size = sum(item.file_size for item in entries)
            if len(entries) > ORDER_IMPORT_MAX_ENTRIES or expanded_size > ORDER_IMPORT_MAX_UNCOMPRESSED_BYTES:
                raise ValueError("File Excel có cấu trúc quá lớn để đọc an toàn")
        strict_analysis = strict_daily_preview(temp, upload.filename, continuous=request.form.get('continuous') == '1')
        sheets = (
            strict_analysis["daySheets"] + strict_analysis["purchaseSheets"]
            if strict_analysis else analyze_workbook(temp)
        )
    except (OSError, ValueError, zipfile.BadZipFile) as exc:
        try:
            temp.unlink()
        except OSError:
            pass
        with PENDING_IMPORT_LOCK:
            PENDING_IMPORTS.pop(token, None)
        message = str(exc) if isinstance(exc, ValueError) else "File Excel bị hỏng hoặc không đúng định dạng"
        return jsonify({"ok": False, "error": message}), 400
    except Exception:
        try:
            temp.unlink()
        except OSError:
            pass
        with PENDING_IMPORT_LOCK:
            PENDING_IMPORTS.pop(token, None)
        raise
    if not sheets:
        temp.unlink(missing_ok=True)
        with PENDING_IMPORT_LOCK:
            PENDING_IMPORTS.pop(token, None)
        return jsonify({
            "ok": False,
            "error": "Không tìm thấy sheet đơn hàng có Mã bếp, Tên/Mã hàng và Số lượng.",
        }), 400
    with PENDING_IMPORT_LOCK:
        PENDING_IMPORTS[token]["state"] = "ready"
        if strict_analysis:
            PENDING_IMPORTS[token]["strict_analysis"] = strict_analysis
    payload = {"ok": True, "token": token, "filename": upload.filename, "sheets": sheets}
    if strict_analysis:
        payload.update(
            strictDaily=True,
            stateHash=strict_analysis["stateHash"],
            detectedWorkDate=strict_analysis["detectedWorkDate"],
            phase=strict_analysis.get("phase", "first_load"),
            batchId=strict_analysis.get("batchId"),
            scopeSelectionRequired=bool(strict_analysis.get("scopeSelectionRequired")),
            ignoredSheets=strict_analysis["ignoredSheets"],
        )
    return jsonify(payload)


@app.post("/api/import/cancel")
def api_import_cancel():
    body = request.get_json(silent=True) or {}
    token = clean_text(body.get("token"))
    with PENDING_IMPORT_LOCK:
        pending = PENDING_IMPORTS.pop(token, None)
    if pending:
        try:
            pending["path"].unlink()
        except OSError:
            pass
    return jsonify({"ok": True, "idempotent": pending is None})


def confirm_strict_daily_import(pending, body):
    """Consume one strict preview token and atomically apply its customer-order scope."""
    analysis = pending["strict_analysis"]
    try:
        raw_sheets = body.get("sheets") or []
        if not isinstance(raw_sheets, list):
            raise DailyImportError(
                "Danh sách sheet không hợp lệ", code="invalid_sheet_list", status=400,
            )
        sheets = sorted({
            str(value) for value in raw_sheets if value is not None and str(value) != ""
        })
        supplied_state_hash = clean_text(body.get("state_hash") or body.get("stateHash")).upper()
        if not supplied_state_hash or supplied_state_hash != analysis["stateHash"]:
            raise DailyImportError(
                "Preview đã cũ hoặc không đúng phiên; vui lòng chọn lại file",
                code="stale_preview",
            )
        if analysis.get("phase") == "finalization":
            return confirm_strict_daily_finalization(pending, body, sheets)
        allowed = {
            item["name"] for item in analysis["daySheets"] if item.get("confirmAvailable")
        }
        if len(sheets) != 1 or sheets[0] not in allowed:
            raise DailyImportError(
                "Chỉ được xác nhận đúng một sheet ngày đã qua preview; sheet đặt hàng và tham chiếu đang bị khóa",
                code="scope_not_confirmable",
                status=400,
            )
        detected_date = analysis.get("detectedWorkDate") or ""
        work_date = valid_iso_date(body.get("work_date") or detected_date, "Ngày phiên đơn")
        if not detected_date or work_date != detected_date:
            raise DailyImportError(
                "Ngày phiên đơn phải trùng ngày đã nhận diện trong workbook",
                code="work_date_mismatch",
                status=400,
            )
        import_key, source_hash, sheets = order_import_fingerprint(
            pending["path"], work_date, sheets,
        )
        if source_hash != analysis["sourceHash"]:
            raise DailyImportError(
                "Nội dung workbook đã thay đổi sau preview",
                code="source_changed",
            )
        orders, skipped = parse_workbook(pending["path"], work_date, sheets)
        if not orders:
            raise DailyImportError(
                "Sheet ngày đã chọn không có dòng đơn để nhập",
                code="empty_daily_sheet",
                status=400,
            )
        contract_rows = strict_daily_contract_rows(orders)
        with db() as conn:
            conn.execute("BEGIN IMMEDIATE")
            existing = valid_order_import_batch(conn, import_key)
            if existing:
                batch_id = int(existing["id"])
                payload = batch_payload(conn, batch_id)
                payload.update(
                    ok=True,
                    skippedSheets=skipped,
                    selectedSheets=sheets,
                    sourceHash=source_hash,
                    importKey=import_key,
                    idempotent=True,
                    dailyImport=strict_daily_version_payload(
                        conn, batch_id, source_hash, sheets[0],
                    ),
                )
                return jsonify(payload)

            conn.execute("DELETE FROM order_import_receipts WHERE import_key=?", (import_key,))
            contract = prepare_daily_import_version(
                conn,
                source_hash=source_hash,
                work_date=work_date,
                day_sheet=sheets[0],
                phase="first_load",
                rows_by_scope={"customer_orders": contract_rows},
                now_iso=now_iso,
            )
            batch_id = int(contract["batch_id"])
            latest_version = int(conn.execute(
                "SELECT MAX(version_no) n FROM daily_import_versions WHERE daily_workday_id=?",
                (int(contract["daily_workday_id"]),),
            ).fetchone()["n"])
            if contract["idempotent"] and int(contract["version_no"]) != latest_version:
                raise DailyImportError(
                    "Bản workbook này đã bị một phiên bản mới hơn thay thế; không thể âm thầm khôi phục dữ liệu cũ",
                    code="superseded_version",
                )
            current_order_count = int(conn.execute(
                "SELECT COUNT(*) n FROM orders WHERE batch_id=?", (batch_id,)
            ).fetchone()["n"])
            is_existing_current = bool(contract["idempotent"] and current_order_count)
            if not is_existing_current:
                if current_order_count:
                    old_ids = [r["id"] for r in conn.execute("SELECT id FROM orders WHERE batch_id=?", (batch_id,))]
                    blocked = batch_mutation_blocker(conn, batch_id)
                    if blocked or strict_order_removal_conflicts(conn, old_ids):
                        raise DailyImportError(blocked or "Bản cũ đã liên kết đặt NCC/hóa đơn; không thể tự thay toàn bộ đơn", code="linked_order_replacement")
                    if any(item["errors"] for item in orders):
                        raise DailyImportError("File mới còn lỗi; giữ nguyên đơn cũ, sửa file rồi nạp lại", code="invalid_replacement", status=400)
                    old_orders = rows_dict(conn.execute('SELECT * FROM orders WHERE batch_id=?', (batch_id,)))
                    previous_keys = [r['import_key'] for r in conn.execute('SELECT import_key FROM order_import_receipts WHERE batch_id=?', (batch_id,))]
                    previous_source = conn.execute('SELECT source_name FROM batches WHERE id=?', (batch_id,)).fetchone()['source_name']
                    conn.execute('''INSERT INTO order_reimport_history(batch_id,previous_source,next_source,rows_json,previous_keys_json,created_at)
                        VALUES(?,?,?,?,?,?)''', (batch_id, previous_source, pending['name'], json.dumps(old_orders, ensure_ascii=False), json.dumps(previous_keys), now_iso()))
                # A changed workbook for the same workday replaces only the
                # lifecycle-owned batch contents. Historical version hashes remain.
                conn.execute("DELETE FROM order_import_receipts WHERE batch_id=?", (batch_id,))
                conn.execute("DELETE FROM orders WHERE batch_id=?", (batch_id,))
                sync_receivable_ledger(conn, timestamp=now_iso())
                conn.execute(
                    "UPDATE batches SET work_date=?,source_name=? WHERE id=?",
                    (work_date, pending["name"], batch_id),
                )
                save_imported_orders(conn, batch_id, orders)
            scope = next(
                item for item in contract["scopes"] if item["scope"] == "customer_orders"
            )
            confirm_daily_import_scope(
                conn,
                version_id=int(contract["version_id"]),
                scope="customer_orders",
                expected_state_hash=scope["state_hash"],
                now_iso=now_iso,
            )
            conn.execute(
                """INSERT INTO order_import_receipts(
                       import_key,source_hash,source_name,work_date,selected_sheets_json,batch_id,created_at
                   ) VALUES(?,?,?,?,?,?,?)""",
                (
                    import_key, source_hash, pending["name"], work_date,
                    json.dumps(sheets, ensure_ascii=False), batch_id, now_iso(),
                ),
            )
            sync_receivable_ledger(conn, timestamp=now_iso())
            daily_import = strict_daily_version_payload(
                conn, batch_id, source_hash, sheets[0],
            )
            if daily_import:
                daily_import["idempotent"] = is_existing_current
            payload = batch_payload(conn, batch_id)
            payload.update(
                ok=True,
                skippedSheets=skipped,
                selectedSheets=sheets,
                sourceHash=source_hash,
                importKey=import_key,
                idempotent=is_existing_current,
                dailyImport=daily_import,
            )
            return jsonify(payload)
    except DailyImportError as exc:
        return jsonify({"ok": False, "error": str(exc), "code": exc.code}), exc.status
    except (OSError, ValueError) as exc:
        return jsonify({"ok": False, "error": str(exc)}), 400
    finally:
        try:
            pending["path"].unlink()
        except OSError:
            pass


@app.post("/api/import/confirm")
def api_import_confirm():
    body = request.get_json(force=True) or {}
    token = clean_text(body.get("token"))
    with PENDING_IMPORT_LOCK:
        pending = PENDING_IMPORTS.pop(token, None)
    if not pending or not pending["path"].exists():
        return jsonify({"ok": False, "error": "Phiên chọn sheet đã hết hạn. Vui lòng chọn lại file."}), 400
    if pending.get("strict_analysis"):
        return confirm_strict_daily_import(pending, body)
    raw_sheets = body.get("sheets") or []
    if not isinstance(raw_sheets, list):
        with PENDING_IMPORT_LOCK:
            PENDING_IMPORTS[token] = pending
        return jsonify({"ok": False, "error": "Danh sách sheet không hợp lệ"}), 400
    # Preserve exact worksheet titles, including legal trailing spaces.
    sheets = sorted({
        str(value) for value in raw_sheets if value is not None and str(value) != ""
    })
    if not sheets:
        with PENDING_IMPORT_LOCK:
            PENDING_IMPORTS[token] = pending
        return jsonify({"ok": False, "error": "Cần chọn ít nhất một sheet"}), 400
    try:
        work_date = valid_iso_date(
            body.get("work_date") or date.today().isoformat(), "Ngày phiên đơn",
        )
        import_key, source_hash, sheets = order_import_fingerprint(
            pending["path"], work_date, sheets,
        )
    except (OSError, ValueError) as exc:
        with PENDING_IMPORT_LOCK:
            PENDING_IMPORTS[token] = pending
        return jsonify({"ok": False, "error": str(exc)}), 400
    try:
        orders, skipped = parse_workbook(pending["path"], work_date, sheets)
    finally:
        try:
            pending["path"].unlink()
        except OSError:
            pass
    if not orders:
        return jsonify({"ok": False, "error": "Các sheet đã chọn không có dòng đơn hợp lệ"}), 400
    with db() as conn:
        # Serialize receipt check + write. The unique key also protects a
        # future multi-process deployment that shares this SQLite database.
        conn.execute("BEGIN IMMEDIATE")
        existing = valid_order_import_batch(conn, import_key)
        if existing:
            payload = batch_payload(conn, existing["id"])
            payload.update(
                ok=True,
                skippedSheets=skipped,
                selectedSheets=sheets,
                sourceHash=source_hash,
                importKey=import_key,
                idempotent=True,
            )
            return jsonify(payload)
        if conn.execute("SELECT 1 FROM order_reimport_history h, json_each(h.previous_keys_json) j WHERE j.value=? LIMIT 1", (import_key,)).fetchone():
            return jsonify(ok=False, error="File này đã được bản mới hơn thay thế; không tự khôi phục đơn cũ. Chọn file hiện hành để tiếp tục"), 409
        conn.execute("DELETE FROM order_import_receipts WHERE import_key=?", (import_key,))
        # A later valid file of the same workday/sheet scope replaces the
        # previous import, not an additional sale. Independent manual batches
        # and other sheet scopes are retained. Never discard linked documents.
        candidates = rows_dict(conn.execute(
            """SELECT DISTINCT b.* FROM batches b JOIN order_import_receipts r ON r.batch_id=b.id
               WHERE r.work_date=? AND r.selected_sheets_json=?
                 AND NOT EXISTS(SELECT 1 FROM daily_workdays d WHERE d.batch_id=b.id)""",
            (work_date, json.dumps(sheets, ensure_ascii=False)),
        ))
        replacement = None
        if len(candidates) > 1:
            return jsonify(ok=False, error="Ngày này có nhiều bản nhập cũ cùng phạm vi; cần đối chiếu trước khi thay, không tự cộng thêm đơn"), 409
        if candidates:
            previous = candidates[0]
            batch_id = previous["id"]
            old_orders = rows_dict(conn.execute("SELECT * FROM orders WHERE batch_id=?", (batch_id,)))
            blocker = batch_mutation_blocker(conn, batch_id)
            if blocker or strict_order_removal_conflicts(conn, [r["id"] for r in old_orders]):
                return jsonify(ok=False, error=blocker or "Bản cũ đã gắn với đặt NCC/hóa đơn; cần đối chiếu các chứng từ trước khi thay"), 409
            if any(item["errors"] for item in orders):
                return jsonify(ok=False, error="File mới còn dòng lỗi; bản cũ được giữ nguyên. Sửa file rồi nạp lại"), 400
            previous_keys = [r['import_key'] for r in conn.execute('SELECT import_key FROM order_import_receipts WHERE batch_id=?', (batch_id,))]
            conn.execute("""INSERT INTO order_reimport_history(batch_id,previous_source,next_source,rows_json,previous_keys_json,created_at)
                VALUES(?,?,?,?,?,?)""", (batch_id, previous["source_name"], pending["name"], json.dumps(old_orders, ensure_ascii=False), json.dumps(previous_keys), now_iso()))
            conn.execute("DELETE FROM order_import_receipts WHERE batch_id=?", (batch_id,))
            conn.execute("DELETE FROM orders WHERE batch_id=?", (batch_id,))
            sync_receivable_ledger(conn, timestamp=now_iso())
            sync_payable_ledger(conn, timestamp=now_iso())
            save_imported_orders(conn, batch_id, orders)
            conn.execute("UPDATE batches SET source_name=?,status='draft',approved_at=NULL WHERE id=?", (pending["name"], batch_id))
            replacement = {"replaced_rows": len(old_orders), "new_rows": len(orders),
                           "message": "Đã thay các dòng thuộc cùng ngày và sheet đã chọn; giữ ngày/sheet khác, danh mục và lịch sử bản cũ."}
        else:
            batch_id = save_imported_batch(conn, orders, work_date, pending["name"])
        conn.execute(
            """INSERT INTO order_import_receipts(
                   import_key,source_hash,source_name,work_date,selected_sheets_json,batch_id,created_at
               ) VALUES(?,?,?,?,?,?,?)""",
            (
                import_key, source_hash, pending["name"], work_date,
                json.dumps(sheets, ensure_ascii=False), batch_id, now_iso(),
            ),
        )
        sync_receivable_ledger(conn, timestamp=now_iso())
        payload = batch_payload(conn, batch_id)
        payload.update(
            ok=True,
            skippedSheets=skipped,
            selectedSheets=sheets,
            sourceHash=source_hash,
            importKey=import_key,
            idempotent=False,
            replacement=replacement,
        )
        return jsonify(payload)


@app.post("/api/batches")
def api_create_batch():
    body = request.get_json(force=True) or {}
    try:
        work_date = valid_iso_date(
            body.get("work_date") or date.today().isoformat(), "Ngày phiên đơn",
        )
    except ValueError as exc:
        return jsonify({"ok": False, "error": str(exc)}), 400
    with db() as conn:
        cur = conn.execute(
            "INSERT INTO batches(work_date,source_name,status,created_at) VALUES(?,?,?,?)",
            (work_date, "Nhập tay", "draft", now_iso()),
        )
        return jsonify({"ok": True, **batch_payload(conn, cur.lastrowid)})


@app.post("/api/batches/<int:batch_id>/approve")
def api_approve_batch(batch_id):
    with db() as conn:
        conn.execute("BEGIN IMMEDIATE")
        batch = conn.execute("SELECT work_date FROM batches WHERE id=?", (batch_id,)).fetchone()
        if not batch:
            return jsonify({"ok": False, "error": "Không tìm thấy phiên đơn"}), 404
        try:
            valid_iso_date(batch["work_date"], "Ngày phiên đơn")
        except ValueError as exc:
            return jsonify({"ok": False, "error": str(exc)}), 400
        mismatched_dates = conn.execute(
            "SELECT COUNT(*) n FROM orders WHERE batch_id=? AND work_date!=?",
            (batch_id, batch["work_date"]),
        ).fetchone()["n"]
        if mismatched_dates:
            return jsonify({
                "ok": False,
                "error": f"Còn {mismatched_dates} dòng có ngày không trùng phiên đơn",
            }), 400
        bad = conn.execute(
            "SELECT COUNT(*) n FROM orders WHERE batch_id=? AND errors!='[]'", (batch_id,)
        ).fetchone()["n"]
        total = conn.execute("SELECT COUNT(*) n FROM orders WHERE batch_id=?", (batch_id,)).fetchone()["n"]
        if total == 0:
            return jsonify({"ok": False, "error": "Đơn đang trống"}), 400
        if bad:
            return jsonify({"ok": False, "error": f"Còn {bad} dòng lỗi cần xử lý"}), 400
        conn.execute(
            "UPDATE batches SET status='approved',approved_at=? WHERE id=?",
            (now_iso(), batch_id),
        )
        post_purchase_list_inventory(conn, batch_id, now_iso)
        payable_ledger = sync_payable_ledger(conn, timestamp=now_iso())
        receivable_ledger = sync_receivable_ledger(conn, timestamp=now_iso())
        return jsonify({
            "ok": True, **batch_payload(conn, batch_id),
            "payable_ledger": payable_ledger,
            "receivable_ledger": receivable_ledger,
        })


@app.post("/api/orders")
def api_add_order():
    body = request.get_json(force=True) or {}
    batch_id = int(body.get("batch_id") or 0)
    with db() as conn:
        conn.execute("BEGIN IMMEDIATE")
        batch = conn.execute("SELECT * FROM batches WHERE id=?", (batch_id,)).fetchone()
        if not batch:
            return jsonify({"ok": False, "error": "Không tìm thấy phiên đơn"}), 404
        blocked = batch_mutation_blocker(conn, batch_id)
        if blocked:
            return jsonify({"ok": False, "error": blocked}), 409
        resolved = resolve_order(conn, body, batch["work_date"], *product_lookup(conn))
        clear_batch_derived_inventory(conn, batch_id)
        cur = conn.execute(
            """INSERT INTO orders(
                batch_id,work_date,contractor,kitchen,product_code,product_name,qty,
                actual_received,actual_delivered,damaged_qty,supplier_return_qty,customer_return_qty,
                unit,supplier,buy_price,sell_price,tax,invoice_nature,
                purchase_list,seller,cccd,note,source_sheet,source_row,errors,warnings,updated_at
            ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (
                batch_id, resolved["work_date"], resolved["contractor"], resolved["kitchen"],
                resolved["product_code"], resolved["product_name"], resolved["qty"],
                resolved["actual_received"], resolved["actual_delivered"], resolved["damaged_qty"],
                resolved["supplier_return_qty"], resolved["customer_return_qty"], resolved["unit"],
                resolved["supplier"], resolved["buy_price"], resolved["sell_price"], resolved["tax"],
                resolved["invoice_nature"],
                resolved["purchase_list"], resolved["seller"], resolved["cccd"], resolved["note"],
                "Nhập tay", 0, json.dumps(resolved["errors"], ensure_ascii=False),
                json.dumps(resolved["warnings"], ensure_ascii=False), now_iso(),
            ),
        )
        conn.execute("UPDATE batches SET status='draft',approved_at=NULL WHERE id=?", (batch_id,))
        sync_payable_ledger(conn, timestamp=now_iso())
        sync_receivable_ledger(conn, timestamp=now_iso())
        return jsonify({"ok": True, "id": cur.lastrowid, **batch_payload(conn, batch_id)})


@app.post("/api/orders/bulk")
def api_add_orders_bulk():
    body = request.get_json(force=True) or {}
    batch_id = int(body.get("batch_id") or 0)
    text = clean_text(body.get("text"))
    if not text:
        return jsonify({"ok": False, "error": "Chưa có dữ liệu để dán"}), 400
    lines = [line for line in text.splitlines() if line.strip()]
    matrix = [[cell.strip() for cell in line.split("\t")] for line in lines]
    if not matrix:
        return jsonify({"ok": False, "error": "Không đọc được dữ liệu đã dán"}), 400

    header = {}
    for idx, value in enumerate(matrix[0]):
        field = canonical_header(value)
        if field and field not in header:
            header[field] = idx
    if len(header) >= 2:
        data_rows = matrix[1:]
    else:
        # Simple paste order: bếp, tên hàng, số lượng, NCC, giá mua, giá bán, thuế, ghi chú.
        fields = ["kitchen", "product_name", "qty", "supplier", "buy_price", "sell_price", "tax", "note"]
        header = {field: idx for idx, field in enumerate(fields)}
        data_rows = matrix

    with db() as conn:
        conn.execute("BEGIN IMMEDIATE")
        batch = conn.execute("SELECT * FROM batches WHERE id=?", (batch_id,)).fetchone()
        if not batch:
            return jsonify({"ok": False, "error": "Không tìm thấy phiên đơn"}), 404
        blocked = batch_mutation_blocker(conn, batch_id)
        if blocked:
            return jsonify({"ok": False, "error": blocked}), 409
        by_code, by_name = product_lookup(conn)
        inserted = 0
        for cells in data_rows:
            raw = {
                field: cells[idx] if idx < len(cells) else ""
                for field, idx in header.items()
            }
            if not any(clean_text(raw.get(key)) for key in ("product_name", "product_code", "kitchen")):
                continue
            resolved = resolve_order(conn, raw, batch["work_date"], by_code, by_name)
            if inserted == 0:
                clear_batch_derived_inventory(conn, batch_id)
            conn.execute(
                """INSERT INTO orders(
                    batch_id,work_date,contractor,kitchen,product_code,product_name,qty,
                    actual_received,actual_delivered,damaged_qty,supplier_return_qty,customer_return_qty,
                    unit,supplier,buy_price,sell_price,tax,invoice_nature,
                    purchase_list,seller,cccd,note,source_sheet,source_row,errors,warnings,updated_at
                ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (
                    batch_id, resolved["work_date"], resolved["contractor"], resolved["kitchen"],
                    resolved["product_code"], resolved["product_name"], resolved["qty"],
                    resolved["actual_received"], resolved["actual_delivered"], resolved["damaged_qty"],
                    resolved["supplier_return_qty"], resolved["customer_return_qty"], resolved["unit"],
                    resolved["supplier"], resolved["buy_price"], resolved["sell_price"], resolved["tax"],
                    resolved["invoice_nature"],
                    resolved["purchase_list"], resolved["seller"], resolved["cccd"], resolved["note"],
                    "Dán từ Excel", inserted + 1,
                    json.dumps(resolved["errors"], ensure_ascii=False),
                    json.dumps(resolved["warnings"], ensure_ascii=False), now_iso(),
                ),
            )
            inserted += 1
        if inserted == 0:
            return jsonify({"ok": False, "error": "Không có dòng hợp lệ để thêm"}), 400
        conn.execute("UPDATE batches SET status='draft',approved_at=NULL WHERE id=?", (batch_id,))
        sync_payable_ledger(conn, timestamp=now_iso())
        sync_receivable_ledger(conn, timestamp=now_iso())
        payload = batch_payload(conn, batch_id)
        payload.update(ok=True, inserted=inserted)
        return jsonify(payload)


@app.put("/api/orders/bulk-update")
def api_update_orders_bulk():
    """Atomically revalidate and update the editable order grid."""
    body = request.get_json(force=True) or {}
    batch_id = int(body.get("batch_id") or 0)
    patches = body.get("items")
    if not batch_id or not isinstance(patches, list) or not patches:
        return jsonify({"ok": False, "error": "Cần phiên đơn và danh sách dòng cần cập nhật"}), 400
    if len(patches) > 5000:
        return jsonify({"ok": False, "error": "Mỗi lần chỉ cập nhật tối đa 5.000 dòng"}), 413
    try:
        order_ids = [int(item.get("id") or 0) for item in patches if isinstance(item, dict)]
    except (TypeError, ValueError):
        return jsonify({"ok": False, "error": "Danh sách có mã dòng không hợp lệ"}), 400
    if len(order_ids) != len(patches) or any(order_id <= 0 for order_id in order_ids):
        return jsonify({"ok": False, "error": "Mỗi dòng cập nhật phải có mã id hợp lệ"}), 400
    if len(set(order_ids)) != len(order_ids):
        return jsonify({"ok": False, "error": "Danh sách cập nhật bị trùng dòng"}), 400

    with db() as conn:
        conn.execute("BEGIN IMMEDIATE")
        batch = conn.execute("SELECT * FROM batches WHERE id=?", (batch_id,)).fetchone()
        if not batch:
            return jsonify({"ok": False, "error": "Không tìm thấy phiên đơn"}), 404
        blocked = batch_mutation_blocker(conn, batch_id)
        if blocked:
            return jsonify({"ok": False, "error": blocked}), 409
        placeholders = ",".join("?" for _ in order_ids)
        current_rows = {
            row["id"]: dict(row)
            for row in conn.execute(
                f"SELECT * FROM orders WHERE batch_id=? AND id IN ({placeholders})",
                (batch_id, *order_ids),
            )
        }
        missing = [order_id for order_id in order_ids if order_id not in current_rows]
        if missing:
            return jsonify({
                "ok": False,
                "error": "Một số dòng không còn thuộc phiên đang mở; hãy tải lại trước khi lưu",
                "missing_ids": missing[:20],
            }), 409
        direct_price_ids = direct_price_change_ids(current_rows, patches)
        if direct_price_ids:
            return jsonify({
                "ok": False,
                "error": "Giá bán chỉ được đổi qua lưới override có lý do, người thực hiện và revision",
                "code": "sell_price_override_required",
                "order_ids": direct_price_ids[:50],
            }), 409

        lookup = product_lookup(conn)
        prepared = []
        for patch_item, order_id in zip(patches, order_ids):
            item = current_rows[order_id]
            for key in ORDER_FIELDS:
                if key in patch_item:
                    item[key] = patch_item[key]
            resolved = resolve_order(
                conn, item, batch["work_date"], *lookup
            )
            for key in ORDER_FIELDS:
                if key in resolved:
                    item[key] = resolved[key]
            prepared.append((order_id, item, resolved["errors"], resolved["warnings"]))

        clear_batch_derived_inventory(conn, batch_id)
        columns = [
            "work_date", "contractor", "kitchen", "product_code", "product_name", "qty",
            "actual_received", "actual_delivered", "unit", "supplier", "buy_price",
            "sell_price", "tax", "invoice_nature", "purchase_list", "seller", "cccd", "note",
            "damaged_qty", "supplier_return_qty", "customer_return_qty",
        ]
        for order_id, item, errors, warnings in prepared:
            conn.execute(
                f"UPDATE orders SET {','.join(f'{column}=?' for column in columns)},"
                "errors=?,warnings=?,updated_at=? WHERE id=? AND batch_id=?",
                tuple(item[column] for column in columns) + (
                    json.dumps(errors, ensure_ascii=False),
                    json.dumps(warnings, ensure_ascii=False), now_iso(), order_id, batch_id,
                ),
            )
        conn.execute("UPDATE batches SET status='draft',approved_at=NULL WHERE id=?", (batch_id,))
        sync_payable_ledger(conn, timestamp=now_iso())
        sync_receivable_ledger(conn, timestamp=now_iso())
        payload = batch_payload(conn, batch_id)
        payload.update(
            ok=True,
            updated=len(prepared),
            error_rows=sum(bool(errors) for _, _, errors, _ in prepared),
            warning_rows=sum(bool(warnings) for _, _, _, warnings in prepared),
        )
        return jsonify(payload)


@app.put("/api/orders/<int:order_id>")
def api_update_order(order_id):
    body = request.get_json(force=True) or {}
    with db() as conn:
        conn.execute("BEGIN IMMEDIATE")
        current = conn.execute("SELECT * FROM orders WHERE id=?", (order_id,)).fetchone()
        if not current:
            return jsonify({"ok": False, "error": "Không tìm thấy dòng đơn"}), 404
        direct_price_ids = direct_price_change_ids(
            {order_id: dict(current)}, [{**body, "id": order_id}],
        )
        if direct_price_ids:
            return jsonify({
                "ok": False,
                "error": "Giá bán chỉ được đổi qua lưới override có lý do, người thực hiện và revision",
                "code": "sell_price_override_required",
            }), 409
        blocked = batch_mutation_blocker(conn, current["batch_id"])
        if blocked:
            return jsonify({"ok": False, "error": blocked}), 409
        item = dict(current)
        for key in ORDER_FIELDS:
            if key in body:
                item[key] = body[key]
        batch = conn.execute(
            "SELECT work_date FROM batches WHERE id=?", (current["batch_id"],),
        ).fetchone()
        resolved = resolve_order(
            conn, item, batch["work_date"], *product_lookup(conn)
        )
        for key in ORDER_FIELDS:
            if key in resolved:
                item[key] = resolved[key]
        errors = resolved["errors"]
        warnings = resolved["warnings"]
        clear_batch_derived_inventory(conn, current["batch_id"])
        columns = [
            "work_date", "contractor", "kitchen", "product_code", "product_name", "qty",
            "actual_received", "actual_delivered", "unit", "supplier", "buy_price",
            "sell_price", "tax", "invoice_nature", "purchase_list", "seller", "cccd", "note",
            "damaged_qty", "supplier_return_qty", "customer_return_qty",
        ]
        conn.execute(
            f"UPDATE orders SET {','.join(f'{col}=?' for col in columns)},errors=?,warnings=?,updated_at=? WHERE id=?",
            tuple(item[col] for col in columns) + (
                json.dumps(errors, ensure_ascii=False),
                json.dumps(warnings, ensure_ascii=False), now_iso(), order_id
            ),
        )
        conn.execute("UPDATE batches SET status='draft',approved_at=NULL WHERE id=?", (current["batch_id"],))
        sync_payable_ledger(conn, timestamp=now_iso())
        sync_receivable_ledger(conn, timestamp=now_iso())
        return jsonify({"ok": True, **batch_payload(conn, current["batch_id"])})


@app.delete("/api/orders/<int:order_id>")
def api_delete_order(order_id):
    with db() as conn:
        conn.execute("BEGIN IMMEDIATE")
        current = conn.execute("SELECT batch_id FROM orders WHERE id=?", (order_id,)).fetchone()
        if not current:
            return jsonify({"ok": False, "error": "Không tìm thấy dòng đơn"}), 404
        blocked = batch_mutation_blocker(conn, current["batch_id"])
        if blocked:
            return jsonify({"ok": False, "error": blocked}), 409
        clear_batch_derived_inventory(conn, current["batch_id"])
        conn.execute("DELETE FROM orders WHERE id=?", (order_id,))
        conn.execute("UPDATE batches SET status='draft',approved_at=NULL WHERE id=?", (current["batch_id"],))
        sync_payable_ledger(conn, timestamp=now_iso())
        sync_receivable_ledger(conn, timestamp=now_iso())
        return jsonify({"ok": True, **batch_payload(conn, current["batch_id"])})


@app.post("/api/payments")
def api_payment():
    body = request.get_json(silent=True) or {}
    kind = clean_text(body.get("kind")).lower()
    party_type = clean_text(body.get("party_type")).lower()
    expected_party_type = {"receipt": "contractor", "payment": "supplier"}.get(kind)
    if expected_party_type is None:
        return jsonify({"ok": False, "error": "Loại thanh toán chỉ nhận thu khách hoặc trả nhà cung cấp"}), 400
    if party_type != expected_party_type:
        return jsonify({"ok": False, "error": "Nhóm đối tượng không khớp loại thu/chi"}), 400
    party_code_input = clean_text(body.get("party_code"))
    if not party_code_input:
        return jsonify({"ok": False, "error": "Thiếu mã nhà thầu/nhà cung cấp"}), 400
    try:
        payment_date = valid_iso_date(body.get("payment_date"), "Ngày thanh toán")
        amount = finite_number(body.get("amount"), "Số tiền")
    except ValueError as exc:
        return jsonify({"ok": False, "error": str(exc)}), 400
    if amount <= 0:
        return jsonify({"ok": False, "error": "Số tiền thanh toán phải lớn hơn 0"}), 400
    if kind == "payment":
        try:
            with db() as conn:
                conn.execute("BEGIN IMMEDIATE")
                result = create_payable_payment(
                    conn, body, timestamp=now_iso(),
                    canonical_party_code=canonical_party_code,
                    audit_event=audit_event,
                )
            return jsonify({"ok": True, **result}), 200 if result["idempotent"] else 201
        except PayablePaymentError as exc:
            return jsonify({"ok": False, "error": str(exc), "code": exc.code}), exc.status
    try:
        with db() as conn:
            conn.execute("BEGIN IMMEDIATE")
            result = customer_receipt(conn, body, round3_context())
        return jsonify({"ok": True, **result}), 200 if result["idempotent"] else 201
    except ValueError as exc:
        return jsonify({"ok": False, "error": str(exc)}), 400


@app.delete("/api/payments/<int:payment_id>")
def api_delete_payment(payment_id):
    with db() as conn:
        current = conn.execute("SELECT kind,party_type FROM payments WHERE id=?", (payment_id,)).fetchone()
        if not current:
            return jsonify({"ok": False, "error": "Không tìm thấy giao dịch thanh toán"}), 404
        supplier = current["kind"] == "payment" and current["party_type"] == "supplier"
        return jsonify({
            "ok": False, "error": "Giao dịch phải được hoàn tác có lý do, không được xóa",
            "code": "payable_reversal_required" if supplier else "receipt_reversal_required",
            "reversal_endpoint": f"/api/debts/{'payables/payments' if supplier else 'receipts'}/{payment_id}/reverse",
        }), 409


@app.post("/api/balances")
def api_balance():
    body = request.get_json(silent=True) or {}
    party_type = clean_text(body.get("party_type")).lower()
    party_code_input = clean_text(body.get("party_code"))
    if party_type not in {"contractor", "supplier"}:
        return jsonify({"ok": False, "error": "Nhóm số dư chỉ nhận nhà thầu hoặc nhà cung cấp"}), 400
    if not party_code_input:
        return jsonify({"ok": False, "error": "Thiếu mã nhà thầu/nhà cung cấp"}), 400
    try:
        as_of_date = valid_iso_date(
            clean_text(body.get("as_of_date")) or "1900-01-01", "Ngày chốt số dư"
        )
        opening = finite_number(body.get("opening"), "Số dư đầu kỳ")
    except ValueError as exc:
        return jsonify({"ok": False, "error": str(exc)}), 400
    with db() as conn:
        try:
            party_code = canonical_party_code(conn, party_type, party_code_input)
        except ValueError as exc:
            return jsonify({"ok": False, "error": str(exc)}), 400
        conn.execute(
            "INSERT INTO balances(party_type,party_code,opening,as_of_date) VALUES(?,?,?,?) "
            "ON CONFLICT(party_type,party_code) DO UPDATE SET "
            "opening=excluded.opening,as_of_date=excluded.as_of_date",
            (party_type, party_code, opening, as_of_date),
        )
        audit_event(
            conn, "balance.upsert", entity_type=party_type, entity_id=party_code,
            metadata={"opening": opening, "as_of_date": as_of_date},
        )
    return jsonify({"ok": True})


@app.post("/api/master/sync")
def api_master_sync():
    sync_master_if_needed(force=True)
    return jsonify({"ok": True, "message": "Đã đồng bộ danh mục từ Em Thành.xlsx"})


@app.get("/api/products/search")
def api_product_search():
    term = request.args.get("q", "").strip()
    with db() as conn:
        candidates = rows_dict(conn.execute(
            """SELECT p.*,COALESCE(o.invoice_name,'') invoice_name
               FROM products p
               LEFT JOIN outgoing_product_names o ON o.product_code=p.code
               ORDER BY p.name,p.code"""
        ))
        if not term:
            rows = candidates[:30]
        else:
            needle = slug(term)
            ranked = []
            for row in candidates:
                code_key = slug(row.get("code"))
                name_keys = [slug(row.get("name")), slug(row.get("invoice_name"))]
                searchable = [code_key, *[key for key in name_keys if key]]
                contains = any(needle and needle in value for value in searchable)
                ratio = max(
                    (SequenceMatcher(None, needle, value).ratio() for value in searchable if value),
                    default=0,
                )
                if contains or ratio >= 0.42:
                    ranked.append((
                        0 if code_key == needle else 1 if contains else 2,
                        -ratio,
                        row["name"],
                        row,
                    ))
            rows = [item[3] for item in sorted(ranked, key=lambda item: item[:3])[:30]]
        return jsonify({"ok": True, "items": rows})


def workbook_bytes(wb: Workbook):
    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream


def set_title(ws, text, subtitle="", end_col=8):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_col)
    cell = ws.cell(1, 1, text)
    cell.font = Font(name="Times New Roman", size=16, bold=True)
    cell.alignment = Alignment(
        horizontal="center", vertical="center", wrap_text=True,
        shrink_to_fit=True,
    )
    ws.row_dimensions[1].height = 30
    if subtitle:
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end_col)
        ws.cell(2, 1, subtitle)
        ws.cell(2, 1).font = Font(name="Times New Roman", size=11, italic=True)
        ws.cell(2, 1).alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True,
            shrink_to_fit=True,
        )


def style_table(ws, header_row, end_col, last_row=None):
    last_row = last_row or ws.max_row
    thin = Side(style="thin", color="000000")
    for cell in ws[header_row][:end_col]:
        cell.font = Font(name="Times New Roman", size=11, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    for row in ws.iter_rows(min_row=header_row + 1, max_row=last_row, min_col=1, max_col=end_col):
        for cell in row:
            cell.font = Font(name="Times New Roman", size=11)
            cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    if last_row >= header_row:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(end_col)}{last_row}"
    ws.freeze_panes = f"A{header_row + 1}"


def autosize(ws, max_width=48, min_row=1):
    """Size columns from tabular rows, optionally ignoring merged page titles.

    A merged title is stored in the first cell of the merged range.  Counting
    that text as column-A data makes narrow columns such as ``STT`` several
    times wider than the actual item-name column when the workbook is printed.
    """
    for col in range(1, ws.max_column + 1):
        width = 8
        for row in range(max(1, min_row), min(ws.max_row, 180) + 1):
            value = ws.cell(row, col).value
            if value is not None:
                lines = str(value).splitlines() or [""]
                width = max(width, max(len(line) for line in lines))
        ws.column_dimensions[get_column_letter(col)].width = min(width + 2, max_width)


def money_format(wb):
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, (int, float)) and cell.column > 1:
                    if any(key in clean_text(ws.cell(4 if ws.max_row >= 4 else 1, cell.column).value).lower()
                           for key in ("giá", "tiền", "thu", "trả", "vốn", "lợi nhuận")):
                        cell.number_format = "#,##0"


def require_batch(conn, batch_id):
    batch = conn.execute("SELECT * FROM batches WHERE id=?", (batch_id,)).fetchone()
    if not batch:
        raise ValueError("Không tìm thấy phiên đơn")
    orders = rows_dict(conn.execute("SELECT * FROM orders WHERE batch_id=? ORDER BY id", (batch_id,)))
    for item in orders:
        item["errors"] = json.loads(item["errors"] or "[]")
        item["warnings"] = json.loads(item["warnings"] or "[]")
    if not orders:
        raise ValueError("Phiên đơn đang trống")
    return dict(batch), orders


def batch_mutation_blocker(conn, batch_id: int) -> str:
    """Explain why an order batch is immutable once downstream documents exist."""
    draft = conn.execute(
        """SELECT status,minvoice_status FROM outgoing_invoice_drafts
           WHERE batch_id=? AND status!='cancelled'
           ORDER BY CASE status WHEN 'issued' THEN 0 ELSE 1 END,id LIMIT 1""",
        (batch_id,),
    ).fetchone()
    if draft:
        if draft["status"] == "issued":
            return "Phiên đã có hóa đơn phát hành nên không được sửa/xóa dòng đơn"
        if draft["minvoice_status"] in {"saved", "saving", "unknown"}:
            return "Dự thảo đã lưu hoặc đang đối soát M-Invoice; xử lý dứt điểm trên M-Invoice trước khi sửa đơn"
        return "Phiên đang có dự thảo hóa đơn giữ tồn; hãy hủy dự thảo trước khi sửa đơn"
    print_job = conn.execute(
        """SELECT status FROM print_jobs
           WHERE batch_id=? AND status NOT IN ('stale','cancelled')
           ORDER BY id DESC LIMIT 1""",
        (batch_id,),
    ).fetchone()
    if print_job:
        if print_job["status"] in {"submitting", "submitted", "submission_unknown", "printed"}:
            return "Bộ chứng từ đã gửi sang máy in nên phiên được khóa để bảo toàn lịch sử"
        return "Phiên đang có bộ PDF in đã chuẩn bị/duyệt; hãy hủy bộ in cũ trước khi sửa đơn"
    return ""


def clear_batch_derived_inventory(conn, batch_id: int) -> None:
    """Purge retired auto-generated BK rows while a batch is being edited."""
    conn.execute(
        "DELETE FROM inventory_transactions WHERE source_type='BK_INPUT' AND source_id=?",
        (str(batch_id),),
    )


def send_xlsx(wb, filename):
    return send_file(
        workbook_bytes(wb),
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def export_supplier_orders(conn, batch, orders):
    """Build the customer's canonical, editable ``đặt hàng`` worksheet.

    Columns A/B/Q are hidden compatibility/identity columns. The visible
    business layout remains the customer's C:P sheet and never exposes or
    accepts a sales price.
    """
    if not DAILY_ORDER_TEMPLATE_SOURCE.is_file():
        raise FileNotFoundError(
            "Thiếu mẫu đơn đặt hàng đã chốt; không tạo file bằng mẫu tự đoán"
        )
    wb = load_workbook(DAILY_ORDER_TEMPLATE_SOURCE, data_only=False, read_only=False)
    try:
        from document_preview import white_print_style
        from document_totals import quantity_cell
    except ImportError:
        from .document_preview import white_print_style
        from .document_totals import quantity_cell
    ws = wb.active
    payload = purchase_order_payload(conn, int(batch["id"]))
    headers = [
        "Mã hàngNCC", "Mã hàng", "Mã bếp", "", "Tên hàng ", "Số lượng",
        "ĐVT", "NCC", "ghi chú", "giá mua", "hỏng", "thêm", "Giảm",
        "thiếu", "SL \nthực té", "Thành tiền", "Mã dòng hệ thống",
    ]
    prototype_styles = [copy.copy(ws.cell(3, column)._style) for column in range(1, 18)]
    if ws.max_row > 3:
        ws.delete_rows(4, ws.max_row - 3)
    for column, label in enumerate(headers, 1):
        ws.cell(2, column, label)
    occurrences = Counter()
    for output_index, item in enumerate(payload["rows"], start=3):
        identity = (
            clean_text(item.get("product_code")).casefold(),
            clean_text(item.get("kitchen")).casefold(),
        )
        occurrences[identity] += 1
        row_key = item.get("row_key") or purchase_business_row_key(
            item.get("work_date") or batch["work_date"], item.get("product_code"),
            item.get("kitchen"), item.get("product_name"), item.get("unit"),
            occurrences[identity],
        )
        row_number = output_index
        base_qty = item.get("demand_qty", item.get("order_qty", 0))
        damaged = item.get("damaged_qty", 0)
        added = item.get("added_qty", 0)
        reduced = item.get("reduced_qty", 0)
        missing = item.get("missing_qty", 0)
        values = [
            f"=B{row_number}&H{row_number}", item.get("product_code"), item.get("kitchen"),
            datetime.strptime(
                item.get("work_date") or batch["work_date"], "%Y-%m-%d"
            ).strftime("%d.%m.%Y"), item.get("product_name"), base_qty,
            item.get("unit"), item.get("supplier"), item.get("note"), item.get("buy_price"),
            damaged, added, reduced, missing,
            f"=F{row_number}+L{row_number}-K{row_number}-M{row_number}-N{row_number}",
            f"=IFERROR(O{row_number}*J{row_number},0)", row_key,
        ]
        for column, value in enumerate(values, 1):
            cell = ws.cell(row_number, column, value)
            cell._style = copy.copy(prototype_styles[column - 1])
            if isinstance(value, str) and column not in (1, 15, 16):
                cell.data_type = 's'
    last_row = max(ws.max_row, 2)
    ws["A1"] = "=B1&H1"
    ws["F1"] = f"=SUBTOTAL(9,F3:F{last_row})"
    ws["P1"] = f"=SUBTOTAL(9,P3:P{last_row})"
    qty_total = quantity_cell([{'quantity':row.get('demand_qty',row.get('order_qty',0)),
        'unit':row.get('unit','')} for row in payload['rows']], 'quantity')
    if isinstance(qty_total,str):
        ws['F1'] = qty_total
        ws['F1'].alignment = Alignment(horizontal='right',vertical='center',wrap_text=True)
        ws.row_dimensions[1].height = max(26, 18 * math.ceil(len(qty_total)/12))
    for row_index in range(3, ws.max_row + 1):
        for column_index in (6, 11, 12, 13, 14, 15):
            # Preserve the customer's accounting formats and only make the
            # quantity precision explicit where the source row has no format.
            ws.cell(row_index, column_index).number_format = "#,##0.######"
        for column_index in (10, 16):
            ws.cell(row_index, column_index).number_format = "#,##0"
        ws.row_dimensions[row_index].height = max(24,18*math.ceil(len(str(ws.cell(row_index,5).value or ''))/32))
        ws.cell(row_index,5).alignment=Alignment(horizontal='left',vertical='center',wrap_text=True)
    for column in ("A", "B", "Q"):
        ws.column_dimensions[column].hidden = True
    ws.auto_filter.ref = f"A2:Q{last_row}"
    ws.freeze_panes = "C3"
    ws.print_area = f"A1:Q{last_row}"
    ws.print_title_rows = "$2:$2"
    ws.sheet_view.showGridLines = True
    try:
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
        wb.calculation.calcMode = "auto"
    except AttributeError:
        pass
    white_print_style(wb)
    return wb


def export_deliveries(conn, batch, orders, kitchen_codes=None):
    selected_kitchens = {
        clean_text(value).upper() for value in (kitchen_codes or []) if clean_text(value)
    }
    groups = defaultdict(list)
    for item in orders:
        delivered = net_delivered(item)
        kitchen = clean_text(item["kitchen"]).upper() or "CHƯA XÁC ĐỊNH"
        if delivered > 0 and (not selected_kitchens or kitchen in selected_kitchens):
            groups[kitchen].append(
                (item, delivered)
            )
    deliveries = []
    for kitchen, rows in sorted(groups.items()):
        meta_row = conn.execute(
            "SELECT code,contractor,name,address FROM kitchens WHERE code=?", (kitchen,)
        ).fetchone()
        meta = dict(meta_row) if meta_row else {
            "code": kitchen, "contractor": "", "name": kitchen, "address": "",
        }
        row_contractors = {
            clean_text(item["contractor"]).upper() for item, _ in rows
            if clean_text(item["contractor"])
        }
        contractor = clean_text(meta.get("contractor")).upper()
        if not contractor and len(row_contractors) == 1:
            contractor = next(iter(row_contractors))
        recipient = clean_text(
            setting_get(conn, f"delivery_recipient_{kitchen}", "")
        ) or clean_text(meta.get("name")) or kitchen
        address = clean_text(
            setting_get(conn, f"delivery_address_{kitchen}", "")
        ) or clean_text(meta.get("address"))
        show_price = delivery_prices_visible(kitchen, contractor)
        items = []
        for item, delivered in rows:
            record = {
                "product_code": clean_text(item["product_code"]).upper(),
                "product_name": clean_text(item["product_name"]),
                "quantity": delivered,
                "unit": clean_text(item["unit"]),
                "note": clean_text(item["note"]),
            }
            if show_price:
                record["sell_price"] = item["sell_price"]
            items.append(record)
        deliveries.append({
            "kitchen": kitchen,
            "contractor": contractor,
            "recipient": recipient,
            "address": address,
            "items": items,
        })
    return build_delivery_workbook(
        deliveries,
        work_date=batch["work_date"],
        template_path=MASTER_SOURCE,
    )


def export_report(conn, batch, orders):
    period, report_rows = collect_monthly_report_rows(
        conn, batch, orders, totals_fn=order_totals,
    )
    present = {clean_text(row["kitchen"]).casefold() for row in report_rows}
    for kitchen in conn.execute("SELECT code,contractor FROM kitchens ORDER BY contractor,code"):
        if clean_text(kitchen["code"]).casefold() not in present and clean_text(kitchen["contractor"]):
            report_rows.append({"kitchen": kitchen["code"], "contractor": kitchen["contractor"],
                                "revenue": 0, "cost": 0, "profit": 0, "total": 0})
    configured_groups = {
        clean_text(row["key"])[len("report_group_"):].upper(): clean_text(row["value"])
        for row in conn.execute(
            "SELECT key,value FROM settings WHERE key LIKE 'report_group_%' ORDER BY key"
        )
        if clean_text(row["value"])
    }
    return build_monthly_report_workbook(
        report_rows,
        period=period,
        template_path=MASTER_SOURCE,
        configured_groups=configured_groups,
    )


def export_purchase_documents(conn, batch, orders):
    """Create the golden purchase summary and one receipt per seller/day."""

    excluded = []
    rows = collect_purchase_summary_rows(conn, batch, orders, excluded_rows=excluded)
    receipt_rows = enrich_receipt_identity_rows(conn, rows)
    workbook = build_purchase_documents_workbook(
        receipt_rows,
        template_path=MASTER_SOURCE,
        date_from=batch["work_date"],
        date_to=batch["work_date"],
        buyer_name=setting_get(conn, "purchase_receipt_buyer_name", ""),
        buyer_title=setting_get(conn, "purchase_receipt_buyer_title", ""),
        company_name=setting_get(conn, "company", ""),
        company_address=setting_get(conn, "company_address", ""),
        location=setting_get(conn, "purchase_receipt_location", "Hải Phòng"),
    )
    if excluded:
        from openpyxl.comments import Comment
        names = ", ".join(sorted({row["seller"] for row in excluded}))
        warning = (f"Không lập bảng kê/biên nhận cho {len(excluded)} dòng của {names}. "
                   "Tổng trên chứng từ chỉ gồm các dòng hợp lệ; dữ liệu gốc, kho và công nợ không bị thay đổi.")
        workbook._tdp_warnings = [warning]
        workbook.worksheets[0]["A1"].comment = Comment(warning, "TĐP")
    return workbook


def export_optional_purchase_documents(conn, batch, orders):
    """Return no workbook only when this batch has no eligible BK rows.

    A mixed print bundle may legitimately contain only sale-side documents.
    Invalid BK data remains a hard error so missing prices or seller identity
    cannot be hidden by silently omitting the purchase summary.
    """

    try:
        return export_purchase_documents(conn, batch, orders)
    except PurchaseSummaryError as error:
        if error.code == "no_purchase_summary_rows":
            return None
        raise


def export_outgoing_statement(conn, batch, orders):
    """Create the human-readable outgoing-goods statement for one order batch.

    This is deliberately a review document, so it can be downloaded before an
    invoice draft is created.  Official invoice files remain protected by the
    approval and inventory-reservation guards in ``api_export``.
    """
    invoice_names = {
        row["product_code"]: row["invoice_name"]
        for row in conn.execute("SELECT product_code,invoice_name FROM outgoing_product_names")
    }
    groups = defaultdict(list)
    for item in orders:
        if net_delivered(item) > 1e-9:
            groups[item["contractor"] or "KHÁC"].append(item)

    wb = Workbook()
    summary = wb.active
    summary.title = "Tổng hợp"
    status_text = "Đã duyệt" if batch["status"] == "approved" else "Bản nháp để kiểm tra"
    set_title(
        summary,
        "BẢNG KÊ HÀNG HÓA ĐẦU RA",
        f"Ngày {display_date_vn(batch['work_date'])} · {status_text} · Theo số thực giao ròng",
        6,
    )
    summary.append([])
    summary.append([
        "Nhà thầu", "Số dòng", "Tiền trước thuế", "Tiền thuế",
        "Tổng thanh toán", "Trạng thái phiên",
    ])

    for contractor, rows in sorted(groups.items()):
        subtotal = 0
        tax_amount = 0
        for item in rows:
            amount = vnd_product(net_delivered(item), vnd_round(number_value(item["sell_price"])))
            vat_percent = invoice_vat_percent(item["tax"])
            subtotal += amount
            if vat_percent > 0:
                tax_amount += vnd_product(amount, vat_percent / 100)
        summary.append([
            contractor, len(rows), subtotal, tax_amount, subtotal + tax_amount, status_text,
        ])

        ws = wb.create_sheet(safe_sheet_name(contractor))
        set_title(
            ws,
            "BẢNG KÊ HÀNG HÓA ĐẦU RA",
            f"Nhà thầu {contractor} · Ngày {display_date_vn(batch['work_date'])} · Theo số thực giao ròng",
            11,
        )
        ws.append([])
        ws.append([
            "STT", "Bếp", "Mã hàng", "Tên hàng xuất hóa đơn", "ĐVT",
            "Số lượng thực giao", "Đơn giá", "Tiền trước thuế", "Thuế suất",
            "Tiền thuế", "Tổng thanh toán",
        ])
        contractor_subtotal = 0
        contractor_tax = 0
        for index, item in enumerate(rows, 1):
            qty = net_delivered(item)
            unit_price = vnd_round(number_value(item["sell_price"]))
            amount = vnd_product(qty, unit_price)
            vat_percent = invoice_vat_percent(item["tax"])
            line_tax = 0 if vat_percent <= 0 else vnd_product(amount, vat_percent / 100)
            tax_label = {-2: "KKKNT", -1: "KCT"}.get(vat_percent, f"{vat_percent:g}%")
            contractor_subtotal += amount
            contractor_tax += line_tax
            ws.append([
                index, item["kitchen"], item["product_code"],
                invoice_names.get(item["product_code"], item["product_name"]),
                item["unit"], qty, unit_price, amount, tax_label, line_tax, amount + line_tax,
            ])
        ws.append([
            "", "", "", "TỔNG CỘNG", "", "", "", contractor_subtotal, "",
            contractor_tax, contractor_subtotal + contractor_tax,
        ])
        style_table(ws, 4, 11)
        for col in (6, 7, 8, 10, 11):
            for row in range(5, ws.max_row + 1):
                ws.cell(row, col).number_format = "#,##0.##" if col == 6 else "#,##0"
        autosize(ws, min_row=4)
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
        ws.page_setup.scale = None
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.print_title_rows = "$4:$4"
        ws.print_area = f"A1:K{ws.max_row}"
        ws.page_margins.left = 0.25
        ws.page_margins.right = 0.25
        ws.page_margins.top = 0.45
        ws.page_margins.bottom = 0.45
        ws.sheet_view.showGridLines = False

    if not groups:
        summary.append(["Không có số thực giao dương", 0, 0, 0, 0, status_text])
    style_table(summary, 4, 6)
    for col in (3, 4, 5):
        for row in range(5, summary.max_row + 1):
            summary.cell(row, col).number_format = "#,##0"
    autosize(summary, min_row=4)
    summary.sheet_properties.pageSetUpPr.fitToPage = True
    summary.page_setup.paperSize = summary.PAPERSIZE_A4
    summary.page_setup.orientation = summary.ORIENTATION_LANDSCAPE
    summary.page_setup.scale = None
    summary.page_setup.fitToWidth = 1
    summary.page_setup.fitToHeight = 0
    summary.print_title_rows = "$4:$4"
    summary.print_area = f"A1:F{summary.max_row}"
    summary.page_margins.left = 0.25
    summary.page_margins.right = 0.25
    summary.page_margins.top = 0.45
    summary.page_margins.bottom = 0.45
    summary.sheet_view.showGridLines = False
    return wb


def invoice_vat_percent(value):
    """Return the exact percentage/code expected by the customer's 13-column template."""

    raw = clean_text(value).upper().replace(" ", "")
    aliases = {
        "KCT": -1,
        "KHÔNGCHỊUTHUẾ": -1,
        "KHONGCHIUTHUE": -1,
        "KKKNT": -2,
        "KHÔNGKÊKHAI": -2,
        "KHONGKEKHAI": -2,
    }
    if raw in aliases:
        return aliases[raw]
    has_percent = raw.endswith("%")
    numeric = number_value(raw.rstrip("%"), 0)
    if not has_percent and 0 < abs(numeric) < 1:
        numeric *= 100
    return round(numeric, 4)


def vnd_round(value):
    """Round VND with the commercial HALF_UP rule used by M-Invoice."""
    try:
        return int(Decimal(str(value)).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
    except (InvalidOperation, TypeError, ValueError):
        return 0


def vnd_product(*values):
    try:
        total = Decimal("1")
        for value in values:
            total *= Decimal(str(value))
        return int(total.quantize(Decimal("1"), rounding=ROUND_HALF_UP))
    except (InvalidOperation, TypeError, ValueError):
        return 0


def invoice_workbook(rows, invoice_names=None):
    invoice_names = invoice_names or {}
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(INVOICE_HEADERS)
    for item in rows:
        qty = net_delivered(item)
        unit_price = vnd_round(number_value(item["sell_price"]))
        amount = vnd_product(qty, unit_price)
        invoice_name = invoice_names.get(item["product_code"], item["product_name"])
        name_key = slug(f"{item['product_name']} {invoice_name} {item.get('note', '')}")
        is_promotion = clean_text(item.get("invoice_nature")) == "2" or (
            "khuyenmai" in name_key and unit_price <= 0
        )
        vat_percent = invoice_vat_percent(item["tax"])
        tax_amount = 0 if vat_percent <= 0 else vnd_product(amount, vat_percent / 100)
        total = amount + tax_amount
        if is_promotion:
            money_cells = [None, None, None, None, None]
            nature = "2"
        else:
            money_cells = [unit_price, amount, amount, tax_amount, total]
            nature = "1"
        ws.append([
            item["product_code"], invoice_name, item["unit"], qty,
            money_cells[0], money_cells[1], None, None, money_cells[2],
            vat_percent, money_cells[3], money_cells[4], nature,
        ])
    style_table(ws, 1, 13)
    for col in (4, 5, 6, 8, 9, 11, 12):
        for row in range(2, ws.max_row + 1):
            ws.cell(row, col).number_format = "#,##0"
    ws.column_dimensions["B"].width = max(ws.column_dimensions["B"].width or 0, 42)
    autosize(ws)
    return wb


def export_invoices_zip(conn, batch, orders):
    try:
        from outgoing_readiness import OutgoingReadinessError, validate_draft_export_stock
    except ImportError:
        from .outgoing_readiness import OutgoingReadinessError, validate_draft_export_stock
    # Export only quantities in the current, locally editable draft rounds.
    # Earlier issued/saved rounds and the unallocated remainder must never be
    # copied back from the full customer order into a new upload file.
    lines = [dict(row) for row in conn.execute(
        """SELECT l.id line_id,l.product_code,l.product_name,l.qty,l.unit,
                  l.unit_price,l.tax,l.invoice_nature,l.amount,
                  d.id draft_id,d.contractor,d.round_no,d.invoice_date,d.minvoice_status
             FROM outgoing_invoice_lines l
             JOIN outgoing_invoice_drafts d ON d.id=l.draft_id
            WHERE d.batch_id=? AND d.status='draft'
              AND COALESCE(d.minvoice_status,'not_sent') NOT IN ('saved','saving','unknown')
            ORDER BY d.contractor,d.round_no,l.id""",
        (batch["id"],),
    )]
    try:
        for draft_id, invoice_date in {(line["draft_id"], line["invoice_date"]) for line in lines}:
            # Tax upload templates contain no issue date. Validate today's
            # allocation here; the actual issue date is checked at confirmation.
            validate_draft_export_stock(conn, draft_id)
    except OutgoingReadinessError as error:
        raise InvoiceTaxExportError(str(error), code=error.code, status=error.status) from None
    for line in lines:
        line["vat_percent"] = invoice_vat_percent(line["tax"])
    return export_invoice_drafts_zip(
        lines,
        work_date=batch["work_date"],
        template_dir=TAX_TEMPLATE_DIR,
    )


def get_quote_rows(conn, contractor: str, batch_id=None, period="", version_id=None):
    row = conn.execute("SELECT * FROM contractors WHERE code=?", (contractor,)).fetchone()
    if not row:
        raise QuoteExportError(
            f"Không tìm thấy nhà thầu {contractor or '(trống)'}",
            code="contractor_not_found",
            status=404,
        )
    group = row["price_group"] or contractor
    mode = row["pricing_mode"]
    rows = []
    resolved_period = clean_text(period)
    if not resolved_period and batch_id:
        batch = conn.execute("SELECT work_date FROM batches WHERE id=?", (batch_id,)).fetchone()
        resolved_period = clean_text(batch["work_date"])[:7] if batch else ""
    if not resolved_period:
        resolved_period = date.today().strftime("%Y-%m")
    if mode == "daily":
        if version_id is not None:
            raise QuoteExportError(
                "Báo giá theo ngày không dùng phiên bản báo giá tháng",
                code="daily_quote_has_no_period_version",
                status=400,
            )
        if not batch_id:
            return mode, rows, {
                "period": resolved_period, "price_group": group, "version": None,
                "daily_source": None, "contractor_name": row["name"],
                "conflicts": [], "excluded_count": 0,
                "output_count": 0, "status_count": 0,
            }
        batch = conn.execute(
            "SELECT id,work_date,source_name FROM batches WHERE id=?", (batch_id,),
        ).fetchone()
        if not batch:
            raise QuoteExportError(
                "Không tìm thấy phiên đơn nguồn cho báo giá theo ngày",
                code="daily_batch_not_found",
                status=404,
            )
        resolved_period = clean_text(batch["work_date"])[:7]
        candidates = conn.execute(
            """SELECT id,product_code,product_name,unit,tax,sell_price,
                      source_sheet,source_row
               FROM orders
               WHERE batch_id=? AND contractor=?
               ORDER BY UPPER(product_code),id""",
            (batch_id, contractor),
        ).fetchall()
        by_code = defaultdict(list)
        conflicts = []
        for item in candidates:
            code = clean_text(item["product_code"]).upper()
            if not code:
                conflicts.append({
                    "product_code": "",
                    "source_rows": [int(item["source_row"])] if item["source_row"] is not None else [],
                    "order_ids": [int(item["id"])],
                    "reasons": ["thiếu mã hàng"],
                })
                continue
            by_code[code].append(item)
        for code, code_rows in by_code.items():
            def text_key(value):
                return re.sub(r"\s+", " ", clean_text(value)).casefold()

            dimensions = {
                "tên hàng khác nhau": {text_key(item["product_name"]) for item in code_rows},
                "ĐVT khác nhau": {text_key(item["unit"]) for item in code_rows},
                "thuế khác nhau": {text_key(item["tax"]) for item in code_rows},
                "giá bán theo ngày khác nhau": {number_value(item["sell_price"]) for item in code_rows},
            }
            reasons = [label for label, values in dimensions.items() if len(values) > 1]
            source_rows = sorted({
                int(item["source_row"]) for item in code_rows if item["source_row"] is not None
            })
            order_ids = [int(item["id"]) for item in code_rows]
            if reasons:
                conflicts.append({
                    "product_code": code,
                    "source_rows": source_rows,
                    "order_ids": order_ids,
                    "reasons": reasons,
                })
                continue
            selected = code_rows[0]
            sell_price = number_value(selected["sell_price"])
            if math.isfinite(float(sell_price)) and float(sell_price).is_integer():
                sell_price = int(sell_price)
            state = "zero" if sell_price == 0 else "numeric"
            rows.append({
                "product_code": code,
                "product_name": selected["product_name"],
                "unit": selected["unit"],
                "tax": selected["tax"],
                "sell_price": sell_price,
                "status": "Giá 0 theo ngày – giữ để xác nhận" if state == "zero" else "Giá theo ngày",
                "price_state": state,
                "exportable": True,
                "source_row": selected["source_row"],
                "source_rows": source_rows,
                "order_ids": order_ids,
                "duplicate_count": max(0, len(code_rows) - 1),
            })
        rows.sort(key=lambda item: (clean_text(item["product_name"]).casefold(), item["product_code"]))
        return mode, rows, {
            "period": resolved_period, "price_group": group, "version": None,
            "daily_source": {
                "batch_id": int(batch["id"]),
                "work_date": clean_text(batch["work_date"]),
                "source_name": clean_text(batch["source_name"]),
            },
            "contractor_name": row["name"],
            "conflicts": conflicts, "excluded_count": 0,
            "output_count": len(rows), "status_count": 0,
        }
    else:
        versioned = quote_rows_for_contractor(
            conn, contractor, resolved_period, version_id=version_id,
        )
        if versioned is not None:
            versioned["daily_source"] = None
            versioned["contractor_name"] = row["name"]
            return mode, versioned["items"], versioned
        if version_id is not None:
            raise QuoteExportError(
                "Không tìm thấy phiên bản báo giá đã xác nhận trong kỳ đã chọn",
                code="quote_version_not_found",
                status=404,
            )
        excluded_count = 0
        for item in conn.execute(
            """SELECT p.code product_code,p.name product_name,p.unit,p.tax,
                      pp.price_value sell_price,pp.price_text
               FROM products p JOIN product_prices pp ON pp.product_code=p.code
               WHERE pp.price_group=? ORDER BY p.name""", (group,)
        ):
            data = dict(item)
            text = clean_text(data.pop("price_text", ""))
            if data["sell_price"] is not None:
                data["price_state"] = "zero" if number_value(data["sell_price"]) == 0 else "numeric"
                data["status"] = "Giá 0 – giữ để xác nhận" if data["price_state"] == "zero" else ""
                data["exportable"] = True
            elif text.lower() == "x" or not text:
                excluded_count += 1
                continue
            else:
                data["price_state"] = "text"
                data["status"] = text
                data["exportable"] = False
            rows.append(data)
        return mode, rows, {
            "period": resolved_period, "price_group": group, "version": None,
            "daily_source": None, "contractor_name": row["name"],
            "conflicts": [], "excluded_count": excluded_count,
            "output_count": sum(bool(item["exportable"]) for item in rows),
            "status_count": sum(not item["exportable"] for item in rows),
        }


def quote_period_arg(raw_value):
    period = clean_text(raw_value)
    if not period:
        return ""
    try:
        if not re.fullmatch(r"\d{4}-\d{2}", period):
            raise ValueError
        datetime.strptime(period + "-01", "%Y-%m-%d")
    except ValueError:
        raise QuoteExportError(
            "Kỳ báo giá phải có dạng YYYY-MM",
            code="quote_period_invalid",
            status=400,
        ) from None
    return period


def quote_version_arg(raw_value):
    value = clean_text(raw_value)
    if not value:
        return None
    try:
        version_id = int(value)
    except (TypeError, ValueError):
        version_id = 0
    if version_id <= 0:
        raise QuoteExportError(
            "Phiên bản báo giá không hợp lệ",
            code="quote_version_invalid",
            status=400,
        )
    return version_id


@app.get("/api/quotes")
def api_quotes():
    contractor = clean_text(request.args.get("contractor")).upper()
    batch_id = request.args.get("batch_id", type=int)
    try:
        period = quote_period_arg(request.args.get("period"))
        version_id = quote_version_arg(request.args.get("version_id"))
        with db() as conn:
            mode, rows, context = get_quote_rows(
                conn, contractor, batch_id, period, version_id=version_id,
            )
            recipient = quote_recipient(
                contractor,
                configured=setting_get(conn, f"quote_recipient_{contractor}", ""),
                fallback_name=context["contractor_name"],
            )
            return jsonify({
                "ok": True, "contractor": contractor, "mode": mode, "items": rows,
                "recipient": recipient,
                "period": context["period"], "price_group": context["price_group"],
                "version": context["version"], "dailySource": context["daily_source"],
                "conflicts": context["conflicts"],
                "excludedCount": context["excluded_count"],
                "outputCount": context["output_count"], "statusCount": context["status_count"],
            })
    except QuoteExportError as exc:
        return jsonify({"ok": False, "error": str(exc), "code": exc.code}), exc.status


@app.get("/api/export/<kind>/<int:batch_id>")
def api_export(kind, batch_id):
    with db() as conn:
        batch, orders = require_batch(conn, batch_id)
        stamp = batch["work_date"]
        if kind == "suppliers":
            return send_xlsx(export_supplier_orders(conn, batch, orders), f"Don_dat_hang_NCC_{stamp}.xlsx")
        if kind == "deliveries":
            try:
                workbook = export_deliveries(conn, batch, orders)
            except DeliveryExportError as error:
                return jsonify({"ok": False, "error": str(error), "code": error.code}), error.status
            return send_xlsx(workbook, f"Phieu_giao_hang_{stamp}.xlsx")
        if kind == "report":
            try:
                workbook = export_report(conn, batch, orders)
            except ReportExportError as error:
                return jsonify({"ok": False, "error": str(error), "code": error.code}), error.status
            return send_xlsx(workbook, f"Bao_cao_tong_hop_{batch['work_date'][:7]}.xlsx")
        if kind == "purchases":
            try:
                workbook = export_purchase_documents(conn, batch, orders)
            except PurchaseSummaryError as error:
                return jsonify({"ok": False, "error": str(error), "code": error.code}), error.status
            return send_xlsx(workbook, f"Bang_ke_bien_nhan_{stamp}.xlsx")
        if kind == "outgoing-statement":
            return send_xlsx(
                export_outgoing_statement(conn, batch, orders),
                f"Bang_ke_hang_hoa_dau_ra_{stamp}.xlsx",
            )
        if kind == "invoices":
            if batch["status"] != "approved":
                return jsonify({"ok": False, "error": "Phải duyệt phiên đơn trước khi tạo file hóa đơn"}), 409
            try:
                payload = export_invoices_zip(conn, batch, orders)
            except InvoiceTaxExportError as error:
                return jsonify({"ok": False, "error": str(error), "code": error.code}), error.status
            return send_file(
                payload, as_attachment=True,
                download_name=f"File_tai_phan_mem_trung_gian_{stamp}.zip",
                mimetype="application/zip",
            )
        return jsonify({"ok": False, "error": "Loại file không hợp lệ"}), 404


@app.post("/api/export/selected-documents")
def api_export_selected_documents():
    """Create one clear ZIP for the dates and document types chosen by the user.

    The UI uses this route for historical reprints.  Selection is explicit: it
    never guesses a date range from the current batch and never mutates source
    data while producing printable Excel files.
    """

    body = request.get_json(silent=True) or {}
    raw_batch_ids = body.get("batch_ids") or []
    raw_documents = body.get("documents") or []
    raw_delivery_selections = body.get("delivery_selections") or []
    if not isinstance(raw_batch_ids, list) or not isinstance(raw_documents, list):
        return jsonify({"ok": False, "error": "Danh sách đơn hoặc giấy tờ không hợp lệ"}), 400
    if not isinstance(raw_delivery_selections, list):
        return jsonify({"ok": False, "error": "Danh sách phiếu giao không hợp lệ"}), 400

    batch_ids = []
    for value in raw_batch_ids:
        try:
            batch_id = int(value)
        except (TypeError, ValueError):
            return jsonify({"ok": False, "error": "Có mã đơn không hợp lệ"}), 400
        if batch_id > 0 and batch_id not in batch_ids:
            batch_ids.append(batch_id)
    if not batch_ids:
        return jsonify({"ok": False, "error": "Hãy chọn ít nhất một đơn cần in"}), 400
    if len(batch_ids) > 100:
        return jsonify({"ok": False, "error": "Mỗi lần chỉ chọn tối đa 100 đơn"}), 400

    allowed_documents = {
        "suppliers": ("Don_hang", export_supplier_orders),
        "deliveries": ("Phieu_giao_hang", export_deliveries),
        "purchases": ("Bang_ke_bien_nhan", export_purchase_documents),
        "report": ("Bao_cao_tong_hop", export_report),
    }
    documents = []
    for value in raw_documents:
        key = clean_text(value).lower()
        if key in allowed_documents and key not in documents:
            documents.append(key)
    if not documents:
        return jsonify({"ok": False, "error": "Hãy chọn ít nhất một loại giấy tờ"}), 400

    delivery_selections = defaultdict(set)
    for item in raw_delivery_selections:
        if not isinstance(item, dict):
            return jsonify({"ok": False, "error": "Có phiếu giao được chọn không hợp lệ"}), 400
        try:
            selected_batch_id = int(item.get("batch_id"))
        except (TypeError, ValueError):
            return jsonify({"ok": False, "error": "Có mã đơn trong phiếu giao không hợp lệ"}), 400
        selected_kitchen = clean_text(item.get("kitchen")).upper()
        if selected_batch_id not in batch_ids or not selected_kitchen:
            return jsonify({"ok": False, "error": "Có phiếu giao không thuộc danh sách đã chọn"}), 400
        delivery_selections[selected_batch_id].add(selected_kitchen)

    try:
        archive_stream = io.BytesIO()
        written = 0
        written_report_periods = set()
        with db() as conn, zipfile.ZipFile(
            archive_stream, mode="w", compression=zipfile.ZIP_DEFLATED,
        ) as archive:
            batches = []
            for batch_id in batch_ids:
                batch, orders = require_batch(conn, batch_id)
                batches.append((batch, orders))
            batches.sort(key=lambda item: (item[0]["work_date"], item[0]["id"]))

            for batch, orders in batches:
                stamp = clean_text(batch["work_date"])
                for document in documents:
                    if document == "report":
                        period = stamp[:7]
                        if period in written_report_periods:
                            continue
                        written_report_periods.add(period)
                    label, builder = allowed_documents[document]
                    try:
                        if document == "deliveries" and raw_delivery_selections:
                            selected_codes = delivery_selections.get(int(batch["id"]), set())
                            if not selected_codes:
                                continue
                            workbook = export_deliveries(conn, batch, orders, selected_codes)
                        else:
                            workbook = builder(conn, batch, orders)
                    except PurchaseSummaryError as error:
                        if document == "purchases" and error.code == "no_purchase_summary_rows":
                            continue
                        raise
                    try:
                        stream = workbook_bytes(workbook)
                        try:
                            payload = stream.getvalue()
                        finally:
                            stream.close()
                    finally:
                        workbook.close()
                    suffix = stamp[:7] if document == "report" else f"{stamp}_don-{int(batch['id'])}"
                    archive.writestr(f"{label}_{suffix}.xlsx", payload)
                    written += 1

        if not written:
            archive_stream.close()
            return jsonify({"ok": False, "error": "Các đơn đã chọn chưa có giấy tờ phù hợp để tải"}), 409
        archive_stream.seek(0)
        return send_file(
            archive_stream,
            as_attachment=True,
            download_name=f"GIAY_TO_DA_CHON_{date.today().strftime('%Y%m%d')}.zip",
            mimetype="application/zip",
        )
    except (ValueError, DeliveryExportError, PurchaseSummaryError, ReportExportError) as error:
        return jsonify({
            "ok": False,
            "error": str(error),
            "code": getattr(error, "code", "selected_document_export_failed"),
        }), int(getattr(error, "status", 422))


@app.get("/api/export/quote/<contractor>")
def api_export_quote(contractor):
    contractor = clean_text(contractor).upper()
    batch_id = request.args.get("batch_id", type=int)
    try:
        period = quote_period_arg(request.args.get("period"))
        version_id = quote_version_arg(request.args.get("version_id"))
        with db() as conn:
            mode, rows, context = get_quote_rows(
                conn, contractor, batch_id, period, version_id=version_id,
            )
            if context["conflicts"]:
                return jsonify({
                    "ok": False,
                    "error": "Báo giá còn mã trùng xung đột; không tự chọn dòng để xuất",
                    "code": "quote_duplicate_conflict",
                    "conflicts": context["conflicts"],
                }), 409
            if mode == "daily" and context["daily_source"] is None:
                raise QuoteExportError(
                    "Phải chọn phiên đơn nguồn trước khi xuất báo giá theo ngày",
                    code="daily_batch_required",
                )
            recipient = quote_recipient(
                contractor,
                configured=setting_get(conn, f"quote_recipient_{contractor}", ""),
                fallback_name=context["contractor_name"],
            )
            workbook = build_contractor_quote_workbook(
                rows,
                contractor=contractor,
                recipient=recipient,
                period=context["period"],
                version=context["version"],
                daily_source=context["daily_source"] if mode == "daily" else None,
            )
            filename = contractor_quote_filename(
                contractor,
                period=context["period"],
                version_no=(context["version"] or {}).get("version_no"),
                daily_source=context["daily_source"] if mode == "daily" else None,
            )
            return send_xlsx(workbook, filename)
    except QuoteExportError as exc:
        return jsonify({"ok": False, "error": str(exc), "code": exc.code}), exc.status


def quote_workbook_payload(conn, contractor, rows, context, mode):
    recipient = quote_recipient(
        contractor,
        configured=setting_get(conn, f"quote_recipient_{contractor}", ""),
        fallback_name=context["contractor_name"],
    )
    workbook = build_contractor_quote_workbook(
        rows,
        contractor=contractor,
        recipient=recipient,
        period=context["period"],
        version=context["version"],
        daily_source=context["daily_source"] if mode == "daily" else None,
    )
    try:
        stream = workbook_bytes(workbook)
        try:
            payload = stream.getvalue()
        finally:
            stream.close()
    finally:
        workbook.close()
    filename = contractor_quote_filename(
        contractor,
        period=context["period"],
        version_no=(context["version"] or {}).get("version_no"),
        daily_source=context["daily_source"] if mode == "daily" else None,
    )
    return filename, payload


@app.get("/api/export/quotes/all")
def api_export_all_quotes():
    """Download one isolated XLSX per contractor in a single ZIP archive."""
    try:
        period = quote_period_arg(request.args.get("period"))
        if not period:
            raise QuoteExportError(
                "Phải chọn kỳ báo giá trước khi tải toàn bộ",
                code="quote_period_required",
                status=400,
            )
        requested_version_id = quote_version_arg(request.args.get("version_id"))
        batch_id = request.args.get("batch_id", type=int)
        with db() as conn:
            if requested_version_id is None:
                version = conn.execute(
                    """SELECT id,version_no FROM quote_versions
                       WHERE effective_period=? AND status='confirmed'
                       ORDER BY version_no DESC LIMIT 1""",
                    (period,),
                ).fetchone()
                missing_code = "quote_period_not_confirmed"
                missing_status = 409
                missing_message = "Kỳ báo giá này chưa có phiên bản đã xác nhận"
            else:
                version = conn.execute(
                    """SELECT id,version_no FROM quote_versions
                       WHERE id=? AND effective_period=? AND status='confirmed'""",
                    (requested_version_id, period),
                ).fetchone()
                missing_code = "quote_version_not_found"
                missing_status = 404
                missing_message = "Không tìm thấy phiên bản báo giá đã xác nhận trong kỳ đã chọn"
            if not version:
                raise QuoteExportError(
                    missing_message, code=missing_code, status=missing_status,
                )

            selected_version_id = int(version["id"])
            contractors = conn.execute(
                """SELECT code,pricing_mode FROM contractors
                   ORDER BY CASE WHEN pricing_mode='daily' THEN 1 ELSE 0 END,code"""
            ).fetchall()
            daily_batch_allowed = False
            if batch_id:
                batch = conn.execute(
                    "SELECT work_date FROM batches WHERE id=?", (batch_id,),
                ).fetchone()
                if not batch:
                    raise QuoteExportError(
                        "Không tìm thấy đơn hàng dùng cho báo giá theo ngày",
                        code="daily_batch_not_found",
                        status=404,
                    )
                daily_batch_allowed = clean_text(batch["work_date"])[:7] == period
                if not daily_batch_allowed:
                    raise QuoteExportError(
                        "Đơn hàng theo ngày không thuộc kỳ báo giá đã chọn",
                        code="daily_batch_period_mismatch",
                    )

            archive_stream = io.BytesIO()
            written = 0
            with zipfile.ZipFile(
                archive_stream, mode="w", compression=zipfile.ZIP_DEFLATED,
            ) as archive:
                for contractor_row in contractors:
                    contractor = clean_text(contractor_row["code"]).upper()
                    is_daily = contractor_row["pricing_mode"] == "daily"
                    if is_daily and not daily_batch_allowed:
                        continue
                    mode, rows, context = get_quote_rows(
                        conn,
                        contractor,
                        batch_id if is_daily else None,
                        period,
                        version_id=None if is_daily else selected_version_id,
                    )
                    if context["conflicts"]:
                        raise QuoteExportError(
                            f"Báo giá {contractor} còn mã trùng xung đột; chưa thể tải toàn bộ",
                            code="quote_duplicate_conflict",
                        )
                    exportable = [item for item in rows if item.get("exportable")]
                    if not exportable:
                        continue
                    filename, payload = quote_workbook_payload(
                        conn, contractor, rows, context, mode,
                    )
                    archive.writestr(filename, payload)
                    written += 1

            if not written:
                raise QuoteExportError(
                    "Phiên bản này chưa có báo giá nào đủ điều kiện để tải",
                    code="quote_bundle_empty",
                )
            archive_stream.seek(0)
            month = int(period[5:7])
            year = int(period[:4])
            filename = f"BAO_GIA_TAT_CA_T{month:02d}-{year}_V{int(version['version_no'])}.zip"
            return send_file(
                archive_stream,
                as_attachment=True,
                download_name=filename,
                mimetype="application/zip",
            )
    except QuoteExportError as exc:
        return jsonify({"ok": False, "error": str(exc), "code": exc.code}), exc.status


@app.get("/api/backup/status")
def api_backup_status():
    return jsonify(backup_status(DATA_DIR))


@app.get("/api/backup")
def api_backup():
    # A consistent SQLite copy while the app is live.
    try:
        hosted_admin = app.config.get("TDP_CLOUD_ADMIN_USER")
        authenticated_admin = bool(hosted_admin and session.get("user") == hosted_admin)
        if not authenticated_admin and not ipaddress.ip_address(request.remote_addr or "").is_loopback:
            return jsonify({
                "ok": False,
                "error": "Chỉ được tải toàn bộ bản sao dữ liệu trực tiếp trên máy chủ",
            }), 403
    except ValueError:
        return jsonify({"ok": False, "error": "Không xác định được máy yêu cầu sao lưu"}), 403
    if not DB_PATH.is_file():
        return jsonify({"ok": False, "error": "Cơ sở dữ liệu chưa sẵn sàng để sao lưu"}), 503
    backup_path = DATA_DIR / (
        f"backup_{datetime.now():%Y%m%d_%H%M%S}_{uuid.uuid4().hex[:10]}.sqlite3"
    )
    source = target = None
    try:
        source = sqlite3.connect(DB_PATH)
        target = sqlite3.connect(backup_path)
        source.backup(target)
    except (OSError, sqlite3.Error):
        try:
            backup_path.unlink(missing_ok=True)
        except OSError:
            pass
        return jsonify({"ok": False, "error": "Không tạo được bản sao dữ liệu nhất quán"}), 503
    finally:
        if target is not None:
            target.close()
        if source is not None:
            source.close()
    response = send_file(
        backup_path, as_attachment=True, download_name=backup_path.name,
        mimetype="application/vnd.sqlite3",
    )

    @response.call_on_close
    def _cleanup():
        try:
            backup_path.unlink()
        except OSError:
            pass

    return response


def local_ip():
    try:
        with socket.socket(socket.AF_INET, socket.SOCK_DGRAM) as sock:
            sock.connect(("8.8.8.8", 80))
            return sock.getsockname()[0]
    except OSError:
        return "127.0.0.1"


def application_url(runtime_port: int) -> str:
    return f"http://127.0.0.1:{runtime_port}"


def tdp_health_available(runtime_port: int, timeout: float = 0.8) -> bool:
    try:
        with urlopen(f"{application_url(runtime_port)}/health", timeout=timeout) as response:
            payload = json.loads(response.read().decode("utf-8"))
        return bool(
            response.status == 200
            and payload.get("ok") is True
            and payload.get("database_ready") is True
            and payload.get("schema_ready") is True
        )
    except (OSError, URLError, ValueError, json.JSONDecodeError):
        return False


def open_browser(runtime_port: int = 8765, wait_seconds: float = 30.0) -> bool:
    deadline = time.monotonic() + max(wait_seconds, 0)
    while time.monotonic() < deadline:
        if tdp_health_available(runtime_port):
            webbrowser.open(application_url(runtime_port))
            return True
        time.sleep(0.25)
    return False


_SINGLE_INSTANCE_HANDLE = None


def single_instance_mutex_name() -> str:
    """Keep production single-instance while allowing an isolated packaged smoke run."""

    production_name = "Local\\ThanhDatPhatDesktopApplication"
    if "--smoke-test-instance" not in sys.argv:
        return production_name

    smoke_id = clean_text(os.environ.get("TDP_SMOKE_INSTANCE_ID", ""))
    if not re.fullmatch(r"[A-Za-z0-9_-]{1,32}", smoke_id):
        raise RuntimeError("TDP_SMOKE_INSTANCE_ID không hợp lệ")
    required_names = ("TDP_DATA_DIR", "TDP_DB_PATH", "TDP_EXPORT_DIR", "TDP_PORT")
    missing = [name for name in required_names if not clean_text(os.environ.get(name, ""))]
    if missing:
        raise RuntimeError("Chế độ smoke thiếu cấu hình cô lập: " + ", ".join(missing))
    try:
        smoke_port = int(os.environ["TDP_PORT"])
    except ValueError as exc:
        raise RuntimeError("Cổng smoke không hợp lệ") from exc
    if smoke_port == 8765 or smoke_port < 1024 or smoke_port > 65535:
        raise RuntimeError("Chế độ smoke phải dùng cổng riêng từ 1024 đến 65535, khác 8765")

    raw_temp_root = clean_text(os.environ.get("TEMP") or os.environ.get("TMP") or "")
    if not raw_temp_root:
        raise RuntimeError("Không xác định được thư mục tạm Windows cho chế độ smoke")
    temp_root = Path(raw_temp_root).resolve()
    isolated_paths = [
        Path(os.environ["TDP_DATA_DIR"]).resolve(),
        Path(os.environ["TDP_DB_PATH"]).resolve(),
        Path(os.environ["TDP_EXPORT_DIR"]).resolve(),
    ]
    for isolated_path in isolated_paths:
        try:
            inside_temp = os.path.commonpath((str(temp_root), str(isolated_path))) == str(temp_root)
        except ValueError:
            inside_temp = False
        if not inside_temp or isolated_path == temp_root:
            raise RuntimeError("Chế độ smoke chỉ được dùng đường dẫn con trong thư mục tạm Windows")

    fingerprint = hashlib.sha256(
        f"{smoke_id}|{isolated_paths[0]}|{smoke_port}".encode("utf-8")
    ).hexdigest()[:16]
    return production_name + "Smoke_" + fingerprint


def acquire_single_instance() -> bool:
    """Return False for a second frozen Windows launch and retain the mutex otherwise."""

    global _SINGLE_INSTANCE_HANDLE
    if not FROZEN or os.name != "nt":
        return True
    import ctypes

    kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
    kernel32.CreateMutexW.argtypes = (ctypes.c_void_p, ctypes.c_bool, ctypes.c_wchar_p)
    kernel32.CreateMutexW.restype = ctypes.c_void_p
    kernel32.CloseHandle.argtypes = (ctypes.c_void_p,)
    handle = kernel32.CreateMutexW(None, False, single_instance_mutex_name())
    if not handle:
        raise ctypes.WinError(ctypes.get_last_error())
    if ctypes.get_last_error() == 183:  # ERROR_ALREADY_EXISTS
        kernel32.CloseHandle(handle)
        return False
    _SINGLE_INSTANCE_HANDLE = handle
    return True


def show_startup_error(message: str):
    if FROZEN and os.name == "nt":
        try:
            import ctypes
            ctypes.windll.user32.MessageBoxW(
                None,
                message,
                "Thành Đạt Phát – Không thể khởi động",
                0x10,
            )
        except (AttributeError, OSError):
            pass


def api_debt_adjustment_guarded():
    """Validated, audited replacement for the contract-module adjustment route."""
    body = request.get_json(silent=True) or {}
    party_type = clean_text(body.get("party_type")).lower()
    party_code_input = clean_text(body.get("party_code"))
    note = clean_text(body.get("note"))
    if party_type not in {"contractor", "supplier"}:
        return jsonify({"ok": False, "error": "Nhóm công nợ không hợp lệ"}), 400
    if not party_code_input:
        return jsonify({"ok": False, "error": "Thiếu mã nhà thầu/nhà cung cấp"}), 400
    if not note:
        return jsonify({"ok": False, "error": "Điều chỉnh công nợ phải có lý do"}), 400
    try:
        adjustment_date = valid_iso_date(body.get("adjustment_date"), "Ngày điều chỉnh")
        amount = finite_number(body.get("amount"), "Số tiền điều chỉnh")
    except ValueError as exc:
        return jsonify({"ok": False, "error": str(exc)}), 400
    if amount == 0:
        return jsonify({"ok": False, "error": "Số tiền điều chỉnh phải khác 0"}), 400
    with db() as conn:
        try:
            party_code = canonical_party_code(conn, party_type, party_code_input)
        except ValueError as exc:
            return jsonify({"ok": False, "error": str(exc)}), 400
        cur = conn.execute(
            """INSERT INTO debt_adjustments(
                   adjustment_date,party_type,party_code,amount,note,created_at
               ) VALUES(?,?,?,?,?,?)""",
            (adjustment_date, party_type, party_code, amount, note, now_iso()),
        )
        audit_event(
            conn, "debt.adjustment.create", entity_type=party_type, entity_id=party_code,
            metadata={"adjustment_id": cur.lastrowid, "adjustment_date": adjustment_date,
                      "amount": amount},
        )
        return jsonify({"ok": True, "id": cur.lastrowid})


def api_attendance_guarded():
    """Validated, audited replacement for manual attendance entry."""
    body = request.get_json(silent=True) or {}
    employee_code = clean_text(body.get("employee_code")).upper()
    if not employee_code:
        return jsonify({"ok": False, "error": "Thiếu mã nhân sự"}), 400
    try:
        work_date = valid_iso_date(body.get("work_date"), "Ngày chấm công")
        hours = {
            key: finite_number(body.get(key), label, default=0)
            for key, label in (
                ("normal_hours", "Giờ thường"),
                ("overtime_hours", "Giờ tăng ca"),
                ("sunday_hours", "Giờ Chủ nhật"),
                ("night_hours", "Giờ ca đêm"),
                ("holiday_hours", "Giờ ngày lễ"),
            )
        }
    except ValueError as exc:
        return jsonify({"ok": False, "error": str(exc)}), 400
    if any(value < 0 or value > 24 for value in hours.values()):
        return jsonify({"ok": False, "error": "Mỗi loại giờ công phải nằm trong khoảng 0–24"}), 400
    primary_hours = sum(hours[key] for key in (
        "normal_hours", "overtime_hours", "sunday_hours", "holiday_hours",
    ))
    if primary_hours > 24:
        return jsonify({"ok": False, "error": "Tổng giờ thường/tăng ca/Chủ nhật/ngày lễ không được vượt 24 giờ/ngày"}), 400

    with db() as conn:
        staff_row = conn.execute(
            "SELECT id FROM staff WHERE employee_code=?", (employee_code,)
        ).fetchone()
        if not staff_row:
            return jsonify({"ok": False, "error": "Nhân sự không tồn tại"}), 404
        conn.execute(
            """INSERT INTO attendance_entries(
                   employee_id,work_date,normal_hours,overtime_hours,sunday_hours,night_hours,
                   holiday_hours,note,source,updated_at
               ) VALUES(?,?,?,?,?,?,?,?,?,?) ON CONFLICT(employee_id,work_date) DO UPDATE SET
                   normal_hours=excluded.normal_hours,overtime_hours=excluded.overtime_hours,
                   sunday_hours=excluded.sunday_hours,night_hours=excluded.night_hours,
                   holiday_hours=excluded.holiday_hours,note=excluded.note,source=excluded.source,
                   updated_at=excluded.updated_at""",
            (
                staff_row["id"], work_date, hours["normal_hours"], hours["overtime_hours"],
                hours["sunday_hours"], hours["night_hours"], hours["holiday_hours"],
                clean_text(body.get("note")), "manual", now_iso(),
            ),
        )
        audit_event(
            conn, "attendance.upsert", entity_type="staff", entity_id=employee_code,
            metadata={"work_date": work_date, **hours},
        )
        return jsonify({"ok": True})


register_contract_routes(app, {
    "db": db,
    "now_iso": now_iso,
    "clean_text": clean_text,
    "number_value": number_value,
    "tax_factor": tax_factor,
    "setting_get": setting_get,
    "setting_set": setting_set,
    "create_minvoice_client": create_minvoice_client,
    "create_msmi_client": create_msmi_client,
    "root": ROOT,
    "data_dir": DATA_DIR,
    "require_batch": require_batch,
    "export_supplier_orders": export_supplier_orders,
    "export_deliveries": export_deliveries,
    "export_purchase_documents": export_purchase_documents,
    "export_optional_purchase_documents": export_optional_purchase_documents,
    "export_report": export_report,
})

register_invoice_workbench_routes(app, {
    "db": db,
    "now_iso": now_iso,
    "setting_get": setting_get,
    "audit_event": audit_event,
    "create_msmi_client": create_msmi_client,
    "create_minvoice_client": create_minvoice_client,
})

register_invoice_input_export_routes(app, {
    "db": db,
    "now_iso": now_iso,
})

register_invoice_mapping_routes(app, {
    "db": db,
    "now_iso": now_iso,
})

register_invoice_inventory_routes(app, {
    "db": db,
    "now_iso": now_iso,
})

register_invoice_valuation_routes(app, {
    "db": db,
})

register_inventory_export_routes(app, {
    "db": db,
    "opening_template_path": OPENING_TEMPLATE_SOURCE,
})

register_inventory_period_close_routes(app, {
    "db": db,
    "now_iso": now_iso,
})

register_bk_import_routes(app, {
    "db": db,
    "now_iso": now_iso,
    "audit_event": audit_event,
})

register_payable_ledger_routes(app, {
    "db": db,
})

register_payable_payment_routes(app, {
    "db": db,
    "now_iso": now_iso,
    "canonical_party_code": canonical_party_code,
    "audit_event": audit_event,
})

register_payable_export_routes(app, {
    "db": db,
    "canonical_party_code": canonical_party_code,
})

register_receivable_ledger_routes(app, {
    "db": db,
})

register_receivable_export_routes(app, {
    "db": db,
    "debt_period_payload": debt_period_payload,
    "tax_factor": tax_factor,
})

register_daily_reference_routes(app, db, now_iso, clean_text)

register_quote_import_routes(app, {
    "db": db,
    "now_iso": now_iso,
    "audit_event": audit_event,
})

register_outgoing_substitution_routes(app, {
    "db": db,
    "now_iso": now_iso,
    "audit_event": audit_event,
})

try:
    from round3_documents import register_round3_routes, customer_receipt
except ImportError:
    from .round3_documents import register_round3_routes, customer_receipt


def round3_context():
    return dict(db=db, clean_text=clean_text, now_iso=now_iso, valid_iso_date=valid_iso_date,
                canonical_party_code=canonical_party_code, audit_event=audit_event,
                export_report=export_report, send_xlsx=send_xlsx)


register_round3_routes(app, round3_context())

try:
    from round4_documents import register_document_routes
except ImportError:
    from .round4_documents import register_document_routes
register_document_routes(app, lambda: globals())
order_worksheet.register(app, globals())

register_physical_inventory_routes(app, {
    "db": db, "now_iso": now_iso, "batch_payload": batch_payload,
})

register_order_price_override_routes(app, {
    "db": db,
    "now_iso": now_iso,
    "clean_text": clean_text,
    "finite_number": finite_number,
    "batch_mutation_blocker": batch_mutation_blocker,
    "product_lookup": product_lookup,
    "resolve_order": resolve_order,
    "clear_batch_derived_inventory": clear_batch_derived_inventory,
    "batch_payload": batch_payload,
    "audit_event": audit_event,
    "sync_payable_ledger": sync_payable_ledger,
    "sync_receivable_ledger": sync_receivable_ledger,
})

# Keep the route topology in the contract module while enforcing financial and
# attendance invariants here, in the same SQLite transaction as their audit log.
app.view_functions["api_debt_adjustment"] = api_debt_adjustment_guarded
app.view_functions["api_save_attendance"] = api_attendance_guarded


def run_application():
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except (AttributeError, OSError):
        pass
    allow_lan = clean_text(os.environ.get("TDP_ALLOW_LAN", "")).lower() in {
        "1", "true", "yes", "co", "có",
    }
    try:
        runtime_port = int(clean_text(os.environ.get("TDP_PORT", "8765")))
    except ValueError as error:
        raise RuntimeError("TDP_PORT phải là số nguyên từ 1 đến 65535") from error
    if runtime_port < 1 or runtime_port > 65535:
        raise RuntimeError("TDP_PORT phải là số nguyên từ 1 đến 65535")
    bind_host = "0.0.0.0" if allow_lan else "127.0.0.1"

    if not acquire_single_instance():
        if "--no-browser" not in sys.argv:
            open_browser(runtime_port, wait_seconds=45)
        return
    if tdp_health_available(runtime_port):
        if "--no-browser" not in sys.argv:
            webbrowser.open(application_url(runtime_port))
        return

    install_seed_database_if_missing()
    init_database()
    if sys.stdout is not None:
        print("\nTHÀNH ĐẠT PHÁT – HỆ THỐNG QUẢN LÝ")
        print(f"Đang chạy tại: {application_url(runtime_port)}")
        if allow_lan:
            print(f"Dùng trong mạng văn phòng: http://{local_ip()}:{runtime_port}")
        else:
            print("Chế độ an toàn: chỉ dùng trên máy này.")
        print("Giữ cửa sổ này mở trong lúc sử dụng. Nhấn Ctrl+C để dừng.\n")
    if "--no-browser" not in sys.argv:
        threading.Thread(target=open_browser, args=(runtime_port,), daemon=True).start()
    backup_stop, backup_worker = start_backup_worker(auto_backup)
    try:
        serve(app, host=bind_host, port=runtime_port, threads=8)
    finally:
        backup_stop.set()
        backup_worker.join(timeout=2)


if __name__ == "__main__":
    try:
        run_application()
    except KeyboardInterrupt:
        pass
    except Exception as error:
        show_startup_error(
            "Không khởi động được phần mềm.\n\n"
            f"Chi tiết: {error}\n\n"
            "Hãy chụp thông báo này và gửi người hỗ trợ."
        )
        raise
