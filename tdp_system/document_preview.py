"""Read-only, immutable workbook previews. Excel remains the print artwork."""
from __future__ import annotations

import ast
import hashlib
import io
import json
import re
import threading
import time
import uuid
from copy import copy
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from html import escape
from pathlib import Path

from openpyxl.styles import PatternFill
import unicodedata
from openpyxl.utils import get_column_letter, range_boundaries, column_index_from_string

PDF_LOCK = threading.Lock()
SNAPSHOT_TTL = 24 * 60 * 60


def white_print_style(workbook):
    """Only generated copies: no template/source writes or numeric changes."""
    for sheet in workbook:
        if sheet.title == 'Đề nghị thanh toán' and workbook.properties.subject == 'Đề nghị thanh toán từ hóa đơn đỏ đã phát hành':
            continue  # Preserve the explicitly supplied customer template artwork.
        receipt = bool(re.fullmatch(r'biên nhận(?:\s+\d+)?', sheet.title.strip().casefold()))
        sheet.conditional_formatting._cf_rules.clear()
        for row in sheet:
            for cell in row:
                cell.fill = PatternFill(fill_type=None)
                font = copy(cell.font)
                font.color = '000000'
                name = unicodedata.normalize('NFC', str(cell.value or '')).strip().casefold()
                font.b = name == 'vũ thị thụy'
                cell.font = font
                border = copy(cell.border)
                for side in ('left', 'right', 'top', 'bottom', 'diagonal', 'vertical', 'horizontal'):
                    original = getattr(border, side)
                    if original and original.style:
                        thin = copy(original)
                        thin.style = 'thin' if receipt else 'hair'
                        thin.color = '000000'
                        setattr(border, side, thin)
                cell.border = border
                if isinstance(cell.value, (int,float,Decimal)) and '#,##0' in cell.number_format:
                    alignment=copy(cell.alignment)
                    alignment.shrinkToFit=True
                    cell.alignment=alignment
                # Excel prints a trailing dot for integer values with ".######".
                # Remove only optional decimal places, never fixed decimals/tax.
                if isinstance(cell.value, (int,float,Decimal)) and not isinstance(cell.value, bool):
                    if Decimal(str(cell.value)) == Decimal(str(cell.value)).to_integral_value():
                        cell.number_format = re.sub(r'\.#+', '', cell.number_format)
    return workbook


def plain_docx_print_style(document):
    """Normalize generated Word copies, including headers and table styles."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    roots = [document.element, document.styles.element]
    for section in document.sections:
        roots.extend([section.header._element, section.footer._element])
    for root in roots:
        for shading in root.xpath('.//w:shd'):
            shading.set(qn('w:fill'), 'FFFFFF')
            shading.set(qn('w:val'), 'clear')
            for attr in ('themeFill', 'themeFillTint', 'themeFillShade'):
                shading.attrib.pop(qn('w:' + attr), None)
        for border in root.xpath('.//w:tblBorders/* | .//w:tcBorders/* | .//w:pBdr/*'):
            if border.get(qn('w:val')) not in ('nil', 'none'):
                border.set(qn('w:val'), 'single')
                border.set(qn('w:sz'), '2')
                border.set(qn('w:color'), '000000')
        for run in root.xpath('.//w:r'):
            text = ''.join(run.xpath('.//w:t/text()')).strip().casefold()
            props = run.get_or_add_rPr()
            for tag, val in [('b', '1' if text == 'vũ thị thụy' else '0'), ('bCs', '0'), ('color', '000000')]:
                element = props.find(qn('w:' + tag))
                if element is None:
                    element = OxmlElement('w:' + tag)
                    props.append(element)
                element.set(qn('w:val'), val)
                if tag == 'color':
                    element.attrib.pop(qn('w:themeColor'), None)
    return document


def display_value(value, number_format='General'):
    if value is None:
        return ''
    if isinstance(value, (date, datetime)):
        return value.strftime('%d/%m/%Y')
    if isinstance(value, bool) or not isinstance(value, (int, float, Decimal)):
        return str(value)
    number = Decimal(str(value))
    if not number.is_finite():
        raise ValueError('Chứng từ có số không hợp lệ')
    if number_format == '@':
        return str(value)
    fmt = number_format.split(';')[0]
    if fmt == 'General':
        return format(number, 'f').rstrip('0').rstrip('.') if '.' in format(number, 'f') else str(value)
    percent = '%' in fmt
    if percent:
        number *= 100
    match = re.search(r'[0#][,0#]*\.([0#]+)', fmt)
    places = len(match[1]) if match else 0
    rounded = number.quantize(Decimal(1).scaleb(-places), rounding=ROUND_HALF_UP)
    text = format(rounded, (',' if ',' in fmt else '') + f'.{places}f')
    if match and '#' in match[1]:
        minimum = len(match[1].rstrip('#'))
        integer, fraction = text.split('.')
        fraction = fraction.rstrip('0')
        fraction += '0' * max(0, minimum - len(fraction))
        text = integer + ('.' + fraction if fraction else '')
    return text + ('%' if percent else '')


def formula_value(sheet, cell, seen=None):
    """Evaluate only the small arithmetic formulas our exporters emit; never eval."""
    if cell.data_type != 'f':
        return cell.value
    seen = set(seen or ())
    if cell.coordinate in seen:
        raise ValueError('Công thức vòng tại ' + cell.coordinate)
    seen.add(cell.coordinate)
    expr = str(cell.value)[1:].replace('$', '')
    match = re.fullmatch(r'(?:SUBTOTAL\(9,|SUM\()([A-Z]+\d+):([A-Z]+\d+)\)', expr)
    if match:
        left, top, right, bottom = range_boundaries(match[1] + ':' + match[2])
        return sum((Decimal(str(formula_value(sheet, sheet.cell(r, c), seen) or 0))
                    for r in range(top, bottom + 1) for c in range(left, right + 1)), Decimal(0))
    if expr.startswith('IFERROR(') and expr.endswith(',0)'):
        expr = expr[8:-3]
    # ROUND is used on generated VND amounts, not on quantities.
    do_round = expr.startswith('ROUND(') and expr.endswith(',0)')
    if do_round:
        expr = expr[6:-3]
    def ref(match):
        value = formula_value(sheet, sheet[match[0]], seen)
        return str(Decimal(str(value or 0)))
    expr = re.sub(r'\b[A-Z]+\d+\b', ref, expr)
    def visit(node):
        if isinstance(node, ast.Constant) and type(node.value) in (int, float):
            return Decimal(str(node.value))
        if isinstance(node, ast.UnaryOp) and isinstance(node.op, (ast.USub, ast.UAdd)):
            return (-1 if isinstance(node.op, ast.USub) else 1) * visit(node.operand)
        if isinstance(node, ast.BinOp):
            a, b = visit(node.left), visit(node.right)
            if isinstance(node.op, ast.Add): return a + b
            if isinstance(node.op, ast.Sub): return a - b
            if isinstance(node.op, ast.Mult): return a * b
            if isinstance(node.op, ast.Div): return a / b
        raise ValueError('Công thức chưa hỗ trợ tại ' + cell.coordinate)
    try:
        result = visit(ast.parse(expr, mode='eval').body)
        return result.quantize(Decimal(1), rounding=ROUND_HALF_UP) if do_round else result
    except (SyntaxError, ArithmeticError, TypeError) as exc:
        raise ValueError('Không đọc được công thức tại ' + cell.coordinate) from exc


def sheet_preview(sheet):
    ranges = re.findall(r'\$?[A-Z]+\$?\d+:\$?[A-Z]+\$?\d+', str(sheet.print_area or ''))
    if len(ranges) > 1:
        raise ValueError('Mẫu có nhiều vùng in; cần tách thành từng sheet trước khi xem')
    bounds = range_boundaries(ranges[0]) if ranges else (1, 1, sheet.max_column, sheet.max_row)
    left, top, right, bottom = bounds
    if (right-left+1) * (bottom-top+1) > 200000:
        raise ValueError('Chứng từ quá lớn để xem một lần; hãy thu hẹp khoảng ngày')
    def column_hidden(c):
        return any(d.hidden and (d.min or column_index_from_string(key)) <= c <=
                   (d.max or d.min or column_index_from_string(key))
                   for key, d in sheet.column_dimensions.items())
    columns = [c for c in range(left, right + 1) if not column_hidden(c)]
    rows = [r for r in range(top, bottom + 1) if not sheet.row_dimensions[r].hidden]
    merged = {}
    skipped = set()
    for area in sheet.merged_cells.ranges:
        cs = [c for c in columns if area.min_col <= c <= area.max_col]
        rs = [r for r in rows if area.min_row <= r <= area.max_row]
        if not cs or not rs:
            continue
        anchor = (rs[0], cs[0])
        merged[anchor] = (len(rs), len(cs), sheet.cell(area.min_row, area.min_col))
        skipped.update((r, c) for r in rs for c in cs if (r, c) != anchor)
    header_rows = [int(x) for x in re.findall(r'\d+', str(sheet.print_title_rows or ''))]
    header = max(header_rows) if header_rows else 0
    widths = [max(24, (sheet.column_dimensions[get_column_letter(c)].width or 12) * 7) for c in columns]
    parts = ['<table class="document-sheet"><colgroup>']
    parts.extend(f'<col style="width:{width:.2f}px">' for width in widths)
    parts.append('</colgroup><tbody>')
    for r in rows:
        parts.append('<tr' + (' class="document-header-row"' if r == header else '') + '>')
        for c in columns:
            if (r, c) in skipped:
                continue
            rowspan, colspan, cell = merged.get((r,c), (1,1,sheet.cell(r,c)))
            value = formula_value(sheet, cell)
            text = display_value(value, cell.number_format)
            align = cell.alignment.horizontal
            if align not in {'left', 'center', 'right'}:
                align = 'right' if isinstance(value, (int,float,Decimal)) else 'left'
            # Excel lets unwrapped labels extend across adjacent empty cells.
            # Give those labels the same space in HTML without changing the
            # workbook, crossing a populated/merged cell or removing borders.
            def bordered(candidate):
                return any(getattr(candidate.border, side) and getattr(candidate.border, side).style
                           for side in ('left', 'right', 'top', 'bottom'))
            if (isinstance(value, str) and value and align == 'left'
                    and not cell.alignment.wrap_text and rowspan == colspan == 1
                    and (r,c) not in merged and not bordered(cell)):
                for next_column in columns[columns.index(c)+1:]:
                    candidate = sheet.cell(r, next_column)
                    if ((r,next_column) in merged or (r,next_column) in skipped
                            or candidate.value is not None or bordered(candidate)):
                        break
                    colspan += 1
                    skipped.add((r,next_column))
            style = [f'text-align:{align}', 'vertical-align:middle',
                     f'font-size:{cell.font.sz or 11}pt',
                     f'font-family:{"Times New Roman" if "Times" in (cell.font.name or "") else "Arial"}',
                     f'height:{sheet.row_dimensions[r].height or 15}pt']
            style.extend(['color:#000', 'background:#fff', 'font-weight:bold' if cell.font.b else 'font-weight:normal'])
            if cell.font.i: style.append('font-style:italic')
            for side in ('left','right','top','bottom'):
                border = getattr(cell.border, side)
                if border and border.style:
                    style.append(f'border-{side}:0.5px solid #000')
            shrink = bool(cell.alignment.shrinkToFit and not cell.alignment.wrap_text)
            nowrap = ' style="white-space:nowrap"' if shrink or (r == header and len(text) <= 4) else ''
            if shrink: nowrap += ' data-shrink="true"'
            parts.append(f'<td rowspan="{rowspan}" colspan="{colspan}" style="{";".join(style)}">'
                         f'<div{nowrap}>{escape(text)}</div></td>')
        parts.append('</tr>')
    parts.append('</tbody></table>')
    return {'name': sheet.title, 'html': ''.join(parts), 'width': sum(widths)}


def create_snapshot(root, workbooks):
    """Keep the exact generated bytes for later Excel/PDF, never regenerate on print."""
    sheets, files, warnings = [], [], []
    for index, (name, workbook) in enumerate(workbooks):
        warnings.extend(getattr(workbook, '_tdp_warnings', []))
        white_print_style(workbook)
        for sheet in workbook:
            if sheet.sheet_state == 'visible':
                preview = sheet_preview(sheet)
                if len(workbooks) > 1:
                    preview['name'] = Path(name).stem.replace('_',' ') + ' · ' + sheet.title
                preview['file_index'] = index
                sheets.append(preview)
        stream = io.BytesIO()
        workbook.save(stream)
        files.append((name, stream.getvalue()))
    if not sheets:
        raise ValueError('Không có chứng từ phù hợp để xem')
    if len(sheets) > 300:
        raise ValueError('Mỗi lần xem tối đa 300 phiếu; hãy thu hẹp bộ lọc')
    # Expire only generated files inside an exact UUID snapshot directory.
    # Unknown files/directories are left untouched, never recursively removed.
    root = Path(root).resolve()
    if root.is_dir():
        for old in root.iterdir():
            if not old.is_dir() or not re.fullmatch(r'[a-f0-9]{32}',old.name) or old.resolve().parent != root:
                continue
            try:
                metadata=json.loads((old/'manifest.json').read_text(encoding='utf-8'))
                if time.time()-metadata['created'] <= SNAPSHOT_TTL: continue
                for file in old.iterdir():
                    if file.is_file() and re.fullmatch(r'(?:manifest\.json|\d+\.xlsx|[a-f0-9]{20}\.(?:pdf|print\.json)|print_[a-f0-9]{20}_\d+\.xlsx)',file.name):
                        file.unlink()
                old.rmdir()
            except (OSError,ValueError,KeyError,TypeError):
                pass
    token = uuid.uuid4().hex
    directory = Path(root) / token
    directory.mkdir(parents=True)
    for index, (_, payload) in enumerate(files):
        (directory / f'{index}.xlsx').write_bytes(payload)
    manifest = {'created': time.time(), 'files': [name for name, _ in files],
                'hashes': [hashlib.sha256(payload).hexdigest() for _,payload in files]}
    (directory / 'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False), encoding='utf-8')
    return {'ok': True, 'token': token, 'sheets': sheets, 'sheet_count': len(sheets), 'warnings': warnings}


def snapshot_files(root, token):
    if not re.fullmatch(r'[a-f0-9]{32}', token):
        raise ValueError('Bản xem không hợp lệ')
    directory = Path(root) / token
    try:
        manifest = json.loads((directory / 'manifest.json').read_text(encoding='utf-8'))
    except (OSError, ValueError) as exc:
        raise ValueError('Bản xem không còn; hãy bấm xem lại') from exc
    if time.time() - manifest['created'] > SNAPSHOT_TTL:
        raise ValueError('Bản xem đã hết hạn; hãy bấm xem lại dữ liệu mới nhất')
    files = [(name, directory / f'{index}.xlsx') for index, name in enumerate(manifest['files'])]
    for index,(_,path) in enumerate(files):
        try:
            valid=hashlib.sha256(path.read_bytes()).hexdigest()==manifest['hashes'][index]
        except (OSError,KeyError,IndexError):
            valid=False
        if not valid:
            raise ValueError('File của bản xem đã thay đổi; hãy mở lại chứng từ trước khi tải hoặc in')
    return directory, files
