import io
import unittest
import zipfile
from unittest.mock import patch
from openpyxl import load_workbook
from . import server, test_outgoing_readiness as fixture
from .outgoing_review import SHEET
from .outgoing_source_refresh import refresh_sources
from .test_invoice_output_sync import documented_minvoice_invoice, OutputFixtureMsmi


class InvoiceReviewTests(unittest.TestCase):
    setUpClass=classmethod(fixture.OutgoingReadinessTests.setUpClass.__func__)
    tearDownClass=classmethod(fixture.OutgoingReadinessTests.tearDownClass.__func__)

    def setUp(self):
        with server.db() as c:
            for table in ('invoice_inventory_ledger','invoice_sync_batch_output_invoices','invoice_mapping_revisions','invoice_line_mappings','outgoing_source_order_scopes'):
                c.execute('DELETE FROM '+table)
        fixture.OutgoingReadinessTests.setUp(self)
        with server.db() as c:
            for table in ('outgoing_order_choices','outgoing_contractor_choices','outgoing_review_notes','outgoing_product_units','outgoing_product_names','outgoing_buyer_profiles'):
                c.execute('DELETE FROM '+table)
            fixture.OutgoingReadinessTests.add_opening(c,100)
            _,self.old=fixture.OutgoingReadinessTests.add_batch(c,'2026-09-01',[{'qty':5}])
            _,self.ids=fixture.OutgoingReadinessTests.add_batch(c,'2026-09-05',[{'qty':3},{'qty':4},{'qty':2,'contractor':'NT-B'}])
            _,self.later=fixture.OutgoingReadinessTests.add_batch(c,'2026-09-10',[{'qty':6}])
        self.period={'from':'2026-09-05','to':'2026-09-07','contractor':'NT-A'}

    def download(self):
        r=self.client.get('/api/outgoing-invoices/review.xlsx',query_string=self.period)
        self.assertEqual(r.status_code,200,r.get_json(silent=True))
        return r.data

    def edited(self,data,change):
        w=load_workbook(io.BytesIO(data));change(w);output=io.BytesIO();w.save(output);w.close();return output.getvalue()

    def preview(self,data):
        return self.client.post('/api/outgoing-invoices/review/preview',data={'file':(io.BytesIO(data),'review.xlsx')})

    def save(self,preview):
        return self.client.put('/api/outgoing-invoices/review/choices',json={'token':preview['token'],'confirmed':True})

    def business(self):
        with server.db() as c:
            return {t:[tuple(r) for r in c.execute('SELECT * FROM '+t)] for t in
                    ('orders','batches','products','invoice_inventory_ledger','receivable_ledger_lines','payable_ledger_lines')}

    def test_deleted_excel_row_persists_and_does_not_exclude_other_dates_or_parties(self):
        before=self.business();data=self.download()
        with server.db() as c:raw_before=c.serialize()
        changed=self.edited(data,lambda w:w[SHEET].delete_rows(2))
        preview=self.preview(changed);self.assertEqual(preview.status_code,200,preview.json)
        self.assertEqual([(r['order_id'],r['enabled']) for r in preview.json['changes']],[(self.ids[0],False)])
        with server.db() as c:self.assertEqual(c.serialize(),raw_before)
        saved=self.save(preview.json);self.assertEqual(saved.status_code,200,saved.json)
        self.assertEqual(self.business(),before)
        w=load_workbook(io.BytesIO(self.download()));self.assertEqual(w[SHEET]['A2'].value,0);self.assertEqual(w[SHEET]['A3'].value,1);w.close()
        r=self.client.get('/api/outgoing-invoices/unissued',query_string=self.period).json
        self.assertEqual([x['order_id'] for x in r['details']],[self.ids[1]])
        with server.db() as c:
            self.assertEqual([tuple(r) for r in c.execute('SELECT order_id,enabled FROM outgoing_order_choices')],[(self.ids[0],0)])
        # A selected worksheet can restore the same order line without recreating a sale.
        restored=self.edited(self.download(),lambda w:setattr(w[SHEET]['A2'],'value',1))
        self.assertEqual(self.save(self.preview(restored).json).status_code,200)
        self.assertEqual(self.business(),before)

    def test_range_applies_to_unissued_workbook_and_final_export(self):
        r=self.client.get('/api/outgoing-invoices/unissued-template.zip',query_string=self.period)
        self.assertEqual(r.status_code,200,r.get_json(silent=True))
        with zipfile.ZipFile(io.BytesIO(r.data)) as z:
            names=[n for n in z.namelist() if n.endswith('.xlsx')]
            self.assertEqual(names,['CHUA_XUAT_NT-A_TU_2026-09-05_DEN_2026-09-07.xlsx'])
            w=load_workbook(io.BytesIO(z.read(names[0])));self.assertEqual(w.active['D2'].value,7);w.close()
        r=self.client.post('/api/export/order-invoices',json={**self.period,'scope':'approved_range'})
        self.assertEqual(r.status_code,200,r.get_json(silent=True))
        with zipfile.ZipFile(io.BytesIO(r.data)) as z:
            w=load_workbook(io.BytesIO(z.read(next(n for n in z.namelist() if n.endswith('.xlsx')))))
            self.assertEqual(w.active['D2'].value,7);w.close()

    def test_changed_quantities_or_duplicate_lines_never_save_choices(self):
        data=self.download()
        for change in [lambda w:setattr(w[SHEET]['I2'],'value',999),lambda w:w[SHEET].append([c.value for c in w[SHEET][2]]),
                       lambda w:setattr(w[SHEET]['A2'],'value','=0'),lambda w:setattr(w['_DoiChieu']['A1'],'value','forged')]:
            with self.subTest(change=change):
                response=self.preview(self.edited(data,change));self.assertEqual(response.status_code,409,response.json)
        with server.db() as c:self.assertEqual(c.execute('SELECT COUNT(*) FROM outgoing_order_choices').fetchone()[0],0)

    def test_current_order_or_catalog_change_invalidates_review_before_save(self):
        changed=self.edited(self.download(),lambda w:w[SHEET].delete_rows(2))
        preview=self.preview(changed).json
        with server.db() as c:c.execute("INSERT INTO outgoing_product_names VALUES('HH-01','Tên mới','now')")
        response=self.save(preview);self.assertEqual(response.status_code,409,response.json)
        with server.db() as c:self.assertEqual(c.execute('SELECT COUNT(*) FROM outgoing_order_choices').fetchone()[0],0)

    def test_notes_survive_reload_are_scoped_and_reject_concurrent_overwrite(self):
        before=self.business()
        first=self.client.get('/api/outgoing-invoices/review-notes',query_string=self.period).json
        data={**self.period,'note':'Đã kiểm tra đến 07/09','revision':first['current']['revision']}
        saved=self.client.put('/api/outgoing-invoices/review-notes',json=data)
        self.assertEqual(saved.status_code,200,saved.json)
        current=self.client.get('/api/outgoing-invoices/review-notes',query_string=self.period).json
        self.assertEqual(current['current']['note'],data['note'])
        another=self.client.get('/api/outgoing-invoices/review-notes',query_string={**self.period,'contractor':'NT-B'}).json
        self.assertEqual(another['current']['note'],'');self.assertEqual(len(another['recent']),1)
        self.assertEqual(self.client.put('/api/outgoing-invoices/review-notes',json={**data,'note':'stale'}).status_code,409)
        self.assertEqual(self.business(),before)

    def test_bad_date_range_rejected_by_every_new_entrypoint(self):
        bad={**self.period,'from':'2026-09-08'}
        for url in ('review.xlsx','review-notes','unissued','unissued-template.zip'):
            r=self.client.get('/api/outgoing-invoices/'+url,query_string=bad)
            self.assertEqual(r.status_code,400,(url,r.json))

    def test_bad_preview_tokens_return_validation_error(self):
        for token in (None,False,{},''):
            r=self.client.put('/api/outgoing-invoices/review/choices',json={'token':token,'confirmed':True})
            self.assertEqual(r.status_code,409,r.json)

    def test_sync_covers_earlier_orders_even_when_reviewing_a_later_period(self):
        with patch('tdp_system.outgoing_source_refresh.refresh_sources',return_value={}) as refresh:
            r=self.client.post('/api/outgoing-invoices/sync-issued',json={**self.period,'scope':'approved_range'})
            self.assertEqual(r.status_code,200,r.json)
            self.assertEqual(refresh.call_args.args[3],'2026-09-01')

    def test_unissued_download_stops_when_signed_source_refresh_fails(self):
        for path in ('unissued-template.zip','unissued.xlsx'):
            with patch('tdp_system.outgoing_source_refresh.refresh_sources',side_effect=ValueError('offline')):
                r=self.client.post('/api/outgoing-invoices/'+path,query_string=self.period)
            self.assertEqual(r.status_code,409,r.get_json(silent=True))
            self.assertIn('chưa tải bảng chưa xuất',r.json['error'])

    def test_unissued_download_refreshes_later_signed_invoice_before_export(self):
        with server.db() as c:
            c.execute("INSERT INTO outgoing_buyer_profiles(contractor,legal_name,tax_code,address,updated_at) VALUES('NT-A','Khách kiểm thử','0209999999','Địa chỉ','now')")
        raw=documented_minvoice_invoice(3)
        raw.update(inv_invoiceIssuedDate='2026-09-08',tgtcthue=140,tgtthue=0,tgtttbso=140,inv_buyerAddressLine='Địa chỉ')
        raw['details'][0].update(inv_itemCode='HH-01',inv_itemName='Hàng hóa 01',inv_unitCode='kg',inv_quantity=7,inv_unitPrice=20,inv_TotalAmountWithoutVat=140,inv_vatAmount=0,inv_TotalAmount=140,ma_thue='0')
        before=self.client.get('/api/outgoing-invoices/unissued',query_string=self.period).json
        self.assertEqual(7,sum(r['unissued_qty'] for r in before['details']))
        with patch.object(server,'create_minvoice_client',return_value=OutputFixtureMsmi([raw])):
            response=self.client.post('/api/outgoing-invoices/unissued-template.zip',query_string=self.period)
        self.assertEqual(response.status_code,200,response.get_json(silent=True))
        with zipfile.ZipFile(io.BytesIO(response.data)) as archive:
            book=load_workbook(io.BytesIO(archive.read(next(n for n in archive.namelist() if n.endswith('.xlsx')))))
            self.assertEqual(5,book.active['D2'].value)
            checked=next(row[1] for row in book['Tong hop'].values if row[0]=='Cập nhật hóa đơn đã ký')
            self.assertNotIn('chưa cập nhật',checked)
            book.close()

    def test_signed_fifo_is_computed_before_filtering_order_dates(self):
        with server.db() as c:
            c.execute("INSERT INTO outgoing_buyer_profiles(contractor,legal_name,tax_code,address,updated_at) VALUES('NT-A','Khách kiểm thử','0209999999','Địa chỉ','now')")
        raw=documented_minvoice_invoice(3)
        raw.update(inv_invoiceIssuedDate='2026-09-08',tgtcthue=140,tgtthue=0,tgtttbso=140,inv_buyerAddressLine='Địa chỉ')
        raw['details'][0].update(inv_itemCode='HH-01',inv_itemName='Hàng hóa 01',inv_unitCode='kg',inv_quantity=7,inv_unitPrice=20,inv_TotalAmountWithoutVat=140,inv_vatAmount=0,inv_TotalAmount=140,ma_thue='0')
        remote=OutputFixtureMsmi([raw]);refresh_sources(server.db,lambda:remote,server.now_iso,'2026-09-01','2026-09-08')
        data=self.client.get('/api/outgoing-invoices/unissued',query_string=self.period).json
        # Five units settle the older 01/09 order; only two settle 05/09.
        self.assertEqual({r['order_id']:r['unissued_qty'] for r in data['details']},{self.ids[0]:1,self.ids[1]:4})


if __name__=='__main__':unittest.main()
