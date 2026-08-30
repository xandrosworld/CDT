from __future__ import annotations

import json
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
import server


SOURCE = Path(r"C:\Users\DELL\Downloads\Đơn hàng 29.08.xlsx")
TODAY = "29.08"
HISTORY_SHEETS = [
    "28.08", "27.08", "26.08", "25.08", "24.08", "23.08", "22.08",
    "21.08", "20.08", "19.08", "18.08", "17.08", "16.8", "15.08",
    "14.08", "13.8", "12.08",
]


def raw_rows(ws):
    header_row, mapping = server.detect_header(ws)
    if not header_row:
        return []
    output = []
    for row in range(header_row + 1, ws.max_row + 1):
        qty = server.number_value(ws.cell(row, mapping["qty"]).value)
        if qty <= 0:
            continue
        item = {
            field: ws.cell(row, col).value
            for field, col in mapping.items()
        }
        item["source_row"] = row
        item["product_code"] = server.clean_text(item.get("product_code")).upper()
        item["contractor"] = server.clean_text(item.get("contractor")).upper()
        item["kitchen"] = server.clean_text(item.get("kitchen")).upper()
        item["supplier"] = server.clean_text(item.get("supplier"))
        item["buy_price"] = server.number_value(item.get("buy_price"))
        item["sell_price"] = server.number_value(item.get("sell_price"))
        item["cccd"] = server.clean_text(item.get("cccd"))
        output.append(item)
    return output


def main():
    server.init_database()
    current, _ = server.parse_workbook(SOURCE, "2026-08-29", [TODAY])
    wb = load_workbook(SOURCE, data_only=True, read_only=False)
    history = defaultdict(list)
    for sheet in HISTORY_SHEETS:
        if sheet not in wb.sheetnames:
            continue
        for item in raw_rows(wb[sheet]):
            history[item["product_code"]].append({
                "sheet": sheet,
                "contractor": item["contractor"],
                "kitchen": item["kitchen"],
                "supplier": item["supplier"],
                "buy_price": item["buy_price"],
                "sell_price": item["sell_price"],
                "cccd": item["cccd"],
            })

    report = []
    for item in current:
        if not item["errors"]:
            continue
        candidates = history.get(item["product_code"], [])
        same_contractor = [
            row for row in candidates
            if row["contractor"] == item["contractor"]
            and (row["sell_price"] > 0 or row["buy_price"] > 0 or row["cccd"])
        ]
        same_kitchen = [
            row for row in same_contractor if row["kitchen"] == item["kitchen"]
        ]
        chosen = (same_kitchen or same_contractor or candidates)[:3]
        report.append({
            "row": item["source_row"],
            "contractor": item["contractor"],
            "kitchen": item["kitchen"],
            "code": item["product_code"],
            "name": item["product_name"],
            "buy": item["buy_price"],
            "sell": item["sell_price"],
            "cccd": item["cccd"],
            "errors": item["errors"],
            "history": chosen,
        })
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
