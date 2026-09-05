"""Versioned, period-scoped import of the customer quotation matrix.

This is deliberately separate from the daily-order importer.  A confirmed
quotation is append-only: a correction creates the next version for the same
period, while orders keep the numeric price snapshots they were created with.
"""

from __future__ import annotations

import hashlib
import io
import json
import math
import re
import threading
import time
import unicodedata
import uuid
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Any

from flask import jsonify, request
from openpyxl import load_workbook


QUOTE_IMPORT_MAX_BYTES = 20 * 1024 * 1024
QUOTE_IMPORT_MAX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
QUOTE_IMPORT_MAX_ENTRIES = 10_000
QUOTE_IMPORT_MAX_ROWS = 100_000
QUOTE_IMPORT_TTL_SECONDS = 15 * 60

PENDING_QUOTE_IMPORTS: dict[str, dict[str, Any]] = {}
QUOTE_IMPORT_LOCK = threading.Lock()


QUOTE_IMPORT_SCHEMA = """
CREATE TABLE IF NOT EXISTS quote_versions (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    effective_period TEXT NOT NULL,
    version_no INTEGER NOT NULL,
    source_hash TEXT NOT NULL,
    source_name TEXT NOT NULL,
    source_sheet TEXT NOT NULL,
    content_hash TEXT NOT NULL,
    product_count INTEGER NOT NULL,
    price_count INTEGER NOT NULL,
    status TEXT NOT NULL DEFAULT 'confirmed',
    created_at TEXT NOT NULL,
    confirmed_at TEXT NOT NULL,
    UNIQUE(effective_period,version_no),
    UNIQUE(effective_period,source_hash),
    CHECK(status='confirmed')
);
CREATE INDEX IF NOT EXISTS idx_quote_versions_period
    ON quote_versions(effective_period,version_no DESC);

CREATE TABLE IF NOT EXISTS quote_version_products (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    version_id INTEGER NOT NULL REFERENCES quote_versions(id) ON DELETE CASCADE,
    source_row INTEGER NOT NULL,
    product_code TEXT NOT NULL,
    product_name TEXT NOT NULL,
    unit TEXT,
    tax TEXT,
    supplier TEXT,
    buy_price REAL,
    buy_price_state TEXT NOT NULL,
    UNIQUE(version_id,source_row),
    CHECK(buy_price_state IN ('numeric','zero','blank','text','formula'))
);
CREATE INDEX IF NOT EXISTS idx_quote_products_lookup
    ON quote_version_products(version_id,product_code,source_row);

CREATE TABLE IF NOT EXISTS quote_version_prices (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    version_id INTEGER NOT NULL REFERENCES quote_versions(id) ON DELETE CASCADE,
    product_line_id INTEGER NOT NULL REFERENCES quote_version_products(id) ON DELETE CASCADE,
    source_row INTEGER NOT NULL,
    source_column INTEGER NOT NULL,
    product_code TEXT NOT NULL,
    price_group TEXT NOT NULL,
    source_header TEXT NOT NULL DEFAULT '',
    price_text TEXT,
    price_value REAL,
    price_state TEXT NOT NULL,
    UNIQUE(product_line_id,price_group),
    CHECK(price_state IN ('numeric','zero','blank','excluded','text','formula'))
);
CREATE INDEX IF NOT EXISTS idx_quote_prices_lookup
    ON quote_version_prices(version_id,product_code,price_group,source_row);
"""


class QuoteImportError(ValueError):
    def __init__(self, message: str, *, code: str = "invalid_quote", status: int = 400):
        super().__init__(message)
        self.code = code
        self.status = status


def init_quote_import_schema(conn) -> None:
    conn.executescript(QUOTE_IMPORT_SCHEMA)
    columns = {row["name"] for row in conn.execute("PRAGMA table_info(quote_version_prices)")}
    if "source_header" not in columns:
        conn.execute(
            "ALTER TABLE quote_version_prices ADD COLUMN source_header TEXT NOT NULL DEFAULT ''"
        )


def _plain(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip())


def _key(value: Any) -> str:
    text = str(value or "").replace("đ", "d").replace("Đ", "D")
    text = unicodedata.normalize("NFKD", text).casefold()
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"[^a-z0-9]+", "", text)


def _hash_json(value: Any) -> str:
    payload = json.dumps(
        value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str,
    ).encode("utf-8")
    return hashlib.sha256(payload).hexdigest().upper()


def _valid_period(value: Any) -> str:
    period = _plain(value)
    if not re.fullmatch(r"\d{4}-\d{2}", period):
        raise QuoteImportError("Kỳ hiệu lực phải có dạng YYYY-MM", code="invalid_period")
    try:
        datetime.strptime(period + "-01", "%Y-%m-%d")
    except ValueError:
        raise QuoteImportError("Kỳ hiệu lực không hợp lệ", code="invalid_period") from None
    return period


def _period_from_date(value: Any) -> str:
    text = _plain(value)
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", text):
        return ""
    try:
        return datetime.strptime(text, "%Y-%m-%d").strftime("%Y-%m")
    except ValueError:
        return ""


def _cell_value(value: Any, *, buy_price: bool = False) -> tuple[str, float | None, str]:
    """Return display text, numeric value and an explicit semantic state."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return "", None, "blank"
    if isinstance(value, str) and value.startswith("="):
        return value, None, "formula"
    if not buy_price and _key(value) == "x":
        return _plain(value), None, "excluded"
    if isinstance(value, bool):
        return _plain(value), None, "text"
    parsed: float | None = None
    if isinstance(value, (int, float)):
        parsed = float(value)
    elif isinstance(value, str):
        candidate = value.strip().replace(" ", "")
        if re.fullmatch(r"[-+]?\d+(?:[.,]\d+)?", candidate):
            # In these workbooks comma is a thousands separator when three
            # digits follow it; otherwise it is treated as a decimal mark.
            if re.fullmatch(r"[-+]?\d{1,3}(?:,\d{3})+", candidate):
                candidate = candidate.replace(",", "")
            else:
                candidate = candidate.replace(",", ".")
            try:
                parsed = float(candidate)
            except ValueError:
                parsed = None
    if parsed is None or not math.isfinite(parsed):
        return _plain(value), None, "text"
    if parsed < 0:
        return _plain(value), None, "text"
    return _plain(value), parsed, "zero" if parsed == 0 else "numeric"


def _source_cell(formula_sheet, value_sheet, row: int, column: int, *, buy_price: bool = False):
    """Read a source cell without confusing cached formulas with plain values.

    The real matrix contains a small number of constant arithmetic formulas
    such as ``=140000/30``.  They are accepted only when the expression is
    arithmetic-only and Excel stored a finite numeric cache.  References,
    functions and external links remain blocked.
    """
    raw = formula_sheet.cell(row, column).value
    if isinstance(raw, str) and raw.startswith("="):
        if not re.fullmatch(r"=[0-9+\-*/(). ]+", raw):
            return raw, None, "formula"
        cached = value_sheet.cell(row, column).value if value_sheet is not None else None
        _text, value, state = _cell_value(cached, buy_price=buy_price)
        if state in {"numeric", "zero"}:
            return raw, value, state
        return raw, None, "formula"
    return _cell_value(raw, buy_price=buy_price)


HEADER_ALIASES = {
    "product_code": {"mahang", "mavt", "mahanghoa"},
    "product_name": {"tenthanhdatphat", "tenhang", "tenhanghoa", "tenvt"},
    "buy_price": {"giamua", "dongiamua"},
    "supplier": {"ncc", "nhacungcap"},
    "unit": {"dvt", "donvitinh"},
    "tax": {"thue", "thuesuat", "thuegtgt"},
}


def _find_price_sheet(workbook):
    matches = [sheet for sheet in workbook.worksheets if _key(sheet.title) == "baogia"]
    if len(matches) != 1:
        raise QuoteImportError(
            "Workbook phải có đúng một sheet BÁO GIÁ",
            code="quote_sheet_missing" if not matches else "quote_sheet_ambiguous",
        )
    return matches[0]


def _find_headers(sheet) -> tuple[int, dict[str, int], list[tuple[int, str]]]:
    for row_no in range(1, min(sheet.max_row, 30) + 1):
        mapping: dict[str, int] = {}
        for column in range(1, sheet.max_column + 1):
            key = _key(sheet.cell(row_no, column).value)
            for field, aliases in HEADER_ALIASES.items():
                if key in aliases and field not in mapping:
                    mapping[field] = column
        if "product_code" not in mapping or "product_name" not in mapping:
            continue
        boundary = mapping.get("tax") or max(mapping.values())
        price_columns: list[tuple[int, str]] = []
        for column in range(boundary + 1, sheet.max_column + 1):
            raw = sheet.cell(row_no, column).value
            if isinstance(raw, str) and raw.startswith("="):
                break
            header = _plain(raw).upper()
            if not header:
                continue
            if _key(header) == "them":
                break
            price_columns.append((column, header))
        if not price_columns:
            raise QuoteImportError(
                "Không tìm thấy cột giá nhà thầu sau cột THUẾ",
                code="contractor_columns_missing",
            )
        duplicate_groups = sorted({
            group for _, group in price_columns
            if sum(candidate == group for _, candidate in price_columns) > 1
        })
        if duplicate_groups:
            raise QuoteImportError(
                "Header nhà thầu bị lặp trong vùng giá: " + ", ".join(duplicate_groups[:10]),
                code="duplicate_contractor_headers",
            )
        return row_no, mapping, price_columns
    raise QuoteImportError(
        "Không nhận diện được header MÃ HÀNG/TÊN THÀNH ĐẠT PHÁT",
        code="quote_header_missing",
    )


def _map_price_columns(conn, columns: list[tuple[int, str]]) -> list[tuple[int, str, str]]:
    """Map each literal source header to one canonical configured price group.

    Matching is exact after whitespace/case normalization.  Column position is
    retained only as provenance; it never decides the contractor identity.
    """
    aliases: dict[str, set[str]] = {}
    for row in conn.execute("SELECT code,price_group FROM contractors"):
        code = _plain(row["code"]).upper()
        canonical = _plain(row["price_group"] or code).upper()
        aliases.setdefault(code, set()).add(canonical)
        aliases.setdefault(canonical, set()).add(canonical)
    mapped: list[tuple[int, str, str]] = []
    used: dict[str, str] = {}
    for source_column, source_header in columns:
        header = _plain(source_header).upper()
        candidates = aliases.get(header, set())
        if not candidates:
            raise QuoteImportError(
                f"Header giá {source_header} chưa khớp mã/nhóm nhà thầu cấu hình",
                code="contractor_header_unknown",
            )
        if len(candidates) != 1:
            raise QuoteImportError(
                f"Header giá {source_header} khớp nhiều nhóm nhà thầu",
                code="contractor_header_ambiguous",
            )
        canonical = next(iter(candidates))
        if canonical in used:
            raise QuoteImportError(
                f"Hai header {used[canonical]} và {source_header} cùng trỏ nhóm {canonical}",
                code="duplicate_contractor_mapping",
            )
        used[canonical] = source_header
        mapped.append((source_column, canonical, source_header))
    return mapped


def _price_signature(price: dict[str, Any]) -> tuple[Any, ...]:
    state = price["price_state"]
    if state in {"numeric", "zero"}:
        return state, float(price["price_value"] or 0)
    return state, _key(price.get("price_text"))


def _buy_explicit_signature(item: dict[str, Any]) -> tuple[Any, ...] | None:
    if item["buy_price_state"] in {"numeric", "zero"}:
        return item["buy_price_state"], float(item["buy_price"] or 0)
    return None


def _duplicate_analysis(items: list[dict[str, Any]], price_groups: list[str]) -> dict[str, Any]:
    by_code: dict[str, list[dict[str, Any]]] = {}
    for item in items:
        by_code.setdefault(item["product_code"], []).append(item)
    conflicts: list[dict[str, Any]] = []
    duplicate_codes = 0
    safe_duplicate_codes = 0
    for code, rows in by_code.items():
        if len(rows) < 2:
            continue
        duplicate_codes += 1
        code_conflicts = 0
        for group in price_groups:
            candidates = []
            for item in rows:
                price = next(price for price in item["prices"] if price["price_group"] == group)
                if price["price_state"] not in {"blank", "excluded"}:
                    candidates.append((item, price))
            if len(candidates) < 2:
                continue
            reasons = []
            if len({_price_signature(price) for _item, price in candidates}) > 1:
                reasons.append("giá/trạng thái khác nhau")
            if len({
                (_key(item["product_name"]), _key(item["unit"]), _key(item["tax"]))
                for item, _price in candidates
            }) > 1:
                reasons.append("tên/ĐVT/thuế khác nhau")
            explicit_buys = {
                signature for item, _price in candidates
                if (signature := _buy_explicit_signature(item)) is not None
            }
            if len(explicit_buys) > 1:
                reasons.append("giá mua khác nhau")
            if not reasons:
                continue
            code_conflicts += 1
            conflicts.append({
                "productCode": code,
                "priceGroup": group,
                "sourceRows": [item["source_row"] for item, _price in candidates],
                "reasons": reasons,
                "candidates": [{
                    "sourceRow": item["source_row"],
                    "buyPrice": item["buy_price"],
                    "buyPriceState": item["buy_price_state"],
                    "priceValue": price["price_value"],
                    "priceState": price["price_state"],
                    "priceText": price["price_text"],
                } for item, price in candidates],
            })
        if not code_conflicts:
            safe_duplicate_codes += 1
    return {
        "duplicate_codes": duplicate_codes,
        "safe_duplicate_codes": safe_duplicate_codes,
        "conflict_count": len(conflicts),
        "conflict_codes": len({item["productCode"] for item in conflicts}),
        "conflicts": conflicts,
    }


def _database_state_hash(conn) -> str:
    queries = (
        "SELECT code,name FROM products ORDER BY code",
        "SELECT code,COALESCE(price_group,''),pricing_mode FROM contractors ORDER BY code",
        "SELECT effective_period,version_no,source_hash,content_hash,status "
        "FROM quote_versions ORDER BY effective_period,version_no",
    )
    digest = hashlib.sha256()
    for sql in queries:
        digest.update(sql.encode("utf-8"))
        digest.update(b"\0")
        for row in conn.execute(sql):
            digest.update(json.dumps(list(row), ensure_ascii=False, separators=(",", ":")).encode("utf-8"))
            digest.update(b"\n")
    return digest.hexdigest().upper()


def _parse_quote(conn, workbook, value_workbook=None) -> dict[str, Any]:
    sheet = _find_price_sheet(workbook)
    value_sheet = _find_price_sheet(value_workbook) if value_workbook is not None else None
    header_row, columns, raw_price_columns = _find_headers(sheet)
    price_columns = _map_price_columns(conn, raw_price_columns)
    known_products = {
        _plain(row["code"]).upper(): _plain(row["name"])
        for row in conn.execute("SELECT code,name FROM products")
    }
    items: list[dict[str, Any]] = []
    row_errors = 0
    warnings = 0
    for row_no in range(header_row + 1, min(sheet.max_row, QUOTE_IMPORT_MAX_ROWS) + 1):
        code_raw = sheet.cell(row_no, columns["product_code"]).value
        name_raw = sheet.cell(row_no, columns["product_name"]).value
        # The customer's matrix has one field-index row (2, 3, 4, ...)
        # immediately below the labels.  It is metadata, not a product whose
        # code happens to be "2".
        if (
            isinstance(code_raw, (int, float)) and not isinstance(code_raw, bool)
            and isinstance(name_raw, (int, float)) and not isinstance(name_raw, bool)
        ):
            continue
        code = _plain(code_raw).upper()
        name = _plain(name_raw)
        if not code and not name:
            continue
        if not code:
            other_columns = [
                column for field, column in columns.items()
                if field not in {"product_code", "product_name"}
            ] + [column for column, _group, _header in price_columns]
            if not any(sheet.cell(row_no, column).value not in (None, "") for column in other_columns):
                # Notes/group labels in the name column are not product rows.
                continue
        errors: list[str] = []
        row_warnings: list[str] = []
        if not code:
            errors.append("Thiếu mã hàng")
        elif code not in known_products:
            errors.append("Mã hàng chưa có trong danh mục")
        if not name and code in known_products:
            name = known_products[code]
            row_warnings.append("Tên hàng trống; dùng tên danh mục")
        elif not name:
            errors.append("Thiếu tên hàng")
        elif code in known_products and _key(name) != _key(known_products[code]):
            row_warnings.append("Tên hàng khác danh mục")
        if "buy_price" in columns:
            buy_text, buy_value, buy_state = _source_cell(
                sheet, value_sheet, row_no, columns["buy_price"], buy_price=True,
            )
        else:
            buy_text, buy_value, buy_state = "", None, "blank"
        if buy_state == "formula":
            errors.append("Công thức giá mua không an toàn hoặc chưa có giá trị lưu")
        elif buy_state == "text":
            row_warnings.append("Giá mua là trạng thái chữ; chỉ mua thực tế mới được bổ sung")
        prices = []
        for source_column, group, source_header in price_columns:
            price_text, price_value, price_state = _source_cell(
                sheet, value_sheet, row_no, source_column,
            )
            if price_state == "formula":
                errors.append(f"Giá {group} không được là công thức")
            elif price_state == "text" and price_text:
                row_warnings.append(f"Giá {group} là trạng thái chữ")
            prices.append({
                "source_column": source_column,
                "price_group": group,
                "source_header": source_header,
                "price_text": price_text,
                "price_value": price_value,
                "price_state": price_state,
            })
        item = {
            "source_row": row_no,
            "product_code": code,
            "product_name": name,
            "unit": _plain(sheet.cell(row_no, columns["unit"]).value) if "unit" in columns else "",
            "tax": _plain(sheet.cell(row_no, columns["tax"]).value) if "tax" in columns else "",
            "supplier": _plain(sheet.cell(row_no, columns["supplier"]).value) if "supplier" in columns else "",
            "buy_price_text": buy_text,
            "buy_price": buy_value,
            "buy_price_state": buy_state,
            "prices": prices,
            "errors": errors,
            "warnings": row_warnings,
        }
        items.append(item)
        row_errors += bool(errors)
        warnings += bool(row_warnings)
    if sheet.max_row > QUOTE_IMPORT_MAX_ROWS:
        raise QuoteImportError("Sheet BÁO GIÁ vượt giới hạn 100.000 dòng", code="too_many_rows", status=413)
    if not items:
        raise QuoteImportError("Sheet BÁO GIÁ không có dòng dữ liệu", code="quote_empty")
    price_groups = [group for _column, group, _header in price_columns]
    duplicate_analysis = _duplicate_analysis(items, price_groups)
    serializable = [{key: value for key, value in item.items() if key != "warnings"} for item in items]
    return {
        "sheet": sheet.title,
        "header_row": header_row,
        "price_groups": price_groups,
        "price_columns": [{
            "sourceColumn": column,
            "sourceHeader": source_header,
            "priceGroup": group,
        } for column, group, source_header in price_columns],
        "items": items,
        "conflicts": duplicate_analysis["conflicts"],
        "content_hash": _hash_json(serializable),
        "counts": {
            "products": len(items),
            "price_groups": len(price_columns),
            "price_cells": len(items) * len(price_columns),
            "rows_with_errors": row_errors,
            "rows_with_warnings": warnings,
            "duplicate_codes": duplicate_analysis["duplicate_codes"],
            "safe_duplicate_codes": duplicate_analysis["safe_duplicate_codes"],
            "conflict_codes": duplicate_analysis["conflict_codes"],
            "conflicts": duplicate_analysis["conflict_count"],
        },
    }


def _public_items(items: list[dict[str, Any]], limit: int = 250) -> list[dict[str, Any]]:
    public = []
    for item in items[:limit]:
        public.append({
            "sourceRow": item["source_row"],
            "productCode": item["product_code"],
            "productName": item["product_name"],
            "unit": item["unit"],
            "tax": item["tax"],
            "supplier": item["supplier"],
            "buyPrice": item["buy_price"],
            "buyPriceState": item["buy_price_state"],
            "prices": [{
                "group": price["price_group"],
                "sourceHeader": price["source_header"],
                "sourceColumn": price["source_column"],
                "value": price["price_value"],
                "state": price["price_state"],
                "text": price["price_text"],
            } for price in item["prices"]],
            "errors": item["errors"],
            "warnings": item["warnings"],
        })
    return public


def active_quote_version(conn, period: str):
    return conn.execute(
        """SELECT * FROM quote_versions
           WHERE effective_period=? AND status='confirmed'
           ORDER BY version_no DESC LIMIT 1""",
        (period,),
    ).fetchone()


def _resolve_stored_candidates(rows) -> dict[str, Any]:
    candidates = [dict(row) for row in rows if row["price_state"] not in {"blank", "excluded"}]
    if not candidates:
        states = {row["price_state"] for row in rows}
        return {
            "selected": None,
            "conflict": False,
            "source_rows": [row["source_row"] for row in rows],
            "excluded": "excluded" in states,
        }
    reasons = []
    if len({_price_signature(row) for row in candidates}) > 1:
        reasons.append("giá/trạng thái khác nhau")
    if len({
        (_key(row["product_name"]), _key(row["unit"]), _key(row["tax"]))
        for row in candidates
    }) > 1:
        reasons.append("tên/ĐVT/thuế khác nhau")
    explicit_buys = {
        (row["buy_price_state"], float(row["buy_price"] or 0))
        for row in candidates if row["buy_price_state"] in {"numeric", "zero"}
    }
    if len(explicit_buys) > 1:
        reasons.append("giá mua khác nhau")
    if reasons:
        return {
            "selected": None,
            "conflict": True,
            "reasons": reasons,
            "source_rows": [row["source_row"] for row in candidates],
            "candidates": candidates,
        }
    selected = candidates[0]
    if explicit_buys:
        buy_signature = next(iter(explicit_buys))
        selected = next(
            row for row in candidates
            if (row["buy_price_state"], float(row["buy_price"] or 0)) == buy_signature
        )
    return {
        "selected": selected,
        "conflict": False,
        "source_rows": [row["source_row"] for row in candidates],
        "duplicate_count": max(len(candidates) - 1, 0),
    }


def _stored_group_rows(conn, version_id: int, product_code: str, price_group: str):
    return conn.execute(
        """SELECT p.id product_line_id,p.source_row,p.product_code,p.product_name,p.unit,p.tax,
                  p.supplier,p.buy_price,p.buy_price_state,q.source_column,q.source_header,
                  q.price_group,q.price_text,q.price_value,q.price_state
           FROM quote_version_products p
           JOIN quote_version_prices q ON q.product_line_id=p.id
           WHERE p.version_id=? AND p.product_code=? AND q.price_group=?
           ORDER BY p.source_row""",
        (version_id, _plain(product_code).upper(), _plain(price_group).upper()),
    ).fetchall()


def quote_rows_for_contractor(
    conn, contractor: str, period: str, version_id: int | None = None,
) -> dict[str, Any] | None:
    """Return one unambiguous row per code for a confirmed period quotation.

    X/blank rows are not returned.  Numeric zero is deliberately returned and
    marked exportable.  Conflicting duplicates are listed and omitted rather
    than choosing a source row silently.
    """
    contractor_code = _plain(contractor).upper()
    contractor_row = conn.execute(
        "SELECT code,price_group,pricing_mode FROM contractors WHERE code=?", (contractor_code,)
    ).fetchone()
    if not contractor_row or contractor_row["pricing_mode"] == "daily":
        return None
    if version_id is None:
        version = active_quote_version(conn, period)
    else:
        version = conn.execute(
            """SELECT * FROM quote_versions
               WHERE id=? AND effective_period=? AND status='confirmed'""",
            (version_id, period),
        ).fetchone()
    if not version:
        return None
    group = _plain(contractor_row["price_group"] or contractor_code).upper()
    rows = conn.execute(
        """SELECT p.id product_line_id,p.source_row,p.product_code,p.product_name,p.unit,p.tax,
                  p.supplier,p.buy_price,p.buy_price_state,q.source_column,q.source_header,
                  q.price_group,q.price_text,q.price_value,q.price_state
           FROM quote_version_products p
           JOIN quote_version_prices q ON q.product_line_id=p.id
           WHERE p.version_id=? AND q.price_group=?
           ORDER BY p.product_code,p.source_row""",
        (version["id"], group),
    ).fetchall()
    by_code: dict[str, list[Any]] = {}
    for row in rows:
        by_code.setdefault(row["product_code"], []).append(row)
    items = []
    conflicts = []
    excluded_count = 0
    for code, code_rows in by_code.items():
        resolution = _resolve_stored_candidates(code_rows)
        if resolution["conflict"]:
            conflicts.append({
                "product_code": code,
                "price_group": group,
                "source_rows": resolution["source_rows"],
                "reasons": resolution["reasons"],
            })
            continue
        selected = resolution["selected"]
        if selected is None:
            excluded_count += 1
            continue
        state = selected["price_state"]
        if state == "zero":
            status = "Giá 0 – giữ để xác nhận"
        elif state == "text":
            status = _plain(selected["price_text"]) or "Trạng thái chữ"
        else:
            status = ""
        items.append({
            "product_code": selected["product_code"],
            "product_name": selected["product_name"],
            "unit": selected["unit"],
            "tax": selected["tax"],
            "sell_price": selected["price_value"],
            "status": status,
            "price_state": state,
            "exportable": state in {"numeric", "zero"},
            "source_row": selected["source_row"],
            "source_rows": resolution["source_rows"],
            "source_column": selected["source_column"],
            "source_header": selected["source_header"],
            "duplicate_count": resolution.get("duplicate_count", 0),
        })
    items.sort(key=lambda item: (_key(item["product_name"]), item["product_code"]))
    return {
        "period": period,
        "price_group": group,
        "version": {
            "id": version["id"],
            "version_no": version["version_no"],
            "source_hash": version["source_hash"],
            "confirmed_at": version["confirmed_at"],
        },
        "items": items,
        "conflicts": conflicts,
        "excluded_count": excluded_count,
        "output_count": sum(bool(item["exportable"]) for item in items),
        "status_count": sum(not item["exportable"] for item in items),
    }


def quote_sell_price(conn, product_code: str, contractor: str, work_date: str) -> dict[str, Any]:
    period = _period_from_date(work_date)
    contractor_code = _plain(contractor).upper()
    contractor_row = conn.execute(
        "SELECT price_group,pricing_mode FROM contractors WHERE code=?", (contractor_code,)
    ).fetchone()
    if contractor_row and contractor_row["pricing_mode"] == "daily":
        return {"has_version": False, "applicable": False, "value": None, "message": ""}
    version = active_quote_version(conn, period) if period else None
    if not version:
        return {"has_version": False, "applicable": True, "value": None, "message": ""}
    group = _plain(contractor_row["price_group"] if contractor_row else contractor_code).upper()
    rows = _stored_group_rows(conn, version["id"], product_code, group)
    context = {
        "has_version": True,
        "applicable": True,
        "version_id": version["id"],
        "version_no": version["version_no"],
        "period": period,
        "value": None,
    }
    if not rows:
        context["message"] = "Chưa có giá cho nhà thầu trong báo giá đúng kỳ"
        return context
    resolution = _resolve_stored_candidates(rows)
    if resolution["conflict"]:
        context["message"] = "Mã hàng bị trùng và xung đột trong báo giá đúng kỳ"
        return context
    row = resolution["selected"]
    if row is None:
        context["message"] = (
            "Mã hàng bị đánh dấu X cho nhà thầu trong báo giá đúng kỳ"
            if resolution["excluded"] else "Giá nhà thầu để trống trong báo giá đúng kỳ"
        )
        return context
    if row["price_state"] == "numeric" and row["price_value"] is not None and row["price_value"] > 0:
        context.update(value=float(row["price_value"]), message="")
    elif row["price_state"] == "zero":
        context["message"] = "Giá nhà thầu bằng 0 trong báo giá đúng kỳ"
    elif row["price_state"] == "excluded":
        context["message"] = "Mã hàng bị đánh dấu X cho nhà thầu trong báo giá đúng kỳ"
    else:
        context["message"] = _plain(row["price_text"]) or "Giá nhà thầu để trống trong báo giá đúng kỳ"
    return context


def quote_buy_price(conn, product_code: str, work_date: str, contractor: str = "") -> dict[str, Any]:
    period = _period_from_date(work_date)
    version = active_quote_version(conn, period) if period else None
    if not version:
        return {"has_version": False, "value": None, "allow_actual_fallback": True}
    contractor_code = _plain(contractor).upper()
    contractor_row = conn.execute(
        "SELECT price_group,pricing_mode FROM contractors WHERE code=?", (contractor_code,)
    ).fetchone() if contractor_code else None
    if contractor_row and contractor_row["pricing_mode"] != "daily":
        group = _plain(contractor_row["price_group"] or contractor_code).upper()
        priced_rows = _stored_group_rows(conn, version["id"], product_code, group)
        resolution = _resolve_stored_candidates(priced_rows)
        if resolution["conflict"]:
            rows = []
            forced_conflict = True
        elif resolution["selected"] is not None:
            rows = [resolution["selected"]]
            forced_conflict = False
        else:
            rows = []
            forced_conflict = False
    else:
        rows = conn.execute(
            """SELECT buy_price,buy_price_state,source_row FROM quote_version_products
               WHERE version_id=? AND product_code=? ORDER BY source_row""",
            (version["id"], _plain(product_code).upper()),
        ).fetchall()
        forced_conflict = False
    result = {
        "has_version": True,
        "version_id": version["id"],
        "version_no": version["version_no"],
        "period": period,
        "value": None,
        "allow_actual_fallback": False,
    }
    if forced_conflict:
        result["message"] = "Mã hàng bị trùng và xung đột trong báo giá đúng kỳ"
        return result
    if not rows:
        result.update(allow_actual_fallback=True, message="Mã hàng không có trong báo giá đúng kỳ")
        return result
    explicit = {
        (row["buy_price_state"], float(row["buy_price"] or 0))
        for row in rows if row["buy_price_state"] in {"numeric", "zero"}
    }
    if len(explicit) > 1:
        result["message"] = "Mã hàng bị trùng và khác giá mua trong báo giá đúng kỳ"
        return result
    if explicit:
        signature = next(iter(explicit))
        row = next(
            row for row in rows
            if (row["buy_price_state"], float(row["buy_price"] or 0)) == signature
        )
    else:
        row = rows[0]
    if row["buy_price_state"] == "numeric" and row["buy_price"] is not None and row["buy_price"] > 0:
        result.update(value=float(row["buy_price"]), message="")
    elif row["buy_price_state"] in {"blank", "text"}:
        result.update(allow_actual_fallback=True, message="Giá mua trống; được bổ sung từ mua thực tế")
    elif row["buy_price_state"] == "zero":
        result["message"] = "Giá mua bằng 0 trong báo giá đúng kỳ"
    else:
        result["message"] = "Giá mua không hợp lệ trong báo giá đúng kỳ"
    return result


def register_quote_import_routes(app, ctx) -> None:
    db_factory = ctx["db"]
    now_iso = ctx["now_iso"]
    audit_event = ctx["audit_event"]

    def error_response(exc: QuoteImportError):
        return jsonify({"ok": False, "error": str(exc), "code": exc.code}), exc.status

    @app.post("/api/quotes/import/preview")
    def api_quote_import_preview():
        try:
            period = _valid_period(request.form.get("effective_period"))
        except QuoteImportError as exc:
            return error_response(exc)
        upload = request.files.get("file")
        if not upload or not upload.filename:
            return jsonify({"ok": False, "error": "Chưa chọn file báo giá", "code": "file_required"}), 400
        if Path(upload.filename).suffix.lower() not in {".xlsx", ".xlsm"}:
            return jsonify({"ok": False, "error": "Chỉ nhận file .xlsx hoặc .xlsm", "code": "invalid_file_type"}), 400
        payload = upload.read(QUOTE_IMPORT_MAX_BYTES + 1)
        if len(payload) > QUOTE_IMPORT_MAX_BYTES:
            return jsonify({"ok": False, "error": "File báo giá vượt giới hạn 20 MB", "code": "file_too_large"}), 413
        try:
            with zipfile.ZipFile(io.BytesIO(payload)) as archive:
                entries = archive.infolist()
                if len(entries) > QUOTE_IMPORT_MAX_ENTRIES or sum(item.file_size for item in entries) > QUOTE_IMPORT_MAX_UNCOMPRESSED_BYTES:
                    raise QuoteImportError("Workbook có cấu trúc quá lớn", code="expanded_too_large", status=413)
        except zipfile.BadZipFile:
            return jsonify({"ok": False, "error": "Workbook bị hỏng hoặc sai định dạng", "code": "invalid_workbook"}), 400

        workbook = None
        value_workbook = None
        try:
            # data_only=False is intentional: formula headers/price cells must
            # be detectable instead of trusting a stale Excel formula cache.
            workbook = load_workbook(io.BytesIO(payload), data_only=False, read_only=False, keep_links=False)
            value_workbook = load_workbook(io.BytesIO(payload), data_only=True, read_only=False, keep_links=False)
            with db_factory() as conn:
                parsed = _parse_quote(conn, workbook, value_workbook)
                database_state_hash = _database_state_hash(conn)
                existing = conn.execute(
                    "SELECT id,version_no FROM quote_versions WHERE effective_period=? AND source_hash=?",
                    (period, hashlib.sha256(payload).hexdigest().upper()),
                ).fetchone()
                latest_no = conn.execute(
                    "SELECT COALESCE(MAX(version_no),0) FROM quote_versions WHERE effective_period=?",
                    (period,),
                ).fetchone()[0]
        except QuoteImportError as exc:
            return error_response(exc)
        except (OSError, ValueError, KeyError, zipfile.BadZipFile):
            return jsonify({"ok": False, "error": "Không đọc được workbook báo giá", "code": "invalid_workbook"}), 400
        finally:
            if workbook is not None:
                workbook.close()
            if value_workbook is not None:
                value_workbook.close()

        source_hash = hashlib.sha256(payload).hexdigest().upper()
        proposed_version = int(existing["version_no"] if existing else latest_no + 1)
        state_hash = hashlib.sha256(
            f"{source_hash}\0{period}\0{database_state_hash}\0{parsed['content_hash']}\0{proposed_version}".encode("utf-8")
        ).hexdigest().upper()
        token = uuid.uuid4().hex
        pending = {
            "created": time.time(),
            "period": period,
            "source_hash": source_hash,
            "source_name": Path(upload.filename).name[:255],
            "state_hash": state_hash,
            "database_state_hash": database_state_hash,
            "proposed_version": proposed_version,
            "existing_id": int(existing["id"]) if existing else None,
            **parsed,
        }
        with QUOTE_IMPORT_LOCK:
            cutoff = time.time() - QUOTE_IMPORT_TTL_SECONDS
            for old_token, item in list(PENDING_QUOTE_IMPORTS.items()):
                if item["created"] < cutoff:
                    PENDING_QUOTE_IMPORTS.pop(old_token, None)
            PENDING_QUOTE_IMPORTS[token] = pending
        return jsonify({
            "ok": True,
            "token": token,
            "effectivePeriod": period,
            "sourceHash": source_hash,
            "stateHash": state_hash,
            "sheet": parsed["sheet"],
            "headerRow": parsed["header_row"],
            "proposedVersion": proposed_version,
            "replay": existing is not None,
            "canConfirm": (
                parsed["counts"]["rows_with_errors"] == 0
                and parsed["counts"]["conflicts"] == 0
            ),
            "priceGroups": parsed["price_groups"],
            "priceColumns": parsed["price_columns"],
            "counts": parsed["counts"],
            "conflicts": parsed["conflicts"][:250],
            "conflictsTruncated": len(parsed["conflicts"]) > 250,
            "items": _public_items(parsed["items"]),
            "itemsTruncated": len(parsed["items"]) > 250,
        })

    @app.post("/api/quotes/import/confirm")
    def api_quote_import_confirm():
        body = request.get_json(force=True) or {}
        token = _plain(body.get("token"))
        with QUOTE_IMPORT_LOCK:
            pending = PENDING_QUOTE_IMPORTS.pop(token, None)
        if not pending or time.time() - pending["created"] > QUOTE_IMPORT_TTL_SECONDS:
            return jsonify({"ok": False, "error": "Phiên preview đã hết hạn", "code": "preview_expired"}), 410
        if body.get("confirmed") is not True:
            return jsonify({"ok": False, "error": "Cần xác nhận rõ trước khi ghi", "code": "confirmation_required"}), 400
        supplied_hash = _plain(body.get("state_hash") or body.get("stateHash")).upper()
        if supplied_hash != pending["state_hash"]:
            return jsonify({"ok": False, "error": "Preview không còn đúng trạng thái", "code": "stale_preview"}), 409
        if pending["counts"]["rows_with_errors"] or pending["counts"]["conflicts"]:
            return jsonify({
                "ok": False,
                "error": "Preview còn lỗi/xung đột mã nên không thể ghi",
                "code": "preview_has_errors",
            }), 400

        with db_factory() as conn:
            try:
                conn.execute("BEGIN IMMEDIATE")
                receipt = conn.execute(
                    "SELECT id,version_no FROM quote_versions WHERE effective_period=? AND source_hash=?",
                    (pending["period"], pending["source_hash"]),
                ).fetchone()
                if receipt:
                    return jsonify({
                        "ok": True,
                        "effectivePeriod": pending["period"],
                        "versionId": receipt["id"],
                        "versionNo": receipt["version_no"],
                        "sourceHash": pending["source_hash"],
                        "counts": pending["counts"],
                        "idempotent": True,
                    })
                if _database_state_hash(conn) != pending["database_state_hash"]:
                    return jsonify({
                        "ok": False,
                        "error": "Dữ liệu giá/danh mục đã đổi sau preview; chưa ghi gì, vui lòng preview lại",
                        "code": "stale_database",
                    }), 409
                next_version = conn.execute(
                    "SELECT COALESCE(MAX(version_no),0)+1 FROM quote_versions WHERE effective_period=?",
                    (pending["period"],),
                ).fetchone()[0]
                if int(next_version) != pending["proposed_version"]:
                    return jsonify({
                        "ok": False,
                        "error": "Đã có phiên bản báo giá mới hơn; vui lòng preview lại",
                        "code": "stale_version",
                    }), 409
                timestamp = now_iso()
                cursor = conn.execute(
                    """INSERT INTO quote_versions(
                           effective_period,version_no,source_hash,source_name,source_sheet,
                           content_hash,product_count,price_count,status,created_at,confirmed_at
                       ) VALUES(?,?,?,?,?,?,?,?, 'confirmed',?,?)""",
                    (
                        pending["period"], next_version, pending["source_hash"], pending["source_name"],
                        pending["sheet"], pending["content_hash"], pending["counts"]["products"],
                        pending["counts"]["price_cells"], timestamp, timestamp,
                    ),
                )
                version_id = cursor.lastrowid
                for item in pending["items"]:
                    product_cursor = conn.execute(
                        """INSERT INTO quote_version_products(
                               version_id,source_row,product_code,product_name,unit,tax,supplier,
                               buy_price,buy_price_state
                           ) VALUES(?,?,?,?,?,?,?,?,?)""",
                        (
                            version_id, item["source_row"], item["product_code"], item["product_name"],
                            item["unit"], item["tax"], item["supplier"], item["buy_price"],
                            item["buy_price_state"],
                        ),
                    )
                    product_line_id = product_cursor.lastrowid
                    for price in item["prices"]:
                        conn.execute(
                            """INSERT INTO quote_version_prices(
                                   version_id,product_line_id,source_row,source_column,product_code,
                                   price_group,source_header,price_text,price_value,price_state
                               ) VALUES(?,?,?,?,?,?,?,?,?,?)""",
                            (
                                version_id, product_line_id, item["source_row"], price["source_column"],
                                item["product_code"], price["price_group"], price["source_header"], price["price_text"],
                                price["price_value"], price["price_state"],
                            ),
                        )
                audit_event(
                    conn,
                    "quote.import.confirm",
                    entity_type="quote_version",
                    entity_id=str(version_id),
                    metadata={
                        "effective_period": pending["period"],
                        "version_no": int(next_version),
                        "source_hash": pending["source_hash"],
                        "content_hash": pending["content_hash"],
                        "product_count": pending["counts"]["products"],
                        "price_count": pending["counts"]["price_cells"],
                    },
                )
            except Exception:
                conn.rollback()
                raise
            return jsonify({
                "ok": True,
                "effectivePeriod": pending["period"],
                "versionId": version_id,
                "versionNo": int(next_version),
                "sourceHash": pending["source_hash"],
                "counts": pending["counts"],
                "idempotent": False,
            })

    @app.get("/api/quotes/versions")
    def api_quote_versions():
        raw_period = _plain(request.args.get("period"))
        try:
            period = _valid_period(raw_period) if raw_period else ""
        except QuoteImportError as exc:
            return error_response(exc)
        with db_factory() as conn:
            sql = (
                "SELECT id,effective_period,version_no,source_hash,source_name,source_sheet,"
                "content_hash,product_count,price_count,status,created_at,confirmed_at "
                "FROM quote_versions"
            )
            params: tuple[Any, ...] = ()
            if period:
                sql += " WHERE effective_period=?"
                params = (period,)
            sql += " ORDER BY effective_period DESC,version_no DESC"
            items = [dict(row) for row in conn.execute(sql, params)]
        return jsonify({"ok": True, "period": period, "items": items})
