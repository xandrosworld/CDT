"""Import the customer's values-only CCCD reference, without guessed addresses."""
from __future__ import annotations

import hashlib
import re
import unicodedata
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


CATALOG_FILENAME = "seller_identities_20260904.xlsx"
VERSION_KEY = "seller_identity_catalog_version"
HEADERS = ("STT", "Tên người bán", "Địa chỉ", "Số CMT nhân dân", "Ngày cấp", "Nơi cấp")
# Explicit customer decision, 05/09/2026. Keep historical rows, never guess IDs.
EXCLUDED_SELLERS = ("Đoàn Văn Giang", "Nguyễn Văn Toại")


def name_key(value):
    # Preserve accents: do not guess that different Vietnamese names are one person.
    return " ".join(unicodedata.normalize("NFC", str(value or "")).split()).casefold()


def is_excluded_seller(value):
    return name_key(value) in {name_key(name) for name in EXCLUDED_SELLERS}


def read_catalog(path):
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        sheet = workbook["CCCD"]
        if tuple(cell.value for cell in next(sheet.iter_rows(max_row=1, max_col=6))) != HEADERS:
            raise ValueError("Danh mục CCCD không đúng tiêu đề nguồn")
        records, names = [], set()
        for row_number, cells in enumerate(sheet.iter_rows(min_row=2, max_col=6), 2):
            if not any(cell.value is not None for cell in cells[1:]):
                continue
            if any(cell.data_type == "f" for cell in cells):
                raise ValueError(f"Danh mục CCCD dòng {row_number}: không nhận công thức")
            _, name, address, identity, issued, place = [cell.value for cell in cells]
            name, address, place = [str(v or "").strip() for v in (name, address, place)]
            if not name or not address or not place or not isinstance(identity, str) or not re.fullmatch(r"(?:\d{9}|\d{12})", identity):
                raise ValueError(f"Danh mục CCCD dòng {row_number}: thiếu hoặc sai thông tin")
            if name_key(name) in names:
                raise ValueError(f"Danh mục CCCD dòng {row_number}: trùng tên người bán")
            names.add(name_key(name))
            try:
                issued = issued if isinstance(issued, date) else datetime.strptime(str(issued).strip(), "%d/%m/%Y")
            except ValueError:
                raise ValueError(f"Danh mục CCCD dòng {row_number}: ngày cấp không hợp lệ") from None
            records.append((name, identity, issued.strftime("%d/%m/%Y"), place, address))
        if not records:
            raise ValueError("Danh mục CCCD trống")
        return records
    finally:
        workbook.close()


def sync_catalog(conn, path):
    """Apply once per source revision; caller owns transaction. Never touch orders.

    Duplicate IDs in the supplied source are retained for correction, but receipt
    enrichment must refuse them. No personal values appear in errors or markers.
    """
    path = Path(path)
    version = "1-" + hashlib.sha256(path.read_bytes()).hexdigest()
    previous = conn.execute("SELECT value FROM settings WHERE key=?", (VERSION_KEY,)).fetchone()
    if previous and previous[0] == version:
        return 0
    records = read_catalog(path)
    existing = {}
    for (name,) in conn.execute("SELECT name FROM people"):
        existing.setdefault(name_key(name), []).append(name)
    # Preflight every match before any writes, including ambiguous existing data.
    for name, *_ in records:
        if len(existing.get(name_key(name), [])) > 1:
            raise ValueError("Danh mục hiện tại có tên người bán trùng nhau; cần kiểm tra trước khi cập nhật CCCD")
    for name, identity, issued, place, address in records:
        matches = existing.get(name_key(name), [])
        target = matches[0] if matches else name
        conn.execute(
            "INSERT INTO people(name,cccd,issue_date,issue_place,address) VALUES(?,?,?,?,?) "
            "ON CONFLICT(name) DO UPDATE SET cccd=excluded.cccd,issue_date=excluded.issue_date,"
            "issue_place=excluded.issue_place,address=excluded.address",
            (target, identity, issued, place, address),
        )
    conn.execute("INSERT INTO settings(key,value) VALUES(?,?) ON CONFLICT(key) DO UPDATE SET value=excluded.value", (VERSION_KEY, version))
    return len(records)
