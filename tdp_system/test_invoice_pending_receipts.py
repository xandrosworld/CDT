"""Carry unresolved invoices across dates; review and post complete invoices only."""
import sqlite3
import unittest
from contextlib import contextmanager

from flask import Flask
from openpyxl import load_workbook

from .invoice_mapping import save_mapping, save_conversion
from .invoice_workbench import register_invoice_workbench_routes
from .invoice_workbench_listing import invoice_range_payload
from .test_invoice_workbench_listing import seed_round1, init_test_database, now_iso


class PendingReceiptTests(unittest.TestCase):
    def setUp(self):
        self.conn = sqlite3.connect(':memory:')
        init_test_database(self.conn)
        self.ids = seed_round1(self.conn)
        self.conn.commit()

        @contextmanager
        def db():
            with self.conn:
                yield self.conn

        app = Flask(__name__)
        register_invoice_workbench_routes(app, {'db': db, 'now_iso': now_iso, 'setting_get': lambda c, k, d: d})
        self.client = app.test_client()
        self.query = '?from=2026-09-01&to=2026-09-02&scope=pending'

    def tearDown(self):
        self.conn.close()

    def listing(self, query=None):
        response = self.client.get('/api/invoice-workbench/invoices' + (query or self.query))
        self.assertEqual(200, response.status_code)
        return response.json

    def test_carry_forward_is_read_only_and_respects_end_date_filters_and_tenant(self):
        before = '\n'.join(self.conn.iterdump())
        self.assertEqual(0, self.listing(self.query.replace('&scope=pending', ''))['totals']['invoice_count'])
        payload = self.listing()
        self.assertEqual({'2', '3'}, {r['invoice_number'] for r in payload['items']})
        self.assertEqual(3, payload['totals']['line_count'])
        self.assertEqual({'2'}, {r['invoice_number'] for r in self.listing('?from=2026-08-01&to=2026-08-02&scope=pending')['items']})
        self.assertEqual(1, self.listing(self.query + '&line_filter=unit_review')['totals']['line_count'])
        other = invoice_range_payload(self.conn, tenant='OTHER', invoice_type='input', date_from='2026-09-01', date_to='2026-09-02', scope='pending')
        self.assertEqual([], other['items'])
        self.assertEqual(before, '\n'.join(self.conn.iterdump()))
        # Output now has the same explicit carry-forward view for unfinished invoices.
        self.assertEqual(200, self.client.get('/api/invoice-workbench/invoices' + self.query + '&invoice_type=output').status_code)
        self.assertEqual(400, self.client.get('/api/invoice-workbench/invoices' + self.query.replace('pending', 'bogus')).status_code)

    def test_review_posts_ready_invoice_once_and_keeps_other_invoice_next_day(self):
        save_mapping(self.conn, direction='input', item_id=self.ids['line_b'], product_code='R1-KG', now_iso=now_iso)
        save_conversion(self.conn, direction='input', item_id=self.ids['line_box'], conversion_factor=12, now_iso=now_iso)
        self.conn.commit()
        before = '\n'.join(self.conn.iterdump())
        preview = self.client.get('/api/invoice-workbench/input-receipts' + self.query)
        self.assertEqual(200, preview.status_code)
        self.assertEqual(before, '\n'.join(self.conn.iterdump()))
        review = preview.json
        self.assertEqual([self.ids['input_ids']['2']], [r['id'] for r in review['items']])
        self.assertEqual([self.ids['input_ids']['3']], [r['id'] for r in review['blocked']])
        self.assertEqual(2, review['items'][0]['line_count'])
        post = self.client.post('/api/invoice-workbench/input-receipts', json={'items': review['items']})
        self.assertEqual(200, post.status_code)
        self.assertEqual(1, post.json['posted_count'])
        remaining = self.listing('?from=2026-09-03&to=2026-09-03&scope=pending')
        self.assertEqual({'3'}, {r['invoice_number'] for r in remaining['items']})
        again = self.client.post('/api/invoice-workbench/input-receipts', json={'items': review['items']})
        self.assertEqual(1, again.json['already_posted_count'])
        self.assertEqual(3, self.conn.execute('SELECT COUNT(*) FROM invoice_inventory_ledger').fetchone()[0])

    def test_pending_excludes_posted_source_conflicts_and_expense_only_invoices(self):
        self.conn.execute("UPDATE msmi_invoices SET sync_status='review_required',error_message='Changed source' WHERE id=?", (self.ids['input_ids']['1'],))
        self.conn.execute("UPDATE msmi_invoices SET receipt_status='not_inventory' WHERE id=?", (self.ids['input_ids']['3'],))
        self.conn.commit()
        self.assertEqual({'2'}, {r['invoice_number'] for r in self.listing()['items']})

    def test_excel_uses_same_pending_scope_and_totals(self):
        response = self.client.get('/api/invoice-workbench/invoices/export' + self.query)
        self.assertEqual(200, response.status_code)
        from io import BytesIO
        workbook = load_workbook(BytesIO(response.data))
        sheet = workbook['Hoa don']
        self.assertIn('CÒN CHƯA NHẬP ĐẾN 2026-09-02', sheet['A1'].value)
        numbers = {sheet.cell(i, 2).value for i in range(3, 6)}
        self.assertEqual({'C26TST / 2', 'C26TST / 3'}, numbers)
        workbook.close()
