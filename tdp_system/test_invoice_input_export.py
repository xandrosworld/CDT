from __future__ import annotations

import io
import sqlite3
import tempfile
import unittest
from contextlib import contextmanager
from pathlib import Path

from flask import Flask
from openpyxl import load_workbook

try:
    from .contract_modules import init_contract_schema
    from .invoice_input_export import (
        InvoiceInputExportError,
        input_invoice_export_data,
        input_invoice_workbook,
        register_invoice_input_export_routes,
    )
    from .invoice_workbench import init_invoice_workbench_schema, prepare_sync_batch
except ImportError:  # pragma: no cover
    from contract_modules import init_contract_schema
    from invoice_input_export import (
        InvoiceInputExportError,
        input_invoice_export_data,
        input_invoice_workbook,
        register_invoice_input_export_routes,
    )
    from invoice_workbench import init_invoice_workbench_schema, prepare_sync_batch


NOW = "2026-09-03T10:00:00"


def now_iso() -> str:
    return NOW


def init_database(conn) -> None:
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys=ON")
    conn.executescript(
        """
        CREATE TABLE settings(key TEXT PRIMARY KEY,value TEXT NOT NULL);
        CREATE TABLE products(code TEXT PRIMARY KEY,name TEXT,unit TEXT);
        CREATE TABLE balances(id INTEGER PRIMARY KEY);
        CREATE TABLE orders(id INTEGER PRIMARY KEY);
        """
    )
    init_contract_schema(conn)
    init_invoice_workbench_schema(conn)


def seed_batch(conn, *, invoice_type="input") -> int:
    batch, _ = prepare_sync_batch(
        conn, tenant="TDP", source="msmi", invoice_type=invoice_type,
        date_from="2026-08-01", date_to="2026-08-31", now_iso=now_iso,
    )
    if invoice_type == "input":
        cursor = conn.execute(
            """INSERT INTO msmi_invoices(
                   remote_id,tenant,invoice_type,seller_tax_code,seller_name,invoice_number,
                   invoice_series,invoice_date,subtotal,tax_amount,total_amount,sync_status,
                   receipt_status,raw_json,synced_at,created_at,updated_at
               ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,'{}',?,?,?)""",
            (
                "REMOTE-SECRET-1", "TDP", "INPUT_ELECTRONIC_INVOICE", "0200000001",
                "=NHÀ CUNG CẤP", "0000123", "1C26TDP", "2026-08-15",
                200000, 16000, 216000, "synced", "pending_mapping", NOW, NOW, NOW,
            ),
        )
        invoice_id = cursor.lastrowid
        conn.execute(
            """INSERT INTO msmi_invoice_items(
                   invoice_id,line_index,source_item_code,source_item_name,source_unit,qty,
                   unit_price,amount,tax_rate,source_nature,inventory_eligible,validation_note,
                   product_code,mapping_status,conversion_factor,stock_qty,stock_unit_price
               ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (
                invoice_id, 1, "+SRC", "Cà rốt", "Kg", 10, 20000, 200000, "8%", "1",
                1, "", "", "unmapped", None, 0, 0,
            ),
        )
        conn.execute(
            "INSERT INTO invoice_sync_batch_invoices(batch_id,invoice_id,linked_at) VALUES(?,?,?)",
            (batch["id"], invoice_id, NOW),
        )
        conn.execute(
            """UPDATE invoice_sync_batches SET status='needs_mapping',fetched_count=1,
                      needs_mapping_count=1,updated_at=? WHERE id=?""",
            (NOW, batch["id"]),
        )
    return int(batch["id"])


class InvoiceInputExportTests(unittest.TestCase):
    def setUp(self):
        self.conn = sqlite3.connect(":memory:")
        init_database(self.conn)
        self.batch_id = seed_batch(self.conn)

    def tearDown(self):
        self.conn.close()

    def test_export_is_available_before_mapping_or_inventory_post(self):
        data = input_invoice_export_data(self.conn, self.batch_id)
        self.assertEqual(1, data["totals"]["invoice_count"])
        self.assertEqual(216000, data["totals"]["total_amount"])
        self.assertEqual("pending_mapping", data["invoices"][0]["receipt_status"])

        workbook = input_invoice_workbook(data)
        try:
            stream = io.BytesIO()
            workbook.save(stream)
        finally:
            workbook.close()
        stream.seek(0)
        loaded = load_workbook(stream, data_only=False)
        try:
            self.assertEqual(
                ["Hóa đơn đầu vào", "Chi tiết hàng hóa", "Thông tin lần tải"],
                loaded.sheetnames,
            )
            summary = loaded["Hóa đơn đầu vào"]
            self.assertEqual("DANH SÁCH HÓA ĐƠN ĐẦU VÀO ĐÃ TẢI", summary["A1"].value)
            self.assertEqual("'=NHÀ CUNG CẤP", summary["F4"].value)
            self.assertEqual(216000, summary["I4"].value)
            detail = loaded["Chi tiết hàng hóa"]
            self.assertEqual("'+SRC", detail["H4"].value)
            self.assertEqual("Hàng hóa, dịch vụ", detail["O4"].value)
            self.assertFalse(any(
                cell.data_type == "f" for sheet in loaded.worksheets
                for row in sheet.iter_rows() for cell in row
            ))
            all_text = "\n".join(
                str(cell.value or "") for sheet in loaded.worksheets
                for row in sheet.iter_rows() for cell in row
            )
            self.assertNotIn("REMOTE-SECRET-1", all_text)
            self.assertNotIn("raw_json", all_text.casefold())
        finally:
            loaded.close()

    def test_output_batch_and_empty_input_batch_are_rejected(self):
        output_id = seed_batch(self.conn, invoice_type="output")
        with self.assertRaises(InvoiceInputExportError) as context:
            input_invoice_export_data(self.conn, output_id)
        self.assertEqual("invoice_input_batch_required", context.exception.code)

        empty, _ = prepare_sync_batch(
            self.conn, tenant="TDP", source="msmi", invoice_type="input",
            date_from="2026-09-01", date_to="2026-09-30", now_iso=now_iso,
        )
        with self.assertRaises(InvoiceInputExportError) as context:
            input_invoice_export_data(self.conn, empty["id"])
        self.assertEqual("invoice_input_batch_empty", context.exception.code)


class InvoiceInputExportRouteTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="tdp_input_export_")
        self.path = Path(self.temp.name) / "test.sqlite3"
        with self.db() as conn:
            self.batch_id = seed_batch(conn)
        app = Flask(__name__)
        app.config["TESTING"] = True
        register_invoice_input_export_routes(app, {"db": self.db, "now_iso": now_iso})
        self.client = app.test_client()

    def tearDown(self):
        self.temp.cleanup()

    @contextmanager
    def db(self):
        conn = sqlite3.connect(self.path)
        init_database(conn) if not self.path.exists() or self.path.stat().st_size == 0 else None
        conn.row_factory = sqlite3.Row
        conn.execute("PRAGMA foreign_keys=ON")
        try:
            yield conn
            conn.commit()
        except Exception:
            conn.rollback()
            raise
        finally:
            conn.close()

    def test_route_downloads_xlsx_and_records_no_inventory_effect(self):
        response = self.client.get(
            f"/api/invoice-workbench/batches/{self.batch_id}/export-input-xlsx"
        )
        self.assertEqual(200, response.status_code)
        self.assertEqual("none", response.headers["X-TDP-Inventory-Effect"])
        self.assertIn(".xlsx", response.headers["Content-Disposition"])
        loaded = load_workbook(io.BytesIO(response.data), read_only=True)
        loaded.close()
        with self.db() as conn:
            self.assertEqual(0, conn.execute(
                "SELECT COUNT(*) FROM invoice_inventory_ledger"
            ).fetchone()[0])
            audit = conn.execute(
                "SELECT metadata_json FROM audit_log WHERE event_type='invoice_input.export'"
            ).fetchone()
            self.assertIsNotNone(audit)
            self.assertNotIn("REMOTE-SECRET", audit[0])


if __name__ == "__main__":
    unittest.main()
