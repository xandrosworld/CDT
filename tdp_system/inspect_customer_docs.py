from __future__ import annotations

import json
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "bosung.30.8.26" / "đề nghị thanh toán"
OUTPUT = Path(__file__).resolve().parent / "data" / "customer_doc_inventory.json"


def clean(value: str) -> str:
    return " ".join((value or "").split())


def inspect(path: Path) -> dict:
    document = Document(path)
    paragraphs = [clean(item.text) for item in document.paragraphs if clean(item.text)]
    tables = []
    for table in document.tables:
        rows = []
        for row in table.rows:
            values = [clean(cell.text) for cell in row.cells]
            if any(values):
                rows.append(values)
        tables.append(rows)
    return {
        "file": path.name,
        "paragraph_count": len(paragraphs),
        "paragraphs": paragraphs,
        "table_count": len(tables),
        "tables": tables,
        "sections": len(document.sections),
    }


def main():
    result = [inspect(path) for path in sorted(SOURCE.glob("*.docx"))]
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"ok": True, "files": len(result), "output": str(OUTPUT)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
