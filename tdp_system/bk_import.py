"""Controlled BK input for goods bought without an electronic invoice.

BK is an independent Excel source supplied by this system. It is never pulled
from mSMI and is never posted merely because an order contains a ``bk`` flag.
Only a validated preview plus an explicit confirmation appends movements to
the canonical inventory ledger used by readiness and TĐK–NXT.
"""

from __future__ import annotations

import hashlib
import io
import json
import math
import re
import sqlite3
import threading
import time
import unicodedata
import uuid
import zipfile
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any, Iterable, Mapping

from flask import jsonify, request, send_file
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.page import PageMargins

try:
    from template_workbook import safe_workbook_bytes
except ImportError:  # pragma: no cover - package import
    from .template_workbook import safe_workbook_bytes


BK_IMPORT_SHEET = "BK_IMPORT"
BK_IMPORT_HEADER_ROW = 3
BK_IMPORT_MAX_BYTES = 10 * 1024 * 1024
BK_IMPORT_MAX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
BK_IMPORT_MAX_ENTRIES = 2_000
BK_IMPORT_MAX_ROWS = 20_000
BK_IMPORT_TTL_SECONDS = 30 * 60
BK_IMPORT_SOURCE_TYPE = "BẢNG KÊ MUA VÀO KHÔNG HÓA ĐƠN"
BK_LEDGER_SOURCE_TABLE = "bk_import_documents"

# Compatibility only. The question is resolved and is not a runtime blocker.
BK_IMPORT_POLICY_QUESTION = "Q-004-RESOLVED"

BK_IMPORT_COLUMNS = (
    ("document_date", "Ngày chứng từ"),
    ("source_type", "Loại nguồn (cố định)"),
    ("source_reference", "Số tham chiếu"),
    ("source_line", "Dòng nguồn"),
    ("product_code", "Mã hàng TĐP"),
    ("product_name", "Tên hàng"),
    ("unit", "ĐVT"),
    ("qty", "Số lượng"),
    ("unit_cost", "Đơn giá vốn"),
    ("amount", "Thành tiền"),
    ("source_party", "Mã/NCC nguồn"),
    ("note", "Ghi chú"),
)

BK_IMPORT_SCHEMA = f"""
CREATE TABLE IF NOT EXISTS bk_import_documents (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    content_hash TEXT NOT NULL UNIQUE,
    source_hash TEXT NOT NULL,
    filename TEXT NOT NULL DEFAULT '',
    source_kind TEXT NOT NULL DEFAULT '{BK_IMPORT_SOURCE_TYPE}',
    status TEXT NOT NULL DEFAULT 'posted',
    row_count INTEGER NOT NULL,
    qty_total REAL NOT NULL,
    amount_total REAL NOT NULL,
    confirmed_at TEXT NOT NULL,
    reversed_at TEXT,
    reversal_date TEXT,
    reversal_reason TEXT NOT NULL DEFAULT '',
    created_at TEXT NOT NULL,
    updated_at TEXT NOT NULL,
    CHECK(source_kind='{BK_IMPORT_SOURCE_TYPE}'),
    CHECK(status IN ('posted','reversed')),
    CHECK(row_count > 0),
    CHECK(qty_total > 0),
    CHECK(amount_total > 0)
);
CREATE TABLE IF NOT EXISTS bk_import_lines (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    document_id INTEGER NOT NULL REFERENCES bk_import_documents(id) ON DELETE RESTRICT,
    source_key TEXT NOT NULL UNIQUE,
    snapshot_hash TEXT NOT NULL,
    source_row INTEGER NOT NULL,
    document_date TEXT NOT NULL,
    source_type TEXT NOT NULL DEFAULT '{BK_IMPORT_SOURCE_TYPE}',
    source_reference TEXT NOT NULL,
    source_line INTEGER NOT NULL,
    product_code TEXT NOT NULL REFERENCES products(code) ON DELETE RESTRICT,
    product_name_snapshot TEXT NOT NULL,
    unit_snapshot TEXT NOT NULL,
    qty REAL NOT NULL,
    unit_cost REAL NOT NULL,
    amount REAL NOT NULL,
    source_party TEXT NOT NULL,
    note TEXT NOT NULL DEFAULT '',
    created_at TEXT NOT NULL,
    CHECK(source_row > 0),
    CHECK(source_type='{BK_IMPORT_SOURCE_TYPE}'),
    CHECK(source_line > 0),
    CHECK(qty > 0),
    CHECK(unit_cost > 0),
    CHECK(amount > 0)
);
CREATE INDEX IF NOT EXISTS idx_bk_import_lines_document
    ON bk_import_lines(document_id,source_line,id);
CREATE INDEX IF NOT EXISTS idx_bk_import_lines_product_date
    ON bk_import_lines(product_code,document_date,id);
"""

_PENDING: dict[str, dict[str, Any]] = {}
_PENDING_LOCK = threading.Lock()


class BKImportError(ValueError):
    def __init__(self, message: str, *, code: str = "invalid_bk_workbook", status: int = 400):
        super().__init__(message)
        self.code = code
        self.status = status


def init_bk_import_schema(conn) -> None:
    conn.executescript(BK_IMPORT_SCHEMA)


def _plain(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip())


def _key(value: Any) -> str:
    text = str(value or "").replace("đ", "d").replace("Đ", "D")
    text = unicodedata.normalize("NFKD", text).casefold()
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"[^a-z0-9]+", "", text)


def _hash_json(value: Any) -> str:
    payload = json.dumps(
        value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str,
    ).encode("utf-8")
    return hashlib.sha256(payload).hexdigest().upper()


def _formula(cell) -> bool:
    return cell.data_type == "f" or (
        isinstance(cell.value, str) and cell.value.lstrip().startswith("=")
    )


def _contains_identity_number(value: Any) -> bool:
    return bool(re.search(r"(?<!\d)(?:\d{9}|\d{12})(?!\d)", str(value or "")))


def _iso_date(value: Any, label: str = "Ngày chứng từ") -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = _plain(value)
    for pattern in ("%Y-%m-%d", "%d/%m/%Y", "%d.%m.%Y"):
        try:
            return datetime.strptime(text, pattern).date().isoformat()
        except ValueError:
            continue
    raise ValueError(f"{label} phải có dạng YYYY-MM-DD hoặc DD/MM/YYYY")


def _excel_date(value: Any) -> date:
    return date.fromisoformat(_iso_date(value))


def _decimal(value: Any, label: str) -> Decimal:
    if isinstance(value, bool) or value in (None, ""):
        raise ValueError(f"{label} không được để trống")
    text = str(value).strip().replace(" ", "")
    if "," in text and "." not in text:
        text = text.replace(",", ".")
    try:
        number = Decimal(text)
    except (InvalidOperation, ValueError):
        raise ValueError(f"{label} phải là số") from None
    if not number.is_finite() or not math.isfinite(float(number)):
        raise ValueError(f"{label} phải là số hữu hạn")
    return number


def _positive_line(value: Any) -> int:
    if isinstance(value, bool) or value in (None, ""):
        raise ValueError("Dòng nguồn không được để trống")
    try:
        number = Decimal(str(value).strip())
    except InvalidOperation:
        raise ValueError("Dòng nguồn phải là số nguyên dương") from None
    if not number.is_finite() or number != number.to_integral_value() or number <= 0:
        raise ValueError("Dòng nguồn phải là số nguyên dương")
    return int(number)


def _number(value: Decimal) -> int | float:
    return int(value) if value == value.to_integral_value() else float(value)


def _purchase_rate(conn) -> Decimal:
    row = conn.execute("SELECT value FROM settings WHERE key='purchase_rate'").fetchone()
    try:
        rate = Decimal(str(row["value"] if row else "0.95"))
    except InvalidOperation:
        rate = Decimal("0.95")
    if not rate.is_finite() or rate <= 0:
        raise BKImportError(
            "Tỷ lệ giá bảng kê phải lớn hơn 0", code="invalid_purchase_rate", status=409,
        )
    return rate


def _template_rows_for_batch(conn, batch_id: int) -> list[dict[str, Any]]:
    """Prefill only confirmed BK candidates from one approved order batch."""
    batch = conn.execute("SELECT * FROM batches WHERE id=?", (batch_id,)).fetchone()
    if not batch:
        raise BKImportError("Không tìm thấy phiên đơn để tạo mẫu BK", code="batch_not_found", status=404)
    if str(batch["status"]) != "approved":
        raise BKImportError(
            "Chỉ tạo mẫu BK tự điền từ phiên đơn đã duyệt", code="batch_not_approved", status=409,
        )
    if conn.execute("SELECT 1 FROM sqlite_master WHERE name='batch_bk_approvals'").fetchone():
        posted = conn.execute("""SELECT l.* FROM batch_bk_approvals a
            JOIN bk_import_documents d ON d.id=a.document_id AND d.status='posted'
            JOIN bk_import_lines l ON l.document_id=d.id WHERE a.batch_id=? ORDER BY l.id""", (batch_id,)).fetchall()
        if posted:
            return [{
                'document_date': r['document_date'], 'source_type': r['source_type'],
                'source_reference': r['source_reference'], 'source_line': r['source_line'],
                'product_code': r['product_code'], 'product_name': r['product_name_snapshot'],
                'unit': r['unit_snapshot'], 'qty': r['qty'], 'unit_cost': r['unit_cost'],
                'amount': r['amount'], 'source_party': r['source_party'], 'note': r['note'],
            } for r in posted]
    rate = _purchase_rate(conn)
    canonical = conn.execute(
        "SELECT 1 FROM purchase_workbook_lines WHERE batch_id=? LIMIT 1", (batch_id,),
    ).fetchone()
    if canonical:
        source_rows = conn.execute(
            """SELECT l.source_row,l.id,l.work_date,l.product_code,l.product_name,l.unit,
                      l.actual_qty qty,l.supplier source_party,l.note,
                      COALESCE(o.purchase_list,p.purchase_list,0) purchase_list,l.status,
                      COALESCE(o.sell_price,0) sell_price
                 FROM purchase_workbook_lines l
                 LEFT JOIN orders o ON o.id=l.order_id
                 LEFT JOIN products p ON p.code=l.product_code
                WHERE l.batch_id=? ORDER BY l.source_row,l.id""",
            (batch_id,),
        ).fetchall()
        rows = [dict(row) for row in source_rows if (
            int(row["purchase_list"] or 0) == 1
            and str(row["status"] or "").casefold() == "confirmed"
            and float(row["qty"] or 0) > 0
            and _key(row["source_party"]) != "kho"
        )]
    else:
        source_rows = conn.execute(
            """SELECT COALESCE(source_row,id) source_row,id,work_date,product_code,
                      product_name,unit,
                      MAX(COALESCE(actual_received,0)-COALESCE(damaged_qty,0)
                          -COALESCE(supplier_return_qty,0),0) qty,
                      sell_price,supplier source_party,note,purchase_list
                 FROM orders WHERE batch_id=? ORDER BY COALESCE(source_row,id),id""",
            (batch_id,),
        ).fetchall()
        rows = [dict(row) for row in source_rows if (
            int(row["purchase_list"] or 0) == 1
            and float(row["qty"] or 0) > 0
            and _key(row["source_party"]) != "kho"
        )]

    result = []
    for index, row in enumerate(rows, start=1):
        qty = Decimal(str(row["qty"] or 0))
        sell_price = Decimal(str(row["sell_price"] or 0))
        if not sell_price.is_finite() or sell_price <= 0:
            raise BKImportError(
                f"Dòng BK nguồn {row['source_row']}: thiếu giá bán dương để tính giá BK",
                code="bk_sales_price_missing", status=409,
            )
        # Customer rule: the prefilled BK cost is exactly the configured rate
        # (95% by default) of the row's sales price. Never substitute buy cost.
        unit_cost = (sell_price * rate).quantize(Decimal("1"), rounding=ROUND_HALF_UP)
        amount = (qty * unit_cost).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        result.append({
            "document_date": _iso_date(row["work_date"]),
            "source_type": BK_IMPORT_SOURCE_TYPE,
            "source_reference": f"TDP-BATCH-{batch_id}",
            # Database row id is unique and stable even when Excel source rows repeat.
            "source_line": int(row["id"] or index),
            "product_code": _plain(row["product_code"]).upper(),
            "product_name": _plain(row["product_name"]),
            "unit": _plain(row["unit"]),
            "qty": _number(qty),
            "unit_cost": _number(unit_cost),
            "amount": _number(amount),
            "source_party": _plain(row["source_party"]),
            "note": (
                f"Giá BK mặc định {rate * 100}% giá bán; dòng nguồn {row['source_row']}. "
                + _plain(row["note"])
            ).strip(),
        })
    if not result:
        raise BKImportError(
            "Phiên đơn không có dòng BK đã chốt, có số lượng thực nhận dương",
            code="batch_has_no_bk_rows", status=409,
        )
    return result


def build_bk_import_template(rows: Iterable[Mapping[str, Any]] = ()) -> bytes:
    """Return the official formula-free BK input template, optionally prefilled."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = BK_IMPORT_SHEET
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(BK_IMPORT_COLUMNS))
    sheet.cell(1, 1).value = "MẪU NHẬP BK – HÀNG MUA VÀO KHÔNG CÓ HÓA ĐƠN"
    sheet.cell(1, 1).font = Font(name="Times New Roman", bold=True, size=14)
    sheet.cell(1, 1).alignment = Alignment(horizontal="center")
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(BK_IMPORT_COLUMNS))
    sheet.cell(2, 1).value = (
        f"Nguồn cố định: “{BK_IMPORT_SOURCE_TYPE}”. Mẫu theo phiên tự điền giá BK bằng tỷ lệ "
        "giá bán đã cấu hình (mặc định 95%); hãy xem/sửa trước khi tải lại. "
        "Không nhập CCCD/CMND và không dùng dữ liệu mSMI trong file này."
    )
    sheet.cell(2, 1).alignment = Alignment(wrap_text=True)
    for column, (_field, label) in enumerate(BK_IMPORT_COLUMNS, start=1):
        cell = sheet.cell(BK_IMPORT_HEADER_ROW, column)
        cell.value = label
        cell.font = Font(name="Times New Roman", size=11, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=4, max_col=len(BK_IMPORT_COLUMNS)):
        for cell in row:
            cell.font = Font(name="Times New Roman", size=11)
    for item in rows:
        sheet.append([
            _excel_date(item.get(field)) if field == "document_date" else item.get(field, "")
            for field, _label in BK_IMPORT_COLUMNS
        ])
    widths = (15, 34, 20, 12, 17, 28, 11, 14, 17, 17, 23, 42)
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[sheet.cell(3, column).column_letter].width = width
    for row_no in range(4, sheet.max_row + 1):
        sheet.cell(row_no, 1).number_format = "dd/mm/yyyy"
        for column in (8, 9, 10):
            sheet.cell(row_no, column).number_format = "#,##0.######"
    sheet.freeze_panes = "A4"
    sheet.auto_filter.ref = f"A3:L{max(sheet.max_row, 3)}"
    sheet.row_dimensions[1].height = 25
    sheet.row_dimensions[2].height = 48
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.orientation = sheet.ORIENTATION_LANDSCAPE
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_margins = PageMargins(
        left=0.2, right=0.2, top=0.35, bottom=0.35, header=0.15, footer=0.15,
    )
    sheet.print_area = f"$A$1:$L${max(sheet.max_row, BK_IMPORT_HEADER_ROW)}"
    sheet.print_title_rows = "$1:$3"
    sheet.print_options.horizontalCentered = True
    workbook.properties.title = "Mẫu nhập BK hàng mua vào không có hóa đơn"
    workbook.properties.subject = "Nguồn Excel BK độc lập; preview và xác nhận trước khi ghi kho"
    try:
        return safe_workbook_bytes(workbook)
    finally:
        workbook.close()


build_bk_draft_template = build_bk_import_template


def _validate_archive(payload: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            entries = archive.infolist()
            if len(entries) > BK_IMPORT_MAX_ENTRIES:
                raise BKImportError(
                    "Workbook có quá nhiều thành phần", code="expanded_too_large", status=413,
                )
            if sum(item.file_size for item in entries) > BK_IMPORT_MAX_UNCOMPRESSED_BYTES:
                raise BKImportError(
                    "Workbook có cấu trúc giải nén quá lớn", code="expanded_too_large", status=413,
                )
            names = {item.filename.casefold() for item in entries}
            if any(name.startswith("xl/externallinks/") for name in names):
                raise BKImportError(
                    "Workbook BK không được chứa liên kết ngoài", code="external_link_forbidden",
                )
            if "xl/vbaproject.bin" in names:
                raise BKImportError("Workbook BK không được chứa macro", code="macro_forbidden")
    except zipfile.BadZipFile:
        raise BKImportError("Workbook bị hỏng hoặc sai định dạng", code="invalid_workbook") from None


def _existing_source_lines(conn) -> dict[str, dict[str, Any]]:
    return {
        str(row["source_key"]): dict(row)
        for row in conn.execute(
            """SELECT l.source_key,l.snapshot_hash,l.document_id,d.status,d.content_hash
                 FROM bk_import_lines l JOIN bk_import_documents d ON d.id=l.document_id"""
        )
    }


def _database_state_hash(conn, rows: Iterable[Mapping[str, Any]]) -> str:
    items = list(rows)
    codes = sorted({str(item["productCode"]) for item in items if item.get("productCode")})
    products = []
    if codes:
        placeholders = ",".join("?" for _ in codes)
        products = [dict(row) for row in conn.execute(
            f"SELECT code,name,COALESCE(unit,'') unit FROM products WHERE code IN ({placeholders}) ORDER BY code",
            codes,
        )]
    existing = _existing_source_lines(conn)
    source_state = sorted([
        {"source_key": item["sourceKey"], "existing": existing.get(item["sourceKey"])}
        for item in items if item.get("sourceKey")
    ], key=lambda item: item["source_key"])
    return _hash_json({"products": products, "sources": source_state})


def parse_bk_preview(conn, payload: bytes) -> dict[str, Any]:
    _validate_archive(payload)
    workbook = None
    try:
        workbook = load_workbook(
            io.BytesIO(payload), data_only=False, read_only=False, keep_links=True,
        )
        if workbook.sheetnames != [BK_IMPORT_SHEET]:
            raise BKImportError(
                f"File BK phải có đúng một sheet {BK_IMPORT_SHEET}",
                code="invalid_sheet_contract",
            )
        sheet = workbook[BK_IMPORT_SHEET]
        if sheet.sheet_state != "visible":
            raise BKImportError("Sheet BK_IMPORT phải hiển thị", code="hidden_import_sheet")
        if sheet.max_row > BK_IMPORT_MAX_ROWS + BK_IMPORT_HEADER_ROW:
            raise BKImportError(
                f"Sheet BK_IMPORT vượt giới hạn {BK_IMPORT_MAX_ROWS:,} dòng",
                code="too_many_rows", status=413,
            )
        actual_headers = tuple(
            _key(sheet.cell(BK_IMPORT_HEADER_ROW, column).value)
            for column in range(1, len(BK_IMPORT_COLUMNS) + 1)
        )
        expected_headers = tuple(_key(label) for _field, label in BK_IMPORT_COLUMNS)
        if actual_headers != expected_headers:
            raise BKImportError(
                "Tiêu đề BK_IMPORT không đúng mẫu; hãy tải lại file mẫu từ hệ thống",
                code="invalid_headers",
            )

        known_products = {
            _plain(row["code"]).upper(): {
                "name": _plain(row["name"]), "unit": _plain(row["unit"]),
            }
            for row in conn.execute("SELECT code,name,unit FROM products ORDER BY code")
        }
        rows: list[dict[str, Any]] = []
        for row_no in range(BK_IMPORT_HEADER_ROW + 1, sheet.max_row + 1):
            cells = [sheet.cell(row_no, column) for column in range(1, len(BK_IMPORT_COLUMNS) + 1)]
            if not any(cell.value not in (None, "") for cell in cells):
                continue
            raw = {
                field: cells[index].value
                for index, (field, _label) in enumerate(BK_IMPORT_COLUMNS)
            }
            errors: list[str] = []
            warnings: list[str] = []
            formula_fields = [
                BK_IMPORT_COLUMNS[index][1] for index, cell in enumerate(cells) if _formula(cell)
            ]
            if formula_fields:
                errors.append("Không nhận công thức tại: " + ", ".join(formula_fields))
            if any(cell.hyperlink for cell in cells):
                errors.append("Không nhận hyperlink trong dòng dữ liệu")

            try:
                document_date = _iso_date(raw["document_date"])
            except ValueError as exc:
                document_date = ""
                errors.append(str(exc))
            supplied_source_type = _plain(raw["source_type"])
            source_type = BK_IMPORT_SOURCE_TYPE
            if supplied_source_type and _key(supplied_source_type) != _key(BK_IMPORT_SOURCE_TYPE):
                errors.append(
                    f"Loại nguồn được cố định là “{BK_IMPORT_SOURCE_TYPE}”; không dùng mSMI cho BK"
                )
            source_reference = _plain(raw["source_reference"])
            product_code = _plain(raw["product_code"]).upper()
            product_name = _plain(raw["product_name"])
            unit = _plain(raw["unit"])
            source_party = _plain(raw["source_party"])
            note = _plain(raw["note"])
            for label, value in (
                ("Số tham chiếu", source_reference),
                ("Mã hàng TĐP", product_code),
                ("Mã/NCC nguồn", source_party),
            ):
                if not value:
                    errors.append(f"{label} không được để trống")
                elif len(value) > 255:
                    errors.append(f"{label} vượt 255 ký tự")
            if len(product_name) > 255 or len(unit) > 100:
                errors.append("Tên hàng hoặc ĐVT vượt giới hạn cho phép")
            if len(note) > 500:
                errors.append("Ghi chú vượt 500 ký tự")
            if _contains_identity_number(source_party) or _contains_identity_number(note):
                errors.append("Không nhập CCCD/CMND vào file BK")
                if _contains_identity_number(source_party):
                    source_party = "[ĐÃ ẨN SỐ ĐỊNH DANH]"
                if _contains_identity_number(note):
                    note = "[ĐÃ ẨN SỐ ĐỊNH DANH]"
            try:
                source_line = _positive_line(raw["source_line"])
            except ValueError as exc:
                source_line = 0
                errors.append(str(exc))
            try:
                qty_decimal = _decimal(raw["qty"], "Số lượng")
                if qty_decimal <= 0:
                    raise ValueError("Số lượng phải lớn hơn 0")
            except ValueError as exc:
                qty_decimal = Decimal("0")
                errors.append(str(exc))
            try:
                unit_cost_decimal = _decimal(raw["unit_cost"], "Đơn giá vốn")
                if unit_cost_decimal <= 0:
                    raise ValueError("Đơn giá vốn phải lớn hơn 0")
            except ValueError as exc:
                unit_cost_decimal = Decimal("0")
                errors.append(str(exc))
            try:
                amount_decimal = _decimal(raw["amount"], "Thành tiền")
                if amount_decimal <= 0:
                    raise ValueError("Thành tiền phải lớn hơn 0")
            except ValueError as exc:
                amount_decimal = Decimal("0")
                errors.append(str(exc))
            expected_amount = (qty_decimal * unit_cost_decimal).quantize(
                Decimal("0.01"), rounding=ROUND_HALF_UP,
            )
            if (
                qty_decimal > 0 and unit_cost_decimal > 0 and amount_decimal > 0
                and abs(amount_decimal - expected_amount) > Decimal("0.01")
            ):
                errors.append("Thành tiền không khớp Số lượng × Đơn giá vốn")

            product = known_products.get(product_code)
            if product_code and product is None:
                errors.append("Mã hàng TĐP chưa có trong danh mục")
            elif product is not None:
                if not product_name:
                    product_name = product["name"]
                    warnings.append("Tên hàng trống; hệ thống dùng tên danh mục")
                elif _key(product_name) != _key(product["name"]):
                    errors.append("Tên hàng không khớp mã hàng trong danh mục TĐP")
                if not product["unit"]:
                    errors.append("Mã hàng trong danh mục TĐP chưa có ĐVT")
                elif not unit:
                    unit = product["unit"]
                    warnings.append("ĐVT trống; hệ thống dùng ĐVT danh mục")
                elif _key(unit) != _key(product["unit"]):
                    errors.append("ĐVT không khớp danh mục TĐP; không được tự quy đổi")

            identity = {
                "document_date": document_date,
                "source_type": _key(source_type),
                "source_reference": source_reference.casefold(),
                "source_line": source_line,
            }
            source_key = _hash_json(identity) if all(identity.values()) else ""
            snapshot = {
                **identity,
                "product_code": product_code, "product_name": product_name, "unit": unit,
                "qty": str(qty_decimal), "unit_cost": str(unit_cost_decimal),
                "amount": str(amount_decimal), "source_party": source_party, "note": note,
            }
            snapshot_hash = _hash_json(snapshot) if source_key else ""
            rows.append({
                "sourceRow": row_no, "sourceKey": source_key, "snapshotHash": snapshot_hash,
                "documentDate": document_date, "sourceType": source_type,
                "sourceReference": source_reference, "sourceLine": source_line,
                "productCode": product_code, "productName": product_name, "unit": unit,
                "qty": float(qty_decimal), "unitCost": float(unit_cost_decimal),
                "amount": float(amount_decimal), "sourceParty": source_party, "note": note,
                "errors": errors, "warnings": warnings,
            })

        if not rows:
            raise BKImportError("Sheet BK_IMPORT chưa có dòng dữ liệu", code="bk_workbook_empty")
        duplicates: dict[str, list[int]] = {}
        for item in rows:
            if item["sourceKey"]:
                duplicates.setdefault(item["sourceKey"], []).append(item["sourceRow"])
        duplicate_groups = {
            key: source_rows for key, source_rows in duplicates.items() if len(source_rows) > 1
        }
        for item in rows:
            if item["sourceKey"] in duplicate_groups:
                item["errors"].append(
                    "Trùng ngày/số tham chiếu/dòng nguồn với dòng "
                    + ", ".join(str(value) for value in duplicate_groups[item["sourceKey"]])
                )

        normalized = sorted([
            {key: value for key, value in item.items()
             if key not in {"errors", "warnings", "sourceRow"}}
            for item in rows
        ], key=lambda item: item["sourceKey"])
        content_hash = _hash_json(normalized)
        existing_document = conn.execute(
            "SELECT id,status FROM bk_import_documents WHERE content_hash=?", (content_hash,),
        ).fetchone()
        if not (existing_document and str(existing_document["status"]) == "posted"):
            earliest_date = min(
                (item["documentDate"] for item in rows if item["documentDate"]),
                default="",
            )
            newer_opening = earliest_date and conn.execute(
                """SELECT 1 FROM inventory_transactions
                    WHERE source_type='OPENING' AND status='posted' AND txn_date>? LIMIT 1""",
                (earliest_date,),
            ).fetchone()
            if newer_opening:
                for item in rows:
                    item["errors"].append(
                        "Ngày BK nằm trước một kỳ tồn đầu đã chốt; cần rebuild kỳ sau"
                    )
        existing_lines = _existing_source_lines(conn)
        for item in rows:
            existing = existing_lines.get(item["sourceKey"])
            if not existing:
                continue
            if existing_document and int(existing["document_id"]) == int(existing_document["id"]):
                if str(existing_document["status"]) == "posted":
                    item["warnings"].append("Dòng đã nhập; xác nhận lại sẽ không cộng lặp")
                else:
                    item["errors"].append(
                        "Bộ BK đã hoàn tác; cần số tham chiếu mới nếu phát sinh mua lại"
                    )
            else:
                item["errors"].append(
                    "Ngày/số tham chiếu/dòng nguồn đã thuộc một bộ BK khác; không được ghi đè"
                )

        counts = {
            "sourceRows": len(rows),
            "readyRows": sum(not item["errors"] for item in rows),
            "errorRows": sum(bool(item["errors"]) for item in rows),
            "warningRows": sum(bool(item["warnings"]) for item in rows),
            "duplicateGroups": len(duplicate_groups),
        }
        return {
            "sheet": sheet.title, "rows": rows, "contentHash": content_hash,
            "alreadyPosted": bool(existing_document and existing_document["status"] == "posted"),
            "existingDocumentId": int(existing_document["id"]) if existing_document else None,
            "counts": counts,
            "totals": {
                "qty": float(sum(Decimal(str(item["qty"])) for item in rows)),
                "amount": float(sum(Decimal(str(item["amount"])) for item in rows)),
            },
            "canConfirm": counts["errorRows"] == 0 and counts["readyRows"] > 0,
        }
    except BKImportError:
        raise
    except (OSError, ValueError, KeyError, zipfile.BadZipFile):
        raise BKImportError("Không đọc được workbook BK", code="invalid_workbook") from None
    finally:
        if workbook is not None:
            workbook.close()


def _confirmation(conn, *, document_id: int, action: str, note: str, timestamp: str) -> int:
    confirmation_key = hashlib.sha256(
        f"input\0{action}\0{BK_LEDGER_SOURCE_TABLE}\0{document_id}".encode("utf-8")
    ).hexdigest()
    conn.execute(
        """INSERT OR IGNORE INTO invoice_inventory_confirmations(
               confirmation_key,direction,source_invoice_table,source_invoice_id,
               action,confirmed,note,created_at
           ) VALUES(?,'input',?,?,?,1,?,?)""",
        (confirmation_key, BK_LEDGER_SOURCE_TABLE, document_id, action, note, timestamp),
    )
    row = conn.execute(
        """SELECT id,direction,source_invoice_table,source_invoice_id,action,note
             FROM invoice_inventory_confirmations WHERE confirmation_key=?""",
        (confirmation_key,),
    ).fetchone()
    if not row or (
        str(row["direction"]), str(row["source_invoice_table"]),
        int(row["source_invoice_id"]), str(row["action"])
    ) != ("input", BK_LEDGER_SOURCE_TABLE, document_id, action):
        raise BKImportError(
            "Khóa xác nhận BK trùng nhưng khác nguồn; dữ liệu chưa được ghi",
            code="bk_confirmation_conflict", status=409,
        )
    if action == "reversal" and str(row["note"]) != note:
        raise BKImportError(
            "Lý do hoàn tác không khớp xác nhận đã có", code="bk_reversal_conflict", status=409,
        )
    return int(row["id"])


def _event_key(event_type: str, document_id: int, line_id: int) -> str:
    return hashlib.sha256(
        f"{event_type}\0{BK_LEDGER_SOURCE_TABLE}\0{document_id}\0{line_id}".encode("utf-8")
    ).hexdigest()


def _verify_posted_document(conn, document) -> int:
    lines = conn.execute(
        "SELECT * FROM bk_import_lines WHERE document_id=? ORDER BY source_line,id",
        (int(document["id"]),),
    ).fetchall()
    events = conn.execute(
        """SELECT * FROM invoice_inventory_ledger
            WHERE direction='input' AND event_type='POST'
              AND source_invoice_table=? AND source_invoice_id=? ORDER BY source_line_index,id""",
        (BK_LEDGER_SOURCE_TABLE, int(document["id"])),
    ).fetchall()
    if len(lines) != int(document["row_count"]) or len(events) != len(lines) or not lines:
        raise BKImportError(
            "Bộ BK đã đánh dấu nhập nhưng thiếu dấu vết sổ kho; cần đối chiếu",
            code="bk_post_trace_conflict", status=409,
        )
    line_map = {int(row["id"]): row for row in lines}
    for event in events:
        line = line_map.get(int(event["source_line_id"]))
        if (
            line is None or str(event["product_code"]) != str(line["product_code"])
            or abs(float(event["qty_delta"]) - float(line["qty"])) > 1e-9
            or abs(float(event["unit_cost"]) - float(line["unit_cost"])) > 1e-9
        ):
            raise BKImportError(
                "Dấu vết sổ kho BK khác snapshot nguồn; không được ghi đè",
                code="bk_post_trace_conflict", status=409,
            )
    return len(events)


def _post_pending(conn, pending: Mapping[str, Any], timestamp: str, audit_event) -> dict[str, Any]:
    existing = conn.execute(
        "SELECT * FROM bk_import_documents WHERE content_hash=?", (pending["content_hash"],),
    ).fetchone()
    if existing:
        if str(existing["status"]) != "posted":
            raise BKImportError(
                "Bộ BK này đã hoàn tác; không được nhập lại cùng dấu vết nguồn",
                code="bk_document_reversed", status=409,
            )
        count = _verify_posted_document(conn, existing)
        return {
            "documentId": int(existing["id"]), "inventoryLines": count,
            "newInventoryLines": 0, "idempotent": True, "status": "posted",
        }
    earliest_date = min(str(item["documentDate"]) for item in pending["rows"])
    newer_opening = conn.execute(
        """SELECT 1 FROM inventory_transactions
            WHERE source_type='OPENING' AND status='posted' AND txn_date>? LIMIT 1""",
        (earliest_date,),
    ).fetchone()
    if newer_opening:
        raise BKImportError(
            "BK nằm trước một kỳ tồn đầu đã chốt; cần đối chiếu/rebuild kỳ sau trước khi nhập",
            code="backdated_before_opening_snapshot", status=409,
        )
    if _database_state_hash(conn, pending["rows"]) != pending["database_state_hash"]:
        raise BKImportError(
            "Danh mục hoặc nguồn BK đã thay đổi sau khi xem trước; hãy chọn lại file",
            code="bk_preview_stale", status=409,
        )

    total_qty = sum(Decimal(str(item["qty"])) for item in pending["rows"])
    total_amount = sum(Decimal(str(item["amount"])) for item in pending["rows"])
    cursor = conn.execute(
        """INSERT INTO bk_import_documents(
               content_hash,source_hash,filename,source_kind,status,row_count,qty_total,
               amount_total,confirmed_at,created_at,updated_at
           ) VALUES(?,?,?,?,'posted',?,?,?,?,?,?)""",
        (
            pending["content_hash"], pending["source_hash"], pending["filename"],
            BK_IMPORT_SOURCE_TYPE, len(pending["rows"]), float(total_qty), float(total_amount),
            timestamp, timestamp, timestamp,
        ),
    )
    document_id = int(cursor.lastrowid)
    confirmation_id = _confirmation(
        conn, document_id=document_id, action="post",
        note="Xác nhận nhập BK hàng mua vào không có hóa đơn", timestamp=timestamp,
    )
    created = 0
    for index, item in enumerate(pending["rows"], start=1):
        product = conn.execute(
            "SELECT name,COALESCE(unit,'') unit FROM products WHERE code=?",
            (item["productCode"],),
        ).fetchone()
        if not product or _key(product["name"]) != _key(item["productName"]):
            raise BKImportError(
                f"Mã {item['productCode']} đã đổi tên sau preview; dữ liệu chưa được ghi",
                code="bk_preview_stale", status=409,
            )
        if _key(product["unit"]) != _key(item["unit"]):
            raise BKImportError(
                f"ĐVT mã {item['productCode']} đã thay đổi sau preview; dữ liệu chưa được ghi",
                code="bk_preview_stale", status=409,
            )
        line_cursor = conn.execute(
            """INSERT INTO bk_import_lines(
                   document_id,source_key,snapshot_hash,source_row,document_date,source_type,
                   source_reference,source_line,product_code,product_name_snapshot,unit_snapshot,
                   qty,unit_cost,amount,source_party,note,created_at
               ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (
                document_id, item["sourceKey"], item["snapshotHash"], item["sourceRow"],
                item["documentDate"], BK_IMPORT_SOURCE_TYPE, item["sourceReference"],
                item["sourceLine"], item["productCode"], item["productName"], item["unit"],
                item["qty"], item["unitCost"], item["amount"], item["sourceParty"],
                item["note"], timestamp,
            ),
        )
        line_id = int(line_cursor.lastrowid)
        conn.execute(
            """INSERT INTO invoice_inventory_ledger(
                   event_key,direction,event_type,source_invoice_table,source_invoice_id,
                   source_line_id,source_line_index,product_code,txn_date,qty_delta,unit_cost,
                   mapping_revision_id,confirmation_id,reverses_event_key,status,created_at
               ) VALUES(?,'input','POST',?,?,?,?,?,?,?,?,NULL,?,'','posted',?)""",
            (
                _event_key("POST", document_id, line_id), BK_LEDGER_SOURCE_TABLE, document_id,
                line_id, index, item["productCode"], item["documentDate"], item["qty"],
                item["unitCost"], confirmation_id, timestamp,
            ),
        )
        created += 1
    if audit_event is not None:
        audit_event(
            conn, "bk_import.inventory_post", entity_type="bk_import_document",
            entity_id=str(document_id), metadata={
                "content_hash": pending["content_hash"], "source_hash": pending["source_hash"],
                "rows": created, "qty_total": float(total_qty),
                "amount_total": float(total_amount), "source_kind": BK_IMPORT_SOURCE_TYPE,
            },
        )
    return {
        "documentId": document_id, "inventoryLines": created,
        "newInventoryLines": created, "idempotent": False, "status": "posted",
    }


def _selected_opening(conn, anchor_date: str) -> tuple[str, str] | None:
    rows = conn.execute(
        """SELECT source_id,MIN(txn_date) txn_date,MAX(txn_date) max_date
             FROM inventory_transactions
            WHERE source_type='OPENING' AND status='posted' AND txn_date<=?
            GROUP BY source_id ORDER BY txn_date DESC,source_id DESC""",
        (anchor_date,),
    ).fetchall()
    if not rows:
        return None
    latest_date = str(rows[0]["txn_date"])
    same_date = [row for row in rows if str(row["txn_date"]) == latest_date]
    if len(same_date) != 1 or str(rows[0]["max_date"]) != latest_date:
        raise BKImportError(
            "Kỳ tồn đầu bị xung đột; chưa thể kiểm tra an toàn hoàn tác BK",
            code="opening_period_conflict", status=409,
        )
    return str(rows[0]["source_id"]), latest_date


def _assert_no_negative_stock_after_reversal(
    conn, product_codes: Iterable[str], reversal_date: str,
) -> None:
    """Replay canonical event order after staged reversals, until the next reset."""
    opening = _selected_opening(conn, reversal_date)
    opening_source = opening[0] if opening else ""
    opening_date = opening[1] if opening else "0001-01-01"
    next_row = conn.execute(
        """SELECT MIN(txn_date) next_date FROM inventory_transactions
            WHERE source_type='OPENING' AND status='posted' AND txn_date>?""",
        (reversal_date,),
    ).fetchone()
    next_opening = str(next_row["next_date"] or "") if next_row else ""
    for product_code in sorted(set(product_codes)):
        balance = Decimal("0")
        if opening:
            row = conn.execute(
                """SELECT COALESCE(SUM(qty_in-qty_out),0) qty FROM inventory_transactions
                    WHERE source_type='OPENING' AND status='posted'
                      AND source_id=? AND product_code=?""",
                (opening_source, product_code),
            ).fetchone()
            balance = Decimal(str(row["qty"] or 0))
        params: list[Any] = [product_code, opening_date]
        end_clause = ""
        if next_opening:
            end_clause = " AND txn_date<?"
            params.append(next_opening)
        events = conn.execute(
            f"""SELECT txn_date,direction,event_type,qty_delta,source_invoice_table,
                       source_invoice_id,source_line_index,id
                  FROM invoice_inventory_ledger
                 WHERE status='posted' AND product_code=? AND txn_date>=?{end_clause}
                 ORDER BY txn_date,
                   CASE
                     WHEN direction='input' AND event_type='POST' THEN 10
                     WHEN direction='input' AND event_type='REVERSAL' THEN 15
                     WHEN direction='output' AND event_type='POST' THEN 20
                     WHEN direction='output' AND event_type='REVERSAL' THEN 30
                     ELSE 90 END,
                   source_invoice_table,source_invoice_id,source_line_index,id""",
            params,
        ).fetchall()
        for event in events:
            balance += Decimal(str(event["qty_delta"] or 0))
            if balance < Decimal("-0.0000005"):
                raise BKImportError(
                    f"Không thể hoàn tác BK vì mã {product_code} sẽ âm kho tại {event['txn_date']}",
                    code="bk_reversal_negative_stock", status=409,
                )


def _assert_reversal_valuation_safe(conn, reversal_date: str) -> None:
    """Run the canonical valuation through the last event before the next reset."""
    next_row = conn.execute(
        """SELECT MIN(txn_date) next_date FROM inventory_transactions
            WHERE source_type='OPENING' AND status='posted' AND txn_date>?""",
        (reversal_date,),
    ).fetchone()
    next_opening = str(next_row["next_date"] or "") if next_row else ""
    params: list[Any] = [reversal_date]
    end_clause = ""
    if next_opening:
        end_clause = " AND txn_date<?"
        params.append(next_opening)
    row = conn.execute(
        f"""SELECT MAX(txn_date) value FROM invoice_inventory_ledger
             WHERE status='posted' AND txn_date>=?{end_clause}""",
        params,
    ).fetchone()
    report_to = max(reversal_date, str(row["value"] or reversal_date))
    try:
        try:
            from invoice_valuation import InvoiceValuationError, moving_average_report
        except ImportError:  # pragma: no cover - package import
            from .invoice_valuation import InvoiceValuationError, moving_average_report
        moving_average_report(
            conn, date_from=reversal_date, date_to=report_to,
            include_zero=True, include_events=False,
        )
    except InvoiceValuationError as error:
        raise BKImportError(
            f"Không thể hoàn tác BK vì sổ NXT sẽ không hợp lệ: {error}",
            code=f"bk_reversal_{error.code}", status=409,
        ) from None


def _reverse_document(
    conn, document_id: int, *, reversal_date: str, reason: str, timestamp: str, audit_event,
) -> dict[str, Any]:
    document = conn.execute(
        "SELECT * FROM bk_import_documents WHERE id=?", (document_id,),
    ).fetchone()
    if not document:
        raise BKImportError("Không tìm thấy bộ BK", code="bk_document_not_found", status=404)
    original_count = _verify_posted_document(conn, document)
    existing_reversals = int(conn.execute(
        """SELECT COUNT(*) n FROM invoice_inventory_ledger
            WHERE direction='input' AND event_type='REVERSAL'
              AND source_invoice_table=? AND source_invoice_id=?""",
        (BK_LEDGER_SOURCE_TABLE, document_id),
    ).fetchone()["n"])
    if str(document["status"]) == "reversed":
        if existing_reversals != original_count:
            raise BKImportError(
                "Trạng thái hoàn tác BK không khớp sổ kho; cần đối chiếu",
                code="bk_reversal_trace_conflict", status=409,
            )
        if (
            str(document["reversal_date"] or "") != reversal_date
            or str(document["reversal_reason"] or "") != reason
        ):
            raise BKImportError(
                "Bộ BK đã hoàn tác bằng ngày hoặc lý do khác; không được sửa dấu vết cũ",
                code="bk_reversal_conflict", status=409,
            )
        return {
            "documentId": document_id, "reversalLines": existing_reversals,
            "newReversalLines": 0, "idempotent": True, "status": "reversed",
        }
    if str(document["status"]) != "posted" or existing_reversals:
        raise BKImportError(
            "Bộ BK có trạng thái hoàn tác dở dang; cần đối chiếu",
            code="bk_reversal_trace_conflict", status=409,
        )
    latest_source_date = conn.execute(
        "SELECT MAX(document_date) value FROM bk_import_lines WHERE document_id=?", (document_id,),
    ).fetchone()["value"]
    if reversal_date < str(latest_source_date):
        raise BKImportError(
            "Ngày hoàn tác không được trước ngày chứng từ BK", code="bk_reversal_date_invalid", status=400,
        )
    latest_opening = conn.execute(
        """SELECT MAX(txn_date) value FROM inventory_transactions
            WHERE source_type='OPENING' AND status='posted'"""
    ).fetchone()["value"]
    if latest_opening and reversal_date < str(latest_opening):
        raise BKImportError(
            "Ngày hoàn tác BK nằm trước kỳ tồn đầu mới nhất; cần rebuild kỳ tồn trước khi hoàn tác",
            code="bk_reversal_before_opening_snapshot", status=409,
        )

    confirmation_id = _confirmation(
        conn, document_id=document_id, action="reversal", note=reason, timestamp=timestamp,
    )
    originals = conn.execute(
        """SELECT * FROM invoice_inventory_ledger
            WHERE direction='input' AND event_type='POST'
              AND source_invoice_table=? AND source_invoice_id=? ORDER BY source_line_index,id""",
        (BK_LEDGER_SOURCE_TABLE, document_id),
    ).fetchall()
    product_codes = []
    for event in originals:
        conn.execute(
            """INSERT INTO invoice_inventory_ledger(
                   event_key,direction,event_type,source_invoice_table,source_invoice_id,
                   source_line_id,source_line_index,product_code,txn_date,qty_delta,unit_cost,
                   mapping_revision_id,confirmation_id,reverses_event_key,status,created_at
               ) VALUES(?,'input','REVERSAL',?,?,?,?,?,?,?,?,NULL,?,?,'posted',?)""",
            (
                _event_key("REVERSAL", document_id, int(event["source_line_id"])),
                BK_LEDGER_SOURCE_TABLE, document_id, int(event["source_line_id"]),
                int(event["source_line_index"]), str(event["product_code"]), reversal_date,
                -float(event["qty_delta"]), float(event["unit_cost"]), confirmation_id,
                str(event["event_key"]), timestamp,
            ),
        )
        product_codes.append(str(event["product_code"]))
    _assert_no_negative_stock_after_reversal(conn, product_codes, reversal_date)
    _assert_reversal_valuation_safe(conn, reversal_date)
    conn.execute(
        """UPDATE bk_import_documents SET status='reversed',reversed_at=?,reversal_date=?,
                  reversal_reason=?,updated_at=? WHERE id=?""",
        (timestamp, reversal_date, reason, timestamp, document_id),
    )
    if conn.execute("SELECT 1 FROM sqlite_master WHERE name='batch_bk_approvals'").fetchone():
        conn.execute("""UPDATE batches SET status='draft',approved_at=NULL WHERE id IN
            (SELECT batch_id FROM batch_bk_approvals WHERE document_id=?)""", (document_id,))
    if audit_event is not None:
        audit_event(
            conn, "bk_import.inventory_reversal", entity_type="bk_import_document",
            entity_id=str(document_id), metadata={
                "reversal_date": reversal_date, "reason": reason,
                "reversal_lines": len(originals), "source_kind": BK_IMPORT_SOURCE_TYPE,
            },
        )
    return {
        "documentId": document_id, "reversalLines": len(originals),
        "newReversalLines": len(originals), "idempotent": False, "status": "reversed",
    }


def _document_payloads(conn, limit: int = 50) -> list[dict[str, Any]]:
    rows = conn.execute(
        """SELECT d.*,
                  (SELECT COUNT(*) FROM invoice_inventory_ledger l
                    WHERE l.source_invoice_table=? AND l.source_invoice_id=d.id
                      AND l.event_type='POST') posted_lines,
                  (SELECT COUNT(*) FROM invoice_inventory_ledger l
                    WHERE l.source_invoice_table=? AND l.source_invoice_id=d.id
                      AND l.event_type='REVERSAL') reversal_lines
             FROM bk_import_documents d ORDER BY d.id DESC LIMIT ?""",
        (BK_LEDGER_SOURCE_TABLE, BK_LEDGER_SOURCE_TABLE, limit),
    ).fetchall()
    return [{
        "id": int(row["id"]), "filename": row["filename"], "sourceKind": row["source_kind"],
        "status": row["status"], "rowCount": int(row["row_count"]),
        "qtyTotal": float(row["qty_total"]), "amountTotal": float(row["amount_total"]),
        "confirmedAt": row["confirmed_at"], "reversedAt": row["reversed_at"],
        "reversalDate": row["reversal_date"], "reversalReason": row["reversal_reason"],
        "postedLines": int(row["posted_lines"]), "reversalLines": int(row["reversal_lines"]),
    } for row in rows]


def register_bk_import_routes(app, ctx) -> None:
    db_factory = ctx["db"]
    now_iso = ctx.get("now_iso") or (lambda: datetime.now().isoformat(timespec="seconds"))
    audit_event = ctx.get("audit_event")

    def error_response(exc: BKImportError):
        return jsonify({"ok": False, "error": str(exc), "code": exc.code}), exc.status

    @app.get("/api/bk-import/template")
    def api_bk_import_template():
        try:
            batch_text = _plain(request.args.get("batch_id"))
            rows: list[dict[str, Any]] = []
            if batch_text:
                if not batch_text.isdigit() or int(batch_text) <= 0:
                    raise BKImportError("Mã phiên đơn không hợp lệ", code="invalid_batch_id")
                with db_factory() as conn:
                    rows = _template_rows_for_batch(conn, int(batch_text))
            suffix = f"_PHIEN_{batch_text}" if batch_text else ""
            return send_file(
                io.BytesIO(build_bk_import_template(rows)), as_attachment=True,
                download_name=f"MAU_NHAP_BK_HANG_MUA_KHONG_HOA_DON{suffix}.xlsx",
                mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        except BKImportError as exc:
            return error_response(exc)

    @app.post("/api/bk-import/preview")
    def api_bk_import_preview():
        upload = request.files.get("file")
        if not upload or not upload.filename:
            return jsonify({"ok": False, "error": "Chưa chọn file BK", "code": "file_required"}), 400
        if Path(upload.filename).suffix.casefold() != ".xlsx":
            return jsonify({
                "ok": False, "error": "Luồng BK chỉ nhận file .xlsx không macro",
                "code": "invalid_file_type",
            }), 400
        payload = upload.read(BK_IMPORT_MAX_BYTES + 1)
        if len(payload) > BK_IMPORT_MAX_BYTES:
            return jsonify({
                "ok": False, "error": "File BK vượt giới hạn 10 MB", "code": "file_too_large",
            }), 413
        try:
            with db_factory() as conn:
                conn.execute("PRAGMA query_only=ON")
                parsed = parse_bk_preview(conn, payload)
                database_state_hash = _database_state_hash(conn, parsed["rows"])
        except BKImportError as exc:
            return error_response(exc)
        source_hash = hashlib.sha256(payload).hexdigest().upper()
        preview_id = _hash_json({
            "sourceHash": source_hash, "contentHash": parsed["contentHash"],
            "sourceKind": BK_IMPORT_SOURCE_TYPE,
        })
        token = uuid.uuid4().hex
        pending = {
            "created": time.time(), "filename": Path(upload.filename).name[:255],
            "source_hash": source_hash, "content_hash": parsed["contentHash"],
            "preview_id": preview_id, "rows": parsed["rows"],
            "database_state_hash": database_state_hash,
            "has_errors": not parsed["canConfirm"], "result": None,
        }
        cutoff = time.time() - BK_IMPORT_TTL_SECONDS
        with _PENDING_LOCK:
            for old_token, item in list(_PENDING.items()):
                if item["created"] < cutoff:
                    _PENDING.pop(old_token, None)
            _PENDING[token] = pending
        return jsonify({
            "ok": True, "token": token, "filename": pending["filename"],
            "sourceHash": source_hash, "previewId": preview_id,
            "previewOnly": True, "writesInventory": False,
            "sourcePolicy": BK_IMPORT_SOURCE_TYPE,
            "policy": (
                "BK là hàng mua vào không có hóa đơn, dùng file Excel riêng do hệ thống cung cấp; "
                "chỉ cộng vào sổ kho chuẩn sau khi người dùng xác nhận."
            ),
            "rowIdentityFields": ["Ngày chứng từ", "Số tham chiếu", "Dòng nguồn"],
            "expiresInMinutes": BK_IMPORT_TTL_SECONDS // 60,
            **parsed,
        })

    @app.post("/api/bk-import/confirm")
    def api_bk_import_confirm():
        body = request.get_json(silent=True) or {}
        if body.get("confirmed") is not True:
            return jsonify({
                "ok": False, "error": "Cần xác nhận rõ trước khi nhập BK",
                "code": "confirmation_required",
            }), 400
        token = _plain(body.get("token"))
        preview_id = _plain(body.get("previewId"))
        with _PENDING_LOCK:
            pending = _PENDING.get(token)
        if not pending or time.time() - pending["created"] > BK_IMPORT_TTL_SECONDS:
            return jsonify({
                "ok": False, "error": "Phiên xem trước BK đã hết hạn; hãy chọn lại file",
                "code": "bk_preview_expired",
            }), 410
        if preview_id != pending["preview_id"]:
            return jsonify({
                "ok": False, "error": "Mã xem trước BK không khớp",
                "code": "bk_preview_mismatch",
            }), 409
        if pending["has_errors"]:
            return jsonify({
                "ok": False, "error": "File BK còn dòng lỗi nên chưa thể nhập",
                "code": "bk_rows_invalid",
            }), 400
        with _PENDING_LOCK:
            cached_result = pending.get("result")
        if cached_result:
            return jsonify({"ok": True, **cached_result, "idempotent": True})
        try:
            with db_factory() as conn:
                conn.execute("BEGIN IMMEDIATE")
                result = _post_pending(conn, pending, now_iso(), audit_event)
        except BKImportError as exc:
            return error_response(exc)
        except sqlite3.IntegrityError:
            return jsonify({
                "ok": False,
                "error": "Khóa nguồn BK đã tồn tại hoặc dữ liệu thay đổi; chưa ghi thêm vào kho",
                "code": "bk_source_conflict",
            }), 409
        except Exception:
            return jsonify({
                "ok": False, "error": "Không ghi được BK; toàn bộ bút toán đã hoàn tác",
                "code": "bk_write_failed",
            }), 409
        with _PENDING_LOCK:
            if _PENDING.get(token) is pending:
                pending["result"] = dict(result)
        return jsonify({"ok": True, **result})

    @app.get("/api/bk-import/documents")
    def api_bk_import_documents():
        try:
            limit = min(max(int(request.args.get("limit", 50)), 1), 200)
        except ValueError:
            limit = 50
        with db_factory() as conn:
            return jsonify({
                "ok": True, "items": _document_payloads(conn, limit),
                "sourcePolicy": BK_IMPORT_SOURCE_TYPE,
            })

    @app.post("/api/bk-import/documents/<int:document_id>/reversal")
    def api_bk_import_reversal(document_id: int):
        body = request.get_json(silent=True) or {}
        if body.get("confirmed") is not True:
            return jsonify({
                "ok": False, "error": "Cần xác nhận rõ trước khi hoàn tác BK",
                "code": "confirmation_required",
            }), 400
        reason = _plain(body.get("reason"))
        if not reason:
            return jsonify({
                "ok": False, "error": "Phải nhập lý do hoàn tác BK",
                "code": "reversal_reason_required",
            }), 400
        if len(reason) > 500:
            return jsonify({
                "ok": False, "error": "Lý do hoàn tác vượt 500 ký tự",
                "code": "reversal_reason_too_long",
            }), 400
        try:
            reversal_date = _iso_date(body.get("reversalDate"), "Ngày hoàn tác")
            with db_factory() as conn:
                conn.execute("BEGIN IMMEDIATE")
                result = _reverse_document(
                    conn, document_id, reversal_date=reversal_date, reason=reason,
                    timestamp=now_iso(), audit_event=audit_event,
                )
            return jsonify({"ok": True, **result})
        except BKImportError as exc:
            return error_response(exc)
        except ValueError as exc:
            return jsonify({"ok": False, "error": str(exc), "code": "invalid_reversal_date"}), 400
        except sqlite3.IntegrityError:
            return jsonify({
                "ok": False, "error": "Hoàn tác BK xung đột dấu vết; dữ liệu chưa thay đổi",
                "code": "bk_reversal_conflict",
            }), 409
        except Exception:
            return jsonify({
                "ok": False, "error": "Không hoàn tác được BK; toàn bộ bút toán đã hoàn tác",
                "code": "bk_reversal_failed",
            }), 409


__all__ = [
    "BKImportError", "BK_IMPORT_COLUMNS", "BK_IMPORT_POLICY_QUESTION",
    "BK_IMPORT_SOURCE_TYPE", "BK_IMPORT_SHEET", "BK_LEDGER_SOURCE_TABLE",
    "build_bk_draft_template", "build_bk_import_template", "init_bk_import_schema",
    "parse_bk_preview", "register_bk_import_routes",
]
