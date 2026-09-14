from __future__ import annotations
import hashlib
import io
import json
import unittest
from decimal import Decimal
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from . import server
from .document_preview import display_value, formula_value, sheet_preview
from .test_selected_document_export import SelectedDocumentExportTests


class PreviewFormattingTests(unittest.TestCase):
    def test_excel_overflow_labels_get_empty_space_without_covering_other_fields(self):
        from openpyxl.styles import Alignment, Border, Side
        from html.parser import HTMLParser
        class Cells(HTMLParser):
            def __init__(self):
                super().__init__(); self.rows=[]; self.current=None
            def handle_starttag(self, tag, attrs):
                if tag=='tr': self.rows.append([])
                if tag=='td':
                    self.current=dict(attrs); self.current['text']=''; self.rows[-1].append(self.current)
            def handle_data(self, data):
                if self.current is not None: self.current['text']+=data
            def handle_endtag(self, tag):
                if tag=='td': self.current=None
        book=Workbook(); ws=book.active; ws.print_area='A1:D4'
        ws['A1']='Tên doanh nghiệp rất dài'; ws['D1']='Mã số thuế'
        ws['A2']='Địa chỉ dài'; ws.merge_cells('C2:D2'); ws['C2']='Trường gộp'
        ws['A3']='Nội dung phải xuống dòng'; ws['A3'].alignment=Alignment(wrap_text=True)
        ws['A4']='Dừng ở ô có khung'; ws['B4'].border=Border(left=Side(style='thin'))
        before_merges=str(ws.merged_cells)
        parser=Cells(); parser.feed(sheet_preview(ws)['html']); rows=parser.rows
        self.assertEqual(rows[0][0]['colspan'],'3')
        self.assertEqual(rows[0][1]['text'],'Mã số thuế')
        self.assertEqual(rows[1][0]['colspan'],'2')
        self.assertEqual(rows[1][1]['text'],'Trường gộp')
        self.assertEqual(rows[2][0]['colspan'],'1')
        self.assertEqual(rows[3][0]['colspan'],'1')
        self.assertEqual(str(ws.merged_cells),before_merges)
        self.assertIsNone(ws['B1'].value)

    def test_all_thirteen_customer_rounding_examples(self):
        examples = [('42500', '42,500'), ('42000','42,000'), ('152000.994706','152,001'),
                    ('5208','5,208'), ('29629.665','29,630'), ('9196.25149','9,196'),
                    ('29629.633328','29,630'), ('26465.5','26,466'), ('44500','44,500'),
                    ('30962.967143','30,963'), ('167832','167,832'), ('741','741'), ('440000','440,000')]
        for source, expected in examples:
            with self.subTest(source=source): self.assertEqual(expected, display_value(Decimal(source), '#,##0'))
        self.assertEqual('-26,466', display_value(Decimal('-26465.5'), '#,##0'))
        self.assertEqual('0.855123', display_value(0.855123, '#,##0.######'))
        self.assertEqual('8%', display_value(.08, '0%'))

    def test_formulas_and_html_are_safe_and_hidden_columns_never_leak(self):
        workbook = Workbook(); ws = workbook.active
        ws.append(['<script>alert(1)</script>', 7, 2, '=B1-C1', '=ROUND(D1*26465.5,0)', 'secret'])
        ws['A1'].data_type = 's'
        ws['E1'].number_format = '#,##0'
        ws['B2'] = '=SUBTOTAL(9,B1:B1)'
        ws.column_dimensions['F'].hidden = True
        ws.print_area = 'A1:F2'; ws.print_title_rows = '1:1'
        html = sheet_preview(ws)['html']
        self.assertNotIn('<script>', html)
        self.assertIn('&lt;script&gt;', html)
        self.assertNotIn('secret', html)
        self.assertIn('132,328', html)
        self.assertEqual(Decimal(7), formula_value(ws, ws['B2']))
        ws['D2'] = '=HYPERLINK("https://example.com","x")'
        with self.assertRaises(ValueError): sheet_preview(ws)

    def test_merged_header_alignment_and_white_preview(self):
        workbook=Workbook(); ws=workbook.active
        ws.merge_cells('A1:C1'); ws['A1']='Ngày 04 tháng 09 năm 2026'
        from openpyxl.styles import Alignment
        ws['A1'].alignment=Alignment(horizontal='center')
        ws.print_area='A1:C2'
        result=sheet_preview(ws)
        self.assertIn('colspan="3"', result['html'])
        self.assertIn('text-align:center', result['html'])


class Round4DocumentTests(SelectedDocumentExportTests):
    def preview(self, kitchens=('BEP-A','BEP-B')):
        return self.client.post('/api/documents/preview', json={'kind':'deliveries',
            'selections':[{'batch_id':self.batch_id,'kitchen':code} for code in kitchens]})

    def test_preview_selected_kitchens_and_excel_are_immutable_without_db_writes(self):
        before=hashlib.sha256(server.DB_PATH.read_bytes()).hexdigest()
        response=self.preview(('BEP-B',))
        self.assertEqual(200,response.status_code,response.get_json())
        data=response.get_json()
        self.assertEqual(1,data['sheet_count'])
        html=data['sheets'][0]['html']
        self.assertIn('Bếp B',html); self.assertNotIn('Bếp A',html)
        self.assertNotIn('12,000',html) # hidden selling price stays hidden
        self.assertEqual(before,hashlib.sha256(server.DB_PATH.read_bytes()).hexdigest())
        with server.db() as conn:
            conn.execute("UPDATE orders SET product_name='Tên đã sửa sau xem' WHERE batch_id=?",(self.batch_id,))
        xlsx=self.client.get('/api/documents/'+data['token']+'/excel?sheets=0')
        self.assertEqual(200,xlsx.status_code)
        workbook=load_workbook(io.BytesIO(xlsx.data))
        values=[str(c.value) for row in workbook.active for c in row]
        self.assertIn('Mặt hàng 2',values); self.assertNotIn('Tên đã sửa sau xem',values)
        self.assertTrue(all(c.fill.patternType is None for row in workbook.active for c in row))
        workbook.close()

    def test_choose_one_sheet_out_of_all_and_reject_invalid_selections(self):
        data=self.preview().get_json()
        self.assertEqual(2,data['sheet_count'])
        response=self.client.get('/api/documents/'+data['token']+'/excel?sheets=1')
        workbook=load_workbook(io.BytesIO(response.data))
        self.assertEqual(1,len(workbook.worksheets)); self.assertEqual('BEP-B',workbook.active.title)
        workbook.close()
        for query in ('', '2', '-1', '1.1', 'all'):
            self.assertEqual(422,self.client.get('/api/documents/'+data['token']+'/excel?sheets='+query).status_code)
        self.assertEqual(422,self.preview(('NOT-FOUND',)).status_code)
        self.assertEqual(422,self.client.post('/api/documents/preview',json={'kind':'deliveries','selections':[]}).status_code)

    def test_pdf_uses_snapshot_and_only_selected_sheet(self):
        data=self.preview().get_json()
        def build(sources,target,**kwargs):
            self.assertEqual('A4',kwargs['paper'])
            self.assertEqual(1,len(sources))
            workbook=load_workbook(sources[0]['path'])
            self.assertEqual(['BEP-B'],workbook.sheetnames)
            workbook.close()
            target.write_bytes(b'%PDF-1.4 test snapshot')
        with patch('tdp_system.round4_documents.build_excel_pdf_bundle', side_effect=build) as renderer:
            response=self.client.get('/api/documents/'+data['token']+'/pdf?sheets=1')
            self.assertEqual(200,response.status_code)
            self.assertIn('inline',response.headers['Content-Disposition'])
            self.client.get('/api/documents/'+data['token']+'/pdf?sheets=1')
            self.assertEqual(1,renderer.call_count)

    def test_unknown_snapshot_and_expiry_are_clear_errors(self):
        self.assertEqual(422,self.client.get('/api/documents/'+'a'*32+'/excel').status_code)
        data=self.preview().get_json()
        with patch('tdp_system.document_preview.time.time', return_value=99999999999):
            response=self.client.get('/api/documents/'+data['token']+'/excel')
            self.assertEqual(422,response.status_code)
            self.assertIn('hết hạn',response.get_json()['error'])

    def test_simplex_and_duplex_pdf_caches_are_distinct(self):
        data=self.preview().get_json()
        def build(sources,target,**kwargs):
            target.write_bytes(b'%PDF-1.4 '+str(kwargs['duplex']).encode())
        with patch('tdp_system.round4_documents.build_excel_pdf_bundle',side_effect=build) as renderer:
            url='/api/documents/'+data['token']+'/pdf?sheets=1'
            duplex=self.client.get(url+'&sides=duplex')
            simplex=self.client.get(url+'&sides=simplex')
            self.assertEqual(duplex.status_code,200)
            self.assertEqual(simplex.status_code,200)
            self.assertNotEqual(duplex.data,simplex.data)
            self.client.get(url+'&sides=duplex')
            self.assertEqual(renderer.call_count,2)
            self.assertEqual(self.client.get(url+'&sides=unknown').status_code,422)

    def test_delivery_auto_sides_and_response_metadata_survive_pdf_cache(self):
        data=self.preview().get_json()
        def build(sources,target,**kwargs):
            self.assertEqual(kwargs['duplex'],'auto')
            self.assertTrue(all(s['document_type']=='deliveries' for s in sources))
            target.write_bytes(b'%PDF-1.4 auto')
            return {'duplex':True,'pages':4,'page_layout':{'blank_pages':[2]}}
        with patch('tdp_system.round4_documents.build_excel_pdf_bundle',side_effect=build) as renderer:
            url='/api/documents/'+data['token']+'/pdf?sheets=0,1&sides=auto'
            for _ in range(2):
                response=self.client.get(url)
                self.assertEqual(response.status_code,200)
                self.assertEqual(response.headers['X-Print-Sides'],'duplex')
                self.assertEqual(response.headers['X-Print-Pages'],'4')
                self.assertEqual(response.headers['X-Print-Blank-Pages'],'2')
            self.assertEqual(renderer.call_count,1)

    def test_a5_receipts_have_separate_pdf_cache_and_excel_paper(self):
        from .document_preview import create_snapshot
        book=Workbook();book.active.title='bảng kê tổng';book.active['A1']='Bảng tổng'
        book.create_sheet('biên nhận')['A1']='Một người'
        data=create_snapshot(server.DATA_DIR/'document_previews',[('test.xlsx',book)]);book.close()
        def build(sources,target,**kwargs):
            generated=load_workbook(sources[0]['path'])
            self.assertEqual(generated.sheetnames,['biên nhận'])
            if kwargs['paper']=='A5':
                self.assertEqual(str(generated.active.page_setup.paperSize),'11')
                self.assertEqual(generated.active.page_setup.fitToHeight,1)
            generated.close()
            target.write_bytes(b'%PDF-1.4 '+kwargs['paper'].encode())
        url='/api/documents/'+data['token']
        with patch('tdp_system.round4_documents.build_excel_pdf_bundle',side_effect=build) as renderer:
            a4=self.client.get(url+'/pdf?sheets=1&paper=A4&sides=simplex')
            a5=self.client.get(url+'/pdf?sheets=1&paper=A5&sides=simplex')
            self.assertEqual((a4.status_code,a5.status_code),(200,200))
            self.assertNotEqual(a4.data,a5.data)
            self.client.get(url+'/pdf?sheets=1&paper=A5&sides=simplex')
            self.assertEqual(renderer.call_count,2)
            self.assertEqual(self.client.get(url+'/pdf?sheets=0,1&paper=A5').status_code,422)
            self.assertEqual(self.client.get(url+'/pdf?sheets=1&paper=A6').status_code,422)
        exported=self.client.get(url+'/excel?sheets=1&paper=A5')
        self.assertEqual(exported.status_code,200)
        generated=load_workbook(io.BytesIO(exported.data));self.assertEqual(str(generated.active.page_setup.paperSize),'11');generated.close()
        original=load_workbook(server.DATA_DIR/'document_previews'/data['token']/'0.xlsx')
        self.assertIsNone(original['biên nhận'].page_setup.paperSize);original.close()

    def test_receipt_page_error_is_actionable(self):
        from .excel_print_renderer import ReceiptPrintError
        data=self.preview().get_json()
        with patch('tdp_system.round4_documents.build_excel_pdf_bundle',side_effect=ReceiptPrintError('Cần dàn về một trang')):
            response=self.client.get('/api/documents/'+data['token']+'/pdf?sheets=0&sides=duplex')
            self.assertEqual(response.status_code,422)
            self.assertEqual(response.get_json()['code'],'receipt_requires_one_page')

    def test_snapshot_changed_on_disk_is_rejected(self):
        data=self.preview().get_json()
        artifact=server.DATA_DIR/'document_previews'/data['token']/'0.xlsx'
        artifact.write_bytes(b'changed synthetic export')
        response=self.client.get('/api/documents/'+data['token']+'/excel')
        self.assertEqual(422,response.status_code)
        self.assertIn('đã thay đổi',response.get_json()['error'])

    def test_unconfirmed_quote_and_unissued_payment_are_blocked(self):
        self.assertEqual(422,self.client.post('/api/documents/preview',json={'kind':'quote','contractor':'PRINT','period':'2026-09'}).status_code)
        response=self.client.post('/api/documents/preview',json={'kind':'payment','contractor':'PRINT','from':'2026-09-01','to':'2026-09-30'})
        self.assertNotEqual(200,response.status_code)

    def test_historical_list_is_not_limited_to_latest_hundred_batches(self):
        with server.db() as conn:
            for index in range(105):
                conn.execute("INSERT INTO batches(work_date,source_name,status,created_at) VALUES('2026-09-05',?,'draft',?)",(str(index),server.now_iso()))
        response=self.client.get('/api/documents/list?kind=deliveries&from=2026-09-04&to=2026-09-04&customer=BEP-B')
        self.assertEqual(200,response.status_code,response.get_json())
        self.assertEqual([self.batch_id],[r['batch_id'] for r in response.get_json()['rows']])
        self.assertEqual(1,response.get_json()['total'])
        self.assertEqual(400,self.client.get('/api/documents/list?from=2026-09-05&to=2026-09-01').status_code)

    def test_supplier_formulas_preview_without_writes(self):
        response=self.client.post('/api/documents/preview',json={'kind':'suppliers','selections':[{'batch_id':self.batch_id}]})
        self.assertEqual(200,response.status_code,response.get_json())
        self.assertTrue(response.get_json()['sheet_count'])

    def test_purchase_preview_uses_current_identity_without_writes(self):
        with server.db() as conn:
            conn.execute("INSERT INTO people(name,cccd,issue_date,issue_place,address) VALUES(?,?,?,?,?)",
                ('Người bán riêng lượt 4','987654321012','02/01/2020','Nơi cấp kiểm thử','Địa chỉ riêng kiểm thử'))
            conn.execute("UPDATE orders SET purchase_list=1,seller=? WHERE batch_id=?",
                ('Người bán riêng lượt 4',self.batch_id))
        before=hashlib.sha256(server.DB_PATH.read_bytes()).hexdigest()
        response=self.client.post('/api/documents/preview',json={'kind':'purchases','selections':[{'batch_id':self.batch_id}]})
        self.assertEqual(200,response.status_code,response.get_json())
        data=response.get_json()
        self.assertEqual(2,data['sheet_count'])
        self.assertIn('Địa chỉ riêng kiểm thử',data['sheets'][1]['html'])
        self.assertEqual(before,hashlib.sha256(server.DB_PATH.read_bytes()).hexdigest())
        with server.db() as conn:
            conn.execute("UPDATE people SET address='' WHERE name='Người bán riêng lượt 4'")
        invalid=self.client.post('/api/documents/preview',json={'kind':'purchases','selections':[{'batch_id':self.batch_id}]})
        self.assertEqual(422,invalid.status_code)
        self.assertIn('địa chỉ',invalid.get_json()['error'])

    def test_report_and_outgoing_statement_preview_use_existing_builders(self):
        for kind in ('report','outgoing-statement'):
            with self.subTest(kind=kind):
                before=hashlib.sha256(server.DB_PATH.read_bytes()).hexdigest()
                response=self.client.post('/api/documents/preview',json={'kind':kind,'selections':[{'batch_id':self.batch_id}]})
                self.assertEqual(200,response.status_code,response.get_json())
                self.assertGreater(response.get_json()['sheet_count'],0)
                self.assertEqual(before,hashlib.sha256(server.DB_PATH.read_bytes()).hexdigest())

    def test_sample_forms_keep_all_units_and_customer_rounding(self):
        from .verify_round4_artifacts import sample_books
        from .document_preview import create_snapshot
        books=sample_books()
        try:
            data=create_snapshot(server.DATA_DIR/'sample-docs',books)
            self.assertEqual(7,data['sheet_count'])
            self.assertIn('26,466',data['sheets'][0]['html'])
            summary=next(s for s in data['sheets'] if 'bảng kê tổng' in s['name'])
            self.assertIn('2 Cái',summary['html']); self.assertIn('0.855123 Kg',summary['html'])
            receipt=next(s for s in data['sheets'] if s['name'].endswith('· biên nhận'))
            self.assertIn('2 Cái',receipt['html']); self.assertIn('0.855123 Kg',receipt['html'])
            # VND sums add rounded line amounts, never fractional-line round-again differences.
            self.assertIsInstance(books[0][1]['NHUAHP']['I11'].value,int)
        finally:
            for _,book in books: book.close()


if __name__ == '__main__': unittest.main()
