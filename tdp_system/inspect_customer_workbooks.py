from __future__ import annotations

import json
import re
import unicodedata
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "bosung.30.8.26"
OUTPUT = Path(__file__).resolve().parent / "data" / "customer_workbook_inventory.json"
KEYWORDS = {
    "bep", "unit", "menu", "thuc don", "dinh luong", "nguyen lieu", "suat an",
    "po", "don gia", "thanh tien", "de nghi thanh toan", "cong no", "doanh thu",
    "cham cong", "luong", "tang ca", "chu nhat", "le", "bhxh", "tam ung",
    "phu cap", "trach nhiem", "thu viec", "ho ten", "ma nhan vien",
}
ERROR_TOKENS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!")


def normalized(value) -> str:
    text = str(value or "").strip().lower()
    text = unicodedata.normalize("NFD", text)
    return "".join(ch for ch in text if unicodedata.category(ch) != "Mn")


def formula_kind(formula: str) -> str:
    match = re.match(r"=\s*([A-Z][A-Z0-9_.]*)", formula.upper())
    return match.group(1) if match else "OTHER"


def inspect_sheet(ws) -> dict:
    formulas = []
    formula_kinds = Counter()
    error_cells = []
    label_hits = []
    max_row = min(ws.max_row or 0, 5000)
    max_col = min(ws.max_column or 0, 120)
    for row in ws.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_col):
        for cell in row:
            value = cell.value
            if isinstance(value, str) and value.startswith("="):
                formula_kinds[formula_kind(value)] += 1
                if len(formulas) < 12:
                    formulas.append({"cell": cell.coordinate, "formula": value[:180]})
            if isinstance(value, str) and any(token in value.upper() for token in ERROR_TOKENS):
                if len(error_cells) < 30:
                    error_cells.append({"cell": cell.coordinate, "value": value[:180]})
            if isinstance(value, str) and not value.startswith("="):
                text = normalized(value)
                if text and len(text) <= 120 and any(keyword in text for keyword in KEYWORDS):
                    clean = " ".join(value.split())
                    hit = {"cell": cell.coordinate, "label": clean[:160]}
                    if hit not in label_hits and len(label_hits) < 50:
                        label_hits.append(hit)
    return {
        "name": ws.title,
        "state": ws.sheet_state,
        "rows": ws.max_row,
        "columns": ws.max_column,
        "merged_ranges": len(ws.merged_cells.ranges),
        "formula_count": sum(formula_kinds.values()),
        "formula_kinds": dict(formula_kinds.most_common(20)),
        "formula_samples": formulas,
        "error_cells": error_cells,
        "label_hits": label_hits,
    }


def inspect_workbook(path: Path) -> dict:
    try:
        wb = load_workbook(path, data_only=False, read_only=False, keep_links=True)
    except Exception as error:  # The inventory must continue across every customer file.
        return {"file": str(path.relative_to(SOURCE)), "error": type(error).__name__}
    result = {
        "file": str(path.relative_to(SOURCE)),
        "size": path.stat().st_size,
        "sheet_count": len(wb.sheetnames),
        "external_links": len(getattr(wb, "_external_links", []) or []),
        "defined_names": len(wb.defined_names),
        "sheets": [inspect_sheet(ws) for ws in wb.worksheets],
    }
    wb.close()
    return result


def main():
    files = sorted(
        path for path in SOURCE.rglob("*.xlsx")
        if not path.name.startswith("~$")
    )
    inventory = {
        "source": SOURCE.name,
        "file_count": len(files),
        "workbooks": [inspect_workbook(path) for path in files],
    }
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(inventory, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({
        "ok": True,
        "files": len(files),
        "workbooks_with_errors": sum("error" in item for item in inventory["workbooks"]),
        "output": str(OUTPUT),
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
