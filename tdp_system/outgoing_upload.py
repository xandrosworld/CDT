"""Invoice-only workbook requests, constrained to remaining approved order lines."""
from collections import defaultdict
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from hashlib import sha256
from io import BytesIO
import json
import time
import uuid
import zipfile

from flask import jsonify, request, send_file
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

try:
    from .outgoing_contractors import assert_order_enabled, selected_orders
    from .outgoing_unissued import issued_allocations
    from .outgoing_line_policy import unit_issues
    from .outgoing_readiness import canonical_available_stock, invoice_order_issues, validate_draft_export_stock
    from .outgoing_consolidation import _write_draft, decimal, export_quantity, money
    from .outgoing_weights import confirmed_weights, invoice_rows, order_snapshot
    from .outgoing_names import invoice_name
    from .stock_tax_policy import exempt_order_codes
    from .invoice_tax_export import build_invoice_workbook, _safe_name
    from .template_workbook import safe_workbook_bytes
except ImportError:
    from outgoing_contractors import assert_order_enabled, selected_orders
    from outgoing_unissued import issued_allocations
    from outgoing_line_policy import unit_issues
    from outgoing_readiness import canonical_available_stock, invoice_order_issues, validate_draft_export_stock
    from outgoing_consolidation import _write_draft, decimal, export_quantity, money
    from outgoing_weights import confirmed_weights, invoice_rows, order_snapshot
    from outgoing_names import invoice_name
    from stock_tax_policy import exempt_order_codes
    from invoice_tax_export import build_invoice_workbook, _safe_name
    from template_workbook import safe_workbook_bytes

SCHEMA = '''CREATE TABLE IF NOT EXISTS outgoing_upload_jobs (
    token TEXT PRIMARY KEY, source_name TEXT NOT NULL, cutoff TEXT NOT NULL,
    contractor TEXT NOT NULL, rows_json TEXT NOT NULL, preview_hash TEXT NOT NULL,
    preview_json TEXT NOT NULL, created_at TEXT NOT NULL, expires_at REAL NOT NULL,
    result_json TEXT NOT NULL DEFAULT '', generated_at TEXT NOT NULL DEFAULT ''
);'''
HEADERS = ['Dòng đơn', 'Nhà thầu', 'Ngày đơn', 'Bếp', 'Mã hàng', 'Tên xuất hóa đơn',
           'ĐVT đơn', 'SL đề nghị xuất', 'ĐVT xuất hóa đơn', 'Đối chiếu']
SHEET = 'Lap hoa don'
MAX_ROWS = 10000


def packed(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(',', ':'), default=str)


def digest(value):
    return sha256(packed(value).encode()).hexdigest()


def source_rows(conn, cutoff, contractor):
    return [dict(r) for r in conn.execute('''SELECT o.*,b.status batch_status,
        COALESCE(NULLIF(u.invoice_unit,''),p.unit) invoice_unit
        FROM orders o JOIN batches b ON b.id=o.batch_id
        LEFT JOIN products p ON p.code=o.product_code
        LEFT JOIN outgoing_product_units u ON u.product_code=o.product_code
        WHERE o.work_date<=? AND (?='' OR o.contractor=?) ORDER BY o.work_date,o.id''',
        (cutoff, contractor, contractor))]


def template(conn, cutoff, contractor):
    issued, _ = issued_allocations(conn)
    wb = Workbook(); ws = wb.active; ws.title = SHEET; ws.append(HEADERS)
    for o in selected_orders(conn, source_rows(conn, cutoff, contractor)):
        remaining = max(o['actual_delivered']-o['customer_return_qty']-issued.get(o['id'], 0), 0)
        if o['batch_status'] != 'approved' or remaining <= 1e-8:
            continue
        ws.append([o['id'], o['contractor'], o['work_date'], o['kitchen'], o['product_code'],
                   invoice_name(conn, o['product_code'], o['product_name']), o['unit'],
                   round(remaining, 8), o['invoice_unit'], digest(order_snapshot(o))])
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
    for i, width in enumerate([12,22,15,22,16,38,14,20,20,16], 1):
        ws.column_dimensions[ws.cell(1,i).column_letter].width = width
    ws.column_dimensions['J'].hidden = True
    for cell in ws[1]: cell.font = Font(bold=True)
    for row in ws.iter_rows(min_row=2):
        row[7].fill = PatternFill('solid', fgColor='FFF2CC')
        for cell in row: cell.alignment = Alignment(vertical='top', wrap_text=True)
    guide = wb.create_sheet('Huong dan')
    for text in [
        'BẢNG ĐỀ NGHỊ LẬP HÓA ĐƠN TỪ ĐƠN ĐÃ DUYỆT',
        'Chỉ sửa SL đề nghị xuất hoặc xóa những dòng không cần đưa vào lần này; số lượng theo ĐVT đơn.',
        'Giữ các thông tin khác và cột Đối chiếu. Tên/ĐVT hóa đơn lấy từ danh mục, giá bán lấy từ đơn đã duyệt.',
        'Dòng cần kg thực tế: lưu số kg trong mục riêng trên web trước khi nhập file; hệ thống tính lượng và đơn giá hóa đơn.',
        'Nhập file rồi xem trước. Chỉ phần đủ tồn và quy đổi được tạo file; phần chưa đủ vẫn chờ.',
        'Không tạo đơn bán mới, không cộng lại doanh thu, giá vốn hay công nợ.',
        'Tải file chưa phải ký hóa đơn. Số chưa xuất chỉ giảm theo hóa đơn đã ký được đối chiếu.',
        'File mới thay các bản chưa ký có cùng dòng đã chọn; không sử dụng đồng thời các file cũ và mới.',
        f'Phạm vi: đến {cutoff}; nhà thầu {contractor or "tất cả đang chọn"}.',
    ]: guide.append([text])
    guide.column_dimensions['A'].width = 110
    output = safe_workbook_bytes(wb); wb.close()
    return output


def parse_workbook(data):
    # A read-only worksheet can fail while iterating, after load_workbook succeeds.
    # Keep malformed uploads at the input boundary instead of returning a server error.
    try:
        return _parse_workbook(data)
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError('Không đọc được Excel. Dùng file mẫu lập hóa đơn của màn này.') from exc


def _parse_workbook(data):
    if not data or len(data) > 10*1024*1024:
        raise ValueError('Chọn file Excel .xlsx không quá 10 MB.')
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 50*1024*1024:
                raise ValueError('Nội dung Excel quá lớn; tách file dưới 10.000 dòng.')
        wb = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
    except (zipfile.BadZipFile, OSError, KeyError) as exc:
        raise ValueError('Không đọc được Excel. Dùng file mẫu lập hóa đơn của màn này.') from exc
    try:
        if SHEET not in wb.sheetnames:
            raise ValueError('File thiếu sheet Lap hoa don. Tải mẫu ngay tại mục Nhập Excel lập hóa đơn.')
        ws = wb[SHEET]
        if ws.max_row and ws.max_row > MAX_ROWS+1:
            raise ValueError('Mỗi file tối đa 10.000 dòng.')
        if [c.value for c in next(ws.iter_rows(max_row=1, max_col=10))] != HEADERS:
            raise ValueError('Tiêu đề file không đúng mẫu. Giữ nguyên các cột của mẫu đã tải.')
        result = []; seen = set()
        for number, cells in enumerate(ws.iter_rows(min_row=2, max_col=10), 2):
            if number > MAX_ROWS+1: raise ValueError('Mỗi file tối đa 10.000 dòng.')
            if all(c.value in (None, '') for c in cells): continue
            if any(c.data_type == 'f' for c in cells):
                raise ValueError(f'Dòng Excel {number}: dùng giá trị, không dùng công thức.')
            values = [c.value for c in cells]
            try:
                oid = Decimal(str(values[0])); qty = Decimal(str(values[7] or 0))
                if not oid.is_finite() or oid != int(oid) or oid <= 0: raise ValueError()
                if not qty.is_finite() or qty < 0 or qty > 100000000 or qty != qty.quantize(Decimal('.00000001')): raise ValueError()
            except (InvalidOperation, ValueError, OverflowError):
                raise ValueError(f'Dòng Excel {number}: dòng đơn hoặc số lượng không hợp lệ.') from None
            oid = int(oid)
            if oid in seen: raise ValueError(f'Dòng đơn {oid} lặp trong file; gộp số lượng về một dòng.')
            seen.add(oid)
            if qty == 0: continue
            result.append({'order_id':oid, 'qty':float(qty), 'values':values, 'excel_row':number})
        if not result: raise ValueError('File chưa có dòng nào có SL đề nghị xuất lớn hơn 0.')
        return result
    finally:
        wb.close()


def build_plan(conn, requests, cutoff, contractor, tax_percent):
    sources = {o['id']:o for o in source_rows(conn, cutoff, contractor)}
    external = {}
    issued, warnings = issued_allocations(conn, external_quantities=external)
    weights = confirmed_weights(conn)
    items = []; eligible = {}
    for req in requests:
        oid = req['order_id']; o = sources.get(oid); v = req['values']
        if not o or o['batch_status'] != 'approved':
            raise ValueError(f'Dòng đơn {oid} không thuộc đơn đã duyệt trong phạm vi chọn.')
        assert_order_enabled(conn, o['contractor'], oid)
        name = invoice_name(conn, o['product_code'], o['product_name'])
        expected = [o['id'],o['contractor'],o['work_date'],o['kitchen'],o['product_code'],name,o['unit']]
        if (any(str(a or '').strip()!=str(b or '').strip() for a,b in zip(v[:7],expected))
                or str(v[8] or '').strip()!=str(o['invoice_unit'] or '').strip()
                or v[9]!=digest(order_snapshot(o))):
            raise ValueError(f'Dòng đơn {oid}: thông tin khác đơn gốc/danh mục hiện tại. Tải mẫu mới, chỉ sửa số lượng.')
        remaining = max(o['actual_delivered']-o['customer_return_qty']-issued.get(oid,0),0)
        if req['qty'] > remaining+1e-8:
            raise ValueError(f'Dòng đơn {oid}: còn {remaining:g} {o["unit"]} sau đối trừ hóa đơn đã ký; file đề nghị {req["qty"]:g}. Tải mẫu mới.')
        reasons = [w['message'] for w in warnings if not w['contractor'] or w['contractor']==o['contractor']]
        units = unit_issues(conn,[o])
        if oid in units: reasons.append(units[oid]['message'])
        for issue in invoice_order_issues([o]): reasons.extend(issue['messages'])
        if float(o['sell_price'] or 0) < 0: reasons.append('Giá bán không hợp lệ; đối chiếu đơn gốc.')
        protected = conn.execute('''SELECT d.id FROM outgoing_order_allocations a JOIN outgoing_invoice_drafts d ON d.id=a.draft_id
            WHERE a.order_id=? AND d.status='draft' AND d.minvoice_status IN ('saved','saving','unknown') LIMIT 1''',(oid,)).fetchone()
        if protected: reasons.append(f'Đang có dự thảo {protected[0]} trên M-Invoice; đối chiếu bản đó trước.')
        item = {'order_id':oid,'contractor':o['contractor'],'date':o['work_date'],'kitchen':o['kitchen'],
            'product_code':o['product_code'],'invoice_name':name,'unit':o['unit'],'invoice_unit':o['invoice_unit'],
            'requested_qty':req['qty'],'remaining_qty':remaining,'ready_qty':0,'waiting_qty':req['qty'],
            'invoice_qty':0,'invoice_price':0,'amount':0,'reason':' · '.join(dict.fromkeys(reasons))}
        items.append(item)
        if not reasons: eligible[oid] = {**o,'order_id':oid,'qty':req['qty'],'_source_price':decimal(o['sell_price'] or 0)}
    ids = json.dumps(sorted(eligible))
    old = [dict(r) for r in conn.execute('''SELECT DISTINCT d.* FROM outgoing_invoice_drafts d
        JOIN outgoing_order_allocations a ON a.draft_id=d.id WHERE d.status='draft'
        AND COALESCE(d.minvoice_status,'') NOT IN ('saved','saving','unknown')
        AND a.order_id IN (SELECT value FROM json_each(?)) ORDER BY d.id''',(ids,))]
    outside = []; released = defaultdict(lambda:Decimal(0)); all_old_rows = []
    for d in old:
        rows = [dict(r) for r in conn.execute('''SELECT a.*,o.batch_id,o.work_date,o.buy_price,o.contractor,
            o.product_code original_code,o.unit original_unit,COALESCE(x.source_unit_price,a.unit_price) source_price
            FROM outgoing_order_allocations a JOIN orders o ON o.id=a.order_id
            LEFT JOIN outgoing_line_allocations x ON x.line_id=a.id AND x.order_id=a.order_id
            WHERE a.draft_id=? ORDER BY a.id,a.order_id''',(d['id'],))]
        for r in rows:
            if r['product_code']!=r['original_code'] or r['unit'].strip().casefold()!=r['original_unit'].strip().casefold():
                raise ValueError(f'Dự thảo {d["id"]} có mã/ĐVT khác đơn gốc; cần đối chiếu luồng cũ trước khi nhập file.')
            all_old_rows.append(r)
            if r['order_id'] in eligible: released[r['product_code']]+=decimal(r['qty'])
            else: outside.append({**r,'_source_price':decimal(r['source_price'])})
    stock = canonical_available_stock(conn)
    capacity = {code:decimal(r['raw_available_qty'])+released[code] for code,r in stock.items()}
    exempt = {party:exempt_order_codes(conn,[o for o in eligible.values() if o['contractor']==party]) for party in {o['contractor'] for o in eligible.values()}}
    groups = defaultdict(list)
    for o in eligible.values():
        group = (o['contractor'],o['product_code'],o['unit'].strip().casefold(),tax_percent(o['tax']),o.get('invoice_nature') or '1',o['_source_price'],o['id'] if o['id'] in weights else 0)
        groups[group].append(o)
    held=defaultdict(lambda:Decimal(0))
    for r in all_old_rows:
        if r['order_id'] in eligible:held[r['order_id']]+=decimal(r['qty'])
    ordered_groups=sorted(groups.items(),key=lambda pair:(min((r['work_date'],r['id']) for r in pair[1]),pair[0]))
    planned={}
    # Preserve each selected buyer's existing hold before assigning free stock.
    # Uploading all remaining rows must not move an old hold to another buyer.
    for group,rows in ordered_groups:
        code=group[1];want=sum((min(decimal(r['qty']),held[r['id']]) for r in rows),Decimal(0))
        have=max(capacity.get(code,Decimal(0)),Decimal(0))
        keep=export_quantity(want if code in exempt[group[0]] else min(want,have),group[2])
        planned[group]=keep;capacity[code]=have-keep
    for group,rows in ordered_groups:
        code=group[1];total=sum((decimal(r['qty']) for r in rows),Decimal(0))
        have=max(capacity.get(code,Decimal(0)),Decimal(0));keep=planned[group]
        can=export_quantity(total if code in exempt[group[0]] else min(total,keep+have),group[2])
        planned[group]=can;capacity[code]=have-(can-keep)
    by_id = {i['order_id']:i for i in items}; ready = []
    for group, rows in ordered_groups:
        can=planned[group];code=group[1]
        rows=sorted(rows,key=lambda r:(held[r['id']]<=0,r['work_date'],r['id']))
        amount_left=money(can*group[5])
        for o in rows:
            take=min(decimal(o['qty']),can);can-=take;item=by_id[o['id']]
            item['ready_qty']=float(take);item['waiting_qty']=float(decimal(o['qty'])-take)
            amount=Decimal(0) if take<=0 else amount_left if can<=0 else min(money(take*o['_source_price']),amount_left)
            amount_left-=amount;item['amount']=float(amount)
            item['invoice_qty']=float(take);item['invoice_price']=float(o['_source_price'])
            if take>0 and o['id'] in weights:
                w=weights[o['id']]
                kg=(take*decimal(w['actual_kg'])/decimal(w['base_qty'])).quantize(Decimal('.000001'),rounding=ROUND_HALF_UP)
                if kg<=0: raise ValueError(f'{code}: kg của phần đề nghị quá nhỏ; gộp thêm số lượng trước khi xuất.')
                price=(decimal(item['amount'])/kg).quantize(Decimal('.000001'),rounding=ROUND_HALF_UP)
                if money(kg*price)!=decimal(item['amount']):raise ValueError(f'{code}: cần đối chiếu kg để giữ đúng thành tiền.')
                item['invoice_qty']=float(kg);item['invoice_price']=float(price)
            if item['waiting_qty']>1e-8:
                item['reason']='Chưa đủ tồn khả dụng hoặc còn phần lẻ theo bước xuất; phần này tiếp tục chờ.'
            if take>0: ready.append({**o,'qty':float(take)})
    # No eligible quantity means no draft or reservation is touched.
    fingerprint=digest({'items':items,'orders':[order_snapshot(sources[i['order_id']]) for i in items],
        'weights':{oid:weights[oid]['revision'] for oid in eligible if oid in weights},
        'old':old,'old_rows':all_old_rows,
        'stock':{code:stock.get(code) for code in {o['product_code'] for o in eligible.values()}}})
    return {'items':items,'ready_rows':ready,'outside':outside,'old':old,'fingerprint':fingerprint,
        'external':{oid:external.get(oid,0) for oid in eligible},
        'ready_count':sum(i['ready_qty']>1e-8 for i in items),'waiting_count':sum(i['waiting_qty']>1e-8 for i in items),
        'amount':sum(i['amount'] for i in items)}


def public_plan(plan):
    return {k:plan[k] for k in ('items','ready_count','waiting_count','amount')}


def create_plan_drafts(conn, plan, timestamp, tax_percent):
    if not plan['ready_rows']:raise ValueError('Chưa có lượng đủ điều kiện. Xem lý do từng dòng; không tạo dự thảo.')
    for old in plan['old']:
        conn.execute("UPDATE outgoing_invoice_drafts SET status='cancelled' WHERE id=?",(old['id'],))
        conn.execute("UPDATE inventory_transactions SET status='cancelled',updated_at=? WHERE source_type='OUTGOING_DRAFT' AND source_id=? AND status='reserved'",(timestamp,str(old['id'])))
    for oid,qty in plan['external'].items():
        conn.execute('''INSERT INTO outgoing_waiting_settlements(order_id,external_issued_qty,updated_at) VALUES(?,?,?)
            ON CONFLICT(order_id) DO UPDATE SET external_issued_qty=excluded.external_issued_qty,updated_at=excluded.updated_at''',(oid,qty,timestamp))
    created=[]; exported=[]
    for rows, kind in ((plan['outside'],'waiting'),(plan['ready_rows'],'invoice_upload')):
        groups=defaultdict(list)
        for r in rows: groups[(r['contractor'],tax_percent(r['tax']))].append(r)
        for (party,_),group in sorted(groups.items()):
            did=_write_draft(conn,party,group,{r['batch_id'] for r in group},tax_percent,timestamp,floor_kg=False,kind=kind)
            if did:
                created.append(did)
                if kind=='invoice_upload':exported.append(did)
    for did in created:validate_draft_export_stock(conn,did)
    return exported


def export_archive(conn, draft_ids, template_dir, tax_percent, timestamp):
    output=BytesIO()
    with zipfile.ZipFile(output,'w',zipfile.ZIP_DEFLATED) as archive:
        for did in draft_ids:
            d=conn.execute('SELECT * FROM outgoing_invoice_drafts WHERE id=?',(did,)).fetchone()
            if not d or d['status']!='draft' or d['minvoice_status'] in ('saved','saving','unknown'):
                raise ValueError('Bộ file này đã được xử lý hoặc có bản trên M-Invoice. Đối chiếu phần đã ký rồi tải mẫu mới.')
            validate_draft_export_stock(conn,did)
            lines=invoice_rows(conn,[{**dict(r),'contractor':d['contractor']} for r in conn.execute('SELECT * FROM outgoing_invoice_lines WHERE draft_id=? ORDER BY id',(did,))],freeze=True,timestamp=timestamp)
            vat=tax_percent(lines[0]['tax']); data,_=build_invoice_workbook(lines,vat_percent=vat,template_dir=template_dir)
            label='KKKNT' if vat==-2 else 'KCT' if vat==-1 else f'VAT{vat:g}'
            archive.writestr(f'Bang_ke_{_safe_name(d["contractor"])}_{label}_{did}.xlsx',data)
        archive.writestr('HUONG_DAN.txt','\ufeff'+
            'File lập hóa đơn từ phần còn lại của đơn đã duyệt. Không tạo đơn bán mới hoặc cộng lại doanh thu, giá vốn, công nợ.\n'
            'Tên và ĐVT hóa đơn lấy theo danh mục; quy đổi chỉ dùng số kg đã xác nhận.\n'
            'Dùng bộ file này thay các file cũ chưa ký chứa cùng các dòng đã chọn. Tải file chưa tính là đã phát hành.\n'
            'Phần chưa đủ điều kiện vẫn ở bảng còn chờ. Đồng bộ hóa đơn đã ký trước lần tải tiếp theo.')
    return output.getvalue()


def register(app, ctx):
    def scope():
        cutoff=ctx['valid_iso_date'](request.values.get('to') or ctx['date'].today().isoformat(),'Đến ngày đơn')
        party=str(request.values.get('contractor') or '').strip().upper()
        with ctx['db']() as conn:
            if party and not conn.execute('SELECT 1 FROM contractors WHERE code=?',(party,)).fetchone():raise ValueError('Không tìm thấy nhà thầu.')
        return cutoff,party

    @app.get('/api/outgoing-invoice-upload/template.xlsx')
    def upload_template():
        try:
            cutoff,party=scope()
            with ctx['db']() as conn:
                conn.execute('PRAGMA query_only=ON');conn.execute('BEGIN')
                data=template(conn,cutoff,party)
            return send_file(BytesIO(data),as_attachment=True,download_name=f'DE_NGHI_LAP_HOA_DON_{party or "TAT_CA"}_{cutoff}.xlsx')
        except ValueError as exc:return jsonify(ok=False,error=str(exc)),400

    @app.post('/api/outgoing-invoice-upload/preview')
    def upload_preview():
        try:
            cutoff,party=scope();file=request.files.get('file')
            if not file or not file.filename.lower().endswith('.xlsx'):raise ValueError('Chọn file .xlsx theo mẫu lập hóa đơn.')
            requested=parse_workbook(file.read(10*1024*1024+1))
            with ctx['db']() as conn:
                conn.execute('BEGIN IMMEDIATE')
                plan=build_plan(conn,requested,cutoff,party,ctx['invoice_tax_percent'])
                token=uuid.uuid4().hex
                conn.execute('INSERT INTO outgoing_upload_jobs(token,source_name,cutoff,contractor,rows_json,preview_hash,preview_json,created_at,expires_at) VALUES(?,?,?,?,?,?,?,?,?)',
                    (token,file.filename[:240],cutoff,party,packed(requested),plan['fingerprint'],packed(public_plan(plan)),ctx['now_iso'](),time.time()+900))
            return jsonify(ok=True,token=token,**public_plan(plan))
        except (ValueError,InvalidOperation) as exc:return jsonify(ok=False,error=str(exc)),409

    @app.post('/api/outgoing-invoice-upload/export')
    def upload_export():
        try:
            body=request.get_json(silent=True) or {}
            if not isinstance(body,dict) or body.get('confirmed') is not True:raise ValueError('Xem trước và xác nhận phần đề nghị trước khi tạo file.')
            token=str(body.get('token') or '')
            with ctx['db']() as conn:
                job=conn.execute('SELECT * FROM outgoing_upload_jobs WHERE token=?',(token,)).fetchone()
                if not job:raise ValueError('Không tìm thấy bản xem trước. Nhập lại file.')
                job=dict(job)
                if not job['result_json'] and time.time()>job['expires_at']:raise ValueError('Bản xem trước đã hết hạn 15 phút. Nhập lại file để kiểm tra số mới.')
                requested=json.loads(job['rows_json'])
                for r in requested:
                    o=conn.execute('SELECT contractor FROM orders WHERE id=?',(r['order_id'],)).fetchone()
                    if not o:raise ValueError('Dòng đơn đã thay đổi. Nhập lại file.')
                    assert_order_enabled(conn,o['contractor'],r['order_id'])
                connected=ctx['setting_get'](conn,'minvoice_active_connection','')
            if connected:
                try:
                    from .outgoing_source_refresh import refresh_sources
                    from .order_export_scope import business_today
                except ImportError:
                    from outgoing_source_refresh import refresh_sources
                    from order_export_scope import business_today
                try:
                    refresh_sources(ctx['db'],ctx['create_minvoice_client'],ctx['now_iso'],
                        min(str(r['values'][2]) for r in requested),max(job['cutoff'],business_today()))
                except (ValueError,ctx['MinvoiceError']) as exc:
                    raise ValueError('Chưa cập nhật đủ hóa đơn đã ký; chưa tạo file. '+str(exc)) from exc
            with ctx['db']() as conn:
                conn.execute('BEGIN IMMEDIATE')
                job=dict(conn.execute('SELECT * FROM outgoing_upload_jobs WHERE token=?',(token,)).fetchone())
                timestamp=ctx['now_iso']()
                if job['result_json']:
                    result=json.loads(job['result_json']);ids=result['draft_ids']
                else:
                    if time.time()>job['expires_at']:raise ValueError('Bản xem trước đã hết hạn. Nhập lại file.')
                    plan=build_plan(conn,requested,job['cutoff'],job['contractor'],ctx['invoice_tax_percent'])
                    if plan['fingerprint']!=job['preview_hash']:
                        raise ValueError('Đơn, lựa chọn, quy đổi, tồn hoặc hóa đơn đã ký vừa thay đổi. Nhập lại file để xem số mới; chưa tạo dự thảo.')
                    ids=create_plan_drafts(conn,plan,timestamp,ctx['invoice_tax_percent'])
                    result={'draft_ids':ids,**public_plan(plan)}
                    conn.execute('UPDATE outgoing_upload_jobs SET result_json=?,generated_at=? WHERE token=?',(packed(result),timestamp,token))
                    ctx['audit_event'](conn,'outgoing.invoice_upload',entity_type='outgoing_invoice',entity_id=token,
                        metadata={'draft_ids':ids,'order_ids':[r['order_id'] for r in requested],
                                  'retired_draft_ids':[d['id'] for d in plan['old']],'sale_recorded_again':False})
                data=export_archive(conn,ids,ctx['TAX_TEMPLATE_DIR'],ctx['invoice_tax_percent'],timestamp)
            response=send_file(BytesIO(data),as_attachment=True,download_name=f'BANG_KE_TU_FILE_{token[:8]}.zip',mimetype='application/zip')
            response.headers['X-Invoice-Files']=str(len(ids))
            response.headers['X-Waiting-Rows']=str(result['waiting_count'])
            return response
        except (ValueError,InvalidOperation) as exc:return jsonify(ok=False,error=str(exc)),409
