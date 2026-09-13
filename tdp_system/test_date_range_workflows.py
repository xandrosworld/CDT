"""Inclusive report ranges and independent supplier plans/statuses across dates."""
import io
import hashlib
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook
from . import server, contract_modules
from .test_round3_documents import seed_round3
from . import test_supplier_order_checklist as checklist_fixtures


class DateRangeWorkflowsTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.original = server.DB_PATH, server.DATA_DIR, server.MASTER_SOURCE, server.app.config['TESTING']
        server.DATA_DIR = Path(self.temp.name)
        server.DB_PATH = server.DATA_DIR / 'range.sqlite3'
        server.MASTER_SOURCE = server.DATA_DIR / 'no-import.xlsx'
        server.app.config['TESTING'] = True
        server.init_database(sync_master=False)
        server.MASTER_SOURCE = Path(__file__).resolve().parent.parent / 'Em Thành.xlsx'
        self.client = server.app.test_client()

    def tearDown(self):
        server.DB_PATH, server.DATA_DIR, server.MASTER_SOURCE, server.app.config['TESTING'] = self.original
        self.temp.cleanup()

    def get(self, path):
        response = self.client.get(path)
        self.assertEqual(response.status_code, 200, response.json)
        return response.json

    def seed_report_dates(self):
        with server.db() as conn:
            first = seed_round3(conn, 1)
            conn.execute("UPDATE batches SET work_date='2026-12-31' WHERE id=?", (first,))
            conn.execute("UPDATE orders SET work_date='2026-12-31' WHERE batch_id=?", (first,))
            source = dict(conn.execute('SELECT * FROM orders WHERE batch_id=?', (first,)).fetchone())
            ids = [first]
            for day, status, qty in [('2027-01-01','approved',2), ('2026-12-30','approved',3),
                                      ('2027-01-02','approved',4), ('2027-01-01','draft',999)]:
                bid = conn.execute("INSERT INTO batches(work_date,source_name,status,created_at) VALUES(?,?,?,?)", (day,day,status,server.now_iso())).lastrowid
                item = {**source, 'batch_id':bid, 'work_date':day, 'qty':qty,'actual_received':qty,'actual_delivered':qty}
                del item['id']
                conn.execute(f"INSERT INTO orders({','.join(item)}) VALUES({','.join('?' for _ in item)})", tuple(item.values()))
                ids.append(bid)
            totals = [0,0,0,0]
            for row in conn.execute('SELECT * FROM orders WHERE batch_id IN (?,?)', ids[:2]):
                totals = [a+b for a,b in zip(totals, server.order_totals(row))]
            return totals

    def test_report_inclusive_cross_year_and_exact_workbook(self):
        expected = self.seed_report_dates()
        query = '?from=2026-12-31&to=2027-01-01'
        data = self.get('/api/reports/summary'+query)
        self.assertEqual(data['rows'][-1][3:], expected)
        self.assertEqual(data['rows'][-1][0], 'TỔNG KỲ')
        self.assertEqual(data['draft_count'], 1)
        self.assertEqual(next(row for row in data['rows'] if row[2]=='K2')[3:], [0,0,0,0])
        response = self.client.get('/api/reports/summary/export'+query)
        self.assertEqual(response.status_code,200)
        with io.BytesIO(response.data) as stream:
            book = load_workbook(stream)
            self.assertEqual([list(row) for row in book.active.iter_rows(min_row=3,values_only=True)], data['rows'])
            self.assertIn('31/12/2026',book.active['A1'].value)
            self.assertIn('01/01/2027',book.active['A1'].value)
            self.assertIn('2026-12-31_2027-01-01',response.headers['Content-Disposition'])
            book.close()

    def test_report_single_day_empty_and_month_compatibility(self):
        self.seed_report_dates()
        day = self.get('/api/reports/summary?from=2026-12-31&to=2026-12-31')
        month = self.get('/api/reports/monthly?period=2026-12')
        whole = self.get('/api/reports/summary?from=2026-12-01&to=2026-12-31')
        self.assertEqual(month['rows'][-1][0], 'TỔNG THÁNG')
        self.assertEqual(month['rows'][-1][3:], whole['rows'][-1][3:])
        self.assertGreater(whole['rows'][-1][3], day['rows'][-1][3])
        empty = self.get('/api/reports/summary?from=2028-02-29&to=2028-02-29')
        self.assertEqual(empty['rows'][-1][3:], [0,0,0,0])
        self.assertEqual(empty['draft_count'],0)

    def test_print_report_keeps_exact_range_and_never_writes_accounts(self):
        expected = self.seed_report_dates()
        before = hashlib.sha256(server.DB_PATH.read_bytes()).hexdigest()
        response = self.client.post('/api/documents/preview', json={
            'kind': 'report', 'from': '2026-12-31', 'to': '2027-01-01'})
        self.assertEqual(response.status_code, 200, response.json)
        sheets = response.json['sheets']
        self.assertEqual(len(sheets), 1)
        self.assertIn('31/12/2026', sheets[0]['html'])
        self.assertIn('01/01/2027', sheets[0]['html'])
        from .document_preview import snapshot_files
        _, files = snapshot_files(server.DATA_DIR / 'document_previews', response.json['token'])
        book = load_workbook(files[0][1], data_only=True)
        self.assertEqual([book.active.cell(book.active.max_row, c).value for c in range(4, 8)], expected)
        book.close()
        for start, end in [('2026-12-31', None), ('2027-01-01', '2026-12-31'), ('2026-02-29', '2026-03-01')]:
            invalid = self.client.post('/api/documents/preview', json={'kind':'report', 'from':start, 'to':end})
            self.assertEqual(invalid.status_code, 422, invalid.json)
        self.assertEqual(before, hashlib.sha256(server.DB_PATH.read_bytes()).hexdigest())

    def test_invalid_ranges_rejected_for_json_and_download(self):
        for path in ('/api/reports/summary','/api/reports/summary/export','/api/supplier-needs'):
            for query in ('', '?from=2026-09-13', '?from=2026-09-13&to=2026-09-12',
                          '?from=2026-02-29&to=2026-03-01','?from=no&to=2026-09-13'):
                response = self.client.get(path+query)
                self.assertEqual(response.status_code,400,(path,query,response.json))

    def supplier_batch(self, conn, day, qty, status='pending'):
        bid = checklist_fixtures.SupplierOrderChecklistTests._create_batch(conn, source_name=day,fixtures=[('HƯƠNG','POT','Thịt heo vai sấn',qty)])
        conn.execute('UPDATE batches SET work_date=? WHERE id=?',(day,bid))
        conn.execute("UPDATE purchase_workbook_lines SET work_date=?,status='confirmed' WHERE batch_id=?",(day,bid))
        if status != 'pending':
            conn.execute("INSERT INTO supplier_order_statuses(batch_id,supplier_key,supplier_label,status,revision,updated_at) VALUES(?,'huong','HƯƠNG',?,1,?)", (bid,status,server.now_iso()))
        return bid

    def test_supplier_range_matches_single_day_source_and_keeps_missing_day(self):
        with server.db() as conn:
            before = self.supplier_batch(conn,'2026-08-30',100)
            first = self.supplier_batch(conn,'2026-08-31',29.5,'ordered')
            second = self.supplier_batch(conn,'2026-09-01',3.25)
            after = self.supplier_batch(conn,'2026-09-02',200)
            # A customer order is intentionally different and cannot replace a plan.
            conn.execute("INSERT INTO orders(batch_id,work_date,product_name,qty,supplier,updated_at) VALUES(?,'2026-09-01','Sai sheet',999,'HƯƠNG',?)",(second,server.now_iso()))
            missing = conn.execute("INSERT INTO batches(work_date,source_name,status,created_at) VALUES('2026-09-01','Thiếu sheet','draft',?)",(server.now_iso(),)).lastrowid
            conn.execute("INSERT INTO orders(batch_id,work_date,product_name,qty,supplier,updated_at) VALUES(?,'2026-09-01','Không lấy đơn khách',999,'HƯƠNG',?)",(missing,server.now_iso()))
        data = self.get('/api/supplier-needs?from=2026-08-31&to=2026-09-01')
        self.assertEqual([d['batch_id'] for d in data['days']],[first,second,missing])
        for item in data['days']:
            single = self.get('/api/supplier-needs/'+str(item['batch_id']))
            del single['ok']
            self.assertEqual(item,single)
        self.assertEqual(data['days'][0]['groups'][0]['total_qty'],29.5)
        self.assertEqual(data['days'][1]['groups'][0]['total_qty'],3.25)
        self.assertEqual(data['days'][0]['checklist_counts']['ordered'],1)
        self.assertEqual(data['days'][1]['checklist_counts']['pending'],1)
        self.assertFalse(data['days'][2]['send_available'])
        self.assertEqual(data['days'][2]['groups'],[])
        self.assertEqual(data['days'][2]['rows'],[])
        self.assertEqual(self.get('/api/supplier-needs?from=2027-01-01&to=2027-01-31')['days'],[])

    def test_supplier_status_change_only_affects_target_day_and_checks_plan_hash(self):
        with server.db() as conn:
            first=self.supplier_batch(conn,'2026-09-01',29.5)
            second=self.supplier_batch(conn,'2026-09-02',3.25)
        data=self.get('/api/supplier-needs?from=2026-09-01&to=2026-09-02')
        url=f'/api/supplier-order-status/{second}/huong'
        body={'status':'ordered','revision':0,'plan_hash':data['days'][0]['plan_hash']}
        self.assertEqual(self.client.put(url,json=body).status_code,409)
        body['plan_hash']=data['days'][1]['plan_hash']
        self.assertEqual(self.client.put(url,json=body).status_code,200)
        current=self.get('/api/supplier-needs?from=2026-09-01&to=2026-09-02')
        self.assertEqual(current['days'][0]['checklist_counts']['pending'],1)
        self.assertEqual(current['days'][1]['checklist_counts']['ordered'],1)
        self.assertEqual(self.client.put(url,json=body).status_code,409)
        self.assertEqual(self.client.put(url,json={'status':'reopened','revision':1}).status_code,200)
        self.assertEqual(self.get('/api/supplier-needs/'+str(first))['checklist_counts']['pending'],1)


if __name__ == '__main__':
    unittest.main()
