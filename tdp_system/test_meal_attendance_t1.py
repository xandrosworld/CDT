import sqlite3
import unittest
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook


SYSTEM_DIR = Path(__file__).resolve().parent
PROJECT_DIR = SYSTEM_DIR.parent

try:
    from .contract_modules import parse_meal_attendance_workbook
except ImportError:  # pragma: no cover - direct file invocation
    from contract_modules import parse_meal_attendance_workbook


REAL_T1_FILE = (
    PROJECT_DIR
    / "bosung.30.8.26"
    / "CHẤM CÔNG+ SUẤT ĂN  2026"
    / "SUẤT ĂN XƯỞNG CƠM 2026"
    / "SUẤT ĂN T1-2026 -.xlsx"
)


def meal_connection():
    connection = sqlite3.connect(":memory:")
    connection.row_factory = sqlite3.Row
    connection.execute(
        """CREATE TABLE meal_attendance(
           work_date TEXT NOT NULL,
           kitchen TEXT NOT NULL,
           shift TEXT NOT NULL,
           actual_count REAL NOT NULL DEFAULT 0,
           ordered_count REAL NOT NULL DEFAULT 0,
           PRIMARY KEY(work_date,kitchen,shift)
           )"""
    )
    return connection


def add_legacy_header(worksheet, header_row=4):
    worksheet.cell(header_row, 1, "STT")
    worksheet.cell(header_row, 2, "Ngày")
    worksheet.cell(header_row, 3, "Ca đêm")
    worksheet.cell(header_row, 4, "Tổng")


class MealAttendanceJanuaryLegacyTests(unittest.TestCase):
    def test_real_t1_uses_explicit_period_and_is_idempotent(self):
        self.assertTrue(REAL_T1_FILE.exists(), REAL_T1_FILE)
        connection = meal_connection()
        workbook = load_workbook(
            REAL_T1_FILE, read_only=False, data_only=True, keep_links=False,
        )
        try:
            preview = parse_meal_attendance_workbook(connection, workbook, "2026-01")
        finally:
            workbook.close()

        self.assertTrue(preview["can_confirm"])
        self.assertEqual(preview["periods"], ["2026-01"])
        self.assertEqual(preview["counts"]["items"], 30)
        self.assertEqual(preview["counts"]["dates"], 30)
        self.assertEqual(preview["counts"]["kitchens"], 1)
        self.assertEqual(preview["totals"]["actual"], 1881)
        self.assertEqual(preview["totals"]["ordered"], 0)
        self.assertEqual({row["kitchen"] for row in preview["rows"]}, {"UNI"})
        self.assertEqual({row["shift"] for row in preview["rows"]}, {"Đêm"})
        self.assertEqual(preview["rows"][0]["work_date"], "2026-01-02")
        self.assertEqual(preview["rows"][-1]["work_date"], "2026-01-31")
        self.assertTrue(all(row["work_date"].startswith("2026-01-") for row in preview["rows"]))

        connection.executemany(
            """INSERT INTO meal_attendance(
               work_date,kitchen,shift,actual_count,ordered_count
               ) VALUES(?,?,?,?,?)""",
            [
                (
                    item["work_date"], item["kitchen"], item["shift"],
                    item["actual_count"], item["ordered_count"],
                )
                for item in preview["items"]
            ],
        )
        workbook = load_workbook(
            REAL_T1_FILE, read_only=False, data_only=True, keep_links=False,
        )
        try:
            repeated = parse_meal_attendance_workbook(connection, workbook, "2026-01")
        finally:
            workbook.close()
            connection.close()
        self.assertEqual(repeated["counts"]["new"], 0)
        self.assertEqual(repeated["counts"]["update"], 0)
        self.assertEqual(repeated["counts"]["unchanged"], 30)
        self.assertEqual(repeated["totals"]["actual"], 1881)

    def test_broken_shift_uses_visible_total_without_guessing_shift(self):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "UNI"
        add_legacy_header(worksheet)
        worksheet["A6"] = 1
        worksheet["B6"] = date(2025, 12, 1)  # stale linked date must be ignored
        worksheet["C6"] = "#REF!"
        worksheet["D6"] = 12
        connection = meal_connection()
        try:
            preview = parse_meal_attendance_workbook(connection, workbook, "2026-01")
        finally:
            connection.close()
            workbook.close()

        self.assertTrue(preview["can_confirm"])
        self.assertEqual(preview["counts"]["items"], 1)
        row = preview["rows"][0]
        self.assertEqual(row["work_date"], "2026-01-01")
        self.assertEqual(row["shift"], "Tổng")
        self.assertEqual(row["actual_count"], 12)
        self.assertIn("không tự phân bổ", row["warnings"][0])

    def test_unresolved_ref_reports_exact_sheet_row_and_cells(self):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "TQ"
        add_legacy_header(worksheet)
        worksheet["A6"] = 1
        worksheet["C6"] = "#REF!"
        worksheet["D6"] = "#REF!"
        connection = meal_connection()
        try:
            with self.assertRaisesRegex(
                ValueError, r"Sheet TQ, dòng 6:.*C6.*D6",
            ):
                parse_meal_attendance_workbook(connection, workbook, "2026-01")
        finally:
            connection.close()
            workbook.close()

    def test_conflicting_visible_total_is_not_silently_accepted(self):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "UNI"
        add_legacy_header(worksheet)
        worksheet["A6"] = 1
        worksheet["C6"] = 4
        worksheet["D6"] = 5
        connection = meal_connection()
        try:
            with self.assertRaisesRegex(
                ValueError, r"Sheet UNI, dòng 6:.*5.*D6.*4.*C6",
            ):
                parse_meal_attendance_workbook(connection, workbook, "2026-01")
        finally:
            connection.close()
            workbook.close()


if __name__ == "__main__":
    unittest.main()
