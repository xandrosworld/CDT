import copy
import unittest

from openpyxl import load_workbook

from . import test_inventory_export as fixtures
from .inventory_export import monthly_customer_model
from .inventory_customer_report import customer_nxt_workbook
from .inventory_customer_report import tax_value
from .inventory_closing_report import customer_closing_workbook, TEMPLATE_PATH
from .invoice_workbench_listing import invoice_range_payload
from .inventory_report_company import DEFAULT_COMPANY
from .inventory_tax_review import compare_tax, display_tax


class ClosingReportTests(unittest.TestCase):
    def test_tax_representations_are_compared_without_false_conflicts(self):
        for source in ('8', '8%', '0.08', 0.08):
            self.assertFalse(compare_tax({'tax':'0.08'},[{'tax_rate':source}])['differs'])
        for value in ('KKKNT', -2, '-2.0'):
            self.assertEqual('KKKNT',display_tax(value))
            self.assertFalse(compare_tax({'tax':'KKKNT'},[{'tax_rate':value}])['differs'])
        self.assertTrue(compare_tax({'tax':'KKKNT'},[{'tax_rate':'8'}])['differs'])
        self.assertFalse(compare_tax({'tax':''},[{'tax_rate':'8'}])['differs'])

    def test_catalog_invoice_mismatch_is_explained_without_changing_tax_or_stock(self):
        item=self.model['items'][0];item['tax']='KKKNT'
        code=item['product_code']
        for line in self.sales['lines']:
            if line.get('product_code')==code:line['tax_rate']='8'
        before_model=copy.deepcopy(self.model);before_sales=copy.deepcopy(self.sales)
        for builder,first in ((customer_closing_workbook,7),(customer_nxt_workbook,10)):
            book=builder(self.model,self.sales)
            try:
                sheet=book['Đối chiếu thuế']
                row=next(r for r in sheet.iter_rows(min_row=3,values_only=True) if r[0]==code)
                self.assertEqual(('KKKNT','8%'),row[2:4])
                self.assertIn(str(self.sales['items'][0]['invoice_number']),row[6])
                self.assertEqual('KKKNT',book.active.cell(first,5).value)
                self.assertIn('8%',book.active.cell(first,5).comment.text)
                self.assertIn('Đối chiếu thuế',book.active.cell(6 if first==7 else 7,1).value)
            finally:book.close()
        self.assertEqual(before_model,self.model);self.assertEqual(before_sales,self.sales)

    def test_stock_tax_category_codes_are_not_rendered_as_negative_percentages(self):
        self.assertEqual('KKKNT', tax_value({'tax': -2}, []))
        self.assertEqual('KCT', tax_value({'tax': '-1'}, []))
        self.assertEqual('KKKNT', tax_value({'tax': ''}, [{'tax_rate': '-2'}]))
        self.assertEqual(0.08, tax_value({'tax': 0.08}, []))

    def setUp(self):
        self.fixture = fixtures.InventoryExportTests()
        self.fixture.setUp()
        self.model = monthly_customer_model(self.fixture.conn, self.fixture._model())
        self.sales = invoice_range_payload(self.fixture.conn, tenant='TDP', invoice_type='output',
                                         date_from='2026-08-01', date_to='2026-08-31')

    def tearDown(self):
        self.fixture.tearDown()

    def test_customer_layout_and_values_match_nxt_including_tax_fallback(self):
        self.model['items'][0]['tax'] = ''
        book = customer_closing_workbook(self.model, self.sales)
        nxt = customer_nxt_workbook(self.model, self.sales)
        template = load_workbook(TEMPLATE_PATH)
        try:
            ws = book.active
            for row in (4, 5):
                for col in range(1, 10):
                    self.assertEqual(template.active.cell(row, col).value, ws.cell(row, col).value)
                    self.assertEqual(str(template.active.cell(row, col).font), str(ws.cell(row, col).font))
            self.assertEqual(9, ws.max_column)
            self.assertEqual([nxt.active.cell(10, c).value for c in (1,2,3,4,5,6,16,17,18)],
                             [ws.cell(7, c).value for c in range(1,10)])
            self.assertEqual('G7', ws.freeze_panes)
            self.assertEqual("'Tồn trong kỳ'!$A$1:$I$8", str(ws.print_area))
            self.assertEqual('8', str(ws.page_setup.paperSize))
            self.assertIn('THÁNG 8/2026', ws['A6'].value)
            self.assertEqual(1600, ws['I8'].value)
        finally:
            book.close(); nxt.close(); template.close()

    def test_zero_negative_stock_mixed_units_and_literal_names_survive(self):
        original = self.model['items'][0]
        original['product_name'] = '=1+1'
        zero = dict(original, product_code='ZERO', closing_qty=0, closing_value=0)
        negative = dict(original, product_code='NEG', unit='Other', closing_qty=-2,
                        closing_value=-200, average_unit_cost=100, valuation_status='negative_closing')
        self.model['items'].extend([zero, negative])
        before = copy.deepcopy(self.model)
        book = customer_closing_workbook(self.model, self.sales)
        try:
            ws = book.active
            self.assertEqual('s', ws['C7'].data_type)
            self.assertEqual('=1+1', ws['C7'].value)
            self.assertEqual(0, ws['G8'].value)
            self.assertEqual(-2, ws['G9'].value)
            self.assertEqual(-200, ws['I9'].value)
            self.assertEqual(1400, ws['I10'].value)
            self.assertIn('2 ĐVT', ws['G10'].value)
            self.assertEqual(('NEG', 'negative_closing'), list(book['Đối chiếu'].values)[1])
            self.assertEqual(before, self.model)
        finally:
            book.close()

    def test_monthly_valuation_is_required(self):
        self.model['valuation_method'] = 'moving_average'
        with self.assertRaises(ValueError):
            customer_closing_workbook(self.model, self.sales)

    def test_reports_use_current_company_settings_instead_of_sample_identity(self):
        for key, value in DEFAULT_COMPANY.items():
            self.fixture.conn.execute('INSERT OR REPLACE INTO settings(key,value) VALUES (?,?)', (key, value))
        model = monthly_customer_model(self.fixture.conn, self.fixture._model())
        for builder in (customer_closing_workbook, customer_nxt_workbook):
            book = builder(model, self.sales)
            try:
                self.assertEqual(DEFAULT_COMPANY['company'], book.active['A1'].value)
                self.assertEqual('Địa chỉ: ' + DEFAULT_COMPANY['company_address'], book.active['A2'].value)
                self.assertEqual('MST: 0202265016', book.active['A3'].value)
                value_col = 8 if builder == customer_closing_workbook else 17
                self.assertEqual(1600, next(row[value_col] for row in list(book.active.values) if row[1] == model['items'][0]['product_code']))
                self.assertNotIn('0201650135', str(list(book.active.values)))
            finally:
                book.close()
        self.fixture.conn.execute("UPDATE settings SET value='Tên công ty đã cập nhật' WHERE key='company'")
        model = monthly_customer_model(self.fixture.conn, self.fixture._model())
        book = customer_closing_workbook(model, self.sales)
        self.assertEqual('Tên công ty đã cập nhật', book.active['A1'].value)
        book.close()
