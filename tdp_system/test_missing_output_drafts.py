import copy
import sqlite3
import unittest
from types import SimpleNamespace
from openpyxl import load_workbook

from .test_invoice_input_sync import init_test_database
from .test_minvoice_portal import document
from .minvoice_portal import normalize_portal_document
from .invoice_workbench import prepare_sync_batch
from .invoice_output_sync import upsert_output_invoice, sync_output_batch, InvoiceOutputSyncError
from .invoice_workbench_listing import invoice_range_payload, range_workbook
from .invoice_output_register import output_sales_workbook


class MissingDraftTests(unittest.TestCase):
    def setUp(self):
        self.c = sqlite3.connect(':memory:')
        init_test_database(self.c)
        self.draft = document()
        self.draft.update(invoiceNumber=None, sendTaxStatus=1)
        self.draft = normalize_portal_document(self.draft)
        self.draft_id = self.insert(self.draft)
        self.issued = copy.deepcopy(self.draft)
        self.issued.update(id='00000000-0000-0000-0000-000000000002', invoiceNumber=817, sendTaxStatus=4)
        self.issued_id = self.insert(self.issued)
        self.batch, _ = prepare_sync_batch(self.c, tenant='TDP', source='minvoice', invoice_type='output',
            date_from='2026-08-01', date_to='2026-08-31', now_iso=lambda:'2026-09-21 16:00:00')
        self.client = SimpleNamespace(config=SimpleNamespace(api_mode='portal'), is_test_environment=False,
            get_invoice_series=lambda:[{'value':'1C26TYY','invoiceYear':26}])
        self.rows = [self.issued]
        self.client.get_outgoing_invoices = lambda *args, **kw: {
            'data':self.rows[kw['start']:kw['start']+kw['count']], 'total':len(self.rows)}

    def tearDown(self):
        self.c.close()

    def insert(self, d, tenant='TDP'):
        return upsert_output_invoice(self.c,d,tenant=tenant,now='2026-09-16 20:00:00',
            status_map={},status_fields=(),reference_fields=())[0]

    def sync(self, **kw):
        return sync_output_batch(self.c,self.client,self.batch['id'],lambda:'2026-09-21 16:00:00',**kw)

    def payload(self, **kw):
        return invoice_range_payload(self.c,tenant='TDP',invoice_type='output',
            date_from='2026-08-01',date_to='2026-08-31',**kw)

    def test_missing_unsigned_draft_is_history_not_total_or_either_export(self):
        before = tuple(self.c.execute('SELECT * FROM outgoing_source_invoices WHERE id=?',(self.draft_id,)).fetchone())
        self.sync()
        p = self.payload()
        self.assertEqual([i['id'] for i in p['items']], [self.issued_id])
        self.assertEqual(p['missing_drafts'][0]['id'],self.draft_id)
        self.assertEqual(p['output_summary']['total_amount'],108000)
        self.assertEqual(p['counts']['draft'],0)
        self.assertEqual(self.payload(scope='pending')['counts']['draft'],0)
        for export in (range_workbook,output_sales_workbook):
            rows=list(load_workbook(export(p),data_only=True).active.values)
            self.assertEqual(sum(r[1]=='1C26TYY / 817' for r in rows),1)
            self.assertFalse(any(r[1]=='1C26TYY / ' for r in rows))
        self.assertEqual(before,tuple(self.c.execute('SELECT * FROM outgoing_source_invoices WHERE id=?',(self.draft_id,)).fetchone()))
        self.assertEqual(self.c.execute('SELECT COUNT(*) FROM invoice_inventory_ledger').fetchone()[0],0)
        self.sync()
        self.assertEqual(self.c.execute('SELECT COUNT(*) FROM outgoing_missing_drafts').fetchone()[0],1)

    def test_reappearing_draft_and_then_signed_document_are_restored(self):
        self.sync()
        self.rows.append(self.draft)
        self.sync()
        self.assertEqual(len(self.payload()['items']),2)
        self.assertEqual(self.payload()['missing_drafts'],[])
        self.rows[1] = dict(self.draft,invoiceNumber=818,sendTaxStatus=4,dateSign='2026-09-21T16:00:00+07:00')
        self.sync()
        self.assertEqual({i['invoice_number'] for i in self.payload()['items']},{'817','818'})

    def test_issued_filter_independent_of_stock_status_and_matches_filtered_excel(self):
        for stock in ['pending_mapping','ready','posted','blocked']:
            self.c.execute('UPDATE outgoing_source_invoices SET stock_status=? WHERE id=?',(stock,self.issued_id))
            p=self.payload(status='issued')
            self.assertEqual([i['id'] for i in p['items']],[self.issued_id])
            self.assertEqual(p['counts']['issued'],1)
            self.assertEqual(p['totals']['invoice_amount'],108000)
            rows=list(load_workbook(range_workbook(p),data_only=True).active.values)
            self.assertEqual(sum(r[1]=='1C26TYY / 817' for r in rows),1)
        self.c.execute("UPDATE outgoing_source_invoices SET source_status_class='cancelled' WHERE id=?",(self.issued_id,))
        self.assertEqual(self.payload(status='issued')['items'],[])

    def test_partial_or_resumed_tail_does_not_prove_absence(self):
        other=dict(self.issued,id='00000000-0000-0000-0000-000000000003',invoiceNumber=818)
        self.rows.append(other)
        self.assertFalse(self.sync(max_pages=1,page_size=1)['complete'])
        self.assertEqual(len(self.payload()['items']),2)
        self.assertTrue(self.sync(max_pages=1,page_size=1)['complete'])
        self.assertEqual(self.payload()['missing_drafts'],[])
        self.sync()
        self.assertEqual(len(self.payload()['missing_drafts']),1)

    def test_failed_request_rolls_back_missing_state(self):
        self.client.get_outgoing_invoices=lambda *args, **kw: (_ for _ in ()).throw(RuntimeError('offline'))
        with self.assertRaises(RuntimeError): self.sync()
        self.assertEqual(self.payload()['missing_drafts'],[])

    def test_missing_issued_invoice_or_empty_series_never_hides_draft(self):
        self.rows=[]
        self.sync()
        self.assertEqual(len(self.payload()['items']),2)
        self.client.get_invoice_series=lambda:[]
        self.sync()
        self.assertEqual(self.payload()['missing_drafts'],[])

    def test_empty_complete_period_can_archive_its_only_unsigned_draft(self):
        self.c.execute("UPDATE outgoing_source_invoices SET invoice_date='2026-07-31' WHERE id=?",(self.issued_id,))
        self.rows=[]
        self.sync()
        self.assertEqual(self.payload()['items'],[])
        self.assertEqual(len(self.payload()['missing_drafts']),1)

    def test_posted_record_and_legacy_connector_are_never_archived(self):
        self.c.execute("UPDATE outgoing_source_invoices SET stock_status='posted' WHERE id=?",(self.draft_id,))
        self.sync()
        self.assertEqual(self.payload()['missing_drafts'],[])
        self.c.execute("UPDATE outgoing_source_invoices SET stock_status='blocked' WHERE id=?",(self.draft_id,))
        self.client.config.api_mode='legacy'
        self.sync()
        self.assertEqual(self.payload()['missing_drafts'],[])

    def test_only_unsigned_unposted_same_tenant_series_and_period_can_be_archived(self):
        for n,changes,tenant in [
            (3,{'invoiceDate':'2026-07-31T00:00:00'},'TDP'),
            (4,{'invoiceSerial':'1C26OTHER','inv_invoiceSeries':'1C26OTHER'},'TDP'),
            (5,{},'OTHER'), (6,{'dateSign':'2026-08-31T20:00:00+07:00'},'TDP'),
            (7,{'invoiceNumber':819},'TDP')]:
            self.insert(dict(self.draft,id=f'00000000-0000-0000-0000-{n:012}',**changes),tenant)
        self.sync()
        self.assertEqual([r[0] for r in self.c.execute('SELECT invoice_id FROM outgoing_missing_drafts')],[self.draft_id])

    def test_changing_total_or_duplicate_pages_do_not_archive(self):
        for duplicate in (False,True):
            def page(*args, **kw):
                if not kw['start']: return {'data':[self.issued],'total':2}
                return {'data':[self.issued],'total':2 if duplicate else 3}
            self.client.get_outgoing_invoices=page
            with self.assertRaises(InvoiceOutputSyncError): self.sync(page_size=1)
            self.assertEqual(self.payload()['missing_drafts'],[])


if __name__=='__main__': unittest.main()
