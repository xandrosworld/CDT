"""Cumulative approved demand less issued invoices; downloads are never issues."""
from collections import defaultdict
from io import BytesIO
from decimal import Decimal, ROUND_HALF_UP
import json

from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill
try:
    from .outgoing_source_scope import resolve_scope
    from .outgoing_line_policy import unit_issues
except ImportError:
    from outgoing_source_scope import resolve_scope
    from outgoing_line_policy import unit_issues


def _identity(r,local=False):
    return (str(r['issued_invoice_series'] if local else r['invoice_series']).strip().upper(),
            str(r['issued_invoice_number'] if local else r['invoice_number']).strip(),
            str((r['issued_invoice_date'] or r['invoice_date']) if local else r['invoice_date']).strip())


def _stock_only_remap_sources(conn):
    """Stock reclassification is not evidence of which sold item was invoiced.

    The Excel deficit workflow records inventory.output.remap. Audited repairs
    of an incorrect item identity use a separate event and keep their existing
    reconciliation behavior. Only still-active stock changes need this hold.
    """
    required={'audit_log','output_stock_remaps','output_stock_remap_parts','invoice_inventory_ledger'}
    tables={r[0] for r in conn.execute("SELECT name FROM sqlite_master WHERE type='table'")}
    if not required <= tables:
        return set()
    ledger_ids=set()
    incomplete_history=False
    for row in conn.execute("SELECT metadata_json FROM audit_log WHERE event_type='inventory.output.remap' AND status='ok'"):
        try:
            metadata=json.loads(row['metadata_json'])
            for change in metadata['changes']:
                ledger_id=int(change['ledger_id'])
                if ledger_id < 0:
                    part=conn.execute('SELECT ledger_id FROM output_stock_remap_parts WHERE id=?',(-ledger_id//2,)).fetchone()
                    if part:ledger_ids.add(part[0])
                    else:incomplete_history=True
                else:ledger_ids.add(ledger_id)
        except (ValueError,TypeError,KeyError):
            incomplete_history=True
    if not ledger_ids and not incomplete_history:
        return set()
    changed=conn.execute("""SELECT l.id,l.source_invoice_id FROM invoice_inventory_ledger l
        JOIN output_stock_remaps m ON m.ledger_id=l.id
        WHERE l.direction='output' AND l.source_invoice_table='outgoing_source_invoices'
          AND l.event_type='POST' AND l.status='posted' AND m.product_code<>l.product_code
        UNION SELECT l.id,l.source_invoice_id FROM invoice_inventory_ledger l
        JOIN output_stock_remap_parts p ON p.ledger_id=l.id
        WHERE l.direction='output' AND l.source_invoice_table='outgoing_source_invoices'
          AND l.event_type='POST' AND l.status='posted' AND p.qty>0 AND p.product_code<>l.product_code""")
    return {r['source_invoice_id'] for r in changed if incomplete_history or r['id'] in ledger_ids}


def _remap_review_warning(invoice,contractor):
    return {'contractor':contractor,'code':'stock_remap_requires_order_review','invoice_id':invoice['id'],
            'message':'Hóa đơn '+invoice['invoice_number']+' đã đổi mã trừ kho nội bộ. '
            'Cần đối chiếu mặt hàng trên hóa đơn với đơn đã bán trước khi trừ phần đã xuất; '
            'chưa tự đối trừ theo mã kho mới.'}


def issued_allocations(conn,asof='9999-12-31',*,external_quantities=None,source_order_allocations=None):
    """Prefer explicit order links; otherwise allocate mapped M-Invoice FIFO once."""
    orders=[dict(r) for r in conn.execute("SELECT o.* FROM orders o JOIN batches b ON b.id=o.batch_id WHERE b.status='approved' AND o.work_date<=? ORDER BY o.work_date,o.id",(asof,))]
    try:
        from .outgoing_amount_settlement import coverage
    except ImportError:
        from outgoing_amount_settlement import coverage
    money_orders, money_invoices, money_warnings = coverage(conn)
    orders = [o for o in orders if o['id'] not in money_orders]
    quantities=defaultdict(float);warnings=list(money_warnings);linked=set(money_invoices)
    remapped_sources=_stock_only_remap_sources(conn)
    sources=[dict(r) for r in conn.execute("SELECT * FROM outgoing_source_invoices WHERE source='minvoice' AND invoice_date<=? ORDER BY invoice_date,id",(asof,))]
    by_identity=defaultdict(list)
    for s in sources:by_identity[_identity(s)].append(s)
    for d in conn.execute("SELECT * FROM outgoing_invoice_drafts WHERE status='issued' AND COALESCE(issued_invoice_date,invoice_date)<=?",(asof,)):
        matching=by_identity.get(_identity(d,True),[])
        if matching and all(s['id'] in money_invoices for s in matching):
            continue
        if matching and any(s['source_status_class'] in ('cancelled','replaced','adjusted') for s in matching):
            warnings.append({'contractor':d['contractor'],'message':'Hóa đơn '+str(d['issued_invoice_number'])+' đã thay đổi trạng thái trên M-Invoice; cần đối chiếu.'})
            linked.update(s['id'] for s in matching);continue
        # A confirmed local issue carries exact order lineage. Same invoice is not counted again.
        if matching:
            warnings.extend(_remap_review_warning(s,d['contractor']) for s in matching if s['id'] in remapped_sources)
            local=defaultdict(float)
            for r in conn.execute('SELECT product_code,qty FROM outgoing_order_allocations WHERE draft_id=?',(d['id'],)):
                local[r['product_code']]+=r['qty']
            for s in matching:
                posted={r['product_code']:r['qty'] for r in conn.execute("""SELECT product_code,-SUM(qty_delta) qty FROM invoice_inventory_effective_ledger
                    WHERE direction='output' AND status='posted' AND source_invoice_table='outgoing_source_invoices'
                    AND source_invoice_id=? GROUP BY product_code""",(s['id'],))}
                if s['source_status_class']!='issued' or s['sync_status']!='synced' or set(local)!=set(posted) or any(abs(q-posted.get(code,0))>1e-8 for code,q in local.items()):
                    warnings.append({'contractor':d['contractor'],'message':'Hóa đơn '+str(d['issued_invoice_number'])+' chưa khớp trạng thái/lượng giữa xác nhận và M-Invoice; số chưa xuất cần đối chiếu.'})
        linked.update(s['id'] for s in matching)
        for r in conn.execute('SELECT order_id,qty FROM outgoing_order_allocations WHERE draft_id=?',(d['id'],)):
            quantities[r['order_id']]+=r['qty']
            if source_order_allocations is not None:
                for s in matching:
                    target=source_order_allocations.setdefault(s['id'],{})
                    target[r['order_id']]=target.get(r['order_id'],0)+r['qty']
    profiles=defaultdict(list)
    for r in conn.execute("SELECT contractor,tax_code FROM outgoing_buyer_profiles WHERE TRIM(COALESCE(tax_code,''))!=''"):
        profiles[r['tax_code'].strip().upper()].append(r['contractor'])
    indexed=defaultdict(list)
    for o in orders:indexed[(o['contractor'],o['product_code'],o['unit'].strip().casefold())].append(o)
    earliest=min((o['work_date'] for o in orders),default=asof)
    seen=set()
    for s in sources:
        if s['id'] in linked or s['invoice_date']<earliest or s['source_status_class'] in ('draft','cancelled','replaced'):continue
        party,scope_error=resolve_scope(conn,s,profiles)
        if scope_error:
            if s['source_status_class']=='issued':warnings.append({'contractor':'','message':scope_error})
            continue
        if party is None:continue
        if not any(o['contractor']==party and o['work_date']<=s['invoice_date'] for o in orders):continue
        identity=_identity(s)
        if identity in seen:
            warnings.append({'contractor':party,'message':'Hóa đơn '+s['invoice_number']+' trùng định danh nguồn; cần đối chiếu.'});continue
        seen.add(identity)
        if s['source_status_class']!='issued' or s['sync_status']!='synced' or s['stock_status'] not in ('posted','not_inventory'):
            warnings.append({'contractor':party,'message':'Hóa đơn '+s['invoice_number']+' chưa đủ đối chiếu mã/lượng để trừ khỏi đơn.'});continue
        if s['id'] in remapped_sources:
            warnings.append(_remap_review_warning(s,party));continue
        # The posted stock ledger already includes reviewed conversions and reversals.
        lines=conn.execute("""SELECT il.product_code,p.unit,-SUM(il.qty_delta) qty FROM invoice_inventory_effective_ledger il
            JOIN products p ON p.code=il.product_code WHERE il.direction='output' AND il.status='posted'
            AND il.source_invoice_table='outgoing_source_invoices' AND il.source_invoice_id=? GROUP BY il.product_code,p.unit""",(s['id'],)).fetchall()
        for line in lines:
            remaining=max(float(line['qty']),0)
            candidates=indexed.get((party,line['product_code'],line['unit'].strip().casefold()),[])
            for o in candidates:
                if o['work_date']>s['invoice_date']:continue
                need=max(o['actual_delivered']-o['customer_return_qty']-quantities[o['id']],0)
                take=min(need,remaining);quantities[o['id']]+=take;remaining-=take
                if take and source_order_allocations is not None:
                    target=source_order_allocations.setdefault(s['id'],{})
                    target[o['id']]=target.get(o['id'],0)+take
                if external_quantities is not None:external_quantities[o['id']]=external_quantities.get(o['id'],0)+take
                if remaining<=1e-8:break
            if remaining>1e-8:
                warnings.append({'contractor':party,'message':'Hóa đơn '+s['invoice_number']+': còn '+format(remaining,'.10g')+' '+line['unit']+' mã '+line['product_code']+' đã xuất chưa khớp đơn đã duyệt. Kiểm tra đúng mã hàng hoặc phạm vi đơn trước khi xuất tiếp.'})
    return dict(quantities),warnings


def unissued_payload(conn,asof,contractor='',*,respect_export_choices=False,start=''):
    # The cutoff selects order dates. An invoice issued later can settle those orders.
    source_allocations = {} if respect_export_choices else None
    issued,warnings=issued_allocations(conn,source_order_allocations=source_allocations)
    orders=[dict(r) for r in conn.execute("""SELECT o.* FROM orders o JOIN batches b ON b.id=o.batch_id
        WHERE b.status='approved' AND o.work_date<=? AND o.work_date>=? AND (?='' OR o.contractor=?) ORDER BY o.work_date,o.id""",(asof,start,contractor,contractor))]
    try:
        from .outgoing_amount_settlement import coverage, history, progress
    except ImportError:
        from outgoing_amount_settlement import coverage, history, progress
    money_orders, _, _ = coverage(conn)
    orders = [o for o in orders if o['id'] not in money_orders]
    try:
        from .outgoing_contractors import excluded_order_ids
    except ImportError:
        from outgoing_contractors import excluded_order_ids
    skipped_ids = excluded_order_ids(conn)
    reconciliation_groups = {}
    for o in orders:
        qty=max(o['actual_delivered']-o['customer_return_qty'],0)
        if qty<=1e-8:continue
        key=(o['contractor'],str(o['tax']))
        g=reconciliation_groups.setdefault(key,{'contractor':key[0],'tax':key[1],
            'order_rows':0,'fully_issued_rows':0,'remaining_rows':0,'invoices':{}})
        g['order_rows']+=1
        full=issued.get(o['id'],0)>=qty-1e-8
        g['fully_issued_rows']+=int(full)
        g['remaining_rows']+=int(not full)
    invoice_refs=defaultdict(list)
    if source_allocations:
        scoped={o['id']:o for o in orders}
        for s in conn.execute('''SELECT id,invoice_number,invoice_series,invoice_date
            FROM outgoing_source_invoices WHERE id IN (SELECT value FROM json_each(?))''',
            (json.dumps(list(source_allocations)),)):
            for oid,qty in source_allocations[s['id']].items():
                if oid not in scoped or qty<=1e-8:continue
                ref={'number':s['invoice_number'],'series':s['invoice_series'],'date':s['invoice_date'],'qty':qty}
                invoice_refs[oid].append(ref)
                o=scoped[oid]
                if (o['contractor'],str(o['tax'])) in reconciliation_groups:
                    reconciliation_groups[(o['contractor'],str(o['tax']))]['invoices'][s['id']]=ref
    excluded=set()
    line_choices=[]
    if respect_export_choices:
        try:
            from .outgoing_contractors import excluded_codes, selected_orders, line_choices_payload
        except ImportError:
            from outgoing_contractors import excluded_codes, selected_orders, line_choices_payload
        excluded=excluded_codes(conn)
        line_choices=line_choices_payload(conn,asof,contractor,orders=orders,issued=issued)
        orders=selected_orders(conn,orders)
        warnings=[w for w in warnings if w['contractor'] not in excluded]
    try:
        from .outgoing_waiting import waiting_readiness
        from .outgoing_readiness import canonical_available_stock
        from .outgoing_pending import explain_pending
    except ImportError:
        from outgoing_waiting import waiting_readiness
        from outgoing_readiness import canonical_available_stock
        from outgoing_pending import explain_pending
    stock=canonical_available_stock(conn)
    invoice_names={r['product_code']:r['invoice_name'] for r in conn.execute("SELECT product_code,invoice_name FROM outgoing_product_names WHERE TRIM(invoice_name)!=''")}
    ready,stock_warnings=waiting_readiness(conn,orders,issued,stock=stock)
    units=unit_issues(conn,orders)
    warnings.extend(stock_warnings)
    drafted={r['order_id']:r['qty'] for r in conn.execute("SELECT l.order_id,SUM(l.qty) qty FROM outgoing_order_allocations l JOIN outgoing_invoice_drafts d ON d.id=l.draft_id WHERE d.status='draft' GROUP BY l.order_id")}
    grouped={};details=[]
    for o in orders:
        q=max(o['actual_delivered']-o['customer_return_qty'],0)
        if q<=1e-8:continue
        done=issued.get(o['id'],0);left=max(q-done,0)
        r={'order_id':o['id'],'work_date':o['work_date'],'contractor':o['contractor'],'product_code':o['product_code'],
           'product_name':o['product_name'],'invoice_name':invoice_names.get(o['product_code'],o['product_name']),'unit':o['unit'],'approved_qty':q,'issued_qty':done,'drafted_qty':drafted.get(o['id'],0),'unissued_qty':left,'unit_price':o['sell_price'],
           'ready_qty':min(ready.get(o['id'],0),left) if not any(not w['contractor'] or w['contractor']==o['contractor'] for w in warnings) else 0,
           'tax':o['tax'],'invoice_nature':str(o.get('invoice_nature') or '1')}
        r['waiting_qty']=max(left-r['ready_qty'],0)
        r['pending_reason']=units[o['id']]['message'] if o['id'] in units else ''
        r['needs_conversion']=o['id'] in units
        r['conversion_reason']=r['pending_reason']
        r['unissued_amount']=int((Decimal(str(left))*Decimal(str(o['sell_price']))).quantize(Decimal('1'),rounding=ROUND_HALF_UP))
        details.append(r)
        key=(o['contractor'],o['product_code'],o['unit'].strip().casefold())
        g=grouped.setdefault(key,{**r,'approved_qty':0,'issued_qty':0,'drafted_qty':0,'unissued_qty':0,'ready_qty':0,'waiting_qty':0,'first_date':o['work_date'],'last_date':o['work_date']})
        for field in ('approved_qty','issued_qty','drafted_qty','unissued_qty','ready_qty','waiting_qty'):g[field]+=r[field]
        g['last_date']=o['work_date']
    pending=explain_pending(conn,orders,details,units,warnings,stock)
    for r in details:
        r['export_skipped'] = r['order_id'] in skipped_ids
        if r['export_skipped']:
            r['pending_reason'] = 'Đã bỏ chọn — chưa xuất. Giữ nguyên doanh thu và công nợ; có thể chọn lại khi cần.'
            r['pending_codes'] = ['user_skipped']
    detail_by_id={r['order_id']:r for r in details}
    for r in line_choices:
        d=detail_by_id.get(r['order_id'],{})
        r.update(pending_reason=d.get('pending_reason',''),pending_codes=d.get('pending_codes',[]),
                 issued_invoices=invoice_refs.get(r['order_id'],[]))
    skipped_details=[]
    for r in line_choices:
        if r['enabled'] or r['order_id'] in detail_by_id:continue
        skipped_details.append({'order_id':r['order_id'],'batch_id':r['batch_id'],'contractor':r['contractor'],
            'work_date':r['date'],'product_code':r['product_code'],'product_name':r['invoice_name'],'invoice_name':r['invoice_name'],
            'unit':r['unit'],'tax':r['tax'],'unissued_qty':r['qty'],'unit_price':r['price'],
            'approved_qty':r['approved_qty'],'issued_qty':r['issued_qty'],'drafted_qty':0,'ready_qty':0,'waiting_qty':r['qty'],
            'unissued_amount':int((Decimal(str(r['qty']))*Decimal(str(r['price']))).quantize(Decimal('1'),rounding=ROUND_HALF_UP)),
            'export_skipped':True,'pending_reason':'Đã bỏ chọn — chưa xuất. Giữ nguyên doanh thu và công nợ; có thể chọn lại khi cần.',
            'pending_codes':['user_skipped']})
    for g in reconciliation_groups.values():
        g['invoices']=sorted(({k:v for k,v in r.items() if k!='qty'} for r in g['invoices'].values()),
                             key=lambda r:(r['date'],r['number']))
    reasons=defaultdict(list)
    for r in details:
        if r['pending_reason']:
            reasons[(r['contractor'],r['product_code'],r['unit'].strip().casefold())].append(r['pending_reason'])
    for key,r in grouped.items():r['pending_reason']=' · '.join(dict.fromkeys(reasons[key]))
    rows=[r for r in grouped.values() if r['unissued_qty']>1e-8]
    rows.sort(key=lambda r:(r['contractor'],r['product_name'],r['product_code']))
    totals=defaultdict(lambda:defaultdict(float))
    for r in grouped.values():
        for field in ('approved_qty','issued_qty','drafted_qty','unissued_qty','ready_qty','waiting_qty'):totals[r['unit'].strip().casefold()][field]+=r[field]
    warnings=[w for w in warnings if not contractor or not w['contractor'] or w['contractor']==contractor]
    try:
        from .outgoing_signed_stock_review import signed_stock_issues
    except ImportError:
        from outgoing_signed_stock_review import signed_stock_issues
    return {'from':start,'asof':asof,'contractor':contractor,'excluded_contractors':sorted(excluded),'rows':rows,'details':[r for r in details if r['unissued_qty']>1e-8],
            'amount_settlements': [r for r in history(conn, contractor) if r['date_from'] <= asof and r['date_to'] >= start],
            'amount_progress': progress(conn, contractor, start, asof) if contractor and start else None,
            'line_choices':line_choices,'skipped_details':skipped_details,
            'reconciliation_groups':list(reconciliation_groups.values()),
            'pending_rows':pending,'pending_order_rows':sum(r['waiting_qty']>1e-8 for r in details),
            'signed_stock_issues':signed_stock_issues(conn,contractor),
            'held_line_issues':[{**units[o['id']],'contractor':o['contractor'],'work_date':o['work_date']} for o in orders if o['id'] in units and o['actual_delivered']-o['customer_return_qty']-issued.get(o['id'],0)>1e-8],
            'totals_by_unit':dict(totals),'source_order_rows':len(details),'unissued_order_rows':sum(r['unissued_qty']>1e-8 for r in details),
            'warnings':warnings,'reconciliation_complete':not warnings,
            'policy':'Cộng dồn đơn đã duyệt đến ngày chọn, trừ lượng hóa đơn đã phát hành được đồng bộ hoặc xác nhận đến hiện tại, kể cả hóa đơn phát hành sau ngày đơn. Tải file và tạo nháp không làm giảm lượng chưa xuất.'}


def unissued_workbook(payload):
    w=Workbook();s=w.active;s.title='Chua xuat cong don'
    if payload.get('portion')=='waiting':
        s.title='Hang con cho'
        s.append(['Nhà thầu','Mã hàng','Tên hàng','ĐVT','Ngày đơn đầu','Ngày đơn cuối','Lượng còn chờ','Lý do còn chờ'])
        for r in payload['pending_rows']:
            s.append([r[k] for k in ('contractor','product_code','product_name','unit','first_date','last_date','waiting_qty','pending_reason')])
        s=w.create_sheet('Doi chieu dong con cho')
    headers=['Nhà thầu','Mã hàng','Tên hàng','ĐVT','Ngày đơn đầu','Ngày đơn cuối','Lượng đã duyệt','Đã phát hành','Tổng lượng đang giữ','Chưa xuất hóa đơn','Đã đủ điều kiện, giữ chờ xuất','Chưa đủ điều kiện / chờ cộng lẻ']
    s.append(headers+['Lý do còn chờ'])
    for r in payload['rows']:s.append([r[k] for k in ('contractor','product_code','product_name','unit','first_date','last_date','approved_qty','issued_qty','drafted_qty','unissued_qty','ready_qty','waiting_qty')]+[r.get('pending_reason','')])
    s=w.create_sheet('Chi tiet theo ngay');s.append(['Dòng đơn','Ngày đơn','Nhà thầu','Mã','Tên xuất hóa đơn','ĐVT đơn','Đã duyệt','Đã phát hành','Tổng lượng đang giữ','Chưa xuất','Giá trên đơn','Đủ điều kiện, giữ chờ xuất','Chưa đủ điều kiện / chờ cộng lẻ'])
    s.cell(1,14,'Lý do còn chờ')
    s.cell(1,15,'Thuế');s.cell(1,16,'Tiền hàng chưa xuất');s.cell(1,17,'Cần quy đổi / đối chiếu ĐVT')
    for r in payload['details']:
        s.append([r.get('invoice_name',r['product_name']) if k=='product_name' else r[k] for k in ('order_id','work_date','contractor','product_code','product_name','unit','approved_qty','issued_qty','drafted_qty','unissued_qty','unit_price','ready_qty','waiting_qty')]+[r.get('pending_reason',''),r['tax'],r['unissued_amount'],r.get('conversion_reason','')])
    note=w.create_sheet('Ghi chu');note.append(['Từ ngày đơn',payload.get('from') or 'Tất cả']);note.append(['Đến ngày đơn',payload['asof']]);note.append(['Cách tính',payload['policy']])
    note.append(['Đối chiếu M-Invoice','Nếu đã ký bên ngoài, đồng bộ hóa đơn hoặc xác nhận đúng số hóa đơn đã phát hành trước khi lập tiếp.'])
    for warning in payload['warnings']:note.append(['Cần đối chiếu',warning['message']])
    for issue in payload.get('held_line_issues',[]):note.append(['Dòng giữ riêng',issue['work_date'],issue['contractor'],issue['message']])
    if payload.get('signed_stock_issues'):
        review=w.create_sheet('Hoa don da ky can doi chieu')
        review.append(['Hóa đơn','Ngày hóa đơn','Thời điểm ký','Nhà thầu','Mã hàng','Tên hàng','ĐVT','Lượng hóa đơn đã ký','Tồn đầu','Đầu vào','Tồn hiện tại','Cần xử lý'])
        for r in payload['signed_stock_issues']:review.append([r[k] for k in ('invoice_number','invoice_date','signed_at','contractor','product_code','product_name','unit','signed_qty','opening_qty','input_qty','closing_qty','message')])
    if payload.get('amount_settlements'):
        settled = w.create_sheet('Da doi tru theo tien')
        settled.append(['Bản đối trừ', 'Nhà thầu', 'Từ ngày đơn', 'Đến ngày đơn', 'Tổng tiền gồm thuế', 'Hóa đơn đã ký', 'Người xác nhận', 'Trạng thái'])
        for r in payload['amount_settlements']:
            settled.append([r['id'], r['contractor'], r['date_from'], r['date_to'], r['amount'],
                            ', '.join(i['invoice_series']+'/'+i['invoice_number'] for i in r['invoices']),
                            r['actor'], 'Cần đối chiếu lại' if r['needs_review'] else 'Đã đối trừ theo tiền; không xác nhận khớp mặt hàng'])
    if payload.get('skipped_details'):
        skipped=w.create_sheet('Da bo chon chua xuat');skipped.append(['Ngày đơn','Nhà thầu','Mã hàng','Tên hàng','ĐVT','Lượng chưa xuất','Tiền hàng','Lý do'])
        for r in payload['skipped_details']:skipped.append([r[k] for k in ('work_date','contractor','product_code','invoice_name','unit','unissued_qty','unissued_amount','pending_reason')])
    for s in w:
        s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
        for cell in s[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='163247')
        for col in s.columns:s.column_dimensions[col[0].column_letter].width=min(55,max(16,max(len(str(c.value or '')) for c in col)+2))
    try:
        from .template_workbook import safe_workbook_bytes
        from .document_preview import white_print_style
    except ImportError:
        from template_workbook import safe_workbook_bytes
        from document_preview import white_print_style
    try:
        white_print_style(w)
        for number,r in enumerate(payload['details'],start=2):
            if r.get('needs_conversion'):
                for cell in w['Chi tiet theo ngay'][number]:cell.font=Font(color='B42318')
        return BytesIO(safe_workbook_bytes(w,apply_print_style=False))
    finally:
        w.close()
