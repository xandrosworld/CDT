from __future__ import annotations

import io
import unittest
import zipfile

from openpyxl import load_workbook

try:
    from invoice_delivery_statement import (
        InvoiceDeliveryStatementError,
        invoice_delivery_statement_workbook,
    )
except ImportError:  # pragma: no cover
    from .invoice_delivery_statement import (
        InvoiceDeliveryStatementError,
        invoice_delivery_statement_workbook,
    )


class InvoiceDeliveryStatementTests(unittest.TestCase):
    @staticmethod
    def draft(draft_id=1, *, contractor="NT-A", status="issued", total=216, number="0000001"):
        return {
            "id": draft_id,
            "contractor": contractor,
            "status": status,
            "issued_invoice_date": f"2026-09-{draft_id:02d}",
            "issued_invoice_series": "1C26TDP",
            "issued_invoice_number": number,
            "subtotal": 200,
            "tax_amount": 16,
            "total_amount": total,
            "verification_source": "synced_issued_source" if draft_id == 2 else "local_issued_confirmation",
            "source_invoice_id": 900 + draft_id if draft_id == 2 else None,
        }

    @staticmethod
    def line(draft_id=1, *, line_id=1, order_id=11, product_name="Hàng giao"):
        return {
            "id": line_id,
            "draft_id": draft_id,
            "order_id": order_id,
            "work_date": f"2026-09-{draft_id:02d}",
            "kitchen": "BẾP A",
            "product_code": f"P-{draft_id}",
            "product_name": product_name,
            "unit": "kg",
            "qty": 2,
            "unit_price": 100,
            "amount": 200,
            "tax": "8%",
        }

    def test_two_issued_rounds_same_contractor_are_static_and_reconciled(self):
        drafts = [self.draft(), self.draft(2, number="0000002")]
        lines = [self.line(), self.line(2, line_id=2, order_id=12)]
        workbook = invoice_delivery_statement_workbook(
            lines, drafts, contractor="nt-a", scope_id="SCOPE-ABC",
        )
        try:
            self.assertEqual(workbook.sheetnames, ["Bảng kê giao hàng", "Đối chiếu hóa đơn"])
            detail = workbook["Bảng kê giao hàng"]
            reconcile = workbook["Đối chiếu hóa đơn"]
            self.assertEqual(detail.max_row, 17)
            self.assertEqual(reconcile.max_row, 5)
            self.assertEqual([reconcile.cell(row, 10).value for row in (4, 5)], [0, 0])
            self.assertEqual(reconcile.cell(5, 11).value, "synced_issued_source")
            self.assertEqual(reconcile.cell(5, 12).value, 902)
            self.assertIn("BẢNG TỔNG HỢP GIAO NHẬN", detail["A5"].value)
            self.assertEqual("STT", detail["A10"].value)
            self.assertEqual("Thanh toán", detail["J10"].value)
            self.assertTrue(detail.column_dimensions["K"].hidden)
            self.assertTrue(reconcile.column_dimensions["L"].hidden)
            self.assertEqual(reconcile.sheet_state, "hidden")
            self.assertEqual(str(detail.page_setup.paperSize), str(detail.PAPERSIZE_A4))
            self.assertEqual(detail.page_setup.orientation, "portrait")
            self.assertEqual(detail.print_title_rows, "$10:$10")
            self.assertFalse(any(
                cell.data_type == "f"
                for sheet in workbook.worksheets for row in sheet.iter_rows() for cell in row
            ))
        finally:
            workbook.close()

    def test_consolidated_line_keeps_each_source_order_once(self):
        from copy import deepcopy
        lines = [{**self.line(order_id=11), 'qty': 0.5, 'amount': 50},
                 {**self.line(order_id=12), 'qty': 1.5, 'amount': 150, 'work_date': '2026-08-31'}]
        before = deepcopy(lines)
        book = invoice_delivery_statement_workbook(lines, [self.draft()])
        try:
            detail = book['Bảng kê giao hàng']
            self.assertEqual(2, sum(detail.cell(r, 5).value for r in (11, 12)))
            self.assertEqual(200, sum(detail.cell(r, 7).value for r in (11, 12)))
            self.assertEqual(216, sum(detail.cell(r, 10).value for r in (11, 12)))
            self.assertEqual({11, 12}, {detail.cell(r, 18).value for r in (11, 12)})
            self.assertEqual(0, book['Đối chiếu hóa đơn'].cell(4, 10).value)
            self.assertEqual(before, lines)
        finally:
            book.close()

    def test_true_duplicate_allocation_is_still_rejected(self):
        line = self.line()
        # Even if totals were also doubled, repeated data must never pass.
        draft = {**self.draft(total=432), 'subtotal': 400, 'tax_amount': 32}
        with self.assertRaisesRegex(InvoiceDeliveryStatementError, 'bị lặp'):
            invoice_delivery_statement_workbook([line, dict(line)], [draft])
        line.pop('order_id')
        with self.assertRaisesRegex(InvoiceDeliveryStatementError, 'bị lặp'):
            invoice_delivery_statement_workbook([line, dict(line)], [draft])

    def test_unissued_and_mixed_contractor_scopes_are_rejected(self):
        with self.assertRaises(InvoiceDeliveryStatementError) as unissued:
            invoice_delivery_statement_workbook([self.line()], [self.draft(status="draft")])
        self.assertEqual(unissued.exception.code, "delivery_statement_unissued_invoice")

        drafts = [self.draft(), self.draft(2, contractor="NT-B", number="0000002")]
        with self.assertRaises(InvoiceDeliveryStatementError) as mixed:
            invoice_delivery_statement_workbook(
                [self.line(), self.line(2, line_id=2, order_id=12)], drafts,
            )
        self.assertEqual(mixed.exception.code, "delivery_statement_contractor_conflict")

    def test_total_mismatch_is_blocked_without_auto_balance(self):
        with self.assertRaises(InvoiceDeliveryStatementError) as mismatch:
            invoice_delivery_statement_workbook([self.line()], [self.draft(total=218)])
        self.assertEqual(mismatch.exception.code, "delivery_statement_total_mismatch")

    def test_formula_like_text_is_literal_and_file_has_no_external_links(self):
        draft = self.draft()
        draft["issued_invoice_series"] = "=DANGEROUS"
        line = self.line(product_name="+SUM(A1:A2)")
        workbook = invoice_delivery_statement_workbook([line], [draft], scope_id="STATIC")
        stream = io.BytesIO()
        try:
            workbook.save(stream)
        finally:
            workbook.close()
        stream.seek(0)
        reopened = load_workbook(stream, data_only=False, keep_links=False)
        try:
            detail = reopened["Bảng kê giao hàng"]
            self.assertEqual(detail.cell(11, 12).value, "'=DANGEROUS")
            self.assertEqual(detail.cell(11, 3).value, "'+SUM(A1:A2)")
            self.assertFalse(any(
                cell.data_type == "f" or cell.hyperlink is not None
                for sheet in reopened.worksheets for row in sheet.iter_rows() for cell in row
            ))
        finally:
            reopened.close()
        with zipfile.ZipFile(io.BytesIO(stream.getvalue())) as archive:
            self.assertFalse(any(name.startswith("xl/externalLinks/") for name in archive.namelist()))


if __name__ == "__main__":
    unittest.main()
