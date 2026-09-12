import os
import tempfile
import unittest
from unittest.mock import patch
from pathlib import Path

from openpyxl import Workbook, load_workbook
from pypdf import PdfReader
from reportlab.lib.pagesizes import A4, landscape, letter
from reportlab.pdfgen.canvas import Canvas

try:
    from .excel_print_renderer import (
        ExcelPrintError,
        build_excel_pdf_bundle,
        verify_excel_pdf,
        _export_libreoffice_sheets,
    )
except ImportError:  # pragma: no cover - direct file invocation
    from excel_print_renderer import ExcelPrintError, build_excel_pdf_bundle, verify_excel_pdf


def _small_pdf(path: Path, pagesize) -> None:
    canvas = Canvas(str(path), pagesize=pagesize)
    canvas.drawString(36, 36, "TDP")
    canvas.save()


class ExcelPrintRendererTests(unittest.TestCase):
    def test_calc_copies_remove_other_print_ranges_but_keep_formula_dependencies(self):
        with tempfile.TemporaryDirectory(prefix='tdp_calc_copy_') as folder:
            root = Path(folder)
            source = root / 'source.xlsx'
            book = Workbook()
            first = book.active; first.title = 'First'
            first['A1'] = '=Internal!A1'; first.print_area = 'A1:D10'
            second = book.create_sheet('Second'); second['A1'] = 'Second'; second.print_area = 'A1:D10'
            hidden = book.create_sheet('Internal'); hidden['A1'] = 42
            hidden.sheet_state = 'hidden'; hidden.print_area = 'A1:D10'; hidden.print_title_rows = '1:2'
            book.save(source); book.close()
            before = source.read_bytes()
            exported = []
            def convert(command, **kwargs):
                path = Path(command[-1]); copy = load_workbook(path)
                try:
                    selected = copy.active
                    exported.append(selected.title)
                    self.assertTrue(selected.print_area)
                    self.assertEqual(copy['First']['A1'].value, '=Internal!A1')
                    self.assertEqual(copy['Internal']['A1'].value, 42)
                    for sheet in copy:
                        if sheet != selected:
                            self.assertEqual(sheet.sheet_state, 'hidden')
                            self.assertFalse(sheet.print_area)
                finally: copy.close()
                _small_pdf(path.with_suffix('.pdf'), A4)
                return type('Result', (), {'returncode': 0})()
            with patch(_export_libreoffice_sheets.__module__ + '.shutil.which', return_value='soffice'), \
                 patch(_export_libreoffice_sheets.__module__ + '.subprocess.run', side_effect=convert):
                result = _export_libreoffice_sheets([{'path':source,'title':'Fixture','document_type':'test'}],
                                                   paper='A4',render_dir=root)
            self.assertEqual(exported, ['First', 'Second'])
            self.assertEqual(len(result), 2)
            self.assertEqual(source.read_bytes(), before)

    def test_pdf_verifier_accepts_portrait_and_landscape_a4(self):
        with tempfile.TemporaryDirectory(prefix="tdp_excel_pdf_verify_") as temp_name:
            root = Path(temp_name)
            portrait = root / "portrait.pdf"
            horizontal = root / "landscape.pdf"
            _small_pdf(portrait, A4)
            _small_pdf(horizontal, landscape(A4))
            self.assertEqual(verify_excel_pdf(portrait, paper="A4")["orientations"], ["portrait"])
            self.assertEqual(verify_excel_pdf(horizontal, paper="A4")["orientations"], ["landscape"])

    def test_pdf_verifier_rejects_wrong_paper(self):
        with tempfile.TemporaryDirectory(prefix="tdp_excel_pdf_wrong_") as temp_name:
            output = Path(temp_name) / "letter.pdf"
            _small_pdf(output, letter)
            with self.assertRaises(ExcelPrintError):
                verify_excel_pdf(output, paper="A4")

    @unittest.skipUnless(os.name == "nt", "Microsoft Excel COM is Windows-only")
    def test_excel_exports_visible_artwork_and_omits_hidden_sheet(self):
        try:
            import win32com.client  # noqa: F401
        except ImportError:
            self.skipTest("pywin32 is not installed")
        with tempfile.TemporaryDirectory(prefix="tdp_excel_print_integration_") as temp_name:
            root = Path(temp_name)
            source = root / "source.xlsx"
            target = root / "bundle.pdf"
            workbook = Workbook()
            visible = workbook.active
            visible.title = "Phiếu in"
            visible["A1"] = "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM"
            visible["A2"] = "Nội dung kiểm tra"
            visible.print_area = "A1:D20"
            visible.page_setup.paperSize = visible.PAPERSIZE_A4
            visible.page_setup.orientation = visible.ORIENTATION_LANDSCAPE
            visible.sheet_properties.pageSetUpPr.fitToPage = True
            visible.page_setup.fitToWidth = 1
            hidden = workbook.create_sheet("Dữ liệu nội bộ")
            hidden["A1"] = "KHÔNG ĐƯỢC IN"
            hidden.sheet_state = "hidden"
            workbook.save(source)
            workbook.close()

            try:
                manifest = build_excel_pdf_bundle(
                    [{"path": source, "document_type": "test", "title": "Phiếu kiểm tra"}],
                    target,
                    paper="A4",
                    generated_at="2026-09-03T00:00:00",
                )
            except ExcelPrintError as exc:
                if "Không mở được Microsoft Excel" in str(exc):
                    self.skipTest(str(exc))
                raise
            self.assertTrue(target.is_file())
            self.assertEqual(manifest["format_version"], "tdp-excel-artwork-pdf-v5-plain-print")
            self.assertEqual(manifest["section_count"], 1)
            self.assertEqual(manifest["sections"][0]["sheet"], "Phiếu in")
            self.assertEqual(len(PdfReader(str(target)).pages), manifest["pages"])


if __name__ == "__main__":
    unittest.main(verbosity=2)
