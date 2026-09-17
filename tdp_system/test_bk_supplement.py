import io
import sqlite3
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook
from . import server, bk_supplement, inventory_period_close as closing
from . import test_bk_import as fixture
from .bk_draft import draft_rows
from .bk_import import build_bk_import_template
from .invoice_inventory import invoice_stock_rows


class SupplementTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        fixture.BKImportTests.setUpClass.__func__(cls)
        server.MASTER_SOURCE=Path(__file__).resolve().parent.parent/'Em Thành.xlsx'

    @classmethod
    def tearDownClass(cls):fixture.BKImportTests.tearDownClass.__func__(cls)

    def setUp(self):
        fixture.BKImportTests.setUp(self)
        with server.db() as c:
            c.execute('DELETE FROM inventory_period_closures')
            c.execute('DELETE FROM people')
            c.execute("INSERT INTO people(name,cccd,address,issue_date,issue_place) VALUES('Seller Test',?,'Test address','01/01/2020','Test place')",('1'*12,))
            c.execute("UPDATE products SET tax='KKKNT'")

    def body(self,reference='SUP-01',qty=2,cost=100):
        return {'from':'2026-08-01','to':'2026-08-31','document_date':'2026-08-01','reference':reference,
            'rows':[{'product_code':'BK-P1','qty':qty,'unit_cost':cost,'source_party':'Seller Test','note':''}]}

    def preview(self,body=None):
        r=self.client.post('/api/bk-import/draft/preview',json=body or self.body())
        self.assertEqual(200,r.status_code,r.get_json());return r.get_json()

    def confirm(self,p,**overrides):
        return self.client.post('/api/bk-import/draft/confirm',json={
            'token':p['import_token'],'confirmed':True,'actor':'Test reviewer',**overrides})

    def snapshot(self):
        with server.db() as c:
            return {t:[tuple(r) for r in c.execute('SELECT * FROM '+t+' ORDER BY rowid')] for t in
                ['orders','batches','receivable_ledger_lines','payable_ledger_lines','outgoing_invoice_drafts']}

    def stock(self):
        with server.db() as c:return {r['product_code']:r['closing_qty'] for r in invoice_stock_rows(c,as_of='2026-09-17')}

    def test_golden_summary_receipt_post_and_excel_retry_without_sales_or_debt(self):
        before=self.snapshot();p=self.preview()
        self.assertEqual(['bảng kê tổng','biên nhận'],[s['name'] for s in p['sheets']])
        self.assertEqual({},self.stock())
        excel=self.client.get('/api/documents/'+p['token']+'/excel')
        w=load_workbook(io.BytesIO(excel.data),data_only=True)
        self.assertEqual(['bảng kê tổng','biên nhận'],w.sheetnames)
        self.assertEqual('Từ ngày 01/08/2026 đến ngày 01/08/2026',w['bảng kê tổng']['A2'].value)
        self.assertNotIn('$52',str(w['bảng kê tổng'].print_area))
        self.assertEqual('Seller Test',w['biên nhận']['D8'].value);w.close()
        self.assertEqual(409,self.confirm(p,confirmed=False).status_code)
        with patch('tdp_system.outgoing_waiting.refresh_after_change') as refresh:
            result=self.confirm(p)
            refresh.assert_not_called()
        self.assertEqual(200,result.status_code,result.get_json())
        self.assertEqual(2,self.stock()['BK-P1'])
        self.assertTrue(self.confirm(p).get_json()['idempotent'])
        upload=self.client.post('/api/bk-import/draft/excel',json=self.body())
        loaded=self.client.post('/api/bk-import/draft/file',data={'file':(io.BytesIO(upload.data),'changed-name.xlsx')})
        self.assertEqual(200,loaded.status_code,loaded.get_json())
        p2=self.preview({**self.body(),**{k:loaded.get_json()[k] for k in ['document_date','reference','rows']}})
        self.assertTrue(p2['already_posted']);self.assertTrue(self.confirm(p2).get_json()['idempotent'])
        self.assertEqual(2,self.stock()['BK-P1']);self.assertEqual(before,self.snapshot())

    def test_identity_and_amount_limit_rechecked_across_documents(self):
        self.assertEqual(200,self.confirm(self.preview(self.body(qty=3,cost=1000000))).status_code)
        r=self.client.post('/api/bk-import/draft/preview',json=self.body('SUP-02',3,1000000))
        self.assertEqual(400,r.status_code);self.assertIn('5.000.000',r.get_json()['error'])
        p=self.body('SUP-03');p['rows'][0]['source_party']='Unknown'
        self.assertEqual(400,self.client.post('/api/bk-import/draft/preview',json=p).status_code)

    def test_second_user_changes_stock_invalidates_confirmation(self):
        p=self.preview();other=self.preview(self.body('SUP-02',1))
        self.assertEqual(200,self.confirm(other).status_code)
        r=self.confirm(p);self.assertEqual(409,r.status_code);self.assertIn('vừa thay đổi',r.get_json()['error'])
        self.assertEqual(1,self.stock()['BK-P1'])

    def create_closed_deficit(self):
        with server.db() as c:
            c.execute("INSERT INTO inventory_transactions(txn_date,product_code,qty_in,qty_out,unit_cost,source_type,source_id,source_line,status,created_at,updated_at) VALUES('2026-08-01','BK-P1',0,2,100,'OPENING','2026-08','BK-P1','posted',?,?)",(fixture.NOW,fixture.NOW))
            p=closing.inventory_period_close_preview(c,'2026-08',today=date(2026,9,17))
            closing.close_inventory_period(c,'2026-08',expected_source_hash=p['source_hash'],expected_target_hash=p['target_hash'],timestamp=fixture.NOW,today=date(2026,9,17),actor='Test')

    def test_august_post_rebuilds_september_opening_and_keeps_ledgers(self):
        self.create_closed_deficit();before=self.snapshot();self.assertEqual(-2,self.stock()['BK-P1'])
        p=self.preview();self.assertEqual(['2026-08'],p['rebuild_periods'])
        r=self.confirm(p);self.assertEqual(200,r.status_code,r.get_json())
        self.assertEqual(['2026-08'],r.get_json()['rebuilt_periods'])
        self.assertEqual(0,self.stock().get('BK-P1',0));self.assertEqual(before,self.snapshot())
        self.assertEqual([{'product_code':'BK-P1','closing_qty':0}],r.get_json()['stock'])
        with server.db() as c:
            self.assertEqual('closed',c.execute("SELECT status FROM inventory_period_closures WHERE period='2026-08'").fetchone()[0])
        self.assertTrue(self.confirm(p).get_json()['idempotent']);self.assertEqual(0,self.stock().get('BK-P1',0))

    def test_failed_carry_rebuild_rolls_back_purchase_and_opening(self):
        self.create_closed_deficit();p=self.preview()
        with patch.object(closing,'close_inventory_period',side_effect=ValueError('Test rebuild failed')):
            r=self.confirm(p);self.assertEqual(409,r.status_code)
        self.assertEqual(-2,self.stock()['BK-P1'])
        with server.db() as c:
            self.assertEqual(0,c.execute('SELECT COUNT(*) FROM bk_import_documents').fetchone()[0])
            self.assertEqual('closed',c.execute("SELECT status FROM inventory_period_closures WHERE period='2026-08'").fetchone()[0])

    def test_manual_opening_cannot_be_overwritten(self):
        self.create_closed_deficit()
        with server.db() as c:c.execute("UPDATE inventory_transactions SET qty_out=9 WHERE source_type='OPENING' AND source_id='2026-09' AND product_code='BK-P1'")
        r=self.client.post('/api/bk-import/draft/preview',json=self.body())
        self.assertEqual(400,r.status_code);self.assertIn('không khớp',r.get_json()['error'])

    def test_multiple_sellers_get_separate_receipts(self):
        with server.db() as c:
            c.execute("INSERT INTO people(name,cccd,address,issue_date,issue_place) VALUES('Seller Two',?,'Test address','01/01/2020','Test place')",('2'*12,))
        body=self.body();body['rows'].append({'product_code':'BK-P2','qty':3,'unit_cost':50,'source_party':'Seller Two'})
        p=self.preview(body)
        self.assertEqual(['bảng kê tổng','biên nhận','biên nhận 02'],[s['name'] for s in p['sheets']])
        self.assertEqual(350,p['totals']['amount'])

    def test_multiple_generated_months_rebuild_in_order(self):
        self.create_closed_deficit()
        with server.db() as c:
            p=closing.inventory_period_close_preview(c,'2026-09',today=date(2026,10,17))
            closing.close_inventory_period(c,'2026-09',expected_source_hash=p['source_hash'],expected_target_hash=p['target_hash'],timestamp=fixture.NOW,today=date(2026,10,17),actor='Test')
        p=self.preview();self.assertEqual(['2026-08','2026-09'],p['rebuild_periods'])
        with patch.object(closing,'date',wraps=date) as calendar:
            calendar.today.return_value=date(2026,10,17)
            r=self.confirm(p)
        self.assertEqual(200,r.status_code,r.get_json())
        with server.db() as c:
            stock={r['product_code']:r['closing_qty'] for r in invoice_stock_rows(c,as_of='2026-10-17',include_zero=True)}
            self.assertEqual(0,stock['BK-P1'])
            self.assertEqual(['closed','closed'],[r[0] for r in c.execute('SELECT status FROM inventory_period_closures ORDER BY period')])


if __name__=='__main__':unittest.main()
