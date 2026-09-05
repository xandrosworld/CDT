"""Dynamic golden-template monthly report by customer group, contractor and kitchen."""

from __future__ import annotations

import re
import unicodedata
from collections import defaultdict
from copy import copy
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Callable, Iterable, Mapping

from openpyxl.cell.cell import MergedCell
from openpyxl.worksheet.page import PageMargins

try:
    from purchase_summary_export import EM_THANH_SHA256
    from template_workbook import assert_workbook_safe, clone_template_workbook, write_literal
except ImportError:  # pragma: no cover - package import path
    from .purchase_summary_export import EM_THANH_SHA256
    from .template_workbook import assert_workbook_safe, clone_template_workbook, write_literal


REPORT_TEMPLATE_SHEET = "báo cáo tổng hợp"
HEADER_ROW = 2
FIRST_DATA_ROW = 3


class ReportExportError(ValueError):
    def __init__(
        self, message: str, *, code: str = "invalid_monthly_report", status: int = 422,
    ):
        super().__init__(message)
        self.code = code
        self.status = status


def _plain(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip())


def _key(value: Any) -> str:
    text = _plain(value).replace("đ", "d").replace("Đ", "D")
    text = unicodedata.normalize("NFKD", text).casefold()
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"[^a-z0-9]+", "", text)


def _row_dict(row: Any) -> dict[str, Any]:
    return dict(row) if not isinstance(row, dict) else row.copy()


def _iso_date(value: Any, *, label: str = "Ngày báo cáo") -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = _plain(value)
    for pattern in ("%Y-%m-%d", "%d/%m/%Y", "%d.%m.%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, pattern).date().isoformat()
        except ValueError:
            continue
    raise ReportExportError(f"{label} không hợp lệ", code="invalid_report_date")


def _number(value: Any, label: str) -> Decimal:
    if isinstance(value, bool):
        raise ReportExportError(f"{label} phải là số hữu hạn")
    try:
        result = Decimal(str(value))
    except (InvalidOperation, TypeError, ValueError) as error:
        raise ReportExportError(f"{label} phải là số hữu hạn") from error
    if not result.is_finite():
        raise ReportExportError(f"{label} phải là số hữu hạn")
    return result


def _excel_number(value: Decimal) -> int | float:
    if value == value.to_integral_value():
        return int(value)
    return float(value)


def collect_monthly_report_rows(
    conn: Any,
    batch: Mapping[str, Any],
    selected_orders: Iterable[Mapping[str, Any]],
    *,
    totals_fn: Callable[[Mapping[str, Any]], tuple[Any, Any, Any, Any]],
) -> tuple[str, list[dict[str, Any]]]:
    """Read approved orders in the selected month plus a selected draft preview.

    The query deliberately excludes payments, balances and debt ledgers.  This
    report is a monthly operating sales/cost view, not a receivable statement.
    """

    batch_data = _row_dict(batch)
    work_date = _iso_date(batch_data.get("work_date"))
    period = work_date[:7]
    date_from = period + "-01"
    if period.endswith("-12"):
        date_to = f"{int(period[:4]) + 1:04d}-01-01"
    else:
        date_to = f"{period[:5]}{int(period[5:7]) + 1:02d}-01"

    sources: dict[int, dict[str, Any]] = {}
    batch_id = int(batch_data.get("id") or 0)
    if batch_id:
        for source in conn.execute(
            """SELECT o.*
                 FROM orders o
                 JOIN batches b ON b.id=o.batch_id
                WHERE b.status='approved' AND b.work_date>=? AND b.work_date<?
                ORDER BY b.work_date,o.id""",
            (date_from, date_to),
        ):
            item = _row_dict(source)
            sources[int(item.get("id") or 0)] = item

    selected = [_row_dict(row) for row in selected_orders]
    if not batch_id:
        sources = {index: item for index, item in enumerate(selected, start=1)}
    elif _plain(batch_data.get("status")).casefold() != "approved":
        for index, item in enumerate(selected, start=1):
            key = int(item.get("id") or 0) or -index
            sources[key] = item

    rows = []
    for source_ref, item in sources.items():
        contractor = _plain(item.get("contractor"))
        kitchen = _plain(item.get("kitchen"))
        if not contractor or not kitchen:
            raise ReportExportError(
                f"Dòng nguồn {source_ref} thiếu nhà thầu hoặc mã bếp",
                code="missing_report_dimension",
            )
        revenue, cost, profit, total = totals_fn(item)
        values = {
            "revenue": _number(revenue, "Doanh số bán"),
            "cost": _number(cost, "Giá vốn"),
            "profit": _number(profit, "Lợi nhuận gộp"),
            "total": _number(total, "Tổng thanh toán"),
        }
        if values["profit"] != values["revenue"] - values["cost"]:
            raise ReportExportError(
                f"Dòng nguồn {source_ref} có lợi nhuận không khớp doanh số trừ giá vốn",
                code="report_profit_mismatch",
            )
        if not any(values.values()):
            continue
        rows.append({
            "contractor": contractor,
            "kitchen": kitchen,
            **{key: _excel_number(value) for key, value in values.items()},
            "source_ref": source_ref,
        })
    return period, rows


def aggregate_monthly_report_rows(
    rows: Iterable[Mapping[str, Any]],
) -> list[dict[str, Any]]:
    groups: dict[str, dict[str, Any]] = {}
    for index, source in enumerate(rows, start=1):
        item = _row_dict(source)
        contractor = _plain(item.get("contractor"))
        kitchen = _plain(item.get("kitchen"))
        if not contractor or not kitchen:
            raise ReportExportError(f"Dòng {index} thiếu nhà thầu hoặc mã bếp")
        kitchen_key = _key(kitchen)
        group = groups.setdefault(kitchen_key, {
            "contractor": contractor,
            "kitchen": kitchen,
            "revenue_decimal": Decimal("0"),
            "cost_decimal": Decimal("0"),
            "profit_decimal": Decimal("0"),
            "total_decimal": Decimal("0"),
            "source_refs": [],
        })
        if _key(group["contractor"]) != _key(contractor):
            raise ReportExportError(
                f"Dòng {index} làm mã bếp thuộc nhiều nhà thầu",
                code="conflicting_report_contractor",
            )
        for field in ("revenue", "cost", "profit", "total"):
            group[f"{field}_decimal"] += _number(item.get(field), f"{field} dòng {index}")
        if item.get("source_ref") not in (None, ""):
            group["source_refs"].append(item.get("source_ref"))

    output = []
    for group in groups.values():
        if group["profit_decimal"] != group["revenue_decimal"] - group["cost_decimal"]:
            raise ReportExportError(
                "Tổng lợi nhuận bếp không khớp doanh số trừ giá vốn",
                code="report_profit_mismatch",
            )
        output.append({
            "contractor": group["contractor"],
            "kitchen": group["kitchen"],
            "revenue": _excel_number(group["revenue_decimal"]),
            "cost": _excel_number(group["cost_decimal"]),
            "profit": _excel_number(group["profit_decimal"]),
            "total": _excel_number(group["total_decimal"]),
            "source_refs": tuple(group["source_refs"]),
        })
    return output


def _row_snapshot(sheet: Any, row: int) -> dict[str, Any]:
    cells = []
    for column in range(1, sheet.max_column + 1):
        cell = sheet.cell(row, column)
        if isinstance(cell, MergedCell):
            cells.append(None)
            continue
        cells.append({
            "style": copy(cell._style),
            "number_format": cell.number_format,
            "font": copy(cell.font),
            "fill": copy(cell.fill),
            "border": copy(cell.border),
            "alignment": copy(cell.alignment),
            "protection": copy(cell.protection),
        })
    return {"height": sheet.row_dimensions[row].height, "cells": cells}


def _apply_row_snapshot(sheet: Any, row: int, snapshot: Mapping[str, Any]) -> None:
    sheet.row_dimensions[row].height = snapshot.get("height")
    for column, source in enumerate(snapshot["cells"], start=1):
        if source is None:
            continue
        cell = sheet.cell(row, column)
        cell.value = None
        cell._style = copy(source["style"])
        cell.number_format = source["number_format"]
        cell.font = copy(source["font"])
        cell.fill = copy(source["fill"])
        cell.border = copy(source["border"])
        cell.alignment = copy(source["alignment"])
        cell.protection = copy(source["protection"])


def _golden_group_metadata(sheet: Any) -> tuple[dict[str, str], list[str], dict[str, int], dict[str, dict[str, Any]]]:
    kitchen_groups: dict[str, str] = {}
    group_order: list[str] = []
    kitchen_rank: dict[str, int] = {}
    group_styles: dict[str, dict[str, Any]] = {}
    for row_index in range(FIRST_DATA_ROW, sheet.max_row + 1):
        kitchen = _plain(sheet.cell(row_index, 3).value)
        contractor = _plain(sheet.cell(row_index, 2).value)
        if not kitchen or _key(sheet.cell(row_index, 1).value) == "tongthang":
            continue
        group = _plain(sheet.cell(row_index, 1).value) or contractor
        kitchen_groups[_key(kitchen)] = group
        kitchen_rank[_key(kitchen)] = len(kitchen_rank)
        group_key = _key(group)
        if group_key not in {_key(value) for value in group_order}:
            group_order.append(group)
            group_styles[group_key] = _row_snapshot(sheet, row_index)
    return kitchen_groups, group_order, kitchen_rank, group_styles


def build_monthly_report_workbook(
    rows: Iterable[Mapping[str, Any]],
    *,
    period: str,
    template_path: str | Path,
    configured_groups: Mapping[str, str] | None = None,
    expected_sha256: str = EM_THANH_SHA256,
) -> Any:
    if not re.fullmatch(r"\d{4}-\d{2}", _plain(period)):
        raise ReportExportError("Kỳ báo cáo phải có dạng YYYY-MM", code="invalid_report_period")
    aggregated = aggregate_monthly_report_rows(rows)
    if not aggregated:
        raise ReportExportError("Không có phát sinh để lập báo cáo tổng hợp", code="empty_monthly_report")
    configured = {_key(key): _plain(value) for key, value in (configured_groups or {}).items() if _plain(value)}

    cloned = clone_template_workbook(
        template_path,
        sheet_names=[REPORT_TEMPLATE_SHEET],
        expected_sha256=expected_sha256,
    )
    workbook = cloned.workbook
    sheet = workbook[REPORT_TEMPLATE_SHEET]
    try:
        kitchen_groups, group_order, kitchen_rank, group_styles = _golden_group_metadata(sheet)
        plain_style = _row_snapshot(sheet, 25)
        subtotal_style = _row_snapshot(sheet, 14)
        for item in aggregated:
            kitchen_key = _key(item["kitchen"])
            item["customer_group"] = (
                configured.get(kitchen_key)
                or kitchen_groups.get(kitchen_key)
                or item["contractor"]
            )

        group_rank = {_key(value): index for index, value in enumerate(group_order)}
        buckets: dict[str, list[dict[str, Any]]] = defaultdict(list)
        group_labels: dict[str, str] = {}
        for item in aggregated:
            group_key = _key(item["customer_group"])
            buckets[group_key].append(item)
            group_labels.setdefault(group_key, item["customer_group"])
        ordered_groups = sorted(
            buckets,
            key=lambda key: (group_rank.get(key, len(group_rank)), _key(group_labels[key])),
        )
        for key in ordered_groups:
            buckets[key].sort(
                key=lambda item: (
                    kitchen_rank.get(_key(item["kitchen"]), len(kitchen_rank)),
                    _key(item["contractor"]),
                    _key(item["kitchen"]),
                )
            )

        for row_index in range(FIRST_DATA_ROW, sheet.max_row + 1):
            for column in range(1, 8):
                sheet.cell(row_index, column).value = None

        output_row = FIRST_DATA_ROW
        month_totals = {field: Decimal("0") for field in ("revenue", "cost", "profit", "total")}
        for group_key in ordered_groups:
            items = buckets[group_key]
            style = group_styles.get(group_key, plain_style)
            group_totals = {field: Decimal("0") for field in month_totals}
            for item in items:
                _apply_row_snapshot(sheet, output_row, style)
                write_literal(sheet, f"A{output_row}", item["customer_group"])
                write_literal(sheet, f"B{output_row}", item["contractor"])
                write_literal(sheet, f"C{output_row}", item["kitchen"])
                for column, field in zip((4, 5, 6, 7), month_totals):
                    value = _number(item[field], field)
                    group_totals[field] += value
                    month_totals[field] += value
                    write_literal(sheet, f"{chr(64 + column)}{output_row}", _excel_number(value))
                output_row += 1
            if len(items) > 1:
                _apply_row_snapshot(sheet, output_row, subtotal_style)
                write_literal(sheet, f"A{output_row}", "TỔNG " + group_labels[group_key])
                for column, field in zip((4, 5, 6, 7), month_totals):
                    write_literal(
                        sheet,
                        f"{chr(64 + column)}{output_row}",
                        _excel_number(group_totals[field]),
                    )
                output_row += 1

        _apply_row_snapshot(sheet, output_row, subtotal_style)
        write_literal(sheet, f"A{output_row}", "TỔNG THÁNG")
        for column, field in zip((4, 5, 6, 7), month_totals):
            write_literal(sheet, f"{chr(64 + column)}{output_row}", _excel_number(month_totals[field]))

        if sheet.max_row > output_row:
            sheet.delete_rows(output_row + 1, sheet.max_row - output_row)
        sheet.freeze_panes = "A3"
        sheet.print_title_rows = "$2:$2"
        sheet.print_area = f"$A$2:$G${output_row}"
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.orientation = sheet.ORIENTATION_LANDSCAPE
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.page_margins = PageMargins(
            left=0.25, right=0.25, top=0.35, bottom=0.35, header=0.15, footer=0.15,
        )
        workbook.active = 0
        sheet.sheet_view.tabSelected = True
        assert_workbook_safe(workbook)
        return workbook
    except Exception:
        workbook.close()
        raise


__all__ = [
    "REPORT_TEMPLATE_SHEET",
    "ReportExportError",
    "aggregate_monthly_report_rows",
    "build_monthly_report_workbook",
    "collect_monthly_report_rows",
]
