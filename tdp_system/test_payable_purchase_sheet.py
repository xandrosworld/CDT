import io
import json
import unittest
from openpyxl import load_workbook
from . import server, contract_modules as cm
from . import test_supplier_plan_source as fixtures
from .payable_ledger import sync_payable_ledger, pending_purchase_sheets
from .payable_export import payable_export_data


class PayablePurchaseSheetTests(fixtures.SupplierPlanSourceTests):
    def setUp(self):
        with server.db() as conn:
            conn.execute('DELETE FROM payable_payment_allocations')
            conn.execute('DELETE FROM payable_payment_revisions')
            conn.execute('DELETE FROM payments')
            conn.execute('DELETE FROM payable_ledger_revisions')
            conn.execute('DELETE FROM payable_ledger_lines')
        super().setUp()

    def test_sheet_rounding_preview_ledger_export_and_payment_remain_consistent(self):
        wb=load_workbook(io.BytesIO(self.file(1,10.5)));ws=wb['đặt hàng']
        for cell in list(ws[3]):ws.cell(4,cell.column,cell.value)
        ws.cell(4,8,20.5);ws.cell(4,14,20.5)
        stream=io.BytesIO();wb.save(stream);wb.close();raw=stream.getvalue()
        bid=self.daily(raw)
        preview=self.preview(bid,raw)
        self.assertEqual((preview['total_amount'],preview['rounding_adjustment']),(31,-1))
        self.approve_for_payable(bid)
        rows=sorted(self.active(),key=lambda r:r['source_row'])
        self.assertEqual([r['amount'] for r in rows],[10,21])
        with server.db() as conn:
            from .payable_ledger import set_payable_allocation_total, payable_ledger_payload
            from .payable_export import payable_workbook
            repeat=sync_payable_ledger(conn,timestamp=server.now_iso())
            self.assertEqual(repeat['updated'],0)
            data=payable_export_data(conn,date_from='2026-09-01',date_to='2026-09-01',supplier='',canonical_party_code=lambda conn,kind,code:code)
            self.assertEqual(sum(r['amount'] for r in data['lines']),31)
            book=payable_workbook(data)
            self.assertTrue(any(c.comment and '-1đ' in c.comment.text for sh in book for line in sh for c in line))
            book.close()
            set_payable_allocation_total(conn,ledger_line_id=rows[0]['id'],paid_amount=10,timestamp=server.now_iso())
            ledger=payable_ledger_payload(conn,date_from='2026-09-01',date_to='2026-09-01',statuses='all')
            self.assertEqual((ledger['summary']['charge_amount'],ledger['summary']['paid_amount'],ledger['summary']['remaining_amount']),(31,10,21))
            self.assertEqual(sync_payable_ledger(conn,timestamp=server.now_iso())['updated'],0)

    def approve_for_payable(self, bid):
        with server.db() as conn:
            conn.execute("UPDATE batches SET status='approved' WHERE id=?", (bid,))
            sync_payable_ledger(conn, timestamp=server.now_iso())

    def active(self):
        with server.db() as conn:
            return [dict(r) for r in conn.execute("SELECT * FROM payable_ledger_lines WHERE status!='reversed'")]

    def test_approved_day_uses_purchase_quantity_price_and_adjustments_not_sales(self):
        w=load_workbook(io.BytesIO(self.file(7, 9000)));s=w['đặt hàng']
        s.cell(3,9,1);s.cell(3,10,2);s.cell(3,13,8);s.cell(3,14,72000)
        stream=io.BytesIO();w.save(stream);w.close()
        bid=self.daily(stream.getvalue())
        with server.db() as conn:
            protected=[tuple(r)for r in conn.execute('SELECT * FROM orders')]
        self.approve_for_payable(bid)
        row=self.active()[0]
        self.assertEqual((row['source_table'],row['actual_qty'],row['buy_price'],row['amount']), ('supplier_plan_sources',8,9000,72000))
        with server.db() as conn:
            self.assertEqual([tuple(r)for r in conn.execute('SELECT * FROM orders')],protected)
            repeat=sync_payable_ledger(conn,timestamp=server.now_iso())
            self.assertEqual(repeat['inserted']+repeat['updated']+repeat['reversed'],0)
            data=payable_export_data(conn,date_from='2026-09-01',date_to='2026-09-01',supplier='',canonical_party_code=lambda conn,kind,code: code)
            line=next(x for x in data['lines']if x['status']!='reversed')
            self.assertEqual((line['base_qty'],line['damaged_qty'],line['added_qty'],line['amount']),(7,1,2,72000))

    def test_no_sales_fallback_and_unpriced_sheet_shows_pending(self):
        bid=self.daily(self.file())
        self.approve_for_payable(bid)
        self.assertEqual(self.active(),[])
        with server.db() as conn:
            pending=pending_purchase_sheets(conn,'2026-09-01','2026-09-01')
            self.assertIn('Chưa có giá mua',str(pending))
            conn.execute('DELETE FROM supplier_plan_sources')
            sync_payable_ledger(conn,timestamp=server.now_iso())
            self.assertEqual(len(pending_purchase_sheets(conn,'2026-09-01','2026-09-01')),1)
        self.assertEqual(self.active(),[])

    def test_complete_sheet_includes_kho_in_details_summary_and_export(self):
        wb = load_workbook(io.BytesIO(self.file(7, 9000)))
        sheet = wb['đặt hàng']
        for cell in list(sheet[3]):
            sheet.cell(4, cell.column, cell.value)
        sheet.cell(4, 4, 6)
        sheet.cell(4, 6, 'kho')
        sheet.cell(4, 8, 14000)
        sheet.cell(4, 13, 6)
        sheet.cell(4, 14, 84000)
        stream = io.BytesIO(); wb.save(stream); wb.close()
        bid = self.daily(stream.getvalue())
        self.approve_for_payable(bid)
        with server.db() as conn:
            rows = [dict(r) for r in conn.execute("SELECT * FROM payable_ledger_lines WHERE status!='reversed'")]
            self.assertEqual(len(rows), 2)
            kho = next(r for r in rows if r['supplier_snapshot'] == 'kho')
            self.assertEqual((kho['status'], kho['amount']), ('open', 84000))
            # Existing deployments have retained this source as a reversed row.
            conn.execute("UPDATE payable_ledger_lines SET status='reversed',reversal_reason='internal_stock' WHERE id=?", (kho['id'],))
        before = self.protected()
        with server.db() as conn:
            result = sync_payable_ledger(conn, timestamp=server.now_iso())
            self.assertEqual(result['reactivated'], 1)
            repeat = sync_payable_ledger(conn, timestamp=server.now_iso())
            self.assertEqual(repeat['inserted'] + repeat['updated'] + repeat['reversed'], 0)
            self.assertEqual(repeat['reactivated'], 0)
            data = payable_export_data(conn, date_from='2026-09-01', date_to='2026-09-01',
                supplier='', canonical_party_code=lambda conn, kind, code: code)
            self.assertEqual(sum(r['amount'] for r in data['lines'] if r['status'] != 'reversed'), 147000)
            restored = dict(conn.execute('SELECT * FROM payable_ledger_lines WHERE id=?', (kho['id'],)).fetchone())
            self.assertEqual(restored['revision'], kho['revision'] + 1)
            self.assertEqual(restored['source_hash'], kho['source_hash'])
        for table, old in before.items():
            if not table.startswith('payable_'):
                self.assertEqual(self.protected()[table], old, table)
        ledger = self.client.get('/api/debts/payables/ledger?from=2026-09-01&to=2026-09-01').get_json()
        self.assertEqual(ledger['summary']['charge_amount'], 147000)
        period = self.client.get('/api/debts?from=2026-09-01&to=2026-09-01').get_json()
        self.assertEqual(sum(r['period_charge'] for r in period['suppliers'].values()), 147000)
        filtered = self.client.get('/api/debts/payables/ledger?from=2026-09-01&to=2026-09-01&supplier=kho').get_json()
        self.assertEqual(filtered['summary']['charge_amount'], 84000)
        self.assertEqual(len(filtered['rows']), 1)

    def test_correcting_priced_sheet_updates_only_payable_and_keeps_revisions(self):
        bid=self.daily(self.file(7,9000));self.approve_for_payable(bid)
        before=self.protected();old=self.active()[0]
        self.confirm(self.preview(bid,self.file(8,11000)))
        row=self.active()[0]
        self.assertEqual((row['id'],row['revision'],row['amount']),(old['id'],old['revision']+1,88000))
        after=self.protected()
        for table in before:
            if not table.startswith('payable_'):self.assertEqual(before[table],after[table],table)

    def test_paid_sheet_correction_is_rejected_and_rolled_back(self):
        bid=self.daily(self.file(7,9000));self.approve_for_payable(bid)
        with server.db() as conn:
            conn.execute("UPDATE payable_ledger_lines SET paid_amount=1000,status='partially_paid' WHERE status='open'")
            before=tuple(conn.execute('SELECT * FROM supplier_plan_sources').fetchone())
        self.confirm(self.preview(bid,self.file(8,9000)),status=409)
        with server.db() as conn:
            self.assertEqual(tuple(conn.execute('SELECT * FROM supplier_plan_sources').fetchone()),before)
        self.assertEqual(self.active()[0]['amount'],63000)

    def test_confirmed_purchase_uses_sheet_price_even_with_existing_quote(self):
        bid=self.daily(self.file(7,9000))
        preview=self.preview(bid,self.file(7,9000),plan=False)
        self.assertEqual(preview['total_amount'],63000)
        self.confirm(preview)
        with server.db() as conn:
            conn.execute("UPDATE purchase_workbook_lines SET price_source='Bảng báo giá',buy_price=10000,amount=70000")
        preview=self.preview(bid,self.file(7,11000),plan=False)
        self.assertEqual(preview['total_amount'],77000)
        self.confirm(preview)
        with server.db() as conn:
            self.assertEqual(tuple(conn.execute('SELECT buy_price,amount,price_source FROM purchase_workbook_lines').fetchone()),(11000,77000,'Sheet đặt hàng chuẩn'))

    def test_paid_sheet_cannot_be_replaced_by_another_source_table(self):
        bid=self.daily(self.file(7,9000));self.approve_for_payable(bid)
        with server.db() as conn:
            conn.execute("UPDATE payable_ledger_lines SET paid_amount=1000,status='partially_paid' WHERE status='open'")
        before=self.protected()
        self.confirm(self.preview(bid,self.file(7,9000),plan=False),status=409)
        self.assertEqual(self.protected(),before)
        self.assertEqual(self.active()[0]['source_table'],'supplier_plan_sources')


if __name__=='__main__':unittest.main()
