"""New catalogue rows travel atomically with the daily workbook."""
import io
import unittest
from openpyxl import load_workbook
from . import server
from . import test_daily_reference_import as fixtures


class DailyCatalogCaptureTests(unittest.TestCase):
    setUpClass = classmethod(fixtures.DailyReferenceImportTests.setUpClass.__func__)
    tearDownClass = classmethod(fixtures.DailyReferenceImportTests.tearDownClass.__func__)

    def setUp(self):
        fixtures.DailyReferenceImportTests.setUp(self)
        with server.db() as conn:
            conn.execute('DROP TRIGGER IF EXISTS fail_daily_catalog_order')
            conn.execute("DELETE FROM outgoing_product_names WHERE product_code!='P1'")
            conn.execute("DELETE FROM outgoing_product_units WHERE product_code!='P1'")
            conn.execute("DELETE FROM products WHERE code!='P1'")

    def workbook(self, *, used=False, title='danh mục hh', conflict=False, invalid=False):
        wb = load_workbook(io.BytesIO(fixtures.DailyReferenceImportTests.workbook_bytes()))
        ws = wb['danh mục hàng hóa']
        ws.title = 'temporary catalog'
        ws.title = title
        # Existing catalogue values in old daily files must not overwrite masters.
        ws.cell(2, 4, 'Old file name')
        ws.cell(2, 6, 'box')
        ws.append([2, 'NEW001', 'G', 'New product', '', '' if invalid else 'Lon', '8%'])
        if conflict:
            ws.append([3, 'NEW001', 'G', 'Different product', '', 'Lon', '8%'])
        if used:
            day = wb['01.09']
            day.cell(3, 4, 'NEW001')
            day.cell(3, 5, 'New product')
            day.cell(3, 7, 'Lon')
            day.cell(3, 8, 'S1')
        stream = io.BytesIO()
        wb.save(stream)
        wb.close()
        return stream.getvalue()

    def preview(self, payload):
        response = self.client.post('/api/import/analyze', data={
            'file': (io.BytesIO(payload), 'Đơn hàng 01.09.2026.xlsx'), 'continuous': '1',
        })
        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        return response.get_json()

    def confirm(self, preview):
        return self.client.post('/api/import/confirm', json={
            'token': preview['token'], 'sheets': ['01.09'],
            'work_date': '2026-09-01', 'state_hash': preview['stateHash'],
        })

    def test_new_unused_catalog_code_is_added_with_order(self):
        with server.db() as conn:
            before = fixtures.DailyReferenceImportTests.protected_snapshot(conn)
        preview = self.preview(self.workbook())
        with server.db() as conn:
            self.assertIsNone(conn.execute("SELECT code FROM products WHERE code='NEW001'").fetchone())
        response = self.confirm(preview)
        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        with server.db() as conn:
            item = conn.execute("SELECT name,unit,tax,supplier,buy_price FROM products WHERE code='NEW001'").fetchone()
            self.assertIsNotNone(item)
            self.assertEqual(tuple(item), ('New product', 'Lon', '0.08', '', 0))
            after = fixtures.DailyReferenceImportTests.protected_snapshot(conn)
            after['products'] = [row for row in after['products'] if row[0] != 'NEW001']
            self.assertEqual(before, after)
        self.assertEqual(response.get_json()['catalogImport']['inserted'], 1)

    def test_code_only_in_price_sheet_is_captured_with_daily_order(self):
        wb=load_workbook(io.BytesIO(self.workbook(used=True)))
        ws=wb['danh mục hh']
        ws.delete_rows(ws.max_row)
        if 'BÁO GIÁ' in wb:
            del wb['BÁO GIÁ']
        price=wb.create_sheet('BÁO GIÁ')
        price.append(['Mã hàng','Tên hàng','ĐVT','Thuế'])
        price.append(['NEW001','New product','Lon','8%'])
        data=io.BytesIO();wb.save(data);wb.close()
        preview=self.preview(data.getvalue())
        self.assertEqual(1,preview['catalogAdditions']['newCount'])
        response=self.confirm(preview)
        self.assertEqual(200,response.status_code,response.json)
        self.assertEqual(0,response.json['summary']['totals']['errors'])

    def test_new_code_used_by_order_is_resolved_before_save(self):
        preview = self.preview(self.workbook(used=True))
        day = next(row for row in preview['sheets'] if row['name'] == '01.09')
        self.assertEqual(day['errorRows'], 0, day)
        response = self.confirm(preview)
        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        self.assertEqual(response.get_json()['orders'][0]['product_code'], 'NEW001')
        self.assertEqual(response.get_json()['summary']['totals']['errors'], 0)

    def test_same_file_replay_does_not_duplicate_products_or_orders(self):
        payload = self.workbook(used=True)
        first = self.confirm(self.preview(payload)).get_json()
        preview = self.preview(payload)
        self.assertEqual(preview['catalogAdditions']['newCount'], 0)
        second = self.confirm(preview).get_json()
        self.assertTrue(second['idempotent'], second)
        self.assertEqual(second['catalogImport']['inserted'], 0)
        self.assertEqual([r['id'] for r in first['orders']], [r['id'] for r in second['orders']])
        with server.db() as conn:
            self.assertEqual(conn.execute("SELECT COUNT(*) FROM products WHERE code='NEW001'").fetchone()[0], 1)

    def test_catalogue_aliases_are_recognized(self):
        for title in ('danh mục hh', 'DANH MỤC HÀNG HÓA', 'danh mục hàng'):
            preview = self.preview(self.workbook(title=title))
            self.assertTrue(preview['strictDaily'])
            self.assertEqual(preview['catalogAdditions']['newCount'], 1)

    def test_new_catalogue_conflict_blocks_confirmation_with_source_rows(self):
        preview = self.preview(self.workbook(conflict=True))
        self.assertFalse(preview['catalogAdditions']['canConfirm'])
        self.assertIn('NEW001', ' '.join(preview['catalogAdditions']['errors']))
        self.assertIn('dòng 3', ' '.join(preview['catalogAdditions']['errors']))
        self.assertEqual(self.confirm(preview).status_code, 400)
        with server.db() as conn:
            self.assertEqual(conn.execute('SELECT COUNT(*) FROM batches').fetchone()[0], 0)
            self.assertIsNone(conn.execute("SELECT * FROM products WHERE code='NEW001'").fetchone())

    def test_invalid_new_row_names_missing_field(self):
        preview = self.preview(self.workbook(invalid=True))
        self.assertFalse(preview['catalogAdditions']['canConfirm'])
        self.assertIn('Thiếu đơn vị tính', ' '.join(preview['catalogAdditions']['errors']))

    def test_cancel_does_not_write_catalogue(self):
        preview = self.preview(self.workbook())
        self.assertEqual(self.client.post('/api/import/cancel', json={'token': preview['token']}).status_code, 200)
        with server.db() as conn:
            self.assertIsNone(conn.execute("SELECT * FROM products WHERE code='NEW001'").fetchone())

    def test_catalogue_change_after_preview_prevents_stale_apply(self):
        preview = self.preview(self.workbook())
        with server.db() as conn:
            conn.execute("UPDATE products SET name='Concurrent edit' WHERE code='P1'")
        response = self.confirm(preview)
        self.assertNotEqual(response.status_code, 200)
        self.assertIn('Danh mục đã thay đổi', response.get_json()['error'])
        with server.db() as conn:
            self.assertIsNone(conn.execute("SELECT * FROM products WHERE code='NEW001'").fetchone())
            self.assertEqual(conn.execute('SELECT COUNT(*) FROM batches').fetchone()[0], 0)

    def test_order_failure_rolls_back_catalogue_and_audit(self):
        preview = self.preview(self.workbook())
        with server.db() as conn:
            audits = conn.execute("SELECT COUNT(*) FROM audit_log WHERE event_type='catalog.daily_additions'").fetchone()[0]
            conn.execute("""CREATE TRIGGER fail_daily_catalog_order BEFORE INSERT ON orders
                BEGIN SELECT RAISE(ABORT, 'forced atomicity check'); END""")
        response = self.confirm(preview)
        self.assertEqual(response.status_code, 500)
        with server.db() as conn:
            self.assertIsNone(conn.execute("SELECT * FROM products WHERE code='NEW001'").fetchone())
            self.assertEqual(conn.execute('SELECT COUNT(*) FROM batches').fetchone()[0], 0)
            self.assertEqual(conn.execute("SELECT COUNT(*) FROM audit_log WHERE event_type='catalog.daily_additions'").fetchone()[0], audits)

    def test_later_daily_file_adds_unused_code(self):
        initial = self.confirm(self.preview(fixtures.DailyReferenceImportTests.workbook_bytes()))
        self.assertEqual(initial.status_code, 200)
        preview = self.preview(self.workbook())
        self.assertEqual(preview['phase'], 'finalization')
        response = self.confirm(preview)
        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        self.assertEqual(response.get_json()['catalogImport']['inserted'], 1)

    def test_new_code_after_formatted_blank_rows_is_not_silently_skipped(self):
        wb = load_workbook(io.BytesIO(self.workbook()))
        ws = wb['danh mục hh']
        values = [cell.value for cell in ws[3]]
        ws.delete_rows(3)
        for col, value in enumerate(values, 1):
            ws.cell(12669, col, value)
        out = io.BytesIO(); wb.save(out); wb.close()
        preview = self.preview(out.getvalue())
        self.assertEqual(preview['catalogAdditions']['newCount'], 1)
        self.assertEqual(self.confirm(preview).status_code, 200)

    def test_locked_day_purchase_update_adds_code_without_changing_sales_or_stock(self):
        from .test_supplier_plan_source import SupplierPlanSourceTests
        payload = SupplierPlanSourceTests.file(self, 7, 9000)
        # This purchase fixture uses the same P1/C1/K1 master keys.
        first = self.confirm(self.preview(payload))
        self.assertEqual(first.status_code, 200, first.get_data(as_text=True))
        bid = first.get_json()['batch']['id']
        with server.db() as conn:
            conn.execute("INSERT INTO batch_bk_approvals VALUES(?,999999,'posted-source','2026-09-01')", (bid,))
            conn.execute("UPDATE batches SET status='approved' WHERE id=?", (bid,))
            tables = ('orders', 'batches', 'receivable_ledger_lines', 'invoice_inventory_ledger', 'batch_bk_approvals')
            before = {t: [tuple(r) for r in conn.execute('SELECT * FROM '+t+' ORDER BY 1')] for t in tables}
        wb = load_workbook(io.BytesIO(payload))
        ws = wb.create_sheet('danh mục hàng hóa')
        ws.append(['STT', 'Mã hàng', 'Nhóm hàng', 'Tên Thành Đạt Phát', 'Tên xuất hóa đơn', 'ĐVT', 'Thuế'])
        ws.append([2, 'NEW001', 'G', 'New product', 'New product', 'Lon', '8%'])
        out = io.BytesIO(); wb.save(out); wb.close()
        preview = self.preview(out.getvalue())
        self.assertTrue(preview['purchasePlanOnly'])
        response = self.client.post('/api/import/confirm', json={
            'token': preview['token'], 'state_hash': preview['stateHash'],
            'sheets': ['đặt hàng'], 'work_date': '2026-09-01',
        })
        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        self.assertEqual(response.get_json()['catalogImport']['inserted'], 1)
        with server.db() as conn:
            after = {t: [tuple(r) for r in conn.execute('SELECT * FROM '+t+' ORDER BY 1')] for t in tables}
        self.assertEqual(before, after)
