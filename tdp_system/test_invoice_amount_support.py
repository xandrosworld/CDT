import json
import sqlite3
import unittest
from contextlib import contextmanager
from copy import deepcopy
from io import BytesIO
from unittest.mock import Mock, patch

from flask import Flask
from openpyxl import load_workbook

from .test_invoice_input_sync import init_test_database
from .test_minvoice_portal import document, client as portal_client
from .minvoice_client import MinvoiceError
from .minvoice_portal import normalize_portal_document
from .invoice_output_sync import upsert_output_invoice
from .invoice_mapping import save_mapping
from .invoice_amount_support import register_amount_support_routes, inspect_amounts

NOW = '2026-09-09T03:00:00'


class AmountSupportTests(unittest.TestCase):
    def setUp(self):
        self.conn = sqlite3.connect(':memory:'); init_test_database(self.conn)
        self.conn.execute("INSERT INTO products(code,name,unit) VALUES('A','Goods','Kg')")
        self.raw = document(); self.raw.update(totalAmountWithoutVAT=120000, totalAmount=128000)
        self.iid = upsert_output_invoice(self.conn, normalize_portal_document(self.raw), tenant='TDP', now=NOW,
                                        status_map={}, status_fields=(), reference_fields=())[0]
        item = self.conn.execute('SELECT id FROM outgoing_source_invoice_items').fetchone()[0]
        save_mapping(self.conn, direction='output', item_id=item, product_code='A', now_iso=lambda: NOW)
        self.remote = Mock(); self.remote.get_outgoing_invoice.side_effect = lambda **kw: normalize_portal_document(deepcopy(self.raw))
        @contextmanager
        def db():
            with self.conn: yield self.conn
        app = Flask(__name__)
        register_amount_support_routes(app, {'db': db, 'now_iso': lambda: NOW,
            'setting_get': lambda conn,key,default: default, 'create_minvoice_client': lambda: self.remote})
        self.client = app.test_client(); self.url = f'/api/invoice-workbench/output/{self.iid}/amount-review'
        self.conn.commit()

    def tearDown(self):
        self.conn.close()

    def snapshot(self):
        tables = [r[0] for r in self.conn.execute("SELECT name FROM sqlite_master WHERE type='table' ORDER BY name")]
        return {t: [tuple(r) for r in self.conn.execute('SELECT * FROM "' + t + '"')] for t in tables}

    def check(self, fresh=False):
        response = self.client.get(self.url + ('?fresh=1' if fresh else ''))
        self.assertEqual(response.status_code, 200, response.json)
        return response.json

    def test_open_recheck_and_export_do_not_write_or_guess_missing_line(self):
        before = self.snapshot()
        report = self.check(); self.assertFalse(report['can_apply'])
        self.remote.get_outgoing_invoice.assert_not_called()
        self.assertEqual(report['comparisons'][0]['difference'], 20000)
        self.assertIn('Chưa xác định', report['message']); self.assertEqual(report['lines'][0]['note'], '')
        self.assertFalse(self.check(True)['can_apply'])
        # Minimal Linux containers do not necessarily know the .xlsx MIME type.
        with patch('mimetypes.guess_type', return_value=(None, None)):
            response = self.client.get(self.url + '?download=1')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.mimetype, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        wb = load_workbook(BytesIO(response.data)); ws = wb.active
        self.assertEqual(ws['D8'].value, 20000); self.assertEqual(ws['F13'].value, 100000)
        self.assertEqual(ws['D8'].font.color.rgb, '00B42318')
        self.assertEqual(self.snapshot(), before)

    def test_apply_fresh_verified_source_restores_mapping_and_never_posts(self):
        before = self.snapshot(); self.raw = document()
        report = self.check(True); self.assertTrue(report['can_apply']); self.assertTrue(report['changes'])
        response = self.client.post(self.url + '/apply', json={k: report[k] for k in ('expected','source_digest')})
        self.assertEqual(response.status_code, 200, response.json)
        row = self.conn.execute('SELECT * FROM outgoing_source_invoices').fetchone()
        self.assertEqual(row['sync_status'], 'synced'); self.assertEqual(row['subtotal'], 100000)
        line = self.conn.execute('SELECT * FROM outgoing_source_invoice_items').fetchone()
        self.assertEqual((line['product_code'],line['mapping_status'],line['stock_qty']), ('A','mapped',2))
        after = self.snapshot()
        for table in ('invoice_inventory_ledger', 'inventory_transactions', 'orders', 'invoice_line_mappings'):
            self.assertEqual(after[table], before[table], table)
        self.assertEqual(self.conn.execute("SELECT COUNT(*) FROM audit_log WHERE event_type='invoice_amount_source_refresh'").fetchone()[0], 1)
        self.assertEqual(self.client.post(self.url + '/apply', json=report).status_code, 409)

    def test_stale_local_or_remote_or_posted_source_cannot_apply(self):
        self.raw = document(); report = self.check(True)
        self.conn.execute("UPDATE outgoing_source_invoice_items SET product_code='' "); self.conn.commit()
        before = self.snapshot()
        self.assertEqual(self.client.post(self.url + '/apply', json=report).status_code, 409)
        self.assertEqual(self.snapshot(), before)
        report = self.check(True); self.raw['invoiceDetail'][0]['productName'] = 'Different product'
        self.assertEqual(self.client.post(self.url + '/apply', json=report).status_code, 409)
        self.assertEqual(self.snapshot(), before)
        self.conn.execute("UPDATE outgoing_source_invoices SET stock_status='posted'"); self.conn.commit()
        before = self.snapshot()
        self.assertEqual(self.client.post(self.url + '/apply', json=report).status_code, 409)
        self.assertEqual(self.snapshot(), before)

    def test_remote_failure_and_identity_conflict_leave_every_table_unchanged(self):
        before = self.snapshot()
        self.remote.get_outgoing_invoice.side_effect = MinvoiceError('offline')
        self.assertEqual(self.client.get(self.url + '?fresh=1').status_code, 502)
        self.assertEqual(self.snapshot(), before)
        wrong = document(); wrong['id'] = '00000000-0000-0000-0000-000000000099'
        self.remote.get_outgoing_invoice.side_effect = lambda **kw: normalize_portal_document(wrong)
        self.assertEqual(self.client.get(self.url + '?fresh=1').status_code, 409)
        self.assertEqual(self.snapshot(), before)

    def test_unresolved_source_or_wrong_status_cannot_be_forced(self):
        before = self.snapshot(); report = self.check(True)
        self.assertEqual(self.client.post(self.url + '/apply', json=report).status_code, 409)
        self.assertEqual(self.snapshot(), before)
        self.raw = document(); self.raw['sendTaxStatus'] = 0
        report = self.check(True); self.assertFalse(report['can_apply'])
        self.assertTrue(report['status_note'])
        self.assertEqual(self.client.post(self.url + '/apply', json=report).status_code, 409)
        self.assertEqual(self.snapshot(), before)

    def test_line_findings_respect_discount_and_export_text_not_formula(self):
        raw = normalize_portal_document(self.raw); row = dict(self.conn.execute('SELECT * FROM outgoing_source_invoices').fetchone())
        raw['invoiceDetail'][0]['unitPrice'] = 60000
        self.assertTrue(inspect_amounts(row, raw)['lines'][0]['note'])
        raw['invoiceDetail'][0]['discountAmount'] = 20000
        self.assertFalse(inspect_amounts(row, raw)['lines'][0]['note'])
        raw['invoiceDetail'][0]['productName'] = '=1+1'
        self.conn.execute('UPDATE outgoing_source_invoices SET raw_json=?', (json.dumps(raw),)); self.conn.commit()
        response = self.client.get(self.url + '?download=1')
        ws = load_workbook(BytesIO(response.data)).active
        self.assertEqual(ws['B13'].value, '=1+1'); self.assertEqual(ws['B13'].data_type, 's')

    def test_single_portal_read_checks_company_number_series_date_and_id(self):
        for key, value in ((None, None), ('id','other'), ('sellerTaxCode','other'), ('invoiceNumber',13),
                           ('invoiceSerial','OTHER'), ('invoiceDate','2026-09-01')):
            c = portal_client(); c._token = 'fixture'; raw = document()
            if key: raw[key] = value
            c._portal_json = Mock(return_value=raw)
            args = dict(remote_id=document()['id'], series='1C26TYY', number='12', invoice_date='2026-08-31')
            if key:
                with self.assertRaises(MinvoiceError): c.get_outgoing_invoice(**args)
            else:
                self.assertEqual(c.get_outgoing_invoice(**args)['invoiceNumber'], 12)
                self.assertEqual(c._portal_json.call_args.args[0], 'GET')

    def test_failure_after_upsert_rolls_back_source_mapping_and_audit(self):
        self.raw = document(); report = self.check(True); before = self.snapshot()
        with patch('tdp_system.invoice_amount_support.refresh_linked_batches', side_effect=RuntimeError('fixture')):
            self.assertEqual(self.client.post(self.url + '/apply', json=report).status_code, 502)
        self.assertEqual(self.snapshot(), before)

    def test_concurrent_change_during_remote_read_and_other_tenant_are_rejected(self):
        self.raw = document(); report = self.check(True)
        def remote(**kwargs):
            self.conn.execute("UPDATE outgoing_source_invoice_items SET product_code='' "); self.conn.commit()
            return normalize_portal_document(self.raw)
        self.remote.get_outgoing_invoice.side_effect = remote
        self.assertEqual(self.client.post(self.url + '/apply', json=report).status_code, 409)
        header = self.conn.execute('SELECT * FROM outgoing_source_invoices').fetchone()
        self.assertEqual(header['subtotal'], 120000)
        self.conn.execute("UPDATE outgoing_source_invoices SET tenant='OTHER'"); self.conn.commit()
        self.assertEqual(self.client.get(self.url).status_code, 409)
        self.assertEqual(self.client.get(self.url+'?fresh=1').status_code, 409)


if __name__ == '__main__': unittest.main()
