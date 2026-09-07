"""Read-only spreadsheet snapshots from the same workbooks used for download."""
from datetime import date, datetime
from decimal import Decimal
import re

from openpyxl.styles.colors import COLOR_INDEX
from openpyxl.utils import coordinate_to_tuple, get_column_letter, range_boundaries


def _color(color, default):
    if color is not None and color.type == 'rgb':
        return '#' + color.rgb[-6:]
    if color is not None and color.type == 'indexed' and 0 <= color.indexed < len(COLOR_INDEX):
        return '#' + COLOR_INDEX[color.indexed][-6:]
    return default


def workbook_preview(workbook, title):
    sheets = {}
    total_cells = 0
    for index, sheet in enumerate(workbook):
        # Internal reconciliation metadata stays hidden, as it is in Excel.
        if sheet.sheet_state != 'visible':
            continue
        ranges = re.findall(r'\$?[A-Z]+\$?\d+:\$?[A-Z]+\$?\d+', str(sheet.print_area or ''))
        bounds = [range_boundaries(r) for r in ranges]
        right = max(b[2] for b in bounds) if bounds else sheet.max_column
        bottom = max(b[3] for b in bounds) if bounds else sheet.max_row
        total_cells += right * bottom
        if total_cells > 300000:
            raise ValueError('Báo cáo quá lớn để xem một lần. Hãy thu hẹp khoảng ngày hoặc tải file ZIP.')
        data = {}
        for row in sheet.iter_rows(max_row=bottom, max_col=right):
            for cell in row:
                if cell.value is None and not cell.has_style:
                    continue
                value = cell.value
                if cell.data_type == 'f':
                    # These exports contain static values. Do not execute new formulas.
                    raise ValueError('Báo cáo có công thức chưa được chuyển thành số liệu để xem.')
                if isinstance(value, (date, datetime)):
                    value = value.strftime('%d/%m/%Y')
                if isinstance(value, Decimal):
                    value = float(value)
                number = isinstance(value, (int, float)) and not isinstance(value, bool)
                pattern = cell.number_format if number else '@'
                if number and float(value).is_integer():
                    # The viewer otherwise prints a trailing dot for optional decimals.
                    pattern = re.sub(r'\.#+', '', pattern)
                style = {
                    'ff': cell.font.name or 'Arial', 'fs': cell.font.sz or 11,
                    'bl': int(bool(cell.font.b)), 'it': int(bool(cell.font.i)),
                    'ht': {'left': 1, 'center': 2, 'right': 3}.get(cell.alignment.horizontal, 3 if number else 1),
                    'vt': {'top': 1, 'center': 2, 'bottom': 3}.get(cell.alignment.vertical, 2),
                    'cl': {'rgb': _color(cell.font.color, '#172b3a')},
                    'bg': {'rgb': _color(cell.fill.fgColor, '#ffffff') if cell.fill.patternType == 'solid' else '#ffffff'},
                    'n': {'pattern': pattern},
                    'tb': 3 if cell.alignment.wrap_text else 1,
                }
                borders = {}
                for source, target in [('left','l'),('right','r'),('top','t'),('bottom','b')]:
                    edge = getattr(cell.border, source)
                    if edge and edge.style:
                        borders[target] = {'s': 1, 'cl': {'rgb': _color(edge.color, '#d1d8dd')}}
                if borders:
                    style['bd'] = borders
                data.setdefault(cell.row - 1, {})[cell.column - 1] = {'v': value if value is not None else '', 't': 2 if number else 1, 's': style}
        frozen_row, frozen_col = coordinate_to_tuple(str(sheet.freeze_panes or 'A1'))
        key = 'report-' + str(index)
        sheets[key] = {
            'id': key, 'name': sheet.title, 'rowCount': max(bottom, 2), 'columnCount': right,
            'cellData': data, 'showGridlines': 1, 'defaultRowHeight': 26, 'defaultColumnWidth': 120,
            'mergeData': [{'startRow': r.min_row - 1, 'endRow': r.max_row - 1,
                           'startColumn': r.min_col - 1, 'endColumn': r.max_col - 1}
                          for r in sheet.merged_cells.ranges if r.max_row <= bottom and r.max_col <= right],
            'freeze': {'xSplit': frozen_col - 1, 'ySplit': frozen_row - 1, 'startRow': frozen_row - 1, 'startColumn': frozen_col - 1},
            'columnData': {c - 1: {'w': max(55, (sheet.column_dimensions[get_column_letter(c)].width or 12) * 7 + 5),
                                    'hd': int(bool(sheet.column_dimensions[get_column_letter(c)].hidden))}
                           for c in range(1, right + 1)},
            'rowData': {r - 1: {'h': max(24, (sheet.row_dimensions[r].height or 18) * 4 / 3),
                                'hd': int(bool(sheet.row_dimensions[r].hidden))} for r in range(1, bottom + 1)},
        }
    return {'name': title, 'sheetOrder': list(sheets), 'sheets': sheets}
