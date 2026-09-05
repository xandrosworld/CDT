from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation


OUTPUT = r"C:\Users\DELL\Downloads\TDP_FULL_TRANSFER_20260901_200101\CHECKLIST_XAC_NHAN_NOI_DUNG_SUA_TDP_20260904.xlsx"

items = [
    "Khi sửa đơn, có nút đi thẳng tới dòng lỗi đầu tiên; lọc riêng lỗi bắt buộc và cảnh báo.",
    "Bảng sửa đơn giữ hàng tiêu đề khi cuộn; các cột quan trọng nằm gần nhau, dễ nhìn.",
    "Màn sửa đơn hiện tổng số dòng, tổng số lượng và tổng tiền theo phần đang lọc.",
    "Nếu nạp nhiều file cùng ngày, dùng file hợp lệ tải sau cùng và báo rõ dữ liệu nào được thay hoặc giữ.",
    "Màn đặt hàng nhà cung cấp chỉ hiện tên nhà cung cấp, số dòng, tổng lượng, tổng tiền và trạng thái; chi tiết mở bên trong.",
    "Sao chép đơn đặt hàng xong tự chuyển sang Đã đặt; nhà cung cấp đã đặt nằm dưới, chưa đặt nằm trên; có nút mở lại.",
    "Bỏ mục Phiếu giao bị trùng nếu chức năng đã có đầy đủ trong Công việc hằng ngày.",
    "Ẩn Suất ăn & đặt hàng bếp/PO và Chấm công & lương khỏi giao diện, nhưng giữ nguyên dữ liệu cũ.",
    "Mỗi màn hình có một câu ngắn nói rõ dùng để làm gì và bước tiếp theo là gì.",
    "Các chứng từ và báo cáo mở ra xem trực tiếp trên phần mềm như bảng Excel, không tự tải file.",
    "Bảng xem trực tiếp có thanh cuộn ngang/dọc, giữ tiêu đề khi cuộn và có thể phóng to/thu nhỏ.",
    "Các cột số lượng, đơn giá, thành tiền căn phải; nơi có số lượng phải có tổng lượng, nơi có tiền phải có tổng tiền.",
    "Bản xem trên phần mềm và file Excel/PDF tải xuống phải cùng số liệu.",
    "Màn in đơn/bảng kê lọc được khách hàng và khoảng ngày; chọn tất cả hoặc chọn riêng nhiều phiếu rồi in một lần.",
    "Có thể xem trước từng phiếu; màn hình hiện tổng số phiếu đang lọc và tổng số phiếu đã chọn.",
    "Bảng và bản in dùng nền trắng, chữ đen, đường kẻ rõ, ít màu và không tốn mực.",
    "Tên hàng dài tự xuống dòng và tăng chiều cao hàng, không đè lên đường viền hoặc lệch cột.",
    "Các số tiền/đơn giá làm tròn đến số nguyên gần nhất; từ 0,5 làm tròn lên; hiển thị dạng #,##0.",
    "Hóa đơn đầu vào lọc theo ngày Việt Nam; tháng 08/2026 phải đủ đúng 266 hóa đơn và tải lại không bị trùng.",
    "Hóa đơn đầu vào chưa ghép mã, thiếu đơn vị hoặc có lỗi tự nằm trên cùng và lọc riêng được.",
    "Sau khi ghép mã phải có bước xác nhận ghi kho; ghi xong mở thẳng được chi tiết nhập trong Báo cáo vật tư hàng hóa.",
    "Hóa đơn đầu ra M-Invoice lọc đúng ngày Việt Nam, không lệch đầu/cuối kỳ và tải lại không bị trùng.",
    "Phân biệt rõ hóa đơn dự thảo, đã phát hành, đã hủy và thất bại; chỉ hóa đơn phát hành hợp lệ mới trừ kho.",
    "Bảng kê & hóa đơn tách rõ phần Được xuất và Chưa được xuất, kèm lý do cụ thể cho phần chưa đủ điều kiện.",
    "Không cho xuất vượt tồn kho hóa đơn; nếu cần 10 mà có 7 thì chỉ xuất 7 và giữ lại 3 để xử lý lần sau.",
    "Phần còn thiếu tự được tính tiếp khi có thêm đầu vào, không mất dòng, không xuất trùng và không làm âm kho.",
    "Màn tổng quan hóa đơn hiện tổng cần, đã dự thảo, đã phát hành, có thể xuất, còn thiếu và giá trị tương ứng.",
    "File đưa lên M-Invoice chỉ chứa phần đủ điều kiện và tách đúng nhà thầu, nhóm thuế, lần xuất.",
    "Chốt tháng hiện tổng mã, tổng lượng, tổng giá trị; chặn chốt nếu còn mã âm hoặc cần kiểm tra.",
    "Tồn cuối tháng 8 chuyển đúng sang tồn đầu tháng 9; bấm lại không cộng trùng và có thể mở lại kỳ khi cần sửa.",
    "Báo cáo tổng hợp tự lấy bếp/nhà thầu hiện có; bếp mới tự xuất hiện; có tổng tháng và tổng theo nhóm/nhà thầu.",
    "Công nợ phải trả lọc được khoảng ngày và nhà cung cấp; xem trực tiếp tổng lượng, tổng tiền; Excel đúng mẫu 14 cột.",
    "Khoản phải trả đã thanh toán được ẩn khỏi danh sách còn nợ nhưng vẫn còn lịch sử; cho hoàn tác khi ghi sai.",
    "Công nợ phải thu có phần theo bếp, phần tổng và phần phải trả; khoản đã thu vẫn giữ lịch sử.",
    "Giấy biên nhận lấy đúng CCCD, ngày cấp, nơi cấp và địa chỉ của cùng một người bán; thiếu dữ liệu thì cảnh báo, không tự đoán.",
    "Phiếu giao và giấy tờ in đúng mẫu khách: ngày, tiêu đề, phần ký đúng vị trí; cân khổ giấy và không chèn ảnh thay chữ.",
    "Mọi thao tác xác nhận được lưu ngay; hệ thống tự sao lưu và hiện lần sao lưu gần nhất cùng trạng thái.",
    "Dữ liệu cũ còn nguyên sau khi nâng cấp; bản EXE mới mở được trên máy khách và dùng lại cơ sở dữ liệu hiện tại.",
    "Làm lại hướng dẫn Word và clip theo giao diện cuối, hướng dẫn theo quy trình công việc và có cách xử lý lỗi thường gặp.",
]

wb = Workbook()
ws = wb.active
ws.title = "Khách hàng kiểm tra"
ws.sheet_view.showGridLines = False

ws.merge_cells("A1:F1")
ws["A1"] = "CHECKLIST XÁC NHẬN NỘI DUNG SỬA PHẦN MỀM TĐP"
ws["A1"].font = Font(name="Arial", size=16, bold=True, color="FFFFFF")
ws["A1"].fill = PatternFill("solid", fgColor="0F766E")
ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[1].height = 32

ws.merge_cells("A2:F2")
ws["A2"] = "Chị xem từng nội dung, chọn ở cột “Khách chốt nội dung”. Sau khi bên em sửa, chị kiểm tra và chọn kết quả."
ws["A2"].font = Font(name="Arial", size=11, italic=True, color="475569")
ws["A2"].fill = PatternFill("solid", fgColor="F0FDFA")
ws["A2"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
ws.row_dimensions[2].height = 36

headers = ["STT", "Nội dung khách yêu cầu", "Khách chốt nội dung", "Tiến độ bên em", "Khách kiểm tra kết quả", "Ý kiến khách"]
for col, value in enumerate(headers, 1):
    cell = ws.cell(4, col, value)
    cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="115E59")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
ws.row_dimensions[4].height = 36

thin = Side(style="thin", color="CBD5E1")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

for idx, item in enumerate(items, 1):
    row = idx + 4
    values = [idx, item, "Chưa chốt", "Chưa làm", "Chưa kiểm tra", ""]
    for col, value in enumerate(values, 1):
        cell = ws.cell(row, col, value)
        cell.font = Font(name="Arial", size=10.5, color="0F172A")
        cell.border = border
        cell.alignment = Alignment(
            horizontal="center" if col in (1, 3, 4, 5) else "left",
            vertical="top",
            wrap_text=True,
        )
        if idx % 2 == 0:
            cell.fill = PatternFill("solid", fgColor="F8FAFC")
    ws.row_dimensions[row].height = 48

last_row = len(items) + 4
dv_confirm = DataValidation(type="list", formula1='"Chưa chốt,Đúng ý chị,Cần viết lại"', allow_blank=False)
dv_progress = DataValidation(type="list", formula1='"Chưa làm,Đang làm,Đã sửa"', allow_blank=False)
dv_result = DataValidation(type="list", formula1='"Chưa kiểm tra,Đạt,Cần sửa thêm"', allow_blank=False)
for dv, rng in ((dv_confirm, f"C5:C{last_row}"), (dv_progress, f"D5:D{last_row}"), (dv_result, f"E5:E{last_row}")):
    ws.add_data_validation(dv)
    dv.add(rng)

ws.column_dimensions["A"].width = 7
ws.column_dimensions["B"].width = 76
ws.column_dimensions["C"].width = 22
ws.column_dimensions["D"].width = 18
ws.column_dimensions["E"].width = 24
ws.column_dimensions["F"].width = 34
ws.freeze_panes = "A5"
ws.auto_filter.ref = f"A4:F{last_row}"
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.print_title_rows = "4:4"
ws.sheet_properties.outlinePr.summaryBelow = True

wb.save(OUTPUT)
print(OUTPUT)
